"""Explicit invented standing artwork in the marked, disposable browser fixture."""
from dataclasses import replace
from hashlib import sha256
from html import escape
from io import BytesIO
import json

from scripts.workspace_history_observations import namespace

MODES = ('wide-tall', 'tall-wide', 'square', 'absent-photo')
COLORS = {'brand_logo': (35, 94, 140), 'client_logo': (124, 43, 89), 'cover_photo': (23, 63, 73)}


def profile():
    from app.contracts import RRH_CONTRACT
    from app.monthly_report_model import synthetic_profiles
    return replace(synthetic_profiles()[0], contract=RRH_CONTRACT, key='synthetic-first-cover')


def artwork(size, role):
    from PIL import Image, ImageDraw
    image = Image.new('RGBA', size, (255, 255, 255, 0))
    draw = ImageDraw.Draw(image)
    radius = min(size) // 4
    cx, cy = size[0] // 2, size[1] // 2
    draw.ellipse((cx-radius, cy-radius, cx+radius, cy+radius), fill=COLORS[role])
    out = BytesIO(); image.save(out, 'PNG')
    return out.getvalue()


def setup(base, token, mode):
    from app import monthly_report_testing as testing, monthly_report_library as library
    from app.monthly_report_model import ResolvedBlock
    _, identity, generation = namespace(base, token)
    if (testing.scope(), testing.generation()) != (identity, generation) or mode not in MODES:
        raise ValueError('Standing-cover setup requires its current isolated namespace.')
    selected = profile()
    if any(item.key == selected.key for item in library.list_profiles(selected.contract)):
        raise ValueError('Existing synthetic standing configuration cannot be replaced.')
    from app import monthly_report_designs as designs
    from app.monthly_report_import import inspect_docx
    from app.monthly_report_model import default_sections
    from docx import Document
    from docx.shared import Inches
    if designs.master_state()['revision'] != 0:
        raise ValueError('Existing synthetic master cannot be replaced.')
    sizes = {'wide-tall': ((400,100), (100,400)), 'tall-wide': ((100,400), (400,100)),
             'square': ((200,200), (200,200)), 'absent-photo': ((400,100), (400,100))}[mode]
    images = {role: artwork(size, role) for role,size in
              zip(('brand_logo','client_logo','cover_photo'), (*sizes,(400,300)))
              if role != 'cover_photo' or mode != 'absent-photo'}
    # An invented standing company master, NOT an uploaded starting report.
    # Deliberately incorrect logo/photo frames let the real renderer establish
    # proportional placement rather than the fixture supplying pre-fixed frames.
    document = Document()
    document.add_paragraph('Synthetic facility')
    document.add_paragraph('October 2026')
    for role, raw in images.items():
        document.add_picture(BytesIO(raw), width=Inches(4 if role=='cover_photo' else 3),
                             height=Inches(2 if role=='cover_photo' else 1))
    for section in default_sections():
        if section.key == 'cover':
            continue
        document.add_page_break()
        document.add_heading(section.title, 1)
        document.add_paragraph('Synthetic master placeholder')
    source = library._root() / 'qa-standing-master.docx'
    source.parent.mkdir(parents=True, exist_ok=True)
    document.save(source)
    inspection = inspect_docx(source)
    image_items = [item for item in inspection.items if item.kind=='image']
    assert len(image_items) == len(images)
    digest = designs.install_master(source, actor='Synthetic Editor', expected_revision=0)
    selected = designs.pin(selected, latest_master=True)
    assert selected.template == designs.MASTER_PREFIX + digest
    state = library.save_profile(selected, expected_revision=0, actor='Synthetic Editor', confirmed=True)
    records = {}
    for (role,raw), item in zip(images.items(), image_items):
        reference = library.asset_reference(raw, 'png')
        state = library.replace_block(selected.contract, selected.key, role,
            ResolvedBlock(role, 'Library', asset_hashes=(reference,),
                          references=(f'docx:{inspection.sha256}:{item.id}',)), assets=((reference, raw),),
            expected_revision=state.revision, actor='Synthetic Editor', confirmed=True, label=role)
        records[role] = {'reference': reference, 'sha256': sha256(raw).hexdigest()}
    marker = library._root() / 'qa-standing-cover.json'
    marker.write_text(json.dumps({'mode': mode, 'artwork': records, 'master': digest}))


def select():
    import streamlit as st
    from app.monthly_report_editor import _signature
    from app.monthly_report_editor import _key
    selected = profile()
    st.session_state['report_contract'] = selected.contract
    st.session_state['report_change_contract'] = False
    st.session_state['report_month_number'] = 10
    st.session_state['report_year'] = 2026
    for facility in selected.facilities:
        st.session_state['report_sites_' + _key(selected.contract) + '_' + facility.key] = True
    prefix = 'report_guided_' + _signature((selected.contract, selected.key, '2026-10'))
    st.session_state[prefix + '_prepared'] = 'Synthetic Editor'


def render(base, token, field, move):
    import streamlit as st
    from app import monthly_report_testing as testing, monthly_report_library as library
    from app.monthly_report_guided import render_guided_workflow
    from app.monthly_report_editor import _signature
    from app.monthly_report_model import ReportPeriod
    testing.data_dir(base)
    _, identity, generation = namespace(base, token)
    mode = st.query_params.get('cover_mode', '')
    if mode not in MODES:
        raise ValueError('Choose an explicit synthetic standing-cover mode.')
    selected = profile()
    marker = library._root() / 'qa-standing-cover.json'
    if not marker.exists():
        if st.button('Create synthetic standing cover'):
            setup(base, token, mode)
            select()
            st.rerun()
        return
    expected = json.loads(marker.read_text())
    if expected['mode'] != mode:
        raise ValueError('Changing mode cannot replace existing synthetic artwork.')
    if 'report_contract' not in st.session_state:
        select()
    render_guided_workflow(token, 'UTC', field, move)
    prefix = 'report_guided_' + _signature((selected.contract, selected.key, '2026-10'))
    draft = st.session_state[prefix + '_draft']
    blocks = {block.key: block for block in draft.blocks}
    observed = {}
    for role, item in expected['artwork'].items():
        block = blocks[role]
        observed[role] = {'references': list(block.asset_hashes), 'source': block.source,
                         'sha256': sha256(library.read_asset(selected.contract, selected.key, block.asset_hashes[0])).hexdigest()}
    if library.load_snapshot(selected.contract, selected.key, ReportPeriod(2026,10)) is not None:
        raise ValueError('First-display test must not create a saved report.')
    if library.load_imported_draft(selected.contract, selected.key) is not None or draft.sources:
        raise ValueError('First-display test must not import an old report.')
    record = {'scope': identity, 'generation': generation, 'mode': mode, 'artwork': observed,
              'photo_present': bool(blocks['cover_photo'].asset_hashes), 'draft': draft.fingerprint,
              'saved_report': False, 'imported_report': False, 'template': draft.profile.template}
    st.html('<span hidden data-enfra-standing-cover="' + escape(json.dumps(record), quote=True) + '"></span>')
