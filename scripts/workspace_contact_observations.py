"""Explicit invented setup and passive receipts for the marked hosted fixture.

Production UI/storage/converters remain in use. The sole fault seam raises one
requested synthetic shared-write error, inside this disposable namespace only.
"""
from dataclasses import asdict, replace
from hashlib import sha256
from html import escape
import json
import os

from scripts.workspace_history_observations import namespace

VENDOR_COLUMNS = ('Name', '', 'Custom detail')
VENDOR_ROWS = (('Invented Vendor Alpha', 'Retain blank-heading value', 'Retain custom value'),
               ('', '', ''), ('Invented Vendor Alpha', 'Retain blank-heading value', 'Retain custom value'))


def _folder(base, token):
    from app import monthly_report_library as library, monthly_report_testing as testing
    target, identity, generation = namespace(base, token)
    if (os.environ.get('EPC_MONTHLY_CONTACT_TASK') != '1'
            or testing.scope() != identity or testing.generation() != generation
            or library._root().resolve() != (target / 'monthly_reports').resolve()):
        raise ValueError('Contact fixture requires its marked cookie-isolated storage root.')
    return library._root() / 'qa-compact-contacts'


def setup(base, token):
    from app import monthly_report_qa_fixture as fixture, monthly_report_library as library
    from app.monthly_report_model import ReportTable
    from docx import Document
    folder = _folder(base, token)
    if folder.exists() or library.load_snapshot(fixture.CONTRACT, fixture.profile().key, fixture.PERIOD):
        raise ValueError('Existing report preserved; contact setup requires an empty isolated namespace.')
    draft = fixture.prepare('prefill')
    # This is the exact saved-directory default, not a manually authored row
    # promoted by title matching. The general fixture remains unchanged.
    blocks = tuple(replace(block, source='Library') if block.key == 'contact_matrix' else
                   replace(block, extra_tables=(ReportTable(VENDOR_COLUMNS, VENDOR_ROWS, 'synthetic:custom-vendors'),))
                   if block.key == 'subcontractor_matrix' else block for block in draft.blocks)
    draft = replace(draft, blocks=blocks)
    library.save_snapshot(draft, expected_revision=1, entered_editor=fixture.ACTOR)
    folder.mkdir()
    doc = Document()
    doc.add_heading('ORGANIZATIONAL CHART', 1)
    table = doc.add_table(rows=2, cols=5)
    for cell, text in zip(table.rows[0].cells, ('Facility','Role','Name','Phone','Email')):
        cell.text = text
    for cell in table.rows[1].cells:
        cell.text = 'OLD SYNTHETIC CONTACT'
    doc.save(folder / 'source.docx')
    (folder / 'setup.json').write_text(json.dumps({
        'source_sha256': sha256((folder / 'source.docx').read_bytes()).hexdigest(),
        'saved_fingerprint': draft.fingerprint,
        'vendor': asdict(next(block for block in draft.blocks if block.key == 'subcontractor_matrix')),
    }))
    (folder / 'attempts.json').write_text('[]')


def install_fault_seam(base, token):
    import streamlit as st
    from app import monthly_report_directory as directory, monthly_report_library as library
    from app import monthly_report_testing as testing
    _folder(base, token)
    if getattr(directory.automatically_keep_directory_contact_edits, '_synthetic_contact_fault', False):
        return
    original = directory.automatically_keep_directory_contact_edits
    def measured(*args, **kwargs):
        from scripts.workspace_browser_observations import require_synthetic_root
        require_synthetic_root(base)
        folder = library._root() / 'qa-compact-contacts'
        if (os.environ.get('EPC_MONTHLY_CONTACT_TASK') != '1' or testing.scope() is None
                or not (folder / 'setup.json').is_file()):
            raise AssertionError('Synthetic contact fault seam cannot access unmarked standing data.')
        target = testing._location(base, testing.scope())
        if (library._root().resolve() != (target / 'monthly_reports').resolve()
                or testing._marker(target / '.enfra-qa-namespace', testing.scope()) != testing.generation()):
            raise AssertionError('Synthetic contact fault seam requires the current isolated generation.')
        mode = st.session_state.pop('_contact_fault_once', '')
        path = folder / 'attempts.json'
        attempts = json.loads(path.read_text())
        attempts.append({'mode': mode or 'ordinary'})
        path.write_text(json.dumps(attempts))
        if mode == 'before':
            raise OSError('Deliberate synthetic failure before shared save')
        result = original(*args, **kwargs)
        if mode == 'after':
            raise OSError('Deliberate synthetic failure after shared save acknowledgement')
        return result
    measured._synthetic_contact_fault = True
    directory.automatically_keep_directory_contact_edits = measured


def snapshot(base, token):
    from app import monthly_report_library as library, monthly_report_qa_fixture as fixture
    from app import monthly_report_directory as directory
    from app.monthly_report_active_work import load
    from app.monthly_report_testing_ui import _working_report
    folder = _folder(base, token)
    _, identity, generation = namespace(base, token)
    saved = library.load_snapshot(fixture.CONTRACT, fixture.profile().key, fixture.PERIOD)
    working = _working_report(saved)
    journal = load(token, fixture.CONTRACT, fixture.profile().key, fixture.PERIOD)
    def state(draft):
        unrelated = replace(draft, blocks=tuple(b for b in draft.blocks
            if b.key not in {'contact_matrix','subcontractor_matrix'}))
        return {'contacts': asdict(next(b for b in draft.blocks if b.key == 'contact_matrix')),
                'vendors': asdict(next(b for b in draft.blocks if b.key == 'subcontractor_matrix')),
                'unrelated': unrelated.fingerprint,
                # Only reached after _folder/namespace validate the marked synthetic fixture.
                # Keep actual field evidence when a strict unrelated-content check fails.
                'unrelated_payload': asdict(unrelated)}
    standing = directory.load_directory(fixture.CONTRACT)
    setup_record = json.loads((folder / 'setup.json').read_text())
    return {'scope': identity, 'generation': generation, 'saved_revision': saved.revision,
            'saved_unchanged': saved.draft.fingerprint == setup_record['saved_fingerprint'],
            'original_unchanged': sha256((folder / 'source.docx').read_bytes()).hexdigest() == setup_record['source_sha256'],
            'working': state(working), 'durable': state(journal.draft if journal else saved.draft),
            'directory': asdict(standing), 'attempts': json.loads((folder / 'attempts.json').read_text())}


def render(base, token, field, move):
    import streamlit as st
    from app import monthly_report_library as library, monthly_report_qa_fixture as fixture
    from app import monthly_report_testing as testing
    from app.monthly_report_testing_ui import _select_fixture, _working_report
    from app.monthly_report_guided import render_guided_workflow
    testing.data_dir(base)
    folder = _folder(base, token)
    st.title('Synthetic compact contacts acceptance')
    saved = library.load_snapshot(fixture.CONTRACT, fixture.profile().key, fixture.PERIOD)
    if saved is None:
        if st.button('Create synthetic compact contact report'):
            setup(base, token)
            _select_fixture()
            st.rerun()
        return
    install_fault_seam(base, token)
    if 'report_contract' not in st.session_state:
        _select_fixture()
    with st.expander('Synthetic fault controls'):
        st.button('Fail next synthetic shared save before commit',
                  on_click=st.session_state.__setitem__, args=('_contact_fault_once','before'))
        st.button('Fail next synthetic shared save after commit',
                  on_click=st.session_state.__setitem__, args=('_contact_fault_once','after'))
    st.button('Observe synthetic passive rerun')
    render_guided_workflow(token, 'UTC', field, move)
    if st.button('Render synthetic native contact outputs'):
        from app.monthly_report_native_layout import build_native_docx
        from app.pdf_converter import convert_to_pdf
        working = _working_report(saved)
        focused = replace(working, sections=tuple(s for s in working.sections if s.key == 'organization'),
                          blocks=tuple(b for b in working.blocks if b.key == 'contact_matrix'))
        raw = build_native_docx(focused, folder / 'source.docx', section_key='organization', master=True)
        path = folder / 'output.docx'
        path.write_bytes(raw)
        pdf = convert_to_pdf(path)
        try:
            st.session_state['_contact_native_outputs'] = (raw, pdf.read_bytes())
        finally:
            pdf.unlink(missing_ok=True)
    if outputs := st.session_state.get('_contact_native_outputs'):
        st.download_button('Download synthetic contact Word', outputs[0], 'synthetic-contacts.docx')
        st.download_button('Download synthetic contact PDF', outputs[1], 'synthetic-contacts.pdf')
    receipt = snapshot(base, token)
    receipt['sequence'] = st.session_state.get('_contact_receipt_sequence', 0) + 1
    st.session_state['_contact_receipt_sequence'] = receipt['sequence']
    st.html('<span hidden data-enfra-contact-receipt="' + escape(json.dumps(receipt), quote=True) + '"></span>')
