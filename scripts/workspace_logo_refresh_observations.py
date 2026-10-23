"""Dedicated setup and passive receipts inside the existing marked QA fixture.

Only explicit setup seeds invented records. Thereafter the production guided
editor owns all report mutation; observations never repair or set its state.
"""
from dataclasses import asdict, replace
from hashlib import sha256
from html import escape
from io import BytesIO
import json
from pathlib import Path

from scripts.workspace_history_observations import namespace, storage_state

MODES = ('imported', 'custom', 'omit')


def artwork(color):
    from PIL import Image
    stream = BytesIO()
    Image.new('RGB', (240, 80), color).save(stream, 'PNG')
    return stream.getvalue()


def setup(base, token, mode):
    from app import monthly_report_testing as testing, monthly_report_library as library
    from app import monthly_report_qa_fixture as fixture, monthly_report_branding as branding
    from app.monthly_report_import import inspect_docx, map_items, ImportMapping
    from docx import Document
    from docx.shared import Inches
    target, identity, generation = namespace(base, token)
    if testing.scope() != identity or testing.generation() != generation or mode not in MODES:
        raise ValueError('Logo refresh setup requires its current synthetic namespace and mode.')
    path = library._root() / 'qa-logo-refresh'
    if path.exists() or library.load_snapshot(fixture.CONTRACT, fixture.profile().key, fixture.PERIOD):
        raise ValueError('Existing synthetic report preserved; setup cannot replace it.')
    if branding.load_branding():
        raise ValueError('Existing synthetic catalog preserved; setup requires an empty catalog.')
    draft = fixture.prepare('prefill')
    path.mkdir()
    old, custom, company = artwork('red'), artwork('purple'), artwork('green')
    document = Document()
    document.add_paragraph('SYNTHETIC LOGO REFRESH ONLY')
    document.add_picture(BytesIO(old), width=Inches(3), height=Inches(1))
    document.add_picture(BytesIO(old), width=Inches(2), height=Inches(2))
    source = path / 'source.docx'
    document.save(source)
    inspection = inspect_docx(source)
    item = next(item for item in inspection.items if item.kind == 'image')
    mapped = map_items(source, inspection, (ImportMapping(item.id, 'brand_logo'),))
    block = mapped.blocks[0]
    assets = dict(mapped.assets)
    if mode == 'custom':
        reference = library.asset_reference(custom, 'png')
        assets[reference] = custom
        block = replace(block, source='Replace once', asset_hashes=(reference,))
    elif mode == 'omit':
        block = replace(block, source='Omit')  # Retained bytes must stay omitted.
    draft = replace(draft, blocks=(*draft.blocks, block))
    library.save_snapshot(draft, expected_revision=1, assets=tuple(assets.items()), entered_editor=fixture.ACTOR)
    reference = library.asset_reference(company, 'png')
    logo = branding.BrandLogo('enfra', 'Synthetic current company', (), reference,
        'https://example.invalid/synthetic', 'https://example.invalid/synthetic.png', '2026-10-22')
    branding.save_branding((logo,), {reference: company}, expected_revision=0,
        actor=fixture.ACTOR, confirmed=True)
    (path / 'setup.json').write_text(json.dumps({'mode': mode, 'source_sha256': sha256(source.read_bytes()).hexdigest()}))


def snapshot(base, token):
    from app import monthly_report_testing as testing, monthly_report_library as library
    from app import monthly_report_qa_fixture as fixture, monthly_report_branding as branding
    from app.monthly_report_active_work import load
    from app.monthly_report_testing_ui import _working_report
    from app.monthly_report_editor import _signature
    import streamlit as st
    target, identity, generation = namespace(base, token)
    if (testing.scope(), testing.generation()) != (identity, generation):
        raise ValueError('Logo receipt requires the current synthetic generation.')
    storage_state(base, token)
    saved = library.load_snapshot(fixture.CONTRACT, fixture.profile().key, fixture.PERIOD)
    if saved is None:
        return None
    working = _working_report(saved)
    journal = load(token, fixture.CONTRACT, fixture.profile().key, fixture.PERIOD)
    prefix = 'report_guided_' + _signature((fixture.CONTRACT, fixture.profile().key, fixture.PERIOD.key))
    memory = st.session_state.get(prefix + '_assets', {})
    def state(draft, assets):
        block = next(b for b in draft.blocks if b.key == 'brand_logo')
        blobs = []
        for ref in block.asset_hashes:
            raw = assets.get(ref)
            if raw is None:
                raw = library.read_asset(fixture.CONTRACT, fixture.profile().key, ref)
            blobs.append({'reference': ref, 'sha256': sha256(raw).hexdigest()})
        return {'block': asdict(block), 'assets': blobs,
            'without_logo': replace(draft, blocks=tuple(b for b in draft.blocks if b.key != 'brand_logo')).fingerprint}
    catalog = branding.load_branding()
    logo = branding.find_logo('', brand=True)
    source = library._root() / 'qa-logo-refresh' / 'source.docx'
    record = json.loads((source.parent / 'setup.json').read_text())
    return {'scope': identity, 'generation': generation, 'mode': record['mode'],
        'working': state(working, memory), 'durable': state(journal.draft, dict(journal.assets)) if journal else None,
        'saved': state(saved.draft, {}), 'saved_revision': saved.revision,
        'company_asset': logo.asset, 'company_sha256': sha256(branding.read_logo(logo)).hexdigest(),
        'catalog_revision': catalog.revision,
        'original_unchanged': sha256(source.read_bytes()).hexdigest() == record['source_sha256']}


def render(base, token, field, move):
    import streamlit as st
    from app import monthly_report_library as library, monthly_report_qa_fixture as fixture
    from app.monthly_report_testing_ui import _select_fixture, _working_report, _qa_asset_loader
    from app.monthly_report_guided import render_guided_workflow
    mode = st.query_params.get('logo_mode', '')
    if mode not in MODES:
        st.error('Choose an explicit synthetic logo fixture mode.'); st.stop()
    from app import monthly_report_testing as testing
    testing.data_dir(base)  # Create only the already-authorized cookie-isolated lease.
    namespace(base, token)
    st.title('Synthetic About company-logo refresh')
    saved = library.load_snapshot(fixture.CONTRACT, fixture.profile().key, fixture.PERIOD)
    if saved is None:
        if st.button('Create synthetic ' + mode + ' logo report'):
            setup(base, token, mode)
            _select_fixture()
            st.rerun()
        return
    if json.loads((library._root() / 'qa-logo-refresh' / 'setup.json').read_text())['mode'] != mode:
        raise ValueError('Changing mode cannot replace an existing report.')
    # Only route selection, never draft values, is restored on a fresh browser render.
    if 'report_contract' not in st.session_state:
        _select_fixture()
    render_guided_workflow(token, 'UTC', field, move)
    if st.button('Render synthetic native logo outputs'):
        from app.monthly_report_native_layout import build_native_docx
        from app.pdf_converter import convert_to_pdf
        working = _working_report(saved)
        source = library._root() / 'qa-logo-refresh' / 'source.docx'
        raw = build_native_docx(working, source, asset_loader=_qa_asset_loader(working, base), section_key='cover')
        # Only disposable fixture artifacts; never a shared version/save action.
        output = source.parent / 'output.docx'
        output.write_bytes(raw)
        pdf = convert_to_pdf(output)
        try:
            st.session_state['_logo_refresh_outputs'] = (raw, pdf.read_bytes())
        finally:
            pdf.unlink(missing_ok=True)
    if outputs := st.session_state.get('_logo_refresh_outputs'):
        st.download_button('Download synthetic native Word', outputs[0], 'synthetic-logo.docx')
        st.download_button('Download synthetic native PDF', outputs[1], 'synthetic-logo.pdf')
    record = snapshot(base, token)
    sequence = st.session_state.get('_logo_refresh_sequence', 0) + 1
    st.session_state['_logo_refresh_sequence'] = sequence
    record['sequence'] = sequence
    st.html('<span hidden data-enfra-logo-refresh="' + escape(json.dumps(record, sort_keys=True), quote=True) + '"></span>')
