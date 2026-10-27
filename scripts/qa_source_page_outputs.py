"""Export actual synthetic source-caption and picture-ownership output evidence.

Called by the existing synthetic output CI step. AppTest controls do not prove
native browser delivery; real saved snapshots and generated download bytes are
retained separately. Photo-grid captions are raster text, so they require the
recorded every-page visual check, not an editable-Word-text assertion.
"""
from dataclasses import asdict, replace
from datetime import datetime, timezone
from hashlib import sha256
from io import BytesIO
import json
from pathlib import Path
import platform
import shutil
import subprocess
import sys
from tempfile import TemporaryDirectory
from zipfile import ZipFile

from docx import Document
import pymupdf as fitz
from pytest import MonkeyPatch

from app import monthly_report_library as library, monthly_report_sources as sources, pdf_converter, web_ui
from app.monthly_report_asset_review import approve_all_assets, pending_asset_indexes
from app.monthly_report_docx import generate_report
from app.monthly_report_model import ReportDraft, ReportPeriod, ResolvedBlock, default_sections
from app.monthly_report_setup import merge_report_block
from app.monthly_report_source_retirement import retire_contributions


def _command(args):
    result = subprocess.run(args, capture_output=True, text=True, check=True, timeout=10)
    return (result.stdout + result.stderr).strip()


def _generated_downloads(app, button, current_draft):
    notices = [widget for widget in app.checkbox if widget.label == "I checked these specific warnings"]
    if notices:
        assert not app.error, [item.value for item in app.error]
        notices[0].check().run()
    create = button(app, "Create Word and PDF")
    assert not create.disabled, [item.value for item in app.error]
    create.click().run(timeout=30)
    assert not app.exception
    packages = [value for key, value in app.session_state.filtered_state.items()
                if key.startswith("report_guided_") and key.endswith("_package")]
    assert len(packages) == 1, [item.value for item in app.error]
    draft = current_draft(app)
    saved = library.load_snapshot(draft.profile.contract, draft.profile.key, draft.period)
    assert saved is not None and saved.draft == draft
    assert {"Download DOCX", "Download PDF"} <= {
        item.label for item in app.get("download_button")}
    return draft, saved, packages[0]


def _record(output, label, draft, saved, package, *, required_text, forbidden_text,
            pictures, route, required_graphic_captions=()):
    assert package.pdf and not package.pdf_error, package.pdf_error
    assert saved.draft == draft
    folder = output / label
    folder.mkdir()
    docx_path = folder / "SYNTHETIC-report.docx"
    docx_path.write_bytes(package.docx)
    (folder / "SYNTHETIC-report.pdf").write_bytes(package.pdf)
    (folder / "saved-draft.json").write_text(json.dumps(asdict(draft), indent=2) + "\n")
    (folder / "saved-snapshot.json").write_text(json.dumps(asdict(saved), indent=2) + "\n")
    originals = []
    original_folder = folder / "synthetic-source-originals"
    for source in draft.sources:
        original = sources.source_bytes(draft.profile, source)
        assert sha256(original).hexdigest() == source.sha256
        original_folder.mkdir(exist_ok=True)
        name = source.sha256 + Path(source.filename).suffix.lower()
        (original_folder / name).write_bytes(original)
        originals.append({"id": source.id, "sha256": source.sha256,
                          "bytes": len(original), "file": "synthetic-source-originals/" + name})
    document = Document(BytesIO(package.docx))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    text += "\n" + "\n".join(" | ".join(cell.text for cell in row.cells)
                              for table in document.tables for row in table.rows)
    assert len(document.inline_shapes) == pictures, (label, len(document.inline_shapes))
    with ZipFile(BytesIO(package.docx)) as archive:
        xml = "\n".join(archive.read(name).decode("utf-8", "replace")
                        for name in archive.namelist() if name.endswith((".xml", ".rels")))
        # Retain exact graphic bytes as an additional inspection surface; they
        # are not substitutes for the final Word/PDF page check.
        media = [name for name in archive.namelist() if name.startswith("word/media/")]
        if media:
            embedded = folder / "embedded-graphics"
            embedded.mkdir()
            for name in media:
                (embedded / Path(name).name).write_bytes(archive.read(name))
    for expected in required_text:
        assert expected in text, (label, "Missing editable Word text", expected)
    for excluded in forbidden_text:
        assert excluded not in xml, (label, "Old caption remains in a Word package part", excluded)

    # Re-open the delivered DOCX from disk and render it through the named
    # LibreOffice path. This is a repeat rendering of saved bytes with the same
    # backend, not a second independent renderer or Microsoft Word acceptance.
    rendered = pdf_converter.convert_to_pdf(docx_path)
    try:
        word_pdf = rendered.read_bytes()
    finally:
        rendered.unlink(missing_ok=True)
    (folder / "SYNTHETIC-delivered-Word-render.pdf").write_bytes(word_pdf)
    direct_folder = folder / "word-render"
    direct_folder.mkdir()
    comparisons = []
    with fitz.open(stream=package.pdf, filetype="pdf") as pdf, fitz.open(stream=word_pdf, filetype="pdf") as word:
        assert 1 <= len(pdf) <= 40, (label, "Unexpected/unbounded page count", len(pdf))
        assert len(pdf) == len(word), (label, "Delivered Word page count differs")
        pdf_text = "\n".join(page.get_text() for page in pdf)
        for expected in required_text:
            assert expected in pdf_text, (label, "Missing PDF text", expected)
        for excluded in forbidden_text:
            assert excluded not in pdf_text, (label, "Old caption remains in PDF text", excluded)
        for number, (pdf_page, word_page) in enumerate(zip(pdf, word), 1):
            page_image = pdf_page.get_pixmap(dpi=150, colorspace=fitz.csRGB, alpha=False)
            word_image = word_page.get_pixmap(dpi=150, colorspace=fitz.csRGB, alpha=False)
            page_image.save(str(folder / f"pdf-page-{number:02}.png"))
            word_image.save(str(direct_folder / f"page-{number:02}.png"))
            equal = ((page_image.width, page_image.height) == (word_image.width, word_image.height)
                     and page_image.samples == word_image.samples)
            comparisons.append(equal)
        pages = len(pdf)
    assert all(comparisons), (label, "Repeat rendering of delivered Word differs from final PDF")
    (folder / "pdf-text.txt").write_text(pdf_text)
    return {
        "case": label, "route": route, "snapshot_revision": saved.revision,
        "saved_fingerprint": draft.fingerprint,
        "design_pin": {"template": draft.profile.template, "imported_from": draft.profile.imported_from},
        "sources": originals,
        "docx": {"bytes": len(package.docx), "sha256": sha256(package.docx).hexdigest()},
        "pdf": {"bytes": len(package.pdf), "sha256": sha256(package.pdf).hexdigest(), "pages": pages},
        "delivered_word_render": {"bytes": len(word_pdf), "sha256": sha256(word_pdf).hexdigest(),
                                  "same_backend_each_page_matches_pdf": comparisons},
        "editable_tables": len(document.tables), "picture_pages_or_occurrences": pictures,
        "required_editable_text": required_text, "excluded_package_and_pdf_text": forbidden_text,
        "required_visible_graphic_captions": required_graphic_captions,
        "excluded_visible_graphic_captions": forbidden_text if required_graphic_captions else (),
        "visual_inspection": "Required for every pdf-page-*.png and corresponding delivered Word render. Not performed by this script.",
    }


def export_source_page_outputs(output):
    output.mkdir(parents=True, exist_ok=True)
    executable = shutil.which("libreoffice") or shutil.which("soffice")
    if executable is None:
        raise RuntimeError("Source-page output evidence requires the actual LibreOffice renderer.")
    manifest = {
        "started_utc": datetime.now(timezone.utc).isoformat(),
        "source_commit": _command(["git", "rev-parse", "HEAD"]),
        "source_tree": _command(["git", "rev-parse", "HEAD^{tree}"]),
        "renderer": _command([executable, "--version"]),
        "rasterizer": "PyMuPDF " + fitz.VersionBind,
        "python": platform.python_version(), "platform": platform.platform(),
        "font": _command(["fc-match", "Arial"]), "dpi": 150,
        "synthetic_only": True, "cases": [],
        "limits": ["AppTest is not native browser delivery", "No owner original or approved master",
                   "No production writes or business sends", "Not Microsoft Word acceptance",
                   "Raster captions require actual page inspection; XML/text absence does not prove omission",
                   "The guided fixture reuses a repository catalog contract label with invented site/source data"],
    }

    def flush():
        (output / "manifest.json").write_text(json.dumps(manifest, indent=2) + "\n")

    # These existing AppTest fixtures import sibling test helpers. Resolve only
    # this checkout's tests, never a copied proof directory with stale fixtures.
    tests_path = str(Path(__file__).resolve().parents[1] / "tests")
    sys.path.insert(0, tests_path)
    try:
        from tests.test_source_caption_manager_journey import prepared_app, page
        from tests.test_monthly_report_upload_ui import button, current_draft
        from tests.test_source_page_occurrences import source_page
    finally:
        sys.path.remove(tests_path)
    flush()
    with TemporaryDirectory(prefix="synthetic-source-output-") as temporary, MonkeyPatch.context() as monkeypatch:
        temporary = Path(temporary)
        monkeypatch.setattr(web_ui, "device_token", lambda cookies: "e" * 32)
        app = prepared_app(monkeypatch, temporary / "caption-data")
        before = current_draft(app)
        original_caption = page(before).asset_captions[0]
        raw = sources.source_bytes(before.profile, before.sources[0])
        for label, caption in (("corrected-caption", "Verified utility chart for September 2026"),
                               ("explicit-blank-caption", "")):
            button(app, "Manage source files").click().run()
            next(widget for widget in app.text_input if widget.label == "Page caption").set_value(caption)
            button(app, "Include page").click().run()
            assert page(current_draft(app)).asset_captions == (caption,)
            assert not pending_asset_indexes(page(current_draft(app)), current_draft(app).sources)
            button(app, "Done with source files").click().run()
            draft, saved, package = _generated_downloads(app, button, current_draft)
            manifest["cases"].append(_record(output, label, draft, saved, package,
                required_text=[caption] if caption else [],
                forbidden_text=[original_caption] + (["Verified utility chart for September 2026"] if not caption else []),
                pictures=1, route="Guided AppTest Include, Done, Create Word and PDF, real shared save and generated download bytes"))
            assert sources.source_bytes(draft.profile, draft.sources[0]) == raw
            flush()

        profile, first_content, first = source_page(monkeypatch, temporary / "ownership-data",
            variant="A", caption="Retired imported chart")
        _, second_content, second = source_page(monkeypatch, temporary / "ownership-data",
            variant="B", caption="Independent source B chart")
        assert first.asset_hashes == second.asset_hashes
        originals = tuple(sources.source_bytes(profile, content.source)
                          for content in (first_content, second_content))
        human = approve_all_assets(ResolvedBlock(first.key, "This month", text="Independent operator context is retained.",
            asset_hashes=first.asset_hashes, asset_captions=("Operator supplied chart",),
            asset_provenance=((first.asset_hashes[0], ()),)))
        combined = merge_report_block(merge_report_block(human, first, None, None), second, None, None)
        assert len(combined.asset_hashes) == 3
        sections = tuple(replace(section, included=section.key == "activity") for section in default_sections())
        initial = ReportDraft(profile, ReportPeriod(2026, 10), "Synthetic Reviewer", sections,
            (ResolvedBlock("activity_summary", "This month", text="Synthetic meter readings total 3240 gallons."), combined),
            sources=(first_content.source, second_content.source), synthetic=True)
        library.save_snapshot(initial, expected_revision=0, entered_editor=initial.prepared_by)
        remaining = retire_contributions(combined, (first,))
        assert remaining.asset_captions == ("Operator supplied chart", "Independent source B chart")
        final = replace(initial, blocks=(initial.blocks[0], remaining))
        saved = library.save_snapshot(final, expected_revision=1, entered_editor=final.prepared_by)
        loaded = library.load_snapshot(profile.contract, profile.key, final.period)
        assert loaded == saved
        assert tuple(sources.source_bytes(profile, source) for source in final.sources) == originals
        package = generate_report(loaded.draft, acknowledged_fingerprint=loaded.draft.fingerprint,
            asset_loader=lambda reference: library.read_asset(profile.contract, profile.key, reference))
        manifest["cases"].append(_record(output, "equal-byte-occurrences", loaded.draft, loaded, package,
            required_text=["3240 gallons", "Independent operator context is retained."],
            forbidden_text=["Retired imported chart"], pictures=2,
            required_graphic_captions=["Operator supplied chart", "Independent source B chart"],
            route="Real synthetic PDF ingest and ownership merge/retirement helpers, actual saved revisions 1/2 and product generation"))
        flush()
    manifest["completed_utc"] = datetime.now(timezone.utc).isoformat()
    manifest["status"] = "Generated and structurally checked; every-page visual inspection remains required"
    flush()


if __name__ == "__main__":
    export_source_page_outputs(Path("test-artifacts/monthly-synthetic/source-pages"))
