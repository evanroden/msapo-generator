"""Actual hidden shared-fixture browser journey, using invented documents only.

Runs run_web.py unchanged on localhost:8521 with a disposable EPC_DATA_DIR.
Each engine uploads through the real form and takes an independent first
attempt. Expected facts are authored here, never derived from the importer.
No Render traffic, credentials, authentic files, or external actions are used.
"""
from __future__ import annotations

from hashlib import sha256
from io import BytesIO
import ipaddress
import json
import os
from pathlib import Path
import re
import secrets
import subprocess
import sys
import tempfile
import time
import traceback
from urllib.parse import urlparse
import urllib.request


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
OUT = ROOT / "test-artifacts" / "shared-fixture-testing"
ORIGIN = "http://127.0.0.1:8521"
URL = ORIGIN + "/?workflow=monthly&monthly_qa=fixtures"
MARKER = "shared-fixture-dom-synthetic-only.marker"
GUARD_LOG = "blocked-external.jsonl"
SOURCE_TEXT = "Inspected the synthetic pump in October 2026; follow-up is open."
MANUAL_TEXT = "My manually corrected synthetic Activity; follow-up remains open."
REENTRY_TEXT = "Synthetic re-entry after deliberate clearing; follow-up remains open."
EXPECTATION_ONLY = "EXPECTATION_ONLY_CANARY_NOT_SOURCE_CONTENT"


def _expected_frontend(static_dir=None):
    """Independently validate pinned inputs/reconstructed output before reading."""
    import streamlit
    from scripts import patch_streamlit_frontend as compat
    root = static_dir or Path(streamlit.__file__).parent / "static"
    # This read-only API reconstructs all81 expected assets from STATIC pins,
    # exact known repairs/config and edges, then compares every selected blob.
    # A selected filename or a digest of unvalidated installed bytes is not proof.
    selected = compat.selected_frontend(root)
    raw = (root / "static/js" / selected["entry"]).read_bytes()
    if sha256(raw).hexdigest() != selected["patched_entry_sha256"]:
        raise RuntimeError("Pinned primary frontend changed after validation.")
    return selected["entry"], raw


def _expected_overlay(static_dir=None):
    import streamlit
    from scripts import patch_streamlit_frontend as compat
    root = static_dir or Path(streamlit.__file__).parent / "static"
    selected = compat.selected_frontend(root)
    raw = (root / "static/js" / selected["overlay_bundle"]).read_bytes()
    if sha256(raw).hexdigest() != selected["overlay_sha256"]:
        raise RuntimeError("Pinned overlay frontend changed after validation.")
    return selected["overlay_bundle"], raw


def _expected_virtualizer(static_dir=None):
    import streamlit
    from scripts import patch_streamlit_frontend as compat
    root = static_dir or Path(streamlit.__file__).parent / "static"
    selected = compat.selected_frontend(root)
    if not selected["virtualizer_resize"]:
        raise RuntimeError("The pinned frontend lacks virtual-list resize compatibility.")
    raw = (root / "static/js" / selected["virtualizer_bundle"]).read_bytes()
    if sha256(raw).hexdigest() != selected["virtualizer_sha256"]:
        raise RuntimeError("Pinned virtual-list frontend changed after validation.")
    return selected["virtualizer_bundle"], raw


def _asset_response_receipt(response, expected):
    """Hash one actual browser response, without another request or raw JS logs."""
    entry, expected_bytes = expected
    receipt = {"entry": Path(urlparse(response.url).path).name,
               "status": response.status, "expected_sha256": sha256(expected_bytes).hexdigest()}
    declared = int(response.headers.get("content-length", "0"))
    if response.status != 200 or declared != len(expected_bytes) or declared > 1024 * 1024:
        return {**receipt, "matches": False, "error": "unexpected response status or size"}, b""
    raw = response.body()
    if len(raw) > 1024 * 1024:
        return {**receipt, "matches": False, "error": "oversized frontend response"}, b""
    receipt.update({"size": len(raw), "sha256": sha256(raw).hexdigest(),
                    "matches": receipt["entry"] == entry and raw == expected_bytes})
    return receipt, raw


def _frontend_response_receipt(response, expected):
    from scripts import patch_streamlit_frontend as compat
    receipt, raw = _asset_response_receipt(response, expected)
    if receipt.get("matches"):
        receipt.update({"full_scope_predicate_count": raw.count(compat.REPAIRED),
                        "original_predicate_count": raw.count(compat.ORIGINAL)})
    return receipt


def _overlay_response_receipt(response, expected):
    """Check only the Tooltip module naturally requested by this real page."""
    from scripts.streamlit_overlay_resize import OVERLAY_ORIGINAL, OVERLAY_REPAIRED
    receipt, raw = _asset_response_receipt(response, expected)
    receipt.update({"resize_hold_count": raw.count(OVERLAY_REPAIRED),
                    "original_resize_hook_count": raw.count(OVERLAY_ORIGINAL)})
    receipt["matches"] = bool(receipt.get("matches") and receipt["resize_hold_count"] == 1
                              and receipt["original_resize_hook_count"] == 0)
    return receipt


def _virtualizer_response_receipt(response, expected):
    from scripts.streamlit_virtualizer_resize import VIRTUALIZER_ORIGINAL, VIRTUALIZER_REPAIRED
    receipt, raw = _asset_response_receipt(response, expected)
    receipt.update({"resize_hold_count": raw.count(VIRTUALIZER_REPAIRED),
                    "original_resize_hook_count": raw.count(VIRTUALIZER_ORIGINAL)})
    receipt["matches"] = bool(receipt.get("matches") and receipt["resize_hold_count"] == 1
                              and receipt["original_resize_hook_count"] == 0)
    return receipt


def _navigation_decode_cancel(engine, role, stage, reloading, message):
    """Only an image decode aborted by OUR deliberate primary-page reload.

    WebKit can reject an image that the previous page was decoding while that
    document unloads. Record it as a navigation cancellation, never as a
    successful/current preview. Any other stage, browser or error still fails.
    """
    return (engine == "playwright-webkit" and role == "primary"
            and reloading and stage.endswith(": browser reload")
            and message == "EncodingError: Loading was canceled.")


REPORT = {"synthetic_only": True, "actual_entrypoint": "run_web.py",
          "route": "?workflow=monthly&monthly_qa=fixtures",
          "browser_evidence": "Chromium and Playwright WebKit, not shipping Safari",
          "access_policy": "shared with every visitor; no authentication",
          "checks": [], "browsers": {}, "failures": []}


def serve():
    """Harness-only outgoing guard around the unchanged production entrypoint."""
    base = Path(os.environ.get("EPC_DATA_DIR", "")).resolve()
    if (base.parent != Path(tempfile.gettempdir()).resolve()
            or not base.name.startswith("shared-fixture-dom-synthetic-")
            or not (base / MARKER).is_file()
            or (base / MARKER).read_text() != "shared-fixture-dom-v1"):
        raise RuntimeError("This harness requires its own marked disposable storage root")
    guard_log = base / GUARD_LOG

    def denied(name):
        with guard_log.open("a", encoding="utf-8") as stream:
            stream.write(json.dumps({"blocked": name}) + "\n")
        raise AssertionError("Shared-fixture browser QA blocked " + name)

    def audit(event, args):
        if event not in {"socket.connect", "socket.getaddrinfo"}:
            return
        address = args[1] if event == "socket.connect" else args[0]
        host = address[0] if isinstance(address, tuple) else address
        if host in {"localhost", "127.0.0.1", "::1", None}:
            return
        try:
            if ipaddress.ip_address(str(host)).is_loopback:
                return
        except ValueError:
            pass
        denied(event + ":non-loopback")

    sys.addaudithook(audit)
    # Fail on invocation, even if a missing API key would prevent charging.
    # These guards live only in this harness process, never production code.
    from app import monthly_report_ai, ocr, smartsheet
    monthly_report_ai.request_json = lambda *_a, **_k: denied("external monthly reader/writer")
    ocr._ocr_request = lambda *_a, **_k: denied("external OCR")
    smartsheet.submit_po = lambda *_a, **_k: denied("Smartsheet submission")
    from streamlit.web import cli
    sys.argv = ["streamlit", "run", str(ROOT / "run_web.py"),
                "--server.address", "127.0.0.1", "--server.port", "8521",
                "--server.headless", "true", "--browser.gatherUsageStats", "false"]
    cli.main()


def check(browser, name, passed, details=None):
    REPORT["checks"].append({"browser": browser, "name": name,
                             "passed": bool(passed), "details": details})
    print(("PASS " if passed else "FAIL ") + browser + " " + name, flush=True)
    if not passed:
        raise AssertionError(name + ": " + str(details))


def synthetic_docx():
    from docx import Document
    document = Document()
    document.add_heading("Synthetic Site A", 0)
    document.add_paragraph("October 2026 Monthly Report")
    document.add_heading("2 Monthly Activity Summary", 1)
    document.add_paragraph(SOURCE_TEXT)
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def run_directory(base, cookie):
    # Read-only diagnosis of the production namespacing protocol; no namespace
    # or report is created by this helper or prepopulated for the browser.
    token = sha256(("shared-fixture:" + cookie).encode("ascii")).hexdigest()[:32]
    scope = sha256(("monthly-qa-v1:" + token).encode("ascii")).hexdigest()[:32]
    return base / "_enfra_monthly_report_qa_v1" / scope


def journal(root):
    paths = list(root.rglob("working.json")) if root.is_dir() else []
    values = [json.loads(path.read_text()) for path in paths]
    return [value for value in values if value.get("kind") == "browser_active_work_v1"]


def durable_activity(page, root, expected):
    # Poll the actual atomic journal acknowledgment, not a fixed rendering
    # sleep, a repeated user action, or values injected into session state.
    deadline = time.monotonic() + 45
    while time.monotonic() < deadline:
        values = journal(root)
        if len(values) == 1:
            current = values[0]
            text = next((block.get("text", "") for block in current["draft"]["blocks"]
                         if block["key"] == "activity_summary"), "")
            if text == expected:
                return {"revision": current["revision"], "generation": current["generation"],
                        "activity_sha256": sha256(text.encode()).hexdigest()}
        page.wait_for_timeout(100)
    # Failure must retain the precise last persisted source of truth and a
    # bounded browser-action receipt. Visible textarea state alone is not a
    # server commit, and retrying the operator click would conceal this bug.
    summary = []
    for current in journal(root)[:2]:
        last_text = next((b.get("text", "") for b in current.get("draft", {}).get("blocks", ())
                          if b.get("key") == "activity_summary"), "")
        summary.append({"revision": current.get("revision"),
                        "generation": current.get("generation"),
                        "activity_sha256": sha256(last_text.encode()).hexdigest()})
    try:
        client = page.evaluate("""() => ({
          native_clear_events: window.__qaNativeClearEvents || [],
          script_state: document.querySelector('[data-testid="stApp"]')
            ?.getAttribute("data-test-script-state") || "missing",
        })""")
    except Exception as exc:
        client = {"browser_diagnostic_error": type(exc).__name__}
    raise AssertionError("The exact Activity edit did not reach its isolated durable journal; "
                         + json.dumps({"last_committed": summary, "browser": client}, sort_keys=True)[:1000])


def idle(page):
    page.wait_for_function("""() => document.querySelector('[data-testid="stApp"]')
      ?.getAttribute('data-test-script-state') === 'notRunning'""", timeout=45000)


def capture(page, name, label):
    page.screenshot(path=str(OUT / f"{name}-{label}.png"), timeout=30000)
    observed = page.evaluate("""() => ({
      headings:[...document.querySelectorAll('h1,h2,h3')].map(n=>n.textContent),
      alerts:[...document.querySelectorAll('[data-testid="stAlert"]')].map(n=>n.textContent),
      textareas:[...document.querySelectorAll('textarea')].map(n=>({label:n.getAttribute('aria-label'),value:n.value})),
      comboboxes:[...document.querySelectorAll('[role="combobox"]')].map(n=>({label:n.getAttribute('aria-label'),tag:n.tagName,value:n.value,text:n.textContent})),
      scriptState:document.querySelector('[data-testid="stApp"]')?.getAttribute('data-test-script-state')
    })""")
    REPORT["browsers"].setdefault(name, {}).setdefault("states", {})[label] = observed
    (OUT / "results.json").write_text(json.dumps(REPORT, indent=2))


def exercise(browser_type, name, base):
    from playwright.sync_api import expect
    expected_frontend = _expected_frontend()
    expected_overlay = _expected_overlay()
    browser = browser_type.launch(headless=True)
    contexts, attempts, errors = [], [], []
    error_details, canceled_on_reload = [], []
    frontend_delivery = []
    overlay_delivery = []
    phase = {"name": "initial navigation"}
    navigation = {"primary": False, "coworker": False}
    reload_count = 0

    def session():
        cookie = secrets.token_hex(16)
        context = browser.new_context(viewport={"width": 1440, "height": 1000}, accept_downloads=True)
        contexts.append(context)
        context.add_cookies([{"name": "epc_device_id", "value": cookie, "url": ORIGIN}])
        context.add_init_script("""window.__sharedFixtureShareAttempts=[];
          Object.defineProperty(navigator,'share',{configurable:true,value:()=>{
            window.__sharedFixtureShareAttempts.push('share');
            return Promise.reject(new Error('Sharing is forbidden in synthetic QA'));
          }});""")

        def route(request):
            parsed = urlparse(request.request.url)
            if parsed.scheme in {"data", "blob"} or (parsed.hostname == "127.0.0.1" and parsed.port == 8521):
                request.continue_()
            else:
                attempts.append(parsed.scheme + ":" + (parsed.hostname or "external"))
                request.abort()
        context.route("**/*", route)
        page = context.new_page()
        role = "primary" if len(contexts) == 1 else "coworker"
        def observe_frontend(response):
            if (role != "primary" or frontend_delivery
                    or not re.fullmatch(r"/static/js/index\.(?:enfra-frontend-[0-9a-f]{12}|bE3scgDe)\.js",
                                        urlparse(response.url).path)):
                return
            # Mark consumption before reading so nested Playwright dispatch
            # cannot read the module twice. Failures stay visible in results.
            frontend_delivery.append({"matches": False, "error": "response observation incomplete"})
            try:
                frontend_delivery[0] = _frontend_response_receipt(response, expected_frontend)
            except Exception as exc:
                frontend_delivery[0] = {"matches": False, "error": type(exc).__name__ + ": " + str(exc)[:300]}
        page.on("response", observe_frontend)
        def observe_overlay(response):
            if role != "primary" or overlay_delivery or not Path(urlparse(response.url).path).name.startswith("Tooltip.DJg7muty"):
                return
            overlay_delivery.append({"matches": False, "error": "response observation incomplete"})
            try:
                overlay_delivery[0] = _overlay_response_receipt(response, expected_overlay)
            except Exception as exc:
                overlay_delivery[0] = {"matches": False, "error": type(exc).__name__ + ": " + str(exc)[:300]}
        page.on("response", observe_overlay)
        page.set_default_timeout(20000)
        def browser_exception(error):
            # Keep the entire first browser error receipt. A browser image
            # decode canceled by a deliberately initiated page unload has no
            # live document to display it on. It is a *recorded cancellation*,
            # not a current-page success or a license to ignore other errors.
            message = str(error)
            expected = _navigation_decode_cancel(name, role, phase["name"],
                                                   navigation[role], message)
            detail = {"browser": role, "stage": phase["name"],
                      "message": message[:500],
                      "stack": str(getattr(error, "stack", ""))[:1600],
                      "classification": "intentional_navigation_cancel" if expected else "app_error"}
            error_details.append(detail)
            if expected:
                canceled_on_reload.append(detail)
            else:
                errors.append(message)
        page.on("pageerror", browser_exception)
        page.goto(URL, wait_until="domcontentloaded")
        expect(page.get_by_role("heading", name="Shared report testing", exact=True)).to_be_visible(timeout=60000)
        if role == "primary":
            check(name, "browser received the exact current repaired Streamlit entry",
                  len(frontend_delivery) == 1 and frontend_delivery[0].get("matches"), frontend_delivery)
            check(name, "browser received the exact repaired native overlay position module",
                  len(overlay_delivery) == 1 and overlay_delivery[0].get("matches"), overlay_delivery)
        return page, cookie

    page = None
    try:
        page, cookie = session()
        root = run_directory(base, cookie)
        expect(page.get_by_text(re.compile(r"^Anyone with this link can upload, view and download"))).to_be_visible()
        original = synthetic_docx()
        digest = sha256(original).hexdigest()
        friendly = "Synthetic partial report " + name
        page.get_by_text("Upload an actual report", exact=True).click()
        upload = page.locator('[data-testid="stFileUploader"]').filter(
            has=page.get_by_text("Original Word report", exact=True))
        upload.locator('input[type="file"]').set_input_files({
            "name": name + "-invented-partial.docx",
            "mimeType": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "buffer": original})
        page.get_by_role("combobox", name="Scenario", exact=True).click()
        page.get_by_role("option", name="Half-finished report, correct month", exact=True).click()
        for label, value in {
            "Fixture name": friendly, "Exact contract": "Synthetic Testing Contract",
            "Source sites": "Synthetic Site A", "Target sites": "Synthetic Site A",
            "Source month (YYYY-MM)": "2026-10", "Target month (YYYY-MM)": "2026-10",
            "Scenario notes": "Independently authored: preserve current incomplete pump work; do not invent completion.",
            "Expected results": "Activity | include | synthetic pump\nActivity | exclude | " + EXPECTATION_ONLY,
        }.items():
            page.get_by_role("textbox", name=label, exact=True).fill(value)
        partial = page.get_by_role("checkbox", name="For the partial-report scenario, I confirm this is the actual half-finished file", exact=True)
        partial.locator("xpath=ancestor::label[1]").click()
        expect(partial).to_be_checked()
        page.get_by_role("button", name="Save shared fixture", exact=True).click()
        shared = page.get_by_role("combobox", name="Shared fixture", exact=True)
        expect(shared).to_be_visible(timeout=45000)
        # Streamlit 1.61.1 renders the selected label as the accessible
        # combobox input value, not as a text node. Assert that precise value
        # and the immutable receipt, rather than waiting on get_by_text.
        expect(shared).to_have_value(re.compile(r"^" + re.escape(friendly) + r" · [0-9a-f]{8}$"), timeout=45000)
        manifests = [json.loads(path.read_text()) for path in
                     (base / "_enfra_shared_fixtures_v1").glob("*/manifest.json")]
        matches = [value for value in manifests if value["metadata"]["friendly_name"] == friendly]
        check(name, "selected upload has one durable immutable receipt", len(matches) == 1
              and matches[0]["original_sha256"] == digest
              and matches[0]["original_size"] == len(original))
        saved = matches[0]
        check(name, "selected fixture identity matches durable receipt",
              shared.input_value() == friendly + " · " + saved["fixture_id"][:8])
        REPORT["browsers"].setdefault(name, {})["uploaded_receipt"] = saved
        page.get_by_text("Confirmed source, target and expected results", exact=True).click()
        expect(page.get_by_text("Original SHA256: " + digest, exact=True)).to_be_visible()
        check(name, "real browser upload saved immutable synthetic original", True,
              {"sha256": digest, "size": len(original), "scenario": "AUTH03", "target": "Synthetic Testing Contract / Synthetic Site A / 2026-10"})

        def download_original(label, current=page):
            # The selected combobox can show its new value before Streamlit
            # commits the rerun that changes the binary download resource.
            # Require the exact fixture/hash-scoped native action to be
            # mounted before the one real download. This is not a retry or a
            # weaker byte assertion: an obsolete or mismatched action fails.
            scope = saved["fixture_id"] + "_" + digest[:16]
            expected_widget = current.locator(".st-key-_shared_fixture_download_" + scope)
            expect(expected_widget).to_be_visible(timeout=45000)
            idle(current)
            with current.expect_download(timeout=30000) as event:
                current.get_by_role("button", name="Download unchanged original", exact=True).click()
            downloaded = event.value
            path = OUT / f"{name}-{label}-original.docx"
            downloaded.save_as(path)
            check(name, label + " original bytes and hash unchanged", path.read_bytes() == original
                  and sha256(path.read_bytes()).hexdigest() == digest)

        download_original("uploaded")
        coworker, coworker_cookie = session()
        coworker_root = run_directory(base, coworker_cookie)
        coworker.get_by_role("combobox", name="Shared fixture", exact=True).click()
        coworker.get_by_role("option", name=re.compile(r"^" + re.escape(friendly) + r" · ")).click()
        expect(coworker.get_by_role("button", name="Download unchanged original", exact=True)).to_be_visible()
        check(name, "second browser sees shared fixture without inheriting first browser run",
              coworker.get_by_role("textbox", name="Activity summary", exact=True).count() == 0
              and root != coworker_root and not journal(coworker_root))
        coworker_generation = (coworker_root / ".enfra-qa-namespace").read_text()
        download_original("coworker", coworker)

        phase["name"] = "start real Starting report importer"
        page.get_by_role("button", name="Start fresh test with this fixture", exact=True).click()
        actor = page.get_by_role("textbox", name="Your name", exact=True)
        expect(actor).to_be_visible(timeout=60000)
        staged_original = root / "_enfra_fixture_start_v1" / "starting-report" / "source.docx"
        lineage = json.loads(staged_original.with_name("lineage.json").read_text())
        check(name, "fresh Starting report uses byte-identical copy with confirmed target lineage",
              staged_original.read_bytes() == original
              and lineage["original_sha256"] == digest
              and lineage["metadata"]["scenario"] == "AUTH03"
              and lineage["metadata"]["target"] == {
                  "contract": "Synthetic Testing Contract", "sites": ["synthetic-site-a"],
                  "month": "2026-10", "site_titles": ["Synthetic Site A"]})
        actor.fill("Synthetic Reviewer")
        actor.press("Tab")
        proceed = page.get_by_role("button", name="Continue to this month’s updates", exact=True)
        expect(proceed).to_be_enabled(timeout=45000)
        capture(page, name, "production-starting-review")
        phase["name"] = "enter ordered report/editor after import"
        proceed.click()
        activity = page.get_by_role("textbox", name="Activity summary", exact=True)
        expect(activity).to_be_visible(timeout=90000)
        expect(activity).to_have_value(re.compile(re.escape(SOURCE_TEXT)), timeout=30000)
        check(name, "production Starting report imported source text rather than expectations",
              SOURCE_TEXT in activity.input_value() and EXPECTATION_ONLY not in activity.input_value())
        capture(page, name, "imported-current-report")

        # The native Streamlit multiline control presents Ctrl+Enter as its
        # explicit apply gesture. Exercise keyboard clearing that way instead
        # of depending only on WebKit's sometimes-interrupted Tab/blur event.
        # Neither action is accepted until its EXACT empty journal is durable.
        for label, wording, mode in (
            ("manual-edit", MANUAL_TEXT, "tab"),
            ("intentional-keyboard-clear", "", "apply"),
            ("reentry", REENTRY_TEXT, "tab"),
            ("explicit-report-clear", "", "button"),
            ("reentry-after-report-clear", REENTRY_TEXT, "tab"),
        ):
            phase["name"] = label + ": native edit"
            if mode == "button":
                # After the prior page reload a visible textarea can arrive
                # before Streamlit finishes wiring actions. Do not dispatch
                # an intentional Clear into that transient render; require
                # the actual app's idle signal, then one native button click.
                idle(page)
                clear = page.get_by_role("button", name="Clear Activity for this report", exact=True)
                expect(clear).to_be_enabled()
                # Observe one trusted native browser click without altering
                # event dispatch or retrying an unacknowledged business action.
                clear.evaluate("""element => {
                  window.__qaNativeClearEvents = [];
                  for (const name of ["pointerdown", "click"]) {
                    element.addEventListener(name, event => {
                      window.__qaNativeClearEvents.push({
                        name, trusted: event.isTrusted
                      });
                    }, {capture: true, once: true});
                  }
                }""")
                clear.click()
            else:
                activity.fill(wording)
                expect(activity).to_have_value(wording)
                if mode == "apply":
                    activity.press("ControlOrMeta+Enter")
                else:
                    activity.press("Tab")
            receipt = durable_activity(page, root, wording)
            idle(page)
            phase["name"] = label + ": browser reload"
            reload_count += 1
            navigation["primary"] = True
            try:
                page.reload(wait_until="domcontentloaded")
                expect(activity).to_be_visible(timeout=60000)
                expect(activity).to_have_value(wording, timeout=30000)
                # The first page is gone; an old image decode may reject in
                # WebKit during unload. The new page MUST still hydrate and
                # acknowledge the precise current Activity before continuing.
                idle(page)
            finally:
                navigation["primary"] = False
            phase["name"] = label + ": loaded and idle"
            check(name, label + " survives ordinary refresh with exact canonical draft", True, receipt)
            check(name, label + " does not reopen Starting report or duplicate editor",
                  page.get_by_role("textbox", name="Your name", exact=True).count() == 0 and activity.count() == 1)
        capture(page, name, "reentry-after-refresh")
        phase["name"] = "coworker reload"
        coworker.reload(wait_until="domcontentloaded")
        expect(coworker.get_by_role("heading", name="Shared report testing", exact=True)).to_be_visible(timeout=60000)
        check(name, "edited report remains isolated from coworker", not journal(coworker_root)
              and coworker.get_by_role("textbox", name="Activity summary", exact=True).count() == 0)
        old_generation = (root / ".enfra-qa-namespace").read_text()
        phase["name"] = "reset primary workspace"
        page.get_by_role("button", name="Reset this test workspace", exact=True).click()
        expect(activity).to_have_count(0, timeout=45000)
        idle(page)
        check(name, "scoped reset replaces run generation and removes editable run", not journal(root)
              and (root / ".enfra-qa-namespace").read_text() != old_generation)
        check(name, "reset leaves coworker namespace unchanged",
              (coworker_root / ".enfra-qa-namespace").read_text() == coworker_generation)
        # Selection is intentionally disposable on reset. Select the original
        # explicitly before proving its shared bytes still exist unchanged.
        phase["name"] = "select and download original after reset"
        page.get_by_role("combobox", name="Shared fixture", exact=True).click()
        page.get_by_role("option", name=re.compile(r"^" + re.escape(friendly) + r" · ")).click()
        download_original("after-reset")
        capture(page, name, "reset-original-retained")
        check(name, "zero external browser requests", not attempts, attempts)
        check(name, "no browser errors outside explicit navigation unload", not errors, errors)
        check(name, "canceled image decode receipts are bounded by actual reloads",
              len(canceled_on_reload) <= reload_count,
              {"canceled": len(canceled_on_reload), "actual_reloads": reload_count})
        check(name, "no sharing action", all(not p.evaluate("window.__sharedFixtureShareAttempts")
                                             for c in contexts for p in c.pages))
    except Exception:
        if page is not None:
            try:
                capture(page, name, "failure")
            except Exception as capture_error:
                REPORT["browsers"].setdefault(name, {})["capture_error"] = str(capture_error)[:500]
        REPORT["browsers"].setdefault(name, {})["traceback"] = traceback.format_exc()
        raise
    finally:
        REPORT["browsers"].setdefault(name, {}).update({
            "external_attempts": attempts, "browser_errors": errors,
            "frontend_delivery": frontend_delivery,
            "overlay_delivery": overlay_delivery,
            "page_error_details": error_details,
            "navigation_cancellations": canceled_on_reload})
        for context in contexts:
            context.close()
        browser.close()


def main():
    from playwright.sync_api import sync_playwright
    OUT.mkdir(parents=True, exist_ok=True)
    completed = False
    with tempfile.TemporaryDirectory(prefix="shared-fixture-dom-synthetic-") as temp:
        base = Path(temp)
        (base / MARKER).write_text("shared-fixture-dom-v1")
        (base / GUARD_LOG).write_text("")
        env = dict(os.environ, EPC_DATA_DIR=temp, ANTHROPIC_API_KEY="", OPENAI_API_KEY="",
                   SMARTSHEET_API_MODE="disabled", PYTHONDONTWRITEBYTECODE="1")
        with (OUT / "streamlit.log").open("w") as log:
            process = subprocess.Popen([sys.executable, str(Path(__file__).resolve()), "--serve"],
                                       cwd=ROOT, env=env, stdout=log, stderr=subprocess.STDOUT)
            try:
                deadline = time.monotonic() + 30
                while time.monotonic() < deadline:
                    if process.poll() is not None:
                        raise RuntimeError("Production shared-fixture app stopped during startup")
                    try:
                        with urllib.request.urlopen(ORIGIN + "/_stcore/health", timeout=1) as response:
                            if response.status == 200:
                                break
                    except OSError:
                        time.sleep(.1)
                else:
                    raise RuntimeError("Production shared-fixture app did not become healthy")
                with sync_playwright() as playwright:
                    for engine, name in ((playwright.chromium, "chromium"), (playwright.webkit, "playwright-webkit")):
                        try:
                            exercise(engine, name, base)
                        except Exception as exc:
                            REPORT["failures"].append(name + ": " + type(exc).__name__ + ": " + str(exc)[:500])
                allowed = {MARKER, GUARD_LOG, "_enfra_monthly_report_qa_v1", "_enfra_shared_fixtures_v1",
                           "_enfra_shared_fixture_runs_v1", "_enfra_fixture_run_admission_v1"}
                check("server", "all persisted data stays in explicit shared/test buckets",
                      not (set(p.name for p in base.iterdir()) - allowed), sorted(p.name for p in base.iterdir()))
                blocked = [json.loads(line) for line in (base / GUARD_LOG).read_text().splitlines()]
                check("server", "zero external reader, OCR, submission or network invocations", not blocked, blocked)
                completed = True
            finally:
                process.terminate()
                try:
                    process.wait(timeout=20)
                except subprocess.TimeoutExpired:
                    process.kill()
                    process.wait(timeout=5)
                REPORT["result"] = "passed" if completed and not REPORT["failures"] else "failed"
                (OUT / "results.json").write_text(json.dumps(REPORT, indent=2))
    if REPORT["failures"]:
        raise AssertionError("Shared fixture browser failures: " + "; ".join(REPORT["failures"]))


if __name__ == "__main__":
    serve() if "--serve" in sys.argv else main()
