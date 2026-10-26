"""Synthetic-only setup and read-only receipts for the real import controls."""
from hashlib import sha256
from html import escape
from io import BytesIO
import json
from pathlib import Path
from zipfile import ZipFile, ZIP_STORED

from scripts.workspace_history_observations import namespace

EDITED = 'September 2026 synthetic browser correction retained.'
ORIGINAL = 'September 2026 synthetic initial work costs $500.'


def make_source(directory):
    """Real DOCX exceeds the unchanged native-art budget, not upload admission."""
    from docx import Document
    from app.monthly_report_model import synthetic_profiles
    directory = Path(directory); directory.mkdir(parents=True, exist_ok=True)
    doc = Document()
    cover_paragraph=doc.add_paragraph()
    # Two different real drawing representations trigger production import_candidate.
    from PIL import Image
    from lxml import etree as ET
    from copy import deepcopy
    from docx.shared import Inches
    from docx.oxml.ns import qn
    def image_bytes(color, format='PNG'):
        raw=BytesIO(); Image.new('RGB',(160,80),color).save(raw,format); return raw.getvalue()
    from docx.oxml import OxmlElement
    wps='{http://schemas.microsoft.com/office/word/2010/wordprocessingShape}'
    drawing=OxmlElement('w:drawing')
    inline=ET.SubElement(drawing,qn('wp:inline'))
    ET.SubElement(inline,qn('wp:extent'),cx='1828800',cy='280670')
    graphic=ET.SubElement(inline,qn('a:graphic'))
    data=ET.SubElement(graphic,qn('a:graphicData'),uri=wps[1:-1])
    shape=ET.SubElement(data,wps+'wsp'); ET.SubElement(shape,wps+'cNvSpPr',txBox='1')
    props=ET.SubElement(shape,wps+'spPr');ET.SubElement(props,qn('a:prstGeom'),prst='rect')
    box=ET.SubElement(ET.SubElement(shape,wps+'txbx'),qn('w:txbxContent'))
    paragraph=ET.SubElement(box,qn('w:p'))
    ET.SubElement(ET.SubElement(paragraph,qn('w:r')),qn('w:t')).text='INSERT CLIENT LOGO'
    ET.SubElement(shape,wps+'bodyPr')
    picture=doc.add_picture(BytesIO(image_bytes('navy')),width=Inches(2))
    image_drawing=picture._inline.getparent();run=image_drawing.getparent();old=run.getparent()
    paragraph.append(run);old.getparent().remove(old)
    mc='{http://schemas.openxmlformats.org/markup-compatibility/2006}'
    wrapper=ET.Element(mc+'AlternateContent',nsmap={'wps':wps[1:-1]})
    for text in ('Prepared by: Synthetic Editor','September 2026',
                 synthetic_profiles()[0].facilities[0].title+'\nOperations and Maintenance Monthly Review'):
        text_frame=deepcopy(drawing)
        text_box=next(text_frame.iter(qn('w:txbxContent')))
        for child in list(text_box): text_box.remove(child)
        for line in text.splitlines():
            para=ET.SubElement(text_box,qn('w:p'))
            ET.SubElement(ET.SubElement(para,qn('w:r')),qn('w:t')).text=line
        cover_paragraph.add_run()._r.append(text_frame)
    cover_paragraph.add_run()._r.append(wrapper)
    ET.SubElement(wrapper,mc+'Choice',Requires='wps').append(drawing)
    fallback=ET.SubElement(wrapper,mc+'Fallback');fallback.append(deepcopy(drawing))
    rid,_=doc.part.get_or_add_image(BytesIO(image_bytes('green','JPEG')))
    next(fallback.iter(qn('a:blip'))).set(qn('r:embed'),rid)
    doc.add_heading('Activity Summary', 1)
    doc.add_paragraph(ORIGINAL)
    path = directory/'synthetic-25mib-source.docx'; doc.save(path)
    with ZipFile(path, 'a') as package:
        package.writestr('synthetic-padding.bin', b'\0'*(25*1024*1024), compress_type=ZIP_STORED)
    return path


def install_template(base, token):
    from app import monthly_report_designs as designs, monthly_report_library as library
    from docx import Document
    from app import monthly_report_testing as testing
    _, identity, generation = namespace(base, token)
    if (testing.scope(), testing.generation()) != (identity, generation):
        raise ValueError('Template setup requires its current isolated namespace.')
    if designs.master_state()['revision']:
        raise ValueError('Synthetic template already exists; refuse replacement.')
    doc=Document(); doc.add_paragraph('Independent synthetic ENFRA template')
    from app.monthly_report_model import default_sections
    for section in default_sections():
        doc.add_heading(section.title, 1); doc.add_paragraph('Master placeholder')
    path=library._root()/'qa-fallback-master.docx'
    path.parent.mkdir(parents=True,exist_ok=True); doc.save(path)
    designs.install_master(path,actor='Synthetic Editor',expected_revision=0)


def render(base, token, field):
    import streamlit as st
    from app import monthly_report_testing as testing, monthly_report_library as library
    from app import monthly_report_designs as designs
    from app.monthly_report_model import synthetic_profiles, ReportPeriod
    from app.monthly_report_section_ui import render_section_setup
    testing.data_dir(base)
    _,identity,generation=namespace(base,token)
    selected=synthetic_profiles()[0]; period=ReportPeriod(2026,9)
    st.title('Synthetic whole-template fallback acceptance')
    if not designs.master_state()['revision']:
        if st.button('Install synthetic independent template'):
            install_template(base, token); st.rerun()
        return
    saved=(library.load_imported_draft(selected.contract,selected.key)
           if any(p.key==selected.key for p in library.list_profiles(selected.contract)) else None)
    if saved is None:
        render_section_setup(selected.contract,period,'Synthetic Editor',field,identity=selected)
    else:
        st.success('Synthetic reviewed report saved')
        if st.button('Build saved synthetic DOCX'):
            from app.monthly_report_docx import _build_docx
            payload=_build_docx(saved,asset_loader=lambda ref: library.read_asset(selected.contract,selected.key,ref))
            (library._root()/'qa-fallback-final.docx').write_bytes(payload)
    saved=(library.load_imported_draft(selected.contract,selected.key)
           if any(p.key==selected.key for p in library.list_profiles(selected.contract)) else None)
    record={'scope':identity,'generation':generation,'saved':saved is not None,
            'profile_count':len(library.list_profiles(selected.contract))}
    plans=next((value for key,value in st.session_state.items() if key.endswith('_section_plans')), {})
    record['staged_texts']=[text for plan in plans.values() for text in plan.get('texts',{}).values()]
    choice=next((value for key,value in st.session_state.items() if key.endswith('_independent_template')), {})
    record['template_chosen']=bool(choice)
    if saved:
        original=library._profile_path(selected.contract,selected.key)/'imports'/(saved.profile.imported_from+'.docx')
        record.update(template=saved.profile.template,source_sha256=sha256(original.read_bytes()).hexdigest(),
                      source_bytes=original.stat().st_size,saved_texts=[b.text for b in saved.blocks])
    output=library._root()/'qa-fallback-final.docx'
    if output.exists():
        from docx import Document
        raw=output.read_bytes(); doc=Document(BytesIO(raw))
        text='\n'.join(p.text for p in doc.paragraphs)
        with ZipFile(BytesIO(raw)) as package:
            xml=package.read('word/document.xml').decode()
            output_media={sha256(package.read(n)).hexdigest() for n in package.namelist() if n.startswith('word/media/')}
        with ZipFile(original) as package:
            source_media={sha256(package.read(n)).hexdigest() for n in package.namelist() if n.startswith('word/media/')}

        record['output']={'sha256':sha256(raw).hexdigest(),'bytes':len(raw),
                          'edited_text_present':EDITED in text,'pricing_absent':'$500' not in xml,
                          'unread_source_art_absent':bool(source_media) and source_media.isdisjoint(output_media) and 'AlternateContent' not in xml}
    st.html('<span hidden data-enfra-template-fallback="'+escape(json.dumps(record),quote=True)+'"></span>')
