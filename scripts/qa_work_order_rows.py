"""Hosted-only native work-order grid journey; invented data and disposable disk.

The fixture uses production editor, active-work journal and native output code.
No browser widget state, callback, or report result is injected by this driver.
"""
from __future__ import annotations

from dataclasses import replace
from hashlib import sha256
from io import BytesIO
import json
import os
from pathlib import Path
import secrets
import subprocess
import sys
import tempfile
import time
from urllib.parse import urlparse
from urllib.request import urlopen

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))
OUT = ROOT / "test-artifacts" / "work-order-rows"
PORT = 8537
URL = f"http://127.0.0.1:{PORT}"
COLUMNS = ("Work orders", "Jan", "Feb", "Mar", "Apr", "May", "Jun")
BLANK = ("",) * 7
ROWS = (
    ("Received", "1", "2", "3", "4", "5", "6"),
    ("", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"),
    ("", "7", "8", "9", "10", "11", "12"),
    ("Closed", "13", "14", "15", "16", "17", "18"),
    ("", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"),
)
EDIT = "Received edited before insertion"
PREFIX = "synthetic_work_order_rows"
RECEIPT = "data-work-order-receipt"


def scenario(base):
    """Known compact data, with an unchanged synthetic eight-row native source."""
    from docx import Document
    from app.monthly_report_import import inspect_docx
    from app.monthly_report_model import (
        Facility, ReportDraft, ReportPeriod, ReportProfile, ReportTable,
        ResolvedBlock, default_sections,
    )
    path = Path(base) / "synthetic-work-orders.docx"
    if not path.exists():
        source = Document()
        source.add_heading("MONTHLY ACTIVITY SUMMARY", 1)
        source.add_paragraph("WORK ORDERS")
        rows = (COLUMNS, BLANK, *ROWS[:3], BLANK, *ROWS[3:])
        table = source.add_table(rows=len(rows), cols=7)
        table.style = "Table Grid"
        for row, values in zip(table.rows, rows):
            for cell, value in zip(row.cells, values):
                cell.text = value
        source.save(path)
    inspection = inspect_docx(path)
    item = next(item for item in inspection.items if item.kind == "table")
    reference = f"docx:{inspection.sha256}:{item.id}"
    profile = ReportProfile("Synthetic row contract", "synthetic-row-site", "Synthetic Site",
                            (Facility("synthetic-row-site", "Synthetic Site"),))
    sections = (next(section for section in default_sections() if section.key == "activity"),)
    draft = ReportDraft(profile, ReportPeriod(2026, 10), "Synthetic Editor", sections,
        (ResolvedBlock("work_orders", "This month", extra_tables=(ReportTable(COLUMNS, ROWS, reference),)),
         ResolvedBlock("activity_summary", "This month", text="Synthetic unrelated note remains unchanged.")),
        synthetic=True)
    return path, draft


def row_state(draft):
    block = next(block for block in draft.blocks if block.key == "work_orders")
    table, = block.extra_tables
    return {"columns": list(table.columns), "rows": [list(row) for row in table.rows],
            "reference": table.reference, "fingerprint": draft.fingerprint,
            "unrelated": replace(draft, blocks=tuple(b for b in draft.blocks if b.key != "work_orders")).fingerprint}


def expected_rows(stage):
    rows = [list(row) for row in ROWS]
    if stage != "initial":
        rows[0][0] = EDIT
    if stage in {"above", "restored"}:
        rows.insert(0, list(BLANK))
    if stage == "restored":
        rows.insert(4, list(BLANK))
    return rows


def assert_receipt(value, rows, baseline=None):
    assert value["working"] == value["durable"], "Native result has not reached the disk journal"
    assert value["working"]["columns"] == list(COLUMNS)
    assert value["working"]["rows"] == rows
    assert value["source_unchanged"] is True
    assert value["revision"] > 0
    if baseline:
        for key in ("reference", "unrelated"):
            assert value["working"][key] == baseline["working"][key]
        assert value["generation"] == baseline["generation"]


def assert_outputs(word, pdf, rows):
    from docx import Document
    import pymupdf
    document = Document(BytesIO(word))
    assert len(document.tables) == 1
    actual = [[cell.text for cell in row.cells] for row in document.tables[0].rows]
    assert actual == [list(COLUMNS), *rows]
    assert all(len(row) == 7 for row in actual)
    assert not any("__work_order_row_" in cell or "Insert blank" in cell for row in actual for cell in row)
    with pymupdf.open(stream=pdf, filetype="pdf") as rendered:
        text = "\n".join(page.get_text() for page in rendered)
        assert EDIT in " ".join(text.split())
        assert "Insert blank" not in text and "__work_order_row_" not in text
        tables = [(table.row_count, table.col_count) for page in rendered for table in page.find_tables().tables]
        assert tables == [(8, 7)]
        pages = len(rendered)
    return {"docx_sha256": sha256(word).hexdigest(), "pdf_sha256": sha256(pdf).hexdigest(),
            "pdf_pages": pages, "native_table": [8, 7]}


OBSERVE = r"""() => {
  window.rowNativeEvents=[];
  const ids=new WeakMap();let next=0;
  const describe=n=>{
    if(!n)return null;
    if(!ids.has(n))ids.set(n,++next);
    const r=n.getBoundingClientRect?.();
    return {id:ids.get(n),tag:n.tagName,role:n.getAttribute?.('role'),
      label:n.getAttribute?.('aria-label'),testid:n.getAttribute?.('data-testid'),
      text:n.closest?.('[role="menuitem"]')?.textContent?.slice(0,90),
      x:r?.x,y:r?.y,width:r?.width,height:r?.height,connected:n.isConnected};
  };
  const snapshot=()=>({menu:describe(document.querySelector('[data-testid="stDataFrameButtonActionMenu"]')),
    grid:describe(document.querySelector('[data-testid="data-grid-canvas"]')),
    input:describe(document.querySelector('.gdg-input')),
    app:document.querySelector('[data-testid="stApp"]')?.getAttribute('data-test-script-state')});
  const record=value=>{
    window.rowNativeEvents.push({...value,at:performance.now(),wall_ms:performance.timeOrigin+performance.now()});
    if(window.rowNativeEvents.length>500)window.rowNativeEvents.shift();
  };
  // This fixture contains only invented data. Observe the whole pointer/focus
  // sequence so a menu closing before pointerup cannot hide the event's target.
  // Node identities and lifecycle snapshots are diagnostic only: no input,
  // widget state, event dispatch, synchronization wait or retry is introduced.
  for(const type of ['pointerdown','pointerup','mousedown','mouseup','click','dblclick','input','keydown','focusin','focusout']) {
    document.addEventListener(type,event=>{
      const n=event.target;
      record({type,trusted:event.isTrusted,key:event.key||null,
        label:n.getAttribute?.('aria-label')||null,testid:n.getAttribute?.('data-testid')||null,
        target:describe(n),item:describe(n.closest?.('[role="menuitem"]')),
        clientX:event.clientX,clientY:event.clientY,...snapshot()});
    },true);
  }
  let previous='';
  new MutationObserver(()=>{
    const current=snapshot(),value=JSON.stringify(current);
    if(value!==previous){previous=value;record({type:'dom-state',...current});}
  }).observe(document,{subtree:true,childList:true,attributes:true,
    attributeFilter:['data-test-script-state']});
}"""


def settled(page, rows, record, stage, baseline=None):
    page.wait_for_function("""wanted => {
      const nodes=document.querySelectorAll('[data-work-order-receipt]');
      if(!nodes.length)return false;
      const v=JSON.parse(nodes[nodes.length-1].getAttribute('data-work-order-receipt'));
      return JSON.stringify(v.working.rows)===JSON.stringify(wanted) &&
        JSON.stringify(v.working)===JSON.stringify(v.durable) &&
        document.querySelector('[data-testid="stApp"]')?.getAttribute('data-test-script-state')==='notRunning';
    }""", arg=rows, timeout=25000)
    value = json.loads(page.locator(f'[{RECEIPT}]').last.get_attribute(RECEIPT))
    assert_receipt(value, rows, baseline)
    record["steps"].append({"stage": stage, "receipt": value,
                            "native_events": page.evaluate("window.rowNativeEvents || []")})
    return value


def native_grid_target(page, row, *, action=True):
    """Locked Streamlit 1.61.1: 32px row marker, fixed 160px pinned Row column.

    Remaining column widths may stretch; target the start of the first data
    column instead of guessing its full width. Row/header height is 2.1875rem.
    Glide draws into a canvas underlay, but its sibling native scroller receives
    pointer input. Target that surface with ordinary locator actionability,
    never force a click through it onto the visual-only canvas.
    A missing/covered grid fails; no rescue click or DOM event injection follows.
    """
    grid = page.get_by_test_id("stDataFrame")
    canvas = grid.get_by_test_id("data-grid-canvas")
    surface = grid.locator(".dvn-scroller.stDataFrameGlideDataEditor")
    # The server receipt can precede React's lazy-loaded grid mount. Await the
    # actual controls before the first gesture, not a delay or a rescue click.
    for node in (grid, canvas, surface):
        node.wait_for(state="visible", timeout=15000)
    assert grid.count() == canvas.count() == surface.count() == 1, "Ambiguous native grid"
    surface.scroll_into_view_if_needed()
    box, target = canvas.bounding_box(), surface.bounding_box()
    assert box and target, "Native grid is not visible"
    rem = page.evaluate("parseFloat(getComputedStyle(document.documentElement).fontSize)")
    height = round(2.1875 * rem)
    point = {"x": box["x"] + 32 + (80 if action else 160 + 25) - target["x"],
             "y": box["y"] + height * (row + 1.5) - target["y"]}
    assert 0 < point["x"] < target["width"] and 0 < point["y"] < target["height"], "Native cell is clipped"
    assert surface.evaluate("""(node, point) => {
      const rect = node.getBoundingClientRect();
      const hit = document.elementFromPoint(rect.left + point.x, rect.top + point.y);
      return !!hit && (hit === node || node.contains(hit));
    }""", point), "Native cell is covered"
    return surface, point


def menu(page, row):
    from playwright.sync_api import expect
    surface, point = native_grid_target(page, row)
    surface.click(position=point)
    options = page.get_by_role("menu", name="Button action menu")
    expect(options).to_be_visible()
    expect(options.get_by_role("menuitem", name="Insert blank above", exact=True)).to_be_visible()
    expect(options.get_by_role("menuitem", name="Insert blank below", exact=True)).to_be_visible()
    return options


def toggle_table_visibility(page, *, visible):
    from playwright.sync_api import expect
    checkbox = page.get_by_role("checkbox", name="Show synthetic table", exact=True)
    expect(checkbox).to_have_count(1)
    expect(checkbox).to_be_checked(checked=not visible)
    # Streamlit's accessible input is clipped to 1px; people click its native
    # visible label. No .check(), force click, keyboard fallback or retry.
    label = checkbox.locator("xpath=ancestor::label")
    expect(label).to_have_count(1)
    expect(label.get_by_role("checkbox")).to_have_count(1)
    expect(label).to_be_visible()
    label.click()
    expect(checkbox).to_be_checked(checked=visible)


def source_checkout_dirty():
    """Do not mistake this driver's generated evidence for changed source."""
    command = ["git", "status", "--porcelain", "--untracked-files=all", "--", "."]
    try:
        evidence = OUT.relative_to(ROOT)
    except ValueError:
        pass  # Evidence outside the checkout cannot affect its status.
    else:
        if evidence.parts:
            command.append(":(top,exclude,literal)" + evidence.as_posix())
    return bool(subprocess.check_output(command, cwd=ROOT, text=True).strip())


def exercise(browser, name, *, delayed_ack=False):
    from playwright.sync_api import expect
    out = OUT / name
    out.mkdir(parents=True, exist_ok=True)
    record = {"synthetic_only": True, "steps": [], "errors": [], "external": [], "wire": [],
              "browser": name, "tested_checkout": subprocess.check_output(["git", "rev-parse", "HEAD"], cwd=ROOT, text=True).strip(),
              "source_checkout_dirty": source_checkout_dirty()}
    context = browser.new_context(viewport={"width": 1200, "height": 900}, accept_downloads=True)
    context.add_cookies([{"name": "epc_device_id", "value": secrets.token_hex(16), "url": URL}])
    if delayed_ack:
        context.add_cookies([{"name": "synthetic_acknowledge_edit", "value": "1", "url": URL}])
    context.add_init_script("(" + OBSERVE + ")()")
    def route(request):
        parsed = urlparse(request.request.url)
        if parsed.scheme in {"data", "blob"} or parsed.hostname == "127.0.0.1" and parsed.port == PORT:
            request.continue_()
        else:
            record["external"].append(parsed.hostname)
            request.abort()
    context.route("**/*", route)
    page = context.new_page()
    page.set_default_timeout(15000)
    page.on("pageerror", lambda error: record["errors"].append(str(error)))
    def wire(direction, payload):
        """Observe the native protocol; never modify or resend a frame."""
        from streamlit.proto.BackMsg_pb2 import BackMsg
        from streamlit.proto.ForwardMsg_pb2 import ForwardMsg
        if not isinstance(payload, bytes):
            return
        message = BackMsg() if direction == "sent" else ForwardMsg()
        message.ParseFromString(payload)
        event = {"direction": direction, "at": time.monotonic()}
        if direction == "sent" and message.HasField("rerun_script"):
            event["widgets"] = [{"id": widget.id,
                "value_type": widget.WhichOneof("value"),
                "value": widget.string_trigger_value.data if widget.HasField("string_trigger_value")
                         else widget.string_value if widget.HasField("string_value") else None}
                for widget in message.rerun_script.widget_states.widgets]
        elif direction == "received" and message.HasField("delta"):
            delta = message.delta
            if not delta.HasField("new_element") or not delta.new_element.HasField("dataframe"):
                return
            grid = delta.new_element.dataframe
            event.update(grid_id=grid.id, action_ids=dict(grid.button_click_widgets))
        elif direction == "received" and message.HasField("script_finished"):
            event["script_finished"] = message.script_finished
        else:
            return
        record["wire"].append(event)
    def observe_socket(socket):
        socket.on("framesent", lambda payload: wire("sent", payload))
        socket.on("framereceived", lambda payload: wire("received", payload))
    page.on("websocket", observe_socket)
    gate = {"armed": False, "finished": False, "buffer": [], "received": 0}
    if delayed_ack:
        def route_socket(socket):
            from streamlit.proto.ForwardMsg_pb2 import ForwardMsg
            remote = socket.connect_to_server()
            gate["socket"] = socket
            def received(payload):
                if not gate["armed"]:
                    socket.send(payload)
                    return
                # Delay unchanged server bytes only. Native browser inputs go
                # straight to the production callback without modification.
                gate["buffer"].append(payload)
                gate["received"] += 1
                if isinstance(payload, bytes):
                    message = ForwardMsg(); message.ParseFromString(payload)
                    if message.HasField("script_finished") and message.script_finished == 0:
                        gate["finished"] = True
            remote.on_message(received)
        page.route_web_socket("**/_stcore/stream", route_socket)
    try:
        page.goto(URL, wait_until="domcontentloaded")
        baseline = settled(page, expected_rows("initial"), record, "compact unchanged")
        menu(page, 0)
        page.keyboard.press("Escape")
        expect(page.get_by_role("menu", name="Button action menu")).to_have_count(0)
        settled(page, expected_rows("initial"), record, "cancel is no-op", baseline)

        surface, point = native_grid_target(page, 0, action=False)
        surface.dblclick(position=point)
        field = page.locator(".gdg-input")
        expect(field).to_have_value("Received")
        field.fill(EDIT)  # No Enter/Tab: the next native menu click must deliver this edit.
        gate["armed"] = delayed_ack
        options = menu(page, 0)
        if delayed_ack:
            deadline = time.monotonic() + 25
            while not gate["finished"] and time.monotonic() < deadline:
                page.wait_for_timeout(10)  # Pump the controlled response gate.
            assert gate["finished"], "Server edit acknowledgement did not reach the response gate"
            record["delayed_ack"] = {"server_finished_before_action": True,
                                      "buffered_messages": gate["received"]}
        options.get_by_role("menuitem", name="Insert blank above", exact=True).click()
        if delayed_ack:
            gate["armed"] = False
            for payload in gate["buffer"]:
                gate["socket"].send(payload)
            gate["buffer"].clear()
        settled(page, expected_rows("above"), record, "pending edit and above", baseline)

        # A real rapid repeated pointer attempt is not fabricated wire replay.
        # The menu retires on first activation; exactly one row may be inserted.
        menu(page, 3).get_by_role("menuitem", name="Insert blank below", exact=True).dblclick(delay=40)
        settled(page, expected_rows("restored"), record, "repeated below attempt inserts once", baseline)
        page.screenshot(path=str(out / "restored-desktop.png"), full_page=True)
        old_menu = menu(page, 0).element_handle()
        toggle_table_visibility(page, visible=False)
        expect(page.get_by_test_id("data-grid-canvas")).to_have_count(0)
        assert old_menu is not None and not old_menu.evaluate("node => node.isConnected")
        settled(page, expected_rows("restored"), record, "retired menu cannot survive unmount", baseline)
        toggle_table_visibility(page, visible=True)
        settled(page, expected_rows("restored"), record, "reopen without replay", baseline)
        page.set_viewport_size({"width": 545, "height": 757})
        page.reload(wait_until="domcontentloaded")
        settled(page, expected_rows("restored"), record, "narrow disk reload", baseline)
        assert page.evaluate("document.documentElement.scrollWidth<=document.documentElement.clientWidth+2")
        menu(page, 0)
        page.keyboard.press("Escape")
        settled(page, expected_rows("restored"), record, "narrow menu reachable and canceled", baseline)
        page.get_by_role("button", name="Generate synthetic Word and PDF", exact=True).click()
        word = pdf = None
        for kind in ("Word", "PDF"):
            button = page.get_by_role("button", name="Download synthetic " + kind, exact=True)
            expect(button).to_be_visible(timeout=120000)
            with page.expect_download(timeout=120000) as download:
                button.click()
            path = out / ("restored.docx" if kind == "Word" else "restored.pdf")
            download.value.save_as(path)
            if kind == "Word":
                word = path.read_bytes()
            else:
                pdf = path.read_bytes()
        record["outputs"] = assert_outputs(word, pdf, expected_rows("restored"))
        settled(page, expected_rows("restored"), record, "export leaves saved table unchanged", baseline)
        page.screenshot(path=str(out / "restored-narrow.png"), full_page=True)
        assert not record["errors"] and not record["external"]
        record["passed"] = True
    except Exception as error:
        record["first_failure"] = str(error)
        try:
            page.screenshot(path=str(out / "first-failure.png"), full_page=True)
            (out / "first-failure.html").write_text(page.content())
        except Exception as capture_error:
            record["capture_error"] = str(capture_error)
        raise
    finally:
        try:
            record["native_events"] = page.evaluate("window.rowNativeEvents || []")
        except Exception as observation_error:
            record["observation_error"] = str(observation_error)
        (out / "result.json").write_text(json.dumps(record, indent=2) + "\n")
        context.close()


def stop_server(process):
    process.terminate()
    try:
        process.wait(timeout=10)
    except subprocess.TimeoutExpired:
        process.kill()
        process.wait(timeout=5)


def main():
    from playwright.sync_api import sync_playwright
    OUT.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="work-order-rows-synthetic-") as directory:
        base = Path(directory)
        (base / "monthly-dom-synthetic-only.marker").write_text("synthetic-monthly-dom-v1", encoding="ascii")
        env = dict(os.environ, EPC_DATA_DIR=str(base), EPC_MONTHLY_DOM_FIXTURE="1",
                   EPC_WORK_ORDER_ROW_FIXTURE="1", ANTHROPIC_API_KEY="", SMARTSHEET_API_MODE="disabled")
        with (OUT / "streamlit.log").open("w") as log:
            process = subprocess.Popen([sys.executable, "-m", "streamlit", "run",
                "tests/fixtures/workspaces/work_order_rows_app.py", "--server.address", "127.0.0.1",
                "--server.port", str(PORT), "--server.headless", "true", "--browser.gatherUsageStats", "false"],
                cwd=ROOT, env=env, stdout=log, stderr=subprocess.STDOUT)
            try:
                deadline = time.monotonic() + 30
                while time.monotonic() < deadline:
                    if process.poll() is not None:
                        raise RuntimeError("Synthetic row server stopped")
                    try:
                        with urlopen(URL + "/_stcore/health", timeout=1) as response:
                            if response.status == 200:
                                break
                    except OSError:
                        time.sleep(.1)
                else:
                    raise RuntimeError("Synthetic row server did not become healthy")
                failures = []
                with sync_playwright() as playwright:
                    for name, engine in (("chromium", playwright.chromium), ("webkit", playwright.webkit)):
                        for delayed_ack in (False, True):
                            case = "delayed-ack" if delayed_ack else "ordinary"
                            browser = engine.launch(headless=True)
                            try:
                                exercise(browser, f"{name}-{case}", delayed_ack=delayed_ack)
                            except Exception as error:
                                failures.append(f"{name}-{case}: {error}")
                            finally:
                                browser.close()
                if failures:
                    raise AssertionError("; ".join(failures))
            finally:
                stop_server(process)
    print("PASS native work-order actions, disk reload and actual Word/PDF on both engines")


if __name__ == "__main__":
    main()
