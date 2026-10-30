"""Hosted-only native exact-warning acceptance in the guarded full editor.

No local browser run, hidden state injection, paid reading, retry-after-failure,
forced clicks or approval shortcuts. Native user decisions own every repair.
"""
from contextlib import ExitStack
from hashlib import sha256
from io import BytesIO
import json,os,re,secrets,shutil,subprocess,sys,tempfile,time,traceback
from pathlib import Path
from urllib.parse import urlparse
from urllib.request import urlopen
from zipfile import ZipFile
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT))
from scripts.workspace_warning_routes import SCENARIOS,CAPTIONS,PENDING,MISSING_FACT,BAD_DATE_FACT
from scripts.qa_about_logo_refresh import controller_environment
from scripts.qa_import_receipts import wait_receipt
from scripts.monthly_sequential_navigation import navigate_sections
from scripts.workspace_history_observations import namespace
OUT=ROOT/'test-artifacts/warning-routes';PORT=8558;URL=f'http://127.0.0.1:{PORT}'
ENGINES=('chromium','playwright-webkit')
MODES={'wide-pointer':{'width':1365,'height':900},'narrow-keyboard':{'width':390,'height':844}}
REPORT={'synthetic_only':True,'physical_device':False,'cases':{},'failures':[]}


def required_cases():return {e+'/'+m+'/'+s for e in ENGINES for m in MODES for s in SCENARIOS}


def all_cases_passed(report):
    return (set(report.get('cases',{}))==required_cases() and not report.get('failures')
        and all(v.get('result')=='passed' and v.get('outputs_verified') and v.get('reopen_verified')
                for v in report['cases'].values()))


def receipt(page,expected=None):return wait_receipt(page,'data-warning-routes',expected,timeout=120000)


def source_date_control(page,owner):
    """Resolve the exact source's labeled native date widget, never a sibling."""
    from playwright.sync_api import expect
    from scripts.qa_financial_pages import source_date_widget
    expect(owner).to_have_count(1)
    widget=source_date_widget(page,owner)
    expect(widget).to_have_count(1)
    control=widget.get_by_test_id('stDateInputField')
    expect(control).to_have_count(1)
    return control


def select_native_option(page,control,label,*,keyboard=False):
    """Choose one exact option through its native control, not option existence."""
    from playwright.sync_api import expect
    if keyboard:
        expect(control).to_have_count(1);expect(control).to_be_enabled()
        control.focus();expect(control).to_be_focused();control.press('ArrowDown')
        expect(control).to_have_attribute('aria-expanded','true')
        popup_id=control.get_attribute('aria-controls')
        assert popup_id and not any(c.isspace() for c in popup_id),'The control must own one native listbox.'
        popup=page.locator('[role="listbox"][id='+json.dumps(popup_id)+']')
        expect(popup).to_have_count(1);expect(popup).to_be_visible()
        option=popup.get_by_role('option',name=label,exact=True)
        expect(option).to_have_count(1);expect(option).to_be_enabled()
        target=option.get_attribute('id')
        identities=popup.get_by_role('option').evaluate_all('nodes=>nodes.map(node=>node.id)')
        active=control.get_attribute('aria-activedescendant')
        assert identities and all(identities) and len(set(identities))==len(identities)
        assert target in identities and active in identities,'The highlighted option must belong to this exact listbox.'
        steps=identities.index(target)-identities.index(active)
        # One finite route derived from the observed native list order. No
        # search-value no-op, guessed extra ArrowDown, retry or pointer fallback.
        for _ in range(abs(steps)):control.press('ArrowDown' if steps>0 else 'ArrowUp')
        expect(control).to_have_attribute('aria-activedescendant',target)
        control.press('Enter')
    else:
        control.click()
        option=page.get_by_role('option',name=label,exact=True)
        expect(option).to_have_count(1);option.click()
    page.keyboard.press('Escape')



def select_source_caption(page,manager,prefix,source,*,keyboard=False):
    """A selection receipt alone cannot authorize typing into another owner."""
    from playwright.sync_api import expect
    selected_key=prefix+'_evidence_selected'
    choose=manager.get_by_role('combobox',name='File to review',exact=True)
    expect(choose).to_have_count(1);expect(choose).to_be_visible();expect(choose).to_be_enabled()
    previous=receipt(page);previous_value=choose.input_value()
    select_native_option(page,choose,source['filename'],keyboard=keyboard)
    expect(choose).to_have_value(source['filename'])
    owner=manager.locator('.st-key-'+prefix+'_evidence_'+source['id']+'_caption_1')
    expect(owner).to_have_count(1)
    caption=owner.get_by_role('textbox',name='Page caption',exact=True)
    expect(caption).to_have_count(1);expect(caption).to_be_visible();expect(caption).to_be_enabled()
    expect(owner.locator('xpath=ancestor-or-self::*[@data-stale="true"]')).to_have_count(0)
    # A genuine same-owner selection can be a no-op. A changed native value or
    # receipt owner, however, requires a newer settled receipt, never the old ACK.
    changed=(previous_value!=source['filename'] or previous['selected'].get(selected_key)!=source['id'])
    state=after(page,previous['sequence']) if changed else receipt(page)
    assert state['sequence']>=previous['sequence']+(1 if changed else 0),'Stale source-selection receipt'
    assert state['selected'].get(selected_key)==source['id'],'Caption receipt belongs to another source'
    # Recheck the native owner after sampling; do not fill an unscoped caption.
    expect(choose).to_have_value(source['filename']);expect(caption).to_be_visible()
    return caption,state



def open_native_expander(expander,label,activate):
    """Open only the owned outer summary; nested summaries are not siblings."""
    from playwright.sync_api import expect
    expect(expander).to_have_count(1)
    details=expander.locator(':scope > details')
    summary=details.locator(':scope > summary')
    expect(details).to_have_count(1);expect(summary).to_have_count(1)
    expect(summary.get_by_text(label,exact=True)).to_have_count(1)
    if not details.evaluate('node=>node.open'):activate(summary)
    expect(details).to_have_attribute('open','')
    return expander


def save_shared_version(page,prefix,commit):
    from playwright.sync_api import expect
    history=page.locator('.st-key-'+prefix+'_history_open')
    expect(history).to_have_count(1)
    # This tracked native disclosure reruns the app. Its ACK precedes looking
    # for the child Save action; opening it is not itself a shared save.
    open_native_expander(history.locator(':scope > [data-testid="stExpander"]'),
                         'Shared report history',commit)
    return commit(history.get_by_role('button',name='Save shared version',exact=True))


def shared_picture_occurrences(state,source,*,reviewed):
    """Read-only exact source/page and destination/occurrence review oracle."""
    from app.monthly_report_library import block_from_dict
    from app.monthly_report_asset_review import asset_fingerprint
    from app.monthly_report_page_reference import page_identity
    assert state['scenario']=='shared'
    owners=[s for s in state['working']['sources'] if s['id']==source['id']]
    assert len(owners)==1 and owners[0]['sha256']==source['sha256']
    assert owners[0]['filename']==source['filename']
    assert state['originals'][source['id']]==source['sha256']
    occurrences=[]
    for key,caption in zip(('improvements','vendor_reports'),CAPTIONS):
        raw=state['working']['blocks'][key]
        assert raw['key']==key and raw['source']=='This month'
        assert len(raw['asset_hashes'])==len(raw['asset_instances'])==1
        assert raw['asset_captions']==[caption]
        instance,=raw['asset_instances']
        assert instance['key'] and instance['asset']==raw['asset_hashes'][0]
        pages=[page_identity(ref) for ref in instance['references'] if not ref.startswith('asset-source-context:')]
        assert pages==[(source['id'],1)],'Picture occurrence belongs to another source/page'
        fingerprint=asset_fingerprint(block_from_dict(raw),0)
        assert raw['client_asset_reviews']==([fingerprint] if reviewed else []),'Each occurrence needs its own explicit current review'
        occurrences.append((key,instance['key'],instance['asset'],caption,tuple(instance['references']),fingerprint))
    assert len({(x[0],x[1]) for x in occurrences})==2
    assert occurrences[0][2]==occurrences[1][2] and occurrences[0][-1]!=occurrences[1][-1]
    return tuple(occurrences)


def confirm_shared_pictures(page,prefix,source,before,commit):
    """One explicit native consent for the two displayed independent pictures."""
    from playwright.sync_api import expect
    expected=shared_picture_occurrences(before,source,reviewed=False)
    expect(page.locator('.epc-section-progress')).to_have_attribute('data-section','report_review')
    owner=page.locator('[class*="st-key-'+prefix+'_picture_component_"]')
    expect(owner).to_have_count(1)
    expect(owner.locator('xpath=ancestor-or-self::*[@data-stale="true"]')).to_have_count(0)
    form=owner.locator('.picture-review');expect(form).to_have_count(1)
    expect(form.locator('figure')).to_have_count(2);expect(form.locator('img')).to_have_count(2)
    urls=[]
    for number,(title,caption) in enumerate(zip(('Progress photos','Vendor service reports'),CAPTIONS),1):
        title=f'{number}. {title} · picture 1'
        figure=form.locator('figure').filter(has=page.get_by_role('heading',name=title,exact=True))
        expect(figure).to_have_count(1)
        context=' · '.join((caption,source['filename'],source['facility'],source['service_date']))
        expect(figure.get_by_text(context,exact=True)).to_have_count(1)
        picture=figure.get_by_role('img',name=title,exact=True)
        expect(picture).to_have_count(1);expect(picture).to_be_visible()
        assert picture.evaluate('img=>img.complete&&img.naturalWidth>0')
        urls.append(picture.get_attribute('src'))
        expect(figure.get_by_role('checkbox',name='Leave '+title+' for correction',exact=True)).not_to_be_checked()
    assert urls[0] and urls[0]==urls[1],'The two displayed occurrences must retain their shared picture'
    confirm=form.get_by_role('button',name='Confirm these 2 pictures are ready (except those left for correction)',exact=True)
    state=commit(confirm)
    assert shared_picture_occurrences(state,source,reviewed=True)==expected
    assert state['working']['sources']==before['working']['sources'] and state['originals']==before['originals']
    assert state['durable']['fingerprint']==state['working']['fingerprint']
    return state


def after(page,sequence):
    page.wait_for_function('''old => {
      const node=document.querySelector('[data-warning-routes]');
      if(!node || document.querySelector('[data-testid="stApp"]')?.getAttribute('data-test-script-state')!=='notRunning')return false;
      try{return JSON.parse(node.dataset.warningRoutes).sequence>old}catch(_){return false}
    }''',arg=sequence,timeout=120000)
    return receipt(page)


def verify_outputs(directory,scenario,state,*,original,baseline_images=()):
    """Actual generated Word text/image parts and real converted PDF text/paint."""
    from docx import Document
    import pymupdf as fitz
    from PIL import Image
    raw={name:(directory/name).read_bytes() for name in ('report.docx','report.pdf')}
    for name,value in raw.items():
        assert len(value)==state['output'][name]['bytes'] and sha256(value).hexdigest()==state['output'][name]['sha256']
    document=Document(BytesIO(raw['report.docx']))
    words='\n'.join([p.text for p in document.paragraphs]+[c.text for t in document.tables for r in t.rows for c in r.cells])
    with fitz.open(stream=raw['report.pdf'],filetype='pdf') as pdf:
        assert len(pdf)>0
        printed='\n'.join(p.get_text() for p in pdf)
    normalized=lambda text:' '.join(text.split())
    assert normalized(PENDING) in normalized(words) and normalized(PENDING) in normalized(printed)
    if scenario=='ai':
        assert 'Replaced 2 seals.' in words and 'Replaced 2 seals.' in printed
        assert 'Replaced 99 seals.' not in words+printed and MISSING_FACT not in words+printed
    if scenario=='shared':
        for caption in CAPTIONS:
            assert normalized(caption) in normalized(words) and normalized(caption) in normalized(printed)
    if scenario in ('shared','financial'):
        # Read INPUT bytes captured before edits, not final model asset claims.
        from collections import Counter
        from lxml import etree
        import posixpath
        with fitz.open(stream=original,filetype='pdf') as source:
            assert len(source)==1
            expected=Image.open(BytesIO(source[0].get_pixmap(dpi=150,alpha=False).tobytes('png'))).convert('RGB')
            expected.thumbnail((1050,1350),Image.Resampling.LANCZOS)
        def identity(image):return image.size,sha256(image.convert('RGB').tobytes()).hexdigest()
        expected_images={identity(Image.open(BytesIO(data))):Image.open(BytesIO(data)).convert('RGB') for data in baseline_images}
        expected_images[identity(expected)]=expected
        wanted=Counter(identity(Image.open(BytesIO(data))) for data in baseline_images)
        wanted[identity(expected)]+=2 if scenario=='shared' else 1
        with ZipFile(BytesIO(raw['report.docx'])) as package:
            rels=etree.fromstring(package.read('word/_rels/document.xml.rels'))
            paths={r.get('Id'):posixpath.normpath(posixpath.join('word',r.get('Target',''))) for r in rels if r.get('TargetMode')!='External'}
            for name in package.namelist():
                if name.startswith('word/media/'):
                    assert identity(Image.open(BytesIO(package.read(name)))) in expected_images,'Unexpected retained Word image resource'
            xml=etree.fromstring(package.read('word/document.xml'))
            ns={'a':'http://schemas.openxmlformats.org/drawingml/2006/main','r':'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}
            actual=Counter(identity(Image.open(BytesIO(package.read(paths[node.get('{'+ns['r']+'}embed')]))))
                for node in xml.findall('.//a:blip',ns) if node.get('{'+ns['r']+'}embed') in paths)
        assert actual==wanted,'Actual Word image occurrences differ from immutable original inputs'
        actual=Counter()
        with fitz.open(stream=raw['report.pdf'],filetype='pdf') as pdf:
            for xref in range(1,pdf.xref_length()):
                if pdf.xref_get_key(xref,'Subtype')[1]=='/Image':
                    resource=pdf.extract_image(xref)
                    assert resource and identity(Image.open(BytesIO(resource['image']))) in expected_images,'Unexpected retained PDF image resource'
            for page in pdf:
                # get_image_info enumerates displayed occurrences, including
                # a reused image appearing twice; unused resources don't count.
                for info in page.get_image_info(xrefs=True):
                    assert info['xref']>0,'Unverifiable inline PDF image'
                    bounds=fitz.Rect(info['bbox'])
                    assert bounds.width>1 and bounds.height>1 and page.rect.contains(bounds),'Clipped or invisible PDF source image'
                    data=pdf.extract_image(info['xref'])['image']
                    image_id=identity(Image.open(BytesIO(data)))
                    assert image_id in expected_images,'Unknown PDF image'
                    actual[image_id]+=1
                    # A correct hidden XObject under a black/white/text overlay
                    # is not original content preserved in the visible report.
                    transform=info['transform']
                    assert transform[0]>0 and transform[3]>0 and abs(transform[1])<1e-7 and abs(transform[2])<1e-7,'Unexpected source image transform'
                    independent=BytesIO();expected_images[image_id].save(independent,format='PNG')
                    reference=fitz.open();clean=reference.new_page(width=page.rect.width,height=page.rect.height)
                    clean.insert_image(bounds,stream=independent.getvalue(),keep_proportion=False)
                    expected_paint=clean.get_pixmap(clip=bounds,dpi=90,alpha=False,colorspace=fitz.csRGB)
                    actual_paint=page.get_pixmap(clip=bounds,dpi=90,alpha=False,colorspace=fitz.csRGB)
                    assert (actual_paint.width,actual_paint.height)==(expected_paint.width,expected_paint.height)
                    # At most three RGB quantization levels per channel. No
                    # average/percentage budget that could hide a changed sum.
                    from PIL import ImageChops
                    difference=ImageChops.difference(Image.frombytes('RGB',(actual_paint.width,actual_paint.height),actual_paint.samples),
                        Image.frombytes('RGB',(expected_paint.width,expected_paint.height),expected_paint.samples))
                    reference.close()
                    assert all(high<=3 for low,high in difference.getextrema()),'Visible PDF source pixels differ from original input'
        assert actual==wanted,'Actual PDF image occurrences differ from immutable original inputs'
    if scenario=='financial':
        assert not re.search(r'vendor\s+invoice\s+total',words+'\n'+printed,re.I)
    return {name:{'sha256':sha256(value).hexdigest(),'bytes':len(value)} for name,value in raw.items()}


def exercise(engine,name,mode,scenario,base):
    from playwright.sync_api import expect
    key=name+'/'+mode+'/'+scenario;output=OUT/name/mode/scenario;output.mkdir(parents=True,exist_ok=True)
    record={'result':'running','events':[],'browser_errors':[],'external_requests':[],
            'outputs_verified':False,'reopen_verified':False,'mode':mode}
    REPORT['cases'][key]=record;keyboard=mode=='narrow-keyboard';token=secrets.token_hex(16)
    browser=engine.launch(headless=True);context=browser.new_context(viewport=MODES[mode])
    context.add_cookies([{'name':'epc_device_id','value':token,'url':URL}])
    context.tracing.start(screenshots=True,snapshots=True,sources=True)
    page=context.new_page();page.set_default_timeout(30000)
    page.on('pageerror',lambda error:record['browser_errors'].append(str(error)))
    def network(route):
        parsed=urlparse(route.request.url)
        if parsed.scheme in {'data','blob'} or parsed.hostname=='127.0.0.1' and parsed.port==PORT:route.continue_()
        else:record['external_requests'].append(route.request.url);route.abort()
    context.route('**/*',network)
    def capture(label):
        state=receipt(page);record['events'].append({'label':label,'receipt':state})
        assert not state['blocked_external_attempts'] and not record['browser_errors'] and not record['external_requests']
        page.screenshot(path=str(output/(str(len(record['events']))+'-'+label+'.png')))
        return state
    def activate(control):
        expect(control).to_have_count(1);expect(control).to_be_visible();expect(control).to_be_enabled()
        if keyboard:control.focus();expect(control).to_be_focused();control.press('Enter')
        else:control.click()
    def commit(control):
        previous=receipt(page)['sequence'];activate(control);return after(page,previous)
    def select(control,label):
        select_native_option(page,control,label,keyboard=keyboard)
    def enter_combobox(control):
        if keyboard:control.focus()
        else:control.click()
    def check(control):
        expect(control).not_to_be_checked();expect(control).to_be_enabled()
        before=receipt(page)['sequence']
        if keyboard:control.focus();control.press('Space')
        else:control.locator('xpath=ancestor::label').click()
        expect(control).to_be_checked();after(page,before)
    def focus_proof(control):
        expect(control).to_have_count(1);expect(control).to_be_focused()
        visible=control.locator('xpath=ancestor::label') if control.get_attribute('type')=='checkbox' else control
        expect(visible).to_be_visible();expect(visible).to_be_in_viewport(ratio=1)
        assert visible.evaluate('''node=>{const r=node.getBoundingClientRect();const hit=document.elementFromPoint(r.x+r.width/2,r.y+r.height/2);return !!hit&&(node===hit||node.contains(hit)||hit.contains(node))}'''),'Correction target is obscured'
    def return_review():
        back=page.locator('.enfra-correction-return:visible,.enfra-source-review-return:visible')
        expect(back).to_have_count(1);activate(back)
        expect(page.locator('.epc-section-progress')).to_have_attribute('data-section','report_review')
    def warning_action(code,block=None,*,source_page=False,ordinal=None):
        state=receipt(page);matches=[c for c in state['working']['checks'] if c['code']==code and (block is None or c['block']==block)]
        assert len(matches)==1,(code,matches)
        targets=matches[0]['targets'];assert targets
        # Explicit fixture paragraph number comes from saved structured metadata,
        # not warning text or a fallback to an arbitrary sibling.
        token=targets[ordinal] if ordinal is not None else targets[0]
        prefix=state['editor_prefix']
        if code.startswith('ai_'):suffix='_ai_correction_'+('source_' if source_page else '')+token
        elif code=='financial_page':suffix='_financial_correction_'+token
        elif code.startswith('source_'):suffix='_correct_source_'+state['working']['sources'][1]['id']+'_'+code
        else:suffix='_page_correction_'+token+'_'+code
        return page.locator('.st-key-'+prefix+suffix+' button')
    def jump(code,control,block=None,*,source_page=False,ordinal=None):
        navigate_sections(page,'report_review',keyboard=keyboard)
        commit(warning_action(code,block,source_page=source_page,ordinal=ordinal))
        focus_proof(control);capture('focused-'+code+('-source' if source_page else ''))
        return_review()
        # Deliberate second activation proves a new action remains usable; it is
        # not a retry of a failed click. The first focus assertion already passed.
        commit(warning_action(code,block,source_page=source_page,ordinal=ordinal));focus_proof(control)
    try:
        page.goto(URL,wait_until='domcontentloaded')
        select(page.get_by_role('combobox',name='Synthetic historical scenario',exact=True),scenario)
        activate(page.get_by_role('button',name='Load synthetic historical report with unresolved warnings',exact=True))
        initial=capture('unresolved-historical-report')
        assert initial['scenario']==scenario,'The installed fixture must match the requested scenario.'
        prefix=initial['editor_prefix'];old,target=initial['working']['sources']
        # The explicit fixture Load enters through production's one-shot Cover
        # landing. Prove that arrival before the first sequential action; merely
        # mounting the full editor below report identity is not an entry receipt.
        expect(page.locator('.epc-section-progress')).to_have_attribute('data-section','monthly_cover')
        expect(page.locator('.epc-section-back')).to_be_disabled()
        expect(page.locator('.epc-section-next')).to_be_in_viewport(ratio=1)
        capture('native-cover-entry')
        directory,_,_=namespace(base,token)
        original=(directory/'warning-input/source-original').read_bytes()
        assert sha256(original).hexdigest()==target['sha256']==initial['originals'][target['id']]
        baseline_images=tuple(p.read_bytes() for p in sorted((directory/'warning-input').glob('baseline-org-*.png')))
        assert len(baseline_images)==1
        (output/'source-original').write_bytes(original)
        for n,raw in enumerate(baseline_images):(output/f'baseline-org-{n}.png').write_bytes(raw)
        assert not initial['output']
        navigate_sections(page,'monthly_activity',keyboard=keyboard)
        text=page.get_by_role('textbox',name='Your other notes' if scenario=='ai' else 'Activity summary',exact=True)
        pending_sequence=receipt(page)['sequence']
        text.fill((text.input_value()+'\n'+PENDING).strip())  # No explicit blur before leaving the section.
        navigate_sections(page,'report_review',keyboard=keyboard)
        pending=after(page,pending_sequence);assert PENDING in pending['working']['blocks']['activity_summary']['text']
        assert pending['durable']['fingerprint']==pending['working']['fingerprint']
        navigate_sections(page,'monthly_activity',keyboard=keyboard)
        manage=page.get_by_role('button',name='Manage source files',exact=True)
        if manage.count():commit(manage)
        manager=page.locator('.st-key-'+prefix+'_source_manager_panel')
        caption,selected=select_source_caption(page,manager,prefix,old,keyboard=keyboard)
        pending_sequence=selected['sequence']
        caption.fill('Synthetic pending old-source caption')
        navigate_sections(page,'report_review',keyboard=keyboard)
        pending=after(page,pending_sequence);assert dict(pending['working']['sources'][0]['captions'])[1]=='Synthetic pending old-source caption'
        assert pending['originals']==initial['originals'];capture('pending-inputs-retained')
        if scenario=='source':
            owner=page.locator('.st-key-'+prefix+'_managed_source_'+target['id'])
            sites=owner.get_by_role('combobox',name='Source sites',exact=True)
            jump('source_facility',sites)
            select(sites,'Demo North')
            commit(owner.get_by_role('button',name='Use these source details',exact=True))
            receipt(page);date_field=source_date_control(page,owner)
            jump('source_date_unknown',date_field)
            date_field.fill('2026-09-03');date_field.press('Tab')
            commit(owner.get_by_role('button',name='Use these source details',exact=True))
            state=receipt(page);assert state['working']['sources'][1]['service_date']=='2026-09-03'
            assert state['working']['sources'][1]['scope_review']
        elif scenario=='ai':
            evidence=page.get_by_role('combobox',name='Evidence for paragraph 2',exact=True)
            jump('ai_evidence',evidence,'activity_summary',ordinal=1)
            previous=receipt(page)['sequence'];enter_combobox(evidence);evidence.press('Backspace');page.keyboard.press('Escape')
            state=after(page,previous);ids=state['working']['blocks']['activity_summary']['ai_paragraphs'][1]['fact_ids']
            assert MISSING_FACT not in ids and BAD_DATE_FACT in ids
            paragraph=page.get_by_role('textbox',name='Report paragraph 2',exact=True)
            jump('ai_number',paragraph,'activity_summary')
            previous=receipt(page)['sequence'];paragraph.fill('Replaced 2 seals.');paragraph.press('Tab');after(page,previous)
            drafting=page.get_by_test_id('stExpander').filter(has=page.locator('summary').get_by_text('Help write this summary',exact=True))
            open_native_expander(drafting,'Help write this summary',activate)
            files=drafting.get_by_role('combobox',name='File to read for work performed',exact=True)
            for source in (target,old):
                select(files,source['filename'])
                receipt(page,{'selected.'+prefix+'_ai_source':source['id']})
                pages=drafting.get_by_role('combobox',name='Pages to read for facts',exact=True)
                enter_combobox(pages);pages.press('Backspace');page.keyboard.press('Escape')
                receipt(page,{'selected.'+prefix+'_ai_read_pages_'+source['id']:[]})
            choices=receipt(page)['selected'];before_sources=receipt(page)['working']['sources']
            pages=drafting.get_by_role('combobox',name='Pages to read for facts',exact=True)
            jump('ai_date',pages,'activity_summary',source_page=True)
            state=receipt(page);assert state['working']['sources']==before_sources
            for key,value in choices.items():
                if '_ai_read_pages_' in key:assert state['selected'].get(key)==value
            assert not page.get_by_role('button',name='Use these facts',exact=True).count()
            jump('ai_date',evidence,'activity_summary')
            previous=receipt(page)['sequence'];enter_combobox(evidence);evidence.press('Backspace');page.keyboard.press('Escape')
            state=after(page,previous);assert state['working']['blocks']['activity_summary']['ai_paragraphs'][1]['fact_ids']==['synthetic-seals']
            check(page.get_by_role('checkbox',name='I checked this wording against the displayed current evidence',exact=True))
        elif scenario=='shared':
            preview=manager.get_by_role('combobox',name='Preview page / item',exact=True)
            jump('client_source_page',preview,'vendor_reports')
            before=receipt(page)
            assert before['selected'][prefix+'_evidence_selected']==target['id']
            commit(manager.get_by_role('button',name='Include page',exact=True))
            navigate_sections(page,'report_review',keyboard=keyboard)
            state=confirm_shared_pictures(page,prefix,target,receipt(page),commit)
            for key,caption in zip(('improvements','vendor_reports'),CAPTIONS):
                assert state['working']['blocks'][key]['asset_captions']==[caption]
                assert state['working']['blocks'][key]['asset_hashes']==before['working']['blocks'][key]['asset_hashes']
        else:
            from scripts.qa_financial_pages import OPEN,SOURCE,VISUAL,INCLUDE
            navigate_sections(page,'monthly_mbcx',keyboard=keyboard)
            pair=page.locator('.st-key-epc_pair_monthly_mbcx')
            files=pair.get_by_role('combobox',name='File to review',exact=True)
            select(files,old['filename']);receipt(page,{'selected.'+prefix+'_evidence_section_mbcx_report_selected':old['id']})
            check(pair.get_by_role('checkbox',name=OPEN,exact=True))
            batch=pair.get_by_role('combobox',name='Pages to review together',exact=True)
            select(batch,'Page 1 · new')
            receipt(page,{'selected.'+prefix+'_evidence_section_mbcx_report_financial_'+old['id']+'_pages':[1]})
            before=receipt(page)
            jump('financial_page',batch,'mbcx_report')
            state=receipt(page);assert state['working']['sources']==before['working']['sources']
            expect(pair.get_by_role('checkbox',name=SOURCE,exact=True)).not_to_be_checked()
            expect(pair.get_by_role('checkbox',name=VISUAL,exact=True)).not_to_be_checked()
            expect(pair.get_by_role('button',name=INCLUDE,exact=True)).to_be_disabled()
            select(batch,'Page 1 · already included')
            receipt(page,{'selected.'+prefix+'_evidence_section_mbcx_report_financial_'+target['id']+'_pages':[1]})
            check(pair.get_by_role('checkbox',name=SOURCE,exact=True))
            check(pair.get_by_role('checkbox',name=VISUAL,exact=True))
            state=commit(pair.get_by_role('button',name=INCLUDE,exact=True))
            assert state['working']['financial_allowed'][target['id']]==[1]
            assert state['working']['blocks']['mbcx_report']['asset_hashes']==before['working']['blocks']['mbcx_report']['asset_hashes']
            select(files,old['filename']);receipt(page,{'selected.'+prefix+'_evidence_section_mbcx_report_selected':old['id']})
            expect(pair.get_by_role('checkbox',name=OPEN,exact=True)).to_be_checked()
            assert receipt(page)['selected'][prefix+'_evidence_section_mbcx_report_financial_'+old['id']+'_pages']==[1]
            expect(pair.get_by_role('checkbox',name=SOURCE,exact=True)).not_to_be_checked()
        corrected=capture('explicit-corrections-complete')
        assert corrected['originals']==initial['originals']
        assert not [c for c in corrected['working']['checks'] if c['blocking']],corrected['working']['checks']
        navigate_sections(page,'report_review',keyboard=keyboard)
        saved=save_shared_version(page,prefix,commit);assert saved['shared']['fingerprint']==saved['working']['fingerprint']
        page.reload(wait_until='domcontentloaded');reopened=receipt(page)
        assert reopened['working']==saved['working'];record['reopen_verified']=True
        navigate_sections(page,'report_review',keyboard=keyboard)
        # Optional advisories, if any, must be acknowledged through the normal
        # user control; mandatory checks above may not be bypassed.
        advisory=page.get_by_role('checkbox',name='I checked these specific warnings',exact=True)
        if advisory.count():check(advisory)
        activate(page.get_by_role('button',name='Create Word and PDF',exact=True))
        page.wait_for_function('''()=>{const n=document.querySelector('[data-warning-routes]');if(!n)return false;try{return !!JSON.parse(n.dataset.warningRoutes).output['report.pdf']}catch(_){return false}}''',timeout=180000)
        final=capture('actual-saved-output')
        directory,_,_=namespace(base,token)
        for filename in ('report.docx','report.pdf'):shutil.copy2(directory/'warning-output'/filename,output/filename)
        record['outputs']=verify_outputs(output,scenario,final,original=original,baseline_images=baseline_images);record['outputs_verified']=True
        assert not record['browser_errors'] and not record['external_requests']
        record['result']='passed'
    except Exception as exc:
        record['result']='failed';record['error']=type(exc).__name__+': '+str(exc);record['traceback']=traceback.format_exc()
        try:page.screenshot(path=str(output/'failure.png'))
        except Exception as capture_error:record['screenshot_error']=str(capture_error)
        raise
    finally:
        for action in (lambda:context.tracing.stop(path=str(output/'trace.zip')),context.close,browser.close):
            try:action()
            except Exception as exc:record.setdefault('cleanup_errors',[]).append(str(exc));record['result']='failed'
        (output/'result.json').write_text(json.dumps(record,indent=2))


def main():
    if os.environ.get('GITHUB_ACTIONS')!='true':raise RuntimeError('This native acceptance driver is authorized for hosted CI only; no local browser launch.')
    from playwright.sync_api import sync_playwright
    OUT.mkdir(parents=True,exist_ok=True)
    REPORT['tested_checkout']=subprocess.check_output(['git','rev-parse','HEAD'],cwd=ROOT,text=True).strip()
    with tempfile.TemporaryDirectory(prefix='warning-native-') as temp,ExitStack() as controller:
        base=Path(temp);(base/'monthly-dom-synthetic-only.marker').write_text('synthetic-monthly-dom-v1')
        controller.enter_context(controller_environment(base))
        env=dict(os.environ,EPC_DATA_DIR=temp,EPC_MONTHLY_DOM_FIXTURE='1',EPC_MONTHLY_WARNING_ROUTES_TASK='1',
            ANTHROPIC_API_KEY='',SMARTSHEET_API_MODE='disabled',PYTHONDONTWRITEBYTECODE='1')
        for key in tuple(env):
            if key.startswith('EPC_MONTHLY_') and (key.endswith('_TASK') or key=='EPC_MONTHLY_DOM_APPTEST') and key!='EPC_MONTHLY_WARNING_ROUTES_TASK':env.pop(key)
        with (OUT/'streamlit.log').open('w') as log:
            process=subprocess.Popen([sys.executable,'-m','streamlit','run','tests/fixtures/workspaces/monthly_app.py',
                '--server.address','127.0.0.1','--server.port',str(PORT),'--server.headless','true','--browser.gatherUsageStats','false'],cwd=ROOT,env=env,stdout=log,stderr=subprocess.STDOUT)
            try:
                deadline=time.monotonic()+30
                while time.monotonic()<deadline:
                    if process.poll() is not None:raise RuntimeError('Synthetic warning server stopped')
                    try:
                        with urlopen(URL+'/_stcore/health',timeout=1) as response:
                            if response.status==200:break
                    except OSError:time.sleep(.1)
                else:raise RuntimeError('Synthetic warning server not ready')
                with sync_playwright() as p:
                    for engine,name in ((p.chromium,ENGINES[0]),(p.webkit,ENGINES[1])):
                        for mode in MODES:
                            for scenario in SCENARIOS:
                                try:exercise(engine,name,mode,scenario,base)
                                except Exception as exc:REPORT['failures'].append(name+'/'+mode+'/'+scenario+': '+str(exc)[:500])
            finally:
                process.terminate()
                try:process.wait(timeout=20)
                except subprocess.TimeoutExpired:process.kill();process.wait(timeout=5)
                REPORT['result']='passed' if all_cases_passed(REPORT) else 'failed'
                (OUT/'results.json').write_text(json.dumps(REPORT,indent=2))
    if not all_cases_passed(REPORT):raise AssertionError('Warning-native acceptance incomplete: '+str(REPORT['failures']))


if __name__=='__main__':
    try:main()
    except Exception as exc:
        REPORT['result']='failed';REPORT['failures'].append(type(exc).__name__+': '+str(exc))
        OUT.mkdir(parents=True,exist_ok=True);(OUT/'results.json').write_text(json.dumps(REPORT,indent=2));raise
