"""One native source-review/save/reload/output journey per browser; no retries."""
from io import BytesIO
from hashlib import sha256
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import time
import traceback
from urllib.parse import urlparse
from urllib.request import urlopen

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
OUT = ROOT / "test-artifacts" / "training-source-review"
URL = "http://127.0.0.1:8537"
RECEIPT_PREFIX = "QA Training review receipt: "


def observed_receipt_predicate(predicate):
    """Observe the unchanged predicate; never change polling or its outcome."""
    return """payload => {
      const original = (""" + predicate + """);
      let answer, originalError, didThrow = false;
      try { answer = original(payload); } catch (error) { originalError = error; didThrow = true; }
      try {
        const key = '__qaTrainingReceiptObservation';
        let slot = globalThis[key];
        if (!slot || slot.payload !== payload) {
          slot = globalThis[key] = {payload, started: performance.now(),
            count: 0, dropped: 0, transitions: [], last: null};
        }
        const text = document.body?.innerText || '';
        const at = text.lastIndexOf(payload.prefix);
        let parsed = null, parseOK = false;
        if (at >= 0) {
          try { parsed = JSON.parse(text.slice(at + payload.prefix.length).trim()); parseOK = true; }
          catch {}
        }
        const state = {
          result: answer === true, threw: didThrow,
          visible: document.visibilityState === 'visible',
          ready: String(document.readyState).slice(0, 32),
          script: String(document.querySelector('[data-testid="stApp"]')?.getAttribute('data-test-script-state') || '').slice(0, 32),
          prefix: at >= 0, parseOK,
          reviewsMatch: parseOK && (payload.wanted.reviews == null || parsed?.reviews === payload.wanted.reviews),
          savedMatch: parseOK && (payload.wanted.saved == null || parsed?.saved_revision === payload.wanted.saved),
          previousMatch: parseOK && (payload.wanted.previous == null || parsed?.review_revision !== payload.wanted.previous),
          includedMatch: parseOK && Array.isArray(parsed?.tables) &&
            (payload.wanted.included == null || parsed.tables.some(
              t => Array.isArray(t.rows) && t.rows.some(r => r[0] === 'Synthetic drill')) === payload.wanted.included)
        };
        slot.count = Math.min(slot.count + 1, 1000000);
        const signature = JSON.stringify(state);
        const observation = {ms: Math.max(0, performance.now() - slot.started), ...state};
        if (signature !== slot.signature) {
          if (slot.transitions.length < 64) slot.transitions.push(observation);
          else slot.dropped = Math.min(slot.dropped + 1, 1000000);
          slot.signature = signature;
        }
        slot.last = observation;
      } catch {}
      if (didThrow) throw originalError;
      return answer;
    }"""


def stop_initial_trace(context, folder, evidence, *, failed):
    """Retain only the first-entry trace, capped at 16 MiB, without hiding failure."""
    try:
        # Stage outside the uploaded artifact directory: a failed exporter may
        # leave an arbitrarily large partial ZIP even when stop() raises.
        with tempfile.TemporaryDirectory(prefix='training-initial-trace-', ignore_cleanup_errors=True) as temporary:
            path = Path(temporary) / 'trace.zip'
            context.tracing.stop(path=str(path))
            size = path.stat().st_size
            if size > 16 * 1024 * 1024:
                evidence['initial_trace'] = {'retained': False, 'bytes': size, 'reason': 'size cap'}
            else:
                (folder / 'initial-receipt-trace.zip').write_bytes(path.read_bytes())
                evidence['initial_trace'] = {'retained': True, 'bytes': size}
    except Exception as error:
        evidence['initial_trace'] = {'retained': False, 'error_type': type(error).__name__}
        if not failed:
            raise


def receipt(page, *, reviews=None, included=None, saved=None, previous_review=None):
    payload = {"prefix": RECEIPT_PREFIX, "wanted": {"reviews": reviews, "included": included,
        "saved": saved, "previous": previous_review}}
    # Page.wait_for_function recursively omits None dictionary fields in pinned
    # Playwright1.58; nullish absence is optional, explicit zero/false stay strict.
    predicate = '''payload => {
      const text=document.body?.innerText||'';
      const at=text.lastIndexOf(payload.prefix);
      if(at<0)return false;
      let v;
      try{v=JSON.parse(text.slice(at+payload.prefix.length).trim())}catch{return false}
      const wanted=payload.wanted;
      return (wanted.reviews==null||v.reviews===wanted.reviews) &&
        (wanted.included==null||v.tables.some(t=>t.rows.some(r=>r[0]==='Synthetic drill'))===wanted.included) &&
        (wanted.saved==null||v.saved_revision===wanted.saved) &&
        (wanted.previous==null||v.review_revision!==wanted.previous) &&
        document.querySelector('[data-testid="stApp"]')?.getAttribute('data-test-script-state')==='notRunning';
    }'''
    try:
        page.wait_for_function(observed_receipt_predicate(predicate), arg=payload, timeout=30000)
    except Exception as exc:
        probe = page.evaluate('''payload => {
          const candidates=[];
          for (const node of document.querySelectorAll('body *')) {
            const text=node.textContent||'';
            const at=text.indexOf(payload.prefix);
            if(at<0)continue;
            const raw=text.slice(at+payload.prefix.length);
            let parsed=null, parse_error=null;
            try{parsed=JSON.parse(raw)}catch(error){parse_error=String(error)}
            candidates.push({
              tag:node.tagName, testid:node.getAttribute('data-testid'),
              raw:raw.slice(0,1000), parse_error,
              parsed:parsed&&{
                reviews:parsed.reviews, saved_revision:parsed.saved_revision,
                review_revision:parsed.review_revision,
                included:Array.isArray(parsed.tables)&&parsed.tables.some(
                  table=>Array.isArray(table.rows)&&table.rows.some(row=>row[0]==='Synthetic drill')),
                table_count:Array.isArray(parsed.tables)?parsed.tables.length:null
              }
            });
            if(candidates.length>=8)break;
          }
          const bodyText=document.body?.innerText||'';
          const receiptAt=bodyText.lastIndexOf(payload.prefix);
          const bodyReceipt=receiptAt<0?null:bodyText.slice(receiptAt+payload.prefix.length).trim();
          let parsedBodyReceipt=null, bodyReceiptError=null;
          try{parsedBodyReceipt=bodyReceipt===null?null:JSON.parse(bodyReceipt)}
          catch(error){bodyReceiptError=String(error)}
          const observation = globalThis.__qaTrainingReceiptObservation;
          return {
            polling_observation: observation ? {
              count: observation.count, dropped: observation.dropped,
              elapsed_ms: Math.max(0, performance.now() - observation.started),
              transitions: observation.transitions, last: observation.last
            } : null,
            script_state:document.querySelector('[data-testid="stApp"]')?.getAttribute('data-test-script-state')||null,
            body_has_prefix:bodyText.includes(payload.prefix),
            body_receipt:bodyReceipt?.slice(0,1200)||null,
            body_receipt_error:bodyReceiptError,
            parsed_body_receipt:parsedBodyReceipt&&{
              reviews:parsedBodyReceipt.reviews, saved_revision:parsedBodyReceipt.saved_revision,
              review_revision:parsedBodyReceipt.review_revision,
              included:Array.isArray(parsedBodyReceipt.tables)&&parsedBodyReceipt.tables.some(
                table=>Array.isArray(table.rows)&&table.rows.some(row=>row[0]==='Synthetic drill')),
              table_count:Array.isArray(parsedBodyReceipt.tables)?parsedBodyReceipt.tables.length:null
            },
            candidates,
            body_tail:bodyText.slice(-2500)
          };
        }''', payload)
        raise AssertionError("Training receipt readiness diagnostics: " +
                             json.dumps(probe, sort_keys=True, ensure_ascii=False)) from exc
    return page.evaluate('''prefix => {
      const text=document.body?.innerText||'';
      const at=text.lastIndexOf(prefix);
      if(at<0)throw new Error('Training receipt prefix disappeared after readiness.');
      return JSON.parse(text.slice(at+prefix.length).trim());
    }''', RECEIPT_PREFIX)


def journey(engine, name):
    from playwright.sync_api import expect
    from docx import Document
    import pymupdf
    folder = OUT / name
    folder.mkdir(parents=True, exist_ok=True)
    browser = engine.launch(headless=True)
    context = browser.new_context(viewport={"width": 1365, "height": 900}, accept_downloads=True)
    errors, external = [], []
    def route(request):
        parsed = urlparse(request.request.url)
        if parsed.scheme in ("data", "blob") or (parsed.hostname == "127.0.0.1" and parsed.port == 8537):
            request.continue_()
        else:
            external.append(request.request.url)
            request.abort()
    context.route("**/*", route)
    page = context.new_page()
    page.set_default_timeout(20000)
    page.on('pageerror', lambda error: errors.append(str(error)))
    evidence = {"browser": name, "synthetic": True, "steps": [], "errors": errors, "external": external}
    try:
        context.tracing.start(screenshots=True, snapshots=True, sources=False)
        initial_failed = True
        try:
            page.goto(URL + '/?case=' + name, wait_until='domcontentloaded')
            # Preserve the existing bounded initial app-script readiness step.
            # It is separate from the strict30s receipt check; optional-argument
            # transport, not late rendering, caused the reproduced first-receipt failure.
            page.wait_for_function("""() =>
              document.querySelector('[data-testid="stApp"]')?.getAttribute('data-test-script-state')==='notRunning'
            """, timeout=45000)
            evidence['steps'].append(receipt(page, reviews=0, included=False))
            initial_failed = False
        finally:
            stop_initial_trace(context, folder, evidence, failed=initial_failed)
        page.get_by_text('Import Training from a saved source', exact=True).click()
        page.get_by_role('button', name='Review this Training table', exact=True).click()
        reviewed = receipt(page, reviews=1, included=False)
        evidence['steps'].append(reviewed)
        hours = page.get_by_role('textbox', name='Hours', exact=True)
        expect(hours).to_have_value('0')
        hours.fill('')
        # One native form submission delivers the pending intentional clear.
        page.get_by_role('button', name='Save Training row review', exact=True).click()
        receipt(page, reviews=1, included=False, previous_review=reviewed['review_revision'])
        expect(page.get_by_role('textbox', name='Hours', exact=True)).to_have_value('')
        page.get_by_role('button', name='Include reviewed Training rows', exact=True).click()
        accepted = receipt(page, reviews=1, included=True)
        event = next(t for t in accepted['tables'] if t['reference'] == 'training:events:v1')
        assert event['rows'] == [['Synthetic drill', '2026-09-12', 'Virtual', 'Alice', '']]
        page.get_by_role('button', name='Save synthetic shared version', exact=True).click()
        evidence['steps'].append(receipt(page, reviews=1, included=True, saved=1))
        page.reload(wait_until='domcontentloaded')
        reopened = receipt(page, reviews=1, included=True, saved=1)
        assert reopened['fingerprint'] == accepted['fingerprint']
        page.set_viewport_size({'width': 390, 'height': 844})
        page.get_by_text('Import Training from a saved source', exact=True).click()
        expect(page.get_by_role('textbox', name='Hours', exact=True)).to_have_value('')
        page.get_by_role('button', name='Include reviewed Training rows', exact=True).click()
        assert receipt(page, reviews=1, included=True, saved=1)['fingerprint'] == reopened['fingerprint']
        page.get_by_role('button', name='Generate synthetic Training output', exact=True).click()
        word_button = page.get_by_role('button', name='Download synthetic Word', exact=True)
        expect(word_button).to_be_visible(timeout=90000)
        with page.expect_download() as download:
            word_button.click()
        word = Path(download.value.path()).read_bytes()
        with page.expect_download() as download:
            page.get_by_role('button', name='Download synthetic PDF', exact=True).click()
        pdf = Path(download.value.path()).read_bytes()
        doc = Document(BytesIO(word))
        rows = [tuple(c.text for c in r.cells) for t in doc.tables for r in t.rows]
        assert rows.count(('Synthetic drill', '2026-09-12', 'Virtual', 'Alice', '')) == 1
        with pymupdf.open(stream=pdf, filetype='pdf') as result:
            text = '\n'.join(p.get_text() for p in result)
            assert text.count('Synthetic drill') == 1 and '2026-09-12' in text
        assert not errors and not external
        (folder/'report.docx').write_bytes(word)
        (folder/'report.pdf').write_bytes(pdf)
        evidence['output'] = {'docx_sha256': sha256(word).hexdigest(), 'pdf_sha256': sha256(pdf).hexdigest()}
        evidence['passed'] = True
        page.screenshot(path=str(folder/'complete.png'), full_page=True)
    except Exception:
        evidence['passed'] = False
        evidence['failure'] = traceback.format_exc()
        page.screenshot(path=str(folder/'failure.png'), full_page=True)
        raise
    finally:
        (folder/'evidence.json').write_text(json.dumps(evidence, indent=2))
        context.close()
        browser.close()


def main():
    from playwright.sync_api import sync_playwright
    OUT.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix='synthetic-training-browser-') as temporary:
        env = {**os.environ, 'EPC_DATA_DIR': temporary, 'EPC_TRAINING_SOURCE_QA': '1',
               'ANTHROPIC_API_KEY': '', 'OPENAI_API_KEY': ''}
        with (OUT/'server.log').open('w') as log:
            server = subprocess.Popen([sys.executable, '-m', 'streamlit', 'run',
                'tests/fixtures/workspaces/training_import_app.py', '--server.port=8537',
                '--server.address=127.0.0.1', '--server.headless=true', '--browser.gatherUsageStats=false'],
                cwd=ROOT, env=env, stdout=log, stderr=subprocess.STDOUT)
            try:
                for _ in range(150):
                    if server.poll() is not None:
                        raise RuntimeError('Synthetic Training server exited before startup')
                    try:
                        with urlopen(URL+'/_stcore/health', timeout=1) as health:
                            if health.status == 200:
                                break
                    except OSError:
                        time.sleep(.2)
                else:
                    raise RuntimeError('Synthetic Training server did not start')
                failures = []
                with sync_playwright() as p:
                    for name in ('chromium', 'webkit'):
                        try:
                            journey(getattr(p, name), name)
                        except Exception as exc:
                            failures.append(name+': '+str(exc))
                if failures:
                    raise SystemExit('Training source review failures: '+'; '.join(failures))
            finally:
                server.terminate()
                try:
                    server.wait(timeout=10)
                except subprocess.TimeoutExpired:
                    server.kill()
                    server.wait(timeout=10)


if __name__ == '__main__':
    main()
