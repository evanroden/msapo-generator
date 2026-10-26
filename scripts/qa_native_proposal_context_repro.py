"""Synthetic-only reproduction of duplicate retained proposal table labels.

Run with the repository's locked Python environment. This diagnostic writes
invented DOCX sources/outputs and a JSON matrix under test-artifacts. It returns
nonzero while any generated heading is duplicated. No production state or real
report is read. This is a diagnostic branch, not a product repair.
"""
from dataclasses import replace
from hashlib import sha256
from io import BytesIO
import json
from pathlib import Path
import runpy
import sys

from docx import Document
from docx.oxml import OxmlElement

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from app.monthly_report_import import _text, inspect_docx
from app.monthly_report_native_layout import build_native_docx
from app.monthly_report_sections import table_without_prices


def main():
    scenario = runpy.run_path(str(ROOT / "tests/test_monthly_report_native_table_context.py"))["scenario"]
    output = ROOT / "test-artifacts/native-proposal-context"
    output.mkdir(parents=True, exist_ok=True)
    results = []
    for blank in (False, True):
        for trailing in (False, True):
            name = f"blank-{int(blank)}-tail-{int(trailing)}"
            folder = output / name
            folder.mkdir(exist_ok=True)
            path, draft = scenario(folder)
            source = Document(path)
            if blank:
                for table in source.tables[:2]:
                    table._tbl.addprevious(OxmlElement("w:p"))
                source.save(path)
            original = sha256(path.read_bytes()).hexdigest()
            inspection = inspect_docx(path)
            tables = tuple(replace(table_without_prices(item)[0],
                reference=f"docx:{inspection.sha256}:{item.id}")
                for item in inspection.items
                if item.kind == "table" and item.section == "proposals")
            block = draft.blocks[0]
            block = replace(block, extra_tables=tables, text=block.text + (
                "\n\nSynthetic trailing technical note." if trailing else ""))
            draft = replace(draft, blocks=(block, draft.blocks[1]))
            payload = build_native_docx(draft, path, section_key="proposals")
            (folder / "output.docx").write_bytes(payload)
            document = Document(BytesIO(payload))
            text = _text(document.element)
            counts = {label: text.count(label)
                      for label in ("Rejected Quotes", "Pending Quotes")}
            result = {"synthetic_only": True, "case": name,
                "blank_before_table": blank, "current_trailing_note": trailing,
                "label_counts": counts, "labels_appear_once": all(n == 1 for n in counts.values()),
                "current_note_retained": not trailing or "Synthetic trailing technical note." in text,
                "source_unchanged": original == sha256(path.read_bytes()).hexdigest(),
                "native_tables": len(document.tables),
                "output_sha256": sha256(payload).hexdigest()}
            assert result["source_unchanged"] and result["current_note_retained"]
            assert result["native_tables"] == 2
            results.append(result)
    (output / "results.json").write_text(json.dumps(results, indent=2) + "\n")
    print(json.dumps(results, indent=2))
    return 0 if all(result["labels_appear_once"] for result in results) else 1


if __name__ == "__main__":
    raise SystemExit(main())
