"""Synthetic >24 MiB starting-report intake through both native application flows.

Run in the existing 512 MiB CI container. No company documents, installed masters,
network readers or production directories are used. This is intake/resource
evidence; renderer fidelity and maximum-128-MiB acceptance remain separate.
"""
from io import BytesIO
from pathlib import Path
from random import Random
from tempfile import TemporaryDirectory
from time import monotonic
from unittest.mock import patch
from zipfile import ZipFile
import gc
import hashlib
import json
import os
import resource
import sys

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from docx import Document
from docx.shared import Inches
from PIL import Image
from streamlit.testing.v1 import AppTest

from app import monthly_report_import_ui as mapping_ui
from app import monthly_report_library as library
from app import monthly_report_section_ui as section_ui
from app.monthly_report_asset_review import pending_asset_indexes
from app.monthly_report_model import synthetic_profiles
from scripts.qa_monthly_report_resources import _memory, _memory_events


class Upload(BytesIO):
    def __init__(self, raw):
        super().__init__(raw)
        self.name, self.size = 'synthetic-large-start.docx', len(raw)


def _fixture(path):
    profile = synthetic_profiles()[0]
    document = Document()
    document.add_heading(profile.facilities[0].title, 0)
    document.add_paragraph('Monthly Report · September 2026')
    document.add_heading('Monthly Activity Summary', 1)
    document.add_paragraph('September 2026 synthetic completed pump inspection.')
    for index in range(9):
        # Nine distinct valid one-megapixel PNGs exceed the narrow native-art
        # source cap while staying far below the ordinary intake/pixel budgets.
        data = Random(index).randbytes(1024 * 1024 * 3)
        image = Image.frombytes('RGB', (1024, 1024), data)
        output = BytesIO()
        image.save(output, format='PNG', compress_level=0)
        image.close()
        document.add_picture(BytesIO(output.getvalue()), width=Inches(4))
        document.add_paragraph(f'Synthetic inspection picture {index + 1}.')
    document.save(path)


PRIMARY = '''
from app.monthly_report_section_ui import render_section_setup
from app.monthly_report_model import synthetic_profiles, ReportPeriod
from app.monthly_report_ui import _field
p = synthetic_profiles()[0]
render_section_setup(p.contract, ReportPeriod(2026, 9), "Synthetic Editor", _field, identity=p)
'''
SECONDARY = '''
from app import monthly_report_library as library
from app.monthly_report_import_ui import render_import
from app.monthly_report_model import synthetic_profiles
from app.monthly_report_ui import _field
p = synthetic_profiles()[0]
render_import(library.load_profile(p.contract, p.key), _field)
'''


def _clean(app):
    assert not app.exception, [str(item) for item in app.exception]
    assert not app.error, [item.value for item in app.error]


def run(directory, record):
    source = directory / 'synthetic-large-start.docx'
    record['stage'] = 'fixture'
    _fixture(source)
    gc.collect()
    original = source.read_bytes()
    digest = hashlib.sha256(original).hexdigest()
    record['source_bytes'] = len(original)
    assert 24 * 1024 * 1024 < len(original) < 128 * 1024 * 1024
    with ZipFile(source) as package:
        infos = package.infolist()
        record['expanded_bytes'] = sum(item.file_size for item in infos)
        record['media_bytes'] = sum(item.file_size for item in infos if item.filename.startswith('word/media/'))
        record['xml_bytes'] = sum(item.file_size for item in infos if item.filename.endswith(('.xml', '.rels')))
        record['package_parts'] = len(infos)
    profile = synthetic_profiles()[0]
    os.environ['EPC_DATA_DIR'] = str(directory / 'primary')
    record['stage'] = 'primary-read'
    start = monotonic()
    upload = Upload(original)
    real_upload = section_ui.st.file_uploader
    def primary_upload(label, *args, **kwargs):
        return upload if label == 'Starting report' else real_upload(label, *args, **kwargs)
    with patch.object(section_ui.st, 'file_uploader', primary_upload):
        app = AppTest.from_string(PRIMARY, default_timeout=120).run()
        _clean(app)
        staged = next(v for k, v in app.session_state.filtered_state.items() if k.endswith('_stage'))
        assert staged[2].sha256 == digest
        assert sum(item.kind == 'image' for item in staged[2].items) == 9
        assert not any('processing limit' in warning.value for warning in app.warning)
        record['stage'] = 'primary-save'
        button = next(b for b in app.button if b.label == 'Continue to this month’s updates')
        assert not button.disabled
        button.click().run()
        _clean(app)
        draft = library.load_imported_draft(profile.contract, profile.key)
        assert any('completed pump inspection' in block.text for block in draft.blocks)
        assert sum(len(block.asset_hashes) for block in draft.blocks) == 9
        assert sum(len(pending_asset_indexes(block)) for block in draft.blocks) == 9
        assert hashlib.sha256(library.imported_original(profile.contract, profile.key).read_bytes()).hexdigest() == digest
        record['primary_seconds'] = round(monotonic() - start, 3)
        record['primary_saved_images'] = 9
        staged[0].cleanup()
    del app, staged, upload, draft
    gc.collect()
    os.environ['EPC_DATA_DIR'] = str(directory / 'secondary')
    library.save_profile(profile, expected_revision=0, actor='Synthetic Editor', confirmed=True)
    record['stage'] = 'secondary-read'
    start = monotonic()
    upload = Upload(original)
    with patch.object(mapping_ui.st, 'file_uploader', lambda *a, **kw: upload):
        app = AppTest.from_string(SECONDARY, default_timeout=120).run()
        next(b for b in app.button if b.label == 'Read DOCX for mapping').click().run()
        _clean(app)
        staged = next(v for k, v in app.session_state.filtered_state.items() if k.endswith('_stage'))
        assert staged[2].sha256 == digest
        assert sum(item.kind == 'image' for item in staged[2].items) == 9
        text = next(item for item in staged[2].items if 'completed pump inspection' in item.text)
        next(w for w in app.selectbox if w.label == 'Item to inspect').set_value(text.id).run()
        next(b for b in app.button if b.label == 'Add confirmed mapping').click().run()
        _clean(app)
        mappings = next(v for k, v in app.session_state.filtered_state.items() if k.endswith('_mappings'))
        assert text.id in mappings
        assert staged[1].read_bytes() == original
        record['secondary_seconds'] = round(monotonic() - start, 3)
        record['secondary_inventory_images'] = 9
        staged[0].cleanup()
    assert hashlib.sha256(source.read_bytes()).hexdigest() == digest
    record['stage'] = 'complete'
    record['source_unchanged'] = True


def main():
    output = Path(os.environ.get('LARGE_INTAKE_OUTPUT', '/tmp/monthly-large-intake'))
    output.mkdir(parents=True, exist_ok=True)
    record = {'status': 'running', 'stage': 'admission'}
    before = _memory_events()
    limit, _ = _memory()
    record['container_memory_limit_bytes'] = limit
    try:
        if os.environ.get('RESOURCE_REQUIRE_512_MB') == '1':
            assert limit == 512 * 1024 * 1024, 'Acceptance requires an actual 512 MiB cgroup cap'
        with TemporaryDirectory(prefix='large-intake-') as temporary:
            run(Path(temporary), record)
        after = _memory_events()
        assert after.get('oom_kill', 0) == before.get('oom_kill', 0), 'Container OOM during intake'
        record['status'] = 'passed'
    except BaseException as exc:
        record['status'] = 'failed'
        record['error_type'] = type(exc).__name__
        raise
    finally:
        record['peak_rss_bytes'] = resource.getrusage(resource.RUSAGE_SELF).ru_maxrss * 1024
        record['memory_events_before'] = before
        record['memory_events_after'] = _memory_events()
        (output / 'manifest.json').write_text(json.dumps(record, indent=2) + '\n')
        print(json.dumps(record, sort_keys=True), flush=True)


if __name__ == '__main__':
    main()
