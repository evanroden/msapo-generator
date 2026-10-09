"""Synthetic native-report stress probe for a CPU/RAM-capped production image.

The fixture contains large source JPEG frames, as a phone-photo report does.
No customer documents, runtime library, contacts or fonts are exported. This is
resource evidence, not a substitute for the native-layout fidelity regressions.
"""
from __future__ import annotations

from dataclasses import replace
from io import BytesIO
from pathlib import Path
import gc
import hashlib
import json
import os
import resource
import sys
import tempfile
from threading import Event, Thread
from time import monotonic

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from docx import Document
from docx.shared import Inches
from PIL import Image
import pymupdf


def _memory():
    for directory, limit, used in ((Path('/sys/fs/cgroup'), 'memory.max', 'memory.current'),
                                  (Path('/sys/fs/cgroup/memory'), 'memory.limit_in_bytes', 'memory.usage_in_bytes')):
        try:
            return int((directory / limit).read_text()), int((directory / used).read_text())
        except (OSError, ValueError):
            continue
    return None, None


def _fixture(directory):
    from app import monthly_report_designs as designs
    from app.monthly_report_asset_review import approve_all_assets
    from app.monthly_report_import import inspect_docx, read_import_image
    from app.monthly_report_model import BlockSpec, Facility, ReportDraft, ReportPeriod, ReportProfile, ResolvedBlock, SectionSpec

    source = directory / 'synthetic-large-photo-report.docx'
    document = Document()
    document.add_paragraph('Synthetic Resource Site')
    document.add_paragraph('September 2026')
    document.add_paragraph('Prepared by: Synthetic Editor')
    document.add_heading('MONTHLY ACTIVITY SUMMARY', 1)
    document.add_paragraph('Synthetic current inspection completed.')
    document.add_paragraph('Improvement Highlights and Pics')
    for index in range(8):
        # Low-resolution texture enlarged to phone dimensions has a realistic
        # small encoded/large decoded footprint, without allocating pixel lists.
        texture = Image.effect_noise((480, 320), 24).convert('RGB')
        texture.paste((40 + index * 18, 95, 130), (index * 30, 25, index * 30 + 30, 100))
        photo = texture.resize((4800, 3200), Image.Resampling.BICUBIC)
        raw = BytesIO()
        photo.save(raw, format='JPEG', quality=86, optimize=False)
        document.add_picture(BytesIO(raw.getvalue()), width=Inches(6.6))
        if index < 7:
            document.add_page_break()
        photo.close(); texture.close()
    document.save(source)
    del document, raw
    gc.collect()
    profile = ReportProfile('Synthetic Resource Contract', 'synthetic-resource', 'Synthetic Resource Site',
                            (Facility('synthetic', 'Synthetic Resource Site'),))
    profile = designs.install(source, profile, actor='Synthetic Editor', expected_revision=0)
    inspection = inspect_docx(source)
    assets, refs, images = {}, [], []
    text_refs = []
    for item in inspection.items:
        if item.kind == 'text' and item.text == 'Synthetic current inspection completed.':
            text_refs.append(f'docx:{inspection.sha256}:{item.id}')
        if item.kind != 'image' or item.suggested_slot != 'improvements':
            continue
        image = read_import_image(source, item, line_art=False)
        ref = hashlib.sha256(image.data).hexdigest() + '.' + image.extension
        assets[ref] = image.data
        images.append(ref)
        refs.append(f'docx:{inspection.sha256}:{item.id}')
    assert len(images) == 8
    blocks = (ResolvedBlock('activity_summary', 'This month', text='Synthetic current inspection completed.', references=tuple(text_refs)),
              approve_all_assets(ResolvedBlock('improvements', 'This month', asset_hashes=tuple(images), references=tuple(refs))))
    sections = (SectionSpec('activity', '1', 'Monthly Activity Summary',
                            (BlockSpec('activity_summary', 'rich_text'), BlockSpec('improvements', 'image_grid'))),)
    draft = ReportDraft(profile, ReportPeriod(2026, 9), 'Synthetic Editor', sections, blocks)
    return draft, assets, source.stat().st_size


def main():
    from app.monthly_report_docx import generate_report
    from app.monthly_report_preview import preview_section
    from app.monthly_report_preview_store import PreviewStore

    limit, _ = _memory()
    if os.environ.get('RESOURCE_REQUIRE_512_MB') == '1' and (limit is None or limit > 512 * 1024 * 1024):
        raise RuntimeError('This acceptance probe must run inside the 512 MB container limit.')
    started, done, samples = monotonic(), Event(), []
    def sample():
        while not done.wait(.1):
            _, used = _memory()
            if used is not None:
                samples.append((monotonic(), used))
    sampler = Thread(target=sample, daemon=True)
    sampler.start()
    result = {'limit_bytes': limit, 'stages': []}
    store = PreviewStore()
    try:
        with tempfile.TemporaryDirectory(prefix='synthetic-report-resource-') as directory:
            directory = Path(directory)
            os.environ['EPC_DATA_DIR'] = str(directory / 'runtime')
            draft, assets, size = _fixture(directory)
            result['source_bytes'] = size
            for owner in ('synthetic-visitor-one', 'synthetic-visitor-two'):
                tick = monotonic()
                preview = preview_section(draft, 'activity', assets.__getitem__)
                ticket = store.put(owner, preview)
                del preview
                visible = store.pages(owner, ticket, (0, 1), dpi=120)
                assert visible and len(visible) == 2
                assert all(raw.startswith(b'\xff\xd8') for _, raw in visible)
                result['stages'].append({'name': owner, 'seconds': round(monotonic()-tick, 3), 'pages': ticket.pages})
                del visible
                gc.collect()
            tick = monotonic()
            package = generate_report(draft, acknowledged_fingerprint=draft.fingerprint, asset_loader=assets.__getitem__)
            if not package.pdf:
                raise RuntimeError(package.pdf_error)
            with pymupdf.open(stream=package.pdf, filetype='pdf') as pdf:
                assert 'Synthetic current inspection completed.' in '\n'.join(page.get_text() for page in pdf)
                assert sum(len(page.get_images()) for page in pdf) >= 8
                result['final_pages'] = len(pdf)
            result['final_docx_bytes'] = len(package.docx)
            result['final_pdf_bytes'] = len(package.pdf)
            result['stages'].append({'name': 'foreground-export', 'seconds': round(monotonic()-tick, 3)})
            result['cache'] = store.stats()
            assert result['cache']['pdf_bytes'] <= store.disk_limit
            assert result['cache']['raster_bytes'] <= store.raster_limit
    finally:
        store.close()
        done.set(); sampler.join(2)
    result['seconds'] = round(monotonic()-started, 3)
    result['sampled_peak_bytes'] = max((used for _, used in samples), default=0)
    result['python_peak_rss_bytes'] = resource.getrusage(resource.RUSAGE_SELF).ru_maxrss * 1024
    result['max_sampler_gap_seconds'] = round(max((b[0]-a[0] for a, b in zip(samples, samples[1:])), default=0), 3)
    print(json.dumps(result, indent=2), flush=True)


if __name__ == '__main__':
    main()
