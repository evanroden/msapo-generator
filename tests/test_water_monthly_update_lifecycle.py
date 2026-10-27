"""Synthetic Water text is monthly work, never silently reused as current."""
from dataclasses import replace

from app import monthly_report_library as library
from app.monthly_report_guided import _monthly_specs, _standing_specs
from app.monthly_report_model import default_sections, synthetic_draft, synthetic_profiles, ReportPeriod, ResolvedBlock
from app.monthly_report_setup import new_month_draft


def water_draft():
    draft = synthetic_draft(synthetic_profiles()[0], ReportPeriod(2026, 9))
    water = next(s for s in default_sections() if s.key == 'water')
    return replace(draft, sections=(water,), blocks=(
        ResolvedBlock('water_monthly_update', 'This month', text='Synthetic September water update.'),
        ResolvedBlock('water_reports', 'This month'),
    ))


def test_water_update_routes_to_monthly_not_standing_editor():
    water = water_draft().sections[0]
    assert [b.key for b in _monthly_specs(water)] == ['water_monthly_update', 'water_reports']
    assert not _standing_specs(water)


def test_same_month_resume_preserves_update_but_deliberate_new_month_starts_empty():
    draft = water_draft()
    before = draft.fingerprint
    assert new_month_draft(draft, draft.period) == draft
    following = new_month_draft(draft, ReportPeriod(2026, 10))
    update = next(b for b in following.blocks if b.key == 'water_monthly_update')
    assert not update.text and not update.references and not update.ai_paragraphs
    assert update.source == 'This month'
    assert draft.fingerprint == before
    assert draft.blocks[0].text == 'Synthetic September water update.'


def test_snapshot_reopen_and_next_month_preserve_original_history(monkeypatch, tmp_path):
    monkeypatch.setenv('EPC_DATA_DIR', str(tmp_path))
    draft = water_draft()
    library.save_profile(draft.profile, expected_revision=0, actor='Synthetic Editor', confirmed=True)
    library.save_snapshot(draft, expected_revision=0, entered_editor='Synthetic Editor')
    reopened = library.load_snapshot(draft.profile.contract, draft.profile.key, draft.period)
    assert reopened.draft == draft
    following = new_month_draft(reopened.draft, ReportPeriod(2026, 10))
    library.save_snapshot(following, expected_revision=0, entered_editor='Synthetic Editor')
    assert library.load_snapshot(draft.profile.contract, draft.profile.key, draft.period).draft == draft
    assert library.load_snapshot_version(draft.profile.contract, draft.profile.key, draft.period, 1).draft == draft
    assert library.load_snapshot(draft.profile.contract, draft.profile.key, following.period).draft == following


def test_legacy_water_section_places_new_monthly_update_before_existing_pages():
    from app.monthly_report_guided import complete_guided_sections
    draft = water_draft()
    legacy_water = replace(draft.sections[0], blocks=draft.sections[0].blocks[1:])
    legacy = replace(draft, sections=(legacy_water,), blocks=draft.blocks[1:])
    completed = complete_guided_sections(legacy)
    water = next(s for s in completed.sections if s.key == 'water')
    assert [b.key for b in water.blocks] == ['water_monthly_update', 'water_reports']
    assert legacy.sections[0] == legacy_water
    assert legacy.blocks == draft.blocks[1:]


def test_missing_water_update_insertion_preserves_custom_blocks_and_is_idempotent():
    from app.monthly_report_guided import complete_guided_sections
    from app.monthly_report_model import BlockSpec
    draft = water_draft()
    custom = (BlockSpec('water_custom_intro', 'rich_text'), draft.sections[0].blocks[1],
              BlockSpec('water_custom_tail', 'rich_text'))
    draft = replace(draft, sections=(replace(draft.sections[0], blocks=custom),))
    completed = complete_guided_sections(draft)
    water = next(s for s in completed.sections if s.key == 'water')
    assert [b.key for b in water.blocks] == ['water_custom_intro', 'water_monthly_update', 'water_reports', 'water_custom_tail']
    assert complete_guided_sections(completed) == completed
    assert draft.sections[0].blocks == custom


def test_existing_custom_water_update_order_is_not_rewritten():
    from app.monthly_report_guided import complete_guided_sections
    draft = water_draft()
    custom = tuple(reversed(draft.sections[0].blocks))
    draft = replace(draft, sections=(replace(draft.sections[0], blocks=custom),))
    water = next(s for s in complete_guided_sections(draft).sections if s.key == 'water')
    assert [b.key for b in water.blocks] == [b.key for b in custom]


def test_left_out_legacy_water_content_is_not_republished_when_slot_is_added():
    from app.monthly_report_guided import complete_guided_sections
    draft = water_draft()
    legacy = replace(draft, sections=(replace(draft.sections[0], included=False, blocks=draft.sections[0].blocks[1:]),))
    completed = complete_guided_sections(legacy)
    assert next(b for b in completed.blocks if b.key == 'water_monthly_update').source == 'Omit'
    assert next(b for b in completed.blocks if b.key == 'water_monthly_update').text == draft.blocks[0].text
    assert all(not b.required for s in completed.sections if s.key == 'water' for b in s.blocks)
    assert legacy.blocks[0].source == 'This month'


def test_saved_legacy_water_profile_order_resolves_update_and_allows_rename(monkeypatch, tmp_path):
    from app.monthly_report_model import profile_sections
    monkeypatch.setenv('EPC_DATA_DIR', str(tmp_path))
    profile = replace(water_draft().profile, section_block_order=(('water', ('water_reports',)),))
    before = profile.section_block_order
    water = next(s for s in profile_sections(profile) if s.key == 'water')
    assert [b.key for b in water.blocks] == ['water_monthly_update', 'water_reports']
    library.save_profile(profile, expected_revision=0, actor='Synthetic Editor', confirmed=True)
    renamed = replace(profile, title='Synthetic renamed Water report', facilities=(
        replace(profile.facilities[0], aliases=('Synthetic site alias',)), *profile.facilities[1:]))
    library.save_profile(renamed, expected_revision=1, actor='Synthetic Editor', confirmed=True)
    assert library.load_profile(profile.contract, profile.key).profile == renamed
    assert profile.section_block_order == before
    restored = library.restore_profile(profile.contract, profile.key, 1, expected_revision=2,
                                       actor='Synthetic Editor', confirmed=True)
    assert restored.profile == profile


def test_water_profile_order_still_rejects_missing_pages_duplicates_and_unknowns(monkeypatch, tmp_path):
    import pytest
    monkeypatch.setenv('EPC_DATA_DIR', str(tmp_path))
    for order in [(), ('water_monthly_update',), ('water_reports', 'water_reports'), ('water_reports', 'unknown')]:
        profile = replace(water_draft().profile, section_block_order=(('water', order),))
        with pytest.raises(library.LibraryError, match='block order'):
            library.save_profile(profile, expected_revision=0, actor='Synthetic Editor', confirmed=True)


def test_new_word_profile_places_missing_water_update_before_mapped_vendor_pages(tmp_path):
    from io import BytesIO
    from docx import Document
    from PIL import Image
    from app.monthly_report_import import inspect_docx
    from app.monthly_report_model import profile_sections
    from app.monthly_report_setup import design_profile, suggested_mappings
    source = Document()
    source.add_heading('Water Treatment Reports', 1)
    raw = BytesIO()
    Image.new('RGB', (100, 150), 'navy').save(raw, 'PNG')
    source.add_picture(BytesIO(raw.getvalue()))
    path = tmp_path / 'synthetic-water-import.docx'
    source.save(path)
    inspection = inspect_docx(path)
    mappings = suggested_mappings(inspection)
    assert any(mapping.slot == 'water_reports' for mapping in mappings)
    profile = design_profile(synthetic_profiles()[0], inspection, mappings)
    assert dict(profile.section_block_order)['water'] == ('water_monthly_update', 'water_reports')
    water = next(section for section in profile_sections(profile) if section.key == 'water')
    assert [block.key for block in water.blocks] == ['water_monthly_update', 'water_reports']
    # Exercise the actual import-to-native-export route, not a hand-built schema.
    from docx.oxml.ns import qn
    from app.monthly_report_import import map_items, imported_draft, _text
    from app.monthly_report_native_layout import build_native_docx
    mapped = map_items(path, inspection, mappings)
    draft = imported_draft(profile, ReportPeriod(2026, 9), 'Synthetic Editor', mapped)
    update = 'Current synthetic update entered after Word import.'
    draft = replace(draft, blocks=(*draft.blocks, ResolvedBlock('water_monthly_update', 'This month', text=update)))
    assets = dict(mapped.assets)
    original = path.read_bytes()
    rendered = Document(BytesIO(build_native_docx(draft, path, asset_loader=assets.__getitem__)))
    children = list(rendered.element.body)
    summary_position = next(i for i, child in enumerate(children) if _text(child) == update)
    picture_positions = [i for i, child in enumerate(children) if list(child.iter(qn('a:blip')))]
    assert picture_positions and summary_position < min(picture_positions)
    assert path.read_bytes() == original


def test_import_keeps_explicitly_mapped_water_update_after_pages():
    from app.monthly_report_import import DocxInspection, ImportItem, ImportMapping
    from app.monthly_report_setup import design_profile
    items = (
        ImportItem('vendor', 'image', 'word/document.xml', 1, 1, 'water', 'water_reports'),
        ImportItem('update', 'text', 'word/document.xml', 2, 1, 'water', 'water_monthly_update', text='Explicit trailing update.'),
    )
    inspection = DocxInspection('a' * 64, items, (), (), 1)
    mappings = tuple(ImportMapping(item.id, item.suggested_slot) for item in items)
    profile = design_profile(synthetic_profiles()[0], inspection, mappings)
    assert dict(profile.section_block_order)['water'] == ('water_reports', 'water_monthly_update')
