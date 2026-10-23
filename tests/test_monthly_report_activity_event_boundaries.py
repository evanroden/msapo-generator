"""PR274 review4318281487: real row boundaries versus a real negative wrap.

The shared helper is inherited from Activity PR273. These new fixtures use
actual synthetic file parsing; the second reader pass is an explicit local
result, never an external request or proof of native browser delivery.
"""
from dataclasses import replace
from email.message import EmailMessage
from hashlib import sha256
from io import BytesIO
import json
import os
from pathlib import Path

from docx import Document
import fitz
import pytest

from app import monthly_report_sources as sources
from app import monthly_report_structured_uploads as structured
from app.monthly_report_model import (
    BlockSpec, ReportDraft, ReportPeriod, SectionSpec, synthetic_profiles,
)
from conftest import requires_libreoffice


ROWS = (
    'Inspected condenser CD-41 and recorded pressure at 124 psi',
    'Replaced filter FL-12 and confirmed normal air flow',
)
OPENING = 'The technician inspected exhaust fan EF-62 and confirmed it'
TAIL = 'did not start'
EXPECTED = (*ROWS, OPENING + ' ' + TAIL)


def actual_source(monkeypatch, tmp_path, kind):
    monkeypatch.setenv('EPC_DATA_DIR', str(tmp_path))
    body = '\n'.join((*ROWS, OPENING, TAIL))
    if kind == 'pdf':
        with fitz.open() as pdf:
            page = pdf.new_page(width=720, height=400)
            for number, line in enumerate(body.splitlines()):
                page.insert_text((30, 40 + number * 20), line, fontsize=11)
            raw = pdf.tobytes()
    elif kind == 'docx':
        # Word rows are independent even without punctuation or visible bullets.
        # The negative tail is a native line break inside the third row's cell.
        document = Document()
        table = document.add_table(rows=3, cols=1)
        for row, words in zip(table.rows, (*ROWS, OPENING + '\n' + TAIL)):
            row.cells[0].text = words
        buffer = BytesIO()
        document.save(buffer)
        raw = buffer.getvalue()
    elif kind == 'eml':
        message = EmailMessage()
        message['Subject'] = 'Synthetic service notes'
        message.set_content(body)
        raw = message.as_bytes()
    else:
        assert kind == 'txt'
        raw = body.encode()
    profile = synthetic_profiles()[0]
    content, _ = sources.ingest(profile, 'synthetic-independent-events.' + kind, raw)
    assert sources.source_bytes(profile, content.source) == raw
    assert content.source.sha256 == sha256(raw).hexdigest()
    assert len(content.source.page_texts) == 1
    for line in (*ROWS, OPENING, TAIL):
        assert line in content.source.page_texts[0]
    return profile, content, raw


def event_paragraphs(text):
    return tuple(paragraph for paragraph in text.split('\n\n')
                 if any(identity in paragraph for identity in ('CD-41', 'FL-12', 'EF-62')))


@pytest.mark.parametrize('kind', ['pdf', 'docx', 'eml', 'txt'])
def test_independent_rows_and_wrapped_negative_keep_their_own_events(monkeypatch, tmp_path, kind):
    _, content, _ = actual_source(monkeypatch, tmp_path, kind)
    native = structured.native_update((content,), 'activity_summary')
    assert event_paragraphs(native.text) == EXPECTED
    assert structured.native_update((content,), 'activity_summary', native) == native
    # A later independently simulated reader cannot turn the same grounded
    # events into duplicates or replace their established paragraph boundaries.
    repeated, _ = structured.reader_update((content,), 'activity_summary', {'items': [
        {'source': content.source.id, 'page': 1, 'text': words, 'row': []}
        for words in EXPECTED]}, native)
    assert repeated == native


@pytest.mark.parametrize('lines,expected', [
    (('Inspected pump A and confirmed it', 'did not start'),
     ('Inspected pump A and confirmed it did not start',)),
    (('Inspected pump A and confirmed it', 'Replaced filter B'),
     ('Inspected pump A and confirmed it', 'Replaced filter B')),
    (('inspected the condenser drive belt', 'replaced the intake filter'),
     ('inspected the condenser drive belt', 'replaced the intake filter')),
    (('Inspected the isolation valve', 'the technician replaced the actuator'),
     ('Inspected the isolation valve', 'the technician replaced the actuator')),
    (('Inspected fan EF-62 and found it', 'locked out'),
     ('Inspected fan EF-62 and found it locked out',)),
    (('Repaired fan EF-62 but the unit was', 'not yet returned to service'),
     ('Repaired fan EF-62 but the unit was not yet returned to service',)),
    (('Inspected pump A and confirmed it', '- Replaced filter B'),
     ('Inspected pump A and confirmed it', '- Replaced filter B')),
])
def test_only_positive_continuation_evidence_joins_physical_lines(lines, expected):
    assert structured._logical_activity_lines(lines) == expected


@requires_libreoffice
def test_separate_activity_events_persist_and_print_as_separate_word_paragraphs(monkeypatch, tmp_path):
    from app import monthly_report_library as library
    from app.monthly_report_docx import generate_report
    from app.monthly_report_source_details import correct_source_details
    profile, content, raw = actual_source(monkeypatch, tmp_path, 'docx')
    source = correct_source_details(content.source, profile, None, profile.facilities[0].key)
    block = structured.native_update((replace(content, source=source),), 'activity_summary')
    draft = ReportDraft(profile, ReportPeriod(2026, 9), 'Synthetic Reviewer',
        (SectionSpec('activity', '2', 'Monthly Activity Summary', (BlockSpec('activity_summary', 'rich_text'),)),),
        (block,), sources=(source,))
    saved = library.save_snapshot(draft, expected_revision=0, entered_editor=draft.prepared_by)
    restored = library.load_snapshot(profile.contract, profile.key, draft.period)
    assert restored.draft == draft and restored.revision == saved.revision
    package = generate_report(restored.draft, acknowledged_fingerprint=restored.draft.fingerprint)
    assert package.docx and package.pdf, package.pdf_error
    output = Path(os.environ.get('INDEPENDENT_OUTPUT', str(tmp_path))) / 'activity-event-boundaries'
    output.mkdir(exist_ok=True, parents=True)
    (output / 'activity_events.docx').write_bytes(package.docx)
    (output / 'activity_events.pdf').write_bytes(package.pdf)
    paragraphs = tuple(paragraph.text for paragraph in Document(BytesIO(package.docx)).paragraphs
                       if any(identity in paragraph.text for identity in ('CD-41', 'FL-12', 'EF-62')))
    with fitz.open(stream=package.pdf, filetype='pdf') as pdf:
        pdf_text = ' '.join(' '.join(page.get_text().split()) for page in pdf)
        for number, page in enumerate(pdf, 1):
            page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5)).save(output / f'activity_events_page_{number}.png')
    (output / 'activity_events.json').write_text(json.dumps({
        'tested_revision': os.environ.get('MSAPO_SHA', 'unrecorded'),
        'source_sha256': sha256(raw).hexdigest(), 'source_fingerprint': source.fingerprint,
        'saved_revision': saved.revision, 'draft_fingerprint': draft.fingerprint,
        'expected_events': EXPECTED, 'actual_word_paragraphs': paragraphs,
        'docx_bytes': len(package.docx), 'docx_sha256': sha256(package.docx).hexdigest(),
        'pdf_bytes': len(package.pdf), 'pdf_sha256': sha256(package.pdf).hexdigest(),
    }, indent=2) + '\n')
    # Retain before-output before asserting the independently desired structure.
    assert event_paragraphs(restored.draft.blocks[0].text) == EXPECTED
    assert paragraphs == EXPECTED
    for words in EXPECTED:
        assert pdf_text.count(words) == 1
