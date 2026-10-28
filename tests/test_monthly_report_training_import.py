"""Synthetic prior Training import, review, rollover and output regressions.

Source data is fabricated. Every Word/CSV/XLSX input is ingested via the
existing bounded reader and the original bytes remain in isolated storage.
"""
from dataclasses import asdict, replace
from datetime import date
from io import BytesIO
import json

from docx import Document
from openpyxl import Workbook
import pytest

from app import monthly_report_library as library
from app import monthly_report_sources as sources
from app.monthly_report_model import (
    Facility, ReportPeriod, ReportProfile, ResolvedBlock, ReportTable,
    default_sections,
)
from app.monthly_report_training import (
    EVENT_REF, MATRIX_REF, add_event, events_for_period, load_training,
    new_month_training, session_row_tokens,
)
from app.monthly_report_training_import import (
    training_tables, review_training_table, correct_training_row,
    exclude_training_row, confirm_training_word_layout,
    reimport_training_review, accepted_training_block,
    training_review_record, restore_training_review,
)

PROFILE = ReportProfile(
    "Invented Training Contract", "multi-site", "Two synthetic hospitals",
    (Facility("north", "Site North", ("North",)),
     Facility("south", "Site South", ("South",))), "multi_site")
PERIOD = ReportPeriod(2026, 9)
CURRENT = (
    b"Training,Date,Format,Participants,Hours,Status,Source Notes\n"
    b"Pump safety,2026-09-13,Virtual,Alice; Bob,1.5,Completed,Actual attendance\n"
    b"Boiler orientation,2025-09-14,In person,Jo,2,Completed,Historical\n"
    b"Equipment basics,,Virtual,,0,Pending,Incomplete source entry\n")
MATRIX_OLD = (
    b"Team member,Site,Safety refresher,Operations update\n"
    b"Alice,North,Completed,\n"
    b"Jo,South,Pending,Not required\n")


@pytest.fixture
def storage(monkeypatch, tmp_path):
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path))
    return tmp_path


def ingest(name, raw, *, profile=PROFILE):
    content, extra = sources.ingest(profile, name, raw)
    assert not extra
    assert sources.source_bytes(profile, content.source) == raw
    return content


def reviewed(content, kind, *, source_year=None, period=PERIOD, profile=PROFILE):
    candidates = training_tables(content, profile)
    result = [review_training_table(value, profile, period, source_year=source_year)
              for value in candidates]
    match = [value for value in result if value.kind == kind]
    assert len(match) == 1, [(value.kind, value.issues) for value in result]
    return match[0]


def test_csv_prior_source_keeps_real_dates_unknown_columns_and_history(storage):
    content = ingest("prior_training.csv", CURRENT)
    review = reviewed(content, "events")
    assert review.original_columns == (
        "Training", "Date", "Format", "Participants", "Hours", "Status", "Source Notes")
    assert review.unresolved_counts() == {
        "source_rows": 4, "accepted_candidates": 1,
        "needs_review": 2, "excluded": 0, "unmatched": 0}
    assert review.rows[0].values == (
        "Pump safety", "2026-09-13", "Virtual", "Alice; Bob", "1.5")
    assert review.unmapped_cells(review.rows[0].key) == (
        (6, "Source Notes", "Actual attendance"),)
    assert "Historical/other-month" in " ".join(review.rows[1].issues)
    assert "Training date is missing" in " ".join(review.rows[2].issues)
    assert review.rows[2].origin.cells[4] == "0"
    assert not review.ready_for_import
    assert sources.source_bytes(PROFILE, content.source) == CURRENT


def test_review_then_explicitly_exclude_historical_and_incomplete_events(storage):
    original = ingest("prior_training.csv", CURRENT)
    review = reviewed(original, "events")
    initial = review
    for row in review.rows[1:]:
        review = exclude_training_row(review, row.key, expected_revision=review.revision)
    assert review.ready_for_import
    block = ResolvedBlock("training_summary", "This month",
                          text="Manually authored Training notes")
    updated = accepted_training_block(
        review, block, expected_block_fingerprint=block.fingerprint)
    assert updated.text == "Manually authored Training notes"
    assert updated.references == (review.source.reference,)
    assert len(updated.extra_tables) == 1
    events, = updated.extra_tables
    assert events.reference == EVENT_REF
    assert events.rows == ((
        "Pump safety", "2026-09-13", "Virtual", "Alice; Bob", "1.5"),)
    assert len(session_row_tokens(events)) == 1
    assert events_for_period(updated, PERIOD).rows == events.rows
    assert events_for_period(updated, ReportPeriod(2026, 10)).rows == ()
    unchanged = accepted_training_block(
        review, updated, expected_block_fingerprint=updated.fingerprint)
    assert unchanged == updated
    with pytest.raises(ValueError, match="Training changed"):
        accepted_training_block(
            initial, updated, expected_block_fingerprint=block.fingerprint)
    assert not load_training(PROFILE.contract)["revision"]
    assert sources.source_bytes(PROFILE, original.source) == CURRENT


def test_review_correction_uses_cas_and_original_cells_survive(storage):
    data = ingest("correct_training.csv", (
        b"Topic,Training Date,Method,Attendees,Duration Hours,Completion Status,Note\n"
        b"Pump safety,09/13/2026,Virtual,Alice,0,Scheduled,Missing signed attendance\n"))
    review = reviewed(data, "events")
    assert "Training date is missing" in " ".join(review.rows[0].issues)
    assert "Hours must be positive" in " ".join(review.rows[0].issues)
    assert "Source status does not show completed" in " ".join(review.rows[0].issues)
    before = review
    updated = correct_training_row(
        review, review.rows[0].key,
        {"Date": "2026-09-13", "Hours": "", "Completion status": "Completed"},
        expected_revision=review.revision, profile=PROFILE, period=PERIOD)
    assert updated.ready_for_import
    assert updated.rows[0].values == (
        "Pump safety", "2026-09-13", "Virtual", "Alice", "")
    assert updated.rows[0].origin.cells[1] == "09/13/2026"
    assert updated.rows[0].origin.cells[4] == "0"
    assert updated.rows[0].origin.cells[5] == "Scheduled"
    assert updated.rows[0].corrections
    assert reimport_training_review(updated, before) == updated
    with pytest.raises(ValueError, match="changed"):
        correct_training_row(updated, updated.rows[0].key, {"Hours": "3"},
            expected_revision=before.revision, profile=PROFILE, period=PERIOD)
    assert sources.source_bytes(PROFILE, data.source).startswith(b"Topic,Training Date")


def test_annual_2025_matrix_requires_explicit_year_and_resets_old_completion(storage):
    source = ingest("annual_requirements.csv", MATRIX_OLD)
    missing_year = reviewed(source, "requirements")
    assert not missing_year.ready_for_import
    assert "source year" in " ".join(missing_year.issues)
    review = reviewed(source, "requirements", source_year=2025)
    assert review.ready_for_import
    assert review.rows[0].values == ("Alice", "Site North", "Completed", "")
    assert review.rows[1].values == ("Jo", "Site South", "Pending", "Not required")
    block = ResolvedBlock("training_summary", "This month", text="Keep this month's note")
    with pytest.raises(ValueError, match="older requirements"):
        accepted_training_block(review, block, expected_block_fingerprint=block.fingerprint)
    updated = accepted_training_block(
        review, block, expected_block_fingerprint=block.fingerprint,
        carry_prior_requirements=True)
    matrix, = updated.extra_tables
    assert matrix.reference == MATRIX_REF
    assert matrix.columns == ("Team member", "Site", "Safety refresher", "Operations update")
    assert matrix.rows == (("Alice", "Site North", "Not recorded", ""),
                           ("Jo", "Site South", "Not recorded", "Not recorded"))
    assert updated.references == (review.source.reference,)
    assert not load_training(PROFILE.contract)["revision"]
    assert sources.source_bytes(PROFILE, source.source) == MATRIX_OLD


def test_explicit_same_year_standing_values_are_preserved_but_not_events(storage):
    content = ingest("current-requirements.csv", MATRIX_OLD)
    review = reviewed(content, "requirements", source_year=2026)
    block = ResolvedBlock("training_summary", "This month")
    updated = accepted_training_block(
        review, block, expected_block_fingerprint=block.fingerprint)
    matrix, = updated.extra_tables
    assert matrix.rows == (("Alice", "Site North", "Completed", ""),
                           ("Jo", "Site South", "Pending", "Not required"))
    assert not any(t.reference == EVENT_REF for t in updated.extra_tables)
    assert events_for_period(updated, PERIOD) is None
    with pytest.raises(ValueError, match="different standing requirements"):
        accepted_training_block(
            replace(review, rows=tuple(replace(row, values=(
                *row.values[:2], "Pending", *row.values[3:]))
                for row in review.rows)),
            updated, expected_block_fingerprint=updated.fingerprint)
    assert sources.source_bytes(PROFILE, content.source) == MATRIX_OLD


def test_import_current_events_preserves_existing_matrix_and_authored_session(storage):
    content = ingest("current.csv", (
        b"Topic,Date,Method,Participants,Hours\n"
        b"Pump safety,2026-09-13,Virtual,Alice; Bob,1.5\n"))
    review = reviewed(content, "events")
    manual = add_event(None, title="Manual equipment review", when=date(2026, 9, 5),
                       mode="In person", participants=("Jo",), hours="", period=PERIOD)
    matrix = ReportTable(("Team member", "Site", "Safety"),
                         (("Jo", "Site South", "Pending"),), MATRIX_REF)
    block = ResolvedBlock("training_summary", "This month", text="Keep manuscript",
                          extra_tables=(matrix, manual), references=("original-note",))
    updated = accepted_training_block(
        review, block, expected_block_fingerprint=block.fingerprint)
    assert updated.extra_tables[0] == matrix
    events = updated.extra_tables[1]
    assert len(events.rows) == 2
    assert events.rows[0] == manual.rows[0]
    assert events.rows[1][0] == "Pump safety"
    assert session_row_tokens(events)[0] == session_row_tokens(manual)[0]
    assert updated.text == block.text
    assert updated.references == ("original-note", review.source.reference)
    next_month = new_month_training(updated)
    next_month_again = new_month_training(next_month)
    assert any(t.reference == MATRIX_REF for t in next_month.extra_tables)
    assert not any(t.reference == EVENT_REF for t in next_month.extra_tables)
    assert not any(t.reference == EVENT_REF for t in next_month_again.extra_tables)
    assert sources.source_bytes(PROFILE, content.source)


def _word_document():
    doc = Document()
    doc.add_heading("1 Organizational Chart", 1)
    unrelated = doc.add_table(rows=2, cols=2)
    unrelated.cell(0, 0).text = "Name"
    unrelated.cell(0, 1).text = "Role"
    unrelated.cell(1, 0).text = "Unrelated person"
    unrelated.cell(1, 1).text = "Technician"
    doc.add_heading("11 Training Summary", 1)
    annual = doc.add_table(rows=4, cols=4)
    annual.rows[0].cells[0].merge(annual.rows[0].cells[-1]).text = "Annual Training Requirements 2025"
    for cell, label in zip(annual.rows[1].cells,
                           ("Team member", "Site", "Safety refresher", "Operations update")):
        cell.text = label
    for n, values in enumerate((("Alice", "North", "Completed", ""),
                                ("Jo", "South", "Pending", "Not required")), 2):
        for cell, value in zip(annual.rows[n].cells, values):
            cell.text = value
    event = doc.add_table(rows=4, cols=7)
    for cell, label in zip(event.rows[0].cells,
                           ("Topic", "Date", "Format", "Participants", "Hours",
                            "Status", "Facilitator")):
        cell.text = label
    for n, values in enumerate((
        ("Pump safety", "2026-09-13", "Virtual", "Alice", "1.5", "Completed", "Doctor A"),
        ("Boiler procedures", "2025-09-14", "In person", "Jo", "2", "Completed", "Doctor B"),
        ("Incomplete", "", "", "", "", "", "")), 1):
        for cell, value in zip(event.rows[n].cells, values):
            cell.text = value
    return doc


def test_actual_prior_word_training_tables_keep_source_items_and_full_native_rows(storage):
    path = storage / "prior-report.docx"
    _word_document().save(path)
    raw = path.read_bytes()
    content = ingest(path.name, raw)
    choices = training_tables(content, PROFILE)
    assert len(choices) == 2
    assert all(choice.native_word and choice.word_section == "training" for choice in choices)
    assert choices[0].number < choices[1].number
    assert all(choice.word_item_id for choice in choices)
    annual = review_training_table(choices[0], PROFILE, PERIOD, source_year=2025)
    sessions = review_training_table(choices[1], PROFILE, PERIOD)
    assert annual.kind == "requirements" and sessions.kind == "events"
    assert len(annual.rows) == 2 and len(sessions.rows) == 3
    assert annual.rows[0].origin.row_number == 3
    assert annual.rows[0].origin.source_reference == choices[0].reference
    assert sessions.rows[0].origin.source_reference == choices[1].reference
    assert sessions.unmapped_cells(sessions.rows[0].key) == (
        (6, "Facilitator", "Doctor A"),)
    assert not annual.ready_for_import and not sessions.ready_for_import
    annual = confirm_training_word_layout(annual, expected_revision=annual.revision)
    sessions = confirm_training_word_layout(sessions, expected_revision=sessions.revision)
    for row in sessions.rows[1:]:
        sessions = exclude_training_row(sessions, row.key, expected_revision=sessions.revision)
    assert annual.ready_for_import and sessions.ready_for_import
    block = ResolvedBlock("training_summary", "This month", text="Authored notes")
    block = accepted_training_block(
        annual, block, expected_block_fingerprint=block.fingerprint,
        carry_prior_requirements=True)
    block = accepted_training_block(
        sessions, block, expected_block_fingerprint=block.fingerprint)
    assert len(block.extra_tables) == 2
    assert [t.reference for t in block.extra_tables] == [MATRIX_REF, EVENT_REF]
    assert len(block.references) == 2
    assert block.extra_tables[1].rows == ((
        "Pump safety", "2026-09-13", "Virtual", "Alice", "1.5"),)
    assert "Boiler procedures" in content.source.page_texts[choices[1].number-1]
    assert sources.source_bytes(PROFILE, content.source) == raw


def test_word_source_rows_cannot_be_published_without_layout_review(storage):
    path = storage / "training-source.docx"
    _word_document().save(path)
    content = ingest(path.name, path.read_bytes())
    src = training_tables(content, PROFILE)[1]
    review = review_training_table(src, PROFILE, PERIOD)
    review = exclude_training_row(review, review.rows[1].key, expected_revision=review.revision)
    review = exclude_training_row(review, review.rows[2].key, expected_revision=review.revision)
    assert not review.ready_for_import
    with pytest.raises(ValueError, match="source layout"):
        accepted_training_block(
            review, ResolvedBlock("training_summary", "This month"),
            expected_block_fingerprint=ResolvedBlock("training_summary", "This month").fingerprint)
    confirmed = confirm_training_word_layout(review, expected_revision=review.revision)
    assert confirmed.ready_for_import
    with pytest.raises(ValueError, match="changed"):
        confirm_training_word_layout(confirmed, expected_revision=review.revision)


def test_saved_review_receipt_and_source_row_fingerprints_survive_round_trip(storage):
    content = ingest("events.csv", CURRENT)
    review = reviewed(content, "events")
    updated = correct_training_row(
        review, review.rows[2].key,
        {"Date": "2026-09-15", "Participants": "Alice", "Hours": "1",
         "Completion status": "Completed"},
        expected_revision=review.revision, profile=PROFILE, period=PERIOD)
    stored = json.loads(json.dumps(training_review_record(updated)))
    restored = restore_training_review(stored)
    assert restored == updated
    assert restored.rows[2].origin.cells[1] == ""
    assert restored.rows[2].origin.cells[4] == "0"
    stored["review"]["rows"][0]["origin"]["cells"][0] = "rewritten"
    with pytest.raises(ValueError, match="integrity"):
        restore_training_review(stored)
    assert sources.source_bytes(PROFILE, content.source) == CURRENT


def test_duplicate_source_sessions_need_review_not_overwrite(storage):
    content = ingest("duplicate.csv", (
        b"Topic,Date,Mode,Attendees,Hours\n"
        b"Pump safety,2026-09-13,Virtual,Alice,1\n"
        b"Pump safety,2026-09-13,Virtual,Alice,2\n"))
    review = reviewed(content, "events")
    assert len(review.rows) == 2
    assert all("Duplicate Training source row" in " ".join(r.issues) for r in review.rows)
    assert not review.ready_for_import
    before = review
    review = exclude_training_row(review, review.rows[1].key,
                                  expected_revision=review.revision)
    assert not review.ready_for_import or review.rows[0].issues
    with pytest.raises(ValueError, match="Review unknown"):
        accepted_training_block(review, ResolvedBlock("training_summary", "This month"),
                expected_block_fingerprint=ResolvedBlock("training_summary", "This month").fingerprint)
    assert reimport_training_review(review, before) == review


def test_event_conflict_with_manual_same_identity_requires_review(storage):
    content = ingest("manual-conflict.csv", (
        b"Training,Date,Format,Participants,Hours\n"
        b"Pump safety,2026-09-13,Virtual,Alice,2\n"))
    review = reviewed(content, "events")
    manual = add_event(None, title="Pump safety", when=date(2026, 9, 13),
                       mode="Virtual", participants=("Alice",), hours="1", period=PERIOD)
    block = ResolvedBlock("training_summary", "This month",
                          extra_tables=(manual,))
    with pytest.raises(ValueError, match="manually edited"):
        accepted_training_block(review, block,
                                expected_block_fingerprint=block.fingerprint)
    assert manual.rows[0][-1] == "1"


def test_unsupported_totals_grid_and_source_replacement_are_safe(storage):
    content = ingest("aggregate.csv", (
        b"Metric,January,February,March,Total\n"
        b"Training count,1,0,,1\n"))
    review = review_training_table(training_tables(content, PROFILE)[0],
                                   PROFILE, PERIOD)
    assert review.kind == "unresolved"
    assert not review.ready_for_import
    assert review.issues
    source = ingest("events.csv", CURRENT)
    before = reviewed(source, "events")
    changed = ingest("other.csv", CURRENT.replace(
        b"Pump safety", b"Different safety"))
    after = reviewed(changed, "events")
    with pytest.raises(ValueError, match="source"):
        reimport_training_review(before, after)
    with pytest.raises(ValueError, match="original"):
        training_tables(replace(source, source=replace(
            source.source, sha256="a" * 64)), PROFILE)


def test_xlsx_event_export_uses_saved_original_and_keeps_blank_values(storage):
    book = Workbook()
    worksheet = book.active
    worksheet.append(("Topic", "Session Date", "Method", "Participants", "Hours"))
    worksheet.append(("Pump safety", "2026-09-13", "Virtual", "Alice", None))
    raw_buffer = BytesIO()
    book.save(raw_buffer)
    raw = raw_buffer.getvalue()
    source = ingest("training.xlsx", raw)
    review = reviewed(source, "events")
    assert review.ready_for_import
    assert review.rows[0].values[-1] == ""
    assert sources.source_bytes(PROFILE, source.source) == raw


def test_report_block_and_draft_roundtrip_preserves_training_sessions_and_notes(storage):
    from app.monthly_report_model import ReportDraft
    content = ingest("current.csv", (
        b"Topic,Date,Method,Participants,Hours\n"
        b"Pump safety,2026-09-13,Virtual,Alice,1.5\n"))
    review = reviewed(content, "events")
    block = ResolvedBlock("training_summary", "This month", text="A human review note")
    approved = accepted_training_block(review, block, expected_block_fingerprint=block.fingerprint)
    section = next(s for s in default_sections() if s.key == "training")
    draft = ReportDraft(PROFILE, PERIOD, "Synthetic Editor", (section,), (approved,),
                        sources=(content.source,), synthetic=True)
    restored = library.draft_from_dict(json.loads(json.dumps(asdict(draft))))
    assert restored == draft
    assert restored.blocks[0].extra_tables[0].reference == EVENT_REF
    assert restored.blocks[0].references == (review.source.reference,)
    later = new_month_training(restored.blocks[0])
    assert later.text == "" and not later.extra_tables
    assert sources.source_bytes(PROFILE, content.source)


def test_accepted_current_event_renders_once_in_generic_word_training_section(storage):
    from app.monthly_report_docx import assemble_docx
    profile = ReportProfile(PROFILE.contract, "single", "Training Site North",
                            (PROFILE.facilities[0],))
    content = ingest("training.csv", (
        b"Topic,Date,Format,Participants,Hours\n"
        b"Safety drill,2026-09-13,Virtual,Alice,1.5\n"), profile=profile)
    review = reviewed(content, "events", profile=profile)
    block = ResolvedBlock("training_summary", "This month")
    approved = accepted_training_block(review, block,
                                      expected_block_fingerprint=block.fingerprint)
    original = replace(content.source, facility=profile.facilities[0].title)
    section = next(s for s in default_sections() if s.key == "training")
    from app.monthly_report_model import ReportDraft
    draft = ReportDraft(profile, PERIOD, "Synthetic Editor",
                        (section,), (approved,), sources=(original,), synthetic=True)
    raw = assemble_docx(draft, acknowledged_fingerprint=draft.fingerprint)
    doc = Document(BytesIO(raw))
    data = "\n".join(cell.text for table in doc.tables
                     for row in table.rows for cell in row.cells)
    assert data.count("Safety drill") == 1
    assert "2026-09-13" in data and "Virtual" in data
    assert "Alice" in data and "1.5" in data
    assert not any("2025-09" in paragraph.text for paragraph in doc.paragraphs)
