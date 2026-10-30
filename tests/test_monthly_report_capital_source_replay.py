"""Exact carried Capital sources do not duplicate tables or revive exclusions."""
from dataclasses import asdict, replace
from pathlib import Path

import pytest
from docx import Document

from app import monthly_report_library as library
from app.monthly_report_import import inspect_docx, imported_draft
from app.monthly_report_model import ReportPeriod, ReportTable, ResolvedBlock, synthetic_profiles
from app.monthly_report_setup import new_month_draft
from app.monthly_report_start_import import (
    analyze_starting_report, automatic_section_plans, build_starting_import,
)
from app.monthly_report_start_replay import merge_starting_drafts

PERIOD = ReportPeriod(2026, 9)
NEXT = ReportPeriod(2026, 10)
KEYS = ("capital_renewal", "end_of_life")


def block(draft, key):
    return next(item for item in draft.blocks if item.key == key)


def change(draft, key, **fields):
    return replace(draft, blocks=tuple(replace(item, **fields) if item.key == key else item
                                       for item in draft.blocks))


def import_source(path, profile, period):
    inspection = inspect_docx(path)
    analysis = analyze_starting_report(inspection, profile, period)
    assert analysis.site_relation == "same"
    plans = automatic_section_plans(inspection, analysis)
    assert all(plan["approved"] for plan in plans.values())
    mapped, _ = build_starting_import(path, inspection, plans.values())
    return imported_draft(profile, period, "Synthetic Editor", mapped)


@pytest.fixture
def source(tmp_path, monkeypatch):
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path / "data"))
    profile = synthetic_profiles()[0]
    doc = Document()
    doc.add_heading(profile.facilities[0].title, 0)
    doc.add_paragraph("Monthly Report September 2026")
    doc.add_heading("9 Priority Capital Renewal List", 1)
    doc.add_paragraph("Source planning context stays with the list.")
    tables = (
        (("Facility", "Priority", "Recommendation", "Status", "Asset", "Source note"),
         ((profile.facilities[0].title, "1", "Renew north pump", "Open", "0007", "First occurrence"),
          (profile.facilities[0].title, "2", "Review roof", "Closed", "0007", "Distinct source row"),
          (profile.facilities[0].title, "", "Review controls", "", "0010", "0"))),
        (("Facility", "Asset", "Expected replacement timing", "Source status"),
         ((profile.facilities[0].title, "CH-001", "Not established", "Open"),
          (profile.facilities[0].title, "CH-002", "", "Closed"))),
    )
    for number, (columns, rows) in enumerate(tables):
        if number:
            doc.add_heading("Asset End of Useful Life Schedule", 2)
        grid = doc.add_table(rows=len(rows) + 1, cols=len(columns))
        for row, values in zip(grid.rows, (columns, *rows)):
            for cell, value in zip(row.cells, values):
                cell.text = value
    path = tmp_path / "synthetic-capital.docx"
    doc.save(path)
    original = path.read_bytes()
    draft = import_source(path, profile, PERIOD)
    assert len(block(draft, KEYS[0]).extra_tables[0].rows) == 3
    assert len(block(draft, KEYS[1]).extra_tables[0].rows) == 2
    return path, original, draft


@pytest.mark.parametrize("key", KEYS)
@pytest.mark.parametrize("edit", ("correct", "remove", "clear", "header"))
def test_conflicting_reimport_in_new_month_is_rejected_without_duplicate(source, key, edit):
    path, raw, initial = source
    table = block(initial, key).extra_tables[0]
    if edit == "correct":
        rows = (tuple("Manual correction" if n == 2 else value for n, value in enumerate(table.rows[0])), *table.rows[1:])
        edited = replace(table, rows=rows)
    elif edit == "remove":
        edited = replace(table, rows=table.rows[1:])
    elif edit == "clear":
        edited = replace(table, rows=())
    else:
        edited = replace(table, columns=(*table.columns[:-1], "Reviewed source field"))
    current = change(initial, key, extra_tables=(edited,), text="Authored planning notes.")
    carried = new_month_draft(current, NEXT)
    incoming = import_source(path, initial.profile, NEXT)
    before = asdict(carried)
    # A new-month receipt does not establish which competing review owns the
    # changed cells. Refuse rather than silently ignoring either choice.
    with pytest.raises(ValueError, match="Capital.*already contains") as error:
        merge_starting_drafts(carried, incoming)
    assert "edit the current list" in str(error.value)
    assert len(block(carried, key).extra_tables) == 1
    assert asdict(carried) == before and path.read_bytes() == raw


@pytest.mark.parametrize("key", KEYS)
def test_reimport_keeps_explicit_report_omission(source, key):
    path, raw, initial = source
    carried = new_month_draft(change(initial, key, source="Omit", text="Retained manual notes."), NEXT)
    incoming = import_source(path, initial.profile, NEXT)
    with pytest.raises(ValueError, match="Capital.*already contains"):
        merge_starting_drafts(carried, incoming)
    assert block(carried, key).source == "Omit"
    assert block(carried, key).text == "Retained manual notes."
    assert path.read_bytes() == raw


def test_reimport_keeps_excluded_capital_section(source):
    path, _, initial = source
    carried = new_month_draft(replace(initial, sections=tuple(
        replace(s, included=False) if s.key == "capital" else s for s in initial.sections)), NEXT)
    with pytest.raises(ValueError, match="Capital.*already contains"):
        merge_starting_drafts(carried, import_source(path, initial.profile, NEXT))
    assert not next(s for s in carried.sections if s.key == "capital").included


def test_known_receipt_still_allows_explicit_review_change(source):
    _, _, initial = source
    key = KEYS[0]
    old = block(initial, key)
    changed = replace(old.extra_tables[0], rows=old.extra_tables[0].rows[1:])
    incoming = change(initial, key, extra_tables=(changed,))
    result = merge_starting_drafts(initial, incoming, initial)
    assert block(result, key).extra_tables == (changed,)


def test_known_receipt_still_refuses_concurrent_edits(source):
    _, _, initial = source
    key = KEYS[0]
    table = block(initial, key).extra_tables[0]
    current = change(initial, key, extra_tables=(replace(table, rows=table.rows[1:]),))
    incoming = change(initial, key, extra_tables=(replace(table, rows=table.rows[:1]),))
    with pytest.raises(ValueError, match="Capital renewal differs"):
        merge_starting_drafts(current, incoming, initial)


def test_different_original_or_physical_table_is_not_deduplicated(source):
    _, _, initial = source
    key = KEYS[0]
    old = block(initial, key)
    for reference in ("docx:" + "b" * 64 + ":item-0004", old.extra_tables[0].reference + "-other"):
        other = replace(old.extra_tables[0], reference=reference)
        incoming = replace(initial, blocks=(replace(old, extra_tables=(other,), references=(reference,)),))
        result = merge_starting_drafts(initial, incoming)
        assert block(result, key).extra_tables == (*old.extra_tables, other)


@pytest.mark.parametrize("reference", ("", "source-1", "docx:not-a-hash:item-0001"))
def test_unproved_reference_keeps_existing_generic_merge(source, reference):
    _, _, initial = source
    key = KEYS[0]
    table = replace(block(initial, key).extra_tables[0], reference=reference)
    current = change(initial, key, extra_tables=(replace(table, rows=table.rows[1:]),), references=(reference,))
    incoming = change(initial, key, extra_tables=(table,), references=(reference,))
    assert len(block(merge_starting_drafts(current, incoming), key).extra_tables) == 2


def test_same_source_replay_never_crosses_contract_or_selected_sites(source):
    path, _, initial = source
    incoming = import_source(path, initial.profile, PERIOD)
    foreign = replace(initial, profile=replace(initial.profile, contract="Different synthetic contract"))
    with pytest.raises(ValueError, match="different sites or a different month"):
        merge_starting_drafts(foreign, incoming)
    other_sites = replace(initial, profile=replace(initial.profile, facilities=synthetic_profiles()[1].facilities,
                                                   scope_type="multi_site"))
    with pytest.raises(ValueError, match="Capital.*sites"):
        merge_starting_drafts(other_sites, incoming)


def test_unchanged_carried_source_is_reused_once_without_reentry(source):
    path, raw, initial = source
    current = new_month_draft(initial, NEXT)
    incoming = import_source(path, initial.profile, NEXT)
    first = merge_starting_drafts(current, incoming)
    second = merge_starting_drafts(first, incoming)
    for key in KEYS:
        assert block(first, key).extra_tables == block(initial, key).extra_tables
        assert block(second, key) == block(first, key)
    assert path.read_bytes() == raw


def test_cleared_notes_cannot_be_restored_by_repeated_source(source):
    path, _, initial = source
    current = change(new_month_draft(initial, NEXT), KEYS[0], text="")
    with pytest.raises(ValueError, match="Capital.*already contains"):
        merge_starting_drafts(current, import_source(path, initial.profile, NEXT))
    assert block(current, KEYS[0]).text == ""


def test_existing_ambiguous_copies_require_review_instead_of_another_append(source):
    _, _, initial = source
    table = block(initial, KEYS[0]).extra_tables[0]
    current = change(initial, KEYS[0], extra_tables=(table, replace(table, rows=table.rows[1:])))
    with pytest.raises(ValueError, match="Capital.*already contains"):
        merge_starting_drafts(current, initial)
    assert len(block(current, KEYS[0]).extra_tables) == 2


def test_new_physical_table_can_append_beside_an_unchanged_consumed_table(source):
    _, _, initial = source
    old = block(initial, KEYS[0])
    ref = old.extra_tables[0].reference.rsplit(":", 1)[0] + ":item-0999"
    extra = replace(old.extra_tables[0], reference=ref, rows=old.extra_tables[0].rows[:1])
    incoming = change(initial, KEYS[0], extra_tables=(*old.extra_tables, extra),
                      references=(*old.references, ref))
    result = merge_starting_drafts(initial, incoming)
    assert block(result, KEYS[0]).extra_tables == (*old.extra_tables, extra)
    assert merge_starting_drafts(result, incoming) == result


def test_table_reference_alone_does_not_invent_a_consumed_receipt(source):
    _, _, initial = source
    key = KEYS[0]
    table = block(initial, key).extra_tables[0]
    current = change(initial, key, extra_tables=(replace(table, rows=table.rows[1:]),), references=())
    result = merge_starting_drafts(current, initial)
    assert len(block(result, key).extra_tables) == 2


def test_first_import_snapshot_reopen_and_two_rollovers_keep_current_values(source):
    path, raw, initial = source
    digest = inspect_docx(path).sha256
    state = library.save_report_setup(initial.profile, initial, path,
        {"sha256": digest, "period": PERIOD.key}, assets=(), expected_revision=0,
        actor="Synthetic Editor", confirmed=True, source_draft=initial)
    key = KEYS[0]
    table = block(initial, key).extra_tables[0]
    edited = replace(table, rows=(table.rows[1], (*table.rows[2][:-1], "Retained manual note")))
    current = change(new_month_draft(initial, NEXT), key, extra_tables=(edited,))
    library.save_snapshot(current, expected_revision=0)
    saved = library.load_snapshot(initial.profile.contract, initial.profile.key, NEXT)
    assert saved.draft == current
    assert library.starting_contribution(initial.profile.contract, initial.profile.key, NEXT, digest) == (None, False)
    with pytest.raises(ValueError, match="Capital.*already contains"):
        merge_starting_drafts(saved.draft, import_source(path, initial.profile, NEXT))
    assert library.load_snapshot(initial.profile.contract, initial.profile.key, NEXT) == saved
    later = new_month_draft(saved.draft, ReportPeriod(2026, 11))
    assert block(later, key).extra_tables == (edited,)
    assert block(initial, key).extra_tables == (table,)
    assert library.load_profile(initial.profile.contract, initial.profile.key) == state
    assert library.imported_original(initial.profile.contract, initial.profile.key).read_bytes() == raw
    foreign = replace(initial.profile, contract="Other synthetic contract")
    library.save_profile(foreign, expected_revision=0, actor="Synthetic Editor", confirmed=True)
    assert library.load_snapshot(foreign.contract, foreign.key, NEXT) is None
    assert library.load_profile(foreign.contract, foreign.key).block(key) is None


def test_real_continue_shows_named_conflict_and_does_not_save(source, monkeypatch):
    from io import BytesIO
    from streamlit.testing.v1 import AppTest
    from app import monthly_report_section_ui as ui

    path, raw, initial = source
    digest = inspect_docx(path).sha256
    library.save_report_setup(initial.profile, initial, path,
        {"sha256": digest, "period": PERIOD.key}, assets=(), expected_revision=0,
        actor="Synthetic Editor", confirmed=True, source_draft=initial)
    key = KEYS[0]
    table = block(initial, key).extra_tables[0]
    current = change(new_month_draft(initial, NEXT), key,
                     extra_tables=(replace(table, rows=table.rows[1:]),))
    library.save_snapshot(current, expected_revision=0)
    saved = library.load_snapshot(initial.profile.contract, initial.profile.key, NEXT)
    state = library.load_profile(initial.profile.contract, initial.profile.key)
    upload = BytesIO(raw)
    upload.name, upload.size = path.name, len(raw)
    real = ui.st.file_uploader
    monkeypatch.setattr(ui.st, "file_uploader", lambda label, *a, **kw:
                        upload if label == "Starting report" else real(label, *a, **kw))
    app = AppTest.from_string('''
import streamlit as st
from app import monthly_report_library as library
from app.monthly_report_section_ui import render_section_setup
from app.monthly_report_ui import _field
current = st.session_state["test_draft"]
state = library.load_profile(current.profile.contract, current.profile.key)
render_section_setup(current.profile.contract, current.period, "Synthetic Editor", _field,
                     state=state, identity=current.profile, working_draft=current,
                     working_revision=1)
''', default_timeout=30)
    app.session_state["test_draft"] = current
    app.run()
    assert not app.exception
    button = next(b for b in app.button if b.label == "Continue to this month’s updates")
    assert not button.disabled
    attempted = []
    monkeypatch.setattr(library, "save_report_setup", lambda *a, **kw: attempted.append(kw))
    button.click().run()
    assert not app.exception and not attempted
    assert any("Capital renewal already contains" in e.value for e in app.error)
    assert any(key.endswith("_stage") for key in app.session_state.filtered_state)
    assert app.session_state["test_draft"] == current
    assert library.load_snapshot(initial.profile.contract, initial.profile.key, NEXT) == saved
    assert library.load_profile(initial.profile.contract, initial.profile.key) == state
    assert library.imported_original(initial.profile.contract, initial.profile.key).read_bytes() == raw
    app.run()
    assert not app.exception and not attempted
    assert app.session_state["test_draft"] == current


def test_unchanged_source_can_replay_beside_added_manual_notes(source):
    _, _, initial = source
    current = change(initial, KEYS[0], text=block(initial, KEYS[0]).text + "\n\nAdditional manual planning note.")
    result = merge_starting_drafts(current, initial)
    assert block(result, KEYS[0]).text == block(current, KEYS[0]).text
    assert block(result, KEYS[0]).extra_tables == block(current, KEYS[0]).extra_tables
