"""Separate capacity tables keep their identity across updates (Refs #286, G9-pr89-4).

A site may keep a Steam table and a Chilled-water table with different columns.
Re-uploading them, or updating only one, must not look like a column change,
and Replace must never delete the other service. Synthetic values only.
"""
from io import BytesIO

from openpyxl import Workbook
import pytest

from app import monthly_report_capacity as capacity, monthly_report_directory as directory
from app.monthly_report_model import Facility, ReportProfile

FACILITIES = (Facility("north", "North Hospital", ("NH",)), Facility("south", "South Hospital", ("SH",)))
PROFILE = ReportProfile("Synthetic Capacity Contract", "north", "North Hospital", FACILITIES[:1])
STEAM = ["Facility", "Required steam (lb/hr)", "Available steam (lb/hr)"]
CHILLED = ["Facility", "Chillers", "Required cooling (tons)", "Available cooling (tons)"]


@pytest.fixture(autouse=True)
def isolated(tmp_path, monkeypatch):
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path))
    directory.save_directory(PROFILE.contract, tuple(directory.DirectorySite(f.key, f.title, f.aliases) for f in FACILITIES),
                             expected_revision=0, actor="Synthetic Reviewer", confirmed=True)


def workbook(sheets):
    book = Workbook()
    book.remove(book.active)
    for name, rows in sheets:
        sheet = book.create_sheet(name)
        for row in rows:
            sheet.append(row)
    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()


BOTH = workbook([("Steam", [STEAM, ["NH", 40000, 50000], ["SH", 20000, 24000]]),
                 ("Chilled water", [CHILLED, ["NH", "CH-1, CH-2", 1200, 1600], ["SH", "CH-3", 800, 1000]])])


def inspect(raw):
    return capacity.inspect_capacity(PROFILE, "capacity.xlsx", raw).tables


def save(raw, revision, mode="merge"):
    tables = inspect(raw)
    return capacity.save_capacity(PROFILE, tables, expected_revision=revision, actor="Synthetic Reviewer",
                                  reviewed_fingerprint=capacity.review_fingerprint(tables), mode=mode)


def rows(state, site):
    return sorted((table.title, row) for table in state.tables
                  for row, key in zip(table.rows, table.site_keys) if key == site)


def test_identical_reupload_of_two_differently_shaped_tables_is_not_a_conflict():
    saved = save(BOTH, 0)
    assert capacity.schema_conflicts(saved, inspect(BOTH)) == ()
    again = save(BOTH, 1)
    assert rows(again, "north") == rows(saved, "north") and rows(again, "south") == rows(saved, "south")


@pytest.mark.parametrize("mode", ["merge", "replace_sites"])
def test_updating_one_service_keeps_the_other_service_and_other_sites(mode):
    saved = save(BOTH, 0)
    steam_only = workbook([("Steam", [STEAM, ["NH", 42000, 52000]])])
    assert capacity.schema_conflicts(saved, inspect(steam_only)) == ()
    updated = save(steam_only, 1, mode=mode)
    assert rows(updated, "north") == [
        ("Chilled water", ("NH", "CH-1, CH-2", "1200", "1600")),
        ("Steam", ("NH", "42000", "52000"))]
    assert rows(updated, "south") == rows(saved, "south")


def test_a_heading_that_changes_its_columns_is_still_a_reviewed_conflict():
    saved = save(BOTH, 0)
    changed = workbook([("Steam", [["Facility", "Steam demand (lb/hr)", "Installed steam (lb/hr)", "Units"],
                                   ["NH", 41000, 55000, "lb/hr"]])])
    assert capacity.schema_conflicts(saved, inspect(changed)) == ("north",)
    with pytest.raises(ValueError, match="different columns"):
        save(changed, 1)
    replaced = save(changed, 1, mode="replace_sites")
    assert rows(replaced, "north") == [
        ("Chilled water", ("NH", "CH-1, CH-2", "1200", "1600")),
        ("Steam", ("NH", "41000", "55000", "lb/hr"))]
    assert rows(replaced, "south") == rows(saved, "south")
