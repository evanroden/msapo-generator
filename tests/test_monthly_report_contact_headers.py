"""Contact headings are schema, never the first person or old client data.

Synthetic geometry exercises both writers independently. These checks do not
approve the owner's native master, its chart-page layout, or Microsoft Word.
"""
from dataclasses import asdict, replace
from io import BytesIO

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches
import pytest

from app.monthly_report_docx import _build_docx
from app.monthly_report_model import BlockSpec, ColumnSpec, ResolvedBlock
from app.monthly_report_native_layout import build_native_docx
from tests.conftest import requires_libreoffice
from tests.test_monthly_report_generic_contact_page import _draft


# Independent expectation: do not import the constant being tested.
HEADINGS = ("Facility", "Role", "Name", "Phone", "Email")


def contact_rows(count, *, long=False):
    return tuple((
        "Synthetic North",
        "Synthetic Regional Facilities Operations Lead" if long else "Engineer",
        f"QA{i:04d} Synthetic Contact",
        f"555-01{i:02d}",
        f"synthetic.person.{i:04d}@example.invalid" if long else "",
    ) for i in range(count))


def contact_draft(count, *, long=False, chart=True):
    value = _draft(chart=chart)
    return replace(value, blocks=tuple(
        replace(block, rows=contact_rows(count, long=long))
        if block.key == "contact_matrix" else block for block in value.blocks))


def native_contact_source(path):
    """A synthetic native table with independent old payload and geometry."""
    document = Document()
    document.sections[0].left_margin = document.sections[0].right_margin = Inches(.75)
    document.add_paragraph("OLD SYNTHETIC SITE")
    document.add_paragraph("August 2026")
    document.add_paragraph("Prepared by: OLD SYNTHETIC EDITOR")
    document.add_heading("ORGANIZATIONAL CHART", 1)
    table = document.add_table(rows=2, cols=5)
    table.style = "Light Shading Accent 1"
    for cell, text in zip(table.rows[0].cells, HEADINGS):
        cell.text = text
    for cell in table.rows[1].cells:
        cell.text = "OLD CONTACT PAYLOAD"
    document.save(path)
    return path


def assert_contact_table(raw, rows):
    document = Document(BytesIO(raw))
    candidates = [table for table in document.tables
                  if any(row[3] in cell.text for row in rows
                         for current in table.rows for cell in current.cells)]
    assert len(candidates) == 1, "Contacts were omitted, duplicated or fragmented into other schemas."
    table = candidates[0]
    values = tuple(tuple(cell.text for cell in row.cells) for row in table.rows)
    assert values == (HEADINGS, *rows), "A header must be inserted; no person may become a header."
    for index, row in enumerate(table.rows):
        headers = row._tr.xpath("./w:trPr/w:tblHeader")
        assert len(headers) == (1 if index == 0 else 0)
        assert row._tr.xpath("./w:trPr/w:cantSplit"), "A contact row must remain intact."
        if index:
            assert all(run.bold is not True for run in row.cells[2].paragraphs[0].runs)
    return document


def assert_contact_pdf(raw, rows):
    import fitz
    with fitz.open(stream=raw, filetype="pdf") as pdf:
        texts = [page.get_text() for page in pdf]
        compact = ["".join(text.split()) for text in texts]
        contact_pages = [index for index, text in enumerate(texts)
                         if any(row[3] in text for row in rows)]
        assert contact_pages
        for index in contact_pages:
            for heading in HEADINGS:
                assert heading in texts[index], f"Missing repeated {heading} heading on page {index + 1}."
        for row in rows:
            assert sum(text.count(row[3]) for text in texts) == 1, "Dropped or duplicated contact."
            assert sum("".join(row[2].split()) in text for text in compact) == 1
            if row[4]:
                assert sum(row[4] in text for text in compact) == 1
        if len(rows) >= 40:
            assert len(contact_pages) > 1, "The overflow case must actually span pages."
        return len(pdf), contact_pages


def test_default_and_legacy_image_capable_contact_schema_is_explicit_and_stable():
    spec = BlockSpec("contact_matrix", "image_page")
    assert tuple(column.title for column in spec.columns) == HEADINGS
    assert spec.type == "image_page"
    assert asdict(replace(spec)) == asdict(spec)
    assert BlockSpec("org_chart", "image_page").columns == ()
    custom = (ColumnSpec("name", "Contact person"), ColumnSpec("extension", "Extension"))
    assert BlockSpec("contact_matrix", "table", columns=custom).columns == custom
    assert BlockSpec("contact_matrix", "image_page", columns=custom).columns == custom


@pytest.mark.parametrize("count", [1, 2, 48])
def test_generic_contact_headers_preserve_every_person_and_repeat_only_real_heading(count):
    value = contact_draft(count, long=count == 48)
    before = value.fingerprint
    raw = _build_docx(value)
    assert_contact_table(raw, contact_rows(count, long=count == 48))
    assert value.fingerprint == before


def test_generic_empty_contacts_do_not_fabricate_a_person_table_or_page():
    document = Document(BytesIO(_build_docx(contact_draft(0))))
    assert not document.tables
    assert not any(p.text == "Facility contacts" for p in document.paragraphs)


@pytest.mark.parametrize("master", [False, True])
@pytest.mark.parametrize("count", [0, 1, 2, 48])
def test_native_writer_uses_current_contact_schema_not_source_person(tmp_path, master, count):
    path = native_contact_source(tmp_path / "synthetic-native.docx")
    raw = build_native_docx(contact_draft(count, long=count == 48, chart=False), path, master=master)
    document = Document(BytesIO(raw))
    assert "OLD CONTACT PAYLOAD" not in document.element.xml
    if count:
        assert_contact_table(raw, contact_rows(count, long=count == 48))


@requires_libreoffice
@pytest.mark.parametrize("native", [False, True])
def test_actual_pdf_contacts_keep_real_headers_across_overflow(tmp_path, monkeypatch, native):
    from app import pdf_converter
    monkeypatch.setattr(pdf_converter, "OUTPUT_DIR", tmp_path)
    rows = contact_rows(48, long=True)
    value = contact_draft(48, long=True, chart=not native)
    raw = (build_native_docx(value, native_contact_source(tmp_path / "native.docx"), master=True)
           if native else _build_docx(value))
    assert_contact_table(raw, rows)
    source = tmp_path / "current-contact-output.docx"
    source.write_bytes(raw)
    output = pdf_converter.convert_to_pdf(source)
    try:
        assert_contact_pdf(output.read_bytes(), rows)
    finally:
        output.unlink(missing_ok=True)
