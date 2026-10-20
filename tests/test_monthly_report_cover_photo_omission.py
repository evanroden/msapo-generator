"""An explicit source-gallery No-photo decision survives native layout and replay.

The actual native assembler below final preflight is inspected here. These
invented inputs are not genuine/master or accepted whole-report output proof.
"""
from dataclasses import replace
from hashlib import sha256
from io import BytesIO
from zipfile import ZipFile

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches
from PIL import Image, ImageDraw
import pytest
from streamlit.testing.v1 import AppTest

from app import monthly_report_designs as designs, monthly_report_library as library
from app.monthly_report_docx import _build_docx
from app.monthly_report_import import inspect_docx, imported_draft
from app.monthly_report_model import ReportDraft, ReportPeriod, ResolvedBlock, profile_sections, synthetic_profiles
from app.monthly_report_start_import import analyze_starting_report, automatic_section_plans, build_starting_import
from app.monthly_report_start_replay import merge_starting_drafts
from tests.test_monthly_report_start_import import start_app, save_start


def cover_source(tmp_path):
    stream = BytesIO()
    image = Image.new("RGB", (600, 400), "navy")
    ImageDraw.Draw(image).rectangle((80, 50, 520, 350), fill="orange")
    image.save(stream, "PNG")
    photo = stream.getvalue()
    stream = BytesIO()
    Image.new("RGB", (450, 70), "green").save(stream, "PNG")
    logo = stream.getvalue()
    document = Document()
    document.add_heading(synthetic_profiles()[0].facilities[0].title, 0)
    document.add_paragraph("September 2026")
    document.add_picture(BytesIO(photo), width=Inches(4))
    shape = document.sections[0].header.paragraphs[0].add_run().add_picture(BytesIO(logo), width=Inches(2))
    # The no-photo flow must retain this identified mark, not infer it from placement.
    shape._inline.docPr.set("descr", "ENFRA company logo")
    document.add_heading("Monthly Activity Summary", 1)
    document.add_paragraph("Synthetic current pump inspection.")
    path = tmp_path / "invented-cover-source.docx"
    document.save(path)
    return path, photo, logo


def photo_omission_plan(path):
    from app.monthly_report_start_import import cover_photo_candidates
    from app.monthly_report_sections import section_reviews
    inspection = inspect_docx(path)
    analysis = analyze_starting_report(inspection, synthetic_profiles()[0], ReportPeriod(2026, 9))
    plans = automatic_section_plans(inspection, analysis)
    cover = plans["cover"]
    candidates = cover_photo_candidates(next(section for section in section_reviews(inspection) if section.key == "cover").items)
    ids = {item.id for item in candidates}
    cover["selected"] = [identity for identity in cover["selected"] if identity not in ids]
    cover["destinations"] = {identity: slot for identity, slot in cover["destinations"].items() if identity not in ids}
    cover.update(cover_photo_omitted=True, cover_photo_source_ids=tuple(item.id for item in candidates), cover_picture_id="")
    return inspection, plans


def referenced_image_hashes(raw, *, headers=False):
    document = Document(BytesIO(raw))
    parts = [part for part in document.part.package.parts
             if str(part.partname).startswith("/word/header") and hasattr(part, "element")] if headers else [document.part]
    return tuple(sha256(part.rels[node.get(qn("r:embed"))].target_part.blob).hexdigest()
                 for part in parts for node in part.element.iter(qn("a:blip")))


def package_image_hashes(raw):
    with ZipFile(BytesIO(raw)) as package:
        return tuple(sha256(package.read(name)).hexdigest() for name in package.namelist() if name.startswith("word/media/"))


def report(photo=None):
    profile = synthetic_profiles()[0]
    return ReportDraft(profile, ReportPeriod(2026, 9), "Synthetic Editor", profile_sections(profile),
                       (photo,) if photo is not None else ())


def omission(source="a"):
    return ResolvedBlock("cover_photo", "Omit", references=("docx:" + source * 64 + ":photo",))


def test_real_gallery_no_photo_save_removes_visible_and_package_photo_from_native_docx(monkeypatch, tmp_path):
    source, photo, logo = cover_source(tmp_path)
    original = source.read_bytes()
    app = start_app(monkeypatch, tmp_path, source)
    next(button for button in app.button if button.label == "No cover photograph").click().run()
    app.run()
    draft = save_start(app)
    assert not app.exception
    omitted = next(block for block in draft.blocks if block.key == "cover_photo")
    assert omitted.source == "Omit" and not omitted.asset_hashes
    assert all(ref.startswith("docx:" + sha256(original).hexdigest() + ":") for ref in omitted.references)
    assert designs.source_for(draft.profile).read_bytes() == original
    raw = _build_docx(draft, asset_loader=lambda ref: library.read_asset(draft.profile.contract, draft.profile.key, ref))
    assert sha256(photo).hexdigest() not in referenced_image_hashes(raw)
    assert sha256(photo).hexdigest() not in package_image_hashes(raw)
    assert referenced_image_hashes(raw, headers=True) == (sha256(logo).hexdigest(),)
    assert library.imported_original(draft.profile.contract, draft.profile.key).read_bytes() == original
    receipt, verified = library.starting_contribution(draft.profile.contract, draft.profile.key, draft.period, sha256(original).hexdigest())
    assert verified and next(block for block in receipt.blocks if block.key == "cover_photo").source == "Omit"


def test_absent_or_default_unselected_gallery_is_not_an_explicit_no_photo(tmp_path):
    source, _, _ = cover_source(tmp_path)
    inspection, plans = photo_omission_plan(source)
    plans["cover"].pop("cover_photo_omitted")
    mapped, _ = build_starting_import(source, inspection, plans.values())
    assert not any(block.key == "cover_photo" for block in mapped.blocks)
    assert not any(block.source == "Omit" for block in mapped.blocks)


@pytest.mark.parametrize("tamper", ["hash", "candidates", "selected"])
def test_no_photo_intent_requires_its_exact_source_and_candidate_scope(tmp_path, tamper):
    source, _, _ = cover_source(tmp_path)
    inspection, plans = photo_omission_plan(source)
    plan = plans["cover"]
    if tamper == "hash":
        plan["source_sha256"] = "f" * 64
    elif tamper == "candidates":
        plan["cover_photo_source_ids"] = ("another-photo",)
    else:
        identity = plan["cover_photo_source_ids"][0]
        plan["selected"].append(identity)
        plan["destinations"][identity] = "cover_photo"
    with pytest.raises(ValueError, match="No cover photograph choice belongs"):
        build_starting_import(source, inspection, plans.values())


@pytest.mark.parametrize("previous", [None, "equal", "changed_bytes"])
def test_no_photo_cannot_delete_populated_human_or_prior_equal_photo(previous):
    photo = ResolvedBlock("cover_photo", "This month", asset_hashes=("1" * 64 + ".png",))
    current = report(photo)
    prior = report(photo) if previous == "equal" else report(omission("b")) if previous == "changed_bytes" else None
    with pytest.raises(ValueError, match="No cover photograph in this report first"):
        merge_starting_drafts(current, report(omission()), prior)
    assert current.blocks == (photo,)


@pytest.mark.parametrize("cleared", [False, True])
@pytest.mark.parametrize("receipt", [False, True])
def test_no_photo_preserves_already_omitted_retained_assets_or_genuinely_cleared_photo(cleared, receipt):
    photo = ResolvedBlock("cover_photo", "This month" if cleared else "Omit",
                          asset_hashes=() if cleared else ("1" * 64 + ".png",),
                          references=("retained:old-photo",))
    current = report(photo)
    incoming = report(omission())
    merged = merge_starting_drafts(current, incoming, incoming if receipt else None)
    assert merged.blocks == (replace(photo, source="Omit"),)
    assert incoming.blocks == (omission(),), "The source contribution cannot acquire report-only retained assets."


def test_unchanged_consumed_no_photo_receipt_preserves_later_manual_photo():
    photo = ResolvedBlock("cover_photo", "Replace once", asset_hashes=("2" * 64 + ".png",))
    current = report(photo)
    incoming = report(omission())
    assert merge_starting_drafts(current, incoming, incoming) == current
    with pytest.raises(ValueError, match="No cover photograph in this report first"):
        merge_starting_drafts(current, incoming, incoming, verified=False)


def test_same_source_photo_reselection_cannot_resurrect_report_omission():
    retained = ResolvedBlock("cover_photo", "Omit", asset_hashes=("2" * 64 + ".png",))
    current = report(retained)
    selected = report(ResolvedBlock("cover_photo", "Last month", asset_hashes=("3" * 64 + ".png",)))
    with pytest.raises(ValueError, match="differs between this report"):
        merge_starting_drafts(current, selected, report(omission()))
    assert current.blocks == (retained,)


def test_whole_foreign_or_current_design_cover_omission_does_not_create_photo_clear(tmp_path):
    source, _, _ = cover_source(tmp_path)
    inspection, plans = photo_omission_plan(source)
    plans["cover"]["omit"] = True
    mapped, _ = build_starting_import(source, inspection, plans.values())
    assert not any(block.key == "cover_photo" for block in mapped.blocks)
    assert imported_draft(synthetic_profiles()[0], ReportPeriod(2026, 9), "Synthetic Editor", mapped).blocks


def test_gallery_choose_after_no_photo_reverses_only_that_unsaved_intention(monkeypatch, tmp_path):
    source, _, _ = cover_source(tmp_path)
    app = start_app(monkeypatch, tmp_path, source)
    next(button for button in app.button if button.label == "No cover photograph").click().run()
    next(button for button in app.button if button.label == "Choose picture 1").click().run()
    plans = next(value for key, value in app.session_state.filtered_state.items() if key.endswith("_section_plans"))
    assert not plans["cover"]["cover_photo_omitted"]
    mapped, _ = build_starting_import(source, inspect_docx(source), plans.values())
    photo = next(block for block in mapped.blocks if block.key == "cover_photo")
    assert photo.source != "Omit" and len(photo.asset_hashes) == 1
    assert not library.list_profiles(synthetic_profiles()[0].contract)


@pytest.mark.parametrize("stale", [False, True])
def test_retained_new_photo_needs_explicit_use_after_no_photo_without_waiving_price_gate(monkeypatch, tmp_path, stale):
    from app import monthly_report_section_ui as ui
    source, _, _ = cover_source(tmp_path)
    document = Document(source)
    document.sections[0].footer.paragraphs[0].text = "Quoted price $2,500"
    document.save(source)
    app = start_app(monkeypatch, tmp_path, source)
    next(button for button in app.button if button.label == "No cover photograph").click().run()
    next(radio for radio in app.radio if radio.label == "What would you like to do?").set_value("Edit text or change pictures").run()
    incoming = BytesIO()
    Image.new("RGB", (600, 400), "purple").save(incoming, "PNG")
    incoming.name, incoming.size = "new-purple-photo.png", len(incoming.getvalue())
    real = ui.st.file_uploader
    monkeypatch.setattr(ui.st, "file_uploader", lambda label, *a, **kw:
                        incoming if label == "New cover photograph" else real(label, *a, **kw))
    app.run()
    assert not app.exception
    plans = next(value for key, value in app.session_state.filtered_state.items() if key.endswith("_section_plans"))
    assert plans["cover"]["cover_photo_omitted"]
    assert plans["cover"]["new_assets"]
    assert any("pricing" in error.value for error in app.error)
    use = next(button for button in app.button if button.label == "Use this new cover photograph")
    if stale:
        incoming.seek(0)
        incoming.truncate()
        Image.new("RGB", (600, 400), "yellow").save(incoming, "PNG")
        incoming.size = len(incoming.getvalue())
    use.click().run()
    assert not app.exception
    captions = [caption for group in app.image for caption in group.captions]
    assert any(("No cover photograph is selected" if stale else "included in this report") in caption
               for caption in captions), captions
    app.run()
    assert not app.exception
    plans = next(value for key, value in app.session_state.filtered_state.items() if key.endswith("_section_plans"))
    assert plans["cover"]["cover_photo_omitted"] is stale
    assert plans["cover"]["new_assets"]
    assert any("pricing" in error.value for error in app.error)


def test_normal_starting_no_photo_cannot_publish_over_current_human_photo(monkeypatch, tmp_path):
    from test_monthly_report_cover_art_recovery import existing_app
    source, _, _ = cover_source(tmp_path)
    original = source.read_bytes()
    app, working, saved, _, assets = existing_app(monkeypatch, tmp_path, source)
    next(button for button in app.button if button.label == "No cover photograph").click().run()
    from test_monthly_report_starting_review_recovery import continue_button
    continue_button(app).click().run()
    assert not app.exception
    assert any("No cover photograph in this report first" in error.value for error in app.error)
    assert app.session_state["working"] == working
    assert library.load_profile(working.profile.contract, working.profile.key) == saved
    assert library.imported_original(working.profile.contract, working.profile.key) is None
    assert source.read_bytes() == original
    assert all(library.read_asset(working.profile.contract, working.profile.key, reference) == raw
               for reference, raw in assets.items())


def test_whole_cover_omission_with_no_body_contribution_keeps_open_report(monkeypatch, tmp_path):
    from test_monthly_report_cover_art_recovery import existing_app
    from test_monthly_report_starting_review_recovery import drawing_report, continue_button
    source, _ = drawing_report(tmp_path, "cover")
    app, working, _, _, _ = existing_app(monkeypatch, tmp_path, source)
    next(radio for radio in app.radio if radio.label == "What would you like to do?").set_value("Leave this section out").run()
    continue_button(app).click().run()
    assert not app.exception and not app.error
    assert library.load_imported_draft(working.profile.contract, working.profile.key) == working
    assert library.imported_original(working.profile.contract, working.profile.key).read_bytes() == source.read_bytes()
