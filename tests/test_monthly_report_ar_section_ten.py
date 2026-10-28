"""Section 10 Accounts Receivables logical identity and verified native numeral.

Synthetic only. This does not change the financial-output policy, a report's
approved section order, or any original installed Word design.
"""
from dataclasses import replace
from io import BytesIO

from docx import Document
from docx.shared import Inches
import pytest

from app.monthly_report_docx import assemble_docx
from app.monthly_report_import import inspect_docx, _heading
from app.monthly_report_model import (
    BlockSpec, ReportPeriod, ResolvedBlock, included_sections,
    optional_sections, synthetic_draft, synthetic_profiles,
)
from app.monthly_report_divider_numerals import fit_proposals_numeral

from test_monthly_report_divider_numerals import (
    _synthetic_doc, _geometry, _docx_native_parts,
)
from conftest import requires_libreoffice


def _twelve_section_draft():
    """The explicitly selected twelve-section variant, never a default override."""
    base = synthetic_draft(synthetic_profiles()[1], ReportPeriod(2026, 9))
    ar = replace(optional_sections()[0], blocks=(
        BlockSpec("accounts_receivable_notes", "rich_text"),))
    sections = (*base.sections[:9], ar, *base.sections[9:])
    blocks = (*base.blocks, ResolvedBlock(
        "accounts_receivable_notes", "This month",
        text="Synthetic nonfinancial status: invoice review is ongoing."))
    return replace(base, sections=sections, blocks=blocks)


def test_ar_at_section_ten_has_unique_contents_divider_and_body_heading():
    draft = _twelve_section_draft()
    ordered = included_sections(draft.sections)
    assert len(ordered) == 12
    assert [(section.key, section.number) for section in ordered][9:] == [
        ("accounts_receivable", "10"),
        ("proposals", "11"),
        ("training", "12"),
    ]
    assert ordered[0].key == "organization" and ordered[0].number == "1"
    source_draft = draft  # Renderer does not write revised numbering to draft.
    raw = assemble_docx(draft, acknowledged_fingerprint=draft.fingerprint)
    assert draft == source_draft
    doc = Document(BytesIO(raw))
    paragraphs = [paragraph.text for paragraph in doc.paragraphs]
    assert paragraphs.count("10. Accounts Receivable") == 2  # contents and body
    assert paragraphs.count("Section 10") == 1
    assert paragraphs.count("Accounts Receivable") == 1  # divider title only
    assert paragraphs.count("1. Organizational Chart") == 2
    assert paragraphs.count("11. Pending & Declined Proposals") == 2
    assert paragraphs.count("12. Training Summary") == 2
    assert paragraphs.count("Section 1") == 1
    assert paragraphs.index("1. Organizational Chart") < paragraphs.index(
        "10. Accounts Receivable") < paragraphs.index(
        "11. Pending & Declined Proposals")
    assert paragraphs.index("Section 10") < paragraphs.index(
        "10. Accounts Receivable")
    assert paragraphs.count("Synthetic nonfinancial status: invoice review is ongoing.") == 1


def test_optional_section_omission_preserves_dynamic_numbering_not_hardcoded_ten():
    draft = _twelve_section_draft()
    # A chosen design omits Water (previously Section 7). AR is now Section 9,
    # Proposals is Section 10. No universal 'AR is Section 10' mutation.
    sections = tuple(replace(section, included=False)
                     if section.key == "water" else section for section in draft.sections)
    revised = replace(draft, sections=sections)
    ordered = included_sections(revised.sections)
    assert len(ordered) == 11
    assert [(section.key, section.number) for section in ordered][8:] == [
        ("accounts_receivable", "9"),
        ("proposals", "10"),
        ("training", "11"),
    ]
    raw = assemble_docx(revised, acknowledged_fingerprint=revised.fingerprint)
    paragraphs = [p.text for p in Document(BytesIO(raw)).paragraphs]
    assert paragraphs.count("9. Accounts Receivable") == 2
    assert paragraphs.count("Section 9") == 1
    assert paragraphs.count("10. Pending & Declined Proposals") == 2
    assert "10. Accounts Receivable" not in paragraphs
    assert all("Water Treatment Reports" not in text for text in paragraphs)


def test_numbered_word_ar_table_never_routes_into_organizational_chart(tmp_path):
    doc = Document()
    doc.add_heading("Section 1 Organizational Chart", 1)
    org = doc.add_table(rows=2, cols=2)
    org.cell(0, 0).text, org.cell(0, 1).text = "Role", "Name"
    org.cell(1, 0).text, org.cell(1, 1).text = "Director", "Synthetic Person"
    # Even a malformed numeral cannot turn an explicitly labeled AR heading
    # into Organization; original numbering is retained for human design review.
    doc.add_heading("Section 1 Accounts Receivables", 1)
    ar = doc.add_table(rows=2, cols=2)
    ar.cell(0, 0).text, ar.cell(0, 1).text = "Invoice", "Status"
    ar.cell(1, 0).text, ar.cell(1, 1).text = "SYN-INV-1", "Open"
    path = tmp_path / "synthetic-misnumbered-ar.docx"
    doc.save(path)
    before = path.read_bytes()
    inspection = inspect_docx(path)
    tables = [item for item in inspection.items if item.kind == "table"]
    assert len(tables) == 2
    assert tables[0].section == "organization"
    assert tables[0].suggested_slot == "contact_matrix"
    assert tables[1].section == "accounts_receivable"
    assert tables[1].suggested_slot == "accounts_receivable"
    assert tables[1].rows[1] == ("SYN-INV-1", "Open")
    assert _heading("Section 10 Accounts Receivables") == "accounts_receivable"
    assert _heading("Section 1 Accounts Receivables") == "accounts_receivable"
    assert path.read_bytes() == before  # No hidden rewrite of bad source art.


def test_accounts_receivable_section_preview_fits_the_exact_measured_ten_once():
    doc = _synthetic_doc()
    doc.paragraphs[-1].text = "SYNTHETIC ACCOUNTS RECEIVABLES DIVIDER"
    original = _geometry(doc)
    assert original[0] == "2574290" and original[1] == "5596885"
    assert fit_proposals_numeral(doc, section_key="accounts_receivable")
    width, offset, fallback = _geometry(doc)
    assert width == str(2574290 + 20 * 12700)
    assert offset == str(5596885 + 5 * 12700)
    assert "width:222.7pt" in fallback
    assert "margin-top:445.7pt" in fallback
    assert not fit_proposals_numeral(doc, section_key="accounts_receivable")
    assert "ACCOUNTS RECEIVABLES" in doc.paragraphs[-1].text


@pytest.mark.parametrize("section", ("cover", "organization", "activity", "training", "capital"))
def test_ar_digit_measurement_does_not_rewrite_unrelated_section_preview(section):
    doc = _synthetic_doc()
    doc.paragraphs[-1].text = "SYNTHETIC ACCOUNTS RECEIVABLES DIVIDER"
    source = BytesIO(); doc.save(source)
    assert not fit_proposals_numeral(doc, section_key=section)
    current = BytesIO(); doc.save(current)
    assert _docx_native_parts(source.getvalue()) == _docx_native_parts(
        current.getvalue())


def test_unknown_or_ambiguous_ar_native_geometry_is_not_silently_rewritten():
    from copy import deepcopy
    from app.monthly_report_divider_numerals import MC, VML, W
    from lxml import etree
    for damage in ("wrong-size", "missing-choice", "two-frames"):
        doc = _synthetic_doc()
        alt = next(doc.element.body.iter(MC + "AlternateContent"))
        if damage == "wrong-size":
            next(alt.iter(W + "spacing")).set(W + "val", "-200")
        elif damage == "missing-choice":
            alt.remove(alt.find(MC + "Choice"))
        else:
            doc.paragraphs[0]._p[0].append(deepcopy(alt))
        before = etree.tostring(doc.element.body)
        assert not fit_proposals_numeral(doc, section_key="accounts_receivable")
        assert etree.tostring(doc.element.body) == before


@requires_libreoffice
def test_ar_section_ten_real_pdf_preserves_titles_after_conversion(tmp_path, monkeypatch):
    from app import pdf_converter
    import pymupdf
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path / "runtime"))
    monkeypatch.setattr(pdf_converter, "OUTPUT_DIR", tmp_path)
    draft = _twelve_section_draft()
    word = assemble_docx(draft, acknowledged_fingerprint=draft.fingerprint)
    source = tmp_path / "synthetic-twelve.docx"
    source.write_bytes(word)
    target = pdf_converter.convert_to_pdf(source)
    try:
        with pymupdf.open(target) as pdf:
            text_by_page = [page.get_text() for page in pdf]
            toc = next(text for text in text_by_page if "Table of contents" in text)
            assert "10. Accounts Receivable" in toc
            assert "11. Pending & Declined Proposals" in toc
            assert "12. Training Summary" in toc
            ar_dividers = [n for n, value in enumerate(text_by_page)
                           if "Section 10" in value and "Accounts Receivable" in value]
            assert len(ar_dividers) == 1
            index = ar_dividers[0]
            assert index + 1 < len(text_by_page)
            assert "10. Accounts Receivable" in text_by_page[index + 1]
            assert "Synthetic nonfinancial status" in text_by_page[index + 1]
            assert all("Synthetic nonfinancial status" not in value for value in
                       text_by_page[:index + 1])
    finally:
        target.unlink(missing_ok=True)


@requires_libreoffice
def test_ar_measured_native_ten_is_visible_in_pdf_raster(tmp_path, monkeypatch):
    from app import pdf_converter
    import pymupdf
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path / "runtime"))
    monkeypatch.setattr(pdf_converter, "OUTPUT_DIR", tmp_path)
    document = _synthetic_doc()
    document.paragraphs[-1].text = "SYNTHETIC ACCOUNTS RECEIVABLES DIVIDER"
    assert fit_proposals_numeral(document, section_key="accounts_receivable")
    source = tmp_path / "synthetic-ar-native-ten.docx"
    document.save(source)
    output = pdf_converter.convert_to_pdf(source)
    try:
        with pymupdf.open(output) as pdf:
            assert len(pdf) == 1
            page = pdf[0]
            assert "SYNTHETIC ACCOUNTS RECEIVABLES" in page.get_text()
            hits = page.search_for("10")
            assert len(hits) == 1
            area = hits[0]
            # The verified native Word anchor intentionally overhangs the
            # left page edge. Verify the visible intersecting glyph pixels,
            # rather than demanding the authored text frame start at x >= 0.
            assert (area.width > 0 and area.height > 0
                    and area.x1 > page.rect.x0 and area.x0 < page.rect.x1
                    and area.y0 >= page.rect.y0 - 1
                    and area.y1 <= page.rect.y1 + 1)
            # Text extraction is not visual proof. Independently inspect
            # rasterized pixels and require that the measured mint text ink
            # is present in BOTH halves of the bounding rectangle.
            clip = pymupdf.Rect(area.x0 - 2, area.y0 - 2,
                                area.x1 + 2, area.y1 + 2) & page.rect
            pix = page.get_pixmap(clip=clip, dpi=200, alpha=False)
            rgba = memoryview(pix.samples)
            middle = pix.width // 2
            painted = [0, 0]
            for y in range(pix.height):
                for x in range(pix.width):
                    pos = (y * pix.width + x) * pix.n
                    red, green, blue = rgba[pos:pos + 3]
                    if (170 <= red <= 255 and 190 <= green <= 255
                            and 10 <= blue <= 155 and green > blue + 45):
                        painted[int(x >= middle)] += 1
            assert painted[0] > 2 and painted[1] > 2, painted
    finally:
        output.unlink(missing_ok=True)
