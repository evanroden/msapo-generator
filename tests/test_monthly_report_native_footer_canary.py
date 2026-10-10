"""Company-domain footer cannot smuggle another client's source text into a new report."""
from dataclasses import replace
from io import BytesIO
from zipfile import ZipFile

from docx import Document

from app import monthly_report_designs as designs
from app.monthly_report_docx import _build_docx
from app.monthly_report_model import ReportPeriod, synthetic_draft, synthetic_profiles
from test_monthly_report_designs import master_reference

CANARY = "OTHER-CLIENT-FOOTER-CANARY-ZX731"


def _fixture(monkeypatch, tmp_path, *, address=""):
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path / "isolated"))
    master = master_reference(tmp_path)
    doc = Document(master)
    doc.sections[0].header.paragraphs[0].text = (
        "ENFRA | enfrasolutions.com | " + CANARY)
    doc.sections[0].footer.paragraphs[0].text = (
        "ENFRA | enfrasolutions.com | " + CANARY)
    doc.save(master)
    designs.install_master(master, actor="Synthetic Editor", expected_revision=0)
    profile = replace(synthetic_profiles()[0], contract="Synthetic Second Client")
    draft = synthetic_draft(designs.pin(profile), ReportPeriod(2026, 9))
    return replace(draft, address_line=address)


def _header_footer_payload(package):
    with ZipFile(BytesIO(package)) as archive:
        return b"\n".join(archive.read(name) for name in archive.namelist()
            if name.startswith(("word/header", "word/footer")) and name.endswith(".xml"))


def test_blank_current_footer_discards_old_client_mixed_with_company_domain(monkeypatch, tmp_path):
    draft = _fixture(monkeypatch, tmp_path)
    output = _build_docx(draft)
    payload = _header_footer_payload(output)
    assert CANARY.encode() not in payload
    assert b"enfrasolutions.com" in payload
    assert b"ENFRA" in payload


def test_explicit_current_corporate_footer_replaces_mixed_source_without_erasing_address(monkeypatch, tmp_path):
    approved = "ENFRA Solutions | 100 Corporate Test Drive | enfrasolutions.com"
    draft = _fixture(monkeypatch, tmp_path, address=approved)
    payload = _header_footer_payload(_build_docx(draft))
    assert CANARY.encode() not in payload
    assert b"100 Corporate Test Drive" in payload
