"""Independent synthetic controls for exact non-brand Starting replacement."""
from dataclasses import replace
from io import BytesIO
from contextlib import contextmanager
from threading import Event, Thread

from docx import Document
from docx.shared import Inches
from PIL import Image
import pytest

from app import monthly_report_library as library
from app.monthly_report_asset_review import pending_asset_indexes
from app.monthly_report_model import ReportDraft, ReportPeriod, ResolvedBlock, default_sections, synthetic_profiles
from app.monthly_report_setup import merge_drafts
from app.monthly_report_start_conflicts import image_conflicts, apply_replacements, snapshot_guard
from app.monthly_report_start_replay import merge_starting_drafts
from tests.test_monthly_report_start_import import start_app, save_start, PERIOD
from tests.test_monthly_report_starting_replay_edits import repeat_saved_start


def block(draft, key="org_chart"):
    return next(b for b in draft.blocks if b.key == key)


def pair(slot="org_chart"):
    profile = synthetic_profiles()[0]
    old = ResolvedBlock(slot, "This month", asset_hashes=("old.png",), references=("human:separate",))
    new_ref = "docx:" + "a" * 64 + ":item-0004"
    new = ResolvedBlock(slot, "This month", asset_hashes=("new.png",), references=(new_ref,),
                        asset_provenance=(("new.png", (new_ref,)),))
    d = ReportDraft(profile, PERIOD, "Synthetic", default_sections(), (old,))
    return d, replace(d, blocks=(new,))


def context(**changes):
    return dict(verified=True, source_sha256="a" * 64, profile_revision=2,
                snapshot_revision=1, incarnation="one-review", **changes)


@pytest.mark.parametrize("slot", ["org_chart", "contact_matrix", "business_hours_workflow", "after_hours_workflow"])
def test_exact_conflict_replaces_only_detached_target_without_review_transfer(slot):
    old, new = pair(slot)
    current = replace(block(old, slot), client_reviewed_fingerprint=block(old, slot).fingerprint,
                      client_asset_reviews=("old-review",), asset_provenance=(("old.png", ("human:separate",)),))
    sibling = replace(current, key="independent_same_bytes")
    old = replace(old, blocks=(current, sibling))
    conflicts = image_conflicts(old, new, None, **context())
    assert len(conflicts) == 1
    with pytest.raises(ValueError, match="Two different images"):
        merge_drafts(old, new)  # Default guard stays unchanged.
    revised = apply_replacements(old, new, None, {slot: conflicts[0].binding}, **context())
    merged = merge_starting_drafts(revised, new)
    assert block(merged, slot) == block(new, slot)
    assert block(merged, "independent_same_bytes") == sibling
    assert old.blocks == (current, sibling)
    assert not block(merged, slot).client_asset_reviews
    assert pending_asset_indexes(block(merged, slot)) == (0,)


@pytest.mark.parametrize("key,value", [("source_sha256", "b" * 64), ("profile_revision", 3),
    ("snapshot_revision", 2), ("incarnation", "new-review"), ("verified", False)])
def test_stale_binding_never_applies(key, value):
    old, new = pair()
    choice = image_conflicts(old, new, None, **context())[0]
    altered = context(); altered[key] = value
    with pytest.raises(ValueError, match="changed after preview"):
        apply_replacements(old, new, None, {choice.slot: choice.binding}, **altered)


@pytest.mark.parametrize("change", ["old_caption", "new_reference", "same_bytes_other_owner", "other_work", "scope", "month"])
def test_every_relevant_content_and_scope_change_invalidates_intent(change):
    old, new = pair()
    choice = image_conflicts(old, new, None, **context())[0]
    if change == "old_caption": old = replace(old, blocks=(replace(block(old), asset_captions=("Changed",)),))
    if change == "new_reference": new = replace(new, blocks=(replace(block(new), references=("another:occurrence",)),))
    if change == "same_bytes_other_owner": old = replace(old, blocks=(replace(block(old), asset_provenance=(("old.png", ("new:owner",)),)),))
    if change == "other_work": old = replace(old, blocks=(*old.blocks, ResolvedBlock("activity_summary", "This month", text="Unsaved work")))
    if change == "scope": new = replace(new, profile=synthetic_profiles()[1])
    if change == "month": new = replace(new, period=ReportPeriod(2026, 10))
    with pytest.raises(ValueError):
        apply_replacements(old, new, None, {choice.slot: choice.binding}, **context())


@pytest.mark.parametrize("slot", ["brand_logo", "client_logo", "cover_photo", "divider_organization", "improvements"])
def test_finite_nonbrand_scope_never_offers_logo_cover_or_multi_picture_recovery(slot):
    old, new = pair(slot)
    assert image_conflicts(old, new, None, **context()) == ()


@pytest.mark.parametrize("change", ["text", "rows", "extra_picture", "omitted"])
def test_mixed_or_ambiguous_old_content_is_not_silently_replaced(change):
    old, new = pair()
    b = block(old)
    b = replace(b, **{"text": {"text": "Retain my note"}, "rows": {"rows": (("Keep",),)},
                      "extra_picture": {"asset_hashes": ("old.png", "another.png")},
                      "omitted": {"source": "Omit"}}[change])
    assert image_conflicts(replace(old, blocks=(b,)), new, None, **context()) == ()


def test_unchanged_consumed_source_does_not_reprompt_after_report_only_edit():
    old, new = pair()
    assert image_conflicts(old, new, new, **context()) == ()
    assert merge_starting_drafts(old, new, new) == old


def source_file(tmp_path, color):
    doc = Document()
    doc.add_heading(synthetic_profiles()[0].facilities[0].title, 0)
    doc.add_paragraph("September 2026")
    doc.add_heading("1 Organizational Chart", 1)
    image = BytesIO(); Image.new("RGB", (240, 160), color).save(image, "PNG")
    doc.add_picture(image, width=Inches(3))
    path = tmp_path / (color + ".docx"); doc.save(path)
    return path


@pytest.fixture
def collision(monkeypatch, tmp_path):
    from app import monthly_report_section_ui as ui
    first, second = source_file(tmp_path, "navy"), source_file(tmp_path, "red")
    initial = save_start(start_app(monkeypatch, tmp_path, first))
    library.save_snapshot(initial, expected_revision=0, entered_editor="Synthetic")
    state = {"raw": second.read_bytes()}
    real = ui.st.file_uploader
    def upload(label, *args, **kwargs):
        if label != "Starting report": return real(label, *args, **kwargs)
        result = BytesIO(state["raw"]); result.name = "reused-filename.docx"; result.size = len(result.getvalue())
        return result
    monkeypatch.setattr(ui.st, "file_uploader", upload)
    return initial, first, second, state, repeat_saved_start()


def button(app, label):
    return next(b for b in app.button if b.label == label)


def manifest(initial):
    return library._profile_path(initial.profile.contract, initial.profile.key) / "manifest.json"


def test_real_ui_preflight_keep_and_explicit_save_reload_replay(collision):
    initial, first, second, upload, app = collision
    before = manifest(initial).read_bytes()
    assert not app.exception and not app.error
    assert button(app, "Continue to this month’s updates").disabled
    assert len(app.image) == 2  # Actual current/incoming previews, not names alone.
    button(app, "Keep current organization chart and return to review").click().run()
    assert manifest(initial).read_bytes() == before
    assert not library.starting_contribution(initial.profile.contract, initial.profile.key, PERIOD,
        __import__("hashlib").sha256(second.read_bytes()).hexdigest())[1]
    button(app, "Use incoming organization chart").click().run()
    assert manifest(initial).read_bytes() == before
    assert not button(app, "Continue to this month’s updates").disabled
    saved = save_start(app)
    assert block(saved).asset_hashes != block(initial).asset_hashes
    assert pending_asset_indexes(block(saved)) == (0,)
    from app.monthly_report_checks import preflight
    assert any(check.code == "client_pages" and check.block_key == "org_chart" for check in preflight(saved))
    # Existing saved report/history/standing art have not been overwritten.
    assert library.load_snapshot(initial.profile.contract, initial.profile.key, PERIOD).draft == initial
    assert library.load_profile(initial.profile.contract, initial.profile.key).block("org_chart").asset_hashes == block(initial).asset_hashes
    assert first.read_bytes() == library._profile_path(initial.profile.contract, initial.profile.key).joinpath("imports", __import__("hashlib").sha256(first.read_bytes()).hexdigest()+".docx").read_bytes()
    reopened = repeat_saved_start()
    assert not any(b.label == "Use incoming organization chart" for b in reopened.button)
    replayed = save_start(reopened)
    assert block(replayed) == block(saved)


def test_save_failure_retires_intent_and_preserves_original_report(collision, monkeypatch):
    initial, _, _, _, app = collision
    before = manifest(initial).read_bytes()
    button(app, "Use incoming organization chart").click().run()
    old_key = button(app, "Use incoming organization chart").key
    atomic = library._atomic_write
    def fail(path, raw):
        if path == manifest(initial): raise OSError("Synthetic final publication failure")
        return atomic(path, raw)
    monkeypatch.setattr(library, "_atomic_write", fail)
    button(app, "Continue to this month’s updates").click().run()
    assert not app.exception and any("publication failure" in e.value for e in app.error)
    assert manifest(initial).read_bytes() == before
    monkeypatch.setattr(library, "_atomic_write", atomic)
    app.run()
    assert button(app, "Continue to this month’s updates").disabled
    assert button(app, "Use incoming organization chart").key != old_key
    button(app, "Use incoming organization chart").click().run()
    assert block(save_start(app)).asset_hashes != block(initial).asset_hashes


def test_saved_revision_change_invalidates_staged_replacement_even_when_bytes_return(collision):
    initial, _, _, _, app = collision
    button(app, "Use incoming organization chart").click().run()
    old_key = button(app, "Use incoming organization chart").key
    library.save_snapshot(replace(initial, prepared_by="Temporary editor"), expected_revision=1, entered_editor="Synthetic")
    library.save_snapshot(initial, expected_revision=2, entered_editor="Synthetic")
    app.run()
    assert button(app, "Continue to this month’s updates").disabled
    assert button(app, "Use incoming organization chart").key != old_key


def test_source_aba_requires_fresh_review(collision):
    _, first, second, upload, app = collision
    button(app, "Use incoming organization chart").click().run()
    old_key = button(app, "Use incoming organization chart").key
    upload["raw"] = first.read_bytes(); app.run()
    upload["raw"] = second.read_bytes(); app.run()
    assert not app.exception
    assert button(app, "Continue to this month’s updates").disabled
    assert button(app, "Use incoming organization chart").key != old_key


def test_snapshot_guard_holds_writer_until_publication_and_releases_after_exception(collision, monkeypatch):
    initial, _, _, _, _ = collision
    entering, finished = Event(), Event()
    failures = []
    original_lock = library._locked
    snapshot_path = library._profile_path(initial.profile.contract, initial.profile.key, snapshots=True)
    @contextmanager
    def observed(path):
        if path == snapshot_path: entering.set()
        with original_lock(path): yield
    def competing_save():
        try: library.save_snapshot(initial, expected_revision=1, entered_editor="Other editor")
        except Exception as exc: failures.append(exc)
        finally: finished.set()
    with pytest.raises(RuntimeError, match="Synthetic save failure"):
        with snapshot_guard(initial.profile, PERIOD, 1, initial.fingerprint):
            monkeypatch.setattr(library, "_locked", observed)
            thread = Thread(target=competing_save); thread.start()
            assert entering.wait(5)
            assert not finished.is_set()
            assert library.load_snapshot(initial.profile.contract, initial.profile.key, PERIOD).revision == 1
            raise RuntimeError("Synthetic save failure")
    assert finished.wait(5); thread.join(5)
    assert not failures
    assert library.load_snapshot(initial.profile.contract, initial.profile.key, PERIOD).revision == 2
    with pytest.raises(ValueError, match="Another report version"):
        with snapshot_guard(initial.profile, PERIOD, 1, initial.fingerprint): pass


def test_final_profile_cas_failure_retires_intent(collision, monkeypatch):
    initial, _, _, _, app = collision
    button(app, "Use incoming organization chart").click().run()
    real_save = library.save_report_setup
    def competing_profile(*args, **kwargs):
        current = library.load_profile(initial.profile.contract, initial.profile.key)
        library.save_profile(current.profile, expected_revision=current.revision,
                             actor="Other editor", confirmed=True)
        return real_save(*args, **kwargs)
    monkeypatch.setattr(library, "save_report_setup", competing_profile)
    button(app, "Continue to this month’s updates").click().run()
    assert not app.exception and app.error
    assert library.load_imported_draft(initial.profile.contract, initial.profile.key).blocks == initial.blocks
    monkeypatch.setattr(library, "save_report_setup", real_save)
    app.run()
    assert button(app, "Continue to this month’s updates").disabled


def test_changed_source_at_final_save_is_not_published(collision, monkeypatch):
    initial, _, _, _, app = collision
    before = manifest(initial).read_bytes()
    button(app, "Use incoming organization chart").click().run()
    real_save = library.save_report_setup
    def changed_source(profile, draft, source, *args, **kwargs):
        # Mutate only the staged synthetic copy, never the immutable old source.
        source.write_bytes(source.read_bytes() + b"Synthetic changed source")
        return real_save(profile, draft, source, *args, **kwargs)
    monkeypatch.setattr(library, "save_report_setup", changed_source)
    button(app, "Continue to this month’s updates").click().run()
    assert not app.exception and any("changed during review" in e.value for e in app.error)
    assert manifest(initial).read_bytes() == before
    assert all(not value for key, value in app.session_state.filtered_state.items()
               if key.endswith("_image_replacements"))


def test_final_snapshot_race_after_preview_fails_without_publication(collision, monkeypatch):
    from app import monthly_report_start_conflicts as recovery
    initial, _, _, _, app = collision
    before = manifest(initial).read_bytes()
    button(app, "Use incoming organization chart").click().run()
    real_guard = recovery.snapshot_guard
    @contextmanager
    def late_save(profile, period, revision, fingerprint):
        library.save_snapshot(initial, expected_revision=1, entered_editor="Other editor")
        with real_guard(profile, period, revision, fingerprint): yield
    monkeypatch.setattr(recovery, "snapshot_guard", late_save)
    button(app, "Continue to this month’s updates").click().run()
    assert not app.exception and any("Another report version" in e.value for e in app.error)
    assert manifest(initial).read_bytes() == before
    assert library.load_snapshot(initial.profile.contract, initial.profile.key, PERIOD).revision == 2


def test_snapshot_same_revision_different_content_is_not_accepted(collision):
    initial, _, _, _, _ = collision
    with pytest.raises(ValueError, match="Another report version"):
        with snapshot_guard(initial.profile, PERIOD, 1, replace(initial, prepared_by="Changed").fingerprint): pass


def test_duplicate_destination_or_unshown_extra_caption_cannot_gain_replacement():
    old, new = pair()
    ambiguous = replace(old, sections=(*old.sections, old.sections[0]))
    with pytest.raises(ValueError, match="ambiguous section"):
        image_conflicts(ambiguous, new, None, **context())
    extra_caption = replace(old, blocks=(replace(block(old), asset_captions=("Visible", "Unshown")),))
    assert image_conflicts(extra_caption, new, None, **context()) == ()


def test_continue_reuses_one_admitted_import_and_never_caches_assets_across_reruns(collision, monkeypatch):
    from app import monthly_report_start_import as importer
    from app import monthly_report_render_jobs as jobs
    _, _, _, _, app = collision
    button(app, "Use incoming organization chart").click().run()
    calls = []
    real = importer.build_starting_import
    def counted(*args, **kwargs):
        assert getattr(jobs._WORK_LOCAL, "active", False)
        calls.append(1)
        return real(*args, **kwargs)
    monkeypatch.setattr(importer, "build_starting_import", counted)
    save_start(app)
    assert calls == [1]
    from app.monthly_report_import import MappedImport
    assert not any(isinstance(value, MappedImport) for value in app.session_state.filtered_state.values())


def test_busy_admission_never_decodes_or_retains_old_action(collision, monkeypatch):
    from app import monthly_report_start_import as importer
    from app import monthly_report_render_jobs as jobs
    initial, _, _, _, app = collision
    before = manifest(initial).read_bytes()
    button(app, "Use incoming organization chart").click().run()
    old_key = button(app, "Use incoming organization chart").key
    real = importer.build_starting_import
    def forbidden(*args, **kwargs): raise AssertionError("Busy preflight must not decode")
    monkeypatch.setattr(importer, "build_starting_import", forbidden)
    monkeypatch.setattr(jobs, "memory_headroom", lambda: 0)
    app.run()
    assert not app.exception and app.error
    assert button(app, "Continue to this month’s updates").disabled
    assert manifest(initial).read_bytes() == before
    monkeypatch.setattr(jobs, "memory_headroom", lambda: None)
    monkeypatch.setattr(importer, "build_starting_import", real)
    app.run()
    assert button(app, "Continue to this month’s updates").disabled
    assert button(app, "Use incoming organization chart").key != old_key


def test_explicit_page_replacement_is_never_a_pricing_waiver():
    from app.monthly_report_checks import preflight
    old, new = pair()
    priced = replace(block(new), asset_captions=("Service charge $500",))
    new = replace(new, blocks=(priced,))
    choice = image_conflicts(old, new, None, **context())[0]
    revised = apply_replacements(old, new, None, {choice.slot: choice.binding}, **context())
    delivered = merge_starting_drafts(revised, new)
    checks = preflight(delivered)
    assert any(check.code == "pricing" and check.block_key == "org_chart" and check.blocking for check in checks)
    assert any(check.code == "client_pages" and check.block_key == "org_chart" and check.blocking for check in checks)


def test_native_preserved_single_page_gets_same_explicit_recovery(collision, monkeypatch, tmp_path):
    from app import monthly_report_visual_preservation as preservation
    from app.monthly_report_import import inspect_docx
    from app.monthly_report_word_pages import WordPage
    from tests.test_monthly_report_word_pages import drawing_report
    initial, _, _, upload, app = collision
    source = drawing_report(tmp_path)
    document = Document(source)
    document.paragraphs[0].text = initial.profile.facilities[0].title + " — September 2026"
    from docx.oxml.ns import qn
    for node in document.element.body.iter(qn("w:t")):
        if node.text in ("Synthetic Lead", "Synthetic Technician"):
            node.text = ""  # Image-only native drawing, with no separately mapped text.
    document.save(source)
    data = BytesIO(); Image.new("RGB", (640, 480), "green").save(data, "PNG")
    # Stub only native conversion. Actual source inspection, section proof,
    # preserved-page records, mapping, UI, commit and replay stay real.
    monkeypatch.setattr(preservation, "preserve_word_pages", lambda *_: (
        WordPage(1, data.getvalue(), 640, 480, "", False),))
    upload["raw"] = source.read_bytes(); app.run()
    assert not app.exception and not app.error
    assert any(b.label == "Use incoming organization chart" for b in app.button)
    assert button(app, "Continue to this month’s updates").disabled
    plan = next(value["organization"] for key, value in app.session_state.filtered_state.items()
                if key.endswith("_section_plans") and value.get("organization", {}).get("auto_native_diagrams"))
    assert plan["preserved_assets"] and not plan.get("new_assets") and not plan.get("destinations")
    assert len(app.image) == 2
    before = manifest(initial).read_bytes()
    button(app, "Use incoming organization chart").click().run()
    assert manifest(initial).read_bytes() == before
    from app import monthly_report_start_import as importer
    from app import monthly_report_render_jobs as jobs
    mapped = []
    real_import = importer.build_starting_import
    def counted(*args, **kwargs):
        assert getattr(jobs._WORK_LOCAL, "active", False)
        mapped.append(1)
        return real_import(*args, **kwargs)
    monkeypatch.setattr(importer, "build_starting_import", counted)
    saved = save_start(app)
    assert mapped == [1]
    monkeypatch.setattr(importer, "build_starting_import", real_import)
    assert block(saved).asset_hashes != block(initial).asset_hashes
    assert block(saved).references == (f"docx:{inspect_docx(source).sha256}:word-sections:2:page:1",)
    assert pending_asset_indexes(block(saved)) == (0,)
    assert library.load_snapshot(initial.profile.contract, initial.profile.key, PERIOD).draft == initial

    replayed = save_start(repeat_saved_start())
    assert block(replayed) == block(saved)
