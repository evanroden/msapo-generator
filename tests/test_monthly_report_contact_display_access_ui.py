"""Entered-name display gates also cover saved report contact content."""
from streamlit.testing.v1 import AppTest

BASE = '''
from dataclasses import replace
import streamlit as st
from app.monthly_report_model import synthetic_draft, synthetic_profiles, ReportPeriod, ResolvedBlock
from app.monthly_report_directory import CONTACT_SPEC
from app.monthly_report_ui import _field
from app.monthly_report_organization_ui import render_organization
from app.monthly_report_visual_ui import edit_contacts
block = ResolvedBlock("contact_matrix", "Library", rows=(("Synthetic site", "Manager", "Synthetic Private Person", "202-555-0190", "private@example.invalid"),))
draft = replace(synthetic_draft(synthetic_profiles()[0], ReportPeriod(2026, 9)), prepared_by=st.session_state.get("actor", ""), blocks=(block,))
'''


def test_organization_hides_saved_contacts_and_change_controls_until_name_entered():
    script = BASE + '''
blocks = {"contact_matrix": block}
specs = {"contact_matrix": CONTACT_SPEC}
st.session_state["result"] = render_organization(draft, blocks, specs, "test", {}, _field, lambda spec, value: value)
'''
    app = AppTest.from_string(script, default_timeout=20).run()
    assert not app.exception and not app.dataframe and not app.button and not app.get("image")
    assert any("Enter your name" in item.value for item in app.info)
    assert app.session_state["result"]["contact_matrix"].rows[0][2] == "Synthetic Private Person"
    app.session_state["actor"] = "Synthetic Editor"
    app.run()
    assert not app.exception
    assert any("private@example.invalid" in str(item.value) for item in app.dataframe)
    app.session_state["actor"] = " "
    app.run()
    assert not app.dataframe and not app.button


def test_contact_editor_does_not_render_saved_rows_without_entered_name():
    app = AppTest.from_string(BASE + '''
st.session_state["result"] = edit_contacts(draft, CONTACT_SPEC, block, "test", {}, _field)
''', default_timeout=20).run()
    assert not app.exception and not app.dataframe and not app.text_input and not app.button
    assert any("Enter your name" in item.value for item in app.info)
    assert app.session_state["result"].rows[0][4] == "private@example.invalid"


def test_saved_original_uses_current_actor_and_hides_staged_content_when_name_cleared(monkeypatch, tmp_path):
    from docx import Document
    from app import monthly_report_library as library
    from app.monthly_report_model import synthetic_profiles
    from test_monthly_report_ui import monthly
    path = tmp_path / 'original.docx'
    document = Document()
    document.add_paragraph(synthetic_profiles()[0].facilities[0].title + ' September 2026')
    document.add_heading('Organizational Chart', 1)
    document.add_paragraph('Synthetic saved private contact')
    document.save(path)
    monkeypatch.setattr(library, 'imported_original', lambda *args: path)
    app = monthly(monkeypatch, tmp_path / 'data')
    original = next(w for w in app.button if w.label == 'Open the saved original for section review')
    assert original.disabled
    next(w for w in app.text_input if w.label == 'Prepared by').set_value('Current Editor').run()
    next(w for w in app.button if w.label == 'Open the saved original for section review').click().run()
    assert not app.exception
    assert not any(w.label == 'Your name' for w in app.text_input)
    assert any(key.endswith('_stage') for key in app.session_state.filtered_state)
    next(w for w in app.text_input if w.label == 'Prepared by').set_value('').run()
    assert not app.exception
    assert next(w for w in app.button if w.label == 'Open the saved original for section review').disabled
    assert any('Enter Prepared by above to review the saved original' in w.value for w in app.info)
    assert not any('Synthetic saved private contact' in w.value for w in app.text)
