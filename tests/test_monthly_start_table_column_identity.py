"""Starting Report grids retain physical columns without changing source schemas."""
from dataclasses import replace
from io import BytesIO
import json

from docx import Document
import pytest
from streamlit.testing.v1 import AppTest

from app import monthly_report_library as library
from app.monthly_report_docx import assemble_docx
from app.monthly_report_import import ImportItem, inspect_docx, imported_draft
from app.monthly_report_model import ReportPeriod, ReportTable, synthetic_profiles
from app.monthly_report_sections import ImportSection, build_section_import, section_reviews, table_without_prices


APP = '''
import streamlit as st
from pathlib import Path
from app.monthly_report_section_ui import _section_card
from app.monthly_report_ui import _field
from app.monthly_report_model import ReportPeriod
plans = st.session_state.setdefault('plans', {})
_section_card(st.session_state.section, st.session_state.path,
              st.session_state.inspection, ReportPeriod(2026, 9),
              'column_identity', _field, plans, True)
'''


def app_for(table, *, seed_review=True):
    item = ImportItem('synthetic-table', 'table', 'word/document.xml', 1, 1,
                      'capital', rows=(table.columns, *table.rows))
    app = AppTest.from_string(APP)
    app.session_state.section = ImportSection('capital', 'Equipment plans', (item,))
    app.session_state.path = 'synthetic.docx'
    app.session_state.inspection = None
    if seed_review:
        app.session_state.plans = {'capital': {'tables': {item.id: table}}}
    app.run()
    assert not app.exception
    return app


def current(app):
    return app.session_state.plans['capital']['tables']['synthetic-table']


def edit(app):
    app.radio[0].set_value('Edit text or change pictures').run()
    assert not app.exception
    return app


def wire(app, delta):
    # AppTest's neutral run omits data_editor state; deliver its actual wire
    # value alongside the radio/other native controls, not a mocked grid.
    frame = app.dataframe[-1]
    message = app._tree.get_widget_states()
    for old in list(message.widgets):
        if old.id == frame.proto.id:
            message.widgets.remove(old)
    message.widgets.add(id=frame.proto.id, string_value=json.dumps(delta))
    app._run(message)
    assert not app.exception


@pytest.mark.parametrize('columns', [
    ('Reading', 'Reading', '1. Reading'),
    ('', '', '1. '),
    ('Reading', 'Reading', 'Reading'),
    ('A', '2. A', 'A'),
])
@pytest.mark.parametrize('mode', ['keep', 'edit'])
def test_every_column_is_visible_without_passive_data_changes(columns, mode):
    table = ReportTable(columns, (('pump P-001', '0', ''), ('pump P-002', '17', 'pending')),
                        'synthetic:source-table')
    app = app_for(table)
    if mode == 'edit':
        edit(app)
    frame = app.dataframe[-1].value
    assert frame.shape == (2, 3)
    assert tuple(map(tuple, frame.values)) == table.rows
    assert len(set(frame.columns)) == 3
    assert current(app) == table
    if mode == 'edit':
        identity = app.dataframe[-1].proto.id
        wire(app, {'edited_rows': {}, 'added_rows': [], 'deleted_rows': []})
        assert app.dataframe[-1].proto.id == identity
        assert current(app) == table


def test_duplicate_cells_edit_independently_and_survive_keep_and_reopen():
    table = ReportTable(('Reading', 'Reading', '1. Reading'),
                        (('original-A', 'original-B', 'original-C'),), 'synthetic:retained')
    app = edit(app_for(table))
    labels = tuple(app.dataframe[-1].value.columns)
    wire(app, {'edited_rows': {'0': {labels[0]: 'corrected-A', labels[1]: ''}},
               'added_rows': [], 'deleted_rows': []})
    expected = replace(table, rows=(('corrected-A', '', 'original-C'),))
    assert current(app) == expected
    # Keep the delivered delta with a sibling radio change.
    app.radio[0].set_value('Keep and review')
    wire(app, {'edited_rows': {'0': {labels[0]: 'corrected-A', labels[1]: ''}},
               'added_rows': [], 'deleted_rows': []})
    assert current(app) == expected
    edit(app)
    assert tuple(map(tuple, app.dataframe[-1].value.values)) == expected.rows
    assert current(app) == expected


def test_added_removed_and_blank_rows_keep_physical_column_order():
    table = ReportTable(('Value', 'Value', 'Unit'), (('A', 'B', 'C'), ('D', 'E', 'F')))
    app = edit(app_for(table))
    labels = tuple(app.dataframe[-1].value.columns)
    wire(app, {'edited_rows': {}, 'added_rows': [dict(zip(labels, ('0', '', 'kW')))],
               'deleted_rows': [0]})
    assert current(app) == replace(table, rows=(('D', 'E', 'F'), ('0', '', 'kW')))


def test_changed_display_keys_do_not_reuse_a_legacy_collapsed_grid_seed():
    from app.monthly_report_editor import _signature
    table = ReportTable(('Reading', 'Reading'), (('first', 'second'),))
    app = app_for(table)
    old_key = 'column_identity_section_capital_table_synthetic-table_' + _signature(table.columns)
    app.session_state[old_key + '_data'] = [{'Reading': 'stale-collapsed'}]
    app.session_state[old_key + '_data_seed'] = [{'Reading': 'stale-collapsed'}]
    edit(app)
    assert current(app) == table
    assert tuple(map(tuple, app.dataframe[-1].value.values)) == table.rows


def test_ordinary_unique_columns_retain_native_key_and_values():
    from app.monthly_report_editor import _signature
    table = ReportTable(('Equipment', 'Condition'), (('pump', 'pending'),))
    app = edit(app_for(table))
    expected = 'column_identity_section_capital_table_synthetic-table_' + _signature(table.columns)
    assert app.dataframe[-1].proto.id.endswith('-' + expected)
    assert tuple(app.dataframe[-1].value.columns) == table.columns
    assert current(app).rows == table.rows


def test_price_in_any_duplicate_column_still_blocks_approval():
    table = ReportTable(('Reading', 'Reading'), (('pump P-001', '$500'),))
    app = edit(app_for(table))
    assert current(app) == table
    assert not app.session_state.plans['capital']['approved']
    assert any('still contains pricing' in node.value for node in app.error)
    labels = tuple(app.dataframe[-1].value.columns)
    wire(app, {'edited_rows': {'0': {labels[1]: 'pending'}}, 'added_rows': [], 'deleted_rows': []})
    assert current(app).rows == (('pump P-001', 'pending'),)
    assert app.session_state.plans['capital']['approved']


def test_keep_preview_remains_bounded_and_edit_retains_all_rows():
    rows = tuple((f'pump-{n}', str(n)) for n in range(301))
    table = ReportTable(('Reading', 'Reading'), rows)
    app = app_for(table)
    assert app.dataframe[-1].value.shape == (300, 2)
    assert current(app) == table
    edit(app)
    assert app.dataframe[-1].value.shape == (301, 2)
    assert current(app) == table


def test_real_docx_collision_correction_save_reopen_and_output(monkeypatch, tmp_path):
    monkeypatch.setenv('EPC_DATA_DIR', str(tmp_path / 'store'))
    doc = Document()
    doc.add_heading('9 Priority Capital Renewal List', 1)
    source = doc.add_table(rows=3, cols=3)
    values = (('Reading', 'Reading', 'Reading (2)'),
              ('pump P-001', '0', 'pending'), ('pump P-002', '17', 'review'))
    for row, cells in zip(source.rows, values):
        for cell, value in zip(row.cells, cells):
            cell.text = value
    path = tmp_path / 'synthetic-column-collision.docx'
    doc.save(path)
    original = path.read_bytes()
    inspection = inspect_docx(path)
    section = next(section for section in section_reviews(inspection) if section.key == 'capital')
    item = next(item for item in section.items if item.kind == 'table')
    cleaned, removed = table_without_prices(item)
    assert not removed and len(set(cleaned.columns)) < len(cleaned.columns)
    app = AppTest.from_string(APP)
    app.session_state.section = section
    app.session_state.path = path
    app.session_state.inspection = inspection
    app.run()
    assert not app.exception
    edit(app)
    assert tuple(map(tuple, app.dataframe[-1].value.values)) == values[1:]
    labels = tuple(app.dataframe[-1].value.columns)
    wire(app, {'edited_rows': {'0': {labels[0]: 'pump P-001 corrected'}},
               'added_rows': [], 'deleted_rows': []})
    plan = app.session_state.plans['capital']
    mapped, _ = build_section_import(path, inspection, (plan,))
    profile = synthetic_profiles()[0]
    accepted = imported_draft(profile, ReportPeriod(2026, 9), 'Synthetic Editor', mapped)
    original_plan = dict(plan, tables={item.id: cleaned})
    initial_mapped, _ = build_section_import(path, inspection, (original_plan,))
    initial = imported_draft(profile, ReportPeriod(2026, 9), 'Synthetic Editor', initial_mapped)
    library.save_snapshot(initial, expected_revision=0, entered_editor='Synthetic Editor')
    saved = library.save_snapshot(accepted, expected_revision=1, entered_editor='Synthetic Editor')
    reopened = library.load_snapshot(profile.contract, profile.key, accepted.period)
    assert reopened.draft == saved.draft == accepted
    assert library.load_snapshot_version(profile.contract, profile.key, accepted.period, 1).draft == initial
    actual = next(b for b in accepted.blocks if b.key == 'capital_renewal').extra_tables[0]
    assert actual.columns == cleaned.columns
    assert actual.rows == (('pump P-001 corrected', '0', 'pending'), values[2])
    assert actual.reference == f'docx:{inspection.sha256}:{item.id}'
    output = Document(BytesIO(assemble_docx(replace(accepted, synthetic=True))))
    exported = [[tuple(cell.text for cell in row.cells) for row in t.rows] for t in output.tables]
    assert [actual.columns, *actual.rows] in exported
    assert path.read_bytes() == original
