"""Synthetic native-note ownership: author prose must never be substring-deleted."""
from dataclasses import replace
from io import BytesIO
from zipfile import ZipFile

from docx import Document
from docx.oxml.ns import qn
from lxml import etree
import pytest

from app.monthly_report_import import ImportMapping, inspect_docx, map_items
from app.monthly_report_model import ResolvedBlock
from app.monthly_report_native_layout import build_native_docx
from test_monthly_report_native_layout import draft

NOTE = 'Synthetic scope remains limited.'
QUALIFIER = 'Current qualifier quotes Synthetic scope remains limited. without adopting it.'


def scene(tmp_path, kind='footnote', mutation=None):
    path = tmp_path / 'synthetic-note.docx'
    doc = Document()
    doc.add_paragraph('Previous Synthetic Site')
    doc.add_paragraph('August 2026')
    doc.add_heading('MONTHLY ACTIVITY SUMMARY', 1)
    p = doc.add_paragraph('Previous substantive activity narrative for replacement.')
    etree.SubElement(etree.SubElement(p._p, qn('w:r')), qn('w:' + kind + 'Reference'), {qn('w:id'): '1'})
    doc.add_heading('TRAINING SUMMARY', 1)
    doc.add_paragraph('Previous training narrative.')
    doc.save(path)
    with ZipFile(path) as z:
        parts = {name: z.read(name) for name in z.namelist()}
    notes = etree.Element(qn('w:' + kind + 's'), nsmap={'w': qn('w:p')[1:].split('}')[0]})
    note = etree.SubElement(notes, qn('w:' + kind), {qn('w:id'): '1'})
    etree.SubElement(etree.SubElement(etree.SubElement(note, qn('w:p')), qn('w:r')), qn('w:t')).text = NOTE
    parts['word/' + kind + 's.xml'] = etree.tostring(notes)
    rels = etree.fromstring(parts['word/_rels/document.xml.rels'])
    etree.SubElement(rels, '{http://schemas.openxmlformats.org/package/2006/relationships}Relationship',
        Id='rSyntheticNotes', Type='http://schemas.openxmlformats.org/officeDocument/2006/relationships/' + kind + 's', Target=kind + 's.xml')
    parts['word/_rels/document.xml.rels'] = etree.tostring(rels)
    types = etree.fromstring(parts['[Content_Types].xml'])
    etree.SubElement(types, '{http://schemas.openxmlformats.org/package/2006/content-types}Override',
        PartName='/word/' + kind + 's.xml', ContentType='application/vnd.openxmlformats-officedocument.wordprocessingml.' + kind + 's+xml')
    parts['[Content_Types].xml'] = etree.tostring(types)
    if mutation:
        mutation(parts, kind)
    with ZipFile(path, 'w') as z:
        for name, content in parts.items():
            z.writestr(name, content)
    inspection = inspect_docx(path)
    item = next(i for i in inspection.items if i.note == kind + ':1')
    mapped = map_items(path, inspection, (ImportMapping(item.id, 'activity_summary'),)).blocks[0]
    return path, mapped


def render(path, block):
    raw = build_native_docx(draft(block), path)
    return raw, [p.text for p in Document(BytesIO(raw)).paragraphs]


@pytest.mark.parametrize('kind', ['footnote', 'endnote'])
@pytest.mark.parametrize('text', [QUALIFIER + '\n' + NOTE, QUALIFIER, NOTE + '\nCurrent afterword.', NOTE + '\n' + NOTE])
def test_only_a_unique_whole_trailing_note_line_may_be_removed(tmp_path, kind, text):
    path, block = scene(tmp_path, kind)
    raw, paragraphs = render(path, replace(block, text=text))
    # Ambiguous repeated words are retained; no author's qualifier is shortened.
    for line in set(text.splitlines()):
        assert paragraphs.count(line) == text.splitlines().count(line)


@pytest.mark.parametrize('kind', ['footnote', 'endnote'])
def test_unique_trailing_note_is_rendered_once_in_native_note(tmp_path, kind):
    path, block = scene(tmp_path, kind)
    raw, paragraphs = render(path, replace(block, text='Current authored narrative.\n' + NOTE))
    assert 'Current authored narrative.' in paragraphs
    assert NOTE not in paragraphs
    with ZipFile(BytesIO(raw)) as z:
        assert NOTE.encode() in z.read('word/' + kind + 's.xml')
        root = etree.fromstring(z.read('word/document.xml'))
        assert len(list(root.iter(qn('w:' + kind + 'Reference')))) == 1


def mutate_note_style(parts, kind, variant):
    root = etree.fromstring(parts['word/' + kind + 's.xml'])
    p = root.find('.//' + qn('w:p'))
    r = p.find(qn('w:r'))
    styles = etree.fromstring(parts['word/styles.xml'])
    if variant == 'direct-hidden':
        props = etree.Element(qn('w:rPr')); r.insert(0, props)
        etree.SubElement(props, qn('w:vanish'))
    elif variant == 'paragraph-mark-hidden':
        props = etree.Element(qn('w:pPr')); p.insert(0, props)
        etree.SubElement(etree.SubElement(props, qn('w:rPr')), qn('w:vanish'))
    else:
        hidden = etree.SubElement(styles, qn('w:style'), {qn('w:type'): 'paragraph', qn('w:styleId'): 'SyntheticHidden'})
        etree.SubElement(hidden, qn('w:name'), {qn('w:val'): 'Synthetic Hidden'})
        etree.SubElement(etree.SubElement(hidden, qn('w:rPr')), qn('w:vanish'))
        props = etree.Element(qn('w:pPr')); p.insert(0, props)
        if variant == 'duplicate-ppr':
            etree.SubElement(props, qn('w:pStyle'), {qn('w:val'): 'Normal'})
            props = etree.Element(qn('w:pPr')); p.insert(1, props)
        if variant == 'duplicate-pstyle':
            etree.SubElement(props, qn('w:pStyle'), {qn('w:val'): 'Normal'})
        identity = 'SyntheticHidden'
        if variant == 'inherited-hidden':
            derived = etree.SubElement(styles, qn('w:style'), {qn('w:type'): 'paragraph', qn('w:styleId'): 'SyntheticDerived'})
            etree.SubElement(derived, qn('w:name'), {qn('w:val'): 'Synthetic Derived'})
            etree.SubElement(derived, qn('w:basedOn'), {qn('w:val'): 'SyntheticHidden'})
            identity = 'SyntheticDerived'
        elif variant == 'unknown-style':
            identity = 'NoSuchSyntheticStyle'
        elif variant == 'wrong-type-style':
            hidden.set(qn('w:type'), 'character')
        elif variant == 'late-pstyle':
            etree.SubElement(props, qn('w:spacing'), {qn('w:after'): '0'})
        elif variant == 'cyclic-style':
            hidden.remove(hidden.find(qn('w:rPr')))
            etree.SubElement(hidden, qn('w:basedOn'), {qn('w:val'): 'SyntheticHidden'})
        etree.SubElement(props, qn('w:pStyle'), {qn('w:val'): identity})
    parts['word/' + kind + 's.xml'] = etree.tostring(root)
    parts['word/styles.xml'] = etree.tostring(styles)


@pytest.mark.parametrize('kind', ['footnote', 'endnote'])
@pytest.mark.parametrize('variant', ['direct-hidden', 'paragraph-mark-hidden', 'inherited-hidden',
    'unknown-style', 'wrong-type-style', 'late-pstyle', 'cyclic-style', 'duplicate-ppr', 'duplicate-pstyle'])
def test_hidden_or_unproved_native_note_cannot_remove_current_line(tmp_path, kind, variant):
    path, block = scene(tmp_path, kind, lambda parts, kind: mutate_note_style(parts, kind, variant))
    raw, paragraphs = render(path, replace(block, text='Current authored narrative.\n' + NOTE))
    assert NOTE in paragraphs
    assert 'Current authored narrative.' in paragraphs


@pytest.mark.parametrize('kind', ['footnote', 'endnote'])
def test_a_note_from_another_block_cannot_remove_current_line(tmp_path, kind):
    path, block = scene(tmp_path, kind)
    raw, paragraphs = render(path, replace(block, key='training_summary', text='Current training narrative.\n' + NOTE))
    assert NOTE in paragraphs
    assert 'Current training narrative.' in paragraphs


@pytest.mark.parametrize('kind', ['footnote', 'endnote'])
def test_duplicate_native_references_cannot_authorize_deletion(tmp_path, kind):
    def mutate(parts, kind):
        root = etree.fromstring(parts['word/document.xml'])
        reference = next(root.iter(qn('w:' + kind + 'Reference')))
        etree.SubElement(reference.getparent(), reference.tag, dict(reference.attrib))
        parts['word/document.xml'] = etree.tostring(root)
    path, block = scene(tmp_path, kind, mutate)
    raw, paragraphs = render(path, replace(block, text='Current authored narrative.\n' + NOTE))
    assert NOTE in paragraphs


@pytest.mark.parametrize('kind', ['footnote', 'endnote'])
def test_missing_or_edited_note_reference_retains_text(tmp_path, kind):
    path, block = scene(tmp_path, kind)
    raw, paragraphs = render(path, replace(block, references=(), text='Current authored narrative.\n' + NOTE))
    assert NOTE in paragraphs
    raw, paragraphs = render(path, replace(block, text='Current authored narrative.\n' + NOTE + ' Corrected.'))
    assert NOTE + ' Corrected.' in paragraphs


def test_line_removal_preserves_all_other_bytes_and_rejects_nonexact_lines():
    from app.monthly_report_native_notes import without_trailing_note
    assert without_trailing_note('  Current text.  \r\n\r\n' + NOTE + '\r\n\r\n', NOTE) == '  Current text.  \r\n\r\n\r\n'
    for text in (' ' + NOTE, NOTE + ' ', NOTE + '\nNot trailing.', NOTE + '\n' + NOTE, 'Quoted ' + NOTE + '\n' + NOTE):
        assert without_trailing_note(text, NOTE) is None


@pytest.mark.parametrize('kind', ['footnote', 'endnote'])
@pytest.mark.parametrize('variant', ['hidden-reference-run', 'duplicate-ppr', 'late-pstyle', 'table-reference'])
def test_unproved_reference_cannot_delete_current_note_line(tmp_path, kind, variant):
    def mutate(parts, kind):
        root = etree.fromstring(parts['word/document.xml'])
        reference = next(root.iter(qn('w:' + kind + 'Reference')))
        run = reference.getparent(); p = run.getparent()
        if variant == 'hidden-reference-run':
            props = etree.Element(qn('w:rPr')); run.insert(0, props)
            etree.SubElement(props, qn('w:vanish'))
        elif variant == 'table-reference':
            # A conditionally styled table is not a plain paragraph proof.
            table = etree.Element(qn('w:tbl'))
            props = etree.SubElement(table, qn('w:tblPr'))
            etree.SubElement(props, qn('w:tblStyle'), {qn('w:val'): 'TableGrid'})
            grid = etree.SubElement(table, qn('w:tblGrid')); etree.SubElement(grid, qn('w:gridCol'), {qn('w:w'): '8640'})
            cell = etree.SubElement(etree.SubElement(table, qn('w:tr')), qn('w:tc'))
            p.addprevious(table); cell.append(p)
        else:
            props = etree.Element(qn('w:pPr')); p.insert(0, props)
            if variant == 'duplicate-ppr':
                etree.SubElement(props, qn('w:pStyle'), {qn('w:val'): 'Normal'})
                props = etree.Element(qn('w:pPr')); p.insert(1, props)
            else:
                etree.SubElement(props, qn('w:spacing'), {qn('w:after'): '0'})
            etree.SubElement(props, qn('w:pStyle'), {qn('w:val'): 'Normal'})
        parts['word/document.xml'] = etree.tostring(root)
    path, block = scene(tmp_path, kind, mutate)
    raw, paragraphs = render(path, replace(block, text='Current authored narrative.\n' + NOTE))
    assert NOTE in paragraphs


@pytest.mark.parametrize('kind', ['footnote', 'endnote'])
@pytest.mark.parametrize('variant', ['safe', 'qualifier', 'direct-hidden', 'duplicate-pstyle'])
def test_actual_pdf_keeps_authored_note_words(tmp_path, monkeypatch, kind, variant):
    from app import pdf_converter
    from conftest import libreoffice_can_convert
    import pymupdf
    assert libreoffice_can_convert(), 'This note-preservation render gate requires installed LibreOffice.'
    mutation = (lambda parts, kind: mutate_note_style(parts, kind, variant)) if variant in ('direct-hidden', 'duplicate-pstyle') else None
    path, block = scene(tmp_path, kind, mutation)
    current = (QUALIFIER if variant == 'qualifier' else 'Current authored narrative.') + '\n' + NOTE
    raw, paragraphs = render(path, replace(block, text=current))
    output = tmp_path / 'current.docx'; output.write_bytes(raw)
    monkeypatch.setenv('EPC_DATA_DIR', str(tmp_path / 'runtime'))
    monkeypatch.setattr(pdf_converter, 'OUTPUT_DIR', tmp_path)
    converted = pdf_converter.convert_to_pdf(output)
    with pymupdf.open(converted) as pdf:
        text = '\n'.join(page.get_text() for page in pdf)
    if variant == 'qualifier':
        assert QUALIFIER in text
        assert text.count(NOTE) >= 2
    else:
        assert 'Current authored narrative.' in text
        assert text.count(NOTE) >= 1
