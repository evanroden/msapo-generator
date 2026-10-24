"""Current prose must not inherit hidden source paragraph/character defaults."""
from dataclasses import replace
from io import BytesIO
from zipfile import ZipFile
from copy import deepcopy
from lxml import etree
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt
import pytest
from test_monthly_report_native_note_preservation import scene, render, NOTE
from app.monthly_report_native_layout import build_native_docx

CURRENT = 'Current synthetic narrative remains readable.'


def source_mutation(parts, kind, variant):
    root = etree.fromstring(parts['word/document.xml'])
    reference = next(root.iter(qn('w:' + kind + 'Reference')))
    paragraph = reference.getparent().getparent()
    styles = etree.fromstring(parts['word/styles.xml'])
    hidden = etree.SubElement(styles, qn('w:style'), {qn('w:type'): 'paragraph', qn('w:styleId'): 'SyntheticHiddenSource'})
    etree.SubElement(hidden, qn('w:name'), {qn('w:val'): 'Synthetic Hidden Source'})
    etree.SubElement(hidden, qn('w:basedOn'), {qn('w:val'): 'Normal'})
    etree.SubElement(etree.SubElement(hidden, qn('w:rPr')), qn('w:vanish'))
    props = etree.Element(qn('w:pPr')); paragraph.insert(0, props)
    identity = 'SyntheticHiddenSource'
    if variant == 'inherited-hidden':
        child = etree.SubElement(styles, qn('w:style'), {qn('w:type'): 'paragraph', qn('w:styleId'): 'SyntheticChild'})
        etree.SubElement(child, qn('w:name'), {qn('w:val'): 'Synthetic Child'})
        etree.SubElement(child, qn('w:basedOn'), {qn('w:val'): identity}); identity = 'SyntheticChild'
    elif variant in ('default-paragraph', 'default-character', 'document-default'):
        identity = 'Normal'
        if variant == 'document-default':
            target = styles.find('./' + qn('w:docDefaults') + '/' + qn('w:rPrDefault') + '/' + qn('w:rPr'))
        else:
            style_kind = 'paragraph' if variant == 'default-paragraph' else 'character'
            target_style = next(s for s in styles.findall(qn('w:style')) if s.get(qn('w:type')) == style_kind and s.get(qn('w:default')) == '1')
            target = target_style.find(qn('w:rPr'))
            if target is None: target = etree.SubElement(target_style, qn('w:rPr'))
        etree.SubElement(target, qn('w:vanish'))
    elif variant == 'duplicate-selector':
        etree.SubElement(props, qn('w:pStyle'), {qn('w:val'): 'Normal'})
    elif variant == 'malformed-chain':
        etree.SubElement(hidden.find(qn('w:basedOn')), qn('w:basedOn'), {qn('w:val'): 'Normal'})
    elif variant == 'duplicate-style':
        styles.append(deepcopy(hidden))
    etree.SubElement(props, qn('w:pStyle'), {qn('w:val'): identity})
    parts['word/document.xml'] = etree.tostring(root)
    parts['word/styles.xml'] = etree.tostring(styles)


VARIANTS = ['source-hidden', 'inherited-hidden', 'default-paragraph', 'default-character',
            'document-default', 'duplicate-selector', 'malformed-chain', 'duplicate-style']


@pytest.mark.parametrize('kind', ['footnote', 'endnote'])
@pytest.mark.parametrize('variant', VARIANTS)
def test_rewritten_current_text_has_visible_independent_format(tmp_path, variant, kind):
    path, block = scene(tmp_path, kind, lambda parts, kind: source_mutation(parts, kind, variant))
    raw, paras = render(path, replace(block, text=CURRENT + '\n' + NOTE))
    assert CURRENT in paras and NOTE in paras
    doc = Document(BytesIO(raw))
    paragraph = next(p for p in doc.paragraphs if p.text == CURRENT)
    assert paragraph.style.style_id.startswith('ENFRACurrentVisibleNarrative')
    assert paragraph.runs[0].font.hidden is False
    assert paragraph.runs[0].font.size == Pt(11)
    assert paragraph.style.base_style is None


@pytest.mark.parametrize('kind', ['footnote', 'endnote'])
@pytest.mark.parametrize('variant', VARIANTS)
def test_required_actual_pdf_current_prose_survives_hidden_inheritance(tmp_path, monkeypatch, kind, variant):
    from conftest import libreoffice_can_convert
    from app import pdf_converter
    import pymupdf
    assert libreoffice_can_convert(), 'Required actual-PDF narrative visibility gate needs LibreOffice.'
    path, block = scene(tmp_path, kind, lambda parts, kind: source_mutation(parts, kind, variant))
    before = path.read_bytes()
    raw, _ = render(path, replace(block, text=CURRENT + '\n' + NOTE))
    output = tmp_path / 'rewritten.docx'; output.write_bytes(raw)
    monkeypatch.setenv('EPC_DATA_DIR', str(tmp_path / 'runtime'))
    monkeypatch.setattr(pdf_converter, 'OUTPUT_DIR', tmp_path)
    pdf = pdf_converter.convert_to_pdf(output)
    with pymupdf.open(pdf) as doc:
        text = '\n'.join(p.get_text() for p in doc)
    assert text.count(CURRENT) == 1
    assert text.count(NOTE) >= 1
    assert path.read_bytes() == before


def test_proved_typography_is_returned_unchanged():
    from app.monthly_report_native_narrative_visibility import narrative_visibility_writer
    doc = Document(); paragraph = doc.add_paragraph('Synthetic current prose.')
    paragraph.paragraph_format.left_indent = Pt(18)
    run = paragraph.runs[0]; run.bold = True; run.italic = True; run.underline = True; run.font.size = Pt(12)
    before, styles = etree.tostring(paragraph._p), etree.tostring(doc.styles.element)
    result = narrative_visibility_writer(doc)(paragraph._p)
    assert result is paragraph._p
    assert etree.tostring(result) == before and etree.tostring(doc.styles.element) == styles


def test_normal_numbering_and_unrelated_bad_styles_do_not_reset_typography():
    from app.monthly_report_native_narrative_visibility import narrative_visibility_writer
    doc = Document(); paragraph = doc.add_paragraph('Synthetic current prose.')
    paragraph.runs[0].font.size = Pt(13)
    props = paragraph._p.get_or_add_pPr()
    num = etree.SubElement(props, qn('w:numPr')); etree.SubElement(num, qn('w:numId'), {qn('w:val'): '0'})
    # An unrelated malformed style cannot force changes to this safe paragraph.
    style = etree.SubElement(doc.styles.element, qn('w:style'), {qn('w:type'): 'paragraph', qn('w:styleId'): 'UnrelatedMalformed'})
    bad = etree.SubElement(etree.SubElement(style, qn('w:pPr')), qn('w:numPr'))
    etree.SubElement(bad, qn('w:rStyle'), {qn('w:val'): 'NoSuchStyle'})
    assert narrative_visibility_writer(doc)(paragraph._p) is paragraph._p


@pytest.mark.parametrize('variant', ['duplicate-numbering', 'late-numbering', 'existing-fallback-id'])
def test_malformed_format_and_fallback_name_collision_are_not_trusted(variant):
    from app.monthly_report_native_narrative_visibility import narrative_visibility_writer
    doc = Document(); paragraph = doc.add_paragraph('Synthetic current prose.')
    props = paragraph._p.get_or_add_pPr()
    if variant == 'late-numbering':
        etree.SubElement(props, qn('w:spacing'), {qn('w:after'): '0'})
    num = etree.SubElement(props, qn('w:numPr')); etree.SubElement(num, qn('w:numId'), {qn('w:val'): '0'})
    if variant == 'duplicate-numbering': props.append(deepcopy(num))
    if variant == 'existing-fallback-id':
        style = etree.SubElement(doc.styles.element, qn('w:style'), {qn('w:type'): 'paragraph', qn('w:styleId'): 'ENFRACurrentVisibleNarrative'})
        etree.SubElement(style, qn('w:name'), {qn('w:val'): 'Source name collision'})
        etree.SubElement(etree.SubElement(style, qn('w:rPr')), qn('w:vanish'))
        props.insert(0, etree.Element(qn('w:pStyle'), {qn('w:val'): 'ENFRACurrentVisibleNarrative'}))
    result = narrative_visibility_writer(doc)(paragraph._p)
    assert result is not paragraph._p
    assert result.find('./' + qn('w:pPr') + '/' + qn('w:pStyle')).get(qn('w:val')).startswith('ENFRACurrentVisibleNarrative')
    if variant == 'existing-fallback-id':
        assert result.find('./' + qn('w:pPr') + '/' + qn('w:pStyle')).get(qn('w:val')) == 'ENFRACurrentVisibleNarrative1'


def test_fallback_never_defines_previously_missing_style_used_in_unchanged_parts():
    from app.monthly_report_native_narrative_visibility import narrative_visibility_writer
    doc = Document()
    untouched = doc.add_paragraph('Unchanged source artwork label.')
    pp = untouched._p.get_or_add_pPr()
    etree.SubElement(pp, qn('w:pStyle'), {qn('w:val'): 'ENFRACurrentVisibleNarrative'})
    header = doc.sections[0].header.paragraphs[0]
    etree.SubElement(header._p.get_or_add_pPr(), qn('w:pStyle'), {qn('w:val'): 'ENFRACurrentVisibleNarrative1'})
    unsafe = doc.add_paragraph('Current authored prose.')
    unsafe.runs[0].font.hidden = True
    source_before, header_before = etree.tostring(untouched._p), etree.tostring(header._p)
    output = narrative_visibility_writer(doc)(unsafe._p)
    selected = output.find('./' + qn('w:pPr') + '/' + qn('w:pStyle')).get(qn('w:val'))
    assert selected == 'ENFRACurrentVisibleNarrative2'
    assert etree.tostring(untouched._p) == source_before
    assert etree.tostring(header._p) == header_before
    assert not any(s.get(qn('w:styleId')) in ('ENFRACurrentVisibleNarrative', 'ENFRACurrentVisibleNarrative1')
                   for s in doc.styles.element.findall(qn('w:style')))


@pytest.mark.parametrize('override', ['paragraph-off', 'character-off', 'direct-run-off'])
def test_proved_visible_override_keeps_typography_and_renders(tmp_path, monkeypatch, override):
    from conftest import libreoffice_can_convert
    from app import pdf_converter
    import pymupdf
    assert libreoffice_can_convert(), 'Required actual-PDF visibility override gate needs LibreOffice.'
    def mutate(parts, kind):
        source_mutation(parts, kind, 'source-hidden' if override == 'direct-run-off' else 'document-default')
        root = etree.fromstring(parts['word/document.xml'])
        reference = next(root.iter(qn('w:' + kind + 'Reference')))
        p = reference.getparent().getparent(); run = p.find(qn('w:r'))
        rp = run.find(qn('w:rPr'))
        if rp is None: rp = etree.Element(qn('w:rPr')); run.insert(0, rp)
        etree.SubElement(rp, qn('w:sz'), {qn('w:val'): '26'})
        if override == 'direct-run-off':
            rp.insert(0, etree.Element(qn('w:vanish'), {qn('w:val'): '0'}))
        else:
            styles = etree.fromstring(parts['word/styles.xml'])
            identity = 'Normal' if override == 'paragraph-off' else 'DefaultParagraphFont'
            target = next(s for s in styles.findall(qn('w:style')) if s.get(qn('w:styleId')) == identity)
            prop = target.find(qn('w:rPr'))
            if prop is None: prop = etree.SubElement(target, qn('w:rPr'))
            etree.SubElement(prop, qn('w:vanish'), {qn('w:val'): '0'})
            parts['word/styles.xml'] = etree.tostring(styles)
        parts['word/document.xml'] = etree.tostring(root)
    path, block = scene(tmp_path, 'footnote', mutate)
    raw, _ = render(path, replace(block, text=CURRENT + '\n' + NOTE))
    document = Document(BytesIO(raw)); p = next(p for p in document.paragraphs if p.text == CURRENT)
    assert not p.style.style_id.startswith('ENFRACurrentVisibleNarrative')
    assert p.runs[0].font.size == Pt(13)
    output = tmp_path / 'current.docx'; output.write_bytes(raw)
    monkeypatch.setenv('EPC_DATA_DIR', str(tmp_path / 'runtime')); monkeypatch.setattr(pdf_converter, 'OUTPUT_DIR', tmp_path)
    pdf = pdf_converter.convert_to_pdf(output)
    with pymupdf.open(pdf) as document:
        text = '\n'.join(p.get_text() for p in document)
    assert text.count(CURRENT) == 1 and text.count(NOTE) >= 1


def test_saturated_dark_blue_is_not_mistaken_for_light_text():
    from docx.shared import RGBColor
    from app.monthly_report_native_narrative_visibility import narrative_visibility_writer
    doc = Document(); paragraph = doc.add_paragraph('Current blue body typography.')
    paragraph.runs[0].font.color.rgb = RGBColor.from_string('0000FF')
    paragraph.runs[0].font.size = Pt(13)
    assert narrative_visibility_writer(doc)(paragraph._p) is paragraph._p


@pytest.mark.parametrize('empty_run', [False, True])
def test_no_text_prototype_checks_the_run_writer_will_create(empty_run):
    from app.monthly_report_native_narrative_visibility import narrative_visibility_writer
    from app.monthly_report_native_text import write_narrative
    from docx.oxml import OxmlElement
    doc = Document(); doc.styles['Normal'].font.hidden = True
    empty = OxmlElement('w:p')
    if empty_run:
        run = OxmlElement('w:r'); empty.append(run)
        props = OxmlElement('w:rPr'); run.append(props)
        hidden = OxmlElement('w:vanish'); hidden.set(qn('w:val'), '0'); props.append(hidden)
    anchor = doc.add_paragraph()._p
    prototype = narrative_visibility_writer(doc)(empty)
    write_narrative(anchor, prototype, CURRENT)
    assert anchor.find('./' + qn('w:pPr') + '/' + qn('w:pStyle')).get(qn('w:val')).startswith('ENFRACurrentVisibleNarrative')
    assert Document(BytesIO(_save(doc))).paragraphs[0].runs[0].font.hidden is False


def _save(document):
    target = BytesIO(); document.save(target); return target.getvalue()


def test_table_only_native_source_produces_visible_authored_prose(tmp_path, monkeypatch):
    from test_monthly_report_native_layout import draft
    from app.monthly_report_model import ResolvedBlock
    from app import pdf_converter
    from conftest import libreoffice_can_convert
    import pymupdf
    assert libreoffice_can_convert(), 'Required actual-PDF table-only gate needs LibreOffice.'
    source = tmp_path / 'table-only.docx'
    doc = Document(); doc.styles['Normal'].font.hidden = True
    doc.add_heading('MONTHLY ACTIVITY SUMMARY', 1)
    table = doc.add_table(rows=2, cols=2)
    for cell, value in zip(table.rows[0].cells, ('Task', 'Status')): cell.text = value
    for cell, value in zip(table.rows[1].cells, ('Previous synthetic task', 'Previous status')): cell.text = value
    doc.save(source)
    raw = build_native_docx(draft(ResolvedBlock('activity_summary', 'This month', text=CURRENT)), source)
    output = tmp_path / 'current.docx'; output.write_bytes(raw)
    monkeypatch.setenv('EPC_DATA_DIR', str(tmp_path / 'runtime')); monkeypatch.setattr(pdf_converter, 'OUTPUT_DIR', tmp_path)
    pdf = pdf_converter.convert_to_pdf(output)
    with pymupdf.open(pdf) as document:
        text = '\n'.join(p.get_text() for p in document)
    assert text.count(CURRENT) == 1
