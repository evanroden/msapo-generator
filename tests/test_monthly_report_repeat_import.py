"""Repeat imports preserve reviewed work and only compare actual table payloads."""

from dataclasses import replace

from docx import Document
import pytest

from app.monthly_report_directory import CONTACT_SPEC, apply_contacts
from app.monthly_report_import import ImportMapping, MappedImport, imported_draft, inspect_docx, map_items
from app.monthly_report_model import BlockSpec, ColumnSpec, ReportPeriod, ResolvedBlock, synthetic_profiles
from app.monthly_report_setup import merge_blocks, merge_drafts


PERIOD = ReportPeriod(2026, 9)
CONTACT = ("Synthetic North", "Manager", "Synthetic Editor", "202-555-0100", "reviewed@example.invalid")


def draft_with_contacts():
    base = imported_draft(synthetic_profiles()[0], PERIOD, "Synthetic Editor", MappedImport((), (), (), ()))
    contacts = ResolvedBlock("contact_matrix", "This month", rows=(CONTACT,), references=("directory:synthetic:1",))
    contacts = replace(contacts, client_reviewed_fingerprint=contacts.fingerprint)
    return apply_contacts(base, contacts)


def block(draft, key="contact_matrix"):
    return next(b for b in draft.blocks if b.key == key)


def spec(draft, key="contact_matrix"):
    return next(b for s in draft.sections for b in s.blocks if b.key == key)


def import_table(tmp_path, profile, columns, rows, key="contact_matrix"):
    doc = Document()
    table = doc.add_table(rows=1, cols=len(columns))
    for cell, title in zip(table.rows[0].cells, columns):
        cell.text = title
    for row in rows:
        for cell, value in zip(table.add_row().cells, row):
            cell.text = value
    path = tmp_path / "synthetic-contacts.docx"
    doc.save(path)
    inspection = inspect_docx(path)
    item = next(i for i in inspection.items if i.kind == "table")
    mapped = map_items(path, inspection, (ImportMapping(item.id, key),))
    return imported_draft(profile, PERIOD, "Synthetic Importer", mapped)


def test_import_after_reviewed_contacts_preserves_table_and_review(tmp_path, monkeypatch):
    from app import monthly_report_library as library

    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path / "data"))
    existing = draft_with_contacts()
    manual = ResolvedBlock("activity_summary", "This month", text="Manual reviewed repair details.")
    existing = replace(existing, blocks=(*existing.blocks, manual))
    saved = library.save_snapshot(existing, expected_revision=0, entered_editor="Synthetic Editor")
    existing = library.load_snapshot(existing.profile.contract, existing.profile.key, PERIOD).draft
    assert saved.draft == existing

    doc = Document()
    doc.add_heading("Monthly Activity Summary", 1)
    doc.add_paragraph("Additional synthetic repair.")
    path = tmp_path / "second-report.docx"
    doc.save(path)
    inspection = inspect_docx(path)
    item = next(i for i in inspection.items if i.text == "Additional synthetic repair.")
    incoming = imported_draft(existing.profile, PERIOD, "Synthetic Importer",
                              map_items(path, inspection, (ImportMapping(item.id, "activity_summary"),)))
    assert spec(incoming).type == "image_page"

    merged = merge_drafts(existing, incoming)
    assert block(merged) == block(existing)
    assert spec(merged) == CONTACT_SPEC
    assert block(merged, "activity_summary").text == manual.text + "\n\n" + item.text
    assert merge_drafts(merged, incoming) == merged


def test_imported_contact_keys_align_with_reviewed_directory_columns(tmp_path):
    existing = draft_with_contacts()
    added = ("Synthetic North", "Technician", "Synthetic Worker", "", "worker@example.invalid")
    incoming = import_table(tmp_path, existing.profile, tuple(c.title for c in CONTACT_SPEC.columns), (CONTACT, added))
    assert spec(incoming).columns[0].key == "imported_0"
    merged = merge_drafts(existing, incoming)
    assert spec(merged) == CONTACT_SPEC
    assert block(merged).rows == (CONTACT, added)
    assert block(merged).references == (*block(existing).references, *block(incoming).references)
    assert block(existing).rows == (CONTACT,)
    assert merge_drafts(merged, incoming) == merged


def test_reordered_contact_headings_preserve_unknown_columns_and_manual_values(tmp_path):
    existing = draft_with_contacts()
    with_extra = replace(CONTACT_SPEC, columns=(*CONTACT_SPEC.columns, ColumnSpec("unknown", "Escalation Notes")))
    existing = replace(existing, sections=tuple(replace(s, blocks=tuple(with_extra if b.key == "contact_matrix" else b for b in s.blocks)) for s in existing.sections),
                       blocks=(replace(block(existing), rows=((*CONTACT, "Manual escalation wording"),)),))
    titles = ("EMAIL", "Name", " Escalation   Notes ", "Phone", "Role", "Facility")
    row = ("new@example.invalid", "Synthetic Worker", "New escalation wording", "202-555-0101", "Technician", "Synthetic North")
    incoming = import_table(tmp_path, existing.profile, titles, (row,))
    merged = merge_drafts(existing, incoming)
    assert spec(merged) == with_extra
    assert block(merged).rows == ((*CONTACT, "Manual escalation wording"),
                                  ("Synthetic North", "Technician", "Synthetic Worker", "202-555-0101", "new@example.invalid", "New escalation wording"))


@pytest.mark.parametrize("titles", [
    ("Facility", "Role", "Name", "Phone", "Email", "Unmapped notes"),
    ("Facility", "Role", "Name", "Phone"),
    ("Facility", "Role", "Name", "Phone", "Contact details"),
    ("Facility", "Role", "Name", "Phone", "Email", " email "),
])
def test_unknown_or_ambiguous_contact_schema_is_rejected_without_dropping_cells(tmp_path, titles):
    existing = draft_with_contacts()
    incoming = import_table(tmp_path, existing.profile, titles, (tuple("Synthetic value" for _ in titles),))
    before = existing
    with pytest.raises(ValueError, match="Imported table columns differ"):
        merge_drafts(existing, incoming)
    assert existing == before
    assert len(block(incoming).rows[0]) == len(titles)


def test_noncontact_table_conflicts_still_require_mapping(tmp_path):
    existing = draft_with_contacts()
    original_spec = BlockSpec("service_calls", "table", columns=(ColumnSpec("description", "Description"),))
    existing = replace(existing, sections=tuple(replace(s, blocks=tuple(original_spec if b.key == "service_calls" else b for b in s.blocks)) for s in existing.sections),
                       blocks=(*existing.blocks, ResolvedBlock("service_calls", "This month", rows=(("Manual details",),))))
    incoming = import_table(tmp_path, existing.profile, ("Description",), (("Imported details",),), "service_calls")
    with pytest.raises(ValueError, match="Imported table columns differ"):
        merge_drafts(existing, incoming)
    assert block(existing, "service_calls").rows == (("Manual details",),)


def test_unused_noncontact_template_spec_does_not_conflict_with_saved_table():
    existing = draft_with_contacts()
    original_spec = BlockSpec("service_calls", "table", columns=(ColumnSpec("description", "Description"),))
    existing = replace(existing, sections=tuple(replace(s, blocks=tuple(original_spec if b.key == "service_calls" else b for b in s.blocks)) for s in existing.sections),
                       blocks=(*existing.blocks, ResolvedBlock("service_calls", "This month", rows=(("Manual details",),))))
    incoming = imported_draft(existing.profile, PERIOD, "Synthetic Importer",
                              MappedImport((ResolvedBlock("activity_summary", "Last month", text="Additional work."),), (), (), ()))
    merged = merge_drafts(existing, incoming)
    assert block(merged, "service_calls") == block(existing, "service_calls")
    assert spec(merged, "service_calls") == original_spec


@pytest.mark.parametrize("caption", [(), ("Approved technical page",)])
@pytest.mark.parametrize("source", ["This month", "Last month"])
def test_approved_page_merge_into_empty_placeholder_keeps_approval(caption, source):
    approved = ResolvedBlock("vendor_reports", source, asset_hashes=("synthetic.png",), asset_captions=caption,
                             references=("source:1",))
    approved = replace(approved, client_reviewed_fingerprint=approved.fingerprint)
    merged = merge_blocks(ResolvedBlock("vendor_reports", "This month"), approved)
    assert merged.client_reviewed_fingerprint == merged.fingerprint
    assert merged.asset_hashes == approved.asset_hashes
    assert merged.references == approved.references


def test_reimport_identical_approved_pages_keeps_approval():
    approved = ResolvedBlock("vendor_reports", "This month", asset_hashes=("synthetic.png",), references=("source:1",))
    approved = replace(approved, client_reviewed_fingerprint=approved.fingerprint)
    incoming = replace(approved, source="Last month", client_reviewed_fingerprint="")
    merged = merge_blocks(approved, incoming)
    assert merged.client_reviewed_fingerprint == merged.fingerprint
    assert merge_blocks(merged, incoming) == merged


@pytest.mark.parametrize("changes", [
    {"text": "Unreviewed page context."},
    {"references": ("source:2",)},
    {"asset_hashes": ("different.png",), "asset_captions": ("Different page",)},
])
def test_new_page_material_invalidates_previous_approval(changes):
    approved = ResolvedBlock("vendor_reports", "This month", asset_hashes=("synthetic.png",),
                             asset_captions=("Approved page",), references=("source:1",))
    approved = replace(approved, client_reviewed_fingerprint=approved.fingerprint)
    incoming = replace(approved, **changes, client_reviewed_fingerprint="")
    merged = merge_blocks(approved, incoming)
    assert merged.client_reviewed_fingerprint == ""


def test_edited_saved_caption_cannot_inherit_old_page_approval():
    approved = ResolvedBlock("vendor_reports", "This month", asset_hashes=("synthetic.png",), asset_captions=("Approved page",))
    approved = replace(approved, client_reviewed_fingerprint=approved.fingerprint)
    edited = replace(approved, asset_captions=("Edited caption",))
    merged = merge_blocks(edited, replace(approved, source="Last month"))
    assert merged.asset_captions == edited.asset_captions
    assert merged.client_reviewed_fingerprint == ""
