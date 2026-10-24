"""Known branding cannot hide an adjacent cover photo or waive owner review."""

from dataclasses import replace
from io import BytesIO

from docx import Document
from docx.shared import Inches
import pytest

from app.monthly_report_import import ImportMapping, inspect_docx, map_items
from app.monthly_report_model import ReportPeriod, synthetic_profiles
from app.monthly_report_start_import import (
    analyze_starting_report,
    automatic_section_plans,
    cover_photo_candidates,
    unresolved_cover_artwork,
)
from test_cover_branding_contracts import mark, photograph


def _source(tmp_path, first, *, small=False, boundary=""):
    document = Document()
    document.add_paragraph(synthetic_profiles()[0].facilities[0].title)
    document.add_paragraph("November 2026")
    if first == "header-copy":
        document.sections[0].header.paragraphs[0].add_run().add_picture(
            BytesIO(mark()), width=Inches(3)
        )
        document.add_picture(BytesIO(mark()), width=Inches(3))
    elif first == "ambiguous-pair":
        paragraph = document.add_paragraph("Logo")
        paragraph.add_run().add_picture(BytesIO(mark()), width=Inches(3))
        paragraph.add_run().add_picture(BytesIO(mark((130, 80, 30, 255))), width=Inches(3))
    elif first != "none":
        labels = {"generic": "Logo", "company": "Company logo", "client": "Client logo"}
        if first in labels:
            document.add_paragraph(labels[first])
        document.add_picture(
            BytesIO(photograph() if first == "photo" else mark()), width=Inches(3)
        )
    document.add_section()
    if boundary == "later-section":
        document.add_section()
    elif boundary:
        document.add_paragraph(boundary)
    document.add_picture(
        BytesIO(photograph((96, 72) if small else (1000, 700))), width=Inches(5)
    )
    document.add_section()
    document.add_paragraph("Table of Contents")
    document.add_heading("Monthly Activity Summary", 1)
    path = tmp_path / "adjacent-branding.docx"
    document.save(path)
    return path


def _plan(inspection):
    return automatic_section_plans(
        inspection,
        analyze_starting_report(inspection, synthetic_profiles()[0], ReportPeriod(2026, 11)),
    )["cover"]


@pytest.mark.parametrize("first", ["generic", "header-copy", "ambiguous-pair"])
@pytest.mark.parametrize("small", [False, True], ids=["photo", "small-artwork"])
def test_unassigned_branding_keeps_adjacent_picture_reachable_without_approval(tmp_path, first, small):
    source = _source(tmp_path, first, small=small)
    original = source.read_bytes()
    inspection = inspect_docx(source)
    body = tuple(item for item in inspection.items
                 if item.kind == "image" and item.part == "word/document.xml")
    branding, adjacent = body[:-1], body[-1]
    assert branding and all(not item.suggested_slot for item in branding)
    assert all(item.image_role == ("header_artwork" if first == "header-copy" else "logo")
               for item in branding)
    candidates = cover_photo_candidates(inspection.items)
    unresolved = unresolved_cover_artwork(inspection.items)
    if small:
        assert candidates == ()
        assert adjacent in unresolved, "Small adjacent artwork must remain recoverable"
    else:
        assert candidates == (adjacent,), "Known branding must not hide the adjacent photograph"
        assert adjacent not in unresolved
    assert all(item in unresolved for item in branding), "Owner ambiguity must remain unresolved"
    plan = _plan(inspection)
    assert not plan["approved"] and plan["questions"]
    assert all(item.id not in plan["destinations"] for item in branding)
    if not small:
        assert plan["destinations"][adjacent.id] == "cover_photo"
    for item in branding:
        with pytest.raises(ValueError, match="identified as a logo|Header artwork"):
            map_items(source, inspection, (ImportMapping(item.id, "cover_photo"),))
    assert source.read_bytes() == original


@pytest.mark.parametrize("first", ["none", "company", "client"])
def test_explicit_branding_and_title_only_cover_keep_existing_adjacent_mapping(tmp_path, first):
    inspection = inspect_docx(_source(tmp_path, first))
    adjacent = next(item for item in reversed(inspection.items)
                    if item.kind == "image" and item.part == "word/document.xml")
    assert cover_photo_candidates(inspection.items) == (adjacent,)
    assert not unresolved_cover_artwork(inspection.items)
    plan = _plan(inspection)
    assert plan["approved"]
    assert plan["destinations"][adjacent.id] == "cover_photo"


@pytest.mark.parametrize("first", ["photo", "unknown"])
def test_independent_first_cover_art_still_prevents_adjacent_promotion(tmp_path, first):
    inspection = inspect_docx(_source(tmp_path, first))
    images = tuple(item for item in inspection.items if item.kind == "image")
    assert len(images) == 2
    assert images[1] not in cover_photo_candidates(inspection.items)
    assert images[1] not in unresolved_cover_artwork(inspection.items)


def test_unsupported_first_drawing_still_prevents_adjacent_promotion(tmp_path):
    inspection = inspect_docx(_source(tmp_path, "unknown"))
    first = next(item for item in inspection.items if item.kind == "image")
    items = tuple(replace(item, kind="unsupported") if item.id == first.id else item
                  for item in inspection.items)
    assert not cover_photo_candidates(items)
    assert not unresolved_cover_artwork(items)


@pytest.mark.parametrize("boundary", ["Table of Contents", "Monthly Activity Summary", "later-section"])
@pytest.mark.parametrize("small", [False, True], ids=["photo", "small-artwork"])
def test_branding_does_not_extend_cover_across_real_boundaries(tmp_path, boundary, small):
    inspection = inspect_docx(_source(tmp_path, "generic", small=small, boundary=boundary))
    body = tuple(item for item in inspection.items
                 if item.kind == "image" and item.part == "word/document.xml")
    assert len(body) == 2
    assert not cover_photo_candidates(inspection.items)
    assert body[0] in unresolved_cover_artwork(inspection.items)
    assert body[1] not in unresolved_cover_artwork(inspection.items)
