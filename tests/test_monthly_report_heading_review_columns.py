"""Keep every source column visible while ambiguous headings await correction."""
from dataclasses import replace
from io import BytesIO

from docx import Document
import pytest
from streamlit.testing.v1 import AppTest

from app import monthly_report_library as library
from app.monthly_report_docx import assemble_docx
from app.monthly_report_model import (
    ReportPeriod, ReportTable, ResolvedBlock, default_sections, synthetic_draft, synthetic_profiles,
)


SCRIPT = '''
import streamlit as st
from app.monthly_report_model import ReportTable
from app.monthly_report_ui import _field
from app.monthly_report_table_review_ui import review_table_headings
from app.monthly_report_content_policy import table_has_pricing
original = ORIGINAL
current = st.session_state.setdefault('saved', original)
st.session_state['saved'] = review_table_headings(current, 'heading_review_test', _field)
st.session_state['original'] = original
st.session_state['blocked'] = table_has_pricing(st.session_state['saved'].columns, st.session_state['saved'].rows)
'''


def _workspace(table):
    app = AppTest.from_string(SCRIPT.replace("ORIGINAL", repr(table))).run()
    assert not app.exception
    return app


def _fingerprint(table):
    return ResolvedBlock("capital_renewal", "This month", extra_tables=(table,)).fingerprint


CASES = (
    ReportTable(("Description Cost", "Description Cost", "Recommendation"),
                (("Pump P-001", "$1200", "Replace seal"),), "synthetic-duplicate-source"),
    ReportTable(("Description Cost", "Description Cost", "1. Description Cost", "Recommendation"),
                (("Pump P-001", "$1200", "Retain separate technical detail", "Replace seal"),),
                "synthetic-literal-prefix-source"),
    ReportTable(("Description Cost", "", "", "Recommendation"),
                (("Pump P-001", "2030", "$1200", "Replace seal"),), "synthetic-blank-heading-source"),
)


@pytest.mark.parametrize("table", CASES, ids=["duplicate", "literal_prefix", "blank_headings"])
def test_native_heading_review_shows_each_source_cell_without_mutation(table):
    original_fingerprint = _fingerprint(table)
    app = _workspace(table)
    assert len(app.dataframe) == 1
    frame = app.dataframe[0].value
    assert frame.shape == (len(table.rows), len(table.columns))
    assert frame.columns.is_unique
    assert tuple(frame.columns) == tuple(f"{index + 1}. {heading}" for index, heading in enumerate(table.columns))
    assert tuple(tuple(row) for row in frame.itertuples(index=False, name=None)) == table.rows
    assert tuple(field.value for field in app.text_input) == table.columns
    assert app.button[0].disabled
    assert app.session_state["blocked"]
    assert app.session_state["saved"] == table == app.session_state["original"]
    assert _fingerprint(app.session_state["saved"]) == original_fingerprint
    app.run()
    assert not app.exception
    assert tuple(app.dataframe[0].value.columns) == tuple(frame.columns)
    assert app.session_state["saved"] == table


def test_unique_source_headings_keep_existing_preview_labels():
    table = ReportTable(("Description Cost", "Column 2", "Recommendation"),
                        (("Pump P-001", "$1200", "Replace seal"),), "synthetic-unique-source")
    app = _workspace(table)
    assert tuple(app.dataframe[0].value.columns) == table.columns
    assert tuple(app.dataframe[0].value.iloc[0]) == table.rows[0]
    assert app.session_state["saved"] == table


def test_guided_proposal_correction_keeps_all_reviewed_values_through_native_reopen():
    table = CASES[0]
    original = ResolvedBlock("proposals", "This month", extra_tables=(table,),
                             references=("synthetic-owned-source",))
    spec = next(block for section in default_sections() for block in section.blocks if block.key == "proposals")
    script = '''
from types import SimpleNamespace
import streamlit as st
from app.monthly_report_guided import _edit_content
from app.monthly_report_model import BlockSpec, ColumnSpec, ReportPeriod, ReportTable, ResolvedBlock, synthetic_draft, synthetic_profiles
from app.monthly_report_ui import _field
draft = st.session_state.setdefault('draft', synthetic_draft(synthetic_profiles()[0], ReportPeriod(2026, 9)))
original = ORIGINAL
block = st.session_state.setdefault('block', original)
spec = SPEC
st.session_state['block'] = _edit_content(spec, block, SimpleNamespace(profile=draft.profile), 'guided_headings', {}, _field, draft)
st.session_state['original'] = original
'''
    app = AppTest.from_string(script.replace("ORIGINAL", repr(original)).replace("SPEC", repr(spec))).run()
    assert not app.exception
    assert tuple(app.dataframe[0].value.iloc[0]) == table.rows[0]
    assert not any(item.label == "Edit proposal" for item in app.button)
    for control, title in zip(app.text_input, ("Equipment", "Cost", "Recommendation")):
        control.set_value(title)
    app.run()
    assert app.session_state["block"] == original
    next(item for item in app.button if item.label == "Use these column names").click().run()
    app.run()
    expected = ReportTable(("Equipment", "Recommendation"), (("Pump P-001", "Replace seal"),), table.reference)
    assert not app.exception
    assert app.session_state["block"].extra_tables == (expected,)
    assert app.session_state["block"].references == original.references
    for attempt in range(2):
        next(item for item in app.button if item.label == "Edit proposal").click().run()
        assert not app.exception
        assert next(item for item in app.text_input if item.label == "Equipment").value == "Pump P-001"
        assert next(item for item in app.text_area if item.label == "Recommendation").value == "Replace seal"
        next(item for item in app.button if item.label == "Close editing").click().run()
    assert app.session_state["block"].extra_tables == (expected,)
    assert app.session_state["original"] == original
    assert app.session_state["original"].fingerprint == original.fingerprint


@pytest.mark.parametrize("count", [12, 13])
def test_duplicate_heading_preview_preserves_existing_twelve_row_bound(count):
    table = replace(CASES[0], rows=tuple(
        (f"Pump P-{index:03}", "$1200", f"Recommendation {index}") for index in range(count)))
    app = _workspace(table)
    frame = app.dataframe[0].value
    assert frame.shape == (12, len(table.columns))
    assert tuple(tuple(row) for row in frame.itertuples(index=False, name=None)) == table.rows[:12]
    assert app.session_state["saved"] == table


@pytest.mark.parametrize("table, corrected, expected", [
    (CASES[0], ("Equipment", "Cost", "Recommendation"),
     ReportTable(("Equipment", "Recommendation"), (("Pump P-001", "Replace seal"),), CASES[0].reference)),
    (CASES[2], ("Equipment", "Replacement timing", "Cost", "Recommendation"),
     ReportTable(("Equipment", "Replacement timing", "Recommendation"),
                 (("Pump P-001", "2030", "Replace seal"),), CASES[2].reference)),
], ids=["duplicate", "blank_headings"])
def test_confirmed_names_preserve_technical_cells_history_and_actual_word(monkeypatch, tmp_path, table, corrected, expected):
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path))
    original_fingerprint = _fingerprint(table)
    block = ResolvedBlock("capital_renewal", "This month", extra_tables=(table,),
                          references=("synthetic-owned-source",))
    draft = synthetic_draft(synthetic_profiles()[0], ReportPeriod(2026, 9))
    section = next(item for item in default_sections() if item.key == "capital")
    original = replace(draft, sections=(section,), blocks=(block,))
    library.save_snapshot(original, expected_revision=0)
    app = _workspace(table)
    # Show the technical cell before asking the operator to identify prices.
    assert tuple(app.dataframe[0].value.iloc[0]) == table.rows[0]

    for control, heading in zip(app.text_input, corrected):
        control.set_value(heading)
    app.run()
    assert not app.exception and not app.button[0].disabled
    assert app.session_state["blocked"]
    assert app.session_state["saved"] == table
    assert _fingerprint(app.session_state["saved"]) == original_fingerprint
    assert tuple(app.dataframe[0].value.iloc[0]) == table.rows[0]

    app.button[0].click().run()
    assert not app.exception
    assert app.session_state["saved"] == expected
    assert app.session_state["original"] == table
    assert _fingerprint(app.session_state["original"]) == original_fingerprint
    assert not app.session_state["blocked"]
    app.run()
    assert not app.exception and not app.text_input
    assert app.session_state["saved"] == expected

    saved = replace(original, blocks=(replace(block, extra_tables=(expected,)),))
    library.save_snapshot(saved, expected_revision=1)
    reopened = library.load_snapshot(saved.profile.contract, saved.profile.key, saved.period)
    assert reopened.draft == saved
    history = library.load_snapshot_version(saved.profile.contract, saved.profile.key, saved.period, 1)
    assert history.draft == original
    assert _fingerprint(history.draft.blocks[0].extra_tables[0]) == original_fingerprint
    reentered = _workspace(reopened.draft.blocks[0].extra_tables[0])
    assert reentered.session_state["saved"] == expected
    assert not reentered.dataframe and not reentered.text_input

    word = assemble_docx(reopened.draft, acknowledged_fingerprint=reopened.draft.fingerprint)
    document = Document(BytesIO(word))
    native = next(item for item in document.tables
                  if tuple(cell.text for cell in item.rows[0].cells) == expected.columns)
    assert tuple(tuple(cell.text for cell in row.cells) for row in native.rows[1:]) == expected.rows
    all_text = "\n".join(cell.text for item in document.tables for row in item.rows for cell in row.cells)
    assert "$1200" not in all_text
    assert "Pump P-001" in all_text and "Replace seal" in all_text
