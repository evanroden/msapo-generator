"""Saved parenthesized price prose must not reach actual shell/native exports.

All fixtures are invented. Original snapshots and source designs stay unchanged;
an explicit human correction creates a separate saved revision for generation.
"""
from dataclasses import replace
from hashlib import sha256
from io import BytesIO
import json
import os
from pathlib import Path

from docx import Document
import pytest

from app import monthly_report_designs as designs, monthly_report_library as library
from app.monthly_report_checks import preflight
from app.monthly_report_docx import generate_report
from app.monthly_report_guided import client_table_draft
from tests.conftest import requires_libreoffice
from tests.test_monthly_price_approximate_outputs import (
    READINGS, SOURCE_CANARY, _draft as _approximate_draft, _texts, _write_package,
)


TECHNICAL_FACTS = (
    "Synthetic valve remains out of service.\n"
    "Pressure estimate (~125 psi). Motor estimate (~3.2 kW).\n"
    "Quote reference (Q-3200)."
)
PARENTHESIZED_READINGS = replace(
    READINGS,
    rows=(("PUMP-3200", "(~125 psi)", "2026-10-20"),
          ("MOTOR-3200", "(~3.2 kW)", "2026-10-21")),
)


@requires_libreoffice
@pytest.mark.parametrize("native", [False, True], ids=["shell", "native"])
@pytest.mark.parametrize("price", ["Repair estimate: (~3.2k)", "Quote (~3.2k)"], ids=["estimate", "quote"])
def test_saved_parenthesized_price_requires_correction_before_actual_outputs(tmp_path, monkeypatch, native, price):
    runtime = tmp_path / "runtime"
    monkeypatch.setenv("EPC_DATA_DIR", str(runtime))
    base = _approximate_draft(tmp_path, native, price)
    draft = replace(base, blocks=(
        replace(base.blocks[0], text=f"{TECHNICAL_FACTS}\n{price}"),
        replace(base.blocks[1], extra_tables=(PARENTHESIZED_READINGS,)),
    ))
    original = library.save_snapshot(draft, expected_revision=0, entered_editor=draft.prepared_by)
    restored = library.load_snapshot(draft.profile.contract, draft.profile.key, draft.period)
    assert restored == original
    original_paths = tuple(runtime.rglob(f"{draft.period.key}-00000001.json"))
    assert len(original_paths) == 1
    original_path = original_paths[0]
    original_bytes = original_path.read_bytes()
    source = designs.source_for(draft.profile)
    source_bytes = source.read_bytes() if source else None

    # Review must not silently remove price prose or technical parentheses.
    working, removed = client_table_draft(restored.draft)
    assert working == restored.draft and not removed
    assert working.blocks[1].extra_tables == (PARENTHESIZED_READINGS,)
    output = (Path(os.environ.get("INDEPENDENT_OUTPUT", str(tmp_path / "evidence")))
              / "price-parenthesized" / ("native" if native else "shell")
              / ("estimate" if price.startswith("Repair") else "quote"))
    output.mkdir(parents=True, exist_ok=True)
    (output / "SYNTHETIC-original-snapshot.json").write_bytes(original_bytes)
    manifest = {
        "synthetic_only": True,
        "tested_revision": os.environ.get("MSAPO_SHA", "unrecorded"),
        "native_design": native,
        "original_revision": original.revision,
        "original_snapshot_sha256": sha256(original_bytes).hexdigest(),
        "original_fingerprint": working.fingerprint,
        "files": {},
    }
    pricing_checks = [check for check in preflight(working) if check.code == "pricing" and check.blocking]
    try:
        unsafe_package = generate_report(working, acknowledged_fingerprint=working.fingerprint)
    except ValueError as exc:
        assert "pricing" in str(exc).lower(), str(exc)
        generation_blocked = True
    else:
        # Preserve real failing-before Word/PDF and every page before asserting.
        _write_package(output, "unsafe-before-correction", unsafe_package, manifest)
        unsafe_xml, unsafe_pdf = _texts(unsafe_package)
        assert price in unsafe_xml and price in unsafe_pdf
        generation_blocked = False
    manifest["uncorrected_generation_blocked"] = generation_blocked

    corrected = replace(working, blocks=(replace(working.blocks[0], text=TECHNICAL_FACTS), working.blocks[1]))
    saved_correction = library.save_snapshot(corrected, expected_revision=original.revision,
                                             entered_editor="Synthetic Correcting Editor")
    reloaded = library.load_snapshot(draft.profile.contract, draft.profile.key, draft.period)
    assert reloaded == saved_correction and reloaded.revision == 2
    assert reloaded.draft.blocks[0].text == TECHNICAL_FACTS
    assert reloaded.draft.blocks[1].extra_tables == (PARENTHESIZED_READINGS,)
    corrected_paths = tuple(runtime.rglob(f"{draft.period.key}-00000002.json"))
    assert len(corrected_paths) == 1
    corrected_bytes = corrected_paths[0].read_bytes()
    (output / "SYNTHETIC-corrected-snapshot.json").write_bytes(corrected_bytes)
    package = generate_report(reloaded.draft, acknowledged_fingerprint=reloaded.draft.fingerprint)
    _write_package(output, "corrected-reloaded", package, manifest)
    manifest["corrected_revision"] = reloaded.revision
    manifest["corrected_snapshot_sha256"] = sha256(corrected_bytes).hexdigest()
    manifest["corrected_fingerprint"] = reloaded.draft.fingerprint
    (output / "manifest.json").write_text(json.dumps(manifest, indent=2) + "\n")

    word = Document(BytesIO(package.docx))
    assert len(word.tables) == 1
    assert tuple(tuple(cell.text for cell in row.cells) for row in word.tables[0].rows) == (
        PARENTHESIZED_READINGS.columns, *PARENTHESIZED_READINGS.rows)
    assert all(line in [paragraph.text for paragraph in word.paragraphs] for line in TECHNICAL_FACTS.splitlines())
    xml, pdf_text = _texts(package)
    for text in (xml, pdf_text):
        for fact in ("Synthetic valve remains out of service.", "Pressure estimate (~125 psi).",
                     "Motor estimate (~3.2 kW).", "Quote reference (Q-3200).",
                     "PUMP-3200", "MOTOR-3200", "2026-10-20", "2026-10-21"):
            assert fact in text
        assert "(~3.2k)" not in text and SOURCE_CANARY not in text
    assert original_path.read_bytes() == original_bytes
    assert corrected_paths[0].read_bytes() == corrected_bytes
    assert library.load_snapshot_version(draft.profile.contract, draft.profile.key, draft.period, 1) == original
    assert library.load_snapshot(draft.profile.contract, draft.profile.key, draft.period) == reloaded
    if source is not None:
        assert source.read_bytes() == source_bytes
    assert generation_blocked and pricing_checks, "Saved parenthesized approximate price prose reached actual Word/PDF before correction."
