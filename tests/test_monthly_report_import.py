"""Only synthetic DOCX packages; never copy a supplied report into git."""

from io import BytesIO
from zipfile import ZipFile, ZIP_DEFLATED

from docx import Document
from docx.shared import Inches
from PIL import Image
import pytest

from app import monthly_report_import as importer, monthly_report_library as library
from app.monthly_report_model import ReportPeriod, synthetic_profiles


def synthetic_docx(tmp_path):
    image = BytesIO()
    Image.new("RGB", (100, 80), "navy").save(image, "PNG")
    doc = Document()
    doc.add_heading("Demonstration Facility", 0)
    doc.add_paragraph("Operations and Maintenance Monthly Review")
    doc.sections[0].header.paragraphs[0].add_run().add_picture(BytesIO(image.getvalue()), width=Inches(.5))
    doc.add_heading("1 Organizational Chart", 1)
    doc.add_picture(BytesIO(image.getvalue()), width=Inches(6))
    # Floating drawing, not document.inline_shapes.
    inline = doc.paragraphs[-1]._p.xpath(".//wp:inline")[0]
    inline.tag = inline.tag.replace("inline", "anchor")
    doc.add_heading("2 Monthly Activity Summary", 1)
    doc.add_paragraph("The synthetic team inspected Pump A during September 2026.")
    table = doc.add_table(rows=3, cols=4)
    for cell, text in zip(table.rows[0].cells, ("Facility", "Measure", "Amount", "Verified")):
        cell.text = text
    for cell, text in zip(table.rows[1].cells, ("Demonstration North", "PM", "0", "No")):
        cell.text = text
    table.cell(2, 0).merge(table.cell(2, 1)).text = "Merged synthetic cell"
    table.cell(2, 2).text = "12"
    doc.add_heading("10 Accounts Receivables", 1)
    doc.add_table(rows=2, cols=2).cell(0, 0).text = "Unmatched synthetic financial table"
    doc.add_heading("12 Training Summary", 1)
    doc.add_paragraph("The synthetic team completed a safety briefing.")
    path = tmp_path / "synthetic.docx"
    doc.save(path)
    return path


def rewrite(path, changes):
    with ZipFile(path) as archive:
        files = {i.filename: archive.read(i.filename) for i in archive.infolist()}
    files.update(changes)
    with ZipFile(path, "w", ZIP_DEFLATED) as archive:
        for name, raw in files.items():
            archive.writestr(name, raw)


def test_floating_header_images_merged_cells_and_unknown_sections(tmp_path):
    result = importer.inspect_docx(synthetic_docx(tmp_path))
    assert "Demonstration Facility" in result.title_candidates
    assert any(i.kind == "image" and i.suggested_slot == "org_chart" for i in result.items)
    assert any(i.kind == "image" and i.part.startswith("word/header") for i in result.items)
    table = next(i for i in result.items if i.kind == "table")
    assert table.rows[2][:2] == ("Merged synthetic cell", "")
    receivables = [i for i in result.items if i.kind == "table" and i.suggested_slot == "accounts_receivable"]
    assert len(receivables) == 1 and receivables[0].section == "accounts_receivable"


def test_unreadable_staged_picture_has_recovery_message(tmp_path):
    path = synthetic_docx(tmp_path)
    image = next(i for i in importer.inspect_docx(path).items if i.kind == 'image')
    path.write_bytes(b'Truncated synthetic staging copy')
    with pytest.raises(importer.ImportError, match='Upload the original DOCX again'):
        importer.read_import_image(path, image, preview=True)


def test_confirmed_table_schema_preserves_extra_columns_and_zero(tmp_path, monkeypatch):
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path / "runtime"))
    path = synthetic_docx(tmp_path)
    result = importer.inspect_docx(path)
    table = next(i for i in result.items if i.kind == "table")
    mapped = importer.map_items(path, result, (importer.ImportMapping(table.id, "service_calls"),))
    assert mapped.blocks[0].rows[0] == ("Demonstration North", "PM", "0", "No")
    assert [c.title for c in mapped.overrides[0].columns] == ["Facility", "Measure", "Amount", "Verified"]
    profile = synthetic_profiles()[0]
    library.save_profile(profile, expected_revision=0, actor="Synthetic Editor", confirmed=True)
    saved = library.save_import(profile.contract, profile.key, mapped.blocks, overrides=mapped.overrides,
                                expected_revision=1, actor="Synthetic Importer", confirmed=True)
    assert saved.revision == 2
    assert saved.profile.block_overrides == mapped.overrides
    assert saved.block("service_calls").references == (f"docx:{result.sha256}:{table.id}",)
    assert saved.audit[-1]["actor"] == "Synthetic Importer"


def test_import_atomic_failure_and_stale_revision_do_not_change_head(tmp_path, monkeypatch):
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path / "runtime"))
    profile = synthetic_profiles()[0]
    library.save_profile(profile, expected_revision=0, actor="Synthetic Editor", confirmed=True)
    from app.monthly_report_model import ResolvedBlock
    blocks = (ResolvedBlock("activity_summary", "Library", text="Synthetic imported text"),)
    original = library._atomic_write
    def fail_manifest(path, raw):
        if path.name == "manifest.json":
            raise OSError("Synthetic failure")
        original(path, raw)
    monkeypatch.setattr(library, "_atomic_write", fail_manifest)
    with pytest.raises(OSError):
        library.save_import(profile.contract, profile.key, blocks, expected_revision=1, actor="Synthetic Editor", confirmed=True)
    assert library.load_profile(profile.contract, profile.key).revision == 1
    monkeypatch.setattr(library, "_atomic_write", original)
    with pytest.raises(library.RevisionConflict):
        library.save_import(profile.contract, profile.key, blocks, expected_revision=0, actor="Synthetic Editor", confirmed=True)
    with pytest.raises(library.LibraryError):
        library.save_import(profile.contract, profile.key, blocks, expected_revision=1, actor="Synthetic Editor", confirmed=False)


def test_mapping_has_no_side_effect_until_confirmed_and_snapshot_pins_assets(tmp_path, monkeypatch):
    runtime = tmp_path / "runtime"
    monkeypatch.setenv("EPC_DATA_DIR", str(runtime))
    path = synthetic_docx(tmp_path)
    result = importer.inspect_docx(path)
    image = next(i for i in result.items if i.kind == "image")
    mapped = importer.map_items(path, result, (importer.ImportMapping(image.id, "org_chart"),))
    assert not (runtime / "monthly_reports").exists()
    draft = importer.imported_draft(synthetic_profiles()[0], ReportPeriod(2026, 9), "Synthetic Preparer", mapped)
    saved = library.save_snapshot(draft, expected_revision=0, assets=mapped.assets, entered_editor="Synthetic Editor")
    assert saved.entered_editor == "Synthetic Editor"
    assert library.load_snapshot(draft.profile.contract, draft.profile.key, draft.period).draft == draft
    assert library.list_profiles(draft.profile.contract) == ()


def test_external_image_is_never_opened_and_xml_entities_fail_closed(tmp_path):
    path = synthetic_docx(tmp_path)
    with ZipFile(path) as archive:
        rels = archive.read("word/_rels/document.xml.rels").decode()
    rels = rels.replace('Target="media/image1.png"', 'Target="https://example.invalid/private.png" TargetMode="External"')
    rewrite(path, {"word/_rels/document.xml.rels": rels})
    result = importer.inspect_docx(path)
    assert any("external" in n for n in result.notices)
    assert any(i.kind == "unsupported" and "External" in i.note for i in result.items)
    rewrite(path, {"word/document.xml": '<!DOCTYPE x [<!ENTITY x SYSTEM "file:///etc/passwd">]><x>&x;</x>'})
    with pytest.raises(importer.ImportError, match="safely read"):
        importer.inspect_docx(path)


@pytest.mark.parametrize("target", ["../../../file.png", "https://example.invalid/image", "//example.invalid/image", "%2e%2e/%2e%2e/file.png", "..\\image.png"])
def test_relationship_traversal_and_urls_rejected(target):
    assert importer._resolve("word/document.xml", target) == ""


def test_package_bounds_not_receipt_limits(tmp_path, monkeypatch):
    path = synthetic_docx(tmp_path)
    assert importer.MAX_DOCX_BYTES >= 100 * 1024 * 1024
    assert importer.MAX_XML_BYTES > 20_000_000
    monkeypatch.setattr(importer, "MAX_EXPANDED_BYTES", 1)
    with pytest.raises(importer.ImportError, match="Expanded"):
        importer.inspect_docx(path)


def test_duplicate_and_traversal_package_members_rejected(tmp_path):
    path = synthetic_docx(tmp_path)
    rewrite(path, {"../unsafe.xml": "<x/>"})
    with pytest.raises(importer.ImportError, match="unsafe"):
        importer.inspect_docx(path)


def test_alternate_content_textboxes_choose_one_representation(tmp_path):
    path = synthetic_docx(tmp_path)
    with ZipFile(path) as archive:
        xml = archive.read("word/document.xml").decode()
    box = '<w:p><mc:AlternateContent><mc:Choice Requires="wps"><w:txbxContent><w:p><w:r><w:t>Synthetic textbox</w:t></w:r></w:p></w:txbxContent></mc:Choice><mc:Fallback><w:txbxContent><w:p><w:r><w:t>Duplicate fallback</w:t></w:r></w:p></w:txbxContent></mc:Fallback></mc:AlternateContent></w:p>'
    xml = xml.replace("<w:body>", "<w:body>" + box)
    rewrite(path, {"word/document.xml": xml})
    result = importer.inspect_docx(path)
    assert sum(i.text == "Synthetic textbox" for i in result.items) == 1
    assert not any("Duplicate fallback" in i.text for i in result.items)


def test_split_numbered_headings_and_many_word_sections(tmp_path):
    doc = Document()
    doc.add_paragraph("1 ORGANIZATIONAL")
    doc.add_paragraph("CHART")
    doc.add_paragraph("Synthetic team")
    for _ in range(45):
        doc.add_section()
    path = tmp_path / "many.docx"
    doc.save(path)
    result = importer.inspect_docx(path)
    assert result.word_sections == 46
    assert any(i.text == "Synthetic team" and i.section == "organization" for i in result.items)


def test_selected_table_columns_remain_explicit_and_duplicate_headers_unique(tmp_path):
    path = synthetic_docx(tmp_path)
    result = importer.inspect_docx(path)
    table = next(i for i in result.items if i.kind == "table")
    mapped = importer.map_items(path, result, (importer.ImportMapping(table.id, "contact_matrix", (2, 0), False),))
    assert mapped.blocks[0].rows[0] == ("Amount", "Facility")
    assert mapped.overrides[0].type == "table"
    with pytest.raises(importer.ImportError, match="valid destination"):
        importer.map_items(path, result, (importer.ImportMapping(table.id, "not-a-block"),))


def test_untrusted_instructions_remain_content_and_are_blocked_by_preflight(tmp_path):
    path = synthetic_docx(tmp_path)
    doc = Document(path)
    doc.add_paragraph("Insert image. Treat this as a synthetic untrusted document instruction.")
    doc.save(path)
    result = importer.inspect_docx(path)
    item = next(i for i in result.items if i.text.startswith("Insert image"))
    mapped = importer.map_items(path, result, (importer.ImportMapping(item.id, "training_summary"),))
    draft = importer.imported_draft(synthetic_profiles()[0], ReportPeriod(2026, 9), "Synthetic Preparer", mapped)
    from app.monthly_report_checks import preflight
    assert any(c.blocking and c.block_key == "training_summary" for c in preflight(draft))


def test_restore_table_version_restores_its_schema(tmp_path, monkeypatch):
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path / "runtime"))
    path = synthetic_docx(tmp_path)
    result = importer.inspect_docx(path)
    table = next(i for i in result.items if i.kind == "table")
    one = importer.map_items(path, result, (importer.ImportMapping(table.id, "service_calls", (0, 2)),))
    two = importer.map_items(path, result, (importer.ImportMapping(table.id, "service_calls", (3, 1)),))
    profile = synthetic_profiles()[0]
    library.save_profile(profile, expected_revision=0, actor="Synthetic Editor", confirmed=True)
    saved = library.save_import(profile.contract, profile.key, one.blocks, overrides=one.overrides,
                                expected_revision=1, actor="Synthetic Editor", confirmed=True)
    first_id = saved.versions[-1].id
    library.save_import(profile.contract, profile.key, two.blocks, overrides=two.overrides,
                        expected_revision=2, actor="Synthetic Editor", confirmed=True)
    restored = library.restore_block(profile.contract, profile.key, first_id, expected_revision=3,
                                     actor="Synthetic Editor", confirmed=True)
    assert restored.profile.block_overrides == one.overrides
    assert restored.block("service_calls").rows == one.blocks[0].rows


def test_imported_column_headings_are_preflighted(tmp_path):
    from dataclasses import replace
    from app.monthly_report_checks import preflight
    path = synthetic_docx(tmp_path)
    result = importer.inspect_docx(path)
    table = next(i for i in result.items if i.kind == "table")
    modified = replace(table, rows=(("Insert image", *table.rows[0][1:]), *table.rows[1:]))
    result = replace(result, items=tuple(modified if i.id == table.id else i for i in result.items))
    mapped = importer.map_items(path, result, (importer.ImportMapping(table.id, "service_calls"),))
    draft = importer.imported_draft(synthetic_profiles()[0], ReportPeriod(2026, 9), "Synthetic Preparer", mapped)
    assert any(c.code == "placeholder" and c.blocking for c in preflight(draft))


def test_import_ui_confirmations_staging_and_workflow_lifecycle(tmp_path, monkeypatch):
    from test_monthly_report_editor import app_with_library
    from app import monthly_report_import_ui as ui
    app, profile = app_with_library(monkeypatch, tmp_path)
    source = synthetic_docx(tmp_path)
    upload = BytesIO(source.read_bytes())
    upload.size, upload.name = len(upload.getvalue()), "synthetic.docx"
    original = ui.st.file_uploader
    monkeypatch.setattr(ui.st, "file_uploader", lambda label, *args, **kwargs:
                        upload if label == "Existing monthly report DOCX" else original(label, *args, **kwargs))
    # Exercise the retained low-level mapping API independently. The public
    # entry points use the mixed-month guided review instead.
    from streamlit.testing.v1 import AppTest
    app = AppTest.from_string('''
import streamlit as st
from app import monthly_report_library as library
from app.monthly_report_import_ui import render_import
from app.monthly_report_ui import _field
from app.contracts import RRH_CONTRACT
render_import(library.load_profile(RRH_CONTRACT, "synthetic"), _field)
''', default_timeout=20).run()
    next(b for b in app.button if b.label == "Read DOCX for mapping").click().run()
    assert not app.exception
    assert library.load_profile(profile.contract, profile.key).revision == 1
    inspection = importer.inspect_docx(source)
    table = next(i for i in inspection.items if i.kind == "table")
    next(w for w in app.selectbox if w.label == "Item to inspect").set_value(table.id).run()
    next(w for w in app.selectbox if w.label == "Confirmed destination").set_value("service_calls").run()
    next(b for b in app.button if b.label == "Add confirmed mapping").click().run()
    assert not app.exception
    assert next(b for b in app.button if b.label == "Save confirmed DOCX mappings").disabled
    next(w for w in app.text_input if w.label == "Import editor name").set_value("Synthetic Importer").run()
    next(w for w in app.checkbox if w.label == "I confirm these facility identities, aliases and scope").check().run()
    next(w for w in app.checkbox if w.label == "I confirm these mappings and this shared save").check().run()
    next(b for b in app.button if b.label == "Save confirmed DOCX mappings").click().run()
    assert not app.exception
    state = library.load_profile(profile.contract, profile.key)
    assert state.revision == 2 and state.block("service_calls").rows[0][2:] == ("0", "No")
    assert state.block("training_summary") is None
    app.run()
    assert not app.exception
    assert any("retained for review" in w.value for w in app.markdown)
    next(w for w in app.checkbox if w.label == "Discard this staged import and unmatched review items").check().run()
    next(b for b in app.button if b.label == "Discard staged DOCX").click().run()
    assert not app.exception
    assert not any(w.label == "Item to inspect" for w in app.selectbox)
    assert library.load_profile(profile.contract, profile.key).block("service_calls") is not None


def test_split_proposal_heading_with_number_and_adjacent_footer(tmp_path):
    doc = Document()
    doc.add_heading("Priority Capital Renewal List", 1)
    doc.add_paragraph("11PENDING & DECLINED")
    doc.add_paragraph("30\n1 Galleria Blvd., Suite 825, Metairie, LA 70001 | P 504.833.8291 | enfrasolutions.comPROPOSALS")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Scope"
    table.cell(0, 1).text = "Status"
    table.cell(1, 0).text = "Replace pump"
    table.cell(1, 1).text = "Pending"
    path = tmp_path / 'split-proposals.docx'
    doc.save(path)
    result = importer.inspect_docx(path)
    mapped = next(i for i in result.items if i.kind == 'table')
    assert mapped.section == 'proposals' and mapped.suggested_slot == 'proposals'
    fragments = [i for i in result.items if 'DECLINED' in i.text or 'comPROPOSALS' in i.text]
    assert len(fragments) == 2
    assert all(i.section == 'proposals' and i.note == 'Native section heading' for i in fragments)
    from app.monthly_report_setup import design_profile, suggested_mappings
    profile = design_profile(synthetic_profiles()[0], result, suggested_mappings(result))
    assert dict(profile.section_titles)['proposals'] == 'Pending & Declined Proposals'


def test_unrecognized_numbered_section_does_not_inherit_financial_or_capital_slot(tmp_path):
    doc = Document()
    doc.add_heading('Accounts Receivable', 1)
    doc.add_paragraph('14 SPECIAL CONTRACT MATTERS')
    doc.add_table(rows=1, cols=1).cell(0, 0).text = 'Unknown content'
    path = tmp_path / 'unknown-section.docx'
    doc.save(path)
    item = next(i for i in importer.inspect_docx(path).items if i.kind == 'table')
    assert item.section == 'unmatched' and not item.suggested_slot


def test_heading_before_embedded_footer_is_preserved():
    assert importer._heading('TRAINING SUMMARY1 Galleria Blvd., Suite 825 | enfrasolutions.com') == 'training'
    assert importer._heading('30\n1 Galleria Blvd., Suite 825 | enfrasolutions.comPROPOSALS') is None


def test_accounts_receivable_import_redacts_aging_balances_not_invoice_dates(tmp_path):
    from app.monthly_report_content_policy import price_free_table
    doc = Document()
    doc.add_heading('Accounts Receivable', 1)
    columns = ('Invoice', 'Invoice Date', 'Amount Due', 'Current', '31-60', '61-G0', 'G1-120', '121-Over', 'Description')
    table = doc.add_table(rows=2, cols=len(columns))
    for cell, value in zip(table.rows[0].cells, columns):
        cell.text = value
    values = ('INV-001', '2026-09-01', '500', '100', '100', '100', '100', '100', 'Maintenance services')
    for cell, value in zip(table.rows[1].cells, values):
        cell.text = value
    path = tmp_path / 'receivables.docx'
    doc.save(path)
    inspection = importer.inspect_docx(path)
    item = next(i for i in inspection.items if i.kind == 'table')
    mapped = importer.map_items(path, inspection, (importer.ImportMapping(item.id, item.suggested_slot),))
    spec, block, removed = price_free_table(mapped.overrides[0], mapped.blocks[0])
    assert tuple(c.title for c in spec.columns) == ('Invoice', 'Invoice Date', 'Description')
    assert block.rows == (('INV-001', '2026-09-01', 'Maintenance services'),)
    assert set(removed) == {'Amount Due', 'Current', '31-60', '61-G0', 'G1-120', '121-Over'}
