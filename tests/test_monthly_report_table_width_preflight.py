"""Preflight must reject the same table geometry that Word generation rejects.

Synthetic only. Tests ensure actionable 1-based row positions and that
original/unknown cells survive correction, omission, snapshot and output.
"""
from dataclasses import replace
from io import BytesIO

from docx import Document
import pytest

from app import monthly_report_library as library
from app.monthly_report_checks import preflight
from app.monthly_report_docx import assemble_docx
from app.monthly_report_model import (
    Facility, ReportDraft, ReportPeriod, ReportProfile, ReportTable,
    ResolvedBlock, default_sections,
)


def report(block, *, included=True, period=None):
    section_key = {
        "service_calls": "maintenance",
        "subcontractor_matrix": "subcontractors",
        "capital_renewal": "capital",
    }[block.key]
    section = next(section for section in default_sections()
                   if section.key == section_key)
    profile = ReportProfile(
        "Synthetic Width Contract", "north", "Synthetic North",
        (Facility("north", "Synthetic North"),))
    return ReportDraft(
        profile, period or ReportPeriod(2026, 9), "Synthetic reviewer",
        (replace(section, included=included),), (block,), synthetic=True)


def failures(draft):
    return tuple(check for check in preflight(draft)
                 if check.code == "table_columns")


def service_row(*, extra=False):
    row = ("000012", "2026-09-14", "Boiler room", "Replace seal", "PMP-1")
    return (*row, "Unmapped source detail") if extra else row


def test_width_mismatch_is_blocked_before_word_generation_and_preserved():
    block = ResolvedBlock("service_calls", "This month", rows=(service_row(extra=True),))
    draft = report(block)
    found = failures(draft)
    assert len(found) == 1
    issue, = found
    assert issue.blocking is True
    assert issue.block_key == "service_calls"
    assert "Row 1" in issue.message
    assert "6 cells" in issue.message
    assert "5 columns" in issue.message
    assert "service calls" in issue.message
    with pytest.raises(ValueError, match="Row 1.*6 cells"):
        assemble_docx(draft, acknowledged_fingerprint=draft.fingerprint)
    assert draft.blocks[0] == block
    assert draft.blocks[0].rows[0][-1] == "Unmapped source detail"


def test_second_source_row_is_named_and_existing_safe_first_row_remains():
    block = ResolvedBlock(
        "service_calls", "This month",
        rows=(service_row(), service_row(extra=True)))
    found = failures(report(block))
    assert len(found) == 1
    assert "Row 2" in found[0].message
    assert "6 cells" in found[0].message
    assert len(block.rows) == 2
    assert block.rows[0] == service_row()


def test_valid_full_width_and_short_rows_still_render_in_one_generic_word_table():
    block = ResolvedBlock("service_calls", "This month",
                          rows=(service_row(), ("000013", "2026-09-15", "Plant")))
    draft = report(block)
    assert failures(draft) == ()
    output = Document(BytesIO(assemble_docx(
        draft, acknowledged_fingerprint=draft.fingerprint)))
    assert len(output.tables) == 1
    table, = output.tables
    assert len(table.columns) == 5
    assert table.rows[1].cells[0].text == "000012"
    assert table.rows[1].cells[4].text == "PMP-1"
    assert table.rows[2].cells[0].text == "000013"
    assert table.rows[2].cells[3].text == ""


def test_imported_table_with_more_cells_than_headings_is_blocked_in_review():
    imported = ReportTable(
        ("Facility", "Item"),
        (("Synthetic North", "Valve inspection", "Original extra cell"),),
        "synthetic-file:2")
    block = ResolvedBlock("subcontractor_matrix", "This month",
                          extra_tables=(imported,), references=("synthetic-file:2",))
    draft = report(block)
    found = failures(draft)
    assert len(found) == 1
    assert found[0].block_key == "subcontractor_matrix"
    assert "Imported table 1, row 1" in found[0].message
    assert "3 cells" in found[0].message
    assert "2 columns" in found[0].message
    with pytest.raises(ValueError, match="Imported table 1"):
        assemble_docx(draft, acknowledged_fingerprint=draft.fingerprint)
    assert block.extra_tables[0] == imported
    assert imported.rows[0][2] == "Original extra cell"


def test_multiple_imported_tables_identify_the_exact_unmapped_one():
    safe = ReportTable(("Facility", "Status"),
                       (("Synthetic North", "Verified"),), "original:1")
    wide = ReportTable(("Name", "Status"),
                       (("Synthetic North", "Open", "Third cell"),), "original:2")
    block = ResolvedBlock(
        "capital_renewal", "This month",
        extra_tables=(safe, wide),
        references=("original:1", "original:2"),
    )
    draft = report(block)
    found = failures(draft)
    assert len(found) == 1
    assert "Imported table 2, row 1" in found[0].message
    assert found[0].block_key == "capital_renewal"
    assert draft.blocks[0].extra_tables == (safe, wide)


def test_short_imported_rows_are_allowed_and_keep_blanks_on_output():
    imported = ReportTable(("Facility", "Trade", "Vendor"),
                           (("Synthetic North",), ("Synthetic North", "HVAC", "Vendor A")),
                           "synthetic:1")
    block = ResolvedBlock("subcontractor_matrix", "This month",
                          extra_tables=(imported,))
    draft = report(block)
    assert failures(draft) == ()
    result = Document(BytesIO(assemble_docx(
        draft, acknowledged_fingerprint=draft.fingerprint)))
    assert len(result.tables) == 1
    data = result.tables[0]
    assert data.rows[1].cells[1].text == ""
    assert data.rows[2].cells[2].text == "Vendor A"


def test_corrected_imported_heading_preserves_extra_value_through_word_output():
    imported = ReportTable(
        ("Facility", "Trade"),
        (("Synthetic North", "Mechanical", "Source tag: P-1"),),
        "synthetic-source:4",
    )
    before = ResolvedBlock("subcontractor_matrix", "This month",
                           extra_tables=(imported,))
    assert failures(report(before))
    resolved = replace(before, extra_tables=(replace(imported,
                        columns=("Facility", "Trade", "Source tag")),))
    draft = report(resolved)
    assert failures(draft) == ()
    result = Document(BytesIO(assemble_docx(
        draft, acknowledged_fingerprint=draft.fingerprint)))
    assert len(result.tables) == 1
    assert result.tables[0].rows[1].cells[2].text == "Source tag: P-1"
    assert imported.columns == ("Facility", "Trade")
    assert imported.rows == resolved.extra_tables[0].rows


def test_omitted_block_and_excluded_section_are_not_falsely_blocked():
    block = ResolvedBlock(
        "service_calls", "This month", rows=(service_row(extra=True),))
    omitted = report(replace(block, source="Omit"))
    excluded = report(block, included=False)
    assert failures(omitted) == ()
    assert failures(excluded) == ()
    assert omitted.blocks[0].rows == excluded.blocks[0].rows == block.rows
    assert failures(report(block))  # The source remains when re-included.


def test_headerless_data_is_blocked_but_empty_unprinted_tables_are_not():
    unprinted = ReportTable((), (("Unexpected cell",),), "unmapped-source")
    empty = ReportTable(("Facility", "Vendor"), (("", "", ""),), "empty-source")
    block = ResolvedBlock(
        "subcontractor_matrix", "This month",
        extra_tables=(unprinted, empty),
    )
    # The generic writer skips both, but the first contains genuine data;
    # Review must expose that omission rather than silently marking it ready.
    missing = tuple(c for c in preflight(report(block))
                    if c.code == "table_missing_columns")
    assert len(missing) == 1 and missing[0].blocking
    assert "Imported table 1" in missing[0].message
    assert missing[0].block_key == "subcontractor_matrix"
    assert failures(report(block)) == ()
    assert block.extra_tables == (unprinted, empty)
    assert not any(c.code == "table_missing_columns" for c in preflight(
        report(replace(block, extra_tables=(empty,)))))


def test_pricing_gate_still_blocks_financial_columns_independently():
    imported = ReportTable(
        ("Facility", "Amount Due"),
        (("Synthetic North", "150.00", "extra source value"),),
        "financial-source",
    )
    draft = report(ResolvedBlock(
        "subcontractor_matrix", "This month", extra_tables=(imported,)))
    checks = preflight(draft)
    assert any(check.code == "table_columns" and check.blocking for check in checks)
    assert any(check.code == "pricing" and check.blocking for check in checks)


def test_snapshot_reopen_preserves_unmapped_cells_and_no_implicit_repair(monkeypatch, tmp_path):
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path))
    imported = ReportTable(
        ("Facility", "Status"),
        (("Synthetic North", "Open", "Unmapped original note"),),
        "synthetic:source:1")
    block = ResolvedBlock("subcontractor_matrix", "This month",
                          extra_tables=(imported,), text="Independent note")
    draft = report(block)
    saved = library.save_snapshot(draft, expected_revision=0)
    loaded = library.load_snapshot(
        draft.profile.contract, draft.profile.key, draft.period)
    assert loaded.revision == 1
    assert loaded.draft.blocks[0] == block
    assert len(failures(loaded.draft)) == 1
    assert loaded.draft.blocks[0].extra_tables[0].rows[0][-1] == "Unmapped original note"
    assert loaded.draft.blocks[0].text == "Independent note"
    assert saved.draft.blocks == loaded.draft.blocks



def test_unheaded_source_content_is_an_actionable_blocker_instead_of_silent_loss():
    original = ReportTable(
        (), (("00123", "Pending renewal", "October context"),),
        "synthetic-immutable-docx:9",
    )
    block = ResolvedBlock("capital_renewal", "This month",
                          text="Keep independently authored note.",
                          extra_tables=(original,))
    draft = report(block)
    warnings = tuple(c for c in preflight(draft)
                     if c.code == "table_missing_columns")
    assert len(warnings) == 1
    assert warnings[0].blocking
    assert warnings[0].block_key == "capital_renewal"
    assert "Imported table 1" in warnings[0].message
    assert "no column headings" in warnings[0].message
    assert "Add headings" in warnings[0].message
    assert failures(draft) == ()  # Not a *width* mismatch.
    with pytest.raises(ValueError, match="no column headings"):
        assemble_docx(draft, acknowledged_fingerprint=draft.fingerprint)
    assert block.extra_tables == (original,)
    assert block.text == "Keep independently authored note."
    assert original.reference == "synthetic-immutable-docx:9"


def test_failing_headerless_second_table_retains_a_good_first_table():
    good = ReportTable(
        ("Asset", "Status"), (("PUMP-01", "Open"),), "original:1")
    missing = ReportTable(
        (), (("BOILER-02", "Needs replacement"),), "original:2")
    block = ResolvedBlock(
        "capital_renewal", "This month",
        extra_tables=(good, missing),
    )
    draft = report(block)
    errors = tuple(c for c in preflight(draft)
                   if c.code == "table_missing_columns")
    assert len(errors) == 1
    assert "Imported table 2" in errors[0].message
    with pytest.raises(ValueError, match="Imported table 2"):
        assemble_docx(draft, acknowledged_fingerprint=draft.fingerprint)
    assert block.extra_tables == (good, missing)
    assert good.rows == (("PUMP-01", "Open"),)


def test_adding_reviewed_headings_prints_full_source_row_and_keeps_good_sibling():
    good = ReportTable(
        ("Asset", "Status"), (("PUMP-01", "Open"),), "original:1")
    raw = ReportTable(
        (), (("BOILER-02", "Needs replacement"),), "original:2")
    old = ResolvedBlock("capital_renewal", "This month",
                        text="Independently written capital context.",
                        extra_tables=(good, raw))
    assert any(c.code == "table_missing_columns" for c in preflight(report(old)))
    reviewed = replace(raw, columns=("Asset", "Recommendation"))
    new = replace(old, extra_tables=(good, reviewed))
    draft = report(new)
    assert not any(c.code in ("table_missing_columns", "table_columns")
                   for c in preflight(draft))
    result = Document(BytesIO(assemble_docx(
        draft, acknowledged_fingerprint=draft.fingerprint)))
    assert len(result.tables) == 2
    assert tuple(c.text for c in result.tables[0].rows[1].cells) == (
        "PUMP-01", "Open")
    assert tuple(c.text for c in result.tables[1].rows[0].cells) == (
        "Asset", "Recommendation")
    assert tuple(c.text for c in result.tables[1].rows[1].cells) == (
        "BOILER-02", "Needs replacement")
    assert "Independently written capital context." in "\n".join(
        paragraph.text for paragraph in result.paragraphs)
    assert reviewed.reference == raw.reference
    assert raw.columns == ()
    assert raw.rows == reviewed.rows


def test_all_empty_or_whitespace_only_headerless_rows_remain_unprinted_and_allowed():
    empty = ReportTable((), (), "empty:1")
    whitespace = ReportTable((), ((" ", "\t"), ("", "")), "empty:2")
    block = ResolvedBlock("subcontractor_matrix", "This month",
                          extra_tables=(empty, whitespace))
    draft = report(block)
    assert not any(c.code == "table_missing_columns" for c in preflight(draft))
    assert block.extra_tables == (empty, whitespace)
    # The writer intentionally skips these blank/no-header tables.
    result = Document(BytesIO(assemble_docx(
        draft, acknowledged_fingerprint=draft.fingerprint)))
    assert len(result.tables) == 0


def test_omit_or_excluded_section_keeps_headerless_source_in_draft_without_review_block():
    missing = ReportTable((), (("Water vendor", "Category: chemical"),),
                          "synthetic:3")
    block = ResolvedBlock("subcontractor_matrix", "This month",
                          extra_tables=(missing,), references=("synthetic:3",))
    included = report(block)
    assert any(c.code == "table_missing_columns" for c in preflight(included))
    omitted = report(replace(block, source="Omit"))
    excluded = report(block, included=False)
    for item in (omitted, excluded):
        assert not any(c.code == "table_missing_columns" for c in preflight(item))
        assert item.blocks[0].extra_tables[0].rows == missing.rows
        assert item.blocks[0].references == ("synthetic:3",)
    assert omitted.blocks[0].source == "Omit"


def test_persisted_headerless_source_keeps_all_values_and_requires_review_after_reopen(
        tmp_path, monkeypatch):
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path))
    raw = ReportTable(
        (), (("North", "Boiler check"), ("South", "Pump service")),
        "source:original:2",
    )
    block = ResolvedBlock("subcontractor_matrix", "This month",
                          text="Report-only instructions remain.",
                          extra_tables=(raw,), references=("source:original:2",))
    draft = report(block)
    assert any(c.code == "table_missing_columns" for c in preflight(draft))
    library.save_snapshot(draft, expected_revision=0)
    stored = library.load_snapshot(draft.profile.contract, draft.profile.key,
                                   draft.period)
    assert stored.revision == 1
    assert stored.draft.blocks == (block,)
    assert stored.draft.blocks[0].text == "Report-only instructions remain."
    assert stored.draft.blocks[0].extra_tables[0].reference == raw.reference
    assert stored.draft.blocks[0].extra_tables[0].rows == raw.rows
    assert any(c.code == "table_missing_columns" for c in preflight(stored.draft))
    with pytest.raises(ValueError, match="no column headings"):
        assemble_docx(stored.draft, acknowledged_fingerprint=stored.draft.fingerprint)


def test_headerless_data_still_runs_independent_financial_preflight():
    # Missing headings is not permission to bypass privacy gates.
    imported = ReportTable((), (("Invoice amount: $1000",),), "source:4")
    draft = report(ResolvedBlock(
        "subcontractor_matrix", "This month",
        extra_tables=(imported,)))
    codes = {c.code for c in preflight(draft) if c.blocking}
    assert "table_missing_columns" in codes
    assert "pricing" in codes
