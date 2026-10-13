"""Shared route UI checks use invented native DOCX bytes and isolated storage."""
from dataclasses import replace
from io import BytesIO
from pathlib import Path

from docx import Document
import pytest
import requests
from streamlit.testing.v1 import AppTest

from app import memory
from app import monthly_report_fixture_runs as runs
from app import monthly_report_library as report_library
from app import monthly_report_shared_testing_ui as ui
from app import monthly_report_testing as testing
from app.monthly_report_fixture_vault import ExpectedSection, FixtureMetadata, ReportIdentity
from app.monthly_report_shared_fixtures import SharedFixtureLibrary
from app.monthly_report_model import ReportPeriod


TOKEN_A = "c" * 32
TOKEN_B = "d" * 32


@pytest.fixture
def sample(tmp_path, monkeypatch):
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path))
    document = Document()
    document.add_heading("Synthetic Site A", 0)
    document.add_paragraph("October 2026 Monthly Report")
    document.add_heading("2 Monthly Activity Summary", 1)
    document.add_paragraph("Inspected the synthetic pump in October 2026; follow-up is open.")
    output = BytesIO()
    document.save(output)
    original = output.getvalue()
    identity = ReportIdentity("Synthetic Testing Contract", ("synthetic-site-a",), "2026-10", ("Synthetic Site A",))
    metadata = FixtureMetadata(
        "Synthetic partial fixture", identity, identity, "AUTH03",
        "Keep current partial wording; do not invent completion.",
        (ExpectedSection("Activity", includes=("synthetic pump",), excludes=("invented completion",)),),
        "2026-10-13T16:00:00+00:00", True,
    )
    library = SharedFixtureLibrary(tmp_path)
    receipt = library.upload(original, metadata)
    return tmp_path, library, receipt, original


def _script(root, token):
    return f'''
from pathlib import Path
import streamlit as st
from app import memory
from app import monthly_report_testing as testing
from app import monthly_report_shared_testing_ui as ui
def field(key, default):
    st.session_state.setdefault(key, default)
    return key
with testing.isolated({token!r}):
    memory._data_dir()
    ui.render({token!r}, "UTC", field, lambda *_: None, Path({str(root)!r}))
'''


def test_delivered_callbacks_keep_first_pending_action_and_capture_reset_generation(monkeypatch):
    state = {}
    monkeypatch.setattr(ui.st, "session_state", state)
    ui._queue_upload()
    first = dict(state["_shared_fixture_pending"])
    ui._queue_upload()
    ui._queue_run("f" * 32, "1" * 32)
    assert state["_shared_fixture_pending"] == first
    state.clear()
    ui._queue_run("f" * 32, "1" * 32)
    request = dict(state["_shared_fixture_pending"])
    ui._queue_run(None, "2" * 32)
    assert state["_shared_fixture_pending"] == request
    assert request["fixture_id"] == "f" * 32 and request["generation"] == "1" * 32
    assert request["phase"] == "reset" and request["action"] == "start"


def test_expectations_remain_independent_explicit_records():
    result = ui._expectations("Activity | include | Original wording\nActivity | exclude | Fabricated work\nTraining | include | NR")
    assert result == (ExpectedSection("Activity", ("Original wording",), ("Fabricated work",)),
                      ExpectedSection("Training", ("NR",), ()))
    for malformed in ("", "Activity | maybe | wording", "Activity | include |", "include wording"):
        with pytest.raises(ValueError):
            ui._expectations(malformed)


def test_upload_metadata_stores_named_keys_titles_and_explicit_partial_confirmation(sample, monkeypatch):
    root, library, receipt, _original = sample
    state = {"_shared_fixture_upload_" + key: value for key, value in {
        "name": "Confirmed synthetic fixture", "contract": receipt.metadata.target.contract,
        "source_sites": "Synthetic Site A; Synthetic Site B", "target_sites": "Synthetic Site A; Synthetic Site B",
        "source_month": "2026-10", "target_month": "2026-10", "scenario": "AUTH03", "partial": True,
        "notes": "Preserve the incomplete work", "expectations": "Activity | exclude | Invented completion",
    }.items()}
    monkeypatch.setattr(ui.st, "session_state", state)
    with testing.isolated(TOKEN_A):
        memory._data_dir()
        metadata = ui._upload_metadata()
        assert metadata.target.sites == ("synthetic-site-a", "synthetic-site-b")
        assert metadata.target.site_titles == ("Synthetic Site A", "Synthetic Site B")
        profile = ui.target_profile(metadata)
        assert tuple(site.key for site in profile.facilities) == metadata.target.sites
        assert profile.scope_type == "multi_site"
        assert metadata.partial_file_confirmed is True


def test_upload_pending_action_commits_shared_original_without_preparing_a_report(sample, monkeypatch):
    root, library, receipt, original = sample

    class Upload:
        size = len(original)

        def getvalue(self):
            return original

    state = {"_shared_fixture_upload_file": Upload()}
    monkeypatch.setattr(ui.st, "session_state", state)
    monkeypatch.setattr(ui, "_upload_metadata", lambda: replace(receipt.metadata, friendly_name="Coworker upload"))
    ui._queue_upload()
    with testing.isolated(TOKEN_A):
        isolated = memory._data_dir()
        ui._process_pending(library, TOKEN_A, root)
        assert runs.read_run_fixture(library, TOKEN_A) is None
        assert {path.name for path in isolated.iterdir()} == {testing._SENTINEL}
    assert "_shared_fixture_pending" not in state
    uploaded = state["_shared_fixture_upload_receipt"]
    assert library.download(uploaded) == original
    assert library.read(uploaded).receipt.metadata.friendly_name == "Coworker upload"


def test_old_reset_request_cannot_delete_a_new_generation(sample, monkeypatch):
    root, library, _receipt, _original = sample
    state = {"_shared_fixture_pending": {"action": "reset", "phase": "reset", "generation": "0" * 32}}
    monkeypatch.setattr(ui.st, "session_state", state)
    with testing.isolated(TOKEN_A):
        isolated = memory._data_dir()
        canary = isolated / "current-work.txt"
        canary.write_text("Synthetic current work")
        ui._process_pending(library, TOKEN_A, root)
        assert canary.read_text() == "Synthetic current work"
    assert "_shared_fixture_pending" not in state
    assert "changed" in state["_shared_fixture_error"]


def test_actual_starting_ui_two_sessions_refresh_and_scoped_reset(sample, monkeypatch):
    root, library, receipt, original = sample

    def no_external(*_args, **_kwargs):
        pytest.fail("Shared native fixture testing must not make external reader calls.")

    monkeypatch.setattr(requests.sessions.Session, "request", no_external)
    app = AppTest.from_string(_script(root, TOKEN_A), default_timeout=30).run()
    assert not app.exception
    assert any("without a login" in notice.value for notice in app.warning)
    app.button(key="_shared_fixture_start").click().run()
    assert not app.exception
    staged = app.session_state["_shared_fixture_staged"]
    assert staged.source_path.read_bytes() == original
    assert any(item.label == "Your name" for item in app.text_input)
    assert any("Sites selected (1): Synthetic Site A" in caption.value for caption in app.caption)
    assert not list(staged.context.data_root.rglob("manifest.json"))  # No implicit import approval.
    assert "synthetic pump" in " ".join(item.text for item in staged.inspection.items)

    refreshed = AppTest.from_string(_script(root, TOKEN_A), default_timeout=30).run()
    assert not refreshed.exception
    assert refreshed.session_state["_shared_fixture_staged"].source_path == staged.source_path
    assert any(item.label == "Your name" for item in refreshed.text_input)

    coworker = AppTest.from_string(_script(root, TOKEN_B), default_timeout=30).run()
    assert not coworker.exception
    coworker.button(key="_shared_fixture_start").click().run()
    assert not coworker.exception
    other = coworker.session_state["_shared_fixture_staged"]
    assert other.context.scope != staged.context.scope
    assert other.source_path != staged.source_path
    assert other.source_path.read_bytes() == original
    refreshed.button(key="_shared_fixture_reset").click().run()
    assert not refreshed.exception
    assert not staged.source_path.exists()
    assert other.source_path.read_bytes() == original
    assert library.download(receipt.fixture_id) == original


def test_hidden_ui_rejects_unbound_or_wrong_run_token_before_shared_access(sample):
    root, _library, _receipt, _original = sample
    with pytest.raises(ValueError, match="bind"):
        ui.render(TOKEN_A, "UTC", lambda key, value: key, lambda *_: None, root)
    with testing.isolated(TOKEN_A):
        memory._data_dir()
        with pytest.raises(ValueError, match="token"):
            ui.render(TOKEN_B, "UTC", lambda key, value: key, lambda *_: None, root)


def test_production_continue_import_and_manual_edits_survive_refresh_without_reimport(sample, monkeypatch):
    root, library, receipt, original = sample

    def no_external(*_args, **_kwargs):
        pytest.fail("Native Starting import and draft recovery must not call external readers.")

    monkeypatch.setattr(requests.sessions.Session, "request", no_external)
    app = AppTest.from_string(_script(root, TOKEN_A), default_timeout=30).run()
    app.button(key="_shared_fixture_start").click().run()
    staged = app.session_state["_shared_fixture_staged"]
    next(item for item in app.text_input if item.label == "Your name").input("Synthetic Reviewer").run()
    next(item for item in app.button if item.label == "Continue to this month’s updates").click().run()
    assert not app.exception and not app.error
    assert app.session_state["_shared_fixture_initial_review_finished"] is True
    assert staged.setup_stage_key not in app.session_state.filtered_state
    activity = next(item for item in app.text_area if item.label == "Activity summary")
    assert "synthetic pump" in activity.value
    activity.input("My manually corrected synthetic Activity; follow-up remains open.").run()
    assert not app.exception and not app.error
    with testing.isolated(TOKEN_A):
        memory._data_dir()
        saved = report_library.load_profile(staged.profile.contract, staged.profile.key)
        imported = report_library.load_imported_draft(staged.profile.contract, staged.profile.key)
        assert saved.profile.key == staged.profile.key
        assert imported.period == ReportPeriod(2026, 10)
        assert report_library.imported_original(staged.profile.contract, staged.profile.key).read_bytes() == original
    refreshed = AppTest.from_string(_script(root, TOKEN_A), default_timeout=30).run()
    assert not refreshed.exception and not refreshed.error
    assert next(item for item in refreshed.text_area if item.label == "Activity summary").value == (
        "My manually corrected synthetic Activity; follow-up remains open.")
    assert not any(item.label == "Your name" for item in refreshed.text_input)
    assert staged.setup_stage_key not in refreshed.session_state.filtered_state
    assert library.download(receipt.fixture_id) == original


def test_second_reset_invalidates_pending_start_exact_generation(sample, monkeypatch):
    root, library, receipt, original = sample
    state = {}
    monkeypatch.setattr(ui.st, "session_state", state)

    class RerunRequested(Exception):
        pass

    def rerun():
        raise RerunRequested()

    monkeypatch.setattr(ui.st, "rerun", rerun)
    with testing.isolated(TOKEN_A):
        memory._data_dir()
        ui._queue_run(receipt.fixture_id, testing.generation())
        with pytest.raises(RerunRequested):
            ui._process_pending(library, TOKEN_A, root)
    pending = dict(state["_shared_fixture_pending"])
    assert pending["expected_generation"] != pending["generation"]
    with testing.isolated(TOKEN_A):
        memory._data_dir()
        assert testing.generation() == pending["expected_generation"]
        testing.reset(root)  # A second tab deliberately resets this new run.
    with testing.isolated(TOKEN_A):
        current = memory._data_dir()
        assert testing.generation() != pending["expected_generation"]
        ui._process_pending(library, TOKEN_A, root)
        assert runs.read_run_fixture(library, TOKEN_A) is None
        assert {path.name for path in current.iterdir()} == {testing._SENTINEL}
    assert "_shared_fixture_pending" not in state
    assert "reset again" in state["_shared_fixture_error"]
    assert library.download(receipt.fixture_id) == original


@pytest.mark.parametrize("change", ["contract", "month", "year", "sites"])
def test_target_drift_requires_explicit_restore_without_refilling_report_content(sample, monkeypatch, change):
    root, _library, receipt, _original = sample
    state = {"report_manual_text": "Keep my exact wording"}
    monkeypatch.setattr(ui.st, "session_state", state)
    with testing.isolated(TOKEN_A):
        memory._data_dir()
        profile = ui.target_profile(receipt.metadata)
        period = ReportPeriod(2026, 10)
        ui._seed_target(profile, period)
        assert ui._target_matches(profile, period)
        changed = {"contract": ("report_contract", "Other Contract"),
                   "month": ("report_month_number", 11), "year": ("report_year", 2027),
                   "sites": ("report_sites_synthetic-testing-contract_new_site", "Other Site")}[change]
        state[changed[0]] = changed[1]
        ui._seed_target(profile, period)
        assert state[changed[0]] == changed[1]  # A routine rerun never silently resets a selection.
        assert not ui._target_matches(profile, period)
        ui._seed_target(profile, period, force=True)  # Explicit Restore fixture target action.
        assert ui._target_matches(profile, period)
        assert state["report_manual_text"] == "Keep my exact wording"
