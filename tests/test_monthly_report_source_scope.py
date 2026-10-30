"""Unknown included scope needs exact source or contribution evidence."""
from dataclasses import asdict, replace
from datetime import date
from hashlib import sha256
from io import BytesIO
import json

from docx import Document
import pytest
from streamlit.testing.v1 import AppTest

from app import monthly_report_library as library
from app.monthly_report_checks import preflight
from app.monthly_report_cmms import CMMSMapping, map_cmms
from app.monthly_report_docx import assemble_docx
from app.monthly_report_guided import generation_block_reason
from app.monthly_report_model import (
    BlockSpec, ReportDraft, ReportPeriod, ResolvedBlock, SectionSpec, synthetic_profiles,
)
from app.monthly_report_source_details import correct_source_details
from app.monthly_report_source_scope import (
    CONTRACT_ARTWORK, MULTIPLE_FACILITIES, USE_PREFIX, contribution_receipt,
    correct_scope, included_scope_valid, reviewed_scope,
)
from app.monthly_report_sources import ingest, source_bytes, source_reference

PERIOD = ReportPeriod(2026, 9)


@pytest.fixture
def sample(monkeypatch, tmp_path):
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path))
    profile = synthetic_profiles()[1]
    raw = b"Inspected a synthetic pump. Follow-up remains pending."
    content, _ = ingest(profile, "unknown-site.txt", raw)
    source = replace(content.source, service_date="2026-09-16", classification="Vendor service")
    block = ResolvedBlock("activity_summary", "This month", text=raw.decode(),
                          references=(source_reference(source, 1),))
    draft = ReportDraft(profile, PERIOD, "Synthetic Editor",
        (SectionSpec("activity", "2", "Monthly activity", (BlockSpec(block.key, "rich_text"),)),),
        (block,), synthetic=True, sources=(source,))
    return draft, raw


def site_blockers(draft):
    return tuple(c for c in preflight(draft) if c.code == "source_facility" and c.blocking)


@pytest.mark.parametrize("classification,selected", [("Vendor service", (1,)), ("Reference only", (1,)), ("Vendor service", ())])
def test_included_unknown_site_cannot_be_waived_by_warning_acknowledgment(sample, classification, selected):
    draft, _ = sample
    draft = replace(draft, sources=(replace(draft.sources[0], classification=classification, selected_pages=selected),))
    assert [c.block_key for c in site_blockers(draft)] == [draft.sources[0].id]
    assert generation_block_reason(conflict=None, checks=preflight(draft), warnings_ok=True)
    with pytest.raises(ValueError, match="Establish the site scope"):
        assemble_docx(draft, acknowledged_fingerprint=draft.fingerprint)


@pytest.mark.parametrize("mode", ["reference", "unused", "omitted", "empty"])
def test_nonincluded_unknown_sources_do_not_block(sample, mode):
    draft, _ = sample
    block = draft.blocks[0]
    if mode in {"reference", "unused"}:
        block = replace(block, references=())
    elif mode == "omitted":
        block = replace(block, source="Omit")
    else:
        block = replace(block, text="")
    draft = replace(draft, blocks=(block,))
    assert not site_blockers(draft)
    assert assemble_docx(draft, acknowledged_fingerprint=draft.fingerprint).startswith(b"PK")


def test_combined_site_correction_is_explicit_persisted_and_keeps_original_bytes(sample):
    draft, raw = sample
    source = correct_source_details(draft.sources[0], draft.profile, date(2026, 9, 16),
                                    MULTIPLE_FACILITIES, chosen_facilities=("north", "south"))
    assert source.facility == "" and reviewed_scope(source, draft.profile)["facilities"] == ["north", "south"]
    corrected = replace(draft, sources=(source,))
    assert not site_blockers(corrected)
    library.save_profile(corrected.profile, expected_revision=0, actor="Synthetic Editor", confirmed=True)
    library.save_snapshot(corrected, expected_revision=0, entered_editor="Synthetic Editor")
    reopened = library.load_snapshot(corrected.profile.contract, corrected.profile.key, PERIOD).draft
    assert reopened == corrected and reopened.fingerprint == corrected.fingerprint
    assert source_bytes(reopened.profile, reopened.sources[0]) == raw
    word = assemble_docx(reopened, acknowledged_fingerprint=reopened.fingerprint)
    assert raw.decode() in "\n".join(p.text for p in Document(BytesIO(word)).paragraphs)


@pytest.mark.parametrize("change", ["bytes", "text", "date", "target", "contract"])
def test_scope_correction_cannot_approve_changed_source_or_foreign_target(sample, change):
    draft, _ = sample
    source = correct_scope(draft.sources[0], draft.profile, ("north", "south"))
    draft = replace(draft, sources=(source,))
    if change == "bytes": source = replace(source, sha256="f" * 64)
    elif change == "text": source = replace(source, page_texts=("Different source content",))
    elif change == "date": source = replace(source, service_date="2026-09-17")
    elif change == "target": draft = replace(draft, profile=synthetic_profiles()[0])
    else: draft = replace(draft, profile=replace(draft.profile, contract="Another contract"))
    assert site_blockers(replace(draft, sources=(source,)))


def test_page_selection_and_visual_review_do_not_require_repeat_site_confirmation(sample):
    draft, _ = sample
    source = correct_scope(draft.sources[0], draft.profile, ("north", "south"))
    changed = replace(source, selected_pages=(), captions=((1, "Human source-page caption"),),
                      client_page_reviews=((1, "visual-review"),), client_page_exclusions=((1, "excluded"),))
    assert reviewed_scope(changed, draft.profile)
    assert not site_blockers(replace(draft, sources=(changed,)))


def test_absent_scope_record_preserves_legacy_source_and_report_fingerprints(sample):
    draft, _ = sample
    def old_hash(record):
        return sha256(json.dumps(record, sort_keys=True, ensure_ascii=False).encode()).hexdigest()
    source_record = asdict(draft.sources[0]); source_record.pop("scope_review")
    source_record.pop("pdf_text_blocks")  # Absent in the historical schema too.
    source_record.pop("financial_source_review")
    source_record.pop("financial_page_reviews")
    draft_record = asdict(draft)
    draft_record.pop("cover_recovery_request")  # Private intent did not exist in old reports.
    draft_record.pop("source_row_reviews", None)
    # The historical report schema predates occurrence metadata, report-owned
    # row-review receipts, as well as
    # source-scope receipts. Current asdict() otherwise invents a new empty
    # field in this independently reconstructed old-format record.
    for block in draft_record["blocks"]:
        block.pop("asset_instances")
        block.pop("cover_original")
    for source in draft_record["sources"]:
        source.pop("scope_review")
        source.pop("pdf_text_blocks")
        source.pop("financial_source_review")
        source.pop("financial_page_reviews")
    assert draft.sources[0].fingerprint == old_hash(source_record)
    assert draft.fingerprint == old_hash(draft_record)
    # Independently evaluated on immutable #635 tree3ef4dd54, before adding
    # cover recovery. Do not derive this expected value from current asdict.
    assert draft.fingerprint == "d527d3e4764a82d65db3da8ccdb51fa1d4ebc9b763a7c97aab2382d1f2246329"
    assert library.draft_from_dict(draft_record) == draft


def test_filename_hint_never_establishes_source_site(sample):
    draft, _ = sample
    filename_only, _ = ingest(draft.profile, "Demo North service report.txt", b"Inspected synthetic pump.")
    assert filename_only.source.classification == "Vendor service"
    assert filename_only.source.facility == ""
    native, _ = ingest(draft.profile, "unlabeled.txt", b"Service report\nSite: Demo North\nInspected synthetic pump.")
    assert native.source.facility == "Demo North"
    mixed, _ = ingest(draft.profile, "Demo North service report.txt",
                      b"Work at Demo North and Demo South: inspection records.")
    assert mixed.source.facility == ""


def test_real_cmms_mapping_retains_exact_row_scope_without_source_wide_guess(sample):
    draft, _ = sample
    raw = (b"ID,Finished,Facility,Task\n001,2026-09-02,Demo North,Inspected pump\n"
           b"002,2026-09-03,Demo South,Checked fan\n003,2026-09-04,Foreign site,Foreign work\n")
    content, _ = ingest(draft.profile, "combined-cmms.csv", raw)
    assert content.source.facility == ""
    mapped = map_cmms(content.tables[0], CMMSMapping("Finished", "ID", facility="Facility", description="Task",
        finish_date_means_completed=True), draft.profile, PERIOD, source_reference(content.source, 1),
        scope_source=content.source)
    mapped_draft = replace(draft, sources=(content.source,), blocks=mapped.blocks,
        sections=(SectionSpec("maintenance", "5", "Onsite service calls", mapped.specs),))
    assert not site_blockers(mapped_draft)
    word = assemble_docx(mapped_draft, acknowledged_fingerprint=mapped_draft.fingerprint)
    text = "\n".join(cell.text for table in Document(BytesIO(word)).tables for row in table.rows for cell in row.cells)
    assert "Inspected pump" in text and "Checked fan" in text and "Foreign work" not in text
    assert library.draft_from_dict(asdict(mapped_draft)) == mapped_draft
    visual_only = replace(content.source, captions=((1, "Edited page caption"),), selected_pages=(),
                          client_page_reviews=((1, "page review"),))
    assert not site_blockers(replace(mapped_draft, sources=(visual_only,)))
    assert site_blockers(replace(mapped_draft, sources=(replace(content.source, page_texts=("Changed data",)),)))
    changed = replace(mapped.blocks[0], rows=(("Foreign site", *mapped.blocks[0].rows[0][1:]), *mapped.blocks[0].rows[1:]))
    assert site_blockers(replace(mapped_draft, blocks=(changed, mapped.blocks[1])))
    raw_only = tuple(replace(b, references=tuple(r for r in b.references if not r.startswith(USE_PREFIX))) for b in mapped.blocks)
    assert site_blockers(replace(mapped_draft, blocks=raw_only)), "Legacy scope is not inferred from a CMMS label."
    assert source_bytes(draft.profile, content.source) == raw


def test_totals_scope_covers_only_accepted_text_and_current_source(sample):
    draft, _ = sample
    source = draft.sources[0]
    accepted = "September synthetic work orders: 3 total, 2 completed, 1 open."
    reference = source_reference(source, 1)
    receipt = contribution_receipt(reference, draft.profile, PERIOD, "activity_summary", text=accepted,
                                   prior_block=replace(draft.blocks[0], text="Independent human note.", references=()))
    block = replace(draft.blocks[0], text="Independent human note.\n\n" + accepted, references=(reference, receipt))
    ready = replace(draft, blocks=(block,))
    assert not site_blockers(ready)
    assert site_blockers(replace(ready, blocks=(replace(block, text=block.text.replace("3 total", "30 total")),)))
    assert site_blockers(replace(ready, sources=(replace(source, findings="New source finding"),)))
    assert site_blockers(replace(ready, period=ReportPeriod(2026, 10)))
    assert source.scope_review == "", "A contribution receipt does not declare the whole source site-correct."


def test_totals_receipt_never_approves_prior_or_later_native_source_narrative(sample):
    draft, _ = sample
    source, original = draft.sources[0], draft.blocks[0]
    reference = source_reference(source, 1)
    totals = "Accepted synthetic totals."
    receipt = contribution_receipt(reference, draft.profile, PERIOD, original.key, text=totals,
                                   source=source, prior_block=original)
    mixed = replace(original, text=original.text + "\n\n" + totals, references=(reference, receipt))
    unresolved = replace(draft, blocks=(mixed,))
    assert site_blockers(unresolved)
    with pytest.raises(ValueError, match="Establish the site scope"):
        assemble_docx(unresolved, acknowledged_fingerprint=unresolved.fingerprint)
    assert not site_blockers(replace(draft, blocks=(replace(mixed, text=totals),)))
    # Independent/known-source text remains usable through an ordinary first
    # acceptance; a new native paragraph reference from this same unknown
    # source afterwards must reopen only this source's scope.
    unrelated = replace(original, text="Independent human note.\n\nOther site-confirmed source wording.",
                        references=("another-source:1:known",))
    receipt = contribution_receipt(reference, draft.profile, PERIOD, original.key, text=totals,
                                   source=source, prior_block=unrelated)
    mixed = replace(unrelated, text=unrelated.text + "\n\n" + totals,
                    references=(*unrelated.references, reference, receipt))
    assert not site_blockers(replace(draft, blocks=(mixed,)))
    later = replace(mixed, text=mixed.text + "\n\nLater native source paragraph.",
                    references=(*mixed.references, source_reference(source, 2)))
    assert site_blockers(replace(draft, blocks=(later,)))
    # A second accepted totals contribution from the same file keeps the
    # previously established independent human wording; it adds no blanket
    # approval of an intervening native paragraph.
    next_totals = "Accepted second-table totals."
    next_receipt = contribution_receipt(source_reference(source, 2), draft.profile, PERIOD, original.key,
        text=next_totals, source=source, prior_block=mixed)
    second = replace(mixed, text=mixed.text + "\n\n" + next_totals,
        references=(*mixed.references, source_reference(source, 2), next_receipt))
    assert not site_blockers(replace(draft, blocks=(second,)))
    bad_next_receipt = contribution_receipt(source_reference(source, 2), draft.profile, PERIOD, original.key,
        text=next_totals, source=source, prior_block=later)
    assert site_blockers(replace(draft, blocks=(replace(second,
        text=later.text + "\n\n" + next_totals, references=(*later.references, bad_next_receipt)),)))


def test_contract_artwork_scope_cannot_approve_operational_content(sample):
    draft, _ = sample
    with pytest.raises(ValueError, match="shared report artwork"):
        correct_scope(draft.sources[0], draft.profile, contract_artwork=True, artwork_keys=("activity_summary",))
    source = correct_source_details(draft.sources[0], draft.profile, date(2026, 9, 16), CONTRACT_ARTWORK,
                                    artwork_keys=("client_logo",))
    artwork = ResolvedBlock("client_logo", "This month", asset_hashes=("logo.png",),
                           references=(source_reference(source, 1),))
    art_draft = replace(draft, sources=(source,), blocks=(artwork,))
    assert included_scope_valid(art_draft, source)
    same_contract_other_site = replace(art_draft, profile=synthetic_profiles()[0])
    assert included_scope_valid(same_contract_other_site, source)
    assert not included_scope_valid(replace(art_draft, blocks=(artwork, draft.blocks[0])), source)
    assert not included_scope_valid(replace(art_draft, profile=replace(draft.profile, contract="Other")), source)


def test_same_site_historical_source_can_be_retained_after_targeted_scope_correction(sample):
    draft, _ = sample
    source = correct_source_details(draft.sources[0], draft.profile, date(2026, 8, 16), "north")
    draft = replace(draft, sources=(source,))
    assert not site_blockers(draft)
    assert any(c.code == "source_period" and not c.blocking for c in preflight(draft))
    assert assemble_docx(draft, acknowledged_fingerprint=draft.fingerprint).startswith(b"PK")


def test_single_site_correction_is_also_bound_to_current_source_revision(sample):
    draft, _ = sample
    source = correct_source_details(draft.sources[0], draft.profile, date(2026, 9, 16), "north")
    assert reviewed_scope(source, draft.profile)["facilities"] == ["north"]
    assert not site_blockers(replace(draft, sources=(source,)))
    assert site_blockers(replace(draft, sources=(replace(source, sha256="f" * 64),)))


def test_mapping_scope_survives_independent_row_edits_and_merges(sample):
    from app.monthly_report_setup import merge_report_block
    draft, _ = sample
    mappings, sources = [], []
    for number, site in ((1, "Demo North"), (2, "Demo South")):
        raw = f"ID,Finished,Facility,Task\n{number},2026-09-02,{site},Inspection {number}\n".encode()
        content, _ = ingest(draft.profile, f"work-{number}.csv", raw)
        # The explicit per-row mapping remains proof even when the source
        # covers multiple sites and has no whole-source facility label.
        source = replace(content.source, facility="")
        mappings.append(map_cmms(content.tables[0], CMMSMapping("Finished", "ID", facility="Facility",
            description="Task", finish_date_means_completed=True), draft.profile, PERIOD, source_reference(source, 1)))
        sources.append(source)
    first, second = mappings
    merged = tuple(merge_report_block(a, b, x, y)
                   for a, b, x, y in zip(first.blocks, second.blocks, first.specs, second.specs))
    combined = replace(draft, sources=tuple(sources), blocks=merged,
        sections=(SectionSpec("maintenance", "5", "Onsite service calls", first.specs),))
    assert not site_blockers(combined)
    row = merged[0].rows[0]
    edited = replace(merged[0], rows=((*row[:4], "Human corrected description", row[-1]),))
    assert not site_blockers(replace(combined, blocks=(edited, merged[1])))
    foreign = replace(edited, rows=(("Foreign site", *edited.rows[0][1:]),))
    assert site_blockers(replace(combined, blocks=(foreign, merged[1])))


def test_scope_receipts_are_internal_metadata_and_never_client_citations(sample):
    from zipfile import ZipFile
    from app.monthly_report_ai import ai_references
    from app.monthly_report_asset_review import asset_context
    from app.monthly_report_model import used_source_ids
    from app.monthly_report_compare_ui import _block_summary
    draft, _ = sample
    source = draft.sources[0]
    reference = source_reference(source, 1)
    text = "Accepted synthetic totals."
    receipt = contribution_receipt(reference, draft.profile, PERIOD, "activity_summary", text=text)
    block = replace(draft.blocks[0], text=text, references=(reference, receipt))
    ready = replace(draft, blocks=(block,))
    assert ai_references(block) == (reference,)
    assert asset_context(replace(block, asset_hashes=("synthetic.png",)), 0) == (reference,)
    assert used_source_ids(ready) == {source.id}
    shown = _block_summary(block, ready.sections[0].blocks[0], ready)
    assert source.filename in shown and receipt not in shown and USE_PREFIX not in shown
    word = assemble_docx(ready, acknowledged_fingerprint=ready.fingerprint)
    with ZipFile(BytesIO(word)) as archive:
        for name in archive.namelist():
            assert receipt.encode() not in archive.read(name) and USE_PREFIX.encode() not in archive.read(name)
    # A valid receipt for a previously unknown source does not excuse a
    # subsequently known foreign site, even with independent human wording.
    foreign = replace(source, facility="Foreign site")
    forged_current = contribution_receipt(source_reference(foreign, 1), draft.profile, PERIOD,
                                          block.key, text=text)
    foreign_block = replace(block, text="Human note.\n\n" + text,
                            references=(source_reference(foreign, 1), forged_current))
    assert site_blockers(replace(ready, sources=(foreign,), blocks=(foreign_block,)))


def test_combined_source_form_requires_explicit_sites_and_rejects_stale_submission(sample):
    draft, _ = sample
    app = AppTest.from_string('''
import streamlit as st
from app.monthly_report_upload_ui import render_source_details
from app.monthly_report_model import synthetic_profiles
def field(key, default):
    st.session_state.setdefault(key, default)
    return key
st.session_state["source"] = render_source_details(st.session_state["source"], synthetic_profiles()[1], "scope-test", field)
''')
    app.session_state["source"] = draft.sources[0]
    app.run()
    assert not app.exception and app.multiselect[0].value == []
    app.multiselect[0].set_value(["north", "south"])
    app.button[0].click().run()
    assert not app.exception
    accepted = app.session_state["source"]
    assert reviewed_scope(accepted, draft.profile)["facilities"] == ["north", "south"]
    app.run()
    assert app.multiselect[0].value == ["north", "south"]
    app.multiselect[0].set_value(["south"])
    app.session_state["source"] = replace(accepted, findings="Changed in another view")
    app.button[0].click().run()
    assert app.session_state["source"].facility == ""
    assert any("corrected elsewhere" in message.value for message in app.info)


def test_unknown_site_totals_use_one_targeted_correction_then_one_acceptance(monkeypatch, tmp_path):
    from test_monthly_report_workorder_upload_ui import _upload
    from test_monthly_report_ui import monthly

    app = monthly(monkeypatch, tmp_path)
    next(widget for widget in app.text_input if widget.label == "Prepared by").set_value("Synthetic Editor").run()
    raw = b"Work Order,Date,Status\n001,2026-09-03,Completed\n002,2026-09-20,Open\n"
    _upload(monkeypatch, raw)
    app.run()
    assert not app.exception
    add = next(button for button in app.button if button.label == "Add totals to summary")
    assert add.disabled
    assert not any("verified that this work-order" in widget.label for widget in app.checkbox)
    site = next(widget for widget in app.selectbox if widget.label == "Source facility")
    assert site.value == ""
    site.set_value("north")
    next(button for button in app.button if button.label == "Use these source details").click().run()
    assert not app.exception
    # The submitted source correction is durable before its next ordinary
    # insertion action. No second submit or injected acceptance repairs it.
    add = next(button for button in app.button if button.label == "Add totals to summary")
    assert not add.disabled
    add.click().run()
    assert not app.exception
    drafts = [value for key, value in app.session_state.filtered_state.items()
              if key.endswith("_draft") and isinstance(value, ReportDraft)]
    current = next(d for d in drafts if any(s.filename == "work_orders.csv" for s in d.sources))
    assert not site_blockers(current)
    activity = next(b for b in current.blocks if b.key == "activity_summary")
    assert activity.text.count("2 total, 1 completed, 1 open") == 1
    assert any(r.startswith(USE_PREFIX) for r in activity.references)
    output = assemble_docx(current, acknowledged_fingerprint=current.fingerprint)
    assert "2 total, 1 completed, 1 open" in "\n".join(p.text for p in Document(BytesIO(output)).paragraphs)
    app.run()
    assert next(w for w in app.text_area if w.label == "Activity summary").value == activity.text
    source = next(s for s in current.sources if s.filename == "work_orders.csv")
    assert source_bytes(current.profile, source) == raw


def test_isolated_replay_scope_correction_has_exact_durable_report_ack(monkeypatch, tmp_path):
    from test_monthly_report_qa_ui import visit, ACTOR_COOKIE
    from app import monthly_report_testing as qa, monthly_report_qa_fixture as fixture
    from app.monthly_report_active_work import load
    app = visit(monkeypatch, tmp_path)
    next(button for button in app.button if button.label == "Replay synthetic source files").click().run()
    assert not app.exception
    with qa.isolated(ACTOR_COOKIE):
        saved = library.load_snapshot(fixture.CONTRACT, fixture.profile().key, fixture.PERIOD)
    assert site_blockers(saved.draft)
    original = next(source for source in saved.draft.sources if source.filename == "synthetic-activity.txt")
    selector = next(widget for widget in app.multiselect if widget.label == "Source sites" and original.id in widget.key)
    assert selector.value == []
    selector.set_value(["qa-north"])
    next(button for button in app.button if button.label == "Use these source details" and original.id in button.key).click().run()
    assert not app.exception
    with qa.isolated(ACTOR_COOKIE):
        active = load(ACTOR_COOKIE, fixture.CONTRACT, fixture.profile().key, fixture.PERIOD)
        assert active is not None and not site_blockers(active.draft)
        assert active.draft.fingerprint != saved.draft.fingerprint
        source = next(source for source in active.draft.sources if source.id == original.id)
        assert source_bytes(active.draft.profile, source) == fixture.ACTIVITY_TXT
        assert source.sha256 == original.sha256
        assert any(check.blocking and "follow-up source changed" in check.message for check in preflight(active.draft))
    # The site correction changes source provenance. It removes only its
    # scope blocker; explicitly replace the issue's stale evidence link.
    selected = next(widget for widget in app.multiselect if widget.label == "Supporting source pages")
    selected.set_value([source_reference(source, 1)]).run()
    next(widget for widget in app.checkbox if widget.label == "Remove unavailable source links").check().run()
    assert not app.exception
    with qa.isolated(ACTOR_COOKIE):
        active = load(ACTOR_COOKIE, fixture.CONTRACT, fixture.profile().key, fixture.PERIOD)
        assert active is not None and not any(check.blocking for check in preflight(active.draft))
        output = assemble_docx(active.draft, acknowledged_fingerprint=active.draft.fingerprint)
        assert output.startswith(b"PK")
    assert any(item.value == "QA durable report SHA256: " + active.draft.fingerprint for item in app.caption)
    assert any(item.value == "QA working report SHA256: " + active.draft.fingerprint for item in app.caption)
