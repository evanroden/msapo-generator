"""Current improvement photographs must not inherit another image's stretch/crop."""
from dataclasses import replace
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches
from PIL import Image

from app.monthly_report_model import Facility, ReportDraft, ReportPeriod, ReportProfile, ResolvedBlock, default_sections
from app.monthly_report_native_layout import build_native_docx
from conftest import requires_libreoffice


def png(size, color):
    data=BytesIO()
    Image.new('RGB',size,color).save(data,format='PNG')
    return data.getvalue()


def fixture(tmp_path, *, cropped=False):
    source=tmp_path/'old-report.docx'
    source_photo=png((600,450),'blue')
    replacement=png((1400,900),'orange')
    doc=Document()
    doc.add_paragraph('Previous Synthetic Client')
    doc.add_heading('MONTHLY ACTIVITY SUMMARY',1)
    doc.add_paragraph('Previous private activity')
    paragraph=doc.add_paragraph()
    frame=paragraph.add_run().add_picture(BytesIO(source_photo),width=Inches(3),height=Inches(2.25))
    if cropped:
        blipfill=next(frame._inline.iter(qn('pic:blipFill')))
        crop=OxmlElement('a:srcRect');crop.set('l','10000');crop.set('t','12000')
        blipfill.insert(1,crop)
    doc.add_heading('TRAINING SUMMARY',1)
    doc.add_paragraph('Previous training')
    doc.save(source)
    profile=ReportProfile('Synthetic Contract','site','Synthetic Current Site',(Facility('site','Synthetic Current Site'),))
    draft=ReportDraft(profile,ReportPeriod(2026,9),'Synthetic Editor',default_sections(),(
        ResolvedBlock('improvements','This month',asset_hashes=('photo.png',),asset_captions=('Current synthetic image',)),
    ))
    return source,draft,source_photo,replacement


def test_photo_replacement_contains_full_original_image_without_stretch_or_crop(tmp_path):
    source,draft,old,new=fixture(tmp_path,cropped=True)
    before=source.read_bytes()
    raw=build_native_docx(draft,source,master=True,section_key='activity',asset_loader=lambda _:new)
    after=Document(BytesIO(raw))
    picture=after.inline_shapes[0]
    assert (picture.width / picture.height)==pytest.approx(1400/900,abs=.003)
    assert float(picture.width) <= float(Inches(3)) + 200
    assert float(picture.height) <= float(Inches(2.25)) + 200
    assert not list(picture._inline.iter(qn('a:srcRect')))
    with ZipFile(BytesIO(raw)) as z:
        values=[z.read(name) for name in z.namelist()]
        assert old not in values and new in values
    assert 'Previous private activity' not in '\n'.join(p.text for p in after.paragraphs)
    assert source.read_bytes()==before


@requires_libreoffice
def test_rendered_photo_matches_its_source_pixel_ratio(tmp_path,monkeypatch):
    import pymupdf
    from app import pdf_converter
    source,draft,old,new=fixture(tmp_path)
    monkeypatch.setenv('EPC_DATA_DIR',str(tmp_path/'runtime'))
    monkeypatch.setattr(pdf_converter,'OUTPUT_DIR',tmp_path)
    path=tmp_path/'current.docx'
    path.write_bytes(build_native_docx(draft,source,master=True,section_key='activity',asset_loader=lambda _:new))
    output=pdf_converter.convert_to_pdf(path)
    try:
        with pymupdf.open(output) as pdf:
            images=[img for page in pdf for img in page.get_image_info()]
            # LibreOffice is permitted to resample embedded pixels, not distort
            # their placement; use the rendered rectangle rather than its xref.
            assert len(images)==1
            bounds=images[0]['bbox']
            width,height=bounds[2]-bounds[0],bounds[3]-bounds[1]
            assert width/height==pytest.approx(1400/900,abs=.01)
            assert width <= 216.05 and height <= 162.05
    finally:
        output.unlink(missing_ok=True)


def test_anchored_photo_preserves_original_envelope_and_current_ratio(tmp_path):
    source,draft,old,new=fixture(tmp_path)
    document=Document(source)
    frame=document.inline_shapes[0]._inline
    frame.tag=qn('wp:anchor')
    document.save(source)
    raw=build_native_docx(draft,source,master=True,section_key='activity',asset_loader=lambda _:new)
    result=Document(BytesIO(raw))
    anchors=list(result.element.body.iter(qn('wp:anchor')))
    assert len(anchors)==1
    extent=anchors[0].find(qn('wp:extent'))
    assert int(extent.get('cx'))/int(extent.get('cy')) == pytest.approx(1400/900,abs=.002)
    assert int(extent.get('cx')) <= int(Inches(3))
    assert int(extent.get('cy')) <= int(Inches(2.25))
    assert not list(anchors[0].iter(qn('a:srcRect')))


@pytest.mark.parametrize('style',('width:3in;height:2.25in','height:162pt;width:216pt'))
def test_vml_fallback_replacement_contains_current_image(style):
    from lxml import etree
    from app.monthly_report_native_layout import _replace_image
    document=Document()
    namespace='{urn:schemas-microsoft-com:vml}'
    paragraph=OxmlElement('w:p')
    pict=OxmlElement('w:pict');paragraph.append(pict)
    shape=etree.SubElement(pict, namespace+'shape', style=style, id='old-frame')
    image=etree.SubElement(shape, namespace+'imagedata', cropleft='.1', croptop='.2')
    # Relinking uses the content address, not the old placeholder relationship.
    image.set(qn('r:id'),'rId9000')
    replacement=png((1400,900),'orange')
    _replace_image(document,paragraph,replacement,contain=True)
    parsed={k:float(v) * (72 if unit=='in' else 1)
            for k,v,unit in __import__('re').findall(r'(width|height):([0-9.]+)(in|pt)',shape.get('style'))}
    assert parsed['width']/parsed['height']==pytest.approx(1400/900,abs=.001)
    assert parsed['width']<=216.01 and parsed['height']<=162.01
    assert shape.get('id')=='old-frame'
    assert not {'cropleft','cropright','croptop','cropbottom'}.intersection(image.attrib)
    assert image.get(qn('r:id'))!='rId9000'


def test_unmeasurable_vml_replacement_fails_before_export():
    from lxml import etree
    from app.monthly_report_native_layout import _replace_image, NativeLayoutError
    document=Document(); paragraph=OxmlElement('w:p');pict=OxmlElement('w:pict');paragraph.append(pict)
    shape=etree.SubElement(pict,'{urn:schemas-microsoft-com:vml}shape',style='width:200px;height:120px')
    image=etree.SubElement(shape,'{urn:schemas-microsoft-com:vml}imagedata')
    image.set(qn('r:id'),'rId9000')
    with pytest.raises(NativeLayoutError,match='dimensions cannot be fitted'):
        _replace_image(document,paragraph,png((1200,700),'blue'),contain=True)
