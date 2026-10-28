"""Replacing a staged mapping source is transactional until decoding succeeds."""
from io import BytesIO

from docx import Document
import pytest
from streamlit.testing.v1 import AppTest

from app import monthly_report_import_ui as ui
from app import monthly_report_library as library
from app.monthly_report_import import ImportMapping
from app.monthly_report_model import synthetic_profiles


APP = '''
from app import monthly_report_library as library
from app.monthly_report_import_ui import render_import
from app.monthly_report_model import synthetic_profiles
from app.monthly_report_ui import _field
profile = synthetic_profiles()[0]
render_import(library.load_profile(profile.contract, profile.key), _field)
'''


def upload(text):
    document = Document()
    document.add_heading("2 Monthly Activity Summary", 1)
    document.add_paragraph(text)
    stream = BytesIO()
    document.save(stream)
    stream.size, stream.name = len(stream.getvalue()), "synthetic.docx"
    return stream


@pytest.fixture
def reviewed(monkeypatch, tmp_path):
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path))
    profile = synthetic_profiles()[0]
    library.save_profile(profile, expected_revision=0, actor="Synthetic Reviewer", confirmed=True)
    selected = [upload("Synthetic prior source finding.")]
    monkeypatch.setattr(ui.st, "file_uploader", lambda *a, **kw: selected[0])
    app = AppTest.from_string(APP, default_timeout=20).run()
    next(button for button in app.button if button.label == "Read DOCX for mapping").click().run()
    assert not app.exception
    prefix = "report_import_" + ui._signature((profile.contract, profile.key))
    stage = app.session_state[prefix + "_stage"]
    item = next(item for item in stage[2].items if item.text == "Synthetic prior source finding.")
    mappings = {item.id: ImportMapping(item.id, "activity_summary")}
    app.session_state[prefix + "_mappings"] = mappings
    app.session_state[prefix + "_saved"] = "Prior confirmed save receipt."
    app.run()
    assert not app.exception
    yield app, profile, selected, prefix, stage, mappings
    final = app.session_state.filtered_state.get(prefix + "_stage")
    if final:
        final[0].cleanup()


@pytest.mark.parametrize("failure", ["malformed", "oversize", "reader"])
def test_failed_replacement_keeps_original_review_and_saved_work(reviewed, monkeypatch, failure):
    app, profile, selected, prefix, stage, mappings = reviewed
    original = stage[1].read_bytes()
    manifest = library._profile_path(profile.contract, profile.key) / "manifest.json"
    saved = manifest.read_bytes()
    if failure == "malformed":
        selected[0] = BytesIO(b"Not a Word package")
        selected[0].size, selected[0].name = len(selected[0].getvalue()), "broken.docx"
    else:
        selected[0] = upload("Synthetic replacement finding.")
        if failure == "oversize":
            monkeypatch.setattr(ui, "MAX_DOCX_BYTES", 10)
        else:
            def unavailable(path):
                raise OSError("Synthetic source reader unavailable")
            monkeypatch.setattr(ui, "inspect_docx", unavailable)
    next(button for button in app.button if button.label == "Read DOCX for mapping").click().run()
    assert not app.exception
    assert app.session_state[prefix + "_stage"] is stage
    assert stage[1].read_bytes() == original
    assert app.session_state[prefix + "_mappings"] == mappings
    assert app.session_state[prefix + "_saved"] == "Prior confirmed save receipt."
    assert manifest.read_bytes() == saved
    assert any("previous source review is unchanged" in error.value for error in app.error)
    app.run()
    assert not app.exception and app.session_state[prefix + "_mappings"] == mappings
    assert stage[1].read_bytes() == original


def test_successful_replacement_changes_review_only_after_valid_decode(reviewed):
    app, profile, selected, prefix, stage, mappings = reviewed
    selected[0] = upload("Synthetic replacement finding.")
    manifest = library._profile_path(profile.contract, profile.key) / "manifest.json"
    before = manifest.read_bytes()
    next(button for button in app.button if button.label == "Read DOCX for mapping").click().run()
    assert not app.exception
    replacement = app.session_state[prefix + "_stage"]
    assert replacement is not stage and replacement[1].read_bytes() == selected[0].getvalue()
    assert not stage[1].exists()
    assert app.session_state[prefix + "_mappings"] == {}
    assert prefix + "_saved" not in app.session_state.filtered_state
    assert manifest.read_bytes() == before


def test_old_temporary_cleanup_failure_does_not_discard_the_valid_replacement(reviewed, monkeypatch):
    app, profile, selected, prefix, stage, mappings = reviewed
    cleanup = stage[0].cleanup
    def unavailable():
        raise OSError("Synthetic old-directory cleanup failure")
    selected[0] = upload("Synthetic replacement after cleanup failure.")
    monkeypatch.setattr(stage[0], "cleanup", unavailable)
    try:
        next(button for button in app.button if button.label == "Read DOCX for mapping").click().run()
        assert not app.exception
        assert app.session_state[prefix + "_stage"][1].read_bytes() == selected[0].getvalue()
        assert app.session_state[prefix + "_mappings"] == {}
        assert any("replacement is ready for review" in item.value for item in app.warning)
        assert not any("previous source review is unchanged" in item.value for item in app.error)
    finally:
        cleanup()
