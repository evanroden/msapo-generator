"""Converter descendants must stop before another visitor can use the slot."""
from pathlib import Path
from time import monotonic, sleep
from io import BytesIO
import os
import signal
import subprocess
import sys

import pytest
from docx import Document

from app import pdf_converter as converter
from app import monthly_report_render_jobs as jobs
from conftest import requires_libreoffice


def _running(pid):
    try:
        # Zombies have no running code or address space and cannot write output.
        stat = Path(f'/proc/{pid}/stat').read_text().rsplit(')', 1)[1].split()
        return stat[0] not in {'Z', 'X'}
    except FileNotFoundError:
        return False


@pytest.fixture
def launch_probe(monkeypatch, tmp_path):
    monkeypatch.setenv('EPC_DATA_DIR', str(tmp_path / 'runtime'))
    out = tmp_path / 'outputs'; out.mkdir()
    monkeypatch.setattr(converter, 'OUTPUT_DIR', out)
    monkeypatch.setattr(converter, 'PDF_BACKEND', 'libreoffice')
    executable = tmp_path / 'synthetic-office'
    ledger = tmp_path / 'owned-pids'
    late = tmp_path / 'late-output'
    executable.write_text(f'''#!{sys.executable}
import os, subprocess, sys, time
from pathlib import Path
args = sys.argv
out = Path(args[args.index('--outdir') + 1])
stem = Path(args[-1]).stem
child = subprocess.Popen([sys.executable, '-c', "import time; from pathlib import Path; time.sleep(0.8); Path({str(late)!r}).write_text('late'); Path(" + repr(str(out / (stem + '.pdf'))) + ").write_bytes(b'%PDF-1.7 late'); time.sleep(10)"], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
Path({str(ledger)!r}).write_text(str(os.getpid()) + ' ' + str(child.pid))
time.sleep(10)
''')
    executable.chmod(0o700)
    monkeypatch.setattr(converter.shutil, 'which', lambda _: str(executable))
    actual_popen = subprocess.Popen
    started = []

    def fast_deadline(args, *a, **kwargs):
        process = actual_popen(args, *a, **kwargs)
        if str(args[0]) == str(executable):
            started.append((process, args))
            communicate, wait = process.communicate, process.wait
            # Shorten only the launcher's existing deadline, not its cleanup.
            def quick_communicate(*a, **kw):
                if kw.get('timeout') == 120:
                    kw['timeout'] = .35
                return communicate(*a, **kw)
            def quick_wait(timeout=None):
                return wait(timeout=.35 if timeout == 120 else timeout)
            process.communicate, process.wait = quick_communicate, quick_wait
        return process

    monkeypatch.setattr(subprocess, 'Popen', fast_deadline)
    source = tmp_path / 'current-report.docx'
    doc = Document(); doc.add_paragraph('Synthetic current report'); doc.save(source)
    try:
        yield source, out, ledger, late, started
    finally:
        if ledger.exists():
            for raw in ledger.read_text().split():
                try:
                    os.kill(int(raw), signal.SIGKILL)
                except ProcessLookupError:
                    pass
        for process, _ in started:
            process.wait(timeout=5)


@pytest.mark.skipif(sys.platform != 'linux', reason='Linux process-tree contract for Render')
def test_writer_timeout_stops_children_before_releasing_the_slot(launch_probe):
    source, out, ledger, late, started = launch_probe
    error = None
    try:
        converter.convert_to_pdf(source)
    except Exception as exc:
        error = exc
    assert ledger.exists(), 'the test launcher must actually start its child'
    with jobs.conversion_slot(wait_seconds=0):
        assert all(not _running(int(pid)) for pid in ledger.read_text().split()), 'slot released with an active converter child'
    assert isinstance(error, converter.PDFConversionError)
    assert 'too long' in str(error).lower()
    assert str(source.parent) not in str(error)
    sleep(.9)
    assert not late.exists() and not (out / (source.stem + '.pdf')).exists()
    assert not (out / (source.stem + '.htm')).exists()
    assert source.exists()



def test_conversion_failure_does_not_return_an_old_output(monkeypatch, tmp_path):
    monkeypatch.setenv('EPC_DATA_DIR', str(tmp_path / 'runtime'))
    monkeypatch.setattr(converter, 'OUTPUT_DIR', tmp_path)
    monkeypatch.setattr(converter, 'PDF_BACKEND', 'libreoffice')
    binary = tmp_path / 'successful-no-op'
    binary.write_text('#!/bin/sh\nexit 0\n'); binary.chmod(0o700)
    monkeypatch.setattr(converter.shutil, 'which', lambda _: str(binary))
    source = tmp_path / 'current.docx'; source.write_bytes(b'placeholder')
    old = tmp_path / 'current.pdf'; old.write_bytes(b'%PDF-1.7 previous visitor')
    with pytest.raises(converter.PDFConversionError):
        converter.convert_to_pdf(source)
    # An existing owner's output must not be deleted or falsely returned.
    assert old.read_bytes() == b'%PDF-1.7 previous visitor'


def test_invalid_new_output_is_not_published(monkeypatch, tmp_path):
    monkeypatch.setenv('EPC_DATA_DIR', str(tmp_path / 'runtime'))
    monkeypatch.setattr(converter, 'OUTPUT_DIR', tmp_path)
    monkeypatch.setattr(converter, 'PDF_BACKEND', 'libreoffice')
    executable = tmp_path / 'bad-pdf'
    executable.write_text(f'''#!{sys.executable}
from pathlib import Path
import sys
out = Path(sys.argv[sys.argv.index('--outdir')+1])
(out / (Path(sys.argv[-1]).stem + '.pdf')).write_bytes(b'<html>not a PDF</html>')
'''); executable.chmod(0o700)
    monkeypatch.setattr(converter.shutil, 'which', lambda _: str(executable))
    source = tmp_path / 'bad.docx'; source.write_bytes(b'placeholder')
    with pytest.raises(converter.PDFConversionError):
        converter.convert_to_pdf(source)
    assert not (tmp_path / 'bad.pdf').exists()


@requires_libreoffice
def test_real_writer_result_survives_temp_cleanup(monkeypatch, tmp_path):
    monkeypatch.setenv('EPC_DATA_DIR', str(tmp_path / 'runtime'))
    monkeypatch.setattr(converter, 'OUTPUT_DIR', tmp_path)
    monkeypatch.setattr(converter, 'PDF_BACKEND', 'libreoffice')
    source = tmp_path / 'normal.docx'
    doc = Document(); doc.add_paragraph('Synthetic current text survives conversion.'); doc.save(source)
    result = converter.convert_to_pdf(source)
    import pymupdf
    with pymupdf.open(result) as pdf:
        assert len(pdf) == 1
        assert 'Synthetic current text survives conversion.' in pdf[0].get_text()
    assert source.exists() and result.parent == tmp_path


def test_pending_cleanup_blocks_then_recovers_without_another_launch(monkeypatch, tmp_path):
    from app import office_process
    monkeypatch.setenv('EPC_DATA_DIR', str(tmp_path / 'runtime'))
    import json
    marker = office_process._marker()
    marker.write_text(json.dumps({'pid': 123456, 'boot': office_process._boot_id()}))
    monkeypatch.setattr(office_process, '_group_running', lambda _: True)
    with pytest.raises(jobs.RenderBusy, match='still stopping'):
        with jobs.conversion_slot(wait_seconds=0):
            pytest.fail('an uncleared process must block all converters')
    assert marker.exists()
    monkeypatch.setattr(office_process, '_group_running', lambda _: False)
    with jobs.conversion_slot(wait_seconds=0):
        assert not marker.exists()


def test_cleanup_deadline_leaves_marker_and_has_bounded_wait(monkeypatch, tmp_path):
    from app import office_process
    monkeypatch.setenv('EPC_DATA_DIR', str(tmp_path / 'runtime'))
    monkeypatch.setattr(office_process, '_group_running', lambda _: True)
    monkeypatch.setattr(office_process, '_CLEANUP_SECONDS', .1)
    tick = monotonic()
    with pytest.raises(office_process.OfficeCleanupPending):
        office_process.run_office([sys.executable, '-c', 'pass'], timeout=5)
    assert monotonic() - tick < 2
    assert office_process._marker().exists()


def test_diagnostics_are_bounded_and_current_exit_status_is_kept(monkeypatch, tmp_path):
    from app import office_process
    monkeypatch.setenv('EPC_DATA_DIR', str(tmp_path / 'runtime'))
    result = office_process.run_office([sys.executable, '-c',
        "import sys; print('x'*40000); sys.stderr.write('y'*40000); sys.exit(7)"], timeout=5)
    assert result.returncode == 7
    assert len(result.stdout) <= 16 * 1024 and len(result.stderr) <= 16 * 1024
    assert not office_process._marker().exists()


def test_timeout_preserves_an_unrelated_process(monkeypatch, tmp_path):
    from app import office_process
    monkeypatch.setenv('EPC_DATA_DIR', str(tmp_path / 'runtime'))
    unrelated = subprocess.Popen([sys.executable, '-c', 'import time; time.sleep(10)'])
    try:
        with pytest.raises(subprocess.TimeoutExpired):
            office_process.run_office([sys.executable, '-c', 'import time; time.sleep(10)'], timeout=.1)
        assert unrelated.poll() is None
        assert not office_process._marker().exists()
    finally:
        unrelated.kill(); unrelated.wait(timeout=5)


def test_monthly_timeout_keeps_docx_and_specific_message(monkeypatch, tmp_path):
    from app import monthly_report_docx as output
    from app.monthly_report_model import synthetic_draft, synthetic_profiles, ReportPeriod
    monkeypatch.setenv('EPC_DATA_DIR', str(tmp_path / 'runtime'))
    monkeypatch.setattr(converter, 'OUTPUT_DIR', tmp_path)
    def fail(_):
        raise converter.PDFConversionTimeout('PDF conversion took too long. The conversion was stopped.')
    monkeypatch.setattr(converter, 'convert_to_pdf', fail)
    draft = synthetic_draft(synthetic_profiles()[0], ReportPeriod(2026, 9))
    package = output.generate_report(draft, acknowledged_fingerprint=draft.fingerprint)
    assert package.docx.startswith(b'PK') and package.pdf is None
    assert 'too long' in package.pdf_error and 'DOCX is ready' in package.pdf_error
    assert 'TimeoutExpired' not in package.pdf_error
    assert Document(BytesIO(package.docx)).paragraphs


@requires_libreoffice
def test_real_writer_timeout_stops_its_process_group(monkeypatch, tmp_path):
    from app import office_process
    monkeypatch.setenv('EPC_DATA_DIR', str(tmp_path / 'runtime'))
    monkeypatch.setattr(converter, 'OUTPUT_DIR', tmp_path)
    monkeypatch.setattr(converter, 'PDF_BACKEND', 'libreoffice')
    source = tmp_path / 'slow.docx'
    doc = Document()
    for index in range(2500):
        doc.add_paragraph(f'Synthetic conversion timeout paragraph {index}.')
    doc.save(source)
    run = office_process.run_office
    pids = []
    record = office_process._record
    def capture(pid):
        pids.append(pid); record(pid)
    monkeypatch.setattr(office_process, '_record', capture)
    monkeypatch.setattr(office_process, 'run_office', lambda command, **kw: run(command, timeout=.03))
    with pytest.raises(converter.PDFConversionError, match='too long'):
        converter.convert_to_pdf(source)
    assert pids and all(not office_process._group_running(pid) for pid in pids)
    assert not office_process._marker().exists()
    assert not (tmp_path / 'slow.pdf').exists()
    with jobs.conversion_slot(wait_seconds=0):
        assert source.exists()


def test_quarantine_is_visible_to_another_python_process(monkeypatch, tmp_path):
    from app import office_process
    monkeypatch.setenv('EPC_DATA_DIR', str(tmp_path / 'runtime'))
    child = subprocess.Popen([sys.executable, '-c', 'import time; time.sleep(10)'], start_new_session=True)
    try:
        office_process._record(child.pid)
        script = '''from app.monthly_report_render_jobs import conversion_slot, RenderBusy
try:
    with conversion_slot(wait_seconds=0):
        raise SystemExit('unexpected conversion admission')
except RenderBusy:
    print('blocked')
'''
        result = subprocess.run([sys.executable, '-c', script], capture_output=True, text=True, timeout=5)
        assert result.returncode == 0 and result.stdout.strip() == 'blocked'
    finally:
        child.kill(); child.wait(timeout=5)
    with jobs.conversion_slot(wait_seconds=0):
        assert not office_process._marker().exists()


def test_invalid_process_marker_blocks_but_previous_boot_does_not(monkeypatch, tmp_path):
    from app import office_process
    monkeypatch.setenv('EPC_DATA_DIR', str(tmp_path / 'runtime'))
    marker = office_process._marker()
    marker.write_text('x' * 2000)
    with pytest.raises(jobs.RenderBusy, match='could not be checked'):
        with jobs.conversion_slot(wait_seconds=0):
            pytest.fail('invalid cleanup state must not be ignored')
    import json
    marker.write_text(json.dumps({'pid': 123456, 'boot': 'previous-container'}))
    monkeypatch.setattr(office_process, '_group_running', lambda _: True)
    monkeypatch.setattr(office_process, '_boot_id', lambda: 'current-container')
    with jobs.conversion_slot(wait_seconds=0):
        assert not marker.exists()


def test_same_named_sources_publish_independent_current_outputs(monkeypatch, tmp_path):
    monkeypatch.setenv('EPC_DATA_DIR', str(tmp_path / 'runtime'))
    output = tmp_path / 'output'; output.mkdir()
    monkeypatch.setattr(converter, 'OUTPUT_DIR', output)
    monkeypatch.setattr(converter, 'PDF_BACKEND', 'libreoffice')
    executable = tmp_path / 'current-pdf'
    executable.write_text(f'''#!{sys.executable}
from pathlib import Path
import sys
source = Path(sys.argv[-1])
out = Path(sys.argv[sys.argv.index('--outdir')+1])
(out / (source.stem + '.pdf')).write_bytes(b'%PDF-1.7 ' + source.read_bytes())
'''); executable.chmod(0o700)
    monkeypatch.setattr(converter.shutil, 'which', lambda _: str(executable))
    legacy = output / 'report.pdf'; legacy.write_bytes(b'%PDF-1.7 old owner')
    sources = []
    for folder in ('one', 'two'):
        directory = tmp_path / folder; directory.mkdir()
        source = directory / 'report.docx'; source.write_bytes(folder.encode())
        sources.append(source)
    first, second = (converter.convert_to_pdf(source) for source in sources)
    assert first != second and first != legacy and second != legacy
    assert first.read_bytes() == b'%PDF-1.7 one'
    assert second.read_bytes() == b'%PDF-1.7 two'
    assert legacy.read_bytes() == b'%PDF-1.7 old owner'
    assert not list(output.glob('.office-*'))
