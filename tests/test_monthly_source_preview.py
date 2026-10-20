"""Only invented DOCX fixtures. Fake conversion is not production fidelity proof."""
from contextlib import contextmanager
from dataclasses import replace
from hashlib import sha256
from io import BytesIO
from pathlib import Path
import json
import subprocess
from zipfile import ZipFile, ZIP_DEFLATED

from docx import Document
import fitz
import pytest

from app import memory
from app import monthly_report_source_render_qa as p
from app import monthly_report_testing as qa

TOKEN = 'a' * 32
OWNER = 'b' * 32


@pytest.fixture
def source():
    document = Document()
    document.add_paragraph('Invented source: estimated cost $3200; NOT approved.')
    output = BytesIO()
    document.save(output)
    return output.getvalue()


@pytest.fixture
def isolated(tmp_path, monkeypatch):
    monkeypatch.setenv('EPC_DATA_DIR', str(tmp_path))
    with qa.isolated(TOKEN):
        memory._data_dir()
        yield tmp_path


def request(raw):
    epoch = p.retire(OWNER)
    return p.new_request(OWNER, epoch, raw, upload_id='fresh-file-1', label='Synthetic original', kind='original')


def fake_render(path, directory):
    assert path.name == 'render-input.docx'
    with fitz.open() as doc:
        doc.new_page().insert_text((72, 72), 'Synthetic preview only')
        output = directory / 'render-input.pdf'
        doc.save(output)
    return output, {'version': 'LibreOffice 24.2.7.2 synthetic stub, NOT render proof'}, 1


def rewrite(raw, mutate):
    with ZipFile(BytesIO(raw)) as source:
        parts = {name: source.read(name) for name in source.namelist()}
    mutate(parts)
    output = BytesIO()
    with ZipFile(output, 'w', ZIP_DEFLATED) as package:
        for name, data in parts.items():
            package.writestr(name, data)
    return output.getvalue()


def test_uses_real_passive_pipeline_and_distinct_immutable_receipts(isolated, source, monkeypatch):
    monkeypatch.setattr(p, '_render', fake_render)
    req = request(source)
    result = p.execute(req, source, upload_id=req.upload_id)
    assert result['status'] == 'rendered' and result['pages'] == 1
    assert result['source_sha256'] == sha256(source).hexdigest()
    assert result['render_input_sha256'] != result['source_sha256']
    assert p.artifact(req, 'source.docx') == source
    passive = p.artifact(req, 'render-input.docx')
    assert sha256(passive).hexdigest() == result['render_input_sha256']
    assert Document(BytesIO(passive)).paragraphs[0].text == Document(BytesIO(source)).paragraphs[0].text
    assert p.artifact(req, 'source-preview.pdf').startswith(b'%PDF-')
    assert json.loads(p.artifact(req, 'receipt.json')) == result
    assert not (isolated / 'monthly_reports').exists()
    assert list(p._root().rglob('*.docx'))
    assert not list(p._root().rglob('.conversion-*'))


def test_no_live_namespace_even_if_token_is_forged(tmp_path, monkeypatch, source):
    monkeypatch.setenv('EPC_DATA_DIR', str(tmp_path))
    with pytest.raises(p.SourcePreviewError, match='active isolated'):
        p.retire(OWNER)
    with qa.isolated(TOKEN):  # Lease not bound yet.
        with pytest.raises(p.SourcePreviewError, match='active isolated'):
            p.retire(OWNER)
    assert not list(tmp_path.iterdir())


def test_replay_changed_upload_and_a_b_a_are_rejected(isolated, source, monkeypatch):
    monkeypatch.setattr(p, '_render', fake_render)
    req = request(source)
    with pytest.raises(p.SourcePreviewError, match='stale'):
        p.execute(req, source, upload_id='another-file')
    with pytest.raises(p.SourcePreviewError, match='stale'):
        p.execute(req, source + b'x', upload_id=req.upload_id)
    p.execute(req, source, upload_id=req.upload_id)
    with pytest.raises(p.SourcePreviewError, match='already attempted'):
        p.execute(req, source, upload_id=req.upload_id)
    p.retire(OWNER)
    with pytest.raises(p.SourcePreviewError, match='stale'):
        p.artifact(req, 'source.docx')
    fresh = request(source)
    assert fresh.epoch != req.epoch
    p.execute(fresh, source, upload_id=fresh.upload_id)
    with pytest.raises(p.SourcePreviewError, match='stale'):
        p.execute(req, source, upload_id=req.upload_id)


def test_two_sessions_and_two_browsers_cannot_reuse_results(isolated, source, monkeypatch):
    monkeypatch.setattr(p, '_render', fake_render)
    req = request(source)
    p.execute(req, source, upload_id=req.upload_id)
    p.retire('c' * 32)
    with pytest.raises(p.SourcePreviewError):
        p.artifact(replace(req, owner='c' * 32), 'source.docx')
    with qa.isolated('d' * 32):
        memory._data_dir()
        with pytest.raises(p.SourcePreviewError, match='session changed'):
            p.artifact(req, 'source.docx')
    assert p.artifact(req, 'source.docx') == source


@pytest.mark.parametrize('payload', ['macro', 'unreachable-active', 'external-image', 'external-template', 'active-shape'])
def test_active_or_external_package_never_reaches_renderer(isolated, source, monkeypatch, payload):
    called = []
    monkeypatch.setattr(p, '_render', lambda *_: called.append(True))
    def mutate(parts):
        if payload == 'macro':
            parts['[Content_Types].xml'] = parts['[Content_Types].xml'].replace(
                b'wordprocessingml.document.main+xml', b'ms-word.document.macroEnabled.main+xml')
        elif payload == 'unreachable-active':
            parts['word/vbaProject.bin'] = b'invented payload'
        elif payload == 'active-shape':
            parts['word/document.xml'] = parts['word/document.xml'].replace(b'<w:body>', b'<w:body><w:object/>')
        else:
            kind = 'image' if payload == 'external-image' else 'attachedTemplate'
            parts['word/_rels/document.xml.rels'] = parts['word/_rels/document.xml.rels'].replace(
                b'</Relationships>', f'<Relationship Id="bad" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/{kind}" Target="https://invalid.example/resource" TargetMode="External"/></Relationships>'.encode())
    raw = rewrite(source, mutate)
    req = request(raw)
    result = p.execute(req, raw, upload_id=req.upload_id)
    assert result['status'] == 'failed' and result['stage'] == 'passive package validation'
    assert not called and 'pdf_sha256' not in result
    assert p.artifact(req, 'source.docx') == raw


def test_macro_in_unreachable_content_type_is_rejected(isolated, source, monkeypatch):
    monkeypatch.setattr(p, '_render', lambda *_: pytest.fail('active package reached renderer'))
    raw = rewrite(source, lambda parts: parts.update({'[Content_Types].xml': parts['[Content_Types].xml'].replace(
        b'</Types>', b'<Override PartName="/unused.bin" ContentType="application/vnd.ms-office.vbaProject"/></Types>')}))
    req = request(raw)
    assert p.execute(req, raw, upload_id=req.upload_id)['status'] == 'failed'


def test_failure_receipt_is_frozen_and_content_diagnostics_are_not_logged(isolated, source, monkeypatch, caplog):
    count = []
    def fail(*_):
        count.append(1)
        raise RuntimeError('PRIVATE SOURCE CONTENT THAT MUST NEVER BE LOGGED')
    monkeypatch.setattr(p, '_render', fail)
    req = request(source)
    result = p.execute(req, source, upload_id=req.upload_id)
    receipt = p.artifact(req, 'receipt.json')
    assert result['status'] == 'failed' and result['stage'] == 'LibreOffice conversion'
    assert 'PRIVATE SOURCE' not in receipt.decode() + caplog.text
    assert p.artifact(req, 'render-input.docx')
    with pytest.raises(p.SourcePreviewError):
        p.execute(req, source, upload_id=req.upload_id)
    assert count == [1] and p.artifact(req, 'receipt.json') == receipt
    assert not list(p._root().rglob('.conversion-*'))


def test_busy_is_one_failure_not_automatic_retry(isolated, source, monkeypatch):
    @contextmanager
    def busy(**_):
        raise p.jobs.RenderBusy('busy')
        yield
    monkeypatch.setattr(p.jobs, 'work_slot', busy)
    req = request(source)
    result = p.execute(req, source, upload_id=req.upload_id)
    assert result['status'] == 'failed' and result['stage'] == 'admission'
    assert 'busy' in result['error']
    assert not (p._folder(OWNER) / 'source.docx').exists()


def test_cancel_during_render_cannot_publish_old_result(isolated, source, monkeypatch):
    def cancel(path, directory):
        result = fake_render(path, directory)
        p.retire(OWNER)
        return result
    monkeypatch.setattr(p, '_render', cancel)
    req = request(source)
    with pytest.raises(p.SourcePreviewError, match='stale'):
        p.execute(req, source, upload_id=req.upload_id)
    assert not list(p._root().rglob('*.pdf'))
    assert not list(p._root().rglob('receipt.json'))


def test_reset_during_render_cannot_recreate_removed_namespace(isolated, source, monkeypatch):
    def reset(path, directory):
        result = fake_render(path, directory)
        qa.reset(isolated)
        return result
    monkeypatch.setattr(p, '_render', reset)
    req = request(source)
    with pytest.raises(OSError, match='reset'):
        p.execute(req, source, upload_id=req.upload_id)
    assert {item.name for item in (isolated / qa._BUCKET).iterdir()} == {'.source-preview-lock'}


def test_byte_limit_and_storage_budget_fail_before_converter(isolated, source, monkeypatch):
    monkeypatch.setattr(p, 'MAX_BYTES', len(source) - 1)
    with pytest.raises(p.SourcePreviewError, match='32 MiB'):
        request(source)
    monkeypatch.setattr(p, 'MAX_BYTES', 32 * 1024 * 1024)
    req = request(source)
    monkeypatch.setattr(p, 'MAX_STORE_BYTES', 1)
    with pytest.raises(p.SourcePreviewError, match='storage is full'):
        p.execute(req, source, upload_id=req.upload_id)
    assert not (p._folder(OWNER) / 'attempt.json').exists()


def test_changed_artifacts_and_symlink_are_rejected(isolated, source, monkeypatch):
    monkeypatch.setattr(p, '_render', fake_render)
    req = request(source)
    p.execute(req, source, upload_id=req.upload_id)
    path = p._folder(OWNER) / 'source.docx'
    assert path.stat().st_mode & 0o222 == 0
    path.chmod(0o600)  # Simulate out-of-band disk corruption, not an app edit.
    path.write_bytes(b'changed')
    with pytest.raises(p.SourcePreviewError, match='checksum'):
        p.artifact(req, 'source.docx')
    path.unlink()
    path.symlink_to(isolated / 'not-a-source')
    with pytest.raises(p.SourcePreviewError, match='unavailable'):
        p.artifact(req, 'source.docx')
    with pytest.raises(p.SourcePreviewError, match='unavailable'):
        p.retire(OWNER)


def test_timeout_retains_first_attempt(isolated, source, monkeypatch):
    def timeout(*_):
        raise subprocess.TimeoutExpired(['office'], 120)
    monkeypatch.setattr(p, '_render', timeout)
    req = request(source)
    result = p.execute(req, source, upload_id=req.upload_id)
    assert result['status'] == 'failed' and 'timed out' in result['error']
    assert 'pdf_sha256' not in result


@pytest.mark.parametrize('result_kind', ['valid', 'failure', 'invalid', 'too-many-pages'])
def test_renderer_uses_qa_output_existing_lease_and_process_cleanup(isolated, tmp_path, monkeypatch, result_kind):
    calls = []
    monkeypatch.setattr(p.shutil, 'which', lambda _: '/usr/bin/libreoffice')
    monkeypatch.setattr(p, '_renderer', lambda _: {'version': 'LibreOffice 24.2.7.2 synthetic stub'})
    @contextmanager
    def slot(**kwargs):
        calls.append(('slot', kwargs))
        yield
    monkeypatch.setattr(p.jobs, 'conversion_slot', slot)
    directory = p._folder(OWNER) / 'bounded-job'
    directory.mkdir()
    source = p._folder(OWNER) / 'render-input.docx'
    source.write_bytes(b'synthetic stub input')
    def office(command, *, timeout):
        calls.append(('office', command, timeout))
        output = directory / 'render-input.pdf'
        if result_kind == 'invalid':
            output.write_bytes(b'not a PDF')
        elif result_kind != 'failure':
            with fitz.open() as pdf:
                for _ in range(2 if result_kind == 'too-many-pages' else 1):
                    pdf.new_page()
                pdf.save(output)
        return subprocess.CompletedProcess(command, 1 if result_kind == 'failure' else 0, 'PRIVATE SOURCE', '')
    monkeypatch.setattr(p, 'run_office', office)
    if result_kind == 'too-many-pages':
        monkeypatch.setattr(p, 'MAX_PAGES', 1)
    if result_kind == 'valid':
        result, _, pages = p._render(source, directory)
        assert result.parent == directory and pages == 1
    else:
        with pytest.raises((p.SourcePreviewError, RuntimeError)):
            p._render(source, directory)
    assert calls[0] == ('slot', {'wait_seconds': 0})
    command, timeout = calls[1][1:]
    assert timeout == 120
    assert command[:3] == [p.sys.executable, '-c', p._LIMIT_EXEC]
    assert 'RLIMIT_FSIZE' in command[2] and 'RLIMIT_AS' in command[2] and 'RLIMIT_CPU' in command[2]
    assert command[command.index('--outdir') + 1] == str(directory)
    assert str(source.resolve()) == command[-1]
    assert (directory / 'profile/user/registrymodifications.xcu').is_file()


@pytest.mark.parametrize('version', ['LibreOffice 26.8.0.0.alpha0 fake', 'LibreOffice 7.6.0.0 fake', 'not office'])
def test_non_production_renderer_rejected_before_fonts_or_conversion(monkeypatch, version):
    monkeypatch.setattr(p, 'run_office', lambda *_a, **_kw: subprocess.CompletedProcess([], 0, version, ''))
    monkeypatch.setattr(p.subprocess, 'run', lambda *_a, **_kw: pytest.fail('wrong engine continued'))
    with pytest.raises(p.SourcePreviewError, match='24.2'):
        p._renderer('/invented/libreoffice')


def test_passive_transform_inventory_is_specific(isolated, source, monkeypatch):
    monkeypatch.setattr(p, '_render', fake_render)
    req = request(source)
    receipt = p.execute(req, source, upload_id=req.upload_id)
    changes = receipt['changed_package_parts']
    assert any(row['part'] == 'docProps/core.xml' and row['render_sha256'] is None for row in changes)
    assert all(set(row) == {'part', 'source_sha256', 'render_sha256'} for row in changes)


def test_office_diagnostics_preserve_evidence_without_private_text():
    result = subprocess.CompletedProcess([], 0, 'PRIVATE SOURCE warning', 'PRIVATE PATH')
    evidence = p._diagnostics(result)
    assert evidence['stdout']['present'] and evidence['stderr']['present']
    assert evidence['stdout']['sha256'] == sha256(result.stdout.encode()).hexdigest()
    assert evidence['stderr']['decoded_capture_bytes'] == len(result.stderr)
    assert 'PRIVATE' not in json.dumps(evidence)


@pytest.mark.skipif(__import__('os').environ.get('GITHUB_ACTIONS') != 'true',
                    reason='Production 24.2 integration runs in GitHub CI; local renderer execution is unavailable here.')
def test_ci_real_24_2_source_render_and_cleanup(isolated, source):
    """Mandatory on GitHub's installed production renderer, never a mock pass."""
    req = request(source)
    result = p.execute(req, source, upload_id=req.upload_id)
    assert result['status'] == 'rendered', result
    assert result['renderer']['version'].startswith('LibreOffice 24.2.')
    assert result['renderer']['fonts'] and result['renderer']['font_inventory']
    assert 'office_diagnostics' in result['renderer']
    assert p.artifact(req, 'source.docx') == source
    passive = p.artifact(req, 'render-input.docx')
    assert sha256(passive).hexdigest() == result['render_input_sha256']
    pdf = p.artifact(req, 'source-preview.pdf')
    assert sha256(pdf).hexdigest() == result['pdf_sha256']
    with fitz.open(stream=pdf, filetype='pdf') as rendered:
        assert rendered.is_pdf and len(rendered) == result['pages'] == 1
        assert '3200' in rendered[0].get_text()  # Source inspection is NOT price approval/removal.
    assert not list(p._root().rglob('.conversion-*'))
    assert not (p.library._root() / 'render-jobs/.active-office.json').exists()
    assert not (isolated / 'monthly_reports').exists()
    p.retire(OWNER)
    assert not list(p._root().rglob('*.docx')) and not list(p._root().rglob('*.pdf'))


def test_identical_bytes_fresh_requests_are_independent_renders_not_cache_hits(isolated, source, monkeypatch):
    calls = []
    def independent(path, directory):
        calls.append(sha256(path.read_bytes()).hexdigest())
        output = directory / 'render-input.pdf'
        with fitz.open() as document:
            document.new_page().insert_text((72, 72), f'Independent synthetic render {len(calls)}')
            document.save(output)
        return output, {'version': 'LibreOffice 24.2.7.2 synthetic stub'}, 1
    monkeypatch.setattr(p, '_render', independent)
    first = request(source)
    result_a = p.execute(first, source, upload_id=first.upload_id)
    receipt_a = p.artifact(first, 'receipt.json')
    with pytest.raises(p.SourcePreviewError, match='already attempted'):
        p.execute(first, source, upload_id=first.upload_id)
    assert len(calls) == 1
    second = request(source)  # Deliberate fresh selection/attempt, same exact bytes.
    result_b = p.execute(second, source, upload_id=second.upload_id)
    assert len(calls) == 2 and calls[0] == calls[1]
    assert first.request_id != second.request_id and first.epoch != second.epoch
    assert result_a['source_sha256'] == result_b['source_sha256']
    assert result_a['pdf_sha256'] != result_b['pdf_sha256']
    assert receipt_a != p.artifact(second, 'receipt.json')


def test_cookie_churn_cannot_escape_global_budget(isolated, source, monkeypatch):
    monkeypatch.setattr(p, '_render', fake_render)
    first = request(source)
    p.execute(first, source, upload_id=first.upload_id)
    before = p.artifact(first, 'receipt.json')
    monkeypatch.setattr(p, 'MAX_GLOBAL_BYTES', p.RESERVE_BYTES + p.SCRATCH_RESERVE_BYTES + 1)
    with qa.isolated('c' * 32):
        memory._data_dir()
        other = request(source)
        with pytest.raises(p.SourcePreviewError, match='storage is full'):
            p.execute(other, source, upload_id=other.upload_id)
        assert not (p._folder(OWNER) / 'attempt.json').exists()
    assert p.artifact(first, 'receipt.json') == before


def test_low_free_disk_and_owner_count_fail_closed(isolated, source, monkeypatch):
    from types import SimpleNamespace
    req = request(source)
    monkeypatch.setattr(p.shutil, 'disk_usage', lambda _p: SimpleNamespace(free=1))
    with pytest.raises(p.SourcePreviewError, match='storage is full'):
        p.execute(req, source, upload_id=req.upload_id)
    monkeypatch.setattr(p, 'MAX_OWNERS', 1)
    with pytest.raises(p.SourcePreviewError, match='storage is full'):
        p.retire('c' * 32)
    assert not (p._root() / ('c' * 32)).exists()


def test_expiry_removes_only_diagnostics_and_preserves_bounded_active_lease(isolated, source, monkeypatch):
    import os
    req = request(source)
    folder = p._folder(OWNER)
    marker = folder / 'active.json'
    active = json.loads(marker.read_text())
    active['deadline'] = p.time.time() + p.ABANDONED_SECONDS
    marker.write_text(json.dumps(active))
    (folder / 'attempt.json').write_text('{}')
    (folder / 'source.docx').write_bytes(source)
    report = p.library._root() / 'keep-editor.json'
    report.write_text('keep me')
    old = p.time.time() - p.RETENTION_SECONDS - 1
    os.utime(marker, (old, old))
    with p._storage_lock():
        rows = p._catalog(prune=True)
    assert folder.exists() and rows[0][1] >= p.RESERVE_BYTES
    active['deadline'] = p.time.time() - 1
    marker.write_text(json.dumps(active)); os.utime(marker, (old, old))
    with p._storage_lock():
        p._catalog(prune=True)
    assert not folder.exists() and report.read_text() == 'keep me'
    assert not (isolated / 'monthly_reports').exists()


def test_reload_recovery_is_browser_local_and_incarnation_bound(isolated, source, monkeypatch):
    monkeypatch.setattr(p, '_render', fake_render)
    old = request(source)
    p.execute(old, source, upload_id=old.upload_id)
    owner = 'c' * 32
    epoch = p.retire(owner)  # Simulate the new Streamlit session after reload.
    snapshot = p.earlier_previews(owner)
    assert snapshot == ((OWNER, old.epoch),)
    fresh = request(source)  # Old same-source tab reopened before stale cleanup.
    p.clear_earlier(owner, epoch, snapshot)
    assert p._active(OWNER)['epoch'] == fresh.epoch
    latest = p.earlier_previews(owner)
    p.clear_earlier(owner, epoch, latest)
    assert not (p._root() / OWNER).exists()
    assert p._active(owner)['epoch'] == epoch


def test_recovery_cannot_remove_pending_attempt_or_other_browser(isolated, source, monkeypatch):
    def during(path, directory):
        current = 'c' * 32
        epoch = p.retire(current)
        assert p.earlier_previews(current) == ()
        p.clear_earlier(current, epoch, ((OWNER, req.epoch),))
        assert path.exists()
        return fake_render(path, directory)
    monkeypatch.setattr(p, '_render', during)
    req = request(source)
    assert p.execute(req, source, upload_id=req.upload_id)['status'] == 'rendered'
    with qa.isolated('d' * 32):
        memory._data_dir()
        epoch = p.retire('e' * 32)
        assert p.earlier_previews('e' * 32) == ()
        p.clear_earlier('e' * 32, epoch, ((OWNER, req.epoch),))
    assert p.artifact(req, 'source.docx') == source


def test_oversized_inventory_fails_before_render_with_downloadable_failure(isolated, source, monkeypatch):
    def extra(parts):
        for i in range(2600):
            parts[f'word/unused/{i}-' + 'x' * 310 + '.bin'] = b'ordinary unused bytes'
    raw = rewrite(source, extra)
    monkeypatch.setattr(p, '_render', lambda *_: pytest.fail('over-budget inventory rendered'))
    req = request(raw)
    result = p.execute(req, raw, upload_id=req.upload_id)
    assert result['status'] == 'failed' and 'inventory' in result['error']
    assert len(p.artifact(req, 'receipt.json')) < p.MAX_RECEIPT_BYTES
    assert p.artifact(req, 'source.docx') == raw
    assert p.artifact(req, 'render-input.docx')
    assert 'pdf_sha256' not in result


def test_disappearing_accounting_file_does_not_consume_request(isolated, source, monkeypatch):
    monkeypatch.setattr(p, '_render', fake_render)
    req = request(source)
    p.retire('c' * 32)
    vanished = p._folder('c' * 32) / 'vanished.tmp'
    vanished.write_bytes(b'invented temporary')
    real = Path.stat
    count = []
    def racing(path, *args, **kwargs):
        if path == vanished:
            count.append(1)
            if len(count) == 2:
                vanished.unlink()
                raise FileNotFoundError(str(path))
        return real(path, *args, **kwargs)
    monkeypatch.setattr(Path, 'stat', racing)
    assert p.execute(req, source, upload_id=req.upload_id)['status'] == 'rendered'
    assert count and p.artifact(req, 'receipt.json')


def test_capacity_recovery_needs_no_new_owner_and_rejects_reopened_snapshot(isolated, source):
    old = request(source)
    snapshot = p.earlier_previews(None)
    p.clear_earlier(None, None, snapshot)
    assert not (p._root() / OWNER).exists()
    newer = request(source)
    p.clear_earlier(None, None, snapshot)
    assert p._active(OWNER)['epoch'] == newer.epoch


def test_expired_deadline_does_not_make_live_execution_safe_to_delete(isolated, source):
    import os
    req = request(source)
    folder = p._folder(OWNER)
    active = folder / 'active.json'
    state = json.loads(active.read_text()); state['deadline'] = p.time.time() - 1
    active.write_text(json.dumps(state)); (folder / 'attempt.json').write_text('{}')
    old = p.time.time() - p.RETENTION_SECONDS - 1
    os.utime(active, (old, old))
    with p._execution_lease(folder):
        with p._storage_lock():
            assert p._catalog(prune=True)
        assert folder.exists()
        p.clear_earlier(None, None, ((OWNER, req.epoch),))
        assert folder.exists()
    with p._storage_lock():
        p._catalog(prune=True)
    assert not folder.exists()


def test_expired_validation_never_starts_converter(isolated, source, monkeypatch):
    real = p._safe_source
    def expire(path):
        result = real(path)
        active = p._folder(OWNER) / 'active.json'
        state = json.loads(active.read_text()); state['deadline'] = p.time.time() - 1
        active.write_text(json.dumps(state))
        return result
    monkeypatch.setattr(p, '_safe_source', expire)
    monkeypatch.setattr(p, '_render', lambda *_: pytest.fail('expired job started converter'))
    req = request(source)
    with pytest.raises(p.SourcePreviewError, match='expired'):
        p.execute(req, source, upload_id=req.upload_id)
    assert not list(p._root().rglob('*.pdf'))


def test_global_scan_is_bounded_and_rejects_symlinked_peer(isolated, source, monkeypatch):
    request(source)
    monkeypatch.setattr(p, 'MAX_SCAN_ENTRIES', 1)
    with p._storage_lock(), pytest.raises(p.SourcePreviewError, match='inventory'):
        p._catalog()
    monkeypatch.setattr(p, 'MAX_SCAN_ENTRIES', 4096)
    link = p._root().parents[2] / ('f' * 32)
    link.symlink_to(isolated / 'outside', target_is_directory=True)
    with p._storage_lock(), pytest.raises(p.SourcePreviewError, match='unavailable'):
        p._catalog(prune=True)
    assert link.is_symlink()


def test_missing_owner_does_not_hide_later_valid_siblings_from_quota(isolated, source):
    owners = {char * 32 for char in '23456'}
    for owner in owners:
        p.retire(owner)
        (p._folder(owner) / 'source.docx').write_bytes(source)
    first = next(folder for folder in p._entries(p._root()) if folder.name in owners)
    (first / 'active.json').unlink()
    (first / 'source.docx').unlink()  # Model interrupted empty-owner initialization.
    with p._storage_lock():
        rows = p._catalog()
    assert {folder.name for folder, _size, _ in rows} == owners - {first.name}
    assert all(size >= len(source) for _folder, size, _ in rows)
    (first / 'source.docx').write_bytes(source)  # Orphaned payload must instead fail closed.
    with p._storage_lock(), pytest.raises(p.SourcePreviewError, match='metadata'):
        p._catalog()
    assert not p.session_current(first.name, 'a' * 32)
    missing = p._root() / ('1' * 32)
    assert not p.session_current('1' * 32, 'a' * 32) and not missing.exists()
