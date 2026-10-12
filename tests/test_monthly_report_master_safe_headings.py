"""Fact-free shape regression for a Word master's 7/9-column tables."""
from io import BytesIO
from zipfile import ZipFile
from docx import Document

from app.monthly_report_import import inspect_docx
from app.monthly_report_master_tables import _heading, from_inspection
from app.monthly_report_model import (
    Facility, ReportDraft, ReportPeriod, ReportProfile, ReportTable,
    ResolvedBlock, default_sections,
)
from app.monthly_report_native_layout import build_native_docx


def test_current_safe_work_order_months_and_engineering_headers_render_in_native_word(tmp_path):
    old = tmp_path / "invented-other-site.docx"
    source = Document()
    source.add_heading("MONTHLY ACTIVITY SUMMARY", 1)
    months = ("WO Received", "Jan", "Feb", "Mar", "Apr", "May", "June")
    grid = source.add_table(rows=2, cols=7)
    for c, value in zip(grid.rows[0].cells, months):
        c.text = value
    for c in grid.rows[1].cells:
        c.text = "OTHER PRIVATE SITE - NEVER REUSE"

    source.add_heading("PENDING & DECLINED PROPOSALS", 1)
    engineering = ("#", "DESCRIPTION", "CAMPUS", "BUILDING", "DESIGNATION",
                   "CAPACITY", "UNITS", "ESTIMATED END OF USEFUL LIFE", "O&M OPTION")
    grid = source.add_table(rows=2, cols=9)
    for c, value in zip(grid.rows[0].cells, engineering):
        c.text = value
    for c in grid.rows[1].cells:
        c.text = "OTHER PRIVATE SITE - NEVER REUSE"
    source.save(old)

    found = from_inspection(inspect_docx(old))
    months_shape, = (s for s in found if s.slot == "work_orders")
    engineering_shape, = (s for s in found if s.slot == "proposals")
    assert (months_shape.columns, engineering_shape.columns) == (months, engineering)
    assert (months_shape.physical_width, engineering_shape.physical_width) == (7, 9)
    assert months_shape.indices == tuple(range(7))
    assert engineering_shape.indices == tuple(range(9))

    profile = ReportProfile("Invented New Client", "qa", "QA Facility",
                            (Facility("qa", "QA Facility"),))
    month_data = ("3", "1", "2", "1", "0", "0", "0")
    engineering_data = ("1", "Synthetic pump", "QA Campus", "QA Building",
                        "PUMP-QA-01", "10", "Count", "2041", "Pending")
    draft = ReportDraft(profile, ReportPeriod(2026, 10), "QA Operator",
                        default_sections(), (
        ResolvedBlock("work_orders", "This month", extra_tables=(
            ReportTable(months, (month_data,), months_shape.reference),)),
        ResolvedBlock("proposals", "This month", extra_tables=(
            ReportTable(engineering, (engineering_data,), engineering_shape.reference),)),
    ), synthetic=True)
    raw = build_native_docx(draft, old, master=True)
    tables = Document(BytesIO(raw)).tables
    assert [(len(t.columns), len(t.rows)) for t in tables] == [(7, 2), (9, 2)]
    assert tuple(c.text for c in tables[0].rows[0].cells) == months
    assert tuple(c.text for c in tables[1].rows[0].cells) == engineering
    assert tuple(c.text for c in tables[0].rows[1].cells) == month_data
    assert tuple(c.text for c in tables[1].rows[1].cells) == engineering_data
    with ZipFile(BytesIO(raw)) as archive:
        package = b"".join(archive.read(name) for name in archive.namelist())
    assert b"OTHER PRIVATE SITE" not in package
    assert b"QA Campus" in package and b"PUMP-QA-01" in package


def test_month_headings_allowed_only_in_work_order_slot_and_other_client_headers_still_rejected():
    assert _heading("May", 4, set(), slot="work_orders") == "May"
    assert _heading("December", 6, set(), slot="work_orders") == "December"
    assert _heading("May", 4, set(), slot="proposals") == "Detail 5"
    assert _heading("Private Medical Center", 3, set(), slot="proposals") == "Detail 4"
    assert _heading("Finance Total Cost", 4, set(), slot="proposals") == "Detail 5"
    assert _heading("MSA", 3, set(), slot="subcontractor_matrix") == "MSA"
    assert _heading("Discipline", 0, set(), slot="subcontractor_matrix") == "Discipline"



def test_old_site_header_named_like_one_month_cannot_escape_unverified_grid(tmp_path):
    # The monthly word alone is ambiguous. A broken Jan-Jun sequence must
    # fail closed even though several remaining headings look familiar.
    source = Document()
    source.add_heading("MONTHLY ACTIVITY SUMMARY", 1)
    columns = ("WO Received", "Jan", "Feb", "Prior Client Campus", "Apr", "May", "June")
    grid = source.add_table(rows=2, cols=7)
    for cell, value in zip(grid.rows[0].cells, columns):
        cell.text = value
    for cell in grid.rows[1].cells:
        cell.text = "PRIVATE SITE OLD DATA"
    path = tmp_path / "unverified-months.docx"
    source.save(path)
    work, = (shape for shape in from_inspection(inspect_docx(path))
             if shape.slot == "work_orders")
    assert work.columns == (
        "WO Received", "Detail 2", "Detail 3", "Detail 4",
        "Detail 5", "Detail 6", "Detail 7")
    assert "Prior Client Campus" not in str(work)
