"""The browser evidence must reject changed/missing native row outputs."""
from dataclasses import replace
from io import BytesIO

from docx import Document
import pymupdf
import pytest

from scripts.qa_report_row_cards import assert_outputs, STATUS, NOTE, TASK, ASSET
from scripts.workspace_row_observations import row_state


def outputs(*, status=STATUS, asset=ASSET, note=NOTE):
    doc=Document()
    for columns, row in [
        (('Project Number','Status','Technical Notes'),('Q-TEST-731',status,note)),
        (('WO#','Task','Asset'),('WO-TEST-1005',TASK,asset)),
    ]:
        table=doc.add_table(rows=2,cols=len(columns))
        for i,value in enumerate(columns):table.cell(0,i).text=value
        for i,value in enumerate(row):table.cell(1,i).text=value
    word=BytesIO();doc.save(word)
    with pymupdf.open() as pdf:
        page=pdf.new_page()
        # Use a Unicode-capable font only when present; PDF assertion keeps the
        # chosen status punctuation independent of the text extraction method.
        from pathlib import Path
        font=Path('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf')
        if not font.is_file():pytest.skip('PDF text fixture needs DejaVu Sans')
        page.insert_font(fontname='qa',fontfile=str(font))
        page.insert_textbox((40,40,550,500),'\n'.join((status,note,TASK,asset)),fontname='qa',fontsize=10)
        raw=pdf.tobytes()
    return word.getvalue(),raw


def test_native_output_assertion_checks_cells_and_pdf():
    result=assert_outputs(*outputs())
    assert result['docx']['bytes']>0 and result['pdf']['pages']==1


@pytest.mark.parametrize('change',[{'status':'Pending'},{'asset':'AHU-WRONG'},{'note':''}])
def test_changed_field_cannot_pass_the_actual_output_assertion(change):
    with pytest.raises(AssertionError):assert_outputs(*outputs(**change))


def test_row_receipt_preserves_all_source_columns_and_unrelated_fingerprint():
    from app.monthly_report_model import synthetic_draft, synthetic_profiles, ReportPeriod, ResolvedBlock, ReportTable
    draft=synthetic_draft(synthetic_profiles()[0],ReportPeriod(2026,9))
    row=ResolvedBlock('service_calls','This month',references=('source-1:1',),extra_tables=(
        ReportTable(('WO#','Task','Asset','Custom'),(('00042','Task','P-1','Unknown retained'),),'source-1:1'),))
    draft=replace(draft,blocks=(*[b for b in draft.blocks if b.key!='service_calls'],row))
    before=row_state(draft)
    table=row.extra_tables[0]
    changed=replace(row,extra_tables=(replace(table,rows=(('00042','Changed task','P-1','Unknown retained'),)),))
    after=row_state(replace(draft,blocks=tuple(changed if b.key=='service_calls' else b for b in draft.blocks)))
    assert before['unrelated']==after['unrelated'] and before['fingerprint']!=after['fingerprint']
    assert after['rows']['service_calls']['tables'][0]['columns']==table.columns
    assert after['rows']['service_calls']['tables'][0]['reference']=='source-1:1'
    assert after['rows']['service_calls']['references']==('source-1:1',)
