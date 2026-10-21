"""Refs #276: parentheses retain the meaning of ASCII approximate amounts.

Synthetic report text exercises preflight, actual Word generation and reviewed
source-page allowance. Original wording remains available for correction.
Unicode approximation and broader language inference are outside this slice.
"""

from dataclasses import replace
from io import BytesIO

from docx import Document
import pytest

from app.monthly_report_checks import preflight
from app.monthly_report_content_policy import page_allowed, page_fingerprint
from app.monthly_report_docx import assemble_docx
from app.monthly_report_guided import client_table_draft
from app.monthly_report_model import (
    ReportDraft, ReportPeriod, ReportSource, ReportTable, ResolvedBlock,
    default_sections, synthetic_profiles,
)


@pytest.fixture(autouse=True)
def isolated_report_data(monkeypatch, tmp_path):
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path / "data"))


def _draft(text="", *, table=None):
    section, key = ("capital", "capital_renewal") if table else ("activity", "activity_summary")
    sections = tuple(replace(item, included=item.key == section) for item in default_sections())
    block = ResolvedBlock(key, "This month", text=text,
                          extra_tables=(table,) if table else ())
    return ReportDraft(synthetic_profiles()[0], ReportPeriod(2026, 11),
                       "Synthetic Price Reviewer", sections, (block,))


@pytest.mark.parametrize("text", (
    pytest.param("Quote (~3.2k)", id="quote-with-parentheses"),
    pytest.param("Repair estimate: (~3.2k)", id="repair-estimate"),
    pytest.param("$(~3200)", id="currency-symbol-prefix"),
    pytest.param("USD (~3.2k)", id="currency-code-prefix"),
    pytest.param("(~3.2k) USD", id="currency-code-suffix"),
    pytest.param("(~3.2k) cost", id="amount-before-label"),
    pytest.param("Quote: (~3k)–(~4k)", id="parenthesized-range-endpoints"),
    pytest.param("Quote: (~3k–~4k)", id="parenthesized-whole-range"),
    pytest.param("Quote received from Synthetic Vendor: (~3.2k)", id="vendor-assignment"),
    pytest.param("Cost\n(~ 3200)", id="extracted-heading-and-amount"),
    pytest.param("Budget: (~2026)", id="year-shaped-budget-amount"),
    pytest.param("Budget for (~2026): (~3.2k)", id="budget-year-and-amount"),
    pytest.param("Quote (3.2)k", id="legacy-multiplier-after-parenthesis"),
    pytest.param("Quote (3200", id="legacy-unmatched-opening-parenthesis"),
    pytest.param("3.2kUSD", id="legacy-compact-currency-suffix"),
    pytest.param("Quote (~ 3.2k )", id="spaces-inside-parentheses"),
    pytest.param("Budget for (2026) k", id="year-shaped-budget-with-magnitude"),
    pytest.param("Budget for (2026)–(2027) k", id="year-range-budget-with-magnitude"),
    pytest.param("(~3.2)k USD", id="multiplier-outside-parenthesis-with-currency"),
))
def test_parenthesized_prices_remain_editable_but_block_preflight_and_word(text):
    original = _draft(text)
    working, removed = client_table_draft(original)
    assert working == original and not removed
    assert working.blocks[0].text == text

    checks = preflight(working)
    assert not [check for check in checks if check.blocking and check.code != "pricing"]
    priced = any(check.code == "pricing" and check.blocking
                 and check.block_key == "activity_summary" for check in checks)
    try:
        raw = assemble_docx(working, acknowledged_fingerprint=working.fingerprint)
    except ValueError as exc:
        blocked = "pric" in str(exc).lower()
    else:
        # Prove the unfixed generator actually publishes the price, rather
        # than mistaking an unrelated or empty-output failure for this defect.
        assert text in "\n".join(p.text for p in Document(BytesIO(raw)).paragraphs)
        blocked = False
    assert (priced, blocked) == (True, True), {
        "text": text, "pricing_preflight": priced, "word_blocked": blocked,
    }
    assert original.blocks[0].text == text


@pytest.mark.parametrize("text", (
    pytest.param("Estimate (~3.2) kW", id="unit-after-parenthesis"),
    pytest.param("Estimate (~3.2 kW)", id="unit-inside-parenthesis"),
    pytest.param("Estimate (~3.2kW)", id="compact-power-unit"),
    pytest.param("Estimate (~3200 kWh)", id="energy-unit"),
    pytest.param("Estimate: (~125–~150) psi", id="whole-pressure-range"),
    pytest.param("Estimate: (~125)–(~150) psi", id="pressure-range-endpoints"),
    pytest.param("Estimate (~125)psi", id="unit-adjacent-to-closing-parenthesis"),
    pytest.param("Estimate (~3.2 ) kW", id="space-before-close-with-unit"),
    pytest.param("Estimate (~125 )–(~150 ) psi", id="spaces-before-range-closures"),
    pytest.param("Estimate (~125) – (~150psi)", id="range-endpoint-with-compact-unit"),
    pytest.param("Estimate (~125) to (~150psi)", id="word-range-with-compact-unit"),
    pytest.param("Not to exceed (~12 A)", id="current-unit"),
    pytest.param("Repair estimate (~48 hours)", id="duration-unit"),
    pytest.param("Utility costs decreased (~8%)", id="percentage"),
    pytest.param("Budget for (~2026) remains pending.", id="budget-year"),
    pytest.param("Budget for (~2026–~2027) remains pending.", id="budget-year-range"),
    pytest.param("Budget for (~2026)–(~2027) remains pending.", id="budget-year-endpoints"),
    pytest.param("Budget for (~2026 - 11 - 14) remains pending.", id="spaced-date-is-not-a-budget-period"),
    pytest.param("Quote date: (~2026-11-14)", id="quote-date"),
    pytest.param("Quote (2026-11-14)", id="date-is-not-a-partial-price-range"),
    pytest.param("Quote (~2026-11-14)", id="approximate-date-is-not-a-partial-price-range"),
    pytest.param("Quote received from Synthetic Vendor: (~2026-11-14)", id="vendor-date"),
    pytest.param("Quote received from Synthetic Vendor: (~2026 - 11 - 14)", id="spaced-vendor-date"),
    pytest.param("Quotation reference: (Q-TEST-3200)", id="quote-reference"),
    pytest.param("Number of Quotes: (~3)", id="reference-count"),
    pytest.param("(~3.2k) readings recorded.", id="unlabeled-quantity"),
    pytest.param("Quote is pending. (~3200) operating hours were recorded.", id="sentence-boundary"),
))
def test_parenthesized_technical_values_periods_and_references_reach_word_unchanged(text):
    original = _draft(text)
    working, removed = client_table_draft(original)
    assert working == original and not removed
    assert not [check for check in preflight(working)
                if check.code == "pricing" and check.blocking]
    raw = assemble_docx(working, acknowledged_fingerprint=working.fingerprint)
    assert text in "\n".join(p.text for p in Document(BytesIO(raw)).paragraphs)
    assert original.blocks[0].text == text


def test_parenthesized_inline_table_price_stays_for_correction_and_blocks_word():
    text = "Inspect the synthetic pump; repair estimate: (~3.2k)"
    table = ReportTable(("Asset", "Recommendation", "Quote reference"),
                        (("SYN-PUMP-1", text, "Q-TEST-3200"),), "synthetic-source")
    original = _draft(table=table)
    working, removed = client_table_draft(original)
    assert working == original and not removed
    assert working.blocks[0].extra_tables == (table,)
    assert any(check.code == "pricing" and check.blocking
               and check.block_key == "capital_renewal" for check in preflight(working))
    with pytest.raises(ValueError, match="[Pp]ric"):
        assemble_docx(working, acknowledged_fingerprint=working.fingerprint)
    assert original.blocks[0].extra_tables == (table,)


@pytest.mark.parametrize("text,caption,reviewed,allowed", (
    pytest.param("Inspected the synthetic pump. Repair estimate: (~3.2k)", "", True, False,
                 id="review-cannot-waive-source-price"),
    pytest.param("Inspected the synthetic pump.", "Quote (~3.2k)", True, False,
                 id="review-cannot-waive-caption-price"),
    pytest.param("Inspected the synthetic pump. Estimate (~3.2) kW", "Quote reference: Q-TEST-3200",
                 True, True, id="reviewed-technical-page"),
    pytest.param("Inspected the synthetic pump. Estimate (~3.2) kW", "", False, False,
                 id="technical-page-still-needs-review"),
))
def test_source_page_review_does_not_override_parenthesized_price_policy(text, caption, reviewed, allowed):
    source = ReportSource("a" * 64, "synthetic-inspection.pdf", "b" * 64, ".pdf",
                          page_texts=(text,), selected_pages=(1,), captions=((1, caption),))
    if reviewed:
        source = replace(source, client_page_reviews=((1, page_fingerprint(source, 1)),))
        assert dict(source.client_page_reviews)[1] == page_fingerprint(source, 1)
    original_fingerprint = source.fingerprint
    assert page_allowed(source, 1) is allowed
    assert source.fingerprint == original_fingerprint
    assert source.page_texts == (text,) and source.captions == ((1, caption),)
