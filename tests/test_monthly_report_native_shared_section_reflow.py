"""Replaced content cannot split a pinned native Word section."""
from dataclasses import replace
from io import BytesIO
from copy import deepcopy
from hashlib import sha256

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.oxml.ns import qn
from docx.shared import Inches
import pytest

from app.monthly_report_native_layout import build_native_docx
from app import monthly_report_library as library
from app.monthly_report_import import _text
from app.monthly_report_native_art_scope import expected_header_scope
from test_monthly_report_native_art_review import sample, setup
from test_monthly_report_native_art_lifecycle import save_sample
from test_monthly_report_source_frame_logo_role import raster


TITLE = "Contract Asset List — Synthetic covered equipment"
SCOPE = "Synthetic regional campuses. Source: the approved example contract schedule."
TRAINING = "Current training information has not been supplied."


@pytest.mark.parametrize("section", [None, "training"])
def test_training_context_retains_original_section_ownership(sample, section):
    path, *_ = sample
    document = Document(path)
    document.add_section(WD_SECTION_START.NEW_PAGE)
    document.add_heading("TRAINING SUMMARY", 1)
    document.add_paragraph("Old monthly training facts")
    document.add_section(WD_SECTION_START.NEW_PAGE)
    document.add_paragraph(TITLE).runs[0].bold = True
    document.add_paragraph(SCOPE).runs[0].italic = True
    table = document.add_table(rows=2, cols=2)
    for row, values in zip(table.rows, (("Description", "Campus"),
                                      ("Synthetic air separator", "Synthetic western campus"))):
        for cell, value in zip(row.cells, values):
            cell.text = value
    document.save(path)
    fixture = setup(path)
    draft, loader = save_sample(fixture)
    draft = replace(draft, blocks=tuple(
        replace(block, text="\n\n".join((TRAINING, TITLE, SCOPE)))
        if block.key == "training_summary" else block for block in draft.blocks))
    output = build_native_docx(draft, path, asset_loader=loader, section_key=section)
    result = Document(BytesIO(output))
    assert result.tables[-1].cell(1, 1).text == "Synthetic western campus"
    assert sum(p.text == TITLE for p in result.paragraphs) == 1
    assert sum(p.text == SCOPE for p in result.paragraphs) == 1
    assert sum(p.text == TRAINING for p in result.paragraphs) == 1


@pytest.mark.parametrize("section", [None, "water"])
def test_appended_water_pages_keep_shared_section_with_next_divider(sample, section):
    path, *_ = sample
    document = Document(path)
    document.add_section(WD_SECTION_START.NEW_PAGE)
    document.add_heading("WATER TREATMENT REPORTS", 1)
    document.add_section(WD_SECTION_START.NEW_PAGE)
    for color in ("red", "green", "blue", "yellow"):
        document.add_picture(BytesIO(raster(color)), width=Inches(4))
    divider = document.add_heading("EQUIPMENT PERFORMANCE ISSUES", 1)
    document.add_section(WD_SECTION_START.NEW_PAGE)
    empty_boundary = document.paragraphs[-1]._p
    boundary = empty_boundary.find(".//" + qn("w:sectPr"))
    divider._p.get_or_add_pPr().append(deepcopy(boundary))
    empty_boundary.getparent().remove(empty_boundary)
    document.add_paragraph("Synthetic current issue observation.")
    document.save(path)
    fixture = setup(path)
    draft, loader = save_sample(fixture)
    water = next(block for block in draft.blocks if block.key == "water_reports")
    assert len(water.asset_hashes) == 4
    added = {}
    for color in ("purple", "pink", "gray"):
        raw = raster(color)
        added[library.asset_reference(raw, "png")] = raw
    water = replace(water, asset_hashes=(*water.asset_hashes, *added),
                    asset_captions=(*water.asset_captions,
                                    *("Current synthetic source page" for _ in added)))
    draft = replace(draft, blocks=tuple(
        water if block.key == water.key else block for block in draft.blocks))
    result = Document(BytesIO(build_native_docx(
        draft, path, asset_loader=lambda key: added[key] if key in added else loader(key),
        section_key=section)))
    scope = expected_header_scope(fixture[1], section)
    assert len(result.sections) == len(scope.routes)
    body = list(result.element.body)
    start = next(i for i, node in enumerate(body) if _text(node) == "WATER TREATMENT REPORTS")
    end = next((i for i, node in enumerate(body)
                if i > start and _text(node) == "EQUIPMENT PERFORMANCE ISSUES"), len(body))
    pictures = [image for node in body[start + 1:end] for image in node.iter(qn("a:blip"))]
    assert [sha256(result.part.related_parts[image.get(qn("r:embed"))].blob).hexdigest()
            for image in pictures] == [ref.split(".")[0] for ref in water.asset_hashes]
    assert sum(p.text == "Current synthetic source page" for p in result.paragraphs) == 3
    if section is None:
        assert "Synthetic current issue observation." in "\n".join(p.text for p in result.paragraphs)
