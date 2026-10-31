"""Repeat-output checks use the final export profile, not a weaker oracle."""
from dataclasses import replace
from io import BytesIO
import random

from docx import Document
from docx.shared import Inches
from PIL import Image
import pymupdf as fitz
import pytest

from app import monthly_report_designs, monthly_report_docx, pdf_converter
from app.monthly_report_library import SavedReport
from app.monthly_report_model import ReportPeriod, synthetic_draft, synthetic_profiles
from scripts.qa_source_page_outputs import _record
from conftest import requires_libreoffice


@pytest.fixture
def output_case(monkeypatch, tmp_path):
    monkeypatch.setenv('EPC_DATA_DIR', str(tmp_path / 'runtime'))
    monkeypatch.setattr(pdf_converter, 'OUTPUT_DIR', tmp_path)
    monkeypatch.setattr(pdf_converter, 'PDF_BACKEND', 'libreoffice')
    monkeypatch.setattr(monthly_report_designs, 'source_for', lambda _: None)
    draft = synthetic_draft(synthetic_profiles()[0], ReportPeriod(2026, 9))
    saved = SavedReport(draft, '2026-09-30T00:00:00+00:00', 1)
    image = Image.frombytes('RGB', (320, 180), random.Random(125).randbytes(320 * 180 * 3))
    png = BytesIO()
    image.save(png, 'PNG')
    document = Document()
    document.add_paragraph('Verified synthetic chart caption')
    document.add_picture(BytesIO(png.getvalue()), width=Inches(2))
    word = BytesIO()
    document.save(word)
    raw = word.getvalue()
    pdf, error = monthly_report_docx._convert_report_pdf(draft, raw)
    assert pdf and not error
    package = monthly_report_docx.ReportPackage(raw, pdf, 'synthetic.docx', 'synthetic.pdf', draft.fingerprint)
    return draft, saved, package


def record_case(tmp_path, output_case):
    draft, saved, package = output_case
    return _record(tmp_path, 'same-profile', draft, saved, package,
                   required_text=['Verified synthetic chart caption'],
                   forbidden_text=['Retired chart caption'], pictures=1,
                   route='Synthetic real final-converter and delivered-Word repeat')


@requires_libreoffice
def test_actual_record_repeats_final_profile_with_exact_every_page_pixels(monkeypatch, tmp_path, output_case):
    calls = []
    real = pdf_converter.convert_to_pdf
    def observe(path, **kwargs):
        calls.append((path.read_bytes(), kwargs))
        return real(path, **kwargs)
    monkeypatch.setattr(pdf_converter, 'convert_to_pdf', observe)
    result = record_case(tmp_path, output_case)
    package = output_case[2]
    assert calls == [(package.docx, {'lossless_images': True})]
    assert result['delivered_word_render']['same_backend_each_page_matches_pdf'] == [True]
    assert (tmp_path / 'same-profile/SYNTHETIC-report.docx').read_bytes() == package.docx
    assert not list(tmp_path.glob('.office-*'))


@requires_libreoffice
def test_default_profile_negative_control_still_fails_exact_pixel_oracle(monkeypatch, tmp_path, output_case):
    real = pdf_converter.convert_to_pdf
    def wrong_profile(path, **kwargs):
        # Deliberately reproduce the former call without changing the oracle.
        return real(path)
    monkeypatch.setattr(pdf_converter, 'convert_to_pdf', wrong_profile)
    with pytest.raises(AssertionError, match='Repeat rendering of delivered Word differs from final PDF'):
        record_case(tmp_path, output_case)
    assert (tmp_path / 'same-profile/SYNTHETIC-report.docx').read_bytes() == output_case[2].docx


@pytest.mark.parametrize('damage', ['pixels', 'geometry', 'page_count'])
@requires_libreoffice
def test_same_profile_does_not_accept_changed_pdf_pixels_geometry_or_page_count(tmp_path, output_case, damage):
    draft, saved, package = output_case
    with fitz.open(stream=package.pdf, filetype='pdf') as pdf:
        if damage == 'pixels':
            pdf[0].draw_rect(fitz.Rect(30, 30, 45, 45), fill=(1, 0, 0), overlay=True)
        elif damage == 'geometry':
            box = pdf[0].mediabox
            pdf[0].set_mediabox(fitz.Rect(box.x0, box.y0, box.x1 + 10, box.y1))
        else:
            pdf.new_page()
        changed = replace(package, pdf=pdf.tobytes())
    message = 'Delivered Word page count differs' if damage == 'page_count' else 'Repeat rendering of delivered Word differs from final PDF'
    with pytest.raises(AssertionError, match=message):
        record_case(tmp_path, (draft, saved, changed))
