"""Desired behavior for the four exact-head PR275 review findings.

Synthetic DOCX import, AppTest uploader delivery, persisted source state and
actual preview/native DOCX bytes. The full-preview converter is a bounded
stub; its input is inspected, so this is not a PDF visual acceptance claim.
"""

from dataclasses import replace
import hashlib
from io import BytesIO
from zipfile import ZipFile

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches
import fitz
from PIL import Image
import pytest
from streamlit.testing.v1 import AppTest

from app import monthly_report_library as library, monthly_report_preview as preview
from app.monthly_report_asset_review import approve_all_assets
from app.monthly_report_cover_ui import replace_cover_picture
from app.monthly_report_import import ImportMapping, inspect_docx, map_items
from app.monthly_report_model import ReportDraft, ReportPeriod, ResolvedBlock, default_sections, synthetic_profiles
from app.monthly_report_native_layout import build_native_docx
from app.monthly_report_sections import build_section_import, section_reviews
from app.monthly_report_start_import import cover_photo_candidates, unresolved_cover_artwork
from test_cover_branding_contracts import assert_mark_preserved, mark, photograph
from test_cover_import_task_ui import current, stage_source


def test_reselecting_committed_source_replaces_another_staged_source(monkeypatch, tmp_path):
    app, first, incoming = stage_source(monkeypatch, tmp_path)
    first_bytes = first.read_bytes()
    next(button for button in app.button if button.label == "Continue to this month’s updates").click().run()
    assert not app.exception
    committed = current(app)
    first_photo = next(block for block in committed.blocks if block.key == "cover_photo")
    assert library.imported_original(committed.profile.contract, committed.profile.key).read_bytes() == first_bytes
    assert not any(key.endswith("_stage") for key in app.session_state.filtered_state)

    # A successful save may leave the same native upload present. This must
    # not reopen it until an actual different selection has been delivered.
    app.run()
    assert not any(key.endswith("_stage") for key in app.session_state.filtered_state)
    second = Document()
    second.add_paragraph(committed.profile.facilities[0].title)
    second.add_paragraph(committed.period.label)
    second.add_picture(BytesIO(photograph((810, 560))), width=Inches(5))
    second.add_section()
    second.add_heading("Monthly Activity Summary", 1)
    second.add_paragraph("Synthetic second-source valve inspection.")
    stream = BytesIO()
    second.save(stream)
    second_upload = BytesIO(stream.getvalue())
    second_upload.name, second_upload.size = "second-source.docx", len(stream.getvalue())
    incoming["Starting report"] = second_upload
    app.run()
    assert not app.exception
    stage_key = next(key for key in app.session_state.filtered_state if key.endswith("_stage"))
    stage_b = app.session_state[stage_key]
    assert stage_b[2].sha256 == hashlib.sha256(stream.getvalue()).hexdigest(), "The second upload must actually be processed"

    first_upload = BytesIO(first_bytes)
    first_upload.name, first_upload.size = first.name, len(first_bytes)
    incoming["Starting report"] = first_upload
    app.run()
    assert not app.exception
    stage_a = app.session_state[stage_key]
    assert stage_a[2].sha256 == hashlib.sha256(first_bytes).hexdigest(), "Review must follow the current upload, not a previous staged file"
    assert stage_a[1].read_bytes() == first_bytes
    assert not stage_b[1].exists(), "Only the superseded temporary stage is retired"
    next(button for button in app.button if button.label == "Continue to this month’s updates").click().run()
    assert not app.exception
    loaded = library.load_imported_draft(committed.profile.contract, committed.profile.key)
    assert library.imported_original(committed.profile.contract, committed.profile.key).read_bytes() == first_bytes
    assert next(block for block in loaded.blocks if block.key == "cover_photo").asset_hashes == first_photo.asset_hashes
    assert not any(key.endswith("_stage") for key in app.session_state.filtered_state)


def _draft(*blocks):
    sections = tuple(replace(section, included=section.key == "activity") for section in default_sections())
    return ReportDraft(synthetic_profiles()[0], ReportPeriod(2026, 11), "Synthetic Reviewer", sections,
                       (*blocks, ResolvedBlock("activity_summary", "This month", text="Synthetic inspection completed.")))


def _preview_input(monkeypatch, tmp_path, report, assets, route):
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path / "data"))
    if route != "report":
        return preview.section_preview_docx(report, route, assets.__getitem__)
    captured = []
    def convert(path):
        captured.append(path.read_bytes())
        result = tmp_path / (path.stem + ".pdf")
        with fitz.open() as pdf:
            pdf.new_page().insert_text((72, 72), "Synthetic converter boundary")
            pdf.save(result)
        return result
    monkeypatch.setattr(preview.pdf_converter, "convert_to_pdf", convert)
    result = preview.preview_report(report, assets.__getitem__)
    assert result.pages == 1 and len(captured) == 1
    return captured[0]


@pytest.mark.parametrize("route,slot", [
    ("report", "brand_logo"), ("report", "client_logo"),
    ("cover", "brand_logo"), ("cover", "client_logo"), ("activity", "brand_logo"),
])
def test_preview_docx_keeps_logo_alpha_and_intentional_white(monkeypatch, tmp_path, route, slot):
    assets = {}
    logo = approve_all_assets(replace_cover_picture(ResolvedBlock(slot, "Omit"), mark(), ".png", assets))
    raw = _preview_input(monkeypatch, tmp_path, _draft(logo), assets, route)
    with ZipFile(BytesIO(raw)) as package:
        images = [package.read(name) for name in package.namelist() if name.startswith("word/media/")]
    assert images, "The actual preview DOCX must contain the supplied logo"
    for image in images:
        assert_mark_preserved(image)


@pytest.mark.parametrize("route", ["report", "cover"])
def test_preview_photograph_keeps_its_existing_opaque_canvas(monkeypatch, tmp_path, route):
    reference = hashlib.sha256(mark()).hexdigest() + ".png"
    block = approve_all_assets(ResolvedBlock("cover_photo", "This month", asset_hashes=(reference,)))
    raw = _preview_input(monkeypatch, tmp_path, _draft(block), {reference: mark()}, route)
    with ZipFile(BytesIO(raw)) as package:
        images = [package.read(name) for name in package.namelist() if name.startswith("word/media/")]
    assert images
    assert all(Image.open(BytesIO(image)).convert("RGBA").getpixel((0, 0)) == (255, 255, 255, 255) for image in images)


def _header_source(tmp_path, evidence):
    document = Document()
    document.add_paragraph(synthetic_profiles()[0].facilities[0].title)
    document.add_paragraph("November 2026")
    header = document.sections[0].header
    paragraph = header.paragraphs[0]
    if evidence == "caption":
        paragraph.add_run("Company logo")
    company = mark((70, 110, 150, 255))
    shape = paragraph.add_run().add_picture(BytesIO(company), width=Inches(2))
    if evidence == "alt":
        shape._inline.docPr.set("descr", "ENFRA company logo")
    if evidence == "caption":
        paragraph = header.add_paragraph()
    unrelated = (photograph((700, 100)), photograph((180, 180)), photograph((640, 80)))
    for raw in unrelated:
        paragraph.add_run().add_picture(BytesIO(raw), width=Inches(1))
    document.add_section()
    document.add_heading("Monthly Activity Summary", 1)
    document.add_paragraph("Synthetic inspection completed.")
    source = tmp_path / "mixed-header.docx"
    document.save(source)
    return source, company, unrelated


def _header_images(raw):
    document = Document(BytesIO(raw))
    return tuple(part.related_parts[node.get(qn("r:embed"))].blob
                 for part in document.part.package.parts
                 if str(part.partname).startswith("/word/header") and hasattr(part, "element")
                 for node in part.element.iter(qn("a:blip")))


@pytest.mark.parametrize("evidence", ["alt", "caption", "mapped"])
@pytest.mark.parametrize("master", [False, True])
def test_brand_replacement_targets_only_the_identified_header_mark(tmp_path, evidence, master):
    source, company, unrelated = _header_source(tmp_path, evidence)
    original = source.read_bytes()
    inspection = inspect_docx(source)
    company_item = next(item for item in inspection.items if item.kind == "image"
                        and item.image_digest == hashlib.sha256(company).hexdigest())
    assets = {}
    logo = approve_all_assets(replace_cover_picture(ResolvedBlock("brand_logo", "Omit"), mark(), ".png", assets))
    if evidence == "mapped":
        mapped = map_items(source, inspection, (ImportMapping(company_item.id, "brand_logo"),))
        logo = replace(logo, references=mapped.blocks[0].references)
    else:
        assert company_item.image_role == "brand_logo"
        assert not any(item.image_role == "brand_logo" for item in inspection.items
                       if item.kind == "image" and item.id != company_item.id), "Header placement cannot identify a company mark"
    raw = build_native_docx(_draft(logo), source, asset_loader=assets.__getitem__, master=master)
    actual = _header_images(raw)
    if master:
        assert len(actual) == 1, "A fresh master cannot turn unapproved client art into duplicate ENFRA logos"
    else:
        assert actual[1:] == unrelated, "Updating the company mark must preserve the other header occurrences"
    assert_mark_preserved(actual[0])
    assert actual[0] != company
    assert source.read_bytes() == original


@pytest.mark.parametrize("master", [False, True])
def test_unlabelled_header_art_is_not_evidence_of_company_identity(tmp_path, master):
    source, company, unrelated = _header_source(tmp_path, "none")
    inspection = inspect_docx(source)
    header_items = tuple(item for item in inspection.items if item.kind == "image" and item.part.startswith("word/header"))
    assert header_items and all(item.suggested_slot != "brand_logo" for item in header_items)
    assert set(header_items) <= set(unresolved_cover_artwork(inspection.items))
    assets = {}
    logo = approve_all_assets(replace_cover_picture(ResolvedBlock("brand_logo", "Omit"), mark(), ".png", assets))
    raw = build_native_docx(_draft(logo), source, asset_loader=assets.__getitem__, master=master)
    assert _header_images(raw) == (() if master else (company, *unrelated))


def test_two_inline_pictures_share_ambiguous_label_and_reachable_repair(tmp_path, monkeypatch):
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path / "data"))
    document = Document()
    document.add_paragraph(synthetic_profiles()[0].facilities[0].title)
    document.add_paragraph("November 2026")
    paragraph = document.add_paragraph("Company logo")
    paragraph.add_run().add_picture(BytesIO(mark()), width=Inches(3))
    paragraph.add_run().add_picture(BytesIO(photograph()), width=Inches(4))
    source = tmp_path / "same-paragraph-label.docx"
    document.save(source)
    original = source.read_bytes()
    inspection = inspect_docx(source)
    section = next(value for value in section_reviews(inspection) if value.key == "cover")
    images = tuple(item for item in section.items if item.kind == "image")
    assert len(images) == 2 and all(item.image_role == "logo" for item in images)
    assert all(not item.suggested_slot for item in images)
    assert not cover_photo_candidates(section.items)
    assert unresolved_cover_artwork(section.items) == images
    app = AppTest.from_string('''
import streamlit as st
from app.monthly_report_section_ui import _section_card
from app.monthly_report_model import ReportPeriod
from app.monthly_report_ui import _field
plans = st.session_state.setdefault('plans', {})
_section_card(st.session_state['section'], st.session_state['source'], st.session_state['inspection'],
 ReportPeriod(2026, 11), 'inline_art', _field, plans, True, cover_design_available=True)
''', default_timeout=20)
    app.session_state["section"], app.session_state["source"], app.session_state["inspection"] = section, source, inspection
    app.run()
    assert not app.exception
    assert any("Repair it in the original Word file and upload it again" in entry.value for entry in app.error)
    assert not app.session_state["plans"]["cover"]["approved"]
    assert not any(button.label.startswith("Choose picture") for button in app.button)
    assert source.read_bytes() == original


def test_per_picture_alt_labels_do_not_claim_a_neighboring_photograph(tmp_path):
    document = Document()
    document.add_paragraph(synthetic_profiles()[0].facilities[0].title)
    document.add_paragraph("November 2026")
    paragraph = document.add_paragraph()
    logo = paragraph.add_run().add_picture(BytesIO(mark()), width=Inches(3))
    logo._inline.docPr.set("descr", "Company logo")
    photo = paragraph.add_run().add_picture(BytesIO(photograph()), width=Inches(4))
    photo._inline.docPr.set("descr", "Facility exterior photograph")
    source = tmp_path / "independent-inline-labels.docx"
    document.save(source)
    inspection = inspect_docx(source)
    images = tuple(item for item in inspection.items if item.kind == "image")
    assert images[0].image_role == "brand_logo"
    assert not images[1].image_role
    assert cover_photo_candidates(inspection.items) == (images[1],)


@pytest.mark.parametrize("route", ["mapping", "section", "manual_ui"])
def test_unidentified_header_and_its_body_copy_never_become_photo_options(tmp_path, monkeypatch, route):
    from app import monthly_report_import_ui as import_ui
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path / "data"))
    profile = synthetic_profiles()[0]
    document = Document()
    document.add_paragraph(profile.facilities[0].title)
    document.add_paragraph("November 2026")
    document.sections[0].header.paragraphs[0].add_run().add_picture(BytesIO(mark()), width=Inches(2))
    document.add_picture(BytesIO(mark()), width=Inches(3))
    document.add_picture(BytesIO(photograph()), width=Inches(5))
    source = tmp_path / "unidentified-header.docx"
    document.save(source)
    original = source.read_bytes()
    inspection = inspect_docx(source)
    restricted = tuple(item for item in inspection.items if item.kind == "image"
                       and item.image_digest == hashlib.sha256(mark()).hexdigest())
    assert len(restricted) == 2 and all(item.suggested_slot != "brand_logo" for item in restricted)
    candidate, = cover_photo_candidates(inspection.items)
    assert candidate.image_digest == hashlib.sha256(photograph()).hexdigest()
    assert set(restricted) <= set(unresolved_cover_artwork(inspection.items))
    if route != "manual_ui":
        for item in restricted:
            with pytest.raises(ValueError, match="[Hh]eader artwork"):
                if route == "mapping":
                    map_items(source, inspection, (ImportMapping(item.id, "cover_photo"),))
                else:
                    build_section_import(source, inspection, ({"key": "cover", "target": "cover", "approved": True,
                        "selected": (item.id,), "destinations": {item.id: "cover_photo"}},))
    else:
        library.save_profile(profile, expected_revision=0, actor="Synthetic Reviewer", confirmed=True)
        upload = BytesIO(original)
        upload.name, upload.size = source.name, len(original)
        real_upload = import_ui.st.file_uploader
        def uploader(label, *args, **kwargs):
            native = real_upload(label, *args, **kwargs)
            return upload if label == "Existing monthly report DOCX" else native
        monkeypatch.setattr(import_ui.st, "file_uploader", uploader)
        app = AppTest.from_string('''
import streamlit as st
from app import monthly_report_library as library, monthly_report_import_ui as import_ui
from app.monthly_report_model import synthetic_profiles
profile = synthetic_profiles()[0]
def field(key, value):
    st.session_state.setdefault(key, value)
    return key
import_ui.render_import(library.load_profile(profile.contract, profile.key), field)
''').run()
        next(button for button in app.button if button.label == "Read DOCX for mapping").click().run()
        assert not app.exception
        for item in restricted:
            next(widget for widget in app.selectbox if widget.label == "Item to inspect").set_value(item.id).run()
            assert not app.exception
            assert "cover_photo" not in next(widget for widget in app.selectbox if widget.label == "Confirmed destination").options
            assert any("Header artwork" in caption.value for caption in app.caption)
        next(widget for widget in app.selectbox if widget.label == "Item to inspect").set_value(candidate.id).run()
        assert "cover_photo" in next(widget for widget in app.selectbox if widget.label == "Confirmed destination").options
    assert source.read_bytes() == original


@pytest.mark.parametrize("same_paragraph", [False, True])
def test_explicit_mapping_does_not_replace_another_equal_byte_header_occurrence(tmp_path, same_paragraph):
    document = Document()
    document.add_paragraph(synthetic_profiles()[0].facilities[0].title)
    document.add_paragraph("November 2026")
    original_image = mark((80, 130, 170, 255))
    header = document.sections[0].header
    header.paragraphs[0].add_run().add_picture(BytesIO(original_image), width=Inches(2))
    second = header.paragraphs[0] if same_paragraph else header.add_paragraph()
    second.add_run().add_picture(BytesIO(original_image), width=Inches(1))
    document.add_section()
    document.add_heading("Monthly Activity Summary", 1)
    document.add_paragraph("Synthetic inspection completed.")
    source = tmp_path / "equal-header-occurrences.docx"
    document.save(source)
    inspection = inspect_docx(source)
    items = tuple(item for item in inspection.items if item.kind == "image" and item.part.startswith("word/header"))
    assert len(items) == 2 and items[0].id != items[1].id
    assert (items[0].position == items[1].position) == same_paragraph
    mapped = map_items(source, inspection, (ImportMapping(items[0].id, "brand_logo"),))
    assets = {}
    logo = replace_cover_picture(mapped.blocks[0], mark(), ".png", assets)
    # Replacement retains the deliberate source occurrence mapping, even
    # though its saved image bytes have changed.
    logo = approve_all_assets(replace(logo, references=mapped.blocks[0].references))
    raw = build_native_docx(_draft(logo), source, asset_loader=assets.__getitem__)
    result = _header_images(raw)
    assert len(result) == 2
    assert_mark_preserved(result[0])
    assert result[1] == original_image, "An independently mapped equal-byte header occurrence must survive"


@pytest.mark.parametrize("master", [False, True])
def test_header_mapping_updates_equivalent_branches_not_independent_duplicate(tmp_path, master):
    from copy import deepcopy
    from lxml import etree

    document = Document()
    document.add_paragraph(synthetic_profiles()[0].facilities[0].title)
    document.add_paragraph("November 2026")
    original_image = mark((80, 130, 170, 255))
    paragraph = document.sections[0].header.paragraphs[0]
    shape = paragraph.add_run().add_picture(BytesIO(original_image), width=Inches(2))
    drawing = shape._inline.getparent()
    run = drawing.getparent()
    mc = "{http://schemas.openxmlformats.org/markup-compatibility/2006}"
    wrapper = etree.Element(mc + "AlternateContent")
    run.replace(drawing, wrapper)
    etree.SubElement(wrapper, mc + "Choice", Requires="w").append(drawing)
    etree.SubElement(wrapper, mc + "Fallback").append(deepcopy(drawing))
    paragraph.add_run().add_picture(BytesIO(original_image), width=Inches(1))
    document.add_section()
    document.add_heading("Monthly Activity Summary", 1)
    document.add_paragraph("Synthetic inspection completed.")
    source = tmp_path / "equivalent-header-occurrences.docx"
    document.save(source)
    original = source.read_bytes()
    inspection = inspect_docx(source)
    items = tuple(item for item in inspection.items if item.kind == "image"
                  and item.part.startswith("word/header"))
    assert len(items) == 2, "Choice and Fallback are one inspected logical image"
    mapped = map_items(source, inspection, (ImportMapping(items[0].id, "brand_logo"),))
    assets = {}
    logo = replace_cover_picture(mapped.blocks[0], mark(), ".png", assets)
    logo = approve_all_assets(replace(logo, references=mapped.blocks[0].references))
    result = _header_images(build_native_docx(_draft(logo), source,
                                             asset_loader=assets.__getitem__, master=master))
    assert len(result) == (2 if master else 3)
    assert_mark_preserved(result[0])
    assert result[0] == result[1] != original_image
    if not master:
        assert result[2] == original_image
    assert source.read_bytes() == original


def test_repeated_image_occurrences_preserve_later_saved_source_ids(tmp_path, monkeypatch):
    document = Document()
    document.add_paragraph(synthetic_profiles()[0].facilities[0].title)
    document.add_paragraph("November 2026")
    first = mark((80, 130, 170, 255))
    later = mark((140, 90, 50, 255))
    header = document.sections[0].header
    paragraph = header.paragraphs[0]
    paragraph.add_run().add_picture(BytesIO(first), width=Inches(2))
    paragraph.add_run().add_picture(BytesIO(first), width=Inches(1))
    header.add_paragraph().add_run().add_picture(BytesIO(later), width=Inches(2))
    source = tmp_path / "historical-image-ids.docx"
    document.save(source)
    inspection = inspect_docx(source)
    images = tuple(item for item in inspection.items if item.kind == "image")
    assert tuple(item.id for item in images) == (
        "item-0003", "item-0003-occurrence-2", "item-0004")
    # item-0004 was the later distinct mark before duplicate occurrences were
    # exposed. The persisted reference must never silently switch to its sibling.
    assets = {}
    logo = approve_all_assets(replace_cover_picture(ResolvedBlock("brand_logo", "Omit"), mark(), ".png", assets))
    logo = replace(logo, references=(f"docx:{inspection.sha256}:item-0004",))
    draft = replace(_draft(logo), blocks=(logo,))
    raw = build_native_docx(draft, source, asset_loader=assets.__getitem__)
    actual = _header_images(raw)
    assert actual[:2] == (first, first)
    assert actual[2] != later
    assert_mark_preserved(actual[2])

    # Safety budgets count actual occurrences, not just legacy primary IDs.
    from app import monthly_report_import as importer
    monkeypatch.setattr(importer, "MAX_ITEMS", 4)
    with pytest.raises(ValueError, match="item/text/table review budget"):
        inspect_docx(source)
