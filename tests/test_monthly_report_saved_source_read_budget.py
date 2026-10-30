"""Saved evidence enforces its byte budget at the actual read, not only stat."""
from dataclasses import replace
from hashlib import sha256
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace

import pytest

from app import monthly_report_library as library, monthly_report_sources as sources
from app.monthly_report_model import ReportPeriod, ResolvedBlock, synthetic_draft
from app.monthly_report_source_recovery import restore_sources, retain_unavailable
from test_monthly_report_sources import profile, pdf_bytes, image_bytes


TEXT = b"Synthetic original: inspected DEMO-01."


def stored(profile, raw=TEXT, suffix=".txt"):
    content, _ = sources.ingest(profile, "synthetic-original" + suffix, raw)
    return content, sources._path(profile, content.source.sha256, suffix)


def watch_read(monkeypatch, target, *, replace_with=None, read_error=False):
    """Observe a real file descriptor; optionally replace a test-owned link at open."""
    native_open = Path.open
    observations = []
    replaced = False

    class Probe:
        def __init__(self, stream):
            self.stream = stream
            self.record = {"requested": [], "returned": [], "stream": stream}
            observations.append(self.record)

        def __enter__(self):
            return self

        def __exit__(self, *exc):
            self.stream.close()

        def read(self, size=-1):
            self.record["requested"].append(size)
            if read_error:
                raise OSError("Synthetic read failure")
            data = self.stream.read(size)
            self.record["returned"].append(len(data))
            return data

    def opening(path, mode="r", *args, **kwargs):
        nonlocal replaced
        if path == target and mode == "rb":
            if replace_with is not None and not replaced:
                # Never modify the content-addressed pool's original hardlink.
                path.unlink()
                with native_open(path, "wb") as stream:
                    stream.write(replace_with)
                replaced = True
            return Probe(native_open(path, mode, *args, **kwargs))
        return native_open(path, mode, *args, **kwargs)

    monkeypatch.setattr(Path, "open", opening)
    return observations


def assert_read(records, limit, size):
    assert len(records) == 1
    assert records[0]["requested"] == [limit + 1]
    assert records[0]["returned"] == [size]
    assert records[0]["stream"].closed


@pytest.mark.parametrize("suffix", [".txt", ".csv", ".pdf", ".png", ".docx"])
def test_valid_saved_formats_use_their_unchanged_limit_and_exact_original(profile, monkeypatch, suffix):
    if suffix == ".docx":
        from docx import Document
        document = Document()
        document.add_paragraph("Synthetic original document.")
        output = BytesIO()
        document.save(output)
        raw = output.getvalue()
    else:
        raw = {".txt": TEXT, ".csv": b"ID,Status\n001,Open\n",
               ".pdf": pdf_bytes(1), ".png": image_bytes()}[suffix]
    content, path = stored(profile, raw, suffix)
    before = path.read_bytes()
    records = watch_read(monkeypatch, path)
    assert sources.source_bytes(profile, content.source) == before == raw
    limit = 128 * 1024 * 1024 if suffix == ".docx" else sources.MAX_FILE_BYTES
    assert_read(records, limit, len(raw))


@pytest.mark.parametrize("headroom", [0, 1])
def test_at_and_below_limit_return_complete_original(profile, monkeypatch, headroom):
    content, path = stored(profile)
    limit = len(TEXT) + headroom
    monkeypatch.setattr(sources, "MAX_FILE_BYTES", limit)
    records = watch_read(monkeypatch, path)
    assert sources.source_bytes(profile, content.source) == TEXT
    assert_read(records, limit, len(TEXT))


def test_already_oversized_source_is_rejected_without_opening(profile, monkeypatch):
    content, path = stored(profile)
    monkeypatch.setattr(sources, "MAX_FILE_BYTES", len(TEXT) - 1)
    records = watch_read(monkeypatch, path)
    with pytest.raises(ValueError, match="Saved source exceeds the file limit"):
        sources.source_bytes(profile, content.source)
    assert not records


@pytest.mark.parametrize("extra", [1, 4096])
def test_growth_after_stat_is_bounded_and_rejected_before_hashing(profile, monkeypatch, extra):
    content, path = stored(profile)
    monkeypatch.setattr(sources, "MAX_FILE_BYTES", len(TEXT))
    records = watch_read(monkeypatch, path, replace_with=TEXT + b"X" * extra)
    # This exact-original prefix would pass SHA if overflow were just truncated.
    def unexpected_hash(raw):
        pytest.fail("An over-budget saved source reached hashing")
    monkeypatch.setattr(sources, "hashlib", SimpleNamespace(sha256=unexpected_hash))
    with pytest.raises(ValueError, match="Saved source exceeds the file limit"):
        sources.source_bytes(profile, content.source)
    assert_read(records, len(TEXT), len(TEXT) + 1)


@pytest.mark.parametrize("replacement", [b"X" * len(TEXT), TEXT[:-1], b"", TEXT + b"X"])
def test_within_budget_changes_still_fail_exact_sha_check(profile, monkeypatch, replacement):
    content, path = stored(profile)
    limit = len(TEXT) + 1
    monkeypatch.setattr(sources, "MAX_FILE_BYTES", limit)
    records = watch_read(monkeypatch, path, replace_with=replacement)
    with pytest.raises(ValueError, match="integrity check"):
        sources.source_bytes(profile, content.source)
    assert_read(records, limit, len(replacement))


def test_failed_stream_read_closes_handle_and_keeps_source_for_retry(profile, monkeypatch):
    content, path = stored(profile)
    with monkeypatch.context() as patch:
        records = watch_read(patch, path, read_error=True)
        with pytest.raises(OSError, match="Synthetic read failure"):
            sources.source_bytes(profile, content.source)
        assert len(records) == 1 and records[0]["stream"].closed
        assert records[0]["requested"] == [sources.MAX_FILE_BYTES + 1]
    assert path.read_bytes() == TEXT
    assert sources.source_bytes(profile, content.source) == TEXT


def test_docx_limit_is_not_replaced_by_the_smaller_evidence_limit(profile, monkeypatch):
    from docx import Document
    output = BytesIO()
    Document().save(output)
    raw = output.getvalue()
    content, path = stored(profile, raw, ".docx")
    monkeypatch.setattr(sources, "MAX_FILE_BYTES", 1)
    records = watch_read(monkeypatch, path)
    assert sources.source_bytes(profile, content.source) == raw
    assert_read(records, 128 * 1024 * 1024, len(raw))


def test_recovery_keeps_siblings_reviews_history_and_allows_same_original_retry(profile, monkeypatch, tmp_path):
    first, path = stored(profile)
    sibling, _ = stored(profile, b"Independent accepted evidence.")
    first = replace(first, source=replace(first.source, captions=((1, "Reviewed original"),)))
    draft = replace(synthetic_draft(profile, ReportPeriod(2026, 9)),
                    blocks=(ResolvedBlock("activity_summary", "This month", text="Keep my manual edit."),),
                    sources=(first.source, sibling.source))
    saved = library.save_snapshot(draft, expected_revision=0, entered_editor="Synthetic Editor")
    before = {str(p): sha256(p.read_bytes()).hexdigest() for p in tmp_path.rglob("*") if p.is_file()}
    with monkeypatch.context() as patch:
        patch.setattr(sources, "MAX_FILE_BYTES", len(TEXT))
        records = watch_read(patch, path, replace_with=TEXT + b"X" * 4096)
        available, missing = restore_sources(profile, draft.sources)
        assert tuple(c.source for c in available) == (sibling.source,)
        assert missing == (first.source,)
        assert set(retain_unavailable(tuple(c.source for c in available), missing)) == set(draft.sources)
        assert_read(records, len(TEXT), len(TEXT) + 1)
    # The simulated out-of-band replacement changes only this test-owned path.
    after = {str(p): sha256(p.read_bytes()).hexdigest() for p in tmp_path.rglob("*") if p.is_file()}
    assert {p: h for p, h in before.items() if p != str(path)} == {
        p: h for p, h in after.items() if p != str(path)}
    assert library.load_snapshot(profile.contract, profile.key, draft.period) == saved
    restored, messages = sources.ingest_batch(profile, ((first.source.filename, TEXT),), (first, sibling))
    assert restored == (first, sibling)
    assert any("Saved source restored" in message for message in messages)
    assert path.read_bytes() == TEXT
    assert restore_sources(profile, draft.sources) == ((first, sibling), ())
    assert library.load_snapshot(profile.contract, profile.key, draft.period) == saved
    assert {str(p): sha256(p.read_bytes()).hexdigest() for p in tmp_path.rglob("*") if p.is_file()} == before


@pytest.mark.parametrize("change", [{"sha256": "../not-a-digest"}, {"suffix": ".exe"}])
def test_invalid_source_identity_never_opens_a_file(profile, monkeypatch, change):
    content, path = stored(profile)
    records = watch_read(monkeypatch, path)
    with pytest.raises(ValueError, match="reference"):
        sources.source_bytes(profile, replace(content.source, **change))
    assert not records


@pytest.mark.parametrize("suffix", [".pdf", ".png"])
def test_page_preview_rejects_grown_source_before_decoding_or_caching(profile, monkeypatch, suffix):
    raw = pdf_bytes(1) if suffix == ".pdf" else image_bytes()
    content, path = stored(profile, raw, suffix)
    prepared = path.parent / "prepared" / content.source.sha256
    assert not prepared.exists()
    with monkeypatch.context() as patch:
        patch.setattr(sources, "MAX_FILE_BYTES", len(raw))
        records = watch_read(patch, path, replace_with=raw + b"X" * 4096)
        def no_decode(*args, **kwargs):
            pytest.fail("An oversized saved source reached image/PDF decoding")
        patch.setattr(sources.fitz, "open", no_decode)
        patch.setattr(sources, "normalize_report_image", no_decode)
        with pytest.raises(ValueError, match="Saved source exceeds the file limit"):
            sources.page_image(profile, content, 1, preview=True)
        assert_read(records, len(raw), len(raw) + 1)
        assert not prepared.exists()
    restored, _ = sources.ingest_batch(profile, ((content.source.filename, raw),), (content,))
    assert restored == (content,)
    image = sources.page_image(profile, content, 1, preview=True)
    assert image.data and image.width > 0 and image.height > 0
    assert sources.source_bytes(profile, content.source) == raw
