"""Synthetic unchanged-photo orientation and representation boundaries."""
from dataclasses import replace
from io import BytesIO
from zipfile import ZipFile

from docx import Document
from docx.oxml.ns import qn
from lxml import etree
from PIL import Image, ImageChops, ImageDraw
import pytest

from app.monthly_report_import import inspect_docx, map_items, ImportMapping
from app.monthly_report_native_layout import _fit_cover_photo, build_native_docx
from test_cover_unchanged_photo_geometry import source_photo, photo_geometry


def oriented_source(tmp_path, orientation, crop):
    path, draft, _, pixels = source_photo(tmp_path, crop)
    photo = Image.open(BytesIO(pixels))
    draw = ImageDraw.Draw(photo)
    for (x, y), color in zip(((200, 80), (200, 220), (180, 150), (220, 150)),
                             ((200, 30, 30), (30, 30, 200), (20, 150, 40), (150, 20, 150))):
        draw.rectangle((x-7, y-7, x+7, y+7), fill=color)
    exif = photo.getexif(); exif[274] = orientation
    out = BytesIO(); photo.save(out, 'JPEG', quality=95, exif=exif)
    photograph = out.getvalue()
    with ZipFile(path) as package:
        parts = {name: package.read(name) for name in package.namelist()}
    parts['word/media/image1.jpg'] = photograph
    del parts['word/media/image1.png']
    parts['word/_rels/document.xml.rels'] = parts['word/_rels/document.xml.rels'].replace(b'image1.png', b'image1.jpg')
    parts['[Content_Types].xml'] = parts['[Content_Types].xml'].replace(b'</Types>', b'<Default Extension="jpg" ContentType="image/jpeg"/></Types>')
    with ZipFile(path, 'w') as package:
        for name, data in parts.items(): package.writestr(name, data)
    inspection = inspect_docx(path)
    item = next(item for item in inspection.items if item.kind == 'image')
    assert (item.image_width, item.image_height) == (400, 300)
    mapped = map_items(path, inspection, (ImportMapping(item.id, 'cover_photo'),))
    return path, replace(draft, blocks=mapped.blocks), dict(mapped.assets), photograph


def color_bounds(image, target, tolerance=15):
    difference = ImageChops.difference(image.convert('RGB'), Image.new('RGB', image.size, target))
    bands = [band.point(lambda value: 255 if value <= tolerance else 0) for band in difference.split()]
    bounds = ImageChops.multiply(ImageChops.multiply(bands[0], bands[1]), bands[2]).getbbox()
    assert bounds, target
    return bounds


from tests.conftest import requires_libreoffice


@requires_libreoffice
@pytest.mark.parametrize('orientation', [5, 6, 7, 8])
def test_actual_pdf_oriented_photo_keeps_original_subject_and_visual_axes(tmp_path, monkeypatch, orientation):
    import fitz
    from app import pdf_converter
    from app.monthly_report_render_compat import rendering_docx
    monkeypatch.setenv('EPC_DATA_DIR', str(tmp_path/'store'))
    monkeypatch.setattr(pdf_converter, 'OUTPUT_DIR', tmp_path)
    path, draft, assets, _ = oriented_source(tmp_path, orientation, {'t': 10000, 'b': 23333})
    result = build_native_docx(draft, path, asset_loader=assets.__getitem__, section_key='cover')
    measurements = []
    for name, raw in (('original', path.read_bytes()), ('current', result)):
        output = tmp_path/(name+'.docx'); output.write_bytes(rendering_docx(raw))
        with fitz.open(pdf_converter.convert_to_pdf(output)) as pdf:
            pix = pdf[0].get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
            actual = Image.frombytes('RGB', (pix.width, pix.height), pix.samples)
        actual.save(tmp_path/(name+'.png'))
        bounds = [color_bounds(actual, color) for color in
                  ((23, 63, 73), (200, 30, 30), (30, 30, 200), (20, 150, 40), (150, 20, 150))]
        left, top, right, bottom = bounds[0]
        measurements.append([(right-left, bottom-top)] +
                            [(sum(b[::2])/2-left, sum(b[1::2])/2-top) for b in bounds[1:]])
    for original, current in zip(*measurements):
        assert current == pytest.approx(original, abs=2), 'Unproved orientation must not alter original native subject geometry or axes'


@pytest.mark.parametrize('orientation', [2, 3, 4, 5, 6, 7, 8])
@pytest.mark.parametrize('crop', [{}, {'t': 25000, 'b': 37500}])
@pytest.mark.parametrize('section', [None, 'cover'])
def test_unchanged_exif_photo_retains_unproved_source_geometry(tmp_path, orientation, crop, section):
    path, draft, assets, photograph = oriented_source(tmp_path, orientation, crop)
    original = path.read_bytes()
    result = build_native_docx(draft, path, asset_loader=assets.__getitem__, section_key=section)
    fitted, frame = photo_geometry(result)
    assert (fitted, frame) == photo_geometry(original)
    assert path.read_bytes() == original
    with ZipFile(BytesIO(result)) as package:
        assert photograph in [package.read(n) for n in package.namelist() if n.startswith('word/media/')]


@pytest.mark.parametrize('crop', [{}, {'t': 10000, 'b': 23333}])
def test_ordinary_jpeg_orientation_still_fits_and_preserves_valid_crop(tmp_path, crop):
    path, draft, assets, photograph = oriented_source(tmp_path, 1, crop)
    result = build_native_docx(draft, path, asset_loader=assets.__getitem__)
    fitted, frame = photo_geometry(result)
    assert 400*(1-fitted['l']-fitted['r'])/(300*(1-fitted['t']-fitted['b'])) == pytest.approx(frame[0]/frame[1], rel=.001)
    if crop:
        assert fitted == {'l': 0, 'r': 0, 't': .1, 'b': .23333}
    with ZipFile(BytesIO(result)) as package:
        assert photograph in [package.read(n) for n in package.namelist() if n.startswith('word/media/')]


@pytest.mark.parametrize('subtype', ['rect', 'roundrect', 'oval', 'image'])
def test_unsupported_unchanged_vml_shape_subtype_is_left_untouched(subtype):
    v = '{urn:schemas-microsoft-com:vml}'
    shape = etree.Element(v + subtype, style='width:4in;height:2in')
    image = etree.SubElement(shape, v+'imagedata')
    before = etree.tostring(shape)
    _fit_cover_photo(image, 400, 300, preserve_crop=True)
    assert etree.tostring(shape) == before


@pytest.mark.parametrize('subtype', ['rect', 'roundrect', 'oval', 'image'])
@pytest.mark.parametrize('section', [None, 'cover'])
def test_native_vml_subtype_preview_and_export_remain_available(tmp_path, subtype, section):
    path, draft, _, photograph = source_photo(tmp_path, {})
    doc = Document(path)
    drawing = next(doc.element.iter(qn('w:drawing')))
    rid = next(drawing.iter(qn('a:blip'))).get(qn('r:embed'))
    parent = drawing.getparent(); parent.remove(drawing)
    pict = etree.SubElement(parent, qn('w:pict'))
    v = '{urn:schemas-microsoft-com:vml}'
    shape = etree.SubElement(pict, v+subtype, style='width:4in;height:2in')
    etree.SubElement(shape, v+'imagedata', {qn('r:id'): rid, 'croptop': '.1'})
    doc.save(path)
    original = path.read_bytes()
    inspection = inspect_docx(path)
    item = next(i for i in inspection.items if i.kind == 'image')
    mapped = map_items(path, inspection, (ImportMapping(item.id, 'cover_photo'),))
    draft = replace(draft, blocks=mapped.blocks)
    result = build_native_docx(draft, path, asset_loader=dict(mapped.assets).__getitem__, section_key=section)
    output = Document(BytesIO(result))
    retained = next(output.element.iter(v+subtype))
    assert retained.get('style') == shape.get('style')
    assert retained.find(v+'imagedata').attrib == shape.find(v+'imagedata').attrib
    assert path.read_bytes() == original
    with ZipFile(BytesIO(result)) as package:
        assert photograph in [package.read(n) for n in package.namelist() if n.startswith('word/media/')]


@pytest.mark.parametrize('coordinates', [None, '400,200', 'bad,200', 'nan,200', 'inf,200', '0,200'])
def test_unproved_vml_group_geometry_is_preserved(coordinates):
    v = '{urn:schemas-microsoft-com:vml}'
    group = etree.Element(v+'group', style='width:4in;height:3in')
    if coordinates is not None:
        group.set('coordsize', coordinates)
    shape = etree.SubElement(group, v+'shape', style='width:400;height:200')
    image = etree.SubElement(shape, v+'imagedata', croptop='.1', cropbottom='.2')
    before = etree.tostring(group)
    _fit_cover_photo(image, 400, 300, preserve_crop=True)
    assert etree.tostring(group) == before


@pytest.mark.parametrize('group_height,vertical_crop', [(2, 1/6), (3, 0)])
def test_changed_photo_keeps_existing_grouped_vml_fitting(group_height, vertical_crop):
    v = '{urn:schemas-microsoft-com:vml}'
    group = etree.Element(v+'group', style=f'width:4in;height:{group_height}in', coordsize='400,200')
    shape = etree.SubElement(group, v+'shape', style='width:400;height:200')
    image = etree.SubElement(shape, v+'imagedata')
    _fit_cover_photo(image, 400, 300, preserve_crop=False)
    fitted = {name: float(image.get(name)) for name in ('cropleft', 'cropright', 'croptop', 'cropbottom')}
    assert fitted == pytest.approx({'cropleft': 0, 'cropright': 0, 'croptop': vertical_crop, 'cropbottom': vertical_crop})


@pytest.mark.parametrize('complexity', ['inset-fill', 'tiled-fill', 'group-scale', 'unsupported-vml', 'vml-group'])
@pytest.mark.parametrize('reverse', [False, True])
def test_unsupported_alternate_picture_representations_are_skipped_together(tmp_path, complexity, reverse):
    path, _, _, _ = source_photo(tmp_path, {})
    doc = Document(path)
    picture = next(doc.element.iter(qn('pic:pic')))
    parent = picture.getparent(); parent.remove(picture)
    mc = '{http://schemas.openxmlformats.org/markup-compatibility/2006}'
    alternate = etree.SubElement(parent, mc+'AlternateContent')
    choice = etree.SubElement(alternate, mc+'Choice', Requires='pic')
    choice.append(picture)
    fallback = etree.SubElement(alternate, mc+'Fallback')
    v = '{urn:schemas-microsoft-com:vml}'
    shape = etree.SubElement(fallback, v+('rect' if complexity == 'unsupported-vml' else 'shape'), style='width:4in;height:2in')
    legacy = etree.SubElement(shape, v+'imagedata')
    blip = next(picture.iter(qn('a:blip')))
    if complexity == 'inset-fill':
        picture.find(qn('pic:blipFill')).find(qn('a:stretch')).find(qn('a:fillRect')).set('b', '25000')
    elif complexity == 'tiled-fill':
        fill = picture.find(qn('pic:blipFill'))
        fill.remove(fill.find(qn('a:stretch'))); etree.SubElement(fill, qn('a:tile'))
    elif complexity == 'group-scale':
        choice.remove(picture)
        wpg = '{http://schemas.microsoft.com/office/word/2010/wordprocessingGroup}'
        group = etree.SubElement(choice, wpg+'wgp')
        properties = etree.SubElement(group, wpg+'grpSpPr')
        transform = etree.SubElement(properties, qn('a:xfrm'))
        etree.SubElement(transform, qn('a:ext'), cx='400', cy='300')
        etree.SubElement(transform, qn('a:chExt'), cx='400', cy='200')
        group.append(picture)
    elif complexity == 'vml-group':
        fallback.remove(shape)
        group = etree.SubElement(fallback, v+'group', style='width:4in;height:3in', coordsize='400,200')
        group.append(shape)
    before = etree.tostring(alternate)
    for node in ((legacy, blip) if reverse else (blip, legacy)):
        _fit_cover_photo(node, 400, 300, preserve_crop=True)
    assert etree.tostring(alternate) == before, 'All representations of an unsupported unchanged picture must remain unchanged'


def test_orientation_on_another_alternate_representation_prevents_partial_recrop(tmp_path):
    from app.monthly_report_native_layout import _fit_unchanged_cover_photo
    path, _, _, _ = source_photo(tmp_path, {})
    item = next(i for i in inspect_docx(path).items if i.kind == 'image')
    doc = Document(path)
    picture = next(doc.element.iter(qn('pic:pic')))
    parent = picture.getparent(); parent.remove(picture)
    mc = '{http://schemas.openxmlformats.org/markup-compatibility/2006}'
    alternate = etree.SubElement(parent, mc+'AlternateContent')
    etree.SubElement(alternate, mc+'Choice', Requires='pic').append(picture)
    fallback = etree.SubElement(alternate, mc+'Fallback')
    v = '{urn:schemas-microsoft-com:vml}'
    shape = etree.SubElement(fallback, v+'shape', style='width:4in;height:2in')
    peer = tmp_path/'oriented-peer'; peer.mkdir()
    _, _, _, oriented = oriented_source(peer, 6, {})
    rid, _ = doc.part.get_or_add_image(BytesIO(oriented))
    legacy = etree.SubElement(shape, v+'imagedata', {qn('r:id'): rid})
    blip = next(picture.iter(qn('a:blip')))
    before = etree.tostring(alternate)
    for node in (blip, legacy): _fit_unchanged_cover_photo(node, path, item, doc.part)
    assert etree.tostring(alternate) == before
