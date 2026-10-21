"""Source-derived header reservations, never generic geometry permission."""
from dataclasses import replace
from copy import deepcopy
from docx import Document
from docx.shared import Inches
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import pytest
from test_monthly_report_native_art_review import sample, setup
from test_monthly_report_native_art_lifecycle import save_sample
from test_monthly_report_native_art_scope import rewrite, xml
from app.monthly_report_native_layout import build_native_docx
from app.monthly_report_preview import section_preview_docx
from app.monthly_report_model import ResolvedBlock
from app.monthly_report_native_art_scope import verify_header_routing, W


def split_header(sample):
    path = sample[0]
    doc = Document(path)
    doc.sections[0].top_margin = Inches(.5)
    inline = next(doc.sections[0].header._element.iter(qn('wp:inline')))
    inline.tag = qn('wp:anchor')
    position = OxmlElement('wp:positionV'); position.set('relativeFrom', 'page')
    offset = OxmlElement('wp:posOffset'); offset.text = str(45*12700)
    position.append(offset); inline.insert(0, position)
    inline.find(qn('wp:extent')).set('cy', str(30*12700))
    heading = next(p for p in doc.paragraphs if p.text == 'MONTHLY ACTIVITY SUMMARY')
    boundary = OxmlElement('w:p'); ppr = OxmlElement('w:pPr'); boundary.append(ppr)
    ppr.append(deepcopy(doc.sections[-1]._sectPr)); heading._p.addprevious(boundary)
    doc.save(path)
    return setup(path)


def render(sample, monkeypatch, *, blank=True, section='activity', edited=True):
    from app import monthly_report_native_art_review as art
    fixture = split_header(sample)
    path = fixture[0]
    draft, loader = save_sample(fixture)
    if edited:
        if blank:
            # The guided UI creates a previously absent/blank notes block.
            draft = replace(draft, blocks=tuple(b for b in draft.blocks if b.key != 'activity_summary') +
                (ResolvedBlock('activity_summary', 'This month', text='Synthetic current-month inputs missing.\n\nTEST DRAFT.'),))
        else:
            draft = replace(draft, blocks=tuple(replace(b, text='Synthetic update.\n\n'+b.text)
                if b.key == 'activity_summary' else b for b in draft.blocks))
    captured = []
    real = art.verify_native_art_output
    def verify(raw, review, **kwargs):
        result = real(raw, review, **kwargs)
        if kwargs.get('source_scope') is not None:
            captured.append((raw, kwargs['source_scope'], kwargs['emission_proof']))
        return result
    monkeypatch.setattr(art, 'verify_native_art_output', verify)
    output = (build_native_docx(draft, path, asset_loader=loader) if section is None
              else section_preview_docx(draft, section, loader))
    assert captured and output == captured[-1][0]
    return captured[-1]


@pytest.mark.parametrize('blank', [False, True])
@pytest.mark.parametrize('section', [None, 'activity'])
def test_actual_note_rewrite_selects_exact_source_header_reservation(sample, monkeypatch, blank, section):
    output, scope, proof = render(sample, monkeypatch, blank=blank, section=section)
    assert proof.reserved_source_sections == (2,)
    index = scope.source_word_sections.index(2)
    assert scope.reserved_top_margins[index] == '1620'
    assert scope.routes[index].properties_sha256 != scope.reserved_routes[index].properties_sha256
    verify_header_routing(output, scope, emission_proof=proof)


def test_unchanged_and_cover_regions_cannot_claim_reservation(sample, monkeypatch):
    output, scope, proof = render(sample, monkeypatch, section='cover', edited=False)
    assert not proof.reserved_source_sections
    verify_header_routing(output, scope, emission_proof=proof)


@pytest.mark.parametrize('case', ['extra_top', 'left', 'page_width', 'section_type', 'first', 'even',
                                 'unrecorded', 'unknown', 'duplicate', 'foreign'])
def test_reservation_cannot_authorize_unrelated_changes(sample, monkeypatch, case):
    output, scope, proof = render(sample, monkeypatch)
    if case == 'unrecorded': proof=replace(proof,reserved_source_sections=())
    elif case == 'unknown': proof=replace(proof,reserved_source_sections=(999,))
    elif case == 'duplicate': proof=replace(proof,reserved_source_sections=(2,2))
    elif case == 'foreign': proof=replace(proof,scope_sha256='0'*64)
    else:
        def mutate(parts):
            def change(root):
                props=root.find('.//'+W+'sectPr')
                if case in ('extra_top','left'):
                    margin=props.find(W+'pgMar');attr=W+('top' if case=='extra_top' else 'left')
                    margin.set(attr,str(int(margin.get(attr))+20))
                elif case=='page_width': props.find(W+'pgSz').set(W+'w','13000')
                elif case=='section_type':
                    kind=props.find(W+'type')
                    if kind is None:kind=OxmlElement('w:type');props.append(kind)
                    kind.set(W+'val','continuous')
                elif case=='first': props.append(OxmlElement('w:titlePg'))
            if case=='even':
                xml(parts,'word/settings.xml',lambda root:root.append(OxmlElement('w:evenAndOddHeaders')))
            else: xml(parts,'word/document.xml',change)
        output=rewrite(output,mutate)
    with pytest.raises(ValueError):verify_header_routing(output,scope,emission_proof=proof)


def test_cover_output_cannot_take_even_a_known_source_reservation(sample, monkeypatch):
    output, scope, proof = render(sample, monkeypatch, section='cover', edited=False)
    def mutate(parts):
        xml(parts,'word/document.xml',lambda root:root.find('.//'+W+'pgMar').set(W+'top','1620'))
    with pytest.raises(ValueError):
        verify_header_routing(rewrite(output,mutate),scope,emission_proof=proof)


def test_noop_header_measurement_keeps_original_margin_exact():
    from app.monthly_report_native_text import running_header_top
    doc=Document(); properties=doc.sections[0]._sectPr
    doc.sections[0].top_margin=Inches(1)
    assert running_header_top(properties,[]) is None
    header=doc.sections[0].header
    from test_monthly_report_narrative_flow import textbox
    # Explicit page-anchored text below the current margin needs no expansion.
    textbox(header.paragraphs[0],['Synthetic header'],y=5,height=10,header=True)
    top=running_header_top(properties,[header._element])
    assert top=='1440'
