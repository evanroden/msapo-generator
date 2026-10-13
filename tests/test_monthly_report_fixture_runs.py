"""Synthetic bytes exercise the real Starting inspector and reset boundary."""
from contextvars import copy_context
from dataclasses import asdict, replace
from hashlib import sha256
from io import BytesIO
import json

from docx import Document
import pytest
import requests

from app import memory
from app import monthly_report_fixture_runs as runs
from app import monthly_report_fixture_vault as vault
from app import monthly_report_import as production_import
from app import monthly_report_library as report_library
from app import monthly_report_testing as testing
from app.monthly_report_editor import _signature
from app.monthly_report_model import Facility, ReportPeriod, ReportProfile
from app.monthly_report_shared_fixtures import SharedFixtureLibrary
from app.monthly_report_start import report_key


TOKEN_A = "a" * 32
TOKEN_B = "b" * 32


@pytest.fixture
def fixture_case(tmp_path, monkeypatch):
    live = tmp_path / "app-storage"
    live.mkdir()
    shared = tmp_path / "shared-originals"
    shared.mkdir()
    monkeypatch.setenv("EPC_DATA_DIR", str(live))
    facilities = (Facility("synthetic-site-a", "Synthetic Site A"),)
    profile = ReportProfile("Synthetic Contract", report_key(facilities), "Synthetic Report", facilities)
    metadata = vault.FixtureMetadata(
        friendly_name="Synthetic source for real parser",
        source=vault.ReportIdentity(profile.contract, (facilities[0].key,), "2026-09"),
        target=vault.ReportIdentity(profile.contract, (facilities[0].key,), "2026-10"),
        scenario="AUTH01", owner_expectations="Preserve standing contacts and independently review old work.",
        expected_content=(vault.ExpectedSection("Activity", excludes=("MANIFEST-ONLY CANARY MUST NOT BE PREFILLED",)),),
        confirmed_at="2026-10-13T16:00:00+00:00",
    )
    document = Document()
    document.add_heading("Synthetic Site A", 0)
    document.add_paragraph("September 2026 Monthly Report")
    document.add_heading("1 Organizational Chart", 1)
    table = document.add_table(rows=2, cols=2)
    for row, values in zip(table.rows, (("Role", "Name"), ("Site Lead", "Synthetic Person"))):
        for cell, value in zip(row.cells, values):
            cell.text = value
    document.add_heading("2 Monthly Activity Summary", 1)
    document.add_paragraph("September 2026: inspected the synthetic pump; follow-up remains open.")
    binary = BytesIO()
    document.save(binary)
    original = binary.getvalue()
    library = SharedFixtureLibrary(shared)
    receipt = library.upload(original, metadata)
    return live, library, receipt, original, profile


def test_two_fresh_runs_use_actual_inspector_and_preserve_independent_metadata(fixture_case, monkeypatch):
    live, library, receipt, original, profile = fixture_case
    calls = []
    inspect = production_import.inspect_docx

    def spy(path):
        calls.append((path, path.read_bytes()))
        return inspect(path)

    def no_network(*_args, **_kwargs):
        pytest.fail("Native DOCX inspection must not invoke an external reader.")

    monkeypatch.setattr(production_import, "inspect_docx", spy)
    monkeypatch.setattr(requests.sessions.Session, "request", no_network)
    results = []
    for token in (TOKEN_A, TOKEN_B):
        with testing.isolated(token):
            memory._data_dir()  # Explicit initial binding belongs to the caller.
            results.append(runs.stage_fixture(library, receipt.fixture_id, token, profile))
    first, second = results
    assert first.context.scope != second.context.scope
    assert first.context.generation != second.context.generation
    assert first.source_path != second.source_path
    assert first.source_path.is_relative_to(first.context.data_root)
    assert second.source_path.is_relative_to(second.context.data_root)
    assert [raw for _, raw in calls] == [original, original]
    assert first.inspection == inspect(first.source_path)
    assert first.inspection.sha256 == receipt.original_sha256 == sha256(original).hexdigest()
    assert any("synthetic pump" in item.text for item in first.inspection.items)
    assert "MANIFEST-ONLY CANARY" not in str(first.inspection)
    assert first.lineage.metadata == receipt.metadata
    assert first.lineage.metadata.expected_content == receipt.metadata.expected_content
    assert first.period == ReportPeriod(2026, 10)
    assert first.profile is profile
    assert first.stage == (None, first.source_path, first.inspection)
    assert first.setup_stage_key == "report_setup_" + _signature((profile.contract, profile.key, "2026-10")) + "_stage"
    assert json.loads(first.source_path.with_name("lineage.json").read_bytes()) == json.loads(
        json.dumps({"schema": 1, **asdict(first.lineage)}))
    assert library.download(receipt.fixture_id) == original


def test_run_reset_removes_only_copy_and_old_jobs_cannot_recreate_or_enter_new_generation(fixture_case):
    live, library, receipt, original, profile = fixture_case
    with testing.isolated(TOKEN_A):
        memory._data_dir()
        staged = runs.stage_fixture(library, receipt.fixture_id, TOKEN_A, profile)
        old_job = copy_context()
        old_callback = testing.capture_namespace()
        testing.reset(live)
        assert not staged.context.data_root.exists()
        with pytest.raises(OSError, match="reset|changed"):
            runs.validate_run(staged.context)
        with pytest.raises(OSError, match="reset|changed"):
            old_job.run(runs.stage_fixture, library, receipt.fixture_id, TOKEN_A, profile)
        assert not staged.context.data_root.exists()
    with testing.isolated(TOKEN_A):
        memory._data_dir()
        replacement = runs.stage_fixture(library, receipt.fixture_id, TOKEN_A, profile)
        assert replacement.context.generation != staged.context.generation
        with pytest.raises(OSError, match="reset|changed"):
            old_job.run(runs.validate_run, staged.context)
        with old_callback(), pytest.raises(OSError, match="reset|changed"):
            runs.validate_run(staged.context)
        assert replacement.source_path.read_bytes() == original
    assert library.download(receipt.fixture_id) == original


def test_reset_during_native_inspection_blocks_late_completion_without_recreating_namespace(fixture_case, monkeypatch):
    live, library, receipt, original, profile = fixture_case
    inspect = production_import.inspect_docx

    def reset_after_read(path):
        result = inspect(path)
        testing.reset(live)
        return result

    monkeypatch.setattr(production_import, "inspect_docx", reset_after_read)
    with testing.isolated(TOKEN_A):
        root = memory._data_dir()
        with pytest.raises(OSError, match="reset"):
            runs.stage_fixture(library, receipt.fixture_id, TOKEN_A, profile)
        assert not root.exists()
    assert library.download(receipt.fixture_id) == original


def test_live_profile_history_and_canaries_remain_unchanged(fixture_case):
    live, library, receipt, original, profile = fixture_case
    report_library.save_profile(profile, expected_revision=0, actor="Synthetic Live", confirmed=True)
    canary = live / "standing-canary.txt"
    canary.write_text("Synthetic live state must not change.")
    saved = report_library.load_profile(profile.contract, profile.key)
    before = {str(path.relative_to(live)): path.read_bytes() for path in live.rglob("*") if path.is_file()}
    with testing.isolated(TOKEN_A):
        memory._data_dir()
        assert report_library.list_profiles(profile.contract) == ()
        staged = runs.stage_fixture(library, receipt.fixture_id, TOKEN_A, profile)
        assert report_library.list_profiles(profile.contract) == ()
        assert not list(staged.context.data_root.rglob("manifest.json"))
        testing.reset(live)
    # Global fixture-copy reservation metadata is intentionally separate from
    # live report/standing data. Reset reclaims its reservation on next admission.
    after = {str(path.relative_to(live)): path.read_bytes() for path in live.rglob("*")
             if path.is_file() and "_enfra_fixture_run_admission_v1" not in path.parts}
    assert before == after
    assert report_library.load_profile(profile.contract, profile.key) == saved
    assert library.download(receipt.fixture_id) == original


@pytest.mark.parametrize("mismatch", ["contract", "sites", "period"])
def test_incorrect_target_rejected_before_any_run_copy(fixture_case, mismatch):
    live, library, receipt, original, profile = fixture_case
    target = (replace(profile, contract="Other Contract") if mismatch == "contract"
              else replace(profile, facilities=(Facility("other-site", "Other Site"),)) if mismatch == "sites"
              else profile)
    period = ReportPeriod(2026, 11) if mismatch == "period" else None
    with testing.isolated(TOKEN_A):
        root = memory._data_dir()
        with pytest.raises(ValueError, match="confirmed target"):
            runs.stage_fixture(library, receipt.fixture_id, TOKEN_A, target, target_period=period)
        assert {path.name for path in root.iterdir()} == {testing._SENTINEL}
    assert library.download(receipt.fixture_id) == original


def test_context_and_bound_generation_are_required_before_library_read(fixture_case, monkeypatch):
    live, library, receipt, _original, profile = fixture_case

    def forbidden_read(*_args, **_kwargs):
        pytest.fail("Invalid fixture scope must be rejected before reading its original.")

    monkeypatch.setattr(library, "read", forbidden_read)
    with pytest.raises(ValueError, match="bind"):
        runs.stage_fixture(library, receipt.fixture_id, TOKEN_A, profile)
    with testing.isolated(TOKEN_A):
        with pytest.raises(ValueError, match="bind"):
            runs.stage_fixture(library, receipt.fixture_id, TOKEN_A, profile)
        memory._data_dir()
        with pytest.raises(ValueError, match="token"):
            runs.stage_fixture(library, receipt.fixture_id, TOKEN_B, profile)


def test_same_namespace_cannot_overwrite_or_accumulate_staged_fixtures(fixture_case):
    live, library, receipt, original, profile = fixture_case
    with testing.isolated(TOKEN_A):
        memory._data_dir()
        staged = runs.stage_fixture(library, receipt.fixture_id, TOKEN_A, profile)
        before = {str(path): path.read_bytes() for path in staged.context.data_root.rglob("*") if path.is_file()}
        with pytest.raises((ValueError, vault.VaultIntegrityError), match="fresh isolated run|already reserved"):
            runs.stage_fixture(library, receipt.fixture_id, TOKEN_A, profile)
        after = {str(path): path.read_bytes() for path in staged.context.data_root.rglob("*") if path.is_file()}
        assert before == after
    assert library.download(receipt.fixture_id) == original


def test_live_root_returned_by_faulty_adapter_is_rejected(fixture_case, monkeypatch):
    live, library, receipt, _original, profile = fixture_case
    with testing.isolated(TOKEN_A):
        memory._data_dir()
        monkeypatch.setattr(memory, "_data_dir", lambda: live)
        with pytest.raises(OSError, match="fall back"):
            runs.stage_fixture(library, receipt.fixture_id, TOKEN_A, profile)
    assert not (live / "_enfra_fixture_start_v1").exists()


def test_mismatched_original_hash_never_reaches_parser_or_disk(fixture_case, monkeypatch):
    live, library, receipt, original, profile = fixture_case
    monkeypatch.setattr(library, "read", lambda _: vault.FixtureOriginal(receipt, original + b"changed"))
    with testing.isolated(TOKEN_A):
        root = memory._data_dir()
        with pytest.raises(vault.VaultIntegrityError, match="hash"):
            runs.stage_fixture(library, receipt.fixture_id, TOKEN_A, profile)
        assert {path.name for path in root.iterdir()} == {testing._SENTINEL}


def test_inspection_failure_retains_only_disposable_copy_and_original(fixture_case, monkeypatch):
    live, library, receipt, original, profile = fixture_case

    def failed_reader(_path):
        raise ValueError("Synthetic bounded native-reader failure")

    monkeypatch.setattr(production_import, "inspect_docx", failed_reader)
    with testing.isolated(TOKEN_A):
        root = memory._data_dir()
        with pytest.raises(ValueError, match="native-reader failure"):
            runs.stage_fixture(library, receipt.fixture_id, TOKEN_A, profile)
        assert next(root.rglob("source.docx")).read_bytes() == original
        testing.reset(live)
    assert library.download(receipt.fixture_id) == original


def test_refresh_recovers_validated_lineage_then_reuses_same_production_inspector(fixture_case, monkeypatch):
    _live, library, receipt, original, profile = fixture_case
    with testing.isolated(TOKEN_A):
        memory._data_dir()
        assert runs.read_run_fixture(library, TOKEN_A) is None
        staged = runs.stage_fixture(library, receipt.fixture_id, TOKEN_A, profile)
    with testing.isolated(TOKEN_A):
        memory._data_dir()
        record = runs.read_run_fixture(library, TOKEN_A)
        assert record.context == staged.context
        assert record.lineage == staged.lineage
        assert record.source_path == staged.source_path
        resumed = runs.resume_fixture(record, profile)
        assert resumed.inspection == staged.inspection
        assert resumed.source_path.read_bytes() == original
        with pytest.raises(ValueError, match="profile identity"):
            runs.resume_fixture(record, replace(profile, key="changed-report-key"))


@pytest.mark.parametrize("changed", ["lineage", "original", "transaction"])
def test_resume_rejects_changed_run_receipts_or_bytes(fixture_case, changed):
    _live, library, receipt, original, profile = fixture_case
    with testing.isolated(TOKEN_A):
        memory._data_dir()
        staged = runs.stage_fixture(library, receipt.fixture_id, TOKEN_A, profile)
        if changed == "lineage":
            path = staged.source_path.with_name("lineage.json")
            value = json.loads(path.read_bytes())
            value["metadata"]["target"]["month"] = "2026-11"
            path.write_text(json.dumps(value))
        elif changed == "original":
            staged.source_path.write_bytes(original + b"changed")
        else:
            staged.source_path.with_name(vault._TRANSACTION).write_text("{}")
        with pytest.raises(vault.VaultIntegrityError):
            runs.read_run_fixture(library, TOKEN_A)
    assert library.download(receipt.fixture_id) == original
