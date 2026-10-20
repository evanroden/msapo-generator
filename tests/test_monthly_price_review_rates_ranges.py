"""Refs #276: financial rates/ranges versus real metadata and measurements.

Synthetic current-report/snapshot and editable package checks. The final case
also retains actual converted PDF pages; it is not native Word or graphic OCR
acceptance. Before-output is written before desired assertions when permitted.
"""
from dataclasses import replace
from hashlib import sha256
from io import BytesIO
import json
import os
from pathlib import Path
from zipfile import ZipFile

from docx import Document
import fitz
import pytest

from app import monthly_report_library as library
from app.monthly_report_checks import preflight
from app.monthly_report_docx import assemble_docx, generate_report
from app.monthly_report_guided import client_table_draft
from app.monthly_report_model import ReportTable, ResolvedBlock
from test_monthly_report_monetary_context import _draft
from conftest import requires_libreoffice


RATES = ('Cost / kg', 'Unit Price / kWh', 'Fee / hour', 'Labor rate / hour')
METADATA = ('Date Quoted', 'Reference Quote', 'Number of Quotes', 'Count of Estimates')
MONEY = ('Repair estimate 5000-10000', 'Quote: 5k-10k',
         'Estimated repair: 6875', 'Estimated replacement: 18675.50',
         'Quote 5000–10000', 'Repair estimate 5,000–10,000',
         'Budget for 2026-2027: 8,640')
OPERATING = ('Quote received from Acme: 2026-11-14', 'Quote date 09/01/2026',
             'Estimated repair time: 48 hours', 'Estimated replacement flow: 125 gpm',
             'Not to exceed 3-4 starts per hour', 'Estimated replacement: 480 V',
             'Budget for 2026-2027 remains under review.',
             'Estimate for 2026–2027 remains pending.',
             'Date Quoted: 2026-09-17; Number of Quotes: 2')


def output_dir(tmp_path):
    result = Path(os.environ.get('INDEPENDENT_OUTPUT', str(tmp_path))) / 'price-rates-ranges'
    result.mkdir(parents=True, exist_ok=True)
    return result


@pytest.mark.parametrize('heading,monetary', [
    *((heading, True) for heading in RATES),
    *((heading, False) for heading in METADATA),
    ('Amount / kg', False), ('Estimate / hours', False),
])
def test_rate_and_metadata_columns_reach_saved_editable_client_output(monkeypatch, tmp_path, heading, monetary):
    monkeypatch.setenv('EPC_DATA_DIR', str(tmp_path / 'data'))
    value = '2026-09-17' if heading == 'Date Quoted' else 'Q-0071' if heading == 'Reference Quote' else '17.25' if monetary else '48'
    table = ReportTable(('Asset', heading, 'Findings'), (('EF-51', value, 'The fan did not start.'),))
    original = _draft('capital', 'capital_renewal', ResolvedBlock(
        'capital_renewal', 'This month', extra_tables=(table,)))
    saved = library.save_snapshot(original, expected_revision=0, entered_editor='Synthetic Reviewer')
    restored = library.load_snapshot(original.profile.contract, original.profile.key, original.period)
    assert restored.draft == original and restored.revision == saved.revision
    working, removed = client_table_draft(restored.draft)
    raw = assemble_docx(working, acknowledged_fingerprint=working.fingerprint)
    actual = Document(BytesIO(raw)).tables[0]
    headers = tuple(cell.text for cell in actual.rows[0].cells)
    cells = tuple(cell.text for cell in actual.rows[1].cells)
    assert removed == ((heading,) if monetary else ())
    assert headers == (('Asset', 'Findings') if monetary else table.columns)
    assert cells == (('EF-51', 'The fan did not start.') if monetary else table.rows[0])
    assert bool([check for check in preflight(original) if check.code == 'pricing' and check.blocking]) == monetary
    assert library.load_snapshot(original.profile.contract, original.profile.key, original.period) == restored
    assert restored.draft.blocks[0].extra_tables == (table,)


@pytest.mark.parametrize('text,monetary', [*((text, True) for text in MONEY), *((text, False) for text in OPERATING)])
def test_price_ranges_and_adjectives_block_actual_generation_without_rewriting(text, monetary):
    original = _draft('activity', 'activity_summary', ResolvedBlock('activity_summary', 'This month', text=text))
    working, removed = client_table_draft(original)
    assert working == original and not removed
    blocked = [check for check in preflight(working) if check.code == 'pricing' and check.blocking]
    assert bool(blocked) == monetary, (text, blocked)
    if monetary:
        with pytest.raises(ValueError, match='[Pp]ric'):
            assemble_docx(working, acknowledged_fingerprint=working.fingerprint)
    else:
        raw = assemble_docx(working, acknowledged_fingerprint=working.fingerprint)
        assert text in '\n'.join(paragraph.text for paragraph in Document(BytesIO(raw)).paragraphs)


@requires_libreoffice
def test_rate_exclusion_and_metadata_survival_in_actual_word_pdf(monkeypatch, tmp_path):
    monkeypatch.setenv('EPC_DATA_DIR', str(tmp_path / 'data'))
    tables = tuple(ReportTable(('Asset', heading, 'Findings'),
        (('EF-' + str(51 + index), str(19000.25 + index), 'Remains out of service.'),))
        for index, heading in enumerate(RATES))
    metadata_rows = (('EF-71', '2026-09-17', 'Q-0071', '2', '3'),)
    tables += (ReportTable(('Asset', *METADATA), metadata_rows),)
    original = _draft('capital', 'capital_renewal', ResolvedBlock('capital_renewal', 'This month', extra_tables=tables))
    before = library.save_snapshot(original, expected_revision=0, entered_editor='Synthetic Reviewer')
    working, removed = client_table_draft(original)
    saved = library.save_snapshot(working, expected_revision=before.revision, entered_editor='Synthetic Reviewer')
    restored = library.load_snapshot(working.profile.contract, working.profile.key, working.period)
    assert restored.draft == working and restored.revision == saved.revision
    package = generate_report(restored.draft, acknowledged_fingerprint=restored.draft.fingerprint)
    assert package.docx and package.pdf, package.pdf_error
    out = output_dir(tmp_path)
    (out / 'rates-and-metadata.docx').write_bytes(package.docx)
    (out / 'rates-and-metadata.pdf').write_bytes(package.pdf)
    with fitz.open(stream=package.pdf, filetype='pdf') as document:
        pdf_text = '\n'.join(page.get_text() for page in document)
        for index, page in enumerate(document, 1):
            page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5)).save(out / f'page-{index:02}.png')
    with ZipFile(BytesIO(package.docx)) as zipped:
        xml = '\n'.join(zipped.read(name).decode('utf-8', errors='replace') for name in zipped.namelist() if name.endswith('.xml'))
    manifest = {'tested_revision': os.environ.get('MSAPO_SHA', 'unrecorded'),
        'original_revision': before.revision, 'accepted_transform_revision': saved.revision,
        'original_fingerprint': original.fingerprint, 'saved_fingerprint': working.fingerprint,
        'removed': removed, 'docx_bytes': len(package.docx), 'docx_sha256': sha256(package.docx).hexdigest(),
        'pdf_bytes': len(package.pdf), 'pdf_sha256': sha256(package.pdf).hexdigest(),
        'scope': 'Synthetic saved-state, actual editable DOCX and LibreOffice PDF; every saved PDF page available for separate visual inspection.'}
    (out / 'manifest.json').write_text(json.dumps(manifest, indent=2) + '\n')
    assert removed == RATES
    for index in range(len(RATES)):
        assert str(19000.25 + index) not in xml and str(19000.25 + index) not in pdf_text
    last = Document(BytesIO(package.docx)).tables[-1]
    assert tuple(cell.text for cell in last.rows[0].cells) == ('Asset', *METADATA)
    assert tuple(cell.text for cell in last.rows[1].cells) == metadata_rows[0]
    assert '2026-09-17' in pdf_text and 'Q-0071' in pdf_text
    assert library.load_snapshot_version(original.profile.contract, original.profile.key, original.period, before.revision).draft == original
