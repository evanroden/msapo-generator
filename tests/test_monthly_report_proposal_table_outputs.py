"""Actual synthetic output preserves proposal decision and scope separately."""
from collections import Counter
from dataclasses import replace
from hashlib import sha256
from io import BytesIO
import json
import os
from pathlib import Path

from docx import Document
import fitz
import pytest

from app.monthly_report_docx import assemble_docx
from app.monthly_report_followups import confirmation, report_text
from app.monthly_report_model import ReportPeriod
from app.monthly_report_native_layout import build_native_docx
from app.monthly_report_setup import new_month_draft
from tests.conftest import requires_libreoffice
from test_monthly_report_proposal_table_carryover import imported_proposals


@requires_libreoffice
@pytest.mark.parametrize("native", [False, True])
def test_imported_proposals_keep_cells_and_decisions_in_word_pdf(tmp_path, monkeypatch, native):
    from app import pdf_converter
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path / "data"))
    monkeypatch.setattr(pdf_converter, "OUTPUT_DIR", tmp_path)
    source, original, prior = imported_proposals(tmp_path)
    current = new_month_draft(prior, ReportPeriod(2026, 10))
    current = replace(current, follow_ups=tuple(replace(item,
        reviewed_fingerprint=confirmation(item, current.period)) for item in current.follow_ups))
    # Full deterministic validation stays mandatory for both output paths.
    plain = assemble_docx(current, acknowledged_fingerprint=current.fingerprint)
    raw = build_native_docx(current, source, section_key="proposals") if native else plain
    document = Document(BytesIO(raw))
    paragraphs = [p.text for p in document.paragraphs]
    expected = Counter(report_text(item) for item in current.follow_ups)
    # The existing renderer expresses a cell's newline as consecutive Word
    # paragraphs; its complete wording and occurrence count must still match.
    body_text = "\n".join(paragraphs)
    for text, count in expected.items():
        assert body_text.count(text) == count
    assert [item.status for item in current.follow_ups] == ["pending", "declined", "pending"]
    assert all("$1200" not in p and "$2500" not in p for p in paragraphs)
    assert not any("1200" in cell.text or "2500" in cell.text
                   for table in document.tables for row in table.rows for cell in row.cells)
    path = tmp_path / "proposals.docx"
    path.write_bytes(raw)
    pdf_path = pdf_converter.convert_to_pdf(path)
    pdf_bytes = pdf_path.read_bytes()
    label = "native" if native else "plain"
    target = Path(os.environ.get("INDEPENDENT_OUTPUT", str(tmp_path))) / "proposal-table-carryover"
    target.mkdir(parents=True, exist_ok=True)
    (target / f"{label}.docx").write_bytes(raw)
    (target / f"{label}.pdf").write_bytes(pdf_bytes)
    with fitz.open(stream=pdf_bytes, filetype="pdf") as pdf:
        text = " ".join(" ".join(page.get_text().split()) for page in pdf)
        for value, count in expected.items():
            assert text.count(" ".join(value.split())) == count
        assert "1200" not in text and "2500" not in text
        for n, page in enumerate(pdf, 1):
            page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5)).save(target / f"{label}-page-{n}.png")
        pages = len(pdf)
    (target / f"{label}-manifest.json").write_text(json.dumps({
        "docx_sha256": sha256(raw).hexdigest(), "pdf_sha256": sha256(pdf_bytes).hexdigest(),
        "pages": pages, "proposal_occurrences": len(current.follow_ups),
        "source_unchanged": source.read_bytes() == original,
    }, indent=2))
    assert source.read_bytes() == original
