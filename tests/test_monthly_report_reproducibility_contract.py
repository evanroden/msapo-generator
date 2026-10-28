"""Three synthetic report variants verify deterministic accepted-state assembly.

These are audits of the current official assembler, not a replacement report
generator or permission to bypass preflight. No real client source is used.
"""
from dataclasses import replace
from io import BytesIO
from zipfile import ZipFile, ZipInfo, ZIP_DEFLATED
from hashlib import sha256
import shutil

import pytest
import pymupdf as fitz
from docx import Document
from PIL import Image

from app.monthly_report_docx import assemble_docx, generate_report
from app.monthly_report_model import (
    BlockSpec, ColumnSpec, ReportPeriod, ReportSource, ReportTable, ResolvedBlock,
    synthetic_draft, synthetic_profiles,
)
from app.monthly_report_reproducibility import (
    canonical_input_manifest, canonical_docx_manifest, verify_repeated_assembly,
)
from app.monthly_report_sources import source_reference


def illustration():
    image = BytesIO()
    Image.new("RGB", (160, 100), "navy").save(image, format="PNG")
    return image.getvalue()


def fixture_variant(kind):
    profiles = synthetic_profiles()
    profile = profiles[{"single": 0, "two_sites": 1, "regional_images": 2}[kind]]
    draft = synthetic_draft(profile, ReportPeriod(2026, 9))
    image = b""
    if kind == "single":
        return draft, None
    if kind == "two_sites":
        sections = tuple(replace(section, included=section.key not in
                        ("water", "training", "rfi")) for section in draft.sections)
        return replace(draft, sections=sections), None
    image = illustration()
    table = BlockSpec("work_orders", "table", columns=(
        ColumnSpec("wo", "WO #"), ColumnSpec("status", "Status")))
    activity = next(section for section in draft.sections if section.key == "activity")
    updated_activity = replace(activity, blocks=(
        BlockSpec("activity_summary", "rich_text"), table),
        divider_asset="synthetic-divider.png")
    sections = tuple(updated_activity if section.key == "activity" else section
                     for section in draft.sections)
    blocks = tuple(block for block in draft.blocks if block.key != "demo_activity") + (
        ResolvedBlock("activity_summary", "This month",
                      text="Synthetic operators verified two work orders."),
        ResolvedBlock("work_orders", "This month",
                      rows=(("0007", "Open"), ("0008", "Completed")),
                      extra_tables=(ReportTable(("Part", "Reading"),
                                                (("Pump", "0"), ("Valve", ""))),)),
    )
    return replace(draft, sections=sections, blocks=blocks), lambda ref: image


@pytest.mark.parametrize("kind", ("single", "two_sites", "regional_images"))
def test_three_report_variants_produce_stable_content_and_selected_sections(kind):
    draft, loader = fixture_variant(kind)
    receipt = verify_repeated_assembly(draft, asset_loader=loader)
    assert len(receipt["input_sha256"]) == len(receipt["docx_content_sha256"]) == 64
    assert receipt["sections"]
    assert receipt["sections"][0] == ("organization", "1")
    if kind == "two_sites":
        assert not any(k in ("water", "training", "rfi") for k, _ in receipt["sections"])
        assert ("proposals", "9") in receipt["sections"]
    if kind == "regional_images":
        assert ("activity", "2") in receipt["sections"]
        original = assemble_docx(draft, acknowledged_fingerprint=draft.fingerprint,
                                 asset_loader=loader)
        doc = Document(BytesIO(original))
        texts = " ".join(p.text for p in doc.paragraphs)
        assert "Synthetic operators verified two work orders." in texts
        assert any("0007" in cell.text and "0008" not in cell.text
                   for table in doc.tables for row in table.rows for cell in row.cells)
        manifest = canonical_docx_manifest(original)["manifest"]
        assert sha256(illustration()).hexdigest() in manifest["media_sha256"]
    assert not draft.sources  # no private or unreviewed source interpretation


def test_content_manifest_ignores_volatile_package_timestamps_but_detects_rows():
    draft, loader = fixture_variant("regional_images")
    original = assemble_docx(draft, acknowledged_fingerprint=draft.fingerprint,
                             asset_loader=loader)
    before = canonical_docx_manifest(original)
    output = BytesIO()
    with ZipFile(BytesIO(original)) as z, ZipFile(output, "w") as copy:
        for item in z.infolist():
            info = ZipInfo(item.filename, date_time=(2025, 1, 1, 0, 0, 0))
            info.compress_type = ZIP_DEFLATED
            copy.writestr(info, z.read(item.filename))
    assert original != output.getvalue()
    assert canonical_docx_manifest(output.getvalue()) == before
    changed = replace(draft, blocks=tuple(
        replace(b, rows=(("0007", "Open"), ("0008", "Open")))
        if b.key == "work_orders" else b for b in draft.blocks))
    other = assemble_docx(changed, acknowledged_fingerprint=changed.fingerprint,
                          asset_loader=loader)
    assert canonical_docx_manifest(other)["sha256"] != before["sha256"]


def test_unrelated_source_interpretation_is_not_promoted_into_accepted_report():
    draft, _ = fixture_variant("single")
    original = assemble_docx(draft, acknowledged_fingerprint=draft.fingerprint)
    incoming = ReportSource("synthetic-one", "unapproved.csv", "a" * 64, ".csv",
                            page_texts=("Proposed unchecked claim.",))
    different = replace(incoming, page_texts=("Other unreviewed interpretation.",))
    first = replace(draft, sources=(incoming,))
    second = replace(draft, sources=(different,))
    one = assemble_docx(first, acknowledged_fingerprint=first.fingerprint)
    two = assemble_docx(second, acknowledged_fingerprint=second.fingerprint)
    assert canonical_docx_manifest(original) == canonical_docx_manifest(one)
    assert canonical_docx_manifest(one) == canonical_docx_manifest(two)
    assert canonical_input_manifest(first) == canonical_input_manifest(second)


def test_selected_source_provenance_is_pinned_without_model_reinterpretation():
    draft, _ = fixture_variant("single")
    source = ReportSource("synthetic-one", "synthetic-evidence.pdf", "a" * 64, ".pdf",
                          page_texts=("Original supporting report page.",),
                          selected_pages=(1,))
    ref = source_reference(source, 1)
    activity = next(s for s in draft.sections if s.key == "activity")
    amended = replace(activity, blocks=(BlockSpec("activity_summary", "rich_text"),))
    revised = replace(draft, sections=tuple(amended if s.key == "activity" else s
                                          for s in draft.sections),
                      blocks=(*draft.blocks, ResolvedBlock(
                          "activity_summary", "This month",
                          text="A human accepted this single synthetic observation.",
                          references=(ref,))),
                      sources=(source,))
    record = canonical_input_manifest(revised)["manifest"]
    assert record["sources"] == [(source.id, source.sha256, source.fingerprint)]
    assert record["sections"][1]["blocks"][0][3]["references"] == (ref,)
    # Original source metadata changes invalidate the provenance identity,
    # but do not automatically revise the human-accepted paragraph.
    changed_source = replace(source, page_texts=("A different proposed claim.",))
    changed = replace(revised, sources=(changed_source,))
    assert canonical_input_manifest(changed)["sha256"] != canonical_input_manifest(revised)["sha256"]
    assert next(b for b in changed.blocks if b.key == "activity_summary").text == (
        "A human accepted this single synthetic observation.")


def test_generation_never_needs_a_new_source_reader_response(monkeypatch):
    import app.monthly_report_sources as sources
    import app.ocr as ocr

    def forbidden(*args, **kwargs):
        pytest.fail("Accepted report assembly called an unapproved source reader")

    monkeypatch.setattr(sources, "ingest", forbidden)
    monkeypatch.setattr(sources, "suggest_facts", forbidden)
    monkeypatch.setattr(ocr, "run_ocr", forbidden, raising=False)
    draft, loader = fixture_variant("regional_images")
    result = verify_repeated_assembly(draft, asset_loader=loader)
    assert result["sections"][1] == ("activity", "2")


@pytest.mark.skipif(not shutil.which("libreoffice"), reason="LibreOffice conversion unavailable")
def test_real_downloadable_word_and_pdf_semantics_reproduce_for_same_report(monkeypatch, tmp_path):
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path))
    draft, loader = fixture_variant("two_sites")
    first = generate_report(draft, acknowledged_fingerprint=draft.fingerprint,
                            asset_loader=loader)
    second = generate_report(draft, acknowledged_fingerprint=draft.fingerprint,
                             asset_loader=loader)
    assert first.pdf and second.pdf, (first.pdf_error, second.pdf_error)
    assert canonical_docx_manifest(first.docx) == canonical_docx_manifest(second.docx)
    with fitz.open(stream=first.pdf, filetype="pdf") as left, fitz.open(
            stream=second.pdf, filetype="pdf") as right:
        assert len(left) == len(right)
        assert tuple(page.get_text() for page in left) == tuple(page.get_text() for page in right)
        assert tuple((page.rect.width, page.rect.height) for page in left) == tuple(
            (page.rect.width, page.rect.height) for page in right)
