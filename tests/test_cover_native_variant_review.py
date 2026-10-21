"""Preserved native company variants must survive fresh-master cleanup."""

from dataclasses import replace
from copy import deepcopy
import hashlib
from io import BytesIO

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches
from lxml import etree
import pytest

from app.monthly_report_asset_review import approve_all_assets
from app.monthly_report_import import ImportMapping, inspect_docx, map_items
from app.monthly_report_model import ReportDraft, ReportPeriod, ResolvedBlock, default_sections, synthetic_profiles
from app.monthly_report_native_layout import build_native_docx
from test_cover_branding_contracts import mark


def _source(tmp_path, *, part_name="header", identified=True, duplicate=False, alternate=False):
    document = Document()
    document.add_paragraph(synthetic_profiles()[0].facilities[0].title)
    document.add_paragraph("November 2026")
    primary, variant = mark(), mark((110, 70, 160, 255))
    cover = document.add_picture(BytesIO(primary), width=Inches(3))
    cover._inline.docPr.set("descr", "ENFRA company logo")
    part = getattr(document.sections[0], part_name)
    shape = part.paragraphs[0].add_run().add_picture(BytesIO(variant), width=Inches(2))
    if identified:
        shape._inline.docPr.set("descr", "ENFRA company logo")
    if alternate:
        drawing = shape._inline.getparent()
        run = drawing.getparent()
        mc = "{http://schemas.openxmlformats.org/markup-compatibility/2006}"
        wrapper = etree.Element(mc + "AlternateContent")
        run.replace(drawing, wrapper)
        etree.SubElement(wrapper, mc + "Choice", Requires="w").append(drawing)
        etree.SubElement(wrapper, mc + "Fallback").append(deepcopy(drawing))
    if duplicate:
        part.add_paragraph().add_run().add_picture(BytesIO(variant), width=Inches(1))
    document.add_section()
    document.add_heading("Monthly Activity Summary", 1)
    document.add_paragraph("Synthetic inspection completed.")
    source = tmp_path / "native-brand-variants.docx"
    document.save(source)
    inspection = inspect_docx(source)
    cover_item = next(item for item in inspection.items
                      if item.kind == "image" and item.part == "word/document.xml")
    header_items = tuple(item for item in inspection.items
                         if item.kind == "image" and item.part.startswith("word/" + part_name))
    mapped = map_items(source, inspection, (ImportMapping(cover_item.id, "brand_logo"),))
    brand = approve_all_assets(mapped.blocks[0])
    sections = tuple(replace(section, included=section.key == "activity") for section in default_sections())
    draft = ReportDraft(
        synthetic_profiles()[0], ReportPeriod(2026, 11), "Synthetic Reviewer", sections,
        (brand, ResolvedBlock("activity_summary", "This month", text="Synthetic inspection completed.")),
    )
    assert hashlib.sha256(variant).hexdigest() != hashlib.sha256(primary).hexdigest()
    return source, inspection, header_items, draft, dict(mapped.assets), variant


def _part_images(raw, part_name):
    document = Document(BytesIO(raw))
    return tuple(part.related_parts[node.get(qn("r:embed"))].blob
                 for part in document.part.package.parts
                 if str(part.partname).startswith("/word/" + part_name) and hasattr(part, "element")
                 for node in part.element.iter(qn("a:blip")))


@pytest.mark.parametrize("part_name", ["header", "footer"])
@pytest.mark.parametrize("master", [False, True])
def test_unchanged_cover_retains_identified_native_variant(tmp_path, part_name, master):
    source, _, items, draft, assets, variant = _source(tmp_path, part_name=part_name)
    original = source.read_bytes()
    assert items[0].image_role == "brand_logo"
    actual = _part_images(build_native_docx(
        draft, source, asset_loader=assets.__getitem__, master=master), part_name)
    assert actual == (variant,), "Preserving a native company variant must survive final master cleanup"
    assert source.read_bytes() == original


@pytest.mark.parametrize("master", [False, True])
def test_explicit_mapping_replaces_variant_even_when_cover_is_unchanged(tmp_path, master):
    source, inspection, items, draft, assets, variant = _source(tmp_path)
    brand = draft.blocks[0]
    brand = replace(brand, references=(*brand.references, f"docx:{inspection.sha256}:{items[0].id}"))
    draft = replace(draft, blocks=(brand, *draft.blocks[1:]))
    actual = _part_images(build_native_docx(
        draft, source, asset_loader=assets.__getitem__, master=master), "header")
    assert actual == (assets[brand.asset_hashes[0]],)
    assert actual != (variant,)


@pytest.mark.parametrize("master", [False, True])
def test_unknown_variant_is_not_approved_by_an_unchanged_cover(tmp_path, master):
    source, _, items, draft, assets, variant = _source(tmp_path, identified=False)
    assert items[0].image_role == "header_artwork"
    actual = _part_images(build_native_docx(
        draft, source, asset_loader=assets.__getitem__, master=master), "header")
    assert actual == (() if master else (variant,))


@pytest.mark.parametrize("mapped_other_role", [False, True])
@pytest.mark.parametrize("alternate", [False, True])
def test_preserving_one_variant_does_not_approve_an_independent_equal_byte_occurrence(tmp_path, mapped_other_role, alternate):
    source, inspection, items, draft, assets, variant = _source(tmp_path, duplicate=True, alternate=alternate)
    assert items[0].image_role == "brand_logo"
    assert items[1].image_role == "header_artwork"
    if mapped_other_role:
        other = ResolvedBlock("client_logo", "Omit",
                              references=(f"docx:{inspection.sha256}:{items[1].id}",))
        draft = replace(draft, blocks=(*draft.blocks, other))
    actual = _part_images(build_native_docx(
        draft, source, asset_loader=assets.__getitem__, master=True), "header")
    assert actual == (variant,) * (2 if alternate else 1), \
        "Equivalent representations survive, but cannot authorize another use of equal bytes"
