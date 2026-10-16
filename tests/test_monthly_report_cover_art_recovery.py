"""Unread imported cover artwork cannot be certified by an optional photo."""

from io import BytesIO
from dataclasses import replace

from PIL import Image
from docx import Document
from streamlit.testing.v1 import AppTest
import pytest

from app import monthly_report_designs as designs, monthly_report_library as library
from app import monthly_report_section_ui as ui
from app.monthly_report_asset_review import approve_all_assets
from app.monthly_report_model import ReportDraft, ReportPeriod, ResolvedBlock, profile_sections, synthetic_profiles
from app.monthly_report_start import save_design_start
from test_monthly_report_starting_review_recovery import drawing_report, continue_button


APP = '''
import streamlit as st
from app import monthly_report_library as library
from app.monthly_report_section_ui import render_section_setup
from app.monthly_report_ui import _field
draft=st.session_state['working']
state=library.load_profile(draft.profile.contract,draft.profile.key)
render_section_setup(draft.profile.contract,draft.period,'Synthetic Editor',_field,
 state=state,identity=draft.profile,working_draft=draft,
 working_assets=tuple(st.session_state['assets'].items()),working_revision=0)
from app.monthly_report_section_ui import render_pending_cover_selection
proposed,_=render_pending_cover_selection(draft,st.session_state['assets'],draft.profile.contract,draft.period,state)
st.session_state['preview_draft']=proposed
'''


def existing_app(monkeypatch, tmp_path, source=None, *, missing_brand=False):
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path / "runtime"))
    if source is None:
        source, _ = drawing_report(tmp_path, "cover")
        document = Document(source)
        document.add_heading("Activity Summary", 1)
        document.add_paragraph("Reviewed source valve update.")
        document.save(source)
    profile = designs.pin(synthetic_profiles()[0])
    assets, blocks = {}, []
    for slot, color in (("brand_logo", "red"), ("client_logo", "green"), ("cover_photo", "purple")):
        if missing_brand and slot == "brand_logo":
            continue
        stream = BytesIO()
        Image.new("RGB", (160, 120), color).save(stream, "PNG")
        raw = stream.getvalue()
        reference = library.asset_reference(raw, "png")
        assets[reference] = raw
        blocks.append(approve_all_assets(ResolvedBlock(slot, "This month", asset_hashes=(reference,))))
    blocks.extend((ResolvedBlock("activity_summary", "This month", text="Current human pump update."),
                   ResolvedBlock("footer_text", "This month", text="Current approved address")))
    working = ReportDraft(profile, ReportPeriod(2026, 9), "Synthetic Editor", profile_sections(profile),
                          tuple(blocks), address_line="Current approved address")
    saved = save_design_start(working, assets, actor="Synthetic Editor", confirmed=True)
    uploaded = BytesIO(source.read_bytes())
    uploaded.name, uploaded.size = source.name, len(uploaded.getvalue())
    real = ui.st.file_uploader
    monkeypatch.setattr(ui.st, "file_uploader", lambda label, *a, **kw:
                        uploaded if label == "Starting report" else real(label, *a, **kw))
    app = AppTest.from_string(APP, default_timeout=30)
    app.session_state["working"] = working
    app.session_state["assets"] = assets
    return app.run(), working, saved, source, assets


@pytest.mark.parametrize("unknown_raster", [False, True])
def test_existing_cover_omission_route_keeps_current_design_and_original(monkeypatch, tmp_path, unknown_raster):
    source = small_art_report(tmp_path) if unknown_raster else None
    app, working, saved, source, assets = existing_app(monkeypatch, tmp_path, source)
    assert not app.exception
    assert continue_button(app).disabled
    action = next(r for r in app.radio if r.label == "What would you like to do?")
    assert "Use this report’s cover design" in action.options
    action.set_value("Leave this section out").run()
    assert not app.exception
    assert not continue_button(app).disabled
    assert app.session_state["working"] == working
    assert library.load_profile(working.profile.contract, working.profile.key) == saved
    continue_button(app).click().run()
    assert not app.exception
    assert not app.error, [e.value for e in app.error]
    imported = library.load_imported_draft(working.profile.contract, working.profile.key)
    assert imported.profile == working.profile
    for slot in ("brand_logo", "client_logo", "cover_photo", "footer_text"):
        assert next(b for b in imported.blocks if b.key == slot) == next(b for b in working.blocks if b.key == slot)
    assert imported.address_line == working.address_line
    assert "Current human pump update." in next(b.text for b in imported.blocks if b.key == "activity_summary")
    assert "Reviewed source valve update." in next(b.text for b in imported.blocks if b.key == "activity_summary")
    assert not any("Synthetic equipment decision diagram" in str(b) for b in imported.blocks)
    assert app.session_state["working"] == working
    original = library.imported_original(working.profile.contract, working.profile.key)
    assert original is not None
    assert original.read_bytes() == source.read_bytes()
    for reference, raw in assets.items():
        assert library.read_asset(working.profile.contract, working.profile.key, reference) == raw
    # Omission is absent from this new source's contribution, never approval
    # of its unread drawing, and never a report-level Clear of the current art.
    from app.monthly_report_import import inspect_docx
    contribution, verified = library.starting_contribution(working.profile.contract, working.profile.key,
                                                         working.period, inspect_docx(source).sha256)
    assert verified and contribution is not None
    assert not any(b.key in {"brand_logo", "client_logo", "cover_photo", "footer_text"} for b in contribution.blocks)


def test_optional_photo_does_not_certify_unread_cover_drawing(monkeypatch, tmp_path):
    from tests.test_monthly_report_start_import import start_app
    source, _ = drawing_report(tmp_path, "cover")
    original = source.read_bytes()
    app = start_app(monkeypatch, tmp_path, source)
    next(r for r in app.radio if r.label == "What would you like to do?").set_value("Edit text or change pictures").run()
    incoming = BytesIO()
    Image.new("RGB", (450, 300), "orange").save(incoming, "PNG")
    incoming.name, incoming.size = "optional-cover-photo.png", len(incoming.getvalue())
    real = ui.st.file_uploader
    monkeypatch.setattr(ui.st, "file_uploader", lambda label, *a, **kw:
                        incoming if label == "New cover photograph" else real(label, *a, **kw))
    app.run()
    assert not app.exception
    assert continue_button(app).disabled, "An optional photograph approved unrelated unread VML."
    plan = next(v["cover"] for k, v in app.session_state.filtered_state.items() if k.endswith("_section_plans"))
    assert not plan["approved"] and not plan.get("unsupported_reviewed")
    assert plan["new_assets"]
    assert source.read_bytes() == original
    assert not library.list_profiles(synthetic_profiles()[0].contract)


def test_fresh_cover_omission_cannot_pin_the_unread_source_as_its_design(monkeypatch, tmp_path):
    from tests.test_monthly_report_start_import import start_app
    source, _ = drawing_report(tmp_path, "cover")
    document = Document(source)
    document.add_heading("Activity Summary", 1)
    document.add_paragraph("Reviewed source valve update.")
    document.save(source)
    original = source.read_bytes()
    app = start_app(monkeypatch, tmp_path, source)
    next(r for r in app.radio if r.label == "What would you like to do?").set_value("Leave this section out").run()
    assert not app.exception
    assert continue_button(app).disabled, "The unread DOCX would still become the new report's native design."
    assert any("original Word" in e.value for e in app.error)
    assert source.read_bytes() == original
    assert not library.list_profiles(synthetic_profiles()[0].contract)


def small_art_report(tmp_path, *, contents=False, header=False):
    from docx.enum.section import WD_SECTION
    from docx.shared import Inches
    document = Document()
    document.add_heading(synthetic_profiles()[0].facilities[0].title, 0)
    document.add_paragraph("September 2026")
    stream = BytesIO()
    Image.new("RGB", (64, 96), "orange").save(stream, "PNG")
    document.add_picture(BytesIO(stream.getvalue()), width=Inches(1))
    if header:
        stream = BytesIO()
        Image.new("RGB", (300, 70), "red").save(stream, "PNG")
        document.sections[0].header.paragraphs[0].add_run().add_picture(BytesIO(stream.getvalue()), width=Inches(2))
    if contents:
        document.add_section(WD_SECTION.NEW_PAGE)
        document.add_heading("Contents", 0)
        stream = BytesIO()
        Image.new("RGB", (700, 400), "green").save(stream, "PNG")
        document.add_picture(BytesIO(stream.getvalue()), width=Inches(4))
    document.add_heading("Activity Summary", 1)
    document.add_paragraph("Reviewed source valve update.")
    source = tmp_path / "small-cover-art.docx"
    document.save(source)
    return source


def test_unassigned_small_raster_has_specific_recovery_instead_of_silent_omission(monkeypatch, tmp_path):
    from tests.test_monthly_report_start_import import start_app
    source = small_art_report(tmp_path)
    original = source.read_bytes()
    app = start_app(monkeypatch, tmp_path, source)
    assert not app.exception
    assert continue_button(app).disabled, "Unidentified visible cover artwork was silently dropped."
    assert any("Cover artwork 1" in e.value for e in app.error)
    assert any(source.name in caption.value for caption in app.caption)
    assert not any(b.label == "No cover photograph" for b in app.button)
    assert source.read_bytes() == original
    assert not library.list_profiles(synthetic_profiles()[0].contract)


def test_contents_image_is_not_promoted_when_earliest_cover_has_only_small_art(monkeypatch, tmp_path):
    from tests.test_monthly_report_start_import import start_app
    source = small_art_report(tmp_path, contents=True)
    app = start_app(monkeypatch, tmp_path, source)
    assert not app.exception
    assert not any(b.label == "No cover photograph" for b in app.button), "A Contents image became the cover-photo gallery."
    assert continue_button(app).disabled


def foreign_cover_report(tmp_path, *, logo=True):
    from docx.shared import Inches
    source, _ = drawing_report(tmp_path, "cover")
    document = Document(source)
    document.paragraphs[0].text = synthetic_profiles()[1].facilities[1].title
    if logo:
        stream = BytesIO()
        Image.new("RGB", (300, 70), "teal").save(stream, "PNG")
        document.sections[0].header.paragraphs[0].add_run().add_picture(BytesIO(stream.getvalue()), width=Inches(2))
    document.add_heading("Activity Summary", 1)
    document.add_paragraph("Foreign site motor work must not reach the target report.")
    document.save(source)
    return source


def test_excluded_foreign_art_keeps_only_recognized_source_branding_with_current_design(monkeypatch, tmp_path):
    from app.monthly_report_asset_review import pending_asset_indexes
    source = foreign_cover_report(tmp_path)
    original = source.read_bytes()
    monkeypatch.setattr(ui, "_site_options", lambda contract: synthetic_profiles()[1].facilities)
    app, working, saved, _, assets = existing_app(monkeypatch, tmp_path, source, missing_brand=True)
    assert not app.exception
    assert not continue_button(app).disabled
    assert not app.radio
    assert library.load_profile(working.profile.contract, working.profile.key) == saved
    continue_button(app).click().run()
    assert not app.exception and not app.error
    imported = library.load_imported_draft(working.profile.contract, working.profile.key)
    assert imported.profile == working.profile
    assert not any("Foreign site motor" in str(b) or "Synthetic equipment decision diagram" in str(b) for b in imported.blocks)
    for slot in ("client_logo", "cover_photo", "footer_text"):
        assert next(b for b in imported.blocks if b.key == slot) == next(b for b in working.blocks if b.key == slot)
    brand = next(b for b in imported.blocks if b.key == "brand_logo")
    assert len(brand.asset_hashes) == 1 and pending_asset_indexes(brand) == (0,)
    assert app.session_state["working"] == working
    assert library.imported_original(working.profile.contract, working.profile.key).read_bytes() == original
    assert all(library.read_asset(working.profile.contract, working.profile.key, key) == raw for key, raw in assets.items())


@pytest.mark.parametrize("logo", [False, True])
def test_fresh_foreign_source_cannot_hide_unread_cover_in_new_native_design(monkeypatch, tmp_path, logo):
    from tests.test_monthly_report_start_import import start_app
    source = foreign_cover_report(tmp_path, logo=logo)
    original = source.read_bytes()
    app = start_app(monkeypatch, tmp_path, source, known_sites=synthetic_profiles()[1].facilities)
    assert not app.exception
    assert continue_button(app).disabled
    assert any("original Word" in error.value and source.name in error.value for error in app.error)
    assert source.read_bytes() == original
    assert not library.list_profiles(synthetic_profiles()[0].contract)


@pytest.mark.parametrize("tamper", ["source", "site", "photo", "new_photo"])
def test_foreign_art_exclusion_receipt_cannot_cover_stale_or_unrecognized_material(tmp_path, tamper):
    from app.monthly_report_import import inspect_docx
    from app.monthly_report_start_import import analyze_starting_report, automatic_section_plans, build_starting_import
    from app.monthly_report_docx import normalize_report_image
    source = foreign_cover_report(tmp_path)
    inspection = inspect_docx(source)
    analysis = analyze_starting_report(inspection, synthetic_profiles()[0], ReportPeriod(2026, 9), synthetic_profiles()[1].facilities)
    assert analysis.site_relation == "other"
    plans = automatic_section_plans(inspection, analysis)
    plan = plans["cover"]
    mapped, _ = build_starting_import(source, inspection, plans.values())
    assert {block.key for block in mapped.blocks} == {"brand_logo"}
    if tamper == "source":
        plan["source_sha256"] = "0" * 64
    elif tamper == "site":
        plan["source_site_relation"] = "same"
    elif tamper == "photo":
        photo = next(i for i in inspection.items if i.kind == "image" and i.suggested_slot != "brand_logo")
        plan["selected"].append(photo.id)
        plan["destinations"][photo.id] = "brand_logo"
    else:
        image = BytesIO()
        Image.new("RGB", (400, 300), "green").save(image, "PNG")
        plan["new_assets"] = [("cover_photo", normalize_report_image(image.getvalue(), ".png"))]
    with pytest.raises(ValueError, match="[Cc]over artwork"):
        build_starting_import(source, inspection, plans.values())


def test_recognized_fixed_header_logo_keeps_its_role_beside_one_gallery(monkeypatch, tmp_path):
    from docx.shared import Inches
    from app.monthly_report_import import inspect_docx
    from tests.test_monthly_report_start_import import start_app
    source, _ = drawing_report(tmp_path, "cover")
    document = Document(source)
    stream = BytesIO()
    Image.new("RGB", (300, 70), "red").save(stream, "PNG")
    document.sections[0].header.paragraphs[0].add_run().add_picture(BytesIO(stream.getvalue()), width=Inches(2))
    document.save(source)
    original = source.read_bytes()
    logo = next(item for item in inspect_docx(source).items if item.suggested_slot == "brand_logo")
    app = start_app(monkeypatch, tmp_path, source)
    for _ in range(2):
        assert not app.exception
        plan = next(v["cover"] for k, v in app.session_state.filtered_state.items() if k.endswith("_section_plans"))
        assert logo.id in plan["selected"]
        assert plan["destinations"][logo.id] == "brand_logo"
        assert sum(b.label == "No cover photograph" for b in app.button) == 1
        assert not any("Other imported cover artwork" == x.value.strip("*") for x in app.markdown)
        assert continue_button(app).disabled
        assert source.read_bytes() == original
        assert not library.list_profiles(synthetic_profiles()[0].contract)
        app.run()


@pytest.mark.parametrize("unread_drawing", [False, True])
def test_direct_builder_cannot_forge_cover_review_with_photo_or_generic_boolean(tmp_path, unread_drawing):
    from app.monthly_report_import import inspect_docx
    from app.monthly_report_sections import build_section_import
    from app.monthly_report_docx import normalize_report_image
    source = drawing_report(tmp_path, "cover")[0] if unread_drawing else small_art_report(tmp_path)
    inspection = inspect_docx(source)
    image = BytesIO()
    Image.new("RGB", (400, 300), "green").save(image, "PNG")
    photo = normalize_report_image(image.getvalue(), ".png")
    plan = {"key": "cover", "target": "cover", "approved": True, "selected": [],
            "new_assets": [("cover_photo", photo)], "unsupported_reviewed": True}
    with pytest.raises(ValueError, match="[Cc]over artwork"):
        build_section_import(source, inspection, (plan,))


def test_saved_cover_choice_does_not_approve_a_later_upload(monkeypatch, tmp_path):
    app, working, saved, source, _ = existing_app(monkeypatch, tmp_path)
    next(r for r in app.radio if r.label == "What would you like to do?").set_value("Leave this section out").run()
    assert not continue_button(app).disabled
    document = Document(source)
    document.add_paragraph("Changed source bytes require another source-bound decision.")
    document.save(source)
    incoming = BytesIO(source.read_bytes())
    incoming.name, incoming.size = source.name, len(incoming.getvalue())
    real = ui.st.file_uploader
    monkeypatch.setattr(ui.st, "file_uploader", lambda label, *a, **kw:
                        incoming if label == "Starting report" else real(label, *a, **kw))
    app.run()
    assert not app.exception
    assert continue_button(app).disabled
    assert next(r for r in app.radio if r.label == "What would you like to do?").value == "Keep and review"
    assert app.session_state["working"] == working
    assert library.load_profile(working.profile.contract, working.profile.key) == saved


def test_current_cover_recovery_keeps_preview_truthful_and_can_be_reversed(monkeypatch, tmp_path):
    app, working, saved, source, _ = existing_app(monkeypatch, tmp_path)
    raw = source.read_bytes()
    current_photo = next(b for b in working.blocks if b.key == "cover_photo")
    next(r for r in app.radio if r.label == "What would you like to do?").set_value("Leave this section out").run()
    app.run()
    assert not app.exception
    preview_photo = next(b for b in app.session_state["preview_draft"].blocks if b.key == "cover_photo")
    assert preview_photo == current_photo, "A deliberately omitted source photo was still proposed in the preview."
    assert not continue_button(app).disabled
    assert not any(b.label == "No cover photograph" for b in app.button)
    action = next(r for r in app.radio if r.label == "What would you like to do?")
    assert action.value == "Leave this section out"
    action.set_value("Keep and review").run()
    app.run()
    assert not app.exception
    assert continue_button(app).disabled
    assert sum(b.label == "No cover photograph" for b in app.button) == 1
    assert source.read_bytes() == raw
    assert app.session_state["working"] == working
    assert library.load_profile(working.profile.contract, working.profile.key) == saved


def test_saved_cover_choice_does_not_waive_price_or_image_review(monkeypatch, tmp_path):
    from app.monthly_report_checks import preflight
    app, working, _, _, _ = existing_app(monkeypatch, tmp_path)
    next(r for r in app.radio if r.label == "What would you like to do?").set_value("Leave this section out").run()
    continue_button(app).click().run()
    assert not app.error and not app.exception
    imported = library.load_imported_draft(working.profile.contract, working.profile.key)
    photo = next(b for b in imported.blocks if b.key == "cover_photo")
    unsafe_photo = replace(photo, client_asset_reviews=(), client_reviewed_fingerprint="", asset_captions=("Quoted price $2,500",))
    unsafe = replace(imported, blocks=tuple(unsafe_photo if b.key == "cover_photo" else b for b in imported.blocks))
    codes = {(c.code, c.block_key) for c in preflight(unsafe) if c.blocking}
    assert ("client_pages", "cover_photo") in codes
    # Layout picture captions are checked during output package preparation;
    # ordinary narrative pricing must still be a preflight blocker here.
    unsafe = replace(unsafe, address_line="Quoted price $2,500")
    assert any(c.blocking and c.code == "pricing" for c in preflight(unsafe))
