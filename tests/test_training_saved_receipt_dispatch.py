"""Mixed saved receipts retain structural integrity without granting authority."""
from copy import deepcopy
import json
from dataclasses import asdict, replace

import pytest
from app import monthly_report_library as library
from app import monthly_report_training_review_flow as training
from app import monthly_report_work_order_review_flow as workorders
from app.monthly_report_row_review_receipts import decoded
from app.monthly_report_model import ReportPeriod
from app.monthly_report_setup import new_month_draft
from test_training_review_flow import report, prepared, selected, publish
from test_monthly_report_work_order_review_flow import draft_source


@pytest.fixture(autouse=True)
def storage(monkeypatch, tmp_path):
    monkeypatch.setenv('EPC_DATA_DIR', str(tmp_path))


def test_actual_mixed_saved_snapshot_reopens_without_receipt_or_source_changes():
    draft, value = selected(*prepared(report()))
    draft = publish(draft, value)
    other, content = draft_source()
    bundle = workorders.prepare_bundle(other, content)
    raw = workorders.encoded(bundle)
    mixed = replace(draft, source_row_reviews=(*draft.source_row_reviews, raw),
                    sources=(*draft.sources, content.source))
    original = asdict(mixed)
    saved = library.save_snapshot(mixed, expected_revision=0, entered_editor='Synthetic Editor')
    reopened = library.load_snapshot(mixed.profile.contract, mixed.profile.key, mixed.period)
    assert reopened == saved
    assert asdict(reopened.draft) == original
    assert library.load_snapshot_version(mixed.profile.contract, mixed.profile.key, mixed.period, 1) == saved
    assert decoded(raw) == bundle
    later = new_month_draft(reopened.draft, ReportPeriod(2026, 10))
    assert later.source_row_reviews == ()
    assert new_month_draft(later, ReportPeriod(2026, 11)).source_row_reviews == ()
    assert reopened.draft.source_row_reviews == mixed.source_row_reviews


def test_stale_intact_training_receipt_is_retained_but_never_authorizes_publication():
    draft, value = prepared(report())
    stale = {**value, 'scope': 'old report incarnation'}
    raw = training.encoded(stale)
    restored = library.draft_from_dict(asdict(replace(draft, source_row_reviews=(raw,))))
    assert restored.source_row_reviews == (raw,)
    assert decoded(raw) == json.loads(raw)
    with pytest.raises(ValueError, match='different report or month'):
        training.validate(restored, stale)
    changed = replace(draft, sources=(replace(draft.sources[0], page_texts=('Changed source',)),))
    with pytest.raises(ValueError, match='changed or was removed'):
        training.validate(changed, value)


@pytest.mark.parametrize('field,bad', [
    ('kind', 'unknown'), ('schema', True), ('table', True), ('table', 0),
    ('history', {}), ('history', ['missing row']), ('history', [1]),
    ('record', []), ('record', {'schema': 1}), ('published', []),
    ('published', {'decision': 'x', 'block': 'y', 'before': {},
                   'replace_source': False, 'replace_matrix': False}),
    ('source_id', 'different source'), ('source_fingerprint', None),
])
def test_malformed_training_receipt_is_controlled_and_never_erases_original(field, bad):
    draft, value = prepared(report())
    value = deepcopy(value)
    value[field] = bad
    raw = training.encoded(value)
    record = asdict(replace(draft, source_row_reviews=(raw,)))
    original = deepcopy(record)
    with pytest.raises(library.LibraryError):
        library.draft_from_dict(record)
    assert record == original


@pytest.mark.parametrize('raw', ['null', '[]', '{broken', '"training"', 'x' * 4_000_001])
def test_dispatch_rejects_nonobject_invalid_json_and_per_receipt_overflow(raw):
    with pytest.raises(ValueError, match='saved row review'):
        decoded(raw)


def test_training_integrity_corruption_and_aggregate_limits_remain_closed():
    draft, value = prepared(report())
    damaged = json.loads(training.encoded(value))
    damaged['record']['review']['rows'][0]['values'][0] = 'Changed without integrity receipt'
    with pytest.raises(ValueError):
        decoded(training.encoded(damaged))
    raw = training.encoded(value)
    for receipts in ((raw,) * 101, ('x' * 4_000_000,) * 5):
        with pytest.raises(library.LibraryError):
            library.draft_from_dict(asdict(replace(draft, source_row_reviews=receipts)))


@pytest.mark.parametrize('target,key,bad', [
    ('record', 'schema', True),
    ('review', 'native_layout_confirmed', 'yes'),
    ('source', 'native_word', 1),
    ('source', 'number', True),
    ('row', 'excluded', 'false'),
])
def test_recomputed_integrity_hash_does_not_bypass_structural_types(target, key, bad):
    from hashlib import sha256
    draft, value = prepared(report())
    value = json.loads(training.encoded(value))
    review = value['record']['review']
    container = {'record': value['record'], 'review': review,
                 'source': review['source'], 'row': review['rows'][0]}[target]
    container[key] = bad
    value['record']['sha256'] = sha256(json.dumps(review, ensure_ascii=False,
                                                sort_keys=True).encode()).hexdigest()
    raw = training.encoded(value)
    with pytest.raises(ValueError):
        decoded(raw)
    with pytest.raises(library.LibraryError):
        library.draft_from_dict(asdict(replace(draft, source_row_reviews=(raw,))))


def test_mixed_inmemory_receipts_allow_each_native_review_without_cross_kind_authority():
    training_draft, training_value = prepared(report())
    training_raw = training_draft.source_row_reviews[0]
    work_draft, content = draft_source()
    mixed_work = replace(work_draft, source_row_reviews=(training_raw,))
    bundle = workorders.prepare_bundle(mixed_work, content)
    stored_work = workorders.store_bundle(mixed_work, bundle)
    assert stored_work.source_row_reviews[0] == training_raw
    assert workorders.saved_bundle(stored_work, content.source.id, 'workorders') == bundle
    replacement = workorders.prepare_bundle(stored_work, content)
    updated_work = workorders.store_bundle(stored_work, replacement,
                                         expected_revision=workorders.revision(bundle))
    assert updated_work.source_row_reviews == stored_work.source_row_reviews
    work_raw = stored_work.source_row_reviews[1]
    mixed_training = replace(training_draft, source_row_reviews=(work_raw, training_raw))
    found = training.saved_review(mixed_training, mixed_training.sources[0], 1)
    assert found == json.loads(training_raw)
    updated = training.store_review(mixed_training, training_value,
                                    training.revision(training_value))
    assert updated.source_row_reviews == mixed_training.source_row_reviews
    # The other kind is structurally retained, never accepted as this source's review.
    assert workorders.saved_bundle(updated, mixed_training.sources[0].id, 'workorders') is None


@pytest.mark.parametrize('path', [
    ('issues',), ('unmatched',), ('rows',), ('original_columns',),
    ('original_rows',), ('normalized_columns',), ('source', 'columns'),
    ('source', 'rows'), ('rows', 0, 'issues'), ('rows', 0, 'corrections'),
    ('rows', 0, 'values'), ('rows', 0, 'origin', 'columns'),
    ('rows', 0, 'origin', 'cells'),
])
def test_raw_review_container_strings_cannot_be_normalized_into_valid_lists(path):
    draft, value = prepared(report())
    value = json.loads(training.encoded(value))
    target = value['record']['review']
    for key in path[:-1]:
        target = target[key]
    target[path[-1]] = ''
    raw = training.encoded(value)
    with pytest.raises(ValueError):
        decoded(raw)
    with pytest.raises(library.LibraryError):
        library.draft_from_dict(asdict(replace(draft, source_row_reviews=(raw,))))


def test_same_original_three_kinds_replace_publish_and_remove_only_their_output():
    from io import BytesIO
    from openpyxl import Workbook
    from app.monthly_report_model import ResolvedBlock, default_sections
    from app.monthly_report_training_review_flow import source_content
    from test_monthly_report_work_order_review_flow import publish as publish_work

    book = Workbook()
    sheet = book.active
    sheet.title = 'Training'
    sheet.append(['Topic', 'Date', 'Method', 'Participants', 'Hours', 'Status'])
    sheet.append(['Current drill', '2026-09-13', 'Virtual', 'Alice', '1', 'Completed'])
    orders = book.create_sheet('Orders')
    orders.append(['Work Order', 'Site', 'Date', 'Status', 'Type', 'Trade'])
    orders.append(['0001', 'Synthetic North', '2026-09-03', 'Open', 'PM', 'HVAC'])
    maintenance = book.create_sheet('Maintenance')
    maintenance.append(['WO #', 'Finish Date', 'Area', 'Task', 'Tag', 'Site', 'Status'])
    maintenance.append(['00041', '2026-09-15', 'Plant', 'Check pump', 'P-41', 'Synthetic North', 'Completed'])
    buffer = BytesIO()
    book.save(buffer)
    draft = report(buffer.getvalue(), 'all-three.xlsx')
    draft = replace(draft, sections=default_sections(), blocks=(*draft.blocks,
        ResolvedBlock('work_orders', 'This month'), ResolvedBlock('service_calls', 'This month')))
    draft, training_value = prepared(draft)
    stale_training = training.encoded({**training_value, 'scope': 'older training report'})
    draft = replace(draft, source_row_reviews=(*draft.source_row_reviews, stale_training))
    content = source_content(draft, draft.sources[0])
    for kind, kept in [('workorders', 1), ('maintenance', 2)]:
        bundle = workorders.prepare_bundle(draft, content, kind=kind)
        assert len(bundle['groups']) == 3
        for index in range(3):
            bundle = workorders.group_decisions(bundle, index, omitted=index != kept,
                                                financial=True, layout=True)
        draft = workorders.store_bundle(draft, bundle)
        draft = publish_work(draft, bundle)
    stale_work = workorders.encoded({**workorders.saved_bundle(draft, draft.sources[0].id, 'workorders'),
                                    'scope': 'older work-order report'})
    draft = replace(draft, source_row_reviews=(*draft.source_row_reviews, stale_work))
    siblings = tuple(raw for raw in draft.source_row_reviews if decoded(raw)['kind'] != 'training')
    assert {decoded(raw)['kind'] for raw in siblings} == {'workorders', 'maintenance'}
    assert {decoded(raw)['source_id'] for raw in draft.source_row_reviews} == {draft.sources[0].id}
    accepted = publish(draft, training_value)
    assert tuple(raw for raw in accepted.source_row_reviews if decoded(raw)['kind'] != 'training') == siblings
    for kind in ('workorders', 'maintenance'):
        found = workorders.saved_bundle(accepted, accepted.sources[0].id, kind)
        assert found['kind'] == kind
        revised = workorders.store_bundle(accepted, found, expected_revision=workorders.revision(found))
        assert revised.source_row_reviews == accepted.source_row_reviews
    # Removing Training output manually never retires or republishes sibling kinds.
    removed = replace(accepted, blocks=tuple(replace(block, extra_tables=())
        if block.key == 'training_summary' else block for block in accepted.blocks))
    saved_training = training.saved_review(removed, removed.sources[0], 1)
    assert publish(removed, saved_training) == removed
    assert next(b for b in removed.blocks if b.key == 'work_orders').extra_tables
    assert next(b for b in removed.blocks if b.key == 'service_calls').extra_tables
    assert stale_training in removed.source_row_reviews and stale_work in removed.source_row_reviews
    for kind, destination in [('workorders', 'work_orders'), ('maintenance', 'service_calls')]:
        without = replace(accepted, blocks=tuple(replace(block, extra_tables=())
            if block.key == destination else block for block in accepted.blocks))
        receipt = workorders.saved_bundle(without, without.sources[0].id, kind)
        with pytest.raises(ValueError, match='edited or removed'):
            publish_work(without, receipt)
        assert without.source_row_reviews == accepted.source_row_reviews
        assert next(b for b in without.blocks if b.key == 'training_summary').extra_tables
        sibling_key = 'service_calls' if kind == 'workorders' else 'work_orders'
        assert next(b for b in without.blocks if b.key == sibling_key).extra_tables
    snapshot = library.save_snapshot(removed, expected_revision=0, entered_editor='Synthetic Editor')
    assert library.load_snapshot(removed.profile.contract, removed.profile.key, removed.period) == snapshot
