from dataclasses import replace
from pathlib import Path
from docx import Document
from streamlit.testing.v1 import AppTest
from app.monthly_report_import import inspect_docx
from app.monthly_report_model import Facility, ReportPeriod, synthetic_profiles
from app.monthly_report_start_import import analyze_starting_report


def test_unheaded_second_page_reference_can_falsely_detect_other_site(tmp_path):
    path = tmp_path / "synthetic.docx"
    doc = Document()
    doc.add_heading("Synthetic North Facility", 0)
    doc.add_paragraph("Monthly Report September 2026")
    doc.add_page_break()
    doc.add_paragraph("Reference only: Synthetic South Facility is outside this report.")
    doc.add_heading("2 Monthly Activity Summary", 1)
    doc.add_paragraph("September 2026 North work.")
    doc.save(path)
    north = Facility("north", "Synthetic North Facility")
    south = Facility("south", "Synthetic South Facility")
    profile = replace(synthetic_profiles()[0], title=north.title, facilities=(north,))
    inspection = inspect_docx(path)
    before = analyze_starting_report(inspection, profile, ReportPeriod(2026, 9), (north, south))
    assert before.site_relation == "other"
    assert {s.key for s in before.source_sites} == {"north", "south"}
    corrected = analyze_starting_report(inspection, profile, ReportPeriod(2026, 9), (north, south), confirmed_sites=(north,))
    assert corrected.site_relation == "same"
    assert profile.facilities == (north,)


def correction_app():
    return AppTest.from_string('''
from types import SimpleNamespace
from dataclasses import replace
import streamlit as st
from app.monthly_report_model import synthetic_profiles, Facility, ReportPeriod
from app.monthly_report_source_sites_ui import correct_detected_sites
from app.monthly_report_ui import _field
north = Facility("north", "Synthetic North")
south = Facility("south", "Synthetic South")
sha = st.text_input("Source hash", "aaa")
target_key = st.selectbox("Target", ["north", "south"])
view = st.text_input("View receipt", "view-one")
target = replace(synthetic_profiles()[0], facilities=(north if target_key == "north" else south,))
result = correct_detected_sites(SimpleNamespace(sha256=sha), target, ReportPeriod(2026,9),
                               (north,south), (north,south), "synthetic", _field,
                               source_view=SimpleNamespace(receipt_sha256=view))
st.session_state["result"] = tuple(s.key for s in result) if result else None
st.session_state["target_sites"] = tuple(s.key for s in target.facilities)
''').run()


def test_correction_requires_confirm_and_changes_do_not_reuse_old_approval():
    app = correction_app()
    assert app.session_state["result"] is None
    app.multiselect[0].set_value(["north"]).run()
    assert app.session_state["result"] is None
    next(b for b in app.button if b.label == "Confirm source sites").click().run()
    assert app.session_state["result"] == ("north",)
    assert app.session_state["target_sites"] == ("north",)
    app.multiselect[0].set_value(["south"]).run()
    assert app.session_state["result"] is None
    app.multiselect[0].set_value(["north"]).run()
    assert app.session_state["result"] is None


def test_changed_source_view_or_target_invalidates_and_revert_does_not_restore():
    for label, new, old in (("Source hash", "bbb", "aaa"), ("View receipt", "view-two", "view-one"), ("Target", "south", "north")):
        app = correction_app()
        app.multiselect[0].set_value(["north"]).run()
        next(b for b in app.button if b.label == "Confirm source sites").click().run()
        assert app.session_state["result"] == ("north",)
        controls = app.selectbox if label == "Target" else app.text_input
        next(w for w in controls if w.label == label).set_value(new).run()
        assert app.session_state["result"] is None
        controls = app.selectbox if label == "Target" else app.text_input
        next(w for w in controls if w.label == label).set_value(old).run()
        assert app.session_state["result"] is None


def test_cancel_and_empty_membership_do_not_grant_import_scope():
    app = correction_app()
    next(b for b in app.button if b.label == "Confirm source sites").click().run()
    next(b for b in app.button if b.label == "Use detected source sites").click().run()
    assert app.session_state["result"] is None
    app.multiselect[0].set_value([]).run()
    assert next(b for b in app.button if b.label == "Confirm source sites").disabled
    assert app.session_state["result"] is None


def test_real_upload_corrects_false_other_without_changing_target_and_reseeds_plans(monkeypatch, tmp_path):
    from test_monthly_report_start_import import start_app, save_start
    north, south = synthetic_profiles()[1].facilities
    path = tmp_path / "synthetic-misdetected.docx"
    doc = Document()
    doc.add_heading(north.title, 0)
    doc.add_paragraph("Monthly Report September 2026")
    doc.add_page_break()
    doc.add_paragraph("Reference only: " + south.title + " is outside this report.")
    doc.add_heading("2 Monthly Activity Summary", 1)
    doc.add_paragraph("September 2026 North inspection completed.")
    doc.save(path)
    app = start_app(monkeypatch, tmp_path, path, known_sites=(north, south))
    assert any("different sites" in e.value for e in app.info)
    selection = next(m for m in app.multiselect if m.label == "Actual sites covered by the uploaded report")
    selection.set_value([north.key]).run()
    assert any("different sites" in e.value for e in app.info)
    next(b for b in app.button if b.label == "Confirm source sites").click().run()
    assert not any("different sites" in e.value for e in app.info)
    assert not app.exception
    draft = save_start(app)
    assert draft.profile.facilities == (north,)
    assert "North inspection completed" in next(b.text for b in draft.blocks if b.key == "activity_summary")


def test_custom_slug_collisions_cannot_impersonate_known_or_merge_distinct_sites():
    for custom in ("North!", "A B; A-B"):
        app = correction_app()
        app.multiselect[0].set_value([]).run()
        next(w for w in app.text_input if w.label == "Source site not listed (optional)").set_value(custom).run()
        assert next(b for b in app.button if b.label == "Confirm source sites").disabled
        assert app.session_state["result"] is None


def test_scope_epoch_rotates_on_change_and_does_not_resurrect_old_widgets():
    app = AppTest.from_string("""
from types import SimpleNamespace
from dataclasses import replace
import streamlit as st
from app.monthly_report_source_sites_ui import source_review_epoch
from app.monthly_report_model import synthetic_profiles, Facility, ReportPeriod
north = Facility("north", "Synthetic North")
south = Facility("south", "Synthetic South")
choice = st.selectbox("Source", ["north", "south"])
site = north if choice == "north" else south
profile = replace(synthetic_profiles()[0], facilities=(north,))
epoch = source_review_epoch(SimpleNamespace(sha256="aaa"), profile, ReportPeriod(2026,9), (site,), "same" if choice == "north" else "other")
st.session_state["epoch"] = epoch
st.text_input("Section edit", key="section_" + epoch)
""").run()
    first = app.session_state["epoch"]
    app.text_input[0].set_value("Private pending edit").run()
    assert app.session_state["epoch"] == first
    app.selectbox[0].set_value("south").run()
    assert app.session_state["epoch"] != first
    app.selectbox[0].set_value("north").run()
    assert app.session_state["epoch"] != first
    assert app.text_input[0].value == ""


def test_invalid_custom_edit_permanently_invalidates_confirmation():
    app = correction_app()
    app.multiselect[0].set_value(["north"]).run()
    next(b for b in app.button if b.label == "Confirm source sites").click().run()
    assert app.session_state["result"] == ("north",)
    custom = lambda: next(w for w in app.text_input if w.label == "Source site not listed (optional)")
    custom().set_value("!!!").run()
    assert app.session_state["result"] is None
    custom().set_value("").run()
    assert app.session_state["result"] is None


def test_unknown_source_transition_clears_section_epoch_before_return():
    app = AppTest.from_string("""
from types import SimpleNamespace
from dataclasses import replace
import streamlit as st
from app.monthly_report_source_sites_ui import correct_detected_sites, source_review_epoch
from app.monthly_report_model import synthetic_profiles, Facility, ReportPeriod
from app.monthly_report_ui import _field
north = Facility("north", "Synthetic North")
known = st.checkbox("Known", value=True)
profile = replace(synthetic_profiles()[0], facilities=(north,))
inspection = SimpleNamespace(sha256="known" if known else "unknown")
correct_detected_sites(inspection, profile, ReportPeriod(2026,9), (north,) if known else (), (north,), "test", _field)
if known:
    epoch = source_review_epoch(inspection, profile, ReportPeriod(2026,9), (north,), "same")
    st.session_state["last_epoch"] = epoch
""").run()
    first = app.session_state["last_epoch"]
    app.checkbox[0].uncheck().run()
    assert "monthly_source_membership_render_scope" not in app.session_state
    app.checkbox[0].check().run()
    assert app.session_state["last_epoch"] != first
