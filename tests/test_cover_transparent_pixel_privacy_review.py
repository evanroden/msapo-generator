"""Synthetic erased-artwork canaries must not remain under exported alpha zero.

Metadata stripping remains pixel-preserving; this distinct privacy boundary
must preserve every visible (including partly transparent) pixel exactly.
"""

from dataclasses import replace
import hashlib
from io import BytesIO
import struct
import zlib
from zipfile import ZipFile

from docx import Document
from docx.shared import Inches
import fitz
from PIL import Image, ImageCms, ImageDraw, PngImagePlugin
import pytest

from app import monthly_report_docx as report_docx
from app.monthly_report_asset_review import approve_all_assets
from app.monthly_report_import import ImportMapping, inspect_docx, map_items
from app.monthly_report_media_privacy import (
    clear_transparent_png_rgb, strip_image_metadata, supports_transparent_png_sanitization,
)
from app.monthly_report_model import (
    ReportDraft, ReportPeriod, ResolvedBlock, default_sections, synthetic_draft, synthetic_profiles,
)
from app.monthly_report_native_layout import build_native_docx


def _hidden_artwork(*, variant=False):
    image = Image.new("RGBA", (96, 40))
    for y in range(image.height):
        for x in range(image.width):
            image.putpixel((x, y), ((x * 11 + 19) % 256, (y * 17 + 29) % 256, 173, 0))
    draw = ImageDraw.Draw(image)
    draw.rectangle((8, 8, 28, 32), fill=(45 if variant else 90, 155, 75, 255))
    draw.rectangle((36, 8, 55, 32), fill=(255, 255, 255, 255))
    draw.rectangle((64, 8, 85, 32), fill=(130, 80, 210, 73))
    metadata = PngImagePlugin.PngInfo()
    metadata.add_text("Author", "SYNTHETIC-HIDDEN-ARTWORK-AUTHOR")
    stream = BytesIO()
    image.save(stream, "PNG", pnginfo=metadata, dpi=(144, 144))
    return stream.getvalue()


def _assert_safe_pixels(original, exported):
    with Image.open(BytesIO(original)) as old, Image.open(BytesIO(exported)) as new:
        assert new.size == old.size
        before, after = old.convert("RGBA"), new.convert("RGBA")
        old_bytes, new_bytes = before.tobytes(), after.tobytes()
        pixels = zip(*(old_bytes[channel::4] for channel in range(4)))
        cleaned_pixels = zip(*(new_bytes[channel::4] for channel in range(4)))
        for pixel, cleaned in zip(pixels, cleaned_pixels):
            assert cleaned[3] == pixel[3], "The original alpha channel must not be flattened"
            if pixel[3]:
                assert cleaned == pixel, "Opaque and partially transparent artwork must be exact"
            else:
                assert cleaned == (0, 0, 0, 0), "Hidden RGB must not reveal erased artwork when alpha is removed"


def test_normalization_keeps_stable_asset_identity_and_export_derivative_clears_hidden_rgb():
    raw = _hidden_artwork()
    normalized = report_docx.normalize_report_image(raw, ".png", preserve_transparency=True)
    assert normalized.extension == "png" and (normalized.width, normalized.height) == (96, 40)
    with Image.open(BytesIO(raw)) as image:
        legacy = BytesIO()
        image.convert("RGBA").save(legacy, format="PNG", optimize=True)
    assert normalized.data == legacy.getvalue(), "Internal normalized hashes must retain their saved-asset meaning"
    _assert_safe_pixels(raw, clear_transparent_png_rgb(normalized.data))


def test_metadata_only_sanitizer_retains_its_pixel_preserving_contract():
    raw = _hidden_artwork()
    clean = strip_image_metadata(raw)
    assert b"SYNTHETIC-HIDDEN-ARTWORK-AUTHOR" not in clean
    with Image.open(BytesIO(raw)) as before, Image.open(BytesIO(clean)) as after:
        assert after.convert("RGBA").tobytes() == before.convert("RGBA").tobytes()


def _palette_artwork():
    image = Image.new("P", (96, 40))
    image.putpalette([61, 97, 203, 25, 180, 90, 255, 255, 255, 140, 80, 210, 33, 67, 113] + [0] * (256 * 3 - 15))
    image.info["transparency"] = bytes((0, 255, 255, 73, 0))
    draw = ImageDraw.Draw(image)
    draw.rectangle((8, 8, 28, 32), fill=1)
    draw.rectangle((36, 8, 55, 32), fill=2)
    draw.rectangle((64, 8, 85, 32), fill=3)
    draw.rectangle((3, 1, 90, 4), fill=4)
    stream = BytesIO()
    image.save(stream, "PNG")
    return stream.getvalue()


@pytest.mark.parametrize("slot", ["brand_logo", "client_logo"])
def test_generic_report_export_sanitizes_company_and_client_logo_only_in_derivative(monkeypatch, slot):
    raw = _hidden_artwork()
    reference = hashlib.sha256(raw).hexdigest() + ".png"
    assets = {reference: raw}
    draft = synthetic_draft(synthetic_profiles()[0], ReportPeriod(2026, 9))
    logo = approve_all_assets(ResolvedBlock(slot, "Library", asset_hashes=(reference,)))
    draft = replace(draft, blocks=tuple(block for block in draft.blocks if block.key != slot) + (logo,))
    monkeypatch.setattr(report_docx, "_convert_report_pdf", lambda _draft, _raw: (None, "Synthetic converter boundary"))
    package = report_docx.generate_report(draft, asset_loader=assets.__getitem__)
    with ZipFile(BytesIO(package.docx)) as archive:
        pictures = [archive.read(name) for name in archive.namelist() if name.startswith("word/media/")]
    assert pictures
    for picture in pictures:
        _assert_safe_pixels(raw, picture)
    assert assets[reference] == raw and hashlib.sha256(raw).hexdigest() + ".png" == reference


def test_normalized_palette_export_derivative_is_safe_but_native_palette_is_explicitly_unsupported():
    raw = _palette_artwork()
    normalized = report_docx.normalize_report_image(raw, ".png", preserve_transparency=True)
    _assert_safe_pixels(raw, clear_transparent_png_rgb(normalized.data))
    assert not supports_transparent_png_sanitization(raw)
    assert clear_transparent_png_rgb(raw) == raw


def _chunks(raw):
    chunks, offset = [], 8
    while offset < len(raw):
        length = int.from_bytes(raw[offset:offset + 4], "big")
        chunks.append((raw[offset + 4:offset + 8], raw[offset:offset + 12 + length]))
        offset += 12 + length
    return chunks


def _chunk(kind, payload):
    return (struct.pack(">I", len(payload)) + kind + payload
            + struct.pack(">I", zlib.crc32(kind + payload) & 0xFFFFFFFF))


def test_native_rgba_primitive_preserves_every_non_idat_chunk_including_profile_orientation_density():
    with Image.open(BytesIO(_hidden_artwork())) as image:
        exif = Image.Exif()
        exif[274], exif[315] = 6, "SYNTHETIC-ORIENTATION-AUTHOR"
        profile = ImageCms.ImageCmsProfile(ImageCms.createProfile("sRGB")).tobytes()
        output = BytesIO()
        image.save(output, "PNG", exif=exif.tobytes(), icc_profile=profile, dpi=(144, 144))
    raw = output.getvalue()
    assert supports_transparent_png_sanitization(raw)
    clean = clear_transparent_png_rgb(raw)
    assert clean != raw
    _assert_safe_pixels(raw, clean)
    assert [(kind, data) for kind, data in _chunks(clean) if kind != b"IDAT"] == [
        (kind, data) for kind, data in _chunks(raw) if kind != b"IDAT"]
    with Image.open(BytesIO(clean)) as image:
        assert image.info["icc_profile"] == profile
        assert image.getexif()[274] == 6 and image.getexif()[315] == "SYNTHETIC-ORIENTATION-AUTHOR"
        assert tuple(round(value) for value in image.info["dpi"]) == (144, 144)
    assert clear_transparent_png_rgb(clean) == clean


def _sixteen_bit_png():
    header = struct.pack(">IIBBBBB", 2, 1, 16, 6, 0, 0, 0)
    samples = struct.pack(">8H", 0x1234, 0xABCD, 0x5678, 0, 0x1234, 0xABCD, 0x5678, 0x1234)
    return b"\x89PNG\r\n\x1a\n" + _chunk(b"IHDR", header) + _chunk(b"IDAT", zlib.compress(b"\0" + samples)) + _chunk(b"IEND", b"")


def _interlaced_png():
    with Image.open(BytesIO(_hidden_artwork())) as image:
        width, height = image.size
        scanlines = bytearray()
        for start_x, start_y, stride_x, stride_y in (
                (0, 0, 8, 8), (4, 0, 8, 8), (0, 4, 4, 8), (2, 0, 4, 4),
                (0, 2, 2, 4), (1, 0, 2, 2), (0, 1, 1, 2)):
            for y in range(start_y, height, stride_y):
                scanlines.append(0)
                for x in range(start_x, width, stride_x):
                    scanlines.extend(image.getpixel((x, y)))
    header = struct.pack(">IIBBBBB", width, height, 8, 6, 0, 0, 1)
    return b"\x89PNG\r\n\x1a\n" + _chunk(b"IHDR", header) + _chunk(b"IDAT", zlib.compress(scanlines)) + _chunk(b"IEND", b"")


def _apng():
    output = BytesIO()
    with Image.open(BytesIO(_hidden_artwork())) as first, Image.open(BytesIO(_hidden_artwork(variant=True))) as second:
        first.save(output, "PNG", save_all=True, append_images=[second], duration=100, loop=0)
    return output.getvalue()


@pytest.mark.parametrize("fixture", [_sixteen_bit_png, _interlaced_png, _apng])
def test_unsupported_native_png_encodings_are_exactly_unchanged_and_not_claimed_sanitized(fixture):
    raw = fixture()
    with Image.open(BytesIO(raw)) as image:
        image.load()  # These are valid images, not only unsupported header stubs.
    assert not supports_transparent_png_sanitization(raw)
    assert clear_transparent_png_rgb(raw) == raw


def test_unsupported_grayscale_other_formats_and_malformed_images_stay_unchanged():
    examples = [b"", b"\x89PNG\r\n\x1a\nbroken", _hidden_artwork()[:-1]]
    for mode, format in (("LA", "PNG"), ("RGB", "JPEG"), ("RGBA", "TIFF"), ("P", "GIF")):
        output = BytesIO()
        Image.new(mode, (8, 8)).save(output, format)
        examples.append(output.getvalue())
    corrupt = bytearray(_hidden_artwork())
    corrupt[-1] ^= 1  # Invalid CRC must never be transformed.
    examples.append(bytes(corrupt))
    for raw in examples:
        assert not supports_transparent_png_sanitization(raw)
        assert clear_transparent_png_rgb(raw) == raw


def test_transparency_bounds_reject_before_decoding(monkeypatch):
    raw = _hidden_artwork()
    chunks = _chunks(raw)
    oversized_header = _chunk(b"IHDR", struct.pack(">IIBBBBB", 4_000_001, 1, 8, 6, 0, 0, 0))
    too_many_pixels = raw[:8] + oversized_header + b"".join(data for kind, data in chunks if kind != b"IHDR")
    too_many_bytes = raw + b"\0" * (30 * 1024 * 1024)
    monkeypatch.setattr(Image, "open", lambda *a, **kw: pytest.fail("Unsupported input must be bounded before decode"))
    for unsupported in (too_many_pixels, too_many_bytes):
        assert not supports_transparent_png_sanitization(unsupported)
        assert clear_transparent_png_rgb(unsupported) == unsupported


def test_transparency_chunk_count_is_bounded_before_decoding(monkeypatch):
    raw = _hidden_artwork()
    chunks = _chunks(raw)
    many_chunks = raw[:8] + chunks[0][1] + _chunk(b"vpAg", b"") * 1024 + b"".join(
        data for kind, data in chunks if kind != b"IHDR")
    assert len(many_chunks) < 30 * 1024 * 1024
    monkeypatch.setattr(Image, "open", lambda *a, **kw: pytest.fail("Excessive chunks must be rejected before decode"))
    assert not supports_transparent_png_sanitization(many_chunks)
    assert clear_transparent_png_rgb(many_chunks) == many_chunks


def test_malformed_decoder_syntax_error_leaves_input_unchanged(monkeypatch):
    raw = _hidden_artwork()
    def malformed(*args, **kwargs):
        raise SyntaxError("Synthetic malformed PNG decoder failure")
    monkeypatch.setattr(Image, "open", malformed)
    assert clear_transparent_png_rgb(raw) == raw


def test_pre_privacy_saved_logo_digest_still_preserves_distinct_native_variant(monkeypatch, tmp_path):
    import test_cover_native_variant_review as native_fixture
    from app import monthly_report_library as library
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path / "saved"))
    monkeypatch.setattr(native_fixture, "mark", lambda color=None: _hidden_artwork(variant=color is not None))
    path, inspection, _, draft, _, variant = native_fixture._source(tmp_path)
    original = path.read_bytes()
    # Reproduce the exact previous normalizer for this already-upright,
    # below-budget PNG: retain alpha/RGB, strip metadata, optimize losslessly.
    with Image.open(BytesIO(_hidden_artwork())) as image:
        legacy = BytesIO()
        image.convert("RGBA").save(legacy, format="PNG", optimize=True)
    legacy_bytes = legacy.getvalue()
    legacy_reference = hashlib.sha256(legacy_bytes).hexdigest() + ".png"
    brand = approve_all_assets(replace(draft.blocks[0], asset_hashes=(legacy_reference,)))
    assert brand.references[0].startswith("docx:" + inspection.sha256 + ":")
    draft = replace(draft, blocks=(brand, *draft.blocks[1:]))
    library.save_snapshot(draft, expected_revision=0, assets=((legacy_reference, legacy_bytes),))
    saved = library.load_snapshot(draft.profile.contract, draft.profile.key, draft.period)
    assert saved.draft == draft
    loader = lambda reference: library.read_asset(draft.profile.contract, draft.profile.key, reference)
    output = build_native_docx(saved.draft, path, asset_loader=loader, master=True)
    assert native_fixture._part_images(output, "header") == (variant,), \
        "Cleaning export pixels must not reinterpret a saved unchanged company asset as a replacement"
    monkeypatch.setattr(report_docx, "assemble_docx", lambda *a, **kw: output)
    monkeypatch.setattr(report_docx, "_convert_report_pdf", lambda _draft, _raw: (None, "Synthetic converter boundary"))
    generated = report_docx.generate_report(saved.draft, asset_loader=loader)
    exported_header = native_fixture._part_images(generated.docx, "header")
    assert len(exported_header) == 1
    _assert_safe_pixels(variant, exported_header[0])
    assert path.read_bytes() == original
    assert loader(legacy_reference) == legacy_bytes


@pytest.mark.parametrize("route", ["full", "section"])
@pytest.mark.parametrize("native", [False, True])
def test_full_and_section_preview_conversion_inputs_clear_hidden_rgb(monkeypatch, tmp_path, route, native):
    from app import monthly_report_designs as designs, monthly_report_preview as preview
    from app import monthly_report_render_jobs as jobs
    from app.monthly_report_render_profile import validate_profile
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path / "preview-data"))
    originals = [_hidden_artwork()]
    source = None
    if native:
        import test_cover_native_variant_review as native_fixture
        monkeypatch.setattr(native_fixture, "mark", lambda color=None: _hidden_artwork(variant=color is not None))
        source, _, _, draft, assets, variant = native_fixture._source(tmp_path)
        source_bytes = source.read_bytes()
        originals.append(variant)
        monkeypatch.setattr(designs, "is_master", lambda _profile: True)
    else:
        reference = hashlib.sha256(originals[0]).hexdigest() + ".png"
        assets = {reference: originals[0]}
        draft = synthetic_draft(synthetic_profiles()[0], ReportPeriod(2026, 9))
        brand = approve_all_assets(ResolvedBlock("brand_logo", "Library", asset_hashes=(reference,)))
        draft = replace(draft, blocks=tuple(block for block in draft.blocks if block.key != "brand_logo") + (brand,))
    monkeypatch.setattr(designs, "source_for", lambda _profile: source)
    monkeypatch.setattr(designs, "render_profile_for", lambda _profile: validate_profile())
    captured = []
    def convert(path, **kwargs):
        captured.append(path.read_bytes())
        target = tmp_path / (path.stem + ".pdf")
        with fitz.open() as pdf:
            pdf.new_page().insert_text((72, 72), "Synthetic preview renderer")
            pdf.save(target)
        return target
    monkeypatch.setattr(jobs, "convert_to_pdf", convert)
    result = (preview.preview_report(draft, assets.__getitem__) if route == "full"
              else preview.preview_section(draft, "cover", assets.__getitem__))
    assert result.pdf and len(captured) == 1
    expected = {}
    for raw in originals:
        with Image.open(BytesIO(raw)) as image:
            expected[image.getpixel((8, 8))] = raw
    with ZipFile(BytesIO(captured[0])) as archive:
        pictures = [archive.read(name) for name in archive.namelist() if name.startswith("word/media/")]
    assert pictures
    for picture in pictures:
        with Image.open(BytesIO(picture)) as image:
            original = expected[image.getpixel((8, 8))]
        _assert_safe_pixels(original, picture)
    if source is not None:
        assert source.read_bytes() == source_bytes


def test_legacy_package_pdf_retry_cleans_conversion_copy_without_rewriting_saved_word(monkeypatch, tmp_path):
    from app import monthly_report_designs as designs, monthly_report_render_jobs as jobs
    raw = _hidden_artwork()
    document = Document()
    document.add_picture(BytesIO(raw), width=Inches(2))
    output = BytesIO()
    document.save(output)
    original_word = output.getvalue()
    draft = synthetic_draft(synthetic_profiles()[0], ReportPeriod(2026, 9))
    package = report_docx.ReportPackage(original_word, None, "existing.docx", "existing.pdf", draft.fingerprint)
    monkeypatch.setattr(designs, "source_for", lambda _profile: None)
    monkeypatch.setattr(report_docx, "assemble_docx", lambda *a, **kw: pytest.fail("Retry must not assemble a new Word file"))
    captured = []
    def convert(path, **kwargs):
        captured.append(path.read_bytes())
        target = tmp_path / "retry.pdf"
        target.write_bytes(b"%PDF-synthetic-retry")
        return target
    monkeypatch.setattr(jobs, "convert_to_pdf", convert)
    result = report_docx.retry_report_pdf(draft, package)
    assert result.pdf == b"%PDF-synthetic-retry" and not result.pdf_error
    assert result.docx == package.docx == original_word, "A retry does not retroactively rewrite a prior Word download"
    with ZipFile(BytesIO(captured[0])) as archive:
        pictures = [archive.read(name) for name in archive.namelist() if name.startswith("word/media/")]
    assert pictures
    for picture in pictures:
        _assert_safe_pixels(raw, picture)


def test_native_preserved_logo_export_clears_hidden_rgb_without_mutating_source_or_geometry(monkeypatch, tmp_path):
    primary, variant = _hidden_artwork(), _hidden_artwork(variant=True)
    document = Document()
    document.add_paragraph(synthetic_profiles()[0].facilities[0].title)
    document.add_paragraph("November 2026")
    shape = document.add_picture(BytesIO(primary), width=Inches(3))
    shape._inline.docPr.set("descr", "ENFRA company logo")
    shape = document.sections[0].header.paragraphs[0].add_run().add_picture(
        BytesIO(variant), width=Inches(2))
    shape._inline.docPr.set("descr", "ENFRA company logo")
    document.add_section()
    document.add_heading("Monthly Activity Summary", 1)
    document.add_paragraph("Synthetic inspection complete.")
    path = tmp_path / "synthetic-native-hidden-artwork.docx"
    document.save(path)
    original = path.read_bytes()
    inspection = inspect_docx(path)
    cover = next(item for item in inspection.items if item.kind == "image" and item.part == "word/document.xml")
    mapped = map_items(path, inspection, (ImportMapping(cover.id, "brand_logo"),))
    assets = dict(mapped.assets)
    asset_hashes = {key: hashlib.sha256(value).hexdigest() for key, value in assets.items()}
    sections = tuple(replace(section, included=section.key == "activity") for section in default_sections())
    draft = ReportDraft(synthetic_profiles()[0], ReportPeriod(2026, 11), "Synthetic Reviewer", sections, (
        approve_all_assets(mapped.blocks[0]),
        ResolvedBlock("activity_summary", "This month", text="Synthetic inspection complete."),
    ))
    native = build_native_docx(draft, path, asset_loader=assets.__getitem__, master=True)
    with ZipFile(BytesIO(native)) as archive:
        native_media = {name: archive.read(name) for name in archive.namelist() if name.startswith("word/media/")}
    assert primary in native_media.values() and variant in native_media.values(), \
        "Exercise the native byte-preserving path, not a replacement approximation"
    # Keep the actual native package; isolate the already-tested preflight and
    # external PDF renderer to inspect the common final export boundary.
    monkeypatch.setattr(report_docx, "assemble_docx", lambda *a, **kw: native)
    pdf_inputs = []
    def capture_pdf(_draft, raw):
        pdf_inputs.append(raw)
        return None, "Synthetic converter boundary"
    monkeypatch.setattr(report_docx, "_convert_report_pdf", capture_pdf)
    package = report_docx.generate_report(draft, asset_loader=assets.__getitem__)
    assert pdf_inputs == [package.docx]
    with ZipFile(BytesIO(native)) as before, ZipFile(BytesIO(package.docx)) as after:
        assert after.namelist() == before.namelist()
        for name in before.namelist():
            if name in native_media:
                _assert_safe_pixels(before.read(name), after.read(name))
                assert b"SYNTHETIC-HIDDEN-ARTWORK-AUTHOR" not in after.read(name)
            else:
                assert after.read(name) == before.read(name), "Export sanitation must not change document geometry"
    assert path.read_bytes() == original
    assert inspect_docx(path).sha256 == inspection.sha256
    assert {key: hashlib.sha256(value).hexdigest() for key, value in assets.items()} == asset_hashes
