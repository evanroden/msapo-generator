"""Current client artwork must reach its exact native header/footer occurrence."""

from copy import deepcopy
from dataclasses import replace
from io import BytesIO

from docx import Document
from docx.shared import Inches
from lxml import etree
import pytest

from app import monthly_report_library as library
from app.monthly_report_asset_review import approve_all_assets
from app.monthly_report_cover_ui import replace_cover_picture
from app.monthly_report_import import ImportMapping, inspect_docx, map_items
from app.monthly_report_model import ResolvedBlock, synthetic_profiles
from app.monthly_report_native_layout import build_native_docx
from test_cover_branding_contracts import mark
from test_cover_native_variant_review import _part_images
from test_cover_review_regressions import _draft


def _source(tmp_path, *, part_name="header", identified=True, alternate=False):
    document = Document()
    document.add_paragraph(synthetic_profiles()[0].facilities[0].title)
    document.add_paragraph("November 2026")
    company, old_client = mark((40, 90, 150, 255)), mark((150, 60, 90, 255))
    current_assets = {}
    current_client = replace_cover_picture(
        ResolvedBlock("client_logo", "Omit"), mark((90, 150, 60, 255)), ".png", current_assets)
    replacement = current_assets[current_client.asset_hashes[0]]
    part = getattr(document.sections[0], part_name)
    company_shape = part.paragraphs[0].add_run().add_picture(BytesIO(company), width=Inches(2))
    company_shape._inline.docPr.set("descr", "ENFRA company logo")
    shape = part.add_paragraph().add_run().add_picture(BytesIO(old_client), width=Inches(1))
    if identified:
        shape._inline.docPr.set("descr", "Client logo")
    if alternate:
        drawing = shape._inline.getparent()
        run = drawing.getparent()
        mc = "{http://schemas.openxmlformats.org/markup-compatibility/2006}"
        wrapper = etree.Element(mc + "AlternateContent")
        run.replace(drawing, wrapper)
        etree.SubElement(wrapper, mc + "Choice", Requires="w").append(drawing)
        etree.SubElement(wrapper, mc + "Fallback").append(deepcopy(drawing))
    # Independent occurrences share both the old and new bytes. Neither is
    # authorized merely because the selected client logo has matching pixels.
    part.add_paragraph().add_run().add_picture(BytesIO(old_client), width=Inches(1))
    part.add_paragraph().add_run().add_picture(BytesIO(replacement), width=Inches(1))
    document.add_section()
    document.add_heading("Monthly Activity Summary", 1)
    document.add_paragraph("Synthetic inspection completed.")
    source = tmp_path / "client-header-logo.docx"
    document.save(source)
    inspection = inspect_docx(source)
    items = tuple(item for item in inspection.items
                  if item.kind == "image" and item.part.startswith("word/" + part_name))
    assert len(items) == 4
    return source, inspection, items, company, old_client, replacement


@pytest.mark.parametrize("part_name", ["header", "footer"])
@pytest.mark.parametrize("master", [False, True])
@pytest.mark.parametrize("evidence", ["mapped", "classified"])
def test_client_replacement_reaches_native_occurrence_after_save_reload(
        monkeypatch, tmp_path, part_name, master, evidence):
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path / "data"))
    source, inspection, items, company, old_client, replacement = _source(
        tmp_path, part_name=part_name, identified=evidence == "classified")
    original = source.read_bytes()
    mapped = map_items(source, inspection, (
        ImportMapping(items[0].id, "brand_logo"), ImportMapping(items[1].id, "client_logo")))
    assets = dict(mapped.assets)
    blocks = {block.key: approve_all_assets(block) for block in mapped.blocks}
    client = replace_cover_picture(blocks["client_logo"], replacement, ".png", assets)
    assert client.asset_hashes != blocks["client_logo"].asset_hashes
    assert client.references == blocks["client_logo"].references
    if evidence == "classified":
        assert items[1].image_role == "client_logo"
        client = replace(client, references=())
    else:
        assert items[1].image_role != "client_logo"
    report = _draft(blocks["brand_logo"], approve_all_assets(client))
    library.save_profile(report.profile, expected_revision=0, actor="Synthetic Reviewer", confirmed=True)
    library.save_snapshot(report, expected_revision=0, assets=tuple(assets.items()), entered_editor="Synthetic Reviewer")
    loaded = library.load_snapshot(report.profile.contract, report.profile.key, report.period).draft
    loader = lambda ref: library.read_asset(report.profile.contract, report.profile.key, ref)
    actual = _part_images(build_native_docx(loaded, source, asset_loader=loader, master=master), part_name)
    company_output = company  # Unchanged native company art retains original bytes.
    client_output = loader(client.asset_hashes[0])
    assert actual[:2] == (company_output, client_output)
    assert client_output != company_output and client_output != old_client
    assert actual[2:] == (() if master else (old_client, replacement))
    assert source.read_bytes() == original


@pytest.mark.parametrize("master", [False, True])
def test_mapped_client_logo_does_not_require_a_company_logo_block(tmp_path, master):
    source, inspection, items, company, old_client, replacement = _source(tmp_path, identified=False)
    mapped = map_items(source, inspection, (ImportMapping(items[1].id, "client_logo"),))
    assets = dict(mapped.assets)
    client = approve_all_assets(replace_cover_picture(mapped.blocks[0], replacement, ".png", assets))
    actual = _part_images(build_native_docx(
        _draft(client), source, asset_loader=assets.__getitem__, master=master), "header")
    assert actual == ((assets[client.asset_hashes[0]],) if master else
                      (company, assets[client.asset_hashes[0]], old_client, replacement))


@pytest.mark.parametrize("master", [False, True])
def test_client_mapping_updates_equivalent_branches_not_independent_equal_bytes(tmp_path, master):
    source, inspection, items, company, old_client, replacement = _source(tmp_path, identified=False, alternate=True)
    mapped = map_items(source, inspection, (ImportMapping(items[1].id, "client_logo"),))
    assets = dict(mapped.assets)
    client = approve_all_assets(replace_cover_picture(mapped.blocks[0], replacement, ".png", assets))
    actual = _part_images(build_native_docx(
        _draft(client), source, asset_loader=assets.__getitem__, master=master), "header")
    current = assets[client.asset_hashes[0]]
    assert actual == ((current, current) if master else
                      (company, current, current, old_client, replacement))


@pytest.mark.parametrize("master", [False, True])
def test_a_client_label_cannot_authorize_company_replacement_without_a_mapping(tmp_path, master):
    source, inspection, items, company, old_client, replacement = _source(tmp_path)
    mapped = map_items(source, inspection, (ImportMapping(items[0].id, "brand_logo"),))
    assets = dict(mapped.assets)
    brand = approve_all_assets(mapped.blocks[0])
    actual = _part_images(build_native_docx(
        _draft(brand), source, asset_loader=assets.__getitem__, master=master), "header")
    assert actual == ((company,) if master else
                      (company, old_client, old_client, replacement))


@pytest.mark.parametrize("master", [False, True])
def test_explicit_other_mapping_overrides_client_classification(tmp_path, master):
    source, inspection, items, _, _, _ = _source(tmp_path)
    mapped = map_items(source, inspection, (ImportMapping(items[1].id, "brand_logo"),))
    assets = dict(mapped.assets)
    brand = approve_all_assets(replace_cover_picture(mapped.blocks[0], mark(), ".png", assets))
    client = approve_all_assets(replace_cover_picture(
        ResolvedBlock("client_logo", "Omit"), mark((15, 45, 75, 255)), ".png", assets))
    actual = _part_images(build_native_docx(
        _draft(brand, client), source, asset_loader=assets.__getitem__, master=master), "header")
    assert actual[0] == actual[1] == assets[brand.asset_hashes[0]]
    assert assets[client.asset_hashes[0]] not in actual


@pytest.mark.parametrize("master", [False, True])
def test_company_variant_exception_does_not_preserve_a_stale_client_variant(tmp_path, master):
    source, _, _, company, old_client, replacement = _source(tmp_path)
    document = Document(source)
    for index, (raw, label) in enumerate(((company, "ENFRA company logo"), (replacement, "Client logo")), 2):
        shape = document.add_picture(BytesIO(raw), width=Inches(2))
        shape._inline.docPr.set("descr", label)
        paragraph = shape._inline.getparent().getparent().getparent()
        paragraph.getparent().remove(paragraph)
        document.element.body.insert(index, paragraph)
    document.save(source)
    inspection = inspect_docx(source)
    cover_items = tuple(item for item in inspection.items
                        if item.kind == "image" and item.part == "word/document.xml")
    mapped = map_items(source, inspection, (
        ImportMapping(cover_items[0].id, "brand_logo"), ImportMapping(cover_items[1].id, "client_logo")))
    assets = dict(mapped.assets)
    blocks = {block.key: approve_all_assets(block) for block in mapped.blocks}
    # Both cover logos are unchanged. Only the company may retain a distinct
    # known native variant; a client header must use the current selected logo.
    actual = _part_images(build_native_docx(
        _draft(*blocks.values()), source, asset_loader=assets.__getitem__, master=master), "header")
    assert actual[:2] == (company, assets[blocks["client_logo"].asset_hashes[0]])
    assert actual[1] != old_client
    assert actual[2:] == (() if master else (old_client, replacement))


@pytest.mark.parametrize("master", [False, True])
def test_conflicting_explicit_roles_do_not_guess_a_header_logo_owner(tmp_path, master):
    source, inspection, items, company, old_client, replacement = _source(tmp_path)
    reference = f"docx:{inspection.sha256}:{items[1].id}"
    assets = {}
    brand = approve_all_assets(replace_cover_picture(
        ResolvedBlock("brand_logo", "Omit"), mark(), ".png", assets))
    client = approve_all_assets(replace_cover_picture(
        ResolvedBlock("client_logo", "Omit"), replacement, ".png", assets))
    brand, client = replace(brand, references=(reference,)), replace(client, references=(reference,))
    actual = _part_images(build_native_docx(
        _draft(brand, client), source, asset_loader=assets.__getitem__, master=master), "header")
    assert actual == ((assets[brand.asset_hashes[0]],) if master else
                      (assets[brand.asset_hashes[0]], old_client, old_client, replacement))
