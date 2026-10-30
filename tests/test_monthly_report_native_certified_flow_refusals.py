"""Unproved synthetic context keeps its authored body order and source boundaries."""
from copy import deepcopy
from dataclasses import replace
from io import BytesIO

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE
from lxml import etree
import pytest

from app import monthly_report_native_table_flow as certified
from app.monthly_report_import import _text
from app.monthly_report_model import BlockSpec, ResolvedBlock
from test_monthly_report_native_certified_flow import (
    CAPACITY_INTRO, CAPACITY_NOTES, CHILLED_LABEL, NATIVE_NOTE, NEW_MARKER, PENDING,
    PROPOSAL_OVERVIEW, REJECTED, STEAM_LABEL, flow_case, rendered,
)


CAVEAT = "Additional synthetic qualification must stay in its authored position."


def _block(draft, key):
    return next(block for block in draft.blocks if block.key == key)


def _replace_block(draft, block):
    return replace(draft, blocks=tuple(block if old.key == block.key else old for old in draft.blocks))


def _recorded_render(flow, monkeypatch):
    """Observe the real writer's exact-table gate, before it mutates source XML."""
    calls = []
    original = certified.supplied_native_table_flows

    def record(*args, **kwargs):
        result = original(*args, **kwargs)
        # Source/style mutation tests use a detached copy of the actual proof
        # inputs. They do not loosen the importer's or writer's separate guards.
        saved = list(args)
        detached = OxmlElement("w:body")
        for node in args[1]:
            detached.append(deepcopy(node))
        saved[1] = tuple(detached)
        saved_kwargs = dict(kwargs, styles=deepcopy(kwargs["styles"]))
        calls.append((tuple(saved), saved_kwargs, result))
        return result

    with monkeypatch.context() as patch:
        patch.setattr(certified, "supplied_native_table_flows", record)
        raw = rendered(flow)
    assert len(calls) == 1
    return raw, calls[0]


def _current_lines(text):
    return [line.strip() for line in text.splitlines() if line.strip() and line.strip() != NATIVE_NOTE]


def _assert_authored_run(raw, current):
    document = Document(BytesIO(raw))
    actual = [line.strip() for paragraph in document.paragraphs
              for line in paragraph.text.splitlines() if line.strip()]
    expected = _current_lines(current)
    starts = [index for index, line in enumerate(actual) if line == expected[0]]
    assert len(starts) == 1
    assert actual[starts[0]:starts[0] + len(expected)] == expected


@pytest.mark.parametrize("key", ["utility_analysis", "proposals"])
@pytest.mark.parametrize("change", ["edited", "reordered", "repeated", "interleaved", "trailing"])
def test_unproved_current_body_remains_one_authored_run(tmp_path, monkeypatch, key, change):
    path, draft, loader, view = flow_case(tmp_path, monkeypatch)
    block = _block(draft, key)
    first, second = (CHILLED_LABEL, STEAM_LABEL) if key == "utility_analysis" else (REJECTED, PENDING)
    text = block.text
    if change == "edited":
        anchor = CAPACITY_INTRO if key == "utility_analysis" else PROPOSAL_OVERVIEW
        text = text.replace(anchor, anchor + " Corrected by the current synthetic author.")
    elif change == "reordered":
        text = text.replace(first, "__synthetic_label_swap__").replace(second, first).replace(
            "__synthetic_label_swap__", second)
    elif change == "repeated":
        text = text.replace(first, first + "\n\n" + first)
    elif change == "interleaved":
        text = text.replace(second, CAVEAT + "\n\n" + second)
    elif key == "proposals":
        # A real native footnote does not excuse an arbitrary trailing body note.
        text = text.replace(NATIVE_NOTE, CAVEAT + "\n\n" + NATIVE_NOTE)
    else:
        text += "\n\n" + CAVEAT
    draft = _replace_block(draft, replace(block, text=text))
    raw, (_, _, result) = _recorded_render((path, draft, loader, view), monkeypatch)
    assert key not in result
    _assert_authored_run(raw, text)


@pytest.mark.parametrize("key, owner", [("utility_analysis", "thermal_capacity"), ("proposals", "proposals")])
def test_edited_current_table_cannot_authorize_source_context(tmp_path, monkeypatch, key, owner):
    path, draft, loader, view = flow_case(tmp_path, monkeypatch)
    block = _block(draft, owner)
    first = block.extra_tables[0]
    edited = replace(first, rows=(("Edited synthetic cell", *first.rows[0][1:]), *first.rows[1:]))
    draft = _replace_block(draft, replace(block, extra_tables=(edited, *block.extra_tables[1:])))
    raw, (_, _, result) = _recorded_render((path, draft, loader, view), monkeypatch)
    assert key not in result
    assert "Edited synthetic cell" in _text(Document(BytesIO(raw)).element.body)
    _assert_authored_run(raw, _block(draft, key).text)


def _with_capacity_spacing(document):
    paragraph = next(p for p in document.paragraphs if p.text == CHILLED_LABEL)
    paragraph.insert_paragraph_before()


@pytest.fixture
def proof_inputs(tmp_path, monkeypatch):
    flow = flow_case(tmp_path, monkeypatch, mutate_source=_with_capacity_spacing)
    _, (args, kwargs, result) = _recorded_render(flow, monkeypatch)
    assert {"utility_analysis", "proposals"} <= result.keys()
    return args, kwargs


def _source_paragraph(args, text):
    return next(node for node in args[1] if node.tag == qn("w:p") and _text(node) == text)


def _inherited_property(paragraph, styles, name):
    parent = OxmlElement("w:style")
    parent.set(qn("w:type"), "paragraph")
    parent.set(qn("w:styleId"), "SyntheticUnsafeSpacerBase")
    properties = OxmlElement("w:rPr" if name == "vanish" else "w:pPr")
    properties.append(OxmlElement("w:" + name))
    parent.append(properties)
    styles.append(parent)
    child = OxmlElement("w:style")
    child.set(qn("w:type"), "paragraph")
    child.set(qn("w:styleId"), "SyntheticUnsafeSpacer")
    base = OxmlElement("w:basedOn")
    base.set(qn("w:val"), "SyntheticUnsafeSpacerBase")
    child.append(base)
    styles.append(child)
    selected = OxmlElement("w:pStyle")
    selected.set(qn("w:val"), "SyntheticUnsafeSpacer")
    paragraph.get_or_add_pPr().append(selected)


@pytest.mark.parametrize("key", ["utility_analysis", "proposals"])
@pytest.mark.parametrize("hazard", ["opaque", "page_before", "inherited_page_before", "hidden", "numbered", "physical_section"])
def test_empty_looking_source_nodes_do_not_prove_spacing(proof_inputs, key, hazard):
    args, kwargs = proof_inputs
    if key == "utility_analysis":
        spacer = _source_paragraph(args, CHILLED_LABEL).getprevious()
    else:
        spacer = _source_paragraph(args, REJECTED).getnext()
    assert spacer.tag == qn("w:p") and not _text(spacer)
    if hazard == "opaque":
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), "PAGE")
        spacer.append(field)
    elif hazard == "page_before":
        spacer.get_or_add_pPr().append(OxmlElement("w:pageBreakBefore"))
    elif hazard == "physical_section":
        spacer.get_or_add_pPr().append(OxmlElement("w:sectPr"))
    else:
        name = {"inherited_page_before": "pageBreakBefore", "hidden": "vanish", "numbered": "numPr"}[hazard]
        _inherited_property(spacer, kwargs["styles"], name)
    assert key not in certified.supplied_native_table_flows(*args, **kwargs)


@pytest.mark.parametrize("hazard", ["page", "column", "unknown_attribute", "wrong_parent"])
def test_context_line_break_requires_plain_wrapping_structure(proof_inputs, hazard):
    args, kwargs = proof_inputs
    paragraph = _source_paragraph(args, CAPACITY_NOTES)
    line_break = next(paragraph.iter(qn("w:br")))
    if hazard in ("page", "column"):
        line_break.set(qn("w:type"), hazard)
    elif hazard == "unknown_attribute":
        line_break.set(qn("w:clear"), "all")
    else:
        line_break.getparent().remove(line_break)
        paragraph.get_or_add_pPr().append(line_break)
    assert "utility_analysis" not in certified.supplied_native_table_flows(*args, **kwargs)


@pytest.mark.parametrize("key", ["utility_analysis", "thermal_capacity", "proposals"])
def test_primary_rows_cannot_be_displaced_by_extra_table_context(proof_inputs, key):
    args, kwargs = proof_inputs
    draft = _replace_block(args[0], replace(_block(args[0], key), rows=(("Synthetic primary-row caveat",),)))
    result = certified.supplied_native_table_flows(draft, *args[1:], **kwargs)
    assert ("proposals" if key == "proposals" else "utility_analysis") not in result


@pytest.mark.parametrize("change", ["foreign_reference", "duplicate_table", "reordered_tables", "competing_owner", "competing_text", "excluded_section"])
def test_capacity_requires_unique_canonical_source_ownership(proof_inputs, change):
    args, kwargs = proof_inputs
    draft = args[0]
    block = _block(draft, "thermal_capacity")
    if change == "foreign_reference":
        table = replace(block.extra_tables[0], reference="docx:" + "f" * 64 + ":synthetic-foreign-table")
        draft = _replace_block(draft, replace(block, extra_tables=(table, *block.extra_tables[1:])))
    elif change == "duplicate_table":
        draft = _replace_block(draft, replace(block, extra_tables=(block.extra_tables[0],) * 2))
    elif change == "reordered_tables":
        draft = _replace_block(draft, replace(block, extra_tables=tuple(reversed(block.extra_tables))))
    elif change == "competing_owner":
        other = _block(draft, "training_summary")
        draft = _replace_block(draft, replace(other, extra_tables=other.extra_tables + block.extra_tables[:1]))
    elif change == "competing_text":
        other = ResolvedBlock("synthetic_competing_context", "This month", text=CAVEAT)
        draft = replace(draft, blocks=draft.blocks + (other,), sections=tuple(
            replace(section, blocks=section.blocks + (BlockSpec(other.key, "narrative"),))
            if section.key == "scorecards" else section for section in draft.sections))
    else:
        draft = replace(draft, sections=tuple(replace(section, included=False)
            if section.key == "scorecards" else section for section in draft.sections))
    assert "utility_analysis" not in certified.supplied_native_table_flows(draft, *args[1:], **kwargs)


@pytest.mark.parametrize("change", ["unreferenced_note", "unproved_note", "earlier_anchor", "duplicate_anchor"])
def test_proposal_suffix_exception_requires_one_proved_later_native_note(proof_inputs, change):
    args, kwargs = proof_inputs
    draft = args[0]
    block = _block(draft, "proposals")
    if change == "unreferenced_note":
        draft = _replace_block(draft, replace(block, references=()))
    elif change == "unproved_note":
        kwargs = dict(kwargs, supplied_notes={})
    else:
        tables = [args[4][table.reference] for table in block.extra_tables]
        later = args[1][tables[-1].position - 1]
        reference = next(later.iter(qn("w:footnoteReference")))
        if change == "earlier_anchor":
            reference.getparent().remove(reference)
            earlier = args[1][tables[0].position - 1]
            next(earlier.iter(qn("w:r"))).append(reference)
        else:
            reference.getparent().append(deepcopy(reference))
    assert "proposals" not in certified.supplied_native_table_flows(draft, *args[1:], **kwargs)


def test_ambiguous_native_note_text_does_not_authorize_single_suffix(tmp_path, monkeypatch):
    def add_second_note(document):
        part = document.part.part_related_by(RELATIONSHIP_TYPE.FOOTNOTES)
        notes = parse_xml(part.blob)
        first = next(note for note in notes if note.get(qn("w:id")) == "1")
        second = deepcopy(first)
        second.set(qn("w:id"), "2")
        notes.append(second)
        part._blob = etree.tostring(notes)
        later = next(table for table in document.tables if NEW_MARKER in _text(table._tbl))
        reference = OxmlElement("w:footnoteReference")
        reference.set(qn("w:id"), "2")
        later.cell(1, 0).paragraphs[0].add_run()._r.append(reference)

    path, draft, loader, view = flow_case(tmp_path, monkeypatch, mutate_source=add_second_note)
    block = _block(draft, "proposals")
    inspection = view.original_inspection
    second = next(item for item in inspection.items if item.note == "footnote:2")
    reference = f"docx:{inspection.sha256}:{second.id}"
    draft = _replace_block(draft, replace(block, references=block.references + (reference,)))
    raw, (_, kwargs, result) = _recorded_render((path, draft, loader, view), monkeypatch)
    assert {"footnote:1", "footnote:2"} <= kwargs["supplied_notes"].keys()
    assert "proposals" not in result
    _assert_authored_run(raw, block.text)

@pytest.mark.parametrize("hazard", ["note_run", "inherited_note_style", "reference_run"])
def test_hidden_source_note_cannot_suppress_visible_author_qualification(
        tmp_path, monkeypatch, hazard):
    """No source-visible note means the authored trailing caveat stays in body."""
    def conceal(document):
        if hazard == "reference_run":
            later = next(table for table in document.tables if NEW_MARKER in _text(table._tbl))
            reference = next(later._tbl.iter(qn("w:footnoteReference")))
            reference_run = next(reference.iterancestors(qn("w:r")))
            reference_run.get_or_add_rPr().append(OxmlElement("w:vanish"))
            return
        part = document.part.part_related_by(RELATIONSHIP_TYPE.FOOTNOTES)
        root = parse_xml(part.blob)
        note = next(node for node in root
                    if node.tag == qn("w:footnote") and node.get(qn("w:id")) == "1")
        paragraph = note[0]
        if hazard == "note_run":
            run = next(node for node in paragraph.findall(qn("w:r"))
                       if NATIVE_NOTE in _text(node))
            properties = OxmlElement("w:rPr")
            properties.append(OxmlElement("w:vanish"))
            run.insert(0, properties)
        else:
            style = document.styles.add_style(
                "Synthetic Hidden Footnote Paragraph", WD_STYLE_TYPE.PARAGRAPH)
            style.font.hidden = True
            properties = OxmlElement("w:pPr")
            selector = OxmlElement("w:pStyle")
            selector.set(qn("w:val"), style.style_id)
            properties.append(selector)
            paragraph.insert(0, properties)
        part._blob = etree.tostring(root)

    flow = flow_case(tmp_path, monkeypatch, mutate_source=conceal)
    block = _block(flow[1], "proposals")
    raw, (_, kwargs, result) = _recorded_render(flow, monkeypatch)
    assert "footnote:1" in kwargs["supplied_notes"]
    assert "proposals" not in result
    assert _text(Document(BytesIO(raw)).element.body).count(NATIVE_NOTE) == 1
    assert NATIVE_NOTE in block.text
