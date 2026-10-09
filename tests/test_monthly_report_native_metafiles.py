"""Raster review proves native vector bytes only for the current mapped image."""
from dataclasses import replace
from hashlib import sha256
from io import BytesIO
import struct
from zipfile import ZipFile, ZIP_DEFLATED

from docx import Document
from docx.shared import Inches
from PIL import Image
import pytest

from app.monthly_report_import import inspect_docx, read_import_image
from app.monthly_report_model import Facility, ReportDraft, ReportPeriod, ReportProfile, ResolvedBlock, default_sections
from app.monthly_report_native_layout import build_native_docx


def vector_fixture(tmp_path, monkeypatch):
    png = BytesIO()
    Image.new('RGB', (100, 50), 'blue').save(png, 'PNG')
    header = bytearray(88)
    struct.pack_into('<II', header, 0, 1, 88)
    struct.pack_into('<iiii', header, 8, 0, 0, 100, 50)
    struct.pack_into('<iiii', header, 24, 0, 0, 2540, 1270)
    header[40:44] = b' EMF'
    struct.pack_into('<II', header, 48, 108, 2)
    struct.pack_into('<iiii', header, 72, 100, 50, 25, 13)
    emf = bytes(header) + struct.pack('<IIIII', 14, 20, 0, 0, 20)
    doc = Document()
    doc.add_heading('Equipment Performance Issues', 1)
    doc.add_paragraph().add_run().add_picture(BytesIO(png.getvalue()), width=Inches(3))
    path = tmp_path / 'vector.docx'
    doc.save(path)
    with ZipFile(path) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    parts['word/media/image1.emf'] = emf
    del parts['word/media/image1.png']
    parts['word/_rels/document.xml.rels'] = parts['word/_rels/document.xml.rels'].replace(b'media/image1.png', b'media/image1.emf')
    parts['[Content_Types].xml'] = parts['[Content_Types].xml'].replace(b'</Types>', b'<Default Extension="emf" ContentType="image/x-emf"/></Types>')
    with ZipFile(path, 'w', ZIP_DEFLATED) as archive:
        for name, raw in parts.items():
            archive.writestr(name, raw)
    from app import monthly_report_metafiles
    monkeypatch.setattr(monthly_report_metafiles, 'rasterize_metafile', lambda raw, suffix: png.getvalue())
    inspection = inspect_docx(path)
    image = next(item for item in inspection.items if item.kind == 'image')
    assert image.section == 'issues'
    reviewed = read_import_image(path, image)
    digest = sha256(reviewed.data).hexdigest() + '.' + reviewed.extension
    assets = {digest: reviewed.data}
    block = ResolvedBlock('equipment_issues_evidence', 'This month', asset_hashes=(digest,),
                          references=(f'docx:{inspection.sha256}:{image.id}',))
    draft = ReportDraft(ReportProfile('Synthetic', 'site', 'Current site', (Facility('site', 'Current site'),)),
                        ReportPeriod(2026, 9), 'Current editor', default_sections(), (block,))
    return path, draft, assets, emf


@pytest.mark.parametrize('change', ['unchanged', 'unmapped', 'changed'])
def test_reviewed_vector_retained_only_with_exact_current_raster_proof(tmp_path, monkeypatch, change):
    path, draft, assets, emf = vector_fixture(tmp_path, monkeypatch)
    block = draft.blocks[0]
    if change == 'unmapped':
        block = replace(block, asset_hashes=())
    elif change == 'changed':
        png = BytesIO()
        Image.new('RGB', (100, 50), 'red').save(png, 'PNG')
        assets['current.png'] = png.getvalue()
        block = replace(block, asset_hashes=('current.png',))
    result = build_native_docx(replace(draft, blocks=(block,)), path, asset_loader=assets.__getitem__)
    with ZipFile(BytesIO(result)) as archive:
        media = [archive.read(name) for name in archive.namelist() if name.startswith('word/media/')]
    assert (emf in media) == (change == 'unchanged')
    assert len(media) == (0 if change == 'unmapped' else 1)
