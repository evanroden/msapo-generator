"""Actual final-PDF destinations and Word caches for generated contents."""
from dataclasses import replace
from io import BytesIO
from zipfile import ZipFile, ZIP_DEFLATED
from types import SimpleNamespace

from docx import Document
from docx.oxml.ns import qn
from PIL import Image
import pymupdf as fitz
import pytest

from app import monthly_report_docx as renderer
from app.monthly_report_contents_pagination import anchor_for, synchronize
from app.monthly_report_model import (
    BlockSpec, ReportDraft, ReportPeriod, ResolvedBlock, SectionSpec,
    included_sections, synthetic_profiles,
)
from tests.conftest import requires_libreoffice


def current_report(*, pages=0, reverse=False):
    from app.monthly_report_asset_review import approve_all_assets
    activity = SectionSpec("activity", "1", "Monthly Activity Summary", (
        BlockSpec("activity_summary", "rich_text"), BlockSpec("evidence", "pdf_pages")))
    training = SectionSpec("training", "2", "Training Summary", (BlockSpec("training_summary", "rich_text"),))
    omitted = SectionSpec("water", "3", "Water Treatment Reports", (), included=False)
    return ReportDraft(synthetic_profiles()[0], ReportPeriod(2026, 10), "Synthetic Editor",
        (training, omitted, activity) if reverse else (activity, omitted, training), (
            ResolvedBlock("activity_summary", "This month", text="Current activity detail."),
            approve_all_assets(ResolvedBlock("evidence", "This month", asset_hashes=tuple(f"page-{i}" for i in range(pages)))),
            ResolvedBlock("training_summary", "This month", text="Current training detail."),
        ))


def image(_):
    data = BytesIO()
    Image.new("RGB", (480, 640), "navy").save(data, "PNG")
    return data.getvalue()


def cached_pages(raw):
    document = Document(BytesIO(raw))
    return {field.get(qn("w:instr")).strip(): "".join(t.text or "" for t in field.iter(qn("w:t")))
            for field in document.element.iter(qn("w:fldSimple"))
            if field.get(qn("w:instr"), "").strip().startswith("PAGEREF mr_contents_")}


def test_assembled_contents_has_internal_targets_for_selected_order_only(tmp_path, monkeypatch):
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path / "data"))
    value = current_report(reverse=True)
    raw = renderer.assemble_docx(value, acknowledged_fingerprint=value.fingerprint)
    document = Document(BytesIO(raw))
    fields = cached_pages(raw)
    assert list(fields) == [f"PAGEREF {anchor_for(s.key)} \\h" for s in included_sections(value.sections)]
    assert set(fields.values()) == {"?"}
    assert len(list(document.element.iter(qn("w:bookmarkStart")))) == 2
    assert "Water Treatment Reports" not in " ".join(document.element.itertext())
    assert synchronize(raw, None, included_sections(value.sections))[1]


@requires_libreoffice
@pytest.mark.parametrize("reverse", [False, True])
def test_real_pdf_and_word_contents_follow_added_pages_omissions_and_order(tmp_path, monkeypatch, reverse):
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path / "data"))
    calls = []
    convert = renderer._convert_report_pdf
    def counted(*args):
        calls.append(args[0].fingerprint)
        return convert(*args)
    monkeypatch.setattr(renderer, "_convert_report_pdf", counted)
    outputs = []
    for pages in (0, 3):
        value = current_report(pages=pages, reverse=reverse)
        result = renderer.generate_report(value, acknowledged_fingerprint=value.fingerprint, asset_loader=image)
        assert result.pdf, result.pdf_error
        assert result.contents_notice == ""
        expected = included_sections(value.sections)
        with fitz.open(stream=result.pdf, filetype="pdf") as pdf:
            outline = pdf.get_toc()
            assert [entry[1] for entry in outline] == [f"{s.number}. {s.title}" for s in expected]
            assert "Water Treatment Reports" not in "\n".join(page.get_text() for page in pdf)
            cache = cached_pages(result.docx)
            for section, entry in zip(expected, outline):
                assert cache[f"PAGEREF {anchor_for(section.key)} \\h"] == str(entry[2])
                assert section.title in pdf[entry[2] - 1].get_text()
            outputs.append(outline)
        # Repeating synchronization is byte-stable, with no conversion.
        updated, notice = synchronize(result.docx, result.pdf, expected)
        assert renderer.canonical_docx(updated) == result.docx and notice == ""
    assert len(calls) == 2  # Exactly one existing conversion per report.
    if not reverse:
        assert outputs[1][-1][2] > outputs[0][-1][2]


def test_non_generated_document_is_byte_identical_and_has_no_claim():
    document = Document(); document.add_paragraph("Unrelated editable report")
    data = BytesIO(); document.save(data)
    assert synchronize(data.getvalue(), b"%PDF-unavailable", ()) == (data.getvalue(), None)


@requires_libreoffice
def test_missing_pdf_links_do_not_publish_unverified_page_cache(tmp_path, monkeypatch):
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path / "data"))
    value = current_report()
    raw = renderer.assemble_docx(value, acknowledged_fingerprint=value.fingerprint)
    pdf, error = renderer._convert_report_pdf(value, raw)
    assert pdf, error
    with fitz.open(stream=pdf, filetype="pdf") as altered:
        for page in altered:
            for link in page.get_links():
                page.delete_link(link)
        raw_pdf = altered.tobytes()
    unchanged, notice = synchronize(raw, raw_pdf, included_sections(value.sections))
    assert unchanged == raw and "could not be verified" in notice
    assert set(cached_pages(unchanged).values()) == {"?"}


@requires_libreoffice
def test_pdf_retry_fills_only_generated_caches_without_reassembly(tmp_path, monkeypatch):
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path / "data"))
    value = current_report()
    convert = renderer._convert_report_pdf
    monkeypatch.setattr(renderer, "_convert_report_pdf", lambda *args: (None, "Synthetic first failure"))
    package = renderer.generate_report(value, acknowledged_fingerprint=value.fingerprint)
    assert package.pdf is None and package.contents_notice
    monkeypatch.setattr(renderer, "assemble_docx", lambda *args, **kwargs: pytest.fail("Retry must not reassemble"))
    monkeypatch.setattr(renderer, "_convert_report_pdf", convert)
    result = renderer.retry_report_pdf(value, package)
    assert result.pdf and not result.pdf_error and not result.contents_notice
    assert set(cached_pages(result.docx).values()) != {"?"}
    from app.monthly_report_contents_pagination import only_generated_cache_changed
    assert only_generated_cache_changed(package.docx, result.docx)
    with ZipFile(BytesIO(package.docx)) as before, ZipFile(BytesIO(result.docx)) as after:
        assert before.namelist() == after.namelist()
        assert all(before.read(name) == after.read(name) for name in before.namelist() if name != "word/document.xml")


@pytest.mark.parametrize("change", ["cache", "body", "target", "other-part"])
def test_cache_only_proof_preserves_failed_save_receipt_without_accepting_content_changes(tmp_path, monkeypatch, change):
    from app import monthly_report_output_ui as output
    from app.monthly_report_contents_pagination import only_generated_cache_changed
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path / "data"))
    value = current_report()
    raw = renderer.assemble_docx(value, acknowledged_fingerprint=value.fingerprint)
    document = Document(BytesIO(raw))
    for field in document.element.iter(qn("w:fldSimple")):
        if field.get(qn("w:instr"), "").strip().startswith("PAGEREF mr_contents_"):
            next(field.iter(qn("w:t"))).text = "4"
    if change == "body":
        document.add_paragraph("An unrelated edit must not inherit a save receipt.")
    elif change == "target":
        next(document.element.iter(qn("w:bookmarkStart"))).set(qn("w:name"), "changed-target")
    elif change == "other-part":
        document.core_properties.title = "Different document"
    stream = BytesIO(); document.save(stream)
    # Production updates document.xml only. python-docx also rewrites the
    # core-properties serialization on save, which is deliberately rejected.
    content = BytesIO()
    with ZipFile(BytesIO(raw)) as original, ZipFile(BytesIO(stream.getvalue())) as edited, ZipFile(content, "w", ZIP_DEFLATED) as target:
        for info in original.infolist():
            selected = info.filename == "word/document.xml" or (change == "other-part" and info.filename == "docProps/core.xml")
            target.writestr(info, (edited if selected else original).read(info.filename))
    updated = renderer.canonical_docx(content.getvalue())
    assert only_generated_cache_changed(raw, updated) == (change == "cache")
    before = renderer.ReportPackage(raw, None, "a.docx", "a.pdf", value.fingerprint, generated_contents=True)
    after = replace(before, docx=updated, pdf=b"%PDF-synthetic")
    state = {"q_package": before, "q_revision": 2}
    fake = SimpleNamespace(session_state=state)
    monkeypatch.setattr(output, "st", fake)
    monkeypatch.setattr(output, "discard_current_session_pdf_preview", lambda *args: None)
    output.record_shared_save("q", before, error="Synthetic shared save failed")
    receipt = dict(state["q_package_save"])
    output.replace_current_session_pdf_package("q", "scope", after)
    if change == "cache":
        assert state["q_package_save"] == {**receipt, "identity": output.package_identity(after)}
    else:
        assert state["q_package_save"] == receipt
    assert state["q_package_save"]["error"] == "Synthetic shared save failed"
