"""Synthetic edited body units retain exact source-owned page chrome."""
from dataclasses import replace
import pytest
from app.monthly_report_native_layout import build_native_docx
from test_monthly_report_native_art_review import sample
from test_monthly_report_native_art_lifecycle import save_sample


def test_multiline_activity_keeps_reviewed_source_region(sample):
    path, view, profile, period, inventory, receipt = sample
    draft, loader = save_sample(sample)
    draft = replace(draft, blocks=tuple(replace(block, text='Synthetic current qualification.\n\n' + block.text)
        if block.key == 'activity_summary' else block for block in draft.blocks))
    assert build_native_docx(draft, path, asset_loader=loader, section_key='activity')

from io import BytesIO
from copy import deepcopy
from docx import Document
from docx.oxml import OxmlElement
from app.monthly_report_native_art_scope import (expected_header_scope, source_emission_proof,
    verify_header_routing, NativeEmissionProof, W, R, _routing, _package)
from app.monthly_report_native_pages import NativePages
from test_monthly_report_native_art_scope import sample as routed_sample, rewrite, xml


def ledger(view):
    document = Document(BytesIO(view.rendition_bytes))
    body = document.element.body
    scope = expected_header_scope(view)
    pages = NativePages(body, tuple(body), dict(scope.source_plan.section_by_position),
                        set(scope.source_plan.headings), set(scope.selected_positions))
    return document, body, scope, pages


def raw(document):
    output = BytesIO(); document.save(output); return output.getvalue()


def test_source_owned_paragraph_and_table_growth_and_shrink(routed_sample):
    _, view = routed_sample
    document, body, scope, pages = ledger(view)
    origin = next(node for node in body if node.tag == W+'p' and node.find('.//'+W+'sectPr') is None)
    paragraph = OxmlElement('w:p'); origin.addnext(paragraph); pages.assigned(paragraph, origin)
    table = OxmlElement('w:tbl'); paragraph.addnext(table); pages.assigned(table, origin)
    proof = source_emission_proof(scope, body, pages)
    verify_header_routing(raw(document), scope, emission_proof=proof)
    body.remove(paragraph); body.remove(table)
    proof = source_emission_proof(scope, body, pages)
    verify_header_routing(raw(document), scope, emission_proof=proof)


@pytest.mark.parametrize('case', ['unknown', 'reordered', 'missing', 'foreign_scope'])
def test_emission_proof_requires_all_roots_and_exact_source_coverage(routed_sample, case):
    _, view = routed_sample
    document, body, scope, pages = ledger(view)
    if case == 'unknown': body.insert(0, OxmlElement('w:p'))
    elif case == 'reordered': body.insert(0, body[-2])
    elif case == 'missing':
        for node in list(body):
            if pages.source_word_sections[pages.owner[node]] == 2: body.remove(node)
    else:
        proof = source_emission_proof(scope, body, pages)
        with pytest.raises(ValueError, match='different source scope'):
            verify_header_routing(raw(document), replace(scope, section_key='activity'), emission_proof=proof)
        return
    with pytest.raises(ValueError, match='owner|order or coverage'):
        source_emission_proof(scope, body, pages)


@pytest.mark.parametrize('case', ['moved', 'extra', 'missing', 'reordered', 'route', 'geometry', 'inserted_root'])
def test_output_changes_cannot_redefine_precomputed_emission_permission(routed_sample, case):
    _, view = routed_sample
    document, body, scope, pages = ledger(view)
    proof = source_emission_proof(scope, body, pages)
    def mutate(parts):
        def change(root):
            body = root.find(W+'body')
            boundaries = [node for node in body if node.find('.//'+W+'sectPr') is not None]
            props = boundaries[0].find('.//'+W+'sectPr')
            if case == 'moved':
                target = next(node for node in body if node.tag == W+'p' and node not in boundaries)
                props.getparent().remove(props)
                ppr = target.find(W+'pPr')
                if ppr is None: ppr=OxmlElement('w:pPr'); target.insert(0,ppr)
                ppr.append(props)
            elif case == 'extra':
                node=OxmlElement('w:p'); ppr=OxmlElement('w:pPr');node.append(ppr);ppr.append(deepcopy(props));body.insert(1,node)
            elif case == 'missing': props.getparent().remove(props)
            elif case == 'reordered':
                other=boundaries[1].find('.//'+W+'sectPr')
                left,right=props.getparent(),other.getparent();left.remove(props);right.remove(other);left.append(other);right.append(props)
            elif case == 'route': props.find(W+'headerReference').set(R+'id','unowned')
            elif case == 'geometry': props.find(W+'pgSz').set(W+'w','20000')
            else: body.insert(0,OxmlElement('w:p'))
        xml(parts,'word/document.xml',change)
    with pytest.raises(ValueError):
        verify_header_routing(rewrite(raw(document),mutate),scope,emission_proof=proof)


def test_empty_terminal_source_section_keeps_its_owner(routed_sample):
    _, view = routed_sample
    document, body, scope, pages = ledger(view)
    final = scope.source_word_sections[-1]
    for node in list(body)[:-1]:
        if pages.source_word_sections[pages.owner[node]] == final: body.remove(node)
    proof = source_emission_proof(scope, body, pages)
    verify_header_routing(raw(document),scope,emission_proof=proof)


@pytest.mark.parametrize('section', [None, 'cover', 'activity', 'issues'])
def test_real_native_issues_multiline_growth_keeps_full_and_preview_scope(sample, section):
    from test_monthly_report_native_art_review import setup
    path, *_ = sample
    document = Document(path)
    document.add_heading('EQUIPMENT PERFORMANCE ISSUES', 1)
    document.add_paragraph('Synthetic pump remains under observation.')
    document.save(path)
    fixture = setup(path)
    draft, loader = save_sample(fixture)
    assert any(b.key == 'equipment_issues' and b.text for b in draft.blocks)
    draft = replace(draft, blocks=tuple(replace(b, text='Synthetic October qualification.\n\n'+b.text)
        if b.key in ('activity_summary','equipment_issues') else b for b in draft.blocks))
    assert build_native_docx(draft,path,asset_loader=loader,section_key=section)


@pytest.mark.parametrize('section', [None, 'issues'])
def test_real_native_issue_table_growth_and_shrink_keep_source_scope(sample, section):
    from test_monthly_report_native_art_review import setup
    from app.monthly_report_model import ReportTable
    path, *_ = sample
    document = Document(path)
    document.add_heading('EQUIPMENT PERFORMANCE ISSUES', 1)
    document.add_paragraph('Synthetic issue notes.')
    table=document.add_table(rows=2, cols=2)
    for row, values in zip(table.rows, [('Equipment','Condition'),('Pump A','Under observation')]):
        for cell,value in zip(row.cells,values):cell.text=value
    document.save(path)
    draft, loader = save_sample(setup(path))
    for count in (3,1):
        tables=tuple(ReportTable(('Equipment','Condition'),((f'Synthetic pump {n}','Under observation'),)) for n in range(count))
        edited=replace(draft,blocks=tuple(replace(b,extra_tables=tables)
            if b.key=='equipment_issues' else b for b in draft.blocks))
        assert build_native_docx(edited,path,asset_loader=loader,section_key=section)
