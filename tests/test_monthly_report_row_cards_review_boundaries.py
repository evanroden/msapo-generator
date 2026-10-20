"""Imported row fields keep readable labels and exact multiline content (Refs #301)."""
from dataclasses import replace
from io import BytesIO

from docx import Document
from streamlit.testing.v1 import AppTest

from app import monthly_report_active_work as active
from app.monthly_report_docx import assemble_docx
from app.monthly_report_model import ReportTable, default_sections
from test_monthly_report_row_cards_ui import SCRIPT, button, example, field


def imported_table(columns, row):
    spec, block = example(extra=True)
    block = replace(block, extra_tables=(ReportTable(columns, (row,), "synthetic-custom-source"),))
    script = SCRIPT.replace("INITIAL", repr(block)).replace("SPEC", repr(spec))
    return AppTest.from_string(script).run(), spec


def assert_saved_word(app, spec, columns, row, tmp_path):
    block = app.session_state["block"]
    table, = block.extra_tables
    assert table.columns == columns and table.rows == (row,)
    assert table.reference == "synthetic-custom-source"
    section = next(s for s in default_sections() if s.key == "proposals")
    draft = replace(app.session_state["draft"], sections=(replace(section, blocks=(spec,)),), blocks=(block,))
    saved = active.save("e" * 32, draft, expected_revision=0, base_snapshot_revision=0)
    reloaded = active.load("e" * 32, draft.profile.contract, draft.profile.key, draft.period)
    assert saved.draft == reloaded.draft == draft
    raw = assemble_docx(reloaded.draft, acknowledged_fingerprint=draft.fingerprint)
    (tmp_path / "custom-row-fields.docx").write_bytes(raw)
    document = Document(BytesIO(raw))
    native = next(t for t in document.tables if t.cell(0, 0).text == "Proposal #")
    assert tuple(c.text for c in native.rows[0].cells) == columns
    assert tuple(c.text for c in native.rows[1].cells) == row


def test_blank_imported_headings_get_display_labels_without_changing_saved_or_word_schema(tmp_path, monkeypatch):
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path))
    columns = ("Proposal #", "", "  ")
    original = ("0007", "Synthetic custom value", "Retain this other field")
    app, spec = imported_table(columns, original)
    button(app, "Edit proposal").click().run()
    assert not app.exception
    assert tuple(item.label for item in app.text_input) == ("Proposal #", "Column 2", "Column 3")
    identity = field(app, "Column 2").proto.id
    field(app, "Column 2").set_value("Synthetic edited custom value")
    button(app, "Close editing").click().run()
    button(app, "Edit proposal").click().run()
    assert field(app, "Column 2").proto.id == identity
    assert field(app, "Column 2").value == "Synthetic edited custom value"
    button(app, "Close editing").click().run()
    assert_saved_word(app, spec, columns, (original[0], "Synthetic edited custom value", original[2]), tmp_path)


def test_multiline_custom_value_keeps_native_editor_identity_and_exact_saved_word_lines(tmp_path, monkeypatch):
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path))
    columns = ("Proposal #", "Task", "Custom field")
    original = ("0007", "Inspect pump\nRetain the second line", "Unchanged source detail")
    app, spec = imported_table(columns, original)
    button(app, "Edit proposal").click().run()
    assert not app.exception
    assert tuple(item.label for item in app.text_area) == ("Task",)
    task = field(app, "Task", "text_area")
    identity = task.proto.id
    assert task.value == original[1]
    task.set_value("Temporary one-line value").run()
    assert field(app, "Task", "text_area").proto.id == identity
    field(app, "Task", "text_area").set_value("").run()
    assert field(app, "Task", "text_area").proto.id == identity
    edited = "Inspect pump after service\nRetain the follow-up\nSynthetic third line"
    field(app, "Task", "text_area").set_value(edited)
    button(app, "Close editing").click().run()
    button(app, "Edit proposal").click().run()
    assert field(app, "Task", "text_area").proto.id == identity
    assert field(app, "Task", "text_area").value == edited
    button(app, "Close editing").click().run()
    assert_saved_word(app, spec, columns, (original[0], edited, original[2]), tmp_path)
