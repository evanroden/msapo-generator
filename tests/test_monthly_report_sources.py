from dataclasses import asdict, replace
from email.message import EmailMessage
from io import BytesIO
import struct

import fitz
from PIL import Image
from openpyxl import Workbook
import pytest

from app import monthly_report_library as library, monthly_report_sources as sources
from app.monthly_report_checks import preflight
from app.monthly_report_model import ReportPeriod, synthetic_draft, synthetic_profiles


@pytest.fixture
def profile(monkeypatch, tmp_path):
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path))
    profile = synthetic_profiles()[0]
    library.save_profile(profile, expected_revision=0, actor="Synthetic Editor", confirmed=True)
    return profile


def pdf_bytes(count=26):
    with fitz.open() as pdf:
        for i in range(count):
            page = pdf.new_page()
            page.insert_text((50, 50), f"Synthetic service report page {i+1}. Completed inspection of DEMO-01.")
        return pdf.tobytes()


def image_bytes():
    output = BytesIO()
    Image.new("RGB", (32, 24), "blue").save(output, "PNG")
    return output.getvalue()


def minimal_msg():
    """A genuine synthetic CFB container, generated rather than a client fixture.

    Property streams use regular FAT storage (4096 bytes) to avoid mini-stream
    machinery. Only plain text and subject are needed to exercise MSG decoding.
    """
    end, free, fat = 0xFFFFFFFE, 0xFFFFFFFF, 0xFFFFFFFD
    header = bytearray(512)
    header[:8] = bytes.fromhex("d0cf11e0a1b11ae1")
    struct.pack_into("<HHHHH", header, 24, 0x3E, 3, 0xFFFE, 9, 6)
    struct.pack_into("<IIIIIIIII", header, 40, 0, 1, 1, 0, 4096, end, 0, end, 0)
    struct.pack_into("<109I", header, 76, 0, *([free] * 108))
    table = [fat, end] + [i+1 for i in range(2, 18)] + [free] * 110
    table[9] = table[17] = end
    directory = bytearray(512)
    def entry(index, name, kind, start, size, child=free, right=free):
        offset = index * 128
        encoded = (name + "\x00").encode("utf-16-le")
        directory[offset:offset+len(encoded)] = encoded
        struct.pack_into("<HBBIII", directory, offset+64, len(encoded), kind, 1, free, right, child)
        struct.pack_into("<IQ", directory, offset+116, start, size)
    entry(0, "Root Entry", 5, end, 0, child=1)
    entry(1, "__substg1.0_0037001F", 2, 2, 4096, right=2)
    entry(2, "__substg1.0_1000001F", 2, 10, 4096)
    subject = "Synthetic service report".encode("utf-16-le").ljust(4096, b"\0")
    body = "Vendor: Example Mechanical\nDate: 2026-09-12\nCompleted: Repaired DEMO-01.".encode("utf-16-le").ljust(4096, b"\0")
    return bytes(header) + struct.pack("<128I", *table) + directory + subject + body


def test_native_26_page_pdf_is_not_truncated_to_vision_limit(profile):
    content, _ = sources.ingest(profile, "synthetic.pdf", pdf_bytes())
    assert len(content.source.page_texts) == 26
    assert content.source.selected_pages == tuple(range(1, 27))
    assert "page 26" in content.source.page_texts[-1]


def test_pdf_page_selection_caption_and_cache(profile, monkeypatch):
    content, _ = sources.ingest(profile, "synthetic.pdf", pdf_bytes(3))
    content = replace(content, source=replace(content.source, selected_pages=(1, 3), captions=((3, "Inspected final page"),)))
    from app.monthly_report_content_policy import page_fingerprint
    content = replace(content, source=replace(content.source, client_page_reviews=tuple((n, page_fingerprint(content.source, n)) for n in (1, 3))))
    blocks = sources.prepare_pages(profile, (content,), ((content.source.id, "vendor_reports"),))
    assert len(blocks[0].asset_hashes) == 2
    assert blocks[0].asset_captions[-1] == "Inspected final page"
    assert all(library.read_asset(profile.contract, profile.key, ref) for ref in blocks[0].asset_hashes)
    monkeypatch.setattr(sources.fitz, "open", lambda *a, **k: pytest.fail("cached page decoded again"))
    assert sources.page_image(profile, content, 3).data


def test_duplicate_digest_suppression_and_size_budget(profile, monkeypatch):
    raw = image_bytes()
    result, messages = sources.ingest_batch(profile, (("photo.png", raw), ("renamed.png", raw)))
    assert len(result) == 1 and "Duplicate skipped" in messages[0]
    monkeypatch.setattr(sources, "MAX_REPORT_BYTES", 1)
    result, messages = sources.ingest_batch(profile, (("photo.png", raw),))
    assert not result and "Report limit" in messages[0]


def test_extraction_cache_uses_saved_source(profile, monkeypatch):
    raw = b"Work order,Date\n001,2026-09-01\n"
    content, _ = sources.ingest(profile, "demo.csv", raw)
    monkeypatch.setattr(sources, "_spreadsheet", lambda *a: pytest.fail("cached CSV reparsed"))
    cached, _ = sources.ingest(profile, "renamed.csv", raw)
    assert cached.source.filename == "renamed.csv"
    assert cached.tables == content.tables


def test_image_validation_and_thumbnails(profile):
    content, _ = sources.ingest(profile, "photo.png", image_bytes())
    assert content.source.needs_vision == (1,)
    assert sources.page_image(profile, content, 1, preview=True).width == 32
    with pytest.raises((ValueError, OSError)):
        sources.ingest(profile, "invalid.png", b"invalid")


def test_csv_preserves_leading_zero_and_xlsx_preserves_false(profile):
    content, _ = sources.ingest(profile, "cmms.csv", b"Work order,Value\n001,0\n002,False\n")
    assert content.tables[0].rows[0] == ("001", "0")
    book = Workbook()
    sheet = book.active
    sheet.append(["Identity", "Number", "Complete", "Formula"])
    sheet.append(["001", 0, False, "=1+1"])
    book.create_sheet("Empty")
    out = BytesIO()
    book.save(out)
    content, _ = sources.ingest(profile, "cmms.xlsx", out.getvalue())
    assert content.tables[0].rows[0] == ("001", "0", "No", "")
    assert len(content.tables) == 1 and "formulas are never executed" in content.source.notices[0]


def test_eml_body_and_attachment_are_read_without_external_fetch(profile):
    message = EmailMessage()
    message["From"] = "service@example.invalid"
    message["Subject"] = "Synthetic service report"
    message.set_content("Vendor: Example Mechanical\nDate: 09/12/2026\nFindings: Inspection complete.")
    message.add_attachment(image_bytes(), maintype="image", subtype="png", filename="../../photo.png")
    result, notices = sources.ingest_batch(profile, (("service.eml", message.as_bytes()),))
    assert len(result) == 2 and not notices
    assert result[0].source.vendor == "Example Mechanical"
    assert result[0].source.service_date == "2026-09-12"
    assert result[1].source.filename == "photo.png"


def test_msg_uses_real_compound_file_reader(profile):
    content, attachments = sources.ingest(profile, "synthetic.msg", minimal_msg())
    assert not attachments
    assert content.source.vendor == "Example Mechanical"
    assert "Repaired DEMO-01" in content.source.actions
    assert content.source.service_date == "2026-09-12"


def test_html_body_is_text_only():
    result = sources._plain_html('<p>Inspection complete.</p><img src="https://example.invalid/a"><script>bad()</script>')
    assert "Inspection complete." in result and "bad()" not in result


def test_native_docx_items_and_images(profile, tmp_path):
    from docx import Document
    document = Document()
    document.add_paragraph("Vendor: Example Mechanical")
    document.add_picture(BytesIO(image_bytes()))
    path = tmp_path / "source.docx"
    document.save(path)
    content, _ = sources.ingest(profile, path.name, path.read_bytes())
    assert content.image_items
    number, _ = content.image_items[0]
    assert sources.page_image(profile, content, number).data
    assert "extracted items" in content.source.notices[-1]


def test_source_fact_date_facility_warnings_pin_snapshot(profile):
    content, _ = sources.ingest(profile, "vendor.csv", b"Vendor,Date\nExample Mechanical,2026-08-01\n")
    source = replace(content.source, classification="Vendor service", service_date="2026-08-01", facility="Other synthetic facility")
    draft = replace(synthetic_draft(profile, ReportPeriod(2026, 9)), sources=(source,))
    assert {c.code for c in preflight(draft)} >= {"source_period", "source_facility"}
    assert replace(draft, sources=(replace(source, service_date="2026-09-01"),)).fingerprint != draft.fingerprint
    assert library.draft_from_dict(asdict(draft)) == draft
    saved = library.save_snapshot(draft, expected_revision=0, entered_editor="Synthetic Editor")
    assert saved.draft.sources == (source,)


def test_page_and_embed_limits(profile, monkeypatch):
    content, _ = sources.ingest(profile, "synthetic.pdf", pdf_bytes(2))
    monkeypatch.setattr(sources, "MAX_EMBED_PAGES", 1)
    with pytest.raises(ValueError, match="150 selected"):
        sources.prepare_pages(profile, (content,), ((content.source.id, "water_reports"),))
    with pytest.raises(ValueError, match="does not exist"):
        sources.page_image(profile, content, 3)


def test_bad_file_errors_do_not_abort_remaining_batch(profile):
    result, notices = sources.ingest_batch(profile, (("bad.pdf", b"not pdf"), ("good.png", image_bytes())))
    assert len(result) == 1 and notices


def test_integrity_and_unsafe_reference_rejected(profile):
    content, _ = sources.ingest(profile, "photo.png", image_bytes())
    sources._path(profile, content.source.sha256, ".png").write_bytes(b"changed")
    with pytest.raises(ValueError, match="integrity"):
        sources.source_bytes(profile, content.source)
    with pytest.raises(ValueError, match="reference"):
        sources._path(profile, "../bad", ".png")
