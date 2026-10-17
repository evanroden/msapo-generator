"""A same-site Starting cover has one photo choice beside separate art recovery."""

from io import BytesIO

from docx import Document
from docx.enum.section import WD_SECTION
from docx.shared import Inches
from PIL import Image

from app import monthly_report_library as library
from app.monthly_report_asset_review import pending_asset_indexes
from app.monthly_report_import import inspect_docx
from app.monthly_report_model import synthetic_profiles
from app.monthly_report_start_import import build_starting_import
from test_monthly_report_starting_review_recovery import drawing_report, continue_button
from tests.test_monthly_report_start_import import start_app


def cover_plan(app):
    return next(value["cover"] for key, value in app.session_state.filtered_state.items()
                if key.endswith("_section_plans") and isinstance(value, dict))


def with_candidates(tmp_path):
    path, _ = drawing_report(tmp_path, "cover")
    document = Document(path)
    stream = BytesIO()
    Image.new("RGB", (450, 300), "teal").save(stream, "PNG")
    document.add_picture(BytesIO(stream.getvalue()), width=Inches(3))
    document.save(path)
    return path


def photo_report(tmp_path):
    document = Document()
    document.add_heading(synthetic_profiles()[0].facilities[0].title, 0)
    document.add_paragraph("September 2026")
    for color, size in (("navy", (600, 400)), ("teal", (450, 300))):
        stream = BytesIO()
        Image.new("RGB", size, color).save(stream, "PNG")
        document.add_picture(BytesIO(stream.getvalue()), width=Inches(3))
    path = tmp_path / "invented-cover-photos.docx"
    document.save(path)
    return path


def test_unresolved_cover_has_one_photo_gallery_and_stays_blocked(monkeypatch, tmp_path):
    path = with_candidates(tmp_path)
    raw = path.read_bytes()
    app = start_app(monkeypatch, tmp_path, path)
    for _ in range(2):
        assert not app.exception
        assert sum(b.label == "No cover photograph" for b in app.button) == 1
        assert sum(x.value == "**Cover photograph**" for x in app.markdown) == 1
        assert any("could not be preserved" in x.value for x in app.error)
        assert continue_button(app).disabled
        assert path.read_bytes() == raw
        assert library.list_profiles(synthetic_profiles()[0].contract) == ()
        app.run()


def test_no_photo_choice_is_not_reversed_by_unresolved_art_card(monkeypatch, tmp_path):
    path = with_candidates(tmp_path)
    raw = path.read_bytes()
    app = start_app(monkeypatch, tmp_path, path)
    # The existing dedicated gallery is the ordinary first visible selection.
    gallery = next(b for b in app.button if b.label == "No cover photograph"
                   and "_cover_gallery_" in b.key)
    gallery.click().run()
    assert not app.exception
    assert cover_plan(app).get("cover_picture_id") == ""
    assert "cover_photo" not in cover_plan(app)["destinations"].values()
    assert not cover_plan(app)["approved"]
    app.run()
    assert "cover_photo" not in cover_plan(app)["destinations"].values()
    assert continue_button(app).disabled
    assert path.read_bytes() == raw
    assert library.list_profiles(synthetic_profiles()[0].contract) == ()


def test_chosen_photo_survives_recovery_edit_and_gallery_close_reopen(monkeypatch, tmp_path):
    path = with_candidates(tmp_path)
    app = start_app(monkeypatch, tmp_path, path)
    omit = next(b for b in app.button if b.label == "No cover photograph")
    native_id = omit.proto.id
    original = cover_plan(app)["cover_picture_id"]
    next(b for b in app.button if b.label == "Choose picture 2").click().run()
    chosen = cover_plan(app)["cover_picture_id"]
    assert chosen and chosen != original
    assert list(cover_plan(app)["destinations"].values()).count("cover_photo") == 1
    next(r for r in app.radio if r.label == "What would you like to do?").set_value("Edit text or change pictures").run()
    assert cover_plan(app)["cover_picture_id"] == chosen
    assert next(b for b in app.button if b.label == "No cover photograph").proto.id == native_id
    assert continue_button(app).disabled

    # AppTest cannot click expander headers. Deliver its native open state to
    # exercise the real closed/open branches; this is not browser focus proof.
    gallery_key = next(key for key in app.session_state.filtered_state if key.endswith("_cover_gallery_expander"))
    app.session_state[gallery_key] = False
    app.run()
    assert not app.exception
    assert not any(b.label == "No cover photograph" for b in app.button)
    assert cover_plan(app)["cover_picture_id"] == chosen
    app.session_state[gallery_key] = True
    app.run()
    assert not app.exception
    assert sum(b.label == "No cover photograph" for b in app.button) == 1
    assert cover_plan(app)["cover_picture_id"] == chosen
    assert continue_button(app).disabled


def test_later_contents_artwork_is_not_promoted_to_a_second_cover_gallery(monkeypatch, tmp_path):
    path, _ = drawing_report(tmp_path, "cover")
    document = Document(path)
    document.add_section(WD_SECTION.NEW_PAGE)
    document.add_heading("Contents", 0)
    stream = BytesIO()
    Image.new("RGB", (700, 400), "orange").save(stream, "PNG")
    document.add_picture(BytesIO(stream.getvalue()), width=Inches(4))
    document.save(path)
    app = start_app(monkeypatch, tmp_path, path)
    assert not app.exception
    assert sum(b.label == "No cover photograph" for b in app.button) == 1
    assert not any(b.label.startswith("Choose picture") for b in app.button)
    assert len(cover_plan(app)["picture_choices"]) == 1
    assert continue_button(app).disabled


def test_gallery_choice_does_not_approve_imported_photo_for_client_export(monkeypatch, tmp_path):
    path = photo_report(tmp_path)
    app = start_app(monkeypatch, tmp_path, path)
    next(b for b in app.button if b.label == "Choose picture 2").click().run()
    assert not app.exception
    assert not continue_button(app).disabled
    plans = next(value for key, value in app.session_state.filtered_state.items()
                 if key.endswith("_section_plans") and isinstance(value, dict))
    mapped, _ = build_starting_import(path, inspect_docx(path), plans.values())
    photo = next(block for block in mapped.blocks if block.key == "cover_photo")
    assert len(photo.asset_hashes) == 1
    assert pending_asset_indexes(photo) == (0,)
    assert library.list_profiles(synthetic_profiles()[0].contract) == ()


def test_no_photo_is_valid_without_unread_art_and_survives_reload(monkeypatch, tmp_path):
    path = photo_report(tmp_path)
    raw = path.read_bytes()
    app = start_app(monkeypatch, tmp_path, path)
    next(b for b in app.button if b.label == "No cover photograph").click().run()
    app.run()
    assert not app.exception
    assert cover_plan(app)["cover_picture_id"] == ""
    assert "cover_photo" not in cover_plan(app)["destinations"].values()
    assert not continue_button(app).disabled
    assert path.read_bytes() == raw
    assert library.list_profiles(synthetic_profiles()[0].contract) == ()
