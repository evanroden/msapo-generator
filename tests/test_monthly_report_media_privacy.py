"""Generated reports must not carry picture authoring metadata (Refs #296).

Every fixture is synthetic. Canary strings stand in for an author, a device, an
editing tool or a comment that a retained design/source picture could carry.
"""
from dataclasses import replace
from functools import lru_cache
from io import BytesIO
import struct
from zipfile import ZipFile

from docx import Document
from docx.shared import Inches
from PIL import Image, ImageCms, ImageOps
from PIL.PngImagePlugin import PngInfo
import pytest

from app.monthly_report_docx import generate_report
from app.monthly_report_media_privacy import sanitize_package_media, strip_image_metadata
from app.monthly_report_model import ReportPeriod, ResolvedBlock, synthetic_draft, synthetic_profiles
from tests.conftest import requires_libreoffice

ARTIST = "SYNTH-ARTIST-CANARY"
XP_AUTHOR = "SYNTH-XPAUTHOR-CANARY"
XMP_CREATOR = "SYNTH-XMP-CREATOR-CANARY"
IPTC_BYLINE = "SYNTH-IPTC-BYLINE-CANARY"
COMMENT = "SYNTH-COMMENT-CANARY"
TRAILER = "SYNTH-TRAILER-CANARY"
PNG_AUTHOR = "SYNTH-PNG-AUTHOR-CANARY"
PNG_ZTXT = "SYNTH-PNG-ZTXT-CANARY"
GIF_XMP = "SYNTH-GIF-XMP-CANARY"
CANARIES = (ARTIST, XP_AUTHOR, XMP_CREATOR, IPTC_BYLINE, COMMENT, TRAILER, PNG_AUTHOR, PNG_ZTXT, GIF_XMP)
XMP = ('<x:xmpmeta xmlns:x="adobe:ns:meta/"><rdf:RDF xmlns:rdf="http://www.w3.org/1999/02/22-rdf-syntax-ns#">'
       '<rdf:Description xmlns:dc="http://purl.org/dc/elements/1.1/"><dc:creator>' + XMP_CREATOR
       + '</dc:creator></rdf:Description></rdf:RDF></x:xmpmeta>')


def _carries_canary(data: bytes) -> list[str]:
    """Canaries found as ASCII/UTF-8, UTF-16LE, or inside zlib-compressed PNG text."""
    import zlib
    found = [c for c in CANARIES if c.encode() in data or c.encode("utf-16-le") in data]
    if data.startswith(b"\x89PNG"):
        pos = 8
        while pos + 12 <= len(data):
            length = struct.unpack(">I", data[pos:pos + 4])[0]
            if data[pos + 4:pos + 8] == b"zTXt":
                body = data[pos + 8:pos + 8 + length]
                text = zlib.decompress(body[body.index(b"\0") + 2:])
                found += [c for c in CANARIES if c.encode() in text]
            pos += 12 + length
    return found


def _segment(marker: int, payload: bytes) -> bytes:
    return bytes((0xFF, marker)) + struct.pack(">H", len(payload) + 2) + payload


@lru_cache(maxsize=1)
def _icc() -> bytes:
    # LittleCMS stamps each new profile with its creation time; build it once
    # so comparisons never depend on which second the test reached.
    return ImageCms.ImageCmsProfile(ImageCms.createProfile("sRGB")).tobytes()


def _picture(size=(48, 32)):
    image = Image.new("RGB", size, "white")
    for x in range(size[0]):
        for y in range(size[1]):
            image.putpixel((x, y), ((x * 5) % 256, (y * 7) % 256, (x * y) % 256))
    return image


def canary_jpeg(*, orientation=1, trailer=True) -> bytes:
    exif = Image.Exif()
    exif[315] = ARTIST
    exif[0x9C9D] = (XP_AUTHOR + "\0").encode("utf-16-le")
    if orientation != 1:
        exif[274] = orientation
    buffer = BytesIO()
    _picture().save(buffer, "JPEG", quality=90, exif=exif.tobytes(), icc_profile=_icc())
    raw = buffer.getvalue()
    extra = (_segment(0xE1, b"http://ns.adobe.com/xap/1.0/\x00" + XMP.encode())
             + _segment(0xED, b"Photoshop 3.0\x008BIM\x04\x04\x00\x00" + IPTC_BYLINE.encode())
             + _segment(0xFE, COMMENT.encode()))
    at = raw.index(b"\xff\xdb")  # after Pillow's own JFIF/EXIF/ICC segments, as in camera files
    raw = raw[:at] + extra + raw[at:]
    if trailer:  # e.g. an MPO/gain-map secondary picture after the primary EOI
        second = BytesIO()
        Image.new("RGB", (8, 8), "black").save(second, "JPEG")
        raw += second.getvalue()[:2] + _segment(0xFE, TRAILER.encode()) + second.getvalue()[2:]
    return raw


def canary_png() -> bytes:
    info = PngInfo()
    info.add_text("Author", PNG_AUTHOR)
    info.add_text("Comment", PNG_ZTXT, zip=True)
    info.add_itxt("XML:com.adobe.xmp", XMP)
    info.add(b"tIME", struct.pack(">HBBBBB", 2026, 10, 20, 12, 0, 0))
    exif = Image.Exif()
    exif[315] = ARTIST
    buffer = BytesIO()
    _picture().save(buffer, "PNG", pnginfo=info, icc_profile=_icc(), dpi=(150, 150), exif=exif.tobytes())
    return buffer.getvalue()


def canary_gif() -> bytes:
    buffer = BytesIO()
    _picture().convert("P").save(buffer, "GIF", comment=COMMENT.encode())
    raw = buffer.getvalue()
    flags = raw[10]
    start = 13 + (3 * (2 << (flags & 0x07)) if flags & 0x80 else 0)
    payload = (b"<x:xmpmeta>" + GIF_XMP.encode() + b"</x:xmpmeta>")
    blocks = b"".join(bytes((len(payload[i:i + 255]),)) + payload[i:i + 255] for i in range(0, len(payload), 255))
    xmp = b"\x21\xff\x0bXMP DataXMP" + blocks + b"\x00"
    return raw[:start] + xmp + raw[start:]


def _pixels(raw: bytes) -> tuple:
    with Image.open(BytesIO(raw)) as image:
        return image.size, image.mode, image.convert("RGBA").tobytes()


def test_jpeg_loses_authoring_metadata_but_keeps_pixels_profile_and_density():
    raw = canary_jpeg()
    assert set(_carries_canary(raw)) == {ARTIST, XP_AUTHOR, XMP_CREATOR, IPTC_BYLINE, COMMENT, TRAILER}
    clean = strip_image_metadata(raw)
    assert _carries_canary(clean) == []
    assert _pixels(clean) == _pixels(raw)
    with Image.open(BytesIO(raw)) as before, Image.open(BytesIO(clean)) as after:
        assert after.info.get("icc_profile") == before.info.get("icc_profile") == _icc()
        for key in ("jfif", "jfif_unit", "jfif_density"):
            assert after.info.get(key) == before.info.get(key)
        assert not after.getexif()
    assert clean.endswith(b"\xff\xd9") and len(clean) < len(raw)


@pytest.mark.parametrize("orientation", [3, 6, 8])
def test_jpeg_keeps_only_its_display_orientation(orientation):
    raw = canary_jpeg(orientation=orientation)
    clean = strip_image_metadata(raw)
    assert _carries_canary(clean) == []
    with Image.open(BytesIO(clean)) as after, Image.open(BytesIO(raw)) as before:
        assert dict(after.getexif()) == {274: orientation}
        assert ImageOps.exif_transpose(after).size == ImageOps.exif_transpose(before).size
        assert ImageOps.exif_transpose(after).tobytes() == ImageOps.exif_transpose(before).tobytes()


def test_png_loses_text_time_and_exif_but_keeps_profile_density_and_pixels():
    raw = canary_png()
    assert {ARTIST, PNG_AUTHOR, PNG_ZTXT, XMP_CREATOR} <= set(_carries_canary(raw))
    clean = strip_image_metadata(raw)
    assert _carries_canary(clean) == []
    assert b"tIME" not in clean and b"eXIf" not in clean
    assert _pixels(clean) == _pixels(raw)
    with Image.open(BytesIO(clean)) as after:
        assert after.info.get("icc_profile") == _icc()
        assert tuple(round(v) for v in after.info["dpi"]) == (150, 150)


def test_gif_loses_comment_and_xmp_but_keeps_pixels():
    raw = canary_gif()
    assert {COMMENT, GIF_XMP} <= set(_carries_canary(raw))
    clean = strip_image_metadata(raw)
    assert _carries_canary(clean) == []
    assert _pixels(clean) == _pixels(raw)


def test_webp_loses_exif_and_xmp_when_the_container_has_them():
    exif = Image.Exif()
    exif[315] = ARTIST
    buffer = BytesIO()
    try:
        _picture().save(buffer, "WEBP", lossless=True, exif=exif.tobytes(), xmp=XMP.encode())
    except (OSError, KeyError, ValueError):
        pytest.skip("This Pillow build cannot write WebP")
    raw = buffer.getvalue()
    if not {ARTIST, XMP_CREATOR} <= set(_carries_canary(raw)):
        pytest.skip("This Pillow build does not write WebP EXIF/XMP chunks")
    clean = strip_image_metadata(raw)
    assert _carries_canary(clean) == []
    assert _pixels(clean) == _pixels(raw)
    assert struct.unpack("<I", clean[4:8])[0] == len(clean) - 8


def test_clean_pictures_and_other_data_are_returned_unchanged():
    plain = BytesIO()
    _picture().save(plain, "JPEG", quality=90)
    assert strip_image_metadata(plain.getvalue()) == plain.getvalue()
    for data in (b"<w:document/>", b"", b"\xff\xd8truncated", b"\x89PNG\r\n\x1a\nshort", b"GIF89a"):
        assert strip_image_metadata(data) == data


def _package_with_pictures() -> bytes:
    """A Word package whose pictures were kept byte-for-byte from a source design."""
    document = Document()
    document.add_paragraph("Synthetic report body")
    for fmt, width in (("JPEG", 1), ("PNG", 1), ("GIF", 0.5)):
        plain = BytesIO()
        _picture().convert("P" if fmt == "GIF" else "RGB").save(plain, fmt)
        document.add_picture(BytesIO(plain.getvalue()), width=Inches(width))
    buffer = BytesIO()
    document.save(buffer)
    retained = {".jpeg": canary_jpeg(), ".jpg": canary_jpeg(), ".png": canary_png(), ".gif": canary_gif()}
    output = BytesIO()
    with ZipFile(BytesIO(buffer.getvalue())) as source, ZipFile(output, "w") as target:
        for info in source.infolist():
            data = source.read(info)
            if info.filename.startswith("word/media/"):
                data = retained[info.filename[info.filename.rindex("."):]]
            target.writestr(info, data)
    return output.getvalue()


def test_package_sanitizer_changes_only_picture_metadata():
    package = _package_with_pictures()
    clean = sanitize_package_media(package)
    with ZipFile(BytesIO(package)) as before, ZipFile(BytesIO(clean)) as after:
        assert after.namelist() == before.namelist()
        pictures = 0
        for info in before.infolist():
            old, new = before.read(info), after.read(info.filename)
            if info.filename.startswith("word/media/"):
                pictures += 1
                assert _carries_canary(old)
            if old.startswith((b"\xff\xd8", b"\x89PNG", b"GIF8")):
                # python-docx's stock docProps thumbnail is a picture too.
                assert _carries_canary(new) == []
                assert _pixels(new) == _pixels(old)
            else:
                assert new == old, info.filename
            assert after.getinfo(info.filename).compress_type == info.compress_type
            assert after.getinfo(info.filename).date_time == info.date_time
        assert pictures == 3
    assert Document(BytesIO(clean)).inline_shapes
    assert sanitize_package_media(clean) == clean


def test_bytes_that_are_not_a_package_pass_through_unchanged():
    for data in (b"native-download", b"", b"PK\x03\x04truncated"):
        assert sanitize_package_media(data) == data


def _draft_with_canary_logo():
    draft = synthetic_draft(synthetic_profiles()[0], ReportPeriod(2026, 9))
    from app.monthly_report_asset_review import approve_all_assets
    logo = approve_all_assets(ResolvedBlock("brand_logo", "Library", asset_hashes=("synthetic-canary-logo",)))
    blocks = tuple(b for b in draft.blocks if b.key != "brand_logo") + (logo,)
    return replace(draft, blocks=blocks)


def _loader(reference):
    assert reference == "synthetic-canary-logo"
    return canary_jpeg()


def test_generated_word_file_carries_no_logo_metadata(monkeypatch, tmp_path):
    from app import pdf_converter
    requested=[]
    def no_pdf(path, *, lossless_images):
        requested.append(lossless_images)
        raise pdf_converter.PDFConversionError("Synthetic converter unavailable")

    monkeypatch.setattr(pdf_converter, "OUTPUT_DIR", tmp_path)
    monkeypatch.setattr(pdf_converter, "convert_to_pdf", no_pdf)
    package = generate_report(_draft_with_canary_logo(), asset_loader=_loader)
    assert requested==[True]
    with ZipFile(BytesIO(package.docx)) as archive:
        media = [name for name in archive.namelist() if name.startswith("word/media/")]
        assert media, "the synthetic logo must reach the generated document"
        assert [(name, _carries_canary(archive.read(name))) for name in media
                if _carries_canary(archive.read(name))] == []


def _pdf_findings(pdf: bytes) -> list:
    """Canaries in any embedded PDF image (decoded stream bytes) or PDF metadata."""
    import fitz
    found = []
    with fitz.open(stream=pdf, filetype="pdf") as document:
        images = {item[0] for page in document for item in page.get_images(full=True)}
        assert images, "the synthetic logo must reach the generated PDF"
        for xref in images:
            found += _carries_canary(document.xref_stream_raw(xref) or b"")
            found += _carries_canary(document.extract_image(xref)["image"])
        found += _carries_canary(repr(document.metadata).encode())
        found += _carries_canary((document.get_xml_metadata() or "").encode())
    return found


@requires_libreoffice
def test_generated_word_and_pdf_carry_no_logo_metadata():
    package = generate_report(_draft_with_canary_logo(), asset_loader=_loader)
    assert package.pdf, package.pdf_error
    with ZipFile(BytesIO(package.docx)) as archive:
        assert [name for name in archive.namelist() if _carries_canary(archive.read(name))] == []
    assert _pdf_findings(package.pdf) == []
