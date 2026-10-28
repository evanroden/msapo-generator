"""Synthetic capital master: reviewed contract/site reuse with pinned report history."""
from dataclasses import asdict, replace
from io import BytesIO
import json

from docx import Document
import pytest

from app import monthly_report_library as library
from app.monthly_report_docx import assemble_docx
from app.monthly_report_model import (
    Facility, ReportDraft, ReportPeriod, ReportProfile, ResolvedBlock,
    default_sections,
)
from app.monthly_report_setup import new_month_draft
from app.monthly_report_capital_master import (
    CapitalItem, CapitalMaster, load_capital_master, save_capital_master,
    select_capital, capital_report_block,
)

SEPTEMBER = ReportPeriod(2026, 9)
OCTOBER = ReportPeriod(2026, 10)


@pytest.fixture
def fixture(monkeypatch, tmp_path):
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path))
    north = Facility("north", "Synthetic North", ("North",))
    south = Facility("south", "Synthetic South", ("South",))
    both = ReportProfile("Synthetic Capital Contract A", "both", "Both sites",
                         (north, south), scope_type="multi_site")
    n = ReportProfile(both.contract, "north", "Site North", (north,))
    s = ReportProfile(both.contract, "south", "Site South", (south,))
    other = ReportProfile("Synthetic Capital Contract B", "north", "Different customer",
                          (Facility("north", "Other North"),))
    return both, n, s, other


def capital():
    return (
        CapitalItem(
            "primary-pump", "High", "Replace aging chilled-water pump",
            "PUMP-01", "Pending", "2026-09", "2028-08-15",
            source_reference="synthetic-source:1:verified",
            extra_fields=(("Condition", "Observed vibration"), ("Owner", "Plant team"))),
        CapitalItem(
            "south-pump", "Critical", "Schedule southern pump replacement",
            "PUMP-01", "Deferred", "2026-08", "",
            site_keys=("south",), override_of="primary-pump",
            source_reference="synthetic-source:2:verified"),
        CapitalItem(
            "north-boiler", "Medium", "Plan east boiler renewal",
            "BOILER-9", "", "", "", site_keys=("north",),
            extra_fields=(("Notes", "NR"),), source_reference="synthetic-source:3:verified"),
    )


def approved(profile, items=None, *, revision=0, removed=()):
    return save_capital_master(
        profile, items if items is not None else capital(),
        expected_revision=revision, removed_ids=removed,
        actor="Synthetic reviewer", confirmed=True)


def report(profile, block, period=SEPTEMBER):
    section = next(s for s in default_sections() if s.key == "capital")
    return ReportDraft(profile, period, "Synthetic reviewer", (section,),
                       (block,), synthetic=True)


def test_two_sites_see_selected_current_default_with_site_override_and_no_duplicate(fixture):
    both, n, s, other = fixture
    master = approved(both)
    assert master.revision == 1
    assert load_capital_master(both.contract) == master
    combined = select_capital(both, master, SEPTEMBER)
    assert [(r.site_key, r.item.key) for r in combined.occurrences] == [
        ("north", "primary-pump"), ("north", "north-boiler"),
        ("south", "south-pump")]
    assert combined.historical_statuses == (
        ("Synthetic South", "south-pump", "2026-08", "Deferred"),)
    rows = capital_report_block(both, combined)
    assert rows.rows == (
        ("Synthetic North", "High", "Replace aging chilled-water pump"),
        ("Synthetic North", "Medium", "Plan east boiler renewal"),
        ("Synthetic South", "Critical", "Schedule southern pump replacement"),
    )
    assert rows.source == "Library"
    assert rows.references == (master.reference,)
    assert len(rows.extra_tables) == 1
    detail_rows = rows.extra_tables[0].rows
    assert ("Synthetic North", "primary-pump", "Condition", "Observed vibration") in detail_rows
    assert ("Synthetic South", "south-pump", "Recorded status", "Deferred") in detail_rows
    assert ("Synthetic South", "south-pump", "Status recorded month", "2026-08") in detail_rows
    assert not any(r.site_title == "Synthetic South" and r.item.key == "primary-pump"
                   for r in combined.occurrences)
    assert [(r.item.key, r.item.asset) for r in select_capital(n, master, SEPTEMBER).occurrences] == [
        ("primary-pump", "PUMP-01"), ("north-boiler", "BOILER-9")]
    assert [r.item.key for r in select_capital(s, master, SEPTEMBER).occurrences] == [
        "south-pump"]
    assert load_capital_master(other.contract).items == ()


def test_cross_contract_content_isolation_and_distinct_revision_namespaces(fixture):
    both, _, _, other = fixture
    a = approved(both)
    assert load_capital_master(other.contract).revision == 0
    with pytest.raises(ValueError, match="cross contract"):
        select_capital(other, a, SEPTEMBER)
    b = approved(other, (CapitalItem("other-only", "High", "Different site asset",
                                      "AHU-2", site_keys=("north",)),))
    assert b.revision == 1
    assert b.reference != a.reference
    assert tuple(item.key for item in load_capital_master(both.contract).items) == (
        "primary-pump", "south-pump", "north-boiler")
    assert tuple(item.key for item in load_capital_master(other.contract).items) == (
        "other-only",)


def test_explicit_confirm_expected_revision_and_immutable_history(fixture):
    both, n, _, _ = fixture
    with pytest.raises(library.LibraryError, match="Confirm"):
        save_capital_master(both, capital(), expected_revision=0,
                            actor="Editor", confirmed=False)
    assert load_capital_master(both.contract).revision == 0
    first = approved(both)
    changed = tuple(replace(item, recommendation="Replace pump after inspection")
                    if item.key == "primary-pump" else item for item in first.items)
    saved = approved(both, changed, revision=1)
    assert saved.revision == 2
    assert saved.actor == "Synthetic reviewer"
    assert load_capital_master(both.contract, 1) == first
    assert load_capital_master(both.contract, 2) == saved
    with pytest.raises(library.RevisionConflict, match="changed in another view"):
        approved(both, capital(), revision=1)
    assert approved(both, changed, revision=2) == saved
    assert load_capital_master(both.contract, 1).items[0].recommendation == (
        "Replace aging chilled-water pump")
    assert load_capital_master(both.contract).revision == 2


def test_report_only_edits_omission_and_history_never_advance_contract_master(fixture):
    both, n, _, _ = fixture
    master = approved(both)
    selection = select_capital(n, master, SEPTEMBER)
    block = capital_report_block(n, selection)
    modified = replace(block, rows=(
        ("Synthetic North", "Urgent", "Report-only pump correction"),
        *block.rows[1:]), text="Manual report-only context")
    with pytest.raises(ValueError, match="report-only capital edits"):
        capital_report_block(n, selection, existing=modified)
    with pytest.raises(ValueError, match="report-only capital edits"):
        capital_report_block(n, selection, existing=replace(block, source="Omit"))
    assert capital_report_block(n, selection, existing=block) is block
    old = library.save_snapshot(report(n, modified), expected_revision=0)
    current = library.save_snapshot(report(n, block), expected_revision=1)
    assert current.revision == 2
    assert library.load_snapshot_version(n.contract, n.key, SEPTEMBER, 1).draft.blocks[0] == modified
    assert library.load_snapshot(n.contract, n.key, SEPTEMBER).draft.blocks[0] == block
    assert load_capital_master(both.contract) == master
    with pytest.raises(library.RevisionConflict):
        library.save_snapshot(report(n, modified), expected_revision=1)
    assert old.draft.blocks[0].rows[0][1] == "Urgent"


def test_report_omit_is_exact_site_row_and_does_not_republish_next_month(fixture):
    both, n, s, _ = fixture
    master = approved(both)
    only_pump = select_capital(
        n, master, SEPTEMBER, omitted=(("north", "north-boiler"),))
    assert [r.item.key for r in only_pump.occurrences] == ["primary-pump"]
    with pytest.raises(ValueError, match="no longer selected"):
        select_capital(n, master, SEPTEMBER, omitted=(("north", "south-pump"),))
    with pytest.raises(ValueError, match="no longer selected"):
        select_capital(s, master, SEPTEMBER, omitted=(("south", "primary-pump"),))
    block = capital_report_block(n, only_pump)
    next_month = new_month_draft(report(n, block), OCTOBER)
    assert next_month.blocks[0] == block
    assert next_month.blocks[0].rows == (
        ("Synthetic North", "High", "Replace aging chilled-water pump"),)
    following = new_month_draft(next_month, ReportPeriod(2026, 11))
    assert following.blocks[0].rows == next_month.blocks[0].rows
    assert following.blocks[0].references == (master.reference,)
    assert load_capital_master(both.contract).items == capital()


def test_approved_retirement_needs_exact_row_and_preserves_untouched_sites(fixture):
    both, n, _, _ = fixture
    initial = approved(both)
    less = tuple(item for item in initial.items if item.key != "north-boiler")
    with pytest.raises(ValueError, match="explicit removal"):
        approved(both, less, revision=1)
    assert load_capital_master(both.contract) == initial
    with pytest.raises(ValueError, match="another site's capital item"):
        approved(n, tuple(item for item in initial.items
                          if item.key != "south-pump"), revision=1,
                 removed=("south-pump",))
    assert load_capital_master(both.contract) == initial
    saved = approved(n, less, revision=1, removed=("north-boiler",))
    assert saved.revision == 2
    assert [item.key for item in saved.items] == ["primary-pump", "south-pump"]
    assert len(load_capital_master(both.contract, 1).items) == 3


def test_foreign_new_site_and_override_without_parent_are_rejected(fixture):
    both, _, _, _ = fixture
    with pytest.raises(ValueError, match="confirmed contract membership"):
        approved(both, (CapitalItem("foreign", "High", "For another site",
                                    site_keys=("outside",)),))
    with pytest.raises(ValueError, match="contract-wide parent"):
        approved(both, (CapitalItem("orphan", "High", "No parent",
                                   site_keys=("north",), override_of="missing"),))
    with pytest.raises(ValueError, match="same site"):
        approved(both, (*capital(), replace(capital()[1], key="south-dupe")))
    assert load_capital_master(both.contract).items == ()


def test_missing_dates_and_unknown_columns_are_preserved_without_inference(fixture):
    both, n, _, _ = fixture
    master = approved(both)
    first = master.items[0]
    assert first.target_date == "2028-08-15"
    assert first.extra_fields == (("Condition", "Observed vibration"),
                                   ("Owner", "Plant team"))
    last = master.items[-1]
    assert last.status == last.status_period == last.target_date == ""
    assert last.extra_fields == (("Notes", "NR"),)
    selected = select_capital(n, master, OCTOBER)
    assert selected.historical_statuses == ((
        "Synthetic North", "primary-pump", "2026-09", "Pending"),)
    block = capital_report_block(n, selected)
    assert ("Synthetic North", "primary-pump", "Target date", "2028-08-15") in block.extra_tables[0].rows
    assert ("Synthetic North", "north-boiler", "Notes", "NR") in block.extra_tables[0].rows
    assert not any(item[2] == "Target date" and item[1] == "north-boiler"
                   for item in block.extra_tables[0].rows)
    assert load_capital_master(both.contract) == master


def test_financial_fields_in_master_are_blocked_before_report_inclusion(fixture):
    both, _, _, _ = fixture
    with pytest.raises(ValueError, match="Financial source columns"):
        approved(both, (CapitalItem(
            "priced", "High", "Replace pump", extra_fields=(("Amount Due", "500"),)),))
    with pytest.raises(ValueError, match="financial content"):
        approved(both, (CapitalItem(
            "priced", "High", "Replace pump for $5,000"),))
    assert load_capital_master(both.contract).revision == 0


@pytest.mark.parametrize("invalid", ("2026-02-30", "09/15/2026", "2026-13-01"))
def test_ambiguous_or_nonexistent_target_date_never_becomes_approved_fact(fixture, invalid):
    both, _, _, _ = fixture
    with pytest.raises(ValueError, match="target date"):
        approved(both, (CapitalItem("one", "High", "New pump",
                                   target_date=invalid),))
    assert load_capital_master(both.contract).revision == 0


def test_saved_report_can_pin_revision_one_even_after_shared_revision_two(fixture):
    both, n, _, _ = fixture
    original = approved(both)
    first = capital_report_block(n, select_capital(n, original, SEPTEMBER))
    report1 = report(n, first)
    library.save_snapshot(report1, expected_revision=0)
    revised = tuple(replace(row, priority="Emergency")
                    if row.key == "primary-pump" else row for row in original.items)
    second = approved(both, revised, revision=1)
    assert second.revision == 2
    pinned = capital_report_block(n, select_capital(n, original, SEPTEMBER))
    assert pinned == first
    assert pinned.rows[0][1] == "High"
    fresh = capital_report_block(n, select_capital(n, second, SEPTEMBER))
    assert fresh.rows[0][1] == "Emergency"
    assert fresh.references != first.references
    assert library.load_snapshot(n.contract, n.key, SEPTEMBER).draft.blocks[0] == first
    assert load_capital_master(both.contract, 1) == original


def test_generic_word_section_nine_prints_approved_rows_and_details_once(fixture):
    both, n, _, _ = fixture
    master = approved(both)
    block = capital_report_block(n, select_capital(n, master, SEPTEMBER))
    docx = assemble_docx(report(n, block),
                         acknowledged_fingerprint=report(n, block).fingerprint)
    doc = Document(BytesIO(docx))
    assert len(doc.tables) == 2
    main, details = doc.tables
    assert [cell.text for cell in main.rows[0].cells] == (
        ["Facility", "Priority", "Recommendation"])
    assert [tuple(cell.text for cell in row.cells)
            for row in main.rows[1:]] == list(block.rows)
    assert [cell.text for cell in details.rows[0].cells] == [
        "Facility", "Capital ID", "Field", "Value"]
    assert details.rows[1].cells[1].text == "primary-pump"
    assert any(cell.text == "2028-08-15"
               for row in details.rows for cell in row.cells)
    assert not any("synthetic-source:" in cell.text
                   for table in doc.tables for row in table.rows for cell in row.cells)
    future = new_month_draft(report(n, block), OCTOBER)
    assert future.blocks[0].rows == block.rows
    again = assemble_docx(future, acknowledged_fingerprint=future.fingerprint)
    doc2 = Document(BytesIO(again))
    assert len(doc2.tables) == 2
    assert [tuple(c.text for c in row.cells) for row in doc2.tables[0].rows[1:]] == list(block.rows)


def test_unapproved_or_forged_master_cannot_publish_a_report(fixture):
    both, n, _, _ = fixture
    from app.monthly_report_capital_master import CapitalSelection
    forged = CapitalMaster(both.contract, 1, capital())
    with pytest.raises((ValueError, library.LibraryError)):
        select_capital(n, forged, SEPTEMBER)
    master = approved(both)
    proper = select_capital(n, master, SEPTEMBER)
    altered = replace(proper, occurrences=tuple(
        replace(entry, item=replace(entry.item, recommendation="Forged report"))
        if entry.item.key == "primary-pump" else entry
        for entry in proper.occurrences))
    with pytest.raises(ValueError, match="differs from"):
        capital_report_block(n, altered)
    with pytest.raises(ValueError, match="not valid"):
        capital_report_block(n, replace(proper, master_reference="capital-master:wrong"))
    assert capital_report_block(n, proper).rows[0][2] == (
        "Replace aging chilled-water pump")



def test_only_current_site_can_change_an_approved_site_specific_item(fixture):
    both, n, s, _ = fixture
    original = approved(both)
    changed = tuple(
        replace(item, recommendation="Unauthorized South correction")
        if item.key == "south-pump" else item for item in original.items)
    with pytest.raises(ValueError, match="another site's capital entry"):
        approved(n, changed, revision=1)
    assert load_capital_master(both.contract) == original
    permitted = approved(s, changed, revision=1)
    assert permitted.revision == 2
    assert load_capital_master(both.contract, 1) == original
    assert permitted.items[0] == original.items[0]


def test_status_and_source_month_must_be_reviewed_together(fixture):
    both, _, _, _ = fixture
    with pytest.raises(ValueError, match="status and its recorded month"):
        approved(both, (CapitalItem(
            "incomplete", "High", "Pump replacement",
            status="Completed", status_period=""),))
    with pytest.raises(ValueError, match="status and its recorded month"):
        approved(both, (CapitalItem(
            "incomplete", "High", "Pump replacement",
            status="", status_period="2026-09"),))
    assert load_capital_master(both.contract).revision == 0


def test_archived_source_fields_and_zero_vs_blank_are_not_inferred(fixture):
    both, n, _, _ = fixture
    original = CapitalItem(
        "legacy", "1", "Replace motor after inspection", "MTR-01",
        extra_fields=(("Observed failures", "0"), ("Target quarter", "")),
        source_reference="synthetic-original:4:hash")
    master = approved(both, (original,))
    assert master.items[0].extra_fields == (
        ("Observed failures", "0"), ("Target quarter", ""))
    selection = select_capital(n, master, SEPTEMBER)
    block = capital_report_block(n, selection)
    assert ("Synthetic North", "legacy", "Observed failures", "0") in block.extra_tables[0].rows
    # A truly blank source cell remains in the approved master; it is never
    # filled as a fabricated estimate or copied into the Word report as 0.
    assert master.items[0].extra_fields[-1] == ("Target quarter", "")
    assert not any(row[2] == "Target quarter" for row in block.extra_tables[0].rows)
    assert block.references == (master.reference,)
    assert load_capital_master(both.contract) == master
