"""CMMS service-call source mapping: synthetic CSV, Word and Word-output safety."""
from dataclasses import asdict, replace
from io import BytesIO
import json

from docx import Document
import pytest

from app import monthly_report_sources as sources
from app.monthly_report_model import (
    ReportPeriod, ResolvedBlock, synthetic_draft, synthetic_profiles,
)
from app.monthly_report_service_call_mapping import (
    confirm_financial_exclusion, confirm_word_layout,
    correct_service_call, exclude_service_call, exclude_unmatched_call, reimport_service_calls,
    restore_service_call_review, review_service_calls, reviewed_service_call_block,
    service_call_review_record, suggest_call_mapping,
)
from app.monthly_report_service_call_word import review_word_maintenance, reviewed_word_maintenance_block

PROFILES = synthetic_profiles()
PROFILE = PROFILES[1]  # two different synthetic facilities, one contract
PERIOD = ReportPeriod(2026, 9)


@pytest.fixture
def storage(monkeypatch, tmp_path):
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path))
    return tmp_path


def read_csv(name, raw, profile=PROFILE):
    original, _ = sources.ingest(profile, name, raw)
    assert sources.source_bytes(profile, original.source) == raw
    return original


def word_fixture(*, two_tables=False, preamble="Service Calls"):
    document = Document()
    document.add_heading("1 Organizational Chart", level=1)
    extra = document.add_table(rows=2, cols=2)
    for c, v in zip(extra.rows[0].cells, ("Name", "Role")):
        c.text = v
    for c, v in zip(extra.rows[1].cells, ("Synthetic person", "Director")):
        c.text = v
    document.add_heading("5 Maintenance Schedule / In-House Maintenance", level=1)
    columns = ("WO #", "Finish Date", "Area", "Task Code-Description",
               "Tag #", "Facility", "Status")
    def table(rows):
        target = document.add_table(rows=len(rows) + 2, cols=len(columns))
        target.rows[0].cells[0].merge(target.rows[0].cells[-1]).text = preamble
        for cell, text in zip(target.rows[1].cells, columns):
            cell.text = text
        for j, row in enumerate(rows):
            for cell, text in zip(target.rows[j + 2].cells, row):
                cell.text = text
    table((("0009", "2026-09-05", "Boiler Room", "Replace seal",
            "PUMP-09", "Demo North", "Closed"),
           ("0010", "", "North Wing", "Inspect AHU",
            "AHU-02", "Demo North", "Open")))
    if two_tables:
        table((("0009", "2026-09-07", "Boiler Room", "Second entry",
                "PUMP-09", "Demo North", "Completed"),))
    document.add_heading("6 Sub-Contractor Status", level=1)
    return document


def read_word(storage, document):
    path = storage / "synthetic-maintenance.docx"
    document.save(path)
    raw = path.read_bytes()
    content, _ = sources.ingest(PROFILE, path.name, raw)
    assert sources.source_bytes(PROFILE, content.source) == raw
    return content, raw


def test_two_csv_layouts_map_to_same_five_editable_fields(storage):
    first = read_csv("maintenance.csv", (
        b"WO #,Finish Date,Area,Task Code-Description,Tag #,Site,Status,Hours,Note\n"
        b"00001,2026-09-05,Plant A,Inspect pump,PMP-01,Demo North,Closed,0,\n"
        b"00002,2026-09-06,Plant B,Repair valve,VLV-02,Demo North,Completed,,Cleared\n"
    ))
    second = read_csv("maintenance-semicolon.csv", (
        b"Ticket Number;Completion Date;Room;Work Description;Asset ID;Facility;WO Status;Hours;Note\n"
        b"00001;2026-09-05;Plant A;Inspect pump;PMP-01;Demo North;Closed;0;\n"
        b"00002;2026-09-06;Plant B;Repair valve;VLV-02;Demo North;Completed;;Cleared\n"
    ))
    a = review_service_calls(first, PROFILE, PERIOD)
    b = review_service_calls(second, PROFILE, PERIOD)
    assert a.ready_for_inclusion and b.ready_for_inclusion
    assert tuple((v.wo, v.finished, v.area, v.task, v.tag, v.site, v.status)
                 for v in a.rows) == tuple((v.wo, v.finished, v.area,
                 v.task, v.tag, v.site, v.status) for v in b.rows)
    assert a.counts() == {
        "original": 2, "reviewed_candidates": 2, "ready": 2,
        "unresolved": 0, "excluded": 0, "unmatched": 0}
    assert a.rows[0].wo == "00001" and a.rows[1].wo == "00002"
    assert a.unmapped(a.rows[0]) == ((7, "Hours", "0"), (8, "Note", ""))
    assert a.unmapped(a.rows[1]) == ((7, "Hours", ""), (8, "Note", "Cleared"))
    block = reviewed_service_call_block(a)
    assert block.rows == (
        ("00001", "2026-09-05", "Demonstration North Facility — Plant A", "Inspect pump", "PMP-01"),
        ("00002", "2026-09-06", "Demonstration North Facility — Plant B", "Repair valve", "VLV-02"),
    )
    assert block.references == (sources.source_reference(first.source, 1),)
    assert sources.source_bytes(PROFILE, first.source).startswith(b"WO #")
    assert reviewed_service_call_block(a, existing=block) == block


def test_unknown_finish_status_and_site_are_not_invented(storage):
    raw = (b"WO Number,Date,Location,Description,Asset Tag,Facility,Status\n"
           b"00003,09/05/2026,North Wing,Service fan,TAG-03,North / South,Scheduled\n"
           b"00004,,Mechanical,Inspect filter,TAG-04,Demo North,Open\n")
    content = read_csv("not-complete.csv", raw)
    review = review_service_calls(content, PROFILE, PERIOD)
    assert not review.ready_for_inclusion
    assert len(review.rows) == 2
    assert not review.rows[0].finished and not review.rows[1].finished
    assert "Site scope needs review." in review.rows[0].issues
    assert "Completion status is not verified." in review.rows[0].issues
    assert review.rows[0].origin.cells[1] == "09/05/2026"
    assert review.rows[1].origin.cells[1] == ""
    with pytest.raises(ValueError, match="Review unresolved"):
        reviewed_service_call_block(review)
    assert sources.source_bytes(PROFILE, content.source) == raw


def test_explicit_corrections_reimport_duplicate_protection_and_safe_omission(storage):
    data = read_csv("partial.csv", (
        b"Work Order Number,Finish Date,Area,Task,Tag,Site,Status\n"
        b"00015,2026-09-05,Plant,Test pump,PMP-1,Demo North,Closed\n"
        b"00015,2026-09-06,Plant,Repeat record,PMP-1,Demo North,Closed\n"
        b"00016,2026-09-07,Boilers,Complete boiler inspection,B-2,Demo North,\n"
    ))
    review = review_service_calls(data, PROFILE, PERIOD)
    assert review.counts()["ready"] == 0
    assert all("Duplicate work-order identity." in r.issues for r in review.rows[:2])
    corrected = correct_service_call(review, review.rows[1].key,
        {"wo": "00017"}, expected_revision=review.revision)
    assert corrected.counts()["ready"] == 2
    corrected = correct_service_call(corrected, review.rows[2].key,
        {"status": "Completed"}, expected_revision=corrected.revision)
    assert corrected.ready_for_inclusion
    assert corrected.counts()["ready"] == 3
    assert corrected.rows[1].origin.cells[0] == "00015"
    assert ("wo", "00017") in corrected.rows[1].corrections
    assert reimport_service_calls(corrected, review) == corrected
    report = reviewed_service_call_block(corrected)
    assert [row[0] for row in report.rows] == ["00015", "00017", "00016"]
    omitted = exclude_service_call(corrected, corrected.rows[0].key,
        expected_revision=corrected.revision)
    assert omitted.counts()["excluded"] == 1
    assert len(reviewed_service_call_block(omitted).rows) == 2
    with pytest.raises(ValueError, match="changed"):
        correct_service_call(omitted, review.rows[1].key,
            {"task": "Stale action"}, expected_revision=review.revision)
    with pytest.raises(ValueError, match="already has edited"):
        reviewed_service_call_block(corrected, existing=replace(
            report, rows=(("MANUAL", "2026-09-02", "X", "Custom change", "Y"),)))
    newer = read_csv("changed.csv", (
        b"Work Order Number,Finish Date,Area,Task,Tag,Site,Status\n"
        b"00015,2026-09-05,Plant,Test pump,PMP-1,Demo North,Closed\n"))
    with pytest.raises(ValueError, match="changed"):
        reimport_service_calls(corrected, review_service_calls(newer, PROFILE, PERIOD))


def test_financial_headers_require_specific_exclusion_review(storage):
    data = read_csv("with-costs.csv", (
        b"WO #,Finish Date,Area,Task,Tag #,Site,Status,Amount Due\n"
        b"00004,2026-09-08,Plant,Test pressure,P-4,Demo North,Completed,250.00\n"))
    review = review_service_calls(data, PROFILE, PERIOD)
    assert review.financial_columns == (7,)
    assert not review.ready_for_inclusion
    with pytest.raises(ValueError, match="financial"):
        reviewed_service_call_block(review)
    accepted = confirm_financial_exclusion(review, expected_revision=review.revision)
    assert accepted.ready_for_inclusion
    output = reviewed_service_call_block(accepted)
    assert output.rows == (("00004", "2026-09-08",
                            "Demonstration North Facility — Plant", "Test pressure", "P-4"),)
    assert not any("250" in value for row in output.rows for value in row)
    assert review.rows[0].origin.cells[-1] == "250.00"
    with pytest.raises(ValueError, match="changed"):
        confirm_financial_exclusion(accepted, expected_revision=review.revision)


def test_mapping_duplicates_and_extra_columns_keep_physical_values(storage):
    content = read_csv("ambiguous.csv", (
        b"WO #,WO Number,Finish Date,Area,Task Code-Description,Tag #,Site,Status,Notes,Notes\n"
        b"00011,00099,2026-09-10,Boiler Room,Repair pump,P-9,Demo North,Closed,0,\n"))
    rows = content.tables[0]
    assert rows.columns[-1] == "Notes (2)"
    mapping = suggest_call_mapping(rows.columns)
    assert "wo" in mapping.ambiguous and mapping.wo is None
    pending = review_service_calls(content, PROFILE, PERIOD)
    assert pending.notes and not pending.ready_for_inclusion
    explicit = suggest_call_mapping(rows.columns, {"wo": 0})
    corrected = review_service_calls(content, PROFILE, PERIOD, mapping=explicit)
    assert corrected.ready_for_inclusion
    assert corrected.unmapped(corrected.rows[0]) == (
        (1, "WO Number", "00099"), (8, "Notes", "0"), (9, "Notes (2)", ""))
    with pytest.raises(ValueError, match="two service-call"):
        suggest_call_mapping(rows.columns, {"wo": 0, "task": 0})
    with pytest.raises(ValueError, match="physical"):
        suggest_call_mapping(rows.columns, {"wo": 99})


def test_word_import_keeps_title_and_partial_rows_without_false_completion(storage):
    source, raw = read_word(storage, word_fixture())
    result = review_word_maintenance(PROFILE, source, PERIOD)
    assert len(result.tables) == 1 and len(result.ignored_item_ids) == 1
    table, = result.tables
    assert table.header_row_number == 2
    assert len(table.title_rows) == 1
    assert table.title_rows[0].cells[0] == "Service Calls"
    assert table.review is not None and not table.review.native_layout_confirmed
    assert table.review.rows[0].wo == "0009"
    assert table.review.rows[0].origin.row_number == 3
    assert table.review.rows[0].origin.table_number == table.item_number
    assert table.review.rows[1].wo == "0010"
    assert table.review.rows[1].status == "Unfinished"
    assert "Finish date missing" in "; ".join(table.review.rows[1].issues)
    assert sources.source_bytes(PROFILE, source.source) == raw
    # Missing completed fields remain candidates; removing that one call is
    # deliberate and leaves the source row accessible.
    excluded = exclude_service_call(table.review, table.review.rows[1].key,
        expected_revision=table.review.revision)
    accepted = confirm_word_layout(excluded, expected_revision=excluded.revision)
    assert accepted.ready_for_inclusion
    assert reviewed_service_call_block(accepted).rows == (
        ("0009", "2026-09-05", "Demonstration North Facility — Boiler Room", "Replace seal", "PUMP-09"),)
    assert sources.source_bytes(PROFILE, source.source) == raw


def test_word_duplicate_across_tables_and_unsupported_source_are_rejected(storage):
    source, raw = read_word(storage, word_fixture(two_tables=True))
    result = review_word_maintenance(PROFILE, source, PERIOD)
    assert len(result.tables) == 2
    assert result.cross_table_duplicates == (("north", "0009"),)
    assert all(not table.review.ready_for_inclusion for table in result.tables)
    assert sources.source_bytes(PROFILE, source.source) == raw
    wrong = read_csv("not-word.csv", b"WO,Finish Date,Task\n0001,2026-09-01,Pump\n")
    with pytest.raises(ValueError, match="Word"):
        review_word_maintenance(PROFILE, wrong, PERIOD)


def test_json_review_roundtrip_preserves_manual_changes_and_source_identity(storage):
    content = read_csv("review.csv", (
        b"WO #,Finish Date,Area,Task,Tag,Site,Status,Notes\n"
        b"00023,2026-09-17,Basement,Oil check,BOILER-3,Demo North,Closed,0\n"))
    base = review_service_calls(content, PROFILE, PERIOD)
    updated = correct_service_call(base, base.rows[0].key, {
        "task": "Synthetic corrected oil check"}, expected_revision=base.revision)
    record = json.loads(json.dumps(service_call_review_record(updated)))
    assert restore_service_call_review(record) == updated
    assert record["review"]["rows"][0]["origin"]["cells"][-1] == "0"
    record["review"]["rows"][0]["origin"]["cells"][0] = "000999"
    with pytest.raises(ValueError, match="integrity"):
        restore_service_call_review(record)


def test_service_calls_render_in_existing_word_schema_with_source_receipts(storage):
    from app.monthly_report_docx import assemble_docx
    from app.monthly_report_model import default_sections
    profile = PROFILES[0]
    data = read_csv("single-site-service.csv", (
        b"WO #,Finish Date,Area,Task,Tag,Site,Status\n"
        b"00031,2026-09-15,Plant,Inspect boiler,B-31,Demo North,Closed\n"), profile=profile)
    tagged_source = replace(data.source, facility="Demo North")
    data = replace(data, source=tagged_source)
    review = review_service_calls(data, profile, PERIOD)
    assert review.ready_for_inclusion
    block = reviewed_service_call_block(review)
    base = synthetic_draft(profile, PERIOD)
    section = next(s for s in default_sections() if s.key == "maintenance")
    draft = replace(base, sections=(section,), blocks=(block,), sources=(tagged_source,))
    output = assemble_docx(draft, acknowledged_fingerprint=draft.fingerprint)
    parsed = Document(BytesIO(output))
    text = "\n".join(cell.text for table in parsed.tables
                     for row in table.rows for cell in row.cells)
    assert "00031" in text and "Inspect boiler" in text
    assert "2026-09-15" in text and "B-31" in text
    assert len(parsed.tables) == 1
    assert sources.source_bytes(profile, tagged_source) == sources.source_bytes(profile, data.source)



def test_unlinked_source_continuation_is_excluded_only_after_explicit_review(storage):
    content = read_csv("continued.csv", (
        b"WO #,Finish Date,Area,Task,Tag,Site,Status\n"
        b"00041,2026-09-15,Plant,Check pump,P-41,Demo North,Completed\n"
        b",,,Additional handwritten note,,,\n"
    ))
    review = review_service_calls(content, PROFILE, PERIOD)
    assert review.counts()["reviewed_candidates"] == 1
    assert review.counts()["unmatched"] == 1
    assert review.unmatched[0].number == 3
    assert not review.ready_for_inclusion
    with pytest.raises(ValueError, match="Review unresolved"):
        reviewed_service_call_block(review)
    excluded = exclude_unmatched_call(
        review, 3, expected_revision=review.revision)
    assert excluded.counts()["unmatched"] == 0
    assert excluded.unmatched[0].cells[3] == "Additional handwritten note"
    assert excluded.ready_for_inclusion
    assert reviewed_service_call_block(excluded).rows[0][0] == "00041"
    with pytest.raises(ValueError, match="changed"):
        exclude_unmatched_call(excluded, 3, expected_revision=review.revision)
    restored = restore_service_call_review(
        json.loads(json.dumps(service_call_review_record(excluded))))
    assert restored == excluded
    assert reimport_service_calls(restored, review) == restored


def test_native_word_multi_table_collisions_require_combined_review(storage):
    content, raw = read_word(storage, word_fixture(two_tables=True))
    report = review_word_maintenance(PROFILE, content, PERIOD)
    assert report.cross_table_duplicates == (("north", "0009"),)
    with pytest.raises(ValueError, match="Duplicate work orders"):
        reviewed_word_maintenance_block(report)
    assert sources.source_bytes(PROFILE, content.source) == raw


def test_native_word_reviewed_table_builds_one_source_attributed_section_block(storage):
    content, original = read_word(storage, word_fixture())
    report = review_word_maintenance(PROFILE, content, PERIOD)
    table, = report.tables
    review = exclude_service_call(table.review, table.review.rows[1].key,
        expected_revision=table.review.revision)
    reviewed = confirm_word_layout(review, expected_revision=review.revision)
    accepted = replace(report, tables=(replace(table, review=reviewed),))
    block = reviewed_word_maintenance_block(accepted)
    assert block.rows == (
        ("0009", "2026-09-05", "Demonstration North Facility — Boiler Room",
         "Replace seal", "PUMP-09"),)
    assert block.references == (reviewed.source_reference,)
    assert reviewed_word_maintenance_block(accepted, existing=block) == block
    assert sources.source_bytes(PROFILE, content.source) == original



def test_xlsx_cmms_is_supported_but_annual_count_grid_is_not_calls(storage):
    from openpyxl import Workbook
    wb = Workbook()
    sheet = wb.active
    sheet.title = "Synthetic maintenance"
    sheet.append(("Work Order Number", "Actual Finish", "Area", "Work Performed",
                  "Equipment Tag", "Site", "State"))
    sheet.append(("000061", "2026-09-09", "Mechanical", "Inspected AHU",
                  "AHU-61", "Demo South", "Closed"))
    buf = BytesIO()
    wb.save(buf)
    raw = buf.getvalue()
    content, _ = sources.ingest(PROFILE, "export.xlsx", raw)
    review = review_service_calls(content, PROFILE, PERIOD)
    assert review.ready_for_inclusion
    assert review.rows[0].wo == "000061"
    assert review.rows[0].site == "south"
    assert review.rows[0].finished == "2026-09-09"
    assert reviewed_service_call_block(review).rows[0][2] == (
        "Demonstration South Facility — Mechanical")
    assert sources.source_bytes(PROFILE, content.source) == raw

    summary = read_csv("annual-grid.csv",
        b"Metric,January,February,March,Total\n"
        b"PM Count,4,0,,4\n"
        b"Open Orders,2,1,,3\n")
    not_calls = review_service_calls(summary, PROFILE, PERIOD)
    assert not_calls.rows and not not_calls.ready_for_inclusion
    assert not_calls.notes
    assert not_calls.counts()["unmatched"] == 2
    assert sources.source_bytes(PROFILE, summary.source).startswith(b"Metric")



def test_excluding_a_duplicate_recalculates_remaining_workorder_without_row_loss(storage):
    original = read_csv("duplicate-exclusion.csv",
        b"WO #,Finish Date,Area,Task,Tag,Site,Status\n"
        b"00091,2026-09-21,Plant,Repair pump,P-91,Demo North,Completed\n"
        b"00091,2026-09-22,Plant,Repeat line,P-91,Demo North,Completed\n")
    review = review_service_calls(original, PROFILE, PERIOD)
    assert len(review.rows) == 2
    assert review.counts()["ready"] == 0
    assert not review.ready_for_inclusion
    accepted = exclude_service_call(review, review.rows[1].key,
        expected_revision=review.revision)
    assert accepted.ready_for_inclusion
    assert len(accepted.rows) == 2
    assert accepted.rows[1].excluded
    assert accepted.rows[0].issues == ()
    assert len(reviewed_service_call_block(accepted).rows) == 1
    assert reviewed_service_call_block(accepted).rows[0][0] == "00091"
    assert reimport_service_calls(accepted, review) == accepted
    assert sources.source_bytes(PROFILE, original.source).endswith(b"Completed\n")
