"""Synthetic acceptance across the independently developed source adapters."""
from dataclasses import replace
from io import BytesIO

from docx import Document
import pytest

from app import monthly_report_sources as sources, monthly_report_library as library
from app import monthly_report_work_order_review_flow as flow
from app.monthly_report_model import synthetic_profiles
from app.monthly_report_source_scope import correct_scope
from test_monthly_report_work_order_review_flow import draft_source, stored, publish, PERIOD
from tests.conftest import requires_libreoffice


def table(document, columns, rows, title=""):
    result = document.add_table(rows=1 + len(rows) + bool(title), cols=len(columns))
    offset = int(bool(title))
    if title:
        result.rows[0].cells[0].merge(result.rows[0].cells[-1]).text = title
    for target, values in zip(result.rows[offset:], (columns, *rows)):
        for cell, value in zip(target.cells, values):
            cell.text = value
    return result


def bytes_for(document):
    result = BytesIO()
    document.save(result)
    return result.getvalue()


def saved(draft):
    library.save_profile(draft.profile, expected_revision=0, actor="Synthetic Editor", confirmed=True)
    library.save_snapshot(draft, expected_revision=0)
    return library.load_snapshot(draft.profile.contract, draft.profile.key, draft.period).draft


def generated(draft, storage, monkeypatch):
    from app import pdf_converter
    from app.monthly_report_docx import generate_report
    monkeypatch.setattr(pdf_converter, "OUTPUT_DIR", storage)
    monkeypatch.setattr(pdf_converter, "PDF_BACKEND", "libreoffice")
    result = generate_report(draft, acknowledged_fingerprint=draft.fingerprint)
    assert result.pdf and not result.pdf_error
    import pymupdf
    with pymupdf.open(stream=result.pdf, filetype="pdf") as pdf:
        text = "\n".join(page.get_text() for page in pdf)
    return result, text


@requires_libreoffice
def test_qualified_site_tables_review_reopen_and_export_beside_annual_totals(tmp_path, monkeypatch):
    from test_monthly_report_annual_site_month_grid import source, inspected, EXPECTED
    from app.monthly_report_annual_work_order_grid import annual_site_month_table
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path / "data"))
    profile = synthetic_profiles()[1]
    annual_path = source(tmp_path / "annual.docx")
    _, annual_item = inspected(annual_path)
    annual = annual_site_month_table(annual_item)
    document = Document()
    document.add_heading("2 Monthly Activity Summary", 1)
    columns = ("Work Order", "Site", "Date", "Status", "Type", "Trade")
    for site, heading in (("Demo North", "WORK ORDER STATUS — SITE A"),
                          ("Demo South", "PM STATUS — SITE B")):
        document.add_paragraph("Improvement Highlights")
        document.add_paragraph("Synthetic prior improvement; never a work-order record.")
        document.add_paragraph(heading)
        table(document, columns, (("000071", site, "2026-09-05", "Open", "PM", "HVAC"),),
              title="Work Order Status")
    original = bytes_for(document)
    draft, content = draft_source(original, "qualified-workorders.docx", profile)
    content = replace(content, source=correct_scope(content.source, profile, ("north", "south")))
    draft = replace(draft, sources=(content.source,),
        sections=tuple(s for s in draft.sections if s.key == "activity"),
        blocks=tuple(replace(b, extra_tables=(annual,)) if b.key == "work_orders" else b for b in draft.blocks))
    draft, bundle = stored(draft, content)
    assert len(bundle["groups"]) == 2
    second = flow.restored_group(bundle, bundle["groups"][1])
    corrected = flow.change_row(bundle, 1, second.rows[0].key, {"status": "Completed"})
    draft = flow.store_bundle(draft, corrected, expected_revision=flow.revision(bundle))
    accepted = publish(draft, corrected)
    reopened = saved(accepted)
    receipt = flow.saved_bundle(reopened, content.source.id, "workorders")
    assert publish(reopened, receipt) == reopened
    block = next(b for b in reopened.blocks if b.key == "work_orders")
    assert block.extra_tables[0] == annual and annual.rows == EXPECTED
    records = [t for t in block.extra_tables if t.columns[0] == "Work order"]
    assert [(t.rows[0][0], t.rows[0][1], t.rows[0][4]) for t in records] == [
        ("000071", "north", "Open"), ("000071", "south", "Completed")]
    result, text = generated(reopened, tmp_path, monkeypatch)
    assert text.count("000071") == 2
    word = Document(BytesIO(result.docx))
    assert any(tuple(c.text for c in row.cells) == EXPECTED[0] for t in word.tables for row in t.rows)
    assert sources.source_bytes(profile, content.source) == original


@pytest.mark.parametrize("suffix", [".csv", ".docx"])
def test_multiline_cmms_tasks_remain_editable_and_keep_original_context(tmp_path, monkeypatch, suffix):
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path))
    columns = ("WO #", "Finish Date", "Area", "Task Code-Description", "Tag #", "Site", "Status", "Notes")
    rows = (("00031", "2026-09-06", "Plant", "Inspect pump\nRecord pressure", "P-31", "North", "Closed", "Line one\nLine two"),
            ("00032", "2026-09-07", "Plant", "Test valve", "V-32", "South", "Completed", ""))
    if suffix == ".csv":
        import csv
        from io import StringIO
        stream = StringIO(newline="")
        csv.writer(stream).writerows((columns, *rows))
        raw = stream.getvalue().encode()
    else:
        document = Document()
        document.add_heading("5 Maintenance Schedule / In-House Maintenance", 1)
        table(document, columns, rows, title="Service calls")
        raw = bytes_for(document)
    draft, content = draft_source(raw, "multiline" + suffix)
    draft, bundle = stored(draft, content, "maintenance")
    record = flow.restored_group(bundle, bundle["groups"][0])
    assert [r.wo for r in record.rows] == ["00031", "00032"]
    assert not record.unmatched
    assert record.rows[0].task == "Inspect pump\nRecord pressure"
    revised = flow.change_row(bundle, 0, record.rows[0].key, {"task": "Inspect pump\nPressure checked"})
    revised = flow.group_decisions(revised, 0, omitted=False, layout=True)
    draft = flow.store_bundle(draft, revised, expected_revision=flow.revision(bundle))
    reopened = saved(publish(draft, revised))
    receipt = flow.saved_bundle(reopened, content.source.id, "maintenance")
    assert publish(reopened, receipt) == reopened
    restored = flow.restored_group(receipt, receipt["groups"][0])
    assert restored.rows[0].origin.cells[3] == "Inspect pump\nRecord pressure"
    assert restored.unmapped(restored.rows[0])[-1] == (7, "Notes", "Line one\nLine two")
    output = next(b for b in reopened.blocks if b.key == "service_calls").extra_tables[0]
    assert [row[0] for row in output.rows] == ["00031", "00032"]
    assert output.rows[0][3] == "Inspect pump\nPressure checked"
    assert sources.source_bytes(reopened.profile, content.source) == raw


def test_multiline_identity_is_still_an_unresolved_native_record(tmp_path, monkeypatch):
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path))
    draft, content = draft_source(b'WO #,Finish Date,Area,Task,Tag,Site,Status\n"001\n002",2026-09-06,Plant,Inspect,P-1,North,Closed\n')
    _, bundle = stored(draft, content, "maintenance")
    review = flow.restored_group(bundle, bundle["groups"][0])
    assert not review.rows and review.unmatched[0].cells[0] == "001\n002"
    with pytest.raises(ValueError):
        publish(draft, bundle)


def test_native_task_editor_preserves_line_breaks_on_correction(tmp_path, monkeypatch):
    from test_monthly_report_work_order_review_ui import app_for, button
    app = app_for(monkeypatch, tmp_path,
        b'WO #,Finish Date,Area,Task,Tag,Site,Status\n00031,2026-09-06,Plant,"Inspect pump\nRecord pressure",P-31,North,Closed\n')
    next(w for w in app.selectbox if w.label == "Use these records for").select("Completed in-house maintenance calls").run()
    button(app, "Review and correct imported rows").click().run()
    assert not app.exception
    task = next(w for w in app.text_area if w.label == "Work performed")
    assert task.value == "Inspect pump\nRecord pressure"
    task.set_value("Inspect pump\nPressure checked")
    button(app, "Save row correction").click().run()
    assert not app.exception
    draft = app.session_state["draft"]
    receipt = flow.saved_bundle(draft, draft.sources[0].id, "maintenance")
    assert flow.restored_group(receipt, receipt["groups"][0]).rows[0].task == "Inspect pump\nPressure checked"


@requires_libreoffice
def test_native_title_band_review_exports_first_last_calls_and_multiline_task(tmp_path, monkeypatch):
    from app.monthly_report_import import inspect_docx
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path / "data"))
    document = Document()
    document.add_heading("5 Maintenance Schedule / In-House Maintenance", 1)
    columns = ("WO #", "Finish Date", "Area", "Task Code-Description", "Tag #")
    rows = (("00081", "2026-09-06", "Plant", "Inspect pump\nRecord pressure", "P-81"),
            ("00082", "2026-09-07", "Plant", "Test last valve", "V-82"))
    table(document, columns, rows, title="Service calls")
    raw = bytes_for(document)
    path = tmp_path / "title-band.docx"
    path.write_bytes(raw)
    inspection = inspect_docx(path)
    assert any(item.text == "Service calls" and item.suggested_slot == "service_calls" for item in inspection.items)
    profile = synthetic_profiles()[0]
    draft, content = draft_source(raw, path.name, profile)
    content = replace(content, source=correct_scope(content.source, profile, ("north",)))
    draft = replace(draft, sources=(content.source,), sections=tuple(s for s in draft.sections if s.key == "maintenance"))
    draft, bundle = stored(draft, content, "maintenance")
    record = flow.restored_group(bundle, bundle["groups"][0])
    assert len(record.rows) == 2
    changed = bundle
    for row in record.rows:
        changed = flow.change_row(changed, 0, row.key, {"site": "north", "status": "Completed"})
    changed = flow.group_decisions(changed, 0, omitted=False, layout=True)
    draft = flow.store_bundle(draft, changed, expected_revision=flow.revision(bundle))
    reopened = saved(publish(draft, changed))
    result, text = generated(reopened, tmp_path, monkeypatch)
    assert text.count("00081") == text.count("00082") == 1
    assert all(value in text for value in ("Inspect pump", "Record pressure", "Test last valve"))
    output = Document(BytesIO(result.docx))
    assert any(tuple(c.text for c in row.cells) == rows[0] for t in output.tables for row in t.rows)
    assert sources.source_bytes(profile, content.source) == raw
