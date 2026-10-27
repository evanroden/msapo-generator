"""Content labels describe inspected structures without promising OCR/editability."""

from app.monthly_report_content_kind import content_kind_description, section_content_description
from app.monthly_report_import import ImportItem, inspect_docx
from test_monthly_report_import import synthetic_docx, rewrite
from zipfile import ZipFile


def test_real_word_inventory_distinguishes_picture_table_and_floating_layout(tmp_path):
    path = synthetic_docx(tmp_path)
    with ZipFile(path) as archive:
        xml = archive.read("word/document.xml").decode()
    box = ('<w:p><w:r><w:pict><v:shape xmlns:v="urn:schemas-microsoft-com:vml"><w:txbxContent><w:p><w:r>'
           '<w:t>Synthetic floating text</w:t></w:r></w:p>'
           '</w:txbxContent></v:shape></w:pict></w:r></w:p>')
    rewrite(path, {"word/document.xml": xml.replace("</w:body>", box + "</w:body>").encode()})
    original = path.read_bytes()
    inspection = inspect_docx(path)
    table = next(i for i in inspection.items if i.kind == "table")
    picture = next(i for i in inspection.items if i.kind == "image")
    floating = next(i for i in inspection.items if i.kind == "unsupported")
    assert content_kind_description(table).startswith("Native Word table")
    assert "not directly editable" in content_kind_description(picture)
    assert "optional" in content_kind_description(picture)
    assert "not an editable table" in content_kind_description(floating)
    assert section_content_description((table, picture, floating)) == (
        "Imported content: 1 native table · 1 picture · 1 layout needing review")
    assert path.read_bytes() == original
    assert inspect_docx(path) == inspection


def test_empty_and_unknown_inventory_does_not_invent_editable_content():
    assert section_content_description(()) == ""
    item = ImportItem("unknown", "future_type", "word/document.xml", 1, 1)
    assert "not recognized" in content_kind_description(item)
    native = ImportItem("native", "native_layout", "word/document.xml", 2, 1)
    assert "not an editable table" in content_kind_description(native)
    assert section_content_description((native,)) == "Imported content: 1 native layout"


def test_section_review_displays_native_table_type_without_changing_cells():
    from streamlit.testing.v1 import AppTest
    app = AppTest.from_string('''
import streamlit as st
from pathlib import Path
from app.monthly_report_model import ReportPeriod
from app.monthly_report_import import ImportItem
from app.monthly_report_sections import ImportSection
from app.monthly_report_section_ui import _section_card
from app.monthly_report_ui import _field
item = ImportItem('table', 'table', 'word/document.xml', 1, 1, 'capital',
                  rows=(('Equipment', 'Timing'), ('Synthetic pump', '2030')))
plans = st.session_state.setdefault('plans', {})
_section_card(ImportSection('capital', 'Equipment plans', (item,)), Path('synthetic.docx'),
              None, ReportPeriod(2026,9), 'report_kind', _field, plans, True)
''', default_timeout=20).run()
    assert not app.exception
    assert any(c.value == "Imported content: 1 native table" for c in app.caption)
    app.radio[0].set_value("Edit text or change pictures").run()
    assert not app.exception
    table = app.session_state['plans']['capital']['tables']['table']
    assert table.rows == (("Synthetic pump", "2030"),)
    assert table.columns == ("Equipment", "Timing")


def test_picture_guidance_reaches_existing_replacement_without_losing_text(tmp_path, monkeypatch):
    from streamlit.testing.v1 import AppTest
    from docx import Document
    from PIL import Image
    from io import BytesIO
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path / "data"))
    picture = BytesIO()
    Image.new("RGB", (500, 600), "navy").save(picture, "PNG")
    document = Document()
    document.add_heading("2 Monthly Activity Summary", 1)
    document.add_paragraph("Synthetic pump inspected.")
    document.add_picture(BytesIO(picture.getvalue()))
    path = tmp_path / "activity.docx"
    document.save(path)
    original = path.read_bytes()
    app = AppTest.from_string('''
import streamlit as st
from pathlib import Path
from app.monthly_report_import import inspect_docx
from app.monthly_report_sections import section_reviews
from app.monthly_report_section_ui import _section_card
from app.monthly_report_model import ReportPeriod
from app.monthly_report_ui import _field
path = Path(SOURCE)
inspection = inspect_docx(path)
section = next(s for s in section_reviews(inspection) if s.key == 'activity')
plans = st.session_state.setdefault('plans', {})
_section_card(section, path, inspection, ReportPeriod(2026,9), 'report_picture_kind', _field, plans, True)
'''.replace('SOURCE', repr(str(path))), default_timeout=20).run()
    assert not app.exception
    assert any("1 picture" in c.value for c in app.caption)
    assert any("Update work, counts or photos” to reach replacement uploads" in c.value
               for c in app.caption)
    next(w for w in app.radio if w.label == "What would you like to do?").set_value(
        "Edit text or change pictures").run()
    assert not app.exception
    assert "New improvement photo" in [w.label for w in app.get("file_uploader")]
    assert "Synthetic pump inspected." in list(app.session_state['plans']['activity']['texts'].values())
    app.run()
    assert not app.exception
    assert path.read_bytes() == original
