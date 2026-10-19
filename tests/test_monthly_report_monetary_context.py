"""Client-report price boundaries, with source/state/editable-output controls.

All inputs are synthetic. This suite does not claim native Word or graphic
privacy acceptance; the document test inspects the actual generated package.
"""
from dataclasses import replace
from io import BytesIO
from hashlib import sha256
import json
import os
from pathlib import Path
from zipfile import ZipFile

import fitz
import pytest
from docx import Document

from app.monthly_report_checks import preflight
from app.monthly_report_content_policy import contains_price, price_column
from app.monthly_report_docx import assemble_docx
from app.monthly_report_followups import seed_followups
from app.monthly_report_guided import client_table_draft
from app.monthly_report_issues import canonical_issues
from app.monthly_report_model import (
    ReportDraft, ReportPeriod, ReportSource, ReportTable, ResolvedBlock,
    default_sections, synthetic_profiles,
)
from app.monthly_report_sources import SourceContent, SourceTable
from app.monthly_report_structured_uploads import mapped_rows
from tests.conftest import requires_libreoffice


@pytest.mark.parametrize("text", (
    "Quote: 18,675.50", "Quotation 18,675.50", "Vendor quoted 4,500 for seals",
    "Repair estimate 6,875", "Est. 6,875", "NTE 29,450", "Not-to-exceed 29,450",
    "Not to exceed: 29,450", "Amount: 7,815.25", "Grand Total 4,500.00",
    "Balance Due 3,200.00", "Invoice Total 2,310.55", "4,500.00 $", "CHF 4500",
    "Quote\n18,675.50", "Repair estimate\n6,875", "NTE\n29,450",
    "Cost to repair: 8,640", "1500 cost", "September 2026 cost: 1500",
    "Quote: 12,500 per unit", "Estimate: 1.5k", "Amount (USD) 3500",
    "Cost for emergency repair: 8,640", "Repair cost, including labor: 8,640",
    "Quote—18,675.50", "Quotation revised to 18,000", "Not-to-exceed limit: 29,450",
    "Repair cost including labor and materials: 8,640",
    "Repair cost, including labor, materials and tax: 8,640",
    "Total cost including tax: 8,640", "Cost of repair: 8,640 including tax",
    "Quote for replacement parts: 8,640", "Repair cost: 2026", "Estimate 2026",
    "Budget for 2026: 8,640", "Estimate for 2026 is 2026",
))
def test_explicit_money_context_is_blocked(text):
    assert contains_price(text), text


@pytest.mark.parametrize("text", (
    "Life-cycle cost analysis for Chiller 2 was reviewed.",
    "Utility costs decreased 8% compared with last year.",
    "Utility costs decreased 8 percent compared with last year.",
    "Labor rates review meeting held with the team.",
    "Energy cost trend chart updated.",
    "Price of failure discussed in safety meeting.",
    "Cost-effective repair completed in 4 hours.",
    "Quote number 18675", "Quote #18675", "Quotation reference Q-18675",
    "Quote date: 2026-11-14", "Repair estimate: 48 hours",
    "Estimated completion: 2026-11-14", "Estimated flow: 125 gpm",
    "Amount of refrigerant: 650 lb", "Amount of water: 240 gallons",
    "Not to exceed 125 psi", "Estimate: 8%", "Estimate: 8 percent",
    "2026 cost analysis is ongoing.",
    "Flow rate 150 gpm; conductivity 3000 uS/cm; temperature 110 F",
    "Budget planning is under review; no figures are approved.",
    "Not to exceed 480 V", "Not to exceed 12 A", "Estimate: 200 F",
    "Cost for 2026 remains unknown.", "Estimate for 2026 remains pending.",
    "Budget for 2026 has not been approved.",
    "Cost estimate for Chiller 2 remains pending.",
    "Repair cost decreased 12 percent after scope correction.",
))
def test_identifiers_measurements_dates_and_unpriced_work_survive(text):
    assert not contains_price(text), text


@pytest.mark.parametrize("heading", (
    "Quote", "Quotation", "Repair estimate", "Est.", "NTE", "Not-to-exceed",
    "Amount ($)", "Amount (USD)", "Grand Total", "Balance Due", "Invoice Total",
))
def test_money_headings_identify_the_price_column(heading):
    assert price_column(heading), heading


@pytest.mark.parametrize("heading", (
    "Quote number", "Quote #", "Quote date", "Quotation reference", "Quote status",
    "Estimate duration (hours)", "Amount of refrigerant (lb)", "Estimated flow",
    "Not-to-exceed pressure (psi)", "Flow rate", "Capacity", "Work Description",
    "Quote ref.", "Estimate (hours)", "Amount (kg)", "Amount (liters)",
))
def test_nonmonetary_headings_are_never_removed_as_prices(heading):
    assert not price_column(heading), heading


def test_technical_quote_reference_and_quantity_cells_survive_client_transform():
    table = ReportTable(("Asset", "Quote ref.", "Estimate (hours)", "Amount (kg)"),
                        (("EF-9", "Q-18765", "48", "50"),))
    draft = _draft("capital", "capital_renewal", ResolvedBlock("capital_renewal", "This month",
                                                            extra_tables=(table,)))
    working, removed = client_table_draft(draft)
    assert not removed
    assert working.blocks[0].extra_tables == (table,)


def _draft(section, key, block):
    sections = tuple(replace(s, included=s.key == section) for s in default_sections())
    return ReportDraft(synthetic_profiles()[0], ReportPeriod(2026, 11),
                       "Synthetic Price Reviewer", sections, (replace(block, key=key),))


@pytest.mark.parametrize("heading,value", (
    ("Quote", "18,675.50"), ("Repair estimate", "6,875"), ("Not-to-exceed", "29,450"),
))
def test_issue_import_reports_exclusion_and_retains_editable_facts(heading, value):
    table = ReportTable(("Asset", "Condition", heading, "Pressure", "Work order"),
                        (("EF-8", "Actuator did not close", value, "125 psi", "WO-8675309"),))
    original = _draft("issues", "equipment_issues",
                      ResolvedBlock("equipment_issues", "This month", extra_tables=(table,)))
    working, removed = client_table_draft(original)
    assert removed == (heading,), "The UI needs the exact exclusion notice."
    assert original.blocks[0].extra_tables == (table,), "Original/history stays unchanged."
    assert working.blocks[0].extra_tables[0].columns == ("Asset", "Condition", "Pressure", "Work order")
    migrated = canonical_issues(working)
    assert not any(c.blocking and c.code == "pricing" for c in preflight(migrated))
    raw = assemble_docx(migrated, acknowledged_fingerprint=migrated.fingerprint)
    with ZipFile(BytesIO(raw)) as z:
        assert not any(value.encode() in z.read(name) for name in z.namelist()), "Price in a visible/hidden package part."
    text = "\n".join(p.text for p in Document(BytesIO(raw)).paragraphs)
    assert all(fact in text for fact in ("EF-8", "Actuator did not close", "125 psi", "WO-8675309"))


@pytest.mark.parametrize("section,key", (("activity", "activity_summary"), ("issues", "equipment_issues")))
def test_unresolved_narrative_price_is_preserved_for_correction_and_blocks_generation(section, key):
    text = "Valve remains out of service. Repair estimate: 6,875."
    original = _draft(section, key, ResolvedBlock(key, "This month", text=text))
    working, removed = client_table_draft(original)
    assert not removed and working.blocks[0].text == text
    checks = preflight(working)
    assert any(c.code == "pricing" and c.blocking and c.block_key == key for c in checks)
    with pytest.raises(ValueError, match="[Pp]ric"):
        assemble_docx(working, acknowledged_fingerprint=working.fingerprint)


def test_proposal_rollover_and_custom_schema_exclude_only_money():
    headings = ("Facility", "Vendor", "Quote", "Work Description", "Status", "Quote reference", "Submitted")
    row = ("Synthetic North", "Synthetic Vendor", "18,675.50", "Replace actuator; test at 125 psi", "Pending", "Q-18675", "2026-11-14")
    source = ReportSource("a" * 64, "synthetic-proposal.csv", "a" * 64, ".csv",
                          page_texts=("\t".join(headings) + "\n" + "\t".join(row),))
    content = SourceContent(source, (SourceTable("Proposal", headings, (row,)),))
    target = ("Vendor", "Work Description", "Quote reference", "Submitted", "Status", "Facility")
    assert mapped_rows(content, "proposals", target) == ((row[1], row[3], row[5], row[6], row[4], row[0]),)
    original = _draft("proposals", "proposals", ResolvedBlock("proposals", "This month",
                       extra_tables=(ReportTable(headings, (row,)),)))
    carried = seed_followups(original)
    assert len(carried) == 1
    assert row[2] not in carried[0].text
    assert all(v in carried[0].text for v in (row[1], row[3], row[5], row[6]))


@requires_libreoffice
def test_saved_price_exclusion_reaches_editable_word_and_real_pdf(monkeypatch, tmp_path):
    """Synthetic editor tables -> two real saved revisions -> normal output.

    This checks the table correction path, not a native file-uploader event or
    unreviewed image privacy. The original private history keeps its evidence.
    """
    from app import monthly_report_library as library
    from app.monthly_report_docx import generate_report

    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path / "data"))
    fixtures = (
        ("equipment_issues", "Quote", "18,675.50", "EF-8", "Actuator did not close"),
        ("capital_renewal", "Repair estimate", "6,875", "CH-12", "Replace worn seals"),
        ("proposals", "Not-to-exceed", "29,450", "P-27", "Inspect isolation valve"),
    )
    sections = tuple(replace(s, included=s.key in {"issues", "capital", "proposals"})
                     for s in default_sections())
    blocks = tuple(ResolvedBlock(key, "This month", extra_tables=(ReportTable(
        ("Asset", "Scope", heading, "Pressure", "Work order"),
        ((asset, scope, value, "125 psi", "WO-8675309"),)),))
        for key, heading, value, asset, scope in fixtures)
    original = ReportDraft(synthetic_profiles()[0], ReportPeriod(2026, 11),
                           "Synthetic Price Reviewer", sections, blocks)
    saved_original = library.save_snapshot(original, expected_revision=0,
                                          entered_editor=original.prepared_by)
    loaded = library.load_snapshot(original.profile.contract, original.profile.key, original.period)
    corrected, excluded = client_table_draft(loaded.draft)
    assert excluded == tuple(row[1] for row in fixtures)
    corrected = canonical_issues(corrected)
    saved = library.save_snapshot(corrected, expected_revision=saved_original.revision,
                                  entered_editor=corrected.prepared_by)
    restored = library.load_snapshot(original.profile.contract, original.profile.key, original.period)
    assert restored.revision == saved.revision == 2
    assert restored.draft == corrected
    history = library.load_snapshot_version(original.profile.contract, original.profile.key,
                                            original.period, saved_original.revision)
    assert history.draft == original
    assert not [c for c in preflight(restored.draft) if c.blocking]
    package = generate_report(restored.draft, acknowledged_fingerprint=restored.draft.fingerprint)
    assert package.docx and package.pdf, package.pdf_error
    document = Document(BytesIO(package.docx))
    editable_cells = [cell.text for table in document.tables for row in table.rows for cell in row.cells]
    assert "Replace worn seals" in editable_cells and "Inspect isolation valve" in editable_cells
    package_text = "\n".join([*(p.text for p in document.paragraphs), *editable_cells])
    with ZipFile(BytesIO(package.docx)) as archive:
        for value in (row[2] for row in fixtures):
            assert not any(value.encode() in archive.read(name) for name in archive.namelist())
    output = Path(os.environ.get("INDEPENDENT_OUTPUT", str(tmp_path)))
    output.mkdir(parents=True, exist_ok=True)
    (output / "price_excluded.docx").write_bytes(package.docx)
    (output / "price_excluded.pdf").write_bytes(package.pdf)
    with fitz.open(stream=package.pdf, filetype="pdf") as pdf:
        pdf_text = " ".join(" ".join(page.get_text().split()) for page in pdf)
        for number, page in enumerate(pdf, 1):
            page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5)).save(output / f"price_excluded_page_{number}.png")
        page_count = len(pdf)
    for _, _, value, asset, scope in fixtures:
        assert value not in pdf_text
        assert asset in package_text and asset in pdf_text
        assert scope in package_text and scope in pdf_text
    assert "125 psi" in package_text and "125 psi" in pdf_text
    assert "WO-8675309" in package_text and "WO-8675309" in pdf_text
    (output / "price_excluded.json").write_text(json.dumps({
        "tested_revision": os.environ.get("MSAPO_SHA", "unrecorded"),
        "fixture": "Synthetic editor tables; no uploaded graphics or approved master",
        "original_fingerprint": original.fingerprint,
        "saved_revision": restored.revision, "draft_fingerprint": restored.draft.fingerprint,
        "excluded_headings": excluded, "pdf_pages": page_count,
        "docx_bytes": len(package.docx), "docx_sha256": sha256(package.docx).hexdigest(),
        "pdf_bytes": len(package.pdf), "pdf_sha256": sha256(package.pdf).hexdigest(),
    }, indent=2) + "\n")
