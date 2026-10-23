"""Explicit row actions survive native Word and actual PDF, with synthetic inputs."""
from dataclasses import replace
from io import BytesIO

from docx import Document
import pymupdf
import pytest

from conftest import requires_libreoffice
from app import pdf_converter
from app.monthly_report_import import inspect_docx
from app.monthly_report_model import (
    Facility, ReportDraft, ReportPeriod, ReportProfile, ReportTable,
    ResolvedBlock, default_sections,
)
from app.monthly_report_native_layout import build_native_docx
from test_work_order_imported_rows import current, launch, table
from test_work_order_row_insertion import ABOVE, BELOW, wire


@pytest.fixture(params=("annual", "custom"))
def native_outputs(request, tmp_path, monkeypatch):
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path / "isolated-data"))
    monkeypatch.setattr(pdf_converter, "OUTPUT_DIR", tmp_path)
    monkeypatch.setattr(pdf_converter, "PDF_BACKEND", "libreoffice")
    if request.param == "annual":
        original, second_anchor, month = table(), 3, 10
    else:
        original = ReportTable(("Ticket", "Asset", "Status"), (
            ("", "", ""), ("WO-001", "Synthetic pump", "Open"),
            ("", "", ""), ("WO-002", "Synthetic fan", "Done"),
        ))
        second_anchor, month = 1, 9
    source = Document()
    source.add_heading("MONTHLY ACTIVITY SUMMARY", 1)
    source.add_paragraph("WORK ORDERS")
    native = source.add_table(rows=1 + len(original.rows), cols=len(original.columns))
    native.style = "Table Grid"
    for row, values in zip(native.rows, (original.columns, *original.rows)):
        for cell, value in zip(row.cells, values):
            cell.text = value
    source_path = tmp_path / "synthetic-source.docx"
    source.save(source_path)
    source_bytes = source_path.read_bytes()
    inspection = inspect_docx(source_path)
    item = next(item for item in inspection.items if item.kind == "table")
    original = replace(original, reference=f"docx:{inspection.sha256}:{item.id}")
    compact = replace(original, rows=tuple(row for row in original.rows if any(row)))

    app = launch(compact)
    assert current(app) == compact  # Opening never guesses the missing rows.
    wire(app, 0, ABOVE)
    wire(app)
    wire(app, second_anchor, BELOW)
    wire(app)
    restored = current(app)
    assert restored == original

    profile = ReportProfile("Synthetic " + request.param, "synthetic-site", "Synthetic Site",
                            (Facility("synthetic-site", "Synthetic Site"),))
    sections = (next(section for section in default_sections() if section.key == "activity"),)
    outputs = []
    for label, value in (("compact-before", compact), ("explicit-after", restored)):
        draft = ReportDraft(profile, ReportPeriod(2026, month), "Synthetic Editor", sections,
                            (ResolvedBlock("work_orders", "This month", extra_tables=(value,)),),
                            synthetic=True)
        raw = build_native_docx(draft, source_path, section_key="activity")
        document = Document(BytesIO(raw))
        assert len(document.tables) == 1
        rows = tuple(tuple(cell.text for cell in row.cells) for row in document.tables[0].rows)
        assert rows == (value.columns, *value.rows)
        assert not any("Insert blank" in cell or "__work_order_row_" in cell
                       for row in rows for cell in row)
        output = tmp_path / (label + ".docx")
        output.write_bytes(raw)
        outputs.append((output, rows))
    assert source_path.read_bytes() == source_bytes
    return outputs


def test_explicit_restoration_preserves_native_word_cells_and_source(native_outputs):
    before, after = native_outputs
    assert len(after[1]) == len(before[1]) + 2
    assert after[1][0] == before[1][0]


@requires_libreoffice
def test_explicit_restoration_retains_blank_table_rows_in_actual_pdf(native_outputs):
    for output, rows in native_outputs:
        pdf_path = pdf_converter.convert_to_pdf(output)
        try:
            with pymupdf.open(pdf_path) as pdf:
                tables = [table for page in pdf for table in page.find_tables().tables]
                assert [(table.row_count, table.col_count) for table in tables] == [
                    (len(rows), len(rows[0]))]
                text = "\n".join(page.get_text() for page in pdf)
                expected_text = "Pending" if len(rows[0]) == 7 else "Synthetic pump"
                assert expected_text in text
                assert "Insert blank" not in text and "__work_order_row_" not in text
        finally:
            pdf_path.unlink(missing_ok=True)
