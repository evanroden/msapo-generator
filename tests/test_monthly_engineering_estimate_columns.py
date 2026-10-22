"""Invented engineering estimates survive import; financial amounts stay excluded."""
from dataclasses import replace
from io import BytesIO

import pytest
from docx import Document
from docx.shared import Inches

from app import monthly_report_library as library
from app.monthly_report_checks import preflight
from app.monthly_report_content_policy import (
    ambiguous_price_columns, contains_price, price_column, price_free_table,
    table_has_pricing, table_price_columns,
)
from app.monthly_report_docx import assemble_docx, generate_report
from app.monthly_report_guided import client_table_draft
from app.monthly_report_import import ImportItem, MappedImport, imported_draft, inspect_docx
from app.monthly_report_model import (
    BlockSpec, ColumnSpec, Facility, ReportPeriod, ReportProfile, ReportTable, ResolvedBlock,
)
from app.monthly_report_sections import build_section_import, table_without_prices
from app.monthly_report_setup import design_profile
from app.monthly_report_start_replay import merge_starting_drafts
from conftest import requires_libreoffice


ENGINEERING_HEADINGS = (
    "Est. Peak Coincident Load (Tons)",
    "Est. Peak Coincident Load (MBH)",
    "Est.\nPeak Coincident\nLoad (Tons)",
    "Est. Peak CoincidentLoad (MBH)",
    "Est PeakCoincidentLoad (Tons)",
    "Estimate Peak Coincident Load (MBH)",
    "Estimates Peak Coincident Load (Tons)",
    "Est.  Peak\tCoincident Load ( MBH )",
)
PROFILE = ReportProfile(
    "Synthetic engineering contract", "synthetic-plants", "Synthetic Plants",
    (Facility("alpha", "Demo Alpha"), Facility("beta", "Demo Beta")), "multi_site",
)
PERIOD = ReportPeriod(2026, 10)
ACTOR = "Synthetic Reviewer"
# These independent, invented values are deliberately unlike any client table.
SOURCE_TABLES = (
    (("Facility", "Available (Tons)", "Est.\nPeak Coincident\nLoad (Tons)", "Unit Price"),
     (("Demo Alpha", "300", "173", "$9300"),
      ("Demo Beta", "400", "296", "$11700"))),
    (("Facility", "Available (MBH)", "Est. Peak CoincidentLoad (MBH)", "Cost"),
     (("Demo Alpha", "5000", "4230", "$14100"),
      ("Demo Beta", "6000", "5680", "$15900"))),
)


@pytest.mark.parametrize("heading", ENGINEERING_HEADINGS)
def test_engineering_estimate_survives_price_free_copy(heading):
    spec = BlockSpec("thermal_capacity", "table", columns=(
        ColumnSpec("site", "Facility"), ColumnSpec("load", heading),
        ColumnSpec("price", "Unit Price"),
    ))
    block = ResolvedBlock("thermal_capacity", "This month",
                          rows=(("Demo Alpha", "173", "$9300"),))
    safe_spec, safe, removed = price_free_table(spec, block)
    assert not price_column(heading)
    assert tuple(column.title for column in safe_spec.columns) == ("Facility", heading)
    assert safe.rows == (("Demo Alpha", "173"),)
    assert removed == ("Unit Price",)
    assert block.rows == (("Demo Alpha", "173", "$9300"),)
    assert price_free_table(safe_spec, safe) == (safe_spec, safe, ())


@pytest.mark.parametrize("heading", (
    "Est.", "Estimate", "Repair estimate", "Est. Peak Coincident Load",
    "Est. Peak Coincident Load (USD)", "Est. Peak Coincident Load (dollars)",
    "Cost of Est. Peak Coincident Load (Tons)",
    "Est. Peak Coincident Load Cost (MBH)",
    "Est. Peak Coincident Load (Tons) / Unit Price",
    "Est. Peak Coincident Load (MBH) Amount",
    "Quote Peak Coincident Load (Tons)", "Amount Peak Coincident Load (MBH)",
    "Cost / tons", "Unit Price / kWh", "Fee / hour",
))
def test_load_units_do_not_waive_financial_headers(heading):
    assert price_column(heading)
    assert table_price_columns(("Facility", heading), (("Demo Alpha", "9300"),)) == (1,)


@pytest.mark.parametrize("heading", (
    "Description / Est. Peak Coincident Load (Tons) / Cost",
    "Est. Peak Coincident Load (MBH) / Recommendation",
))
def test_mixed_estimate_headers_still_require_review(heading):
    columns = ("Facility", heading)
    rows = (("Demo Alpha", "9300"),)
    assert ambiguous_price_columns(columns, rows) == (1,)
    assert not price_column(heading)
    assert table_has_pricing(columns, rows)


def test_currency_type_and_money_inside_an_engineering_column_still_block():
    heading = ENGINEERING_HEADINGS[0]
    assert price_column(heading, "currency")
    assert table_has_pricing(("Facility", heading), (("Demo Alpha", "$9300"),))
    assert contains_price("Repair estimate: (~3.2k)")
    assert contains_price("Est. 9300")


@pytest.mark.parametrize("unit", ("Tons", "MBH"))
@pytest.mark.parametrize("representation", ("extra_table", "schema_rows"))
def test_adjacent_header_lines_keep_load_through_copy_preflight_and_word(unit, representation):
    item = ImportItem("synthetic-table", "table", "word/document.xml", 1, 1,
        "scorecards", "thermal_capacity", rows=(
            ("Facility", "Est.", "Unit"),
            ("", f"Peak Coincident Load ({unit})", "Price"),
            ("Demo Alpha", "173", "$9300"),
        ))
    safe, removed = table_without_prices(item)
    assert removed == (2,)
    assert safe.columns == ("Facility", "Est.")
    assert safe.rows == (("", f"Peak Coincident Load ({unit})"), ("Demo Alpha", "173"))
    block = ResolvedBlock("thermal_capacity", "This month", extra_tables=(safe,))
    overrides = ()
    if representation == "schema_rows":
        block = replace(block, extra_tables=(), rows=safe.rows)
        overrides = (BlockSpec("thermal_capacity", "table", columns=(
            ColumnSpec("site", "Facility"), ColumnSpec("load", "Est."))),)
    draft = imported_draft(PROFILE, PERIOD, ACTOR, MappedImport((block,), overrides, (), ()))
    working, removed_again = client_table_draft(draft)
    assert working == draft and not removed_again
    assert not [check for check in preflight(working) if check.blocking]
    output = Document(BytesIO(assemble_docx(working, acknowledged_fingerprint=working.fingerprint)))
    assert any(tuple(tuple(cell.text for cell in row.cells) for row in table.rows) ==
               (safe.columns, *safe.rows) for table in output.tables)


def test_adjacent_engineering_header_cannot_waive_explicit_currency_type():
    spec = BlockSpec("thermal_capacity", "table", columns=(
        ColumnSpec("site", "Facility"), ColumnSpec("typed", "Est.", "currency"),
    ))
    block = ResolvedBlock("thermal_capacity", "This month", rows=(
        ("", "Peak Coincident Load (Tons)"), ("Demo Alpha", "173"),
    ))
    draft = imported_draft(PROFILE, PERIOD, ACTOR, MappedImport((block,), (spec,), (), ()))
    assert any(check.code == "pricing" and check.blocking for check in preflight(draft))
    safe_spec, safe, removed = price_free_table(spec, block)
    assert tuple(column.title for column in safe_spec.columns) == ("Facility",)
    assert safe.rows == (("",), ("Demo Alpha",))
    assert removed == ("Est.",)


@pytest.mark.parametrize("heading", ("Quote", "Amount", "Cost", "Unit Price"))
def test_other_header_continuations_do_not_waive_raw_financial_guards(heading):
    rows = (("", "Status"), ("Demo Alpha", "9300"))
    table = ReportTable(("Facility", heading), rows)
    block = ResolvedBlock("thermal_capacity", "This month", extra_tables=(table,))
    draft = imported_draft(PROFILE, PERIOD, ACTOR, MappedImport((block,), (), (), ()))
    assert any(check.code == "pricing" and check.blocking for check in preflight(draft))
    spec = BlockSpec("thermal_capacity", "table", columns=(
        ColumnSpec("site", "Facility"), ColumnSpec("amount", heading),
    ))
    safe_spec, safe, removed = price_free_table(spec, replace(block, extra_tables=(), rows=rows))
    assert tuple(column.title for column in safe_spec.columns) == ("Facility",)
    assert safe.rows == (("",), ("Demo Alpha",))
    assert removed == (heading,)


def _source_file(tmp_path):
    document = Document()
    document.add_heading("MONTHLY SCORECARDS", 1)
    for columns, rows in SOURCE_TABLES:
        table = document.add_table(rows=len(rows) + 1, cols=len(columns))
        for native_row, values in zip(table.rows, (columns, *rows)):
            for cell, value in zip(native_row.cells, values):
                cell.text = value
        for column, width in zip(table.columns, (1.7, 1.2, 2.3, 1.1)):
            column.width = Inches(width)
    path = tmp_path / "synthetic-capacity.docx"
    document.save(path)
    return path


def _import(path):
    inspection = inspect_docx(path)
    items = tuple(item for item in inspection.items if item.kind == "table")
    assert len(items) == 2 and all(item.section == "scorecards" for item in items)
    plans = ({"key": "scorecards", "target": "scorecards", "approved": True,
              "selected": [item.id for item in items]},)
    mapped, mappings = build_section_import(path, inspection, plans)
    profile = design_profile(PROFILE, inspection, mappings)
    return inspection, imported_draft(profile, PERIOD, ACTOR, mapped)


def _assert_engineering_tables(draft):
    block = next(block for block in draft.blocks if block.key == "thermal_capacity")
    assert len(block.extra_tables) == len(SOURCE_TABLES)
    for actual, (columns, rows) in zip(block.extra_tables, SOURCE_TABLES):
        assert tuple(" ".join(title.split()) for title in actual.columns) == tuple(
            " ".join(title.split()) for title in columns[:3])
        assert actual.rows == tuple(row[:3] for row in rows)


def _assert_native_tables(raw, path):
    output, source = Document(BytesIO(raw)), Document(path)
    assert len(output.tables) == len(source.tables) == 2
    for actual, original in zip(output.tables, source.tables):
        assert tuple(column.width for column in actual.columns) == tuple(
            column.width for column in original.columns)
        assert len(actual.rows) == len(original.rows)
        for actual_row, original_row in zip(actual.rows, original.rows):
            assert tuple(cell.text for cell in actual_row.cells) == tuple(
                cell.text if index != 3 else ""
                for index, cell in enumerate(original_row.cells))


def _save_import(draft, path, inspection, *, expected_revision=0, source_draft=None):
    return library.save_report_setup(draft.profile, draft, path,
        {"sha256": inspection.sha256, "period": PERIOD.key}, assets=(),
        expected_revision=expected_revision, actor=ACTOR, confirmed=True,
        source_draft=source_draft or draft)


def test_source_import_save_reload_and_native_word_keep_engineering_cells(monkeypatch, tmp_path):
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path / "data"))
    path = _source_file(tmp_path)
    original = path.read_bytes()
    inspection, draft = _import(path)
    _save_import(draft, path, inspection)
    saved = library.save_snapshot(draft, expected_revision=0, entered_editor=ACTOR)
    restored = library.load_snapshot(PROFILE.contract, PROFILE.key, PERIOD)
    assert restored == saved
    assert not [check for check in preflight(restored.draft) if check.blocking]
    raw = assemble_docx(restored.draft, acknowledged_fingerprint=restored.draft.fingerprint)
    # Keep first failing native output before checking the expected cells.
    (tmp_path / "generated.docx").write_bytes(raw)
    _assert_native_tables(raw, path)
    _assert_engineering_tables(restored.draft)
    assert library.load_imported_draft(PROFILE.contract, PROFILE.key) == draft
    assert path.read_bytes() == original
    assert library.imported_original(PROFILE.contract, PROFILE.key).read_bytes() == original


@requires_libreoffice
def test_actual_native_pdf_keeps_load_values_without_prices(monkeypatch, tmp_path):
    import fitz

    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path / "data"))
    path = _source_file(tmp_path)
    inspection, draft = _import(path)
    _save_import(draft, path, inspection)
    package = generate_report(draft, acknowledged_fingerprint=draft.fingerprint)
    assert package.docx and package.pdf, package.pdf_error
    (tmp_path / "generated.docx").write_bytes(package.docx)
    (tmp_path / "generated.pdf").write_bytes(package.pdf)
    with fitz.open(stream=package.pdf, filetype="pdf") as pdf:
        text = "\n".join(page.get_text() for page in pdf)
    for _, rows in SOURCE_TABLES:
        for row in rows:
            assert row[2] in text
            assert row[3] not in text
    assert "Unit Price" not in text and "Cost" not in text
    _assert_native_tables(package.docx, path)


def test_explicit_saved_source_reinterpretation_restores_only_unchanged_contribution(monkeypatch, tmp_path):
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path / "data"))
    path = _source_file(tmp_path)
    original = path.read_bytes()
    inspection, incoming = _import(path)
    block = next(block for block in incoming.blocks if block.key == "thermal_capacity")
    # Model an old acknowledged contribution that had lost the load columns.
    # The missing values are recovered only from this exact original source.
    legacy = replace(block, extra_tables=tuple(replace(table,
        columns=table.columns[:2], rows=tuple(row[:2] for row in table.rows))
        for table in block.extra_tables))
    previous = replace(incoming, blocks=(legacy,))
    setup = _save_import(previous, path, inspection)
    note = ResolvedBlock("utility_analysis", "This month", text="Synthetic reviewed note kept.")
    current = replace(previous, blocks=(legacy, note))
    before = library.save_snapshot(current, expected_revision=0, entered_editor=ACTOR)
    acknowledged, verified = library.starting_contribution(
        PROFILE.contract, PROFILE.key, PERIOD, inspection.sha256)
    assert verified and acknowledged == previous

    corrected = merge_starting_drafts(current, incoming, acknowledged, verified=verified)
    _save_import(corrected, path, inspection, expected_revision=setup.revision,
                 source_draft=incoming)
    after = library.save_snapshot(corrected, expected_revision=before.revision, entered_editor=ACTOR)
    restored = library.load_snapshot(PROFILE.contract, PROFILE.key, PERIOD)
    assert restored == after and restored.revision == 2
    _assert_engineering_tables(restored.draft)
    assert note in restored.draft.blocks
    assert library.load_snapshot_version(PROFILE.contract, PROFILE.key, PERIOD, 1) == before
    acknowledged_after, verified_after = library.starting_contribution(
        PROFILE.contract, PROFILE.key, PERIOD, inspection.sha256)
    assert verified_after and acknowledged_after == incoming
    assert merge_starting_drafts(corrected, incoming, acknowledged_after) == corrected
    raw = assemble_docx(restored.draft, acknowledged_fingerprint=restored.draft.fingerprint)
    _assert_native_tables(raw, path)
    assert path.read_bytes() == original
    assert library.imported_original(PROFILE.contract, PROFILE.key).read_bytes() == original

    # Neither a stale/unverified contribution nor an intervening manual edit
    # can authorize restoring source content over the existing report.
    with pytest.raises(ValueError, match="Thermal capacity differs"):
        merge_starting_drafts(current, incoming, acknowledged, verified=False)
    changed_table = replace(legacy.extra_tables[0], rows=(("Demo Alpha", "999"), *legacy.extra_tables[0].rows[1:]))
    changed = replace(current, blocks=(replace(legacy,
        extra_tables=(changed_table, *legacy.extra_tables[1:])), note))
    with pytest.raises(ValueError, match="Thermal capacity differs"):
        merge_starting_drafts(changed, incoming, acknowledged)
