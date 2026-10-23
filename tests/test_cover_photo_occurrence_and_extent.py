"""Synthetic authority and displayed-frame boundaries for unchanged cover repair."""
from dataclasses import replace
from copy import deepcopy
from io import BytesIO
from zipfile import ZipFile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches
from lxml import etree
import pytest

from app.monthly_report_import import ImportMapping, inspect_docx, map_items
from app.monthly_report_model import BlockSpec, SectionSpec
from app.monthly_report_native_layout import build_native_docx
from test_cover_unchanged_photo_geometry import source_photo, photo_geometry


def geometry(picture):
    crop = picture.find('./'+qn('pic:blipFill')+'/'+qn('a:srcRect'))
    properties = picture.find(qn('pic:spPr'))
    frame = next((n for n in picture.iterancestors() if n.tag in (qn('wp:inline'), qn('wp:anchor'))), None)
    outer = frame.find(qn('wp:extent')) if frame is not None else None
    return (dict(crop.attrib) if crop is not None else None,
            [(n.tag, dict(n.attrib)) for n in properties.iter()] if properties is not None else None,
            dict(outer.attrib) if outer is not None else None)


def repeated_source(tmp_path, same_paragraph, selected):
    path, draft, _, photograph = source_photo(tmp_path, {})
    doc = Document(path)
    first = next(doc.element.iter(qn('wp:inline')))
    second = doc.add_picture(BytesIO(photograph), width=Inches(3), height=Inches(3))._inline
    crop = OxmlElement('a:srcRect'); crop.attrib.update({'l': '10000', 'r': '20000', 't': '5000', 'b': '10000'})
    next(second.iter(qn('a:blip'))).addnext(crop)
    if same_paragraph:
        run = second.getparent().getparent(); paragraph = run.getparent()
        doc.element.body.remove(paragraph)
        first.getparent().getparent().getparent().append(run)
    doc.add_heading('MONTHLY ACTIVITY SUMMARY', 1)
    third = doc.add_picture(BytesIO(photograph), width=Inches(2), height=Inches(3))
    crop = OxmlElement('a:srcRect'); crop.attrib.update({'t': '20000', 'b': '10000'})
    next(third._inline.iter(qn('a:blip'))).addnext(crop)
    doc.save(path)
    inspection = inspect_docx(path)
    images = [i for i in inspection.items if i.kind == 'image']
    assert len(images) == 3 and len({i.id for i in images}) == 3
    assert len({i.image_part for i in images}) == 1
    mapped = map_items(path, inspection, (ImportMapping(images[selected].id, 'cover_photo'), ImportMapping(images[2].id, 'improvements')))
    sections = (SectionSpec('activity', '2', 'Monthly Activity Summary', (BlockSpec('improvements', 'image_grid'),)),)
    return path, replace(draft, sections=sections, blocks=mapped.blocks), dict(mapped.assets), photograph


@pytest.mark.parametrize('same_paragraph', [False, True])
@pytest.mark.parametrize('selected', [0, 1])
@pytest.mark.parametrize('section', [None, 'cover'])
def test_only_selected_occurrence_is_recropped_when_media_is_reused(tmp_path, same_paragraph, selected, section):
    path, draft, assets, photograph = repeated_source(tmp_path, same_paragraph, selected)
    original = path.read_bytes()
    before = list(Document(path).element.iter(qn('pic:pic')))
    result = build_native_docx(draft, path, asset_loader=assets.__getitem__, section_key=section)
    after = list(Document(BytesIO(result)).element.iter(qn('pic:pic')))
    assert len(after) == (3 if section is None else 2)
    for index, picture in enumerate(after):
        if index != selected:
            assert geometry(picture) == geometry(before[index]), 'Same bytes do not authorize recropping another occurrence'
        else:
            crop, _, extent = geometry(picture)
            remaining = {k: 1-sum(float(crop.get(a, 0))/100000 for a in axes)
                         for k, axes in [('width', ('l', 'r')), ('height', ('t', 'b'))]}
            assert 400*remaining['width']/(300*remaining['height']) == pytest.approx(int(extent['cx'])/int(extent['cy']), rel=.001)
    assert path.read_bytes() == original
    with ZipFile(BytesIO(result)) as package:
        assert photograph in [package.read(n) for n in package.namelist() if n.startswith('word/media/')]


@pytest.mark.parametrize('missing_proof', ['no-reference', 'unbound-occurrence'])
def test_missing_exact_occurrence_proof_retains_legacy_geometry(tmp_path, monkeypatch, missing_proof):
    from app import monthly_report_native_layout as native
    path, draft, assets, _ = repeated_source(tmp_path, True, 0)
    if missing_proof == 'no-reference':
        draft = replace(draft, blocks=tuple(replace(b, references=()) if b.key == 'cover_photo' else b for b in draft.blocks))
    else:
        inspection = inspect_docx(path)
        images = [i for i in inspection.items if i.kind == 'image']
        inspection = replace(inspection, items=tuple(i for i in inspection.items if i.id != images[1].id))
        monkeypatch.setattr(native, 'inspect_docx', lambda _: inspection)
    before = [geometry(p) for p in Document(path).element.iter(qn('pic:pic'))]
    result = build_native_docx(draft, path, asset_loader=assets.__getitem__)
    assert [geometry(p) for p in Document(BytesIO(result)).element.iter(qn('pic:pic'))] == before


@pytest.mark.parametrize('selected', [0, 1])
def test_exact_photo_role_is_not_overwritten_by_same_media_client_logo(tmp_path, selected):
    path, draft, _, _ = repeated_source(tmp_path, True, selected)
    inspection = inspect_docx(path)
    images = [i for i in inspection.items if i.kind == 'image']
    mapped = map_items(path, inspection, (ImportMapping(images[selected].id, 'cover_photo'),
                                        ImportMapping(images[1-selected].id, 'client_logo'),
                                        ImportMapping(images[2].id, 'improvements')))
    before = list(Document(path).element.iter(qn('pic:pic')))
    result = build_native_docx(replace(draft, blocks=mapped.blocks), path, asset_loader=dict(mapped.assets).__getitem__)
    after = list(Document(BytesIO(result)).element.iter(qn('pic:pic')))
    # A separately selected logo now receives its own contain policy, never the photo crop.
    logo_crop, _, logo_extent = geometry(after[1-selected])
    if logo_crop is None:
        assert int(logo_extent["cx"]) / int(logo_extent["cy"]) == pytest.approx(400 / 300)
    else:
        assert geometry(after[1-selected]) == geometry(before[1-selected])
    assert geometry(after[2]) == geometry(before[2])
    crop, _, extent = geometry(after[selected])
    ratio = 400*(1-sum(float(crop.get(k, 0))/100000 for k in ('l', 'r')))
    ratio /= 300*(1-sum(float(crop.get(k, 0))/100000 for k in ('t', 'b')))
    assert ratio == pytest.approx(int(extent['cx'])/int(extent['cy']), rel=.001)


def test_conflicting_roles_for_same_occurrence_do_not_authorize_recrop(tmp_path):
    path, draft, assets, _ = source_photo(tmp_path, {})
    photo = next(b for b in draft.blocks if b.key == 'cover_photo')
    draft = replace(draft, blocks=draft.blocks+(replace(photo, key='client_logo'),))
    before = geometry(next(Document(path).element.iter(qn('pic:pic'))))
    result = build_native_docx(draft, path, asset_loader=assets.__getitem__)
    assert geometry(next(Document(BytesIO(result)).element.iter(qn('pic:pic')))) == before


@pytest.mark.parametrize('condition', ['missing-transform', 'missing-inner', 'missing-outer', 'mismatched', 'invalid-inner', 'invalid-outer'])
@pytest.mark.parametrize('section', [None, 'cover'])
def test_unproved_displayed_extent_is_unchanged_and_still_exports(tmp_path, condition, section):
    path, draft, _, photograph = source_photo(tmp_path, {'t': 10000, 'b': 20000})
    doc = Document(path)
    picture = next(doc.element.iter(qn('pic:pic')))
    transform = picture.find('./'+qn('pic:spPr')+'/'+qn('a:xfrm'))
    inner = transform.find(qn('a:ext'))
    outer = next(doc.element.iter(qn('wp:extent')))
    if condition == 'missing-transform': transform.getparent().remove(transform)
    elif condition == 'missing-inner': transform.remove(inner)
    elif condition == 'missing-outer': outer.getparent().remove(outer)
    elif condition == 'mismatched': outer.set('cy', str(Inches(3)))
    elif condition == 'invalid-inner': inner.set('cx', 'oops')
    else: outer.set('cy', 'NaN')
    doc.save(path)
    before = geometry(next(Document(path).element.iter(qn('pic:pic'))))
    inspection = inspect_docx(path)
    item = next(i for i in inspection.items if i.kind == 'image')
    mapped = map_items(path, inspection, (ImportMapping(item.id, 'cover_photo'),))
    result = build_native_docx(replace(draft, blocks=mapped.blocks), path,
                              asset_loader=dict(mapped.assets).__getitem__, section_key=section)
    assert geometry(next(Document(BytesIO(result)).element.iter(qn('pic:pic')))) == before
    with ZipFile(BytesIO(result)) as package:
        assert photograph in [package.read(n) for n in package.namelist() if n.startswith('word/media/')]


@pytest.mark.parametrize('zero', ['0', '0.0', '0%', '0.00%', '-0', '00', '+0.0', '-0.000%'])
def test_equivalent_numeric_zero_destination_insets_still_get_proportional_fit(tmp_path, zero):
    path, draft, _, _ = source_photo(tmp_path, {})
    doc = Document(path)
    rectangle = next(doc.element.iter(qn('a:fillRect')))
    for edge in ('l', 't', 'r', 'b'): rectangle.set(edge, zero)
    doc.save(path)
    inspection = inspect_docx(path)
    item = next(i for i in inspection.items if i.kind == 'image')
    mapped = map_items(path, inspection, (ImportMapping(item.id, 'cover_photo'),))
    result = build_native_docx(replace(draft, blocks=mapped.blocks), path, asset_loader=dict(mapped.assets).__getitem__)
    crop, frame = photo_geometry(result)
    assert 400*(1-crop['l']-crop['r'])/(300*(1-crop['t']-crop['b'])) == pytest.approx(frame[0]/frame[1], rel=.001)
    assert dict(next(Document(BytesIO(result)).element.iter(qn('a:fillRect'))).attrib) == dict(rectangle.attrib)


@pytest.mark.parametrize('fallback_kind', ['modern', 'vml'])
@pytest.mark.parametrize('unproved', [False, True])
def test_selected_alternate_occurrence_is_atomic_without_authorizing_independent_sibling(tmp_path, fallback_kind, unproved):
    path, draft, _, photograph = source_photo(tmp_path, {})
    doc = Document(path)
    drawing = next(doc.element.iter(qn('w:drawing')))
    mc = '{http://schemas.openxmlformats.org/markup-compatibility/2006}'
    alternate = etree.Element(mc+'AlternateContent', nsmap={'pic': qn('pic:pic').split('}')[0][1:]})
    drawing.getparent().replace(drawing, alternate)
    etree.SubElement(alternate, mc+'Choice', Requires='pic').append(drawing)
    fallback = etree.SubElement(alternate, mc+'Fallback')
    if fallback_kind == 'modern': fallback.append(deepcopy(drawing))
    else:
        pict = etree.SubElement(fallback, qn('w:pict'))
        shape = etree.SubElement(pict, '{urn:schemas-microsoft-com:vml}shape', style='width:4in;height:2in')
        etree.SubElement(shape, '{urn:schemas-microsoft-com:vml}imagedata', {qn('r:id'): next(drawing.iter(qn('a:blip'))).get(qn('r:embed'))})
    if unproved: next(drawing.iter(qn('wp:extent'))).set('cy', str(Inches(3)))
    doc.add_picture(BytesIO(photograph), width=Inches(3), height=Inches(3))
    doc.save(path)
    inspection = inspect_docx(path)
    images = [i for i in inspection.items if i.kind == 'image']
    assert len(images) == 2
    mapped = map_items(path, inspection, (ImportMapping(images[0].id, 'cover_photo'),))
    before = Document(path)
    result = build_native_docx(replace(draft, blocks=mapped.blocks), path, asset_loader=dict(mapped.assets).__getitem__)
    after = Document(BytesIO(result))
    assert geometry(list(after.element.iter(qn('pic:pic')))[-1]) == geometry(list(before.element.iter(qn('pic:pic')))[-1])
    previous = next(before.element.iter(mc+'AlternateContent'))
    current = next(after.element.iter(mc+'AlternateContent'))
    for old, new in zip(previous.iter(qn('pic:pic')), current.iter(qn('pic:pic'))):
        if unproved: assert geometry(new) == geometry(old)
        else:
            crop = geometry(new)[0]
            assert float(crop['t'])/100000 == pytest.approx(1/6, abs=.00001)
    vml = '{urn:schemas-microsoft-com:vml}imagedata'
    for old, new in zip(previous.iter(vml), current.iter(vml)):
        if unproved: assert dict(new.attrib) == dict(old.attrib)
        else: assert float(new.get('croptop')) == pytest.approx(1/6, abs=.00001)


def test_header_shared_media_retains_intake_gate_and_separate_occurrence_bindings(tmp_path):
    from app.monthly_report_import import ImportError
    from app.monthly_report_native_layout import _header_image_occurrences
    path, draft, _, photograph = source_photo(tmp_path, {})
    doc = Document(path)
    header = doc.sections[0].header.paragraphs[0].add_run().add_picture(BytesIO(photograph), width=Inches(3), height=Inches(1))
    crop = OxmlElement('a:srcRect'); crop.attrib.update({'l': '10000', 'r': '30000'})
    next(header._inline.iter(qn('a:blip'))).addnext(crop)
    doc.save(path)
    inspection = inspect_docx(path)
    body_item = next(i for i in inspection.items if i.kind == 'image' and i.part == 'word/document.xml')
    header_item = next(i for i in inspection.items if i.kind == 'image' and i.part.startswith('word/header'))
    assert body_item.image_part == header_item.image_part
    body_bindings = _header_image_occurrences(doc.part, inspection, elements=doc.element.body)
    header_bindings = _header_image_occurrences(doc.sections[0].header.part, inspection)
    assert {item.id for _, _, item in body_bindings} == {body_item.id}
    assert {item.id for _, _, item in header_bindings} == {header_item.id}
    assert not {node for node, _, _ in body_bindings} & {node for node, _, _ in header_bindings}
    with pytest.raises(ImportError, match='Header artwork has no confirmed photograph role'):
        map_items(path, inspection, (ImportMapping(body_item.id, 'cover_photo'),))


@pytest.mark.parametrize('unsupported', ['nonzero', 'malformed', 'nonfinite', 'rotated', 'flipped', 'nonrectangular', 'custom'])
def test_unproved_drawing_geometry_preserves_original_crop(tmp_path, unsupported):
    from app.monthly_report_native_layout import _fit_cover_photo
    path, _, _, _ = source_photo(tmp_path, {})
    doc = Document(path)
    picture = next(doc.element.iter(qn('pic:pic')))
    properties = picture.find(qn('pic:spPr'))
    if unsupported in ('nonzero', 'malformed', 'nonfinite'):
        next(picture.iter(qn('a:fillRect'))).set('b', {'nonzero': '1%', 'malformed': '0%%', 'nonfinite': 'NaN'}[unsupported])
    elif unsupported == 'rotated': properties.find(qn('a:xfrm')).set('rot', '5400000')
    elif unsupported == 'flipped': properties.find(qn('a:xfrm')).set('flipH', 'true')
    elif unsupported == 'nonrectangular': properties.find(qn('a:prstGeom')).set('prst', 'ellipse')
    else: etree.SubElement(properties, qn('a:custGeom'))
    before = doc.element.xml
    _fit_cover_photo(next(picture.iter(qn('a:blip'))), 400, 300, preserve_crop=True)
    assert doc.element.xml == before


@pytest.mark.parametrize('unsupported', ['rotation', 'flip', 'skew', 'path'])
def test_unproved_vml_transforms_remain_unchanged(unsupported):
    from app.monthly_report_native_layout import _fit_cover_photo
    v = '{urn:schemas-microsoft-com:vml}'
    shape = etree.Element(v+'shape', style='width:4in;height:2in')
    image = etree.SubElement(shape, v+'imagedata')
    if unsupported == 'rotation': shape.set('style', shape.get('style')+';rotation:30')
    elif unsupported == 'flip': shape.set('style', shape.get('style')+';flip:x')
    elif unsupported == 'skew': etree.SubElement(shape, v+'skew', on='t', matrix='1,0,0.5,1,0,0')
    else: shape.set('path', 'm0,0l400,0,200,200xe')
    before = etree.tostring(shape)
    _fit_cover_photo(image, 400, 300, preserve_crop=True)
    assert etree.tostring(shape) == before


@pytest.mark.parametrize('style', ['', 'height:2in', 'width:4in', 'width:50%;height:2in',
                                 'width:400;height:200', 'width:4pc;height:2pc',
                                 'width:NaNpt;height:2in', 'width:infpt;height:2in',
                                 'width:0in;height:2in', 'width:-4in;height:2in',
                                 'width:4.0.0in;height:2in', 'Width:4in;height:2in'])
def test_unproved_standalone_vml_dimensions_preserve_original(style):
    from app.monthly_report_native_layout import _fit_cover_photo
    v = '{urn:schemas-microsoft-com:vml}'
    shape = etree.Element(v+'shape', style=style)
    image = etree.SubElement(shape, v+'imagedata', croptop='.1')
    before = etree.tostring(shape)
    _fit_cover_photo(image, 400, 300, preserve_crop=True)
    assert etree.tostring(shape) == before


@pytest.mark.parametrize('style', ['width:4in;height:2in', 'width:288pt;height:144pt',
                                 'width:10.16cm;height:50.8mm', 'width:384px;height:192px'])
def test_proved_absolute_vml_units_remain_supported(style):
    from app.monthly_report_native_layout import _fit_cover_photo
    v = '{urn:schemas-microsoft-com:vml}'
    shape = etree.Element(v+'shape', style=style)
    image = etree.SubElement(shape, v+'imagedata')
    _fit_cover_photo(image, 400, 300, preserve_crop=True)
    assert float(image.get('croptop')) == pytest.approx(1/6, abs=.00001)
    assert shape.get('style') == style


@pytest.mark.parametrize('reference', ['#customShape', '#_x0000_t75'])
def test_unresolved_vml_shapetype_reference_is_not_assumed_rectangular(reference):
    from app.monthly_report_native_layout import _fit_cover_photo
    v = '{urn:schemas-microsoft-com:vml}'
    pict = etree.Element(qn('w:pict'))
    etree.SubElement(pict, v+'shapetype', id=reference[1:], path='m0,0l400,0,200,200xe')
    shape = etree.SubElement(pict, v+'shape', style='width:4in;height:2in', type=reference)
    image = etree.SubElement(shape, v+'imagedata')
    before = etree.tostring(pict)
    _fit_cover_photo(image, 400, 300, preserve_crop=True)
    assert etree.tostring(pict) == before


@pytest.mark.parametrize('kind', ['drawingml', 'vml'])
@pytest.mark.parametrize('reverse', [False, True])
def test_finite_dimensions_with_unproved_derived_ratio_are_unchanged(tmp_path, kind, reverse):
    from app.monthly_report_native_layout import _fit_cover_photo
    large, small = '1'+'0'*300, '0.'+'0'*299+'1'
    width, height = (small, large) if reverse else (large, small)
    if kind == 'drawingml':
        path, _, _, _ = source_photo(tmp_path, {})
        doc = Document(path)
        root = doc.element
        for extent in list(root.iter(qn('a:ext'))) + list(root.iter(qn('wp:extent'))):
            extent.set('cx', width); extent.set('cy', height)
        image = next(root.iter(qn('a:blip')))
    else:
        v = '{urn:schemas-microsoft-com:vml}'
        root = etree.Element(v+'shape', style=f'width:{width}pt;height:{height}pt')
        image = etree.SubElement(root, v+'imagedata')
    before = etree.tostring(root)
    _fit_cover_photo(image, 400, 300, preserve_crop=True)
    assert etree.tostring(root) == before
