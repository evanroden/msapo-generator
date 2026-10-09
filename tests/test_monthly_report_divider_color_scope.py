"""Global style changes require evidence for every rendered source use."""
from copy import deepcopy
from zipfile import ZipFile, ZIP_DEFLATED

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from lxml import etree

from app.monthly_report_render_profile import _divider_title_evidence
from test_monthly_report_render_profile import color_pair


def rewrite(source, change, extra=None):
    with ZipFile(source) as archive:
        entries = {item.filename: archive.read(item) for item in archive.infolist()}
    root = etree.fromstring(entries['word/document.xml'])
    change(root)
    entries['word/document.xml'] = etree.tostring(root)
    entries.update(extra or {})
    with ZipFile(source, 'w', ZIP_DEFLATED) as archive:
        for name, raw in entries.items():
            archive.writestr(name, raw)


@pytest.mark.parametrize('part', ['header1.xml', 'footer1.xml', 'footnotes.xml', 'endnotes.xml'])
@pytest.mark.parametrize('derived', [False, True])
def test_other_story_affected_text_disables_global_color_bridge(tmp_path, part, derived):
    source, _ = color_pair(tmp_path)
    assert _divider_title_evidence(source) == ['monthlyscorecards']
    selected = 'Divider'
    if derived:
        document = Document(source)
        style = document.styles.add_style('InheritedDivider', WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = document.styles['Divider']
        document.save(source)
        selected = 'InheritedDivider'
    story = etree.Element(qn('w:hdr'))
    paragraph = etree.SubElement(story, qn('w:p'))
    props = etree.SubElement(paragraph, qn('w:pPr'))
    etree.SubElement(props, qn('w:pStyle')).set(qn('w:val'), selected)
    run = etree.SubElement(paragraph, qn('w:r'))
    etree.SubElement(run, qn('w:t')).text = 'Unproved contact information'
    rewrite(source, lambda _: None, {'word/' + part: etree.tostring(story)})
    assert _divider_title_evidence(source) == []


@pytest.mark.parametrize('placement', ['independent', 'same_branch', 'alternatives', 'different_alternatives'])
def test_title_occurrence_proof_distinguishes_fallback_from_live_duplicate(tmp_path, placement):
    source, _ = color_pair(tmp_path)
    def change(root):
        pict = next(root.iter(qn('w:pict')))
        parent = pict.getparent()
        if placement == 'independent':
            parent.append(deepcopy(pict))
            return
        mc = '{http://schemas.openxmlformats.org/markup-compatibility/2006}'
        alternate = etree.SubElement(parent, mc + 'AlternateContent')
        choice = etree.SubElement(alternate, mc + 'Choice')
        choice.append(pict)
        fallback = etree.SubElement(alternate, mc + 'Fallback')
        copy = deepcopy(pict)
        (choice if placement == 'same_branch' else fallback).append(copy)
        if placement == 'different_alternatives':
            next(copy.iter(qn('w:t'))).text = 'TRAINING SUMMARY'
    rewrite(source, change)
    expected = ['monthlyscorecards'] if placement == 'alternatives' else []
    assert _divider_title_evidence(source) == expected


def test_inherited_character_style_inside_hyperlink_is_not_missed(tmp_path):
    source, _ = color_pair(tmp_path)
    document = Document(source)
    style = document.styles.add_style('InheritedCharacterDivider', WD_STYLE_TYPE.CHARACTER)
    base = etree.SubElement(style.element, qn('w:basedOn'))
    base.set(qn('w:val'), 'Divider')
    paragraph = document.add_paragraph()
    link = etree.SubElement(paragraph._p, qn('w:hyperlink'))
    run = etree.SubElement(link, qn('w:r'))
    props = etree.SubElement(run, qn('w:rPr'))
    etree.SubElement(props, qn('w:rStyle')).set(qn('w:val'), style.style_id)
    etree.SubElement(run, qn('w:t')).text = 'Unproved linked text'
    document.save(source)
    assert _divider_title_evidence(source) == []
