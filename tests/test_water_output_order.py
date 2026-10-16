"""The Water browser gate checks actual DOCX occurrences and caption placement."""

import ast
from hashlib import sha256
from io import BytesIO
from pathlib import Path
import posixpath
from zipfile import ZipFile, ZIP_DEFLATED

from docx import Document
from docx.shared import Inches
from lxml import etree
from PIL import Image
import pytest


def verifier():
    # Required tests do not install browser engines/Playwright. Execute the
    # actual pure output assertion, with only its standard-library globals.
    source = Path(__file__).resolve().parents[1] / "scripts/qa_water_page_journey.py"
    tree = ast.parse(source.read_text())
    fn = next(n for n in tree.body if isinstance(n, ast.FunctionDef) and n.name == "assert_docx_water_pages")
    namespace = {"sha256": sha256, "BytesIO": BytesIO, "ZipFile": ZipFile, "posixpath": posixpath}
    exec(compile(ast.Module(body=[fn], type_ignores=[]), str(source), "exec"), namespace)
    return namespace[fn.name]


def png(color):
    stream = BytesIO()
    Image.new("RGB", (24, 36), color).save(stream, "PNG")
    return stream.getvalue()


IMAGES = {1: png("red"), 2: png("blue")}
CAPTIONS = {1: "Synthetic conductivity check; follow-up remains pending.",
            2: "Invented water source.pdf · page/item 2"}


def expected(captions=CAPTIONS):
    return [{"asset": sha256(IMAGES[number]).hexdigest() + ".png",
             "source_pages": [["synthetic-source", number]],
             "caption_sha256": sha256(captions[number].encode()).hexdigest()}
            for number in (2, 1)]


def output(order=(2, 1), captions=CAPTIONS, *, extra_caption=""):
    document = Document()
    document.add_picture(BytesIO(png("green")), width=Inches(1))
    document.add_heading("7. Water Treatment Reports", level=1)
    for number in order:
        document.add_picture(BytesIO(IMAGES[number]), width=Inches(1))
        if captions[number]:
            document.add_paragraph(captions[number], style="Caption")
        document.add_page_break()
    document.add_heading("8. Equipment Performance Issues", level=1)
    if extra_caption:
        document.add_paragraph(extra_caption)
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def rewrite(raw, mutate):
    with ZipFile(BytesIO(raw)) as original:
        parts = {name: original.read(name) for name in original.namelist()}
    mutate(parts)
    stream = BytesIO()
    with ZipFile(stream, "w", ZIP_DEFLATED) as changed:
        for name, data in parts.items():
            changed.writestr(name, data)
    return stream.getvalue()


def test_real_inline_occurrences_keep_working_order_and_adjacent_captions():
    proof = verifier()(output(), expected())
    assert [p["asset_sha256"] for p in proof] == [p["asset"].split(".")[0] for p in expected()]
    assert [p["caption_sha256"] for p in proof] == [p["caption_sha256"] for p in expected()]
    assert len({p["relationship"] for p in proof}) == 2
    assert all(p["part"].startswith("word/media/") for p in proof)


def test_both_media_present_does_not_make_wrong_document_order_pass():
    wrong = output(order=(1, 2))
    with ZipFile(BytesIO(wrong)) as package:
        media = {sha256(package.read(n)).hexdigest() for n in package.namelist() if n.startswith("word/media/")}
        assert all(p["asset"].split(".")[0] in media for p in expected())
    with pytest.raises(AssertionError, match="order/count"):
        verifier()(wrong, expected())


def test_relationship_rebinding_is_checked_instead_of_media_filename_order():
    def swap(parts):
        rels = etree.fromstring(parts["word/_rels/document.xml.rels"])
        images = [r for r in rels if r.get("Type").endswith("/image")]
        images[-2].set("Target", "media/image3.png")
        images[-1].set("Target", "media/image2.png")
        parts["word/_rels/document.xml.rels"] = etree.tostring(rels)
    with pytest.raises(AssertionError, match="order/count"):
        verifier()(rewrite(output(), swap), expected())


def test_duplicate_picture_occurrence_cannot_hide_behind_deduplicated_media():
    with pytest.raises(AssertionError, match="order/count"):
        verifier()(output(order=(2, 1, 1)), expected())


def test_unreferenced_package_media_cannot_replace_a_missing_inline_page():
    def unused(parts):
        parts["word/media/unused-water-page.png"] = IMAGES[1]
    with pytest.raises(AssertionError, match="order/count"):
        verifier()(rewrite(output(order=(2,)), unused), expected())


@pytest.mark.parametrize("captions", ({1: CAPTIONS[2], 2: CAPTIONS[1]},
                                      {1: "", 2: CAPTIONS[2]}))
def test_caption_must_follow_its_own_image_not_only_exist_elsewhere(captions):
    with pytest.raises(AssertionError, match="caption is missing, changed or attached"):
        verifier()(output(captions=captions), expected())


def test_duplicate_caption_elsewhere_cannot_hide_behind_one_correct_copy():
    with pytest.raises(AssertionError, match="caption occurrence count"):
        verifier()(output(extra_caption=CAPTIONS[1]), expected())


def test_explicit_blank_caption_needs_no_paragraph_or_filename_fallback():
    blank = {**CAPTIONS, 1: ""}
    proof = verifier()(output(captions=blank), expected(blank))
    assert proof[1]["caption_sha256"] == sha256(b"").hexdigest()
    with pytest.raises(AssertionError, match="caption is missing, changed or attached"):
        verifier()(output(captions={**CAPTIONS, 1: "Invented water source.pdf · page/item 1"}), expected(blank))


def test_caption_text_keeps_native_line_breaks_and_tabs():
    multiline = {**CAPTIONS, 1: "Invented finding\nPending\tsample"}
    verifier()(output(captions=multiline), expected(multiline))


@pytest.mark.parametrize("invalid", ("missing", "external", "floating"))
def test_missing_external_or_floating_pictures_are_not_inline_page_proof(invalid):
    def change(parts):
        if invalid == "floating":
            xml = etree.fromstring(parts["word/document.xml"])
            inline = xml.xpath('//*[local-name()="inline"]')[-1]
            inline.tag = inline.tag.replace("inline", "anchor")
            parts["word/document.xml"] = etree.tostring(xml)
        else:
            rels = etree.fromstring(parts["word/_rels/document.xml.rels"])
            image = [r for r in rels if r.get("Type").endswith("/image")][-1]
            image.set("Target", "media/missing.png" if invalid == "missing" else "https://example.invalid/water.png")
            if invalid == "external":
                image.set("TargetMode", "External")
            parts["word/_rels/document.xml.rels"] = etree.tostring(rels)
    with pytest.raises(AssertionError, match="missing|internal|inline"):
        verifier()(rewrite(output(), change), expected())
