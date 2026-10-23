"""Catalog replacements fit existing authorized header/footer occurrences."""
from dataclasses import replace
from io import BytesIO

import pytest
from PIL import Image, ImageDraw
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches

from app import monthly_report_library as library
from app.monthly_report_model import ReportDraft, ReportPeriod, ResolvedBlock, synthetic_profiles
from app.monthly_report_native_layout import build_native_docx


def png(size, alpha, color):
    image = Image.new('RGBA', size, (255, 255, 255, 0 if alpha else 255))
    draw = ImageDraw.Draw(image)
    cx, cy = size[0] // 2, size[1] // 2
    radius = min(size) // 3
    draw.ellipse((cx-radius, cy-radius, cx+radius, cy+radius), fill=color)
    out = BytesIO(); image.save(out, 'PNG')
    return out.getvalue()


def fixture(tmp_path, size, alpha, role='brand_logo'):
    old = png((600, 100), False, 'red')
    new = png(size, alpha, 'blue')
    doc = Document(); doc.add_paragraph('Synthetic facility'); doc.add_paragraph('October 2026')
    label = 'ENFRA company logo' if role == 'brand_logo' else 'Client logo'
    doc.add_picture(BytesIO(old), width=Inches(3))._inline.docPr.set('descr', label)
    for part in (doc.sections[0].header, doc.sections[0].footer):
        part.paragraphs[0].add_run().add_picture(BytesIO(old), width=Inches(2))._inline.docPr.set('descr', label)
    path = tmp_path / 'synthetic-master.docx'; doc.save(path)
    reference = library.asset_reference(new, 'png')
    block = ResolvedBlock(role, 'Library', asset_hashes=(reference,))
    draft = ReportDraft(synthetic_profiles()[0], ReportPeriod(2026, 10), 'Synthetic reviewer', (), (block,))
    return path, draft, {reference: new}, new


def header_pictures(raw):
    doc = Document(BytesIO(raw))
    for part in doc.part.package.parts:
        if not str(part.partname).startswith(('/word/header', '/word/footer')):
            continue
        for picture in part.element.iter(qn('pic:pic')):
            outer = next(node for node in picture.iterancestors() if node.tag in (qn('wp:inline'), qn('wp:anchor')))
            ext = outer.find(qn('wp:extent'))
            inner = picture.find('./'+qn('pic:spPr')+'/'+qn('a:xfrm')+'/'+qn('a:ext'))
            assert dict(ext.attrib) == dict(inner.attrib)
            blip = next(picture.iter(qn('a:blip')))
            yield part.rels[blip.get(qn('r:embed'))].target_part.blob, int(ext.get('cx')), int(ext.get('cy'))


@pytest.mark.parametrize('size', [(500, 100), (100, 500), (200, 200)])
@pytest.mark.parametrize('alpha', [False, True])
@pytest.mark.parametrize('section', [None, 'cover'])
@pytest.mark.parametrize('role', ['brand_logo', 'client_logo'])
def test_catalog_replacement_without_refs_fits_authorized_header_and_footer(tmp_path, size, alpha, section, role):
    path, draft, assets, new = fixture(tmp_path, size, alpha, role)
    source = path.read_bytes()
    assert not draft.blocks[0].references
    raw = build_native_docx(draft, path, asset_loader=assets.__getitem__, master=True, section_key=section)
    pictures = list(header_pictures(raw))
    assert len(pictures) == 2
    for actual, width, height in pictures:
        assert actual == new
        assert width / height == pytest.approx(size[0] / size[1], rel=1e-5)
        assert width <= Inches(2) and height <= Inches(1/3)
    assert path.read_bytes() == source


def test_unchanged_unmapped_header_artwork_keeps_native_geometry(tmp_path):
    path, draft, assets, _ = fixture(tmp_path, (600, 100), False)
    doc = Document(path)
    old = next(doc.sections[0].header.part.element.iter(qn('a:blip')))
    original = doc.sections[0].header.part.rels[old.get(qn('r:embed'))].target_part.blob
    ref = library.asset_reference(original, 'png')
    draft = replace(draft, blocks=(replace(draft.blocks[0], asset_hashes=(ref,)),))
    result = build_native_docx(draft, path, asset_loader={ref: original}.__getitem__, master=True)
    assert [(width, height) for _, width, height in header_pictures(result)] == [(Inches(2), Inches(1/3))] * 2


from tests.conftest import requires_libreoffice


@requires_libreoffice
@pytest.mark.parametrize('role', ['brand_logo', 'client_logo'])
@pytest.mark.parametrize('size', [(500, 100), (100, 500), (200, 200)])
def test_replaced_header_and_footer_circles_render_round(tmp_path, monkeypatch, role, size):
    import fitz
    from PIL import ImageChops
    from app import pdf_converter
    from app.monthly_report_render_compat import rendering_docx
    monkeypatch.setenv('EPC_DATA_DIR', str(tmp_path / 'store'))
    monkeypatch.setattr(pdf_converter, 'OUTPUT_DIR', tmp_path)
    path, draft, assets, _ = fixture(tmp_path, size, True, role)
    raw = build_native_docx(draft, path, asset_loader=assets.__getitem__, master=True, section_key='cover')
    current = tmp_path / 'current.docx'; current.write_bytes(rendering_docx(raw))
    with fitz.open(pdf_converter.convert_to_pdf(current)) as pdf:
        assert len(pdf) == 1
        pixels = pdf[0].get_pixmap(matrix=fitz.Matrix(3, 3), alpha=False)
        image = Image.frombytes('RGB', (pixels.width, pixels.height), pixels.samples)
    red, green, blue = image.split()
    mask = ImageChops.multiply(Image.eval(red, lambda x: 255 if x < 10 else 0),
                              Image.eval(green, lambda x: 255 if x < 10 else 0))
    mask = ImageChops.multiply(mask, Image.eval(blue, lambda x: 255 if x > 245 else 0))
    rows = [i for i, present in enumerate(mask.getprojection()[1]) if present]
    bands = []
    for row in rows:
        if not bands or row > bands[-1][-1] + 1:
            bands.append([row])
        else:
            bands[-1].append(row)
    assert len(bands) >= 2
    for band in (bands[0], bands[-1]):
        bbox = mask.crop((0, band[0], image.width, band[-1] + 1)).getbbox()
        width, height = bbox[2] - bbox[0], bbox[3] - bbox[1]
        assert abs(width - height) <= 2


@pytest.mark.parametrize('unsupported', ['rotation', 'flip', 'tile'])
def test_default_header_replacement_preserves_unsupported_geometry(tmp_path, unsupported):
    from docx.oxml import OxmlElement
    path, draft, assets, new = fixture(tmp_path, (100, 500), True)
    doc = Document(path)
    for part in (doc.sections[0].header.part, doc.sections[0].footer.part):
        picture = next(part.element.iter(qn('pic:pic')))
        xfrm = picture.find('./'+qn('pic:spPr')+'/'+qn('a:xfrm'))
        if unsupported == 'rotation':
            xfrm.set('rot', '900000')
        elif unsupported == 'flip':
            xfrm.set('flipH', '1')
        else:
            fill = picture.find(qn('pic:blipFill'))
            fill.remove(fill.find(qn('a:stretch')))
            fill.append(OxmlElement('a:tile'))
    doc.save(path)
    raw = build_native_docx(draft, path, asset_loader=assets.__getitem__, master=True)
    pictures = list(header_pictures(raw))
    assert len(pictures) == 2
    assert all(actual == new for actual, _, _ in pictures)
    assert [(width, height) for _, width, height in pictures] == [(Inches(2), Inches(1/3))] * 2


@pytest.mark.parametrize('role', ['brand_logo', 'client_logo'])
def test_raw_identical_unmapped_catalog_logo_keeps_valid_crop(tmp_path, role):
    from PIL.PngImagePlugin import PngInfo
    from docx.oxml import OxmlElement
    path, draft, _, _ = fixture(tmp_path, (500, 100), True, role)
    doc = Document(path)
    source_part = next(part for part in doc.part.package.parts if str(part.partname).startswith('/word/media/'))
    image = Image.open(BytesIO(source_part.blob))
    metadata = PngInfo(); metadata.add_text('Synthetic fixture', 'Raw source metadata is not a normalized image identity')
    encoded = BytesIO(); image.save(encoded, 'PNG', pnginfo=metadata)
    original = encoded.getvalue(); source_part._blob = original
    for part in (doc.sections[0].header.part, doc.sections[0].footer.part):
        picture = next(part.element.iter(qn('pic:pic')))
        fill = picture.find(qn('pic:blipFill'))
        crop = OxmlElement('a:srcRect'); crop.set('l', '10000'); crop.set('r', '10000')
        fill.insert(1, crop)
        outer = next(node for node in picture.iterancestors() if node.tag == qn('wp:inline'))
        for ext in (outer.find(qn('wp:extent')), picture.find('./'+qn('pic:spPr')+'/'+qn('a:xfrm')+'/'+qn('a:ext'))):
            ext.set('cy', '381000')
    doc.save(path)
    ref = library.asset_reference(original, 'png')
    draft = replace(draft, blocks=(replace(draft.blocks[0], asset_hashes=(ref,)),))
    raw = build_native_docx(draft, path, asset_loader={ref: original}.__getitem__, master=True)
    pictures = list(header_pictures(raw))
    assert len(pictures) == 2
    assert all(actual == original and (width, height) == (1828800, 381000) for actual, width, height in pictures)
    generated = Document(BytesIO(raw))
    for part in (generated.sections[0].header.part, generated.sections[0].footer.part):
        crop = next(part.element.iter(qn('a:srcRect')))
        assert dict(crop.attrib) == {'l': '10000', 'r': '10000'}
