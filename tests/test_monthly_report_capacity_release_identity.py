"""Persisted safety cases for the combined #323/#324 capacity release (Refs #286)."""
from io import BytesIO

from docx import Document
from openpyxl import Workbook
import pytest

from app import monthly_report_capacity as capacity
from app.monthly_report_docx import assemble_docx
from app.monthly_report_model import ReportDraft, ReportPeriod, default_sections
from test_monthly_report_capacity import profile as capacity_profile


GENERIC = ("Facility", "Required capacity", "Available capacity")
STEAM = ("Facility", "Required steam (lb/hr)", "Available steam (lb/hr)")
CHILLED = ("Facility", "Required cooling (tons)", "Available cooling (tons)")


@pytest.fixture
def profile(tmp_path, monkeypatch):
    return capacity_profile.__wrapped__(tmp_path, monkeypatch)


def workbook(sheets):
    book = Workbook()
    book.remove(book.active)
    for title, rows in sheets:
        sheet = book.create_sheet(title)
        for row in rows:
            sheet.append(row)
    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def word(sheets):
    document = Document()
    for title, rows in sheets:
        document.add_paragraph(title)
        table = document.add_table(rows=len(rows), cols=len(rows[0]))
        for target, row in zip(table.rows, rows):
            for cell, value in zip(target.cells, row):
                cell.text = str(value)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def save(profile, raw, revision=0, *, filename="capacity.xlsx", mode="merge"):
    tables = capacity.inspect_capacity(profile, filename, raw).tables
    return capacity.save_capacity(
        profile, tables, expected_revision=revision, actor="Synthetic Reviewer",
        reviewed_fingerprint=capacity.review_fingerprint(tables), mode=mode)


def rows(state, site):
    return sorted((table.title, row) for table in state.tables
                  for row, key in zip(table.rows, table.site_keys) if key == site)


def test_named_same_schema_replace_keeps_the_other_service_and_its_history(profile):
    raw = workbook([
        ("Steam", [GENERIC, ("NH", 40000, 50000), ("SH", 20000, 24000)]),
        ("Chilled water", [GENERIC, ("NH", 1200, 1600), ("SH", 800, 1000)]),
    ])
    original = save(profile, raw)
    updated = workbook([("Steam", [GENERIC, ("NH", 42000, 52000)])])
    save(profile, updated, 1, mode="replace_sites")

    reloaded = capacity.load_capacity(profile.contract)
    assert rows(reloaded, "north") == [
        ("Chilled water", ("NH", "1200", "1600")),
        ("Steam", ("NH", "42000", "52000")),
    ]
    assert rows(reloaded, "south") == rows(original, "south")
    assert capacity.load_capacity(profile.contract, 1) == original
    assert (capacity._root(profile.contract) / "sources" /
            (original.tables[0].source_sha256 + ".xlsx")).read_bytes() == raw


def test_another_services_schema_does_not_hide_a_named_layout_change(profile):
    original = save(profile, workbook([
        ("Steam", [STEAM, ("NH", 40000, 50000), ("SH", 20000, 24000)]),
        ("Chilled water", [GENERIC, ("NH", 1200, 1600), ("SH", 800, 1000)]),
    ]))
    changed = workbook([("Steam", [GENERIC, ("NH", 42000, 52000)])])
    tables = capacity.inspect_capacity(profile, "capacity.xlsx", changed).tables
    assert capacity.schema_conflicts(original, tables) == ("north",)
    with pytest.raises(ValueError, match="different columns"):
        save(profile, changed, 1)
    assert capacity.load_capacity(profile.contract) == original

    save(profile, changed, 1, mode="replace_sites")
    reloaded = capacity.load_capacity(profile.contract)
    assert rows(reloaded, "north") == [
        ("Chilled water", ("NH", "1200", "1600")),
        ("Steam", ("NH", "42000", "52000")),
    ]
    assert rows(reloaded, "south") == rows(original, "south")
    assert capacity.load_capacity(profile.contract, 1) == original


def test_untitled_matching_shape_replace_keeps_other_untitled_shapes(profile):
    original = save(profile, workbook([
        ("Sheet1", [STEAM, ("NH", 40000, 50000)]),
        ("Sheet2", [CHILLED, ("NH", 1200, 1600)]),
    ]))
    changed = workbook([("Sheet1", [STEAM, ("NH", 42000, 52000)])])
    save(profile, changed, 1, mode="replace_sites")
    reloaded = capacity.load_capacity(profile.contract)
    assert rows(reloaded, "north") == [
        ("Sheet1", ("NH", "42000", "52000")),
        ("Sheet2", ("NH", "1200", "1600")),
    ]
    assert capacity.load_capacity(profile.contract, 1) == original


def test_untitled_new_shape_cannot_destructively_choose_between_saved_shapes(profile):
    original = save(profile, workbook([
        ("Sheet1", [STEAM, ("NH", 40000, 50000)]),
        ("Sheet2", [CHILLED, ("NH", 1200, 1600)]),
    ]))
    changed = workbook([("Sheet1", [GENERIC, ("NH", 42000, 52000)])])
    with pytest.raises(ValueError, match="heading"):
        save(profile, changed, 1, mode="replace_sites")
    assert capacity.load_capacity(profile.contract) == original
    assert not (capacity._root(profile.contract) / "history" / "00000002.json").exists()

    kept = save(profile, changed, 1, mode="keep_both")
    assert all(table in kept.tables for table in original.tables)
    assert len(kept.tables) == 3


@pytest.mark.parametrize("mode", ["merge", "replace_sites"])
def test_word_partial_service_update_keeps_stable_titles_and_other_services(profile, mode):
    original = save(profile, word([
        ("Steam", [GENERIC, ("NH", 40000, 50000), ("SH", 20000, 24000)]),
        ("Chilled water", [GENERIC, ("NH", 1200, 1600), ("SH", 800, 1000)]),
    ]), filename="capacity.docx")
    changed = word([("Chilled water", [GENERIC, ("NH", 1250, 1650)])])
    save(profile, changed, 1, filename="capacity.docx", mode=mode)

    reloaded = capacity.load_capacity(profile.contract)
    assert rows(reloaded, "north") == [
        ("Chilled water", ("NH", "1250", "1650")),
        ("Steam", ("NH", "40000", "50000")),
    ]
    assert rows(reloaded, "south") == [
        ("Chilled water", ("SH", "800", "1000")),
        ("Steam", ("SH", "20000", "24000")),
    ]
    assert capacity.load_capacity(profile.contract, 1) == original
    report = capacity.capacity_block(profile, reloaded)
    assert sorted(row for table in report.extra_tables for row in table.rows) == [
        ("Chilled water", "NH", "1250", "1650"),
        ("Steam", "NH", "40000", "50000"),
    ]
    sections = tuple(section for section in default_sections() if section.key == "scorecards")
    draft = ReportDraft(profile, ReportPeriod(2026, 10), "Synthetic Reviewer", sections, (report,))
    document = Document(BytesIO(assemble_docx(draft, acknowledged_fingerprint=draft.fingerprint)))
    assert sorted(tuple(cell.text for cell in row.cells)
                  for table in document.tables for row in table.rows[1:]) == [
        ("Chilled water", "NH", "1250", "1650"),
        ("Steam", "NH", "40000", "50000"),
    ]
    assert all(tuple(cell.text for cell in table.rows[0].cells) == ("Source table", *GENERIC)
               for table in document.tables)


def test_word_table_site_context_does_not_include_the_preceding_tables_site(profile):
    columns = ("Service", "Required capacity", "Available capacity")
    raw = word([
        ("North Hospital steam", [columns, ("Steam", 40000, 50000)]),
        ("South Hospital chilled water", [columns, ("Chilled water", 800, 1000)]),
    ])
    save(profile, raw, filename="capacity.docx")
    reloaded = capacity.load_capacity(profile.contract)
    assert rows(reloaded, "north") == [
        ("North Hospital steam", ("Steam", "40000", "50000")),
    ]
    assert rows(reloaded, "south") == [
        ("South Hospital chilled water", ("Chilled water", "800", "1000")),
    ]
