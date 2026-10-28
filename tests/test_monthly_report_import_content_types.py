"""Synthetic import capability guidance never changes accepted source content."""
from dataclasses import replace
from io import BytesIO

from docx import Document
from docx.oxml import parse_xml
from PIL import Image
from streamlit.testing.v1 import AppTest

from app import monthly_report_section_ui as ui
from app.monthly_report_import import inspect_docx
from app.monthly_report_sections import section_reviews, build_section_import


def _image(color="navy"):
    raw = BytesIO()
    Image.new("RGB", (480, 320), color).save(raw, "PNG")
    raw.name = "synthetic-picture.png"
    return raw


def _source(tmp_path, *, frame=False):
    doc = Document()
    doc.add_heading("2 Monthly Activity Summary", 1)
    doc.add_paragraph("Synthetic pump inspection completed.")
    table = doc.add_table(rows=2, cols=2)
    for cell, value in zip(table.rows[0].cells, ("Task", "Count")):
        cell.text = value
    for cell, value in zip(table.rows[1].cells, ("Inspections", "3")):
        cell.text = value
    doc.add_picture(_image())
    if frame:
        doc.add_paragraph()._p.append(parse_xml('''
<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
     xmlns:v="urn:schemas-microsoft-com:vml"><w:pict>
<v:shape id="synthetic-frame" style="width:120pt;height:40pt"><v:textbox>
<w:txbxContent><w:p><w:r><w:t>Floating synthetic label</w:t></w:r></w:p></w:txbxContent>
</v:textbox></v:shape></w:pict></w:r>'''))
    path = tmp_path / "synthetic-content-types.docx"
    doc.save(path)
    inspected = inspect_docx(path)
    section = next(s for s in section_reviews(inspected) if s.key == "activity")
    return path, inspected, section


APP = '''
import streamlit as st
from app.monthly_report_section_ui import _section_card
from app.monthly_report_ui import _field
from app.monthly_report_model import ReportPeriod
plans = st.session_state.setdefault("plans", {})
_section_card(st.session_state["section"], st.session_state["path"],
              st.session_state["inspection"], ReportPeriod(2026, 9),
              "content_type", _field, plans, True)
'''


def _app(tmp_path, monkeypatch, *, frame=False):
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path / "runtime"))
    path, inspection, section = _source(tmp_path, frame=frame)
    app = AppTest.from_string(APP, default_timeout=30)
    app.session_state["path"] = path
    app.session_state["inspection"] = inspection
    app.session_state["section"] = section
    app.run()
    assert not app.exception
    return app, path, inspection, section


def test_real_word_section_distinguishes_editable_table_from_image(tmp_path, monkeypatch):
    app, path, inspection, section = _app(tmp_path, monkeypatch)
    original = path.read_bytes()
    summary = [c.value for c in app.caption if c.value.startswith("Content types:")]
    assert summary == ["Content types: Editable text: 1 · Editable tables: 1 · Image-only pictures: 1"]
    assert any("not editable" in c.value and "Leave out" in c.value for c in app.caption)
    # The shared pre-review entry point runs before the native table/picture controls.
    flattened = list(app)
    summary_index = next(n for n, e in enumerate(flattened)
                         if e.type == "caption" and e.value == summary[0])
    assert summary_index < next(n for n, e in enumerate(flattened) if e.type == "dataframe")
    assert len(app.dataframe) == 1
    assert any(r.label == "Use this picture" for r in app.radio)
    plan = app.session_state["plans"]["activity"]
    assert plan["approved"]
    before = {key: plan[key] for key in ("selected", "texts", "tables", "picture_roles")}
    app.run()
    assert not app.exception
    assert {key: app.session_state["plans"]["activity"][key] for key in before} == before
    assert path.read_bytes() == original
    assert inspect_docx(path) == inspection


def test_floating_frame_text_is_not_advertised_as_editable(tmp_path, monkeypatch):
    app, path, inspection, section = _app(tmp_path, monkeypatch, frame=True)
    assert any(item.kind == "unsupported" for item in section.items)
    summary = next(c.value for c in app.caption if c.value.startswith("Content types:"))
    assert "Editable text: 1" in summary and "Editable text: 2" not in summary
    assert "Layout review needed: 1" in summary
    assert any("extracted labels" in c.value for c in app.caption)
    assert not app.session_state["plans"]["activity"]["approved"]
    assert any("could not be preserved" in e.value for e in app.error)


def test_picture_has_source_purpose_in_native_caption(tmp_path, monkeypatch):
    app, _, _, section = _app(tmp_path, monkeypatch)
    def captions():
        return [image.caption for element in app.get("image") for image in element.proto.imgs]
    assert "Picture 1 · Imported picture in Monthly Activity Summary · Image-only" in captions()
    # Existing extracted caption text can provide a more specific purpose; the
    # UI must not run OCR or invent one when this metadata is absent.
    app.session_state["section"] = replace(section, items=tuple(
        replace(i, text="Synthetic pump inspection chart") if i.kind == "image" else i
        for i in section.items))
    app.run()
    assert not app.exception
    assert "Picture 1 · Synthetic pump inspection chart · Image-only" in captions()


def test_replacement_keeps_adjacent_editable_text_and_table(tmp_path, monkeypatch):
    app, path, inspection, section = _app(tmp_path, monkeypatch)
    original = path.read_bytes()
    picture = next(i for i in section.items if i.kind == "image")
    next(r for r in app.radio if r.label == "Use this picture").set_value("keep").run()
    assert not app.exception
    before = app.session_state["plans"]["activity"]
    assert picture.id in before["selected"]
    old_text, old_tables = dict(before["texts"]), dict(before["tables"])
    action = next(r for r in app.radio if r.label == "What would you like to do?")
    action.set_value("Edit text or change pictures").run()
    assert not app.exception
    next(r for r in app.radio if r.label == "Use this picture").set_value("").run()
    replacement = _image("green")
    real_upload = ui.st.file_uploader
    monkeypatch.setattr(ui.st, "file_uploader", lambda label, *a, **kw:
                        replacement if label == "New improvement photo" else real_upload(label, *a, **kw))
    app.run()
    assert not app.exception
    plan = app.session_state["plans"]["activity"]
    assert plan["texts"] == old_text and plan["tables"] == old_tables
    assert plan["approved"]
    assert picture.id not in plan["selected"]
    assert len(plan["new_assets"]) == 1
    mapped, _ = build_section_import(path, inspection, (plan,))
    blocks = {b.key: b for b in mapped.blocks}
    assert blocks["activity_summary"].text == "Synthetic pump inspection completed."
    assert any(t.rows == (("Inspections", "3"),) for b in mapped.blocks for t in b.extra_tables)
    assert len(blocks["improvements"].asset_hashes) == 1
    app.run()
    assert not app.exception
    assert app.session_state["plans"]["activity"]["texts"] == old_text
    assert app.session_state["plans"]["activity"]["tables"] == old_tables
    assert path.read_bytes() == original


def test_inventory_uses_source_location_and_ignores_unused_package_art():
    from app.monthly_report_import import ImportItem
    from app.monthly_report_sections import ImportSection
    from app.monthly_report_import_content_types import content_type_summary

    items = (
        ImportItem("text", "text", "word/document.xml", 1, 1, text="Editable"),
        ImportItem("table", "table", "word/document.xml", 2, 1),
        ImportItem("unresolved", "unsupported", "word/document.xml", 2, 1),
        ImportItem("duplicate-drawing", "unsupported", "word/document.xml", 2, 1),
        # Same position in a different part is not part of the blocked table.
        ImportItem("footer", "text", "word/footer1.xml", 2, 1, text="Footer"),
        ImportItem("unused-image", "image", "word/media/image1.png", 0, 0),
        ImportItem("unused-object", "unsupported", "word/embeddings/object1.bin", 0, 0),
        ImportItem("native-design", "native_layout", "word/document.xml", 3, 1),
    )
    section = ImportSection("activity", "Activity", items)
    assert content_type_summary(section) == "Editable text: 2 · Layout review needed: 1"
    assert section.items is items


def test_cover_only_advertises_its_actual_editable_footer():
    from app.monthly_report_import import ImportItem
    from app.monthly_report_sections import ImportSection
    from app.monthly_report_import_content_types import content_type_summary

    items = (
        ImportItem("old-title", "text", "word/document.xml", 1, 1, text="Old month"),
        ImportItem("old-table", "table", "word/document.xml", 2, 1),
        ImportItem("footer", "text", "word/footer1.xml", 1, 1,
                   suggested_slot="footer_text", text="Address"),
        ImportItem("photo", "image", "word/document.xml", 3, 1),
    )
    assert content_type_summary(ImportSection("cover", "Cover", items)) == (
        "Editable text: 1 · Image-only pictures: 1")
    assert content_type_summary(ImportSection("activity", "Activity", ())) == ""


def test_descriptions_are_bounded_literal_and_never_become_source_content():
    from app.monthly_report_import import ImportItem
    from app.monthly_report_import_content_types import image_review_caption, image_replacement_guidance
    from app.monthly_report_ui_text import literal

    authored = "[Synthetic](https://example.invalid)\n<b>chart</b>"
    item = ImportItem("picture", "image", "word/document.xml", 1, 1, text=authored)
    assert image_review_caption(item, 2, "activity") == (
        "Picture 2 · " + literal(" ".join(authored.split())) + " · Image-only")
    assert item.text == authored
    long = replace(item, text="x" * 200)
    assert "x" * 177 + "…" in image_review_caption(long, 1, "activity")
    assert "x" * 178 not in image_review_caption(long, 1, "activity")
    assert "choose a destination" in image_replacement_guidance("other").lower()
    assert "chart or contact page" in image_replacement_guidance("organization")
    assert "OCR" not in image_review_caption(replace(item, text=""), 1, "other")


def test_complete_chart_page_caption_and_decisions_remain_explicit(tmp_path, monkeypatch):
    from app import monthly_report_word_pages_ui as page_ui
    from app.monthly_report_import import DocxInspection, ImportItem
    from app.monthly_report_sections import ImportSection
    from app.monthly_report_word_pages import WordPage

    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path / "runtime"))
    raw = _image().getvalue()
    monkeypatch.setattr(page_ui, "preserve_word_pages", lambda *a:
                        (WordPage(1, raw, 480, 320, "Synthetic organization chart", False),))
    items = (ImportItem("frame", "unsupported", "word/document.xml", 1, 1, "organization"),)
    app = AppTest.from_string(APP, default_timeout=30)
    app.session_state["path"] = tmp_path / "unused.docx"
    app.session_state["inspection"] = DocxInspection("synthetic", items, (), (), 1)
    app.session_state["section"] = ImportSection("organization", "Organization", items)
    app.run()
    assert not app.exception
    next(b for b in app.button if b.label == "Show original chart pages").click().run()
    next(c for c in app.checkbox if c.label == "Use complete pages for this section").check().run()
    assert not app.exception
    captions = [i.caption for e in app.get("image") for i in e.proto.imgs]
    assert "Original chart/contact page 1 · Image-only" in captions
    plan = app.session_state["plans"]["organization"]
    assert not plan["approved"]
    assert not plan["preserved_assets"]
    assert next(c for c in app.checkbox if c.label == "Include page 1").value is False
