"""Numeric source text must not become a different value merely to open its row."""
from dataclasses import replace
from datetime import date
from decimal import Decimal
from io import BytesIO
from hashlib import sha256
import json
import os
from pathlib import Path

from docx import Document
import pytest
from streamlit.testing.v1 import AppTest

from app import monthly_report_library as library
from app.monthly_report_docx import assemble_docx, generate_report
from tests.conftest import requires_libreoffice
from app.monthly_report_model import BlockSpec, ColumnSpec, ResolvedBlock
from app.monthly_report_row_cards_ui import _field_default
from test_monthly_report_row_cards_ui import SCRIPT, button, field


LOSSY_VALUES = (
    "9007199254740993", "-9007199254740993", "9007199254740992", "1e20",
    "0.1234567890123456789", "0.100000000000000005", "  +1.2345678901234567890  ",
    "1e-500", "-1e-500", "2.4703282292062327e-324",
)


@pytest.fixture(autouse=True)
def isolated(tmp_path, monkeypatch):
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path))


def workspace(value):
    spec = BlockSpec("proposals", "table", columns=(
        ColumnSpec("task", "Task"), ColumnSpec("reading", "Reading", "number")))
    block = ResolvedBlock("proposals", "This month", rows=(
        ("Synthetic reading", value), ("Unaffected reading", "2.5")),
        references=("synthetic-source-occurrence",))
    app = AppTest.from_string(SCRIPT.replace("INITIAL", repr(block)).replace("SPEC", repr(spec))).run()
    assert not app.exception
    return app, spec, block


@pytest.mark.parametrize("kind", ("number", "currency"))
@pytest.mark.parametrize("value", LOSSY_VALUES)
def test_source_values_outside_lossless_numeric_input_remain_exact_text(value, kind):
    result = _field_default(value, kind)
    assert isinstance(result, str)
    assert result == value


@pytest.mark.parametrize("value", ("0", "-0", "0.10", "001.2500", "-3.5", "1e3",
                                     "9007199254740991", "-9007199254740991"))
def test_representable_source_values_keep_existing_numeric_defaults(value):
    result = _field_default(value, "number")
    assert isinstance(result, float)
    assert Decimal(str(result)) == Decimal(value)


def test_blank_unknown_date_and_text_defaults_are_unchanged():
    assert _field_default("", "number") is None
    assert _field_default("NR", "number") == "NR"
    assert _field_default("2026-09-12", "date") == date(2026, 9, 12)
    assert _field_default("Unknown", "date") == "Unknown"
    assert _field_default("9007199254740993", "text") == "9007199254740993"


@pytest.mark.parametrize("value", ("9007199254740993", "0.1234567890123456789", "1e-500"))
def test_exact_source_survives_sibling_edit_close_reopen_and_literal_correction(value):
    app, spec, original = workspace(value)
    button(app, "Edit proposal").click().run()
    assert not app.exception
    control = field(app, "Reading")
    identity = control.proto.id
    assert control.value == value and not app.number_input
    # A sibling edit delivered with Close must not adopt a rounded numeric default.
    field(app, "Task").set_value("Edited description")
    button(app, "Close editing").click().run()
    assert not app.exception
    edited = app.session_state["block"]
    assert edited.rows == (("Edited description", value), original.rows[1])
    assert edited.references == original.references
    button(app, "Edit proposal").click().run()
    assert field(app, "Reading").value == value
    assert field(app, "Reading").proto.id == identity
    field(app, "Reading").set_value("0.123456789012345678901").run()
    assert field(app, "Reading").proto.id == identity and not app.number_input
    field(app, "Reading").set_value("").run()
    assert field(app, "Reading").proto.id == identity
    field(app, "Reading").set_value("3.25")
    button(app, "Close editing").click().run()
    corrected = app.session_state["block"]
    assert corrected.rows == (("Edited description", "3.25"), original.rows[1])
    assert corrected.references == original.references

    draft = app.session_state["draft"]
    section = next(s for s in draft.sections if s.key == "proposals")
    draft = replace(draft, sections=(replace(section, blocks=(spec,)),), blocks=(edited,))
    library.save_snapshot(draft, expected_revision=0, entered_editor=draft.prepared_by)
    before = library.load_snapshot(draft.profile.contract, draft.profile.key, draft.period)
    assert before.draft == draft
    newer = replace(draft, blocks=(corrected,))
    library.save_snapshot(newer, expected_revision=before.revision, entered_editor=draft.prepared_by)
    assert library.load_snapshot(draft.profile.contract, draft.profile.key, draft.period).draft == newer
    assert library.load_snapshot_version(
        draft.profile.contract, draft.profile.key, draft.period, before.revision).draft == draft
    with pytest.raises(library.RevisionConflict):
        library.save_snapshot(draft, expected_revision=before.revision, entered_editor=draft.prepared_by)
    assert library.load_snapshot(draft.profile.contract, draft.profile.key, draft.period).draft == newer
    assert before.draft.blocks == (edited,)

    reopened, _, _ = workspace(value)
    reopened.session_state["block"] = corrected
    reopened.run()
    button(reopened, "Edit proposal").click().run()
    assert field(reopened, "Reading", "number_input").value == 3.25
    # Opening the untouched saved copy in a new session still exposes exact text.
    prior, _, _ = workspace(value)
    prior.session_state["block"] = before.draft.blocks[0]
    prior.run()
    button(prior, "Edit proposal").click().run()
    assert field(prior, "Reading").value == value and not prior.number_input
    doc = Document(BytesIO(assemble_docx(before.draft)))
    output = [tuple(cell.text for cell in row.cells) for table in doc.tables for row in table.rows]
    assert output.count(("Edited description", value)) == 1
    assert output.count(original.rows[1]) == 1


def test_exact_text_edit_retains_existing_same_row_conflict_recovery():
    app, _, original = workspace("9007199254740993")
    button(app, "Edit proposal").click().run()
    field(app, "Reading").set_value("9007199254740997")
    newer = replace(original, rows=(("Newer description", "9007199254740995"), original.rows[1]))
    app.session_state["block"] = newer
    app.run()
    assert not app.exception and app.session_state["block"] == newer
    assert any("9007199254740997" in str(item.value.values) for item in app.dataframe)
    button(app, "Keep current table").click().run()
    assert not app.exception and app.session_state["block"] == newer


@requires_libreoffice
def test_saved_exact_readings_reach_actual_word_and_pdf_once(tmp_path):
    import pymupdf

    values = ("9007199254740993", "0.1234567890123456789", "1e-500")
    app, spec, original = workspace(values[0])
    original = replace(original, rows=tuple((f"Recorded reading {n}", value)
                                           for n, value in enumerate(values, 1)))
    app.session_state["block"] = original
    app.run()
    for index, value in enumerate(values):
        button(app, "Edit proposal", index).click().run()
        assert not app.exception and field(app, "Reading").value == value
        field(app, "Task").set_value(f"Reviewed reading {index + 1}")
        button(app, "Close editing").click().run()
        assert not app.exception
    draft = app.session_state["draft"]
    section = next(s for s in draft.sections if s.key == "proposals")
    draft = replace(draft, sections=(replace(section, blocks=(spec,)),),
                    blocks=(app.session_state["block"],))
    library.save_snapshot(draft, expected_revision=0, entered_editor=draft.prepared_by)
    restored = library.load_snapshot(draft.profile.contract, draft.profile.key, draft.period)
    assert restored.draft == draft
    package = generate_report(restored.draft, acknowledged_fingerprint=restored.draft.fingerprint)
    assert package.docx and package.pdf, package.pdf_error
    document = Document(BytesIO(package.docx))
    rows = [tuple(c.text for c in row.cells) for table in document.tables for row in table.rows]
    for index, value in enumerate(values, 1):
        assert rows.count((f"Reviewed reading {index}", value)) == 1
    out = Path(os.environ.get("INDEPENDENT_OUTPUT", str(tmp_path))) / "numeric-row-precision"
    out.mkdir(parents=True, exist_ok=True)
    (out / "exact-readings.docx").write_bytes(package.docx)
    (out / "exact-readings.pdf").write_bytes(package.pdf)
    with pymupdf.open(stream=package.pdf, filetype="pdf") as pdf:
        text = " ".join(" ".join(page.get_text().split()) for page in pdf)
        for value in values:
            assert text.count(value) == 1
        for index, page in enumerate(pdf, 1):
            page.get_pixmap(matrix=pymupdf.Matrix(1.5, 1.5)).save(out / f"page-{index}.png")
        pages = len(pdf)
    (out / "manifest.json").write_text(json.dumps({
        "source_sha": os.environ.get("MSAPO_SHA"), "fixture": "synthetic exact numeric row values",
        "docx_sha256": sha256(package.docx).hexdigest(), "pdf_sha256": sha256(package.pdf).hexdigest(),
        "pages": pages, "values": values, "saved_revision": restored.revision,
        "renderer": "Local LibreOffice; not native Microsoft Word or served acceptance",
    }, indent=2))
    assert library.load_snapshot(draft.profile.contract, draft.profile.key, draft.period) == restored
