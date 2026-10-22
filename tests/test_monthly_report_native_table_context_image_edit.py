"""Retained table artwork does not prove that current table cells are unchanged."""
from dataclasses import replace
from hashlib import sha256
from io import BytesIO

import pytest
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches
from PIL import Image

from app.monthly_report_import import _text, inspect_docx, read_import_image
from app.monthly_report_native_layout import build_native_docx
from app.monthly_report_sections import table_without_prices
from test_monthly_report_native_table_context import (
    SCOPE, TITLE, TRAINING, before, scenario,
)


def mixed_context(tmp_path, change):
    path, draft = scenario(tmp_path)
    source = Document(path)
    image = BytesIO()
    Image.new("RGB", (24, 24), "blue").save(image, "PNG")
    source.tables[-1].cell(1, 1).paragraphs[0].add_run().add_picture(
        BytesIO(image.getvalue()), width=Inches(.3))
    source.save(path)
    inspection = inspect_docx(path)
    tables = {}
    assets, image_refs = {}, []
    for item in inspection.items:
        reference = f"docx:{inspection.sha256}:{item.id}"
        if item.kind == "table":
            table, removed = table_without_prices(item)
            assert not removed
            tables.setdefault(item.section, []).append(replace(table, reference=reference))
        elif item.kind == "image" and item.section == "training":
            imported = read_import_image(path, item, line_art=True)
            assets[sha256(imported.data).hexdigest() + "." + imported.extension] = imported.data
            image_refs.append(reference)
    assert len(assets) == len(image_refs) == 1
    proposals, training = draft.blocks
    table = tables["training"][0]
    if change == "row":
        table = replace(table, rows=(("Revised synthetic separator", *table.rows[0][1:]),))
    elif change == "column":
        table = replace(table, columns=("Current description", *table.columns[1:]))
    training = replace(training, extra_tables=(table,), asset_hashes=tuple(assets),
                       references=tuple(image_refs))
    proposals = replace(proposals, extra_tables=tuple(tables["proposals"]))
    return path, replace(draft, blocks=(proposals, training)), assets


@pytest.mark.parametrize("section", [None, "training"])
@pytest.mark.parametrize("change", [None, "row", "column"])
def test_mixed_table_context_requires_exact_current_cell_values(tmp_path, section, change):
    path, draft, assets = mixed_context(tmp_path, change)
    source_bytes = path.read_bytes()
    source_table = Document(path).tables[-1]
    output = Document(BytesIO(build_native_docx(
        draft, path, asset_loader=assets.__getitem__, section_key=section)))
    current_table = output.tables[-1]
    assert len(list(current_table._tbl.iter(qn("a:blip")))) == 1
    assert current_table._tbl.tblGrid.xml == source_table._tbl.tblGrid.xml
    assert current_table.cell(0, 0).text == (
        "Current description" if change == "column" else "Description")
    assert current_table.cell(1, 0).text == (
        "Revised synthetic separator" if change == "row" else "Synthetic air separator")
    text = _text(output.element)
    assert text.count(TITLE) == text.count(SCOPE) == 1
    if change is None:
        # Exact current cells and current artwork may preserve adjacent context.
        assert [_text(p) for p in before(current_table._tbl)[-2:]] == [TITLE, SCOPE]
    else:
        # A successful in-place image-table edit is only a placement receipt.
        # The current words stay together in authored order, not in old context.
        lines = [p.text for p in output.paragraphs if p.text.strip()]
        start = lines.index(TRAINING)
        assert lines[start:start + 3] == [TRAINING, TITLE, SCOPE]
        assert not any(_text(p) in (TITLE, SCOPE) for p in before(current_table._tbl))
    assert path.read_bytes() == source_bytes
