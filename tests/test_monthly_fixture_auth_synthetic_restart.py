"""Synthetic AUTH identity and process-reopen gates; NEVER authentic/Render proof.

All DOCX bytes, storage and approval expectations exist only in pytest tmp_path.
A fresh Python interpreter reopens the immutable shared fixture after an isolated
report run is reset. This is local volume simulation, NOT a Render deployment.
"""
from dataclasses import replace
from hashlib import sha256
from io import BytesIO
from pathlib import Path
import subprocess
import sys

from docx import Document
import pytest

from app import memory
from app import monthly_report_fixture_runs as fixture_runs
from app import monthly_report_testing as testing
from app.monthly_report_fixture_vault import ExpectedSection, FixtureMetadata, ReportIdentity
from app.monthly_report_model import ReportPeriod
from app.monthly_report_shared_fixtures import SharedFixtureLibrary
from app.monthly_report_shared_testing_ui import target_profile


_RUN_TOKEN = "6" * 32
_CONTRACT = "Synthetic Release Acceptance Contract"
_SITE_A = ("synthetic-north", "Synthetic North")
_SITE_B = ("synthetic-south", "Synthetic South")


def _bytes(source_title, source_month):
    document = Document()
    document.add_heading(source_title + " TEST DOCUMENT", 0)
    document.add_paragraph("Original report month: " + source_month)
    document.add_heading("2 Monthly Activity Summary", level=1)
    document.add_paragraph(
        "Inspected synthetic cooling pump. Bearing repair remains pending.")
    result = BytesIO()
    document.save(result)
    return result.getvalue()


def _metadata(scenario):
    source_site = _SITE_A
    source_month = "2026-09" if scenario == "AUTH01" else "2026-10"
    target_site = _SITE_B if scenario == "AUTH02" else _SITE_A
    return FixtureMetadata(
        friendly_name="Invented " + scenario + " acceptance case",
        source=ReportIdentity(_CONTRACT, (source_site[0],), source_month,
                              (source_site[1],)),
        target=ReportIdentity(_CONTRACT, (target_site[0],), "2026-10",
                              (target_site[1],)),
        scenario=scenario,
        owner_expectations="SYNTHETIC evidence only, no real owner approval.",
        expected_content=(
            ExpectedSection("Activity",
                includes=("MANIFEST-ONLY CANARY MUST NEVER APPEAR IN SOURCE",),
                excludes=("invented completion",)),
        ),
        confirmed_at="2026-10-14T03:00:00+00:00",
        partial_file_confirmed=(scenario == "AUTH03"),
    )


@pytest.mark.parametrize("scenario", ["AUTH01", "AUTH02", "AUTH03"])
def test_all_three_synthetic_auth_identities_stage_through_real_starting_inspector(
        monkeypatch, tmp_path, scenario):
    root = tmp_path / "simulated-mounted-root"
    root.mkdir()
    monkeypatch.setenv("EPC_DATA_DIR", str(root))
    metadata = _metadata(scenario)
    original = _bytes(metadata.source.site_titles[0], metadata.source.month)
    library = SharedFixtureLibrary(root)
    receipt = library.upload(original, metadata)
    assert receipt.original_sha256 == sha256(original).hexdigest()

    with testing.isolated(_RUN_TOKEN):
        memory._data_dir()
        profile = target_profile(metadata)
        assert profile.contract == metadata.target.contract
        assert tuple(site.key for site in profile.facilities) == metadata.target.sites
        stage = fixture_runs.stage_fixture(
            library, receipt.fixture_id, _RUN_TOKEN, profile,
            target_period=ReportPeriod(2026, 10))
        assert stage.period == ReportPeriod(2026, 10)
        assert stage.profile == profile
        assert stage.inspection.sha256 == receipt.original_sha256
        assert stage.stage[1].read_bytes() == original
        readback = fixture_runs.read_run_fixture(library, _RUN_TOKEN)
        assert readback is not None and readback.lineage.metadata == metadata
        assert readback.lineage.original_sha256 == receipt.original_sha256
        item_text = "\n".join(item.text for item in stage.inspection.items)
        assert "Inspected synthetic cooling pump." in item_text
        assert "MANIFEST-ONLY CANARY" not in item_text
        assert "Bearing repair remains pending." in item_text
        # The source's identity is never silently relabelled as the target.
        assert metadata.source.site_titles[0] in item_text
        if scenario == "AUTH02":
            assert metadata.target.sites != metadata.source.sites
            assert metadata.target.site_titles[0] not in item_text
        if scenario == "AUTH03":
            assert metadata.partial_file_confirmed is True
            assert metadata.source.month == metadata.target.month
        testing.reset(root)

    # Ordinary isolated reset invalidates only the run, not any shared original.
    reopened = SharedFixtureLibrary(root)
    assert reopened.list().fixtures == (receipt,)
    assert reopened.download(receipt.fixture_id) == original


_CHILD_CHECK = """
import sys
from hashlib import sha256
from pathlib import Path
from app.monthly_report_shared_fixtures import SharedFixtureLibrary

root, fixture_id, expected_hash, mode = sys.argv[1:]
library = SharedFixtureLibrary(Path(root))
if mode == 'absent':
    try:
        library.read(fixture_id)
    except (ValueError, OSError):
        print('ROOT_ISOLATED')
    else:
        raise AssertionError('Unknown mount silently fell back to existing data')
else:
    listing = library.list()
    found = [item for item in listing.fixtures if item.fixture_id == fixture_id]
    assert len(found) == 1
    readback = library.read(fixture_id)
    assert readback.receipt == found[0]
    assert readback.receipt.original_sha256 == expected_hash
    assert sha256(readback.original).hexdigest() == expected_hash
    assert library.download(fixture_id) == readback.original
    print('REOPEN_VERIFIED')
"""


def _new_python(root, receipt, mode):
    result = subprocess.run(
        [sys.executable, "-c", _CHILD_CHECK,
         str(root), receipt.fixture_id, receipt.original_sha256, mode],
        cwd=Path(__file__).resolve().parents[1],
        capture_output=True, text=True, timeout=30, check=True,
    )
    return result.stdout.strip()


def test_shared_original_survives_fresh_process_and_disposable_run_reset(
        monkeypatch, tmp_path):
    root = tmp_path / "isolated-volume"
    root.mkdir()
    monkeypatch.setenv("EPC_DATA_DIR", str(root))
    metadata = _metadata("AUTH01")
    original = _bytes(metadata.source.site_titles[0], metadata.source.month)
    library = SharedFixtureLibrary(root)
    receipt = library.upload(original, metadata)
    assert _new_python(root, receipt, "present") == "REOPEN_VERIFIED"

    # Exercise a report reset between independent interpreter read processes.
    with testing.isolated(_RUN_TOKEN):
        memory._data_dir()
        profile = target_profile(metadata)
        stage = fixture_runs.stage_fixture(library, receipt.fixture_id, _RUN_TOKEN, profile)
        assert stage.source_path.read_bytes() == original
        testing.reset(root)
    assert _new_python(root, receipt, "present") == "REOPEN_VERIFIED"

    # Replacing/missing the mounted directory MUST fail closed; no fallback
    # to some live user data root, and no automatic rehydration from this test.
    unmounted = tmp_path / "wrong-volume"
    unmounted.mkdir()
    assert _new_python(unmounted, receipt, "absent") == "ROOT_ISOLATED"
    assert not list(unmounted.iterdir())
    assert SharedFixtureLibrary(root).download(receipt.fixture_id) == original
