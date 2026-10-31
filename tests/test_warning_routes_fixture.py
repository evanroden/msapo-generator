"""Local fixture contracts, never claims of native browser acceptance."""
import pytest
from app import monthly_report_testing as testing,monthly_report_library as library
from scripts.workspace_warning_routes import SCENARIOS,prepare_scenario,MISSING_FACT


@pytest.mark.parametrize('scenario',SCENARIOS)
def test_historical_fixture_retains_unresolved_identity_bound_warnings(monkeypatch,tmp_path,scenario):
    monkeypatch.setenv('EPC_DATA_DIR',str(tmp_path))
    with testing.isolated('a'*32):
        draft=prepare_scenario(scenario)
        assert draft.synthetic
        assert library.load_snapshot(draft.profile.contract,draft.profile.key,draft.period).draft==draft
        from app.monthly_report_sources import source_bytes
        from hashlib import sha256
        assert all(sha256(source_bytes(draft.profile,s)).hexdigest()==s.sha256 for s in draft.sources)
        from app.monthly_report_checks import preflight
        codes={c.code for c in preflight(draft)}
        if scenario=='source':
            assert {'source_facility','source_date_unknown'}<=codes
            assert not draft.sources[1].facility and not draft.sources[1].service_date
        elif scenario=='shared':
            assert 'client_source_page' in codes
            from app.monthly_report_content_policy import page_allowed
            assert not page_allowed(draft.sources[1],1)
            a,b=[next(b for b in draft.blocks if b.key==key) for key in ('improvements','vendor_reports')]
            assert a.asset_hashes==b.asset_hashes and a.asset_captions!=b.asset_captions
        elif scenario=='ai':
            block=next(b for b in draft.blocks if b.key=='activity_summary')
            assert not block.reviewed and MISSING_FACT in block.ai_paragraphs[1].fact_ids
            assert {'ai_evidence','ai_date'}<=codes
        else:
            assert 'financial_page' in codes and not draft.sources[1].financial_source_review
            from app.monthly_report_content_policy import page_allowed
            assert not page_allowed(draft.sources[1],1,profile=draft.profile,destination='mbcx_report')


def test_no_live_namespace_fixture_installation():
    assert testing.scope() is None
    with pytest.raises(ValueError,match='isolated'):prepare_scenario('source')


@pytest.mark.parametrize('scenario',SCENARIOS)
def test_full_guarded_fixture_loads_actual_warning_buttons_without_approval(tmp_path,scenario):
    import os,subprocess,sys
    from pathlib import Path
    root=Path(__file__).resolve().parents[1]
    (tmp_path/'monthly-dom-synthetic-only.marker').write_text('synthetic-monthly-dom-v1')
    script=r'''
from pathlib import Path
import json,sys
from streamlit.testing.v1 import AppTest
from app.monthly_report_checks import preflight
app=AppTest.from_file(str(Path.cwd()/'tests/fixtures/workspaces/monthly_app.py'),default_timeout=60)
app.session_state['_monthly_dom_app_test_token']='a'*32
def landing():
 return next(json.loads(component.proto.json)['landing'] for component in app.get('bidi_component') if component.proto.component_name=='enfra_workspace_layout')
app.run();assert not app.exception,[e.message for e in app.exception]
next(w for w in app.selectbox if w.label=='Synthetic historical scenario').set_value(sys.argv[1]).run()
next(b for b in app.button if b.label=='Load synthetic historical report with unresolved warnings').click().run()
assert not app.exception,[e.message for e in app.exception]
nonce,section=landing()
assert nonce and section=='monthly_cover'
drafts=[v for k,v in app.session_state.filtered_state.items() if k.startswith('report_guided_') and k.endswith('_draft') and hasattr(v,'blocks')]
assert len(drafts)==1
codes={c.code for c in preflight(drafts[0])}
expected={'source':'source_facility','shared':'client_source_page','ai':'ai_evidence','financial':'financial_page'}[sys.argv[1]]
assert expected in codes,(expected,codes)
labels=[b.label for b in app.button]
assert any(s in label for label in labels for s in {'source':['Correct site','Correct date'],'shared':['Open source page'],'ai':['Review evidence for paragraph'],'financial':['Review financial source page']}[sys.argv[1]]),labels
assert next(b for b in app.button if b.label=='Create Word and PDF').disabled
assert not any(k.endswith('_package') for k in app.session_state.filtered_state)
assert not any(b.label=='Continue to this month’s updates' for b in app.button), 'A saved historical report has no setup review to continue'
before=drafts[0]
app.run();assert not app.exception,[e.message for e in app.exception]
assert landing() is None, 'Ordinary reruns must not pull the reader back to Cover'
after=next(v for k,v in app.session_state.filtered_state.items() if k.startswith('report_guided_') and k.endswith('_draft') and hasattr(v,'blocks'))
assert after==before, 'The entry request is display-only, not a report correction or approval'
fresh=AppTest.from_file(str(Path.cwd()/'tests/fixtures/workspaces/monthly_app.py'),default_timeout=60)
fresh.session_state['_monthly_dom_app_test_token']='a'*32;fresh.run()
assert not fresh.exception,[e.message for e in fresh.exception]
assert next(json.loads(component.proto.json)['landing'] for component in fresh.get('bidi_component') if component.proto.component_name=='enfra_workspace_layout') is None, 'Reopening must not manufacture a new entry action'
'''
    env={**os.environ,'EPC_DATA_DIR':str(tmp_path),'EPC_MONTHLY_DOM_FIXTURE':'1',
        'EPC_MONTHLY_WARNING_ROUTES_TASK':'1','EPC_MONTHLY_DOM_APPTEST':'1','ANTHROPIC_API_KEY':'','SMARTSHEET_API_MODE':'disabled'}
    result=subprocess.run([sys.executable,'-c',script,scenario],cwd=root,env=env,capture_output=True,text=True,timeout=100)
    assert result.returncode==0,result.stdout+result.stderr
    assert not (tmp_path/'monthly_reports').exists()


def test_fixture_entry_matches_the_actual_production_continue_contract(monkeypatch,tmp_path):
    """The same guided workspace emits Cover once after real setup Continue."""
    from test_cover_import_task_ui import stage_source
    from test_monthly_report_continue_landing import _landing
    app,_,_=stage_source(monkeypatch,tmp_path)
    assert _landing(app) is None
    next(b for b in app.button if b.label=='Continue to this month’s updates').click().run()
    assert not app.exception
    nonce,section=_landing(app)
    assert nonce and section=='monthly_cover'
    app.run()
    assert not app.exception and _landing(app) is None


@pytest.mark.parametrize('scenario',('source','ai','financial','shared'))
def test_full_fixture_corrections_use_existing_controls_and_reopen(tmp_path,scenario):
    import os,subprocess,sys
    from pathlib import Path
    root=Path(__file__).resolve().parents[1]
    (tmp_path/'monthly-dom-synthetic-only.marker').write_text('synthetic-monthly-dom-v1')
    script=r'''
from pathlib import Path
from datetime import date
import sys
from streamlit.testing.v1 import AppTest
from app.monthly_report_checks import preflight
from scripts.workspace_warning_routes import MISSING_FACT,BAD_DATE_FACT,CAPTIONS
app=AppTest.from_file(str(Path.cwd()/'tests/fixtures/workspaces/monthly_app.py'),default_timeout=60)
app.session_state['_monthly_dom_app_test_token']='b'*32;app.run()
next(w for w in app.selectbox if w.label=='Synthetic historical scenario').set_value(sys.argv[1]).run()
next(b for b in app.button if b.label=='Load synthetic historical report with unresolved warnings').click().run()
def current():return next(v for k,v in app.session_state.filtered_state.items() if k.startswith('report_guided_') and k.endswith('_draft') and hasattr(v,'blocks'))
def button(label):
 next(b for b in app.button if b.label==label).click().run();assert not app.exception,[e.message for e in app.exception]
def target_button(code,ordinal=0):
 c=next(c for c in preflight(current()) if c.code==code)
 assert c.targets, (code, current().sources, next(b for b in current().blocks if b.key=='activity_summary'))
 token=c.targets[ordinal].token
 next(b for b in app.button if (b.key or '').endswith('_'+token) or (b.key or '').endswith('_'+token+'_'+code)).click().run()
 assert not app.exception,[e.message for e in app.exception]
prefix=next(k[:-6] for k,v in app.session_state.filtered_state.items() if k.startswith('report_guided_') and k.endswith('_draft') and hasattr(v,'blocks'))
scenario=sys.argv[1];target=current().sources[1];before=current()
if scenario=='source':
 button('Correct site')
 next(w for w in app.multiselect if w.key==prefix+'_evidence_'+target.id+'_source_details_sites').set_value(['qa-north'])
 next(w for w in app.date_input if w.key==prefix+'_evidence_'+target.id+'_source_details_date').set_value(date(2026,9,3))
 next(b for b in app.button if (b.key or '').startswith('FormSubmitter:'+prefix+'_evidence_'+target.id+'_source_details_form')).click().run()
elif scenario=='ai':
 target_button('ai_evidence',1)
 next(w for w in app.multiselect if w.label=='Evidence for paragraph 2').set_value(['synthetic-seals',BAD_DATE_FACT]).run()
 target_button('ai_number')
 next(w for w in app.text_area if w.label=='Report paragraph 2').set_value('Replaced 2 seals.').run()
 button('Open cited source page 1')
 next(w for w in app.multiselect if w.label=='Evidence for paragraph 2').set_value(['synthetic-seals']).run()
 next(w for w in app.checkbox if w.label=='I checked this wording against the displayed current evidence').check().run()
elif scenario=='financial':
 button('Review financial source page 1')
 next(w for w in app.multiselect if w.label=='Pages to review together').set_value([1]).run()
 next(w for w in app.checkbox if w.label=='This is an ENFRA report intended for the client').check().run()
 next(w for w in app.checkbox if w.label.startswith('I reviewed these pages:')).check().run()
 button('Include reviewed financial pages')
else:
 c=next(c for c in preflight(current()) if c.code=='client_source_page' and c.block_key=='vendor_reports')
 next(b for b in app.button if (b.key or '').endswith('_'+c.targets[0].token+'_client_source_page')).click().run()
 next(b for b in app.button if (b.key or '').startswith(prefix+'_evidence_'+target.id+'_keep_page_')).click().run()
 for key,caption in zip(('improvements','vendor_reports'),CAPTIONS):
  assert next(b for b in current().blocks if b.key==key).asset_captions==(caption,)
 # Custom native picture-group confirmation is intentionally NOT faked by AppTest.
 assert any(c.code=='client_pages' for c in preflight(current()))
assert not app.exception,[e.message for e in app.exception]
if scenario!='shared':assert not [c for c in preflight(current()) if c.blocking],[(c.code,c.message) for c in preflight(current())]
button('Save shared version');saved=current()
fresh=AppTest.from_file(str(Path.cwd()/'tests/fixtures/workspaces/monthly_app.py'),default_timeout=60)
fresh.session_state['_monthly_dom_app_test_token']='b'*32;fresh.run();assert not fresh.exception,[e.message for e in fresh.exception]
reopened=next(v for k,v in fresh.session_state.filtered_state.items() if k.startswith('report_guided_') and k.endswith('_draft') and hasattr(v,'blocks'))
assert reopened==saved
'''
    env={**os.environ,'EPC_DATA_DIR':str(tmp_path),'EPC_MONTHLY_DOM_FIXTURE':'1',
        'EPC_MONTHLY_WARNING_ROUTES_TASK':'1','EPC_MONTHLY_DOM_APPTEST':'1','ANTHROPIC_API_KEY':'','SMARTSHEET_API_MODE':'disabled'}
    result=subprocess.run([sys.executable,'-c',script,scenario],cwd=root,env=env,capture_output=True,text=True,timeout=110)
    assert result.returncode==0,result.stdout+result.stderr
    assert not (tmp_path/'monthly_reports').exists()


def test_gate_requires_all_engines_modes_cases_outputs_and_reopen():
    from scripts.qa_warning_routes import all_cases_passed,required_cases
    cases={key:{'result':'passed','outputs_verified':True,'reopen_verified':True} for key in required_cases()}
    assert len(cases)==16 and all_cases_passed({'cases':cases,'failures':[]})
    assert not all_cases_passed({'cases':cases,'failures':['original failure']})
    for key in cases:
        assert not all_cases_passed({'cases':{k:v for k,v in cases.items() if k!=key},'failures':[]})
        for flag,value in (('result','running'),('outputs_verified',False),('reopen_verified',False)):
            assert not all_cases_passed({'cases':{**cases,key:{**cases[key],flag:value}},'failures':[]})


def test_native_driver_refuses_local_launch_before_importing_browser(monkeypatch):
    from scripts.qa_warning_routes import main
    monkeypatch.delenv('GITHUB_ACTIONS',raising=False)
    with pytest.raises(RuntimeError,match='hosted CI only'):main()


def test_hosted_gate_has_real_outputs_financial_tests_and_no_acceptance_bypass():
    from pathlib import Path
    root=Path(__file__).resolve().parents[1]
    workflow=(root/'.github/workflows/monthly-warning-routes.yml').read_text()
    driver=(root/'scripts/qa_warning_routes.py').read_text()
    assert 'libreoffice-writer' in workflow and 'libreoffice-calc' in workflow
    assert 'chromium webkit' in workflow and "assert not root.findall('.//skipped')" in workflow
    assert 'python scripts/qa_financial_pages.py' in workflow
    assert 'test_monthly_report_mbcx_financial_pages.py' in workflow
    assert 'force=True' not in driver and 'dispatch_event' not in driver and 'session_state' not in driver
    assert 'add_init_script' not in driver and 'approve_asset(' not in driver
    assert 'page.reload' in driver and 'verify_outputs' in driver and 'context.tracing.stop' in driver
    assert 'continue-on-error' not in workflow


@pytest.mark.parametrize('scenario',('source','ai','shared','financial'))
def test_real_output_oracle_rejects_missing_or_wrong_material(tmp_path,scenario):
    """Handcrafted oracle probes, never claims of production renderer acceptance."""
    from io import BytesIO
    from hashlib import sha256
    from docx import Document
    from docx.shared import Inches
    from PIL import Image
    import fitz
    from scripts.qa_warning_routes import verify_outputs
    from scripts.workspace_warning_routes import PENDING,CAPTIONS,_pdf
    original=_pdf('ENFRA client report. Verified savings: $125.',financial=True)
    def pixels(raw):
        with fitz.open(stream=raw,filetype='pdf') as source:
            image=Image.open(BytesIO(source[0].get_pixmap(dpi=150,alpha=False).tobytes('png'))).convert('RGB')
            image.thumbnail((1050,1350),Image.Resampling.LANCZOS)
        out=BytesIO();image.save(out,format='PNG');return out.getvalue()
    picture=pixels(original)
    changed=pixels(_pdf('ENFRA client report. Verified savings: $999.',financial=True))
    def files(*,pending=True,wrong_number=False,missing_caption=False,missing_image=False,
              pdf_only_missing=False,wrong_pdf=False,wrong_both=False,extra_image=False,black_overlay=False,unused_image=False,unused_pdf=False):
        doc=Document();lines=([PENDING] if pending else [])
        if scenario=='ai':lines+=['Replaced 99 seals.' if wrong_number else 'Replaced 2 seals.']
        if scenario=='shared':lines+=list(CAPTIONS[:1] if missing_caption else CAPTIONS)
        for line in lines:doc.add_paragraph(line)
        count=2 if scenario=='shared' else 1
        if scenario in ('shared','financial'):
            for _ in range(count-(1 if missing_image else 0)):
                doc.add_picture(BytesIO(changed if wrong_both else picture),width=Inches(1))
            if extra_image:doc.add_picture(BytesIO(changed),width=Inches(1))
        doc.save(tmp_path/'report.docx')
        if unused_image:
            from zipfile import ZipFile
            with ZipFile(tmp_path/'report.docx','a') as package:package.writestr('word/media/unexpected.png',changed)
        pdf=fitz.open();page=pdf.new_page();page.insert_text((30,30),'\n'.join(lines),fontsize=9)
        if scenario in ('shared','financial'):
            for n in range(count-(1 if missing_image or pdf_only_missing else 0)):
                page.insert_image(fitz.Rect(30+180*n,150,180+180*n,370),stream=changed if wrong_pdf or wrong_both else picture)
            if extra_image:page.insert_image(fitz.Rect(30,450,180,670),stream=changed)
            if black_overlay:page.draw_rect(fitz.Rect(30,150,180,370),fill=(0,0,0),overlay=True)
            if unused_image or unused_pdf:
                content=b'\n'.join(pdf.xref_stream(x) for x in page.get_contents())
                page.insert_image(fitz.Rect(30,450,180,670),stream=changed)
                stream=pdf.get_new_xref();pdf.update_object(stream,'<<>>');pdf.update_stream(stream,content);page.set_contents(stream)
        pdf.save(tmp_path/'report.pdf');pdf.close()
        # Deliberately forge perfectly consistent output/model hashes: these
        # are not independent proof that the original source was preserved.
        state={'output':{},'working':{'blocks':{key:{'asset_hashes':[sha256(changed if wrong_both else picture).hexdigest()+'.png']}
            for key in ('improvements','vendor_reports','mbcx_report')}}}
        for name in ('report.docx','report.pdf'):
            raw=(tmp_path/name).read_bytes();state['output'][name]={'bytes':len(raw),'sha256':sha256(raw).hexdigest()}
        return state
    def verify(**changes):return verify_outputs(tmp_path,scenario,files(**changes),original=original)
    assert verify()
    with pytest.raises(AssertionError):verify(pending=False)
    if scenario=='ai':
        with pytest.raises(AssertionError):verify(wrong_number=True)
    if scenario=='shared':
        with pytest.raises(AssertionError):verify(missing_caption=True)
    if scenario in ('shared','financial'):
        for mutant in ('missing_image','pdf_only_missing','wrong_pdf','wrong_both','extra_image','black_overlay','unused_image','unused_pdf'):
            with pytest.raises(AssertionError):verify(**{mutant:True})


def test_every_hosted_contract_and_native_test_exists_in_the_checkout():
    from pathlib import Path
    import re
    root=Path(__file__).resolve().parents[1]
    workflow=(root/'.github/workflows/monthly-warning-routes.yml').read_text()
    paths=set(re.findall(r'tests/[A-Za-z0-9_./-]+\.(?:py|cjs)',workflow))
    assert paths
    assert all((root/path).is_file() for path in paths),sorted(path for path in paths if not (root/path).is_file())
