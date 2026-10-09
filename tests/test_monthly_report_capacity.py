from dataclasses import replace
from io import BytesIO
import hashlib

import fitz
from docx import Document
from openpyxl import Workbook
from PIL import Image
import pytest

from app import monthly_report_capacity as capacity
from app import monthly_report_directory as directory, monthly_report_library as library
from app.monthly_report_model import Facility, ReportDraft, ReportPeriod, ReportProfile, ReportTable, ResolvedBlock, default_sections


@pytest.fixture
def profile(tmp_path, monkeypatch):
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path))
    facilities = (Facility("north", "North Hospital", ("NH",)), Facility("south", "South Hospital", ("SH",)))
    profile = ReportProfile("Synthetic Capacity Contract", "north", "North Hospital", facilities[:1])
    directory.save_directory(profile.contract, tuple(directory.DirectorySite(f.key, f.title, f.aliases) for f in facilities),
                             expected_revision=0, actor="Synthetic Reviewer", confirmed=True)
    return profile


CSV = b"Facility,Service,Required capacity,Available capacity,Units\nNH,Chilled water,1200,1600,tons\nNH,Steam,40000,50000,lb/hr\nSH,Chilled water,800,1000,tons\nSH,Steam,20000,24000,lb/hr\n"


def saved(profile, raw=CSV, revision=0, filename="capacity.csv", **kwargs):
    inspection = capacity.inspect_capacity(profile, filename, raw)
    return capacity.save_capacity(profile, inspection.tables, expected_revision=revision, actor="Synthetic Reviewer",
                                  reviewed_fingerprint=capacity.review_fingerprint(inspection.tables), **kwargs)


def draft(profile, block=None):
    return ReportDraft(profile, ReportPeriod(2026, 9), "Synthetic Editor", default_sections(),
                       (block or ResolvedBlock("thermal_capacity", "Library"),))


def test_one_upload_populates_first_reports_for_all_contract_sites(profile):
    state = saved(profile)
    assert state.revision == 1
    assert state.tables[0].site_keys == ("north", "north", "south", "south")
    north = capacity.resolve_capacity(draft(profile))
    south_profile = replace(profile, key="south", title="South Hospital", facilities=(Facility("south-report-key", "SH"),))
    south = capacity.resolve_capacity(draft(south_profile))
    assert len(north.blocks[0].extra_tables[0].rows) == 2
    assert north.blocks[0].extra_tables[0].rows[0] == ("NH", "Chilled water", "1200", "1600", "tons")
    assert south.blocks[0].extra_tables[0].rows[0] == ("SH", "Chilled water", "800", "1000", "tons")
    assert "Required capacity" in south.blocks[0].extra_tables[0].columns
    assert "Available capacity" in south.blocks[0].extra_tables[0].columns
    assert capacity._root(profile.contract).joinpath("sources", hashlib.sha256(CSV).hexdigest() + ".csv").read_bytes() == CSV


def test_updated_steam_preserves_chilled_water_and_other_sites(profile):
    saved(profile)
    changed = b"Facility,Service,Required capacity,Available capacity,Units\nNH,Steam,42000,52000,lb/hr\n"
    state = saved(profile, changed, revision=1)
    rows = [row for t in state.tables for row in t.rows]
    assert len(rows) == 4
    assert ("NH", "Steam", "42000", "52000", "lb/hr") in rows
    assert ("NH", "Steam", "40000", "50000", "lb/hr") not in rows
    assert ("NH", "Chilled water", "1200", "1600", "tons") in rows
    assert len(capacity.load_capacity(profile.contract, 1).tables[0].rows) == 4


def test_existing_snapshot_and_manual_content_never_autofollow_shared_updates(profile):
    saved(profile)
    old = capacity.resolve_capacity(draft(profile))
    saved(profile, CSV.replace(b"1200", b"1300"), revision=1)
    assert capacity.resolve_capacity(old) is old
    manual = draft(profile, ResolvedBlock("thermal_capacity", "This month", rows=(("NH", "Steam", "7", "units"),)))
    assert capacity.resolve_capacity(manual) is manual
    new = capacity.resolve_capacity(draft(profile))
    assert new.blocks[0].extra_tables[0].rows[0][2] == "1300"


def test_no_cross_contract_leak(profile):
    saved(profile)
    other = replace(profile, contract="Another Synthetic Contract")
    original = draft(other)
    assert capacity.resolve_capacity(original) is original


def test_unknown_or_ambiguous_site_cannot_be_shared(profile):
    inspection = capacity.inspect_capacity(profile, "unknown.csv", CSV.replace(b"NH,", b"Unity,"))
    assert inspection.tables[0].site_keys[:2] == ("", "")
    with pytest.raises(ValueError, match="Match each"):
        capacity.save_capacity(profile, inspection.tables, expected_revision=0, actor="Editor",
                               reviewed_fingerprint=capacity.review_fingerprint(inspection.tables))
    assert capacity.load_capacity(profile.contract) is None
    assert capacity.match_site("Unity", (Facility("a", "Unity Hospital", ("Unity",)), Facility("b", "Unity Specialty", ("Unity",)))) == ""


def test_review_and_actor_and_revision_are_required(profile):
    inspection = capacity.inspect_capacity(profile, "capacity.csv", CSV)
    with pytest.raises(library.LibraryError, match="Confirm"):
        capacity.save_capacity(profile, inspection.tables, expected_revision=0, actor="", reviewed_fingerprint=capacity.review_fingerprint(inspection.tables))
    with pytest.raises(ValueError, match="review changed"):
        capacity.save_capacity(profile, inspection.tables, expected_revision=0, actor="Editor", reviewed_fingerprint="stale")
    saved(profile)
    with pytest.raises(library.RevisionConflict):
        saved(profile)


def test_schema_conflict_requires_explicit_resolution(profile):
    saved(profile)
    raw = b"Facility,Service,Capacity,Units\nNH,Steam,55000,lb/hr\n"
    with pytest.raises(ValueError, match="different columns"):
        saved(profile, raw, revision=1)
    state = saved(profile, raw, revision=1, mode="replace_sites")
    rows = [row for table in state.tables for row in table.rows]
    assert len(rows) == 3
    assert ("NH", "Steam", "55000", "lb/hr") in rows
    assert ("SH", "Steam", "20000", "24000", "lb/hr") in rows


def test_distinct_plants_and_required_available_are_preserved(profile):
    raw = b"Facility,Plant,Service,Required capacity,Available capacity,Units\nNH,East,Steam,10,20,lb/hr\nNH,West,Steam,30,40,lb/hr\n"
    state = saved(profile, raw)
    assert len(state.tables[0].rows) == 2
    ambiguous = raw.replace(b",West,", b",East,")
    with pytest.raises(ValueError, match="same site and labels"):
        saved(profile, ambiguous, revision=1)


def test_xlsx_bounded_parser_finds_real_heading_and_preserves_all_columns(profile, monkeypatch):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Capacity"
    sheet.append(["Thermal requirements, all sites"])
    sheet.append(["Facility", "Chilled water required (tons)", "Steam required (lb/hr)", "Available steam (lb/hr)"])
    sheet.append(["NH", 1200, 40000, 50000])
    sheet.append(["SH", 800, 20000, 24000])
    stream = BytesIO()
    workbook.save(stream)
    # This upload path must not instantiate openpyxl's XML object graph.
    import app.monthly_report_sources as sources
    monkeypatch.setattr(sources, "load_workbook", lambda *a, **k: pytest.fail("unbounded workbook path"))
    inspection = capacity.inspect_capacity(profile, "capacity.xlsx", stream.getvalue())
    assert inspection.tables[0].site_keys == ("north", "south")
    assert inspection.tables[0].rows[1] == ("SH", "800", "20000", "24000")
    state = capacity.save_capacity(profile, inspection.tables, expected_revision=0, actor="Editor", reviewed_fingerprint=capacity.review_fingerprint(inspection.tables))
    assert state.tables[0].columns[-1] == "Available steam (lb/hr)"


def test_docx_native_table(profile):
    document = Document()
    document.add_heading("Thermal capacity", 1)
    table = document.add_table(rows=3, cols=4)
    for target, cells in zip(table.rows, (("Facility", "Service", "Capacity", "Units"), ("NH", "Steam", "200", "lb/hr"), ("SH", "Chilled water", "300", "tons"))):
        for cell, text in zip(target.cells, cells):
            cell.text = text
    stream = BytesIO()
    document.save(stream)
    inspection = capacity.inspect_capacity(profile, "capacity.docx", stream.getvalue())
    assert len(inspection.tables) == 1
    assert inspection.tables[0].site_keys == ("north", "south")
    assert inspection.tables[0].rows[1][2:] == ("300", "tons")


def test_pdf_native_table(profile):
    with fitz.open() as document:
        page = document.new_page()
        xs, ys = [40, 150, 280, 390, 490], [40, 70, 100, 130]
        for x in xs:
            page.draw_line((x, ys[0]), (x, ys[-1]))
        for y in ys:
            page.draw_line((xs[0], y), (xs[-1], y))
        rows = [("Facility", "Service", "Capacity", "Units"), ("NH", "Steam", "200", "lb/hr"), ("SH", "Chilled water", "300", "tons")]
        for j, row in enumerate(rows):
            for i, value in enumerate(row):
                page.insert_text((xs[i] + 3, ys[j] + 18), value, fontsize=9)
        raw = document.tobytes()
    inspection = capacity.inspect_capacity(profile, "capacity.pdf", raw)
    assert inspection.tables[0].site_keys == ("north", "south")
    assert inspection.tables[0].rows[1][2:] == ("300", "tons")


def test_image_uses_reviewed_structured_reader_and_never_guesses(profile):
    stream = BytesIO()
    Image.new("RGB", (200, 100), "white").save(stream, "PNG")
    inspection = capacity.inspect_capacity(profile, "capacity.png", stream.getvalue())
    assert not inspection.tables
    payload = capacity.reader_request(profile, inspection)
    assert any(part["type"] == "image" for part in payload)
    suggested = {"tables": [{"title": "Thermal capacity", "page": 1, "columns": ["Facility", "Service", "Capacity", "Units"], "rows": [["NH", "Steam", "[unclear]", "lb/hr"]]}]}
    result = capacity.reader_result(profile, inspection, suggested)
    assert result.used_reader
    with pytest.raises(ValueError, match="unreadable"):
        capacity.save_capacity(profile, result.tables, expected_revision=0, actor="Editor", reviewed_fingerprint=capacity.review_fingerprint(result.tables))
    assert capacity.load_capacity(profile.contract) is None


def test_native_reader_rejects_invented_numeric_value(profile):
    inspection = capacity.inspect_capacity(profile, "capacity.txt", b"North Hospital thermal steam capacity required 1200 tons")
    suggested = {"tables": [{"title": "North Hospital", "page": 1, "columns": ["Service", "Capacity", "Units"], "rows": [["Steam", "9999", "tons"]]}]}
    with pytest.raises(ValueError, match="not in the source"):
        capacity.reader_result(profile, inspection, suggested)


def test_failed_manifest_write_keeps_prior_capacity(profile, monkeypatch):
    old = saved(profile)
    atomic = library._atomic_write
    def fail(path, raw):
        if path == capacity._root(profile.contract) / "manifest.json":
            raise OSError("Synthetic capacity write failure")
        return atomic(path, raw)
    monkeypatch.setattr(library, "_atomic_write", fail)
    with pytest.raises(OSError):
        saved(profile, CSV.replace(b"1200", b"1300"), revision=1)
    assert capacity.load_capacity(profile.contract) == old


def test_unsupported_file_rejected_before_any_import(profile):
    with pytest.raises(ValueError, match="Upload a PDF"):
        capacity.inspect_capacity(profile, "capacity.exe", b"bad")


def test_contract_casing_reuses_shared_snapshot(profile):
    saved(profile)
    assert capacity.load_capacity(profile.contract.lower()).revision == 1


def test_separate_steam_and_chilled_water_tables_with_same_columns(profile):
    def workbook_bytes(steam_value):
        workbook = Workbook()
        workbook.remove(workbook.active)
        for title, value, units in (("Chilled water", "1200", "tons"), ("Steam", steam_value, "lb/hr")):
            sheet = workbook.create_sheet(title)
            sheet.append(["Facility", "Required", "Available", "Units"])
            sheet.append(["NH", value, value, units])
            sheet.append(["SH", "800", "1000", units])
        stream = BytesIO()
        workbook.save(stream)
        return stream.getvalue()
    state = saved(profile, workbook_bytes("40000"), filename="capacity.xlsx")
    assert {table.title for table in state.tables} == {"Chilled water", "Steam"}
    updated = saved(profile, workbook_bytes("50000"), revision=1, filename="capacity.xlsx")
    assert len(updated.tables) == 2
    steam = next(t for t in updated.tables if t.title == "Steam")
    assert steam.rows[0][1] == "50000"
    current = capacity.capacity_block(profile)
    assert current.extra_tables[0].columns[0] == "Source table"
    assert {table.rows[0][0] for table in current.extra_tables} == {"Chilled water", "Steam"}


def test_prior_shared_copy_can_refresh_but_manual_changes_cannot(profile):
    saved(profile)
    block = capacity.capacity_block(profile)
    assert capacity.can_refresh_capacity(profile, block)
    assert not capacity.can_refresh_capacity(profile, replace(block, text="Manually added plant note."))
    assert not capacity.can_refresh_capacity(profile, ResolvedBlock("thermal_capacity", "This month", rows=(("Manual",),)))


def test_foreign_report_empty_schema_and_source_reference_do_not_block_shared_capacity(profile):
    saved(profile)
    foreign = ResolvedBlock("thermal_capacity", "This month", extra_tables=(ReportTable(("Foreign site", "Steam requirement"), (), "old-source"),), references=("old-source",))
    filled = capacity.resolve_capacity(draft(profile, foreign))
    assert filled.blocks[0].extra_tables[0].rows[0][0] == "NH"
    assert "old-source" not in filled.blocks[0].references


def test_mixed_native_and_scanned_pages_merge_before_review(profile):
    native = capacity.inspect_capacity(profile, "capacity.csv", CSV)
    source = replace(native.content.source, page_texts=(*native.content.source.page_texts, ""), needs_vision=(2,))
    mixed = replace(native, content=replace(native.content, source=source), pending_pages=(2,))
    suggestion = {"tables": [{"title": "Steam secondary plant", "page": 2, "columns": ["Facility", "Service", "Capacity", "Units"], "rows": [["NH", "Steam", "30000", "lb/hr"]]}]}
    result = capacity.reader_result(profile, mixed, suggestion)
    assert not result.pending_pages
    assert len(result.tables) == 2
    assert result.tables[0] == native.tables[0]
    assert result.tables[1].rows[0][2] == "30000"


def test_native_reader_rejects_invented_unit(profile):
    inspection = capacity.inspect_capacity(profile, "capacity.txt", b"North Hospital Service Steam Capacity 1200 Units lb/hr")
    suggestion = {"tables": [{"title": "North Hospital", "page": 1, "columns": ["Service", "Capacity", "Units"], "rows": [["Steam", "1200", "tons"]]}]}
    with pytest.raises(ValueError, match="label or unit"):
        capacity.reader_result(profile, inspection, suggestion)


@pytest.mark.parametrize("label", ["Capacity type", "Capacity service", "Capacity category"])
def test_capacity_dimension_header_preserves_heating_when_cooling_is_added(profile, label):
    heading = f"Facility,{label},Capacity (kW)\n".encode()
    saved(profile, heading + b"NH,Heating,1000\n")
    state = saved(profile, heading + b"NH,Cooling,500\n", revision=1)
    assert {row for table in state.tables for row in table.rows} == {
        ("NH", "Heating", "1000"), ("NH", "Cooling", "500"),
    }
    changed = saved(profile, heading + b"NH,Heating,1200\n", revision=2)
    assert {row for table in changed.tables for row in table.rows} == {
        ("NH", "Heating", "1200"), ("NH", "Cooling", "500"),
    }


def test_numeric_plant_ids_are_distinct_capacity_dimensions(profile):
    heading = b"Facility,Plant,Service,Capacity (kW)\n"
    saved(profile, heading + b"NH,1,Heating,1000\n")
    state = saved(profile, heading + b"NH,2,Heating,500\n", revision=1)
    assert {row for table in state.tables for row in table.rows} == {
        ("NH", "1", "Heating", "1000"), ("NH", "2", "Heating", "500"),
    }
