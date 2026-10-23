"""An unchanged photograph must not inherit a stretched source frame.

The source photo bytes and the native frame stay intact. Valid asymmetric
source crops remain authoritative; absent/invalid crops need proportional fill.
All artwork here is synthetic.
"""
from io import BytesIO
from zipfile import ZipFile

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches
from PIL import Image, ImageDraw

from app.monthly_report_import import inspect_docx, map_items, ImportMapping
from app.monthly_report_model import Facility, ReportDraft, ReportPeriod, ReportProfile
from app.monthly_report_native_layout import build_native_docx


def source_photo(tmp_path, crop):
    image = Image.new('RGB', (400, 300), '#dae7ed')
    draw = ImageDraw.Draw(image)
    draw.ellipse((150, 100, 250, 200), fill='#173f49')
    draw.rectangle((25, 25, 75, 75), fill='#7c2b59')
    data = BytesIO(); image.save(data, 'PNG'); raw = data.getvalue()
    doc = Document()
    doc.add_paragraph('Synthetic facility')
    doc.add_paragraph('August 2026')
    shape = doc.add_picture(BytesIO(raw), width=Inches(4), height=Inches(2))
    rectangle = OxmlElement('a:srcRect')
    for key, value in crop.items(): rectangle.set(key, str(value))
    shape._inline.graphic.graphicData.pic.blipFill.blip.addnext(rectangle)
    path = tmp_path / 'synthetic-source.docx'; doc.save(path)
    inspected = inspect_docx(path)
    item = next(item for item in inspected.items if item.kind == 'image')
    mapped = map_items(path, inspected, (ImportMapping(item.id, 'cover_photo'),))
    profile = ReportProfile('Synthetic contract', 'synthetic-site', 'Synthetic facility',
                            (Facility('synthetic-site', 'Synthetic facility'),))
    draft = ReportDraft(profile, ReportPeriod(2026, 10), 'Synthetic reviewer', (), mapped.blocks)
    return path, draft, dict(mapped.assets), raw


def photo_geometry(raw, image_bytes=None):
    doc = Document(BytesIO(raw))
    picture = next(picture for picture in doc.element.iter(qn('pic:pic'))
                   if image_bytes is None or doc.part.related_parts[
                       next(picture.iter(qn('a:blip'))).get(qn('r:embed'))].blob == image_bytes)
    fill = picture.find(qn('pic:blipFill'))
    crop = fill.find(qn('a:srcRect'))
    values = {key: float(crop.get(key, '0')) / 100000 for key in ('l', 't', 'r', 'b')}
    extent = picture.find(qn('pic:spPr')).find(qn('a:xfrm')).find(qn('a:ext'))
    frame = (int(extent.get('cx')), int(extent.get('cy')))
    return values, frame


@pytest.mark.parametrize('section', [None, 'cover'])
@pytest.mark.parametrize('master', [False, True])
@pytest.mark.parametrize('crop', [{}, {'t': 5000, 'b': 5000}])
def test_unchanged_selected_photo_is_proportional_in_native_preview_and_export(tmp_path, section, master, crop):
    path, draft, assets, photograph = source_photo(tmp_path, crop)
    original = path.read_bytes()
    result = build_native_docx(draft, path, asset_loader=assets.__getitem__, section_key=section, master=master)
    fitted, frame = photo_geometry(result)
    effective = 400 * (1 - fitted['l'] - fitted['r']) / (300 * (1 - fitted['t'] - fitted['b']))
    assert effective == pytest.approx(frame[0] / frame[1], rel=.001), 'An unchanged photo must not bypass proportional placement'
    assert frame == (Inches(4), Inches(2)), 'Fix only the image crop, not the approved page frame'
    assert path.read_bytes() == original
    with ZipFile(BytesIO(result)) as package:
        assert photograph in [package.read(n) for n in package.namelist() if n.startswith('word/media/')]


@pytest.mark.parametrize('section', [None, 'cover'])
def test_unchanged_valid_asymmetric_source_crop_is_not_recentered(tmp_path, section):
    crop = {'t': 10000, 'b': 23333}
    path, draft, assets, _ = source_photo(tmp_path, crop)
    result = build_native_docx(draft, path, asset_loader=assets.__getitem__, section_key=section)
    fitted, _ = photo_geometry(result)
    assert fitted == {'l': 0, 't': .1, 'r': 0, 'b': .23333}


@pytest.mark.parametrize('section', [None, 'cover'])
@pytest.mark.parametrize('missing', [(0, 0), (400, 0), (0, 300)])
def test_unchanged_photo_uses_decoded_size_when_inspection_size_is_unavailable(tmp_path, monkeypatch, section, missing):
    from dataclasses import replace
    from app import monthly_report_native_layout as native
    path, draft, assets, photograph = source_photo(tmp_path, {})
    original = path.read_bytes()
    inspection = inspect_docx(path)
    inspection = replace(inspection, items=tuple(
        replace(item, image_width=missing[0], image_height=missing[1]) if item.kind == 'image' else item
        for item in inspection.items))
    # Header-only inspection may not have registered a supported decoder yet.
    # The actual normalization/digest path remains real and reads the source.
    monkeypatch.setattr(native, 'inspect_docx', lambda _: inspection)
    result = build_native_docx(draft, path, asset_loader=assets.__getitem__, section_key=section)
    fitted, frame = photo_geometry(result)
    assert 400 * (1 - fitted['l'] - fitted['r']) / (300 * (1 - fitted['t'] - fitted['b'])) == pytest.approx(frame[0] / frame[1], rel=.001)
    assert frame == (Inches(4), Inches(2))
    assert path.read_bytes() == original
    with ZipFile(BytesIO(result)) as package:
        assert photograph in [package.read(n) for n in package.namelist() if n.startswith('word/media/')]


@pytest.mark.parametrize('section', [None, 'cover'])
def test_source_view_photo_is_fitted_without_rewriting_preserved_art(tmp_path, monkeypatch, section):
    from test_monthly_report_source_frame_logo_role import source
    from test_monthly_report_native_art_review import setup
    from test_monthly_report_native_art_lifecycle import save_sample
    from app.monthly_report_preview import section_preview_docx
    monkeypatch.setenv('EPC_DATA_DIR', str(tmp_path / 'store'))
    _, _, _, photograph = source_photo(tmp_path, {})
    path = source(tmp_path, divergent=True, shared_outside=False)
    document = Document(path)
    # One explicitly reviewed, unrelated native header occurrence exercises
    # the source-view preservation guard while the photo remains editable.
    from test_monthly_report_source_frame_logo_role import raster
    document.sections[0].header.paragraphs[0].add_run().add_picture(BytesIO(raster('orange')), width=Inches(1))
    shape = document.add_picture(BytesIO(photograph), width=Inches(4), height=Inches(2))
    paragraph = shape._inline.getparent().getparent().getparent()
    document.element.body.remove(paragraph)
    document.element.body.insert(1, paragraph)
    document.save(path)
    original = path.read_bytes()
    sample = setup(path)
    draft, loader = save_sample(sample)
    result = (build_native_docx(draft, path, asset_loader=loader) if section is None
              else section_preview_docx(draft, section, loader))
    fitted, frame = photo_geometry(result, photograph)
    assert 400 * (1 - fitted['l'] - fitted['r']) / (300 * (1 - fitted['t'] - fitted['b'])) == pytest.approx(frame[0] / frame[1], rel=.001)
    assert path.read_bytes() == original


from tests.conftest import requires_libreoffice


@requires_libreoffice
@pytest.mark.parametrize('section', [None, 'cover'])
def test_actual_pdf_keeps_synthetic_circle_round(tmp_path, monkeypatch, section):
    import fitz
    from app import pdf_converter
    from app.monthly_report_render_compat import rendering_docx
    monkeypatch.setenv('EPC_DATA_DIR', str(tmp_path / 'store'))
    monkeypatch.setattr(pdf_converter, 'OUTPUT_DIR', tmp_path)
    path, draft, assets, _ = source_photo(tmp_path, {})
    raw = build_native_docx(draft, path, asset_loader=assets.__getitem__, section_key=section)
    rendered = tmp_path / 'current.docx'; rendered.write_bytes(rendering_docx(raw))
    pdf_path = pdf_converter.convert_to_pdf(rendered)
    with fitz.open(pdf_path) as pdf:
        pixmap = pdf[0].get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
        pixels = Image.frombytes('RGB', (pixmap.width, pixmap.height), pixmap.samples)
        points = [(x, y) for y in range(pixels.height) for x in range(pixels.width)
                  if all(abs(channel - target) < 5 for channel, target in zip(pixels.getpixel((x, y)), (23, 63, 73)))]
    assert points, 'The actual PDF must contain the selected synthetic photo'
    width = max(x for x, _ in points) - min(x for x, _ in points) + 1
    height = max(y for _, y in points) - min(y for _, y in points) + 1
    assert abs(width - height) <= 2, 'A round source subject must remain round in the PDF'


def test_unchanged_vml_crop_fits_and_keeps_valid_fixed_point_or_outset_geometry():
    from lxml import etree
    from app.monthly_report_native_layout import _fit_cover_photo
    v = '{urn:schemas-microsoft-com:vml}'
    shape = etree.Element(v + 'shape', style='width:4in;height:2in')
    image = etree.SubElement(shape, v + 'imagedata')
    _fit_cover_photo(image, 400, 300, preserve_crop=True)
    assert float(image.get('croptop')) == pytest.approx(1/6, abs=1e-8)
    assert float(image.get('cropbottom')) == pytest.approx(1/6, abs=1e-8)
    image.attrib.update({'croptop': '8192f', 'cropbottom': '24576f', 'cropleft': '0', 'cropright': '0'})
    before = dict(image.attrib)
    _fit_cover_photo(image, 400, 400, preserve_crop=True)
    assert dict(image.attrib) == before
    shape.set('style', 'width:4in;height:4.5in')
    image.attrib.update({'croptop': '-.25', 'cropbottom': '-.25', 'cropleft': '0', 'cropright': '0'})
    before = dict(image.attrib)
    _fit_cover_photo(image, 400, 300, preserve_crop=True)
    assert dict(image.attrib) == before, 'An intentional proportional outset is valid source geometry'


@pytest.mark.parametrize('unsupported', ['group-scale', 'inset-fill', 'tiled-fill'])
def test_unchanged_complex_destination_is_not_recropped_from_local_extent(tmp_path, unsupported):
    from lxml import etree
    from app.monthly_report_native_layout import _fit_cover_photo
    path, _, _, _ = source_photo(tmp_path, {})
    doc = Document(path)
    picture = next(doc.element.iter(qn('pic:pic')))
    blip = next(picture.iter(qn('a:blip')))
    if unsupported == 'group-scale':
        parent = picture.getparent(); parent.remove(picture)
        group = etree.SubElement(parent, '{http://schemas.microsoft.com/office/word/2010/wordprocessingGroup}wgp')
        properties = etree.SubElement(group, '{http://schemas.microsoft.com/office/word/2010/wordprocessingGroup}grpSpPr')
        transform = etree.SubElement(properties, qn('a:xfrm'))
        etree.SubElement(transform, qn('a:ext'), cx='400', cy='300')
        etree.SubElement(transform, qn('a:chExt'), cx='400', cy='200')
        group.append(picture)
    elif unsupported == 'inset-fill':
        picture.find(qn('pic:blipFill')).find(qn('a:stretch')).find(qn('a:fillRect')).set('b', '25000')
    else:
        fill = picture.find(qn('pic:blipFill'))
        fill.remove(fill.find(qn('a:stretch'))); etree.SubElement(fill, qn('a:tile'))
    before = doc.element.xml
    _fit_cover_photo(blip, 400, 300, preserve_crop=True)
    assert doc.element.xml == before, 'Effective group/fill geometry is outside this bounded repair'


@pytest.mark.parametrize('invalid', ['NaN', 'inf', '0', '-1', 'oops'])
def test_replacement_nonfinite_or_nonpositive_frame_is_rejected_without_image_mutation(tmp_path, invalid):
    from app.monthly_report_native_layout import _fit_cover_photo, NativeLayoutError
    path, _, _, _ = source_photo(tmp_path, {})
    doc = Document(path)
    picture = next(doc.element.iter(qn('pic:pic')))
    extent = picture.find(qn('pic:spPr')).find(qn('a:xfrm')).find(qn('a:ext'))
    extent.set('cx', invalid)
    blip = next(picture.iter(qn('a:blip')))
    before = doc.element.xml
    with pytest.raises(NativeLayoutError, match='frame dimensions'):
        _fit_cover_photo(blip, 400, 300, preserve_crop=False)
    assert doc.element.xml == before


def test_unchanged_percentage_lexical_source_crop_retains_asymmetry(tmp_path):
    from app.monthly_report_native_layout import _fit_cover_photo
    path, _, _, _ = source_photo(tmp_path, {'t': '10%', 'b': '23.333%'})
    doc = Document(path)
    blip = next(doc.element.iter(qn('a:blip')))
    before = doc.element.xml
    _fit_cover_photo(blip, 400, 300, preserve_crop=True)
    assert doc.element.xml == before
