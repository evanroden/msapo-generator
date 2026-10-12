"""Bounded thermal XLSX two-line headers with merged service headings.

All records and site names are invented. Unrecognized candidate data rows
remain in review rather than being treated as header or silently discarded.
"""
from io import BytesIO

import pytest
from openpyxl import Workbook

from app import monthly_report_capacity as capacity
from app import monthly_report_directory as directory
from app import monthly_report_library as library
from app.monthly_report_model import Facility, ReportProfile


@pytest.fixture
def profile(tmp_path, monkeypatch):
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path))
    sites = (Facility("north", "Synthetic North", ("SN",)),
             Facility("south", "Synthetic South", ("SS",)))
    report = ReportProfile("Synthetic Two-row Thermal", "north", "Synthetic North", sites, "multi_site")
    directory.save_directory(
        report.contract,
        tuple(directory.DirectorySite(f.key, f.title, f.aliases) for f in sites),
        expected_revision=0, actor="Synthetic Reviewer", confirmed=True)
    return report


def two_row_workbook(*, title=False, ambiguous=False):
    document = Workbook()
    sheet = document.active
    sheet.title = "Invented thermal requirements"
    if title:
        sheet.append(["September 2026 technical capacity overview"])
    first = sheet.max_row + 1
    sheet.cell(first, 1, "Facility")
    sheet.cell(first, 2, "Chilled water")
    sheet.merge_cells(start_row=first, start_column=2, end_row=first, end_column=3)
    sheet.cell(first, 4, "Steam")
    sheet.merge_cells(start_row=first, start_column=4, end_row=first, end_column=5)
    sheet.append(["", "Required (tons)", "Available (%)",
                  "Required (lb/hr)", "Available (lb/hr)"])
    if ambiguous:
        sheet.append(["", "Maintenance required", "NR", "", ""])
    sheet.append(["SN", 1200, "NR", 40000, 50000])
    sheet.append(["SS", 800, "85%", 20000, "NR"])
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


@pytest.mark.parametrize("title", [False, True])
def test_two_tier_merged_header_keeps_all_technical_values_and_nr(profile, title):
    raw = two_row_workbook(title=title)
    inspection = capacity.inspect_capacity(profile, "synthetic.xlsx", raw)
    assert len(inspection.tables) == 1
    table = inspection.tables[0]
    assert table.site_column == 0
    assert table.site_keys == ("north", "south")
    assert table.columns == (
        "Facility",
        "Chilled water / Required (tons)",
        "Chilled water / Available (%)",
        "Steam / Required (lb/hr)",
        "Steam / Available (lb/hr)",
    )
    assert table.rows == (
        ("SN", "1200", "NR", "40000", "50000"),
        ("SS", "800", "85%", "20000", "NR"),
    )
    assert len(table.rows) == 2  # The continuation header was not stored as data.
    first = capacity.first_fill_tables(None, inspection.tables)
    assert first == inspection.tables
    saved = capacity.save_capacity(profile, first, expected_revision=0,
                                   actor="Synthetic Reviewer",
                                   reviewed_fingerprint=capacity.review_fingerprint(first))
    assert saved.tables[0].rows == table.rows
    assert capacity.load_capacity(profile.contract, 1) == saved


def test_ambiguous_blank_site_row_remains_for_review_not_silently_lost(profile):
    raw = two_row_workbook(ambiguous=True)
    inspection = capacity.inspect_capacity(profile, "synthetic.xlsx", raw)
    assert len(inspection.tables) == 1
    table = inspection.tables[0]
    assert table.site_keys == ("", "north", "south")
    assert len(table.rows) == 3
    assert table.rows[0][1] == "Maintenance required"
    assert table.rows[0][2] == "NR"
    with pytest.raises(ValueError, match="Match each capacity row"):
        capacity.save_capacity(profile, inspection.tables, expected_revision=0,
                               actor="Synthetic Reviewer",
                               reviewed_fingerprint=capacity.review_fingerprint(inspection.tables))
    assert capacity.load_capacity(profile.contract) is None
