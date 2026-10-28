"""Do not assign reusable thermal capacity to a model-invented facility title.

The reader output in these tests is a local deterministic dictionary; no AI
request or paid source reading is made. Site scope requires exact text on the
same original page, not the reader's free-form title alone.
"""
from dataclasses import replace
from io import BytesIO

from docx import Document
import pytest

from app import monthly_report_capacity as capacity
from app import monthly_report_directory as directory
from app import monthly_report_library as library
from app import monthly_report_sources as sources
from app.monthly_report_model import (
    Facility, ReportDraft, ReportPeriod, ReportProfile, ReportSource,
    ResolvedBlock, default_sections,
)
from app.monthly_report_sources import SourceContent, SourceTable


PERIOD = ReportPeriod(2026, 9)
UNSCOPED = (
    "Thermal capacity\nService: Steam\nRequired: 1200 tons\n"
    "Available: 1400 tons\n"
)
HEADERS = ["Service", "Required", "Available"]
VALUES = [["Steam", "1200", "1400"]]


@pytest.fixture
def profile(monkeypatch, tmp_path):
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path))
    north = Facility("north", "North Hospital", ("NH",))
    south = Facility("south", "South Hospital", ("SH",))
    result = ReportProfile("Synthetic Capacity Scope Contract", "north-report",
                           "North report", (north,))
    directory.save_directory(
        result.contract,
        tuple(directory.DirectorySite(f.key, f.title, f.aliases) for f in (north, south)),
        expected_revision=0, actor="Synthetic Editor", confirmed=True,
    )
    return result


def inspection(text, *, image=False):
    source = ReportSource(
        "synthetic-capacity-text", "synthetic-capacity.txt", "b" * 64, ".txt",
        page_texts=(text,), selected_pages=(1,),
        needs_vision=(1,) if image else (),
    )
    return capacity.CapacityInspection(SourceContent(source), ())


def proposed(title, *, columns=None, rows=None):
    return {"tables": [{
        "title": title, "page": 1,
        "columns": list(HEADERS if columns is None else columns),
        "rows": [list(row) for row in (VALUES if rows is None else rows)],
    }]}


def read(profile, text, title, *, image=False):
    return capacity.reader_result(
        profile, inspection(text, image=image), proposed(title))


@pytest.mark.parametrize("title", (
    "North Hospital thermal capacity",
    "NH thermal capacity",
    "South Hospital steam capacity",
    "SH steam capacity",
))
def test_native_reader_cannot_invent_site_from_title_not_on_cited_page(profile, title):
    before = directory.load_directory(profile.contract)
    with pytest.raises(ValueError, match="title names a site not on the cited page"):
        read(profile, UNSCOPED, title)
    assert directory.load_directory(profile.contract) == before
    assert capacity.load_capacity(profile.contract) is None


def test_before_case_shows_how_a_fabricated_title_would_become_a_saved_site_key(profile):
    facilities = capacity.contract_facilities(profile)
    source = inspection(UNSCOPED).content.source
    # The previous reader_result passed this exact unverified title directly
    # to _candidate and silently populated the site from that invented title.
    table = SourceTable("North Hospital steam capacity", tuple(HEADERS),
                        tuple(tuple(row) for row in VALUES))
    candidate = capacity._candidate(table, source, 1, facilities, UNSCOPED)
    assert candidate is not None
    assert candidate.site_keys == ("north",), (
        "This is the falsifiable misclassification the new admission blocks."
    )
    with pytest.raises(ValueError, match="site not on the cited page"):
        capacity.reader_result(profile, inspection(UNSCOPED), proposed(table.name))


@pytest.mark.parametrize("source_name,title,expected_key", (
    ("North Hospital", "North Hospital steam capacity", "north"),
    ("North Hospital", "NH steam capacity", "north"),
    ("NH", "North Hospital steam capacity", "north"),
    ("South Hospital", "SH steam capacity", "south"),
    ("North Hospital", "NORTH HOSPITAL STEAM CAPACITY", "north"),
))
def test_same_page_exact_site_or_alias_supports_real_table_scope(
        profile, source_name, title, expected_key):
    result = read(profile, f"{source_name}\n{UNSCOPED}", title)
    assert len(result.tables) == 1
    assert result.tables[0].site_keys == (expected_key,)
    assert result.tables[0].rows == (("Steam", "1200", "1400"),)
    assert result.tables[0].title == title
    assert result.used_reader
    assert capacity.load_capacity(profile.contract) is None


@pytest.mark.parametrize("source_identity,title", (
    ("North Hospitals", "North Hospital thermal capacity"),
    ("NHS", "NH steam capacity"),
    ("South Hospital", "North Hospital steam capacity"),
    ("North Hospitalish", "North Hospital steam capacity"),
))
def test_prefix_or_different_site_name_is_not_evidence(profile, source_identity, title):
    with pytest.raises(ValueError, match="site not on the cited page"):
        read(profile, f"{source_identity}\n{UNSCOPED}", title)


def test_multiple_explicit_site_names_do_not_get_conflated_into_one_site(profile):
    page = f"North Hospital\nSouth Hospital\n{UNSCOPED}"
    title = "North Hospital / South Hospital steam capacity"
    result = read(profile, page, title)
    table, = result.tables
    assert table.site_keys == ("",)
    assert table.site_column == -1
    # No automatic reusable site value can be saved without explicit mapping.
    with pytest.raises(ValueError, match="Match each capacity row"):
        capacity.save_capacity(
            profile, result.tables, expected_revision=0,
            actor="Synthetic Reviewer",
            reviewed_fingerprint=capacity.review_fingerprint(result.tables))
    assert capacity.load_capacity(profile.contract) is None


def test_title_without_any_known_site_keeps_capacity_unmapped_for_review(profile):
    result = read(profile, UNSCOPED, "Steam capacity")
    table, = result.tables
    assert table.site_keys == ("",)
    assert table.rows == (("Steam", "1200", "1400"),)


def test_model_title_length_is_bounded_even_if_cells_and_pages_are_valid(profile):
    with pytest.raises(ValueError, match="invalid table details"):
        read(profile, UNSCOPED, "Thermal capacity " + "x" * 4000)
    assert capacity.load_capacity(profile.contract) is None


def test_source_page_and_evidence_controls_still_fail_closed(profile):
    result = inspection("North Hospital\n" + UNSCOPED)
    invalid_page = proposed("North Hospital steam capacity")
    invalid_page["tables"][0]["page"] = 2
    with pytest.raises(ValueError, match="invalid table details"):
        capacity.reader_result(profile, result, invalid_page)
    different_number = proposed("North Hospital steam capacity", rows=[["Steam", "9999", "1400"]])
    with pytest.raises(ValueError, match="value was not in the source"):
        capacity.reader_result(profile, result, different_number)
    assert capacity.load_capacity(profile.contract) is None


def test_site_named_on_a_different_page_cannot_authorize_this_table(profile):
    native = ReportSource(
        "synthetic-two-page", "synthetic-two-page.pdf", "a" * 64, ".pdf",
        page_texts=(UNSCOPED, "North Hospital has other equipment."),
        selected_pages=(1, 2),
    )
    inspected = capacity.CapacityInspection(SourceContent(native), ())
    suggestion = proposed("North Hospital steam capacity")
    with pytest.raises(ValueError, match="site not on the cited page"):
        capacity.reader_result(profile, inspected, suggestion)
    assert not inspected.tables
    assert capacity.load_capacity(profile.contract) is None


def test_vision_page_keeps_existing_explicit_human_review_path(profile):
    # Visual content cannot be verified by native text extraction. Do not
    # pretend blank OCR text establishes the image's site or silently drop
    # its reviewed rows; the existing site/source approval still applies.
    result = read(profile, "", "North Hospital thermal capacity", image=True)
    table, = result.tables
    assert table.site_keys == ("north",)
    assert result.used_reader


def test_valid_native_source_saves_and_reopens_correct_site_without_cross_site_fill(
        profile, tmp_path):
    text = "North Hospital\n" + UNSCOPED
    raw = text.encode("utf-8")
    content, _ = sources.ingest(profile, "synthetic-capacity.txt", raw)
    assert sources.source_bytes(profile, content.source) == raw
    inspected = capacity.CapacityInspection(content, ())
    result = capacity.reader_result(profile, inspected,
                                   proposed("North Hospital steam capacity"))
    table, = result.tables
    assert table.site_keys == ("north",)
    state = capacity.save_capacity(
        profile, (table,), expected_revision=0,
        actor="Synthetic Reviewer",
        reviewed_fingerprint=capacity.review_fingerprint((table,)),
    )
    assert state.revision == 1
    assert capacity.load_capacity(profile.contract) == state
    assert capacity.load_capacity(profile.contract, 1) == state

    draft = ReportDraft(
        profile, PERIOD, "Synthetic Reviewer", default_sections(),
        (ResolvedBlock("thermal_capacity", "Library"),),
    )
    resolved = capacity.resolve_capacity(draft)
    capacity_block = next(b for b in resolved.blocks if b.key == "thermal_capacity")
    assert capacity_block.extra_tables[0].rows == (
        ("North Hospital steam capacity", "Steam", "1200", "1400"),)
    snapshot = library.save_snapshot(
        resolved, expected_revision=0, entered_editor="Synthetic Reviewer")
    reloaded = library.load_snapshot(profile.contract, profile.key, PERIOD)
    assert reloaded.revision == snapshot.revision
    assert reloaded.draft == resolved
    assert sources.source_bytes(profile, content.source) == raw

    south = replace(profile, key="south-report", title="South report",
                    facilities=(Facility("south", "South Hospital", ("SH",)),))
    assert capacity.capacity_block(south, state) is None
    assert capacity.load_capacity(profile.contract).tables[0].site_keys == ("north",)

    from app.monthly_report_docx import _build_docx
    word = Document(BytesIO(_build_docx(resolved)))
    xml = "\n".join(cell.text for table in word.tables for row in table.rows
                    for cell in row.cells)
    assert "North Hospital steam capacity" in xml
    assert "1200" in xml and "1400" in xml
