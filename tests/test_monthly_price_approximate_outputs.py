"""Synthetic saved price correction through actual shell and native Word/PDF.

These artifacts exercise local generators, immutable history, and an invented
design. They do not certify production delivery or a genuine client design.
"""
from dataclasses import replace
from hashlib import sha256
from io import BytesIO
import json
import os
from pathlib import Path
import re
from zipfile import ZipFile

from docx import Document
import fitz
import pytest

from app import monthly_report_designs as designs, monthly_report_library as library
from app.monthly_report_checks import preflight
from app.monthly_report_docx import generate_report
from app.monthly_report_model import (
    BlockSpec, ColumnSpec, Facility, ReportDraft, ReportPeriod, ReportProfile, ReportTable,
    ResolvedBlock, SectionSpec,
)
from app.monthly_report_guided import client_table_draft
from tests.conftest import requires_libreoffice


TECHNICAL_FACTS = (
    "Synthetic valve remains out of service.\n"
    "Pump pressure: ~125 psi. Motor output: ~3.2 kW.\n"
    "Quote reference Q-3200."
)
SOURCE_CANARY = "SYNTHETIC_SOURCE_ACTIVITY_CANARY"
READINGS = ReportTable(
    ("Asset code", "Reading", "Service date"),
    (("PUMP-3200", "~125 psi", "2026-10-20"),
     ("MOTOR-3200", "~3.2 kW", "2026-10-21")),
    "synthetic-operational-readings",
)


def _draft(tmp_path, native, price):
    profile = ReportProfile("Synthetic Price Review", "north", "Synthetic North",
                            (Facility("north", "Synthetic North"),))
    if native:
        source = tmp_path / "SYNTHETIC-price-design.docx"
        document = Document()
        document.add_paragraph("Operations and Maintenance Monthly Review")
        document.add_paragraph("Synthetic North")
        document.add_paragraph("September 2026")
        document.add_heading("MONTHLY ACTIVITY SUMMARY", 1)
        document.add_paragraph(SOURCE_CANARY)
        table = document.add_table(rows=2, cols=len(READINGS.columns))
        table.style = "Light Shading Accent 1"
        for cell, value in zip(table.rows[0].cells, READINGS.columns):
            cell.text = value
        for cell in table.rows[1].cells:
            cell.text = SOURCE_CANARY
        document.save(source)
        profile = designs.install(source, profile, actor="Synthetic Editor", expected_revision=0)
    section = SectionSpec("activity", "2", "Monthly Activity Summary",
                          (BlockSpec("activity_summary", "rich_text"),
                           BlockSpec("work_orders", "table", columns=tuple(
                               ColumnSpec(f"reading_{index}", title)
                               for index, title in enumerate(READINGS.columns)))))
    block = ResolvedBlock("activity_summary", "This month",
                          text=f"{TECHNICAL_FACTS}\n{price}")
    readings = ResolvedBlock("work_orders", "This month", extra_tables=(READINGS,))
    return ReportDraft(profile, ReportPeriod(2026, 10), "Synthetic Editor", (section,), (block, readings))


def _write_package(output, name, package, manifest):
    assert package.docx and package.pdf, package.pdf_error
    for suffix, raw in (("docx", package.docx), ("pdf", package.pdf)):
        filename = f"SYNTHETIC-{name}.{suffix}"
        (output / filename).write_bytes(raw)
        manifest["files"][filename] = {"bytes": len(raw), "sha256": sha256(raw).hexdigest()}
    with fitz.open(stream=package.pdf, filetype="pdf") as pdf:
        for index, page in enumerate(pdf, 1):
            page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5)).save(output / f"{name}-page-{index:02}.png")


def _texts(package):
    with ZipFile(BytesIO(package.docx)) as archive:
        xml = "\n".join(archive.read(name).decode("utf-8") for name in archive.namelist()
                        if name.endswith(".xml"))
    with fitz.open(stream=package.pdf, filetype="pdf") as pdf:
        pdf_text = re.sub(r"\s+", " ", "\n".join(page.get_text() for page in pdf))
    return xml, pdf_text


@requires_libreoffice
@pytest.mark.parametrize("native", [False, True], ids=["shell", "native"])
@pytest.mark.parametrize("price", ["Repair estimate: ~3.2k", "Quote ~3.2k"], ids=["estimate", "quote"])
def test_saved_approximate_price_requires_explicit_correction_before_outputs(tmp_path, monkeypatch, native, price):
    runtime = tmp_path / "runtime"
    monkeypatch.setenv("EPC_DATA_DIR", str(runtime))
    draft = _draft(tmp_path, native, price)
    original = library.save_snapshot(draft, expected_revision=0, entered_editor=draft.prepared_by)
    restored = library.load_snapshot(draft.profile.contract, draft.profile.key, draft.period)
    assert restored == original
    original_paths = tuple(runtime.rglob(f"{draft.period.key}-00000001.json"))
    assert len(original_paths) == 1
    original_path = original_paths[0]
    original_bytes = original_path.read_bytes()
    source = designs.source_for(draft.profile)
    source_bytes = source.read_bytes() if source else None

    # Price review preserves human prose and this non-financial custom schema.
    working, removed = client_table_draft(restored.draft)
    assert working == restored.draft and not removed
    assert working.blocks[1].extra_tables == (READINGS,)
    output = (Path(os.environ.get("INDEPENDENT_OUTPUT", str(tmp_path / "evidence")))
              / "price-approximate" / ("native" if native else "shell")
              / ("estimate" if price.startswith("Repair") else "quote"))
    output.mkdir(parents=True, exist_ok=True)
    (output / "SYNTHETIC-original-snapshot.json").write_bytes(original_bytes)
    manifest = {
        "synthetic_only": True,
        "tested_revision": os.environ.get("MSAPO_SHA", "unrecorded"),
        "native_design": native,
        "original_revision": original.revision,
        "original_snapshot_sha256": sha256(original_bytes).hexdigest(),
        "original_fingerprint": working.fingerprint,
        "files": {},
    }
    pricing_checks = [check for check in preflight(working) if check.code == "pricing" and check.blocking]
    try:
        unsafe_package = generate_report(working, acknowledged_fingerprint=working.fingerprint)
    except ValueError as exc:
        assert "pricing" in str(exc).lower(), str(exc)
        generation_blocked = True
    else:
        # Keep actual failing-before exports before the desired safety assertion.
        _write_package(output, "unsafe-before-correction", unsafe_package, manifest)
        unsafe_xml, unsafe_pdf = _texts(unsafe_package)
        assert price in unsafe_xml and price in unsafe_pdf
        generation_blocked = False
    manifest["uncorrected_generation_blocked"] = generation_blocked

    # A named human edit creates revision 2; no review helper rewrites revision 1.
    corrected = replace(working, blocks=(replace(working.blocks[0], text=TECHNICAL_FACTS), working.blocks[1]))
    saved_correction = library.save_snapshot(corrected, expected_revision=original.revision,
                                             entered_editor="Synthetic Correcting Editor")
    reloaded = library.load_snapshot(draft.profile.contract, draft.profile.key, draft.period)
    assert reloaded == saved_correction and reloaded.revision == 2
    assert reloaded.draft.blocks[0].text == TECHNICAL_FACTS
    assert reloaded.draft.blocks[1].extra_tables == (READINGS,)
    package = generate_report(reloaded.draft, acknowledged_fingerprint=reloaded.draft.fingerprint)
    _write_package(output, "corrected-reloaded", package, manifest)
    manifest["corrected_revision"] = reloaded.revision
    manifest["corrected_fingerprint"] = reloaded.draft.fingerprint
    (output / "manifest.json").write_text(json.dumps(manifest, indent=2) + "\n")

    word = Document(BytesIO(package.docx))
    assert len(word.tables) == 1
    assert tuple(tuple(cell.text for cell in row.cells) for row in word.tables[0].rows) == (READINGS.columns, *READINGS.rows)
    xml, pdf_text = _texts(package)
    for text in (xml, pdf_text):
        for fact in ("Synthetic valve remains out of service.", "~125 psi", "~3.2 kW",
                     "Q-3200", "PUMP-3200", "MOTOR-3200", "2026-10-20", "2026-10-21"):
            assert fact in text
        assert "~3.2k" not in text and SOURCE_CANARY not in text
    assert original_path.read_bytes() == original_bytes
    assert library.load_snapshot_version(draft.profile.contract, draft.profile.key, draft.period, 1) == original
    assert library.load_snapshot(draft.profile.contract, draft.profile.key, draft.period) == reloaded
    if source is not None:
        assert source.read_bytes() == source_bytes
    assert generation_blocked and pricing_checks, "Saved approximate price prose reached actual Word/PDF before a human correction."
