"""A reused global DOCX master cannot import old-client WordArt watermark text."""
from io import BytesIO
from zipfile import ZipFile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

from app.monthly_report_native_layout import build_native_docx
from test_monthly_report_native_layout import source, draft

VM = "{urn:schemas-microsoft-com:vml}"
MARKER = "OTHER-CLIENT-WATERMARK-ZX731"


def watermark(path, *, text=MARKER):
    document = Document(path)
    header = document.sections[0].header.paragraphs[0]
    pict = OxmlElement("w:pict")
    shape = etree.SubElement(pict, VM + "shape",
                             id="neutral-frame", style="width:80pt;height:20pt",
                             fillcolor="#557f7f")
    etree.SubElement(shape, VM + "textpath", string=text, style="font-size:24pt")
    header.add_run()._r.append(pict)
    document.save(path)


def test_foreign_wordart_textpath_not_preserved_from_global_master(tmp_path):
    path = source(tmp_path)
    watermark(path)
    output = build_native_docx(draft(), path, master=True)
    with ZipFile(BytesIO(output)) as archive:
        header = archive.read("word/header1.xml")
        whole = b"\n".join(archive.read(name) for name in archive.namelist()
                           if name.endswith(".xml"))
    assert MARKER.encode() not in whole
    assert b'neutral-frame' in header
    assert b'width:80pt;height:20pt' in header
    assert b'fillcolor="#557f7f"' in header


def test_current_nonmaster_wordart_not_silently_removed(tmp_path):
    path = source(tmp_path)
    watermark(path, text="CURRENT REPORT WORDART")
    output = build_native_docx(draft(), path, master=False)
    with ZipFile(BytesIO(output)) as archive:
        header = archive.read("word/header1.xml")
    # Master-reuse pruning must not arbitrarily edit a versioned,
    # user-confirmed source design in another report.
    assert b'CURRENT REPORT WORDART' in header
