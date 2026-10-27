"""Reflow must keep source-bound annual table labels, not source narrative."""
from dataclasses import replace
from hashlib import sha256
from io import BytesIO

import pytest
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from app.monthly_report_import import _text, imported_draft, inspect_docx
from app.monthly_report_model import Facility, ReportPeriod, ReportProfile, ResolvedBlock
from app.monthly_report_native_layout import build_native_docx
from app.monthly_report_start_import import analyze_starting_report, automatic_section_plans, build_starting_import


def scenario(tmp_path, *, odd_label=False):
    source = tmp_path / 'synthetic-annual-tables.docx'
    doc = Document()
    doc.add_paragraph('Synthetic Site')
    doc.add_paragraph('August 2026')
    doc.add_section(WD_SECTION_START.NEW_PAGE)
    doc.add_heading('MONTHLY ACTIVITY SUMMARY', 1)
    doc.add_paragraph('Old activity narrative')
    for index, title in enumerate(('WORK ORDER STATUS', 'PM STATUS')):
        label = doc.add_paragraph(title)
        label.runs[0].bold = True
        if odd_label and index == 0:
            symbol = OxmlElement('w:sym')
            symbol.set(qn('w:font'), 'Arial')
            symbol.set(qn('w:char'), '0031')
            label.runs[0]._r.append(symbol)
        doc.add_paragraph('Old private table introduction')
        table = doc.add_table(rows=4, cols=3)
        for cell, value in zip(table.rows[0].cells, ('WO Received', 'Jan', 'Feb')):
            cell.text = value
        for cell, value in zip(table.rows[1].cells, ('Completed', str(12 + index), '14')):
            cell.text = value
        for cell, value in zip(table.rows[2].cells, ('Pending', '2', '3')):
            cell.text = value
    doc.add_paragraph('RATE TREND')
    doc.add_paragraph('INSERT IMAGE HERE')
    doc.add_paragraph('OLD PRIVATE PAYLOAD')
    doc.save(source)
    inspection = inspect_docx(source)
    profile = ReportProfile('Synthetic contract', 'site', 'Synthetic Site', (Facility('site', 'Synthetic Site'),))
    period = ReportPeriod(2026, 9)
    analysis = analyze_starting_report(inspection, profile, period, confirmed_sites=profile.facilities)
    plans = automatic_section_plans(inspection, analysis)
    # An old cover alone cannot certify undated prose as historical. This
    # renderer fixture deliberately reviews it out, retaining only table data;
    # source-bound table-label preservation is the behavior under test below.
    with pytest.raises(ValueError, match="remaining starting-report questions"):
        build_starting_import(source, inspection, plans.values())
    text_ids = {item.id for item in inspection.items if item.kind == "text"}
    for plan in plans.values():
        if plan.get("period_review_items"):
            plan["selected"] = [identity for identity in plan["selected"] if identity not in text_ids]
            plan["texts"] = {}
            plan["questions"] = ()
            plan["approved"] = True
    mapped, _ = build_starting_import(source, inspection, plans.values())
    draft = imported_draft(profile, period, 'Synthetic Editor', mapped)
    block = next(b for b in draft.blocks if b.key == 'work_orders')
    assert not block.text  # The explicit review excludes old monthly prose.
    # Match the existing grid's no-edit normalization of empty imported rows.
    block = replace(block, extra_tables=tuple(replace(t, rows=tuple(
        row for row in t.rows if any(cell.strip() for cell in row))) for t in block.extra_tables))
    draft = replace(draft, blocks=(block, ResolvedBlock('activity_summary', 'This month', text='Current activity narrative')))
    return source, draft


def test_reflow_keeps_each_annual_label_with_its_current_table(tmp_path):
    source, draft = scenario(tmp_path)
    before = sha256(source.read_bytes()).hexdigest()
    result = Document(BytesIO(build_native_docx(draft, source)))
    text = _text(result.element)
    assert text.count('WORK ORDER STATUS') == 1
    assert text.count('PM STATUS') == 1
    assert len(result.tables) == 2
    assert result.tables[0].cell(1, 1).text == '12'
    assert result.tables[1].cell(1, 1).text == '13'
    ordered = [('table' if e.tag == qn('w:tbl') else _text(e)) for e in result.element.body]
    assert ordered.index('WORK ORDER STATUS') < ordered.index('table') < ordered.index('PM STATUS')
    for bad in ('Old private', 'OLD PRIVATE', 'INSERT IMAGE HERE', 'RATE TREND'):
        assert bad not in text
    assert 'Current activity narrative' in text
    assert sha256(source.read_bytes()).hexdigest() == before


def test_omitted_annual_tables_do_not_leave_orphan_labels(tmp_path):
    source, draft = scenario(tmp_path)
    draft = replace(draft, blocks=tuple(b for b in draft.blocks if b.key != 'work_orders'))
    result = Document(BytesIO(build_native_docx(draft, source)))
    assert not result.tables
    assert 'WORK ORDER STATUS' not in _text(result.element)
    assert 'PM STATUS' not in _text(result.element)


def test_unknown_glyph_is_not_accepted_as_plain_table_label(tmp_path):
    source, draft = scenario(tmp_path, odd_label=True)
    result = Document(BytesIO(build_native_docx(draft, source)))
    assert 'WORK ORDER STATUS' not in _text(result.element)
    assert not list(result.element.iter(qn('w:sym')))


@pytest.mark.parametrize('mode', ['wrong_source', 'swapped_table_references'])
def test_unbound_tables_cannot_borrow_source_labels(tmp_path, mode):
    source, draft = scenario(tmp_path)
    block = draft.blocks[0]
    tables = block.extra_tables
    if mode == 'wrong_source':
        tables = tuple(replace(t, reference='docx:' + '0' * 64 + ':' + t.reference.rsplit(':', 1)[-1]) for t in tables)
    else:
        tables = tuple(replace(t, reference=tables[1 - index].reference) for index, t in enumerate(tables))
    draft = replace(draft, blocks=(replace(block, extra_tables=tables), *draft.blocks[1:]))
    result = Document(BytesIO(build_native_docx(draft, source)))
    assert 'WORK ORDER STATUS' not in _text(result.element)
    assert 'PM STATUS' not in _text(result.element)


def test_label_with_hidden_unknown_property_is_not_reused(tmp_path):
    source, draft = scenario(tmp_path)
    doc = Document(source)
    label = next(p for p in doc.paragraphs if p.text == 'WORK ORDER STATUS')
    secret = OxmlElement('w:instrText')
    secret.text = 'PRIVATE PRICE $99999'
    label._p.get_or_add_pPr().append(secret)
    doc.save(source)
    # Refresh only the immutable source binding; this deliberately supplied
    # synthetic table does not certify the unknown label property.
    digest = sha256(source.read_bytes()).hexdigest()
    block = draft.blocks[0]
    tables = tuple(replace(t, reference='docx:' + digest + ':' + t.reference.rsplit(':', 1)[-1]) for t in block.extra_tables)
    draft = replace(draft, blocks=(replace(block, extra_tables=tables, references=()), *draft.blocks[1:]))
    from app.monthly_report_import import ImportError
    with pytest.raises(ImportError, match='unbalanced Word field'):
        build_native_docx(draft, source)
