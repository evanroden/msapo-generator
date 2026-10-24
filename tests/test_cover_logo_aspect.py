"""Generic selected-logo placement, with no client-specific artwork or rules."""
from dataclasses import replace
from io import BytesIO
from zipfile import ZipFile
import pytest
from PIL import Image, ImageDraw
from docx import Document
from docx.shared import Inches
from docx.oxml.ns import qn
from app.monthly_report_import import ImportMapping, inspect_docx, map_items
from app.monthly_report_model import ReportDraft, ReportPeriod, synthetic_profiles
from app.monthly_report_native_layout import build_native_docx


def source(tmp_path, size, role, alpha, repeated=False):
    image = Image.new('RGBA', size, (255,255,255,0 if alpha else 255))
    draw = ImageDraw.Draw(image)
    radius = min(size)//4
    cx, cy = size[0]//2, size[1]//2
    draw.ellipse((cx-radius,cy-radius,cx+radius,cy+radius), fill='#235e8c')
    out=BytesIO(); image.save(out,'PNG'); raw=out.getvalue()
    d=Document(); d.add_paragraph('Synthetic facility');d.add_paragraph('August 2026')
    d.add_picture(BytesIO(raw), width=Inches(3), height=Inches(1))
    if repeated: d.add_picture(BytesIO(raw), width=Inches(2), height=Inches(2))
    path=tmp_path/'source.docx'; d.save(path)
    inspection=inspect_docx(path)
    items=[i for i in inspection.items if i.kind=='image']
    mapped=map_items(path,inspection,(ImportMapping(items[0].id,role),))
    draft=ReportDraft(synthetic_profiles()[0],ReportPeriod(2026,10),'Synthetic reviewer',(),mapped.blocks)
    return path,draft,dict(mapped.assets),raw


def frames(raw):
    d=Document(BytesIO(raw)); answer=[]
    for pic in d.element.iter(qn('pic:pic')):
        ext=pic.find('./'+qn('pic:spPr')+'/'+qn('a:xfrm')+'/'+qn('a:ext'))
        outer=next(a for a in pic.iterancestors() if a.tag in (qn('wp:inline'),qn('wp:anchor'))).find(qn('wp:extent'))
        assert dict(ext.attrib)==dict(outer.attrib)
        answer.append((int(ext.get('cx')),int(ext.get('cy'))))
    return answer


@pytest.mark.parametrize('size',[(400,100),(100,400),(200,200)])
@pytest.mark.parametrize('role',['brand_logo','client_logo'])
@pytest.mark.parametrize('alpha',[False,True])
@pytest.mark.parametrize('section',[None,'cover'])
def test_selected_unchanged_logos_are_contained_without_stretch(tmp_path,size,role,alpha,section):
    path,draft,assets,raw=source(tmp_path,size,role,alpha)
    original=path.read_bytes()
    result=build_native_docx(draft,path,asset_loader=assets.__getitem__,section_key=section)
    width,height=frames(result)[0]
    assert width/height==pytest.approx(size[0]/size[1], rel=1e-5)
    assert width<=Inches(3) and height<=Inches(1)
    assert path.read_bytes()==original
    with ZipFile(BytesIO(result)) as z:
        expected = Image.open(BytesIO(raw)).convert('RGBA')
        images = [Image.open(BytesIO(z.read(n))).convert('RGBA') for n in z.namelist() if n.startswith('word/media/')]
        assert any(image.size == expected.size and image.tobytes() == expected.tobytes() for image in images)


@pytest.mark.parametrize('role',['brand_logo','client_logo'])
def test_equal_byte_unselected_logo_occurrence_keeps_geometry(tmp_path,role):
    path,draft,assets,_=source(tmp_path,(400,100),role,True,repeated=True)
    result=build_native_docx(draft,path,asset_loader=assets.__getitem__)
    assert frames(result)[0][0]/frames(result)[0][1]==pytest.approx(4)
    assert frames(result)[1]==(Inches(2),Inches(2))


@pytest.mark.parametrize('role',['brand_logo','client_logo'])
@pytest.mark.parametrize('size',[(500,100),(100,500),(200,200)])
def test_changed_selected_logo_is_contained(tmp_path,role,size):
    from app.monthly_report_cover_ui import replace_cover_picture
    from app.monthly_report_asset_review import approve_all_assets
    path,draft,assets,_=source(tmp_path,(400,100),role,True)
    data=BytesIO(); Image.new('RGB',size,'purple').save(data,'PNG')
    block=approve_all_assets(replace_cover_picture(draft.blocks[0],data.getvalue(),'.png',assets))
    result=build_native_docx(replace(draft,blocks=(block,)),path,asset_loader=assets.__getitem__)
    width,height=frames(result)[0]
    assert width/height==pytest.approx(size[0]/size[1],rel=1e-5)
    assert width<=Inches(3) and height<=Inches(1)


@pytest.mark.parametrize('crop',[{'l':'10000','r':'15000'}, {'l':'-10000','r':'35000'}])
def test_valid_asymmetric_logo_crop_is_preserved(tmp_path,crop):
    from docx.oxml import OxmlElement
    from app.monthly_report_native_layout import _contain_cover_logo
    path,_,_,_=source(tmp_path,(400,100),'client_logo',True)
    doc=Document(path); node=next(doc.element.iter(qn('a:blip')))
    rect=OxmlElement('a:srcRect'); rect.attrib.update(crop); node.addnext(rect)
    before=doc.element.xml
    _contain_cover_logo(node,400,100,doc.part,preserve_crop=True)
    assert doc.element.xml==before


@pytest.mark.parametrize('mutation',['rotated','flipped','mismatched','tiled'])
def test_unsupported_logo_geometry_is_not_guessed(tmp_path,mutation):
    from docx.oxml import OxmlElement
    from app.monthly_report_native_layout import _contain_cover_logo
    path,_,_,_=source(tmp_path,(400,100),'client_logo',True)
    doc=Document(path); pic=next(doc.element.iter(qn('pic:pic'))); node=next(pic.iter(qn('a:blip')))
    transform=pic.find(qn('pic:spPr')).find(qn('a:xfrm'))
    if mutation=='rotated': transform.set('rot','5400000')
    elif mutation=='flipped': transform.set('flipH','1')
    elif mutation=='mismatched': transform.find(qn('a:ext')).set('cx','100000')
    else:
        fill=pic.find(qn('pic:blipFill'));fill.remove(fill.find(qn('a:stretch')));fill.append(OxmlElement('a:tile'))
    before=doc.element.xml
    _contain_cover_logo(node,400,100,doc.part,preserve_crop=True)
    assert doc.element.xml==before


from tests.conftest import requires_libreoffice

@requires_libreoffice
@pytest.mark.parametrize('role',['brand_logo','client_logo'])
@pytest.mark.parametrize('size',[(400,100),(100,400),(200,200)])
@pytest.mark.parametrize('section',[None,'cover'])
def test_pdf_logo_circle_is_round(tmp_path,monkeypatch,role,size,section):
    import fitz
    from app import pdf_converter
    from app.monthly_report_render_compat import rendering_docx
    monkeypatch.setenv('EPC_DATA_DIR',str(tmp_path/'store'))
    monkeypatch.setattr(pdf_converter,'OUTPUT_DIR',tmp_path)
    path,draft,assets,_=source(tmp_path,size,role,True)
    raw=build_native_docx(draft,path,asset_loader=assets.__getitem__,section_key=section)
    rendered=tmp_path/'current.docx';rendered.write_bytes(rendering_docx(raw))
    pdf_path=pdf_converter.convert_to_pdf(rendered)
    with fitz.open(pdf_path) as pdf:
        pix=pdf[0].get_pixmap(matrix=fitz.Matrix(3,3),alpha=False)
        im=Image.frombytes('RGB',(pix.width,pix.height),pix.samples)
    mask=Image.eval(im.getchannel('R'),lambda x:255 if 30<=x<=40 else 0)
    # Exact blue synthetic artwork, excluding text and antialiasing.
    points=[(x,y) for y in range(im.height) for x in range(im.width)
            if mask.getpixel((x,y)) and all(abs(a-b)<5 for a,b in zip(im.getpixel((x,y)),(35,94,140)))]
    assert points
    w=max(x for x,y in points)-min(x for x,y in points)+1
    h=max(y for x,y in points)-min(y for x,y in points)+1
    assert abs(w-h)<=2


@pytest.mark.parametrize('section',[None,'cover'])
def test_reviewed_source_view_logo_replacement_is_contained(tmp_path,monkeypatch,section):
    from test_monthly_report_source_frame_logo_role import source as framed_source
    from test_monthly_report_native_art_review import setup
    from test_monthly_report_native_art_lifecycle import save_sample
    from app.monthly_report_preview import section_preview_docx
    monkeypatch.setenv('EPC_DATA_DIR',str(tmp_path/'store'))
    path=framed_source(tmp_path,alternate=True,divergent=True,shared_outside=False)
    doc=Document(path)
    for pic in doc.element.iter(qn('pic:pic')):
        inner=pic.find(qn('pic:spPr')).find(qn('a:xfrm')).find(qn('a:ext'))
        outer=next(a for a in pic.iterancestors() if a.tag==qn('wp:inline')).find(qn('wp:extent'))
        for extent in (inner,outer): extent.set('cx',str(Inches(3)));extent.set('cy',str(Inches(1)))
    from test_monthly_report_source_frame_logo_role import raster
    doc.sections[0].header.paragraphs[0].add_run().add_picture(BytesIO(raster('orange')),width=Inches(1))
    doc.save(path);original=path.read_bytes()
    sample=setup(path);draft,loader=save_sample(sample)
    result=build_native_docx(draft,path,asset_loader=loader) if section is None else section_preview_docx(draft,section,loader)
    logo=next(b for b in draft.blocks if b.key=='client_logo')
    im=Image.open(BytesIO(loader(logo.asset_hashes[0])))
    width,height=frames(result)[0]
    assert width/height==pytest.approx(im.width/im.height,rel=1e-5)
    assert path.read_bytes()==original


@pytest.mark.parametrize('role',['brand_logo','client_logo'])
def test_exact_selected_header_logo_is_contained(tmp_path,role):
    path,_,_,raw=source(tmp_path,(400,100),role,True)
    doc=Document();doc.add_paragraph('Synthetic facility')
    doc.sections[0].header.paragraphs[0].add_run().add_picture(BytesIO(raw),width=Inches(3),height=Inches(1))
    doc.save(path)
    inspection=inspect_docx(path);item=next(i for i in inspection.items if i.kind=='image')
    mapped=map_items(path,inspection,(ImportMapping(item.id,role),))
    draft=ReportDraft(synthetic_profiles()[0],ReportPeriod(2026,10),'Synthetic reviewer',(),mapped.blocks)
    output=Document(BytesIO(build_native_docx(draft,path,asset_loader=dict(mapped.assets).__getitem__)))
    pic=next(output.sections[0].header._element.iter(qn('pic:pic')))
    extent=pic.find(qn('pic:spPr')).find(qn('a:xfrm')).find(qn('a:ext'))
    assert int(extent.get('cx'))/int(extent.get('cy'))==pytest.approx(4)


@pytest.mark.parametrize('size',[(400,100),(100,400),(200,200)])
def test_fresh_contracts_choose_their_own_catalog_logos_and_contain(tmp_path,monkeypatch,size):
    from app import monthly_report_branding as branding, monthly_report_library as library
    from app.monthly_report_start import design_seed
    from app.monthly_report_docx import _picture
    monkeypatch.setenv('EPC_DATA_DIR',str(tmp_path/'store'))
    profiles=[replace(synthetic_profiles()[0],contract='Synthetic contract '+key,key='synthetic-'+key) for key in ('alpha','beta')]
    assets={};logos=[]
    for key,contracts,color in [('enfra',(),'green'),('alpha',(profiles[0].contract,),'blue'),('beta',(profiles[1].contract,),'purple')]:
        out=BytesIO();Image.new('RGBA',size,color).save(out,'PNG');raw=out.getvalue()
        ref=library.asset_reference(raw,'png');assets[ref]=raw
        logos.append(branding.BrandLogo(key,'Synthetic '+key,contracts,ref,'https://example.invalid/brand','https://example.invalid/brand.png','2026-10-22'))
    branding.save_branding(tuple(logos),assets,expected_revision=0,actor='Synthetic reviewer',confirmed=True)
    for index,profile in enumerate(profiles):
        draft,selected=design_seed(profile,ReportPeriod(2026,10),'Synthetic reviewer')
        blocks={b.key:b for b in draft.blocks}
        assert blocks['client_logo'].asset_hashes==(logos[index+1].asset,)
        assert blocks['brand_logo'].asset_hashes==(logos[0].asset,)
        assert not blocks['cover_photo'].asset_hashes
        assert logos[2-index].asset not in selected
        # The no-import assembler's shared image primitive uses contain too.
        doc=Document();_picture(doc,selected[logos[index+1].asset],width=2.5,height=.75)
        out=BytesIO();doc.save(out)
        width,height=frames(out.getvalue())[0]
        assert width/height==pytest.approx(size[0]/size[1],rel=1e-5)


@pytest.mark.parametrize('role',['brand_logo','client_logo'])
@pytest.mark.parametrize('master',[False,True])
@pytest.mark.parametrize('size',[(500,100),(100,500),(200,200)])
@pytest.mark.parametrize('section',[None,'cover'])
def test_legacy_catalog_replacement_without_import_references_is_contained(tmp_path,role,master,size,section):
    from app.monthly_report_cover_ui import replace_cover_picture
    from app.monthly_report_asset_review import approve_all_assets
    path,draft,assets,_=source(tmp_path,(600,100) if role=='brand_logo' else (400,100),role,True)
    out=BytesIO();Image.new('RGB',size,'purple').save(out,'PNG')
    block=approve_all_assets(replace_cover_picture(draft.blocks[0],out.getvalue(),'.png',assets))
    block=replace(block,references=())
    result=build_native_docx(replace(draft,blocks=(block,)),path,asset_loader=assets.__getitem__,master=master,section_key=section)
    width,height=frames(result)[0]
    assert width/height==pytest.approx(size[0]/size[1],rel=1e-5)


@pytest.mark.parametrize('size',[(400,100),(100,400),(200,200)])
@pytest.mark.parametrize('section',[None,'cover'])
def test_photo_orientation_matrix_retains_proportions_by_crop(tmp_path,size,section):
    path,draft,assets,_=source(tmp_path,size,'cover_photo',False)
    result=build_native_docx(draft,path,asset_loader=assets.__getitem__,section_key=section)
    doc=Document(BytesIO(result));pic=next(doc.element.iter(qn('pic:pic')))
    crop=pic.find(qn('pic:blipFill')).find(qn('a:srcRect'))
    values={k:float(crop.get(k,'0'))/100000 for k in ('l','r','t','b')}
    effective=size[0]*(1-values['l']-values['r'])/(size[1]*(1-values['t']-values['b']))
    width,height=frames(result)[0]
    assert effective==pytest.approx(width/height,rel=.001)
    assert (width,height)==(Inches(3),Inches(1))


def test_standing_photo_defaults_are_exact_site_and_contract_scoped(tmp_path,monkeypatch):
    from app import monthly_report_library as library
    from app.monthly_report_guided import _initial_draft
    from app.monthly_report_model import ResolvedBlock
    monkeypatch.setenv('EPC_DATA_DIR',str(tmp_path/'store'))
    refs=[]
    for index,contract in enumerate(('Synthetic contract alpha','Synthetic contract beta')):
        profile=replace(synthetic_profiles()[0],contract=contract,key='synthetic-site')
        state=library.save_profile(profile,expected_revision=0,actor='Synthetic reviewer',confirmed=True)
        stream=BytesIO();Image.new('RGB',(400,300),'blue' if index==0 else 'purple').save(stream,'PNG')
        raw=stream.getvalue();ref=library.asset_reference(raw,'png');refs.append(ref)
        state=library.replace_block(contract,profile.key,'cover_photo',ResolvedBlock('cover_photo','Library',asset_hashes=(ref,)),
            assets=((ref,raw),),expected_revision=state.revision,actor='Synthetic reviewer',confirmed=True,label='Cover photo')
        draft=_initial_draft(state,ReportPeriod(2026,10),'Synthetic reviewer')
        assert next(b for b in draft.blocks if b.key=='cover_photo').asset_hashes==(ref,)
        other=replace(profile,key='different-site')
        other_state=library.save_profile(other,expected_revision=0,actor='Synthetic reviewer',confirmed=True)
        other_draft=_initial_draft(other_state,ReportPeriod(2026,10),'Synthetic reviewer')
        assert not next(b for b in other_draft.blocks if b.key=='cover_photo').asset_hashes
    assert refs[0]!=refs[1]
