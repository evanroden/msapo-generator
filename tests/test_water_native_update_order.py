"""A current Water summary must precede retained native vendor payloads."""
from dataclasses import replace
from io import BytesIO

from docx import Document
from docx.oxml.ns import qn
import pytest

from app.monthly_report_import import inspect_docx, _text
from app.monthly_report_model import (ReportDraft, ReportPeriod, ReportTable,
                                    ResolvedBlock, default_sections, synthetic_profiles)
from app.monthly_report_native_layout import build_native_docx
from conftest import requires_libreoffice

UPDATE = "Current synthetic Water update before vendor results."


def native_case(tmp_path, *, reverse=False):
    source = Document()
    source.add_heading("WATER TREATMENT REPORTS", 1)
    for number in (1, 2):
        table = source.add_table(rows=2, cols=2)
        table.style = "Table Grid"
        for row, values in zip(table.rows, (("Vendor result", "Reading"),
                                          (f"Synthetic vendor {number}", str(number)))):
            for cell, value in zip(row.cells, values):
                cell.text = value
    path = tmp_path / "synthetic-water-native.docx"
    source.save(path)
    inspected = inspect_docx(path)
    tables = tuple(ReportTable(item.rows[0], item.rows[1:], f"docx:{inspected.sha256}:{item.id}")
                   for item in inspected.items if item.kind == "table")
    assert len(tables) == 2
    section = next(s for s in default_sections() if s.key == "water")
    if reverse:
        section = replace(section, blocks=tuple(reversed(section.blocks)))
    draft = ReportDraft(synthetic_profiles()[0], ReportPeriod(2026, 9), "Synthetic Editor",
                        (section,), (ResolvedBlock("water_monthly_update", "This month", text=UPDATE),
                                     ResolvedBlock("water_reports", "This month", extra_tables=tables)))
    return path, draft


@pytest.mark.parametrize("scope", [None, "water"])
@pytest.mark.parametrize("reverse_blocks", [False, True])
def test_current_update_precedes_both_unchanged_native_vendor_tables(tmp_path, scope, reverse_blocks):
    path, draft = native_case(tmp_path)
    if reverse_blocks:
        draft = replace(draft, blocks=tuple(reversed(draft.blocks)))
    original = path.read_bytes()
    result = Document(BytesIO(build_native_docx(draft, path, section_key=scope)))
    payloads = [_text(node) for node in result.element.body]
    assert payloads.index(UPDATE) < next(i for i, text in enumerate(payloads) if "Synthetic vendor 1" in text)
    assert sum(text == UPDATE for text in payloads) == 1
    assert [t._tbl.xml for t in result.tables] == [t._tbl.xml for t in Document(path).tables]
    assert path.read_bytes() == original


def test_existing_custom_update_after_pages_is_not_reordered(tmp_path):
    path, draft = native_case(tmp_path, reverse=True)
    result = Document(BytesIO(build_native_docx(draft, path)))
    payloads = [_text(node) for node in result.element.body]
    assert next(i for i, text in enumerate(payloads) if "Synthetic vendor 1" in text) < payloads.index(UPDATE)


@pytest.mark.parametrize("source,text", [("Omit", UPDATE), ("This month", "")])
def test_absent_update_does_not_insert_a_new_native_paragraph(tmp_path, source, text):
    path, draft = native_case(tmp_path)
    baseline = replace(draft, blocks=draft.blocks[1:])
    draft = replace(draft, blocks=(replace(draft.blocks[0], source=source, text=text), draft.blocks[1]))
    actual = Document(BytesIO(build_native_docx(draft, path)))
    expected = Document(BytesIO(build_native_docx(baseline, path)))
    assert actual.element.body.xml == expected.element.body.xml


def test_update_precedes_retained_native_vendor_images_even_with_reversed_input(tmp_path):
    from hashlib import sha256
    from PIL import Image
    from docx.shared import Inches
    from app.monthly_report_import import read_import_image
    source = Document()
    source.add_heading("WATER TREATMENT REPORTS", 1)
    for color in ("navy", "teal"):
        raw = BytesIO()
        Image.new("RGB", (500, 600), color).save(raw, "PNG")
        source.add_paragraph().add_run().add_picture(BytesIO(raw.getvalue()), width=Inches(2))
    path = tmp_path / "synthetic-water-images.docx"
    source.save(path)
    inspection = inspect_docx(path)
    assets, refs, captions = {}, [], []
    for item in inspection.items:
        if item.kind == "image":
            image = read_import_image(path, item, line_art=True)
            assets[sha256(image.data).hexdigest() + "." + image.extension] = image.data
            refs.append(f"docx:{inspection.sha256}:{item.id}")
            captions.append(item.text)
    assert len(assets) == 2
    section = next(s for s in default_sections() if s.key == "water")
    draft = ReportDraft(synthetic_profiles()[0], ReportPeriod(2026, 9), "Synthetic Editor", (section,), (
        ResolvedBlock("water_reports", "This month", asset_hashes=tuple(assets),
                      references=tuple(refs), asset_captions=tuple(captions)),
        ResolvedBlock("water_monthly_update", "This month", text=UPDATE)))
    original = path.read_bytes()
    result = Document(BytesIO(build_native_docx(draft, path, asset_loader=assets.__getitem__)))
    children = list(result.element.body)
    summary = next(i for i, element in enumerate(children) if _text(element) == UPDATE)
    pictures = [i for i, element in enumerate(children) if list(element.iter(qn("a:blip")))]
    assert len(pictures) == 2 and summary < min(pictures)
    assert path.read_bytes() == original


@requires_libreoffice
def test_native_combined_pdf_places_update_before_vendor_results(tmp_path, monkeypatch):
    import pymupdf
    from app import pdf_converter
    path, draft = native_case(tmp_path)
    output = tmp_path / "current-water.docx"
    output.write_bytes(build_native_docx(draft, path))
    monkeypatch.setattr(pdf_converter, "OUTPUT_DIR", tmp_path)
    monkeypatch.setattr(pdf_converter, "PDF_BACKEND", "libreoffice")
    pdf_path = pdf_converter.convert_to_pdf(output)
    with pymupdf.open(pdf_path) as pdf:
        text = "\n".join(page.get_text() for page in pdf)
        assert text.count(UPDATE) == 1
        assert text.index(UPDATE) < text.index("Synthetic vendor 1") < text.index("Synthetic vendor 2")
