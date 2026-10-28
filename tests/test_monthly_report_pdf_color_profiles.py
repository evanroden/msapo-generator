"""Color and density survive the disposable PDF input, never rewriting Word."""
from functools import lru_cache
from hashlib import sha256
from io import BytesIO
import struct
from zipfile import ZipFile

from docx import Document
from docx.shared import Inches
from PIL import Image, ImageCms, ImageDraw
import pytest

from app.monthly_report_media_privacy import sanitize_package_media
from app.monthly_report_pdf_image_budget import compact_native_photos_for_pdf


CANARY = b"SYNTHETIC-PDF-PHOTO-AUTHOR"
DPI = (144, 96)


@lru_cache(maxsize=1)
def _icc():
    return ImageCms.ImageCmsProfile(ImageCms.createProfile("sRGB")).tobytes()


def _photo(*, mode="RGB", display_metadata=True, quality=95, optimize=False):
    # A deterministic patterned photo. Fine gray detail ensures the old RGB
    # conversion actually shrinks an L image instead of passing by coincidence.
    picture = Image.new("RGB", (2400, 1600), "white")
    draw = ImageDraw.Draw(picture)
    step = 23 if mode == "L" else 40
    for y in range(0, 1600, step):
        for x in range(0, 2400, step):
            if mode == "L":
                shade = ((x * 17) ^ (y * 31)) % 256
                color = (shade, shade, shade)
            else:
                color = ((x // 7) % 256, (y // 5) % 256, ((x + y) // 3) % 256)
            draw.rectangle((x, y, x + step - 1, y + step - 1), fill=color)
    picture = picture.convert(mode)
    exif = Image.Exif()
    exif[274] = 1
    exif[315] = CANARY.decode()
    options = {"exif": exif.tobytes()}
    if display_metadata:
        options["dpi"] = DPI
        if mode == "RGB":
            options["icc_profile"] = _icc()
    encoded = BytesIO()
    picture.save(encoded, "JPEG", quality=quality, optimize=optimize, **options)
    return encoded.getvalue()


def _document(photo):
    document = Document()
    document.add_paragraph("Synthetic native report color proof")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Equipment"
    table.cell(0, 1).text = "Observed status"
    table.cell(1, 0).text = "Synthetic pump"
    table.cell(1, 1).text = "Reviewed"
    document.add_picture(BytesIO(photo), width=Inches(5))
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _parts(raw):
    with ZipFile(BytesIO(raw)) as archive:
        return {info.filename: archive.read(info) for info in archive.infolist()}


def _media(parts):
    name, = (name for name in parts if name.startswith("word/media/"))
    return name


def _assert_geometry_and_other_parts_unchanged(before, after):
    old, new = _parts(before), _parts(after)
    media = _media(old)
    assert old.keys() == new.keys()
    assert {name: raw for name, raw in old.items() if name != media} == {
        name: raw for name, raw in new.items() if name != media
    }
    before_shape, = Document(BytesIO(before)).inline_shapes
    after_shape, = Document(BytesIO(after)).inline_shapes
    assert (after_shape.width, after_shape.height) == (before_shape.width, before_shape.height)


def _assert_profiled_shrink(before, after):
    old, new = _parts(before), _parts(after)
    media = _media(old)
    assert len(new[media]) < len(old[media])
    with Image.open(BytesIO(new[media])) as photo:
        assert photo.mode == "RGB" and photo.size == (1400, 933)
        assert photo.info.get("icc_profile") == _icc()
        assert photo.info.get("dpi") == DPI
        assert 315 not in photo.getexif()
    assert CANARY not in new[media]
    _assert_geometry_and_other_parts_unchanged(before, after)


@pytest.mark.parametrize("sanitized", [False, True], ids=["raw-pinned", "sanitized-allowlist"])
def test_eligible_rgb_preserves_exact_profile_density_and_docx_geometry(tmp_path, sanitized):
    original = _document(_photo())
    master = tmp_path / "synthetic-color-master.docx"
    master.write_bytes(original)
    raw = sanitize_package_media(original) if sanitized else original
    before = _parts(raw)
    assert (CANARY not in before[_media(before)]) == sanitized
    with Image.open(BytesIO(before[_media(before)])) as photo:
        assert photo.info["icc_profile"] == _icc() and photo.info["dpi"] == DPI
    compacted = compact_native_photos_for_pdf(raw, pinned_source=master)
    _assert_profiled_shrink(raw, compacted)
    assert master.read_bytes() == original
    assert _parts(raw) == before


@pytest.mark.parametrize("mode", ["CMYK", "L"])
def test_other_color_modes_remain_byte_identical_in_pdf_copy(tmp_path, mode):
    original = _document(_photo(mode=mode))
    master = tmp_path / (mode + "-master.docx")
    master.write_bytes(original)
    clean = sanitize_package_media(original)
    media = _parts(clean)
    with Image.open(BytesIO(media[_media(media)])) as photo:
        assert photo.mode == mode and photo.size == (2400, 1600)
    assert compact_native_photos_for_pdf(clean, pinned_source=master) == clean
    assert master.read_bytes() == original


def test_rgb_without_profile_or_density_keeps_jpeg_defaults(tmp_path):
    original = _document(_photo(display_metadata=False))
    master = tmp_path / "unprofiled-master.docx"
    master.write_bytes(original)
    clean = sanitize_package_media(original)
    compacted = compact_native_photos_for_pdf(clean, pinned_source=master)
    before, after = _parts(clean), _parts(compacted)
    media = _media(before)
    assert len(after[media]) < len(before[media])
    with Image.open(BytesIO(before[media])) as old, Image.open(BytesIO(after[media])) as new:
        assert new.size == (1400, 933)
        assert old.info.get("icc_profile") is new.info.get("icc_profile") is None
        assert old.info.get("dpi") is new.info.get("dpi") is None
        assert (new.info.get("jfif_unit"), new.info.get("jfif_density")) == (
            old.info.get("jfif_unit"), old.info.get("jfif_density"),
        )
    assert CANARY not in after[media]


def test_already_small_jpeg_is_not_replaced_by_larger_reencoding(tmp_path):
    # Full-size, low-quality source is smaller than the normal bounded encode.
    # Its eligibility must not override the separate only-if-smaller guard.
    original = _document(_photo(quality=5, optimize=True))
    master = tmp_path / "compressed-master.docx"
    master.write_bytes(original)
    clean = sanitize_package_media(original)
    assert compact_native_photos_for_pdf(clean, pinned_source=master) == clean
    assert master.read_bytes() == original


@pytest.mark.parametrize("unit,density", [
    (2, (76, 77)), (2, (65535, 65535)), (1, (0, 0)), (1, (0, 144)),
], ids=["fractional-dpi", "large-dpi", "zero-dpi", "zero-axis-dpi"])
def test_density_not_exactly_encodable_as_jpeg_dpi_keeps_original(tmp_path, unit, density):
    # JFIF also supports pixels per centimeter. Pillow exposes those values
    # as DPI, but its JPEG writer only encodes integer 16-bit DPI values.
    source = _document(_photo())
    parts = _parts(source)
    media_name = _media(parts)
    photo = bytearray(parts[media_name])
    jfif = photo.index(b"JFIF\x00")
    photo[jfif + 7] = unit
    photo[jfif + 8:jfif + 12] = struct.pack(">HH", *density)
    # Patch the media after Word has explicit drawing extents, so zero-density
    # fixtures do not depend on python-docx deriving an intrinsic image size.
    output = BytesIO()
    with ZipFile(BytesIO(source)) as before, ZipFile(output, "w") as after:
        for info in before.infolist():
            after.writestr(info, bytes(photo) if info.filename == media_name else before.read(info))
    original = output.getvalue()
    master = tmp_path / "unusual-density-master.docx"
    master.write_bytes(original)
    clean = sanitize_package_media(original)
    media = _parts(clean)
    with Image.open(BytesIO(media[_media(media)])) as image:
        assert image.info["jfif_unit"] == unit
        assert image.info["jfif_density"] == density
        factor = 2.54 if unit == 2 else 1
        assert image.info["dpi"] == pytest.approx(tuple(value * factor for value in density))
    assert compact_native_photos_for_pdf(clean, pinned_source=master) == clean
    assert master.read_bytes() == original


def test_generation_sends_profiled_copy_to_backend_and_keeps_native_word(monkeypatch, tmp_path):
    """Capture the conversion boundary; this does not claim real PDF rendering."""
    from app import monthly_report_docx as engine, monthly_report_designs as designs
    from app import monthly_report_render_jobs as jobs
    from app.monthly_report_model import ReportDraft, ReportPeriod, ReportProfile, Facility

    original = _document(_photo())
    master = tmp_path / "pinned-master.docx"
    master.write_bytes(original)
    master_hash = sha256(original).hexdigest()
    clean = sanitize_package_media(original)
    profile = ReportProfile("Synthetic Contract", "color", "Synthetic Site",
                            (Facility("s", "Synthetic Site"),))
    draft = ReportDraft(profile, ReportPeriod(2026, 10), "Synthetic Editor", (), ())
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path / "runtime"))
    monkeypatch.setattr(engine, "assemble_docx", lambda *args, **kwargs: original)
    monkeypatch.setattr(designs, "render_profile_for", lambda *_: None)
    monkeypatch.setattr(designs, "source_for", lambda *_: master)
    monkeypatch.setattr(engine.pdf_converter, "OUTPUT_DIR", tmp_path)
    captured = []

    def capture_converter_input(path, **kwargs):
        captured.append(path.read_bytes())
        pdf = tmp_path / "synthetic-backend.pdf"
        pdf.write_bytes(b"%PDF-synthetic-backend-capture")
        return pdf

    monkeypatch.setattr(jobs, "convert_to_pdf", capture_converter_input)
    package = engine._generate_report(draft)
    assert package.pdf == b"%PDF-synthetic-backend-capture" and not package.pdf_error
    assert package.docx == clean
    assert len(captured) == 1
    _assert_profiled_shrink(package.docx, captured[0])
    media = _parts(package.docx)
    with Image.open(BytesIO(media[_media(media)])) as photo:
        assert photo.size == (2400, 1600) and photo.info["icc_profile"] == _icc()
    assert CANARY not in media[_media(media)]
    assert sha256(master.read_bytes()).hexdigest() == master_hash
