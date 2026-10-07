"""Synthetic regression reports; private samples never become fixtures."""

from dataclasses import replace
from io import BytesIO

from docx import Document
from docx.enum.section import WD_SECTION_START
import pytest

from app.monthly_report_checks import preflight
from app.monthly_report_content_policy import contains_price, page_status, page_fingerprint, price_column, table_has_pricing
from app.monthly_report_docx import assemble_docx
from app.monthly_report_import import inspect_docx, imported_draft, ImportItem
from app.monthly_report_model import ReportPeriod, ResolvedBlock, synthetic_draft, synthetic_profiles
from app.monthly_report_sections import section_reviews, build_section_import, table_without_prices, apply_section_omissions
from app.monthly_report_setup import design_profile, merge_blocks, new_month_draft, merge_drafts
from app import monthly_report_sources as sources, monthly_report_library as library
from test_monthly_report_setup import make_docx


@pytest.mark.parametrize("text", ["Price: 1200", "USD 30", "30 dollars", "Quoted amount: 450.00", "$1,000.00", "Labour rate 75", "Unit price\n\n1,200.00"])
def test_detects_client_pricing(text):
    assert contains_price(text)
    assert page_status(text)[0] == "pricing"


def test_technical_measurements_are_not_prices():
    assert not contains_price("Flow rate 150 gpm; conductivity 3000 uS/cm; temperature 110 F")
    assert not price_column("Flow rate", "number")
    assert not price_column("Capacity", "number")


@pytest.mark.parametrize("text,status", [("", "blank"), ("Authorized by: Synthetic Role\nSignature:", "signature"),
                                       ("Terms and conditions. Limitation of liability. Governing law.", "legal")])
def test_unneeded_pages_are_identified(text, status):
    assert page_status(text)[0] == status
    assert page_status("", unreadable=True)[0] == "review"


def test_price_free_table_import_keeps_every_schema_and_zero(tmp_path, monkeypatch):
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path / "data"))
    doc = Document()
    doc.add_paragraph("=")
    doc.add_paragraph("TABLE OF CONTENTS")
    doc.add_paragraph("Section 7: Water Treatment Reports")
    doc.add_section(WD_SECTION_START.NEW_PAGE)
    doc.add_heading("2 Monthly Activity Summary", 1)
    doc.add_paragraph("Inspected a synthetic pump during September 2026.")
    for headers, row in ((('Measure', 'Count', 'Unit Price'), ('Inspections', '0', '500')),
                         (('Asset', 'Verified'), ('Synthetic pump', 'No'))):
        table = doc.add_table(rows=2, cols=len(headers))
        for c, value in zip(table.rows[0].cells, headers):
            c.text = value
        for c, value in zip(table.rows[1].cells, row):
            c.text = value
    doc.add_heading("7 WATER TREATMENT", 1)
    doc.add_paragraph("Synthetic treatment was completed.")
    path = tmp_path / "synthetic.docx"
    doc.save(path)
    original = path.read_bytes()
    inspection = inspect_docx(path)
    sections = section_reviews(inspection)
    assert [s.key for s in sections] == ['cover', 'activity', 'water']
    assert not any(i.text == '=' for s in sections for i in s.items)
    activity = next(s for s in sections if s.key == 'activity')
    table = next(i for i in activity.items if i.kind == 'table')
    cleaned, removed = table_without_prices(table)
    assert removed == (2,) and cleaned.rows == (('Inspections', '0'),)
    plans = ({'key': 'activity', 'target': 'activity', 'approved': True, 'selected': [i.id for i in activity.items]},)
    mapped, mappings = build_section_import(path, inspection, plans)
    profile = design_profile(synthetic_profiles()[0], inspection, mappings)
    draft = imported_draft(profile, ReportPeriod(2026, 9), 'Synthetic Editor', mapped)
    raw = assemble_docx(draft)
    assert assemble_docx(draft) == raw
    output = Document(BytesIO(raw))
    assert len(output.tables) == 2
    values = [[c.text for row in table.rows for c in row.cells] for table in output.tables]
    assert '500' not in str(values) and 'Unit Price' not in str(values)
    assert '0' in values[0] and 'No' in values[1]
    library.save_report_setup(profile, draft, path, {'sha256': inspection.sha256}, assets=mapped.assets,
                              expected_revision=0, actor='Synthetic Editor', confirmed=True)
    assert library.load_imported_draft(profile.contract, profile.key) == draft
    assert path.read_bytes() == original
    following = new_month_draft(draft, ReportPeriod(2026, 10))
    old_tables = next(b.extra_tables for b in draft.blocks if b.extra_tables)
    new_tables = next(b.extra_tables for b in following.blocks if b.extra_tables)
    assert [t.columns for t in new_tables] == [t.columns for t in old_tables]
    assert all(not t.rows for t in new_tables)


def test_selected_image_known_to_contain_price_cannot_be_confirmed_away(tmp_path):
    path = make_docx(tmp_path)
    inspection = inspect_docx(path)
    item = next(i for i in inspection.items if i.kind == 'image')
    plan = {'key': item.section, 'target': item.section, 'approved': True, 'selected': [item.id],
            'image_notes': {item.image_part: 'Service performed. Total price USD 400.'}}
    with pytest.raises(ValueError, match='pricing'):
        build_section_import(path, inspection, (plan,))


def test_pricing_and_image_review_enforced_at_export_and_invalidated():
    draft = synthetic_draft(synthetic_profiles()[0], ReportPeriod(2026, 9))
    block = replace(draft.blocks[0], text='Completed repair for $250.')
    draft = replace(draft, blocks=tuple(block if b.key == block.key else b for b in draft.blocks))
    assert any(c.code == 'pricing' and c.blocking for c in preflight(draft))
    with pytest.raises(ValueError, match='pricing'):
        assemble_docx(draft)
    page = ResolvedBlock('org_chart', 'Library', asset_hashes=('a.png',))
    page = replace(page, client_reviewed_fingerprint=page.fingerprint)
    merged = merge_blocks(page, replace(page, asset_hashes=('b.png',)))
    assert merged.asset_hashes == ('a.png', 'b.png')
    assert merged.client_reviewed_fingerprint != merged.fingerprint


def test_pdf_defaults_exclude_prices_legal_blank_and_signature_pages(tmp_path, monkeypatch):
    import fitz
    monkeypatch.setenv('EPC_DATA_DIR', str(tmp_path))
    profile = synthetic_profiles()[0]
    with fitz.open() as pdf:
        for text in ('Completed inspection of synthetic pump. No defects found.', 'Repair quotation: $750.00',
                     'Terms and conditions. Limitation of liability. Governing law.', 'Signature: Synthetic Role', ''):
            page = pdf.new_page()
            page.insert_text((50, 50), text)
        raw = pdf.tobytes()
    content, _ = sources.ingest(profile, 'synthetic.pdf', raw)
    assert content.source.selected_pages == (1,)
    source = replace(content.source, selected_pages=(2,))
    source = replace(source, client_page_reviews=((2, page_fingerprint(source, 2)),))
    with pytest.raises(ValueError, match='Prices'):
        sources.prepare_pages(profile, (replace(content, source=source),), ((source.id, 'vendor_reports'),))
    source = replace(content.source, selected_pages=(1,))
    source = replace(source, client_page_reviews=((1, page_fingerprint(source, 1)),))
    prepared = sources.prepare_pages(profile, (replace(content, source=source),), ((source.id, 'vendor_reports'),))
    assert prepared[0].client_reviewed_fingerprint == prepared[0].fingerprint
    source = replace(source, captions=((1, 'Price: 150'),))
    with pytest.raises(ValueError, match='Check page'):
        sources.prepare_pages(profile, (replace(content, source=source),), ((source.id, 'vendor_reports'),))


def test_multiline_table_headers_and_separate_total_cells_cannot_leak_prices():
    item = ImportItem('table','table','word/document.xml',1,1,'activity','work_orders',
                      rows=(('Description','Unit'),('','Price'),('Synthetic task','500')))
    table,removed=table_without_prices(item)
    assert removed==(1,) and all('500' not in row for row in table.rows)
    assert table_has_pricing(item.rows[0],item.rows[1:])
    assert table_has_pricing(('Description','Value'),(('Total','500'),))
    assert table_has_pricing(('Description','Value'),(('Price','500'),))
    assert not table_has_pricing(('Measure','Count'),(('Total','500'),))


def test_numbered_body_heading_can_end_contents_in_same_word_section(tmp_path):
    doc=Document()
    doc.add_paragraph('Table of contents')
    doc.add_paragraph('Section 1 Organizational Chart')
    doc.add_paragraph('Section 2 Monthly Activity Summary')
    doc.add_heading('Section 1 Organizational Chart',1)
    doc.add_paragraph('Synthetic organizational details.')
    doc.add_heading('Section 2 Monthly Activity Summary',1)
    doc.add_paragraph('Synthetic current inspection was completed.')
    path=tmp_path/'synthetic-toc.docx'; doc.save(path)
    inspection=inspect_docx(path)
    assert next(i for i in inspection.items if i.text=='Synthetic organizational details.').section=='organization'
    assert next(i for i in inspection.items if i.text=='Synthetic current inspection was completed.').section=='activity'


def test_footer_remains_a_layout_destination_and_unsupported_requires_decision(tmp_path):
    path=make_docx(tmp_path)
    doc=Document(path); doc.sections[0].footer.paragraphs[0].text='100 Example Way | example.invalid'
    doc.save(path)
    inspection=inspect_docx(path)
    footer=next(i for s in section_reviews(inspection) for i in s.items if i.suggested_slot=='footer_text')
    plan={'key':'cover','target':'cover','approved':True,'selected':[footer.id],'destinations':{footer.id:'footer_text'}}
    mapped,_=build_section_import(path,inspection,(plan,))
    draft=imported_draft(synthetic_profiles()[0],ReportPeriod(2026,9),'Synthetic Editor',mapped)
    assert draft.address_line==footer.text
    unsupported=ImportItem('drawing','unsupported','word/document.xml',99,1,'activity')
    inspection=replace(inspection,items=(*inspection.items,unsupported))
    text=next(i for i in inspection.items if i.suggested_slot=='activity_summary')
    plan={'key':'activity','target':'activity','approved':True,'selected':[text.id]}
    with pytest.raises(ValueError,match='unread drawings'):
        build_section_import(path,inspection,(plan,))
    mapped,_=build_section_import(path,inspection,(plan|{'unsupported_reviewed':True},))
    assert mapped.blocks


def test_explicit_omission_wins_after_preserving_a_partial_report():
    draft=synthetic_draft(synthetic_profiles()[0],ReportPeriod(2026,9))
    incoming=replace(draft,sections=tuple(replace(s,included=False) if s.key=='activity' else s for s in draft.sections),blocks=tuple(b for b in draft.blocks if b.key!=draft.sections[1].blocks[0].key))
    merged=merge_drafts(draft,incoming)
    assert next(s for s in merged.sections if s.key=='activity').included
    result=apply_section_omissions(merged,({'key':'activity','approved':True,'omit':True},))
    assert not next(s for s in result.sections if s.key=='activity').included
    assert next(s for s in result.sections if s.key=='organization').included
    assert draft.sections[1].included  # Prior immutable draft is preserved.
