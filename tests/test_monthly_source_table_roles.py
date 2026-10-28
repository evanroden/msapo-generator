"""Explicit table instructions/examples require a source-bound review choice."""
from dataclasses import replace

from docx import Document
import pytest

from app.monthly_report_import import inspect_docx
from app.monthly_report_model import synthetic_profiles
from app.monthly_report_start_import import analyze_starting_report, automatic_section_plans, build_starting_import
from test_monthly_report_start_import import PERIOD, report_file
from test_monthly_report_source_content_roles import role_app, STANDING


def source(tmp_path, cue, *, header=False, month="September 2026"):
    path = report_file(tmp_path, month=month, activity=(STANDING,), contact=False)
    doc = Document(path)
    table = doc.add_table(rows=2, cols=2)
    for row, values in zip(table.rows, ((cue if header else "Task", "Recorded date"),
                                       ("Synthetic inspection" if header else cue, "September 4, 2026"))):
        for cell, value in zip(row.cells, values):
            cell.text = value
    doc.save(path)
    inspection = inspect_docx(path)
    analysis = analyze_starting_report(inspection, synthetic_profiles()[0], PERIOD)
    plans = automatic_section_plans(inspection, analysis)
    item = next(item for item in inspection.items if item.kind == "table")
    return path, inspection, analysis, plans, item


@pytest.mark.parametrize("cue", ["Author instructions: Enter this month's work.",
                                 "Example: Synthetic maintenance", "[Insert task here]"])
def test_explicit_table_cues_cannot_become_automatic_report_facts(tmp_path, cue):
    path, inspection, analysis, plans, item = source(tmp_path, cue, month="July 2026")
    assert item.id not in analysis.current_items
    assert not plans["activity"]["approved"] and item.id not in plans["activity"]["selected"]
    with pytest.raises(ValueError, match="starting-report questions"):
        build_starting_import(path, inspection, plans.values())


def test_table_example_can_be_excluded_without_losing_standing_neighbor(tmp_path):
    path, inspection, _, plans, item = source(tmp_path, "Example: Synthetic task")
    original = path.read_bytes()
    app = role_app(path, inspection, plans)
    choice = next(r for r in app.radio if r.label == "Use source table 1")
    choice.set_value("exclude").run()
    assert not app.exception and app.session_state["plans"]["activity"]["approved"]
    mapped, _ = build_starting_import(path, inspection, app.session_state["plans"].values())
    assert any(block.text == STANDING for block in mapped.blocks)
    assert not any(block.extra_tables for block in mapped.blocks)
    reopened = role_app(path, inspection, app.session_state["plans"])
    assert next(r for r in reopened.radio if r.label == "Use source table 1").value == "exclude"
    assert path.read_bytes() == original


def test_verified_table_edits_have_exact_receipt_and_source_identity(tmp_path):
    path, inspection, _, plans, item = source(tmp_path, "[Insert task here]")
    app = role_app(path, inspection, plans)
    next(r for r in app.radio if r.label == "Use source table 1").set_value("include").run()
    assert not app.session_state["plans"]["activity"]["approved"]
    key = next(key for key, value in app.session_state.filtered_state.items()
               if "_role_table_" in key and isinstance(value, dict) and "edited_rows" in value)
    app.session_state[key] = {"edited_rows": {0: {"Task": "Verified synthetic inspection"}},
                              "added_rows": [], "deleted_rows": []}
    app.run()
    assert not app.exception and app.session_state["plans"]["activity"]["approved"]
    reviewed = app.session_state["plans"]
    mapped, _ = build_starting_import(path, inspection, reviewed.values())
    table_block = next(block for block in mapped.blocks if block.extra_tables)
    assert table_block.extra_tables[0].rows[0][0] == "Verified synthetic inspection"
    assert f"docx:{inspection.sha256}:{item.id}" in table_block.references
    reviewed["activity"]["tables"][item.id] = replace(reviewed["activity"]["tables"][item.id],
                                                     rows=(("Unreviewed change", "September 4, 2026"),))
    with pytest.raises(ValueError, match="verified table"):
        build_starting_import(path, inspection, reviewed.values())


def test_header_placeholder_has_a_real_edit_control(tmp_path):
    path, inspection, _, plans, _ = source(tmp_path, "[Enter task heading]", header=True)
    app = role_app(path, inspection, plans)
    next(r for r in app.radio if r.label == "Use source table 1").set_value("include").run()
    next(t for t in app.text_input if t.label == "Verified column 1 heading").set_value("Task").run()
    assert not app.exception and app.session_state["plans"]["activity"]["approved"]
    mapped, _ = build_starting_import(path, inspection, app.session_state["plans"].values())
    assert next(b for b in mapped.blocks if b.extra_tables).extra_tables[0].columns[0] == "Task"


def test_ordinary_sample_word_and_standing_table_remain_unchanged(tmp_path):
    path, inspection, analysis, plans, item = source(tmp_path, "A sample was collected for routine inspection.")
    assert plans["activity"]["approved"] and item.id in analysis.current_items
    mapped, _ = build_starting_import(path, inspection, plans.values())
    assert next(b for b in mapped.blocks if b.extra_tables).extra_tables[0].rows[0][0] == item.rows[1][0]


def test_direct_selection_and_changed_source_cannot_bypass_table_receipt(tmp_path):
    from app.monthly_report_sections import table_without_prices
    path, inspection, _, plans, item = source(tmp_path, "Example: Synthetic task")
    plans["activity"].update(approved=True, selected=[item.id], tables={item.id: table_without_prices(item)[0]})
    with pytest.raises(ValueError, match="source table instructions"):
        build_starting_import(path, inspection, plans.values())
    app = role_app(path, inspection, plans)
    next(r for r in app.radio if r.label == "Use source table 1").set_value("exclude").run()
    reviewed = app.session_state["plans"]
    with pytest.raises(ValueError, match="source table instructions"):
        build_starting_import(path, replace(inspection, sha256="f" * 64), reviewed.values())
    reviewed["activity"]["selected"].append(item.id)
    with pytest.raises(ValueError, match="excluded source table"):
        build_starting_import(path, inspection, reviewed.values())


def test_verified_edits_survive_exclusion_and_same_source_reentry(tmp_path):
    path, inspection, _, plans, _ = source(tmp_path, "[Insert task here]")
    app = role_app(path, inspection, plans)
    next(r for r in app.radio if r.label == "Use source table 1").set_value("include").run()
    key = next(key for key, value in app.session_state.filtered_state.items()
               if "_role_table_" in key and isinstance(value, dict) and "edited_rows" in value)
    app.session_state[key] = {"edited_rows": {0: {"Task": "Verified standing task"}}, "added_rows": [], "deleted_rows": []}
    app.run()
    next(r for r in app.radio if r.label == "Use source table 1").set_value("exclude").run()
    reopened = role_app(path, inspection, app.session_state["plans"])
    next(r for r in reopened.radio if r.label == "Use source table 1").set_value("include").run()
    assert not reopened.exception and reopened.session_state["plans"]["activity"]["approved"]
    mapped, _ = build_starting_import(path, inspection, reopened.session_state["plans"].values())
    assert next(b for b in mapped.blocks if b.extra_tables).extra_tables[0].rows[0][0] == "Verified standing task"


def test_excluded_table_stays_out_of_saved_native_word_output(tmp_path, monkeypatch):
    from io import BytesIO
    from app import monthly_report_library as library
    from app.monthly_report_docx import assemble_docx
    from test_monthly_report_start_import import start_app, save_start
    cue = "Example: Synthetic table task"
    path, _, _, _, _ = source(tmp_path, cue)
    original = path.read_bytes()
    app = start_app(monkeypatch, tmp_path, path)
    next(r for r in app.radio if r.label == "Use source table 1").set_value("exclude").run()
    draft = save_start(app)
    assert library.load_imported_draft(draft.profile.contract, draft.profile.key) == draft
    assert library.imported_original(draft.profile.contract, draft.profile.key).read_bytes() == original
    draft = replace(draft, sections=tuple(replace(s, included=s.key == "activity") for s in draft.sections))
    document = Document(BytesIO(assemble_docx(draft, acknowledged_fingerprint=draft.fingerprint)))
    text = "\n".join(document.element.body.itertext())
    assert cue not in text and STANDING in text
    assert path.read_bytes() == original


def test_organization_table_example_offers_table_review_instead_of_page_copy(tmp_path):
    from streamlit.testing.v1 import AppTest
    from test_monthly_report_source_content_roles import APP
    path, _, _, _, _ = source(tmp_path, "Example: Synthetic role")
    doc = Document(path)
    next(p for p in doc.paragraphs if p.text.startswith("2 Monthly")).text = "1 Organizational Chart"
    doc.save(path)
    inspection = inspect_docx(path)
    analysis = analyze_starting_report(inspection, synthetic_profiles()[0], PERIOD)
    plans = automatic_section_plans(inspection, analysis)
    app = AppTest.from_string(APP.replace("s.key == 'activity'", "s.key == 'organization'"))
    app.session_state["path"] = path
    app.session_state["inspection"] = inspection
    app.session_state["plans"] = plans
    app.run()
    assert not app.exception
    assert any(r.label == "Use source table 1" for r in app.radio)
    assert not any(b.label == "Show original chart pages" for b in app.button)
    next(r for r in app.radio if r.label == "Use source table 1").set_value("exclude").run()
    assert app.session_state["plans"]["organization"]["approved"]


def test_price_only_marked_table_has_a_working_exclusion(tmp_path):
    path = report_file(tmp_path, activity=(STANDING,), contact=False)
    doc = Document(path)
    table = doc.add_table(rows=2, cols=1)
    table.cell(0, 0).text = "Cost"
    table.cell(1, 0).text = "Example: $500"
    doc.save(path)
    inspection = inspect_docx(path)
    plans = automatic_section_plans(inspection, analyze_starting_report(inspection, synthetic_profiles()[0], PERIOD))
    app = role_app(path, inspection, plans)
    choice = next(r for r in app.radio if r.label == "Use source table 1")
    assert "Use verified table edits" not in choice.options
    choice.set_value("exclude").run()
    assert not app.exception and app.session_state["plans"]["activity"]["approved"]


def test_verified_table_replacement_reaches_saved_native_output(tmp_path, monkeypatch):
    from io import BytesIO
    from app import monthly_report_library as library
    from app.monthly_report_docx import assemble_docx
    from test_monthly_report_start_import import start_app, save_start
    path, _, _, _, _ = source(tmp_path, "[Insert task here]")
    original = path.read_bytes()
    app = start_app(monkeypatch, tmp_path, path)
    next(r for r in app.radio if r.label == "Use source table 1").set_value("include").run()
    key = next(key for key, value in app.session_state.filtered_state.items()
               if "_role_table_" in key and isinstance(value, dict) and "edited_rows" in value)
    app.session_state[key] = {"edited_rows": {0: {"Task": "Verified synthetic inspection"}}, "added_rows": [], "deleted_rows": []}
    app.run()
    draft = save_start(app)
    assert library.load_imported_draft(draft.profile.contract, draft.profile.key) == draft
    draft = replace(draft, sections=tuple(replace(s, included=s.key == "activity") for s in draft.sections))
    document = Document(BytesIO(assemble_docx(draft, acknowledged_fingerprint=draft.fingerprint)))
    text = "\n".join(document.element.body.itertext())
    assert "[Insert task here]" not in text and "Verified synthetic inspection" in text and STANDING in text
    assert path.read_bytes() == original
