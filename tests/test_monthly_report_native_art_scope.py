"""Synthetic pre-render section chrome closure and output-routing negatives."""
from collections import defaultdict
from copy import deepcopy
from dataclasses import asdict
from io import BytesIO
from zipfile import ZipFile, ZIP_DEFLATED

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree as ET
import pytest

from app.monthly_report_compatibility_view import derive_compatibility_view, CompatibilityError
from app.monthly_report_import import _heading, _is_native_divider_background, _text
from app.monthly_report_native_art_scope import (
    expected_header_scope, verify_header_routing, source_section_plan,
    _routing, _package, W, R, REL, KINDS,
)
from test_monthly_report_source_frame_logo_role import source


@pytest.fixture
def sample(tmp_path):
    path = source(tmp_path, divergent=True)
    doc = Document(path)
    trailing = [p._p for p in doc.paragraphs[-2:]]
    for node in trailing: node.getparent().remove(node)
    first = doc.sections[0]
    first.header.paragraphs[0].text = 'Synthetic default one'
    first.even_page_header.paragraphs[0].text = 'Synthetic even one'
    first.first_page_header.paragraphs[0].text = 'Synthetic first one'
    first.footer.paragraphs[0].text = 'Synthetic footer one'
    first.different_first_page_header_footer = True
    doc.settings.odd_and_even_pages_header_footer = True
    second = doc.add_section()
    second.header.is_linked_to_previous = False
    second.header.paragraphs[0].text = 'Synthetic default two'
    second.different_first_page_header_footer = False
    for node in trailing: doc.element.body.insert(len(doc.element.body) - 1, node)
    third = doc.add_section()
    third.header.is_linked_to_previous = False
    third.header.paragraphs[0].text = 'Synthetic default three'
    doc.add_heading('TRAINING SUMMARY', 1)
    doc.add_paragraph('Synthetic training completed.')
    doc.save(path)
    view = derive_compatibility_view(path.read_bytes())
    return path, view


def rewrite(raw, mutate):
    with ZipFile(BytesIO(raw)) as z: parts = {n:z.read(n) for n in z.namelist()}
    mutate(parts)
    stream=BytesIO()
    with ZipFile(stream,'w',ZIP_DEFLATED) as z:
        for n,v in parts.items(): z.writestr(n,v)
    return stream.getvalue()


def xml(parts, name, mutate):
    root=ET.fromstring(parts[name]); mutate(root); parts[name]=ET.tostring(root)


def _preview(view, scope):
    """Construct source-only selected slice, not expected scope from its output."""
    def mutate(parts):
        root=ET.fromstring(parts['word/document.xml']); body=root.find(W+'body'); original=tuple(body)
        with _package(view.rendition_bytes) as archive: _,routes,ends,_=_routing(archive)
        inherited={}
        # Same explicit inheritance materialization as the native writer.
        for node in original:
            props=node if node.tag==W+'sectPr' else node.find('.//'+W+'sectPr')
            if props is None: continue
            present={(child.tag,child.get(W+'type','default')):child for child in props
                     if child.tag in (W+'headerReference',W+'footerReference')}
            for key,child in inherited.items():
                if key not in present: props.insert(0,deepcopy(child))
            inherited.update({key:deepcopy(child) for key,child in present.items()})
        closing=next((deepcopy(p) for node in original[max(scope.selected_positions)-1:]
                      if (p := node if node.tag==W+'sectPr' else node.find('.//'+W+'sectPr')) is not None),None)
        for node in list(body): body.remove(node)
        for position in scope.selected_positions: body.append(original[position-1])
        for props in list(body[-1].iter(W+'sectPr')): props.getparent().remove(props)
        body.append(closing)
        parts['word/document.xml']=ET.tostring(root)
    return rewrite(view.rendition_bytes,mutate)


def _legacy_selector(original, inspection):
    items=defaultdict(list)
    for item in inspection.items:
        if item.part=='word/document.xml':items[item.position].append(item)
    section_by_pos,section,headings={},'cover',set()
    for position,element in enumerate(original,1):
        group=items[position]
        detected=next((i.section for i in group if i.section),'')
        if detected:section=detected
        section_by_pos[position]=section
        if any(i.note=='Native divider artwork' for i in group) and _is_native_divider_background(element,_text(element)):
            headings.add(position)
        if any(i.note=='Native section heading' or (i.kind=='text' and not i.suggested_slot and _heading(i.text)) for i in group):
            headings.add(position)
    for heading_position in sorted(headings):
        for position in range(heading_position-1,0,-1):
            element=original[position-1]
            if element.find('.//'+qn('w:sectPr')) is not None:break
            if _is_native_divider_background(element,_text(element)):
                headings.add(position)
                for previous in range(position,heading_position):section_by_pos[previous]=section_by_pos[heading_position]
                break
    return section_by_pos,headings


def test_shared_selector_is_exact_existing_algorithm(sample):
    _,view=sample
    with _package(view.rendition_bytes) as archive:units,*_=_routing(archive)
    plan=source_section_plan(units,view.rendition_inspection)
    mapping,headings=_legacy_selector(units,view.rendition_inspection)
    assert dict(plan.section_by_position)==mapping
    assert set(plan.headings)==headings
    assert plan.selected_positions('cover')==tuple(p for p in plan.nonterminal_positions
        if mapping[p]=='cover' and p<=plan.first_word_end)


def test_source_selected_scope_has_inherited_dormant_variants_and_excludes_unrelated(sample):
    _,view=sample
    complete=expected_header_scope(view)
    cover=expected_header_scope(view,'cover')
    activity=expected_header_scope(view,'activity')
    training=expected_header_scope(view,'training')
    assert cover.source_word_sections==(1,)
    assert activity.source_word_sections==(2,)
    assert training.source_word_sections==(3,)
    assert complete.source_word_sections==(1,2,3)
    assert activity.routes[0].parts[0] != cover.routes[0].parts[0]
    assert activity.routes[0].parts[1:] == cover.routes[0].parts[1:]
    assert cover.routes[0].title_page and not activity.routes[0].title_page
    assert all(scope.even_and_odd for scope in (complete,cover,activity,training))
    assert training.routes[0].parts[0] not in activity.expected_parts
    assert expected_header_scope(view,'activity')==activity
    verify_header_routing(view.rendition_bytes,complete)
    for scope in (cover,activity,training): verify_header_routing(_preview(view,scope),scope)


@pytest.mark.parametrize('case',['swap','remove','extra','type','title','even','boundary','external','wrong_kind','missing','traversal','duplicate_ref','duplicate_rid','root'])
def test_output_routing_mutations_rejected_even_when_parts_remain(sample,case):
    _,view=sample; scope=expected_header_scope(view,'activity'); raw=_preview(view,scope)
    def mutate(parts):
        root=ET.fromstring(parts['word/document.xml']);props=root.find('.//'+W+'sectPr')
        refs=props.findall(W+'headerReference')
        relations=ET.fromstring(parts['word/_rels/document.xml.rels'])
        if case=='swap':
            refs[0].set(R+'id',refs[1].get(R+'id'));refs[1].set(R+'id',scope.routes[0].parts[0])
        elif case=='remove':props.remove(refs[0])
        elif case=='extra':
            relation=next(r for r in relations if r.get('Type','').endswith('/header') and 'word/'+r.get('Target') not in scope.expected_parts)
            refs[0].set(R+'id',relation.get('Id'))
        elif case=='type':refs[0].set(W+'type','mystery')
        elif case=='title':props.append(OxmlElement('w:titlePg'))
        elif case=='even':xml(parts,'word/settings.xml',lambda r:r.remove(r.find(W+'evenAndOddHeaders')))
        elif case=='boundary':root.find(W+'body').insert(0,props)
        elif case=='duplicate_ref':props.append(deepcopy(refs[0]))
        else:
            relation=next(r for r in relations if r.get('Id')==refs[0].get(R+'id'))
            if case=='external':relation.set('TargetMode','External')
            elif case=='wrong_kind':relation.set('Type','https://invalid.example/header')
            elif case=='missing':relations.remove(relation)
            elif case=='traversal':relation.set('Target','../word/'+relation.get('Target'))
            elif case=='duplicate_rid':relations.append(deepcopy(relation))
            elif case=='root':parts['word/'+relation.get('Target')]=b'<bad/>'
        parts['word/document.xml']=ET.tostring(root)
        parts['word/_rels/document.xml.rels']=ET.tostring(relations)
    with pytest.raises(CompatibilityError):verify_header_routing(rewrite(raw,mutate),scope)


def test_swap_valid_ids_with_same_part_set_rejected(sample):
    _,view=sample;scope=expected_header_scope(view,'activity');raw=_preview(view,scope)
    def mutate(parts):
        def change(root):
            refs=root.find('.//'+W+'sectPr').findall(W+'headerReference')
            left,right=refs[0].get(R+'id'),refs[1].get(R+'id')
            refs[0].set(R+'id',right);refs[1].set(R+'id',left)
        xml(parts,'word/document.xml',change)
    with pytest.raises(CompatibilityError,match='routing|boundary'):verify_header_routing(rewrite(raw,mutate),scope)


def test_source_missing_section_and_wrong_typed_view_fail(sample):
    _,view=sample
    with pytest.raises(CompatibilityError):expected_header_scope(view,'does-not-exist')
    with pytest.raises(CompatibilityError):expected_header_scope(None,'cover')


@pytest.mark.parametrize('case',['page_width','header_distance','page_number_restart','break_type','extra_identical_section'])
def test_page_context_or_identical_route_section_count_cannot_change(sample,case):
    _,view=sample;scope=expected_header_scope(view,'activity');raw=_preview(view,scope)
    def mutate(parts):
        def change(root):
            body=root.find(W+'body');props=body.find(W+'sectPr')
            if case=='page_width':props.find(W+'pgSz').set(W+'w','15000')
            elif case=='header_distance':props.find(W+'pgMar').set(W+'header','999')
            elif case=='page_number_restart':
                node=OxmlElement('w:pgNumType');node.set(W+'start','5');props.append(node)
            elif case=='break_type':
                node=props.find(W+'type')
                if node is None:node=OxmlElement('w:type');props.append(node)
                node.set(W+'val','continuous')
            else:
                paragraph=OxmlElement('w:p');properties=OxmlElement('w:pPr')
                properties.append(deepcopy(props));paragraph.append(properties)
                body.insert(0,paragraph)
        xml(parts,'word/document.xml',change)
    with pytest.raises(CompatibilityError,match='routing|boundary'):verify_header_routing(rewrite(raw,mutate),scope)


def test_real_native_full_and_section_outputs_match_input_scope(sample,tmp_path,monkeypatch):
    from dataclasses import replace
    from app.monthly_report_model import synthetic_profiles, Facility, ReportPeriod
    from app.monthly_report_native_layout import build_native_docx
    from test_monthly_report_native_art_lifecycle import save_sample
    monkeypatch.setenv('EPC_DATA_DIR',str(tmp_path/'store'))
    path,view=sample
    profile=replace(synthetic_profiles()[0], imported_from=view.original_inspection.sha256,
        title='Synthetic Regional Community Health Hospital',
        facilities=(Facility('synthetic','Synthetic Regional Community Health Hospital'),))
    draft,loader=save_sample((path,view,profile,ReportPeriod(2026,9),None,None))
    for section_key in (None,'cover','activity','training'):
        scope=expected_header_scope(view,section_key)
        output=build_native_docx(draft,path,asset_loader=loader,source_view=view,section_key=section_key)
        verify_header_routing(output,scope)


def test_moved_nonterminal_boundary_same_routes_and_parts_is_rejected(sample):
    _,view=sample;scope=expected_header_scope(view)
    def mutate(parts):
        def change(root):
            body=root.find(W+'body')
            original=next(p for p in body if p.tag==W+'p' and p.find('.//'+W+'sectPr') is not None)
            target=next(p for p in body if p.tag==W+'p' and p is not original and p.find('.//'+W+'sectPr') is None)
            properties=original.find('.//'+W+'sectPr');properties.getparent().remove(properties)
            ppr=target.find(W+'pPr')
            if ppr is None:ppr=OxmlElement('w:pPr');target.insert(0,ppr)
            ppr.append(properties)
        xml(parts,'word/document.xml',change)
    with pytest.raises(CompatibilityError,match='boundary'):verify_header_routing(rewrite(view.rendition_bytes,mutate),scope)


def test_second_body_cannot_hide_unexamined_routing(sample):
    _,view=sample;scope=expected_header_scope(view)
    def mutate(parts):
        xml(parts,'word/document.xml',lambda root:root.append(deepcopy(root.find(W+'body'))))
    with pytest.raises(CompatibilityError,match='scaffold'):verify_header_routing(rewrite(view.rendition_bytes,mutate),scope)
