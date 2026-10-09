"""Audit F15: a current table's headings travel with its data, not the chart."""
from dataclasses import replace
from io import BytesIO

import pymupdf
import pytest
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from PIL import Image, ImageDraw

from conftest import requires_libreoffice
from app import pdf_converter
from app.monthly_report_import import ImportMapping, inspect_docx, map_items
from app.monthly_report_model import (
    BlockSpec, ColumnSpec, Facility, ReportDraft, ReportPeriod, ReportProfile,
    ResolvedBlock, SectionSpec,
)
from app.monthly_report_native_layout import _table, build_native_docx

COLUMNS = ('Contact name', 'Contact role')


def source_and_draft(tmp_path, count=4, image_size=(1396, 1536)):
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Inches(11), Inches(8.5)
    section.left_margin = section.right_margin = Inches(.5)
    section.top_margin = section.bottom_margin = Inches(.5)
    document.add_heading('ORGANIZATIONAL CHART', 1)
    document.add_section(WD_SECTION_START.NEW_PAGE)
    image = Image.new('RGB', image_size, 'white')
    draw = ImageDraw.Draw(image)
    draw.rectangle((1, 1, image_size[0]-2, image_size[1]-2), outline='black', width=12)
    draw.line((10, image_size[1]//2, image_size[0]-10, image_size[1]//2), fill='black', width=10)
    data = BytesIO(); image.save(data, format='PNG')
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(0)
    if image_size[0] > image_size[1]:
        paragraph.add_run().add_picture(BytesIO(data.getvalue()), width=Inches(9))
    else:
        paragraph.add_run().add_picture(BytesIO(data.getvalue()), height=Inches(6.85))
    table = document.add_table(rows=2, cols=2)
    table.style = 'Table Grid'
    for index, values in enumerate((COLUMNS, ('Previous person', 'Previous duty'))):
        for cell, text in zip(table.rows[index].cells, values):
            cell.text = text
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = p.paragraph_format.space_before = Pt(0)
            p.runs[0].font.size = Pt(10)
        table.rows[index].height = Pt(18 if index == 0 else 43)
    source = tmp_path / 'source.docx'; document.save(source)
    inspection = inspect_docx(source)
    picture = next(item for item in inspection.items if item.kind == 'image')
    mapped = map_items(source, inspection, (ImportMapping(picture.id, 'org_chart'),))
    profile = ReportProfile('Synthetic Contract', 'north', 'Synthetic North', (Facility('north', 'Synthetic North'),))
    sections = (SectionSpec('organization', '1', 'Organizational Chart', (
        BlockSpec('org_chart', 'image_page'),
        BlockSpec('contact_matrix', 'table', columns=tuple(ColumnSpec(str(n), title) for n, title in enumerate(COLUMNS))),
    )),)
    rows = tuple((f'Synthetic Person {n:02}', f'Facilities duty {n:02}') for n in range(count))
    current = ReportDraft(profile, ReportPeriod(2026, 9), 'Synthetic Editor', sections,
        (*mapped.blocks, ResolvedBlock('contact_matrix', 'This month', rows=rows)))
    return source, current, dict(mapped.assets), inspection


def render(tmp_path, monkeypatch, count, image_size=(1396, 1536)):
    monkeypatch.setenv('EPC_DATA_DIR', str(tmp_path / 'runtime'))
    monkeypatch.setattr(pdf_converter, 'OUTPUT_DIR', tmp_path)
    monkeypatch.setattr(pdf_converter, 'PDF_BACKEND', 'libreoffice')
    source, draft, assets, _ = source_and_draft(tmp_path, count, image_size)
    original = source.read_bytes()
    raw = build_native_docx(draft, source, section_key='organization', asset_loader=assets.__getitem__)
    path = tmp_path / 'current.docx'; path.write_bytes(raw)
    output = pdf_converter.convert_to_pdf(path)
    try:
        pdf = output.read_bytes()
    finally:
        output.unlink(missing_ok=True)
    assert source.read_bytes() == original
    return raw, pdf


@pytest.mark.parametrize('count', [1, 4, 40])
def test_rebuilt_table_marks_only_current_header_and_preserves_geometry(tmp_path, count):
    source, draft, assets, _ = source_and_draft(tmp_path, count)
    original = Document(source).tables[0]
    table = Document(BytesIO(build_native_docx(draft, source, section_key='organization', asset_loader=assets.__getitem__))).tables[0]
    assert table._tbl.tblGrid.xml == original._tbl.tblGrid.xml
    assert table.style.style_id == original.style.style_id
    assert len(table.rows) == count + 1
    assert table.rows[0]._tr.trPr.find(qn('w:tblHeader')) is not None
    assert table.rows[0]._tr.trPr.find(qn('w:cantSplit')) is not None
    assert all(p.paragraph_format.keep_with_next is True for c in table.rows[0].cells for p in c.paragraphs)
    assert all(row._tr.trPr.find(qn('w:tblHeader')) is None for row in table.rows[1:])
    assert all(p.paragraph_format.keep_with_next is False for row in table.rows[1:] for c in row.cells for p in c.paragraphs)
    assert table.rows[1]._tr.trPr.find(qn('w:cantSplit')) is not None
    assert table.rows[0].height == Pt(18)
    assert table.rows[1].height == Pt(43)


@requires_libreoffice
@pytest.mark.parametrize('count', [1, 4, 40])
def test_header_never_stranded_and_repeats_above_current_rows(tmp_path, monkeypatch, count):
    raw, pdf = render(tmp_path, monkeypatch, count)
    with pymupdf.open(stream=pdf, filetype='pdf') as document:
        all_text = ''.join(page.get_text() for page in document)
        row_pages = [p for p in document if 'Synthetic Person' in p.get_text()]
        assert row_pages
        for page in document:
            text = page.get_text()
            assert ('Contact name' in text) == ('Synthetic Person' in text)
            if 'Synthetic Person' in text:
                assert 'Contact role' in text
            assert text.strip() or page.get_image_info(), 'no blank page introduced'
        for index in range(count):
            assert all_text.count(f'Synthetic Person {index:02}') == 1
            assert all_text.count(f'Facilities duty {index:02}') == 1
        assert 'Previous person' not in all_text
        images = [image for page in document for image in page.get_image_info()]
        assert len(images) == 1
        box = pymupdf.Rect(images[0]['bbox'])
        assert box.height == pytest.approx(6.85*72, abs=.5)
        assert box.width / box.height == pytest.approx(1396/1536, abs=.001)
        if count < 10:
            assert len(document) == 3
        else:
            assert len(document) == 6
            assert row_pages[0].get_text().count('Synthetic Person') >= 9
            assert len(row_pages) >= 3
    table = Document(BytesIO(raw)).tables[0]
    assert tuple(tuple(c.text for c in row.cells) for row in table.rows[1:]) == tuple((f'Synthetic Person {n:02}', f'Facilities duty {n:02}') for n in range(count))


@requires_libreoffice
def test_wide_chart_and_one_row_can_share_the_page(tmp_path, monkeypatch):
    _, pdf = render(tmp_path, monkeypatch, 1, (1536, 862))
    with pymupdf.open(stream=pdf, filetype='pdf') as document:
        assert len(document) == 2
        page = document[-1]
        assert 'Contact name' in page.get_text() and 'Synthetic Person 00' in page.get_text()
        image = page.get_image_info()[0]
        box = pymupdf.Rect(image['bbox'])
        assert box.width == pytest.approx(9*72, abs=.5)
        assert box.width / box.height == pytest.approx(1536/862, abs=.001)
        words = page.get_text('words')
        assert min(w[1] for w in words if w[4] == 'Contact') > box.y1


def test_empty_contacts_do_not_add_header_or_claims(tmp_path):
    source, draft, assets, _ = source_and_draft(tmp_path, 0)
    output = Document(BytesIO(build_native_docx(draft, source, section_key='organization', asset_loader=assets.__getitem__)))
    assert not output.tables
    assert len(output.inline_shapes) == 1


def test_unchanged_imported_table_is_not_repaginated(tmp_path):
    source, draft, assets, inspection = source_and_draft(tmp_path)
    item = next(i for i in inspection.items if i.kind == 'table')
    mapped = map_items(source, inspection, (ImportMapping(item.id, 'contact_matrix'),))
    current = replace(draft, blocks=(draft.blocks[0], *mapped.blocks))
    before = Document(source).tables[0]._tbl.xml
    after = Document(BytesIO(build_native_docx(current, source, section_key='organization', asset_loader=assets.__getitem__))).tables[0]._tbl.xml
    assert after == before


def test_header_only_prototype_does_not_make_current_data_header_rows():
    doc = Document(); table = doc.add_table(rows=1, cols=2)
    props = table.rows[0]._tr.get_or_add_trPr()
    props.append(OxmlElement('w:tblHeader'))
    for cell, name in zip(table.rows[0].cells, COLUMNS):
        cell.text = name
        cell.paragraphs[0].paragraph_format.keep_with_next = True
    _table(table._tbl, COLUMNS, (('Synthetic Person', 'Current duty'),))
    assert table.rows[1]._tr.trPr.find(qn('w:tblHeader')) is None
    assert all(p.paragraph_format.keep_with_next is False for c in table.rows[1].cells for p in c.paragraphs)


def test_row_properties_follow_schema_order():
    doc = Document(); table = doc.add_table(rows=2, cols=2)
    table.rows[0].height = Pt(18)
    props = table.rows[0]._tr.get_or_add_trPr()
    props.append(OxmlElement('w:jc'))
    _table(table._tbl, COLUMNS, (('Synthetic Person', 'Current duty'),))
    names = [n.tag.rsplit('}',1)[-1] for n in table.rows[0]._tr.trPr]
    assert names == ['cantSplit', 'trHeight', 'tblHeader', 'jc']
    assert table.rows[0]._tr[0].tag == qn('w:trPr')
    assert all(p._p[0].tag == qn('w:pPr') for c in table.rows[0].cells for p in c.paragraphs)


def test_inferred_source_header_is_not_promoted_to_repeating_headings():
    doc = Document(); table = doc.add_table(rows=2, cols=2)
    for row in table.rows:
        for cell in row.cells:
            cell.text = 'Previous technical value'
    _table(table._tbl, (), (('Current value', 'Current units'),))
    assert not list(table._tbl.iter(qn('w:tblHeader')))
    assert tuple(c.text for c in table.rows[1].cells) == ('Current value', 'Current units')


def test_multiline_header_keeps_all_its_paragraphs_with_the_first_row():
    doc = Document(); table = doc.add_table(rows=2, cols=2)
    _table(table._tbl, ('Contact name\nHospital team', 'Contact role'), (('Synthetic Person', 'Current duty'),))
    assert table.rows[0].cells[0].text == 'Contact name\nHospital team'
    assert all(p.paragraph_format.keep_with_next and p.paragraph_format.keep_together
               for c in table.rows[0].cells for p in c.paragraphs)
    assert table.rows[1].cells[0].text == 'Synthetic Person'


@requires_libreoffice
def test_empty_current_contacts_leave_only_the_chart_not_a_header_page(tmp_path, monkeypatch):
    raw, pdf = render(tmp_path, monkeypatch, 0)
    assert not Document(BytesIO(raw)).tables
    with pymupdf.open(stream=pdf, filetype='pdf') as document:
        assert len(document) == 2
        assert 'Contact name' not in ''.join(p.get_text() for p in document)
        assert len(document[-1].get_image_info()) == 1


def test_current_additional_contact_table_is_not_treated_as_empty(tmp_path):
    from app.monthly_report_model import ReportTable
    source, draft, assets, _ = source_and_draft(tmp_path, 0)
    current = replace(draft.blocks[-1], extra_tables=(ReportTable(COLUMNS, (('Synthetic Person', 'Current duty'),)),))
    draft = replace(draft, blocks=(*draft.blocks[:-1], current))
    output = Document(BytesIO(build_native_docx(draft, source, section_key='organization', asset_loader=assets.__getitem__)))
    assert len(output.tables) == 1
    assert tuple(c.text for c in output.tables[0].rows[1].cells) == ('Synthetic Person', 'Current duty')
    assert output.tables[0].rows[0]._tr.trPr.find(qn('w:tblHeader')) is not None
