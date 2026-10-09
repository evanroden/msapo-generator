from dataclasses import asdict, replace
from datetime import date

import pytest

from app import monthly_report_library as library
from app.monthly_report_directory import DirectoryContact, DirectorySite, DirectoryState
from app.monthly_report_model import Facility, ReportPeriod, ReportProfile, ReportTable, ResolvedBlock
from app.monthly_report_training import (EVENT_REF, MATRIX_REF, add_event, apply_tables,
    load_training, save_training, seed_matrix, validate_matrix)


@pytest.fixture
def profile(tmp_path, monkeypatch):
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path))
    return ReportProfile("Training Test", "both", "Both sites", (Facility("a", "Site A"), Facility("b", "Site B")), "multi_site")


def matrix():
    return ReportTable(("Team member", "Site", "Safety", "Operations"),
                       (("Sam", "Site A", "Completed", "Pending"), ("Jo", "Site B", "Not required", "Not recorded")), MATRIX_REF)


def test_automatic_standing_storage_is_scoped_versioned_and_conflict_safe(profile):
    saved = save_training(profile, matrix(), expected_revision=0, actor="Editor")
    assert saved["revision"] == 1
    assert seed_matrix(profile, ResolvedBlock("training_summary", "This month"), load_training(profile.contract)) == matrix()
    single = replace(profile, facilities=profile.facilities[:1], scope_type="individual")
    only = replace(matrix(), rows=(("Sam", "Site A", "Pending", "Pending"),))
    save_training(single, only, expected_revision=1, actor="Editor")
    loaded = load_training(profile.contract)
    assert loaded["sites"]["b"] == saved["sites"]["b"]
    with pytest.raises(library.RevisionConflict):
        save_training(profile, matrix(), expected_revision=1, actor="Editor")
    other = replace(profile, contract="Different Contract")
    assert not seed_matrix(other, ResolvedBlock("training_summary", "This month"), load_training(other.contract)).rows
    assert len(list((library._root() / "training").rglob("history/*.json"))) == 2


def test_only_known_hospital_people_selected_sites_no_invented_completion(profile):
    directory = DirectoryState(profile.contract, 1, (
        DirectorySite("a", "Site A", contacts=(DirectoryContact("ENFRA Technician", "Sam"), DirectoryContact("Client contact", "Client"))),
        DirectorySite("other", "Other", contacts=(DirectoryContact("ENFRA operator", "Elsewhere"),))), "Editor", "today")
    result = seed_matrix(profile, ResolvedBlock("training_summary", "This month"), {"sites": {"a": {"columns": ["Safety"]}}}, directory)
    assert result.rows == (("Client", "Site A", "Not recorded"),)


def test_wizard_modes_hours_month_and_no_duplicate_event():
    period = ReportPeriod(2026, 9)
    args = dict(title="Safety", when=date(2026, 9, 3), mode="Virtual", participants=["Sam"], hours="", period=period)
    table = add_event(None, **args)
    assert table.rows[0][-1] == ""
    assert add_event(table, **args) == table
    for mode in ("In person", "Asynchronous"):
        assert add_event(None, **(args | {"mode": mode, "hours": "1.5"})).rows[0][2] == mode
    for value in ("nan", "-1", "unknown", "inf", "0"):
        with pytest.raises(ValueError):
            add_event(None, **(args | {"hours": value}))
    with pytest.raises(ValueError, match="reporting month"):
        add_event(None, **(args | {"when": date(2026, 8, 31)}))


def test_report_mapping_preserves_imported_design_content_and_roundtrips():
    source_table = ReportTable(("Topic",), (("Source training",),), "source:1")
    block = ResolvedBlock("training_summary", "This month", text="Original narrative", extra_tables=(source_table,), asset_hashes=("photo",))
    event = add_event(None, title="Safety", when=date(2026, 9, 3), mode="In person", participants=["Sam"], hours="", period=ReportPeriod(2026, 9))
    updated = apply_tables(block, matrix(), event)
    assert updated.text == block.text and updated.asset_hashes == block.asset_hashes
    assert updated.extra_tables == (source_table, matrix(), event)
    assert library.block_from_dict(asdict(updated)) == updated
    assert seed_matrix(None, updated, {}) == matrix()
    assert apply_tables(updated, matrix(), event) == updated
    assert updated.extra_tables[-1].reference == EVENT_REF


def test_invalid_status_duplicate_people_and_foreign_site_rejected(profile):
    with pytest.raises(ValueError):
        validate_matrix(replace(matrix(), rows=(("Sam", "Site A", "Probably", "Pending"),)))
    with pytest.raises(ValueError):
        validate_matrix(replace(matrix(), rows=matrix().rows * 2))
    with pytest.raises(ValueError, match="selected sites"):
        save_training(profile, replace(matrix(), rows=(("Sam", "Elsewhere", "Completed", "Pending"),)), expected_revision=0, actor="Editor")


def test_ui_first_render_no_write_wizard_and_resume(profile):
    from streamlit.testing.v1 import AppTest
    from app.monthly_report_model import ReportDraft, default_sections
    app = AppTest.from_string('''
import streamlit as st
from app.monthly_report_training_ui import render_training
def field(key, value):
    if key not in st.session_state:
        st.session_state[key] = value
    return key
st.session_state['blocks'] = render_training(st.session_state['draft'], st.session_state['blocks'], 'test', field)
''', default_timeout=30)
    app.session_state['draft'] = ReportDraft(profile, ReportPeriod(2026, 9), 'Editor', default_sections(), ())
    app.session_state['blocks'] = {}
    app.run()
    assert not app.exception
    assert load_training(profile.contract)['revision'] == 0
    next(w for w in app.text_input if w.label == 'Add a training requirement').set_value('Safety')
    next(w for w in app.button if w.label == 'Add training column').click().run()
    assert not app.exception
    assert load_training(profile.contract)['revision'] == 1
    app.run()
    assert load_training(profile.contract)['revision'] == 1
    next(w for w in app.text_input if w.label == 'Training topic').set_value('Site safety')
    next(w for w in app.multiselect if w.label == 'Participants').set_value(['Sam'])
    next(w for w in app.button if w.label == 'Add completed training').click().run()
    assert not app.exception
    result = app.session_state['blocks']['training_summary']
    assert result.extra_tables[-1].reference == EVENT_REF
    assert result.extra_tables[-1].rows[0][0] == 'Site safety'
    app.run()
    assert app.session_state['blocks']['training_summary'] == result


def test_monthly_events_do_not_roll_forward_with_standing_status():
    from app.monthly_report_training import events_for_period
    event = add_event(None, title="Safety", when=date(2026, 9, 3), mode="Virtual", participants=["Sam"], hours="", period=ReportPeriod(2026, 9))
    block = apply_tables(ResolvedBlock("training_summary", "This month"), matrix(), event)
    carried = events_for_period(block, ReportPeriod(2026, 10))
    assert carried.rows == ()
    october = apply_tables(block, matrix(), carried)
    assert october.extra_tables == (matrix(),)


def test_ui_conflict_reload_preserves_monthly_notes_and_events(profile, monkeypatch):
    from streamlit.testing.v1 import AppTest
    from app.monthly_report_model import ReportDraft, default_sections
    from app import monthly_report_training_ui as ui
    event = add_event(None, title="Safety", when=date(2026, 9, 3), mode="Virtual", participants=["Sam"], hours="", period=ReportPeriod(2026, 9))
    block = apply_tables(ResolvedBlock("training_summary", "This month", text="Monthly notes"), matrix(), event)
    save_training(profile, matrix(), expected_revision=0, actor="Editor")
    app = AppTest.from_string('''
import streamlit as st
from app.monthly_report_training_ui import render_training
def field(key, value):
    if key not in st.session_state:
        st.session_state[key] = value
    return key
st.session_state['blocks'] = render_training(st.session_state['draft'], st.session_state['blocks'], 'test', field)
''', default_timeout=30)
    app.session_state['draft'] = ReportDraft(profile, ReportPeriod(2026, 9), 'Editor', default_sections(), ())
    app.session_state['blocks'] = {'training_summary': block}
    app.run()
    newer = replace(matrix(), rows=(("Sam", "Site A", "Pending", "Completed"), matrix().rows[1]))
    save_training(profile, newer, expected_revision=1, actor="Other Editor")
    original_grid = ui._grid
    monkeypatch.setattr(ui, '_grid', lambda key, rows, **kwargs: [dict(row, Safety="Not required") for row in rows])
    app.run()
    assert app.error and not app.exception
    assert app.session_state['blocks']['training_summary'] == block
    monkeypatch.setattr(ui, '_grid', original_grid)
    next(w for w in app.button if w.label == 'Reload saved training matrix').click().run()
    assert not app.exception and not app.error
    refreshed = app.session_state['blocks']['training_summary']
    assert refreshed.text == 'Monthly notes'
    assert event in refreshed.extra_tables
    assert newer in refreshed.extra_tables
    app.run()
    assert not app.exception and not app.error
    assert load_training(profile.contract)['revision'] == 2


def test_unnamed_visitor_never_reads_shared_training_or_directory(profile, monkeypatch):
    from streamlit.testing.v1 import AppTest
    from app.monthly_report_model import ReportDraft, default_sections
    from app import monthly_report_training_ui as ui
    from app import monthly_report_directory as directory
    def forbidden(*args, **kwargs):
        raise AssertionError('Anonymous shared-data access')
    monkeypatch.setattr(ui, 'load_training', forbidden)
    monkeypatch.setattr(directory, 'load_directory', forbidden)
    app = AppTest.from_string('''
import streamlit as st
from app.monthly_report_training_ui import render_training
render_training(st.session_state['draft'], {}, 'anonymous', lambda k, v: k)
''', default_timeout=30)
    app.session_state['draft'] = ReportDraft(profile, ReportPeriod(2026, 9), '', default_sections(), ())
    app.run()
    assert not app.exception
    assert app.info and not app.dataframe


def test_training_directory_alias_binding_ignores_ambiguous_and_inactive(profile):
    single = replace(profile, facilities=(Facility('legacy-a', 'Site A', ('North',)),), scope_type='individual')
    contact = DirectoryContact('Hospital Technician', 'Sam')
    directory = DirectoryState(profile.contract, 1, (DirectorySite('directory-a', 'North', contacts=(contact,)),), 'Editor', 'today')
    block = ResolvedBlock('training_summary', 'This month')
    assert seed_matrix(single, block, {}, directory).rows == (('Sam', 'Site A'),)
    ambiguous = replace(directory, sites=(*directory.sites, DirectorySite('other', 'Site A', contacts=(contact,))))
    assert not seed_matrix(single, block, {}, ambiguous).rows
    inactive = replace(directory, sites=(replace(directory.sites[0], active=False),))
    assert not seed_matrix(single, block, {}, inactive).rows


def test_native_training_style_preserves_geometry_font_header_and_colors_status():
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from app.monthly_report_training_style import style_training_table
    document = Document()
    table = document.add_table(rows=2, cols=3)
    for cells, values in zip(table.rows, [('Team member', 'Site', 'Safety'), ('Sam', 'Site A', 'Completed')]):
        for cell, value in zip(cells.cells, values):
            cell.text = value
            cell.paragraphs[0].runs[0].font.name = 'Arial'
    height = OxmlElement('w:trHeight')
    height.set(qn('w:val'), '9000')
    table.rows[1]._tr.get_or_add_trPr().append(height)
    grid = table._tbl.find(qn('w:tblGrid')).xml
    header = table.cell(0, 2)._tc.xml
    style_training_table(table._tbl, MATRIX_REF)
    assert table._tbl.find(qn('w:tblGrid')).xml == grid
    assert not list(table._tbl.iter(qn('w:trHeight')))
    assert table.cell(0, 2).text == 'Safety'
    assert table.cell(0, 2).paragraphs[0].runs[0].font.name == 'Arial'
    assert table.cell(1, 2)._tc.find('./' + qn('w:tcPr') + '/' + qn('w:shd')).get(qn('w:fill')) == 'D8F3DC'
    assert 'Safety' in header
    unmodified = table._tbl.xml
    style_training_table(table._tbl, 'source:original')
    assert table._tbl.xml == unmodified
