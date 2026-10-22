"""Source-view import and finite page replacement share one atomic publication."""
from dataclasses import replace
from io import BytesIO

from docx import Document
from docx.shared import Inches
from streamlit.testing.v1 import AppTest

from app import monthly_report_library as library, monthly_report_section_ui as ui
from test_monthly_report_native_art_review import sample, setup, rewrite
from test_monthly_report_native_art_lifecycle import save_sample
from test_monthly_report_source_frame_logo_role import raster


def test_different_view_import_never_silently_replaces_saved_design(sample, monkeypatch, tmp_path):
    path = sample[0]
    doc = Document(path)
    doc.add_heading('1 Organizational Chart', 1)
    doc.add_picture(BytesIO(raster('navy')), width=Inches(2))
    image_part = str(doc.part.related_parts[doc.inline_shapes[-1]._inline.graphic.graphicData.pic.blipFill.blip.embed].partname).lstrip('/')
    doc.save(path)
    initial_sample = setup(path)
    initial, _ = save_sample(initial_sample)
    library.save_snapshot(initial, expected_revision=0, entered_editor='Synthetic Operator')
    incoming = tmp_path/'incoming-view.docx'
    incoming.write_bytes(rewrite(path.read_bytes(), lambda parts: parts.__setitem__(image_part, raster('red'))))
    incoming_sample = setup(incoming)
    assert incoming_sample[1].original_inspection.sha256 != initial_sample[1].original_inspection.sha256
    raw = incoming.read_bytes()
    upload = BytesIO(raw); upload.name = 'synthetic.docx'; upload.size = len(raw)
    real = ui.st.file_uploader
    monkeypatch.setattr(ui.st, 'file_uploader', lambda label,*a,**k: upload if label=='Starting report' else real(label,*a,**k))
    app = AppTest.from_string(f'''
from pathlib import Path
from test_monthly_report_native_art_review import setup
from app import monthly_report_library as library
from app.monthly_report_section_ui import render_section_setup
from app.monthly_report_ui import _field
path, view, profile, period, inventory, receipt = setup(Path({str(incoming)!r}))
state = library.load_profile(profile.contract, profile.key)
render_section_setup(profile.contract, period, 'Synthetic Operator', _field, state=state, identity=state.profile)
''', default_timeout=90).run()
    assert not app.exception and not app.error
    assert any('different starting-report file' in item.value for item in app.info)
    assert not any(b.label == 'Continue to this month’s updates' for b in app.button)
    assert library.load_imported_draft(initial.profile.contract, initial.profile.key) == initial
    assert library.imported_original(initial.profile.contract, initial.profile.key).read_bytes() == path.read_bytes()
    assert library.load_snapshot(initial.profile.contract, initial.profile.key, initial.period).draft == initial
    assert not library.source_view_replay_recorded(initial.profile.contract, initial.profile.key,
                                                  initial.period, incoming_sample[1].original_inspection.sha256)


def test_same_view_preflight_preserves_report_only_page_and_replays(sample, monkeypatch):
    path = sample[0]
    doc = Document(path)
    doc.add_heading('1 Organizational Chart', 1)
    doc.add_picture(BytesIO(raster('navy')), width=Inches(2))
    doc.save(path)
    initial_sample = setup(path)
    initial, _ = save_sample(initial_sample)
    chart = next(b for b in initial.blocks if b.key == 'org_chart')
    other = next(b for b in initial.blocks if b.key == 'client_logo').asset_hashes[0]
    edited_chart = replace(chart, asset_hashes=(other,), references=('human:report-only',),
                           asset_provenance=((other, ('human:report-only',)),))
    edited = replace(initial, blocks=tuple(edited_chart if b.key == 'org_chart' else b for b in initial.blocks))
    library.save_snapshot(edited, expected_revision=0, entered_editor='Synthetic Operator')
    raw = path.read_bytes()
    upload = BytesIO(raw); upload.name='synthetic.docx'; upload.size=len(raw)
    real=ui.st.file_uploader
    monkeypatch.setattr(ui.st,'file_uploader',lambda label,*a,**k:upload if label=='Starting report' else real(label,*a,**k))
    script=f'''
from pathlib import Path
from test_monthly_report_native_art_review import setup
from app import monthly_report_library as library
from app.monthly_report_section_ui import render_section_setup
from app.monthly_report_ui import _field
path, view, profile, period, inventory, receipt = setup(Path({str(path)!r}))
state = library.load_profile(profile.contract, profile.key)
render_section_setup(profile.contract, period, 'Synthetic Operator', _field, state=state, identity=state.profile)
'''
    for _ in range(2):
        app=AppTest.from_string(script,default_timeout=90).run()
        assert not app.exception and not app.error, [e.value for e in app.error]
        assert not any(b.label.startswith('Use incoming') for b in app.button)
        button=next(b for b in app.button if b.label=='Continue to this month’s updates')
        assert not button.disabled
        button.click().run()
        assert not app.exception and not app.error, [e.value for e in app.error]
        loaded=library.load_imported_draft(initial.profile.contract, initial.profile.key)
        assert next(b for b in loaded.blocks if b.key=='org_chart') == edited_chart
        assert library.load_snapshot(initial.profile.contract, initial.profile.key, initial.period).draft == edited
        assert library.imported_native_art_review(initial.profile.contract, initial.profile.key, initial.period, source_view=initial_sample[1])
        assert library.imported_original(initial.profile.contract, initial.profile.key).read_bytes() == raw
