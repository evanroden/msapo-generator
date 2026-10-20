"""Word capacity cells stay whole when a cell holds several paragraphs (Refs #286, G9-pr89-1).

A units line under a heading, or a second plant line, used to become its own
table row and shift every later value. Synthetic values only.
"""
from io import BytesIO

from docx import Document
import pytest

from app import monthly_report_capacity as capacity, monthly_report_directory as directory
from app.monthly_report_model import Facility, ReportProfile

FACILITIES = (Facility("north", "North Hospital", ("NH",)), Facility("south", "South Hospital", ("SH",)))
PROFILE = ReportProfile("Synthetic Capacity Contract", "north", "North Hospital", FACILITIES[:1])


@pytest.fixture(autouse=True)
def isolated(tmp_path, monkeypatch):
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path))
    directory.save_directory(PROFILE.contract, tuple(directory.DirectorySite(f.key, f.title, f.aliases) for f in FACILITIES),
                             expected_revision=0, actor="Synthetic Reviewer", confirmed=True)


def _table(document, rows):
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            first, *rest = value.split("\n")
            cell = table.cell(r, c)
            cell.text = first
            for line in rest:
                cell.add_paragraph(line)


def _inspect(document):
    buffer = BytesIO()
    document.save(buffer)
    return capacity.inspect_capacity(PROFILE, "capacity.docx", buffer.getvalue())


def test_multi_paragraph_header_and_data_cells_stay_in_their_columns():
    document = Document()
    document.add_heading("Thermal capacity", 1)
    _table(document, [["Facility", "Service", "Required capacity\n(lb/hr or tons)", "Available capacity\n(lb/hr or tons)"],
                      ["NH", "Steam", "40000", "50000"],
                      ["SH", "Chilled water\n(Plant B)", "1200", "1600"]])
    tables = _inspect(document).tables
    assert len(tables) == 1
    table = tables[0]
    assert table.columns == ("Facility", "Service", "Required capacity (lb/hr or tons)",
                             "Available capacity (lb/hr or tons)")
    assert table.rows == (("NH", "Steam", "40000", "50000"), ("SH", "Chilled water (Plant B)", "1200", "1600"))
    assert table.site_keys == ("north", "south")
    assert table.title == "Thermal capacity"
    assert table.page == 2  # the Word table item, after its heading item


def test_separately_titled_word_tables_keep_their_headings_and_rows():
    document = Document()
    document.add_paragraph("Steam")
    _table(document, [["Facility", "Required steam\n(lb/hr)", "Available steam\n(lb/hr)"], ["NH", "40000", "50000"]])
    document.add_paragraph("Chilled water")
    _table(document, [["Facility", "Required cooling\n(tons)", "Available cooling\n(tons)"], ["NH", "1200", "1600"]])
    tables = _inspect(document).tables
    assert [(t.title, t.columns[1], t.rows) for t in tables] == [
        ("Steam", "Required steam (lb/hr)", (("NH", "40000", "50000"),)),
        ("Steam / Chilled water", "Required cooling (tons)", (("NH", "1200", "1600"),))]
    assert all(t.site_keys == ("north",) for t in tables)


def test_a_table_without_a_site_column_still_takes_its_site_from_the_heading():
    document = Document()
    document.add_paragraph("North Hospital thermal capacity")
    _table(document, [["Service", "Required capacity\n(units)", "Available capacity"], ["Steam", "40000", "50000"]])
    table = _inspect(document).tables[0]
    assert table.columns[1] == "Required capacity (units)"
    assert table.site_keys == ("north",)
