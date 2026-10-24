"""Synthetic diagnostic: table caption styling is not new narrative styling."""
from io import BytesIO

import pytest
from docx import Document
from docx.shared import Pt

from app.monthly_report_model import Facility, ReportDraft, ReportPeriod, ReportProfile, ResolvedBlock, default_sections
from app.monthly_report_native_layout import build_native_docx


def caption_scene(tmp_path, *, caption=True, body=False, total=False):
    document = Document()
    document.add_heading('PRIORITY CAPITAL RENEWAL LIST', 1)
    document.add_heading('Synthetic Equipment Planning', 2)
    paragraph = document.add_paragraph('Synthetic Mechanical Recommendations' if caption else 'Review remains pending.')
    paragraph.runs[0].bold = paragraph.runs[0].underline = True
    paragraph.runs[0].font.size = Pt(10)
    if caption:
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = 'Equipment Description'
        table.cell(0, 1).text = 'End of Useful Life'
        table.cell(1, 0).text = 'Synthetic pump'
        table.cell(1, 1).text = '2040'
    if total:
        paragraph = document.add_paragraph('Total Recommendations\tN/A')
        paragraph.paragraph_format.left_indent = Pt(100)
        paragraph.runs[0].bold = True
    if body:
        paragraph = document.add_paragraph('An ordinary synthetic planning narrative was reviewed.')
        paragraph.runs[0].font.size = Pt(12)
    source = tmp_path / 'synthetic-caption.docx'
    document.save(source)
    draft = ReportDraft(ReportProfile('Synthetic', 'example', 'Synthetic Facility',
        (Facility('example', 'Synthetic Facility'),)), ReportPeriod(2026, 10),
        'Synthetic Editor', default_sections(),
        (ResolvedBlock('capital_renewal', 'This month', text='Historical context remains unconfirmed for this synthetic review.'),))
    return source, draft


@pytest.mark.parametrize('section', [None, 'capital'])
def test_new_prose_does_not_borrow_table_caption_typography(tmp_path, section):
    source, draft = caption_scene(tmp_path, body=True)
    before = source.read_bytes()
    output = Document(BytesIO(build_native_docx(draft, source, section_key=section)))
    paragraph = next(p for p in output.paragraphs if p.text.startswith('Historical context'))
    assert paragraph.runs[0].bold is not True
    assert paragraph.runs[0].underline is not True
    assert paragraph.runs[0].font.size == Pt(12)
    assert source.read_bytes() == before


@pytest.mark.parametrize('total', [False, True])
def test_caption_only_role_uses_document_default_not_totals(tmp_path, total):
    source, draft = caption_scene(tmp_path, total=total)
    output = Document(BytesIO(build_native_docx(draft, source, section_key='capital')))
    paragraph = next(p for p in output.paragraphs if p.text.startswith('Historical context'))
    assert paragraph.runs[0].bold is not True
    assert paragraph.runs[0].underline is not True
    assert paragraph.paragraph_format.left_indent is None
    assert paragraph.runs[0].font.size is None


def test_intentionally_emphasized_prose_is_not_blanket_stripped(tmp_path):
    source, draft = caption_scene(tmp_path, caption=False)
    output = Document(BytesIO(build_native_docx(draft, source, section_key='capital')))
    paragraph = next(p for p in output.paragraphs if p.text.startswith('Historical context'))
    assert paragraph.runs[0].bold is True
    assert paragraph.runs[0].underline is True


@pytest.mark.parametrize('case', ['no_table', 'no_outline', 'sentence', 'mixed_emphasis', 'other_role', 'section_boundary'])
def test_unproved_caption_keeps_existing_typography(tmp_path, case):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from app.monthly_report_narrative_table_roles import table_body_roles
    source, _ = caption_scene(tmp_path)
    document = Document(source)
    caption = document.paragraphs[-1]
    table = document.tables[0]._tbl
    original = list(document.element.body)
    scopes = {element: ('example', 'planning') for element in original}
    if case == 'no_table':
        table.getparent().remove(table)
    elif case == 'no_outline':
        document.paragraphs[-2].style = document.styles['Normal']
    elif case == 'sentence':
        caption.runs[0].text = 'Review remains pending.'
    elif case == 'mixed_emphasis':
        caption.add_run(' Ordinary prose').bold = False
    elif case == 'other_role':
        scopes[table] = ('example', 'other')
    elif case == 'section_boundary':
        boundary = OxmlElement('w:p')
        properties = OxmlElement('w:pPr')
        properties.append(OxmlElement('w:sectPr'))
        boundary.append(properties)
        caption._p.addprevious(boundary)
        scopes[boundary] = scopes[caption._p]
    excluded, neutral = table_body_roles(list(document.element.body),
        {style.style_id: style.element for style in document.styles}, document.styles.element, scopes)
    assert caption._p not in excluded
    assert not neutral


@pytest.mark.parametrize('prompt', ['', 'Location / Campus / Building:', 'Activity Summary'])
def test_neutral_fallback_does_not_borrow_other_role_scope(tmp_path, prompt):
    source, draft = caption_scene(tmp_path, total=True)
    document = Document(source)
    document.styles['Normal'].font.name = 'Arial'
    document.styles['Normal'].font.size = Pt(11)
    if prompt:
        document.add_paragraph(prompt)
    label = document.add_paragraph('Contract Asset End of Useful Life Schedule')
    label.runs[0].bold = label.runs[0].underline = True
    scope = document.add_paragraph('Synthetic reference scope belongs to its own schedule.')
    scope.runs[0].italic = True
    scope.runs[0].font.size = Pt(7)
    document.add_table(rows=1, cols=1).cell(0, 0).text = 'Description'
    document.save(source)
    output = Document(BytesIO(build_native_docx(draft, source, section_key='capital')))
    paragraph = next(p for p in output.paragraphs if p.text.startswith('Historical context'))
    assert paragraph.runs[0].italic is not True
    assert paragraph.runs[0].font.size is None
    assert paragraph.style.font.name == 'Arial'
    assert paragraph.style.font.size == Pt(11)
    assert paragraph.runs[0].bold is not True
    assert paragraph.paragraph_format.left_indent is None


def test_unchanged_source_captions_and_headings_remain_exact(tmp_path):
    from dataclasses import replace
    from app.monthly_report_import import inspect_docx
    from app.monthly_report_sections import table_without_prices
    source, draft = caption_scene(tmp_path, total=True)
    inspection = inspect_docx(source)
    selected = [item for item in inspection.items if item.kind == 'text'
        and item.section == 'capital' and item.note != 'Native section heading']
    tables = tuple(replace(table_without_prices(item)[0],
        reference=f'docx:{inspection.sha256}:{item.id}') for item in inspection.items if item.kind == 'table')
    block = replace(draft.blocks[0], text='\n\n'.join(item.text.strip() for item in selected),
        references=tuple(f'docx:{inspection.sha256}:{item.id}' for item in selected), extra_tables=tables)
    original = Document(source)
    before = source.read_bytes()
    output = Document(BytesIO(build_native_docx(replace(draft, blocks=(block,)), source, section_key='capital')))
    for paragraph in original.paragraphs:
        current = next(p for p in output.paragraphs if p.text == paragraph.text)
        assert current._p.xml == paragraph._p.xml
    assert output.tables[0]._tbl.xml == original.tables[0]._tbl.xml
    assert source.read_bytes() == before


@pytest.mark.parametrize('body', ['', 'before', 'after'])
def test_qualifier_stays_before_retained_table(tmp_path, body):
    from dataclasses import replace
    from app.monthly_report_import import inspect_docx
    from app.monthly_report_sections import table_without_prices
    source, draft = caption_scene(tmp_path, total=True, body=bool(body))
    if body == 'before':
        document = Document(source)
        document.paragraphs[1]._p.addprevious(document.paragraphs[-1]._p)
        document.save(source)
    inspection = inspect_docx(source)
    item = next(item for item in inspection.items if item.kind == 'table')
    table = replace(table_without_prices(item)[0], reference=f'docx:{inspection.sha256}:{item.id}')
    block = replace(draft.blocks[0], extra_tables=(table,))
    output = Document(BytesIO(build_native_docx(replace(draft, blocks=(block,)), source, section_key='capital')))
    paragraph = next(p for p in output.paragraphs if p.text.startswith('Historical context'))
    elements = list(output.element.body)
    assert elements.index(paragraph._p) < elements.index(output.tables[0]._tbl)
    assert output.tables[0].cell(1, 0).text == 'Synthetic pump'
    assert paragraph.runs[0].bold is not True
