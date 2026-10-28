"""Saved-source reads must stay bounded across stat/open races.

All bytes are synthetic and held in a temporary report profile. The test
checks the *actual read(n) request*, not only whether the final hash rejects
a changed file after an unlimited allocation has already happened.
"""
from hashlib import sha256
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace

from docx import Document
import pytest

from app import monthly_report_library as library
from app import monthly_report_sources as sources
from app.monthly_report_model import Facility, ReportProfile


@pytest.fixture
def profile(monkeypatch, tmp_path):
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path))
    report = ReportProfile(
        "Synthetic Source Read Contract", "north", "Synthetic North",
        (Facility("north", "Synthetic North"),))
    library.save_profile(report, expected_revision=0,
                         actor="Synthetic Editor", confirmed=True)
    return report


def ingest_csv(profile, name="capacity.csv", *, row="NH,Steam"):
    raw = ("Site,Service\n" + row + "\n").encode("utf-8")
    content, attachments = sources.ingest(profile, name, raw)
    assert attachments == ()
    stored = sources._path(profile, content.source.sha256, content.source.suffix)
    assert stored.read_bytes() == raw
    return content.source, stored, raw


def track_binary_reads(monkeypatch, path, *, expected_bound):
    """Reject a hidden read(-1) or Path.read_bytes() after a harmless stat."""
    calls = []
    original_open = Path.open
    original_read_bytes = Path.read_bytes

    class CheckedStream:
        def __init__(self, stream):
            self.stream = stream

        def __enter__(self):
            self.stream.__enter__()
            return self

        def __exit__(self, *args):
            return self.stream.__exit__(*args)

        def read(self, n=-1):
            calls.append(n)
            if n != expected_bound:
                raise AssertionError("Saved-source read was not bounded by its configured size.")
            return self.stream.read(n)

    def open_guard(self, *args, **kwargs):
        raw = original_open(self, *args, **kwargs)
        if self != path:
            return raw
        if not args or args[0] != "rb":
            raw.close()
            raise AssertionError("Saved-source read must open the original in binary mode.")
        return CheckedStream(raw)

    def forbid_unbounded(self):
        if self == path:
            raise AssertionError("Unrestricted Path.read_bytes() is not a safe source read.")
        return original_read_bytes(self)

    monkeypatch.setattr(Path, "open", open_guard)
    monkeypatch.setattr(Path, "read_bytes", forbid_unbounded)
    return calls


def fake_stale_stat(monkeypatch, path, visible_size):
    original_stat = Path.stat

    def old_size(self, *args, **kwargs):
        if self == path:
            return SimpleNamespace(st_size=visible_size)
        return original_stat(self, *args, **kwargs)

    monkeypatch.setattr(Path, "stat", old_size)


def test_unchanged_csv_source_reads_entire_original_and_keeps_review_identity(profile, monkeypatch):
    source, path, raw = ingest_csv(profile)
    fingerprint = source.fingerprint
    monkeypatch.setattr(sources, "MAX_FILE_BYTES", len(raw) + 9)
    calls = track_binary_reads(monkeypatch, path, expected_bound=len(raw) + 10)
    assert sources.source_bytes(profile, source) == raw
    assert calls == [len(raw) + 10]
    assert source.fingerprint == fingerprint
    assert path.is_file()


def test_exact_current_type_limit_is_accepted_without_truncation(profile, monkeypatch):
    source, path, raw = ingest_csv(profile)
    monkeypatch.setattr(sources, "MAX_FILE_BYTES", len(raw))
    calls = track_binary_reads(monkeypatch, path, expected_bound=len(raw) + 1)
    assert sources.source_bytes(profile, source) == raw
    assert calls == [len(raw) + 1]
    assert sha256(raw).hexdigest() == source.sha256


def test_early_over_limit_refusal_never_opens_the_file(profile, monkeypatch):
    source, path, raw = ingest_csv(profile)
    monkeypatch.setattr(sources, "MAX_FILE_BYTES", len(raw) - 1)
    original_open = Path.open
    opened = []

    def denied(self, *args, **kwargs):
        if self == path:
            opened.append(True)
            raise AssertionError("Over-limit file must fail at the stat boundary.")
        return original_open(self, *args, **kwargs)

    monkeypatch.setattr(Path, "open", denied)
    with pytest.raises(ValueError, match="Saved source exceeds the file limit"):
        sources.source_bytes(profile, source)
    assert opened == []
    assert source.sha256 == sha256(raw).hexdigest()


def test_growth_after_small_stat_is_rejected_after_at_most_limit_plus_one_bytes(profile, monkeypatch):
    source, path, original = ingest_csv(profile)
    limit = len(original) + 12
    # Only the saved synthetic test file changes. The reviewed source identity
    # and original bytes remain distinct from the replaced on-disk blob.
    changed = b"X" * (limit + 211)
    library._atomic_write(path, changed)
    monkeypatch.setattr(sources, "MAX_FILE_BYTES", limit)
    fake_stale_stat(monkeypatch, path, len(original))
    calls = track_binary_reads(monkeypatch, path, expected_bound=limit + 1)
    with pytest.raises(ValueError, match="Saved source exceeds the file limit"):
        sources.source_bytes(profile, source)
    assert calls == [limit + 1]
    assert source.sha256 == sha256(original).hexdigest()
    assert source.sha256 != sha256(changed).hexdigest()


def test_exactly_limit_plus_one_after_stat_growth_does_not_pass(profile, monkeypatch):
    source, path, original = ingest_csv(profile)
    limit = len(original)
    library._atomic_write(path, b"R" * (limit + 1))
    monkeypatch.setattr(sources, "MAX_FILE_BYTES", limit)
    fake_stale_stat(monkeypatch, path, limit)
    calls = track_binary_reads(monkeypatch, path, expected_bound=limit + 1)
    with pytest.raises(ValueError, match="Saved source exceeds the file limit"):
        sources.source_bytes(profile, source)
    assert calls == [limit + 1]


def test_same_length_modified_blob_still_fails_immutable_hash_check(profile, monkeypatch):
    source, path, raw = ingest_csv(profile)
    replaced_raw = raw.replace(b"NH,Steam", b"SH,Steam")
    assert len(replaced_raw) == len(raw) and replaced_raw != raw
    library._atomic_write(path, replaced_raw)
    monkeypatch.setattr(sources, "MAX_FILE_BYTES", len(raw) + 10)
    calls = track_binary_reads(monkeypatch, path, expected_bound=len(raw) + 11)
    with pytest.raises(ValueError, match="integrity check"):
        sources.source_bytes(profile, source)
    assert calls == [len(raw) + 11]
    assert source.sha256 == sha256(raw).hexdigest()


def test_one_bad_source_does_not_change_independent_saved_sibling(profile, monkeypatch):
    first, path_one, raw_one = ingest_csv(profile, "north.csv", row="NH,Steam")
    second, path_two, raw_two = ingest_csv(profile, "south.csv", row="SH,Water")
    assert path_one != path_two
    original_two_sha = second.sha256
    library._atomic_write(path_one, raw_one + b"Unapproved appended bytes" * 11)
    monkeypatch.setattr(sources, "MAX_FILE_BYTES", len(raw_one) + 5)
    fake_stale_stat(monkeypatch, path_one, len(raw_one))
    with pytest.raises(ValueError, match="Saved source exceeds the file limit"):
        sources.source_bytes(profile, first)
    assert sources.source_bytes(profile, second) == raw_two
    assert second.sha256 == original_two_sha
    assert path_two.read_bytes() == raw_two


def test_real_small_docx_uses_128_mib_byte_cap_without_allocating_it(profile, monkeypatch):
    doc = Document()
    doc.add_heading("Synthetic prior-month report", 1)
    doc.add_paragraph("No private report content is present.")
    out = BytesIO()
    doc.save(out)
    raw = out.getvalue()
    content, attachments = sources.ingest(profile, "example.docx", raw)
    assert attachments == ()
    source = content.source
    path = sources._path(profile, source.sha256, source.suffix)
    expected = 128 * 1024 * 1024 + 1
    calls = track_binary_reads(monkeypatch, path, expected_bound=expected)
    assert sources.source_bytes(profile, source) == raw
    assert calls == [expected]
    assert source.sha256 == sha256(raw).hexdigest()


def test_missing_source_remains_a_missing_source_not_an_empty_success(profile):
    source, path, raw = ingest_csv(profile)
    path.unlink()
    with pytest.raises(FileNotFoundError):
        sources.source_bytes(profile, source)
    assert source.sha256 == sha256(raw).hexdigest()
