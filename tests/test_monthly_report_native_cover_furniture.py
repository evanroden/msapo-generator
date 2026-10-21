"""Exact supported cover furniture is preserved, never classified as a logo."""
from copy import deepcopy
from dataclasses import replace
from io import BytesIO

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches
from lxml import etree as ET
from PIL import Image
import pytest

from app import monthly_report_native_art_review as art
from app.monthly_report_native_design import A, W, WPS, WPG, MC
from test_monthly_report_source_frame_logo_role import source
from test_monthly_report_native_art_review import setup, rewrite, change_xml
from test_monthly_report_native_art_lifecycle import save_sample, plans_for

V = '{urn:schemas-microsoft-com:vml}'


def band_source(tmp_path, *, mutation=None):
    path = source(tmp_path, divergent=True)
    doc = Document(path)
    raw = BytesIO(); Image.new('RGB', (640,120), 'orange').save(raw,'PNG')
    drawing = doc.add_picture(BytesIO(raw.getvalue()), width=Inches(3))._inline.getparent()
    paragraph = drawing.getparent().getparent()
    data = next(drawing.iter(A+'graphicData'))
    picture = data[0]; data.remove(picture); data.set('uri',WPG[1:-1])
    group = ET.SubElement(data,WPG+'wgp')
    ET.SubElement(group,WPG+'cNvGrpSpPr'); ET.SubElement(group,WPG+'grpSpPr')
    polygon = ET.SubElement(group,WPS+'wsp'); ET.SubElement(polygon,WPS+'cNvSpPr')
    props=ET.SubElement(polygon,WPS+'spPr'); transform=ET.SubElement(props,A+'xfrm')
    ET.SubElement(transform,A+'off',x='0',y='0');ET.SubElement(transform,A+'ext',cx='4000000',cy='900000')
    ET.SubElement(props,A+'prstGeom',prst='rect')
    ET.SubElement(ET.SubElement(props,A+'solidFill'),A+'srgbClr',val='082b23')
    ET.SubElement(polygon,WPS+'bodyPr');group.append(picture)
    run=drawing.getparent();run.remove(drawing)
    wrapper=ET.SubElement(run,MC+'AlternateContent',nsmap={'wpg':WPG[1:-1]})
    ET.SubElement(wrapper,MC+'Choice',Requires='wpg').append(drawing)
    fallback=ET.SubElement(wrapper,MC+'Fallback');pict=ET.SubElement(fallback,W+'pict')
    legacy=ET.SubElement(pict,V+'group',style='width:300pt;height:70pt',coordsize='4000,900')
    ET.SubElement(legacy,V+'rect',style='width:4000;height:900',fillcolor='#082b23',stroked='f')
    image=ET.SubElement(legacy,V+'shape',style='width:3000;height:600')
    ET.SubElement(image,V+'imagedata',{qn('r:id'):next(picture.iter(A+'blip')).get(qn('r:embed'))})
    ET.SubElement(ET.SubElement(paragraph,W+'r'),W+'t').text='='
    doc.element.body.remove(paragraph);doc.element.body.insert(0,paragraph)
    if mutation: mutation(doc,paragraph,wrapper)
    doc.save(path)
    return path


@pytest.fixture
def band_sample(tmp_path,monkeypatch):
    monkeypatch.setenv('EPC_DATA_DIR',str(tmp_path/'store'))
    return setup(band_source(tmp_path))


def test_band_inventory_exact_branches_has_no_role_authority(band_sample):
    path,view,profile,period,inventory,receipt=band_sample
    bands=[o for o in inventory.occurrences if o.preservation_kind=='cover_band']
    assert len(bands)==1
    band=bands[0]
    assert len(band.original_representations)==len(band.rendition_representations)==2
    item=next(i for i in view.original_inspection.items if i.id==band.original_id)
    assert item.image_role=='' and item.suggested_slot=='cover_photo'
    assert band.original_position==band.rendition_position==band.original_word_section==1
    assert all(p.media_sha256==item.image_digest for p in band.original_representations)
    art.verify_native_art_output(view.rendition_bytes,art.verify_native_art_review(path,view,profile,period,receipt),stage='rendition')


def test_band_photo_exclusion_does_not_approve_preservation(band_sample):
    from app.monthly_report_start_import import cover_photo_candidates, unresolved_cover_artwork
    path,view,profile,period,inventory,receipt=band_sample
    ids=art.native_cover_furniture_ids(view.original_inspection.items)
    assert len(ids)==1
    assert ids <= {i.id for i in cover_photo_candidates(view.original_inspection.items)}
    assert not ids.intersection(i.id for i in cover_photo_candidates(view.original_inspection.items,native_furniture_ids=ids))
    assert ids <= {i.id for i in unresolved_cover_artwork(view.original_inspection.items,native_furniture_ids=ids)}
    assert not plans_for(band_sample,False)['cover']['approved']
    assert plans_for(band_sample)['cover']['approved']


def test_band_saved_routine_full_cover_activity_preservation_and_replay(band_sample):
    from app.monthly_report_native_layout import build_native_docx
    from app.monthly_report_preview import section_preview_docx
    from app.monthly_report_native_art_scope import expected_header_scope
    from app import monthly_report_library as library
    path,view,profile,period,inventory,receipt=band_sample
    before=path.read_bytes();draft,loader=save_sample(band_sample)
    context=art.verify_native_art_review(path,view,profile,period,receipt)
    assert library.load_imported_draft(profile.contract,profile.key)==draft
    for key in (None,'cover','activity'):
        output=(build_native_docx(draft,path,asset_loader=loader) if key is None
                else section_preview_docx(draft,key,loader))
        art.verify_native_art_output(output,context,source_scope=expected_header_scope(view,key))
    replay,verified=library.starting_contribution(profile.contract,profile.key,period,view.original_inspection.sha256)
    assert verified and replay==draft
    assert path.read_bytes()==before


@pytest.mark.parametrize('damage', ['position', 'section', 'role', 'second-image', 'other-native-kind'])
def test_structural_photo_exclusion_has_exact_bounds(band_sample, damage):
    _,view,_,_,_,_=band_sample
    items=list(view.original_inspection.items)
    identity=next(iter(art.native_cover_furniture_ids(items)))
    image=next(i for i in items if i.id==identity)
    if damage=='second-image': items.append(replace(image,id='independent-same-byte'))
    elif damage=='other-native-kind':
        items=[replace(i,note='Editable native text frames') if i.kind=='native_layout' and i.position==1 else i for i in items]
    else:
        values={'position':{'position':2},'section':{'word_section':2},'role':{'image_role':'client_logo'}}[damage]
        items=[replace(i,**values) if i.id==identity else i for i in items]
    assert not art.native_cover_furniture_ids(items)


@pytest.mark.parametrize('damage', ['duplicate-picture', 'extra-text', 'nested-wrapper', 'command', 'conflicting-role'])
def test_raw_branch_ambiguity_never_gets_preservation_authority(tmp_path,monkeypatch,damage):
    from app.monthly_report_compatibility_view import derive_compatibility_view
    from app.monthly_report_model import synthetic_profiles,Facility,ReportPeriod
    monkeypatch.setenv('EPC_DATA_DIR',str(tmp_path/'store'))
    def mutate(doc,paragraph,wrapper):
        if damage=='duplicate-picture':
            branch=wrapper[1];node=next(branch.iter(V+'shape'));node.getparent().append(deepcopy(node))
        elif damage=='extra-text':
            for branch in wrapper:
                ET.SubElement(ET.SubElement(branch,W+'r'),W+'t').text='Extra source content'
        elif damage=='nested-wrapper':
            branch=wrapper[1];node=branch[0];branch.remove(node)
            nested=ET.SubElement(branch,MC+'AlternateContent');ET.SubElement(nested,MC+'Choice',Requires='wpg').append(node)
        elif damage=='command':
            ET.SubElement(ET.SubElement(paragraph,W+'r'),W+'instrText').text='INCLUDETEXT forbidden'
        else:
            for node in wrapper.iter():
                if node.tag in (qn('wp:docPr'),qn('pic:cNvPr')):node.set('descr','Client logo')
    path=band_source(tmp_path,mutation=mutate)
    try:
        view=derive_compatibility_view(path.read_bytes())
        profile=replace(synthetic_profiles()[0],imported_from=view.original_inspection.sha256,
            title='Synthetic Regional Community Health Hospital',facilities=(Facility('synthetic','Synthetic Regional Community Health Hospital'),))
        inventory=art.native_art_inventory(path,view,profile,ReportPeriod(2026,9))
    except ValueError:
        return
    assert not any(o.preservation_kind=='cover_band' for o in inventory.occurrences)


def test_same_byte_outside_band_is_not_preserved_or_reclassified(band_sample,tmp_path):
    from app.monthly_report_compatibility_view import derive_compatibility_view
    from app.monthly_report_start_import import cover_photo_candidates
    from zipfile import ZipFile
    path,view,profile,period,_,_=band_sample
    original=next(i for i in view.original_inspection.items if i.id in art.native_cover_furniture_ids(view.original_inspection.items))
    with ZipFile(path) as package:raw=package.read(original.image_part)
    doc=Document(path);doc.add_picture(BytesIO(raw),width=Inches(3));doc.save(path)
    new=derive_compatibility_view(path.read_bytes());profile=replace(profile,imported_from=new.original_inspection.sha256)
    inventory=art.native_art_inventory(path,new,profile,period)
    band=next(o for o in inventory.occurrences if o.preservation_kind=='cover_band')
    matching=[i for i in new.original_inspection.items if i.kind=='image' and i.image_digest==original.image_digest]
    assert len(matching)==2
    assert {o.original_id for o in inventory.occurrences}=={band.original_id}
    assert all(i.image_role=='' for i in matching)


@pytest.mark.parametrize('damage',['remove','move','geometry','media','wrapper-context'])
def test_output_band_placement_and_geometry_are_not_normalized_away(band_sample,damage):
    from app.monthly_report_native_layout import build_native_docx
    from app.monthly_report_native_art_scope import expected_header_scope
    path,view,profile,period,inventory,receipt=band_sample
    draft,loader=save_sample(band_sample);output=build_native_docx(draft,path,asset_loader=loader)
    context=art.verify_native_art_review(path,view,profile,period,receipt)
    proof=next(o for o in inventory.occurrences if o.preservation_kind=='cover_band').rendition_representations[0]
    def corrupt(parts):
        if damage=='media':parts[proof.media_part]+=b'changed';return
        def edit(root):
            from app.monthly_report_compatibility_view import _path_node
            drawing=_path_node(root,proof.drawing_path)
            if damage=='remove':drawing.getparent().remove(drawing)
            elif damage=='move':drawing.getparent().getparent().append(drawing)
            elif damage=='geometry':next(drawing.iter(qn('wp:extent'))).set('cx','999999')
            else:drawing.getparent().set('unknown','changed context')
        change_xml(parts,'word/document.xml',edit)
    with pytest.raises(ValueError):
        art.verify_native_art_output(rewrite(output,corrupt),context,source_scope=expected_header_scope(view))


@pytest.mark.parametrize('action',['map','omit','other-block'])
def test_exact_preserved_band_mapping_or_omission_refused_before_save_and_render(band_sample,action):
    from app.monthly_report_start_import import build_starting_import
    from app.monthly_report_native_layout import build_native_docx
    from app.monthly_report_model import ResolvedBlock
    path,view,profile,period,inventory,receipt=band_sample
    identity=next(o.original_id for o in inventory.occurrences if o.preservation_kind=='cover_band')
    plans=plans_for(band_sample)
    plans['cover']['selected'].append(identity);plans['cover']['destinations'][identity]='cover_photo'
    with pytest.raises(ValueError,match='preservation conflicts'):
        build_starting_import(path,view.original_inspection,tuple(plans.values()),source_view=view,profile=profile,period=period)
    draft,loader=save_sample(band_sample)
    block=ResolvedBlock('client_logo' if action=='other-block' else 'cover_photo',
        'Omit' if action=='omit' else 'Manual',references=(f'docx:{profile.imported_from}:{identity}',))
    changed=replace(draft,blocks=tuple(b for b in draft.blocks if b.key!=block.key)+(block,))
    with pytest.raises(ValueError):build_native_docx(changed,path,asset_loader=loader)


def test_routine_ui_excludes_only_band_from_gallery_before_review_and_saves_after(band_sample,monkeypatch):
    from streamlit.testing.v1 import AppTest
    from app import monthly_report_section_ui as ui,monthly_report_picture_cards as cards,monthly_report_library as library
    path,view,profile,period,inventory,receipt=band_sample
    upload=BytesIO(path.read_bytes());upload.name='synthetic-band.docx';upload.size=len(upload.getvalue())
    real_upload=ui.st.file_uploader
    monkeypatch.setattr(ui.st,'file_uploader',lambda label,*a,**k:upload if label=='Starting report' else real_upload(label,*a,**k))
    offered=[];real_gallery=cards.review_cover_gallery
    def gallery(*args,**kwargs):
        offered.extend(i.id for i in args[0])
        return real_gallery(*args,**kwargs)
    monkeypatch.setattr(cards,'review_cover_gallery',gallery)
    app=AppTest.from_string(f'''
from pathlib import Path
from dataclasses import replace
from test_monthly_report_native_art_review import setup
from app.monthly_report_section_ui import render_section_setup
from app.monthly_report_ui import _field
path,view,profile,period,inventory,_=setup(Path({str(path)!r}))
render_section_setup(profile.contract,period,'Synthetic Operator',_field,identity=replace(profile,imported_from=''))
''',default_timeout=60).run()
    assert not app.exception
    band=next(o.original_id for o in inventory.occurrences if o.preservation_kind=='cover_band')
    assert offered and band not in offered
    assert any('no confirmed' in e.value for e in app.error)
    for checkbox in app.checkbox:
        if checkbox.label.startswith('I reviewed this artwork and want to keep it unchanged'):checkbox.check()
    app.run();assert not app.exception
    button=next(b for b in app.button if b.label=='Continue to this month’s updates')
    assert not button.disabled
    button.click().run();assert not app.exception and not app.error
    loaded=library.load_imported_draft(profile.contract,profile.key)
    assert loaded.profile.imported_from==profile.imported_from


def test_omit_actual_photo_does_not_target_preserved_band(band_sample):
    from app.monthly_report_native_layout import build_native_docx
    from app.monthly_report_native_art_scope import expected_header_scope
    path,view,profile,period,inventory,receipt=band_sample
    draft,loader=save_sample(band_sample)
    photo=next(b for b in draft.blocks if b.key=='cover_photo')
    band=next(o for o in inventory.occurrences if o.preservation_kind=='cover_band')
    assert photo.references and all(band.original_id not in ref for ref in photo.references)
    changed=replace(draft,blocks=tuple(replace(b,source='Omit') if b.key=='cover_photo' else b for b in draft.blocks))
    output=build_native_docx(changed,path,asset_loader=loader)
    context=art.verify_native_art_review(path,view,profile,period,receipt)
    art.verify_native_art_output(output,context,source_scope=expected_header_scope(view))


def test_header_only_old_policy_receipt_cannot_authorize_new_band(band_sample):
    path,view,profile,period,inventory,receipt=band_sample
    old=deepcopy(receipt);old['policy']='native-source-art-preservation-v2-experimental'
    with pytest.raises(ValueError):art.verify_native_art_review(path,view,profile,period,old)
