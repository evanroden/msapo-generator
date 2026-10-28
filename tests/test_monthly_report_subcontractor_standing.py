"""Subcontractor contact/status reuse is contract-scoped, versioned and report-pinned.

These tests use only synthetic contacts under an isolated temporary library.
They verify actual official Section 6 Word output and snapshot persistence,
without touching a real master, source import, or production data.
"""
from dataclasses import asdict, replace
from io import BytesIO
import json

from docx import Document
import pytest

from app import monthly_report_library as library
from app.monthly_report_docx import assemble_docx
from app.monthly_report_model import (
    Facility, ReportDraft, ReportPeriod, ReportProfile, ResolvedBlock,
    default_sections,
)
from app.monthly_report_setup import new_month_draft
from app.monthly_report_subcontractor_standing import (
    VendorStandingRow, load_vendor_roster, roster_version, save_vendor_roster,
    effective_vendor_rows, report_vendor_block, rollover_subcontractor_block,
)

SEPTEMBER = ReportPeriod(2026, 9)
OCTOBER = ReportPeriod(2026, 10)


@pytest.fixture
def contracts(monkeypatch, tmp_path):
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path))
    north = Facility("north", "Site North", ("North",))
    south = Facility("south", "Site South", ("South",))
    both = ReportProfile("Synthetic Contract A", "both", "Contract A",
                         (north, south), "multi_site")
    a_north = ReportProfile(both.contract, "north", "Contract A North",
                            (north,))
    a_south = ReportProfile(both.contract, "south", "Contract A South",
                            (south,))
    other = ReportProfile("Synthetic Contract B", "north", "Contract B North",
                          (Facility("north", "Other North"),))
    return both, a_north, a_south, other


def rows():
    return (
        VendorStandingRow(
            "shared-vendor", "Mechanical", "HVAC", "Synthetic Mechanical",
            "Shared technician", "555-0101", "Agreement on file",
            when_to_contact="Call for scheduled equipment service.",
            status="Available", status_period="2026-09",
            source_reference="synth-source:1:exact"),
        VendorStandingRow(
            "south-exception", "Mechanical", "HVAC", "Synthetic Mechanical",
            "Southern technician", "555-0102", "Agreement on file",
            sites=("south",), override_of="shared-vendor",
            when_to_contact="Call southern service coordinator.",
            status="On call", status_period="2026-09",
            source_reference="synth-source:2:exact"),
        VendorStandingRow(
            "north-water", "Water treatment", "Chemical",
            "Synthetic Water Vendor", "Chemical specialist", "555-0103",
            sites=("north",), when_to_contact="Contact for site water schedule.",
            source_reference="synth-source:3:exact"),
    )


def approved(profile):
    return save_vendor_roster(
        profile, rows(), expected_revision=0, actor="Synthetic editor", confirmed=True)


def report(profile, block, period=SEPTEMBER):
    section = next(s for s in default_sections() if s.key == "subcontractors")
    return ReportDraft(profile, period, "Synthetic Editor", (section,), (block,),
                       synthetic=True)


def test_approved_roster_reuses_contact_and_trade_for_two_sites_without_duplicates(contracts):
    both, north, south, _ = contracts
    master = approved(both)
    assert master.revision == 1
    assert master == load_vendor_roster(both.contract)
    assert len(master.rows) == 3
    selected = effective_vendor_rows(both, master, SEPTEMBER)
    assert len(selected.selected) == 3
    assert [(v.site_key, v.row.id) for v in selected.selected] == [
        ("north", "shared-vendor"), ("north", "north-water"),
        ("south", "south-exception")]
    assert selected.historical_statuses == ()
    block = report_vendor_block(both, selected)
    assert block.key == "subcontractor_matrix" and block.source == "Library"
    assert block.rows == (
        ("Site North", "Mechanical / HVAC", "Synthetic Mechanical",
         "Shared technician", "555-0101", "Agreement on file"),
        ("Site North", "Water treatment / Chemical", "Synthetic Water Vendor",
         "Chemical specialist", "555-0103", ""),
        ("Site South", "Mechanical / HVAC", "Synthetic Mechanical",
         "Southern technician", "555-0102", "Agreement on file"),
    )
    # Category band is an attribute, not another fake contact row.
    assert not any(c[2] == "" for c in block.rows)
    assert block.references == (master.reference,)
    assert len(block.extra_tables) == 1
    assert len(block.extra_tables[0].rows) == 3
    assert selected.current_statuses == (
        ("Site North", "Synthetic Mechanical", "Available"),
        ("Site South", "Synthetic Mechanical", "On call"))
    assert report_vendor_block(both, selected, existing=block) == block
    assert len(effective_vendor_rows(north, master, SEPTEMBER).selected) == 2
    assert [(v.row.id, v.row.contact) for v in effective_vendor_rows(
        south, master, SEPTEMBER).selected] == [
        ("south-exception", "Southern technician")]


def test_two_contracts_cannot_access_or_inherit_another_approved_roster(contracts):
    both, _, _, other = contracts
    first = approved(both)
    assert load_vendor_roster(other.contract).revision == 0
    assert load_vendor_roster(other.contract).rows == ()
    with pytest.raises(ValueError, match="different contract"):
        effective_vendor_rows(other, first, SEPTEMBER)
    approved_other = save_vendor_roster(
        other,
        (VendorStandingRow("contract-b", "Mechanical", "Electrical",
                           "Contract B only", sites=("north",)),),
        expected_revision=0, actor="B editor", confirmed=True)
    assert approved_other.revision == 1
    assert approved_other.reference != first.reference
    assert [r.vendor for r in load_vendor_roster(both.contract).rows] == [
        "Synthetic Mechanical", "Synthetic Mechanical", "Synthetic Water Vendor"]
    assert [r.vendor for r in load_vendor_roster(other.contract).rows] == [
        "Contract B only"]


def test_standing_updates_require_confirmation_and_exact_revision(contracts):
    both, _, _, _ = contracts
    with pytest.raises(library.LibraryError, match="Confirm"):
        save_vendor_roster(both, rows(), expected_revision=0,
                           actor="Synthetic editor", confirmed=False)
    assert load_vendor_roster(both.contract).revision == 0
    current = approved(both)
    changed = tuple(
        replace(row, phone="555-0999") if row.id == "north-water" else row
        for row in current.rows)
    second = save_vendor_roster(
        both, changed, expected_revision=1, actor="Current editor",
        confirmed=True)
    assert second.revision == 2
    assert roster_version(both.contract, 1) == current
    assert roster_version(both.contract, 2) == second
    assert next(r for r in current.rows if r.id == "north-water").phone == "555-0103"
    with pytest.raises(library.RevisionConflict, match="updated elsewhere"):
        save_vendor_roster(both, rows(), expected_revision=1,
                           actor="Stale editor", confirmed=True)
    assert save_vendor_roster(
        both, changed, expected_revision=2, actor="Editor",
        confirmed=True) == second
    assert roster_version(both.contract, 1).reference != second.reference


def test_removing_approved_vendor_requires_explicit_identity(contracts):
    both, _, _, _ = contracts
    current = approved(both)
    less = current.rows[:2]
    with pytest.raises(ValueError, match="explicit removal"):
        save_vendor_roster(both, less, expected_revision=1,
                           actor="Editor", confirmed=True)
    assert load_vendor_roster(both.contract) == current
    updated = save_vendor_roster(
        both, less, expected_revision=1, actor="Editor", confirmed=True,
        removed_ids=("north-water",))
    assert len(updated.rows) == 2 and updated.revision == 2
    assert len(roster_version(both.contract, 1).rows) == 3
    assert len(load_vendor_roster(both.contract).rows) == 2


def test_unconfirmed_or_ambiguous_site_exceptions_cannot_enter_master(contracts):
    both, _, _, _ = contracts
    with pytest.raises(ValueError, match="contract row"):
        save_vendor_roster(
            both, (VendorStandingRow("orphan", "Electrical", "Lighting",
                                      "Synthetic", sites=("north",),
                                      override_of="absent"),),
            expected_revision=0, actor="Editor", confirmed=True)
    with pytest.raises(ValueError, match="same site"):
        save_vendor_roster(
            both, (*rows(), replace(rows()[1], id="south-duplicate")),
            expected_revision=0, actor="Editor", confirmed=True)
    with pytest.raises(ValueError, match="confirmed contract membership"):
        save_vendor_roster(
            both, (VendorStandingRow("wrong-site", "Mechanical", "HVAC",
                                      "Synthetic", sites=("unknown",)),),
            expected_revision=0, actor="Editor", confirmed=True)
    with pytest.raises(ValueError, match="confirmation month"):
        save_vendor_roster(
            both, (replace(rows()[0], status_period=""),),
            expected_revision=0, actor="Editor", confirmed=True)
    assert load_vendor_roster(both.contract).revision == 0


def test_exact_site_report_omission_leaves_all_standing_rows_untouched(contracts):
    both, north, south, _ = contracts
    roster = approved(both)
    selected = effective_vendor_rows(north, roster, SEPTEMBER,
                                     omitted=(("north", "north-water"),))
    assert [(r.site_key, r.row.id) for r in selected.selected] == [
        ("north", "shared-vendor")]
    block = report_vendor_block(north, selected)
    assert len(block.rows) == 1
    assert len(effective_vendor_rows(south, roster, SEPTEMBER).selected) == 1
    assert len(load_vendor_roster(both.contract).rows) == 3


def test_manual_report_only_contact_edits_and_omit_never_promote_master(contracts):
    both, north, _, _ = contracts
    original = approved(both)
    selected = effective_vendor_rows(north, original, SEPTEMBER)
    block = report_vendor_block(north, selected)
    local = replace(block, rows=(
        (*block.rows[0][:3], "Report-only contact", *block.rows[0][4:]),
        block.rows[1]), text="Reviewed in this report only.")
    assert local.rows[0][3] == "Report-only contact"
    with pytest.raises(ValueError, match="edits or an omission"):
        report_vendor_block(north, selected, existing=local)
    with pytest.raises(ValueError, match="edits or an omission"):
        report_vendor_block(north, selected, existing=replace(block, source="Omit"))
    assert load_vendor_roster(both.contract) == original
    assert report_vendor_block(north, selected, existing=block) == block


def test_historical_vendor_status_is_not_relabelled_current_next_month(contracts):
    both, north, _, _ = contracts
    roster = approved(both)
    september = effective_vendor_rows(north, roster, SEPTEMBER)
    october = effective_vendor_rows(north, roster, OCTOBER)
    assert september.current_statuses == ((
        "Site North", "Synthetic Mechanical", "Available"),)
    assert october.current_statuses == ()
    assert october.historical_statuses == (
        ("Site North", "shared-vendor", "2026-09", "Available"),)
    block = report_vendor_block(north, september)
    old = report(north, block)
    next_month = new_month_draft(old, OCTOBER)
    future = next_month.blocks[0]
    assert future.rows == block.rows
    assert future.references == block.references
    assert future.extra_tables[0].rows[0][4] == "Call for scheduled equipment service."
    assert future.extra_tables[0].rows[0][5:] == ("", "")
    assert new_month_draft(next_month, ReportPeriod(2026, 11)).blocks[0] == future
    assert old.blocks[0].extra_tables[0].rows[0][5:] == ("Available", "2026-09")
    assert load_vendor_roster(both.contract) == roster


def test_snapshot_versions_keep_exact_report_only_row_and_old_master_pin(contracts):
    both, north, _, _ = contracts
    roster = approved(both)
    block = report_vendor_block(north, effective_vendor_rows(north, roster, SEPTEMBER))
    report_only = replace(block, rows=(
        (*block.rows[0][:3], "Report edited", *block.rows[0][4:]),
        *block.rows[1:]))
    saved = library.save_snapshot(report(north, report_only), expected_revision=0)
    assert saved.revision == 1
    updated = library.save_snapshot(
        report(north, block), expected_revision=1)
    assert updated.revision == 2
    version1 = library.load_snapshot_version(
        north.contract, north.key, SEPTEMBER, 1)
    assert version1.draft.blocks[0] == report_only
    assert library.load_snapshot(north.contract, north.key, SEPTEMBER).draft.blocks[0] == block
    changed = tuple(replace(r, phone="555-0222") if r.id == "shared-vendor" else r
                    for r in roster.rows)
    newer_master = save_vendor_roster(
        both, changed, expected_revision=1, actor="Editor", confirmed=True)
    assert newer_master.revision == 2
    assert version1.draft.blocks[0].references == (roster.reference,)
    assert roster_version(both.contract, 1) == roster
    assert library.load_snapshot(
        north.contract, north.key, SEPTEMBER).draft.blocks[0].rows == block.rows
    with pytest.raises(library.RevisionConflict):
        library.save_snapshot(report(north, report_only), expected_revision=1)


def test_real_section_six_word_output_has_contacts_and_distinct_status_table(contracts):
    both, north, _, _ = contracts
    roster = approved(both)
    block = report_vendor_block(north, effective_vendor_rows(north, roster, SEPTEMBER))
    docx = assemble_docx(report(north, block),
                         acknowledged_fingerprint=report(north, block).fingerprint)
    document = Document(BytesIO(docx))
    assert len(document.tables) == 2
    primary, status = document.tables
    assert tuple(c.text for c in primary.rows[0].cells) == (
        "Facility", "Discipline", "Vendor", "Contact", "Phone", "MSA")
    assert len(primary.rows) == 3
    assert tuple(c.text for c in primary.rows[1].cells) == block.rows[0]
    assert tuple(c.text for c in status.rows[0].cells) == block.extra_tables[0].columns
    assert status.rows[1].cells[5].text == "Available"
    assert status.rows[1].cells[6].text == "2026-09"
    assert "Shared technician" in primary.rows[1].cells[3].text
    assert primary.rows[2].cells[2].text == "Synthetic Water Vendor"
    assert all("synth-source" not in cell.text
               for table in document.tables for row in table.rows for cell in row.cells)
    next_month = new_month_draft(report(north, block), OCTOBER)
    future = assemble_docx(next_month, acknowledged_fingerprint=next_month.fingerprint)
    doc = Document(BytesIO(future))
    assert len(doc.tables) == 2
    assert doc.tables[1].rows[1].cells[5].text == ""
    assert doc.tables[1].rows[1].cells[6].text == ""
    assert doc.tables[0].rows[1].cells[3].text == "Shared technician"


def test_report_source_category_and_provenance_remain_in_standing_history(contracts):
    both, north, _, _ = contracts
    roster = approved(both)
    source = roster.rows[0]
    assert source.category == "Mechanical" and source.discipline == "HVAC"
    assert source.source_reference == "synth-source:1:exact"
    old = roster_version(both.contract, 1)
    assert old == roster
    current = effective_vendor_rows(north, old, SEPTEMBER)
    assert current.selected[0].row.source_reference == source.source_reference
    assert report_vendor_block(north, current).rows[0][1] == "Mechanical / HVAC"
    raw = library._read(
        library._root() / "subcontractor_rosters" /
        library._contract_directory(both.contract) / "manifest.json")
    assert raw["schema"] == 1 and raw["contract"] == both.contract
    assert raw["rows"][0]["source_reference"] == source.source_reference



def test_stale_site_omission_does_not_target_global_parent_or_ghost_row(contracts):
    both, _, south, _ = contracts
    roster = approved(both)
    # The South site has an exception; the global row ID is not the current
    # displayed South contact and must not be a misleading omit target.
    with pytest.raises(ValueError, match="not active"):
        effective_vendor_rows(south, roster, SEPTEMBER,
                              omitted=(("south", "shared-vendor"),))
    with pytest.raises(ValueError, match="not active"):
        effective_vendor_rows(south, roster, SEPTEMBER,
                              omitted=(("south", "nonexistent"),))
    selected = effective_vendor_rows(
        south, roster, SEPTEMBER, omitted=(("south", "south-exception"),))
    assert selected.selected == ()
    assert load_vendor_roster(both.contract) == roster


def test_contact_status_and_instruction_prices_fail_before_word_output(contracts):
    both, north, _, _ = contracts
    base = rows()[0]
    expensive = replace(base, when_to_contact="Invoice $500 for annual service")
    roster = save_vendor_roster(
        both, (expensive,), expected_revision=0, actor="Editor", confirmed=True)
    review = effective_vendor_rows(north, roster, SEPTEMBER)
    with pytest.raises(ValueError, match="pricing"):
        report_vendor_block(north, review)
    assert load_vendor_roster(both.contract).rows[0].when_to_contact == (
        "Invoice $500 for annual service")
    # The exact original needs explicit report-safe correction, not a rewrite
    # of approved standing or a silently discarded source field.


def test_report_only_omission_carries_without_readding_approved_vendor(contracts):
    both, north, _, _ = contracts
    roster = approved(both)
    partial = effective_vendor_rows(north, roster, SEPTEMBER,
                                     omitted=(("north", "north-water"),))
    block = report_vendor_block(north, partial)
    old = report(north, block)
    new = new_month_draft(old, OCTOBER)
    assert len(new.blocks[0].rows) == 1
    assert new.blocks[0].rows[0][2] == "Synthetic Mechanical"
    assert new.blocks[0].references == (roster.reference,)
    assert len(load_vendor_roster(north.contract).rows) == 3
