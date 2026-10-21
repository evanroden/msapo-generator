"""Native-design data rows copy a source data row, never the header (Refs #304).

In master mode each master table was blanked to its header before current rows
were written, so every Sub-Contractor row took the header's dark fill and bold
text. A source table with only a header row now yields plain data rows.
Synthetic values only.
"""
from dataclasses import replace
from io import BytesIO
from zipfile import ZipFile

from docx import Document
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import RGBColor
from lxml import etree

from app import monthly_report_designs as designs
from app.monthly_report_docx import _build_docx
from app.monthly_report_model import (ReportPeriod, ResolvedBlock, default_sections, synthetic_draft,
                                      synthetic_profiles)
from app.monthly_report_native_layout import _table

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
HEADER_FILL, DATA_FILL = "082B23", "D2E7DF"
VENDORS = (("Synthetic North", "Controls", "Synthetic Vendor A", "Synthetic Contact", "202-555-0101", "Yes"),
           ("Synthetic South", "", "Synthetic Vendor B", "", "", ""))


def _shade(cell, fill):
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shading)


def _master(tmp_path, *, data_fill=None, header_only=False):
    """A synthetic company master whose Sub-Contractor table has a dark, bold, white header."""
    path = tmp_path / "master.docx"
    document = Document()
    document.add_paragraph("Operations and Maintenance Monthly Review")
    for section in default_sections():
        document.add_heading(section.title, 1)
        document.add_paragraph("Source-only content")
        if section.key == "subcontractors":
            columns = section.blocks[0].columns
            table = document.add_table(rows=1 if header_only else 2, cols=len(columns))
            for cell, column in zip(table.rows[0].cells, columns):
                cell.text = column.title
                _shade(cell, HEADER_FILL)
                run = cell.paragraphs[0].runs[0]
                run.bold = True
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            if not header_only:
                table.rows[1].cells[0].text = "Synthetic Old Vendor"
                if data_fill:
                    for cell in table.rows[1].cells:
                        _shade(cell, data_fill)
    document.save(path)
    return path


def _vendor_rows(monkeypatch, tmp_path, **master):
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path / "data"))
    designs.install_master(_master(tmp_path, **master), actor="Synthetic Editor", expected_revision=0)
    draft = synthetic_draft(designs.pin(replace(synthetic_profiles()[0], contract="Synthetic Client")),
                            ReportPeriod(2026, 9))
    standard = next(section for section in default_sections() if section.key == "subcontractors")
    draft = replace(draft, sections=tuple(standard if section.key == standard.key else section
                                          for section in draft.sections),
                    blocks=tuple(block for block in draft.blocks if block.key != "demo_subcontractors")
                    + (ResolvedBlock(standard.blocks[0].key, "This month", rows=VENDORS),))
    body = etree.fromstring(ZipFile(BytesIO(_build_docx(draft))).read("word/document.xml"))
    table = next(table for table in body.iter(W + "tbl")
                 if "Synthetic Vendor A" in "".join(node.text or "" for node in table.iter(W + "t")))
    return table.findall(W + "tr")


def _look(row):
    fills = {node.get(W + "fill") for node in row.iter(W + "shd")}
    return {"fills": fills, "bold": any(True for _ in row.iter(W + "b")),
            "colors": {node.get(W + "val") for node in row.iter(W + "color")}}


def test_master_data_rows_do_not_copy_the_header_fill(monkeypatch, tmp_path):
    header, *data = _vendor_rows(monkeypatch, tmp_path)
    assert _look(header)["fills"] == {HEADER_FILL} and _look(header)["bold"]
    assert len(data) == 2
    assert all(_look(row) == {"fills": set(), "bold": False, "colors": set()} for row in data)


def test_master_data_rows_keep_the_master_data_row_look(monkeypatch, tmp_path):
    _, *data = _vendor_rows(monkeypatch, tmp_path, data_fill=DATA_FILL)
    assert all(_look(row)["fills"] == {DATA_FILL} and not _look(row)["bold"] for row in data)


def test_a_header_only_source_table_yields_plain_data_rows(monkeypatch, tmp_path):
    header, *data = _vendor_rows(monkeypatch, tmp_path, header_only=True)
    assert _look(header) == {"fills": {HEADER_FILL}, "bold": True, "colors": {"FFFFFF"}}
    assert all(_look(row) == {"fills": set(), "bold": False, "colors": set()} for row in data)
    assert ["".join(node.text or "" for node in row.iter(W + "t")) for row in data] == [
        "".join(values) for values in VENDORS]


def _xml_table(*rows):
    body = "".join("<w:tr>" + "".join(f'<w:tc><w:tcPr><w:shd w:fill="{fill}"/></w:tcPr><w:p><w:r><w:t>{text}</w:t></w:r></w:p></w:tc>'
                                      for text, fill in row) + "</w:tr>" for row in rows)
    return parse_xml(f'<w:tbl xmlns:w="{W[1:-1]}"><w:tblGrid><w:gridCol w:w="4000"/><w:gridCol w:w="4000"/></w:tblGrid>{body}</w:tbl>')


def test_a_two_level_header_is_not_reused_as_a_data_row():
    table = _xml_table((("Synthetic Group", HEADER_FILL),), (("Name", HEADER_FILL), ("Detail", HEADER_FILL)))
    _table(table, ("Name", "Detail"), (("Synthetic A", "One"),))
    data = table.findall(W + "tr")[-1]
    assert "".join(node.text for node in data.iter(W + "t")) == "Synthetic AOne"
    assert not list(data.iter(W + "shd"))


def test_a_continuation_frame_still_copies_its_own_data_row():
    table = _xml_table((("Synthetic A", DATA_FILL), ("One", DATA_FILL)))
    _table(table, ("Name", "Detail"), (("Synthetic B", "Two"),), continuation=True)
    rows = table.findall(W + "tr")
    assert len(rows) == 1 and {node.get(W + "fill") for node in rows[0].iter(W + "shd")} == {DATA_FILL}
