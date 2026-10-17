"""Automatic logo reuse is distinct from explicit Starting deselection."""
from io import BytesIO
from hashlib import sha256
from dataclasses import replace

from docx import Document
from docx.shared import Inches
from PIL import Image
import pytest
from streamlit.testing.v1 import AppTest

from app import monthly_report_library as library
from app.monthly_report_import import inspect_docx, imported_draft
from app.monthly_report_model import ReportPeriod, synthetic_profiles
from app.monthly_report_start_import import analyze_starting_report, automatic_section_plans, build_starting_import
from app.monthly_report_start_replay import merge_starting_drafts, retain_automatic_logo_contributions
from tests.test_monthly_report_start_import import start_app, save_start


def source_with_header_logo(tmp_path):
    photo = BytesIO()
    Image.new("RGB", (600, 400), "navy").save(photo, "PNG")
    logo = BytesIO()
    Image.new("RGB", (450, 70), "green").save(logo, "PNG")
    document = Document()
    document.add_heading(synthetic_profiles()[0].facilities[0].title, 0)
    document.add_paragraph("September 2026")
    document.add_picture(BytesIO(photo.getvalue()), width=Inches(4))
    document.sections[0].header.paragraphs[0].add_run().add_picture(
        BytesIO(logo.getvalue()), width=Inches(2))
    document.add_heading("Monthly Activity Summary", 1)
    document.add_paragraph("Synthetic September inspection completed.")
    path = tmp_path / "invented-header-logo.docx"
    document.save(path)
    return path


def repeat_existing_app():
    return AppTest.from_string('''
from app import monthly_report_library as library
from app.monthly_report_section_ui import render_section_setup
from app.monthly_report_model import synthetic_profiles, ReportPeriod
from app.monthly_report_ui import _field
p = synthetic_profiles()[0]
state = library.load_profile(p.contract, p.key)
render_section_setup(p.contract, ReportPeriod(2026, 9), "Synthetic Editor", _field,
                     state=state, identity=state.profile)
''', default_timeout=30).run()


def test_same_bytes_no_photo_repeat_consumes_automatic_preferred_logo_skip(monkeypatch, tmp_path):
    source = source_with_header_logo(tmp_path)
    original = source.read_bytes()
    app = start_app(monkeypatch, tmp_path, source)
    next(button for button in app.button if button.label == "No cover photograph").click().run()
    first = save_start(app)
    original_receipt = library.starting_contribution(
        first.profile.contract, first.profile.key, first.period,
        sha256(original).hexdigest())
    assert original_receipt[1]
    assert next(block for block in first.blocks if block.key == "brand_logo").asset_hashes
    standing = library.load_profile(first.profile.contract, first.profile.key)
    manifest = library._profile_path(first.profile.contract, first.profile.key) / "manifest.json"
    before = manifest.read_bytes()
    repeat = repeat_existing_app()
    next(button for button in repeat.button if button.label == "No cover photograph").click().run()
    assert manifest.read_bytes() == before, "Opening/changing the unsaved review cannot publish anything."
    next(button for button in repeat.button if button.label == "Continue to this month’s updates").click().run()
    assert not repeat.exception
    assert not repeat.error
    repeated = library.load_imported_draft(first.profile.contract, first.profile.key)
    assert repeated == first
    assert library.starting_contribution(
        first.profile.contract, first.profile.key, first.period,
        sha256(original).hexdigest()) == original_receipt
    assert library.imported_original(first.profile.contract, first.profile.key).read_bytes() == original
    after = library.load_profile(first.profile.contract, first.profile.key)
    assert after.versions == standing.versions and after.current == standing.current


@pytest.mark.parametrize("change", ["human", "omitted", "cleared", "removed"])
def test_actual_repeat_keeps_snapshot_logo_edit_past_standing_and_branding_fallbacks(monkeypatch, tmp_path, change):
    from app import monthly_report_branding as branding
    source = source_with_header_logo(tmp_path)
    app = start_app(monkeypatch, tmp_path, source)
    next(button for button in app.button if button.label == "No cover photograph").click().run()
    initial = save_start(app)
    digest = sha256(source.read_bytes()).hexdigest()
    source_receipt = library.starting_contribution(
        initial.profile.contract, initial.profile.key, initial.period, digest)
    prior = next(block for block in initial.blocks if block.key == "brand_logo")
    stream = BytesIO()
    Image.new("RGB", (450, 70), "purple").save(stream, "PNG")
    raw = stream.getvalue()
    reference = library.asset_reference(raw, "png")
    current_logo = replace(prior, source="This month", asset_hashes=(reference,),
                           references=("manual:report",), asset_captions=(), asset_provenance=(),
                           client_asset_reviews=(), client_reviewed_fingerprint="", reviewed_fingerprint="")
    if change == "omitted":
        current_logo = replace(prior, source="Omit")
    elif change == "cleared":
        current_logo = replace(current_logo, asset_hashes=(), references=())
    current = replace(initial, blocks=tuple(
        current_logo if block.key == "brand_logo" else block for block in initial.blocks
        if not (change == "removed" and block.key == "brand_logo")))
    snapshot = library.save_snapshot(current, expected_revision=0, assets=((reference, raw),),
                                     entered_editor="Synthetic correction")
    # A newly available collection must not refill a consumed source's explicit
    # report Clear/removal. Other reports still retain ordinary logo defaults.
    brand = branding.BrandLogo("enfra", "Synthetic brand", (), reference,
        "https://example.invalid/brand", "https://example.invalid/logo.png", "2026-10-16")
    branding_before = branding.save_branding((brand,), {reference: raw}, expected_revision=0,
        actor="Synthetic branding editor", confirmed=True)
    standing = library.load_profile(initial.profile.contract, initial.profile.key)
    repeat = repeat_existing_app()
    next(button for button in repeat.button if button.label == "No cover photograph").click().run()
    next(button for button in repeat.button if button.label == "Continue to this month’s updates").click().run()
    assert not repeat.error and not repeat.exception
    result = library.load_imported_draft(initial.profile.contract, initial.profile.key)
    assert result == current
    assert library.load_snapshot(initial.profile.contract, initial.profile.key, initial.period) == snapshot
    after = library.load_profile(initial.profile.contract, initial.profile.key)
    assert after.versions == standing.versions and after.current == standing.current
    assert branding.load_branding() == branding_before
    assert library.starting_contribution(
        initial.profile.contract, initial.profile.key, initial.period, digest) == source_receipt


def source_context(tmp_path):
    source = source_with_header_logo(tmp_path)
    inspection = inspect_docx(source)
    profile, period = synthetic_profiles()[0], ReportPeriod(2026, 9)
    analysis = analyze_starting_report(inspection, profile, period)
    initial = automatic_section_plans(inspection, analysis)
    mapped, _ = build_starting_import(source, inspection, initial.values())
    previous = imported_draft(profile, period, "Synthetic Editor", mapped)
    plans = automatic_section_plans(inspection, analysis, preferred_logos=("brand_logo",))
    mapped, _ = build_starting_import(source, inspection, plans.values())
    incoming = imported_draft(profile, period, "Synthetic Editor", mapped)
    assert any(plan.get("automatic_logo_skips", {}).get("brand_logo") for plan in plans.values())
    assert not any(block.key == "brand_logo" for block in incoming.blocks)
    return source, inspection, previous, incoming, plans


def test_opening_actual_cover_card_preserves_preferred_logo_receipt_and_skip(tmp_path):
    from app.monthly_report_sections import section_reviews
    path, inspection, previous, incoming, plans = source_context(tmp_path)
    cover = next(section for section in section_reviews(inspection) if section.key == 'cover')
    old_skip = plans['cover']['automatic_logo_skips']
    app = AppTest.from_string('''
import streamlit as st
from app.monthly_report_section_ui import _section_card
from app.monthly_report_model import ReportPeriod
from app.monthly_report_ui import _field
if 'plans' in st.session_state:
 _section_card(st.session_state['section'], st.session_state['path'], st.session_state['inspection'],
               ReportPeriod(2026,9), 'synthetic_logo', _field, st.session_state['plans'], True)
''').run()
    for key,value in {'plans':plans,'section':cover,'path':path,'inspection':inspection}.items():
        app.session_state[key]=value
    app.run()
    assert not app.exception
    revised = app.session_state['plans']
    assert revised['cover'].get('automatic_logo_skips') == old_skip
    assert not set(old_skip['brand_logo']) & set(revised['cover']['selected'])
    retained = retain_automatic_logo_contributions(incoming, previous, verified=True,
        inspection=inspection, plans=revised.values(), preferred_logos=('brand_logo',))
    assert retained == previous
    assert merge_starting_drafts(previous, retained, previous) == previous


@pytest.mark.parametrize("change", ["human", "omitted", "cleared", "removed"])
def test_consumed_logo_receipt_preserves_current_report_changes(tmp_path, change):
    _, inspection, previous, incoming, plans = source_context(tmp_path)
    prior = next(block for block in previous.blocks if block.key == "brand_logo")
    corrected = replace(prior, asset_hashes=("c" * 64 + ".png",), references=("manual:report",))
    if change == "omitted":
        corrected = replace(corrected, source="Omit")
    elif change == "cleared":
        corrected = replace(prior, asset_hashes=(), references=())
    current = replace(previous, blocks=tuple(
        corrected if block.key == "brand_logo" else block for block in previous.blocks
        if not (change == "removed" and block.key == "brand_logo")))
    retained = retain_automatic_logo_contributions(
        incoming, previous, verified=True, inspection=inspection,
        plans=plans.values(), preferred_logos=("brand_logo",))
    assert retained == previous
    assert merge_starting_drafts(current, retained, previous) == current


@pytest.mark.parametrize("boundary", [
    "unverified", "missing-receipt", "profile", "month", "changed-bytes", "wrong-role",
    "foreign-reference", "missing-reference", "manual-review", "whole-section-omit",
    "selected-again", "not-preferred", "changed-plan-source", "no-prior-assets", "not-automatic",
])
def test_automatic_logo_carry_requires_exact_verified_source_role_and_scope(tmp_path, boundary):
    _, inspection, previous, incoming, plans = source_context(tmp_path)
    verified, preferred = True, ("brand_logo",)
    plan = next(plan for plan in plans.values() if plan.get("automatic_logo_skips"))
    item_id = plan["automatic_logo_skips"]["brand_logo"][0]
    if boundary == "unverified":
        verified = False
    elif boundary == "missing-receipt":
        previous = None
    elif boundary == "profile":
        previous = replace(previous, profile=replace(previous.profile, key="another-profile"))
    elif boundary == "month":
        previous = replace(previous, period=ReportPeriod(2026, 8))
    elif boundary == "changed-bytes":
        inspection = replace(inspection, sha256="d" * 64)
        plan["source_sha256"] = inspection.sha256
    elif boundary == "wrong-role":
        inspection = replace(inspection, items=tuple(
            replace(item, suggested_slot="cover_photo") if item.id == item_id else item
            for item in inspection.items))
    elif boundary in ("foreign-reference", "missing-reference", "no-prior-assets"):
        previous = replace(previous, blocks=tuple(
            replace(block, **({"references": ("docx:" + "d" * 64 + ":" + item_id,)}
                             if boundary == "foreign-reference" else {"references": ()}
                             if boundary == "missing-reference" else {"asset_hashes": ()}))
            if block.key == "brand_logo" else block for block in previous.blocks))
    elif boundary == "manual-review":
        plan.pop("automatic_logo_skips")
    elif boundary == "whole-section-omit":
        plan["omit"] = True
    elif boundary == "selected-again":
        plan["selected"].append(item_id)
    elif boundary == "not-preferred":
        preferred = ()
    elif boundary == "changed-plan-source":
        plan["source_sha256"] = "d" * 64
    elif boundary == "not-automatic":
        plan["automatic"] = False
    assert retain_automatic_logo_contributions(
        incoming, previous, verified=verified, inspection=inspection,
        plans=plans.values(), preferred_logos=preferred) == incoming


@pytest.mark.parametrize("manual", [False, True])
def test_explicit_logo_deselection_retains_normal_conservative_recovery(tmp_path, manual):
    _, inspection, previous, incoming, plans = source_context(tmp_path)
    for plan in plans.values():
        plan.pop("automatic_logo_skips", None)
        if manual:
            plan["automatic"] = False
    kept = retain_automatic_logo_contributions(
        incoming, previous, verified=True, inspection=inspection,
        plans=plans.values(), preferred_logos=("brand_logo",))
    with pytest.raises(ValueError, match="Brand logo is still in this report"):
        merge_starting_drafts(previous, kept, previous)
