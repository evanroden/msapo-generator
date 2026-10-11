"""Native RFI continuation is not another old-client heading in real PDF.

Only invented source rows and current report values are used; the private
authored reference is never copied into this test or any public artifact.
"""
from dataclasses import replace
from io import BytesIO

from docx import Document
import pymupdf

from conftest import requires_libreoffice
from app import pdf_converter
from app.monthly_report_import import inspect_docx
from app.monthly_report_master_tables import from_inspection
from app.monthly_report_model import (
    Facility, ReportDraft, ReportPeriod, ReportProfile, ReportTable,
    ResolvedBlock, default_sections,
)
from app.monthly_report_native_layout import build_native_docx


@requires_libreoffice
def test_two_native_rfi_frames_render_without_repeated_continuation_header(tmp_path, monkeypatch):
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path / "test-only-storage"))
    monkeypatch.setattr(pdf_converter, "OUTPUT_DIR", tmp_path)
    monkeypatch.setattr(pdf_converter, "PDF_BACKEND", "libreoffice")

    source = Document()
    source.add_heading("RFI MATRIX", 1)
    first = source.add_table(rows=3, cols=3)
    for cells, values in zip(first.rows, (
        ("#", "Requested Items", "Status"),
        ("14", "PREVIOUS CLIENT private note", "Closed"),
        ("15", "PREVIOUS CLIENT old file", "Closed"),
    )):
        for cell, value in zip(cells.cells, values):
            cell.text = value
    source.add_paragraph("Continuation")
    second = source.add_table(rows=3, cols=3)
    for cells, values in zip(second.rows, (
        ("16", "PREVIOUS CLIENT $77,777.77", "Closed"),
        ("17", "PREVIOUS CLIENT old contractor", "Closed"),
        ("18", "PREVIOUS CLIENT old reference", "Closed"),
    )):
        for cell, value in zip(cells.cells, values):
            cell.text = value
    path = tmp_path / "synthetic-prior-rfi.docx"
    source.save(path)

    shapes = from_inspection(inspect_docx(path))
    assert len(shapes) == 2 and shapes[1].continuation
    assert shapes[0].columns == shapes[1].columns == (
        "#", "Requested Items", "Status")
    sections = (replace(next(s for s in default_sections() if s.key == "rfi"),
                        included=True),)
    profile = ReportProfile("New Synthetic Contract", "new-rfi", "New Synthetic Site",
                            (Facility("new-rfi", "New Synthetic Site"),))
    current = (
        ReportTable(shapes[0].columns,
                    (("1", "Replace invented air filter", "Open"),), shapes[0].reference),
        ReportTable(shapes[1].columns,
                    (("2", "Inspect current pump", "Pending"),), shapes[1].reference),
    )
    draft = ReportDraft(
        profile, ReportPeriod(2026, 10), "Synthetic Current Editor",
        sections, (ResolvedBlock("rfi_matrix", "This month", extra_tables=current),),
        synthetic=True)
    raw = build_native_docx(draft, path, master=True)
    created = Document(BytesIO(raw))
    assert len(created.tables) == 2
    assert len(created.tables[0].rows) == 2  # heading + current data
    assert len(created.tables[1].rows) == 1  # current data, NO old header

    output = tmp_path / "synthetic-current-rfi.docx"
    output.write_bytes(raw)
    pdf_path = pdf_converter.convert_to_pdf(output)
    try:
        with pymupdf.open(pdf_path) as pdf:
            text = "\n".join(page.get_text() for page in pdf)
            assert "Replace invented air filter" in text
            assert "Inspect current pump" in text
            assert text.count("Requested Items") == 1
            assert text.count("Inspect current pump") == 1
            assert "PREVIOUS CLIENT" not in text
            assert "$77,777.77" not in text
    finally:
        pdf_path.unlink(missing_ok=True)
