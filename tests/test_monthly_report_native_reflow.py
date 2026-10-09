"""New monthly payloads must not inherit the previous report's empty pages."""
from dataclasses import replace
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

import pytest
from conftest import requires_libreoffice
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from PIL import Image, ImageDraw

from app.monthly_report_import import _heading, inspect_docx
from app.monthly_report_model import Facility, ReportDraft, ReportPeriod, ReportProfile, ResolvedBlock, default_sections
from app.monthly_report_native_layout import build_native_docx


def chart():
    image = Image.new('RGB', (1200, 675), 'white')
    draw = ImageDraw.Draw(image)
    draw.rectangle((0, 0, 1199, 674), outline='black', width=10)
    for xy in ((5, 5), (1150, 5), (5, 625), (1150, 625)):
        draw.rectangle((*xy, xy[0] + 40, xy[1] + 40), fill='black')
    draw.text((450, 320), 'SYNTHETIC CURRENT CHART', fill='black')
    data = BytesIO(); image.save(data, 'PNG')
    return data.getvalue()


def scenario(tmp_path):
    document = Document()
    document.add_paragraph('Synthetic old site')
    document.add_paragraph('August 2026')
    document.add_section(WD_SECTION_START.NEW_PAGE)
    document.add_heading('ORGANIZATIONAL CHART', 1)
    section = document.add_section(WD_SECTION_START.NEW_PAGE)
    section.page_width, section.page_height = Inches(11), Inches(8.5)
    section.left_margin = section.right_margin = Inches(.5)
    section.top_margin = section.bottom_margin = Inches(.5)
    section.header.paragraphs[0].text = 'ENFRA'
    document.add_paragraph().add_run().add_picture(BytesIO(chart()), width=Inches(3))
    document.add_page_break()
    document.add_paragraph('OLD PRIVATE CHART NOTE')
    section = document.add_section(WD_SECTION_START.NEW_PAGE)
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    document.add_heading('MONTHLY SCORECARDS', 1)
    document.add_section(WD_SECTION_START.NEW_PAGE)
    document.add_paragraph('OLD UTILITY VALUES')
    document.add_page_break()
    document.add_paragraph('RATE TREND')
    document.add_section(WD_SECTION_START.NEW_PAGE)
    document.add_heading('MBCX REPORTS', 1)
    document.add_section(WD_SECTION_START.NEW_PAGE)
    document.add_paragraph('The following section includes imported MBCx reports. The report from ENFRA Connect should be imported and included in the following pages.')
    for _ in range(12): document.add_paragraph()
    document.add_paragraph('OLD PRIVATE STATUS')
    document.add_page_break()
    document.add_paragraph()
    source = tmp_path / 'synthetic-source.docx'; document.save(source)
    profile = ReportProfile('Synthetic contract', 'north', 'Synthetic North', (Facility('north', 'Synthetic North'),))
    draft = ReportDraft(profile, ReportPeriod(2026, 10), 'Synthetic Editor', default_sections(), (
        ResolvedBlock('org_chart', 'This month', asset_hashes=('current.png',)),
        ResolvedBlock('mbcx_status', 'This month', text='Monitoring-based commissioning reporting has not started.'),
    ))
    return source, draft


@pytest.mark.parametrize('text', [
    'The following section includes imported MBCx reports. The report from ENFRA Connect should be imported and included in the following pages.',
    'Monthly scorecards will be available after the review.',
    'We reviewed the organizational chart and corrected a reporting line.',
    'Training summary for the team will follow next month.',
    'Accounts receivable will be reviewed next month.',
])
def test_section_name_in_narrative_is_not_a_native_heading(text):
    assert _heading(text) is None


@pytest.mark.parametrize('text,section', [
    ('Account Receivable Summary', 'accounts_receivable'), ('Appendix G: RFI Matrix', 'rfi'), ('MBCX REPORTS', 'mbcx'), ('SECTION 3: MONTHLY SCORECARDS', 'scorecards'),
    ('MAINTENANCE SCHEDULE / IN-HOUSE MAINTENANCE', 'maintenance'),
    ('ORGANIZATIONAL CHART ORGANIZATIONAL CHART 1 1', 'organization'),
])
def test_real_and_compatibility_branch_headings_remain_recognized(text, section):
    assert _heading(text) == section


def test_new_chart_uses_usable_landscape_width_without_crop(tmp_path):
    source, draft = scenario(tmp_path)
    result = Document(BytesIO(build_native_docx(draft, source, section_key='organization', master=True, asset_loader=lambda _: chart())))
    assert len(result.inline_shapes) == 1
    shape = result.inline_shapes[0]
    assert shape.width / 914400 == pytest.approx(10, abs=.05)
    assert shape.width / shape.height == pytest.approx(1200/675, abs=.001)
    assert not list(shape._inline.iter(qn('a:srcRect')))
    assert result.sections[-1].page_width == Inches(11)
    assert 'ENFRA' in result.sections[-1].header.paragraphs[0].text


def test_status_replaces_instructions_and_source_padding(tmp_path):
    source, draft = scenario(tmp_path)
    raw = build_native_docx(draft, source, section_key='mbcx', master=True)
    document = Document(BytesIO(raw))
    text = '\n'.join(p.text for p in document.paragraphs)
    assert 'should be imported' not in text
    assert 'has not started' in text
    assert 'OLD PRIVATE' not in text
    assert len([p for p in document.paragraphs if not p.text.strip()]) <= 3


def test_empty_scorecards_does_not_emit_unpopulated_template_pages(tmp_path):
    source, draft = scenario(tmp_path)
    raw = build_native_docx(draft, source, section_key='scorecards', master=True)
    document = Document(BytesIO(raw))
    assert len(document.sections) == 1
    assert 'MONTHLY SCORECARDS' in '\n'.join(p.text for p in document.paragraphs)
    assert not any(n.get(qn('w:type')) == 'page' for n in document.element.iter(qn('w:br')))


@requires_libreoffice
@pytest.mark.parametrize('section,pages', [('organization', 2), ('scorecards', 1), ('mbcx', 2)])
def test_rewritten_sections_render_without_blank_tail_pages(tmp_path, section, pages):
    from app.pdf_converter import convert_to_pdf
    import fitz
    source, draft = scenario(tmp_path)
    raw = build_native_docx(draft, source, section_key=section, master=True, asset_loader=lambda _: chart())
    path = tmp_path / (section + '.docx'); path.write_bytes(raw)
    with fitz.open(convert_to_pdf(path)) as pdf:
        assert len(pdf) == pages
        assert all(page.get_text().strip() or page.get_image_info() for page in pdf)
        if section == 'organization':
            image = max(pdf[-1].get_image_info(), key=lambda x: x['width'] * x['height'])
            x0, y0, x1, y1 = image['bbox']
            assert x1 - x0 == pytest.approx(720, abs=2)
            assert (x1-x0)/(y1-y0) == pytest.approx(1200/675, abs=.01)
            assert 0 <= x0 < x1 <= pdf[-1].rect.width and 0 <= y0 < y1 <= pdf[-1].rect.height
        if section == 'mbcx':
            block = next(b for b in pdf[-1].get_text('blocks') if 'has not started' in b[4])
            assert block[1] < 150
            assert 'should be imported' not in ''.join(p.get_text() for p in pdf)


def test_split_heading_ignores_unrelated_prior_fragment(tmp_path):
    doc = Document()
    doc.add_heading('MONTHLY ACTIVITY SUMMARY', 1)
    doc.add_paragraph('An unrelated prior paragraph')
    doc.add_paragraph('11PENDING & DECLINED')
    doc.add_paragraph('PROPOSALS')
    doc.add_paragraph('Current synthetic proposal')
    source = tmp_path / 'split-title.docx'; doc.save(source)
    inspection = inspect_docx(source)
    selected = [i for i in inspection.items if i.kind == 'text' and 'Current synthetic proposal' in i.text]
    assert selected[0].section == 'proposals'
    headings = [i for i in inspection.items if i.note == 'Native section heading']
    assert {i.text for i in headings if i.section == 'proposals'} == {'11PENDING & DECLINED', 'PROPOSALS'}
    assert not any('unrelated' in i.text for i in headings)


def test_additional_chart_pages_use_the_same_section_geometry(tmp_path):
    source, draft = scenario(tmp_path)
    charts = replace(draft.blocks[0], asset_hashes=('first.png', 'second.png', 'third.png'))
    draft = replace(draft, blocks=(charts, *draft.blocks[1:]))
    result = Document(BytesIO(build_native_docx(draft, source, section_key='organization', master=True, asset_loader=lambda _: chart())))
    assert len(result.inline_shapes) == 3
    assert all(shape.width / 914400 == pytest.approx(10, abs=.05) for shape in result.inline_shapes)


@requires_libreoffice
def test_multiple_technical_pages_have_no_overlap_or_trailing_blank(tmp_path):
    from app.pdf_converter import convert_to_pdf
    import fitz
    source, draft = scenario(tmp_path)
    charts = replace(draft.blocks[0], asset_hashes=('first.png', 'second.png', 'third.png'))
    draft = replace(draft, blocks=(charts, *draft.blocks[1:]))
    path = tmp_path / 'technical-pages.docx'
    path.write_bytes(build_native_docx(draft, source, section_key='organization', master=True, asset_loader=lambda _: chart()))
    pdf_path = convert_to_pdf(path)
    try:
        with fitz.open(pdf_path) as pdf:
            assert len(pdf) == 4
            for page in list(pdf)[1:]:
                images = page.get_image_info()
                assert len(images) == 1
                box = fitz.Rect(images[0]['bbox'])
                assert box.width == pytest.approx(720, abs=2)
                assert page.rect.contains(box)
                pix = page.get_pixmap(alpha=False)
                # All four synthetic corner markers survive the actual PDF.
                for x, y in ((.02,.03),(.975,.03),(.02,.95),(.975,.95)):
                    color = pix.pixel(round(box.x0 + box.width*x), round(box.y0 + box.height*y))
                    assert max(color) < 80
    finally:
        pdf_path.unlink(missing_ok=True)


def test_reflow_keeps_boundary_attached_to_current_text(tmp_path):
    from copy import deepcopy
    doc = Document()
    doc.add_heading('MONTHLY ACTIVITY SUMMARY', 1)
    note = doc.add_paragraph('Old activity')
    doc.add_section(WD_SECTION_START.NEW_PAGE)
    empty_boundary = doc.paragraphs[-1]._p
    boundary = empty_boundary.find('.//' + qn('w:sectPr'))
    note._p.get_or_add_pPr().append(deepcopy(boundary))
    empty_boundary.getparent().remove(empty_boundary)
    doc.add_heading('TRAINING SUMMARY', 1)
    doc.add_paragraph('Old training')
    source = tmp_path / 'attached-boundary.docx'; doc.save(source)
    profile = ReportProfile('Synthetic contract', 'north', 'Synthetic North', (Facility('north', 'Synthetic North'),))
    draft = ReportDraft(profile, ReportPeriod(2026, 10), 'Synthetic Editor', default_sections(), (
        ResolvedBlock('activity_summary', 'This month', text='Current activity'),
        ResolvedBlock('training_summary', 'This month', text='Current training'),
    ))
    result = Document(BytesIO(build_native_docx(draft, source, master=True)))
    assert len(result.sections) == 2
    assert 'Current activity' in '\n'.join(p.text for p in result.paragraphs)
    assert 'Current training' in '\n'.join(p.text for p in result.paragraphs)


def test_chart_can_be_added_when_the_master_has_no_picture_frame(tmp_path):
    doc = Document()
    doc.add_paragraph('Synthetic source site')
    doc.add_section(WD_SECTION_START.NEW_PAGE)
    doc.add_heading('ORGANIZATIONAL CHART', 1)
    doc.add_section(WD_SECTION_START.NEW_PAGE)
    doc.add_paragraph('Old organizational text')
    path = tmp_path / 'no-pictures.docx'; doc.save(path)
    profile = ReportProfile('Synthetic contract', 'north', 'Synthetic North', (Facility('north', 'Synthetic North'),))
    draft = ReportDraft(profile, ReportPeriod(2026, 10), 'Synthetic Editor', default_sections(),
                        (ResolvedBlock('org_chart', 'This month', asset_hashes=('chart.png',)),))
    result = Document(BytesIO(build_native_docx(draft, path, section_key='organization', master=True, asset_loader=lambda _: chart())))
    assert len(result.inline_shapes) == 1
    assert result.inline_shapes[0].width / result.inline_shapes[0].height == pytest.approx(1200 / 675, abs=.001)
    assert not any('Old organizational' in p.text for p in result.paragraphs)
