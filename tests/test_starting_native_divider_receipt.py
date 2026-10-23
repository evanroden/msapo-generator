"""A native-owned divider remains an acknowledged source contribution on replay."""
from dataclasses import replace
from io import BytesIO
from docx import Document
from docx.shared import Inches
from PIL import Image
import pytest

from app import monthly_report_start_replay as replay
from app.monthly_report_import import inspect_docx, imported_draft
from app.monthly_report_model import ReportPeriod, ResolvedBlock, ReportTable, synthetic_profiles
from app.monthly_report_start_import import (
    analyze_starting_report, automatic_section_plans, build_starting_import,
)


def context(tmp_path):
    picture = BytesIO()
    Image.new("RGB", (600, 800), "green").save(picture, "PNG")
    document = Document()
    document.add_heading(synthetic_profiles()[0].facilities[0].title, 0)
    document.add_paragraph("September 2026")
    heading = document.add_heading("1 Organizational Chart", 1)
    heading.add_run().add_picture(BytesIO(picture.getvalue()), width=Inches(4))
    document.add_heading("2 Monthly Activity Summary", 1)
    document.add_paragraph("Synthetic September inspection remains under review.")
    path = tmp_path / "invented-divider-report.docx"
    document.save(path)
    inspection = inspect_docx(path)
    profile, period = synthetic_profiles()[0], ReportPeriod(2026, 9)
    analysis = analyze_starting_report(inspection, profile, period)
    original_plans = automatic_section_plans(inspection, analysis)
    mapped, _ = build_starting_import(path, inspection, original_plans.values())
    previous = imported_draft(profile, period, "Synthetic Editor", mapped)
    plans = automatic_section_plans(inspection, analysis, native_design=True)
    mapped, _ = build_starting_import(path, inspection, plans.values())
    incoming = imported_draft(profile, period, "Synthetic Editor", mapped)
    assert any(block.key == "divider_organization" for block in previous.blocks)
    assert not any(block.key == "divider_organization" for block in incoming.blocks)
    return inspection, previous, incoming, plans


def retain(inspection, previous, incoming, plans, *, verified=True, native_design=True):
    return replay.retain_automatic_divider_contributions(
        incoming, previous, verified=verified, inspection=inspection,
        plans=plans.values(), native_design=native_design)


@pytest.mark.parametrize("change", ["unchanged", "human", "omitted", "cleared", "removed"])
def test_native_divider_skip_preserves_report_edits(tmp_path, change):
    inspection, previous, incoming, plans = context(tmp_path)
    prior = next(block for block in previous.blocks if block.key == "divider_organization")
    corrected = prior
    if change == "human":
        corrected = replace(prior, asset_hashes=("c" * 64 + ".png",), references=("manual:report",))
    elif change == "omitted":
        corrected = replace(prior, source="Omit")
    elif change == "cleared":
        corrected = replace(prior, asset_hashes=(), references=())
    current = replace(previous, blocks=tuple(
        corrected if block.key == prior.key else block for block in previous.blocks
        if not (change == "removed" and block.key == prior.key)))
    kept = retain(inspection, previous, incoming, plans)
    assert kept == previous
    assert replay.merge_starting_drafts(current, kept, previous) == current


@pytest.mark.parametrize("boundary", [
    "unverified", "missing-receipt", "contract", "month", "changed-bytes",
    "wrong-role", "wrong-kind", "unknown-divider", "foreign-reference",
    "missing-reference", "manual-review", "whole-section-omit",
    "selected-again", "not-native", "changed-plan-source",
    "no-prior-assets", "not-automatic", "other-site", "unknown-site",
    "no-prior-divider", "prior-omitted",
])
def test_divider_requires_exact_verified_same_source_receipt(tmp_path, boundary):
    inspection, previous, incoming, plans = context(tmp_path)
    verified, native_design = True, True
    plan = plans["organization"]
    item_id = next(item.id for item in inspection.items
                   if item.suggested_slot == "divider_organization")
    if boundary == "unverified":
        verified = False
    elif boundary == "missing-receipt":
        previous = None
    elif boundary == "contract":
        previous = replace(previous, profile=replace(previous.profile, contract="another-synthetic-scope"))
    elif boundary == "month":
        previous = replace(previous, period=ReportPeriod(2026, 8))
    elif boundary == "changed-bytes":
        inspection = replace(inspection, sha256="d" * 64)
        plan["source_sha256"] = inspection.sha256
    elif boundary in ("wrong-role", "wrong-kind"):
        inspection = replace(inspection, items=tuple(replace(item, **(
            {"suggested_slot": "org_chart"} if boundary == "wrong-role" else {"kind": "text"}))
            if item.id == item_id else item for item in inspection.items))
    elif boundary == "unknown-divider":
        inspection = replace(inspection, items=tuple(
            replace(item, suggested_slot="divider_unknown") if item.id == item_id else item
            for item in inspection.items))
        previous = replace(previous, blocks=tuple(
            replace(block, key="divider_unknown") if block.key == "divider_organization" else block
            for block in previous.blocks))
        plan["automatic_divider_skips"] = {"divider_unknown": [item_id]}
    elif boundary in ("foreign-reference", "missing-reference", "no-prior-assets", "prior-omitted"):
        fields = {"references": ("docx:" + "d" * 64 + ":" + item_id,)}
        if boundary == "missing-reference":
            fields = {"references": ()}
        elif boundary == "no-prior-assets":
            fields = {"asset_hashes": ()}
        elif boundary == "prior-omitted":
            fields = {"source": "Omit"}
        previous = replace(previous, blocks=tuple(
            replace(block, **fields) if block.key == "divider_organization" else block
            for block in previous.blocks))
    elif boundary == "manual-review":
        plan.pop("automatic_divider_skips", None)
    elif boundary == "whole-section-omit":
        plan["omit"] = True
    elif boundary == "selected-again":
        plan["selected"].append(item_id)
    elif boundary == "not-native":
        native_design = False
    elif boundary == "changed-plan-source":
        plan["source_sha256"] = "d" * 64
    elif boundary == "not-automatic":
        plan["automatic"] = False
    elif boundary in ("other-site", "unknown-site"):
        plan["source_site_relation"] = boundary.split("-")[0]
    elif boundary == "no-prior-divider":
        previous = replace(previous, blocks=tuple(
            block for block in previous.blocks if block.key != "divider_organization"))
    assert retain(inspection, previous, incoming, plans, verified=verified,
                  native_design=native_design) == incoming


@pytest.mark.parametrize("manual", [False, True])
def test_explicit_deselection_still_blocks(tmp_path, manual):
    inspection, previous, incoming, plans = context(tmp_path)
    plans["organization"].pop("automatic_divider_skips", None)
    if manual:
        plans["organization"]["automatic"] = False
    kept = retain(inspection, previous, incoming, plans)
    with pytest.raises(ValueError, match="Divider organization is still in this report"):
        replay.merge_starting_drafts(previous, kept, previous)


def test_divider_retention_allows_unchanged_capacity_source_correction(tmp_path):
    inspection, previous, incoming, plans = context(tmp_path)
    old_table = ReportTable(("Chiller", "Capacity"), (("Synthetic A", "250"),))
    capacity = ResolvedBlock("thermal_capacity", "Last month", extra_tables=(old_table,))
    previous = replace(previous, blocks=(*previous.blocks, capacity))
    new_table = ReportTable(("Chiller", "Capacity", "Estimated Peak Coincident Load"),
                            (("Synthetic A", "250", "180"),))
    corrected = replace(capacity, extra_tables=(new_table,))
    incoming = replace(incoming, blocks=(*incoming.blocks, corrected))
    current = replace(previous, blocks=tuple(
        replace(block, text="Current manual activity.")
        if block.key == "activity_summary" else block for block in previous.blocks)
        + (ResolvedBlock("training_summary", "This month", text="Current training note."),))
    kept = retain(inspection, previous, incoming, plans)
    expected = replace(current, blocks=tuple(
        corrected if block.key == "thermal_capacity" else block for block in current.blocks))
    assert replay.merge_starting_drafts(current, kept, previous) == expected
