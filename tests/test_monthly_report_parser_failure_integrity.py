"""A malformed source fails by file while successful siblings remain usable.

PNG and damaged OOXML fixtures are real, bounded bytes. The PDF iterator,
workbook schema and image-frame tests separately model native decoder errors;
they are not presented as naturally malformed PDF/TIFF reproductions.
"""
from dataclasses import replace
from hashlib import sha256
from io import BytesIO
from pathlib import Path
import struct
from zipfile import ZipFile, ZIP_DEFLATED
import zlib

from docx import Document
from openpyxl import Workbook
from PIL import Image
import pytest

from app import monthly_report_docx as renderer
from app import monthly_report_library as library
from app import monthly_report_sources as sources
from app.monthly_report_model import ReportPeriod, synthetic_draft, synthetic_profiles


FIRST = b'Technician inspected primary pump P-61; bearing temperature was 76 F.'
LAST = b'Technician verified reserve fan F-82 remained locked out.'


@pytest.fixture
def profile(monkeypatch, tmp_path):
    monkeypatch.setenv('EPC_DATA_DIR', str(tmp_path))
    return synthetic_profiles()[0]


def oversized_png_header():
    # Only 65 bytes, with valid CRCs. No raster is ever created or decoded.
    def chunk(name, data):
        return (struct.pack('>I', len(data)) + name + data
                + struct.pack('>I', zlib.crc32(name + data) & 0xffffffff))
    return (b'\x89PNG\r\n\x1a\n'
            + chunk(b'IHDR', struct.pack('>IIBBBBB', 19000, 21000, 8, 2, 0, 0, 0))
            + chunk(b'IDAT', zlib.compress(b'')) + chunk(b'IEND', b''))


def workbook_bytes():
    workbook = Workbook()
    workbook.active.append(('Equipment', 'Finding'))
    workbook.active.append(('P-64', 'Discharge valve did not close'))
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()


def corrupted_package(suffix):
    if suffix == '.xlsx':
        original = workbook_bytes()
    else:
        document = Document()
        document.add_paragraph('Inspected synthetic pump P-69.')
        stream = BytesIO()
        document.save(stream)
        original = stream.getvalue()
    stream = BytesIO()
    with ZipFile(BytesIO(original)) as old, ZipFile(stream, 'w', ZIP_DEFLATED) as new:
        for name in old.namelist():
            new.writestr(name, old.read(name))
    result = bytearray(stream.getvalue())
    entry = '[Content_Types].xml' if suffix == '.xlsx' else 'word/document.xml'
    with ZipFile(BytesIO(result)) as package:
        info = package.getinfo(entry)
        assert info.compress_type == ZIP_DEFLATED
        offset = info.header_offset
        assert result[offset:offset + 4] == b'PK\x03\x04'
        name_size, extra_size = struct.unpack_from('<HH', result, offset + 26)
        data_offset = offset + 30 + name_size + extra_size
        # Reserved DEFLATE block type; directory/expanded sizes stay bounded.
        result[data_offset] = (result[data_offset] & ~6) | 6
    with ZipFile(BytesIO(result)) as package, pytest.raises(zlib.error):
        package.read(entry)
    return bytes(result)


def missing_workbook_part():
    stream = BytesIO()
    with ZipFile(stream, 'w') as archive:
        archive.writestr('synthetic-note.txt', 'Required workbook parts are absent.')
    return stream.getvalue()


def require_siblings(profile, filename, bad, *, expected_message):
    contents, notices = sources.ingest_batch(profile, (
        ('first.txt', FIRST), (filename, bad), ('last.txt', LAST)))
    assert tuple(content.source.filename for content in contents) == ('first.txt', 'last.txt')
    assert tuple(sources.source_bytes(profile, content.source) for content in contents) == (FIRST, LAST)
    assert len(notices) == 1 and notices[0].startswith('Not read: ' + Path(filename).name + '.')
    assert expected_message in notices[0]
    assert '/private/' not in notices[0]
    for content, original in zip(contents, (FIRST, LAST)):
        assert content.source.id == sha256(original).hexdigest()
        assert sources._path(profile, content.source.id, '.json').exists()
    return contents, notices


def test_png_bomb_has_recoverable_image_validation_error_without_decoding(monkeypatch):
    raw = oversized_png_header()
    assert len(raw) < 100
    monkeypatch.setattr(Image.Image, 'load', lambda *a, **k: pytest.fail('Oversized pixels were decoded.'))
    with pytest.raises(ValueError, match='too large to decode safely'):
        renderer.normalize_report_image(raw, '.png')


def test_actual_png_header_rejection_keeps_both_siblings_and_durable_source_inventory(profile):
    raw = oversized_png_header()
    contents, _ = require_siblings(profile, 'large-dimensions.png', raw,
                                  expected_message='too large to decode safely')
    draft = replace(synthetic_draft(profile, ReportPeriod(2026, 10)),
                    sources=tuple(content.source for content in contents))
    saved = library.save_snapshot(draft, expected_revision=0, entered_editor='Synthetic Parser Reviewer')
    readback = library.load_snapshot(profile.contract, profile.key, draft.period)
    assert readback.revision == saved.revision and readback.draft.sources == draft.sources
    assert not sources._path(profile, sha256(raw).hexdigest(), '.png').exists()
    assert not sources._path(profile, sha256(raw).hexdigest(), '.json').exists()


@pytest.mark.parametrize('suffix', ['.xlsx', '.docx'])
def test_real_corrupt_deflate_package_does_not_abandon_valid_siblings(profile, suffix):
    raw = corrupted_package(suffix)
    assert len(raw) < 50000
    contents, _ = require_siblings(profile, 'damaged-compression' + suffix, raw,
                                  expected_message='Compressed file data is damaged')
    assert len(contents) == 2
    assert not sources._path(profile, sha256(raw).hexdigest(), suffix).exists()
    assert not sources._path(profile, sha256(raw).hexdigest(), '.json').exists()


def test_missing_workbook_part_stays_contained_with_truthful_recovery(profile):
    _, notices = require_siblings(profile, 'missing-part.xlsx', missing_workbook_part(),
                                 expected_message='missing required document data')
    assert 'encoding' not in notices[0]


def test_modeled_pdf_page_tree_error_is_contained_during_enumeration(profile, monkeypatch):
    class BrokenPageTree:
        is_encrypted = False
        def __enter__(self): return self
        def __exit__(self, *args): return False
        def __len__(self): return 1
        def __iter__(self):
            raise sources.fitz.mupdf.FzErrorFormat('/private/server/model: malformed page tree')
    monkeypatch.setattr(sources.fitz, 'open', lambda *a, **k: BrokenPageTree())
    require_siblings(profile, 'modeled-page-tree.pdf', b'%PDF-SYNTHETIC-ITERATOR-ERROR',
                     expected_message='PDF pages could not be decoded')


def test_modeled_workbook_schema_type_error_gets_file_specific_recovery(profile, monkeypatch):
    def broken_workbook(*args, **kwargs):
        raise TypeError('/private/server/model: invalid workbook property')
    monkeypatch.setattr(sources, 'load_workbook', broken_workbook)
    require_siblings(profile, 'modeled-workbook-schema.xlsx', workbook_bytes(),
                     expected_message='Workbook structure could not be read')


def test_modeled_image_frame_type_error_is_a_recoverable_image_validation_error(monkeypatch):
    class BrokenFrameMetadata:
        width = 64
        height = 32
        def __enter__(self): return self
        def __exit__(self, *args): return False
        @property
        def n_frames(self):
            raise TypeError('/private/server/model: Missing dimensions')
    monkeypatch.setattr(Image, 'open', lambda *a, **k: BrokenFrameMetadata())
    with pytest.raises(ValueError, match='Image frame information could not be read'):
        renderer.normalize_report_image(b'SYNTHETIC-FRAME-ERROR', '.tiff')


def test_unexpected_product_type_error_is_not_hidden_as_a_file_failure(profile, monkeypatch):
    def programming_error(*args, **kwargs):
        raise TypeError('Synthetic programming fault outside the decoder boundary')
    monkeypatch.setattr(sources, 'suggest_facts', programming_error)
    with pytest.raises(TypeError, match='programming fault outside'):
        sources.ingest_batch(profile, (('valid.txt', FIRST),))


def test_preexisting_original_and_other_scope_survive_rejection(profile, monkeypatch, tmp_path):
    raw = oversized_png_header()
    digest = sha256(raw).hexdigest()
    previous = sources._path(profile, digest, '.png')
    library._atomic_write(previous, raw)
    other_profile = replace(profile, key='independent-synthetic-profile')
    independent, _ = sources.ingest(other_profile, 'independent.txt', b'Independent accepted work remains.')
    before = previous.stat()
    require_siblings(profile, 'same-original.png', raw, expected_message='too large to decode safely')
    assert previous.read_bytes() == raw and previous.stat().st_ino == before.st_ino
    assert sources.source_bytes(other_profile, independent.source) == b'Independent accepted work remains.'


def test_valid_small_image_and_workbook_are_processed_and_reimport_is_idempotent(profile):
    stream = BytesIO()
    Image.new('RGB', (96, 64), (35, 110, 82)).save(stream, 'PNG')
    originals = (('small.png', stream.getvalue()), ('small.xlsx', workbook_bytes()))
    contents, notices = sources.ingest_batch(profile, originals)
    assert not notices and len(contents) == 2
    assert (contents[1].tables[0].columns, contents[1].tables[0].rows) == (
        ('Equipment', 'Finding'), (('P-64', 'Discharge valve did not close'),))
    image = sources.page_image(profile, contents[0], 1, preview=True)
    assert (image.width, image.height) == (96, 64)
    again, notices = sources.ingest_batch(profile, originals, contents)
    assert again == contents and all(notice.startswith('Duplicate skipped:') for notice in notices)
    assert tuple(sources.source_bytes(profile, content.source) for content in contents) == tuple(raw for _, raw in originals)


def test_guided_bad_image_batch_keeps_delivered_activity_and_processed_siblings(monkeypatch, tmp_path):
    import streamlit as st
    from app import monthly_report_active_work as active_work
    from test_monthly_report_activity_intake_commit import setup_app, editor, draft, activity, TOKEN
    app, profile = setup_app(monkeypatch, tmp_path)
    originals = (('first.txt', FIRST), ('large-dimensions.png', oversized_png_header()), ('last.txt', LAST))
    original_uploader = st.file_uploader
    def upload(label, *args, **kwargs):
        if label == 'Add activity files (optional)':
            values = []
            for name, raw in originals:
                stream = BytesIO(raw)
                stream.name = name
                values.append(stream)
            return values
        return original_uploader(label, *args, **kwargs)
    monkeypatch.setattr(st, 'file_uploader', upload)
    manual = 'Human edit delivered together with this batch remains unchanged.'
    editor(app).set_value(manual).run()
    assert not app.exception
    expected_ids = {sha256(raw).hexdigest() for raw in (FIRST, LAST)}
    assert {source.id for source in draft(app).sources} == expected_ids, 'Both valid uploads must have processed.'
    native_ids = set().union(*(set(value) for key, value in app.session_state.filtered_state.items()
                            if key.endswith('_native_ids')))
    assert expected_ids <= native_ids
    assert manual in activity(app).text
    assert FIRST.decode() in activity(app).text and LAST.decode() in activity(app).text
    notices = '\n'.join(item.value for item in (*app.warning, *app.info, *app.error))
    assert 'large-dimensions.png' in notices and 'too large to decode safely' in notices
    journal = active_work.load(TOKEN, profile.contract, profile.key, ReportPeriod(2026, 10))
    assert journal and next(block for block in journal.draft.blocks if block.key == 'activity_summary') == activity(app)
