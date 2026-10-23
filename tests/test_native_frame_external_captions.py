"""External captions survive frame-local evidence isolation."""
from io import BytesIO
from docx import Document
from docx.oxml.ns import qn
import pytest
from app.monthly_report_import import inspect_docx
from app.monthly_report_start_import import cover_photo_candidates
from test_monthly_report_native_design_import import frame
from test_monthly_report_source_frame_logo_role import raster, _picture_in_frame


def make_source(tmp_path, *, caption='Company logo', inner='Unrelated note', placement='same', framed_picture=False):
    doc = Document()
    p = doc.add_paragraph()
    p.add_run()._r.append(frame(inner))
    if framed_picture:
        _picture_in_frame(doc, p._p, raster('green'))
    if placement == 'same':
        p.add_run(caption)
        p.add_run().add_picture(BytesIO(raster('navy')))
    elif placement == 'after':
        p.add_run(caption)
        doc.add_picture(BytesIO(raster('navy')))
    elif placement == 'before':
        p.add_run(caption)
        picture = doc.add_paragraph()
        picture.add_run().add_picture(BytesIO(raster('navy')))
        p._p.addprevious(picture._p)
    path = tmp_path / 'caption.docx'
    doc.save(path)
    return path


@pytest.mark.parametrize('placement', ['same', 'before', 'after'])
@pytest.mark.parametrize('caption,role', [('Company logo', 'brand_logo'), ('Client logo', 'client_logo'), ('Logo', 'logo')])
def test_external_caption_beside_unrelated_frame_is_not_discarded(tmp_path, placement, caption, role):
    path = make_source(tmp_path, placement=placement, caption=caption)
    inspection = inspect_docx(path)
    images = [x for x in inspection.items if x.kind == 'image']
    assert len(images) == 1
    assert images[0].image_role == role
    assert not cover_photo_candidates(inspection.items)


@pytest.mark.parametrize('placement', ['same', 'before', 'after'])
def test_inner_label_does_not_escape_frame_to_external_picture(tmp_path, placement):
    path = make_source(tmp_path, placement=placement, caption='', inner='INSERT CLIENT LOGO')
    images = [x for x in inspect_docx(path).items if x.kind == 'image']
    assert len(images) == 1
    assert images[0].image_role == ''


def test_external_company_caption_does_not_taint_closed_client_frame(tmp_path):
    path = make_source(tmp_path, inner='INSERT CLIENT LOGO', framed_picture=True)
    images = [x for x in inspect_docx(path).items if x.kind == 'image']
    assert len(images) == 2
    assert images[0].image_role == 'client_logo'
    assert images[1].image_role == 'brand_logo'


@pytest.mark.parametrize('tag', ['shape', 'rect'])
def test_external_caption_survives_unrelated_legacy_frame(tmp_path, tag):
    from lxml import etree as ET
    doc = Document()
    p = doc.add_paragraph('Company logo')
    picture = ET.SubElement(p.add_run()._r, qn('w:pict'))
    shape = ET.SubElement(picture, '{urn:schemas-microsoft-com:vml}' + tag)
    box = ET.SubElement(ET.SubElement(shape, '{urn:schemas-microsoft-com:vml}textbox'), qn('w:txbxContent'))
    ET.SubElement(ET.SubElement(ET.SubElement(box, qn('w:p')), qn('w:r')), qn('w:t')).text = 'Client logo'
    p.add_run().add_picture(BytesIO(raster()))
    path = tmp_path / 'legacy-frame.docx'
    doc.save(path)
    images = [x for x in inspect_docx(path).items if x.kind == 'image']
    assert len(images) == 1
    assert images[0].image_role == 'brand_logo'


def test_deleted_external_caption_does_not_supply_role(tmp_path):
    from lxml import etree as ET
    path = make_source(tmp_path, caption='Company logo', inner='INSERT CLIENT LOGO')
    doc = Document(path)
    p = doc.paragraphs[0]._p
    caption = next(run for run in p.findall(qn('w:r')) if run.find(qn('w:t')) is not None)
    p.remove(caption)
    deletion = ET.SubElement(p, qn('w:del'), {qn('w:id'): '1'})
    text = caption.find(qn('w:t'))
    text.tag = qn('w:delText')
    deletion.append(caption)
    doc.save(path)
    images = [x for x in inspect_docx(path).items if x.kind == 'image']
    assert images[0].image_role == ''


def test_external_caption_multiple_ordinary_images_remains_ambiguous(tmp_path):
    path = make_source(tmp_path)
    doc = Document(path)
    doc.paragraphs[0].add_run().add_picture(BytesIO(raster('red')))
    doc.save(path)
    images = [x for x in inspect_docx(path).items if x.kind == 'image']
    assert len(images) == 2
    assert {x.image_role for x in images} == {'logo'}
    assert {x.suggested_slot for x in images} == {''}


def test_unrecognized_frame_inner_label_stays_local_but_external_caption_survives(tmp_path):
    from lxml import etree as ET
    path = make_source(tmp_path, inner='INSERT CLIENT LOGO', framed_picture=True)
    doc = Document(path)
    shape = next(doc.element.iter('{http://schemas.microsoft.com/office/word/2010/wordprocessingShape}wsp'))
    ET.SubElement(shape, '{urn:synthetic-unknown}payload')
    doc.save(path)
    images = [x for x in inspect_docx(path).items if x.kind == 'image']
    assert [x.image_role for x in images] == ['', 'brand_logo']


@pytest.mark.parametrize('placement', ['same', 'before', 'after'])
def test_ordinary_vml_picture_keeps_external_caption(tmp_path, placement):
    from lxml import etree as ET
    doc = Document()
    p = doc.add_paragraph()
    rid, _ = doc.part.get_or_add_image(BytesIO(raster()))
    shape = ET.SubElement(ET.SubElement(p.add_run()._r, qn('w:pict')), '{urn:schemas-microsoft-com:vml}shape')
    ET.SubElement(shape, '{urn:schemas-microsoft-com:vml}imagedata', {qn('r:id'): rid})
    if placement == 'same':
        p.add_run('Client logo')
    else:
        caption = doc.add_paragraph('Client logo')
        if placement == 'before':
            p._p.addprevious(caption._p)
    path = tmp_path / 'ordinary-vml.docx'
    doc.save(path)
    images = [x for x in inspect_docx(path).items if x.kind == 'image']
    assert len(images) == 1
    assert images[0].image_role == 'client_logo'
