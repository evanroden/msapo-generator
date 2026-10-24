"""Refs #276: an ASCII tilde does not make an explicit price publishable.

All wording is invented. Exercise public client-output and source-page gates,
while keeping the original wording available for correction. Native controls,
saved report history and converted PDF acceptance are covered separately.
"""

from dataclasses import replace
from io import BytesIO

from docx import Document
import pytest

from app.monthly_report_checks import preflight
from app.monthly_report_content_policy import page_allowed, page_fingerprint
from app.monthly_report_docx import assemble_docx
from app.monthly_report_guided import client_table_draft
from app.monthly_report_model import (
    ReportDraft, ReportPeriod, ReportSource, ReportTable, ResolvedBlock,
    default_sections, synthetic_profiles,
)


@pytest.fixture(autouse=True)
def isolated_report_data(monkeypatch, tmp_path):
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path / "data"))


def _draft(text="", *, table=None):
    section, key = ("capital", "capital_renewal") if table else ("activity", "activity_summary")
    sections = tuple(replace(item, included=item.key == section) for item in default_sections())
    block = ResolvedBlock(key, "This month", text=text,
                          extra_tables=(table,) if table else ())
    return ReportDraft(synthetic_profiles()[0], ReportPeriod(2026, 11),
                       "Synthetic Price Reviewer", sections, (block,))


@pytest.mark.parametrize("text", (
    pytest.param("Repair estimate: ~3.2k", id="compact-amount"),
    pytest.param("Cost: ~ 3200", id="spaced-amount"),
    pytest.param("Quote received from Synthetic Vendor: ~3.2k", id="vendor-assignment"),
    pytest.param("Quote for valve replacement is ~ 3200", id="scope-assignment"),
    pytest.param("Repair estimate: ~3000-4000", id="approximate-range-start"),
    pytest.param("Quote: ~3k to ~4k", id="approximate-range-both-ends"),
    pytest.param("Cost\n~ 3200", id="extracted-heading-and-amount"),
    pytest.param("~3.2k cost", id="amount-before-label"),
    pytest.param("USD ~3.2k", id="currency-code"),
    pytest.param("$ ~3200", id="currency-symbol"),
    pytest.param("Budget: ~2026", id="year-shaped-budget-amount"),
    pytest.param("Budget for ~2026: ~3.2k", id="approximate-year-and-budget-amount"),
))
def test_approximate_prices_stay_editable_but_block_preflight_and_word(text):
    original = _draft(text)
    working, removed = client_table_draft(original)
    assert working == original and not removed
    assert working.blocks[0].text == text

    checks = preflight(working)
    assert not [check for check in checks if check.blocking and check.code != "pricing"]
    priced = any(check.code == "pricing" and check.blocking
                 and check.block_key == "activity_summary" for check in checks)
    try:
        raw = assemble_docx(working, acknowledged_fingerprint=working.fingerprint)
    except ValueError as exc:
        blocked = "pric" in str(exc).lower()
    else:
        # On the unfixed policy, verify that generation actually includes the
        # offending wording; an unrelated empty-output failure is no proof.
        actual = "\n".join(p.text for p in Document(BytesIO(raw)).paragraphs)
        assert text in actual
        blocked = False
    assert (priced, blocked) == (True, True), {
        "text": text, "pricing_preflight": priced, "word_blocked": blocked,
    }
    assert original.blocks[0].text == text


@pytest.mark.parametrize("text", (
    pytest.param("Estimate: ~3.2 kW", id="power"),
    pytest.param("Estimate: ~3200 kWh", id="energy"),
    pytest.param("Estimate: ~3.2kW", id="compact-power"),
    pytest.param("Estimate: ~3.2kWh", id="compact-energy"),
    pytest.param("Not to exceed ~125 psi", id="pressure"),
    pytest.param("Estimate: ~125–~150 psi", id="approximate-pressure-range"),
    pytest.param("Not to exceed ~12 A", id="current"),
    pytest.param("Repair estimate: ~48 hours", id="duration"),
    pytest.param("Utility costs decreased ~8%", id="percentage"),
    pytest.param("Quote date: 2026-11-14", id="quote-date"),
    pytest.param("Quote date: ~2026-11-14", id="approximate-quote-date"),
    pytest.param("Quote received from Synthetic Vendor: 2026-11-14", id="vendor-date"),
    pytest.param("Quotation reference: Q-TEST-3200", id="quote-reference"),
    pytest.param("Number of Quotes: ~3", id="reference-count"),
    pytest.param("~3.2k readings recorded.", id="unlabeled-reading-count"),
    pytest.param("~ 3200", id="unlabeled-value"),
    pytest.param("Quote is pending. ~3200 operating hours were recorded.", id="sentence-boundary"),
    pytest.param("Cost estimate for Chiller 2 remains pending.", id="asset-number"),
    pytest.param("Budget for ~2026 remains pending.", id="approximate-budget-year"),
    pytest.param("Budget for ~2026-~2027 remains pending.", id="approximate-budget-years"),
))
def test_approximate_measurements_and_reference_metadata_reach_word_unchanged(text):
    original = _draft(text)
    working, removed = client_table_draft(original)
    assert working == original and not removed
    assert not [check for check in preflight(working)
                if check.code == "pricing" and check.blocking]
    raw = assemble_docx(working, acknowledged_fingerprint=working.fingerprint)
    assert text in "\n".join(p.text for p in Document(BytesIO(raw)).paragraphs)
    assert original.blocks[0].text == text


def test_approximate_price_in_recommendation_stays_for_correction_and_blocks_word():
    text = "Inspect the synthetic pump; repair estimate: ~3.2k"
    table = ReportTable(("Asset", "Recommendation", "Quote reference"),
                        (("SYN-PUMP-1", text, "Q-TEST-3200"),), "synthetic-source")
    original = _draft(table=table)
    working, removed = client_table_draft(original)
    assert working == original and not removed
    assert working.blocks[0].extra_tables == (table,)
    assert any(check.code == "pricing" and check.blocking
               and check.block_key == "capital_renewal" for check in preflight(working))
    with pytest.raises(ValueError, match="[Pp]ric"):
        assemble_docx(working, acknowledged_fingerprint=working.fingerprint)
    assert original.blocks[0].extra_tables == (table,)


@pytest.mark.parametrize("text,caption,reviewed,allowed", (
    pytest.param("Inspected the synthetic pump. Quote: ~3.2k", "", True, False,
                 id="review-cannot-waive-price-in-source"),
    pytest.param("Inspected the synthetic pump.", "Repair estimate: ~ 3200", True, False,
                 id="review-cannot-waive-price-in-caption"),
    pytest.param("Inspected the synthetic pump. Estimate: ~3.2 kW", "Quote reference: Q-TEST-3200",
                 True, True, id="reviewed-technical-page"),
    pytest.param("Inspected the synthetic pump. Estimate: ~3.2 kW", "", False, False,
                 id="technical-page-still-needs-review"),
))
def test_source_page_approval_keeps_the_approximate_price_gate(text, caption, reviewed, allowed):
    source = ReportSource("a" * 64, "synthetic-inspection.pdf", "b" * 64, ".pdf",
                          page_texts=(text,), selected_pages=(1,), captions=((1, caption),))
    if reviewed:
        source = replace(source, client_page_reviews=((1, page_fingerprint(source, 1)),))
        assert dict(source.client_page_reviews)[1] == page_fingerprint(source, 1)
    original_fingerprint = source.fingerprint
    assert page_allowed(source, 1) is allowed
    assert source.fingerprint == original_fingerprint
    assert source.page_texts == (text,) and source.captions == ((1, caption),)
