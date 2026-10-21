"""Synthetic exact-occurrence source-art preservation; never asset approval."""
from copy import deepcopy
from dataclasses import replace
from io import BytesIO
import json
from zipfile import ZipFile, ZIP_DEFLATED

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches
from lxml import etree as ET
import pytest

from app import monthly_report_native_art_review as art
from app.monthly_report_compatibility_view import CompatibilityError, derive_compatibility_view
from app.monthly_report_model import Facility, ReportPeriod, synthetic_profiles
from test_monthly_report_source_frame_logo_role import source, raster, A, MC, WPS


@pytest.fixture
def sample(tmp_path, monkeypatch):
    monkeypatch.setenv('EPC_DATA_DIR', str(tmp_path / 'store'))
    path = source(tmp_path, divergent=True)
    doc = Document(path)
    # Same-byte independent occurrences are deliberately not deduplicated.
    for _ in range(2):
        doc.sections[0].header.paragraphs[0].add_run().add_picture(
            BytesIO(raster('orange')), width=Inches(1))
    doc.sections[0].footer.paragraphs[0].add_run().add_picture(
        BytesIO(raster('teal')), width=Inches(1))
    doc.save(path)
    return setup(path)


def setup(path):
    view = derive_compatibility_view(path.read_bytes())
    profile = replace(synthetic_profiles()[0], imported_from=view.original_inspection.sha256,
        title='Synthetic Regional Community Health Hospital',
        facilities=(Facility('synthetic', 'Synthetic Regional Community Health Hospital'),))
    period = ReportPeriod(2026, 9)
    inventory = art.native_art_inventory(path, view, profile, period)
    ids = tuple(x.original_id for x in inventory.occurrences)
    receipt = art.make_native_art_review(inventory, actor='Synthetic Operator',
        keep_native_ids=ids, visually_reviewed_ids=ids, suitable_ids=ids)
    return path, view, profile, period, inventory, receipt


def reviewed(sample):
    path, view, profile, period, _, receipt = sample
    return art.verify_native_art_review(path, view, profile, period, receipt)


def rewrite(raw, fn):
    with ZipFile(BytesIO(raw)) as package:
        parts = {n: package.read(n) for n in package.namelist()}
    fn(parts)
    stream = BytesIO()
    with ZipFile(stream, 'w', ZIP_DEFLATED) as package:
        for name, data in parts.items():
            package.writestr(name, data)
    return stream.getvalue()


def change_xml(parts, part, mutate):
    root = ET.fromstring(parts[part])
    mutate(root)
    parts[part] = ET.tostring(root)


def test_exact_inventory_roundtrip_keeps_all_instances_without_role_authority(sample):
    path, view, profile, period, inventory, receipt = sample
    original = path.read_bytes()
    assert len(inventory.occurrences) == 3
    first, second = [o for o in inventory.occurrences if o.original_part.startswith('word/header')]
    footer = next(o for o in inventory.occurrences if o.original_part.startswith('word/footer'))
    assert first.original_id != second.original_id
    assert first.original_media_sha256 == second.original_media_sha256
    assert first.preview_group_key == second.preview_group_key
    assert first.sha256 != second.sha256
    assert first.original_representations[0].image_path != second.original_representations[0].image_path
    assert footer.original_part.startswith('word/footer')
    assert all(o.original_part != 'word/document.xml' for o in inventory.occurrences)
    context = reviewed(sample)
    assert context.receipt == json.loads(json.dumps(receipt))
    assert context.inventory == inventory
    assert context.digest != inventory.sha256
    assert receipt['visual_review']['kind'] == 'operator_source_art_inspection'
    assert all(d['action'] == 'keep_native' for d in receipt['dispositions'])
    assert not {'approved', 'image_role', 'asset_approvals', 'pricing_review'} & set(receipt)
    art.verify_native_art_output(view.rendition_bytes, context, stage='rendition')
    assert path.read_bytes() == original
    assert art.native_art_inventory(path, view, profile, period) == inventory


@pytest.mark.parametrize('case', ['missing', 'extra', 'reordered', 'duplicate', 'one_equal_byte_only', 'unseen'])
def test_complete_exact_occurrence_set_required(sample, case):
    *_, inventory, _ = sample
    ids = [x.original_id for x in inventory.occurrences]
    visual = list(ids)
    if case == 'missing': ids.pop()
    elif case == 'extra': ids.append('item-unlisted')
    elif case == 'reordered': ids.reverse()
    elif case == 'duplicate': ids[1] = ids[0]
    elif case == 'one_equal_byte_only': ids = ids[:1]
    elif case == 'unseen': visual.pop()
    with pytest.raises(CompatibilityError):
        art.make_native_art_review(inventory, actor='Operator', keep_native_ids=ids, visually_reviewed_ids=visual, suitable_ids=ids)


@pytest.mark.parametrize('actor', ['', '  ', '\nOperator', 'Operator\x7f', 'x' * 121, None, 1])
def test_named_operator_evidence_required(sample, actor):
    *_, inventory, _ = sample
    ids = [x.original_id for x in inventory.occurrences]
    with pytest.raises(CompatibilityError):
        art.make_native_art_review(inventory, actor=actor, keep_native_ids=ids, visually_reviewed_ids=ids, suitable_ids=ids)


@pytest.mark.parametrize('case', ['source', 'view', 'period', 'contract', 'profile', 'template', 'facility',
                                  'alias', 'occurrence', 'path', 'media', 'geometry', 'actor', 'action', 'schema'])
def test_stale_or_forged_record_rejected(sample, case):
    path, view, profile, period, _, original = sample
    receipt = deepcopy(original)
    if case == 'source': receipt['scope']['original_sha256'] = '0' * 64
    elif case == 'view': receipt['scope']['source_view_receipt_sha256'] = '0' * 64
    elif case == 'period': period = ReportPeriod(2026, 10)
    elif case == 'contract': profile = replace(profile, contract='Other')
    elif case == 'profile': profile = replace(profile, key='other')
    elif case == 'template': profile = replace(profile, template='other')
    elif case == 'facility': profile = replace(profile, facilities=(Facility('foreign', 'Foreign Medical Centre'),))
    elif case == 'alias': profile = replace(profile, facilities=(replace(profile.facilities[0], aliases=('New alias',)),))
    elif case == 'occurrence': receipt['dispositions'][0]['original_id'] = 'item-unlisted'
    elif case == 'path': receipt['dispositions'][0]['image_path'] = [1, 2]
    elif case == 'media': receipt['dispositions'][0]['occurrence_sha256'] = '0' * 64
    elif case == 'geometry': receipt['inventory_sha256'] = '0' * 64
    elif case == 'actor': receipt['visual_review']['actor'] = ''
    elif case == 'action': receipt['dispositions'][0]['action'] = 'approve_logo'
    elif case == 'schema': receipt['schema'] = True
    with pytest.raises(CompatibilityError):
        art.verify_native_art_review(path, view, profile, period, receipt)


def test_missing_receipt_and_source_pin_fail(sample):
    path, view, profile, period, _, _ = sample
    with pytest.raises(CompatibilityError, match='missing'):
        art.verify_native_art_review(path, view, profile, period, None)
    with pytest.raises(CompatibilityError, match='identity'):
        art.native_art_inventory(path, view, replace(profile, imported_from='f' * 64), period)


def test_public_entries_rederive_and_reject_view_tamper(sample, monkeypatch):
    import app.monthly_report_operational_source as operations
    path, view, profile, period, _, receipt = sample
    original = operations.rederive_compatibility_view
    calls = []
    def count(*args, **kwargs):
        calls.append(1)
        return original(*args, **kwargs)
    monkeypatch.setattr(operations, 'rederive_compatibility_view', count)
    art.native_art_inventory(path, view, profile, period)
    assert len(calls) == 1
    art.verify_native_art_review(path, view, profile, period, receipt)
    assert len(calls) == 2
    bad = replace(view, rendition_bytes=view.rendition_bytes + b'changed')
    with pytest.raises(CompatibilityError):
        art.native_art_inventory(path, bad, profile, period)
    assert len(calls) == 3


@pytest.mark.parametrize('case', ['move', 'remove', 'crop', 'rotation', 'geometry', 'media', 'external', 'wrong_type', 'duplicate_rel'])
def test_output_exact_physical_placement_and_media_required(sample, case):
    _, view, _, _, inventory, _ = sample
    context = reviewed(sample)
    occurrence = next(o for o in inventory.occurrences if o.original_part == 'word/header1.xml')
    proof = occurrence.rendition_representations[0]
    def mutate(parts):
        if case == 'media':
            parts[proof.media_part] = raster('red')
            return
        if case in ('external', 'wrong_type', 'duplicate_rel'):
            def relationship(root):
                rel = next(n for n in root if n.get('Id') == proof.relationship_id)
                if case == 'external': rel.set('TargetMode', 'External')
                elif case == 'wrong_type': rel.set('Type', 'https://invalid.example/image')
                else: root.append(deepcopy(rel))
            change_xml(parts, 'word/_rels/header1.xml.rels', relationship)
            return
        def drawing(root):
            node = art._path_node(root, proof.drawing_path)
            if case == 'move':
                parent = node.getparent(); parent.remove(node); parent.append(node)
                # Moving its run changes the exact path even for a sole drawing.
                parent.getparent().insert(0, parent)
            elif case == 'remove': node.getparent().remove(node)
            elif case == 'crop': ET.SubElement(next(node.iter(qn('pic:blipFill'))), A + 'srcRect', l='200')
            elif case == 'rotation': next(node.iter(A + 'xfrm')).set('rot', '60000')
            elif case == 'geometry': next(node.iter(qn('wp:extent'))).set('cx', '1')
        change_xml(parts, occurrence.rendition_part, drawing)
    output = rewrite(view.rendition_bytes, mutate)
    with pytest.raises(CompatibilityError): art.verify_native_art_output(output, context, stage='rendition')


def test_all_equivalent_header_branch_nodes_pinned(sample):
    path, *_ = sample
    doc = Document(path)
    run = doc.sections[0].header.paragraphs[0].runs[0]._r
    drawing = next(run.iter(qn('w:drawing')))
    wrapper = ET.Element(MC + 'AlternateContent', nsmap={'wps': WPS[1:-1]})
    run.replace(drawing, wrapper)
    ET.SubElement(wrapper, MC + 'Choice', Requires='wps').append(drawing)
    ET.SubElement(wrapper, MC + 'Fallback').append(deepcopy(drawing))
    doc.save(path)
    new = setup(path)
    inventory = new[4]
    assert len(inventory.occurrences) == 3
    occurrence = next(o for o in inventory.occurrences if len(o.original_representations) == 2)
    assert occurrence.original_part == 'word/header1.xml'
    assert len(occurrence.rendition_representations) == 2
    context = reviewed(new)
    art.verify_native_art_output(new[1].rendition_bytes, context, stage='rendition')
    proof = occurrence.rendition_representations[1]
    def mutate(parts):
        change_xml(parts, 'word/header1.xml', lambda root:
            art._path_node(root, proof.image_path).set(qn('r:embed'), 'missing'))
    with pytest.raises(CompatibilityError):
        art.verify_native_art_output(rewrite(new[1].rendition_bytes, mutate), context, stage='rendition')


def test_crop_contexts_are_not_preview_grouped(sample):
    path, *_ = sample
    doc = Document(path)
    pictures = tuple(doc.sections[0].header._element.iter(qn('pic:blipFill')))
    ET.SubElement(pictures[1], A + 'srcRect', l='200')
    doc.save(path)
    new = setup(path)
    first, second = [o for o in new[4].occurrences if o.original_part.startswith('word/header')]
    assert first.preview_group_key != second.preview_group_key


@pytest.mark.parametrize('tag', ['fldSimple', 'ins', 'del', 'vanish', 'instrText', 'hyperlink'])
def test_raw_unit_commands_revisions_and_hidden_content_not_erased_into_review(sample, tag):
    path, view, profile, period, _, _ = sample
    doc = Document(path)
    unit = doc.sections[0].header.paragraphs[0]._p
    ET.SubElement(unit, qn('w:' + tag))
    doc.save(path)
    # The original wrapper-specific view may itself refuse unsafe content;
    # otherwise preservation must refuse it before any passive erasure.
    with pytest.raises(ValueError):
        changed = derive_compatibility_view(path.read_bytes())
        art.native_art_inventory(path, changed,
            replace(profile, imported_from=changed.original_inspection.sha256), period)


def test_empty_inventory_is_legacy_unchanged(tmp_path, monkeypatch):
    monkeypatch.setenv('EPC_DATA_DIR', str(tmp_path / 'store'))
    path = source(tmp_path, divergent=True)
    view = derive_compatibility_view(path.read_bytes())
    profile = replace(synthetic_profiles()[0], facilities=(Facility('synthetic',
        'Synthetic Regional Community Health Hospital'),))
    assert art.verify_native_art_review(path, view, profile, ReportPeriod(2026,9), None) is None
    art.verify_native_art_output(b'', None)


def test_budget_and_corrupt_output_rejected(sample, monkeypatch):
    context = reviewed(sample)
    with pytest.raises(CompatibilityError): art.verify_native_art_output(b'not a ZIP', context)
    monkeypatch.setattr(art, 'MAX_VIEW_SOURCE_BYTES', 4)
    with pytest.raises(CompatibilityError): art.verify_native_art_output(sample[1].rendition_bytes, context)
    with pytest.raises(CompatibilityError): art._native_art_inventory(*sample[:4])


def test_authoritative_saved_profile_memberships_join_scope_recognition(sample, monkeypatch):
    from app import monthly_report_library as library, monthly_report_start_import as starting
    path, view, profile, period, _, _ = sample
    other = Facility('independent', 'Independent Synthetic Hospital')
    monkeypatch.setattr(library, 'list_profiles', lambda contract: (replace(profile, facilities=(other,)),))
    original = starting.source_sites
    observed = []
    def recognize(inspection, known):
        known = tuple(known)
        observed.extend(known)
        return (*original(inspection, known), other)
    monkeypatch.setattr(starting, 'source_sites', recognize)
    with pytest.raises(CompatibilityError, match='same-site'):
        art.native_art_inventory(path, view, profile, period)
    assert other in observed


def test_missing_view_empty_facility_and_scope_tamper_fail(sample):
    path, view, profile, period, _, _ = sample
    with pytest.raises(CompatibilityError): art.native_art_inventory(path, None, profile, period)
    with pytest.raises(ValueError):
        art.native_art_inventory(path, view, replace(profile, facilities=()), period)
    path.write_bytes(path.read_bytes() + b'changed-original')
    with pytest.raises(CompatibilityError, match='original changed'):
        art._native_art_inventory(path, view, profile, period)


def test_empty_inventory_does_not_add_profile_gates(tmp_path):
    path = source(tmp_path, divergent=True)
    view = derive_compatibility_view(path.read_bytes())
    assert art.verify_native_art_review(path, view, None, None, None) is None


def test_ineligible_role_cannot_gain_preservation_from_matching_media(sample):
    # A role belongs to exact inspector occurrences, never to same-byte siblings.
    path, *_ = sample
    doc = Document(path)
    first = next(doc.sections[0].header._element.iter(qn('wp:docPr')))
    first.set('descr', 'Company logo')
    doc.save(path)
    new = setup(path)
    assert len(new[4].occurrences) == 2
    assert len([o for o in new[4].occurrences if o.original_part.startswith('word/header')]) == 1


def test_inventory_and_representation_limits(sample, monkeypatch):
    path, view, profile, period, _, _ = sample
    monkeypatch.setattr(art, 'MAX_OCCURRENCES', 2)
    with pytest.raises(CompatibilityError, match='budget'):
        art.native_art_inventory(path, view, profile, period)
    monkeypatch.setattr(art, 'MAX_OCCURRENCES', 128)
    monkeypatch.setattr(art, 'MAX_REPRESENTATIONS', 1)
    with pytest.raises(CompatibilityError, match='budget'):
        art.native_art_inventory(path, view, profile, period)


@pytest.mark.parametrize('persisted', [False, True])
def test_explicit_or_persisted_master_has_no_ordinary_preservation_authority(sample, monkeypatch, persisted):
    from app import monthly_report_library as library
    path, view, profile, period, _, _ = sample
    master = replace(profile, template='enfra_master:' + 'a' * 64)
    if persisted:
        monkeypatch.setattr(library, 'list_profiles', lambda contract: (master,))
    else:
        profile = master
    with pytest.raises(CompatibilityError, match='company-master'):
        art.native_art_inventory(path, view, profile, period)


def test_only_exact_existing_description_scrub_is_allowed_at_output(sample):
    from app.monthly_report_native_layout import _clear_source_drawing_metadata
    context = reviewed(sample)
    doc = Document(BytesIO(sample[1].rendition_bytes))
    _clear_source_drawing_metadata(doc)
    output = BytesIO(); doc.save(output)
    art.verify_native_art_output(output.getvalue(), context)
    occurrence = next(o for o in context.inventory.occurrences if o.rendition_part == 'word/header1.xml')
    proof = occurrence.rendition_representations[0]
    for name, value in [('descr', 'New unsanctioned hidden text'), ('mystery', 'unknown'), ('name', 'Other')]:
        def mutate(parts):
            def drawing(root):
                node = art._path_node(root, proof.drawing_path)
                next(node.iter(qn('wp:docPr'))).set(name, value)
            change_xml(parts, occurrence.rendition_part, drawing)
        with pytest.raises(CompatibilityError):
            art.verify_native_art_output(rewrite(output.getvalue(), mutate), context)


def test_output_privacy_stage_is_explicit_not_interchangeable(sample):
    from app.monthly_report_native_layout import _clear_source_drawing_metadata
    context = reviewed(sample)
    assert context.inventory.scope['drawing_privacy_policy'] == art.DRAWING_PRIVACY_POLICY
    with pytest.raises(CompatibilityError):
        art.verify_native_art_output(sample[1].rendition_bytes, context)
    doc = Document(BytesIO(sample[1].rendition_bytes))
    _clear_source_drawing_metadata(doc)
    output = BytesIO(); doc.save(output)
    art.verify_native_art_output(output.getvalue(), context)
    with pytest.raises(CompatibilityError):
        art.verify_native_art_output(output.getvalue(), context, stage='rendition')
    with pytest.raises(CompatibilityError):
        art.verify_native_art_output(output.getvalue(), context, stage='unspecified')


def test_output_cannot_change_inherited_rendering_context(sample):
    context = reviewed(sample)
    def mutate(parts):
        change_xml(parts, 'word/header1.xml', lambda root: root.set(MC + 'Ignorable', 'wps'))
    with pytest.raises(CompatibilityError):
        art.verify_native_art_output(rewrite(sample[1].rendition_bytes, mutate), context, stage='rendition')


def test_corrupt_output_member_fails_as_preservation_error(sample):
    from zipfile import ZIP_STORED
    context = reviewed(sample)
    media = context.inventory.occurrences[0].rendition_representations[0].media_part
    stream = BytesIO()
    with ZipFile(BytesIO(sample[1].rendition_bytes)) as source, ZipFile(stream, 'w', ZIP_STORED) as output:
        for name in source.namelist():
            output.writestr(name, source.read(name))
    with ZipFile(BytesIO(stream.getvalue())) as package:
        member = package.getinfo(media)
        offset = member.header_offset + 30 + len(member.filename.encode()) + len(member.extra)
    damaged = bytearray(stream.getvalue()); damaged[offset] ^= 1
    with pytest.raises(CompatibilityError):
        art.verify_native_art_output(bytes(damaged), context, stage='rendition')


def test_keep_inspection_alone_does_not_authorize_visible_content(sample):
    path, view, profile, period, inventory, receipt = sample
    ids = tuple(item.original_id for item in inventory.occurrences)
    with pytest.raises(CompatibilityError, match='suitability'):
        art.make_native_art_review(inventory, actor='Synthetic Operator', keep_native_ids=ids,
                                   visually_reviewed_ids=ids)
    missing = deepcopy(receipt); missing.pop('visible_content_review')
    with pytest.raises(CompatibilityError):
        art.verify_native_art_review(path, view, profile, period, missing)
    altered = deepcopy(receipt); altered['visible_content_review']['checks'].remove('no_visible_pricing')
    with pytest.raises(CompatibilityError):
        art.verify_native_art_review(path, view, profile, period, altered)
