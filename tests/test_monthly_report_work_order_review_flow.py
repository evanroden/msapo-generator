from dataclasses import asdict, replace
from io import BytesIO
import json

from docx import Document
import pytest

from app import monthly_report_sources as sources, monthly_report_library as library
from app import monthly_report_work_order_review_flow as flow
from app.monthly_report_work_order_review_ui import apply_request
from app.monthly_report_model import Facility, ReportProfile, ReportPeriod, ReportDraft, ResolvedBlock, default_sections

PROFILE = ReportProfile("Synthetic reviewed records", "review-flow", "Two campuses", (
    Facility("north", "Synthetic North", ("North",)), Facility("south", "Synthetic South", ("South",))), scope_type="multi_site")
PERIOD = ReportPeriod(2026, 9)
CSV = b"Work Order,Site,Date,Status,Type,Trade\n0001,North,2026-09-03,Open,PM,HVAC\n0002,South,2026-09-04,Closed,Repair,Electrical\n"


@pytest.fixture
def storage(tmp_path, monkeypatch):
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path))
    return tmp_path


def draft_source(raw=CSV, name="rows.csv", profile=PROFILE):
    content, _ = sources.ingest(profile, name, raw)
    draft = ReportDraft(profile, PERIOD, "Synthetic Editor", default_sections(), (
        ResolvedBlock("work_orders", "This month"), ResolvedBlock("service_calls", "This month"),
        ResolvedBlock("activity_summary", "This month", text="Keep my authored wording.")),
        synthetic=True, sources=(content.source,))
    return draft, content


def stored(draft, content, kind="workorders"):
    bundle = flow.prepare_bundle(draft, content, kind=kind)
    return flow.store_bundle(draft, bundle), bundle


def publish(draft, bundle, **options):
    key = "work_orders" if bundle["kind"] == "workorders" else "service_calls"
    block = next(b for b in draft.blocks if b.key == key)
    return flow.publish_bundle(draft, bundle, expected_revision=flow.revision(bundle),
        expected_block=block.fingerprint, **options)


def test_corrections_survive_snapshot_active_recovery_and_exact_reimport(storage):
    from app import monthly_report_active_work as active
    draft, content = draft_source()
    draft, initial = stored(draft, content)
    row = flow.restored_group(initial, initial["groups"][0]).rows[0]
    revised = flow.change_row(initial, 0, row.key, {"status": "Completed", "trade": "Plumbing"})
    draft = flow.store_bundle(draft, revised, expected_revision=flow.revision(initial))
    library.save_profile(PROFILE, expected_revision=0, actor="Synthetic Editor", confirmed=True)
    library.save_snapshot(draft, expected_revision=0)
    reopened = library.load_snapshot(PROFILE.contract, PROFILE.key, PERIOD).draft
    assert flow.saved_bundle(reopened, content.source.id, "workorders") == revised
    token = "1" * 32
    active.save(token, reopened, expected_revision=0, base_snapshot_revision=1)
    restored = active.load(token, PROFILE.contract, PROFILE.key, PERIOD).draft
    assert restored == reopened
    again, _ = sources.ingest(PROFILE, "rows.csv", CSV)
    assert flow.saved_bundle(restored, again.source.id, "workorders") == revised
    assert sources.source_bytes(PROFILE, content.source) == CSV
    with pytest.raises(library.RevisionConflict):
        library.save_snapshot(replace(restored, prepared_by="Stale editor"), expected_revision=0)


def test_maintenance_review_decisions_are_reversible_and_rebind_receipt(storage):
    raw = b"WO #,Finish Date,Area,Task,Tag,Site,Status,Cost\n00041,2026-09-15,Plant,Check pump,P-41,North,Completed,12\n,,,Additional handwritten note,,,,\n"
    draft, content = draft_source(raw, "reversible-maintenance.csv")
    draft, bundle = stored(draft, content, "maintenance")
    accepted = flow.group_decisions(bundle, 0, omitted=False, omit_unmatched=True, financial=True, layout=True)
    record = flow.restored_group(accepted, accepted["groups"][0])
    assert record.ready_for_inclusion and record.unmatched[0].excluded
    draft = flow.store_bundle(draft, accepted, expected_revision=flow.revision(bundle))
    reversed_decisions = flow.group_decisions(accepted, 0, omitted=False, omit_unmatched=False, financial=False, layout=False)
    record = flow.restored_group(reversed_decisions, reversed_decisions["groups"][0])
    assert not record.unmatched[0].excluded
    assert not record.financial_exclusion_confirmed and not record.native_layout_confirmed
    assert not record.ready_for_inclusion
    assert flow.revision(reversed_decisions) != flow.revision(accepted)
    draft = flow.store_bundle(draft, reversed_decisions, expected_revision=flow.revision(accepted))
    with pytest.raises(ValueError):
        publish(draft, reversed_decisions)

def test_explicit_publication_is_repeat_safe_preserves_manual_content_and_recomputes(storage):
    draft, content = draft_source()
    draft, original = stored(draft, content)
    accepted = publish(draft, original, include_trades=True)
    saved = flow.saved_bundle(accepted, content.source.id, "workorders")
    block = next(b for b in accepted.blocks if b.key == "work_orders")
    assert len(block.extra_tables) == 3
    assert block.extra_tables[1].rows[0][-3:] == ("2", "1", "1")
    assert publish(accepted, saved, include_trades=True) == accepted
    row = flow.restored_group(saved, saved["groups"][0]).rows[0]
    revised = flow.change_row(saved, 0, row.key, {"status": "Completed"})
    corrected = flow.store_bundle(accepted, revised, expected_revision=flow.revision(saved))
    updated = publish(corrected, revised, include_trades=True)
    changed = next(b for b in updated.blocks if b.key == "work_orders")
    assert len(changed.extra_tables) == 3 and changed.extra_tables[1].rows[0][-3:] == ("2", "2", "0")
    assert next(b for b in updated.blocks if b.key == "activity_summary").text == "Keep my authored wording."
    assert accepted.blocks != updated.blocks and draft.blocks != accepted.blocks


def test_manual_table_edits_or_omissions_cannot_be_overwritten(storage):
    draft, content = draft_source()
    draft, initial = stored(draft, content)
    accepted = publish(draft, initial)
    saved = flow.saved_bundle(accepted, content.source.id, "workorders")
    work = next(b for b in accepted.blocks if b.key == "work_orders")
    edited = replace(work, extra_tables=(replace(work.extra_tables[0], rows=(("Manual table",),)), *work.extra_tables[1:]))
    newer = replace(accepted, blocks=tuple(edited if b.key == work.key else b for b in accepted.blocks))
    with pytest.raises(ValueError, match="edited or removed"):
        publish(newer, saved)
    with pytest.raises(ValueError, match="table changed"):
        flow.publish_bundle(newer, saved, expected_revision=flow.revision(saved), expected_block=work.fingerprint)
    omitted = replace(draft, blocks=tuple(replace(b, source="Omit") if b.key == work.key else b for b in draft.blocks))
    with pytest.raises(ValueError, match="deliberately omitted"):
        publish(omitted, initial)


def test_mapping_numeric_dates_and_unresolved_rows_require_explicit_review(storage):
    draft, content = draft_source(b"Ticket,Building,Reported,State,Type\n0003,North,09/04/2026,Done,PM\n", "ambiguous.csv")
    draft, original = stored(draft, content)
    with pytest.raises(ValueError):
        publish(draft, original)
    mapped = flow.remap_group(draft, content, original, 0,
        {"work_order": 0, "site": 1, "work_date": 2, "status": 3, "work_type": 4,
         "completed_date": None, "trade": None}, "MDY")
    fixed = flow.store_bundle(draft, mapped, expected_revision=flow.revision(original))
    result = publish(fixed, mapped)
    table = next(b for b in result.blocks if b.key == "work_orders").extra_tables[0]
    assert table.rows[0][:3] == ("0003", "north", "2026-09-04")
    assert flow.restored_group(mapped, mapped["groups"][0]).original_rows[0][2] == "09/04/2026"


def test_stale_scope_source_and_row_requests_preserve_current_values(storage):
    draft, content = draft_source()
    draft, original = stored(draft, content)
    request = {"scope": flow.scope(draft), "source": content.source.fingerprint,
        "kind": "workorders", "expected": flow.revision(original), "action": "start"}
    assert apply_request(draft, content, request) == draft
    with pytest.raises(ValueError, match="report or source changed"):
        apply_request(replace(draft, period=ReportPeriod(2026, 10)), content, request)
    with pytest.raises(ValueError, match="row review changed"):
        apply_request(draft, content, {**request, "expected": "older"})
    newer_content = replace(content, source=replace(content.source, vendor="Edited source metadata"))
    updated = replace(draft, sources=(newer_content.source,))
    assert flow.saved_bundle(updated, content.source.id, "workorders") is None
    next_review = flow.prepare_bundle(updated, newer_content)
    retained = flow.store_bundle(updated, next_review)
    assert len(retained.source_row_reviews) == 2
    assert json.loads(retained.source_row_reviews[0]) == original


def test_new_month_drops_active_review_receipts_without_changing_history(storage):
    from app.monthly_report_setup import new_month_draft
    draft, content = draft_source()
    draft, bundle = stored(draft, content)
    after = new_month_draft(draft, ReportPeriod(2026, 10))
    assert not after.source_row_reviews and draft.source_row_reviews
    legacy = asdict(replace(draft, source_row_reviews=()))
    legacy.pop("source_row_reviews")
    assert library.draft_from_dict(legacy).fingerprint == replace(draft, source_row_reviews=()).fingerprint


def test_pdf_pages_persist_corrected_rows_and_scans_need_explicit_omission(storage):
    from test_monthly_report_work_order_pdf_review import _mixed_pdf
    draft, content = draft_source(_mixed_pdf(scanned=True), "native.pdf")
    draft, bundle = stored(draft, content)
    assert len(bundle["groups"]) == 3
    with pytest.raises(ValueError, match="unsupported"):
        publish(draft, bundle)
    accepted = flow.group_decisions(bundle, 2, omitted=True)
    draft = flow.store_bundle(draft, accepted, expected_revision=flow.revision(bundle))
    ready = publish(draft, accepted)
    assert sum(len(t.rows) for t in next(b for b in ready.blocks if b.key == "work_orders").extra_tables
               if t.columns[0] == "Work order") == 7
    assert library.draft_from_dict(asdict(ready)) == ready


def test_native_word_rows_retain_original_provenance_through_actual_word_output(storage):
    from test_monthly_report_work_order_docx_review import _normal_document
    stream = BytesIO()
    _normal_document().save(stream)
    raw = stream.getvalue()
    draft, content = draft_source(raw, "native.docx")
    from app.monthly_report_source_scope import correct_scope
    content = replace(content, source=correct_scope(content.source, PROFILE, ("north", "south")))
    draft = replace(draft, sources=(content.source,),
                    sections=tuple(s for s in draft.sections if s.key == "activity"))
    draft, bundle = stored(draft, content)
    result = publish(draft, bundle)
    from app.monthly_report_docx import assemble_docx
    output = assemble_docx(result, acknowledged_fingerprint=result.fingerprint)
    document = Document(BytesIO(output))
    text = "\n".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
    assert "000071" in text and "000072" in text
    assert text.count("000071") == 1
    assert sources.source_bytes(PROFILE, content.source) == raw


def test_maintenance_rows_are_visible_correctable_and_inserted_without_overwriting_existing_rows(storage):
    raw = b"Work Order Number,Actual Finish,Area,Work Performed,Equipment Tag,Site,State\n000061,2026-09-09,Mechanical,Inspected AHU,AHU-61,South,Closed\n"
    draft, content = draft_source(raw, "maintenance.csv")
    manual = ResolvedBlock("service_calls", "This month", rows=(("MANUAL", "2026-09-01", "Room", "Preserve", ""),))
    draft = replace(draft, blocks=tuple(manual if b.key == manual.key else b for b in draft.blocks))
    draft, bundle = stored(draft, content, "maintenance")
    accepted = publish(draft, bundle)
    block = next(b for b in accepted.blocks if b.key == "service_calls")
    assert block.rows == manual.rows
    assert block.extra_tables[0].rows[0][:3] == ("000061", "2026-09-09", "Synthetic South — Mechanical")
    assert library.draft_from_dict(asdict(accepted)) == accepted


def test_actual_workbook_review_keeps_leading_zeros_and_locale_decisions(storage):
    from openpyxl import Workbook
    wb = Workbook()
    wb.active.append(("Work Order", "Site", "Date", "Status"))
    wb.active.append(("0007", "North", "09/03/2026", "Open"))
    stream = BytesIO()
    wb.save(stream)
    draft, content = draft_source(stream.getvalue(), "orders.xlsx")
    draft, bundle = stored(draft, content)
    mapped = flow.remap_group(draft, content, bundle, 0,
        {"work_order": 0, "site": 1, "work_date": 2, "status": 3}, "MDY")
    draft = flow.store_bundle(draft, mapped, expected_revision=flow.revision(bundle))
    result = publish(draft, mapped)
    assert next(b for b in result.blocks if b.key == "work_orders").extra_tables[0].rows[0][0] == "0007"


def test_changed_source_details_replace_only_exact_prior_publication(storage):
    draft, content = draft_source()
    draft, first = stored(draft, content)
    accepted = publish(draft, first)
    updated_content = replace(content, source=replace(content.source, vendor="Corrected descriptive metadata"))
    updated = replace(accepted, sources=(updated_content.source,))
    updated, current = stored(updated, updated_content)
    result = publish(updated, current)
    assert len(next(b for b in result.blocks if b.key=="work_orders").extra_tables) == 2


def test_different_overlapping_exports_cannot_double_count_accepted_source_rows(storage):
    draft, content = draft_source()
    draft, first = stored(draft, content)
    accepted = publish(draft, first)
    other, _ = sources.ingest(PROFILE, "second.csv", CSV+b"\n")
    combined = replace(accepted, sources=(*accepted.sources, other.source))
    combined, review = stored(combined, other)
    with pytest.raises(ValueError, match="Another included source"):
        publish(combined, review)


def test_native_maintenance_layout_and_incomplete_rows_need_review(storage):
    from test_monthly_report_service_call_mapping import word_fixture, PROFILE as CALL_PROFILE
    raw = BytesIO()
    word_fixture().save(raw)
    draft, content = draft_source(raw.getvalue(), "calls.docx", CALL_PROFILE)
    draft, bundle = stored(draft, content, "maintenance")
    record = flow.restored_group(bundle, bundle["groups"][0])
    assert record.rows and not record.native_layout_confirmed
    with pytest.raises(ValueError):
        publish(draft, bundle)
    revised = flow.change_row(bundle, 0, record.rows[1].key, {}, excluded=True)
    revised = flow.group_decisions(revised, 0, omitted=False, layout=True)
    draft = flow.store_bundle(draft, revised, expected_revision=flow.revision(bundle))
    accepted = publish(draft, revised)
    assert next(b for b in accepted.blocks if b.key=="service_calls").extra_tables[0].rows[0][0] == "0009"


from conftest import requires_libreoffice


@requires_libreoffice
def test_reviewed_source_rows_and_totals_survive_actual_word_pdf_export(storage, monkeypatch):
    import pymupdf
    from app import pdf_converter
    from app.monthly_report_docx import generate_report
    from app.monthly_report_source_scope import correct_scope
    draft, content = draft_source()
    content = replace(content, source=correct_scope(content.source, PROFILE, ("north", "south")))
    draft = replace(draft, sources=(content.source,),
                    sections=tuple(s for s in draft.sections if s.key == "activity"))
    draft, bundle = stored(draft, content)
    accepted = publish(draft, bundle, include_trades=True)
    monkeypatch.setattr(pdf_converter, "OUTPUT_DIR", storage)
    monkeypatch.setattr(pdf_converter, "PDF_BACKEND", "libreoffice")
    output = generate_report(accepted, acknowledged_fingerprint=accepted.fingerprint)
    assert output.docx and output.pdf and not output.pdf_error
    with pymupdf.open(stream=output.pdf, filetype="pdf") as pdf:
        text = "\n".join(page.get_text() for page in pdf)
        assert "0001" in text and "0002" in text
        assert "Recorded" in text and "Unclassified" in text
        assert text.count("0001") == 1 and text.count("0002") == 1
    assert sources.source_bytes(PROFILE, content.source) == CSV



