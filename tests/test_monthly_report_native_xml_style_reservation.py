from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.packuri import PackURI
import pytest

@pytest.mark.parametrize('name',['/word/header.custom','/word/header.XML'])
def test_xml_header_content_type_is_reserved_regardless_of_name(name):
    from app.monthly_report_native_narrative_visibility import narrative_visibility_writer
    doc=Document();header=doc.sections[0].header
    header.part.partname=PackURI(name)
    p=header.paragraphs[0];p.text='Unchanged header witness.'
    old=p._p.get_or_add_pPr().find(qn('w:pStyle'))
    if old is not None: old.getparent().remove(old)
    sel=OxmlElement('w:pStyle');sel.set(qn('w:val'),'ENFRACurrentVisibleNarrative');p._p.get_or_add_pPr().append(sel)
    source=doc.add_paragraph('Hidden source');source.runs[0].font.hidden=True
    result=narrative_visibility_writer(doc)(source._p)
    identity=result.find('./'+qn('w:pPr')+'/'+qn('w:pStyle')).get(qn('w:val'))
    assert identity!='ENFRACurrentVisibleNarrative'
