"""Upload guidance and rejection use the route that actually reads the file."""
from io import BytesIO
import hashlib

from docx import Document
from openpyxl import Workbook
import pymupdf
import pytest
from streamlit.testing.v1 import AppTest

from app import monthly_report_sources as sources
from app import monthly_report_import as importer
from app import monthly_report_capacity as capacity
from app import monthly_report_directory as directory
from app.monthly_report_model import synthetic_profiles
from app.monthly_report_upload_limits import MIB


@pytest.fixture
def profile(tmp_path, monkeypatch):
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path))
    return synthetic_profiles()[0]


def docx_bytes():
    document = Document()
    document.add_paragraph("Inspected the synthetic pump.")
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def xlsx_bytes():
    workbook = Workbook()
    workbook.active.append(("Facility", "Service", "Capacity", "Units"))
    workbook.active.append(("Synthetic North", "Steam", "1200", "lb/hr"))
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def pdf_bytes(count):
    with pymupdf.open() as pdf:
        for number in range(count):
            page = pdf.new_page()
            page.insert_text((30, 50), f"Inspected synthetic pump {number + 1}; readings checked.")
        return pdf.tobytes()


@pytest.mark.parametrize("size", [31, 32])
def test_text_at_or_below_exact_limit_keeps_original(profile, monkeypatch, size):
    monkeypatch.setattr(sources, "MAX_FILE_BYTES", 32)
    raw = b"Inspection: " + b"x" * (size - 12)
    content, _ = sources.ingest(profile, "synthetic.txt", raw)
    assert sources.source_bytes(profile, content.source) == raw
    assert content.source.page_texts == (raw.decode(),)


@pytest.mark.parametrize("name", ["synthetic.txt", "synthetic.pdf", "synthetic.csv"])
def test_oversize_names_binding_file_limit_before_storage(profile, monkeypatch, name):
    monkeypatch.setattr(sources, "MAX_FILE_BYTES", 32)
    monkeypatch.setattr(sources, "_path", lambda *args: pytest.fail("Oversize input reached storage."))
    with pytest.raises(ValueError, match="32 bytes") as error:
        sources.ingest(profile, name, b"x" * 33)
    assert "smaller" in str(error.value)


def test_empty_and_unsupported_inputs_are_not_mislabeled_oversize(profile):
    with pytest.raises(ValueError, match="empty"):
        sources.ingest(profile, "synthetic.txt", b"")
    with pytest.raises(ValueError, match="Unsupported file type"):
        sources.ingest(profile, "synthetic.exe", b"not executed")


@pytest.mark.parametrize("allowance", [0, 1])
def test_actual_docx_at_and_below_import_limit(profile, monkeypatch, allowance):
    raw = docx_bytes()
    monkeypatch.setattr(importer, "MAX_DOCX_BYTES", len(raw) + allowance)
    content, _ = sources.ingest(profile, "synthetic.DOCX", raw)
    assert sources.source_bytes(profile, content.source) == raw
    assert any("Inspected the synthetic pump." in text for text in content.source.page_texts)


def test_evidence_docx_rejects_actual_import_cap_before_storing(profile, monkeypatch):
    raw = docx_bytes()
    limit = len(raw) - 1
    monkeypatch.setattr(importer, "MAX_DOCX_BYTES", limit)
    monkeypatch.setattr(sources, "_path", lambda *args: pytest.fail("Rejected DOCX reached storage."))
    with pytest.raises(ValueError, match=f"{limit:,} bytes"):
        sources.ingest(profile, "synthetic.docx", raw)


def test_saved_docx_read_uses_same_current_limit(profile, monkeypatch):
    raw = docx_bytes()
    content, _ = sources.ingest(profile, "synthetic.docx", raw)
    monkeypatch.setattr(importer, "MAX_DOCX_BYTES", len(raw) - 1)
    with pytest.raises(ValueError, match=f"{len(raw)-1:,} bytes"):
        sources.source_bytes(profile, content.source)
    assert sources._path(profile, content.source.sha256, ".docx").read_bytes() == raw


@pytest.mark.parametrize("allowance", [0, 1])
def test_capacity_xlsx_preserves_exact_boundary_and_original(profile, monkeypatch, allowance):
    raw = xlsx_bytes()
    monkeypatch.setattr(directory, "MAX_UPLOAD_BYTES", len(raw) + allowance)
    inspected = capacity.inspect_capacity(profile, "synthetic.xlsx", raw)
    assert sources.source_bytes(profile, inspected.content.source) == raw
    assert inspected.content.source.sha256 == hashlib.sha256(raw).hexdigest()
    assert capacity.load_capacity(profile.contract) is None


def test_capacity_xlsx_names_its_own_limit_before_parser(profile, monkeypatch):
    monkeypatch.setattr(directory, "MAX_UPLOAD_BYTES", 32)
    monkeypatch.setattr(directory, "inspect_workbook", lambda *a, **kw: pytest.fail("Oversize reached workbook parser."))
    with pytest.raises(ValueError, match="32 bytes"):
        capacity.inspect_capacity(profile, "synthetic.xlsx", b"x" * 33)
    assert capacity.load_capacity(profile.contract) is None


def test_failed_sibling_and_retry_keep_accepted_source(profile, monkeypatch):
    monkeypatch.setattr(sources, "MAX_FILE_BYTES", 32)
    accepted_raw = b"Inspected pump."
    accepted, _ = sources.ingest_batch(profile, (("good.txt", accepted_raw),))
    again, notices = sources.ingest_batch(profile, (("oversize.txt", b"x" * 33),), accepted)
    assert again == accepted
    assert any("32 bytes" in notice for notice in notices)
    assert sources.source_bytes(profile, accepted[0].source) == accepted_raw
    rejected_hash = hashlib.sha256(b"x" * 33).hexdigest()
    assert not sources._path(profile, rejected_hash, ".txt").exists()
    retry_raw = b"Checked valve."
    retried, _ = sources.ingest_batch(profile, (("oversize.txt", retry_raw),), again)
    assert len(retried) == 2 and retried[0] == accepted[0]
    assert sources.source_bytes(profile, retried[1].source) == retry_raw


def test_pdf_page_error_reports_configured_constraint(profile, monkeypatch):
    monkeypatch.setattr(sources, "MAX_PAGES", 1)
    with pytest.raises(ValueError, match="1 pages"):
        sources.ingest(profile, "synthetic.pdf", pdf_bytes(2))


FIELD = '''
import streamlit as st
from app.monthly_report_model import ReportPeriod, synthetic_profiles
profile = synthetic_profiles()[0]
def field(key, value):
    if key not in st.session_state:
        st.session_state[key] = value
    return key
'''


@pytest.mark.parametrize("route", ["source_manager", "section", "structured", "capacity"])
def test_native_uploader_ceiling_help_and_identity_match_route(monkeypatch, tmp_path, route):
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path))
    monkeypatch.setattr(importer, "MAX_DOCX_BYTES", 3 * MIB)
    monkeypatch.setattr(sources, "MAX_FILE_BYTES", 2 * MIB)
    monkeypatch.setattr(directory, "MAX_UPLOAD_BYTES", MIB)
    scripts = {
        "source_manager": 'from app.monthly_report_upload_ui import render_uploads\nrender_uploads(profile, ReportPeriod(2026, 9), "limits", field)',
        "section": 'from app.monthly_report_upload_ui import render_section_uploads\nrender_section_uploads(profile, ReportPeriod(2026, 9), "limits", field, "water_reports")',
        "structured": 'from app.monthly_report_upload_ui import render_structured_uploads\nrender_structured_uploads(profile, ReportPeriod(2026, 9), "limits", field, "activity_summary")',
        "capacity": 'from app.monthly_report_capacity_ui import _capacity_file\n_capacity_file("limits")',
    }
    app = AppTest.from_string(FIELD + scripts[route], default_timeout=20).run()
    assert not app.exception
    picker, = app.get("file_uploader")
    assert picker.proto.max_upload_size_mb == 3
    assert "DOCX: up to 3 MiB" in picker.proto.help
    assert "2 MiB" in picker.proto.help
    expected_detail = "1 MiB" if route == "capacity" else "500"
    assert expected_detail in picker.proto.help
    first_id = picker.id
    app.run()
    assert not app.exception
    assert app.get("file_uploader")[0].id == first_id
    assert not app.session_state.filtered_state.get("limits_evidence")


@pytest.mark.parametrize("optional", [False, True])
def test_primary_and_secondary_starting_report_guidance(monkeypatch, tmp_path, optional):
    from app import monthly_report_compatibility_view as view
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path))
    monkeypatch.setattr(importer, "MAX_DOCX_BYTES", 3 * MIB)
    monkeypatch.setattr(view, "MAX_VIEW_SOURCE_BYTES", MIB)
    script = FIELD + '''
from app.monthly_report_section_ui import render_section_setup
render_section_setup(profile.contract, ReportPeriod(2026, 9), "Synthetic Editor", field,
                     identity=profile, working_draft=object() if OPTIONAL else None)
'''.replace("OPTIONAL", repr(optional))
    app = AppTest.from_string(script, default_timeout=20).run()
    assert not app.exception
    picker, = app.get("file_uploader")
    assert picker.proto.max_upload_size_mb == 3
    assert "3 MiB" in picker.proto.help and "1 MiB" in picker.proto.help
    assert "artwork" in picker.proto.help
    first_id = picker.id
    app.run()
    assert not app.exception and app.get("file_uploader")[0].id == first_id


def test_capacity_xlsx_also_respects_shared_original_read_limit(profile, monkeypatch):
    monkeypatch.setattr(directory, "MAX_UPLOAD_BYTES", 64)
    monkeypatch.setattr(sources, "MAX_FILE_BYTES", 32)
    monkeypatch.setattr(directory, "inspect_workbook", lambda *a, **kw: pytest.fail("Oversize reached workbook parser."))
    with pytest.raises(ValueError, match="32 bytes"):
        capacity.inspect_capacity(profile, "synthetic.xlsx", b"x" * 33)


@pytest.mark.parametrize("constraint", ["count", "bytes", "text"])
def test_batch_error_identifies_binding_budget_without_changing_existing(profile, monkeypatch, constraint):
    raw = b"Inspected pump."
    accepted, _ = sources.ingest_batch(profile, (("accepted.txt", raw),))
    if constraint == "count":
        monkeypatch.setattr(sources, "MAX_SOURCES", 1)
        message = "1 source files"
    elif constraint == "bytes":
        monkeypatch.setattr(sources, "MAX_REPORT_BYTES", len(raw))
        message = f"{len(raw)} bytes"
    else:
        monkeypatch.setattr(sources, "MAX_REPORT_TEXT", len(raw))
        message = f"{len(raw)} characters"
    result, notices = sources.ingest_batch(profile, (("next.txt", b"Checked valve."),), accepted)
    assert result == accepted
    assert any(message in notice for notice in notices)
    assert sources.source_bytes(profile, result[0].source) == raw


@pytest.mark.parametrize("constraint", ["expanded", "part", "parts"])
def test_xlsx_parser_error_names_exact_binding_budget(profile, monkeypatch, constraint):
    raw = xlsx_bytes()
    if constraint == "expanded":
        monkeypatch.setattr(sources, "MAX_WORKBOOK_BYTES", 32)
        expected = "Expanded workbook exceeds the 32 bytes"
    elif constraint == "part":
        monkeypatch.setattr(sources, "MAX_WORKBOOK_PART_BYTES", 32)
        expected = "A workbook part exceeds the 32 bytes"
    else:
        monkeypatch.setattr(sources, "MAX_WORKBOOK_PARTS", 1)
        expected = "Workbook exceeds 1 package parts"
    monkeypatch.setattr(sources, "load_workbook", lambda *a, **kw: pytest.fail("Oversize package reached openpyxl."))
    result, notices = sources.ingest_batch(profile, (("synthetic.xlsx", raw),))
    assert result == () and any(expected in notice for notice in notices)
    assert not sources._path(profile, hashlib.sha256(raw).hexdigest(), ".xlsx").exists()


@pytest.mark.parametrize("count", [1, 2])
def test_capacity_page_budget_is_separate_from_evidence_ingestion(profile, monkeypatch, count):
    monkeypatch.setattr(capacity, "MAX_INPUT_ITEMS", 1)
    monkeypatch.setattr(sources, "MAX_PAGES", 2)
    raw = pdf_bytes(count)
    if count == 1:
        result = capacity.inspect_capacity(profile, "synthetic.pdf", raw)
        assert len(result.content.source.page_texts) == 1
    else:
        with pytest.raises(ValueError, match="1 extracted pages/items"):
            capacity.inspect_capacity(profile, "synthetic.pdf", raw)
    assert capacity.load_capacity(profile.contract) is None
    assert sources._path(profile, hashlib.sha256(raw).hexdigest(), ".pdf").read_bytes() == raw
