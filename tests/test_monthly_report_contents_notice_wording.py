"""Generated reports say "this report" in the master's confidentiality notice (Refs #304).

Owner decision, 21 October 2026: the contents-page notice copied from the
master design said "within this proposal". Only that word changes, in every
retained copy, and a report carrying the new wording is still recognized when
it is imported as next month's Starting Report. Synthetic values only.
"""
from dataclasses import replace
from io import BytesIO
from zipfile import ZipFile

from docx import Document
from lxml import etree

from app import monthly_report_designs as designs
from app.monthly_report_docx import _build_docx
from app.monthly_report_model import ReportPeriod, default_sections, synthetic_draft, synthetic_profiles
from app.monthly_report_native_design import (A, CONTENTS_NOTICE, REPORT_NOTICE, W, WPS,
                                              name_report_in_contents_notice, native_layout_role)

NS = f'xmlns:w="{W[1:-1]}" xmlns:a="{A[1:-1]}" xmlns:wps="{WPS[1:-1]}"'
ADDRESS = "1 Galleria Blvd., Synthetic Suite | enfrasolutions.com"
SPLIT = CONTENTS_NOTICE.index("this pro") + len("this pro")


def _runs(*parts, italic_last=False):
    runs = []
    for index, text in enumerate(parts):
        style = "<w:rPr><w:i/></w:rPr>" if italic_last and index == len(parts) - 1 else ""
        runs.append(f'<w:r>{style}<w:t xml:space="preserve">{text}</w:t></w:r>')
    return "<w:p>" + "".join(runs) + "</w:p>"


def _text(paragraph):
    return "".join(node.text or "" for node in paragraph.iter(W + "t"))


def test_every_copy_of_the_notice_names_a_report_and_keeps_its_runs():
    split = _runs(CONTENTS_NOTICE[:SPLIT], CONTENTS_NOTICE[SPLIT:], italic_last=True)
    root = etree.fromstring(
        f'<w:body {NS}><w:txbxContent>{split}</w:txbxContent>'
        f'<w:txbxContent>{_runs(" " + CONTENTS_NOTICE)}</w:txbxContent>'
        f'{_runs("Include this proposal in the review.")}{_runs(CONTENTS_NOTICE + " Extra synthetic text.")}</w:body>')
    name_report_in_contents_notice(root)
    first, fallback, other, longer = root.iter(W + "p")
    assert _text(first) == REPORT_NOTICE and len(first.findall(W + "r")) == 2
    assert first.findall(W + "r")[1].find(W + "rPr/" + W + "i") is not None
    assert _text(fallback) == " " + REPORT_NOTICE
    assert _text(other) == "Include this proposal in the review."
    assert _text(longer) == CONTENTS_NOTICE + " Extra synthetic text."


def _contents_frame(notice):
    lines = "".join(_runs(line) for line in ("TABLE OF CONTENTS", notice, ADDRESS))
    return etree.fromstring(
        f'<w:p {NS}><w:r><a:graphic><a:graphicData uri="{WPS[1:-1]}"><wps:wsp>'
        f'<wps:spPr><a:prstGeom prst="rect"/></wps:spPr><wps:txbx><w:txbxContent>{lines}</w:txbxContent></wps:txbx>'
        f'<wps:bodyPr/></wps:wsp></a:graphicData></a:graphic></w:r></w:p>')


def test_a_report_with_the_new_wording_is_still_recognized_on_import():
    for notice in (CONTENTS_NOTICE, REPORT_NOTICE):
        text = "\n".join(("TABLE OF CONTENTS", notice, ADDRESS))
        role = native_layout_role(_contents_frame(notice), text, part="word/document.xml",
                                  section="cover", word_section=2)
        assert role == "Native contents notice and artwork"


def test_generated_master_report_names_itself_a_report(monkeypatch, tmp_path):
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path / "data"))
    path = tmp_path / "master.docx"
    document = Document()
    document.add_paragraph("Operations and Maintenance Monthly Review")
    for section in default_sections():
        document.add_heading(section.title, 1)
        document.add_paragraph("Source-only content")
    document.save(path)
    body = etree.fromstring(ZipFile(path).read("word/document.xml"))
    first = body.find(W + "body").find(W + "p")
    first.addnext(_contents_frame(CONTENTS_NOTICE))
    with ZipFile(path) as source:
        parts = {item: source.read(item) for item in source.namelist()}
    parts["word/document.xml"] = etree.tostring(body, xml_declaration=True, encoding="UTF-8", standalone=True)
    with ZipFile(path, "w") as target:
        for name, raw in parts.items():
            target.writestr(name, raw)
    designs.install_master(path, actor="Synthetic Editor", expected_revision=0)
    draft = synthetic_draft(designs.pin(replace(synthetic_profiles()[0], contract="Synthetic Client")),
                            ReportPeriod(2026, 9))
    output = ZipFile(BytesIO(_build_docx(draft))).read("word/document.xml").decode("utf-8")
    assert "within this report" in output and "within this proposal" not in output
