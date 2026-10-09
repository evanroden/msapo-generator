"""Legacy page-backed divider canvases stay artwork, not repeated report photos."""
from io import BytesIO
from hashlib import sha256

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches
from lxml import etree
from PIL import Image
import pytest

from app.monthly_report_import import _is_native_divider_background, inspect_docx, read_import_image
from app.monthly_report_model import Facility, ReportDraft, ReportPeriod, ReportProfile, ResolvedBlock, default_sections
from app.monthly_report_native_layout import build_native_docx

V = '{urn:schemas-microsoft-com:vml}'


def fixture(tmp_path):
    doc = Document()
    doc.add_heading('Organizational Chart', 1)
    pictures = []
    for colour in ('blue', 'red', 'purple'):
        data = BytesIO(); Image.new('RGB', (50, 30), colour).save(data, 'PNG')
        pictures.append(data.getvalue())
    doc.add_paragraph().add_run().add_picture(BytesIO(pictures[0]), width=Inches(2))
    doc.add_section(WD_SECTION_START.NEW_PAGE)
    paragraph = doc.add_paragraph()
    pict = OxmlElement('w:pict'); paragraph.add_run()._r.append(pict)
    group = etree.SubElement(pict, V + 'group', style='position:absolute;width:612pt;height:791.3pt;mso-position-horizontal-relative:page;mso-position-vertical-relative:page')
    etree.SubElement(group, V + 'rect', fillcolor='#527c7c', style='width:612pt;height:100pt')
    etree.SubElement(group, V + 'rect', fillcolor='#d4ed49', style='width:612pt;height:30pt')
    for raw in pictures[1:]:
        rid, _ = doc.part.get_or_add_image(BytesIO(raw))
        shape = etree.SubElement(group, V + 'shape', style='width:100pt;height:60pt')
        etree.SubElement(shape, V + 'imagedata', {qn('r:id'): rid})
    doc.add_paragraph('1.')
    doc.add_paragraph('2.')
    for _ in range(50):
        doc.add_paragraph()
    doc.add_paragraph('MONTHLY ACTIVITY')
    doc.add_paragraph('SUMMARY')
    path = tmp_path / 'legacy.docx'; doc.save(path)
    return path


def test_vml_art_before_split_heading_keeps_correct_section_after_long_blank_gap(tmp_path):
    inspection = inspect_docx(fixture(tmp_path))
    artwork = [item for item in inspection.items if item.image_part and item.position > 2]
    assert len(artwork) == 2
    assert all(item.section == 'activity' and item.suggested_slot == 'divider_activity'
               and item.note == 'Native divider artwork' for item in artwork)
    fragments = [item for item in inspection.items if item.text in ('MONTHLY ACTIVITY', 'SUMMARY')]
    assert len(fragments) == 2 and all(item.note == 'Native section heading' for item in fragments)


@pytest.mark.parametrize('green', ['#527c7c', '#537d7d', '#547e7e', '#557f7f'])
def test_legacy_green_palette_variants_keep_the_same_closed_divider_geometry(tmp_path, green):
    doc = Document(fixture(tmp_path))
    paragraph = next(p for p in doc.element.iter(qn('w:p')) if p.find('.//' + V + 'group') is not None)
    for node in paragraph.iter():
        if node.get('fillcolor') == '#527c7c': node.set('fillcolor', green)
    assert _is_native_divider_background(paragraph)


@pytest.mark.parametrize('payload', ['Private client narrative', '$500', '顧客情報'])
def test_narrative_between_art_and_title_is_not_reclassified_as_decoration(tmp_path, payload):
    path = fixture(tmp_path)
    doc = Document(path)
    paragraph = next(p for p in doc.paragraphs if p.text == '1.')
    paragraph.text = payload
    doc.save(path)
    assert not any(item.note == 'Native divider artwork' for item in inspect_docx(path).items
                   if item.position < 20)


@pytest.mark.parametrize('change', ['ordinary_photo', 'small', 'paragraph_relative', 'private_text', 'wordart', 'opaque'])
def test_vml_background_requires_closed_native_art_context(tmp_path, change):
    doc = Document(fixture(tmp_path))
    paragraph = next(p for p in doc.element.iter(qn('w:p')) if p.find('.//' + V + 'group') is not None)
    group = paragraph.find('.//' + V + 'group')
    if change == 'ordinary_photo':
        for node in group.iter(): node.attrib.pop('fillcolor', None)
    elif change == 'small':
        group.set('style', group.get('style').replace('612pt', '100pt'))
    elif change == 'paragraph_relative':
        group.set('style', group.get('style').replace('relative:page', 'relative:paragraph'))
    elif change == 'wordart':
        etree.SubElement(group, V + 'textpath', string='Private client name')
    elif change == 'opaque':
        etree.SubElement(group, qn('w:object'))
    assert not _is_native_divider_background(paragraph, 'Private client note' if change == 'private_text' else '')


def test_old_org_mapping_does_not_duplicate_native_divider_canvas(tmp_path):
    path = fixture(tmp_path); inspection = inspect_docx(path)
    images = [item for item in inspection.items if item.kind == 'image']
    numerals = [item for item in inspection.items if item.text in ('1.', '2.')]
    assets = {}
    for item in images:
        image = read_import_image(path, item)
        assets[sha256(image.data).hexdigest() + '.' + image.extension] = image.data
    # A persisted draft from the previous inspector mapped the entire canvas
    # into org_chart; provenance still binds these exact immutable source images.
    block = ResolvedBlock('org_chart', 'This month', text='1.\n\n2.', asset_hashes=tuple(assets),
                          asset_captions=('',) * len(assets),
                          references=tuple(f'docx:{inspection.sha256}:{item.id}' for item in images + numerals))
    draft = ReportDraft(ReportProfile('Synthetic', 'site', 'Current site', (Facility('site', 'Current site'),)),
                        ReportPeriod(2026, 9), 'Current editor', default_sections(), (block,))
    result = Document(BytesIO(build_native_docx(draft, path, asset_loader=assets.__getitem__)))
    source = Document(path)
    assert len(result.paragraphs) == len(source.paragraphs)
    assert len(list(result.element.iter(V + 'group'))) == 1
    assert len(list(result.element.iter(V + 'imagedata'))) == 2
