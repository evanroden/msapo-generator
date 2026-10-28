"""Standalone floating Word text stays in layout review, not editable prose.

This differs from PR #506's table-contained floating-drawing boundary.
Only synthetic words and generated DOCX packages are used here.
"""
from pathlib import Path

from docx import Document
from docx.oxml import parse_xml
import pytest

from app import monthly_report_import as importer


def _frame(text, kind="vml"):
    if kind == "vml":
        return parse_xml(
            '<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
            'xmlns:v="urn:schemas-microsoft-com:vml">'
            '<w:pict><v:shape id="synthetic_1"><v:textbox><w:txbxContent>'
            '<w:p><w:r><w:t>' + text + '</w:t></w:r></w:p>'
            '</w:txbxContent></v:textbox></v:shape></w:pict></w:r>'
        )
    return parse_xml(
        '<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
        'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
        'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">'
        '<w:drawing><wp:anchor><a:graphic><a:graphicData '
        'uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">'
        '<wps:wsp><wps:txbx><w:txbxContent><w:p><w:r><w:t>'
        + text +
        '</w:t></w:r></w:p></w:txbxContent></wps:txbx></wps:wsp>'
        '</a:graphicData></a:graphic></wp:anchor></w:drawing></w:r>'
    )


def _report(tmp_path, *, kind="vml", outside="", labels=("Synthetic framed detail",)):
    document = Document()
    document.add_heading("Monthly Activity Summary", 1)
    paragraph = document.add_paragraph(outside)
    for label in labels:
        paragraph._p.append(_frame(label, kind))
    target = tmp_path / "synthetic.docx"
    document.save(target)
    original = target.read_bytes()
    inspection = importer.inspect_docx(target)
    assert target.read_bytes() == original
    return inspection


def _items(inspection, kind):
    return [item for item in inspection.items
            if item.part == "word/document.xml" and item.kind == kind]


@pytest.mark.parametrize("kind", ["vml", "modern"])
def test_standalone_textbox_is_review_only_not_a_plain_paragraph(tmp_path, kind):
    marker = "Synthetic boxed task without ordinary paragraph words"
    inspection = _report(tmp_path, kind=kind, labels=(marker,))
    assert not any(marker in item.text for item in _items(inspection, "text"))
    reviews = _items(inspection, "unsupported")
    assert any(marker in item.text for item in reviews)
    assert any("review" in item.note.lower() or "drawing" in item.note.lower()
               for item in reviews if marker in item.text)


@pytest.mark.parametrize("kind", ["vml", "modern"])
def test_real_sibling_paragraph_text_survives_without_box_leak(tmp_path, kind):
    marker = "Synthetic framed work detail"
    outside = "Visible synthetic service status remains editable"
    inspection = _report(tmp_path, kind=kind, outside=outside, labels=(marker,))
    paragraphs = _items(inspection, "text")
    assert any(item.text == outside for item in paragraphs)
    assert not any(marker in item.text for item in paragraphs)
    assert any(marker in item.text for item in _items(inspection, "unsupported"))


def test_two_textboxes_do_not_enter_the_plain_text_mapping(tmp_path):
    markers = ("Synthetic first floating note", "Synthetic second floating note")
    outside = "Visible synthetic line"
    inspection = _report(tmp_path, outside=outside, labels=markers)
    paragraph_items = _items(inspection, "text")
    assert any(item.text == outside for item in paragraph_items)
    assert all(not any(marker in item.text for item in paragraph_items) for marker in markers)
    assert all(any(marker in item.text for item in _items(inspection, "unsupported"))
               for marker in markers)


def test_ordinary_unframed_paragraph_stays_editable(tmp_path):
    document = Document()
    document.add_heading("Monthly Activity Summary", 1)
    document.add_paragraph("Synthetic regular work summary")
    target = tmp_path / "plain.docx"
    document.save(target)
    before = target.read_bytes()
    inspected = importer.inspect_docx(target)
    assert target.read_bytes() == before
    assert any(item.text == "Synthetic regular work summary"
               for item in _items(inspected, "text"))


def test_native_section_title_inside_frame_keeps_legacy_heading_route(tmp_path):
    document = Document()
    document.add_paragraph()._p.append(_frame("Monthly Activity Summary"))
    document.add_paragraph("Synthetic plain source note")
    target = tmp_path / "native_title.docx"
    document.save(target)
    inspected = importer.inspect_docx(target)
    # Section/title recognition is not reinterpreted by this paragraph-only
    # safeguard. The floating drawing remains independently reviewable.
    assert any(item.kind == "text" and item.text == "Monthly Activity Summary"
               for item in inspected.items)
    assert any(item.kind == "unsupported" and "Monthly Activity Summary" in item.text
               for item in inspected.items)
