"""The hosted row journey cannot accept a missing row, edit, save or output."""
from copy import deepcopy
from dataclasses import replace
from io import BytesIO
import json
from pathlib import Path
import subprocess
import sys
from types import ModuleType
from unittest.mock import Mock

from docx import Document
import pymupdf
import pytest
from streamlit.testing.v1 import AppTest
from conftest import requires_libreoffice

from scripts import qa_work_order_rows as qa

ROOT = Path(__file__).resolve().parents[1]
FIXTURE = ROOT / "tests/fixtures/workspaces/work_order_rows_app.py"


@pytest.mark.parametrize("change,dirty", [
    ("evidence", False), ("tracked", True), ("staged", True),
    ("untracked-source", True), ("other-evidence", True), ("tracked-evidence", True),
])
def test_checkout_provenance_excludes_only_this_drivers_generated_evidence(tmp_path, monkeypatch, change, dirty):
    def git(*args):
        return subprocess.run(["git", *args], cwd=tmp_path, check=True, capture_output=True)
    git("init", "-q")
    source = tmp_path / "source.py"
    source.write_text("original\n")
    evidence = tmp_path / "test-artifacts" / "work-order-rows"
    evidence.mkdir(parents=True)
    baseline = evidence / "tracked-baseline.txt"
    baseline.write_text("original\n")
    git("add", ".")
    git("-c", "user.name=Synthetic Test", "-c", "user.email=test@example.invalid",
        "commit", "-qm", "Synthetic baseline")
    (evidence / "streamlit.log").write_text("Synthetic evidence\n")
    monkeypatch.setattr(qa, "ROOT", tmp_path)
    monkeypatch.setattr(qa, "OUT", evidence)
    if change in {"tracked", "staged"}:
        source.write_text("changed\n")
        if change == "staged":
            git("add", "source.py")
    elif change == "untracked-source":
        (tmp_path / "new_source.py").write_text("new\n")
    elif change == "other-evidence":
        (evidence.parent / "unrelated.txt").write_text("unowned evidence\n")
    elif change == "tracked-evidence":
        baseline.write_text("changed\n")
    assert qa.source_checkout_dirty() is dirty


def mark(monkeypatch, tmp_path):
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path))
    monkeypatch.setenv("EPC_MONTHLY_DOM_FIXTURE", "1")
    monkeypatch.setenv("EPC_WORK_ORDER_ROW_FIXTURE", "1")
    (tmp_path / "monthly-dom-synthetic-only.marker").write_text("synthetic-monthly-dom-v1")


def test_source_and_compact_fixture_are_distinct_and_stable(tmp_path):
    path, draft = qa.scenario(tmp_path)
    original = path.read_bytes()
    assert len(Document(path).tables[0].rows) == 8
    assert len(qa.row_state(draft)["rows"]) == 5  # Native heading + 5 data rows = 6.
    assert qa.scenario(tmp_path) == (path, draft)
    assert path.read_bytes() == original


@pytest.mark.parametrize("guard", ["marker", "flag", "token", "query"])
def test_fixture_rejects_unmarked_or_unscoped_entry_before_data_writes(monkeypatch, tmp_path, guard):
    mark(monkeypatch, tmp_path)
    if guard == "marker":
        (tmp_path / "monthly-dom-synthetic-only.marker").unlink()
    elif guard == "flag":
        monkeypatch.delenv("EPC_WORK_ORDER_ROW_FIXTURE")
    app = AppTest.from_file(FIXTURE)
    if guard == "query":
        app.query_params["report"] = "forbidden"
    app.run()
    assert len(app.exception) == 1
    assert not list(tmp_path.glob("monthly_reports/**/*"))
    assert not list(tmp_path.glob("*.docx"))


def test_fixture_uses_real_editor_and_disk_save_with_no_browser(monkeypatch, tmp_path):
    import streamlit as st
    from app import monthly_report_active_work as active
    from test_work_order_row_insertion import wire, ABOVE, BELOW
    mark(monkeypatch, tmp_path)
    # Unit boundary only; the hosted driver supplies an actual browser cookie.
    monkeypatch.setattr(type(st.context), "cookies", property(lambda self: {"epc_device_id": "c" * 32}))
    app = AppTest.from_file(FIXTURE).run()
    assert not app.exception
    config = json.loads(app.dataframe[0].proto.columns)
    action = config["__work_order_row_action"]
    assert action["width"] == 160 and action["pinned"] is True
    assert st.__version__ == "1.61.1"  # Driver coordinates target this locked native renderer.
    initial = app.session_state.draft
    wire(app, 0, ABOVE, edited={0: {"Work orders": qa.EDIT}})
    wire(app, 3, BELOW)
    changed = app.session_state.draft
    loaded = active.load("c" * 32, changed.profile.contract, changed.profile.key, changed.period)
    assert loaded.draft == changed
    assert qa.row_state(changed)["rows"] == qa.expected_rows("restored")
    assert qa.row_state(changed)["unrelated"] == qa.row_state(initial)["unrelated"]
    reopened = AppTest.from_file(FIXTURE).run()
    assert not reopened.exception
    assert reopened.session_state.draft == changed



def test_fixture_continues_after_a_run_stops_between_its_disk_save_and_acknowledgement(monkeypatch, tmp_path):
    """WebKit delivers the cell edit and the row action as two runs.

    Streamlit yields to a pending rerun on every session-state access, so the
    first run can stop after its journal write but before recording it.
    """
    import streamlit as st
    from streamlit.runtime.scriptrunner_utils.exceptions import StopException
    from app import monthly_report_active_work as active
    from test_work_order_row_insertion import wire, ABOVE
    mark(monkeypatch, tmp_path)
    monkeypatch.setattr(type(st.context), "cookies", property(lambda self: {"epc_device_id": "c" * 32}))
    app = AppTest.from_file(FIXTURE).run()
    assert not app.exception
    initial = active.load("c" * 32, app.session_state.draft.profile.contract,
                          app.session_state.draft.profile.key, app.session_state.draft.period)
    real_save, stopped = active.save, []

    def stop_after_write(*args, **kwargs):
        saved = real_save(*args, **kwargs)
        if not stopped:
            stopped.append(saved.revision)
            raise StopException()  # The session never records this revision.
        return saved

    monkeypatch.setattr(active, "save", stop_after_write)
    wire(app, edited={0: {"Work orders": qa.EDIT}})
    assert stopped == [initial.revision + 1]
    wire(app, 0, ABOVE)
    draft = app.session_state.draft
    durable = active.load("c" * 32, draft.profile.contract, draft.profile.key, draft.period)
    assert qa.row_state(draft)["rows"] == qa.expected_rows("above")
    assert durable.draft == draft and durable.generation == initial.generation


def receipt():
    state = {"columns": list(qa.COLUMNS), "rows": qa.expected_rows("restored"),
             "reference": "synthetic:source", "fingerprint": "same", "unrelated": "same-other"}
    return {"working": state, "durable": deepcopy(state), "revision": 3,
            "generation": "generation", "source_unchanged": True}


def test_matching_native_and_disk_receipt_passes():
    qa.assert_receipt(receipt(), qa.expected_rows("restored"), receipt())


@pytest.mark.parametrize("fault", ["missing_row", "lost_edit", "helper_column", "disk", "source", "reference", "unrelated", "generation"])
def test_receipt_rejects_each_lost_or_unproved_result(fault):
    before = receipt()
    value = deepcopy(before)
    if fault == "missing_row":
        value["working"]["rows"].pop(0)
    elif fault == "lost_edit":
        value["working"]["rows"][1][0] = "Received"
    elif fault == "helper_column":
        value["working"]["columns"].append("__work_order_row_id")
    elif fault == "disk":
        value["durable"]["fingerprint"] = "old"
    elif fault == "source":
        value["source_unchanged"] = False
    elif fault == "generation":
        value["generation"] = "other"
    else:
        value["working"][fault] = "other"
        value["durable"][fault] = "other"
    if fault in {"missing_row", "lost_edit", "helper_column"}:
        # Matching persistence is not enough: prove the row/schema oracle,
        # rather than the earlier working-versus-disk equality guard, rejects it.
        value["durable"] = deepcopy(value["working"])
        assert value["working"] == value["durable"]
    with pytest.raises(AssertionError):
        qa.assert_receipt(value, qa.expected_rows("restored"), before)


def fake_outputs(*, remove_row=False, helper=False, lost_edit=False, wrong_pdf=False):
    rows = qa.expected_rows("restored")
    if remove_row:
        rows.pop(0)
    if lost_edit:
        rows[1][0] = "Received"
    columns = list(qa.COLUMNS)
    if helper:
        columns.append("__work_order_row_id")
        rows = [row + ["private-ui-id"] for row in rows]
    document = Document()
    table = document.add_table(rows=len(rows) + 1, cols=len(columns))
    for row, values in zip(table.rows, (columns, *rows)):
        for cell, value in zip(row.cells, values):
            cell.text = value
    stream = BytesIO()
    document.save(stream)
    # Geometry is genuine PDF vector content; no external renderer dependency.
    with pymupdf.open() as pdf:
        page = pdf.new_page()
        for n in range(8 if wrong_pdf else 9):
            page.draw_line((40, 40 + n * 30), (530, 40 + n * 30))
        for n in range(8):
            page.draw_line((40 + n * 70, 40), (40 + n * 70, 250 if wrong_pdf else 280))
        for n in range(7):
            page.insert_text((45 + n * 70, 60), str(n))
        page.insert_text((45, 340), qa.EDIT)
        pdf_bytes = pdf.tobytes()
    return stream.getvalue(), pdf_bytes


def test_output_oracle_checks_native_cells_and_actual_pdf_geometry():
    assert qa.assert_outputs(*fake_outputs(), qa.expected_rows("restored"))["native_table"] == [8, 7]


@requires_libreoffice
def test_hosted_output_oracle_accepts_real_native_renderer_bytes(monkeypatch, tmp_path):
    from app import pdf_converter
    from app.monthly_report_native_layout import build_native_docx
    source, draft = qa.scenario(tmp_path)
    block, other = draft.blocks
    table, = block.extra_tables
    block = replace(block, extra_tables=(replace(table, rows=tuple(
        tuple(row) for row in qa.expected_rows("restored"))),))
    word = build_native_docx(replace(draft, blocks=(block, other)), source, section_key="activity")
    output = tmp_path / "native-output.docx"
    output.write_bytes(word)
    monkeypatch.setattr(pdf_converter, "OUTPUT_DIR", tmp_path)
    monkeypatch.setattr(pdf_converter, "PDF_BACKEND", "libreoffice")
    pdf_path = pdf_converter.convert_to_pdf(output)
    try:
        assert qa.assert_outputs(word, pdf_path.read_bytes(), qa.expected_rows("restored"))["native_table"] == [8, 7]
    finally:
        pdf_path.unlink(missing_ok=True)


@pytest.mark.parametrize("fault", [{"remove_row": True}, {"helper": True}, {"lost_edit": True}, {"wrong_pdf": True}])
def test_output_oracle_rejects_missing_blank_edit_or_helper_data(fault):
    with pytest.raises(AssertionError):
        qa.assert_outputs(*fake_outputs(**fault), qa.expected_rows("restored"))


def test_new_ci_job_preserves_existing_budgets_and_requires_both_engines():
    from test_ci_native_test_budget import parse_jobs, validate_workflow
    text = (ROOT / ".github/workflows/adaptive-workspaces.yml").read_text()
    validate_workflow(text)
    job = "".join(parse_jobs(text)["synthetic-work-order-rows"])
    for required in ("timeout-minutes: 45", "timeout-minutes: 20", "timeout-minutes: 11",
                     "--minutes 10", "python scripts/qa_work_order_rows.py", "if: always()",
                     "timeout-minutes: 3", "playwright==1.58.0", "chromium webkit"):
        assert required in job
    assert "continue-on-error" not in job
    assert text.count("- 'scripts/qa_work_order_rows.py'") == 2
    driver = (ROOT / "scripts/qa_work_order_rows.py").read_text()
    for forbidden in ("dispatch_event(", "setStringTriggerValue", "_run(message)", "force=True"):
        assert forbidden not in driver
    assert '"chromium", playwright.chromium' in driver and '"webkit", playwright.webkit' in driver


def test_server_shutdown_is_bounded():
    class Slow:
        calls = []
        def terminate(self): self.calls.append("terminate")
        def kill(self): self.calls.append("kill")
        def wait(self, timeout):
            self.calls.append(("wait", timeout))
            if timeout == 10:
                raise subprocess.TimeoutExpired("synthetic", timeout)
    process = Slow()
    qa.stop_server(process)
    assert process.calls == ["terminate", ("wait", 10), "kill", ("wait", 5)]


class NativeGrid:
    """Locator-only contract; no local browser is launched."""
    def __init__(self):
        self.calls = []
        self.counts = {"grid": 1, "canvas": 1, "surface": 1}
        self.boxes = {kind: {"x": 80, "y": 208, "width": 1038, "height": 245}
                      for kind in ("canvas", "surface")}
        self.hit = True

    def get_by_test_id(self, testid):
        assert testid == "stDataFrame"
        return self.Node(self, "grid")

    def evaluate(self, expression):
        assert expression == "parseFloat(getComputedStyle(document.documentElement).fontSize)"
        return 16

    class Node:
        def __init__(self, owner, kind):
            self.owner, self.kind = owner, kind

        def get_by_test_id(self, testid):
            assert self.kind == "grid" and testid == "data-grid-canvas"
            return self.__class__(self.owner, "canvas")

        def locator(self, selector):
            assert self.kind == "grid" and selector == ".dvn-scroller.stDataFrameGlideDataEditor"
            return self.__class__(self.owner, "surface")

        def count(self): return self.owner.counts[self.kind]
        def bounding_box(self): return self.owner.boxes[self.kind]
        def scroll_into_view_if_needed(self): self.owner.calls.append(("scroll", self.kind))

        def wait_for(self, **options):
            assert options == {"state": "visible", "timeout": 15000}
            self.owner.calls.append(("visible", self.kind))

        def evaluate(self, expression, point):
            assert self.kind == "surface"
            assert "elementFromPoint" in expression and "node.contains(hit)" in expression
            self.owner.calls.append(("hit", point))
            return self.owner.hit


@pytest.mark.parametrize("action,x", [(True, 112), (False, 217)])
def test_native_grid_targets_input_surface_not_visual_canvas(action, x):
    page = NativeGrid()
    # Translate coordinates; do not assume the two DOM elements share an origin.
    page.boxes["surface"]["x"] -= 2
    page.boxes["surface"]["y"] -= 3
    target, point = qa.native_grid_target(page, 0, action=action)
    assert target.kind == "surface"
    assert point == {"x": x + 2, "y": 55.5}
    assert page.calls == [("visible", kind) for kind in ("grid", "canvas", "surface")] + [
        ("scroll", "surface"), ("hit", point)]


@pytest.mark.parametrize("fault", ["missing", "ambiguous", "hidden", "clipped", "covered"])
def test_native_grid_refuses_unproved_pointer_target(fault):
    page = NativeGrid()
    if fault == "missing":
        page.counts["canvas"] = 0
    elif fault == "ambiguous":
        page.counts["surface"] = 2
    elif fault == "hidden":
        page.boxes["surface"] = None
    elif fault == "clipped":
        page.boxes["surface"]["height"] = 40
    else:
        page.hit = False
    with pytest.raises(AssertionError):
        qa.native_grid_target(page, 0)


@pytest.mark.parametrize("visible", [False, True])
@pytest.mark.parametrize("fault", [None, "input count", "initial state", "label count", "label inputs", "hidden", "final state"])
def test_visibility_toggle_uses_one_proved_native_label(monkeypatch, visible, fault):
    page, checkbox, label, inputs = (Mock() for _ in range(4))
    page.get_by_role.return_value = checkbox
    checkbox.locator.return_value = label
    label.get_by_role.return_value = inputs
    states = []

    def expectation(node):
        result = Mock()
        def count(number):
            assert number == 1
            reason = {id(checkbox): "input count", id(label): "label count", id(inputs): "label inputs"}[id(node)]
            if fault == reason:
                raise AssertionError(reason)
        def checked(*, checked):
            assert node is checkbox
            assert checked is (not visible if not states else visible)
            if fault == "initial state" and not states:
                raise AssertionError(fault)
            if fault == "final state" and states:
                raise AssertionError(fault)
            states.append(checked)
        def shown():
            assert node is label
            if fault == "hidden":
                raise AssertionError(fault)
        result.to_have_count.side_effect = count
        result.to_be_checked.side_effect = checked
        result.to_be_visible.side_effect = shown
        return result

    # Pytest's locked dependencies deliberately omit browser tooling. Stub only
    # the assertion API; this contract must not need Playwright installation.
    browser_api = ModuleType("playwright.sync_api")
    browser_api.expect = expectation
    monkeypatch.setitem(sys.modules, "playwright", ModuleType("playwright"))
    monkeypatch.setitem(sys.modules, "playwright.sync_api", browser_api)
    if fault:
        with pytest.raises(AssertionError, match=fault):
            qa.toggle_table_visibility(page, visible=visible)
        if fault == "final state":
            label.click.assert_called_once_with()
        else:
            label.click.assert_not_called()
    else:
        qa.toggle_table_visibility(page, visible=visible)
        label.click.assert_called_once_with()
        assert states == [not visible, visible]
        checkbox.locator.assert_called_once_with("xpath=ancestor::label")
        label.get_by_role.assert_called_once_with("checkbox")
    page.get_by_role.assert_called_once_with("checkbox", name="Show synthetic table", exact=True)
    checkbox.click.assert_not_called()
    checkbox.check.assert_not_called()
    checkbox.uncheck.assert_not_called()
    checkbox.press.assert_not_called()
