from dataclasses import replace
from io import BytesIO
from hashlib import sha256

import pytest
from docx import Document
from docx.shared import Inches
from docx.oxml.ns import qn
from PIL import Image

from app.monthly_report_asset_layout import document_page_indexes, picture_runs
from app.monthly_report_model import ReportSource, ResolvedBlock


def image(size, color):
    out = BytesIO(); Image.new('RGB', size, color).save(out, 'PNG')
    raw = out.getvalue()
    return sha256(raw).hexdigest()+'.png', raw


def payload():
    assets = dict((image((600,180),'navy'), image((850,1100),'white'), image((400,200),'green')))
    refs = tuple(assets)
    source = ReportSource('pdf-source', 'report.pdf', 'a'*64, '.pdf',
                          page_texts=('Synthetic service evidence',), selected_pages=(1,))
    reference = f'{source.id}:1:{source.fingerprint}'
    block = ResolvedBlock('improvements','This month',asset_hashes=refs,
                          asset_captions=('Photo A','Service page','Photo B'),
                          references=(reference,),asset_provenance=((refs[1],(reference,)),),photos_per_page=3)
    return block, source, assets


def test_exact_mapping_only_and_ordered_mixed_runs():
    block, source, _ = payload()
    assert document_page_indexes(block,(source,)) == {1}
    assert tuple(picture_runs(block,(source,))) == ((False,(0,)),(True,(1,)),(False,(2,)))
    assert document_page_indexes(replace(block, asset_hashes=tuple(reversed(block.asset_hashes))),(source,)) == {1}


@pytest.mark.parametrize('change', ['missing','other_asset','duplicate_map','duplicate_reference','outside_block','stale_source','jpeg','unselected','out_of_range','marker_only','asset_changed','duplicate_source'])
def test_untrusted_or_missing_mapping_does_not_expand(change):
    block, source, _ = payload(); sources=(source,)
    ref=block.references[0]; asset=block.asset_hashes[1]
    if change=='missing': block=replace(block,asset_provenance=())
    elif change=='other_asset': block=replace(block,asset_provenance=(('0'*64+'.png',(ref,)),))
    elif change=='duplicate_map': block=replace(block,asset_provenance=block.asset_provenance*2)
    elif change=='duplicate_reference': block=replace(block,asset_provenance=((asset,(ref,ref)),))
    elif change=='outside_block': block=replace(block,references=())
    elif change=='stale_source': sources=(replace(source,service_date='2026-10-02'),)
    elif change=='jpeg': sources=(replace(source,suffix='.jpg'),)
    elif change=='unselected': sources=(replace(source,selected_pages=()),)
    elif change=='out_of_range':
        ref=f'{source.id}:2:{source.fingerprint}';block=replace(block,references=(ref,),asset_provenance=((asset,(ref,)),))
    elif change=='marker_only': block=replace(block,asset_provenance=((asset,('asset-source-context:v1:pdf-source:1:'+'a'*64,)),))
    elif change=='asset_changed': block=replace(block,asset_hashes=(block.asset_hashes[0],'b'*64+'.png',block.asset_hashes[2]))
    elif change=='duplicate_source': sources=(source,source)
    assert document_page_indexes(block,sources)==set()


def test_native_mixed_pages_keep_photo_frames_order_captions_and_input(tmp_path):
    from tests.test_monthly_report_native_layout import draft
    from app.monthly_report_native_layout import build_native_docx
    block, source, assets=payload()
    document=Document();document.add_heading('MONTHLY ACTIVITY SUMMARY',1)
    document.add_paragraph('Improvement Highlights and Pics')
    document.add_picture(BytesIO(next(iter(assets.values()))),width=Inches(6),height=Inches(1.8))
    document.add_picture(BytesIO(next(iter(assets.values()))),width=Inches(4),height=Inches(2))
    path=tmp_path/'source.docx';document.save(path)
    current=replace(draft(block),sources=(source,));before=current.fingerprint
    for section in (None,'activity',None):
        raw=build_native_docx(current,path,asset_loader=assets.__getitem__,section_key=section)
        output=Document(BytesIO(raw)); shapes=list(output.inline_shapes)
        assert len(shapes)==3
        assert shapes[0].width==Inches(6) and shapes[0].height==Inches(1.8)
        assert shapes[1].width>Inches(5) and shapes[1].height>Inches(6)
        assert shapes[2].width==Inches(4) and shapes[2].height==Inches(2)
        pictures=[p for p in output.paragraphs if p._p.find('.//'+qn('wp:inline')) is not None]
        assert pictures[1].paragraph_format.page_break_before
        assert pictures[2].paragraph_format.page_break_before
        assert [p.text for p in output.paragraphs if p.text in block.asset_captions]==list(block.asset_captions)
        assert current.fingerprint==before
        assert [sha256(output.part.related_parts[s._inline.graphic.graphicData.pic.blipFill.blip.embed].blob).hexdigest()+'.png' for s in shapes]==list(block.asset_hashes)


def test_generic_mixed_run_emits_original_document_not_grid_and_keeps_receipts():
    from tests.test_monthly_report_native_layout import draft
    from app.monthly_report_docx import _build_docx, outline
    block, source, assets=payload()
    current=replace(draft(block),sources=(source,));before=current.fingerprint
    raw=_build_docx(current,asset_loader=assets.__getitem__)
    output=Document(BytesIO(raw))
    matches=[s for s in output.inline_shapes if output.part.related_parts[s._inline.graphic.graphicData.pic.blipFill.blip.embed].blob==assets[block.asset_hashes[1]]]
    assert len(matches)==1 and matches[0].width>Inches(5) and matches[0].height>Inches(6)
    assert current.fingerprint==before
    assert outline(current)[0][2]==4  # divider estimate + photo / document / photo
    assert 'Service page' in [p.text for p in output.paragraphs]


def test_pending_document_still_blocks_generation():
    from tests.test_monthly_report_native_layout import draft
    from app.monthly_report_checks import preflight
    block, source, assets=payload()
    assert document_page_indexes(block,(source,))=={1}
    checks=preflight(replace(draft(block),sources=(source,)))
    assert any(c.code=='client_pages' and c.blocking for c in checks)


def test_identical_portrait_photo_without_exact_page_mapping_stays_photo():
    block, source, _=payload()
    unmapped=replace(block,asset_provenance=(),references=(),photos_per_page=3)
    assert tuple(picture_runs(unmapped,(source,)))==((False,(0,1,2)),)


def test_document_only_accepts_native_section_without_photo_frame(tmp_path):
    from tests.test_monthly_report_native_layout import draft
    from app.monthly_report_native_layout import build_native_docx
    block, source, assets=payload()
    block=replace(block,asset_hashes=(block.asset_hashes[1],),asset_captions=('Service page',))
    document=Document(); document.add_heading('MONTHLY ACTIVITY SUMMARY',1)
    document.add_paragraph('Old narrative')
    path=tmp_path/'no-photo.docx';document.save(path)
    output=Document(BytesIO(build_native_docx(replace(draft(block),sources=(source,)),path,asset_loader=assets.__getitem__)))
    assert len(output.inline_shapes)==1 and output.inline_shapes[0].width>Inches(5)
    picture=next(p for p in output.paragraphs if p._p.find('.//'+qn('wp:inline')) is not None)
    assert not picture.paragraph_format.page_break_before


def test_retained_original_native_art_is_not_reflowed(tmp_path):
    from tests.test_monthly_report_native_layout import draft
    from app.monthly_report_native_layout import build_native_docx
    from app.monthly_report_import import inspect_docx, map_items, ImportMapping
    block, source, assets=payload();raw=assets[block.asset_hashes[1]]
    document=Document();document.add_heading('MONTHLY ACTIVITY SUMMARY',1)
    document.add_paragraph('Improvement Highlights and Pics')
    document.add_picture(BytesIO(raw),width=Inches(1.4),height=Inches(1.8))
    path=tmp_path/'retained.docx';document.save(path)
    inspection=inspect_docx(path);item=next(i for i in inspection.items if i.kind=='image')
    mapped=map_items(path,inspection,(ImportMapping(item.id,'improvements'),))
    original=mapped.blocks[0];reference=f'{source.id}:1:{source.fingerprint}'
    original=replace(original,references=(*original.references,reference),asset_provenance=((original.asset_hashes[0],(reference,)),))
    assert document_page_indexes(original,(source,))=={0}
    output=Document(BytesIO(build_native_docx(replace(draft(original),sources=(source,)),path,asset_loader=dict(mapped.assets).__getitem__)))
    assert len(output.inline_shapes)==1
    assert output.inline_shapes[0].width==Inches(1.4) and output.inline_shapes[0].height==Inches(1.8)


def test_multiple_document_pages_landscape_section_keep_order(tmp_path):
    from tests.test_monthly_report_native_layout import draft
    from app.monthly_report_native_layout import build_native_docx
    block, source, assets=payload();source=replace(source,page_texts=('One','Two'),selected_pages=(1,2))
    another,raw=image((1100,850),'purple');assets[another]=raw
    refs=(f'{source.id}:1:{source.fingerprint}',f'{source.id}:2:{source.fingerprint}')
    block=replace(block,asset_hashes=(block.asset_hashes[1],another),asset_captions=('Page one','Page two'),references=refs,
                  asset_provenance=((block.asset_hashes[1],(refs[0],)),(another,(refs[1],))))
    document=Document();document.sections[0].page_width=Inches(11);document.sections[0].page_height=Inches(8.5)
    document.add_heading('MONTHLY ACTIVITY SUMMARY',1);document.add_paragraph('Old narrative')
    path=tmp_path/'landscape.docx';document.save(path)
    output=Document(BytesIO(build_native_docx(replace(draft(block),sources=(source,)),path,asset_loader=assets.__getitem__)))
    assert len(output.inline_shapes)==2
    assert output.inline_shapes[1].width>Inches(6)
    assert all(s.height<Inches(7) for s in output.inline_shapes)
    assert [p.text for p in output.paragraphs if p.text in block.asset_captions]==list(block.asset_captions)


@pytest.mark.parametrize('sequence', [(0,1),(1,2),(0,2,1),(1,)])
def test_native_edge_order_and_single_available_frame(tmp_path,sequence):
    from tests.test_monthly_report_native_layout import draft
    from app.monthly_report_native_layout import build_native_docx
    block,source,assets=payload()
    block=replace(block,asset_hashes=tuple(block.asset_hashes[i] for i in sequence),
                  asset_captions=tuple(block.asset_captions[i] for i in sequence))
    document=Document();document.add_heading('MONTHLY ACTIVITY SUMMARY',1)
    document.add_paragraph('Improvement Highlights and Pics')
    document.add_picture(BytesIO(next(iter(assets.values()))),width=Inches(6),height=Inches(1.8))
    path=tmp_path/'single.docx';document.save(path)
    raw=build_native_docx(replace(draft(block),sources=(source,)),path,asset_loader=assets.__getitem__)
    output=Document(BytesIO(raw));shapes=list(output.inline_shapes)
    assert len(shapes)==len(sequence)
    assert [sha256(output.part.related_parts[s._inline.graphic.graphicData.pic.blipFill.blip.embed].blob).hexdigest()+'.png' for s in shapes]==list(block.asset_hashes)
    assert [p.text for p in output.paragraphs if p.text in block.asset_captions]==list(block.asset_captions)
    assert shapes[sequence.index(1)].height>Inches(6)


def test_duplicate_asset_occurrences_do_not_infer_document_intent():
    block,source,_=payload()
    block=replace(block,asset_hashes=(block.asset_hashes[1],block.asset_hashes[1]),asset_captions=('One','Other'))
    assert not document_page_indexes(block,(source,))


def test_shared_photo_and_document_bytes_keep_document_preview_fidelity():
    from tests.test_monthly_report_native_layout import draft
    from app.monthly_report_asset_layout import photo_preview_references
    block,source,_=payload()
    cover=ResolvedBlock('cover_photo','This month',asset_hashes=(block.asset_hashes[1],))
    current=replace(draft(block,cover),sources=(source,))
    assert photo_preview_references(current)=={block.asset_hashes[0],block.asset_hashes[2]}
    assert document_page_indexes(cover,(source,))==set()


def test_trailing_empty_native_paragraph_budget_preserves_boundary_and_stops_at_content():
    from app.monthly_report_native_pages import NativePages
    from docx.shared import Pt
    from docx.enum.text import WD_BREAK
    document=Document();origin=document.add_paragraph('Current picture')
    empty=document.add_paragraph();empty.paragraph_format.space_after=Pt(9)
    empty.add_run().font.size=Pt(12)
    boundary=document.add_paragraph();boundary.add_run().add_break(WD_BREAK.PAGE)
    document.add_paragraph('Next section content')
    before=document._element.xml
    # No page owner mutation is involved in this read-only sizing operation.
    budget=NativePages.trailing_picture_space(None,document,origin._p)
    assert budget >= (12*1.2+9)*20
    assert document._element.xml==before
    assert NativePages.trailing_picture_space(None,document,boundary._p)==0


from tests.conftest import requires_libreoffice


@requires_libreoffice
def test_rendered_document_tail_has_no_new_blank_page(tmp_path):
    import fitz
    from docx.enum.text import WD_BREAK
    from tests.test_monthly_report_native_layout import draft
    from app.monthly_report_native_layout import build_native_docx
    from app.pdf_converter import convert_to_pdf
    block,source,assets=payload()
    block=replace(block,asset_hashes=block.asset_hashes[:2],asset_captions=block.asset_captions[:2])
    document=Document()
    from docx.shared import Pt
    section=document.sections[0]
    section.page_width=Inches(8.5);section.page_height=Inches(11)
    section.top_margin=section.bottom_margin=Inches(.7)
    section.left_margin=section.right_margin=Inches(.8)
    document.styles['Normal'].font.name='Arial'
    document.styles['Normal'].font.size=Pt(11)
    document.styles['Normal'].paragraph_format.space_after=Pt(7)
    document.add_heading('MONTHLY ACTIVITY SUMMARY',1)
    document.add_paragraph('Retained activity narrative')
    document.add_paragraph('Improvement Highlights and Pics')
    document.add_picture(BytesIO(next(iter(assets.values()))),width=Inches(6),height=Inches(1.8))
    document.add_paragraph();document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    document.add_heading('TRAINING SUMMARY',1);document.add_paragraph('Old training')
    path=tmp_path/'tail-source.docx';document.save(path)
    from app.monthly_report_import import inspect_docx,map_items,ImportMapping
    inspection=inspect_docx(path)
    text=next(item for item in inspection.items if item.kind=='text' and item.text=='Retained activity narrative')
    retained=map_items(path,inspection,(ImportMapping(text.id,'activity_summary'),)).blocks[0]
    current=replace(draft(retained,block,ResolvedBlock('training_summary','This month',text='Current training')),sources=(source,))
    for scope,expected in (('activity',2),(None,3)):
        raw=build_native_docx(current,path,asset_loader=assets.__getitem__,section_key=scope)
        out=tmp_path/f'tail-{scope}.docx';out.write_bytes(raw)
        output=convert_to_pdf(out)
        try:
            with fitz.open(output) as pdf:
                assert len(pdf)==expected, f"{scope}: expected {expected} pages, got {len(pdf)}"
                assert 'Service page' in pdf[1].get_text()
                if scope is None: assert 'Current training' in pdf[2].get_text()
        finally:
            output.unlink(missing_ok=True)


@pytest.mark.parametrize('inherited',[False,True])
def test_tail_budget_stops_before_effective_next_page_paragraph(inherited):
    from app.monthly_report_native_pages import NativePages
    from docx.shared import Pt
    from docx.enum.style import WD_STYLE_TYPE
    document=Document();origin=document.add_paragraph('Current picture')
    boundary=document.add_paragraph()
    if inherited:
        style=document.styles.add_style('Synthetic next page',WD_STYLE_TYPE.PARAGRAPH)
        style.paragraph_format.page_break_before=True
        style.paragraph_format.space_before=Pt(600)
        boundary.style=style
    else:
        boundary.paragraph_format.page_break_before=True
        boundary.paragraph_format.space_before=Pt(600)
    document.add_paragraph().paragraph_format.space_after=Pt(600)
    before=document._element.xml
    assert NativePages.trailing_picture_space(None,document,origin._p)==0
    assert document._element.xml==before


def test_legitimate_page_caption_rereview_retains_document_intent_without_marker_only_trust():
    from app.monthly_report_page_controls import revise_included_page
    from app.monthly_report_asset_review import approve_all_assets,refresh_asset_source_context
    from app.monthly_report_content_policy import page_fingerprint
    block,source,_=payload()
    source=replace(source,client_page_reviews=((1,page_fingerprint(source,1)),))
    reference=f'{source.id}:1:{source.fingerprint}'
    block=replace(block,references=(reference,),asset_provenance=tuple((asset,(reference,) if i==1 else ()) for i,asset in enumerate(block.asset_hashes)))
    block=approve_all_assets(refresh_asset_source_context(block,(source,)))
    edited,revised=revise_included_page(block,source,1,'caption',caption='Updated service page',all_sources=(source,))
    assert not document_page_indexes(edited,(revised,))
    reviewed,current=revise_included_page(edited,revised,1,'review',all_sources=(revised,))
    assert document_page_indexes(reviewed,(current,))=={1}
    assert not document_page_indexes(replace(reviewed,client_asset_reviews=block.client_asset_reviews),(current,))
    bad=replace(reviewed,asset_provenance=tuple((asset,tuple(value+'x' if value.startswith('asset-source-context:') else value for value in refs)) for asset,refs in reviewed.asset_provenance))
    assert not document_page_indexes(bad,(current,))
    assert reviewed.asset_hashes==block.asset_hashes


@pytest.mark.parametrize('caption',['Long service caption '*80,'\n'.join(['line']*5),'W'*500])
def test_standalone_document_caption_overflow_blocks_native_generic_and_preflight(tmp_path,caption):
    from tests.test_monthly_report_native_layout import draft
    from app.monthly_report_checks import preflight
    from app.monthly_report_docx import _build_docx
    from app.monthly_report_native_layout import build_native_docx
    block,source,assets=payload();block=replace(block,asset_captions=('Photo A',caption,'Photo B'))
    current=replace(draft(block),sources=(source,))
    assert any(c.code=='document_caption' and c.blocking for c in preflight(current))
    with pytest.raises(ValueError,match='caption'): _build_docx(current,asset_loader=assets.__getitem__)
    document=Document();document.add_heading('MONTHLY ACTIVITY SUMMARY',1)
    document.add_picture(BytesIO(next(iter(assets.values()))),width=Inches(6),height=Inches(1.8))
    path=tmp_path/'caption.docx';document.save(path)
    with pytest.raises(ValueError,match='caption'):build_native_docx(current,path,asset_loader=assets.__getitem__)


def table_source(tmp_path,assets,columns=2):
    document=Document();document.add_heading('MONTHLY ACTIVITY SUMMARY',1)
    document.add_paragraph('Improvement Highlights and Pics')
    table=document.add_table(rows=1,cols=columns)
    for cell in table.rows[0].cells:
        cell.paragraphs[0].add_run().add_picture(BytesIO(next(iter(assets.values()))),width=Inches(2.4),height=Inches(.72))
    path=tmp_path/'grid.docx';document.save(path)
    return path


def test_appended_pdf_leaves_native_photo_grid_and_preserves_photos(tmp_path):
    from tests.test_monthly_report_native_layout import draft
    from app.monthly_report_native_layout import build_native_docx
    block,source,assets=payload()
    block=replace(block,asset_hashes=(block.asset_hashes[0],block.asset_hashes[2],block.asset_hashes[1]),asset_captions=('Photo A','Photo B','Service page'))
    path=table_source(tmp_path,assets)
    output=Document(BytesIO(build_native_docx(replace(draft(block),sources=(source,)),path,asset_loader=assets.__getitem__)))
    assert len(output.tables)==1
    assert len(list(output.tables[0]._tbl.iter(qn('wp:inline'))))==2
    page=next(p for p in output.paragraphs if p._p.find('.//'+qn('wp:inline')) is not None)
    assert page._p.getparent() is output._element.body
    assert output.inline_shapes[-1].width>Inches(5)
    assert output.inline_shapes[0].width==Inches(2.4)
    assert [sha256(output.part.related_parts[s._inline.graphic.graphicData.pic.blipFill.blip.embed].blob).hexdigest()+'.png' for s in output.inline_shapes]==list(block.asset_hashes)


def test_interleaved_pdf_cannot_silently_reorder_or_flatten_native_grid(tmp_path):
    from tests.test_monthly_report_native_layout import draft
    from app.monthly_report_native_layout import build_native_docx
    block,source,assets=payload();path=table_source(tmp_path,assets)
    before=path.read_bytes();current=replace(draft(block),sources=(source,));fingerprint=current.fingerprint
    with pytest.raises(ValueError,match='Move the source page after its photos'):
        build_native_docx(current,path,asset_loader=assets.__getitem__)
    assert path.read_bytes()==before and current.fingerprint==fingerprint


def test_obsolete_document_anchor_break_removed_even_in_protected_region(tmp_path):
    from tests.test_monthly_report_native_layout import draft
    from app.monthly_report_native_layout import build_native_docx
    from app.monthly_report_import import inspect_docx,map_items,ImportMapping
    block,source,assets=payload();block=replace(block,asset_hashes=(block.asset_hashes[1],),asset_captions=('Service page',))
    document=Document();document.add_heading('MONTHLY ACTIVITY SUMMARY',1)
    document.add_paragraph('Retained narrative')
    frame=document.add_paragraph();frame.paragraph_format.page_break_before=True
    frame.add_run().add_picture(BytesIO(next(iter(assets.values()))),width=Inches(6),height=Inches(1.8))
    path=tmp_path/'protected-anchor.docx';document.save(path)
    inspection=inspect_docx(path);item=next(i for i in inspection.items if i.kind=='text' and i.text=='Retained narrative')
    retained=map_items(path,inspection,(ImportMapping(item.id,'activity_summary'),)).blocks[0]
    output=Document(BytesIO(build_native_docx(replace(draft(retained,block),sources=(source,)),path,asset_loader=assets.__getitem__)))
    assert 'Retained narrative' in [p.text for p in output.paragraphs]
    assert len(output.inline_shapes)==1
    assert not any(p.paragraph_format.page_break_before for p in output.paragraphs if not p.text and p._p.find('.//'+qn('wp:inline')) is None)


@requires_libreoffice
def test_rendered_appended_grid_page_is_body_level_and_readable(tmp_path):
    import fitz
    from tests.test_monthly_report_native_layout import draft
    from app.monthly_report_native_layout import build_native_docx
    from app.pdf_converter import convert_to_pdf
    block,source,assets=payload()
    block=replace(block,asset_hashes=(block.asset_hashes[0],block.asset_hashes[2],block.asset_hashes[1]),asset_captions=('Photo A','Photo B','Service page'))
    source_path=table_source(tmp_path,assets)
    raw=build_native_docx(replace(draft(block),sources=(source,)),source_path,asset_loader=assets.__getitem__)
    output_path=tmp_path/'grid-page.docx';output_path.write_bytes(raw)
    pdf=convert_to_pdf(output_path)
    try:
        with fitz.open(pdf) as output:
            assert len(output)==2
            assert 'Photo A' in output[0].get_text() and 'Photo B' in output[0].get_text()
            assert 'Service page' in output[1].get_text()
            images=output[1].get_image_info()
            assert len(images)==1 and images[0]['bbox'][2]-images[0]['bbox'][0]>5*72
    finally:pdf.unlink(missing_ok=True)


def test_multiple_appended_grid_documents_keep_page_and_caption_sequence(tmp_path):
    from tests.test_monthly_report_native_layout import draft
    from app.monthly_report_native_layout import build_native_docx
    block,source,assets=payload();another,raw=image((850,1100),'purple');assets[another]=raw
    source=replace(source,page_texts=('One','Two'),selected_pages=(1,2))
    refs=(f'{source.id}:1:{source.fingerprint}',f'{source.id}:2:{source.fingerprint}')
    page1=block.asset_hashes[1]
    block=replace(block,asset_hashes=(block.asset_hashes[0],block.asset_hashes[2],page1,another),
                  asset_captions=('Photo A','Photo B','Page one','Page two'),references=refs,
                  asset_provenance=((page1,(refs[0],)),(another,(refs[1],))))
    path=table_source(tmp_path,assets,columns=3)
    output=Document(BytesIO(build_native_docx(replace(draft(block),sources=(source,)),path,asset_loader=assets.__getitem__)))
    assert len(output.inline_shapes)==4
    assert [sha256(output.part.related_parts[s._inline.graphic.graphicData.pic.blipFill.blip.embed].blob).hexdigest()+'.png' for s in output.inline_shapes]==list(block.asset_hashes)
    assert [p.text for p in output.paragraphs if p.text in ('Page one','Page two')]==['Page one','Page two']
    assert len(list(output.tables[0]._tbl.iter(qn('wp:inline'))))==2


def test_source_page_caption_refuses_unproved_narrow_native_area(tmp_path):
    from tests.test_monthly_report_native_layout import draft
    from app.monthly_report_native_layout import build_native_docx
    block,source,assets=payload();block=replace(block,asset_hashes=(block.asset_hashes[1],),asset_captions=('Service page',))
    document=Document();document.sections[0].left_margin=document.sections[0].right_margin=Inches(2.5)
    document.add_heading('MONTHLY ACTIVITY SUMMARY',1);document.add_paragraph('Old narrative')
    path=tmp_path/'narrow.docx';document.save(path)
    with pytest.raises(ValueError,match='too narrow'):
        build_native_docx(replace(draft(block),sources=(source,)),path,asset_loader=assets.__getitem__)


def test_anchor_cleanup_never_discards_retained_text_fields_or_boundaries():
    from app.monthly_report_native_layout import _discardable_picture_anchor
    from docx.oxml import OxmlElement
    document=Document();empty=document.add_paragraph();empty.paragraph_format.page_break_before=True
    assert _discardable_picture_anchor(empty._p)
    empty.add_run('Retained current text');assert not _discardable_picture_anchor(empty._p)
    field=document.add_paragraph();field._p.append(OxmlElement('w:bookmarkStart'))
    assert not _discardable_picture_anchor(field._p)
    boundary=document.add_paragraph();boundary._p.get_or_add_pPr().append(OxmlElement('w:sectPr'))
    assert not _discardable_picture_anchor(boundary._p)


def test_document_caption_uses_measured_font_without_inherited_indentation():
    from app.monthly_report_asset_layout import style_document_caption
    from docx.shared import Pt
    document=Document();style=document.styles['Caption'];style.font.bold=True;style.font.italic=True
    style.paragraph_format.left_indent=Inches(2);style.paragraph_format.right_indent=Inches(1)
    paragraph=document.add_paragraph('Full caption retained',style='Caption')
    style_document_caption(paragraph)
    assert paragraph.text=='Full caption retained'
    assert paragraph.paragraph_format.left_indent==paragraph.paragraph_format.right_indent==0
    assert paragraph.runs[0].font.bold is False and paragraph.runs[0].font.italic is False
    assert paragraph.runs[0].font.size==Pt(10)


def test_mapped_native_grid_append_keeps_mapped_photo_bytes_and_body_document(tmp_path):
    from app.monthly_report_import import inspect_docx, map_items, ImportMapping
    from app.monthly_report_native_layout import build_native_docx
    from tests.test_monthly_report_native_layout import draft
    block, source, assets = payload()
    path = table_source(tmp_path, assets)
    inspection = inspect_docx(path)
    mapped = map_items(path, inspection, tuple(
        ImportMapping(item.id, 'improvements') for item in inspection.items if item.kind == 'image'))
    original = mapped.blocks[0]
    current = replace(original, asset_hashes=(*original.asset_hashes, block.asset_hashes[1]),
                      asset_captions=(*original.asset_captions, 'Service page'),
                      references=(*original.references, *block.references),
                      asset_provenance=(*original.asset_provenance, *block.asset_provenance))
    assets.update(mapped.assets)
    before = Document(path)
    output = Document(BytesIO(build_native_docx(replace(draft(current), sources=(source,)),
                                              path, asset_loader=assets.__getitem__)))
    assert len(output.inline_shapes) == 3
    assert len(list(output.tables[0]._tbl.iter(qn('wp:inline')))) == 2
    # Appending to the block invalidates whole-block unchanged replay; the
    # existing importer-normalized photo bytes, not original PNGs, are emitted.
    for old, new in zip(before.inline_shapes, output.inline_shapes):
        assert (new.width, new.height) == (old.width, old.height)
        new_part = output.part.related_parts[new._inline.graphic.graphicData.pic.blipFill.blip.embed]
        assert new_part.blob == assets[original.asset_hashes[0]]
    assert output.inline_shapes[-1].width > Inches(5)
    assert output.inline_shapes[-1]._inline.getparent().getparent().getparent().getparent() is output._element.body


def test_section_preview_tracks_review_only_document_geometry_not_unrelated_consent():
    from app.monthly_report_preview import section_preview_fingerprint
    from app.monthly_report_page_controls import revise_included_page
    from app.monthly_report_asset_review import approve_all_assets, refresh_asset_source_context
    from app.monthly_report_content_policy import page_fingerprint
    from tests.test_monthly_report_native_layout import draft
    block, source, _ = payload()
    source = replace(source, client_page_reviews=((1, page_fingerprint(source, 1)),))
    reference = f'{source.id}:1:{source.fingerprint}'
    block = replace(block, references=(reference,), asset_provenance=tuple(
        (asset, (reference,) if i == 1 else ()) for i, asset in enumerate(block.asset_hashes)))
    block = approve_all_assets(refresh_asset_source_context(block, (source,)))
    pending, changed = revise_included_page(block, source, 1, 'caption',
                                           caption='Updated service page', all_sources=(source,))
    reviewed, current = revise_included_page(pending, changed, 1, 'review', all_sources=(changed,))
    assert pending.asset_hashes == reviewed.asset_hashes
    assert pending.asset_captions == reviewed.asset_captions
    assert pending.references == reviewed.references
    assert document_page_indexes(pending, (changed,)) == set()
    assert document_page_indexes(reviewed, (current,)) == {1}
    before = replace(draft(pending), sources=(changed,))
    after = replace(draft(reviewed), sources=(current,))
    fingerprint = section_preview_fingerprint(after, 'activity')
    assert section_preview_fingerprint(before, 'activity') != fingerprint
    # Revoking the exact current image receipt changes the layout back, too.
    revoked = replace(after, blocks=(replace(reviewed, client_asset_reviews=()),))
    assert section_preview_fingerprint(revoked, 'activity') != fingerprint
    assert section_preview_fingerprint(before, 'training') == section_preview_fingerprint(after, 'training')
    unrelated = replace(after, blocks=(replace(reviewed, reviewed_fingerprint='unrelated'),))
    assert section_preview_fingerprint(unrelated, 'activity') == fingerprint


@pytest.mark.parametrize('count', [1, 2])
def test_document_only_native_grid_uses_body_pages_in_source_order(tmp_path, count):
    from tests.test_monthly_report_native_layout import draft
    from app.monthly_report_native_layout import build_native_docx
    block, source, assets = payload()
    refs = [block.asset_hashes[1]]
    source = replace(source, page_texts=tuple('Service page' for _ in range(count)), selected_pages=tuple(range(1,count+1)))
    if count == 2:
        ref, raw = image((850,1100), 'purple'); assets[ref] = raw; refs.append(ref)
    references = tuple(f'{source.id}:{i+1}:{source.fingerprint}' for i in range(count))
    block = replace(block, asset_hashes=tuple(refs), asset_captions=tuple(f'Service page {i+1}' for i in range(count)),
                    references=references, asset_provenance=tuple((ref,(references[i],)) for i,ref in enumerate(refs)))
    path = table_source(tmp_path, assets)
    output = Document(BytesIO(build_native_docx(replace(draft(block),sources=(source,)),path,asset_loader=assets.__getitem__)))
    assert len(output.inline_shapes) == count
    assert not output.tables
    for shape, ref in zip(output.inline_shapes, refs):
        assert shape.width > Inches(5)
        assert shape._inline.getparent().getparent().getparent().getparent() is output._element.body
        assert output.part.related_parts[shape._inline.graphic.graphicData.pic.blipFill.blip.embed].blob == assets[ref]
    assert [p.text for p in output.paragraphs if p.text.startswith('Service page')] == list(block.asset_captions)


@requires_libreoffice
@pytest.mark.parametrize('count', [1, 2])
@pytest.mark.parametrize('retained', [False, True])
def test_rendered_document_only_native_grid_has_no_blank_page(tmp_path, count, retained):
    import fitz
    from app.pdf_converter import convert_to_pdf
    from app.monthly_report_import import inspect_docx, map_items, ImportMapping
    from app.monthly_report_native_layout import build_native_docx
    from tests.test_monthly_report_native_layout import draft
    block, source, assets = payload()
    refs = [block.asset_hashes[1]]
    if count == 2:
        ref, raw = image((850,1100), 'purple'); assets[ref] = raw; refs.append(ref)
    source = replace(source, page_texts=tuple('Service page' for _ in refs), selected_pages=tuple(range(1,count+1)))
    references = tuple(f'{source.id}:{i+1}:{source.fingerprint}' for i in range(count))
    block = replace(block, asset_hashes=tuple(refs), asset_captions=tuple(f'Service page {i+1}' for i in range(count)),
                    references=references, asset_provenance=tuple((ref,(references[i],)) for i,ref in enumerate(refs)))
    path = table_source(tmp_path, assets)
    blocks = [block]
    if retained:
        document = Document(path)
        document.tables[0]._tbl.addprevious(document.add_paragraph('Retained activity narrative')._p)
        document.save(path)
        inspection = inspect_docx(path)
        item = next(item for item in inspection.items if item.kind == 'text' and item.text == 'Retained activity narrative')
        blocks.insert(0, map_items(path, inspection, (ImportMapping(item.id,'activity_summary'),)).blocks[0])
    raw = build_native_docx(replace(draft(*blocks),sources=(source,)),path,asset_loader=assets.__getitem__)
    generated = Document(BytesIO(raw))
    assert bool(generated.tables) == retained
    if retained:
        assert generated.tables[0]._tbl.tblGrid.xml == Document(path).tables[0]._tbl.tblGrid.xml
    output_path = tmp_path/'document-only.docx'; output_path.write_bytes(raw)
    pdf = convert_to_pdf(output_path)
    try:
        with fitz.open(pdf) as output:
            assert len(output) == count + int(retained)
            for i in range(count):
                page = output[i+int(retained)]
                assert f'Service page {i+1}' in page.get_text()
                images = page.get_image_info()
                assert len(images) == 1 and images[0]['bbox'][2]-images[0]['bbox'][0] > 5*72
            if retained: assert 'Retained activity narrative' in output[0].get_text()
    finally: pdf.unlink(missing_ok=True)
