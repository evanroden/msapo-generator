"""Synthetic analogues of report furniture, with adversarial content controls."""
from copy import deepcopy
from io import BytesIO
from zipfile import ZipFile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree as ET
import pytest

from app.monthly_report_import import inspect_docx
from app.monthly_report_sections import section_reviews
from app.monthly_report_start_import import StartingReportAnalysis, automatic_section_plans, build_starting_import
from app.monthly_report_native_design import clear_native_design_placeholders

WPS = "{http://schemas.microsoft.com/office/word/2010/wordprocessingShape}"
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def frame(text="", *, fill=False, geometry="rect", linked=False):
    drawing = OxmlElement("w:drawing")
    inline = ET.SubElement(drawing, qn("wp:inline"))
    ET.SubElement(inline, qn("wp:extent"), cx="1828800", cy="280670")
    graphic = ET.SubElement(inline, A + "graphic")
    data = ET.SubElement(graphic, A + "graphicData", uri=WPS[1:-1])
    shape = ET.SubElement(data, WPS + "wsp")
    ET.SubElement(shape, WPS + "cNvSpPr", txBox="1")
    props = ET.SubElement(shape, WPS + "spPr")
    ET.SubElement(props, A + "prstGeom", prst=geometry)
    if fill:
        ET.SubElement(ET.SubElement(props, A + "solidFill"), A + "srgbClr", val="FF0000")
    box = ET.SubElement(ET.SubElement(shape, WPS + "txbx"), qn("w:txbxContent"))
    for line in text.splitlines() or ("",):
        paragraph = ET.SubElement(box, qn("w:p"))
        if line:
            ET.SubElement(ET.SubElement(paragraph, qn("w:r")), qn("w:t")).text = line
    if linked:
        ET.SubElement(shape, WPS + "linkedTxbx", id="42", seq="1")
    ET.SubElement(shape, WPS + "bodyPr")
    return drawing


def report(tmp_path, drawings, *, heading="EQUIPMENT PERFORMANCE ISSUES", header=False):
    document = Document()
    document.add_paragraph("Synthetic Hospital")
    document.add_paragraph("August 2026")
    if heading:
        document.add_heading(heading, 1)
    paragraph = document.sections[0].header.paragraphs[0] if header else document.add_paragraph()
    for drawing in drawings:
        paragraph.add_run()._r.append(drawing)
    if heading == "PENDING & DECLINED PROPOSALS":
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text, table.cell(0, 1).text = "Proposal", "Status"
        table.cell(1, 0).text, table.cell(1, 1).text = "Replace pump", "Pending"
    source = tmp_path / "synthetic.docx"
    document.save(source)
    return source


def plans(inspection):
    return automatic_section_plans(inspection, StartingReportAnalysis((), "same", (), (), ()))


def test_native_plain_text_frame_is_editable_content_not_unread_diagram(tmp_path):
    source = report(tmp_path, [frame("Synthetic motor repaired.\nInspect the coupling next month.")])
    info = inspect_docx(source)
    plan = plans(info)["issues"]
    assert plan["approved"], plan["questions"]
    mapped, _ = build_starting_import(source, info, (plan,))
    block = next(b for b in mapped.blocks if b.key == "equipment_issues")
    assert "Synthetic motor repaired." in block.text
    assert "Inspect the coupling next month." in block.text
    assert block.references
    assert any(i.kind == "native_layout" and i.note == "Editable native text frames" for i in info.items)


@pytest.mark.parametrize("count", [0, 1, 21, 50])
def test_empty_unstyled_proposal_frames_do_not_block_native_tables(tmp_path, count):
    source = report(tmp_path, [frame() for _ in range(count)], heading="PENDING & DECLINED PROPOSALS")
    original = source.read_bytes()
    info = inspect_docx(source)
    plan = plans(info)["proposals"]
    assert plan["approved"], plan["questions"]
    mapped, _ = build_starting_import(source, info, (plan,))
    assert next(b for b in mapped.blocks if b.key == "proposals").extra_tables[0].rows == (("Replace pump", "Pending"),)
    if count:
        assert any(i.kind == "native_layout" and i.note == "Empty unstyled native text frames" for i in info.items)
    assert source.read_bytes() == original


@pytest.mark.parametrize("mutation", ["fill", "non_rect", "linked", "drawing_text", "image", "chart", "unknown", "symbol", "field"])
def test_empty_or_textual_appearance_is_not_proof_of_harmless_native_art(tmp_path, mutation):
    drawing = frame(fill=mutation == "fill", geometry="ellipse" if mutation == "non_rect" else "rect",
                    linked=mutation == "linked")
    shape = drawing.find(".//" + WPS + "wsp")
    if mutation == "drawing_text":
        ET.SubElement(shape, A + "t").text = "Source-only confidential figure $550"
    elif mutation == "image":
        ET.SubElement(shape, A + "blip", {qn("r:link"): "external"})
    elif mutation == "chart":
        drawing.find(".//" + A + "graphicData").set("uri", "http://schemas.openxmlformats.org/drawingml/2006/chart")
    elif mutation == "unknown":
        ET.SubElement(shape, "{urn:synthetic-unknown}payload").text = "unread"
    elif mutation == "symbol":
        run = ET.SubElement(shape.find('.//' + qn('w:p')), qn('w:r'))
        ET.SubElement(run, qn('w:sym'), {qn('w:font'): 'Arial', qn('w:char'): '0031'})
    elif mutation == "field":
        ET.SubElement(shape.find('.//' + qn('w:p')), qn('w:fldSimple'), {qn('w:instr'): 'INCLUDETEXT secret.docx'})
    info = inspect_docx(report(tmp_path, [drawing], heading="PENDING & DECLINED PROPOSALS"))
    assert not plans(info)["proposals"]["approved"]
    assert any(i.kind == "unsupported" for s in section_reviews(info) for i in s.items)


def test_known_utility_header_frames_are_preserved_design_not_cover_questions(tmp_path):
    info = inspect_docx(report(tmp_path, [frame("Client Utility Rate Analysis"), frame("Author:")], header=True))
    assert not any(i.kind == "unsupported" and i.part.startswith("word/header") for i in info.items)
    assert any(i.kind == "native_layout" and i.note == "Native running-header text frames" for i in info.items)


def test_unknown_header_facts_still_require_recovery(tmp_path):
    info = inspect_docx(report(tmp_path, [frame("Unexpected private account balance 45000")], header=True))
    assert not plans(info)["cover"]["approved"]


def test_several_text_frames_can_encode_substantive_layout_and_stay_unresolved(tmp_path):
    info = inspect_docx(report(tmp_path, [frame("Supervisor"), frame("Technician")]))
    assert not plans(info)["issues"]["approved"]


def test_cover_text_frame_is_not_blanket_permission_for_unknown_cover_facts(tmp_path):
    info = inspect_docx(report(tmp_path, [frame("Synthetic equipment decision diagram")], heading=""))
    assert not plans(info)["cover"]["approved"]


def band(color="547E7E", *, width=3000000, height=700000):
    drawing = frame()
    shape = drawing.find(".//" + WPS + "wsp")
    shape.remove(shape.find(WPS + "txbx"))
    props = shape.find(WPS + "spPr")
    props.remove(props.find(A + "prstGeom"))
    transform = ET.SubElement(props, A + "xfrm")
    ET.SubElement(transform, A + "ext", cx=str(width), cy=str(height))
    geometry = ET.SubElement(props, A + "custGeom")
    path = ET.SubElement(ET.SubElement(geometry, A + "pathLst"), A + "path", w=str(width), h=str(height))
    for tag, x, y in (("moveTo", width, 0), ("lnTo", 0, 0), ("lnTo", 0, height),
                      ("lnTo", width - 300000, height), ("lnTo", width, 0)):
        ET.SubElement(ET.SubElement(path, A + tag), A + "pt", x=str(x), y=str(y))
    ET.SubElement(path, A + "close")
    ET.SubElement(ET.SubElement(props, A + "solidFill"), A + ("schemeClr" if color == "accent2" else "srgbClr"), val=color)
    return drawing


def mint_report(tmp_path, *, numbers="1.\n2.\n3.\n4.\n5.", broken=False, legacy=False):
    from PIL import Image
    from docx.shared import Inches
    document = Document()
    document.add_paragraph("Synthetic Hospital")
    document.add_paragraph("August 2026")
    document.add_heading("MONTHLY SCORECARDS", 1)
    document.add_paragraph("Synthetic thermal capacity remains unchanged.")
    paragraph = document.add_paragraph()
    for color in ("accent2", "547E7E", "D2E7DF"):
        paragraph.add_run()._r.append(band(color))
    data = BytesIO(); Image.new("RGB", (800, 900), "navy").save(data, "PNG")
    picture = document.add_picture(BytesIO(data.getvalue()), width=Inches(8.5), height=Inches(9.2))
    drawing = picture._inline.getparent()
    old_paragraph = drawing.getparent().getparent()
    inline = picture._inline
    inline.tag = qn("wp:anchor")
    ET.SubElement(inline, qn("wp:positionV"), relativeFrom="page")
    if broken:
        next(inline.iter(A + "blip")).set(qn("r:embed"), "missing-image-id")
    paragraph.add_run()._r.append(drawing)
    old_paragraph.getparent().remove(old_paragraph)
    paragraph.add_run()._r.append(frame(numbers))
    if legacy:
        mc = "{http://schemas.openxmlformats.org/markup-compatibility/2006}"
        for drawing in list(paragraph._p.iter(qn("w:drawing"))):
            parent = drawing.getparent()
            wrapper = ET.Element(mc + "AlternateContent")
            parent.replace(drawing, wrapper)
            ET.SubElement(wrapper, mc + "Choice", Requires="wps").append(drawing)
            ET.SubElement(wrapper, mc + "Fallback").append(deepcopy(drawing))
    document.add_heading("MBCx REPORTS", 1)
    document.add_paragraph("Commissioning inspection completed.")
    source = tmp_path / "mint.docx"; document.save(source)
    return source


@pytest.mark.parametrize("legacy", [False, True])
def test_mint_divider_routes_forward_without_duplicating_it_as_a_scorecard(tmp_path, legacy):
    from app.monthly_report_import import _is_native_divider_background, _text, imported_draft
    from app.monthly_report_model import synthetic_profiles, ReportPeriod
    from app.monthly_report_native_layout import build_native_docx
    source = mint_report(tmp_path, legacy=legacy)
    original = source.read_bytes()
    info = inspect_docx(source)
    picture = next(i for i in info.items if i.kind == "image")
    assert picture.section == "mbcx" and picture.suggested_slot == "divider_mbcx"
    assert not any(i.kind == "unsupported" for s in section_reviews(info) for i in s.items)
    selected = plans(info)
    mapped, _ = build_starting_import(source, info, selected.values())
    draft = imported_draft(synthetic_profiles()[0], ReportPeriod(2026, 10), "Current Editor", mapped)
    output = Document(BytesIO(build_native_docx(draft, source, asset_loader=dict(mapped.assets).__getitem__)))
    output_text = _text(output.element)
    assert "1.\n2.\n3.\n4.\n5." not in output_text
    assert "Synthetic thermal capacity remains unchanged." in output_text
    assert "Commissioning inspection completed." in output_text
    assert any(_is_native_divider_background(p._p, _text(p._p)) for p in output.paragraphs)
    assert len(list(output.element.iter(A + "blip"))) == (2 if legacy else 1)
    assert source.read_bytes() == original


@pytest.mark.parametrize("numbers", ["1.\n2.\n3.", "101\n202", "1.\n2.\n3.\n4.\n5.\nMotor failures"])
def test_other_numbers_or_narrative_are_not_legacy_placeholder_proof(tmp_path, numbers):
    source = mint_report(tmp_path, numbers=numbers)
    info = inspect_docx(source)
    assert not plans(info)["scorecards"]["approved"]
    paragraph = next(p for p in Document(source).paragraphs if p._p.find(".//" + WPS + "wsp") is not None)
    before = ET.tostring(paragraph._p)
    clear_native_design_placeholders(paragraph._p)
    assert ET.tostring(paragraph._p) == before


def test_missing_divider_image_still_blocks_instead_of_becoming_furniture(tmp_path):
    info = inspect_docx(mint_report(tmp_path, broken=True))
    assert not plans(info)["scorecards"]["approved"]


def test_cover_placeholder_is_cleared_only_in_proved_native_cover(tmp_path):
    from app.monthly_report_native_layout import build_native_docx
    from app.monthly_report_import import _text
    from app.monthly_report_model import ReportDraft, ReportPeriod, synthetic_profiles, default_sections
    document = Document()
    paragraph = document.add_paragraph()
    for value in ("Prepared by: Prior Editor", "August 2026", "Synthetic Hospital\nOperations and Maintenance Monthly Review", "INSERT CLIENT LOGO"):
        paragraph.add_run()._r.append(frame(value))
    document.add_heading("MONTHLY ACTIVITY SUMMARY", 1)
    source = tmp_path / "cover.docx"; document.save(source)
    info = inspect_docx(source)
    assert any(i.kind == "native_layout" and i.note == "Native cover title and artwork" for i in info.items)
    draft = ReportDraft(synthetic_profiles()[0], ReportPeriod(2026, 10), "Current Editor", default_sections(), ())
    output = Document(BytesIO(build_native_docx(draft, source)))
    text = _text(output.element)
    assert "INSERT CLIENT LOGO" not in text and "August 2026" not in text
    assert "October 2026" in text and "Current Editor" in text
    assert text.count(draft.profile.title) == 1
    ordinary = frame("INSERT CLIENT LOGO")
    before = ET.tostring(ordinary)
    clear_native_design_placeholders(ordinary)
    assert ET.tostring(ordinary) == before


def test_recognized_native_report_can_continue_through_real_starting_ui(monkeypatch, tmp_path):
    from app import monthly_report_library as library
    from app.monthly_report_model import synthetic_profiles
    from tests.test_monthly_report_start_import import start_app
    from tests.test_monthly_report_starting_review_recovery import continue_button
    document = Document()
    paragraph = document.add_paragraph()
    title = synthetic_profiles()[0].facilities[0].title
    for value in ("Prepared by: Prior Editor", "September 2026", title + "\nOperations and Maintenance Monthly Review", "INSERT CLIENT LOGO"):
        paragraph.add_run()._r.append(frame(value))
    document.add_heading("EQUIPMENT PERFORMANCE ISSUES", 1)
    document.add_paragraph().add_run()._r.append(frame("Current synthetic coupling inspection completed."))
    document.add_heading("PENDING & DECLINED PROPOSALS", 1)
    document.add_paragraph().add_run()._r.append(frame())
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Proposal", "Status"
    table.cell(1, 0).text, table.cell(1, 1).text = "Repair coupling", "Pending"
    source = tmp_path / "native-start.docx"; document.save(source)
    original = source.read_bytes()
    app = start_app(monkeypatch, tmp_path, source)
    assert not continue_button(app).disabled
    app.run()
    assert not continue_button(app).disabled
    continue_button(app).click().run()
    assert not app.exception and not app.error
    profile = synthetic_profiles()[0]
    saved = library.load_imported_draft(profile.contract, profile.key)
    assert saved is not None
    assert any("coupling inspection" in b.text for b in saved.blocks)
    assert any(t.rows == (("Repair coupling", "Pending"),) for b in saved.blocks for t in b.extra_tables)
    assert source.read_bytes() == original


@pytest.mark.parametrize('payload', ['text', 'wordart', 'drawing_text', 'unknown', 'external', 'field', 'geometry', 'symbol'])
def test_divergent_fallback_cannot_acquire_native_text_frame_approval(tmp_path, payload):
    mc = '{http://schemas.openxmlformats.org/markup-compatibility/2006}'
    source = report(tmp_path, [frame('Current routine narrative.')])
    document = Document(source)
    drawing = next(document.element.iter(qn('w:drawing')))
    parent = drawing.getparent()
    wrapper = ET.Element(mc + 'AlternateContent'); parent.replace(drawing, wrapper)
    ET.SubElement(wrapper, mc + 'Choice', Requires='wps').append(drawing)
    fallback = ET.SubElement(wrapper, mc + 'Fallback'); fallback.append(deepcopy(drawing))
    if payload == 'text':
        next(fallback.iter(qn('w:t'))).text = 'SECRET prior invoice $99999'
    elif payload == 'wordart':
        shape = ET.SubElement(fallback, '{urn:schemas-microsoft-com:vml}shape')
        ET.SubElement(shape, '{urn:schemas-microsoft-com:vml}textpath', string='SECRET prior invoice $99999')
    elif payload == 'unknown':
        ET.SubElement(fallback, '{urn:unknown-graphic}shape').text = 'uninspected content'
    elif payload == 'drawing_text':
        ET.SubElement(fallback, A + 't').text = 'SECRET prior invoice $99999'
    elif payload == 'geometry':
        fallback.find('.//' + A + 'prstGeom').set('prst', 'star5')
    elif payload == 'symbol':
        ET.SubElement(fallback, qn('w:sym'), {qn('w:font'): 'Arial', qn('w:char'): '0031'})
    elif payload == 'external':
        ET.SubElement(fallback, A + 'blip', {qn('r:link'): 'external-id'})
    else:
        ET.SubElement(fallback, qn('w:fldSimple'), {qn('w:instr'): 'INCLUDETEXT secret.docx'})
    document.save(source); original = source.read_bytes()
    info = inspect_docx(source); selected = plans(info)
    assert not selected['issues']['approved']
    assert any(i.kind == 'unsupported' and i.note.startswith('Alternate Word') for i in info.items)
    with pytest.raises(ValueError):
        build_starting_import(source, info, selected.values())
    assert source.read_bytes() == original


def test_legacy_band_parser_does_not_accept_arbitrary_vector_plots():
    from app.monthly_report_native_design import _legacy_band_path
    assert _legacy_band_path('m3057385,l,,,771474r2731655,l3057385,xe')
    assert not _legacy_band_path('m1,l0,0,0,1,2,3,4,5,6,7,8,9,1,0xe')
    assert not _legacy_band_path('m1,l0,0,0,4r4,l1,xem2,l0,0,0,5r5,l2,xe')


@pytest.mark.parametrize('payload', ['text', 'image'])
def test_divergent_fallback_cannot_acquire_native_divider_approval(tmp_path, payload):
    from PIL import Image
    mc = '{http://schemas.openxmlformats.org/markup-compatibility/2006}'
    source = mint_report(tmp_path, legacy=True)
    document = Document(source)
    if payload == 'text':
        fallback = next(f for f in document.element.iter(mc + 'Fallback') if next(f.iter(qn('w:t')), None) is not None)
        next(fallback.iter(qn('w:t'))).text = 'SECRET prior invoice $99999'
    else:
        fallback = next(f for f in document.element.iter(mc + 'Fallback') if next(f.iter(A + 'blip'), None) is not None)
        data = BytesIO(); Image.new('RGB', (800, 900), 'orange').save(data, 'PNG')
        rid, _ = document.part.get_or_add_image(BytesIO(data.getvalue()))
        next(fallback.iter(A + 'blip')).set(qn('r:embed'), rid)
    document.save(source); original = source.read_bytes()
    info = inspect_docx(source); selected = plans(info)
    assert not selected['scorecards']['approved']
    assert any(i.kind == 'unsupported' and i.note.startswith('Alternate Word') for i in info.items)
    with pytest.raises(ValueError):
        build_starting_import(source, info, selected.values())
    assert source.read_bytes() == original


def test_selected_choice_screenshot_cannot_approve_hidden_legacy_content(tmp_path, monkeypatch):
    from app import monthly_report_visual_preservation as preservation
    from app.monthly_report_native_layout import build_native_docx, NativeLayoutError
    from app.monthly_report_model import ReportDraft, ReportPeriod, synthetic_profiles, default_sections
    mc = '{http://schemas.openxmlformats.org/markup-compatibility/2006}'
    source = report(tmp_path, [frame('Current routine narrative.')])
    document = Document(source)
    drawing = next(document.element.iter(qn('w:drawing')))
    parent = drawing.getparent(); wrapper = ET.Element(mc + 'AlternateContent'); parent.replace(drawing, wrapper)
    ET.SubElement(wrapper, mc + 'Choice', Requires='wps').append(drawing)
    fallback = ET.SubElement(wrapper, mc + 'Fallback'); fallback.append(deepcopy(drawing))
    next(fallback.iter(qn('w:t'))).text = 'SECRET prior invoice $99999'
    document.save(source)
    info = inspect_docx(source)
    section = next(s for s in section_reviews(info) if s.key == 'issues')
    monkeypatch.setattr(preservation, 'preserve_word_pages', lambda *_: pytest.fail('One rendered Choice cannot prove the unrendered Fallback'))
    with pytest.raises(preservation.DiagramPreservationError, match='modern and legacy'):
        preservation.preserve_section_diagrams(source, info, section)
    # Even a forged/old saved mapping cannot reuse the hidden branch unchanged.
    draft = ReportDraft(synthetic_profiles()[0], ReportPeriod(2026, 10), 'Current Editor', default_sections(), ())
    with pytest.raises(NativeLayoutError, match='hidden alternate'):
        build_native_docx(draft, source)


@pytest.mark.parametrize('location', ['between', 'heading'])
def test_forward_divider_routing_cannot_erase_alternate_content_diagnostic(tmp_path, location):
    from app.monthly_report_native_layout import build_native_docx, NativeLayoutError
    from app.monthly_report_model import ReportDraft, ReportPeriod, synthetic_profiles, default_sections
    mc = '{http://schemas.openxmlformats.org/markup-compatibility/2006}'
    source = mint_report(tmp_path)
    document = Document(source)
    heading = next(p for p in document.paragraphs if p.text == 'MBCx REPORTS')
    if location == 'heading':
        paragraph = heading._p
        for child in list(paragraph):
            if child.tag != qn('w:pPr'):paragraph.remove(child)
    else:
        paragraph = OxmlElement('w:p');heading._p.addprevious(paragraph)
    wrapper = ET.SubElement(ET.SubElement(paragraph, qn('w:r')), mc + 'AlternateContent')
    ET.SubElement(wrapper, mc + 'Choice', Requires='wps').append(frame('MBCx REPORTS' if location == 'heading' else ''))
    ET.SubElement(wrapper, mc + 'Fallback').append(frame('SECRET prior invoice $99999'))
    document.save(source)
    info = inspect_docx(source); selected = plans(info)
    assert any(i.kind == 'unsupported' and i.note.startswith('Alternate Word') for i in info.items)
    assert any(not p['approved'] for p in selected.values())
    with pytest.raises(ValueError):build_starting_import(source, info, selected.values())
    draft = ReportDraft(synthetic_profiles()[0], ReportPeriod(2026, 10), 'Current Editor', default_sections(), ())
    with pytest.raises(NativeLayoutError, match='hidden alternate'):build_native_docx(draft, source)
