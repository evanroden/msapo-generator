"""Synthetic-only experimental source-view proof. No production activation."""
from copy import deepcopy
from dataclasses import asdict
import hashlib
import json
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED

from docx import Document
from docx.oxml.ns import qn
from lxml import etree as ET
import pytest

from app.monthly_report_compatibility_view import (
    CompatibilityError, derive_compatibility_view, rederive_compatibility_view,
    _selection, _branch, POLICY,
)
from app.monthly_report_import import inspect_docx, _relations, _package
from test_monthly_report_source_frame_logo_role import source, mutate, raster, _picture_in_frame, MC, WPS, A


def test_default_off_and_deterministic_original_audit_and_occurrence_map(tmp_path):
    path = source(tmp_path, divergent=True)
    before = path.read_bytes()
    original = inspect_docx(path)
    assert any(i.note.startswith('Alternate Word') for i in original.items)
    view = derive_compatibility_view(before)
    again = derive_compatibility_view(before)
    assert view == again
    assert view.original_inspection == original
    assert path.read_bytes() == before
    assert view.original_inspection.sha256 == hashlib.sha256(before).hexdigest()
    assert view.rendition_inspection.sha256 == hashlib.sha256(view.rendition_bytes).hexdigest()
    assert [b.original_id for b in view.bindings] == [i.id for i in original.items]
    assert len({b.rendition_id for b in view.bindings if b.rendition_id}) == len(view.rendition_inspection.items)
    assert any(b.status == 'resolved' for b in view.bindings)
    assert any(i.kind == 'unsupported' for i in view.original_inspection.items)
    assert not any(i.note.startswith('Alternate Word') for i in view.rendition_inspection.items)
    from io import BytesIO
    with ZipFile(path) as src, ZipFile(BytesIO(view.rendition_bytes)) as dst:
        record = json.loads(view.receipt_bytes)['resolutions'][0]
        selected, fallback = record['branches']
        assert dst.read(selected['media_part']) == src.read(selected['media_part'])
        assert fallback['media_part'] not in dst.namelist()
        assert b'AlternateContent' not in dst.read('word/document.xml')
    # Normal inspector is unchanged and still blocks the immutable original.
    assert inspect_docx(path) == original


def test_rederive_rejects_receipt_or_original_changes(tmp_path):
    path = source(tmp_path, divergent=True)
    view = rederive_compatibility_view(path)
    assert rederive_compatibility_view(path, view.receipt_bytes) == view
    for key, changed in [('policy','other'), ('original_sha256','0'*64), ('rendition_sha256','0'*64), ('bindings',[])]:
        receipt = json.loads(view.receipt_bytes); receipt[key] = changed
        with pytest.raises(CompatibilityError, match='receipt'):
            rederive_compatibility_view(path, json.dumps(receipt).encode())
    document = Document(path); document.add_paragraph('Different original source'); document.save(path)
    with pytest.raises(CompatibilityError, match='receipt'):
        rederive_compatibility_view(path, view.receipt_bytes)


def wrapper(requires='wps', namespace=WPS[1:-1], parent=None):
    node = ET.SubElement(parent, MC+'AlternateContent', nsmap={'wps':namespace}) if parent is not None else ET.Element(MC+'AlternateContent', nsmap={'wps':namespace})
    ET.SubElement(node, MC+'Choice', Requires=requires)
    ET.SubElement(node, MC+'Fallback')
    return node


def test_uri_alias_not_prefix_spelling():
    root = ET.Element('root', nsmap={'shape':WPS[1:-1]})
    node = ET.SubElement(root, MC+'AlternateContent')
    ET.SubElement(node, MC+'Choice', Requires='shape'); ET.SubElement(node, MC+'Fallback')
    assert _selection(node)[1] == (WPS[1:-1],)


@pytest.mark.parametrize('requires,namespace', [('wps','urn:unknown'), ('unbound',WPS[1:-1]), ('',WPS[1:-1]), ('wps wps',WPS[1:-1]), ('wps unknown',WPS[1:-1]), ('wps:bad',WPS[1:-1])])
def test_requires_rejected(requires, namespace):
    with pytest.raises(CompatibilityError): _selection(wrapper(requires,namespace))


@pytest.mark.parametrize('change', ['ancestor_override','reordered','extra_choice','missing_fallback','wrapper_attribute'])
def test_structure_does_not_fall_through(change):
    root = ET.Element('root'); node = wrapper(parent=root)
    if change == 'ancestor_override': root.set(MC+'ProcessContent','wps:*')
    if change == 'reordered': node.insert(0,node[-1])
    if change == 'extra_choice': node.insert(1,deepcopy(node[0]))
    if change == 'missing_fallback': node.remove(node[-1])
    if change == 'wrapper_attribute': node.set('unknown','payload')
    with pytest.raises(CompatibilityError): _selection(node)


@pytest.mark.parametrize('kind', ['two_images','nested_frame','sibling_picture','bodypr_picture','unscaffolded_picture','wrapper_inside_frame',
    'structural_text','extra_text','command','instruction','unknown','hidden_text','effect','extra_layer','company_label','external','missing',
    'fallback_text','fallback_extra_image','fallback_command','nested_alternate'])
def test_all_branch_unknown_content_rejected(tmp_path, kind):
    path = source(tmp_path, divergent=True, mutation=mutate(kind))
    with pytest.raises(CompatibilityError): derive_compatibility_view(path.read_bytes())


@pytest.mark.parametrize('kind', ['unknown_attribute','action','foreign_attribute','wrong_extension','invalid_extent','deleted_command','nonrect','hidden_fallback','wrong_kind','custom_image_kind','duplicate_relationship','misplaced_dpi','duplicate_ppr','duplicate_rpr','competing_line_fills'])
def test_additional_closed_grammar_negatives(tmp_path, kind):
    def mutation(document, slot, wrapper, fallback):
        if kind == 'unknown_attribute': next(slot.iter(WPS+'wsp')).set('mystery','data')
        elif kind == 'action': next(slot.iter(qn('wp:docPr'))).set('onclick','command')
        elif kind == 'foreign_attribute': next(slot.iter(WPS+'wsp')).set('{urn:unknown}payload','data')
        elif kind == 'wrong_extension': ET.SubElement(ET.SubElement(next(slot.iter(A+'blip')), A+'extLst'), A+'ext', uri='urn:unknown')
        elif kind == 'invalid_extent': next(slot.iter(qn('wp:extent'))).set('cx','99999999999')
        elif kind == 'deleted_command': ET.SubElement(ET.SubElement(fallback, qn('w:del')), qn('w:instrText')).text='DDEAUTO command'
        elif kind == 'nonrect': next(slot.iter(A+'prstGeom')).set('prst','ellipse')
        elif kind == 'misplaced_dpi': ET.SubElement(slot.find('.//'+A+'xfrm/'+A+'ext'), '{http://schemas.microsoft.com/office/drawing/2010/main}useLocalDpi',val='0')
        elif kind == 'duplicate_ppr':
            p=next(slot.iter(qn('w:p')));ET.SubElement(p,qn('w:pPr'));ET.SubElement(p,qn('w:pPr'))
        elif kind == 'duplicate_rpr':
            r=next(slot.iter(qn('w:r')));ET.SubElement(r,qn('w:rPr'));ET.SubElement(r,qn('w:rPr'))
        elif kind == 'competing_line_fills':
            line=ET.SubElement(next(slot.iter(WPS+'spPr')),A+'ln',w='6350')
            ET.SubElement(ET.SubElement(line,A+'solidFill'),A+'schemeClr',val='lt1')
            ET.SubElement(ET.SubElement(line,A+'solidFill'),A+'prstClr',val='black')
        elif kind == 'hidden_fallback': ET.SubElement(ET.SubElement(next(fallback.iter(qn('w:r'))), qn('w:rPr')), qn('w:vanish'))
    path = source(tmp_path, divergent=True, mutation=mutation)
    if kind in ('wrong_kind','custom_image_kind','duplicate_relationship'):
        with ZipFile(path) as archive: parts={n:archive.read(n) for n in archive.namelist()}
        name='word/_rels/document.xml.rels'; root=ET.fromstring(parts[name]); image=next(n for n in root if n.get('Type','').endswith('/image'))
        if kind == 'wrong_kind': image.set('Type','urn:unknown')
        elif kind == 'custom_image_kind': image.set('Type','https://unrecognized.invalid/image')
        else: root.append(deepcopy(image))
        parts[name]=ET.tostring(root)
        with ZipFile(path,'w',ZIP_DEFLATED) as archive:
            for name,raw in parts.items(): archive.writestr(name,raw)
    with pytest.raises((CompatibilityError,ValueError)): derive_compatibility_view(path.read_bytes())


def test_contents_frame_keeps_native_role_and_both_original_records(tmp_path):
    from test_monthly_report_native_design_import import frame
    path = source(tmp_path, divergent=True, shared_outside=False)
    document = Document(path)
    # Remove the later body heading/prose; the new Word section is native contents.
    for paragraph in list(document.paragraphs)[1:]:
        paragraph._p.getparent().remove(paragraph._p)
    document.add_section()
    p = document.add_paragraph()
    for value in ('ENFRA', 'Section 1: Work Summary'):
        p.add_run()._r.append(frame(value))
    slot=frame(''); p.add_run()._r.append(slot)
    _picture_in_frame(document,slot,raster('orange'))
    parent=slot.getparent(); ac=ET.Element(MC+'AlternateContent',nsmap={'wps':WPS[1:-1]});parent.replace(slot,ac)
    ET.SubElement(ac,MC+'Choice',Requires='wps').append(slot)
    fallback=ET.SubElement(ac,MC+'Fallback');fallback.append(deepcopy(slot))
    rid,_=document.part.get_or_add_image(__import__('io').BytesIO(raster('green','JPEG')))
    next(fallback.iter(A+'blip')).set(qn('r:embed'),rid)
    document.save(path)
    view=derive_compatibility_view(path.read_bytes())
    assert len([b for b in view.bindings if b.status=='resolved']) == 2
    original_images=[i for i in view.original_inspection.items if i.kind=='image']
    assert original_images[0].image_role=='client_logo'
    assert original_images[1].image_role==''
    assert [b.original_id for b in view.bindings] == [i.id for i in view.original_inspection.items]


def test_legacy_cache_is_bounded_audited_never_decoded_or_forwarded(tmp_path):
    from io import BytesIO
    from app.monthly_report_compatibility_view import O,V
    def legacy(document,slot,wrapper,fallback):
        box=next(fallback.iter(qn('w:txbxContent')))
        fallback.remove(fallback[0])
        shape=ET.SubElement(ET.SubElement(fallback,qn('w:pict')),V+'shape',
            type='#_x0000_t202',id='Text Box 1',style='position:absolute;width:100pt;height:50pt')
        shape.set(O+'gfxdata','dGhpcyBpcyBu b3QgYSB6aXA=\n')
        ET.SubElement(shape,V+'textbox').append(box)
        from test_monthly_report_static_siblings import legacy as static_legacy
        declaration = next(static_legacy('horizontal-textframe', declaration=True).iter(V+'shapetype'))
        ET.SubElement(document.add_paragraph().add_run()._r, qn('w:pict')).append(declaration)
    path=source(tmp_path,divergent=True,mutation=legacy)
    view=derive_compatibility_view(path.read_bytes())
    receipt=json.loads(view.receipt_bytes)
    assert receipt['resolutions'][0]['branches'][1]['opaque_caches']
    with ZipFile(BytesIO(view.rendition_bytes)) as archive:
        assert b'gfxdata' not in archive.read('word/document.xml')
    with ZipFile(path) as archive:
        assert b'gfxdata' in archive.read('word/document.xml')


def test_duplicate_zip_member_fails_before_derivation(tmp_path):
    path=source(tmp_path,divergent=True)
    with ZipFile(path,'a') as archive:
        with pytest.warns(UserWarning): archive.writestr('word/document.xml',archive.read('word/document.xml'))
    with pytest.raises(ValueError): derive_compatibility_view(path.read_bytes())


def test_known_static_paragraph_and_picture_styling_is_retained(tmp_path):
    def static_style(document,slot,wrapper,fallback):
        for branch in wrapper:
            paragraph=next(branch.iter(qn('w:p')))
            props=ET.SubElement(paragraph,qn('w:pPr'))
            ET.SubElement(ET.SubElement(props,qn('w:rPr')),qn('w:b'))
            picprops=next(branch.iter(qn('pic:spPr')))
            ET.SubElement(picprops,A+'noFill')
            ET.SubElement(ET.SubElement(picprops,A+'ln',w='6350'),A+'noFill')
    path=source(tmp_path,divergent=True,mutation=static_style)
    view=derive_compatibility_view(path.read_bytes())
    assert any(b.status=='resolved' for b in view.bindings)


@pytest.mark.parametrize("budget", ["MAX_VIEW_SOURCE_BYTES", "MAX_VIEW_EXPANDED_BYTES", "MAX_VIEW_XML_BYTES", "MAX_VIEW_MEDIA_BYTES", "MAX_VIEW_XML_NODES"])
def test_capability_budgets_precede_inspection_and_image_work(tmp_path, monkeypatch, budget):
    import app.monthly_report_compatibility_view as module
    path=source(tmp_path,divergent=True)
    monkeypatch.setattr(module,budget,1)
    monkeypatch.setattr(module,'inspect_docx',lambda path: (_ for _ in ()).throw(AssertionError('inspection started before budget')))
    with pytest.raises(CompatibilityError,match='budget'):
        module.derive_compatibility_view(path.read_bytes())


def test_equivalent_sibling_cannot_hide_unknown_attribute(tmp_path):
    def mutation(document, slot, wrapper, fallback):
        sibling=deepcopy(wrapper)
        sibling[-1].clear()
        sibling[-1].append(deepcopy(sibling[0][0]))
        for branch in sibling:
            next(branch.iter(WPS+'wsp')).set('mystery','unrecognized payload')
        wrapper.getparent().append(sibling)
    path=source(tmp_path,divergent=True,mutation=mutation)
    with pytest.raises(CompatibilityError,match='unknown or invalid attribute'):
        derive_compatibility_view(path.read_bytes())


def test_subtree_digest_binds_parent_child_structure():
    from app.monthly_report_compatibility_view import _subtree_digest
    first=ET.fromstring('<root><node><leaf/></node></root>')
    second=ET.fromstring('<root><node/><leaf/></root>')
    assert _subtree_digest(first)!=_subtree_digest(second)
