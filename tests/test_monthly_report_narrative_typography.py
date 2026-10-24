"""Synthetic regressions for source lead-in formatting bleeding into new prose."""
from dataclasses import replace
from io import BytesIO

import pytest
from docx import Document
from docx.shared import Pt

from app.monthly_report_import import ImportMapping, inspect_docx, map_items
from app.monthly_report_model import (Facility, ReportDraft, ReportPeriod,
    ReportProfile, ResolvedBlock, default_sections)
from app.monthly_report_native_layout import build_native_docx


def scene(tmp_path, *, body='The previous month included a synthetic inspection and a reviewed repair.',
          underlined=None, mixed=False):
    doc = Document()
    doc.add_heading('MONTHLY ACTIVITY SUMMARY', 1)
    label = doc.add_paragraph('Activity Summary')
    label.runs[0].underline = True
    lead = doc.add_paragraph('Location/Campus/Building Equipment Synthetic Plant:')
    lead.runs[0].underline = True
    lead.runs[0].font.size = Pt(9)
    if body:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Pt(6)
        if mixed:
            prefix = paragraph.add_run('Note: ')
            prefix.underline = prefix.bold = True
            prefix.font.size = Pt(12)
        run = paragraph.add_run(body)
        run.font.size = Pt(12)
        run.underline = underlined
    source = tmp_path / 'synthetic-activity.docx'
    doc.save(source)
    draft = ReportDraft(ReportProfile('Synthetic', 'site', 'Synthetic Site',
        (Facility('site', 'Synthetic Site'),)), ReportPeriod(2026, 10),
        'Synthetic Editor', default_sections(),
        (ResolvedBlock('activity_summary', 'This month',
            text='Current synthetic inspection was completed.\n\nCurrent synthetic follow-up remains pending.'),))
    return source, draft


@pytest.mark.parametrize('section', [None, 'activity'])
@pytest.mark.parametrize('body', ['Completed.', 'The previous month included a synthetic inspection and a reviewed repair.'])
def test_rewritten_body_does_not_inherit_lead_in_underline_or_size(tmp_path, section, body):
    source, draft = scene(tmp_path, body=body)
    before = source.read_bytes()
    output = Document(BytesIO(build_native_docx(draft, source, section_key=section)))
    current = [p for p in output.paragraphs if p.text.startswith('Current synthetic')]
    assert len(current) == 2
    assert all(p.runs[0].underline is not True for p in current)
    assert all(p.runs[0].font.size == Pt(12) for p in current)
    assert all(p.paragraph_format.left_indent == Pt(6) for p in current)
    label = next(p for p in output.paragraphs if p.text == 'Activity Summary')
    assert label.runs[0].underline is True
    assert source.read_bytes() == before


def test_legitimate_underlined_body_is_not_blanket_stripped(tmp_path):
    source, draft = scene(tmp_path, underlined=True)
    output = Document(BytesIO(build_native_docx(draft, source, section_key='activity')))
    current = [p for p in output.paragraphs if p.text.startswith('Current synthetic')]
    assert len(current) == 2
    assert all(p.runs[0].underline is True and p.runs[0].font.size == Pt(12) for p in current)


@pytest.mark.parametrize('body', [
    'The following work was completed this month:', 'Inspection complete.',
    'Building/equipment checks completed this month:',
    'Campus/building work included the following:',
])
def test_real_prose_is_not_replaced_by_a_later_longer_label(tmp_path, body):
    source, draft = scene(tmp_path, body=body)
    document = Document(source)
    label = document.add_paragraph('An unusually long synthetic location label without punctuation')
    label.runs[0].underline = True
    label.runs[0].font.size = Pt(9)
    document.save(source)
    output = Document(BytesIO(build_native_docx(draft, source, section_key='activity')))
    current = [p for p in output.paragraphs if p.text.startswith('Current synthetic')]
    assert len(current) == 2
    assert all(p.runs[0].underline is not True and p.runs[0].font.size == Pt(12) for p in current)
    assert all(p.paragraph_format.left_indent == Pt(6) for p in current)


def test_fragmented_body_does_not_promote_longer_inline_emphasis(tmp_path):
    source, draft = scene(tmp_path, body='Plain opening. ')
    document = Document(source)
    paragraph = document.paragraphs[-1]
    for text in ('plain middle. ', 'plain ending. '):
        paragraph.add_run(text).font.size = Pt(12)
    emphasized = paragraph.add_run('Emphasized inline phrase')
    emphasized.underline = emphasized.bold = True
    emphasized.font.size = Pt(12)
    document.save(source)
    output = Document(BytesIO(build_native_docx(draft, source, section_key='activity')))
    current = [p for p in output.paragraphs if p.text.startswith('Current synthetic')]
    assert len(current) == 2
    assert all(p.runs[0].underline is not True and p.runs[0].bold is not True for p in current)


def test_label_only_source_uses_neutral_fallback(tmp_path):
    source, draft = scene(tmp_path, body='')
    output = Document(BytesIO(build_native_docx(draft, source, section_key='activity')))
    current = [p for p in output.paragraphs if p.text.startswith('Current synthetic')]
    assert len(current) == 2
    assert all(p.runs[0].underline is not True for p in current)


def test_proven_unchanged_source_keeps_labels_and_inline_formatting(tmp_path):
    source, draft = scene(tmp_path, mixed=True)
    inspection = inspect_docx(source)
    mappings = tuple(ImportMapping(i.id, 'activity_summary') for i in inspection.items
        if i.kind == 'text' and i.section == 'activity' and i.note != 'Native section heading')
    mapped = map_items(source, inspection, mappings)
    before = source.read_bytes()
    output = Document(BytesIO(build_native_docx(replace(draft, blocks=mapped.blocks), source, section_key='activity')))
    originals = [p for p in Document(source).paragraphs if p.text and p.text != 'MONTHLY ACTIVITY SUMMARY']
    for original in originals:
        current = next(p for p in output.paragraphs if p.text == original.text)
        assert current._p.xml == original._p.xml
    assert source.read_bytes() == before
