"""First-use and returning reports keep one editing home per section."""

from dataclasses import replace

import pytest
from streamlit.testing.v1 import AppTest

from app import monthly_report_library as library
from app.contracts import RRH_CONTRACT
from app.monthly_report_guided import STEPS, complete_guided_sections, review_destination
from app.monthly_report_model import (BlockSpec, ColumnSpec, ReportDraft, ReportFollowUp, ReportPeriod, ReportSource,
                                     ResolvedBlock, default_sections, synthetic_draft, synthetic_profiles)
from app.monthly_report_sources import SourceContent
from test_monthly_report_ui import monthly, step


def choose_section(app, key):
    assert any(section.key == key for section in working_state(app)[1].sections)
    app.run()
    assert not app.exception


def working_state(app):
    key = next(key for key in app.session_state.filtered_state if key.startswith("report_guided_") and key.endswith("_draft") and isinstance(app.session_state[key], ReportDraft))
    return key, app.session_state[key]


def include(app, *titles):
    # Every standard section is now present, including in older saved designs.
    step(app, 1)
    sections = working_state(app)[1].sections
    assert all(any(section.title == title and section.included for section in sections) for title in titles)
    assert not any(value.label.startswith("Include ") for value in app.checkbox)


def test_first_template_has_every_section_visible_with_one_editing_home(monkeypatch, tmp_path):
    app = monthly(monkeypatch, tmp_path, saved=False)
    next(value for value in app.text_input if value.label == "Site name").set_value("Synthetic First Site").run()
    next(value for value in app.button if value.label == "Use ENFRA template").click().run()
    next(value for value in app.text_input if value.label == "Your name").set_value("Synthetic Editor").run()
    next(value for value in app.button if value.label == "Start this report").click().run()
    step(app, 2)
    assert sum(value.label == "Activity summary" for value in app.text_area) == 1
    assert not any("logo" in value.label.lower() for value in app.get("file_uploader"))
    sections = working_state(app)[1].sections
    headings = [value.value for value in app.subheader]
    assert all(headings.count(section.title) == 1 for section in sections)
    assert [title for title in headings if title in {section.title for section in sections}] == [section.title for section in sections]
    assert not any(value.label in ("Section to update", "Site information to update") for value in app.selectbox)
    assert any(value.value == "Cover and report details" for value in app.subheader)


def test_unsaved_text_and_source_state_survive_section_switches_and_direct_save(monkeypatch, tmp_path):
    app = monthly(monkeypatch, tmp_path)
    include(app, "Training Summary", "Monthly Scorecards")
    next(value for value in app.text_input if value.label == "Prepared by").set_value("Synthetic Editor").run()
    step(app, 2)
    assert not any(value.label == "Save this draft for others on this report to continue" for value in app.checkbox)
    assert not any(value.label == "Show writing help" for value in app.toggle)
    next(value for value in app.text_area if value.label == "Activity summary").set_value("Unsaved pump repair.").run()
    draft_key, _ = working_state(app)
    source = ReportSource("a" * 64, "synthetic.txt", "a" * 64, ".txt", page_texts=("Checked the pump.",))
    app.session_state[draft_key.removesuffix("_draft") + "_evidence"] = (SourceContent(source),)
    choose_section(app, "training")
    assert all(any(value.label == label for value in app.text_area) for label in ("Activity summary", "Utility analysis"))
    next(value for value in app.text_area if value.label == "Training notes").set_value("Completed a controls workshop.").run()
    choose_section(app, "scorecards")
    assert any(value.label == "Utility analysis" for value in app.text_area)
    assert sum(value.label == "Training notes" for value in app.text_area) == 1
    choose_section(app, "activity")
    assert next(value for value in app.text_area if value.label == "Activity summary").value == "Unsaved pump repair."
    assert working_state(app)[1].sources == (source,)
    next(value for value in app.button if value.label == "Save progress").click().run()
    saved = library.load_snapshot(RRH_CONTRACT, "synthetic-guided", ReportPeriod(2026, 9))
    blocks = {block.key: block for block in saved.draft.blocks}
    assert blocks["activity_summary"].text == "Unsaved pump repair."
    assert blocks["training_summary"].text == "Completed a controls workshop."
    assert saved.draft.sources == (source,)


def test_legacy_excluded_active_section_is_restored_without_losing_content(monkeypatch, tmp_path):
    app = monthly(monkeypatch, tmp_path)
    include(app, "Training Summary")
    next(value for value in app.text_input if value.label == "Prepared by").set_value("Synthetic Editor").run()
    step(app, 2)
    choose_section(app, "training")
    next(value for value in app.text_area if value.label == "Training notes").set_value("Completed safety training.").run()
    draft_key, draft = working_state(app)
    app.session_state[draft_key] = replace(draft, sections=tuple(
        replace(section, included=False) if section.key == "training" else section for section in draft.sections
    ))
    app.run()
    assert not any(value.label == "Section to update" for value in app.selectbox)
    assert all(section.included for section in working_state(app)[1].sections)
    assert next(value for value in app.text_area if value.label == "Training notes").value == "Completed safety training."


def test_capacity_lives_with_utility_results_without_standing_confirmation(monkeypatch, tmp_path):
    app = monthly(monkeypatch, tmp_path)
    include(app, "Monthly Scorecards")
    step(app, 3)
    assert not any(value.label == "I checked the standing information for these sites" for value in app.checkbox)
    assert not any(value.label == "Known equipment tags" for value in app.text_area)
    step(app, 2)
    choose_section(app, "scorecards")
    assert any(value.label == "Utility analysis" for value in app.text_area)
    capacity_key = next(key for key in app.session_state.filtered_state
                        if "_edit_thermal_capacity_table_" in key and not key.endswith(("_data", "_seed")))
    app.session_state[capacity_key] = {"edited_rows": {}, "added_rows": [{"Facility": "Synthetic North", "Service": "Heating", "Capacity": 200.0, "Units": "kW"}], "deleted_rows": []}
    app.run()
    assert not app.exception
    updated = working_state(app)[1]
    capacity = next(block for block in updated.blocks if block.key == "thermal_capacity")
    assert capacity.rows and "200" in capacity.rows[0][2]
    step(app, 3)
    assert not any(value.label == "I checked the standing information for these sites" for value in app.checkbox)
    assert any("_edit_thermal_capacity_table_" in (value.key or "") for value in app.dataframe)


def test_visible_followups_still_block_download_and_each_updates_independently(monkeypatch, tmp_path):
    app = monthly(monkeypatch, tmp_path)
    include(app, "Equipment Performance Issues", "Pending & Declined Proposals")
    next(value for value in app.text_input if value.label == "Prepared by").set_value("Synthetic Editor").run()
    draft_key, draft = working_state(app)
    issue = ReportFollowUp("follow-synthetic-pump", "issue", "Pump vibration.", "2026-08")
    proposal = ReportFollowUp("follow-synthetic-controls", "proposal", "Controls renewal.", "2026-08", status="pending")
    app.session_state[draft_key] = replace(draft, follow_ups=(issue, proposal))
    step(app, 2)
    assert all(any(value.label == label for value in app.button) for label in ("Still open", "Update proposal"))
    step(app, 4)
    assert next(value for value in app.button if value.label == "Generate DOCX and PDF").disabled
    assert review_destination(working_state(app)[1], issue.key) == (STEPS[1], "issues")
    assert any(value.value == "Equipment Performance Issues" for value in app.subheader)
    assert any(value.label == "Still open" for value in app.button)
    assert any(value.label == "Update proposal" for value in app.button)
    next(value for value in app.button if value.label == "Still open").click().run()
    assert working_state(app)[1].follow_ups[1] == proposal


def test_save_keeps_name_and_revision_gates_without_extra_confirmation(monkeypatch, tmp_path):
    app = monthly(monkeypatch, tmp_path)
    assert next(value for value in app.button if value.label == "Save progress").disabled
    next(value for value in app.text_input if value.label == "Prepared by").set_value("Synthetic Editor").run()
    assert not next(value for value in app.button if value.label == "Save progress").disabled
    _, draft = working_state(app)
    library.save_snapshot(draft, expected_revision=0, entered_editor="Another synthetic editor")
    app.run()
    assert next(value for value in app.button if value.label == "Save progress").disabled
    assert any("Another version was saved" in value.value for value in app.warning)


@pytest.mark.parametrize("allowed", ({"activity_summary"}, {"utility_analysis"}))
def test_completed_source_read_survives_when_writing_help_is_hidden(monkeypatch, allowed):
    from app import monthly_report_ai_ui as ai_ui

    def finish_read(draft, prefix):
        import streamlit as st
        source = replace(draft.sources[0], page_texts=("Newly read pump inspection.",), needs_vision=())
        st.session_state[prefix + "_evidence"] = (SourceContent(source),)

    monkeypatch.setattr(ai_ui, "_finish_job", finish_read)
    app = AppTest.from_string('''
import streamlit as st
from app.monthly_report_ai_ui import render_drafting
from app.monthly_report_model import ReportDraft, ReportPeriod, ReportSource, synthetic_profiles
from app.monthly_report_ui import _field
source = ReportSource("a" * 64, "synthetic.pdf", "a" * 64, ".pdf", page_texts=("",), needs_vision=(1,))
draft = ReportDraft(synthetic_profiles()[0], ReportPeriod(2026, 9), "Synthetic Editor", (), (), sources=(source,))
st.session_state["result"], _ = render_drafting(draft, {}, "synthetic", _field, allowed_keys=ALLOWED)
'''.replace("ALLOWED", repr(allowed))).run()
    assert not app.exception
    assert app.session_state["result"].sources[0].page_texts == ("Newly read pump inspection.",)


def test_partial_legacy_sections_restore_standard_headings_without_rewriting_content():
    profile = replace(synthetic_profiles()[0], section_order=("training", "activity"),
                      excluded_sections=("training", "activity"),
                      block_overrides=(BlockSpec("activity_summary", "rich_text", required=True),))
    draft = synthetic_draft(profile, ReportPeriod(2026, 9))
    originals = {section.key: section for section in default_sections()}
    custom = BlockSpec("training_summary", "table", required=True, columns=(ColumnSpec("topic", "Topic"),))
    content = ResolvedBlock("training_summary", "This month", rows=(("Synthetic Controls Workshop",),))
    draft = replace(draft, sections=(replace(originals["training"], blocks=(custom,), included=False),
                                     replace(originals["activity"], included=False)), blocks=(content,))
    normalized = complete_guided_sections(draft)
    assert [section.key for section in normalized.sections[:2]] == ["training", "activity"]
    assert {section.key for section in normalized.sections} == set(originals)
    assert all(section.included for section in normalized.sections)
    assert all(not block.required for section in normalized.sections for block in section.blocks)
    assert normalized.sections[0].blocks[0].columns == custom.columns
    assert normalized.blocks == draft.blocks
    assert normalized.profile == draft.profile
    assert complete_guided_sections(normalized) == normalized


def test_live_previews_receive_each_section_and_current_edits(monkeypatch, tmp_path):
    from app import monthly_report_preview as whole_preview, monthly_report_section_preview_ui as section_preview
    previews, full_previews = [], []
    app = monthly(monkeypatch, tmp_path)
    next(item for item in app.text_input if item.label == "Prepared by").set_value("Synthetic Editor").run()
    monkeypatch.setattr(section_preview, "render_section_preview", lambda draft, key, assets, prefix, **kwargs:
                        previews.append((draft, key, kwargs)))
    monkeypatch.setattr(whole_preview, "render_preview", lambda draft, assets, prefix, field:
                        full_previews.append(draft))
    app.run()
    keys = [key for _, key, _ in previews]
    expected = ["cover", *(section.key for section in working_state(app)[1].sections)]
    assert keys == expected
    assert all(options.get("deferred") for _, _, options in previews)
    assert not full_previews
    assert not any("parts have content" in item.value or "sections selected" in item.value for item in app.caption)
    assert not any(item.label.startswith("Include ") for item in app.checkbox)
    assert sum(item.label == "Starting report" for item in app.get("file_uploader")) == 1
    previews.clear()
    next(item for item in app.text_area if item.label == "Activity summary").set_value("Synthetic current repair.").run()
    assert [key for _, key, _ in previews] == expected
    activity = next(draft for draft, key, _ in previews if key == "activity")
    assert next(block for block in activity.blocks if block.key == "activity_summary").text == "Synthetic current repair."
    assert not full_previews
    assert next(block for block in working_state(app)[1].blocks if block.key == "activity_summary").text == "Synthetic current repair."


def test_blank_sections_generate_without_empty_or_standing_confirmations_but_pricing_still_blocks(monkeypatch, tmp_path):
    from io import BytesIO
    from docx import Document
    from app import monthly_report_guided as guided
    from app.monthly_report_docx import ReportPackage, assemble_docx
    generated = []

    def generate(draft, **kwargs):
        result = assemble_docx(draft, **kwargs)
        generated.append((draft, result))
        return ReportPackage(result, b"%PDF-synthetic", "synthetic.docx", "synthetic.pdf", draft.fingerprint)

    monkeypatch.setattr(guided, "generate_report", generate)
    app = monthly(monkeypatch, tmp_path)
    next(item for item in app.text_input if item.label == "Prepared by").set_value("Synthetic Editor").run()
    draft_key, draft = working_state(app)
    app.session_state[draft_key] = replace(draft, blocks=())
    step(app, 4)
    assert not any("included but empty" in item.value or "standing site information" in item.value for item in app.warning)
    assert not any(item.label in ("I checked these specific warnings", "I checked the standing information for these sites") for item in app.checkbox)
    assert not next(item for item in app.button if item.label == "Generate DOCX and PDF").disabled
    next(item for item in app.button if item.label == "Generate DOCX and PDF").click().run()
    assert not app.exception
    generated_draft, raw = generated[0]
    paragraphs = [paragraph.text for paragraph in Document(BytesIO(raw)).paragraphs]
    assert len(generated_draft.sections) == len(default_sections())
    assert all(any(section.title in text for text in paragraphs) for section in generated_draft.sections)
    draft_key, draft = working_state(app)
    # The section editor is always mounted; exercise the actual edit rather than
    # changing its underlying draft behind an existing widget value.
    next(item for item in app.text_area if item.label == "Activity summary").set_value("Synthetic service cost $100.").run()
    assert next(item for item in app.button if item.label == "Generate DOCX and PDF").disabled
    assert any("pricing" in item.value.lower() for item in app.error)


def test_saved_contacts_fill_once_after_preparer_name_without_overwriting_later_edits(monkeypatch, tmp_path):
    from app import monthly_report_directory as directory
    attempts = []
    original = directory.apply_defaults

    def apply_defaults(draft):
        attempts.append(draft.prepared_by)
        return original(draft)

    monkeypatch.setattr(directory, "apply_defaults", apply_defaults)
    app = monthly(monkeypatch, tmp_path)
    assert not attempts
    contact = directory.DirectoryContact("Asset Manager", "Synthetic Contract Lead", email="lead@example.invalid")
    directory.save_directory(RRH_CONTRACT, (), contract_contacts=(contact,), expected_revision=0,
                             actor="Synthetic Directory Editor", confirmed=True)
    next(item for item in app.text_input if item.label == "Prepared by").set_value("Synthetic Editor").run()
    assert attempts == ["Synthetic Editor"]
    draft_key, draft = working_state(app)
    contacts = next(block for block in draft.blocks if block.key == "contact_matrix")
    assert contacts.rows[0][2] == contact.name
    original_facilities = draft.profile.facilities
    app.session_state[draft_key] = replace(draft, blocks=tuple(
        replace(block, rows=()) if block.key == "contact_matrix" else block for block in draft.blocks
    ))
    step(app, 2)
    assert attempts == ["Synthetic Editor"]
    assert next(block for block in working_state(app)[1].blocks if block.key == "contact_matrix").rows == ()
    assert working_state(app)[1].profile.facilities == original_facilities


def test_saved_cleared_contacts_and_capacity_stay_blank_after_reopen_and_rollover(monkeypatch, tmp_path):
    from app import monthly_report_capacity as capacity, monthly_report_directory as directory
    from test_monthly_report_ui import ROOT, choose_report

    app = monthly(monkeypatch, tmp_path)
    draft_key, draft = working_state(app)
    profile = draft.profile
    contact = directory.DirectoryContact("Asset Manager", "Synthetic Contract Lead", email="lead@example.invalid")
    directory.save_directory(profile.contract, (), contract_contacts=(contact,), expected_revision=0,
                             actor="Synthetic Directory Editor", confirmed=True)
    raw = ("Facility,Service,Required capacity,Available capacity,Units\n"
           + profile.facilities[0].title + ",Steam,120,180,lb/hr\n").encode()
    inspection = capacity.inspect_capacity(profile, "synthetic-capacity.csv", raw)
    capacity.save_capacity(profile, inspection.tables, expected_revision=0, actor="Synthetic Editor",
                           reviewed_fingerprint=capacity.review_fingerprint(inspection.tables))
    del app.session_state[draft_key]
    app.run()
    next(item for item in app.text_input if item.label == "Prepared by").set_value("Synthetic Editor").run()
    _, draft = working_state(app)
    seeded = {block.key: block for block in draft.blocks}
    assert seeded["thermal_capacity"].extra_tables
    assert seeded["contact_matrix"].rows
    cleared = {"thermal_capacity", "contact_matrix"}
    app.session_state[draft_key] = replace(draft, blocks=tuple(
        replace(block, source="This month", text="", rows=(), extra_tables=()) if block.key in cleared else block
        for block in draft.blocks
    ))
    next(item for item in app.button if item.label == "Save progress").click().run()
    assert not app.exception
    saved = library.load_snapshot(profile.contract, profile.key, ReportPeriod(2026, 9))
    assert all(not block.rows and not block.extra_tables for block in saved.draft.blocks if block.key in cleared)

    reopened = AppTest.from_file(ROOT / "run_web.py", default_timeout=20).run()
    reopened.segmented_control[0].set_value("Monthly report").run()
    choose_report(reopened, profile.facilities[0].title)
    next(item for item in reopened.text_input if item.label == "Prepared by").set_value("Synthetic Returning Editor").run()
    assert not reopened.exception
    _, continued = working_state(reopened)
    assert all(not block.rows and not block.extra_tables for block in continued.blocks if block.key in cleared)
    reopened.selectbox("report_month_number").set_value(10).run()
    assert not reopened.exception
    current = next(value for key, value in reopened.session_state.filtered_state.items()
                   if key.startswith("report_guided_") and key.endswith("_draft") and isinstance(value, ReportDraft) and value.period == ReportPeriod(2026, 10))
    assert all(not block.rows and not block.extra_tables for block in current.blocks if block.key in cleared)
    assert current.profile.facilities == profile.facilities
