"""Invented missing readings must not turn engineering columns into prices."""
from dataclasses import asdict, replace
from io import BytesIO

import pytest
from docx import Document

from app.monthly_report_content_policy import (
    _table_headings, ambiguous_price_columns, price_free_table,
    schema_price_columns, table_has_pricing, table_price_columns,
)
from app.monthly_report_import import inspect_docx
from app.monthly_report_master_tables import from_inspection
from app.monthly_report_model import (
    BlockSpec, ColumnSpec, Facility, ReportDraft, ReportPeriod, ReportProfile,
    ReportTable, ResolvedBlock, default_sections,
)
from app.monthly_report_native_layout import build_native_docx
from conftest import requires_libreoffice


@pytest.mark.parametrize('dash', ('-', '–', '—', '  —  ', '\t-\n'))
@pytest.mark.parametrize('unit', ('Tons', 'MBH'))
def test_missing_reading_does_not_change_complete_engineering_heading(dash, unit):
    columns = ('Facility', f'Est. Peak Coincident Load ({unit})', 'Cost')
    rows = (('Demo Missing', dash, dash), ('Demo Reading', '2837', '9517'))
    before = (columns, rows)
    assert _table_headings(columns, rows)[1] == columns[1]
    assert table_price_columns(columns, rows) == (2,)
    assert schema_price_columns(columns, rows) == (2,)
    assert table_has_pricing(columns, rows)
    spec = BlockSpec('thermal_capacity', 'table', columns=tuple(
        ColumnSpec(str(i), title) for i, title in enumerate(columns)))
    block = ResolvedBlock('thermal_capacity', 'This month', rows=rows,
                          extra_tables=(ReportTable(columns, rows),))
    original = asdict(block)
    safe_spec, safe, removed = price_free_table(spec, block)
    assert tuple(c.title for c in safe_spec.columns) == columns[:2]
    assert safe.rows == tuple(row[:2] for row in rows)
    assert safe.extra_tables[0].rows == safe.rows
    assert removed == ('Cost',)
    assert asdict(block) == original and (columns, rows) == before


def test_split_headings_remain_inferred_around_missing_value_markers():
    columns = ('Facility', 'Est.', 'Unit')
    rows = (('', 'Peak Coincident Load (MBH)', 'Price'),
            ('Demo Missing', '—', '—'), ('Demo Reading', '2837', '9517'))
    assert table_price_columns(columns, rows) == (2,)
    assert schema_price_columns(columns, rows) == ()
    assert _table_headings(columns, rows)[1] == 'Est. Peak Coincident Load (MBH)'


@pytest.mark.parametrize('value', ('-12', '12–18', '$9517', 'Cost', 'Unit Price', '— Cost'))
def test_nonstandalone_dash_and_financial_continuations_are_not_ignored(value):
    columns = ('Facility', 'Est. Peak Coincident Load (MBH)')
    rows = (('Demo Missing', value), ('Demo Reading', '2837'))
    if any(c.isdigit() for c in value):
        assert _table_headings(columns, rows)[1] == columns[1]
        assert table_has_pricing(columns, rows) == value.startswith('$')
    else:
        assert value in _table_headings(columns, rows)[1]
        assert table_price_columns(columns, rows) == (1,)


@pytest.mark.parametrize('price', ('9517', '$9517', '—'))
def test_mixed_readings_never_hide_a_real_price_column(price):
    columns = ('Facility', 'Est. Peak Coincident Load (MBH)', 'Unit Price')
    rows = (('Demo Mixed', '—', price), ('Demo Reading', '2837', '9518'))
    assert table_price_columns(columns, rows) == (2,)
    assert schema_price_columns(columns, rows) == (2,)
    assert table_has_pricing(columns, rows)


def test_raw_currency_and_ambiguous_financial_guards_remain():
    engineering = ('Facility', 'Est. Peak Coincident Load (MBH)')
    rows = (('Demo Missing', '—'), ('Demo Reading', '2837'))
    assert schema_price_columns(engineering, rows, kinds=('text', 'currency')) == (1,)
    mixed = ('Facility', 'Description / Est. Peak Coincident Load (MBH) / Cost')
    assert ambiguous_price_columns(mixed, rows) == (1,)
    assert table_has_pricing(mixed, rows)
    for title in ('Cost / tons', 'Unit Price / kWh', 'Estimate', 'Amount', 'Rate'):
        assert table_price_columns(('Facility', title), rows) == (1,)


def _native_grid(tmp_path):
    columns = ('Campus', 'Capacity (MBH)', 'Available (MBH)',
               'Est. Peak Coincident Load (MBH)', 'Pressure (psi)', 'Units', 'Status')
    rows = (('Demo Missing', '—', '—', '—', '—', '—', '—'),) + tuple(
        (f'Demo Campus {i}', str(6000+i), str(5000+i), str(2837+i), str(100+i), '2', 'Ready')
        for i in range(1, 7))
    source = Document()
    source.add_heading('MONTHLY SCORECARDS', 1)
    source.add_paragraph('Steam')
    grid = source.add_table(rows=8, cols=7)
    for native_row, values in zip(grid.rows, (columns, *rows)):
        for cell, value in zip(native_row.cells, values):
            cell.text = value
    path = tmp_path / 'synthetic-steam.docx'
    source.save(path)
    original = path.read_bytes()
    inspection = inspect_docx(path)
    item, = (item for item in inspection.items if item.kind == 'table')
    reference = f'docx:{inspection.sha256}:{item.id}'
    profile = ReportProfile('Synthetic contract', 'demo', 'Demo Plant',
                            (Facility('demo', 'Demo Plant'),))
    draft = ReportDraft(profile, ReportPeriod(2026, 10), 'Synthetic Editor', default_sections(), (
        ResolvedBlock('thermal_capacity', 'This month', extra_tables=(
            ReportTable(columns, rows, reference),)),), synthetic=True)
    before = asdict(draft)
    # Model the final client projection, not a destructive edit to the saved draft.
    _, client_block, _ = price_free_table(
        BlockSpec('thermal_capacity', 'table'), draft.blocks[0])
    raw = build_native_docx(replace(draft, blocks=(client_block,)), path)
    (tmp_path / 'generated.docx').write_bytes(raw)
    assert asdict(draft) == before and path.read_bytes() == original
    return raw, (columns, *rows)


def _assert_grid(raw, expected):
    tables = Document(BytesIO(raw)).tables
    assert len(tables) == 1
    actual = tuple(tuple(cell.text for cell in row.cells) for row in tables[0].rows)
    assert actual == expected  # All 56 cells, including heading, dash and six loads.


def test_saved_complete_grid_survives_native_projection_without_mutation(tmp_path):
    raw, expected = _native_grid(tmp_path)
    _assert_grid(raw, expected)
    shape, = from_inspection(inspect_docx(tmp_path / 'synthetic-steam.docx'))
    assert shape.indices == tuple(range(7))


@requires_libreoffice
def test_native_pdf_preserves_six_loads_and_missing_reading(tmp_path):
    import fitz
    from app.pdf_converter import convert_to_pdf

    raw, expected = _native_grid(tmp_path)
    pdf_path = convert_to_pdf(tmp_path / 'generated.docx')
    pdf_bytes = pdf_path.read_bytes()
    (tmp_path / 'generated.pdf').write_bytes(pdf_bytes)
    with fitz.open(stream=pdf_bytes, filetype='pdf') as pdf:
        text = '\n'.join(page.get_text() for page in pdf)
    for row in expected[2:]:
        assert row[3] in text
    assert '—' in text
    _assert_grid(raw, expected)
