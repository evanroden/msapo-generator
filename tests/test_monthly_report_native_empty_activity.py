"""Empty Activity panels must disappear without disturbing native tables/art."""
from copy import deepcopy
from dataclasses import replace
from io import BytesIO
from types import SimpleNamespace

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import pytest

from conftest import requires_libreoffice
from app.monthly_report_import import _text, inspect_docx
from app.monthly_report_model import (
    Facility, ReportDraft, ReportPeriod, ReportProfile, ReportTable,
    ResolvedBlock, default_sections,
)
from app.monthly_report_native_empty_activity import empty_activity_panels
from app.monthly_report_native_layout import build_native_docx
from test_monthly_report_native_reflow import chart


TITLE = "Improvement Highlights and Pics"
CURRENT = "Current synthetic activity: tested the replacement pump."


@pytest.fixture
def report(tmp_path):
    source = Document()
    source.sections[0].header.paragraphs[0].text = "ENFRA"
    source.add_heading("MONTHLY ACTIVITY SUMMARY", 1)
    source.add_paragraph("Work Order Status")
    table = source.add_table(rows=4, cols=2)
    table.style = "Table Grid"
    cells = (("Ticket", "Status"), ("WO-001", "Open"), ("", ""), ("WO-002", "Done"))
    for row, values in zip(table.rows, cells):
        for cell, value in zip(row.cells, values):
            cell.text = value
    source.add_paragraph("Activity Summary")
    for number in range(25):
        source.add_paragraph(f"Prior synthetic activity paragraph {number}.")
    heading = source.add_heading(TITLE, 2)
    heading.add_run()._r.append(OxmlElement("w:lastRenderedPageBreak"))
    note = source.add_paragraph("Prior improvement description.")
    note.paragraph_format.page_break_before = True
    section = source.add_section(WD_SECTION_START.NEW_PAGE)
    section.header.is_linked_to_previous = False
    section.header.paragraphs[0].text = "Client Utility Rate Analysis"
    source.add_heading("TRAINING SUMMARY", 1)
    source.add_paragraph("Prior synthetic training.")
    path = tmp_path / "synthetic-source.docx"
    source.save(path)
    inspected = inspect_docx(path)
    item = next(item for item in inspected.items if item.kind == "table")
    native = ReportTable(cells[0], cells[1:], f"docx:{inspected.sha256}:{item.id}")
    profile = ReportProfile("Synthetic contract", "synthetic", "Synthetic Site",
                            (Facility("synthetic", "Synthetic Site"),))
    draft = ReportDraft(profile, ReportPeriod(2026, 10), "Synthetic Editor", default_sections(), (
        ResolvedBlock("work_orders", "This month", extra_tables=(native,)),
        ResolvedBlock("activity_summary", "This month", text=CURRENT),
        ResolvedBlock("training_summary", "This month", text="Current synthetic training."),
    ), synthetic=True)
    return path, draft


@pytest.mark.parametrize("scope", [None, "activity"])
@pytest.mark.parametrize("improvements", [None, ResolvedBlock("improvements", "This month"),
                                          ResolvedBlock("improvements", "Omit", text="Omitted")])
def test_empty_panel_does_not_leave_old_padding_beside_unchanged_work_orders(report, scope, improvements):
    path, draft = report
    original = path.read_bytes()
    source = Document(path)
    if improvements is not None:
        draft = replace(draft, blocks=(*draft.blocks, improvements))
    result = Document(BytesIO(build_native_docx(draft, path, section_key=scope)))
    assert TITLE not in [p.text for p in result.paragraphs]
    assert sum(p.text == CURRENT for p in result.paragraphs) == 1
    assert "Prior " not in "\n".join(p.text for p in result.paragraphs)
    activity = next(p._p for p in result.paragraphs if p.text == CURRENT)
    # No cleared source narrative or orphan page-break paragraph remains between
    # current activity and the original physical section boundary.
    following = activity.getnext()
    if scope == "activity":
        # Section-only export moves this boundary's properties to the terminal
        # sectPr; its original empty paragraph remains in the same position.
        assert following.tag == qn("w:p") and not _text(following)
        following = following.getnext()
    assert following.tag == qn("w:sectPr") or following.find(".//" + qn("w:sectPr")) is not None
    assert result.tables[0]._tbl.xml == source.tables[0]._tbl.xml
    if scope is None:
        assert [s._sectPr.xml for s in result.sections] == [s._sectPr.xml for s in source.sections]
        assert [s.header.paragraphs[0].text for s in result.sections] == [
            "ENFRA", "Client Utility Rate Analysis"]
        assert "Current synthetic training." in [p.text for p in result.paragraphs]
    assert path.read_bytes() == original


@requires_libreoffice
@pytest.mark.parametrize("scope,expected_pages", [(None, 2), ("activity", 1)])
def test_empty_panel_leaves_no_blank_page_in_actual_pdf(report, tmp_path, monkeypatch, scope, expected_pages):
    import pymupdf
    from app import pdf_converter
    path, draft = report
    monkeypatch.setattr(pdf_converter, "OUTPUT_DIR", tmp_path)
    monkeypatch.setattr(pdf_converter, "PDF_BACKEND", "libreoffice")
    output = tmp_path / "current-synthetic.docx"
    output.write_bytes(build_native_docx(draft, path, section_key=scope))
    with pymupdf.open(pdf_converter.convert_to_pdf(output)) as pdf:
        assert len(pdf) == expected_pages
        text = "\n".join(page.get_text() for page in pdf)
        assert TITLE not in text and CURRENT in text and "Prior " not in text
        assert all(page.get_text().strip() for page in pdf)
        tables = [table for page in pdf for table in page.find_tables().tables]
        assert [(table.row_count, table.col_count) for table in tables] == [(4, 2)]
        assert tables[0].extract()[2] == ["", ""]
        if scope is None:
            assert "Current synthetic training." in pdf[1].get_text()


@pytest.mark.parametrize("block", [
    ResolvedBlock("improvements", "This month", text="Current synthetic improvement."),
    ResolvedBlock("improvements", "This month", asset_hashes=("current.png",),
                  asset_captions=("Current synthetic inspection photo",)),
])
def test_current_improvement_content_keeps_its_panel(report, block):
    path, draft = report
    if block.asset_hashes:
        source = Document(path)
        heading = next(p._p for p in source.paragraphs if p.text == TITLE)
        photo = source.add_paragraph()
        photo.add_run().add_picture(BytesIO(chart()))
        heading.addnext(photo._p)
        source.save(path)
        inspected = inspect_docx(path)
        item = next(item for item in inspected.items if item.kind == "table")
        draft = replace(draft, blocks=tuple(
            replace(value, extra_tables=tuple(replace(table, reference=f"docx:{inspected.sha256}:{item.id}")
                                               for table in value.extra_tables))
            if value.key == "work_orders" else value for value in draft.blocks))
    draft = replace(draft, blocks=(*draft.blocks, block))
    result = Document(BytesIO(build_native_docx(draft, path, asset_loader=lambda _: chart())))
    assert TITLE in [p.text for p in result.paragraphs]
    if block.text:
        assert block.text in [p.text for p in result.paragraphs]
    else:
        assert len(result.inline_shapes) == 1
        assert block.asset_captions[0] in [p.text for p in result.paragraphs]


def panel():
    document = Document()
    before = document.add_paragraph("Prior narrative")._p
    heading = document.add_paragraph(TITLE)._p
    tail = document.add_paragraph("Prior improvement description")._p
    body = document.element.body
    original = list(body)
    sections = {i: "activity" for i in range(1, len(original) + 1)}
    draft = SimpleNamespace(blocks=())
    pages = SimpleNamespace(live=set(), support=set(), headings=set())
    return body, original, sections, draft, pages, before, heading, tail


@pytest.mark.parametrize("tag", ["w:bookmarkStart", "w:fldChar", "w:drawing", "w:footnoteReference",
                                "w:sdt", "w:ins", "w:pict"])
@pytest.mark.parametrize("location", ["heading", "tail"])
def test_uninspected_source_markup_cannot_become_an_empty_panel_after_redaction(tag, location):
    body, original, sections, draft, _, _, heading, tail = panel()
    target = heading if location == "heading" else tail
    target.append(OxmlElement(tag))
    assert empty_activity_panels(draft, original, sections, set(), set()) == ()


@pytest.mark.parametrize("change", ["picture", "caption", "live", "text", "section", "unmapped"])
def test_final_content_or_changed_ownership_keeps_panel(change):
    body, original, sections, draft, pages, before, heading, tail = panel()
    plans = empty_activity_panels(draft, original, sections, set(), set())
    assert len(plans) == 1
    for node in (before, tail):
        for text in node.iter(qn("w:t")):
            text.text = ""
    if change in ("picture", "caption", "unmapped"):
        inserted = OxmlElement("w:p")
        if change == "picture":
            inserted.append(OxmlElement("w:drawing"))
        elif change == "caption":
            run = OxmlElement("w:r"); text = OxmlElement("w:t")
            text.text = "Current caption"; run.append(text); inserted.append(run)
        heading.addnext(inserted)
    elif change == "live":
        pages.live.add(tail)
    elif change == "text":
        next(tail.iter(qn("w:t"))).text = "Current note"
    else:
        tail.get_or_add_pPr().append(deepcopy(body[-1]))
    plans[0].remove_if_empty(body, pages, {before, tail})
    assert heading.getparent() is body


@pytest.mark.parametrize("case", ["other-section", "retained-heading", "retained-tail", "native-heading"])
def test_unproved_panel_ownership_is_left_unchanged(case):
    _, original, sections, draft, _, _, _, _ = panel()
    unchanged, headings = set(), set()
    if case == "other-section":
        sections = {i: "training" for i in sections}
    elif case == "retained-heading":
        unchanged.add(2)
    elif case == "retained-tail":
        unchanged.add(3)
    else:
        headings.add(2)
    assert empty_activity_panels(draft, original, sections, unchanged, headings) == ()


def test_authored_blank_spacing_before_panel_is_not_deleted():
    body, original, sections, draft, pages, before, heading, tail = panel()
    blank = OxmlElement("w:p")
    heading.addprevious(blank)
    original = list(body)
    sections = {i: "activity" for i in range(1, len(original) + 1)}
    plan, = empty_activity_panels(draft, original, sections, set(), set())
    for node in (before, tail):
        next(node.iter(qn("w:t"))).text = ""
    plan.remove_if_empty(body, pages, {before, tail})
    assert heading.getparent() is None
    assert blank.getparent() is body and before.getparent() is body
