"""Preview downsampling retains color interpretation and existing privacy gates.

Refs #491 and #304. All images, ICC profiles and Word packages are synthetic.
"""

from io import BytesIO
import struct
from zipfile import ZipFile
import zlib

from docx import Document
from docx.shared import Inches
from PIL import Image, ImageCms, ImageDraw, ImageStat
import pytest

from app import monthly_report_preview_images as preview
from app.monthly_report_media_privacy import sanitize_package_media


AUTHOR = "Synthetic author metadata canary"


@pytest.fixture(scope="module")
def icc_profile():
    return ImageCms.ImageCmsProfile(ImageCms.createProfile("sRGB")).tobytes()


def _image(mode, size, *, profile=None, fmt="JPEG"):
    image = Image.new("RGB", size, (180, 90, 40)).convert(mode)
    if mode == "RGBA":
        image.paste((200, 10, 5, 0), (0, 0, *size))
        ImageDraw.Draw(image).rectangle(
            (size[0] // 4, size[1] // 4, 3 * size[0] // 4, 3 * size[1] // 4),
            fill=(10, 140, 200, 255))
    exif = Image.Exif()
    exif[274] = 6  # Orientation stays metadata; do not rotate the saved pixels.
    exif[315] = AUTHOR
    output = BytesIO()
    image.save(output, fmt, icc_profile=profile, exif=exif.tobytes(), dpi=(144, 144))
    return output.getvalue()


def _package(image):
    document = Document()
    document.add_paragraph("Synthetic preview color fixture")
    document.add_picture(BytesIO(image), width=Inches(4))
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _parts(package):
    with ZipFile(BytesIO(package)) as source:
        return {info.filename: source.read(info) for info in source.infolist()}


def _preview_package(raw, tmp_path):
    original = tmp_path / "original.docx"
    original.write_bytes(raw)
    # The export privacy pass intentionally retains color profiles while
    # stripping author/device metadata before this preview-only derivative.
    sanitized = sanitize_package_media(raw)
    target = tmp_path / "preview.docx"
    preview.write_preview_docx(sanitized, target)
    assert original.read_bytes() == raw
    before, after = _parts(sanitized), _parts(target.read_bytes())
    media, = (name for name in before if name.startswith("word/media/"))
    assert before.keys() == after.keys()
    assert {name: data for name, data in before.items() if name != media} == {
        name: data for name, data in after.items() if name != media}
    assert AUTHOR.encode() not in before[media] and AUTHOR.encode() not in after[media]
    # Reopen the actual Word package and retain its display geometry.
    source_shape, = Document(BytesIO(sanitized)).inline_shapes
    target_shape, = Document(target).inline_shapes
    assert (target_shape.width, target_shape.height) == (source_shape.width, source_shape.height)
    return before[media], after[media]


@pytest.mark.parametrize("mode,has_profile,size", [
    ("RGB", True, (2400, 1600)),
    ("RGB", True, (1700, 1700)),  # pixel budget alone triggers resizing
    ("RGB", True, (1900, 20)),   # edge budget alone triggers resizing
    ("RGB", False, (2400, 1600)),
    ("CMYK", False, (2400, 1600)),
])
def test_resized_jpeg_retains_existing_color_profile_orientation_and_density(
        mode, has_profile, size, icc_profile, tmp_path):
    profile = icc_profile if has_profile else None
    before, after = _preview_package(_package(_image(mode, size, profile=profile)), tmp_path)
    with Image.open(BytesIO(before)) as original, Image.open(BytesIO(after)) as reduced:
        assert original.info.get("icc_profile") == profile
        assert reduced.info.get("icc_profile") == profile
        assert reduced.format == original.format == "JPEG"
        assert reduced.mode == original.mode == mode
        assert reduced.getexif()[274] == original.getexif()[274] == 6
        assert reduced.info["dpi"] == original.info["dpi"] == (144, 144)
        assert reduced.size != original.size
        assert max(reduced.size) <= preview.MAX_EDGE
        assert reduced.width * reduced.height <= preview.MAX_PIXELS


def test_small_jpeg_keeps_exact_sanitized_bytes_and_profile(icc_profile, tmp_path):
    before, after = _preview_package(
        _package(_image("RGB", (160, 100), profile=icc_profile)), tmp_path)
    assert after == before
    with Image.open(BytesIO(after)) as image:
        assert image.info["icc_profile"] == icc_profile


def test_resized_png_keeps_profile_and_clears_only_invisible_color(icc_profile, tmp_path):
    before, after = _preview_package(
        _package(_image("RGBA", (2400, 1600), profile=icc_profile, fmt="PNG")), tmp_path)
    with Image.open(BytesIO(before)) as original, Image.open(BytesIO(after)) as reduced:
        assert original.getpixel((0, 0)) == (200, 10, 5, 0)
        assert reduced.info["icc_profile"] == original.info["icc_profile"] == icc_profile
        assert reduced.mode == "RGBA" and reduced.format == "PNG"
        assert reduced.getexif()[274] == original.getexif()[274] == 6
        assert reduced.info["dpi"] == original.info["dpi"]
        assert max(reduced.size) <= preview.MAX_EDGE
        assert reduced.width * reduced.height <= preview.MAX_PIXELS
        hidden = reduced.getchannel("A").point([255] + [0] * 255)
        assert hidden.getbbox() is not None
        assert all(high == 0 for _, high in ImageStat.Stat(reduced, hidden).extrema[:3])
        assert reduced.getpixel((reduced.width // 2, reduced.height // 2)) == (10, 140, 200, 255)


def test_existing_decode_limit_rejects_large_dimensions_before_loading_pixels():
    # A tiny PNG header exercises the refusal without allocating a 40M-pixel image.
    def chunk(kind, data):
        return (struct.pack(">I", len(data)) + kind + data
                + struct.pack(">I", zlib.crc32(kind + data) & 0xFFFFFFFF))

    raw = (b"\x89PNG\r\n\x1a\n"
           + chunk(b"IHDR", struct.pack(">IIBBBBB", 8001, 5000, 8, 2, 0, 0, 0))
           + chunk(b"IDAT", zlib.compress(b"")) + chunk(b"IEND", b""))
    with pytest.raises(ValueError, match="too large to prepare a preview safely"):
        preview._smaller_image(raw)
