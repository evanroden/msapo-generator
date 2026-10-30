"""Proposal table cells retain their roles and literal wording on carryover."""
from dataclasses import replace
from io import BytesIO

from docx import Document
import pytest

from app.monthly_report_followups import (
    PROPOSAL_STATUSES, confirmation, normalize_proposal, report_text, seed_followups,
)
from app.monthly_report_model import (
    ReportDraft, ReportFollowUp, ReportPeriod, ReportTable, ResolvedBlock,
    default_sections, synthetic_profiles,
)
from app.monthly_report_setup import new_month_draft


def table_draft(columns, rows):
    section = next(s for s in default_sections() if s.key == "proposals")
    block = ResolvedBlock("proposals", "This month", extra_tables=(ReportTable(columns, rows),))
    return ReportDraft(synthetic_profiles()[0], ReportPeriod(2026, 9),
                       "Synthetic Reviewer", (section,), (block,))


@pytest.mark.parametrize("scope,status,expected", [
    ("Replace actuator; Status: Approved", "Pending", "pending"),
    ("Replace actuator\nStatus: Declined", "Pending", "pending"),
    ("Replace actuator; Decision: Pending", "Pending", "pending"),
    ("Check controls\nRetain existing interlock", "Declined", "declined"),
])
def test_decision_comes_from_its_cell_not_embedded_scope(scope, status, expected):
    draft = table_draft(("Proposal ID", "Scope", "Status"), (("001", scope, status),))
    item, = seed_followups(draft)
    assert item.status == expected
    assert item.text == "Proposal ID: 001; Scope: " + scope
    assert normalize_proposal(item) == item
    assert report_text(item).endswith(item.text)
    assert draft.blocks[0].extra_tables[0].rows == (("001", scope, status),)


@pytest.mark.parametrize("columns,row", [
    (("Proposal ID", "Scope"), ("001", "Replace actuator; Status: Declined")),
    (("Proposal ID", "Scope", "Status"), ("001", "Replace actuator; Status: Approved", "")),
    (("Proposal ID", "Scope", "Status", "Decision"), ("001", "Replace actuator", "Pending", "Awaiting board")),
    (("Proposal ID", "Scope", "Status", "Decision"), ("001", "Replace actuator", "Pending", "")),
])
def test_missing_or_unresolved_decision_is_not_inferred_from_other_fields(columns, row):
    item, = seed_followups(table_draft(columns, (row,)))
    assert item.status == "not_confirmed"
    assert item.text == "; ".join(f"{name}: {value}" for name, value in zip(columns, row) if value.strip())
    assert normalize_proposal(item) == item


@pytest.mark.parametrize("status", PROPOSAL_STATUSES)
def test_explicit_saved_decision_never_reinterprets_authored_scope(status):
    item = ReportFollowUp("saved", "proposal", "Work scope; Status: Declined\nRetain wiring",
                          "2026-09", status=status)
    assert normalize_proposal(item) == item
    assert report_text(item).endswith(item.text)


@pytest.mark.parametrize("status,expected", [
    ("Pending", "pending"), ("Declined", "declined"), ("Rejected", "declined"),
    ("Approved", "approved"), ("On hold", "on_hold"), ("Not confirmed", "not_confirmed"),
])
def test_ordinary_decisions_keep_existing_row_identity_and_price_filter(status, expected):
    from app.monthly_report_followups import _digest
    columns = ("Proposal ID", "Scope", "Amount", "Status", "Unknown field")
    row = ("001", "Replace actuator", "$1200", status, "Retain spare")
    item, = seed_followups(table_draft(columns, (row,)))
    original = f"Proposal ID: 001; Scope: Replace actuator; Status: {status}; Unknown field: Retain spare"
    assert item.key == "follow-" + _digest(("proposal", original))[:20]
    assert item.status == expected and "$1200" not in item.text
    assert item.text == "Proposal ID: 001; Scope: Replace actuator; Unknown field: Retain spare"


def test_conflicting_decision_columns_remain_visible():
    item, = seed_followups(table_draft(("Scope", "Status", "Decision"),
                                      (("Replace actuator", "Pending", "Declined"),)))
    assert item.status == "not_confirmed"
    assert item.text == "Scope: Replace actuator; Status: Pending; Decision: Declined"


def test_legacy_untyped_proposal_still_migrates_explicit_decision():
    item = ReportFollowUp("legacy", "proposal", "Scope: Replace actuator; Status: Declined", "2026-09")
    migrated = normalize_proposal(item)
    assert migrated.status == "declined" and migrated.text == "Scope: Replace actuator"
    assert migrated.key == item.key and item.status == "ongoing"


def imported_proposals(tmp_path):
    from app.monthly_report_import import inspect_docx, imported_draft
    from app.monthly_report_start_import import (
        analyze_starting_report, automatic_section_plans, build_starting_import,
    )
    profile = synthetic_profiles()[0]
    doc = Document()
    doc.add_heading(profile.facilities[0].title, 0)
    doc.add_paragraph("Monthly Report · September 2026")
    doc.add_heading("11 Pending & Declined Proposals", 1)
    columns = ("Proposal ID", "Vendor", "Scope", "Status", "Submitted", "Review note", "Amount")
    pending = ("001", "Synthetic Controls", "Replace actuator; Status: Approved", "Pending",
               "2026-09-08", "Retain spare", "$1200")
    declined = ("001", "Synthetic Airflow", "Inspect damper\nDecision: Pending", "Declined",
                "2026-09-09", "Retain wiring", "$2500")
    rows = (pending, declined, pending)
    table = doc.add_table(rows=len(rows) + 1, cols=len(columns))
    for target, values in zip(table.rows, (columns, *rows)):
        for cell, value in zip(target.cells, values):
            cell.text = value
    path = tmp_path / "synthetic-proposals.docx"
    doc.save(path)
    original = path.read_bytes()
    inspection = inspect_docx(path)
    analysis = analyze_starting_report(inspection, profile, ReportPeriod(2026, 9), ())
    plans = automatic_section_plans(inspection, analysis)
    assert all(plan["approved"] for plan in plans.values())
    mapped, _ = build_starting_import(path, inspection, plans.values())
    draft = imported_draft(profile, ReportPeriod(2026, 9), "Synthetic Reviewer", mapped)
    return path, original, draft


def test_actual_import_reopen_manual_edit_reimport_and_two_rollovers(tmp_path, monkeypatch):
    from app import monthly_report_library as library
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path / "data"))
    path, original, prior = imported_proposals(tmp_path)
    block = next(b for b in prior.blocks if b.key == "proposals")
    table, = block.extra_tables
    assert table.rows[0] == ("001", "Synthetic Controls", "Replace actuator; Status: Approved",
                            "Pending", "2026-09-08", "Retain spare")
    assert table.reference.startswith("docx:")
    library.save_snapshot(prior, expected_revision=0)
    reopened = library.load_snapshot(prior.profile.contract, prior.profile.key, prior.period).draft
    current = new_month_draft(reopened, ReportPeriod(2026, 10))
    assert [item.status for item in current.follow_ups] == ["pending", "declined", "pending"]
    assert len({item.key for item in current.follow_ups}) == 3
    first, second, twin = current.follow_ups
    first = replace(first, text=first.text + "; Manual note: retain bracket", status="on_hold")
    current = replace(current, follow_ups=(first, second, twin))
    library.save_snapshot(current, expected_revision=0)
    current = library.load_snapshot(current.profile.contract, current.profile.key, current.period).draft
    # Re-present the same imported table beside the saved, edited carryover.
    repeated = replace(current, blocks=prior.blocks)
    following = new_month_draft(repeated, ReportPeriod(2026, 11))
    assert len(following.follow_ups) == 3
    assert tuple(item.key for item in following.follow_ups) == tuple(item.key for item in current.follow_ups)
    assert tuple((item.text, item.status) for item in following.follow_ups) == tuple(
        (item.text, item.status) for item in current.follow_ups)
    for item in following.follow_ups:
        assert "$1200" not in item.text and "$2500" not in item.text
        assert "Submitted: 2026-09-" in item.text and "Review note:" in item.text
    assert library.load_snapshot(prior.profile.contract, prior.profile.key, prior.period).draft == prior
    assert path.read_bytes() == original


def test_native_proposal_editor_retains_status_like_scope_across_confirmation(tmp_path, monkeypatch):
    from streamlit.testing.v1 import AppTest
    from test_monthly_report_followup_integrity import APP
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path / "data"))
    scope = "Replace actuator; Status: Approved\nRetain wiring"
    prior = table_draft(("Proposal ID", "Scope", "Status"), (("001", scope, "Pending"),))
    current = new_month_draft(prior, ReportPeriod(2026, 10))
    app = AppTest.from_string(APP, default_timeout=20)
    app.session_state["draft"] = current
    app.session_state["category"] = "proposal"
    app.run()
    assert not app.exception
    assert app.session_state["draft"].follow_ups[0].text.endswith(scope)
    next(button for button in app.button if button.label == "Confirm unchanged").click().run()
    assert not app.exception
    item, = app.session_state["draft"].follow_ups
    assert item.status == "pending" and item.text.endswith(scope)
    assert item.reviewed_fingerprint == confirmation(item, current.period)
    app.run()
    assert app.session_state["draft"].follow_ups == (item,)


def test_retained_scope_pricing_is_still_blocked():
    from app.monthly_report_docx import assemble_docx
    prior = table_draft(("Proposal ID", "Scope", "Status"),
                        (("001", "Replace actuator; Status: Approved; $1200", "Pending"),))
    current = new_month_draft(prior, ReportPeriod(2026, 10))
    current = replace(current, follow_ups=tuple(replace(item,
        reviewed_fingerprint=confirmation(item, current.period)) for item in current.follow_ups))
    with pytest.raises(ValueError, match="(?i)(pric|monetary|amount)"):
        assemble_docx(current, acknowledged_fingerprint=current.fingerprint)


@pytest.mark.parametrize("second", ["Declined", "", "Awaiting board"])
def test_disambiguated_duplicate_status_columns_require_review(second):
    from types import SimpleNamespace
    from app.monthly_report_sections import table_without_prices
    source = SimpleNamespace(table_header_row=0, rows=(("Scope", "Status", "Status"),
                                   ("Replace actuator", "Pending", second)))
    table, _ = table_without_prices(source)
    assert table.columns == ("Scope", "Status", "Status (3)")
    item, = seed_followups(table_draft(table.columns, table.rows))
    assert item.status == "not_confirmed"
    assert "Status: Pending" in item.text
    if second:
        assert "Status (3): " + second in item.text


@pytest.mark.parametrize("serialized", [False, True])
def test_real_legacy_import_item_defaults_to_first_header_and_retains_rows(serialized):
    import json
    from app.monthly_report_import import ImportItem
    from app.monthly_report_sections import table_without_prices
    rows = (("Scope", "Status", "Status"),
            ("First actuator", "Pending", "Declined"),
            ("Last actuator", "Pending", "Awaiting board"))
    legacy = dict(id="legacy-table", kind="table", part="word/document.xml",
                  position=1, word_section=0, rows=rows)
    # Old constructor/decoded-record fields omit the newly added optional
    # header offset. Exercise the actual dataclass, not a duck-typed stub.
    if serialized:
        legacy = json.loads(json.dumps(legacy))
    source = ImportItem(**legacy)
    assert source.table_header_row == 0
    assert tuple(tuple(row) for row in source.rows) == rows
    table, removed = table_without_prices(source)
    assert not removed
    assert table.columns == ("Scope", "Status", "Status (3)")
    assert table.rows == rows[1:]
    assert tuple(tuple(row) for row in source.rows) == rows


def test_short_row_with_missing_decision_stays_unconfirmed():
    item, = seed_followups(table_draft(("Scope", "Status", "Decision"),
                                      (("Replace actuator", "Pending"),)))
    assert item.status == "not_confirmed" and "Status: Pending" in item.text


def test_consistent_multiple_decision_columns_are_not_extra_proposals():
    item, = seed_followups(table_draft(("Scope", "Status", "Approval Decision"),
                                      (("Replace actuator", "Declined", "Rejected"),)))
    assert item.status == "declined" and item.text == "Scope: Replace actuator"
