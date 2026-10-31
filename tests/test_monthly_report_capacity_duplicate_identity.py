"""Repeated capacity identities must never become two approved standing rows.

All fixtures are invented; the real bounded CSV source importer retains the
original bytes. These checks cover first-fill, explicit shared saves, conflict
history, distinct plant/site controls, and actual generic Word table output.
"""
from dataclasses import replace
from io import BytesIO

from docx import Document
import pytest

from app import monthly_report_capacity as capacity
from app import monthly_report_directory as directory
from app import monthly_report_library as library
from app import monthly_report_sources as sources
from app.monthly_report_docx import _build_docx
from app.monthly_report_model import (
    Facility, ReportDraft, ReportPeriod, ReportProfile, default_sections,
)

HEAD = b"Facility,Plant,Service,Required capacity,Available capacity,Units\n"
DUPLICATE = (
    HEAD
    + b"NH,East,Steam,0,NR,lb/hr\n"
    + b"NH,East,Steam,0,NR,lb/hr\n"
)
GOOD = (
    HEAD
    + b"NH,East,Steam,0,NR,lb/hr\n"
    + b"SH,East,Steam,0,NR,lb/hr\n"
)


@pytest.fixture
def profile(tmp_path, monkeypatch):
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path))
    sites = (
        Facility("north", "North Hospital", ("NH",)),
        Facility("south", "South Hospital", ("SH",)),
    )
    profile = ReportProfile("Synthetic Capacity Duplicate Contract", "north",
                            "North Hospital", sites[:1])
    directory.save_directory(
        profile.contract,
        tuple(directory.DirectorySite(site.key, site.title, site.aliases)
              for site in sites),
        expected_revision=0, actor="Synthetic Reviewer", confirmed=True,
    )
    return profile


def inspect(profile, raw, name="synthetic-capacity.csv"):
    inspection = capacity.inspect_capacity(profile, name, raw)
    assert len(inspection.tables) == 1
    assert sources.source_bytes(profile, inspection.content.source) == raw
    return inspection


def save(profile, tables, *, revision=0, mode="merge"):
    return capacity.save_capacity(
        profile, tables, expected_revision=revision, actor="Synthetic Reviewer",
        reviewed_fingerprint=capacity.review_fingerprint(tables), mode=mode,
    )


def test_identical_source_rows_are_not_saved_twice_even_with_matching_values(profile):
    inspection = inspect(profile, DUPLICATE)
    table, = inspection.tables
    assert table.site_keys == ("north", "north")
    assert table.rows[0] == table.rows[1]
    assert capacity.duplicate_capacity_rows(inspection.tables) == (
        ((1, 1), (1, 2), "north"),)
    assert capacity.first_fill_tables(None, inspection.tables) == ()
    with pytest.raises(ValueError, match="same site and labels"):
        save(profile, inspection.tables)
    assert capacity.load_capacity(profile.contract) is None
    assert sources.source_bytes(profile, inspection.content.source) == DUPLICATE


def test_repeated_identical_rows_across_two_source_tables_need_review(profile):
    inspection = inspect(profile, HEAD + b"NH,East,Steam,0,NR,lb/hr\n")
    original, = inspection.tables
    other = replace(original, id="second-table", title=original.title)
    candidates = (original, other)
    assert capacity.duplicate_capacity_rows(candidates) == (
        ((1, 1), (2, 1), "north"),)
    assert capacity.first_fill_tables(None, candidates) == ()
    with pytest.raises(ValueError, match="table 1 row 1, table 2 row 1"):
        save(profile, candidates)
    assert capacity.load_capacity(profile.contract) is None


def test_same_site_identity_with_conflicting_numbers_remains_blocked(profile):
    raw = (
        HEAD + b"NH,East,Steam,100,200,lb/hr\n"
        + b"NH,East,Steam,100,250,lb/hr\n"
    )
    inspection = inspect(profile, raw)
    assert len(capacity.duplicate_capacity_rows(inspection.tables)) == 1
    assert capacity.first_fill_tables(None, inspection.tables) == ()
    with pytest.raises(ValueError, match="same site and labels"):
        save(profile, inspection.tables)
    assert capacity.load_capacity(profile.contract) is None
    assert sources.source_bytes(profile, inspection.content.source) == raw


def test_same_readings_for_distinct_sites_and_plants_are_not_duplicates(profile):
    raw = (
        HEAD + b"NH,East,Steam,0,NR,lb/hr\n"
        + b"SH,East,Steam,0,NR,lb/hr\n"
        + b"NH,West,Steam,0,NR,lb/hr\n"
        + b"SH,West,Steam,0,NR,lb/hr\n"
    )
    inspection = inspect(profile, raw)
    assert capacity.duplicate_capacity_rows(inspection.tables) == ()
    assert capacity.first_fill_tables(None, inspection.tables) == inspection.tables
    state = save(profile, inspection.tables)
    assert state.revision == 1
    assert len(state.tables) == 1 and len(state.tables[0].rows) == 4
    assert state.tables[0].site_keys == ("north", "south", "north", "south")
    assert state.tables[0].rows[0][-2:] == ("NR", "lb/hr")
    assert state.tables[0].rows[0][3] == "0"
    assert capacity.load_capacity(profile.contract) == state


def test_repeated_rows_with_one_valid_sibling_do_not_auto_publish_partial_source(profile):
    raw = (
        DUPLICATE + b"SH,West,Chilled water,800,1000,tons\n"
    )
    inspection = inspect(profile, raw)
    assert len(inspection.tables[0].rows) == 3
    assert capacity.first_fill_tables(None, inspection.tables) == ()
    with pytest.raises(ValueError, match="same site and labels"):
        save(profile, inspection.tables)
    assert capacity.load_capacity(profile.contract) is None
    # A corrected *new source*, rather than an invisible in-memory truncation,
    # is the only input that may be published.
    corrected_raw = (
        HEAD + b"NH,East,Steam,0,NR,lb/hr\n"
        + b"SH,West,Chilled water,800,1000,tons\n"
    )
    corrected = inspect(profile, corrected_raw, "corrected-capacity.csv")
    assert capacity.duplicate_capacity_rows(corrected.tables) == ()
    assert capacity.first_fill_tables(None, corrected.tables) == corrected.tables
    state = save(profile, corrected.tables)
    assert state.revision == 1
    assert len(state.tables[0].rows) == 2
    assert sources.source_bytes(profile, inspection.content.source) == raw
    assert sources.source_bytes(profile, corrected.content.source) == corrected_raw


def test_duplicate_import_cannot_change_other_sites_or_existing_history(profile):
    baseline = inspect(profile, HEAD + b"SH,East,Chilled water,800,1000,tons\n")
    initial = save(profile, baseline.tables)
    assert initial.revision == 1
    raw = (
        HEAD + b"NH,East,Steam,40000,50000,lb/hr\n"
        + b"NH,East,Steam,40000,50000,lb/hr\n"
        + b"SH,West,Steam,20000,24000,lb/hr\n"
    )
    duplicate = inspect(profile, raw)
    assert capacity.first_fill_tables(initial, duplicate.tables) == ()
    with pytest.raises(ValueError, match="same site and labels"):
        save(profile, duplicate.tables, revision=1)
    assert capacity.load_capacity(profile.contract) == initial
    assert capacity.load_capacity(profile.contract, 1) == initial
    assert initial.tables[0].rows == baseline.tables[0].rows
    assert sources.source_bytes(profile, duplicate.content.source) == raw
    corrected = inspect(profile, raw.replace(
        b"NH,East,Steam,40000,50000,lb/hr\n"
        b"NH,East,Steam,40000,50000,lb/hr\n",
        b"NH,East,Steam,40000,50000,lb/hr\n"),
        "corrected-steam.csv")
    first_fill = capacity.first_fill_tables(initial, corrected.tables)
    assert sum(len(table.rows) for table in first_fill) == 2
    saved = save(profile, first_fill, revision=1)
    assert saved.revision == 2
    assert sum(len(table.rows) for table in saved.tables) == 3
    assert capacity.load_capacity(profile.contract, 1) == initial
    with pytest.raises(library.RevisionConflict):
        save(profile, first_fill, revision=1)


def test_existing_site_capacity_still_allows_distinct_new_service_first_fill(profile):
    original = inspect(profile, HEAD + b"NH,East,Chilled water,1200,1600,tons\n")
    saved = save(profile, original.tables)
    incoming = inspect(profile, (
        HEAD + b"NH,East,Chilled water,1200,1600,tons\n"
        + b"NH,East,Steam,40000,50000,lb/hr\n"
    ), "two-services.csv")
    first_fill = capacity.first_fill_tables(saved, incoming.tables)
    assert len(first_fill) == 1
    assert first_fill[0].rows == (("NH", "East", "Steam", "40000", "50000", "lb/hr"),)
    updated = save(profile, first_fill, revision=1)
    assert sum(len(t.rows) for t in updated.tables) == 2
    assert capacity.load_capacity(profile.contract, 1) == saved
    assert sources.source_bytes(profile, incoming.content.source).startswith(HEAD)


def test_valid_corrected_capacity_prints_once_in_actual_generic_word_table(profile):
    inspection = inspect(profile, GOOD)
    state = save(profile, inspection.tables)
    block = capacity.capacity_block(profile, state)
    assert block is not None
    assert len(block.extra_tables) == 1
    assert block.extra_tables[0].rows == (("NH", "East", "Steam", "0", "NR", "lb/hr"),)
    section = next(s for s in default_sections() if s.key == "scorecards")
    draft = ReportDraft(profile, ReportPeriod(2026, 9), "Synthetic Editor",
                        (section,), (block,), synthetic=True)
    raw = _build_docx(draft)
    document = Document(BytesIO(raw))
    rows = [tuple(cell.text for cell in row.cells)
            for table in document.tables for row in table.rows]
    assert rows.count(("NH", "East", "Steam", "0", "NR", "lb/hr")) == 1
    assert not any("SH" in cell for row in rows for cell in row)
    assert sources.source_bytes(profile, inspection.content.source) == GOOD


def test_duplicate_guard_does_not_modify_a_previous_report_only_capacity_block(profile):
    source = inspect(profile, HEAD + b"NH,East,Steam,10,20,lb/hr\n")
    state = save(profile, source.tables)
    block = capacity.capacity_block(profile, state)
    before = block.fingerprint
    duplicated = inspect(profile, (
        HEAD + b"NH,East,Steam,10,20,lb/hr\n"
        + b"NH,East,Steam,10,20,lb/hr\n"
    ), "duplicate-retry.csv")
    with pytest.raises(ValueError, match="same site and labels"):
        save(profile, duplicated.tables, revision=1)
    assert block.fingerprint == before
    assert capacity.capacity_block(profile, capacity.load_capacity(profile.contract)) == block
    assert capacity.load_capacity(profile.contract).revision == 1


@pytest.mark.parametrize("value", ("10", "15"))
def test_keep_both_cannot_duplicate_a_saved_identity_even_from_one_new_row(profile, value):
    first = inspect(profile, HEAD + b"NH,East,Steam,10,20,lb/hr\n")
    old = save(profile, first.tables)
    incoming = inspect(profile, (
        HEAD + f"NH,East,Steam,{value},20,lb/hr\n".encode()
    ), "repeated-standing-identity.csv")
    assert capacity.duplicate_capacity_rows(incoming.tables) == ()
    with pytest.raises(ValueError, match="Keep both would duplicate"):
        save(profile, incoming.tables, revision=1, mode="keep_both")
    assert capacity.load_capacity(profile.contract) == old
    assert capacity.load_capacity(profile.contract, 1) == old
    assert sources.source_bytes(profile, incoming.content.source).endswith(b"lb/hr\n")


def test_keep_both_remains_available_for_genuinely_different_capacity_schema(profile):
    original = inspect(profile, HEAD + b"NH,East,Steam,10,20,lb/hr\n")
    first = save(profile, original.tables)
    alternate = inspect(profile, (
        b"Facility,Plant,Service,Capacity,Units\n"
        b"NH,East,Steam,30,lb/hr\n"
    ), "alternative-service-layout.csv")
    assert capacity.schema_conflicts(first, alternate.tables) == ("north",)
    with pytest.raises(ValueError, match="different columns"):
        save(profile, alternate.tables, revision=1)
    kept = save(profile, alternate.tables, revision=1, mode="keep_both")
    assert kept.revision == 2
    assert len(kept.tables) == 2
    assert [table.columns for table in kept.tables] == [
        original.tables[0].columns, alternate.tables[0].columns]
    assert capacity.load_capacity(profile.contract, 1) == first
    assert sources.source_bytes(profile, alternate.content.source).startswith(b"Facility,Plant")


def test_normal_merge_still_updates_one_matching_standing_row_without_duplicates(profile):
    initial = inspect(profile, HEAD + b"NH,East,Steam,10,20,lb/hr\n")
    old = save(profile, initial.tables)
    revised = inspect(profile, HEAD + b"NH,East,Steam,15,20,lb/hr\n",
                      "revised-steam.csv")
    current = save(profile, revised.tables, revision=1, mode="merge")
    assert current.revision == 2
    assert sum(len(table.rows) for table in current.tables) == 1
    assert current.tables[0].rows[0][3] == "15"
    assert capacity.load_capacity(profile.contract, 1) == old
    assert sources.source_bytes(profile, initial.content.source).endswith(b"10,20,lb/hr\n")



def test_site_aliases_resolve_to_one_identity_and_do_not_double_count(profile):
    raw = (
        HEAD
        + b"NH,East,Steam,100,200,lb/hr\n"
        + b"North Hospital,East,Steam,100,200,lb/hr\n"
    )
    inspection = inspect(profile, raw)
    assert inspection.tables[0].site_keys == ("north", "north")
    assert capacity.duplicate_capacity_rows(inspection.tables) == (
        ((1, 1), (1, 2), "north"),)
    assert capacity.first_fill_tables(None, iter(inspection.tables)) == ()
    with pytest.raises(ValueError, match="same site and labels"):
        save(profile, inspection.tables)
    assert capacity.load_capacity(profile.contract) is None
    assert sources.source_bytes(profile, inspection.content.source) == raw
