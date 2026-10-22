"""Synthetic slot labels bind only their enclosing native image occurrence."""
from copy import deepcopy
import hashlib
from io import BytesIO
from zipfile import ZipFile, ZIP_DEFLATED

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches
from lxml import etree as ET
from PIL import Image
import pytest

from app.monthly_report_import import inspect_docx, _logo_label_role
from app.monthly_report_native_logo_roles import client_logo_frame_pictures
from test_monthly_report_native_design_import import frame, WPS, A

MC = "{http://schemas.openxmlformats.org/markup-compatibility/2006}"


def raster(color="navy", format="PNG"):
    output = BytesIO()
    Image.new("RGB", (160, 80), color).save(output, format)
    return output.getvalue()


def _picture_in_frame(document, shape, raw):
    picture = document.add_picture(BytesIO(raw), width=Inches(2))
    drawing = picture._inline.getparent()
    run, old_paragraph = drawing.getparent(), drawing.getparent().getparent()
    shape.find(".//" + qn("w:txbxContent") + "/" + qn("w:p")).append(run)
    old_paragraph.getparent().remove(old_paragraph)
    return next(drawing.iter(A + "blip"))


def source(tmp_path, *, alternate=True, divergent=False, shared_outside=True, label="INSERT CLIENT LOGO", mutation=None):
    document = Document()
    paragraph = document.add_paragraph()
    for text in ("Prepared by: Synthetic Editor", "September 2026",
                 "Synthetic Regional Community Health Hospital\nOperations and Maintenance Monthly Review"):
        paragraph.add_run()._r.append(frame(text))
    slot = frame(label)
    paragraph.add_run()._r.append(slot)
    _picture_in_frame(document, slot, raster())
    if alternate:
        parent = slot.getparent()
        wrapper = ET.Element(MC + "AlternateContent", nsmap={"wps": WPS[1:-1]})
        parent.replace(slot, wrapper)
        ET.SubElement(wrapper, MC + "Choice", Requires="wps").append(slot)
        fallback = ET.SubElement(wrapper, MC + "Fallback")
        fallback.append(deepcopy(slot))
        if divergent:
            rid, _ = document.part.get_or_add_image(BytesIO(raster("green", "JPEG")))
            next(fallback.iter(A + "blip")).set(qn("r:embed"), rid)
    else:
        wrapper, fallback = None, None
    if mutation:
        mutation(document, slot, wrapper, fallback)
    if shared_outside:
        # Exact same media part is used independently; it must not inherit role.
        document.add_picture(BytesIO(raster()), width=Inches(2))
    document.add_heading("MONTHLY ACTIVITY SUMMARY", 1)
    document.add_paragraph("Synthetic completed inspection.")
    path = tmp_path / "source-frame.docx"
    document.save(path)
    return path


def image_items(path):
    return tuple(item for item in inspect_docx(path).items if item.kind == "image")


@pytest.mark.parametrize("alternate", [False, True])
@pytest.mark.parametrize("label", ["INSERT CLIENT LOGO", "  insert  client logo  "])
def test_exact_nearest_frame_labels_only_its_own_occurrence(tmp_path, alternate, label):
    path = source(tmp_path, alternate=alternate, label=label)
    original = path.read_bytes()
    items = image_items(path)
    assert len(items) == 2 and items[0].image_part == items[1].image_part
    assert items[0].image_role == items[0].suggested_slot == "client_logo"
    assert items[1].image_role == "" and items[1].suggested_slot == "cover_photo"
    assert inspect_docx(path).sha256 == hashlib.sha256(original).hexdigest()
    assert path.read_bytes() == original


def test_role_does_not_approve_or_clear_divergent_alternative_bytes(tmp_path):
    from app.monthly_report_native_layout import build_native_docx, NativeLayoutError
    from app.monthly_report_model import ReportDraft, ReportPeriod, synthetic_profiles, default_sections
    path = source(tmp_path, divergent=True)
    original = path.read_bytes()
    inspection = inspect_docx(path)
    assert image_items(path)[0].image_role == "client_logo"
    assert any(item.kind == "unsupported" and item.note.startswith("Alternate Word") for item in inspection.items)
    draft = ReportDraft(synthetic_profiles()[0], ReportPeriod(2026, 10), "Synthetic Editor", default_sections(), ())
    with pytest.raises(NativeLayoutError, match="hidden alternate"):
        build_native_docx(draft, path)
    assert path.read_bytes() == original


def mutate(kind):
    def change(document, slot, wrapper, fallback):
        if kind == "two_images":
            _picture_in_frame(document, slot, raster("red"))
        elif kind == "nested_frame":
            slot.find(".//" + qn("w:txbxContent") + "/" + qn("w:p")).append(frame(""))
        elif kind == "sibling_picture":
            drawing = next(slot.iter(A + "blip"))
            while drawing.tag != qn("w:drawing"):
                drawing = drawing.getparent()
            slot.getparent().append(drawing)
        elif kind == "unscaffolded_picture":
            slot.find(".//" + qn("w:txbxContent")).append(next(slot.iter(qn("pic:pic"))))
        elif kind == "bodypr_picture":
            picture = next(slot.iter(qn("pic:pic")))
            slot.find(".//" + WPS + "bodyPr").append(picture)
        elif kind == "wrapper_inside_frame":
            outer = frame("")
            parent = wrapper.getparent()
            parent.replace(wrapper, outer)
            outer.find(".//" + qn("w:txbxContent") + "/" + qn("w:p")).append(wrapper)
        elif kind == "structural_text":
            slot.find(".//" + WPS + "spPr").text = "hidden alternate text"
        elif kind == "extra_text":
            next(slot.iter(qn("w:t"))).text += " extra note"
        elif kind == "command":
            ET.SubElement(slot, qn("w:fldSimple"), {qn("w:instr"): "INCLUDETEXT hidden.docx"})
        elif kind == "instruction":
            ET.SubElement(slot, qn("w:instrText")).text = "DDEAUTO hidden"
        elif kind == "unknown":
            ET.SubElement(slot, "{urn:unknown}payload")
        elif kind == "hidden_text":
            ET.SubElement(next(slot.iter(qn("w:r"))), qn("w:vanish"))
        elif kind == "effect":
            ET.SubElement(ET.SubElement(slot.find(".//" + WPS + "spPr"), A + "effectLst"), A + "outerShdw")
        elif kind == "extra_layer":
            ET.SubElement(next(slot.iter(A + "blip")), A + "blip", {qn("r:embed"): next(slot.iter(A + "blip")).get(qn("r:embed"))})
        elif kind == "company_label":
            next(slot.iter(qn("wp:docPr"))).set("descr", "Company logo")
        elif kind == "external":
            picture = next(slot.iter(A + "blip"))
            picture.set(qn("r:link"), picture.attrib.pop(qn("r:embed")))
        elif kind == "missing":
            next(slot.iter(A + "blip")).set(qn("r:embed"), "missing")
        elif kind == "fallback_text":
            next(fallback.iter(qn("w:t"))).text = "Other image"
        elif kind == "fallback_extra_image":
            fallback.append(deepcopy(next(fallback.iter(A + "blip"))))
        elif kind == "fallback_command":
            ET.SubElement(fallback, qn("w:fldSimple"), {qn("w:instr"): "INCLUDETEXT hidden.docx"})
        elif kind == "nested_alternate":
            ET.SubElement(slot, MC + "AlternateContent")
    return change


@pytest.mark.parametrize("kind", ["two_images", "nested_frame", "sibling_picture", "bodypr_picture", "unscaffolded_picture", "wrapper_inside_frame", "structural_text", "extra_text", "command", "instruction",
    "unknown", "hidden_text", "effect", "extra_layer", "company_label", "external", "missing", "fallback_text",
    "fallback_extra_image", "fallback_command", "nested_alternate"])
def test_nonclosed_or_disagreeing_frames_do_not_supply_slot_evidence(tmp_path, kind):
    path = source(tmp_path, mutation=mutate(kind), shared_outside=False)
    document = Document(path)
    def identity(rid):
        relation = document.part.rels.get(rid)
        return hashlib.sha256(relation.target_part.blob).hexdigest() if relation and not relation.is_external else None
    assert not client_logo_frame_pictures(document.element, identity, _logo_label_role)
    assert all(item.image_role != "client_logo" for item in image_items(path))


def test_role_is_recomputed_for_changed_or_moved_source_occurrence(tmp_path):
    path = source(tmp_path)
    before = inspect_docx(path)
    assert image_items(path)[0].image_role == "client_logo"
    document = Document(path)
    for text in document.element.iter(qn("w:t")):
        if text.text == "INSERT CLIENT LOGO":
            text.text = ""
    document.save(path)
    after = inspect_docx(path)
    assert before.sha256 != after.sha256
    assert all(item.image_role != "client_logo" for item in after.items)
    assert [item.id for item in before.items] == [item.id for item in after.items]


def test_unlabelled_contents_image_cannot_inherit_cover_role(tmp_path):
    path = source(tmp_path, shared_outside=False)
    document = Document(path)
    document.add_section()
    slot = frame("")
    document.add_paragraph().add_run()._r.append(slot)
    _picture_in_frame(document, slot, raster())
    document.save(path)
    items = image_items(path)
    assert items[0].image_role == "client_logo"
    assert items[-1].image_role != "client_logo"


def test_saved_contract_client_logo_is_reused_without_new_source_approval(tmp_path, monkeypatch):
    from app import monthly_report_branding as branding, monthly_report_library as library
    from app.monthly_report_model import ReportDraft, ReportPeriod, synthetic_profiles, default_sections
    from app.monthly_report_start_import import automatic_section_plans, StartingReportAnalysis
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path / "private-test-data"))
    profile = synthetic_profiles()[0]
    correct, other = raster("purple"), raster("orange")
    correct_ref, other_ref = library.asset_reference(correct, "png"), library.asset_reference(other, "png")
    logos = (
        branding.BrandLogo("synthetic-client", "Synthetic contract logo", (profile.contract,), correct_ref,
            "https://example.com/client", "https://example.com/client.png", "2026-10-20"),
        branding.BrandLogo("other-client", "Unrelated logo", ("OTHER",), other_ref,
            "https://example.com/other", "https://example.com/other.png", "2026-10-20"),
    )
    branding.save_branding(logos, {correct_ref: correct, other_ref: other}, expected_revision=0,
                          actor="Synthetic Reviewer", confirmed=True)
    path = source(tmp_path)
    inspection = inspect_docx(path)
    item = next(item for item in inspection.items if item.image_role == "client_logo")
    plans = automatic_section_plans(inspection, StartingReportAnalysis((), "same", (), (), ()), preferred_logos=("client_logo",))
    assert plans["cover"]["automatic_logo_skips"]["client_logo"] == [item.id]
    assert item.id not in plans["cover"]["selected"]
    draft = ReportDraft(profile, ReportPeriod(2026, 10), "Synthetic Editor", default_sections(), ())
    assets = {}
    updated = branding.apply_defaults(draft, assets)
    selected = next(block for block in updated.blocks if block.key == "client_logo")
    assert selected.source == "Library" and selected.asset_hashes == (correct_ref,)
    assert assets[correct_ref] == correct and other_ref not in assets
    assert selected.references == (), "Standing asset review is not a source-occurrence approval"



def test_unsupported_media_cannot_be_hidden_by_preferred_logo_skip(tmp_path):
    from app.monthly_report_start_import import automatic_section_plans, StartingReportAnalysis
    path = source(tmp_path, shared_outside=False)
    with ZipFile(path) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    media = next(name for name in parts if name.startswith("word/media/"))
    renamed = media.rsplit(".", 1)[0] + ".unsupported"
    parts[renamed] = parts.pop(media)
    rels = "word/_rels/document.xml.rels"
    parts[rels] = parts[rels].replace(media.removeprefix("word/").encode(), renamed.removeprefix("word/").encode())
    with ZipFile(path, "w", ZIP_DEFLATED) as archive:
        for name, raw in parts.items():
            archive.writestr(name, raw)
    inspection = inspect_docx(path)
    unsupported = [item for item in inspection.items if item.kind == "unsupported"]
    assert unsupported and all(not item.image_role and not item.suggested_slot for item in unsupported)
    plans = automatic_section_plans(inspection, StartingReportAnalysis((), "same", (), (), ()), preferred_logos=("client_logo",))
    assert not plans["cover"]["approved"] and plans["cover"]["questions"]


@pytest.mark.parametrize("conflict", [False, True])
def test_closed_legacy_frame_agrees_without_transferring_asset_approval(tmp_path, conflict):
    def legacy(document, slot, wrapper, fallback):
        old = fallback[0]
        box = next(old.iter(qn("w:txbxContent")))
        fallback.remove(old)
        shape = ET.SubElement(ET.SubElement(fallback, qn("w:pict")), "{urn:schemas-microsoft-com:vml}shape", type="#_x0000_t202")
        ET.SubElement(shape, "{urn:schemas-microsoft-com:vml}textbox").append(box)
        if conflict:
            shape.set("alt", "Company logo")
    path = source(tmp_path, mutation=legacy, shared_outside=False)
    assert (image_items(path)[0].image_role == "client_logo") is not conflict


@pytest.mark.parametrize("alternate", [False, True])
def test_outer_frame_descriptor_conflict_cannot_supply_client_role(tmp_path, alternate):
    def conflict(document, slot, wrapper, fallback):
        ET.SubElement(slot.find(qn("wp:inline")), qn("wp:docPr"),
                      id="999", name="Outer frame", descr="Company logo")
    path = source(tmp_path, alternate=alternate, mutation=conflict, shared_outside=False)
    assert all(item.image_role != "client_logo" for item in image_items(path))


def test_header_frame_caption_preserves_existing_explicit_logo_classification(tmp_path):
    document = Document()
    document.add_paragraph("Synthetic report")
    slot = frame("Company logo")
    document.sections[0].header.paragraphs[0].add_run()._r.append(slot)
    # Resolve the picture through the header's own relationships.
    picture = document.sections[0].header.paragraphs[0].add_run().add_picture(BytesIO(raster()))
    run = picture._inline.getparent().getparent()
    slot.find(".//" + qn("w:txbxContent") + "/" + qn("w:p")).append(run)
    path = tmp_path / "header.docx"
    document.save(path)
    assert image_items(path)[0].image_role == "brand_logo"
