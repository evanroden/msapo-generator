"""Full synthetic DOCX/PDF generation with grounded WO counts and safe proposal mapping.

This is not the private RRH authentic design and is not a browser acceptance test.
It exercises the normal assembly/conversion path and inspects actual output.
"""
from io import BytesIO
from zipfile import ZipFile

from docx import Document
import fitz
import pytest

from app import monthly_report_sources as sources
from app.monthly_report_docx import generate_report
from app.monthly_report_model import (
    BlockSpec, ColumnSpec, ReportDraft, ReportPeriod, ResolvedBlock,
    SectionSpec, synthetic_profiles,
)
from app.monthly_report_structured_uploads import mapped_rows
from app.monthly_report_work_order_csv import work_order_evidence
from tests.conftest import requires_libreoffice


WO = (b"Work Order,Date,Status,Facility,Asset\n"
      b"WO-TEST-1001,2026-09-03,Completed,Demo North,PUMP-TEST-02\n"
      b"WO-TEST-1002,2026-09-20,Completed,Demo North,AHU-TEST-07\n"
      b"WO-TEST-1003,2026-09-24,Open,Demo North,VAV-TEST-03\n")
PRICES = ("$12,345.67", "$987.65", "$13,333.32", "$456.78", "$222.22")
SCHEMA = ("Project Number", "Facility", "Vendor", "Work Description", "Status", "Technical Notes")


def approved_synthetic_draft(monkeypatch, tmp_path):
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path))
    profile = synthetic_profiles()[0]
    period = ReportPeriod(2026, 9)
    work, _ = sources.ingest(profile, "synthetic-workorders.csv", WO)
    findings, = work_order_evidence(work, period, profile)
    assert (findings.total, findings.completed, findings.open) == (3, 2, 1)
    csv = ("Proposal Total,Project Number,Facility,Vendor,Work Description,Status,Technical Notes\n"
           ""$12,345.67",Q-TEST-731,Demo North,Synthetic Mechanical,Inspect supply valve,Pending,45 gpm\n"
           ""$987.65",Q-TEST-732,Demo North,Synthetic Mechanical,Check bearing,Open,1.20 psi\n"
           ""$13,333.32",Q-TEST-733,Demo North,Synthetic Electrical,Inspect switchboard,Pending,230 V\n"
           ""$456.78",Q-TEST-734,Demo North,Synthetic Electrical,Check controls,Open,240 V\n"
           ""$222.22",Q-TEST-735,Demo North,Synthetic Mechanical,Test flow,Pending,60 gpm\n").encode()
    priced, _ = sources.ingest(profile, "synthetic-proposals.csv", csv)
    rows = mapped_rows(priced, "proposals", SCHEMA)
    assert len(rows) == 5
    assert all(len(r) == len(SCHEMA) for r in rows)
    sections = (
        SectionSpec("activity", "2", "Monthly Activity Summary",
                    (BlockSpec("activity_summary", "rich_text"),)),
        SectionSpec("proposals", "10", "Pending and Declined Proposals",
                    (BlockSpec("proposals", "table", columns=tuple(
                        ColumnSpec(f"column-{i}", name) for i, name in enumerate(SCHEMA))),)),
    )
    blocks = (ResolvedBlock("activity_summary", "This month", text=findings.text),
              ResolvedBlock("proposals", "This month", rows=rows))
    return ReportDraft(profile, period, "Synthetic Editor", sections, blocks)


@requires_libreoffice
def test_real_release_docx_pdf_include_source_workorders_once_and_no_prices(monkeypatch, tmp_path):
    draft = approved_synthetic_draft(monkeypatch, tmp_path)
    package = generate_report(draft, acknowledged_fingerprint=draft.fingerprint)
    assert package.docx and package.pdf, package.pdf_error
    assert package.docx.startswith(b"PK") and package.pdf.startswith(b"%PDF-")
    word = Document(BytesIO(package.docx))
    paragraphs = "\n".join(p.text for p in word.paragraphs)
    assert paragraphs.count("3 total, 2 completed, 1 open") == 1
    assert "66.7%" in paragraphs
    assert len(word.tables) == 1
    table = word.tables[0]
    assert tuple(cell.text for cell in table.rows[0].cells) == SCHEMA
    assert len(table.rows) == 6
    assert table.cell(1, 0).text == "Q-TEST-731"
    assert table.cell(1, 4).text == "Pending"
    assert table.cell(2, 5).text == "1.20 psi"
    with ZipFile(BytesIO(package.docx)) as archive:
        xml = b"\n".join(archive.read(path) for path in archive.namelist()
                         if path.endswith(".xml")).decode("utf-8", errors="ignore")
        for token in PRICES:
            assert token not in xml
    with fitz.open(stream=package.pdf, filetype="pdf") as pdf:
        pdf_text = "\n".join(page.get_text() for page in pdf)
        assert pdf_text.count("3 total, 2 completed, 1 open") == 1
        assert "66.7%" in pdf_text and "Q-TEST-731" in pdf_text
        assert "Pending" in pdf_text
        for token in PRICES:
            assert token not in pdf_text
        assert len(pdf) >= 4


def test_unsafe_manual_proposal_price_cannot_be_generated_even_with_warning_ack(monkeypatch, tmp_path):
    draft = approved_synthetic_draft(monkeypatch, tmp_path)
    from dataclasses import replace
    from app.monthly_report_docx import assemble_docx
    bad = replace(draft, blocks=(draft.blocks[0],
        replace(draft.blocks[1], rows=(("Q-TEST-731", "Demo North", "Vendor",
                "Inspect valve $12,345.67", "Pending", "45 gpm"),))))
    with pytest.raises(ValueError):
        assemble_docx(bad, acknowledged_fingerprint=bad.fingerprint)

def test_malformed_csv_unquoted_thousands_separator_is_rejected_not_remapped(monkeypatch, tmp_path):
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path))
    profile = synthetic_profiles()[0]
    # This exact formerly accepted line shifted 345.67 into Project Number,
    # turning a price fragment into apparent client-facing prose.
    malformed = (b"Proposal Total,Project Number,Facility,Vendor,Work Description,Status,Technical Notes\n"
                 b"$12,345.67,Q-TEST-731,Demo North,Synthetic Mechanical,Inspect supply valve,Pending,45 gpm\n")
    with pytest.raises(ValueError, match="CSV rows have different column counts"):
        sources.ingest(profile, "synthetic-malformed.csv", malformed)
    accepted, notices = sources.ingest_batch(profile, (
        ("bad-quote.csv", malformed), ("valid-workorders.csv", WO)))
    assert len(accepted) == 1 and accepted[0].source.filename == "valid-workorders.csv"
    assert any("Quote any values containing commas" in note for note in notices)
    assert work_order_evidence(accepted[0], ReportPeriod(2026, 9), profile)[0].total == 3
