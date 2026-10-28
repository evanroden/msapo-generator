"""Invented native table flows retain supplied words and their original anchors."""
from dataclasses import replace
from io import BytesIO
from zipfile import ZipFile

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION_START
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.opc.constants import CONTENT_TYPE, RELATIONSHIP_TYPE
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.shared import Inches, Pt
from lxml import etree
import pytest

from app.monthly_report_import import _text
from app.monthly_report_model import ResolvedBlock
from app.monthly_report_native_layout import build_native_docx
from app.monthly_report_sections import table_without_prices
from test_monthly_report_native_art_review import setup
from test_monthly_report_native_art_lifecycle import save_sample
from test_monthly_report_source_frame_logo_role import raster, source
from test_monthly_report_static_siblings import modern, A, WP, WPS


CAPACITY_TITLE = "Thermal Services Requirements and Capacity"
CAPACITY_INTRO = "Synthetic delivery matrix for the invented North and South annexes."
CHILLED_LABEL = "Synthetic chilled-water service schedule"
STEAM_LABEL = "Synthetic steam service schedule"
CAPACITY_NOTES = "Synthetic capacity notes retain the source measurement basis.\nNo current operating result is established by this example."
UTILITY_CURRENT = "Current synthetic utility readings have not been supplied."
UTILITY_TITLE = "Client Utility Rate Analysis"
PROPOSAL_OVERVIEW = "Synthetic equipment proposal overview"
PROPOSAL_CURRENT = "Current synthetic proposal decisions have not been supplied."
REJECTED = "Rejected Quotes"
PENDING = "Pending Quotes"
NATIVE_NOTE = "Synthetic service life begins at the later accepted installation milestone."
TRAINING_TITLE = "Synthetic covered-equipment appendix"
TRAINING_SCOPE = "Invented annex equipment only; this example is not a current training record."
TRAINING_CURRENT = "Current synthetic training completion remains unconfirmed."
CHILLED_MARKER = "Synthetic chilled annex"
STEAM_MARKER = "Synthetic steam annex"
REJECTED_MARKER = "Synthetic declined pump"
PENDING_MARKER = "Synthetic pending drive"
NEW_MARKER = "Synthetic replacement motor"
TRAINING_MARKER = "Synthetic air separator"


def _section(document, *, landscape=False):
    section = document.add_section(WD_SECTION_START.NEW_PAGE)
    section.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
    section.page_width = Inches(11 if landscape else 8.5)
    section.page_height = Inches(8.5 if landscape else 11)
    section.top_margin = section.bottom_margin = Inches(.7)
    section.left_margin = section.right_margin = Inches(.7)
    return section


def _table(document, rows):
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    for row, values in zip(table.rows, rows):
        for cell, value in zip(row.cells, values):
            cell.text = value
    for cell in table.rows[0].cells:
        for run in cell.paragraphs[0].runs:
            run.bold = True
    return table


def _add_native_note(document, table):
    notes = OxmlElement("w:footnotes")
    for identity, tag in (("-1", "separator"), ("0", "continuationSeparator")):
        note = OxmlElement("w:footnote"); note.set(qn("w:id"), identity)
        note.set(qn("w:type"), tag)
        paragraph = OxmlElement("w:p"); run = OxmlElement("w:r")
        run.append(OxmlElement("w:" + tag)); paragraph.append(run); note.append(paragraph)
        notes.append(note)
    note = OxmlElement("w:footnote"); note.set(qn("w:id"), "1")
    paragraph = OxmlElement("w:p"); run = OxmlElement("w:r")
    run.append(OxmlElement("w:footnoteRef")); paragraph.append(run)
    run = OxmlElement("w:r"); text = OxmlElement("w:t")
    text.text = " " + NATIVE_NOTE; text.set(qn("xml:space"), "preserve")
    run.append(text); paragraph.append(run); note.append(paragraph); notes.append(note)
    part = Part(PackURI("/word/footnotes.xml"), CONTENT_TYPE.WML_FOOTNOTES,
                etree.tostring(notes), document.part.package)
    document.part.relate_to(part, RELATIONSHIP_TYPE.FOOTNOTES)
    run = table.cell(0, 2).paragraphs[0].add_run()
    run.font.superscript = True
    reference = OxmlElement("w:footnoteReference"); reference.set(qn("w:id"), "1")
    run._r.append(reference)


def _readable_cover(document):
    """Use proved synthetic anchors so audit pages have legible cover furniture."""
    paragraph = document.paragraphs[0]
    for index, drawing in enumerate(list(paragraph._p.findall(qn("w:r") + "/" + qn("w:drawing")))):
        lines = tuple((node.text or "",) for node in drawing.iter(qn("w:t")))
        replacement = modern("horizontal-textframe", paragraphs=lines)[0]
        anchor = replacement[0]
        anchor.find(WP + "positionH/" + WP + "posOffset").text = str(Inches(.7))
        anchor.find(WP + "positionV/" + WP + "posOffset").text = str(Inches(1 + .65 * index))
        for extent in (anchor.find(WP + "extent"), replacement.find(".//" + A + "xfrm/" + A + "ext")):
            extent.set("cx", str(Inches(6.7))); extent.set("cy", str(Inches(.9)))
        anchor.find(WP + "docPr").set("id", str(100 + index))
        drawing.getparent().replace(drawing, replacement)
    for wrapper in paragraph._p.iter("{http://schemas.openxmlformats.org/markup-compatibility/2006}AlternateContent"):
        for branch in wrapper:
            anchor = branch[0][0]
            anchor.tag = WP + "anchor"
            anchor.attrib.update(dict(simplePos="0", relativeHeight="2", behindDoc="0", locked="0", layoutInCell="1", allowOverlap="1"))
            simple = OxmlElement("wp:simplePos"); simple.set("x", "0"); simple.set("y", "0")
            anchor.insert(0, simple)
            for position, axis, value, relative in ((1, "H", 0, "column"), (2, "V", 2.7, "paragraph")):
                node = OxmlElement("wp:position" + axis); node.set("relativeFrom", relative)
                offset = OxmlElement("wp:posOffset"); offset.text = str(Inches(value)); node.append(offset)
                anchor.insert(position, node)
            extent = anchor.find(WP + "extent")
            extent.set("cx", str(Inches(2.5))); extent.set("cy", str(Inches(1.4)))
            anchor.insert(4, OxmlElement("wp:wrapNone"))
    document.paragraphs[1].paragraph_format.space_before = Pt(350)


def flow_case(tmp_path, monkeypatch, *, mutate_source=None):
    """A real synthetic source-art import; mutations precede fresh source refs."""
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path / "store"))
    document = Document(source(tmp_path, divergent=True, shared_outside=False))
    _readable_cover(document)
    normal = document.styles["Normal"]
    normal.font.name = "Arial"; normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(6)
    for section in document.sections:
        section.top_margin = section.bottom_margin = Inches(.7)
        section.left_margin = section.right_margin = Inches(.7)
    document.sections[0].header.paragraphs[0].add_run().add_picture(
        BytesIO(raster("teal")), width=Inches(.5))
    _section(document)
    document.add_heading("MONTHLY SCORECARDS", 1)
    _section(document, landscape=True)
    document.add_paragraph(UTILITY_TITLE)
    _section(document)
    document.add_paragraph(CAPACITY_TITLE).runs[0].bold = True
    document.add_paragraph(CAPACITY_INTRO)
    document.add_paragraph(CHILLED_LABEL)
    _table(document, (("Campus", "Current Peak Load (Tons)", "Annual Demand (Ton-Hrs)"),
                      (CHILLED_MARKER, "112", "3456")))
    document.add_paragraph(STEAM_LABEL)
    _table(document, (("Campus", "Current Peak Load (MBH)", "Annual Demand (MMBtu)"),
                      (STEAM_MARKER, "2870", "5930")))
    document.add_paragraph(CAPACITY_NOTES)
    _section(document)
    document.add_heading("PENDING & DECLINED PROPOSALS", 1)
    _section(document)
    document.add_paragraph(PROPOSAL_OVERVIEW)
    document.add_paragraph(REJECTED).runs[0].bold = True
    document.add_paragraph()
    _table(document, (("Equipment", "End of Useful Life", "Cost", "Scope"),
                      (REJECTED_MARKER, "2039", "$7300", "Synthetic earlier decision")))
    document.add_paragraph()
    document.add_paragraph(PENDING).runs[0].bold = True
    document.add_paragraph()
    _table(document, (("Equipment", "End of Useful Life", "Cost", "Scope"),
                      (PENDING_MARKER, "2041", "$8100", "Synthetic review pending")))
    _section(document, landscape=True)
    third = _table(document, (("Equipment", "Campus", "Estimated End of Useful Life"),
                             (NEW_MARKER, "Synthetic North annex", "2044")))
    _add_native_note(document, third)
    _section(document)
    document.add_heading("TRAINING SUMMARY", 1)
    document.add_paragraph("Synthetic historical training statement.")
    _section(document)
    document.add_paragraph(TRAINING_TITLE).runs[0].bold = True
    document.add_paragraph(TRAINING_SCOPE).runs[0].italic = True
    _table(document, (("Description", "Campus", "Units"),
                      (TRAINING_MARKER, "Synthetic South annex", "Count")))
    if mutate_source is not None:
        mutate_source(document)
    path = tmp_path / "invented-complete-native-flow.docx"
    document.save(path)
    sample = setup(path)
    draft, loader = save_sample(sample)
    view = sample[1]
    inspection = view.original_inspection

    def source_ref(item):
        return f"docx:{inspection.sha256}:{item.id}"

    def tables(section):
        return tuple(replace(table_without_prices(item)[0], reference=source_ref(item))
                     for item in inspection.items if item.kind == "table" and item.section == section)

    notes = tuple(source_ref(item) for item in inspection.items if item.note == "footnote:1")
    current = {
        "utility_analysis": ResolvedBlock("utility_analysis", "This month", text="\n\n".join((
            UTILITY_CURRENT, CAPACITY_TITLE, "Historical context: " + CAPACITY_INTRO,
            CHILLED_LABEL, STEAM_LABEL, "Historical context: " + CAPACITY_NOTES))),
        "thermal_capacity": ResolvedBlock("thermal_capacity", "This month", extra_tables=tables("scorecards")),
        "proposals": ResolvedBlock("proposals", "This month", text="\n\n".join((
            PROPOSAL_CURRENT, PROPOSAL_OVERVIEW, REJECTED, PENDING, NATIVE_NOTE)),
            extra_tables=tables("proposals"), references=notes),
        "training_summary": ResolvedBlock("training_summary", "This month", text="\n\n".join((
            TRAINING_CURRENT, TRAINING_TITLE, TRAINING_SCOPE)), extra_tables=tables("training")),
    }
    draft = replace(draft, blocks=tuple(block for block in draft.blocks if block.key not in current)
                    + tuple(current.values()))
    return path, draft, loader, view


def rendered(flow, *, section=None):
    path, draft, loader, view = flow
    return build_native_docx(draft, path, asset_loader=loader, source_view=view, section_key=section)


def before(table):
    node = table._tbl.getprevious()
    while node is not None:
        if node.tag != qn("w:p") or node.find(".//" + qn("w:sectPr")) is not None:
            return ""
        if _text(node):
            return _text(node)
        node = node.getprevious()
    return ""


def paragraph_section(document, text):
    """Return the physical section index for an exact top-level paragraph."""
    index = 0
    for node in document.element.body:
        if node.tag == qn("w:p") and _text(node) == text:
            return index
        if node.tag == qn("w:p") and node.find(".//" + qn("w:sectPr")) is not None:
            index += 1
    raise AssertionError(f"No source-positioned paragraph: {text}")


@pytest.mark.parametrize("section", [None, "scorecards", "proposals", "training"])
def test_complete_current_flow_stays_at_exact_native_tables(tmp_path, monkeypatch, section):
    flow = flow_case(tmp_path, monkeypatch)
    raw = rendered(flow, section=section)
    document = Document(BytesIO(raw))
    text = _text(document.element.body)
    if section in (None, "scorecards"):
        assert text.count(CAPACITY_TITLE) == text.count(CHILLED_LABEL) == text.count(STEAM_LABEL) == 1
        assert "Historical context: " + CAPACITY_INTRO in text
        assert "Historical context: " + CAPACITY_NOTES in text
        chilled = next(t for t in document.tables if CHILLED_MARKER in _text(t._tbl))
        steam = next(t for t in document.tables if STEAM_MARKER in _text(t._tbl))
        assert before(chilled) == CHILLED_LABEL
        assert before(steam) == STEAM_LABEL
        body = list(document.element.body)
        trailing = next(p._p for p in document.paragraphs
                        if p.text == "Historical context: " + CAPACITY_NOTES)
        assert body.index(steam._tbl) < body.index(trailing)
        assert UTILITY_CURRENT in text and UTILITY_TITLE in text
        utility_index = paragraph_section(document, UTILITY_TITLE)
        utility_section = document.sections[utility_index]
        assert utility_section.page_width > utility_section.page_height
        assert utility_index < paragraph_section(document, CAPACITY_TITLE)
    if section in (None, "proposals"):
        assert text.count(REJECTED) == text.count(PENDING) == 1
        assert before(next(t for t in document.tables if REJECTED_MARKER in _text(t._tbl))) == REJECTED
        assert before(next(t for t in document.tables if PENDING_MARKER in _text(t._tbl))) == PENDING
        third = next(t for t in document.tables if NEW_MARKER in _text(t._tbl))
        refs = list(third.cell(0, 2)._tc.iter(qn("w:footnoteReference")))
        assert len(refs) == 1 and refs[0].get(qn("w:id")) == "1"
        assert NATIVE_NOTE not in text
        with ZipFile(BytesIO(raw)) as archive:
            note_text = _text(parse_xml(archive.read("word/footnotes.xml")))
        assert note_text.count(NATIVE_NOTE) == 1
        assert "$7300" not in text and "$8100" not in text and "Cost" not in text
        assert PROPOSAL_CURRENT in text
    if section in (None, "training"):
        assert text.count(TRAINING_TITLE) == text.count(TRAINING_SCOPE) == 1
        table = next(t for t in document.tables if TRAINING_MARKER in _text(t._tbl))
        assert before(table) == TRAINING_SCOPE
        assert TRAINING_CURRENT in text
    if section is None:
        source = Document(flow[0])
        assert len(document.sections) == len(source.sections)
        assert len(document.tables) == len(source.tables) == 6
        assert paragraph_section(document, UTILITY_TITLE) == paragraph_section(source, UTILITY_TITLE)
        for label in ("MONTHLY SCORECARDS", "PENDING & DECLINED PROPOSALS", "TRAINING SUMMARY"):
            assert text.count(label) == 1
