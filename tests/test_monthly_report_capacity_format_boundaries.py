"""Release review: display formats must not misround, lose units, or block imports."""
from io import BytesIO
from zipfile import ZipFile

from openpyxl import Workbook
from openpyxl.utils.datetime import CALENDAR_MAC_1904
import pytest

from app import monthly_report_capacity as capacity, monthly_report_directory as directory
from app.monthly_report_model import Facility, ReportProfile


@pytest.mark.parametrize(("value", "code", "shown"), [
    ("2.5", '0 "tons"', "3 tons"),
    ("-2.5", "0", "-3"),
    ("-2.5", '"$"0', "-$3"),
    ("-0.001", "0.00", "0.00"),
    ("1.005", "0.00", "1.01"),
    ("0.745", "0%", "75%"),
    ("12", '0000 "tons"', "0012 tons"),
    ("1200", "[$$-409]#,##0", "$1,200"),
    ("1200", "[$USD-409]#,##0", "USD1,200"),
    ("12.5", '0 "x" 0', "12.5"),
    ("12.5", "0?0", "12.5"),
    ("0.75", "mm:ss.0", "0.75"),
    ("0.75", "0%%", "0.75"),
    ("12.5", r"0\;0", "12.5"),
    ("12.345", "0.#0", "12.345"),
])
def test_number_format_boundaries_keep_values_and_units(value, code, shown):
    cell = directory.DirectoryCell(1, 1, value, number_format=code)
    assert directory.formatted_value(cell) == shown


def test_out_of_range_1904_date_keeps_raw_value_without_blocking_capacity(tmp_path, monkeypatch):
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path))
    profile = ReportProfile("Synthetic Format Boundaries", "north", "Demo North",
                            (Facility("north", "Demo North"),))
    workbook = Workbook()
    workbook.epoch = CALENDAR_MAC_1904
    sheet = workbook.active
    sheet.append(["Facility", "Required capacity", "Last tested"])
    sheet.append(["Demo North", 1200, 2958465])
    sheet["C2"].number_format = "yyyy-mm-dd"
    output = BytesIO()
    workbook.save(output)
    table, = capacity.inspect_capacity(profile, "synthetic-format-boundaries.xlsx", output.getvalue()).tables
    assert table.rows == (("Demo North", "1200", "2958465"),)


def test_oversized_number_format_id_keeps_raw_values():
    workbook = Workbook()
    workbook.active.append([1200])
    workbook.active["A1"].number_format = '0 "tons"'
    output = BytesIO()
    workbook.save(output)
    modified = BytesIO()
    with ZipFile(BytesIO(output.getvalue())) as source, ZipFile(modified, "w") as target:
        for info in source.infolist():
            data = source.read(info)
            if info.filename == "xl/styles.xml":
                data = data.replace(b'<xf numFmtId="164"', b'<xf numFmtId="' + b"9" * 5000 + b'"')
            target.writestr(info, data)
    inspected = directory.inspect_workbook(modified.getvalue(), number_formats=True)
    assert inspected.sheets[0].values() == {(1, 1): "1200"}


@pytest.mark.parametrize("number_formats", [False, True])
def test_oversized_cell_style_index_keeps_raw_values(number_formats):
    workbook = Workbook()
    workbook.active.append([1200])
    workbook.active["A1"].number_format = '0 "tons"'
    output = BytesIO()
    workbook.save(output)
    modified = BytesIO()
    with ZipFile(BytesIO(output.getvalue())) as source, ZipFile(modified, "w") as target:
        for info in source.infolist():
            data = source.read(info)
            if info.filename == "xl/worksheets/sheet1.xml":
                data = data.replace(b's="1"', b's="' + b"9" * 5000 + b'"')
            target.writestr(info, data)
    inspected = directory.inspect_workbook(modified.getvalue(), number_formats=number_formats)
    assert inspected.sheets[0].values() == {(1, 1): "1200"}
