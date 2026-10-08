"""Synthetic-only regression cases for mixed-month setup and safe continuation."""

from dataclasses import asdict, replace
from datetime import date
from io import BytesIO

from docx import Document
from PIL import Image
import pytest

from app import monthly_report_library as library
from app.monthly_report_import import ImportItem, DocxInspection, inspect_docx, map_items, imported_draft
from app.monthly_report_model import ReportPeriod, ResolvedBlock, ReportSource, synthetic_profiles, profile_sections
from app.monthly_report_setup import (suggested_period, suggested_mappings, item_findings,
                                      initial_decision, design_profile, merge_blocks, merge_drafts, new_month_draft)


@pytest.mark.parametrize("today,expected", [(date(2026, 10, 1), (2026, 9)), (date(2026, 10, 10), (2026, 9)),
                                           (date(2026, 10, 11), (2026, 10)), (date(2026, 9, 25), (2026, 9)),
                                           (date(2027, 1, 7), (2026, 12))])
def test_month_suggestion(today, expected):
    assert suggested_period(today) == ReportPeriod(*expected)


def test_cover_never_dates_an_embedded_image():
    image = ImportItem("image", "image", "word/document.xml", 3, 1, "maintenance", "vendor_reports", image_part="word/media/page.png")
    period = ReportPeriod(2026, 9)
    assert initial_decision(image, True, period) == "Needs review"
    assert "not been read" in item_findings(image, period)[0]
    findings = item_findings(image, period, "Template July 2024. Service performed 09/28/2026. Prior visit August 2026.")
    assert any("July 2024" in f and "August 2026" in f for f in findings)
    assert not any("09/28/2026" in f for f in findings)
    assert initial_decision(replace(image, text="September 2026"), True, period) == "Needs review"


def test_ambiguous_org_charts_are_not_automatically_selected():
    items = tuple(ImportItem(str(n), "image", "word/document.xml", n, 1, "organization", "org_chart", image_part=f"word/media/{n}.png") for n in (1, 2))
    inspection = DocxInspection("a" * 64, items, ("July 2024",), (), 1)
    assert suggested_mappings(inspection) == ()


def test_unmatched_undated_work_cannot_default_to_reference_only():
    item = ImportItem("new-work", "text", "word/document.xml", 3, 1, text="The synthetic pump repair was completed.")
    assert initial_decision(item, False, ReportPeriod(2026, 9)) == "Needs review"
    heading = replace(item, text="2 Monthly Activity Summary", section="activity")
    assert initial_decision(heading, False, ReportPeriod(2026, 9)) == "Save for reference"


def make_docx(tmp_path):
    doc = Document()
    doc.add_heading("Synthetic report July 2024", 0)
    doc.add_heading("2 Monthly Activity Summary", 1)
    doc.add_paragraph("Synthetic manager already completed a September 2026 repair.")
    doc.add_heading("5 Maintenance Schedule", 1)
    image = BytesIO()
    Image.new("RGB", (200, 300), "white").save(image, "PNG")
    doc.add_picture(image)
    doc.add_heading("1 Organizational Chart", 1)
    image2 = BytesIO()
    Image.new("RGB", (300, 200), "navy").save(image2, "PNG")
    doc.add_picture(image2)
    path = tmp_path / "synthetic-mixed.docx"
    doc.save(path)
    return path


def prepared_setup(tmp_path):
    path = make_docx(tmp_path)
    inspection = inspect_docx(path)
    mappings = suggested_mappings(inspection)
    mapped = map_items(path, inspection, mappings)
    profile = design_profile(synthetic_profiles()[0], inspection, mappings, mapped.overrides)
    draft = imported_draft(profile, ReportPeriod(2026, 9), "Synthetic Editor", mapped)
    return path, inspection, mapped, profile, draft


def test_setup_persists_original_design_and_monthly_seed_but_not_monthly_defaults(tmp_path, monkeypatch):
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path / "data"))
    path, inspection, mapped, profile, draft = prepared_setup(tmp_path)
    state = library.save_report_setup(profile, draft, path, {"sha256": inspection.sha256, "intent": "partial"},
                                      assets=mapped.assets, expected_revision=0, actor="Synthetic Editor", confirmed=True)
    assert state.revision == 1
    assert state.block("vendor_reports") is None
    assert state.block("org_chart").asset_hashes
    assert library.load_imported_draft(profile.contract, profile.key) == draft
    assert library.imported_original(profile.contract, profile.key).read_bytes() == path.read_bytes()
    assert profile.section_order[:3] == ("activity", "maintenance", "organization")
    assert dict(profile.section_titles)["maintenance"] == "Maintenance Schedule"
    assert [s.key for s in profile_sections(profile)][:3] == list(profile.section_order[:3])
    assert library.profile_from_dict(asdict(profile)) == profile
    with pytest.raises(library.RevisionConflict):
        library.save_report_setup(profile, draft, path, {"sha256": inspection.sha256}, assets=(), expected_revision=0, actor="Synthetic Editor", confirmed=True)
    assert library.load_profile(profile.contract, profile.key).revision == 1


def test_failed_setup_head_does_not_expose_half_a_profile(tmp_path, monkeypatch):
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path / "data"))
    path, inspection, mapped, profile, draft = prepared_setup(tmp_path)
    real_write = library._atomic_write
    def fail_head(target, raw):
        if target.name == "manifest.json":
            raise OSError("synthetic disk failure")
        real_write(target, raw)
    monkeypatch.setattr(library, "_atomic_write", fail_head)
    with pytest.raises(OSError):
        library.save_report_setup(profile, draft, path, {"sha256": inspection.sha256}, assets=mapped.assets,
                                  expected_revision=0, actor="Synthetic Editor", confirmed=True)
    assert library.list_profiles(profile.contract) == ()


def test_partial_merge_preserves_text_and_pages_and_is_idempotent():
    old = ResolvedBlock("activity_summary", "This month", text="Synthetic colleague's finished work.", references=("old",))
    incoming = ResolvedBlock("activity_summary", "Last month", text="A newly completed repair.", references=("new",))
    combined = merge_blocks(old, incoming)
    assert combined.text == old.text + "\n\n" + incoming.text
    assert merge_blocks(combined, incoming) == combined
    old_pages = ResolvedBlock("vendor_reports", "This month", asset_hashes=("one",), asset_captions=("First",))
    new_pages = ResolvedBlock("vendor_reports", "This month", asset_hashes=("two",), asset_captions=("Second",))
    assert merge_blocks(old_pages, new_pages).asset_hashes == ("one", "two")
    assert merge_blocks(old_pages, new_pages).asset_captions == ("First", "Second")


def test_only_explicit_new_month_clears_period_specific_content(tmp_path):
    _, _, _, _, draft = prepared_setup(tmp_path)
    merged = merge_drafts(draft, draft)
    assert next(b for b in merged.blocks if b.key == "vendor_reports").asset_hashes
    following = new_month_draft(draft, ReportPeriod(2026, 10))
    assert not next(b for b in following.blocks if b.key == "vendor_reports").asset_hashes
    assert next(b for b in following.blocks if b.key == "org_chart").asset_hashes
    assert next(b for b in draft.blocks if b.key == "vendor_reports").asset_hashes  # original unchanged


def test_reviewed_actions_append_without_converting_recommendations_into_work():
    from app.monthly_report_guided import activity_from_sources
    original = ResolvedBlock("activity_summary", "This month", text="Synthetic colleague's existing paragraph.")
    source = ReportSource("a" * 64, "synthetic.pdf", "a" * 64, ".pdf", "Vendor service", vendor="Synthetic Vendor",
                          actions="Inspected pump.", recommendations="Replace it next year.", page_texts=("Inspected pump.",), selected_pages=(1,))
    combined = activity_from_sources((source,), original)
    assert original.text in combined.text and "Synthetic Vendor: Inspected pump." in combined.text
    assert "next year" not in combined.text
    assert len(combined.references) == 1
    assert activity_from_sources((source,), combined) == combined
    contact_source = replace(source, vendor="Synthetic Vendor contact@example.invalid", actions="Called 202-555-0100; inspected pump.")
    safe = activity_from_sources((contact_source,), original)
    assert "contact@example.invalid" not in safe.text and "202-555-0100" not in safe.text


def test_browser_preferences_do_not_cross_devices_or_preparer_profiles(tmp_path, monkeypatch):
    from app import memory
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path))
    memory.record_report_preferences("synthetic-browser-a", "Synthetic Contract", "north")
    memory.record_report_preparer("synthetic-browser-a", "Synthetic Contract", "north", "Synthetic Editor")
    assert memory.remembered_report_preferences("synthetic-browser-a") == ("Synthetic Contract", "north")
    assert memory.remembered_report_preferences("synthetic-browser-b") == ("", "")
    assert memory.remembered_report_preparer("synthetic-browser-a", "Synthetic Contract", "south") == ""
    assert memory.remembered_report_preferences("") == ("", "")


def test_image_review_is_bounded_cached_and_prepared_on_caller(tmp_path, monkeypatch):
    from app import monthly_report_image_review as review
    from app.receipt_jobs import start_receipt
    from threading import get_ident
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path / "data"))
    source = make_docx(tmp_path)
    item = next(i for i in inspect_docx(source).items if i.kind == "image")
    caller = get_ident()
    prepared_on, read_on = [], []
    original = review.prepare_image
    def prepare_image(*args):
        prepared_on.append(get_ident())
        return original(*args)
    monkeypatch.setattr(review, "prepare_image", prepare_image)
    def network(content):
        read_on.append(get_ident())
        assert content[-1]["text"].find("untrusted") >= 0
        return "Template July 2024; service date September 2026 [unclear]"
    future = start_receipt(lambda: review.prepare_review("synthetic-scope", source, item), lambda value: network(value[2]))
    result = future.result(timeout=10)
    assert prepared_on == [caller] and read_on[0] != caller
    digest, _, _ = review.prepare_review("synthetic-scope", source, item)
    review.complete_review("synthetic-scope", digest, result)
    assert review.remaining_reviews("synthetic-scope") == 19
    assert review.prepare_review("synthetic-scope", source, item) == (digest, result, None)
    # Exhaustion is persistent across sessions and does not discard the source.
    directory = review.review_directory("synthetic-scope")
    library._atomic_write(directory / "budget.json", library._json({"schema": 1, "images": [str(n) for n in range(20)]}))
    monkeypatch.setattr(review, "prepare_image", lambda *a: ("b" * 64, []))
    with pytest.raises(ValueError, match="20-image"):
        review.prepare_review("synthetic-scope", source, item)
    assert source.exists()


def test_same_month_single_image_conflict_cannot_hide_one_version(tmp_path):
    _, _, _, _, draft = prepared_setup(tmp_path)
    chart = next(b for b in draft.blocks if b.key == "org_chart")
    incoming = replace(draft, blocks=(replace(chart, asset_hashes=("different-image",)),))
    with pytest.raises(ValueError, match="single-image"):
        merge_drafts(draft, incoming)


def test_import_title_does_not_include_floating_footer_or_decorative_number(tmp_path):
    from app.monthly_report_import import inspect_docx, ImportMapping
    from app.monthly_report_model import synthetic_profiles
    doc = Document()
    doc.add_paragraph('100 Example Road | example.invalid\n3MONTHLY SCORECARDS')
    doc.add_paragraph('Synthetic utility analysis.')
    doc.add_paragraph('MONTHLY ACTIVITY SUMMARY\n2')
    doc.add_paragraph('Synthetic completed inspection.')
    path = tmp_path / 'synthetic-floating-headings.docx'
    doc.save(path)
    inspection = inspect_docx(path)
    mappings = tuple(ImportMapping(i.id, i.suggested_slot) for i in inspection.items if i.suggested_slot)
    profile = design_profile(synthetic_profiles()[0], inspection, mappings)
    assert dict(profile.section_titles)['scorecards'] == 'Monthly Scorecards'
    assert dict(profile.section_titles)['activity'] == 'MONTHLY ACTIVITY SUMMARY'
    assert all('example.invalid' not in title and '\n' not in title for _, title in profile.section_titles)


def test_replacement_only_section_keeps_original_position(tmp_path):
    from app.monthly_report_import import ImportMapping, MappedImport
    doc = Document()
    doc.add_heading('1 Organizational Chart', 1)
    doc.add_heading('2 Monthly Activity Summary', 1)
    doc.add_paragraph('Synthetic completed inspection.')
    path = tmp_path / 'synthetic-replacement-only.docx'
    doc.save(path)
    inspection = inspect_docx(path)
    item = next(i for i in inspection.items if i.text == 'Synthetic completed inspection.')
    profile = design_profile(synthetic_profiles()[0], inspection, (ImportMapping(item.id, 'activity_summary'),))
    assert profile.section_order[:2] == ('organization', 'activity')
    assert 'organization' in profile.excluded_sections
    mapped = MappedImport((ResolvedBlock('org_chart','Replace once',asset_hashes=('synthetic.png',)),
                           ResolvedBlock('activity_summary','Last month',text=item.text)),(),(),(item.id,))
    draft = imported_draft(profile,ReportPeriod(2026,9),'Synthetic Editor',mapped)
    assert draft.sections[0].key == 'organization' and draft.sections[0].included
