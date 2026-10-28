"""Synthetic floating table contents must not become a lossy editable grid."""
from hashlib import sha256

from docx import Document
from docx.oxml import parse_xml
import pytest

from app import monthly_report_import as importer

W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
A = 'http://schemas.openxmlformats.org/drawingml/2006/main'
V = 'urn:schemas-microsoft-com:vml'


def fixture(tmp_path, drawing):
    doc = Document()
    doc.add_heading('Monthly Activity Summary', 1)
    table = doc.add_table(rows=2, cols=2)
    for cell, text in zip(table.rows[0].cells, ('Area', 'Status')):
        cell.text = text
    table.cell(1, 0).text = 'Pump A'
    table.cell(1, 1).text = 'Reviewed'
    if drawing:
        table.cell(1, 1).paragraphs[0]._p.append(parse_xml(drawing))
    doc.add_paragraph('Adjacent narrative remains editable.')
    path = tmp_path / 'synthetic-floating-table.docx'
    doc.save(path)
    return path


BOX = f'<w:txbxContent xmlns:w="{W}"><w:p><w:r><w:t>Floating note</w:t></w:r></w:p></w:txbxContent>'
DRAWINGS = [
    BOX,
    f'<v:shape xmlns:v="{V}"><v:textbox>{BOX}</v:textbox></v:shape>',
    f'<w:drawing xmlns:w="{W}" xmlns:a="{A}"><a:graphic><a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">{BOX}</a:graphicData></a:graphic></w:drawing>',
    f'<w:drawing xmlns:w="{W}" xmlns:a="{A}"><a:graphic><a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/chart"/></a:graphic></w:drawing>',
    f'<v:group xmlns:v="{V}"><v:shape/></v:group>',
]


@pytest.mark.parametrize('drawing', DRAWINGS)
def test_floating_table_requires_review_before_mapping(tmp_path, monkeypatch, drawing):
    monkeypatch.setenv('EPC_DATA_DIR', str(tmp_path / 'runtime'))
    path = fixture(tmp_path, drawing)
    original = path.read_bytes()
    result = importer.inspect_docx(path)
    assert not [i for i in result.items if i.kind == 'table']
    blocked = [i for i in result.items if i.kind == 'unsupported' and 'manual layout review' in i.note]
    assert len(blocked) == 1
    item = blocked[0]
    assert item.part == 'word/document.xml' and item.position == 2
    assert 'Pump A' in item.text and 'Reviewed' in item.text
    if 'Floating note' in drawing:
        assert item.text.count('Floating note') == 1
    assert 'manual layout review' in item.note
    assert not item.rows
    from app.monthly_report_sections import section_reviews
    assert any(item in section.items for section in section_reviews(result))
    assert result.sha256 == sha256(original).hexdigest()
    assert path.read_bytes() == original
    assert any(i.kind == 'text' and i.text == 'Adjacent narrative remains editable.' for i in result.items)
    with pytest.raises(importer.ImportError, match='Unsupported items'):
        importer.map_items(path, result, (importer.ImportMapping(item.id, 'service_calls'),))
    from app.monthly_report_sections import build_section_import
    plan = {'key': 'activity', 'target': 'activity', 'approved': True,
            'selected': [item.id], 'destinations': {item.id: 'service_calls'}}
    with pytest.raises(ValueError, match='unread drawings'):
        build_section_import(path, result, (plan,))
    with pytest.raises(ValueError, match='needs a replacement'):
        build_section_import(path, result, ({**plan, 'unsupported_reviewed': True},))


def test_real_table_remains_editable_and_source_referenced(tmp_path, monkeypatch):
    monkeypatch.setenv('EPC_DATA_DIR', str(tmp_path / 'runtime'))
    path = fixture(tmp_path, '')
    result = importer.inspect_docx(path)
    item = next(i for i in result.items if i.kind == 'table')
    assert item.rows == (('Area', 'Status'), ('Pump A', 'Reviewed'))
    mapped = importer.map_items(path, result, (importer.ImportMapping(item.id, 'service_calls'),))
    assert mapped.blocks[0].rows == (('Pump A', 'Reviewed'),)
    assert mapped.blocks[0].references == (f'docx:{result.sha256}:{item.id}',)


def test_deleted_textbox_does_not_block_current_table(tmp_path, monkeypatch):
    monkeypatch.setenv('EPC_DATA_DIR', str(tmp_path / 'runtime'))
    path = fixture(tmp_path, f'<w:del xmlns:w="{W}" w:id="1">{BOX}</w:del>')
    result = importer.inspect_docx(path)
    table = next(i for i in result.items if i.kind == 'table')
    assert table.rows == (('Area', 'Status'), ('Pump A', 'Reviewed'))
    assert not any('Floating note' in i.text for i in result.items)


def test_bare_textbox_keeps_source_occurrence_identity(tmp_path):
    path = fixture(tmp_path, '')
    before = importer.inspect_docx(path)
    path = fixture(tmp_path, BOX)
    after = importer.inspect_docx(path)
    assert [(i.id, i.part, i.position) for i in before.items] == [
        (i.id, i.part, i.position) for i in after.items]
