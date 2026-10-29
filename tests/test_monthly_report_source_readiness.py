"""Accepted data gets a truthful summary without approving source pages."""
from dataclasses import asdict, replace
from io import BytesIO

from docx import Document
from PIL import Image
import pytest

from app import monthly_report_library as library, monthly_report_sources as sources
from app.monthly_report_cmms import CMMSMapping, map_cmms
from app.monthly_report_content_policy import page_fingerprint
from app.monthly_report_model import BlockSpec, ReportDraft, ReportPeriod, ResolvedBlock, SectionSpec, synthetic_profiles
from app.monthly_report_source_scope import contribution_receipt
from app.monthly_report_source_status import source_manager_status
from app.monthly_report_work_order_csv import work_order_evidence


PERIOD = ReportPeriod(2026, 9)


@pytest.fixture(params=(".csv", ".xlsx"))
def totals(monkeypatch, tmp_path, request):
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path))
    profile = synthetic_profiles()[0]
    raw = b"Work Order,Date,Status,Facility\n001,2026-09-04,Completed,Demo North\n002,2026-09-05,Open,Demo North\n"
    if request.param == ".xlsx":
        import csv
        from openpyxl import Workbook
        book = Workbook()
        for row in csv.reader(raw.decode().splitlines()): book.active.append(row)
        output = BytesIO(); book.save(output); raw = output.getvalue()
    content, _ = sources.ingest(profile, "synthetic-workorders" + request.param, raw)
    report, = work_order_evidence(content, PERIOD, profile)
    initial = ResolvedBlock("activity_summary", "This month", text="Independent human wording.")
    receipt = contribution_receipt(report.reference, profile, PERIOD, initial.key, text=report.text,
                                   source=content.source, prior_block=initial)
    block = replace(initial, text=initial.text + "\n\n" + report.text, references=(report.reference, receipt))
    draft = ReportDraft(profile, PERIOD, "Synthetic editor",
        (SectionSpec("activity", "2", "Monthly Activity", (BlockSpec(block.key, "rich_text"),)),),
        (block,), sources=(content.source,), synthetic=True)
    return draft, content, raw


def status(draft, content):
    result, = source_manager_status(draft, (content,))
    return result


def test_accepted_totals_are_ready_after_reload_without_a_file_date_ritual(totals):
    draft, content, raw = totals
    assert content.source.service_date == "", "Rows carry their accepted dates; the file has no date label."
    value = status(draft, content)
    assert value.resolved and not value.needs_correction and value.reason == ""
    assert value.summary == "Totals accepted · Monthly Activity"
    assert value.pages == 0 and value.destinations == ("activity_summary",)
    restored = library.draft_from_dict(asdict(draft))
    assert status(restored, replace(content, source=restored.sources[0])) == value
    assert sources.source_bytes(draft.profile, content.source) == raw


def test_unknown_reading_notice_remains_actionable_beside_accepted_totals(totals):
    draft, content, _ = totals
    changed = replace(content.source, notices=(*content.source.notices, "An unreadable source element needs correction."))
    value = status(replace(draft, sources=(changed,)), replace(content, source=changed))
    assert value.needs_correction and not value.resolved and "reading notices" in value.reason
    assert "unreadable" in changed.notices[-1]


@pytest.mark.parametrize("change", ["text", "source", "period"])
def test_changed_accepted_task_reopens_only_affected_source(totals, change):
    draft, content, _ = totals
    other, _ = sources.ingest(draft.profile, "synthetic-reference.txt", b"An unused reference.")
    draft = replace(draft, sources=(*draft.sources, other.source))
    if change == "text":
        draft = replace(draft, blocks=(replace(draft.blocks[0], text=draft.blocks[0].text.replace("2 total", "20 total")),))
    elif change == "source":
        content = replace(content, source=replace(content.source, page_texts=("Changed extraction",)))
        draft = replace(draft, sources=(content.source, other.source))
    else:
        draft = replace(draft, period=ReportPeriod(2026, 10))
    affected, unused = source_manager_status(draft, (content, other))
    assert not affected.resolved and affected.needs_correction
    assert "changed totals" in affected.reason
    assert not unused.resolved and not unused.needs_correction and unused.summary == "Not included"


def test_native_text_inclusion_is_never_mislabeled_reviewed(totals):
    draft, _, _ = totals
    content, _ = sources.ingest(draft.profile, "synthetic-notes.txt",
                              b"Facility: Demo North\nService date: 2026-09-04\nCompleted a synthetic inspection.")
    block = replace(draft.blocks[0], text="Completed a synthetic inspection.",
                    references=(sources.source_reference(content.source, 1),))
    current = replace(draft, sources=(content.source,), blocks=(block,))
    value = status(current, content)
    assert not value.resolved and not value.needs_correction and value.reason == ""
    assert value.summary == "Included content · Monthly Activity"
    from app.monthly_report_ai import evidence_fingerprint, reviewed_block
    suggested = replace(block, ai_written=True,
                        ai_evidence_fingerprint=evidence_fingerprint(current.sources, block.references))
    reviewed = reviewed_block(suggested, current.sources)
    accepted = status(replace(current, blocks=(reviewed,)), content)
    assert accepted.resolved and accepted.summary == "Content reviewed · Monthly Activity"
    changed = status(replace(current, blocks=(replace(reviewed, text="Changed wording"),)), content)
    assert changed.needs_correction and "suggested wording" in changed.reason


@pytest.mark.parametrize("classification", ("Utility results", "Reference only", "CMMS export", "Proposal / quote"))
def test_mixed_docx_keeps_page_review_separate_from_accepted_tables(monkeypatch, tmp_path, classification):
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path))
    profile = synthetic_profiles()[0]
    document = Document()
    document.add_paragraph("Facility: Demo North\nService date: 2026-09-04\nUtility results")
    table = document.add_table(rows=2, cols=4)
    for row, values in zip(table.rows, (("ID", "Finished", "Facility", "Task"),
                                      ("001", "2026-09-04", "Demo North", "Synthetic inspection"))):
        for cell, value in zip(row.cells, values): cell.text = value
    image = BytesIO(); Image.new("RGB", (40, 30), "blue").save(image, "PNG"); image.seek(0)
    document.add_picture(image)
    original = BytesIO(); document.save(original)
    content, _ = sources.ingest(profile, "synthetic-utility-report.docx", original.getvalue())
    assert content.source.facility == "Demo North" and content.source.service_date == "2026-09-04"
    numbers = sources.image_numbers(content)
    assert numbers and len(content.source.page_texts) > len(numbers)
    source = replace(content.source, classification=classification,
                     selected_pages=tuple(range(1, len(content.source.page_texts) + 1)))
    content = replace(content, source=source)
    # DOCX intake retains the table as a native item rather than an XLSX
    # worksheet. Read that unchanged table for this mapping-model check; this
    # does not claim the DOCX uploader exposes a CMMS worksheet control.
    native_table = Document(BytesIO(original.getvalue())).tables[0]
    rows = tuple(tuple(cell.text for cell in row.cells) for row in native_table.rows)
    extracted = sources.SourceTable("Native source table", rows[0], rows[1:])
    mapped = map_cmms(extracted, CMMSMapping("Finished", "ID", facility="Facility", description="Task",
        finish_date_means_completed=True), profile, PERIOD, sources.source_reference(source, 2), scope_source=source)
    sections = (SectionSpec("maintenance", "5", "Onsite service calls", mapped.specs),
                SectionSpec("scorecards", "3", "Scorecards", (BlockSpec("utility_analysis", "image_page"),)))
    draft = ReportDraft(profile, PERIOD, "Synthetic editor", sections, mapped.blocks,
                        sources=(source,), synthetic=True)
    initial = status(draft, content)
    assert not initial.resolved
    assert "Tables accepted" in initial.summary
    if classification == "Utility results":
        assert initial.needs_correction
        assert initial.reason == "Review the selected page and its caption."
    reviewed_source = replace(source, client_page_reviews=tuple((n, page_fingerprint(source, n)) for n in numbers))
    reviewed = replace(content, source=reviewed_source)
    pages, = sources.prepare_pages(profile, (reviewed,), ((source.id, "utility_analysis"),))
    accepted = replace(draft, sources=(reviewed_source,), blocks=(*mapped.blocks, pages))
    complete = status(accepted, reviewed)
    if classification != "Utility results":
        # A current, accepted table and reviewed image bytes do not establish
        # a page use for a reference, spreadsheet or proposal classification.
        assert not complete.resolved
        assert complete.reason == "Review this file's use in the report."
        assert "Tables accepted" in complete.summary and complete.pages == len(numbers)
        assert sources.source_bytes(profile, reviewed_source) == original.getvalue()
        return
    assert complete.resolved and not complete.needs_correction
    assert complete.reason == ""
    assert "Tables accepted" in complete.summary
    assert complete.pages == len(numbers)
    changed = replace(reviewed_source, captions=((numbers[0], "New caption"),))
    stale = status(replace(accepted, sources=(changed,)), replace(reviewed, source=changed))
    assert not stale.resolved and stale.needs_correction
    assert "Tables accepted" in stale.summary
    assert sources.source_bytes(profile, reviewed_source) == original.getvalue()


def test_unsupported_native_docx_drawing_retains_an_actionable_review_notice(monkeypatch, tmp_path):
    from docx.oxml import parse_xml
    from app.monthly_report_import import inspect_docx
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path))
    profile = synthetic_profiles()[0]
    document = Document()
    document.add_paragraph("Facility: Demo North\nService date: 2026-09-04\nSynthetic native drawing")
    paragraph = document.add_paragraph()
    paragraph._p.append(parse_xml('''<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
        xmlns:v="urn:schemas-microsoft-com:vml"><w:pict><v:group><v:shape id="synthetic-shape">
        <v:textbox><w:txbxContent><w:p><w:r><w:t>Synthetic native artwork</w:t></w:r></w:p></w:txbxContent>
        </v:textbox></v:shape></v:group></w:pict></w:r>'''))
    original = BytesIO(); document.save(original)
    content, _ = sources.ingest(profile, "synthetic-native-drawing.docx", original.getvalue())
    inspection = inspect_docx(sources._path(profile, content.source.sha256, content.source.suffix))
    assert any(item.kind == "unsupported" and "Native Word drawing" in item.note for item in inspection.items)
    block = ResolvedBlock("activity_summary", "This month", text="Synthetic native drawing",
        references=(sources.source_reference(content.source, 1),))
    draft = ReportDraft(profile, PERIOD, "Synthetic editor",
        (SectionSpec("activity", "2", "Monthly Activity", (BlockSpec(block.key, "rich_text"),)),),
        (block,), sources=(content.source,), synthetic=True)
    value = status(draft, content)
    assert not value.resolved and value.needs_correction
    assert value.reason == "Resolve the source reading notices."
    assert any(notice.startswith("Item ") and "Native Word drawing" in notice for notice in content.source.notices)
    from app.monthly_report_source_readiness import unresolved_reading_notices
    legacy_notice = "DOCX selections are extracted items, not Word page numbers. Native vector objects remain in the original for review."
    assert unresolved_reading_notices(replace(content.source, notices=(legacy_notice,))) == (legacy_notice,)
    assert sources.source_bytes(profile, content.source) == original.getvalue()


def test_ready_summary_does_not_mutate_read_or_render_originals(totals, monkeypatch):
    from app import monthly_report_ai as ai
    draft, content, _ = totals
    before = asdict(draft), content
    def forbidden(*args, **kwargs): pytest.fail("Readiness must not perform business work.")
    for module, name in ((sources, "source_bytes"), (sources, "page_image"), (library, "_atomic_write"),
                         (ai, "cached"), (ai, "reserve_call")):
        monkeypatch.setattr(module, name, forbidden)
    assert status(draft, content).resolved
    assert (asdict(draft), content) == before


def test_accepted_totals_summary_is_quiet_and_reopens_in_the_one_manager(monkeypatch, tmp_path):
    from test_monthly_report_ui import monthly
    from test_monthly_report_workorder_upload_ui import _upload
    app = monthly(monkeypatch, tmp_path)
    next(widget for widget in app.text_input if widget.label == "Prepared by").set_value("Synthetic editor").run()
    _upload(monkeypatch, b"Work Order,Date,Status,Facility\n001,2026-09-04,Completed,Demo North\n")
    app.run()
    next(button for button in app.button if button.label == "Add totals to summary").click().run()
    assert not app.exception
    # The selected-file summary and its native form stay mounted inside the
    # closed workspace, so closing does not discard browser-only form input.
    assert any("_source_manager_panel{display:none" in item.value for item in app.markdown)
    managers = [button for button in app.button if button.label == "Manage source files"]
    assert len(managers) == 1
    managers[0].click().run()
    assert not app.exception
    assert sum("work_orders.csv · Totals accepted" in item.value for item in app.caption) == 1
    assert any("_source_manager_panel{display:block" in item.value for item in app.markdown)
    assert not any("Review the extracted text or spreadsheet content" in item.value for item in app.caption)
    app.run()
    assert sum("work_orders.csv · Totals accepted" in item.value for item in app.caption) == 1
    assert any("_source_manager_panel{display:block" in item.value for item in app.markdown)
