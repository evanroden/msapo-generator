"""Synthetic title/header/narrative tables retain every original source cell."""
from dataclasses import replace

import pytest
from streamlit.testing.v1 import AppTest

from app.monthly_report_model import ReportTable
from app.monthly_report_row_cards_ui import _onsite_projection
from test_monthly_report_row_cards_ui import SCRIPT, example, button, field

HEADER = ('WO #', 'Finish Date', 'Area', 'Task Code-Description', 'Tag #')
TITLE = ('October 2026 Service Calls', 'Column 2', 'Column 3', 'Column 4', 'Column 5')
NOTE = ('Work-order counts are pending CMMS availability.', '', '', '', '')
BLANK = ('', '', '', '', '')
CALL = ('00042', '2026-10-12', 'Synthetic plant', 'Inspect pump', 'P-0007')


def workspace(rows, columns=TITLE):
    spec, block = example('service_calls', extra=True)
    block = replace(block, extra_tables=(ReportTable(columns, rows, 'synthetic-table-reference'),))
    return AppTest.from_string(SCRIPT.replace('INITIAL', repr(block)).replace('SPEC', repr(spec))).run(), block


def test_title_header_and_pending_context_are_not_calls_and_open_is_lossless():
    app, original = workspace((HEADER, NOTE))
    assert not app.exception
    assert not [b for b in app.button if b.label in ('Edit onsite call', 'Remove onsite call')]
    assert TITLE[0] in [c.value for c in app.caption]
    assert NOTE[0] in [c.value for c in app.caption]
    assert app.session_state['result'] == original
    app.run()
    assert app.session_state['result'] == original


def test_empty_placeholder_preserved_and_real_first_cell_only_call_not_hidden():
    partial = ('Awaiting assigned work-order number', '', '', '', '')
    app, original = workspace((HEADER, BLANK, NOTE, partial))
    assert len([b for b in app.button if b.label == 'Edit onsite call']) == 1
    button(app, 'Edit onsite call').click().run()
    assert field(app, 'WO #').value == partial[0]
    field(app, 'Area').set_value('Synthetic annex').run()
    rows = app.session_state['result'].extra_tables[0].rows
    assert rows[:3] == original.extra_tables[0].rows[:3]
    assert rows[3] == (partial[0], '', 'Synthetic annex', '', '')


def test_mixed_context_edit_close_reopen_remove_add_roundtrip():
    app, original = workspace((HEADER, NOTE, CALL, BLANK))
    assert len([b for b in app.button if b.label == 'Edit onsite call']) == 1
    button(app, 'Edit onsite call').click().run()
    field(app, 'Task Code-Description', 'text_area').set_value('Verified synthetic repair')
    button(app, 'Close editing').click().run()
    changed = app.session_state['result'].extra_tables[0]
    assert changed.columns == TITLE and changed.reference == original.extra_tables[0].reference
    assert changed.rows[:2] == (HEADER, NOTE) and changed.rows[3] == BLANK
    button(app, 'Edit onsite call').click().run()
    assert field(app, 'Task Code-Description', 'text_area').value == 'Verified synthetic repair'
    button(app, 'Remove onsite call').click().run()
    assert app.session_state['result'].extra_tables[0].rows == (HEADER, NOTE, BLANK)
    button(app, 'Add onsite call').click().run()
    field(app, 'WO #').set_value('00043').run()
    button(app, 'Close editing').click().run()
    assert app.session_state['result'].extra_tables[0].rows == (HEADER, NOTE, BLANK, ('00043', '', '', '', ''))
    assert not app.exception


def test_unknown_layout_is_neutral_and_editable_without_relabeling_source():
    columns = ('Source label', 'Source note')
    rows = (('Header-like source value', 'Keep exact'), ('Partial ID', ''))
    app, original = workspace(rows, columns)
    assert not [b for b in app.button if 'onsite call' in b.label]
    assert len([b for b in app.button if b.label == 'Edit table row']) == 2
    assert app.session_state['result'] == original
    button(app, 'Edit table row', 1).click().run()
    field(app, 'Source note').set_value('Synthetic note').run()
    assert app.session_state['result'].extra_tables[0].rows == (rows[0], ('Partial ID', 'Synthetic note'))


def test_context_recognition_is_narrow_and_canonical_data_never_dropped():
    for partial in [('Work-order counts pending review', '', '', '', ''),
                    ('Work-order counts are pending CMMS availability.', '', 'Real area', '', ''),
                    ('00042', '', '', '', '')]:
        _, context, recognized = _onsite_projection(TITLE, (HEADER, partial))
        assert recognized and context == {0: TITLE[0]}
    for columns, rows in [(TITLE, (('WO', *HEADER[1:]), NOTE)),
                          (('Not a monthly title', *TITLE[1:]), (HEADER, NOTE))]:
        returned, context, recognized = _onsite_projection(columns, rows)
        assert not recognized and returned == columns and not context


def test_synthetic_guided_import_and_native_export_preserve_title_and_note(tmp_path, monkeypatch):
    from io import BytesIO
    from docx import Document
    from app.monthly_report_import import inspect_docx, imported_draft
    from app.monthly_report_sections import build_section_import
    from app.monthly_report_setup import design_profile
    from app.monthly_report_model import ReportPeriod, synthetic_profiles, profile_sections
    from app import monthly_report_library as library
    from app.monthly_report_docx import assemble_docx

    monkeypatch.setenv('EPC_DATA_DIR', str(tmp_path / 'data'))
    document = Document()
    document.add_heading('5 Maintenance Schedule / In-House Maintenance', 1)
    table = document.add_table(rows=3, cols=5)
    table.cell(0, 0).merge(table.cell(0, 4)).text = TITLE[0]
    for cell, value in zip(table.rows[1].cells, HEADER):
        cell.text = value
    table.cell(2, 0).merge(table.cell(2, 4)).text = NOTE[0]
    path = tmp_path / 'synthetic-title-and-context.docx'
    document.save(path)
    original = path.read_bytes()
    inspection = inspect_docx(path)
    item = next(item for item in inspection.items if item.kind == 'table')
    plan = {'key': item.section, 'target': item.section, 'approved': True, 'selected': [item.id]}
    mapped, mappings = build_section_import(path, inspection, (plan,))
    profile = design_profile(synthetic_profiles()[0], inspection, mappings)
    draft = imported_draft(profile, ReportPeriod(2026, 10), 'Synthetic Editor', mapped)
    library.save_report_setup(profile, draft, path, {'sha256': inspection.sha256}, assets=mapped.assets,
                              expected_revision=0, actor='Synthetic Editor', confirmed=True)
    block = next(block for block in draft.blocks if block.key == 'service_calls')
    spec = next(spec for section in profile_sections(profile) for spec in section.blocks if spec.key == 'service_calls')
    app = AppTest.from_string(SCRIPT.replace('INITIAL', repr(block)).replace('SPEC', repr(spec))).run()
    assert not app.exception and app.session_state['result'] == block
    output = Document(BytesIO(assemble_docx(draft)))
    assert output.tables[0]._tbl.xml == table._tbl.xml
    assert path.read_bytes() == original


def test_pending_call_edit_conflicts_with_changed_context_without_overwriting_either():
    app, original = workspace((HEADER, NOTE, CALL))
    button(app, 'Edit onsite call').click().run()
    app.session_state['commit'] = False
    field(app, 'Area').set_value('Pending synthetic area').run()
    pending = app.session_state['result']
    changed_note = ('Work-order counts are pending CMMS availability; verified context changed.', '', '', '', '')
    current = replace(original, extra_tables=(replace(original.extra_tables[0], rows=(HEADER, changed_note, CALL)),))
    app.session_state['block'] = current
    app.run()
    assert app.warning and app.session_state['result'] == current
    assert pending.extra_tables[0].rows == (HEADER, NOTE, (CALL[0], CALL[1], 'Pending synthetic area', CALL[3], CALL[4]))
    button(app, 'Keep current table').click().run()
    assert not app.exception and app.session_state['result'] == current


def test_delivered_pre_projection_header_edit_cannot_replay_into_new_editor():
    app, original = workspace((HEADER, NOTE))
    current_key = next(key for key in app.session_state.filtered_state
                       if '_onsite_context_v1_extra_' in key and isinstance(app.session_state[key], dict))
    old_key = current_key.replace('_onsite_context_v1', '')
    source_rows = original.extra_tables[0].rows
    # Exact old callback record, where the old editor wrongly treated header
    # row zero as a call. A later rerun must not acknowledge this as new work.
    bad_rows = (('Old delivered header edit', *HEADER[1:]), NOTE)
    app.session_state[old_key] = {'pending': {
        'baseline': source_rows, 'expected': source_rows, 'result': bad_rows,
        'tokens': ('old-header', 'old-note'), 'active': 'old-header',
        'draft_token': None, 'draft_index': len(source_rows)}}
    app.run()
    assert not app.exception and app.session_state['result'] == original
    assert not [b for b in app.button if b.label == 'Edit onsite call']
    assert app.session_state[old_key]['pending']['result'] == bad_rows


@pytest.mark.parametrize("extra", [False, True])
def test_standard_facility_prefixed_cmms_schema_keeps_all_six_call_fields(extra):
    columns = ('Facility', *HEADER)
    row = ('Synthetic facility', *CALL)
    if extra:
        app, original = workspace((row,), columns)
    else:
        from app.monthly_report_model import ColumnSpec
        spec, original = example('service_calls', (row,))
        spec = replace(spec, columns=tuple(ColumnSpec(str(n), title) for n, title in enumerate(columns)))
        app = AppTest.from_string(SCRIPT.replace('INITIAL', repr(original)).replace('SPEC', repr(spec))).run()
    assert not app.exception
    assert len([b for b in app.button if b.label == 'Edit onsite call']) == 1
    assert not [b for b in app.button if 'table row' in b.label]
    button(app, 'Edit onsite call').click().run()
    assert {f.label for kind in ('text_input', 'text_area') for f in getattr(app, kind)} == set(columns)
    assert field(app, 'Facility').value == row[0]
    assert field(app, 'Tag #').value == row[5]
    field(app, 'Area').set_value('Synthetic CMMS annex')
    button(app, 'Close editing').click().run()
    result = app.session_state['result']
    if extra:
        table = result.extra_tables[0]
        assert table.columns == columns and table.reference == original.extra_tables[0].reference
        rows = table.rows
    else:
        assert result.references == original.references
        rows = result.rows
    assert rows == ((row[0], row[1], row[2], 'Synthetic CMMS annex', row[4], row[5]),)


def test_existing_native_journey_onsite_schema_retains_original_alias_fields():
    columns = ('WO #', 'Date', 'Area', 'Task', 'Asset')
    app, original = workspace((CALL,), columns)
    assert not app.exception
    button(app, 'Edit onsite call').click().run()
    assert {f.label for kind in ('text_input', 'text_area') for f in getattr(app, kind)} == set(columns)
    field(app, 'Asset').set_value('SYNTHETIC-ASSET-02')
    button(app, 'Close editing').click().run()
    table = app.session_state['result'].extra_tables[0]
    assert table.columns == columns and table.reference == original.extra_tables[0].reference
    assert table.rows == ((*CALL[:4], 'SYNTHETIC-ASSET-02'),)
    assert not app.exception
