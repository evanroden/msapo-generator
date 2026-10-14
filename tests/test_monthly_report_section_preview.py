from dataclasses import replace
from io import BytesIO

from docx import Document
import fitz
from PIL import Image
import pytest
from streamlit.testing.v1 import AppTest

from app import monthly_report_preview as engine
from app import monthly_report_section_preview_ui as ui
from app.monthly_report_docx import assemble_docx
from app.monthly_report_model import (
    ReportPeriod, ResolvedBlock, default_sections, synthetic_draft, synthetic_profiles,
)
from tests.conftest import requires_libreoffice
from report_preview_test_support import immediate_worker, enable_view, payload, completed_entry


def draft():
    base = synthetic_draft(synthetic_profiles()[0], ReportPeriod(2026, 9))
    return replace(base, sections=default_sections(), prepared_by="", blocks=(
        ResolvedBlock("activity_summary", "This month", text="Inspected the synthetic pumps."),
        ResolvedBlock("training_summary", "This month", text="Completed synthetic safety training.", ai_written=True),
    ))


def fake_converter(monkeypatch, tmp_path):
    calls = []

    def convert(path):
        calls.append(Document(path))
        output = tmp_path / (path.stem + ".pdf")
        with fitz.open() as pdf:
            for number in range(2):
                page = pdf.new_page(width=612, height=792)
                page.insert_text((72, 72), "Live vector section page " + str(number + 1))
            pdf.save(output)
        return output

    monkeypatch.setattr(engine.pdf_converter, "convert_to_pdf", convert)
    return calls


def text(raw):
    return "\n".join(paragraph.text for paragraph in Document(BytesIO(raw)).paragraphs)


def test_section_docx_contains_only_selected_divider_and_content_with_original_number():
    raw = engine.section_preview_docx(draft(), "training")
    visible = text(raw)
    assert "Section 11" in visible
    assert "11. Training Summary" in visible
    assert "Completed synthetic safety training." in visible
    assert "Table of contents" not in visible
    assert "Prepared by:" not in visible
    assert "Inspected the synthetic pumps." not in visible
    assert "Monthly Activity Summary" not in visible
    document = Document(BytesIO(raw))
    assert len(document.sections) == 2
    assert document.sections[0].left_margin == 0
    assert document.sections[1].left_margin.inches == 0.75
    assert "ENFRA" in document.sections[1].header.paragraphs[0].text
    # Previewing incomplete work never supplies an approval to final generation.
    with pytest.raises(ValueError, match="prepared|Review|required"):
        assemble_docx(draft())


def test_cover_preview_is_only_cover_and_keeps_cover_geometry():
    raw = engine.section_preview_docx(draft(), "cover")
    visible = text(raw)
    assert draft().profile.title in visible
    assert "September 2026" in visible
    assert "Prepared by:" in visible
    assert "Table of contents" not in visible
    assert "Completed synthetic safety training." not in visible
    assert len(Document(BytesIO(raw)).sections) == 1
    assert Document(BytesIO(raw)).sections[0].left_margin == 0


def test_empty_and_unfinished_template_sections_can_be_previewed():
    report = draft()
    assert "Water Treatment Reports" in text(engine.section_preview_docx(report, "water"))
    unfinished = replace(report, blocks=(ResolvedBlock("training_summary", "This month", text="Insert image here"),))
    assert "Insert image here" in text(engine.section_preview_docx(unfinished, "training"))
    with pytest.raises(ValueError, match="template instructions"):
        assemble_docx(unfinished)


def test_preview_does_not_resolve_other_sections_or_cover_art():
    report = draft()
    image = BytesIO()
    Image.new("RGB", (1200, 1200), "#123456").save(image, "PNG")
    report = replace(report, blocks=(*report.blocks,
        ResolvedBlock("water_reports", "This month", asset_hashes=("water.png",)),
        ResolvedBlock("org_chart", "This month", asset_hashes=("unavailable-org.png",)),
        ResolvedBlock("cover_photo", "This month", asset_hashes=("unavailable-cover.png",)),
    ))
    loaded = []

    def loader(ref):
        loaded.append(ref)
        assert ref == "water.png"
        return image.getvalue()

    assert "Section 7" in text(engine.section_preview_docx(report, "water", loader))
    assert loaded == ["water.png"]


def test_visible_fingerprint_ignores_other_work_and_approvals_but_tracks_real_edits():
    report = draft()
    fingerprint = engine.section_preview_fingerprint(report, "training")
    changed = replace(report, blocks=tuple(
        replace(block, text="Unrelated new maintenance.") if block.key == "activity_summary"
        else replace(block, reviewed_fingerprint=block.fingerprint, ai_written=False)
        for block in report.blocks
    ))
    assert engine.section_preview_fingerprint(changed, "training") == fingerprint
    changed = replace(changed, blocks=tuple(
        replace(block, text="Updated training statement.") if block.key == "training_summary" else block
        for block in changed.blocks
    ))
    assert engine.section_preview_fingerprint(changed, "training") != fingerprint
    assert engine.section_preview_fingerprint(replace(report, address_line="New footer"), "training") != fingerprint
    assert engine.section_preview_fingerprint(replace(report, prepared_by="Synthetic Editor"), "cover") != engine.section_preview_fingerprint(report, "cover")


def test_vector_pdf_and_sharp_pages_keep_bounds_and_clean_temporary_outputs(monkeypatch, tmp_path):
    calls = fake_converter(monkeypatch, tmp_path)
    preview = engine.preview_section(draft(), "training")
    assert preview.pages == 2
    with fitz.open(stream=preview.pdf, filetype="pdf") as pdf:
        assert "Live vector section page" in pdf[0].get_text()
    image = Image.open(BytesIO(engine.preview_page(preview, 0)))
    assert image.size == (1275, 1650)
    assert "Table of contents" not in "\n".join(p.text for p in calls[0].paragraphs)
    assert not list(tmp_path.glob("section-preview-*.pdf"))
    monkeypatch.setattr(engine, "MAX_PREVIEW_PAGES", 1)
    with pytest.raises(ValueError, match="limited"):
        engine.preview_section(draft(), "training")
    assert not list(tmp_path.glob("section-preview-*.pdf"))


def preview_app(monkeypatch, tmp_path):
    immediate_worker(monkeypatch)
    calls = fake_converter(monkeypatch, tmp_path)
    app = AppTest.from_string('''
import streamlit as st
from dataclasses import replace
from app.monthly_report_model import ReportPeriod, ResolvedBlock, synthetic_draft, synthetic_profiles, default_sections
from app.monthly_report_section_preview_ui import render_section_preview
base = synthetic_draft(synthetic_profiles()[0], ReportPeriod(2026, 9))
section = st.selectbox("Section", ("activity", "training"))
value = st.text_area("Current text", "Synthetic current edit")
unrelated = st.text_input("Unrelated", "")
draft = replace(base, prepared_by="", sections=default_sections(), blocks=(
    ResolvedBlock("activity_summary", "This month", text=value),
    ResolvedBlock("training_summary", "This month", text="Synthetic training"),
    ResolvedBlock("equipment_issues", "This month", text=unrelated)))
st.session_state["preview"] = render_section_preview(draft, section, {}, "synthetic")
''', default_timeout=15).run()
    assert not app.exception
    assert not calls
    enable_view(app, 'synthetic', 'activity')
    return app, calls


def test_ui_updates_on_edits_is_scrollable_and_reuses_only_unchanged_section(monkeypatch, tmp_path):
    app, calls = preview_app(monkeypatch, tmp_path)
    assert len(calls) == 1
    assert not app.button and not app.get("download_button")
    frame = app.get("bidi_component")[0].proto
    assert 'overflow:auto' in frame.css_content
    assert payload(app)['ready'] and payload(app)['pages']
    assert '<a ' not in frame.html_content
    app.text_input[0].set_value("Unrelated edit").run()
    assert len(calls) == 1
    app.text_area[0].set_value("New selected section text").run()
    assert len(calls) == 2
    app.selectbox[0].set_value("training").run()
    assert len(calls) == 2  # An offscreen pane is not rendered speculatively.
    enable_view(app, 'synthetic', 'training')
    assert len(calls) == 3
    app.selectbox[0].set_value("activity").run()
    assert len(calls) == 3
    assert not app.exception


def test_failed_rerender_labels_previous_page_not_ready_and_retry_preserves_edit(monkeypatch, tmp_path):
    app, _ = preview_app(monkeypatch, tmp_path)

    def fail(*args, **kwargs):
        raise RuntimeError("Synthetic renderer unavailable")

    monkeypatch.setattr(ui, "preview_section", fail)
    app.text_area[0].set_value("Still here after rendering failure").run()
    assert not app.exception
    assert payload(app)['pages'], 'Last good preview should remain visually inspectable'
    assert not payload(app)['ready'], 'An earlier image cannot count as this revision ready'
    assert 'Showing older preview' in payload(app)['status']
    assert 'latest changes not shown' in payload(app)['status']
    assert any("Your edits are still here" in value.value for value in app.warning)
    app.run()
    assert app.button[0].label == "Retry preview"
    monkeypatch.setattr(ui, "preview_section", engine.preview_section)
    app.button[0].click().run()
    assert not app.exception and payload(app)["pages"]
    assert app.text_area[0].value == "Still here after rendering failure"



def test_ticket_evicted_between_availability_and_raster_read_recovers_without_blank_success(monkeypatch, tmp_path):
    # The process-wide cache can be reclaimed by another visitor after a
    # successful STORE.available check. The previous code still published
    # "Current · N pages" with zero images until another unrelated rerun.
    app, _ = preview_app(monkeypatch, tmp_path)
    state = app.session_state['synthetic_ordered_previews']
    old_ticket = state['cache']['activity']['ticket']
    original = ui.STORE.pages
    calls = [0]

    def evict_once(owner, ticket, indexes, *, dpi=120):
        if calls[0] == 0:
            calls[0] += 1
            ui.STORE.discard(owner, ticket)
            return None
        return original(owner, ticket, indexes, dpi=dpi)

    monkeypatch.setattr(ui.STORE, 'pages', evict_once)
    app.run()
    assert not app.exception
    after_miss = payload(app)
    assert not after_miss['ready'], 'An evicted preview is not current output'
    assert not after_miss['pages'], 'An evicted preview must not keep inaccessible page placeholders'
    assert not after_miss['status'].startswith('Current')
    assert state['cache'].get('activity', {}).get('ticket') != old_ticket
    assert after_miss['pending'] or app.button, 'Must automatically rebuild or offer relevant Retry'

    app.run()
    assert not app.exception
    assert payload(app)['ready'] and payload(app)['pages']
    assert payload(app)['status'].startswith('Current')
    assert calls[0] == 1


def test_unreadable_current_page_has_specific_retry_not_current_blank_page(monkeypatch, tmp_path):
    app, _ = preview_app(monkeypatch, tmp_path)
    original = ui.STORE.pages

    def fail_raster(*args, **kwargs):
        raise ValueError('Synthetic page raster exceeds budget')

    monkeypatch.setattr(ui.STORE, 'pages', fail_raster)
    app.run()
    assert not app.exception
    assert not payload(app)['ready'] and not payload(app)['pages']
    assert payload(app)['status'] == 'Preview unavailable — retry'
    assert any('Reduce large images or retry' in warning.value for warning in app.warning)
    assert [button.label for button in app.button] == ['Retry preview']

    monkeypatch.setattr(ui.STORE, 'pages', original)
    app.button[0].click().run()
    app.run()
    assert not app.exception
    assert payload(app)['ready'] and payload(app)['pages']


@requires_libreoffice
def test_real_section_render_excludes_cover_toc_and_other_sections_and_keeps_text():
    report = draft()
    preview = engine.preview_section(report, "training")
    with fitz.open(stream=preview.pdf, filetype="pdf") as pdf:
        content = "\n".join(page.get_text() for page in pdf)
        assert len(pdf) == 2
        assert "Section 11" in content
        assert "Completed synthetic safety training." in content
        assert "Table of contents" not in content
        assert "Prepared by:" not in content
        assert "Inspected the synthetic pumps." not in content
    cover = engine.preview_section(report, "cover")
    with fitz.open(stream=cover.pdf, filetype="pdf") as pdf:
        content = "\n".join(page.get_text() for page in pdf)
        assert "September 2026" in content
        assert "Training Summary" not in content


def test_ordered_queue_runs_one_job_and_discards_a_result_after_edit(monkeypatch):
    from concurrent.futures import Future
    from threading import BoundedSemaphore
    jobs = []
    class Worker:
        def submit(self, fn, *args):
            future = Future()
            jobs.append((future, args))
            return future
    monkeypatch.setattr(ui, '_PREVIEW_WORKER', Worker())
    slot = BoundedSemaphore(1)
    monkeypatch.setattr(ui, '_PREVIEW_SLOT', slot)
    state = {'wanted': {'activity': 'old', 'training': 'train'},
             'queue': {'activity': ('old', draft(), {}, 'Activity'),
                       'training': ('train', draft(), {}, 'Training')},
             'cache': {}, 'running': None}
    ui._advance_preview_queue(state)
    ui._advance_preview_queue(state)
    assert len(jobs) == 1 and state['running'][0] == 'activity'
    state['wanted']['activity'] = 'new'
    old_entry = completed_entry(state, 'OUTDATED PAGE')
    jobs[0][0].set_result(old_entry)
    slot.release()
    ui._advance_preview_queue(state)
    assert 'activity' not in state['cache']
    assert len(jobs) == 2 and state['running'][0] == 'training'
    current_entry = completed_entry(state, 'Current training page')
    jobs[1][0].set_result(current_entry)
    slot.release()
    ui._advance_preview_queue(state)
    assert state['cache']['training']['ticket'] == current_entry['ticket']
    assert not ui.STORE.available(state['owner'], old_entry['ticket'])
    assert ui.STORE.available(state['owner'], current_entry['ticket'])
    assert state['running'] is None


def test_ordered_preview_failure_releases_conversion_slot(monkeypatch):
    from threading import BoundedSemaphore
    slot = BoundedSemaphore(1)
    assert slot.acquire(False)
    monkeypatch.setattr(ui, '_PREVIEW_SLOT', slot)
    def fail(*args):
        raise RuntimeError('converter failed')
    monkeypatch.setattr(ui, 'preview_section', fail)
    result = ui._build_preview_entry(draft(), 'activity', {}, 'Activity')
    assert 'error' in result and 'edits are still here' in result['error']
    assert slot.acquire(False)


def test_busy_preview_retries_snapshot_after_poll_without_caching_failure(monkeypatch):
    from concurrent.futures import Future
    from threading import BoundedSemaphore
    jobs = []
    class Worker:
        def submit(self, fn, *args):
            future = Future()
            jobs.append((future, args))
            return future
    clock = [10.0]
    monkeypatch.setattr(ui, 'monotonic', lambda: clock[0])
    monkeypatch.setattr(ui, '_PREVIEW_WORKER', Worker())
    slot = BoundedSemaphore(1)
    monkeypatch.setattr(ui, '_PREVIEW_SLOT', slot)
    snapshot = ('current', draft(), {'picture': b'pixels'}, 'Activity')
    state = {'wanted': {'activity': 'current'}, 'queue': {'activity': snapshot},
             'cache': {}, 'running': None}
    ui._advance_preview_queue(state)
    jobs[0][0].set_result({'busy': True})
    slot.release()
    ui._advance_preview_queue(state)
    assert state['cache'] == {} and state['queue']['activity'] == snapshot
    ui._advance_preview_queue(state)
    assert len(jobs) == 1
    clock[0] += 2
    ui._advance_preview_queue(state)
    assert len(jobs) == 2 and jobs[0][1] == jobs[1][1]
    current_entry = completed_entry(state, 'Current page')
    jobs[1][0].set_result(current_entry)
    slot.release()
    ui._advance_preview_queue(state)
    assert state['cache']['activity']['ticket'] == current_entry['ticket']
    assert not state['queue'] and state['running'] is None


def test_busy_conversion_is_transient_and_releases_preview_slot(monkeypatch):
    from threading import BoundedSemaphore
    slot = BoundedSemaphore(1)
    assert slot.acquire(False)
    monkeypatch.setattr(ui, '_PREVIEW_SLOT', slot)
    def busy(*args):
        raise ui.RenderBusy('Another export is running')
    monkeypatch.setattr(ui, 'preview_section', busy)
    assert ui._build_preview_entry(draft(), 'activity', {}, 'Activity') == {'busy': True}
    assert slot.acquire(False)


def test_ordered_preview_pane_stays_visible_and_editor_is_usable_while_rendering(monkeypatch):
    from concurrent.futures import Future
    from threading import BoundedSemaphore
    jobs = []
    class Worker:
        def submit(self, fn, *args):
            future = Future()
            jobs.append(future)
            return future
    monkeypatch.setattr(ui, '_PREVIEW_WORKER', Worker())
    slot = BoundedSemaphore(1)
    monkeypatch.setattr(ui, '_PREVIEW_SLOT', slot)
    monkeypatch.setattr(ui, 'DEBOUNCE_SECONDS', 0)
    app = AppTest.from_string('''
import streamlit as st
from dataclasses import replace
from app.monthly_report_model import synthetic_draft, synthetic_profiles, ReportPeriod, ResolvedBlock, default_sections
from app.monthly_report_section_preview_ui import render_section_preview
value = st.text_area("Activity", "Current activity")
draft = synthetic_draft(synthetic_profiles()[0], ReportPeriod(2026, 9))
draft = replace(draft, sections=default_sections(), blocks=(ResolvedBlock("activity_summary", "This month", text=value),))
render_section_preview(draft, "activity", {}, "ordered", deferred=True)
''', default_timeout=15).run()
    assert not app.exception
    assert app.get('bidi_component') and not payload(app)['pages'] and not jobs
    enable_view(app, 'ordered', 'activity')
    assert len(jobs) == 1
    app.text_area[0].set_value('A saved new edit').run()
    assert not app.exception and app.text_area[0].value == 'A saved new edit'
    assert len(jobs) == 1
    state = app.session_state['ordered_ordered_previews']
    old_entry = completed_entry(state, 'Old preview')
    jobs[0].set_result(old_entry)
    slot.release()
    app.run()
    assert not payload(app)['pages'] and len(jobs) == 2
    assert not ui.STORE.available(state['owner'], old_entry['ticket'])
    current = completed_entry(state, 'Current preview')
    jobs[1].set_result(current)
    slot.release()
    app.run()
    assert not app.exception and payload(app)['pages']
    assert state['cache']['activity']['ticket'] == current['ticket']
    assert app.text_area[0].value == 'A saved new edit'


def test_missing_preview_source_does_not_crash_the_editor(monkeypatch, tmp_path):
    def unavailable(*args):
        raise OSError('source unavailable')
    monkeypatch.setattr(ui, 'section_preview_fingerprint', unavailable)
    app, calls = preview_app(monkeypatch, tmp_path)
    assert not app.exception and not calls
    assert any('edits are still here' in value.value for value in app.warning)
    app.text_area[0].set_value('Still editable').run()
    assert not app.exception and app.text_area[0].value == 'Still editable'


def test_section_preview_refreshes_when_report_month_changes():
    from app.monthly_report_model import ReportPeriod
    original = draft()
    changed = replace(original, period=ReportPeriod(2026, 10))
    assert engine.section_preview_fingerprint(original, 'activity') != engine.section_preview_fingerprint(changed, 'activity')


def test_previous_preview_is_bounded_and_cleaned_after_success_or_report_switch():
    def pdf_preview(pages):
        with fitz.open() as document:
            for index in range(pages):
                document.new_page().insert_text((50, 50), 'Synthetic old preview ' + str(index))
            return engine.ReportPreview(document.tobytes(), pages, 'test', 'September 2026')
    owner = 'test-preview-owner'
    state = {'owner': owner, 'cache': {}, 'last_good': {}, 'wanted': {'activity': 'first'},
             'queue': {}, 'running': None}
    first = {'fingerprint': 'first', 'ticket': ui.STORE.put(owner, pdf_preview(1))}
    ui._remember_previous_preview(state, 'activity', first)
    assert state['last_good']['activity']['ticket'] == first['ticket']
    second = {'fingerprint': 'second', 'ticket': ui.STORE.put(owner, pdf_preview(2))}
    ui._remember_previous_preview(state, 'training', second)
    assert set(state['last_good']) == {'training'}  # No 12-section stale-page hoard.
    assert not ui.STORE.available(owner, first['ticket'])
    ui._reset_report_preview_identity(state, ('contract-A', 'sites', '2026-09', 'pin'))
    assert not ui.STORE.available(owner, second['ticket'])
    assert not state['last_good'] and not state['cache']
    assert state['report_identity'][0] == 'contract-A'


def test_stale_result_must_not_repopulate_another_report_after_scope_change(monkeypatch):
    from concurrent.futures import Future
    from threading import BoundedSemaphore

    owner='synthetic-multi-report-owner'
    with fitz.open() as document:
        document.new_page().insert_text((50, 50), 'Old protected preview')
        ticket = ui.STORE.put(owner, engine.ReportPreview(document.tobytes(), 1, 'old', 'Sep'))
    finished=Future()
    finished.set_result({'ticket': ticket, 'pages': 1})
    state={'owner':owner,'wanted':{'activity':'old'},'queue':{},
           'cache':{},'last_good':{},'running':('activity','old',finished)}
    ui._reset_report_preview_identity(state, ('different-client','new-site','2026-10','other-pin'))
    ui._advance_preview_queue(state)
    assert 'activity' not in state['cache']
    assert not ui.STORE.available(owner,ticket), 'The previous report PDF must be discarded'


def test_pending_and_completed_preview_do_not_force_an_extra_full_editor_run():
    # The previous timer-registration path called st.rerun(scope='app') on
    # both transitions. Count actual AppTest script executions; no renderer
    # timing, source-string assertion or fabricated browser behavior is used.
    app = AppTest.from_string('''
import streamlit as st
from concurrent.futures import Future
from app.monthly_report_section_preview_ui import _ordered_preview_state, _render_ordered_preview
st.session_state['full_runs'] = st.session_state.get('full_runs', 0) + 1
state = _ordered_preview_state('synthetic_status')
if not state['wanted']:
    state['wanted']['activity'] = 'current'
    state['visible']['activity'] = True
    state['running'] = ('activity', 'current', Future())
_render_ordered_preview('activity', 'synthetic_status', 'Synthetic Activity')
st.text_input('Unrelated editor', key='untouched')
''').run()
    assert not app.exception
    assert app.session_state['full_runs'] == 1
    assert payload(app)['pending'] is True
    state = app.session_state['synthetic_status_ordered_previews']
    state['running'][2].set_result(completed_entry(state))
    app.run()
    assert not app.exception
    assert app.session_state['full_runs'] == 2
    assert payload(app)['pending'] is False and payload(app)['ready'] is True


def test_isolated_preview_widgets_keep_actual_framework_fragment_identity(monkeypatch, tmp_path):
    from app import monthly_report_testing as testing
    from streamlit.runtime.scriptrunner_utils.script_run_context import ThreadState
    seen = []
    original = ui.render_view

    def observed(**kwargs):
        seen.append((ThreadState.get().fragment_id, testing.scope(), testing.generation()))
        return original(**kwargs)

    monkeypatch.setenv('EPC_DATA_DIR', str(tmp_path))
    monkeypatch.setattr(ui, 'render_view', observed)
    app = AppTest.from_string('''
import os
from pathlib import Path
from app.monthly_report_testing import isolated, data_dir
from app.monthly_report_model import synthetic_draft, synthetic_profiles, ReportPeriod
from app.monthly_report_section_preview_ui import render_section_preview
with isolated('a'*32):
    data_dir(Path(os.environ['EPC_DATA_DIR']))
    draft=synthetic_draft(synthetic_profiles()[0],ReportPeriod(2026,9))
    render_section_preview(draft,'activity',{},'synthetic')
    render_section_preview(draft,'training',{},'synthetic')
''').run()
    assert not app.exception
    assert len(seen) == 2
    assert all(fragment and namespace and lease for fragment, namespace, lease in seen)
    assert seen[0][0] != seen[1][0], 'each section must own a distinct actual fragment'
    assert seen[0][1:] == seen[1][1:], 'both previews retain the same isolated lease'
    for component in app.get('bidi_component'):
        metadata = app.session_state._state._new_widget_state.widget_metadata[component.proto.id]
        assert metadata.fragment_id in {item[0] for item in seen}
    assert testing.scope() is None and testing.generation() is None


def test_captured_namespace_preserves_current_runtime_context_and_rejects_reset(tmp_path):
    from contextvars import ContextVar
    from app import monthly_report_testing as testing
    framework = ContextVar('synthetic_current_fragment', default='outer')
    with testing.isolated('b'*32):
        target = testing.data_dir(tmp_path)
        captured = testing.capture_namespace()
        old = (testing.scope(), testing.generation())
    framework.set('current-fragment')
    with testing.isolated('c'*32):
        sibling = testing.data_dir(tmp_path)
        sibling_identity = (testing.scope(), testing.generation())
        with captured():
            assert framework.get() == 'current-fragment'
            assert (testing.scope(), testing.generation()) == old
            assert testing.data_dir(tmp_path) == target
            testing.reset(tmp_path)
        assert (testing.scope(), testing.generation()) == sibling_identity
        assert testing.data_dir(tmp_path) == sibling
    with captured():
        with pytest.raises(OSError, match='reset'):
            testing.data_dir(tmp_path)
    assert not target.exists() and sibling.exists()
    assert testing.scope() is None and testing.generation() is None


def test_preview_callback_carries_namespace_before_cache_access_and_rejects_reset(monkeypatch, tmp_path):
    from app import monthly_report_testing as testing
    monkeypatch.setenv('EPC_DATA_DIR', str(tmp_path))
    seen = []
    with testing.isolated('d'*32):
        target = testing.data_dir(tmp_path)
        expected = (testing.scope(), testing.generation())
        callback = ui._scoped_callback(lambda value: seen.append((value, testing.scope(), testing.generation())), 'view')
    callback()
    assert seen == [('view', *expected)]
    assert testing.scope() is None
    with testing.isolated('d'*32):
        testing.data_dir(tmp_path)
        testing.reset(tmp_path)
    with pytest.raises(OSError, match='reset'):
        callback()
    assert seen == [('view', *expected)]
    assert not target.exists(), 'stale callback cannot recreate its erased namespace'
