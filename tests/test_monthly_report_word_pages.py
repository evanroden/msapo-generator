"""Synthetic native drawings only; supplied chart XML must never enter git."""

from dataclasses import replace
from io import BytesIO
import hashlib
from zipfile import ZipFile

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.oxml import parse_xml
from PIL import Image
import pytest

from app import monthly_report_word_pages as pages, monthly_report_library as library
from app.monthly_report_docx import ReportImage
from app.monthly_report_import import inspect_docx, imported_draft, W
from app.monthly_report_model import ReportPeriod, synthetic_profiles
from app.monthly_report_sections import build_section_import
from test_monthly_report_import import rewrite
from tests.conftest import requires_libreoffice


def drawing_report(tmp_path):
    doc = Document()
    doc.add_heading("Synthetic cover", 0)
    doc.add_section(WD_SECTION_START.NEW_PAGE)
    doc.add_heading("1 Organizational Chart", 1)
    xml = f'''<w:p xmlns:w="{W[1:-1]}" xmlns:v="urn:schemas-microsoft-com:vml">
      <w:r><w:pict><v:group style="position:relative;width:450pt;height:220pt" coordsize="450,220">
      <v:rect style="position:absolute;left:40;top:20;width:170;height:70" fillcolor="#D9EFDF">
        <v:textbox><w:txbxContent><w:p><w:r><w:t>Synthetic Lead</w:t></w:r></w:p></w:txbxContent></v:textbox>
      </v:rect><v:rect style="position:absolute;left:230;top:130;width:170;height:70" fillcolor="#DAEEB0">
        <v:textbox><w:txbxContent><w:p><w:r><w:t>Synthetic Technician</w:t></w:r></w:p></w:txbxContent></v:textbox>
      </v:rect><v:line from="125,90" to="315,130" strokecolor="#092B24"/>
      </v:group></w:pict></w:r></w:p>'''
    doc._element.body.insert(-1, parse_xml(xml))
    doc.add_section(WD_SECTION_START.NEW_PAGE)
    doc.add_heading("2 Monthly Activity Summary", 1)
    doc.add_paragraph("UNSELECTED MONTHLY TEXT")
    path = tmp_path / "synthetic-drawings.docx"
    doc.save(path)
    return path


def test_passive_subset_retains_geometry_and_drops_unselected_content(tmp_path):
    path = drawing_report(tmp_path)
    before = path.read_bytes()
    raw = pages.safe_section_docx(path, (2,))
    assert path.read_bytes() == before
    with ZipFile(BytesIO(raw)) as z:
        xml = z.read("word/document.xml")
        assert b"Synthetic Technician" in xml and b"UNSELECTED MONTHLY TEXT" not in xml
        assert b"coordsize" in xml and b"220pt" in xml
        assert not any(n.startswith(("customXml/", "word/embeddings/")) for n in z.namelist())
        assert len(Document(BytesIO(raw)).sections) == 1


def test_links_and_fields_are_not_opened_and_embedded_objects_reject(tmp_path):
    path = drawing_report(tmp_path)
    with ZipFile(path) as z:
        xml = z.read("word/document.xml").decode()
    inject = '<w:r><w:instrText>INCLUDETEXT https://example.invalid/secret</w:instrText></w:r>'
    changed = xml.replace('<w:t>Synthetic Lead</w:t>', '<w:t>Synthetic Lead</w:t></w:r>' + inject + '<w:r>')
    rewrite(path, {"word/document.xml": changed})
    raw = pages.safe_section_docx(path, (2,))
    with ZipFile(BytesIO(raw)) as z:
        assert b"INCLUDETEXT" not in z.read("word/document.xml")
        assert b"example.invalid" not in z.read("word/document.xml")
    rewrite(path, {"word/document.xml": xml.replace('<w:t>Synthetic Lead</w:t>', '<w:object/>')})
    with pytest.raises(ValueError, match="embedded object"):
        pages.safe_section_docx(path, (2,))


def test_dependency_entities_depth_and_bounds_fail_closed(tmp_path, monkeypatch):
    path = drawing_report(tmp_path)
    with pytest.raises(ValueError, match="one and eight"):
        pages.safe_section_docx(path, tuple(range(1, 10)))
    with pytest.raises(ValueError, match="changed"):
        pages.safe_section_docx(path, (99,))
    with pytest.raises(ValueError, match="complexity"):
        pages._bounded_dependency(BytesIO(b"<x>" * 90 + b"</x>" * 90))
    monkeypatch.setattr(pages, "MAX_XML", 20)
    with pytest.raises(ValueError, match="XML budget"):
        pages.safe_section_docx(path, (2,))


def test_external_picture_and_macro_parts_never_reach_converter(tmp_path):
    path = drawing_report(tmp_path)
    with ZipFile(path) as z:
        xml = z.read("word/document.xml").decode()
    image = '<w:drawing><a:graphic xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><a:blip r:link="rIdMissing"/></a:graphic></w:drawing>'
    rewrite(path, {"word/vbaProject.bin": b"SYNTHETIC MACRO", "word/document.xml": xml.replace('<w:t>Synthetic Lead</w:t>', image)})
    with pytest.raises(ValueError, match="linked picture"):
        pages.safe_section_docx(path, (2,))
    rewrite(path, {"word/document.xml": xml})
    with ZipFile(BytesIO(pages.safe_section_docx(path, (2,)))) as z:
        assert not any('vba' in n.lower() for n in z.namelist())


def test_smartart_document_scoped_drawing_dependency_is_closed(tmp_path):
    path = drawing_report(tmp_path)
    with ZipFile(path) as z:
        xml = z.read("word/document.xml").decode()
        rels = z.read("word/_rels/document.xml.rels").decode()
    diagram = 'http://schemas.openxmlformats.org/drawingml/2006/diagram'
    drawing = 'http://schemas.microsoft.com/office/drawing/2008/diagram'
    reference = f'<w:drawing><a:graphic xmlns:a="{pages.A[1:-1]}"><a:graphicData uri="{diagram}"><d:relIds xmlns:d="{diagram}" r:dm="rIdData"/></a:graphicData></a:graphic></w:drawing>'
    xml = xml.replace('<w:t>Synthetic Lead</w:t>', reference)
    rels = rels.replace('</Relationships>', f'<Relationship Id="rIdData" Type="{pages._OFFICE}/diagramData" Target="diagrams/data1.xml"/><Relationship Id="rIdDrawing" Type="http://schemas.microsoft.com/office/2007/relationships/diagramDrawing" Target="diagrams/drawing1.xml"/></Relationships>')
    rewrite(path, {"word/document.xml": xml, "word/_rels/document.xml.rels": rels,
                   "word/diagrams/data1.xml": f'<d:dataModel xmlns:d="{diagram}" xmlns:x="{drawing}"><x:dataModelExt relId="rIdDrawing"/></d:dataModel>',
                   "word/diagrams/drawing1.xml": f'<x:drawing xmlns:x="{drawing}"/>',
                   "word/diagrams/unused.xml": '<unused/>'})
    with ZipFile(BytesIO(pages.safe_section_docx(path, (2,)))) as z:
        assert "word/diagrams/drawing1.xml" in z.namelist()
        assert "word/diagrams/unused.xml" not in z.namelist()
        assert b'rIdDrawing' in z.read('word/_rels/document.xml.rels')
        assert b'TargetMode="External"' not in b''.join(z.read(n) for n in z.namelist() if n.endswith('.rels'))
    rewrite(path, {"word/_rels/document.xml.rels": rels.replace('Target="diagrams/drawing1.xml"', 'Target="https://example.invalid/drawing.xml" TargetMode="External"')})
    with pytest.raises(ValueError, match="external, missing or unsupported"):
        pages.safe_section_docx(path, (2,))


@requires_libreoffice
def test_real_native_page_renders_and_stage_cleans_up(tmp_path, monkeypatch):
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path / "runtime"))
    path = drawing_report(tmp_path)
    result = pages.preserve_word_pages(path, (2,))
    assert len(result) == 1
    assert "Synthetic Lead" in result[0].text and "Synthetic Technician" in result[0].text
    assert "UNSELECTED" not in result[0].text
    image = Image.open(BytesIO(result[0].data)).convert("RGB")
    assert image.width > 1000 and image.getextrema()[0][0] < 200
    assert not list((tmp_path / "runtime/monthly_reports/imports").glob("chart-*"))


def plan_for(image, text="Synthetic chart", reviewed=True):
    return {"key": "organization", "target": "organization", "page_layout": True,
            "approved": True, "preserved_assets": [("org_chart", image)],
            "word_page_records": [{"page": 1, "slot": "org_chart", "sections": (2,),
                                   "sha256": hashlib.sha256(image.data).hexdigest(), "text": text, "reviewed": reviewed}]}


def test_page_review_gates_provenance_and_saved_image_survive_reload(tmp_path, monkeypatch):
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path / "runtime"))
    path = drawing_report(tmp_path)
    inspection = inspect_docx(path)
    raw = BytesIO(); Image.new("RGB", (400, 600), "navy").save(raw, "PNG")
    image = ReportImage(raw.getvalue(), "png", 400, 600)
    for text, reviewed in [("Price $500", True), ("Insert image", True), ("Synthetic chart", False)]:
        with pytest.raises(ValueError, match="Review it again"):
            build_section_import(path, inspection, (plan_for(image, text, reviewed),))
    plan = plan_for(image)
    changed = dict(plan, preserved_assets=[("org_chart", replace(image, data=b"changed"))])
    with pytest.raises(ValueError, match="Review it again"):
        build_section_import(path, inspection, (changed,))
    mapped, _ = build_section_import(path, inspection, (plan,))
    draft = imported_draft(synthetic_profiles()[0], ReportPeriod(2026, 9), "Synthetic Editor", mapped)
    saved = library.save_snapshot(draft, expected_revision=0, assets=mapped.assets, entered_editor="Synthetic Editor")
    loaded = library.load_snapshot(draft.profile.contract, draft.profile.key, draft.period)
    assert loaded == saved
    block = loaded.draft.blocks[0]
    assert block.references == (f"docx:{inspection.sha256}:word-sections:2:page:1",)
    assert block.asset_hashes[0] == hashlib.sha256(image.data).hexdigest() + ".png"


def test_failed_render_releases_preparation_slot(tmp_path, monkeypatch):
    path = drawing_report(tmp_path)
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path / "runtime"))
    def fail(*_):
        raise ValueError("Synthetic conversion failed")
    monkeypatch.setattr(pages, "_render", fail)
    with pytest.raises(ValueError, match="Synthetic conversion failed"):
        pages.preserve_word_pages(path, (2,))
    assert pages._RENDER_LOCK.acquire(blocking=False)
    pages._RENDER_LOCK.release()
    assert not list((tmp_path / "runtime/monthly_reports/imports").glob("chart-*"))


@requires_libreoffice
def test_guided_native_page_selection_save_reopen_and_review_invalidation(tmp_path, monkeypatch):
    from streamlit.testing.v1 import AppTest
    from streamlit.testing.v1.element_tree import ElementTree
    from app import monthly_report_section_ui as ui
    from app.monthly_report_setup import new_month_draft
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path / "runtime"))
    path = drawing_report(tmp_path)
    upload = BytesIO(path.read_bytes())
    upload.name, upload.size = path.name, len(upload.getvalue())
    real_upload = ui.st.file_uploader
    monkeypatch.setattr(ui.st, "file_uploader", lambda label, *a, **kw:
                        upload if label == "Older or partially completed report" else real_upload(label, *a, **kw))
    original_states = ElementTree.get_widget_states
    def states(tree):
        values = original_states(tree)
        for key, identity in tree.session_state._state._key_id_mapper._key_id_mapping.items():
            if key.startswith("report_setup_") and key.endswith("_open"):
                values.widgets.add(id=identity, bool_value=bool(tree.session_state[key]))
        return values
    monkeypatch.setattr(ElementTree, "get_widget_states", states)
    app = AppTest.from_string('''
from app.monthly_report_section_ui import render_section_setup
from app.monthly_report_model import synthetic_profiles, ReportPeriod
from app.monthly_report_ui import _field
p = synthetic_profiles()[0]
render_section_setup(p.contract, ReportPeriod(2026,9), "Synthetic Editor", _field, identity=p)
''', default_timeout=30).run()
    next(b for b in app.button if b.label == "Analyze report").click().run()
    inspection = inspect_docx(path)
    stage_key = next(k for k in app.session_state.filtered_state if k.endswith("_stage"))
    prefix = stage_key.removesuffix("_stage") + "_" + inspection.sha256[:16]
    app.session_state[prefix + "_section_organization_open"] = True
    app.run()
    next(b for b in app.button if b.label == "Prepare chart and contact pages").click().run()
    assert not app.exception
    next(w for w in app.checkbox if w.label == "Use complete pages for this section").check().run()
    assert any("at least one complete page" in e.value for e in app.error)
    next(w for w in app.checkbox if w.label == "Include page 1").check().run()
    assert any("Check page 1 against" in e.value for e in app.error)
    next(w for w in app.checkbox if w.label.startswith("Page 1 matches")).check().run()
    org_ready = lambda: next(w for w in app.checkbox if w.key.startswith(prefix + "_section_organization_ready_"))
    org_ready().check().run()
    assert app.session_state[prefix + "_section_plans"]["organization"]["approved"]
    # A different purpose invalidates the specific page confirmation.
    next(w for w in app.selectbox if w.label == "Page 1 contains").set_value("contact_matrix").run()
    assert not app.session_state[prefix + "_section_plans"]["organization"]["approved"]
    next(w for w in app.selectbox if w.label == "Page 1 contains").set_value("org_chart").run()
    next(w for w in app.checkbox if w.label.startswith("Page 1 matches")).check().run()
    org_ready().check().run()
    for key in ("cover", "activity"):
        app.session_state[prefix + "_section_" + key + "_open"] = True
        app.run()
        next(w for w in app.radio if w.key == prefix + "_section_" + key + "_action").set_value("Leave this section out").run()
        next(w for w in app.checkbox if w.key.startswith(prefix + "_section_" + key + "_ready_")).check().run()
    next(w for w in app.checkbox if w.label == "Save this report and its reusable design for these sites").check().run()
    next(b for b in app.button if b.label == "Continue to this month’s updates").click().run()
    assert not app.exception
    profile = synthetic_profiles()[0]
    draft = library.load_imported_draft(profile.contract, profile.key)
    chart = next(b for b in draft.blocks if b.key == "org_chart")
    assert len(chart.asset_hashes) == 1 and chart.references
    later = new_month_draft(draft, ReportPeriod(2026, 10))
    assert next(b for b in later.blocks if b.key == "org_chart").asset_hashes == chart.asset_hashes
    assert library.imported_original(profile.contract, profile.key).read_bytes() == path.read_bytes()
