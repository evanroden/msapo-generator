"""Imported cover dates and native long-title formatting stay source-specific."""
from dataclasses import replace
from io import BytesIO

import pytest
from docx import Document
from docx.shared import Pt

from app.monthly_report_import import inspect_docx
from app.monthly_report_model import (
    BlockSpec, Facility, ReportDraft, ReportPeriod, ReportProfile, ResolvedBlock, SectionSpec,
)
from app.monthly_report_native_layout import build_native_docx


def cover_source(tmp_path, *, combined=False):
    document = Document()
    title = 'Synthetic Hospital With A Long Facility Name'
    document.add_paragraph(title).runs[0].font.size = Pt(27)
    document.add_paragraph('September 2026')
    document.add_paragraph('Prepared by: Source Editor' + (' Date: October 6th, 2026' if combined else ''))
    if not combined:
        document.add_paragraph('Date: 10/06/2026')
    document.add_heading('MONTHLY ACTIVITY SUMMARY', 1)
    document.add_paragraph('Completed technical inspection.')
    path = tmp_path / 'cover.docx'
    document.save(path)
    inspection = inspect_docx(path)
    item = next(item for item in inspection.items if item.kind == 'text' and item.text == 'Completed technical inspection.')
    block = ResolvedBlock('activity_summary', 'This month', text=item.text,
                          references=(f'docx:{inspection.sha256}:{item.id}',))
    draft = ReportDraft(ReportProfile('Synthetic contract', 'site', title, (Facility('site', title),)),
                        ReportPeriod(2026, 9), 'Source Editor',
                        (SectionSpec('activity', '2', 'Monthly Activity Summary',
                                     (BlockSpec('activity_summary', 'rich_text'),)),), (block,))
    return path, draft


@pytest.mark.parametrize('combined', [False, True])
def test_same_report_keeps_issue_date_and_long_title_font(tmp_path, combined):
    path, draft = cover_source(tmp_path, combined=combined)
    actual = Document(BytesIO(build_native_docx(draft, path)))
    assert actual.paragraphs[0].text == draft.profile.title
    assert actual.paragraphs[0].runs[0].font.size == Pt(27)
    text = '\n'.join(p.text for p in actual.paragraphs)
    assert ('Date: October 6th, 2026' if combined else 'Date: 10/06/2026') in text
    assert text.count(draft.profile.title) == 1


@pytest.mark.parametrize('change', ['master', 'month', 'client', 'editor', 'no_import'])
def test_fresh_or_changed_report_does_not_inherit_source_issue_date(tmp_path, change):
    path, draft = cover_source(tmp_path)
    if change == 'month':
        draft = replace(draft, period=ReportPeriod(2026, 10))
    elif change == 'client':
        draft = replace(draft, profile=replace(draft.profile, title='Current other client'))
    elif change == 'editor':
        draft = replace(draft, prepared_by='Current Editor')
    elif change == 'no_import':
        draft = replace(draft, blocks=())
    actual = Document(BytesIO(build_native_docx(draft, path, master=change == 'master')))
    text = '\n'.join(p.text for p in actual.paragraphs)
    assert '10/06/2026' not in text
    assert text.count(draft.profile.title) == 1
