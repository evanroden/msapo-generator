"""Current caption receipts retain page layout without sharing occurrence ownership."""
from dataclasses import replace

import pytest

from app.monthly_report_asset_instances import (
    append_human_asset, new_asset_instance, valid_instances, with_asset_instances,
)
from app.monthly_report_asset_layout import document_page_indexes, picture_runs
from app.monthly_report_asset_review import current_source_page_review_matches
from test_current_source_page_status import reviewed


def with_equal_photo(reviewed):
    draft, _, _, _ = reviewed
    block = draft.blocks[0]
    return append_human_asset(block, block.asset_hashes[0], 'Independent human photo'), draft.sources[0]


def test_corrected_source_caption_keeps_page_layout_beside_equal_human_photo(reviewed):
    block, source = with_equal_photo(reviewed)
    before = block.fingerprint
    assert valid_instances(block)
    assert current_source_page_review_matches(block, 0, source, 1)
    assert not current_source_page_review_matches(block, 1, source, 1)
    assert document_page_indexes(block, (source,)) == {0}
    assert tuple(picture_runs(block, (source,))) == ((True, (0,)), (False, (1,)))
    assert block.fingerprint == before


def test_equal_bytes_from_other_source_keep_independent_page_identity(reviewed):
    draft, _, _, _ = reviewed
    block, source = draft.blocks[0], draft.sources[0]
    other = replace(source, id='another-synthetic-source')
    reference = f'{other.id}:1:{other.fingerprint}'
    block = with_asset_instances(replace(block,
        asset_hashes=block.asset_hashes * 2,
        asset_captions=(*block.asset_captions, 'Other source page'),
        references=(*block.references, reference)),
        (*block.asset_instances, new_asset_instance(block.asset_hashes[0], (reference,))))
    assert valid_instances(block)
    assert current_source_page_review_matches(block, 0, source, 1)
    assert not current_source_page_review_matches(block, 1, source, 1)
    assert document_page_indexes(block, (source, other)) == {0, 1}


@pytest.mark.parametrize('change', [
    'missing_records', 'missing_record', 'duplicate_keys', 'empty_key',
    'wrong_asset', 'changed_bytes', 'changed_caption', 'wrong_page',
    'stale_context', 'missing_asset_review', 'missing_page_review',
    'duplicate_marker', 'missing_block_reference', 'duplicate_mapping',
])
def test_equal_byte_occurrences_never_relax_review_or_identity_guards(reviewed, change):
    block, source = with_equal_photo(reviewed)
    first, second = block.asset_instances
    if change == 'missing_records':
        block = replace(block, asset_instances=())
    elif change == 'missing_record':
        block = replace(block, asset_instances=(first,))
    elif change == 'duplicate_keys':
        block = replace(block, asset_instances=(first, replace(second, key=first.key)))
    elif change == 'empty_key':
        block = replace(block, asset_instances=(replace(first, key=''), second))
    elif change == 'wrong_asset':
        block = replace(block, asset_instances=(replace(first, asset='0' * 64 + '.png'), second))
    elif change == 'changed_bytes':
        changed = '0' * 64 + '.png'
        block = with_asset_instances(replace(block, asset_hashes=(changed, changed)),
            (replace(first, asset=changed), replace(second, asset=changed)))
    elif change == 'changed_caption':
        block = replace(block, asset_captions=('Unreviewed new caption', block.asset_captions[1]))
    elif change == 'wrong_page':
        first = replace(first, references=tuple(ref.replace(':1:', ':2:') for ref in first.references))
        block = with_asset_instances(block, (first, second))
    elif change == 'stale_context':
        source = replace(source, vendor='Changed source vendor')
    elif change == 'missing_asset_review':
        block = replace(block, client_asset_reviews=(), client_reviewed_fingerprint='')
    elif change == 'missing_page_review':
        source = replace(source, client_page_reviews=())
    elif change == 'duplicate_marker':
        marker = next(ref for ref in first.references if ref.startswith('asset-source-context:'))
        block = with_asset_instances(block, (replace(first, references=(*first.references, marker)), second))
    elif change == 'missing_block_reference':
        block = replace(block, references=())
    elif change == 'duplicate_mapping':
        block = replace(block, asset_provenance=block.asset_provenance * 2)
    assert not current_source_page_review_matches(block, 0, source, 1)
    assert document_page_indexes(block, (source,)) == set()


def test_corrected_equal_byte_page_retains_native_document_geometry(reviewed, tmp_path):
    from io import BytesIO
    from docx import Document
    from docx.shared import Inches
    from PIL import Image
    from app.monthly_report_library import read_asset
    from app.monthly_report_native_layout import build_native_docx
    from test_monthly_report_native_layout import draft

    original, _, _, _ = reviewed
    block, source = with_equal_photo(reviewed)
    raw = read_asset(original.profile.contract, original.profile.key, block.asset_hashes[0])
    template = Document()
    template.add_heading('MONTHLY ACTIVITY SUMMARY', 1)
    template.add_paragraph('Improvement Highlights and Pics')
    template.add_picture(BytesIO(raw), width=Inches(4), height=Inches(2))
    path = tmp_path / 'synthetic-native.docx'
    template.save(path)
    current = replace(draft(block), sources=(source,))
    before = current.fingerprint
    result = Document(BytesIO(build_native_docx(current, path,
        asset_loader=lambda asset: raw if asset == block.asset_hashes[0] else None)))
    assert len(result.inline_shapes) == 2
    assert result.inline_shapes[0].height > Inches(6)
    # Ordinary photos fit within the original frame without stretching.
    photo = result.inline_shapes[1]
    assert photo.width <= Inches(4) and photo.height <= Inches(2)
    width, height = Image.open(BytesIO(raw)).size
    assert photo.width / photo.height == pytest.approx(width / height, abs=0.00001)
    for shape in result.inline_shapes:
        reference = shape._inline.graphic.graphicData.pic.blipFill.blip.embed
        assert result.part.related_parts[reference].blob == raw
    assert [p.text for p in result.paragraphs if p.text in block.asset_captions] == list(block.asset_captions)
    assert current.fingerprint == before
