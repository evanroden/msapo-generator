"""Contents prose is not a cover boundary; real headings still are."""

from io import BytesIO

from docx import Document
from docx.shared import Inches
import pytest

from app.monthly_report_import import inspect_docx
from app.monthly_report_model import ReportPeriod, synthetic_profiles
from app.monthly_report_start_import import (
    analyze_starting_report,
    automatic_section_plans,
    cover_photo_candidates,
    unresolved_cover_artwork,
)
from test_cover_branding_contracts import photograph


def _source(tmp_path, text, *, adjacent=False, small=False):
    document = Document()
    document.add_paragraph(synthetic_profiles()[0].facilities[0].title)
    document.add_paragraph("November 2026")
    document.add_paragraph(text)
    if adjacent:
        document.add_section()
    document.add_picture(
        BytesIO(photograph((96, 72) if small else (720, 480))), width=Inches(4)
    )
    source = tmp_path / "contents-boundary.docx"
    document.save(source)
    return source


@pytest.mark.parametrize("text", [
    "The contents of this report are confidential.",
    "See the table of contents for the report sections.",
])
@pytest.mark.parametrize("adjacent", [False, True], ids=["same-section", "adjacent-panel"])
def test_contents_prose_keeps_cover_photo_selectable(tmp_path, text, adjacent):
    source = _source(tmp_path, text, adjacent=adjacent)
    original = source.read_bytes()
    inspection = inspect_docx(source)
    image = next(item for item in inspection.items if item.kind == "image")
    assert cover_photo_candidates(inspection.items) == (image,)
    analysis = analyze_starting_report(inspection, synthetic_profiles()[0], ReportPeriod(2026, 11))
    plan = automatic_section_plans(inspection, analysis)["cover"]
    assert plan["approved"]
    assert plan["destinations"][image.id] == "cover_photo"
    assert source.read_bytes() == original


@pytest.mark.parametrize("text", [
    "The contents of this report are confidential.",
    "See the table of contents for the report sections.",
])
@pytest.mark.parametrize("adjacent", [False, True], ids=["same-section", "adjacent-panel"])
def test_contents_prose_keeps_unidentified_cover_artwork_recoverable(tmp_path, text, adjacent):
    inspection = inspect_docx(_source(tmp_path, text, adjacent=adjacent, small=True))
    image = next(item for item in inspection.items if item.kind == "image")
    assert not cover_photo_candidates(inspection.items)
    assert unresolved_cover_artwork(inspection.items) == (image,)
    analysis = analyze_starting_report(inspection, synthetic_profiles()[0], ReportPeriod(2026, 11))
    plan = automatic_section_plans(inspection, analysis)["cover"]
    assert not plan["approved"]
    assert "Other cover artwork has no confirmed photograph or logo role." in plan["questions"]


@pytest.mark.parametrize("heading", [
    "Contents",
    "TABLE OF CONTENTS",
    "Table of Contents:",
    "Table of\nContents",
    "Table of Contents\n1 Monthly Activity Summary\t3",
    "Monthly Activity Summary",
])
@pytest.mark.parametrize("small", [False, True], ids=["photo", "small-artwork"])
def test_real_contents_and_body_headings_end_cover_window(tmp_path, heading, small):
    inspection = inspect_docx(_source(tmp_path, heading, small=small))
    assert not cover_photo_candidates(inspection.items)
    assert not unresolved_cover_artwork(inspection.items)
