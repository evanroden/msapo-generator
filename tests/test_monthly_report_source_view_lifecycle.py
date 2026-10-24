"""Explicit original/rendition lifecycle, with real temporary local persistence."""
from dataclasses import replace
from io import BytesIO
from zipfile import ZipFile

from docx import Document
from docx.oxml.ns import qn
import pytest

from app import monthly_report_library as library
from app.monthly_report_asset_review import approve_all_assets
from app.monthly_report_compatibility_view import derive_compatibility_view, CompatibilityError
from app.monthly_report_import import ImportMapping, map_items, imported_draft
from app.monthly_report_model import Facility, ReportPeriod, ResolvedBlock, synthetic_profiles
from app.monthly_report_native_layout import build_native_docx, NativeLayoutError
from app.monthly_report_start_import import automatic_section_plans, StartingReportAnalysis, build_starting_import
from test_monthly_report_source_frame_logo_role import source, raster, A


def _body_images(raw):
    document=Document(BytesIO(raw))
    return tuple(document.part.rels[n.get(qn('r:embed'))].target_part.blob for n in document.element.body.iter(A+'blip'))


def _stored(tmp_path, monkeypatch, *, contents=False, static=False):
    monkeypatch.setenv('EPC_DATA_DIR',str(tmp_path/'store'))
    if static:
        from test_monthly_report_source_view_closure import static_source
        path=static_source(tmp_path,shared_outside=True)
    else:
        path=source(tmp_path,divergent=True)
    if contents:
        _add_contents(path)
    view=derive_compatibility_view(path.read_bytes()); audit=view.original_inspection
    profile=replace(synthetic_profiles()[0],imported_from=audit.sha256,
                    title='Synthetic Regional Community Health Hospital',
                    facilities=(Facility('synthetic','Synthetic Regional Community Health Hospital'),))
    analysis=StartingReportAnalysis((), 'same', (), (), ())
    plans=automatic_section_plans(audit,analysis,preferred_logos=('client_logo',),source_view=view)
    assert all(p['approved'] for p in plans.values()),plans
    mapped,_=build_starting_import(path,audit,tuple(plans.values()),source_view=view)
    # Registry review supplies replacement bytes independently of source identity.
    replacement=raster('purple'); ref=library.asset_reference(replacement,'png')
    client=approve_all_assets(ResolvedBlock('client_logo','Library',asset_hashes=(ref,)))
    mapped=replace(mapped,blocks=(*mapped.blocks,client),assets=(*mapped.assets,(ref,replacement)))
    draft=imported_draft(profile,ReportPeriod(2026,9),'Synthetic Editor',mapped)
    library.save_profile(profile,expected_revision=0,actor='Synthetic Reviewer',confirmed=True)
    review={'sha256':audit.sha256,'period':draft.period.key,'image_review_scope':'synthetic'}
    library.save_report_setup(profile,draft,path,review,assets=mapped.assets,expected_revision=1,
                             actor='Synthetic Reviewer',confirmed=True,source_draft=draft,source_view=view)
    loaded=library.load_imported_draft(profile.contract,profile.key,source_view=view)
    loader=lambda key: library.read_asset(profile.contract,profile.key,key)
    return path,view,loaded,loader,replacement


@pytest.fixture(params=(False,True), ids=("picture-only","finite-static-siblings"))
def stored(tmp_path,monkeypatch,request):
    return _stored(tmp_path,monkeypatch,static=request.param)


def _add_contents(path):
    from copy import deepcopy
    from lxml import etree as ET
    from test_monthly_report_native_design_import import frame
    from test_monthly_report_source_frame_logo_role import _picture_in_frame,MC,WPS
    document=Document(path)
    trailing=[p._p for p in document.paragraphs[2:]]
    for node in trailing:node.getparent().remove(node)
    document.add_section()
    paragraph=document.add_paragraph()
    for text in ('ENFRA','Section 1: Work Summary'):
        paragraph.add_run()._r.append(frame(text))
    slot=frame('');paragraph.add_run()._r.append(slot)
    _picture_in_frame(document,slot,raster('orange'))
    fill=next(slot.iter(qn('pic:blipFill')))
    ET.SubElement(fill,A+'srcRect',l='100',r='200',t='300',b='400')
    wrapper=ET.Element(MC+'AlternateContent',nsmap={'wps':WPS[1:-1]})
    slot.getparent().replace(slot,wrapper)
    ET.SubElement(wrapper,MC+'Choice',Requires='wps').append(slot)
    fallback=ET.SubElement(wrapper,MC+'Fallback');fallback.append(deepcopy(slot))
    rid,_=document.part.get_or_add_image(BytesIO(raster('green','JPEG')))
    next(fallback.iter(A+'blip')).set(qn('r:embed'),rid)
    for node in trailing:document.element.body.insert(len(document.element.body)-1,node)
    document.save(path)


def test_import_save_reload_native_full_and_section_keep_exact_occurrences(stored):
    path,view,draft,loader,replacement=stored
    original=path.read_bytes(); references=tuple(b.references for b in draft.blocks)
    result=build_native_docx(draft,path,asset_loader=loader,source_view=view)
    images=_body_images(result)
    assert images[0] == replacement
    assert images[1] != replacement, 'same-byte independent source picture is not the client slot'
    assert path.read_bytes()==original
    assert tuple(b.references for b in draft.blocks)==references
    assert all(not any(view.rendition_inspection.sha256 in ref for ref in b.references) for b in draft.blocks)
    from app.monthly_report_docx import _build_docx
    from app.monthly_report_preview import section_preview_docx
    assert _body_images(_build_docx(draft,asset_loader=loader,source_view=view))[0]==replacement
    assert _body_images(section_preview_docx(draft,'cover',loader,source_view=view))[0]==replacement
    replay,verified=library.starting_contribution(draft.profile.contract,draft.profile.key,draft.period,
                                                 view.original_inspection.sha256,source_view=view)
    assert verified and replay==draft
    assert _body_images(build_native_docx(draft,path,asset_loader=loader))[0] == replacement
    with pytest.raises(NativeLayoutError,match='master'):
        build_native_docx(draft,path,asset_loader=loader,source_view=view,master=True)


def test_each_import_operation_rederives_once_not_per_image(tmp_path,monkeypatch):
    import app.monthly_report_operational_source as operations
    path=source(tmp_path,divergent=True)
    view=derive_compatibility_view(path.read_bytes());audit=view.original_inspection
    images=[i for i in audit.items if i.kind=='image']
    original=operations.rederive_compatibility_view; calls=[]
    def count(*args,**kwargs):calls.append(1);return original(*args,**kwargs)
    monkeypatch.setattr(operations,'rederive_compatibility_view',count)
    mapped=map_items(path,audit,(ImportMapping(images[0].id,'client_logo'),ImportMapping(images[1].id,'cover_photo')),source_view=view)
    assert len(calls)==1
    assert all(ref.startswith('docx:'+audit.sha256+':') for b in mapped.blocks for ref in b.references)
    calls.clear()
    plans=automatic_section_plans(audit,StartingReportAnalysis((),'same',(),(),()),source_view=view)
    build_starting_import(path,audit,tuple(plans.values()),source_view=view)
    assert len(calls)==1


def test_expected_view_tamper_cannot_reach_native_or_import(stored):
    path,view,draft,loader,_=stored
    bad=replace(view,rendition_bytes=view.rendition_bytes+b'tampered')
    with pytest.raises(ValueError):build_native_docx(draft,path,asset_loader=loader,source_view=bad)
    with pytest.raises(CompatibilityError):map_items(path,view.original_inspection,(),source_view=bad)


def test_exact_client_omit_leaves_independent_same_bytes(stored):
    path,view,draft,loader,_=stored
    blocks=tuple(replace(b,source='Omit',asset_hashes=(),references=()) if b.key=='client_logo' else b for b in draft.blocks)
    result=build_native_docx(replace(draft,blocks=blocks),path,asset_loader=loader,source_view=view)
    assert len(_body_images(result))==1


def test_proved_client_requires_current_decision_and_conflicts_fail(stored):
    path,view,draft,loader,_=stored
    missing=replace(draft,blocks=tuple(b for b in draft.blocks if b.key!='client_logo'))
    with pytest.raises(NativeLayoutError,match='no current artwork decision'):
        build_native_docx(missing,path,asset_loader=loader,source_view=view)
    image=next(i for i in view.original_inspection.items if i.image_role=='client_logo')
    ref=f'docx:{view.original_inspection.sha256}:{image.id}'
    conflicting=replace(draft,blocks=tuple(replace(b,references=(*b.references,ref)) if b.key=='cover_photo' else b for b in draft.blocks))
    with pytest.raises(NativeLayoutError,match='Conflicting exact'):
        build_native_docx(conflicting,path,asset_loader=loader,source_view=view)


def test_final_generate_and_full_section_preview_share_view_and_keep_gates(stored,monkeypatch,tmp_path):
    """Exercise actual public builders/storage; PDF adapter stub is NOT layout QA."""
    from app import monthly_report_docx as word, monthly_report_preview as preview
    from app import monthly_report_render_jobs as jobs, pdf_converter
    import fitz
    path,view,draft,loader,replacement=stored
    final=replace(draft,blocks=tuple(approve_all_assets(b) for b in draft.blocks))
    original=path.read_bytes();converted=[]
    monkeypatch.setattr(pdf_converter,'OUTPUT_DIR',tmp_path)
    def convert(source,**kwargs):
        converted.append(source.read_bytes())
        target=tmp_path/f'converted-{len(converted)}.pdf'
        with fitz.open() as pdf:
            pdf.new_page().insert_text((72,72),'Synthetic adapter output, not visual QA')
            pdf.save(target)
        return target
    monkeypatch.setattr(jobs,'convert_to_pdf',convert)
    calls=[];verify=library._verified_source_view
    def count(*args,**kwargs):calls.append(1);return verify(*args,**kwargs)
    monkeypatch.setattr(library,'_verified_source_view',count)
    package=word.generate_report(final,asset_loader=loader,source_view=view,
                                 acknowledged_fingerprint=final.fingerprint)
    assert package.pdf and len(calls)==1
    assert _body_images(package.docx)[0]==replacement
    assert _body_images(converted[-1])[0]==replacement
    calls.clear()
    assert preview.preview_report(final,loader,source_view=view).pages==1
    from PIL import Image
    assert len(calls)==1
    assert Image.open(BytesIO(_body_images(converted[-1])[0])).convert("RGB").getpixel((0,0)) == (128,0,128)
    calls.clear()
    assert preview.preview_section(final,'cover',loader,source_view=view).pages==1
    from PIL import Image
    assert len(calls)==1
    assert Image.open(BytesIO(_body_images(converted[-1])[0])).convert("RGB").getpixel((0,0)) == (128,0,128)
    assert path.read_bytes()==original
    with pytest.raises(ValueError,match='picture'):
        word.generate_report(draft,asset_loader=loader,source_view=view,acknowledged_fingerprint=draft.fingerprint)
    bad=replace(final,blocks=tuple(replace(b,text='Cost: $500') if b.key=='activity_summary' else b for b in final.blocks))
    with pytest.raises(ValueError,match='(?i)pric|cost'):
        word.generate_report(bad,asset_loader=loader,source_view=view,acknowledged_fingerprint=bad.fingerprint)


def test_public_import_signatures_do_not_accept_verification_context():
    import inspect
    from app.monthly_report_import import read_import_image
    from app.monthly_report_sections import build_section_import
    for public in (read_import_image,map_items,build_section_import,build_starting_import):
        assert '_operation' not in inspect.signature(public).parameters


@pytest.mark.parametrize("static",(False,True))
def test_two_resolved_units_lifecycle_retains_contents_geometry_and_source(tmp_path,monkeypatch,static):
    from lxml import etree as ET
    from app.monthly_report_preview import section_preview_docx
    path,view,draft,loader,replacement=_stored(tmp_path,monkeypatch,contents=True,static=static)
    before=path.read_bytes()
    assert len([b for b in view.bindings if b.status=='resolved'])==2
    contents=next(i for i in view.original_inspection.items if i.kind=='image' and i.word_section==2)
    assert contents.image_role==''
    def geometry(raw):
        document=Document(BytesIO(raw))
        node=next(n for n in document.element.body.iter(A+'blip')
                  if document.part.rels[n.get(qn('r:embed'))].target_part.blob==raster('orange'))
        picture=next(n for n in node.iterancestors() if n.tag==qn('pic:pic'))
        return (ET.tostring(picture.find(qn('pic:spPr')),method='c14n'),
                dict(picture.find(qn('pic:blipFill')).find(A+'srcRect').attrib))
    result=build_native_docx(draft,path,asset_loader=loader,source_view=view)
    assert _body_images(result)==(replacement,raster(),raster('orange'))
    assert geometry(result)==geometry(view.rendition_bytes)
    assert path.read_bytes()==before
    assert _body_images(section_preview_docx(draft,'cover',loader,source_view=view))==(replacement,raster())
    assert not _body_images(section_preview_docx(draft,'activity',loader,source_view=view))
    replay,verified=library.starting_contribution(draft.profile.contract,draft.profile.key,draft.period,
                                                 view.original_inspection.sha256,source_view=view)
    assert verified and replay==draft


@pytest.mark.parametrize('decision',['Omit','This month'])
def test_actual_photo_ui_ref_clearing_never_silently_retains_target(stored,decision):
    path,view,draft,loader,replacement=stored
    changed=replace(draft,blocks=tuple(replace(b,source=decision,asset_hashes=() if decision=='Omit' else ('custom.png',),
                                               asset_captions=(),references=()) if b.key=='cover_photo' else b for b in draft.blocks))
    with pytest.raises(library.LibraryError,match='lost its exact source placement'):
        build_native_docx(changed,path,asset_loader=loader,source_view=view)


def test_explicit_retry_revalidates_once_and_never_reassembles_word(stored,monkeypatch):
    from app import monthly_report_docx as word
    path,view,draft,loader,_=stored
    final=replace(draft,blocks=tuple(approve_all_assets(b) for b in draft.blocks))
    monkeypatch.setattr(word,'_convert_report_pdf',lambda *_:(None,'Synthetic adapter unavailable'))
    package=word.generate_report(final,asset_loader=loader,source_view=view,acknowledged_fingerprint=final.fingerprint)
    assert package.source_view_digest==view.receipt_sha256
    import hashlib
    assert package.source_view_docx_sha256==hashlib.sha256(package.docx).hexdigest()
    calls=[];verify=library._verified_source_view
    def count(*args,**kwargs):calls.append(1);return verify(*args,**kwargs)
    monkeypatch.setattr(library,'_verified_source_view',count)
    monkeypatch.setattr(word,'assemble_docx',lambda *a,**k:pytest.fail('Retry reassembled Word'))
    converted=[]
    def convert(draft,raw):converted.append(raw);return b'%PDF-synthetic',''
    monkeypatch.setattr(word,'_convert_report_pdf',convert)
    result=word.retry_report_pdf(final,package,source_view=view)
    assert len(calls)==1 and converted==[package.docx] and result.docx==package.docx
    assert result.source_view_digest==package.source_view_digest
    assert word.retry_report_pdf(final,result).docx == package.docx
    for expected,damaged in ((None,replace(package,source_view_digest='')),(view,replace(package,docx=package.docx+b'changed')),
                             (view,replace(package,source_view_digest='0'*64)),
                             (view,replace(package,source_view_digest='',source_view_docx_sha256=''))):
        with pytest.raises(ValueError,match='exact expected'):
            word.retry_report_pdf(final,damaged,source_view=expected)
    assert len(converted)==1
    saved_original=library.imported_original(final.profile.contract,final.profile.key)
    saved_original.chmod(0o600)  # Synthetic temporary-store tamper, never a real original.
    saved_original.write_bytes(saved_original.read_bytes()+b'changed source')
    with pytest.raises(ValueError):word.retry_report_pdf(final,package,source_view=view)
    assert len(converted)==1
