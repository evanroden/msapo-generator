from dataclasses import replace
from datetime import date
from pathlib import Path

from streamlit.testing.v1 import AppTest

from app import monthly_report_guided as guided, monthly_report_library as library
from app.contracts import RRH_CONTRACT
from app.monthly_report_docx import ReportPackage, assemble_docx
from app.monthly_report_model import ReportPeriod, synthetic_profiles, ResolvedBlock, default_sections

ROOT = Path(__file__).resolve().parents[1]


def monthly(monkeypatch, tmp_path, *, saved=True):
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path))
    monkeypatch.setattr(guided, "operator_today", lambda tz: date(2026, 10, 7))
    if saved:
        profile = replace(synthetic_profiles()[0], contract=RRH_CONTRACT, key="synthetic-guided",
                          excluded_sections=tuple(s.key for s in default_sections() if s.key != "activity"))
        library.save_profile(profile, expected_revision=0, actor="Synthetic Editor", confirmed=True)
        library.replace_block(profile.contract, profile.key, "activity_summary", ResolvedBlock("activity_summary", "Library", text="Synthetic initial activity."),
                              expected_revision=1, actor="Synthetic Editor", confirmed=True, label="Synthetic activity")
    app = AppTest.from_file(ROOT / "run_web.py", default_timeout=20).run()
    app.segmented_control[0].set_value("Monthly report").run()
    choose_report(app, profile.facilities[0].title if saved else None)
    assert not app.exception
    return app


def choose_report(app, site=None):
    card = next((b for b in app.button if b.label == RRH_CONTRACT), None)
    if card:
        card.click().run()
    if site:
        box = next(w for w in app.checkbox if w.label == site)
        if not box.value:
            box.check().run()
    assert not app.exception


def step(app, number):
    next(r for r in app.radio if r.label == "Report steps").set_value(guided.STEPS[number-1]).run()
    assert not app.exception


def test_empty_contract_selects_sites_and_starting_point_without_demo_or_day_selector(monkeypatch, tmp_path):
    app = monthly(monkeypatch, tmp_path, saved=False)
    assert not app.date_input
    assert not any(t.key == "report_use_library" for t in app.toggle)
    assert app.selectbox("report_month_number").value == 9
    assert app.selectbox("report_year").value == 2026
    assert not any(w.label == "Site / report" for w in app.selectbox)
    next(w for w in app.checkbox if w.key.startswith("report_sites_")).check().run()
    next(w for w in app.button if w.key and "report_start_card_upload" in w.key).click().run()
    assert any(w.label == "Older or partially completed report" for w in app.get("file_uploader"))
    assert next(b for b in app.button if b.label == "Analyze report").disabled


def test_guided_text_survives_steps_and_workflow_switches(monkeypatch, tmp_path):
    app = monthly(monkeypatch, tmp_path)
    step(app, 2)
    next(w for w in app.text_area if w.label == "Activity summary").set_value("Synthetic reviewed repair.").run()
    step(app, 3)
    step(app, 2)
    assert next(w for w in app.text_area if w.label == "Activity summary").value == "Synthetic reviewed repair."
    for workflow in ("Expense reimbursement", "Purchase order"):
        app.segmented_control[0].set_value(workflow).run()
        app.segmented_control[0].set_value("Monthly report").run()
        assert not app.exception
        assert next(w for w in app.text_area if w.label == "Activity summary").value == "Synthetic reviewed repair."


def test_guided_progress_is_persistent_and_resumes_same_month(monkeypatch, tmp_path):
    app = monthly(monkeypatch, tmp_path)
    next(w for w in app.text_input if w.label == "Prepared by").set_value("Synthetic Editor").run()
    step(app, 2)
    next(w for w in app.text_area if w.label == "Activity summary").set_value("Synthetic colleague's partial work.").run()
    next(b for b in app.button if b.label == "Save progress").click().run()
    assert not app.exception
    saved = library.load_snapshot(RRH_CONTRACT, "synthetic-guided", ReportPeriod(2026, 9))
    assert saved.entered_editor == "Synthetic Editor"
    assert next(b for b in saved.draft.blocks if b.key == "activity_summary").text == "Synthetic colleague's partial work."
    new = AppTest.from_file(ROOT / "run_web.py", default_timeout=20).run()
    new.segmented_control[0].set_value("Monthly report").run()
    choose_report(new, synthetic_profiles()[0].facilities[0].title)
    step(new, 2)
    assert next(w for w in new.text_area if w.label == "Activity summary").value == "Synthetic colleague's partial work."


def test_confirmed_import_resumes_even_with_an_older_saved_snapshot(monkeypatch, tmp_path):
    from test_monthly_report_setup import make_docx
    from app.monthly_report_import import inspect_docx
    monthly(monkeypatch, tmp_path)
    profile = library.list_profiles(RRH_CONTRACT)[0]
    state = library.load_profile(RRH_CONTRACT, profile.key)
    period = ReportPeriod(2026, 9)
    initial = guided._initial_draft(state, period, "Synthetic Editor")
    library.save_snapshot(initial, expected_revision=0, entered_editor="Synthetic Editor")
    incoming = replace(initial, blocks=tuple(replace(b, text="Synthetic imported partial work.") if b.key == "activity_summary" else b for b in initial.blocks))
    source = make_docx(tmp_path)
    library.save_report_setup(profile, incoming, source, {"sha256": inspect_docx(source).sha256}, assets=(),
                              expected_revision=state.revision, actor="Synthetic Editor", confirmed=True, snapshot_revision=1)
    new = AppTest.from_file(ROOT / "run_web.py", default_timeout=20).run()
    new.segmented_control[0].set_value("Monthly report").run()
    choose_report(new, synthetic_profiles()[0].facilities[0].title)
    step(new, 2)
    assert next(w for w in new.text_area if w.label == "Activity summary").value == "Synthetic imported partial work."


def test_guided_placeholder_warning_download_and_review_gates(monkeypatch, tmp_path):
    monkeypatch.setattr(guided, "generate_report", lambda draft, **kwargs: ReportPackage(
        assemble_docx(draft, **kwargs), b"%PDF-synthetic", "synthetic.docx", "synthetic.pdf", draft.fingerprint))
    app = monthly(monkeypatch, tmp_path)
    next(w for w in app.text_input if w.label == "Prepared by").set_value("Synthetic Editor").run()
    step(app, 2)
    next(w for w in app.text_area if w.label == "Activity summary").set_value("Insert image here").run()
    step(app, 3)
    next(w for w in app.checkbox if w.label == "I checked the standing information for these sites").check().run()
    step(app, 4)
    assert next(b for b in app.button if b.label == "Generate DOCX and PDF").disabled
    assert any("template instructions" in e.value for e in app.error)
    assert any("This month’s work → Monthly Activity Summary" in e.value for e in app.error)
    assert not any("activity_summary" in e.value for e in app.error)
    step(app, 2)
    next(w for w in app.text_area if w.label == "Activity summary").set_value("August 2026 activity.").run()
    step(app, 4)
    assert next(b for b in app.button if b.label == "Generate DOCX and PDF").disabled
    next(w for w in app.checkbox if w.label == "I checked these specific warnings").check().run()
    next(b for b in app.button if b.label == "Generate DOCX and PDF").click().run()
    assert not app.exception
    assert [b.label for b in app.get("download_button")] == ["Download DOCX", "Download PDF"]
    step(app, 2)
    next(w for w in app.text_area if w.label == "Activity summary").set_value("July 2026 activity.").run()
    step(app, 4)
    assert not app.get("download_button")
    assert not next(w for w in app.checkbox if w.label == "I checked these specific warnings").value


def test_completion_messages_name_the_action_and_destination():
    from app.monthly_report_checks import ReportCheck
    from app.monthly_report_model import synthetic_draft
    draft = synthetic_draft(synthetic_profiles()[0], ReportPeriod(2026, 9))
    draft = replace(draft, sections=default_sections())
    text = guided.review_message(ReportCheck("required", "Resolve required block: org_chart.", True, "org_chart"), draft)
    assert text == "Add organizational chart in Site information → Organizational chart."
    text = guided.review_message(ReportCheck("required", "Resolve required block: work_orders.", True, "work_orders"), draft)
    assert "work_orders" not in text
    assert "work-order summary" in text and "This month’s work → Monthly Activity Summary" in text
    text = guided.review_message(ReportCheck("client_pages", "internal", True, "water_reports"), draft)
    assert "Water treatment reports" in text and "no prices" in text
    text = guided.review_message(ReportCheck("ai_number", "internal", True, "activity_summary"), draft)
    assert "unsupported numbers" in text and "wording and linked evidence" in text


def test_upload_setup_reviews_sections_and_preserves_partial_work(monkeypatch, tmp_path):
    from io import BytesIO
    from test_monthly_report_setup import make_docx
    from app import monthly_report_setup_ui as setup_ui
    from app.monthly_report_import import inspect_docx
    # AppTest 1.61 does not serialize tracked expander state yet. Preserve its
    # actual session value alongside the ordinary widget events it does support.
    from streamlit.testing.v1.element_tree import ElementTree
    original_states = ElementTree.get_widget_states
    def states(tree):
        values = original_states(tree)
        for key, identity in tree.session_state._state._key_id_mapper._key_id_mapping.items():
            if key.startswith("report_setup_") and key.endswith("_open"):
                values.widgets.add(id=identity, bool_value=bool(tree.session_state[key]))
        return values
    monkeypatch.setattr(ElementTree, "get_widget_states", states)
    path = make_docx(tmp_path)
    from docx import Document
    doc = Document(path)
    doc.sections[0].footer.paragraphs[0].text = "100 Example Way | example.invalid"
    doc.save(path)
    original_bytes = path.read_bytes()
    upload = BytesIO(original_bytes)
    upload.size, upload.name = len(original_bytes), "synthetic-mixed.docx"
    real_upload = setup_ui.st.file_uploader
    monkeypatch.setattr(setup_ui.st, "file_uploader", lambda label, *a, **kw:
                        upload if label == "Older or partially completed report" else real_upload(label, *a, **kw))
    app = monthly(monkeypatch, tmp_path / "data", saved=False)
    next(w for w in app.text_input if w.label == "Site name").set_value("Synthetic North; Synthetic South").run()
    next(w for w in app.text_input if w.label == "Name for this group (optional)").set_value("Synthetic Region").run()
    next(w for w in app.checkbox if w.label == "This is a regional report").check().run()
    next(w for w in app.text_input if w.label == "Other names for Synthetic North").set_value("Synthetic Legacy North").run()
    next(w for w in app.button if w.key and "report_start_card_upload" in w.key).click().run()
    next(b for b in app.button if b.label == "Analyze report").click().run()
    assert not app.exception
    assert not library.list_profiles(RRH_CONTRACT)
    assert not any(w.label == "What are you preparing?" for w in app.radio)
    next(w for w in app.text_input if w.label == "Your name").set_value("Synthetic Editor").run()
    assert next(b for b in app.button if b.label == "Continue to this month’s updates").disabled
    assert not any(w.label in ("Use this item", "Where it belongs", "Content to review") for w in app.selectbox)
    # A stale July cover must not discard September content or unread images.
    from app.monthly_report_sections import section_reviews
    inspection = inspect_docx(path)
    stage_key = next(k for k in app.session_state.filtered_state if k.startswith("report_setup_") and k.endswith("_stage"))
    prefix = stage_key.removesuffix("_stage") + "_" + inspection.sha256[:16]
    for section in section_reviews(inspection):
        app.session_state[prefix + "_section_" + section.key + "_open"] = True
        app.run()
        next(w for w in app.checkbox if w.key.startswith(prefix + "_section_" + section.key + "_ready_")).check().run()
        assert not app.exception
        assert app.session_state[prefix + "_section_plans"][section.key]["approved"], (section.key, [w.value for w in app.error])
        app.session_state[prefix + "_section_" + section.key + "_open"] = False
        app.run()
    # A closed/reopened section retains its content-bound confirmation.
    app.session_state[prefix + "_section_activity_open"] = True
    app.run()
    assert next(w for w in app.checkbox if w.key.startswith(prefix + "_section_activity_ready_")).value
    next(w for w in app.checkbox if w.label == "Save this report and its reusable design for these sites").check().run()
    assert not next(b for b in app.button if b.label == "Continue to this month’s updates").disabled, [w.value for w in (*app.info, *app.error)]
    next(b for b in app.button if b.label == "Continue to this month’s updates").click().run()
    assert not app.exception
    profile = library.list_profiles(RRH_CONTRACT)[0]
    assert profile.scope_type == "regional" and len(profile.facilities) == 2
    assert profile.facilities[0].aliases == ("Synthetic Legacy North",)
    assert library.imported_original(RRH_CONTRACT, profile.key).read_bytes() == original_bytes
    draft = library.load_imported_draft(RRH_CONTRACT, profile.key)
    assert draft.period == ReportPeriod(2026, 9)
    assert draft.address_line == "100 Example Way | example.invalid"
    assert "September 2026 repair" in next(b.text for b in draft.blocks if b.key == "activity_summary")
    assert next(b for b in draft.blocks if b.key == "vendor_reports").asset_hashes
    assert next(r for r in app.radio if r.label == "Report steps").value == guided.STEPS[1]


def test_section_choices_preserve_content_and_saved_status_across_new_sessions(monkeypatch, tmp_path):
    app = monthly(monkeypatch, tmp_path)
    next(w for w in app.text_input if w.label == "Prepared by").set_value("Synthetic Editor").run()
    box = next(w for w in app.checkbox if w.label == "Include Monthly Activity Summary")
    box.uncheck().run()
    next(b for b in app.button if b.label == "Save progress").click().run()
    assert any("Saved · version 1" in w.value for w in app.success)
    saved = library.load_snapshot(RRH_CONTRACT, "synthetic-guided", ReportPeriod(2026, 9))
    assert not next(s for s in saved.draft.sections if s.key == "activity").included
    assert next(b for b in saved.draft.blocks if b.key == "activity_summary").text == "Synthetic initial activity."
    new = AppTest.from_file(ROOT / "run_web.py", default_timeout=20).run()
    new.segmented_control[0].set_value("Monthly report").run()
    choose_report(new, synthetic_profiles()[0].facilities[0].title)
    assert not next(w for w in new.checkbox if w.label == "Include Monthly Activity Summary").value
    next(w for w in new.checkbox if w.label == "Include Monthly Activity Summary").check().run()
    assert any("changes to save" in w.value for w in new.info)
    next(b for b in new.button if b.label == "Continue to this month’s work").click().run()
    assert next(w for w in new.text_area if w.label == "Activity summary").value == "Synthetic initial activity."


def test_importing_partial_report_preserves_unsaved_work(monkeypatch, tmp_path):
    from io import BytesIO
    from docx import Document
    from app import monthly_report_setup_ui as setup_ui
    from streamlit.testing.v1.element_tree import ElementTree
    original_states = ElementTree.get_widget_states
    def states(tree):
        values = original_states(tree)
        for key, identity in list(tree.session_state._state._key_id_mapper._key_id_mapping.items()):
            if key.startswith("report_setup_") and key.endswith("_open"):
                values.widgets.add(id=identity, bool_value=bool(tree.session_state[key]))
        return values
    monkeypatch.setattr(ElementTree, "get_widget_states", states)
    doc = Document()
    doc.add_heading("2 Monthly Activity Summary", 1)
    doc.add_paragraph("Synthetic colleague completed the pump inspection.")
    upload = BytesIO()
    doc.save(upload)
    upload.size, upload.name = len(upload.getvalue()), "synthetic-partial.docx"
    real_upload = setup_ui.st.file_uploader
    monkeypatch.setattr(setup_ui.st, "file_uploader", lambda label, *a, **kw:
                        upload if label == "Older or partially completed report" else real_upload(label, *a, **kw))
    app = monthly(monkeypatch, tmp_path)
    next(w for w in app.text_input if w.label == "Prepared by").set_value("Synthetic Editor").run()
    step(app, 2)
    next(w for w in app.text_area if w.label == "Activity summary").set_value("Synthetic unsaved filter replacement.").run()
    step(app, 1)
    next(b for b in app.button if b.label == "Analyze report").click().run()
    next(w for w in app.checkbox if w.label.startswith("This section is ready")).check().run()
    next(w for w in app.checkbox if w.label == "Save this report and its reusable design for these sites").check().run()
    next(b for b in app.button if b.label == "Continue to this month’s updates").click().run()
    assert not app.exception
    value = next(w for w in app.text_area if w.label == "Activity summary").value
    assert "Synthetic unsaved filter replacement." in value
    assert "Synthetic colleague completed the pump inspection." in value
    imported = library.load_imported_draft(RRH_CONTRACT, "synthetic-guided")
    assert next(b for b in imported.blocks if b.key == "activity_summary").text == value


def test_shared_report_never_identifies_a_new_browser_as_its_previous_editor(monkeypatch, tmp_path):
    from app import web_ui
    device = ["synthetic-browser-first"]
    monkeypatch.setattr(web_ui, "device_token", lambda cookies: device[0])
    app = monthly(monkeypatch, tmp_path)
    next(w for w in app.text_input if w.label == "Prepared by").set_value("Synthetic First Editor").run()
    step(app, 2)
    next(w for w in app.text_area if w.label == "Activity summary").set_value("Saved work stays available to colleagues.").run()
    next(b for b in app.button if b.label == "Save progress").click().run()
    returning = AppTest.from_file(ROOT / "run_web.py", default_timeout=20).run()
    returning.segmented_control[0].set_value("Monthly report").run()
    assert not returning.exception
    assert next(w for w in returning.text_input if w.label == "Prepared by").value == "Synthetic First Editor"
    device[0] = "synthetic-browser-new"
    newcomer = AppTest.from_file(ROOT / "run_web.py", default_timeout=20).run()
    newcomer.segmented_control[0].set_value("Monthly report").run()
    choose_report(newcomer, synthetic_profiles()[0].facilities[0].title)
    assert next(w for w in newcomer.text_input if w.label == "Prepared by").value == ""
    step(newcomer, 2)
    assert next(w for w in newcomer.text_area if w.label == "Activity summary").value == "Saved work stays available to colleagues."
    assert next(b for b in newcomer.button if b.label == "Save progress").disabled
    device[0] = "synthetic-browser-first"
    returning.selectbox("report_month_number").set_value(10).run()
    assert next(w for w in returning.text_input if w.label == "Prepared by").value == "Synthetic First Editor"
    assert any("saved September 2026" in w.value for w in returning.info)
    step(returning, 2)
    assert next(w for w in returning.text_area if w.label == "Activity summary").value == ""


def test_new_site_reuses_contract_design_without_another_sites_activity(monkeypatch, tmp_path):
    app = monthly(monkeypatch, tmp_path)
    next(w for w in app.checkbox if w.label == synthetic_profiles()[0].facilities[0].title).uncheck().run()
    next(w for w in app.text_input if w.label == "Site name").set_value("Synthetic New Site").run()
    next(b for b in app.button if b.label == "Use a contract report").click().run()
    assert not any(w.label == "Report design to reuse" for w in app.radio)
    assert any("What you’ll need for this first report" in w.value for w in app.markdown)
    next(w for w in app.text_input if w.label == "Your name").set_value("Synthetic New Editor").run()
    next(w for w in app.checkbox if w.label == "Save this design for these sites so we can use it next month").check().run()
    next(b for b in app.button if b.label == "Start this report").click().run()
    assert not app.exception
    step(app, 2)
    assert next(w for w in app.text_area if w.label == "Activity summary").value == ""
    assert library.load_profile(RRH_CONTRACT, "synthetic-guided").block("activity_summary").text == "Synthetic initial activity."
