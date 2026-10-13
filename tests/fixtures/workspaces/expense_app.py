"""Actual expense UI with invented inputs and forbidden external operations."""
import sys
import json
import os
from io import BytesIO
from uuid import uuid4
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parents[3]))
from datetime import date
from hashlib import sha256
import fitz
import streamlit as st
import requests
from openpyxl import load_workbook
from app import expense_ui, expense_report, memory, receipt_jobs, reference_preview
from app.expense_report import build_expense_package as production_build_expense_package
from app.monthly_report_preview_view import render_view as original_render_view
from app.receipt_analyzer import ReceiptAnalysis

root = Path(os.environ.get('EPC_DATA_DIR', ''))
if (not root.is_absolute() or not root.name.startswith('expense-adaptive-synthetic-')
        or not root.is_dir() or not (root / '.synthetic-expense-fixture').is_file()):
    raise RuntimeError('This fixture requires its marked disposable expense test directory')
if memory._data_dir().resolve() != root.resolve():
    raise RuntimeError('Expense test storage was not isolated')
expense_report.PDF_BACKEND = 'libreoffice'
st.set_page_config(layout='wide')
# A hidden fixture-only acknowledgement lets browser checks wait for a genuine
# completed fragment render before comparing the full-app run counter.
def observed_render_view(**kwargs):
    key = kwargs['key']
    counts = st.session_state.setdefault('_fixture_preview_renders', {})
    counts[key] = counts.get(key, 0) + 1
    result = original_render_view(**kwargs)
    from html import escape
    value = st.session_state.get(key, {})
    request = (value.get('viewport') or {}) if isinstance(value, dict) else {}
    pages = ','.join(str(page) for page in request.get('pages', ()))
    st.html(f'<span hidden data-qa-preview-ack="{escape(key)}" data-count="{counts[key]}" data-request-pages="{escape(pages)}"></span>')
    return result
reference_preview.render_view = observed_render_view
def forbidden(*args,**kwargs):
    raise AssertionError('External action forbidden in synthetic expense fixture')
for name in ('start_receipt','prepare_receipt_content','analyze_prepared_receipt','_render_ios_mail_share','_IOS_MAIL_SHARE_COMPONENT'):
    setattr(expense_ui,name,forbidden)
receipt_jobs.start_receipt = forbidden
requests.sessions.Session.request = forbidden


def generated(details, items, *, mileage_items):
    # This wrapper observes the real workbook/template/Calc conversion. It does
    # not replace calculations, validation, receipt inclusion or output bytes.
    if memory._data_dir().resolve() != root.resolve():
        raise RuntimeError('Expense generation escaped its disposable storage')
    st.session_state['_fixture_generation_count'] = st.session_state.get('_fixture_generation_count', 0) + 1
    package = production_build_expense_package(details, items, mileage_items=mileage_items)
    if not package.pdf_bytes or package.pdf_error:
        raise RuntimeError('Actual expense PDF generation failed: ' + package.pdf_error)
    with fitz.open(stream=package.pdf_bytes, filetype='pdf') as document:
        text = '\n'.join(page.get_text() for page in document)
        pages = len(document)
    if 'REIMBURSEMENT OF EXPENSES' not in text or details.employee_name not in text:
        raise RuntimeError('The rendered expense packet lost its signed form')
    workbook = load_workbook(BytesIO(package.workbook_bytes), data_only=False)
    evidence = root / 'evidence'
    evidence.mkdir(exist_ok=True)
    digest = sha256(package.pdf_bytes).hexdigest()
    (evidence / (digest + '.pdf')).write_bytes(package.pdf_bytes)
    (evidence / (digest + '.xlsx')).write_bytes(package.workbook_bytes)
    for item in items:
        (evidence / ('original-' + sha256(item.file_bytes).hexdigest() + '.pdf')).write_bytes(item.file_bytes)
    record = {
        'synthetic_only': True, 'builder': 'app.expense_report.build_expense_package',
        'template_sha256': sha256(expense_report.EXPENSE_TEMPLATE_PATH.read_bytes()).hexdigest(),
        'output_sha256': digest, 'workbook_sha256': sha256(package.workbook_bytes).hexdigest(),
        'backend': expense_report.PDF_BACKEND, 'pages': pages,
        'total': str(package.total), 'receipt_count': package.receipt_count,
        'mileage_count': package.mileage_count, 'line_count': len(items),
        'receipt_hashes': list(dict.fromkeys(sha256(item.file_bytes).hexdigest() for item in items)),
        'workbook_sheets': workbook.sheetnames,
        'receipt_image_count': len(workbook['RECEIPTS']._images) if 'RECEIPTS' in workbook else 0,
        'editable_total_formula': workbook['EXPENSE REIMBURSEMENT']['Q60'].value,
    }
    (evidence / (digest + '.json')).write_text(json.dumps(record, indent=2))
    st.session_state['_fixture_generated_evidence'] = record
    return package


expense_ui.build_expense_package = generated
expense_ui.account_director=lambda account:('Synthetic Director','director@example.invalid')
expense_ui.director_config_problems=lambda accounts:[]
if not st.session_state.get('_fixture_initialized'):
    account='Rochester Regional Health'
    token=sha256(account.encode()).hexdigest()[:10]
    st.session_state['expense_draft_snapshot']={
        f'expense_employee_name_{token}':'Synthetic Employee',
        f'expense_employee_number_{token}':'00042',
        f'expense_report_date_{token}':date(2026,9,30),
        f'expense_approver_name_{token}':'Synthetic Director',
        f'expense_approver_email_{token}':'director@example.invalid',
        f'expense_manager_name_{token}':'Synthetic Manager',
        f'expense_manager_email_{token}':'manager@example.invalid',
    }
    files=[]
    if st.query_params.get('mileage_only')!='1':
        for index in (1,2):
            with fitz.open() as doc:
                for page in (1,2):
                    doc.new_page().insert_text((72,72),f'Synthetic receipt {index}; original page {page}')
                raw=doc.tobytes()
            digest=sha256(raw).hexdigest()
            files.append((f'synthetic-receipt-{index}.pdf',raw,'application/pdf'))
            st.session_state[f'expense_receipt_analysis_{digest}']=ReceiptAnalysis(
                merchant_name=f'Synthetic merchant {index}',transaction_date=date(2026,9,10),
                total_amount='24.50',currency='USD',suggested_description='Synthetic business supplies',confidence='high')
    st.session_state['expense_receipt_files']=files
    st.session_state['_fixture_browser_token']='synthetic-'+uuid4().hex
    st.session_state['_fixture_initialized']=True
st.session_state['_fixture_full_runs']=st.session_state.get('_fixture_full_runs',0)+1
st.caption(f"Synthetic full app runs: {st.session_state['_fixture_full_runs']}")
expense_ui.render_expense_workflow(st.session_state['_fixture_browser_token'],'America/New_York')
expense_ui.preserve_expense_draft_state()

st.caption('Synthetic expense generation count: ' + str(st.session_state.get('_fixture_generation_count', 0)))
