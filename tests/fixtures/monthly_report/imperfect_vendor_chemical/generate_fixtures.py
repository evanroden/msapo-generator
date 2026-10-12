"""Create intentionally imperfect, entirely fictional ENFRA monthly-report QA sources.

Run: python generate_fixtures.py
Requirements: python-docx, Pillow, PyMuPDF, LibreOffice CLI.
All output documents contain invented sites, vendors, people and assets. No authentic
client report or vendor artwork is used. Generated PDFs are real LibreOffice
conversions of the DOCX counterparts, except image-only scan QA PDF.
"""
from __future__ import annotations
from dataclasses import dataclass
from datetime import datetime, timezone
from hashlib import sha256
from pathlib import Path
import json, shutil, subprocess, zipfile

import fitz
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parent
FIXED_DATE = datetime(2026, 9, 30, tzinfo=timezone.utc)

def _font(size=20, bold=False):
    candidates = [
        f"/usr/share/fonts/truetype/dejavu/DejaVuSans{'-Bold' if bold else ''}.ttf",
        f"/usr/share/fonts/truetype/liberation2/LiberationSans{'-Bold' if bold else '-Regular'}.ttf",
    ]
    for p in candidates:
        if Path(p).exists():
            return ImageFont.truetype(p, size)
    return ImageFont.load_default()


def create_chart():
    image = Image.new("RGB", (960, 330), "white")
    d = ImageDraw.Draw(image)
    f = _font(22)
    small = _font(17)
    d.text((38, 19), "SYNTHETIC NEW TOWER - CHLORIDES (ppm)", fill="#174341", font=f)
    d.line((80, 260, 880, 260), fill="#859ba1", width=3)
    d.line((80, 90, 80, 260), fill="#859ba1", width=3)
    d.line((80, 145, 880, 145), fill="#ad543a", width=2)
    d.text((690, 113), "Mock upper limit: 300", fill="#ad543a", font=small)
    data = ((190, 170), (380, 165), (570, 118), (760, 185))
    d.line(data, fill="#255c88", width=5)
    for (x,y), day in zip(data, ("Sep 1", "Sep 7", "Sep 14", "Sep 21")):
        d.ellipse((x-7, y-7, x+7, y+7), fill="#255c88")
        d.text((x-32, 274), day, fill="#23343b", font=small)
    d.text((22, 165), "300", fill="#ad543a", font=small)
    d.text((35, 238), "0", fill="#23343b", font=small)
    out = ROOT / "_synthetic_tower_chart.png"
    image.save(out)
    return out


def create_scan_image():
    image = Image.new("RGB", (1100, 1420), "white")
    d = ImageDraw.Draw(image)
    d.rounded_rectangle((30, 50, 1070, 1370), radius=8, outline="#b5b5b5", width=3)
    d.text((86, 110), "IMAGE-ONLY - SYNTHETIC LAB SCAN", fill="#183c3b", font=_font(32, True))
    d.text((86, 215), "Facility: Demo North", fill="black", font=_font(25))
    d.text((86, 260), "Water treatment - cooling tower analysis", fill="black", font=_font(25))
    d.text((86, 305), "Sample: QA-SCAN-NORTH-09-2026", fill="black", font=_font(25))
    d.text((86, 350), "Collected: 2026-09-21", fill="black", font=_font(25))
    d.line((85, 430, 1000, 430), fill="#777", width=3)
    d.text((86, 480), "Parameter", fill="black", font=_font(24, True))
    d.text((510, 480), "Result", fill="black", font=_font(24, True))
    d.text((770, 480), "Range", fill="black", font=_font(24, True))
    for j,(name,val,ran) in enumerate((("Conductivity", "1,980 uS/cm", "1,300-2,000"),
                                        ("Chlorides", "340 ppm", "100-300"),
                                        ("pH", "8.1", "7.0-9.0"),
                                        ("Legionella", "Pending", "Not established"))):
        y=550+j*85
        d.line((85,y+45,1000,y+45),fill="#bbbbbb",width=1)
        d.text((86,y),name,fill="black",font=_font(23))
        d.text((510,y),val,fill="#a13622" if j==1 else "black",font=_font(22))
        d.text((770,y),ran,fill="black",font=_font(21))
    d.text((86, 975), "These data are fictional. Scan intentionally has no text layer.", fill="#555", font=_font(21))
    d.text((86, 1030), "Must request reviewed OCR/vision or source re-upload.", fill="#555", font=_font(21))
    d.text((86, 1150), "SYNTHETIC TEST ONLY - NOT FOR CLIENT USE", fill="#9d4038", font=_font(24,True))
    out = ROOT / "_synthetic_scanned_lab_page.png"
    image.save(out)
    return out


def setup_doc():
    doc = Document()
    sec=doc.sections[0]
    sec.page_height=Inches(11)
    sec.page_width=Inches(8.5)
    sec.top_margin=Inches(.72)
    sec.bottom_margin=Inches(.68)
    sec.left_margin=Inches(.72)
    sec.right_margin=Inches(.72)
    style=doc.styles["Normal"]
    style.font.name="Aptos"
    style.font.size=Pt(10)
    style.paragraph_format.space_after=Pt(5)
    style.paragraph_format.line_spacing=1.06
    doc.core_properties.author="Synthetic QA Fixture Generator"
    doc.core_properties.title="Fictional test record - not a real vendor report"
    doc.core_properties.subject="Regression only; intentionally flawed pages"
    doc.core_properties.created=FIXED_DATE.replace(tzinfo=None)
    doc.core_properties.modified=FIXED_DATE.replace(tzinfo=None)
    return doc


def title(doc, heading, sub):
    p=doc.add_paragraph()
    p.style=doc.styles["Normal"]
    p.alignment=WD_ALIGN_PARAGRAPH.LEFT
    r=p.add_run("FAKE DOCUMENT  |  SYNTHETIC QA ONLY")
    r.bold=True; r.font.size=Pt(10); r.font.color.rgb=RGBColor(155,50,41)
    p=doc.add_paragraph()
    r=p.add_run(heading)
    r.bold=True; r.font.size=Pt(19); r.font.color.rgb=RGBColor(22,72,69)
    p=doc.add_paragraph(sub)
    p.runs[0].italic=True
    p.runs[0].font.color.rgb=RGBColor(105,105,105)
    doc.add_paragraph()


def fact_table(doc, pairs):
    t=doc.add_table(rows=0,cols=2)
    t.style="Light Shading Accent 1"
    t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for k,v in pairs:
        c=t.add_row().cells
        c[0].text=str(k);c[1].text=str(v)
        c[0].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    doc.add_paragraph()


def readings_table(doc, heads, rows):
    t=doc.add_table(rows=1, cols=len(heads))
    t.style="Light Shading Accent 1"
    for i,h in enumerate(heads):
        t.rows[0].cells[i].text=h
    for row in rows:
        for i,c in enumerate(t.add_row().cells):
            c.text=str(row[i])
    doc.add_paragraph()


def bullet(doc,text):
    doc.add_paragraph(text,style="List Bullet")


def page(doc):
    doc.add_page_break()


def signature_image():
    im=Image.new("RGB",(650,120),"white")
    d=ImageDraw.Draw(im)
    d.line(((10,73),(36,32),(75,80),(115,18),(147,80),(184,28),(231,70),(280,25),(309,82),(367,38),(420,75),(470,43),(520,70),(600,47)),fill="#235d82",width=5)
    p=ROOT/"_synthetic_signature.png"
    im.save(p)
    return p


def vendor_mixed(chart,sign):
    d=setup_doc()
    title(d,"Vendor maintenance visit - Chiller inspection","Service report | fictional contractor: SampleWorks Mechanical")
    fact_table(d,[("Facility", "Demo North"),("Site", "Demo North"),
        ("Vendor", "SampleWorks Mechanical"),("Service Date", "2026-09-16"),
        ("Work Order", "QA-V-0916-041"),("Equipment Tag", "CH-QA-02"),
        ("Job Status", "Completed - inspection only")])
    d.add_heading("Technician observations",level=2)
    bullet(d,"Inspected Chiller CH-QA-02 while operating normally.")
    bullet(d,"Recorded evaporator approach 0.4 F and condenser approach 0.6 F.")
    bullet(d,"Recorded oil pressure differential 25.1 PSID during inspection.")
    bullet(d,"Collected an oil sample; laboratory results are still pending.")
    d.add_paragraph("No replacement or repair was approved as part of this service visit.")
    d.add_paragraph("A proposed follow-up has a separate price-bearing attachment; that proposal is not a completed repair.")
    # TRUE page 2 blank. No running header/footer across the document.
    page(d)
    page(d)
    title(d,"GENERAL TERMS AND CONDITIONS","Deliberately irrelevant contract language; exclude this page")
    d.add_heading("Sample service terms",level=2)
    for t in (
        "1. Governing law. This invented form uses the fictional jurisdiction of Demo County and creates no real legal obligations.",
        "2. Entire agreement. All prior discussions and representations are superseded solely for this synthetic exercise.",
        "3. Limitation of liability. The parties hypothetically disclaim indirect, incidental and consequential damages.",
        "4. Indemnification. Illustrative terms only: the parties would allocate risks of third-party claims and loss.",
        "5. Arbitration. A fictional arbitrator would resolve disputes in an invented forum under hypothetical rules.",
        "6. Severability. Unenforceable provisions would be removed without affecting the rest of this fictitious form.",
        "7. No technical service results, measurements or verified repairs are on this page.",
    ):
        d.add_paragraph(t)
    page(d)
    d.add_paragraph("SIGNATURE PAGE",style="Heading 1")
    d.add_paragraph("SampleWorks Mechanical - sample sign-off only")
    d.add_picture(str(sign),width=Inches(3))
    d.add_paragraph("Synthetic Technician | 2026-09-16")
    page(d)
    title(d,"Separate quotation - NOT COMPLETED WORK","Commercial information; intentionally excluded from client report")
    fact_table(d,[("Facility","Demo North"),("Quote Ref","QA-Q-0916-09"),
                  ("Possible task","Replace auxiliary chiller sensor"),
                  ("Quote Amount","$4,875.00"),("Approval","Pending"),
                  ("Labor Estimate","$1,350.00"),("Taxes","$390.00")])
    d.add_paragraph("SCOPE IS A PROPOSAL, NOT AN ACTION PERFORMED. Do not include this quoted page in customer monthly output.")
    return d


def water_mixed(chart,sign):
    d=setup_doc()
    title(d,"Water treatment - Cooling tower analysis","Technical chemistry sample | Example Aquatic Analytics")
    fact_table(d,[("Facility","Demo North"),("Vendor","Example Aquatic Analytics"),
                  ("Service Date","2026-09-14"),("Sample ID","QA-CT-0914-211"),
                  ("Equipment Tags","CT-QA-01, CT-QA-02"),
                  ("Analysis Status","Partial - Legionella culture pending")])
    d.add_heading("Tower measurements",level=2)
    readings_table(d,("Location","Metric","Measured","Site Limit","Assessment"),(
        ("Tower CT-QA-01", "Chlorides (ppm)","295","100-300","Within"),
        ("Tower CT-QA-02", "Chlorides (ppm)","340","100-300","Above range"),
        ("Tower CT-QA-02", "Conductivity (uS/cm)","1,980","1,300-2,000","Within"),
        ("Tower CT-QA-02", "pH", "8.1", "7.0-9.0", "Within"),
        ("Tower CT-QA-02", "Legionella culture", "Pending", "Not available", "UNKNOWN"),
    ))
    d.add_paragraph("Finding: Chlorides for CT-QA-02 were above the listed range. Sample collection is confirmed, but the culture result was not available as of report date.")
    d.add_paragraph("Action: Inspected feed pump operation and collected follow-up water samples. No chemical correction is claimed as complete.")
    page(d)
    title(d,"Technical annex - sampling trend","Keep meaningful chart and its technical legend; not a signature or legal page")
    d.add_picture(str(chart),width=Inches(6.9))
    d.add_paragraph("Figure 1: fictional measured chloride trend for CT-QA-02. Upper reference limit is 300 ppm; the 9/14 point is high.")
    d.add_paragraph("Suggested follow-up: Re-sample tower water and check bleed control. This is a recommendation, not verified completed work.")
    page(d)
    title(d,"STANDARD CONTRACT TERMS","Legal-only insert; exclude even though delivered with water measurements")
    for s in (
        "Governing law: fictional sample territory and no binding parties.",
        "Arbitration: hypothetical disputes would be resolved under a fictional set of rules.",
        "Indemnification: illustrative allocation of risks among non-existent parties.",
        "Limitation of liability: sample provision excluding certain indirect losses.",
        "Confidentiality: sample disclosure and record-keeping obligations.",
        "No chemical readings, analytical findings or technician actions appear on this page.",
    ):
        d.add_paragraph(s)
    page(d)
    d.add_paragraph("RECEIPT & SIGNATURE ONLY",style="Heading 1")
    d.add_picture(str(sign),width=Inches(3))
    d.add_paragraph("Example Aquatic Analytics | Fictional inspector")
    d.add_paragraph("This page does not establish chemical compliance.")
    page(d)
    d.add_paragraph("\u00a0")  # invisible trailing placeholder forces a real blank page
    return d


def vendor_ambiguous(_chart, _sign):
    d=setup_doc()
    title(d,"Vendor inspection request - status unclear","Maintenance work order with conflicting date fields")
    fact_table(d,[("Facility","Demo North"),("Vendor","Invented Drive Maintenance LLC"),
                  ("Service Date","2026-09-18"),("Completion Date","Invalid date"),
                  ("Work Order","QA-V-0918-088"),("Equipment Tag","VFD-QA-17"),
                  ("Job Status","Planned - follow-up not completed")])
    d.add_heading("Findings",level=2)
    bullet(d,"Observed intermittent overtemperature indication on VFD-QA-17.")
    bullet(d,"Diagnostic trend review recommended; no repair was performed during the planned visit.")
    bullet(d,"Replacement fan requested for evaluation but remains unapproved.")
    d.add_paragraph("IMPORTANT: Invalid completion date must not be silently replaced by service date or treated as proof of completion.")
    page(d)
    title(d,"Conditional follow-up notice","Technical follow-up, not a completion certification")
    d.add_paragraph("The next test is scheduled for 2026-10-03 (future relative to service date). The unit has not been declared repaired.")
    d.add_paragraph("Technician signature is absent. Work-order status stays planned pending confirmation.")
    return d


def chemical_cross_site(_chart,_sign):
    d=setup_doc()
    title(d,"Boiler water chemistry test - different site","Invented source from same synthetic QA contract; test wrong-site boundaries")
    fact_table(d,[("Facility","Demo South"),("Vendor","Fictional Boiler Chemistry Co."),
                  ("Service Date","2026-09-22"),("Report Date","2026-09-24"),
                  ("Sample ID","QA-SOUTH-BOIL-322"),("Equipment Tag","B-QA-03")])
    readings_table(d,("Test","Measured","Expected","Assessment"),(
        ("Conductivity (uS/cm)","3,650","3,000-4,000","Within"),
        ("P alkalinity (ppm)","245","250-350","Below range"),
        ("Molybdate (ppm)","12","10-20","Within"),
        ("Boiler status","Offline","No measured operation","Do not infer online"),
    ))
    d.add_paragraph("Action: Collected a boiler feed sample. The unit was offline and no functional restoration is claimed.")
    d.add_paragraph("The target for cross-site QA is Demo North. These Demo South values and figures must NOT migrate to Demo North.")
    page(d)
    title(d,"Cost and chemical order attachment","Non-technical purchasing page; price-bearing")
    fact_table(d,[("Facility","Demo South"),("Supply","Fictional chemical reagent, 5 gallons"),
                  ("Quote Total","$1,260.00"),("Purchase Order","Unapproved"),
                  ("Payment terms","Net 30"),("Status","Awaiting approval")])
    page(d)
    d.add_paragraph("\u00a0")  # invisible trailing placeholder forces a real blank page
    return d


SPECS = (
    ("QA_Vendor_Service_MixedPages",vendor_mixed,5,[1],[2,3,4,5],"vendor maintenance; blank/legal/signature/price"),
    ("QA_Water_Treatment_MixedPages",water_mixed,5,[1,2],[3,4,5],"chemical results, technical chart; legal/signature/blank"),
    ("QA_Vendor_AmbiguousDate",vendor_ambiguous,2,[1,2],[],"planned work with invalid completion date"),
    ("QA_Chemical_DifferentSite",chemical_cross_site,3,[],[1,2,3],"other site technical reading, priced page, blank"),
)


def render_pdf(input_docx):
    p=subprocess.run(["libreoffice","-env:UserInstallation=file:///tmp/enfra-fixture-lo-qa",
        "--headless","--convert-to","pdf","--outdir",str(ROOT),str(input_docx)],
        capture_output=True,text=True,timeout=45,check=True)
    out=input_docx.with_suffix(".pdf")
    if not out.exists():
        raise RuntimeError("DOCX PDF conversion did not create output: "+p.stdout+p.stderr)
    return out


def image_only_scan():
    src=create_scan_image()
    doc=fitz.open()
    sheet=doc.new_page(width=612,height=792)
    sheet.insert_image(fitz.Rect(20,12,592,780),filename=str(src))
    doc.new_page(width=612,height=792)  # genuinely blank second page
    dst=ROOT/"QA_Water_Treatment_ImageOnlyScan.pdf"
    doc.save(str(dst),garbage=4,deflate=True)
    doc.close()
    return dst


def inspection(path):
    if path.suffix==".pdf":
        doc=fitz.open(str(path))
        result=[{"page":i+1,"chars":len(pg.get_text().strip()),"images":len(pg.get_images(full=True))} for i,pg in enumerate(doc)]
        doc.close()
        return result
    return []


def main():
    chart=create_chart();sign=signature_image()
    manifest={"purpose":"Public synthetic QA fixtures, never private authentic evidence", 
              "generated_at":"2026-09-30T00:00:00Z fixture reference date (not an actual build timestamp)",
              "testing_target":{"contract":"QA Demonstration - Synthetic Only","site_for_in_scope":"Demo North","period":"2026-09"},
              "never_do":["Treat synthetic results as authentic parity acceptance","Copy price/legal/signature-only pages into a client report", "Turn pending/invalid date into completed status","Import Demo South readings into Demo North","Silently discard valid diagrams"],
              "files":[]}
    for name, builder, expected_pages, keep, omit, purpose in SPECS:
        d=builder(chart,sign)
        docx=ROOT/(name+".docx")
        d.save(str(docx))
        pdf=render_pdf(docx)
        pages=inspection(pdf)
        if len(pages)!=expected_pages:
            raise AssertionError(f"{name} has {len(pages)} pages, expected {expected_pages}: {pages}")
        for n in range(1,expected_pages+1):
            if n in set(omit) and n in (2,) and "Vendor_Service" in name:
                assert pages[n-1]["chars"]==0, "Vendor page 2 must be truly blank"
        for path in (docx,pdf):
            raw=path.read_bytes()
            manifest["files"].append({"name":path.name,"bytes":len(raw),"sha256":sha256(raw).hexdigest(),
               "scenario":name,"format":path.suffix[1:],"pages_in_rendered_pdf":expected_pages,
               "expected_relevant_pages_for_north":keep,"exclude_or_scope_block_pages":omit,
               "purpose":purpose,"pdf_page_inspection":pages if path.suffix==".pdf" else None})
    scan=image_only_scan()
    raw=scan.read_bytes()
    p=inspection(scan)
    assert len(p)==2 and all(x["chars"]==0 for x in p),"Scanned PDF must have zero extractable text"
    manifest["files"].append({"name":scan.name,"bytes":len(raw),"sha256":sha256(raw).hexdigest(),
                              "scenario":"image-only water report","format":"pdf", "pages_in_rendered_pdf":2,
                              "expected_relevant_pages_for_north":[],"exclude_or_scope_block_pages":[2],
                              "purpose":"First page needs reviewed OCR/vision or source re-upload; second page truly blank",
                              "pdf_page_inspection":p})
    (ROOT/"MANIFEST.json").write_text(json.dumps(manifest,indent=2)+"\n")
    (ROOT/"README.txt").write_text('''ENFRA - SYNTHETIC IMPERFECT DOCUMENT PACK\n========================================\nAll people, sites, companies and measurements are invented. None of these\nfiles are real vendors, actual chemical sampling results or client evidence.\n\nEach editable DOCX also has a native-text PDF companion converted with\nLibreOffice. The DOCX may expose document items differently from PDF pages,\nso expected PDF page numbers are a visual oracle, not necessarily the DOCX\nsource-item numbering used by the importer.\n\nVENDOR SERVICE MIXED PAGES - Demo North\n  page 1: real technical content (inspect; no replacement claimed)\n  page 2: entirely blank page\n  page 3: legal-only terms\n  page 4: signature-only\n  page 5: price-bearing quotation; work not performed\n\nWATER TREATMENT MIXED PAGES - Demo North\n  page 1: chemical measurements; chlorides over range, culture pending\n  page 2: chart and technical recommendation that should be retained\n  page 3: legal-only boilerplate\n  page 4: signature-only\n  page 5: blank\n\nVENDOR AMBIGUOUS DATE - Demo North\n  page 1: planned inspection; completion field says Invalid date\n  page 2: future follow-up; not a completed repair\n\nCHEMICAL DIFFERENT SITE - Demo South\n  page 1: measurements and sampled asset for another site\n  page 2: price-bearing chemical order; unapproved\n  page 3: blank\n\nWATER TREATMENT IMAGE-ONLY SCAN - Demo North (PDF ONLY)\n  page 1: visual technical data, no PDF text layer: must require review\n  page 2: entirely blank\n\nTEST EXPECTATIONS\n* Place vendor work findings under Monthly Activity Summary, not onsite calls.\n* Place water-treatment readings and relevant page images under Water Treatment.\n* Never include pricing, legal-only, blank or signature-only pages in output.\n* Never infer that a pending culture is negative or a scheduled repair complete.\n* Never reuse Demo South data in a Demo North report.\n* Never silently drop the real chart because another page was disallowed.\n* Track original hashes, decisions and page-to-section lineage.\n* Exclusions and deliberate removals must persist across refresh and replay.\n* Synthetic tests are not proof of authentic-client Word/PDF parity.\n\nUse MANIFEST.json to check exact hashes and page scenarios.\n''')
    z=ROOT.parent/"ENFRA_Synthetic_Imperfect_Vendor_Chemical_QA_Pack.zip"
    with zipfile.ZipFile(z,"w",zipfile.ZIP_DEFLATED,compresslevel=8) as out:
        for path in sorted(ROOT.iterdir()):
            if path.suffix in (".docx",".pdf") or path.name in ("README.txt","MANIFEST.json","generate_fixtures.py"):
                out.write(path,arcname=path.name)
    print(f"Created {z} ({z.stat().st_size:,} bytes)")
    for f in manifest['files']:
        print(f"{f['name']}: {f['bytes']:,} bytes, {f['pages_in_rendered_pdf']} PDF pages. {f['purpose']}")
        if f['pdf_page_inspection'] is not None:
            print("  text chars: ",[x['chars'] for x in f['pdf_page_inspection']])

if __name__ == "__main__":
    main()
