"""One source-cover task and first-submit durable cover application (AppTest).

Uploader bytes enter through the AppTest uploader seam; this is not native
browser input delivery or a production import.
"""

from dataclasses import replace
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.shared import Inches
from streamlit.testing.v1 import AppTest
import pytest

from app import monthly_report_guided, monthly_report_library as library, monthly_report_section_ui as section_ui
from app.monthly_report_model import ReportDraft
from test_cover_branding_contracts import photograph
import test_monthly_report_ui as report_ui_test
from test_monthly_report_ui import choose_report, monthly

ROOT = Path(monthly_report_guided.__file__).resolve().parents[1]


def current(app):
    return next(value for key, value in app.session_state.filtered_state.items()
                if key.startswith("report_guided_") and key.endswith("_draft") and isinstance(value, ReportDraft))


def stage_source(monkeypatch, tmp_path):
    # The same regression can be loaded alongside a read-only older checkout.
    # Bind its AppTest entrypoint to the actual imported product module.
    monkeypatch.setattr(report_ui_test, "ROOT", ROOT)
    app = monthly(monkeypatch, tmp_path / "data")
    next(w for w in app.text_input if w.label == "Prepared by").set_value("Synthetic Reviewer").run()
    draft = current(app)
    doc = Document()
    doc.add_paragraph(draft.profile.facilities[0].title)
    doc.add_paragraph(draft.period.label)
    doc.add_picture(BytesIO(photograph()), width=Inches(5))
    doc.add_section()
    doc.add_heading("Monthly Activity Summary", 1)
    doc.add_paragraph("Synthetic fan inspection completed.")
    source = tmp_path / "source.docx"
    doc.save(source)
    upload = BytesIO(source.read_bytes())
    upload.name, upload.size = source.name, len(upload.getvalue())
    incoming = {"Starting report": upload}
    real_upload = section_ui.st.file_uploader
    def uploader(label, *args, **kwargs):
        native = real_upload(label, *args, **kwargs)
        return incoming.get(label, native)
    monkeypatch.setattr(section_ui.st, "file_uploader", uploader)
    app.run()
    assert not app.exception
    assert any(key.endswith("_stage") for key in app.session_state.filtered_state), "The actual source importer must process the upload"
    return app, source, incoming


def test_staged_source_has_one_photo_task_and_no_conflicting_current_upload(monkeypatch, tmp_path):
    app, source, _ = stage_source(monkeypatch, tmp_path)
    assert sum(e.label == "Cover photograph" for e in app.expander) == 1
    assert not any(e.label == "Change cover photograph" for e in app.expander), "The pending source gallery must own the cover task"
    assert not any(w.label == "New photograph for this report" for w in app.get("file_uploader"))
    assert not any(c.value == "No cover photograph in this report (optional)." for c in app.caption)
    assert any(c.value == "Cover preview includes the selected photo. It is not saved until you continue." for c in app.caption)
    assert library.imported_original(current(app).profile.contract, current(app).profile.key) is None
    assert source.exists(), "Review must retain the original input"


def test_one_continue_commits_photo_retires_staging_and_reloads_durable_choice(monkeypatch, tmp_path):
    app, source, incoming = stage_source(monkeypatch, tmp_path)
    original = source.read_bytes()
    button = next(b for b in app.button if b.label == "Continue to this month’s updates")
    assert not button.disabled
    button.click().run()
    assert not app.exception
    # Applying a source photo is not its separate client/privacy approval.
    # The imported photo must be durable on this first action while the real
    # unreviewed-image generation guard remains in force.
    assert all("Review the new or changed pictures in Cover photo" in e.value for e in app.error)
    draft = current(app)
    photo = next(b for b in draft.blocks if b.key == "cover_photo")
    assert len(photo.asset_hashes) == 1
    assert not photo.client_reviewed_fingerprint
    saved = library.load_imported_draft(draft.profile.contract, draft.profile.key)
    assert next(b for b in saved.blocks if b.key == "cover_photo") == photo
    assert library.imported_original(draft.profile.contract, draft.profile.key).read_bytes() == original
    assert not any(key.endswith("_stage") for key in app.session_state.filtered_state), "Successful import must not leave a second unsaved cover preview"
    assert sum(e.label == "Change cover photograph" for e in app.expander) == 1
    assert not any(e.label == "Cover photograph" for e in app.expander)
    assert not any("not saved until you continue" in c.value for c in app.caption)
    for _ in range(2):
        app.run()
        assert not app.exception
        assert next(b for b in current(app).blocks if b.key == "cover_photo") == photo
        assert not any(key.endswith("_stage") for key in app.session_state.filtered_state)
    incoming.clear()
    fresh = AppTest.from_file(ROOT / "run_web.py", default_timeout=30).run()
    fresh.segmented_control[0].set_value("Monthly report").run()
    choose_report(fresh, draft.profile.facilities[0].title)
    assert not fresh.exception
    assert next(b for b in current(fresh).blocks if b.key == "cover_photo") == photo
    assert source.read_bytes() == original


def test_failed_import_keeps_source_choice_and_retry_then_explicit_reopen(monkeypatch, tmp_path):
    app, source, incoming = stage_source(monkeypatch, tmp_path)
    stage_key = next(key for key in app.session_state.filtered_state if key.endswith("_stage"))
    staged_path = app.session_state[stage_key][1]
    before = source.read_bytes()
    real_save = library.save_report_setup
    def failing_save(*args, **kwargs):
        raise OSError("Synthetic controlled save failure")
    monkeypatch.setattr(library, "save_report_setup", failing_save)
    next(b for b in app.button if b.label == "Continue to this month’s updates").click().run()
    assert not app.exception
    assert any(e.value == "Synthetic controlled save failure" for e in app.error)
    assert app.session_state[stage_key][1] == staged_path and staged_path.exists()
    assert library.imported_original(current(app).profile.contract, current(app).profile.key) is None
    assert not any("_committed_upload_digest" in key for key in app.session_state.filtered_state)
    monkeypatch.setattr(library, "save_report_setup", real_save)
    next(b for b in app.button if b.label == "Continue to this month’s updates").click().run()
    assert not app.exception and stage_key not in app.session_state.filtered_state
    assert not staged_path.exists()
    photo = next(b for b in current(app).blocks if b.key == "cover_photo")
    incoming.clear()
    app.run()
    assert not any("_committed_upload_digest" in key for key in app.session_state.filtered_state)
    next(b for b in app.button if b.label == "Open the saved original for section review").click().run()
    assert not app.exception and stage_key in app.session_state.filtered_state
    assert next(b for b in current(app).blocks if b.key == "cover_photo") == photo
    assert source.read_bytes() == before


def test_pending_photo_uses_only_current_render_source_and_report_scope(monkeypatch, tmp_path):
    from app.monthly_report_editor import _signature
    app, _, _ = stage_source(monkeypatch, tmp_path)
    draft = current(app)
    state = library.load_profile(draft.profile.contract, draft.profile.key)
    values = dict(app.session_state.filtered_state)
    prefix = 'report_setup_' + _signature((draft.profile.contract, draft.profile.key, draft.period.key))
    active = values[prefix + '_active_section_review']
    assert active[1] + '_section_plans' in values
    monkeypatch.setattr(section_ui.st, 'session_state', values)
    from app import monthly_report_import
    monkeypatch.setattr(monthly_report_import, 'read_import_image',
                        lambda *a, **k: pytest.fail('Inactive source photo must not be read'))
    for other in (replace(draft, period=replace(draft.period, month=draft.period.month % 12 + 1)),
                  replace(draft, profile=replace(draft.profile, contract='Another synthetic contract'))):
        assert section_ui.render_pending_cover_selection(other, {}, other.profile.contract, other.period, None) == (other, {})
    values[prefix + '_active_section_review'] = ('0' * 64, active[1])
    assert section_ui.render_pending_cover_selection(draft, {}, draft.profile.contract, draft.period, state) == (draft, {})
    values.pop(prefix + '_active_section_review')
    assert section_ui.render_pending_cover_selection(draft, {}, draft.profile.contract, draft.period, state) == (draft, {})


def test_import_admission_clears_pending_preview_scope_before_busy_return(monkeypatch):
    from app.monthly_report_editor import _signature
    from app.monthly_report_model import synthetic_profiles, ReportPeriod
    from app import monthly_report_render_jobs as jobs, monthly_report_busy_ui
    profile, period = synthetic_profiles()[0], ReportPeriod(2026, 9)
    key = 'report_setup_' + _signature((profile.contract, profile.key, period.key)) + '_active_section_review'
    values = {key: ('old source', 'old scope')}
    monkeypatch.setattr(section_ui.st, 'session_state', values)
    def busy(*a, **k):
        raise jobs.RenderBusy('Synthetic capacity wait')
    monkeypatch.setattr(jobs, 'work_slot', busy)
    monkeypatch.setattr(monthly_report_busy_ui, 'show_busy_retry', lambda: None)
    @section_ui._source_view_admission
    def render(contract, period, *, identity=None, source_view=None):
        pytest.fail('Busy admission must not render an old scope')
    assert render(profile.contract, period, identity=profile, source_view=object()) is True
    assert key not in values
