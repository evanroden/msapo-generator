"""Synthetic native layouts only; private source reports never enter fixtures."""
from io import BytesIO
from zipfile import ZipFile, ZIP_DEFLATED

from docx import Document
from lxml import etree
from PIL import Image
import pytest

from app.monthly_report_import import ImportError
from app.monthly_report_native_package import passive_docx

R = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
REL = 'http://schemas.openxmlformats.org/package/2006/relationships'


def sample(tmp_path):
    document = Document()
    document.sections[0].header.paragraphs[0].text = 'Synthetic header'
    document.sections[0].footer.paragraphs[0].text = 'Synthetic footer'
    paragraph = document.add_paragraph('Native body')
    paragraph.runs[0].bold = True
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = 'Capacity'
    table.cell(1, 0).text = 'Steam'
    image = BytesIO()
    Image.new('RGB', (8, 8), 'green').save(image, format='PNG')
    document.add_picture(BytesIO(image.getvalue()))
    path = tmp_path / 'synthetic.docx'
    document.save(path)
    return path


def rewrite(path, mutate):
    with ZipFile(path) as source:
        parts = {n: source.read(n) for n in source.namelist()}
    mutate(parts)
    with ZipFile(path, 'w', ZIP_DEFLATED) as output:
        for name, raw in parts.items():
            output.writestr(name, raw)


def body_change(parts, change):
    root = etree.fromstring(parts['word/document.xml'])
    change(root.find('{' + W + '}body'))
    parts['word/document.xml'] = etree.tostring(root)


def add_relation(parts, rid, kind, target, external=False):
    root = etree.fromstring(parts['word/_rels/document.xml.rels'])
    attrs = dict(Id=rid, Type=R + '/' + kind, Target=target)
    if external:
        attrs['TargetMode'] = 'External'
    etree.SubElement(root, '{' + REL + '}Relationship', **attrs)
    parts['word/_rels/document.xml.rels'] = etree.tostring(root)


def test_preserves_page_chrome_styles_geometry_and_original_raster(tmp_path):
    path = sample(tmp_path)
    raw = path.read_bytes()
    result = passive_docx(path)
    assert path.read_bytes() == raw
    with ZipFile(BytesIO(raw)) as old, ZipFile(BytesIO(result)) as new:
        for name in ('word/document.xml', 'word/styles.xml', 'word/numbering.xml', 'word/fontTable.xml', 'word/webSettings.xml', 'word/theme/theme1.xml', 'word/header1.xml', 'word/footer1.xml'):
            # Structural comparison avoids serializer declaration differences.
            assert etree.tostring(etree.fromstring(new.read(name)), method='c14n') == etree.tostring(etree.fromstring(old.read(name)), method='c14n')
        assert new.read('word/media/image1.png') == old.read('word/media/image1.png')
        assert not any(n.startswith('customXml/') for n in new.namelist())
    generated = Document(BytesIO(result))
    assert generated.sections[0].header.paragraphs[0].text == 'Synthetic header'
    assert generated.tables[0].cell(1, 0).text == 'Steam'


def test_external_hyperlink_retains_visible_text_without_network_relation(tmp_path):
    path = sample(tmp_path)
    def mutate(parts):
        add_relation(parts, 'rLink', 'hyperlink', 'https://invalid.example/', True)
        def append(body):
            p = etree.SubElement(body, '{' + W + '}p')
            h = etree.SubElement(p, '{' + W + '}hyperlink', {'{' + R + '}id': 'rLink'})
            etree.SubElement(etree.SubElement(h, '{' + W + '}r'), '{' + W + '}t').text = 'Visible hyperlink label'
        body_change(parts, append)
    rewrite(path, mutate)
    with ZipFile(BytesIO(passive_docx(path))) as output:
        assert b'Visible hyperlink label' in output.read('word/document.xml')
        assert b'https://invalid.example' not in output.read('word/_rels/document.xml.rels')
        assert b'rLink' not in output.read('word/document.xml')


@pytest.mark.parametrize('name', ['object', 'altChunk', 'control'])
def test_visible_active_content_fails_explicitly(tmp_path, name):
    path = sample(tmp_path)
    rewrite(path, lambda parts: body_change(parts, lambda body: etree.SubElement(body, '{' + W + '}' + name)))
    with pytest.raises(ImportError, match='unsupported visible'):
        passive_docx(path)


def test_missing_or_external_visible_image_fails(tmp_path):
    path = sample(tmp_path)
    def mutate(parts):
        root = etree.fromstring(parts['word/_rels/document.xml.rels'])
        for rel in root:
            if rel.get('Type').endswith('/image'):
                rel.set('Target', 'https://invalid.example/image.png')
                rel.set('TargetMode', 'External')
        parts['word/_rels/document.xml.rels'] = etree.tostring(root)
    rewrite(path, mutate)
    with pytest.raises(ImportError, match='external image'):
        passive_docx(path)


def test_fields_keep_page_number_and_flatten_unsafe_cached_value(tmp_path):
    path = sample(tmp_path)
    def append(body):
        for instruction, text in [('PAGE', '2'), ('INCLUDETEXT "https://invalid.example"', 'Cached text')]:
            p = etree.SubElement(body, '{' + W + '}p')
            for kind, value in [('fldChar', 'begin'), ('instrText', instruction), ('fldChar', 'separate'), ('t', text), ('fldChar', 'end')]:
                node = etree.SubElement(etree.SubElement(p, '{' + W + '}r'), '{' + W + '}' + kind)
                if kind == 'fldChar':
                    node.set('{' + W + '}fldCharType', value)
                else:
                    node.text = value
    rewrite(path, lambda parts: body_change(parts, append))
    with ZipFile(BytesIO(passive_docx(path))) as output:
        root = etree.fromstring(output.read('word/document.xml'))
        assert [n.text for n in root.iter('{' + W + '}instrText')] == ['PAGE']
        assert 'Cached text' in ''.join(root.itertext())


def test_compatibility_choice_and_fallback_remain_wrapped(tmp_path):
    path = sample(tmp_path)
    mc = 'http://schemas.openxmlformats.org/markup-compatibility/2006'
    def append(body):
        parent = etree.SubElement(body, '{' + mc + '}AlternateContent', nsmap={'mc': mc, 'w': W})
        for branch in ('Choice', 'Fallback'):
            node = etree.SubElement(parent, '{' + mc + '}' + branch)
            if branch == 'Choice':
                node.set('Requires', 'w')
            etree.SubElement(etree.SubElement(etree.SubElement(node, '{' + W + '}p'), '{' + W + '}r'), '{' + W + '}t').text = branch
    rewrite(path, lambda parts: body_change(parts, append))
    with ZipFile(BytesIO(passive_docx(path))) as output:
        root = etree.fromstring(output.read('word/document.xml'))
        wrapper = root.find('.//{' + mc + '}AlternateContent')
        assert wrapper is not None and len(wrapper) == 2


def test_unreferenced_macro_and_external_template_never_copied(tmp_path):
    path = sample(tmp_path)
    def mutate(parts):
        parts['word/vbaProject.bin'] = b'synthetic unused macro'
        add_relation(parts, 'macro', 'vbaProject', 'vbaProject.bin')
        add_relation(parts, 'template', 'attachedTemplate', 'https://invalid.example/template', True)
    rewrite(path, mutate)
    with ZipFile(BytesIO(passive_docx(path))) as output:
        assert 'word/vbaProject.bin' not in output.namelist()
        assert b'invalid.example' not in output.read('word/_rels/document.xml.rels')


def test_static_emf_bounds_and_driver_escape_rejection():
    import struct
    from app.monthly_report_native_package import _static_emf
    def metafile(extra=b''):
        header = bytearray(88)
        struct.pack_into('<II', header, 0, 1, 88)
        header[40:44] = b' EMF'
        eof = struct.pack('<IIIII', 14, 20, 0, 0, 20)
        struct.pack_into('<II', header, 48, len(header) + len(extra) + len(eof), 2 + bool(extra))
        return bytes(header) + extra + eof
    assert _static_emf(metafile())
    assert not _static_emf(metafile(struct.pack('<II', 105, 8)))
    assert not _static_emf(metafile()[:-1])
    assert not _static_emf(metafile(struct.pack('<II', 2, 0)))


def test_complexity_and_dtd_are_rejected(tmp_path):
    path = sample(tmp_path)
    rewrite(path, lambda parts: parts.update({'word/document.xml': b'<!DOCTYPE x [<!ENTITY x "unsafe">]><x>&x;</x>'}))
    with pytest.raises(ImportError, match='unsafe XML'):
        passive_docx(path)


def test_hdphoto_edit_history_uses_existing_rendered_raster(tmp_path):
    path = sample(tmp_path)
    def mutate(parts):
        add_relation(parts, 'rEditHistory', 'hdphoto', 'media/original.wdp')
        parts['word/media/original.wdp'] = b'unsupported editing history'
        image = BytesIO()
        Image.new('RGB', (300, 300), 'green').save(image, format='PNG')
        parts['word/media/image1.png'] = image.getvalue()
        root = etree.fromstring(parts['word/document.xml'])
        drawing = 'http://schemas.openxmlformats.org/drawingml/2006/main'
        editing = 'http://schemas.microsoft.com/office/drawing/2010/main'
        blip = root.find('.//{' + drawing + '}blip')
        extensions = etree.SubElement(blip, '{' + drawing + '}extLst')
        extension = etree.SubElement(extensions, '{' + drawing + '}ext', uri='synthetic-image-editing')
        props = etree.SubElement(extension, '{' + editing + '}imgProps')
        etree.SubElement(props, '{' + editing + '}imgLayer', {'{' + R + '}embed': 'rEditHistory'})
        parts['word/document.xml'] = etree.tostring(root)
    rewrite(path, mutate)
    with ZipFile(path) as source, ZipFile(BytesIO(passive_docx(path))) as output:
        assert output.read('word/media/image1.png') == source.read('word/media/image1.png')
        assert 'word/media/original.wdp' not in output.namelist()
        assert b'rEditHistory' not in output.read('word/_rels/document.xml.rels')


@pytest.mark.parametrize('keep_reference', [True, False])
def test_orphan_notes_are_not_hidden_in_output(tmp_path, keep_reference):
    path = sample(tmp_path)
    def mutate(parts):
        add_relation(parts, 'rFootnotes', 'footnotes', 'footnotes.xml')
        notes = etree.Element('{' + W + '}footnotes', nsmap={'w': W})
        for identity, text in [('1', 'Referenced note'), ('2', 'Unreferenced old note')]:
            note = etree.SubElement(notes, '{' + W + '}footnote', {'{' + W + '}id': identity})
            etree.SubElement(etree.SubElement(etree.SubElement(note, '{' + W + '}p'), '{' + W + '}r'), '{' + W + '}t').text = text
        parts['word/footnotes.xml'] = etree.tostring(notes)
        types = etree.fromstring(parts['[Content_Types].xml'])
        etree.SubElement(types, '{http://schemas.openxmlformats.org/package/2006/content-types}Override', PartName='/word/footnotes.xml', ContentType='application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml')
        parts['[Content_Types].xml'] = etree.tostring(types)
        if keep_reference:
            body_change(parts, lambda body: etree.SubElement(etree.SubElement(etree.SubElement(body, '{' + W + '}p'), '{' + W + '}r'), '{' + W + '}footnoteReference', {'{' + W + '}id': '1'}))
    rewrite(path, mutate)
    with ZipFile(BytesIO(passive_docx(path))) as output:
        assert ('word/footnotes.xml' in output.namelist()) == keep_reference
        if keep_reference:
            assert b'Referenced note' in output.read('word/footnotes.xml')
            assert b'Unreferenced old note' not in output.read('word/footnotes.xml')


def test_native_raster_image_layers_are_visible_content_not_edit_history(tmp_path):
    path = sample(tmp_path)
    def mutate(parts):
        drawing = 'http://schemas.openxmlformats.org/drawingml/2006/main'
        editing = 'http://schemas.microsoft.com/office/drawing/2010/main'
        root = etree.fromstring(parts['word/document.xml'])
        blip = root.find('.//{' + drawing + '}blip')
        props = etree.SubElement(etree.SubElement(etree.SubElement(blip, '{' + drawing + '}extLst'), '{' + drawing + '}ext', uri='synthetic-layer'), '{' + editing + '}imgProps')
        etree.SubElement(props, '{' + editing + '}imgLayer', {'{' + R + '}embed': blip.get('{' + R + '}embed')})
        parts['word/document.xml'] = etree.tostring(root)
    rewrite(path, mutate)
    with ZipFile(BytesIO(passive_docx(path))) as output:
        assert b'imgLayer' in output.read('word/document.xml')


def test_smartart_drawing_relationship_can_be_document_scoped(tmp_path):
    path = sample(tmp_path)
    def mutate(parts):
        add_relation(parts, 'rData', 'diagramData', 'diagrams/data.xml')
        add_relation(parts, 'rDrawing', 'diagramDrawing', 'diagrams/drawing.xml')
        body_change(parts, lambda body: etree.SubElement(body, '{http://schemas.openxmlformats.org/drawingml/2006/diagram}relIds', {'{' + R + '}dm': 'rData'}))
        parts['word/diagrams/data.xml'] = b'<d:data xmlns:d="http://schemas.openxmlformats.org/drawingml/2006/diagram"><d:extension relId="rDrawing"/><d:layout relId=""/></d:data>'
        parts['word/diagrams/drawing.xml'] = b'<d:drawing xmlns:d="http://schemas.openxmlformats.org/drawingml/2006/diagram"/>'
        types = etree.fromstring(parts['[Content_Types].xml'])
        for name, kind in [('data', 'diagramData'), ('drawing', 'diagramDrawing')]:
            etree.SubElement(types, '{http://schemas.openxmlformats.org/package/2006/content-types}Override', PartName=f'/word/diagrams/{name}.xml', ContentType='application/vnd.openxmlformats-officedocument.drawingml.' + kind + '+xml')
        parts['[Content_Types].xml'] = etree.tostring(types)
    rewrite(path, mutate)
    with ZipFile(BytesIO(passive_docx(path))) as output:
        assert 'word/diagrams/drawing.xml' in output.namelist()
        assert b'rDrawing' in output.read('word/_rels/document.xml.rels')


def test_preserves_only_nonempty_original_package_directories(tmp_path):
    path = sample(tmp_path)
    def mutate(parts):
        parts['word/'] = b''
        parts['word/media/'] = b''
        parts['customXml/'] = b''
        parts['unused/'] = b''
    rewrite(path, mutate)
    with ZipFile(BytesIO(passive_docx(path))) as output:
        assert 'word/' in output.namelist()
        assert 'word/media/' in output.namelist()
        assert 'customXml/' not in output.namelist()
        assert 'unused/' not in output.namelist()


def test_missing_referenced_note_cannot_disappear_silently(tmp_path):
    path = sample(tmp_path)
    rewrite(path, lambda parts: body_change(parts, lambda body: etree.SubElement(etree.SubElement(etree.SubElement(body, '{' + W + '}p'), '{' + W + '}r'), '{' + W + '}footnoteReference', {'{' + W + '}id': '99'})))
    with pytest.raises(ImportError, match='missing note part'):
        passive_docx(path)


def test_drawing_roundtrip_cache_is_removed_without_decoding(tmp_path, monkeypatch):
    import base64
    path = sample(tmp_path)
    office = 'urn:schemas-microsoft-com:office:office'
    vml = 'urn:schemas-microsoft-com:vml'
    def append(body):
        pict = etree.SubElement(etree.SubElement(etree.SubElement(body, '{' + W + '}p'), '{' + W + '}r'), '{' + W + '}pict')
        etree.SubElement(pict, '{' + vml + '}shape', {'{' + office + '}gfxdata': 'intentionally-invalid-hidden-old-image-cache', 'id': 'live-vector', 'fillcolor': '#557F7F', 'style': 'width:100pt;height:20pt'})
    rewrite(path, lambda parts: body_change(parts, append))
    def forbidden(*args, **kwargs):
        raise AssertionError('Drawing caches must never be decoded')
    monkeypatch.setattr(base64, 'b64decode', forbidden)
    with ZipFile(BytesIO(passive_docx(path))) as output:
        raw = output.read('word/document.xml')
        assert b'gfxdata' not in raw and b'hidden-old-image-cache' not in raw
        assert b'live-vector' in raw and b'#557F7F' in raw


def test_extended_properties_keep_only_application_compatibility_identity(tmp_path):
    path = sample(tmp_path)
    namespace = 'http://schemas.openxmlformats.org/officeDocument/2006/extended-properties'
    def mutate(parts):
        root = etree.fromstring(parts['docProps/app.xml'])
        for name in ('TitlesOfParts', 'HeadingPairs', 'Company', 'Manager', 'HyperlinkBase'):
            for existing in list(root.findall('{' + namespace + '}' + name)):
                root.remove(existing)
            etree.SubElement(root, '{' + namespace + '}' + name).text = 'old-private-contact@example.invalid'
        app = root.find('{' + namespace + '}Application')
        app.set('sourceFact', 'old-private-contact@example.invalid')
        etree.SubElement(app, 'hidden').text = 'old-private-contact@example.invalid'
        parts['docProps/app.xml'] = etree.tostring(root)
    rewrite(path, mutate)
    with ZipFile(path) as original:
        before = etree.fromstring(original.read('docProps/app.xml'))
    with ZipFile(BytesIO(passive_docx(path))) as output:
        root = etree.fromstring(output.read('docProps/app.xml'))
        assert {etree.QName(child).localname for child in root} == {'Application', 'AppVersion'}
        for name in ('Application', 'AppVersion'):
            assert root.findtext('{' + namespace + '}' + name) == before.findtext('{' + namespace + '}' + name)
        assert all(b'old-private-contact' not in output.read(name) for name in output.namelist() if name.endswith('.xml'))
