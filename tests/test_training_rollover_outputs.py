"""Synthetic saved Training rollover through the actual Word/PDF generators.

The native fixture is invented. These checks do not certify a genuine design,
Microsoft Word, browser download delivery or a deployed application.
"""
from dataclasses import replace
from datetime import date
from hashlib import sha256
from io import BytesIO
import json
import os
from pathlib import Path
import re
from zipfile import ZipFile

from docx import Document
import fitz
import pytest

from app import monthly_report_designs as designs, monthly_report_library as library
from app.monthly_report_docx import generate_report
from app.monthly_report_model import (
    BlockSpec, Facility, ReportDraft, ReportPeriod, ReportProfile, ReportTable,
    ResolvedBlock, SectionSpec,
)
from app.monthly_report_setup import new_month_draft
from app.monthly_report_training import MATRIX_REF, add_event, apply_tables
from tests.conftest import requires_libreoffice


MATRIX = ReportTable(
    ("Team member", "Site", "Fire safety", "Operations", "Orientation"),
    (("Synthetic Trainee", "Synthetic North", "Completed", "Pending", ""),
     ("Synthetic Colleague", "Synthetic North", "Not recorded", "Not required", "")),
    MATRIX_REF,
)


def _save_and_reload(draft, readbacks=None):
    saved = library.save_snapshot(draft, expected_revision=0, entered_editor=draft.prepared_by)
    loaded = library.load_snapshot(draft.profile.contract, draft.profile.key, draft.period)
    if readbacks is None:
        assert loaded == saved
    else:
        readbacks.append((saved, loaded))
    return loaded.draft


def _initial_report(tmp_path, native):
    profile = ReportProfile("Synthetic Training", "north", "Synthetic North",
                            (Facility("north", "Synthetic North"),))
    if native:
        source = tmp_path / "synthetic-training-design.docx"
        document = Document()
        document.add_paragraph("Operations and Maintenance Monthly Review")
        document.add_paragraph("Synthetic North")
        document.add_paragraph("September 2026")
        document.add_heading("TRAINING SUMMARY", 1)
        document.add_paragraph("Synthetic source-only training note")
        table = document.add_table(rows=2, cols=len(MATRIX.columns))
        table.style = "Light Shading Accent 1"
        for cell, value in zip(table.rows[0].cells, MATRIX.columns):
            cell.text = value
        for cell in table.rows[1].cells:
            cell.text = "Synthetic source-only value"
        document.save(source)
        profile = designs.install(source, profile, actor="Synthetic Editor", expected_revision=0)
    period = ReportPeriod(2026, 9)
    event = add_event(None, title="Prior safety briefing", when=date(2026, 9, 4),
                      mode="Virtual", participants=["Synthetic Trainee"], hours="0.75", period=period)
    block = apply_tables(ResolvedBlock("training_summary", "This month"), MATRIX, event)
    section = SectionSpec("training", "11", "Training Summary", (BlockSpec("training_summary", "rich_text"),))
    return ReportDraft(profile, period, "Synthetic Editor", (section,), (block,))


@requires_libreoffice
@pytest.mark.parametrize("native", [False, True], ids=["shell", "native"])
def test_saved_rollovers_keep_standing_matrix_in_actual_word_and_pdf(tmp_path, monkeypatch, native):
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path / "runtime"))
    initial = _save_and_reload(_initial_report(tmp_path, native))
    original = library.load_snapshot_version(initial.profile.contract, initial.profile.key, initial.period, 1)
    output = Path(os.environ.get("INDEPENDENT_OUTPUT", str(tmp_path / "evidence"))) / "training-rollover" / ("native" if native else "shell")
    output.mkdir(parents=True, exist_ok=True)
    reports, readbacks = [], []
    current = initial
    manifest = {"synthetic_only": True, "native_design": native, "files": {}}
    for month in (10, 11):
        current = _save_and_reload(new_month_draft(current, ReportPeriod(2026, month)), readbacks)
        package = generate_report(current, acknowledged_fingerprint=current.fingerprint)
        assert package.docx and package.pdf, package.pdf_error
        # Retain real failing-before output before making desired assertions.
        for suffix, raw in (("docx", package.docx), ("pdf", package.pdf)):
            name = f"SYNTHETIC-Training-{current.period.key}.{suffix}"
            (output / name).write_bytes(raw)
            manifest["files"][name] = {"bytes": len(raw), "sha256": sha256(raw).hexdigest()}
        reports.append((current, package))
    (output / "manifest.json").write_text(json.dumps(manifest, indent=2) + "\n")
    assert library.load_snapshot_version(initial.profile.contract, initial.profile.key, initial.period, 1) == original
    for current, package in reports:
        word = Document(BytesIO(package.docx))
        assert len(word.tables) == 1, "The new month needs exactly the standing matrix, with no old completion table."
        values = tuple(tuple(cell.text for cell in row.cells) for row in word.tables[0].rows)
        assert values == (MATRIX.columns, *MATRIX.rows)
        with ZipFile(BytesIO(package.docx)) as archive:
            xml = "\n".join(archive.read(name).decode("utf-8") for name in archive.namelist() if name.endswith(".xml"))
        for stale in ("Prior safety briefing", "2026-09-04", "0.75", "Synthetic source-only"):
            assert stale not in xml
        with fitz.open(stream=package.pdf, filetype="pdf") as pdf:
            # A native cell can wrap a name without changing its content.
            text = re.sub(r"\s+", " ", "\n".join(page.get_text() for page in pdf))
            assert text.count("Synthetic Trainee") == 1
            assert text.count("Synthetic Colleague") == 1
            for value in ("Fire safety", "Operations", "Orientation", "Completed", "Pending", "Not recorded", "Not required"):
                assert value in text
            assert "Prior safety briefing" not in text and "2026-09-04" not in text
            assert "Synthetic source-only" not in text
        assert current.blocks[0].extra_tables == (MATRIX,)
    assert all(saved == loaded for saved, loaded in readbacks)
