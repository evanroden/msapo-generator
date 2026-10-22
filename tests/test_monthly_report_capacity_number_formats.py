"""XLSX capacity values read as their number formats show them (Refs #286, G9-pr89-5).

Stored values were shown raw: a date became a serial number, a percentage a
fraction, and a "1,200 tons" unit format was dropped. Synthetic values only.
"""
import datetime
from io import BytesIO
from zipfile import ZipFile

from openpyxl import Workbook
from openpyxl.utils.datetime import CALENDAR_MAC_1904
import pytest

from app import monthly_report_capacity as capacity, monthly_report_directory as directory
from app.monthly_report_model import Facility, ReportProfile

FACILITIES = (Facility("north", "North Hospital", ("NH",)),)
PROFILE = ReportProfile("Synthetic Capacity Contract", "north", "North Hospital", FACILITIES)


@pytest.fixture(autouse=True)
def isolated(tmp_path, monkeypatch):
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path))
    directory.save_directory(PROFILE.contract, (directory.DirectorySite("north", "North Hospital", ("NH",)),),
                             expected_revision=0, actor="Synthetic Reviewer", confirmed=True)


def capacity_workbook(epoch=None):
    book = Workbook()
    if epoch:
        book.epoch = epoch
    sheet = book.active
    sheet.title = "Capacity"
    sheet.append(["Facility", "Service", "Required capacity", "Available capacity", "Utilization", "Last tested"])
    sheet.append(["NH", "Chilled water", 1200, 1600.25, 0.75, datetime.date(2026, 9, 15)])
    for cell, code in (("C2", '#,##0 "tons"'), ("D2", '#,##0 "tons"'), ("E2", "0%"), ("F2", "mm/dd/yyyy")):
        sheet[cell].number_format = code
    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()


@pytest.mark.parametrize("epoch", [None, CALENDAR_MAC_1904])
def test_capacity_cells_keep_units_percentages_and_dates(epoch):
    table, = capacity.inspect_capacity(PROFILE, "capacity.xlsx", capacity_workbook(epoch)).tables
    assert table.rows == (("NH", "Chilled water", "1,200 tons", "1,600 tons", "75%", "2026-09-15"),)
    assert table.site_keys == ("north",)


def test_directory_reads_stay_unformatted_unless_requested():
    raw = capacity_workbook()
    plain = directory.inspect_workbook(raw).sheets[0].values()
    assert (plain[(2, 3)], plain[(2, 5)]) == ("1200", "0.75")
    assert all(not cell.number_format for cell in directory.inspect_workbook(raw).sheets[0].cells)
    formatted = {(c.row, c.column): c.number_format for c in directory.inspect_workbook(raw, number_formats=True).sheets[0].cells}
    assert (formatted[(2, 3)], formatted[(2, 5)], formatted[(1, 1)]) == ('#,##0 "tons"', "0%", "")


def test_an_unreadable_styles_part_keeps_stored_values():
    source, output = ZipFile(BytesIO(capacity_workbook())), BytesIO()
    with ZipFile(output, "w") as target:
        for info in source.infolist():
            target.writestr(info, b"<styleSheet" if info.filename == "xl/styles.xml" else source.read(info))
    table, = capacity.inspect_capacity(PROFILE, "capacity.xlsx", output.getvalue()).tables
    assert table.rows[0][2:5] == ("1200", "1600.25", "0.75")


@pytest.mark.parametrize(("value", "code", "shown"), [
    ("1200", '#,##0 "tons"', "1,200 tons"),
    ("1600.25", "#,##0.00", "1,600.25"),
    ("0.755", "0.0%", "75.5%"),
    ("1.5", "#.##", "1.5"),
    ("2", "0.0#", "2.0"),
    ("1200", '"$"#,##0', "$1,200"),
    ("1200", "[Red]#,##0", "1,200"),
    ("1200", '#,##0 "tons;peak"', "1,200 tons;peak"),
    ("1200", "General", "1200"),
    ("46280.5", "m/d/yy h:mm", "2026-09-15 12:00"),
    ("0.75", "h:mm:ss", "18:00:00"),
    ("-5", "#,##0;(#,##0)", "-5"),
    ("0", '#,##0;(#,##0);"-"', "0"),
    ("1200", '[>=1000]#,##0;0', "1200"),
    ("12000", '#,##0,"k"', "12000"),
    ("0.25", "# ?/?", "0.25"),
    ("1200", "0.00E+00", "1200"),
    ("1.5", "[h]:mm", "1.5"),
    ("1200", "", "1200"),
    ("Synthetic note", "#,##0", "Synthetic note"),
])
def test_formatted_value_applies_only_supported_formats(value, code, shown):
    assert directory.formatted_value(directory.DirectoryCell(1, 1, value, number_format=code)) == shown


def test_1904_dates_use_their_own_epoch():
    cell = directory.DirectoryCell(1, 1, "0", number_format="yyyy-mm-dd")
    assert directory.formatted_value(cell, date1904=True) == "1904-01-01"
    assert directory.formatted_value(directory.DirectoryCell(1, 1, "1", number_format="yyyy-mm-dd")) == "1900-01-01"
