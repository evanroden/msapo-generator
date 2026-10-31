"""Synthetic year-wide multi-site Work Order grids: exact source and output data."""
from dataclasses import replace
from io import BytesIO

from docx import Document
import pytest

from app.monthly_report_annual_work_order_grid import (
    annual_grid_candidate, annual_site_month_table, annual_grid_matches_profile,
    annual_header_cells,
)
from app.monthly_report_checks import _annual_work_order_month_cells, preflight
from app.monthly_report_docx import _build_docx
from app.monthly_report_import import inspect_docx, imported_draft
from app.monthly_report_model import ReportPeriod, synthetic_profiles
from app.monthly_report_sections import table_without_prices
from app.monthly_report_setup import new_month_draft
from app.monthly_report_start_import import (
    analyze_starting_report, automatic_section_plans, build_starting_import,
)


PERIOD = ReportPeriod(2026, 9)
MONTHS = ("Jan", "Feb", "Mar", "Apr", "May", "Jun",
          "Jul", "Aug", "Sep", "Oct", "Nov", "Dec")
SITES = tuple(site.title for site in synthetic_profiles()[1].facilities)
WO_NORTH = ("0", "", "2", "0", "", "6", "", "8", "9", "", "", "")
WO_SOUTH = ("5", "0", "", "", "", "", "7", "", "11", "", "", "")
PM_NORTH = ("", "1", "", "", "0", "", "", "8", "", "", "", "")
PM_SOUTH = ("1", "", "", "", "", "", "", "0", "", "", "", "")
EXPECTED = (
    ("Work Order Status", SITES[0], "Received", "count", *WO_NORTH),
    ("Work Order Status", SITES[1], "Received", "count", *WO_SOUTH),
    ("PM Status", SITES[0], "Completed", "count", *PM_NORTH),
    ("PM Status", SITES[1], "Completed", "count", *PM_SOUTH),
)


def source(path, *, first_band=False, malformed=False, foreign=False, repeated=False):
    profile = synthetic_profiles()[1]
    document = Document()
    document.add_heading(profile.title, 0)
    document.add_paragraph("Monthly Report · September 2026")
    document.add_heading("2 Monthly Activity Summary", 1)
    document.add_paragraph("WORK ORDER STATUS")
    header = ("Site", "Metric", "Unit", *MONTHS)
    if malformed:
        header = (*header[:-1], "Nov")  # Two November columns, no December.
    north, south = SITES if not foreign else (SITES[0], "Unrelated Client Facility")
    body = (
        ("Work Order Status", *("",) * 14),
        (north, "Received", "count", *WO_NORTH),
        (south, "Received", "count", *WO_SOUTH),
        ("PM Status", *("",) * 14),
        (north, "Completed", "count", *PM_NORTH),
        (south, "Completed", "count", *PM_SOUTH),
    )
    rows = ((*body[:1], header, *body[1:]) if first_band
            else (header, *body))
    if repeated:
        rows = (*rows[:5], header, *rows[5:])
    grid = document.add_table(rows=len(rows), cols=15)
    for index, data in enumerate(rows):
        if data[0] in ("Work Order Status", "PM Status"):
            grid.rows[index].cells[0].merge(grid.rows[index].cells[-1]).text = data[0]
        else:
            for cell, value in zip(grid.rows[index].cells, data):
                cell.text = value
    document.save(path)
    return path


def inspected(path):
    inspection = inspect_docx(path)
    item = next(i for i in inspection.items if i.kind == "table"
                and any("Received" in cell for row in i.rows for cell in row))
    return inspection, item


@pytest.mark.parametrize("first_band", [False, True])
@pytest.mark.parametrize("repeated", [False, True])
def test_all_months_and_site_category_rows_remain_exact_after_source_import(
        tmp_path, first_band, repeated):
    path = source(tmp_path / f"grid-{first_band}-{repeated}.docx",
                  first_band=first_band, repeated=repeated)
    original = path.read_bytes()
    _, item = inspected(path)
    assert annual_grid_candidate(item)
    annual, removed = table_without_prices(item)
    assert removed == () and annual.columns == ("Category", "Site", "Metric", "Unit", *MONTHS)
    assert annual.rows == EXPECTED
    assert annual_grid_matches_profile(annual, synthetic_profiles()[1])
    assert not annual_grid_matches_profile(annual, synthetic_profiles()[0])
    assert annual_header_cells(annual) == frozenset((0, n) for n in range(4, 16))
    assert _annual_work_order_month_cells("work_orders", annual) == annual_header_cells(annual)
    assert path.read_bytes() == original


def test_one_month_cell_is_changed_without_retyping_any_other_site_or_blank(tmp_path):
    path = source(tmp_path / "grid-edit.docx")
    _, item = inspected(path)
    table = annual_site_month_table(item)
    october = table.columns.index("Oct")
    record = list(table.rows[1])
    assert record[october] == ""  # Source blank is not numeric zero.
    record[october] = "7"
    edited = replace(table, rows=(*table.rows[:1], tuple(record), *table.rows[2:]))
    assert edited.rows[1][october] == "7"
    assert edited.rows[0] == table.rows[0]
    assert edited.rows[2:] == table.rows[2:]
    assert edited.rows[1][table.columns.index("Feb")] == "0"
    assert table.rows[1][october] == ""
    assert annual_header_cells(edited) == annual_header_cells(table)


def test_auto_import_and_word_output_have_only_four_grouped_metric_rows(tmp_path):
    profile = synthetic_profiles()[1]
    path = source(tmp_path / "annual-work-order.docx")
    original = path.read_bytes()
    inspection, item = inspected(path)
    analysis = analyze_starting_report(
        inspection, profile, PERIOD, confirmed_sites=profile.facilities)
    plans = automatic_section_plans(inspection, analysis, profile=profile, period=PERIOD)
    assert all(plan["approved"] for plan in plans.values()), {
        key: plan["questions"] for key, plan in plans.items() if plan["questions"]
    }
    mapped, _ = build_starting_import(path, inspection, plans.values(),
                                     profile=profile, period=PERIOD)
    block = next(b for b in mapped.blocks if b.key == "work_orders")
    assert len(block.extra_tables) == 1
    table = block.extra_tables[0]
    assert table.rows == EXPECTED and table.reference.endswith(item.id)
    draft = imported_draft(profile, PERIOD, "Synthetic Reviewer", mapped)
    word = Document(BytesIO(_build_docx(draft)))
    matching = [t for t in word.tables
                if any(cell.text == "Work Order Status"
                       for row in t.rows for cell in row.cells)]
    assert len(matching) == 1
    assert len(matching[0].rows) == 5
    values = tuple(tuple(c.text for c in row.cells) for row in matching[0].rows)
    assert values == (table.columns, *EXPECTED)
    assert path.read_bytes() == original
    assert not any(ch.code == "period" and ch.block_key == "work_orders"
                   for ch in preflight(draft))


def test_month_rollover_keeps_original_report_but_creates_no_new_month_values(tmp_path):
    profile = synthetic_profiles()[1]
    path = source(tmp_path / "annual-rollover.docx")
    inspection, _ = inspected(path)
    analysis = analyze_starting_report(inspection, profile, PERIOD,
                                       confirmed_sites=profile.facilities)
    plans = automatic_section_plans(inspection, analysis, profile=profile, period=PERIOD)
    mapped, _ = build_starting_import(path, inspection, plans.values(),
                                     profile=profile, period=PERIOD)
    previous = imported_draft(profile, PERIOD, "Synthetic Reviewer", mapped)
    after = new_month_draft(previous, ReportPeriod(2026, 10))
    assert next(b for b in previous.blocks if b.key == "work_orders").extra_tables[0].rows == EXPECTED
    later = next(b for b in after.blocks if b.key == "work_orders")
    assert not any(t.rows for t in later.extra_tables)
    assert later.extra_tables[0].columns == ("Category", "Site", "Metric", "Unit", *MONTHS)
    assert later.extra_tables[0].reference == ""
    assert after.period.month == 10 and previous.period.month == 9


@pytest.mark.parametrize("malformed,foreign", [(True, False), (False, True)])
def test_ambiguous_month_or_unapproved_site_requires_review(tmp_path, malformed, foreign):
    path = source(tmp_path / "review.docx", malformed=malformed, foreign=foreign)
    inspection, item = inspected(path)
    assert annual_grid_candidate(item)
    profile = synthetic_profiles()[1]
    analysis = analyze_starting_report(inspection, profile, PERIOD,
                                       confirmed_sites=profile.facilities)
    plans = automatic_section_plans(inspection, analysis, profile=profile, period=PERIOD)
    activity = plans["activity"]
    assert not activity["approved"]
    assert any("annual work-order" in question for question in activity["questions"])
    assert item.id not in activity["selected"]
    # A manual reviewer can still see every original cell, including the
    # wrong month heading or unrelated facility; no source was destroyed.
    assert item.rows and ("Nov" in item.rows[0] if malformed else
                         any("Unrelated Client Facility" in row for row in item.rows))


def test_single_site_profile_cannot_silently_import_two_site_counts(tmp_path):
    path = source(tmp_path / "wrong-profile.docx")
    inspection, item = inspected(path)
    one_site = synthetic_profiles()[0]
    analysis = analyze_starting_report(inspection, one_site, PERIOD,
                                       confirmed_sites=one_site.facilities)
    plans = automatic_section_plans(inspection, analysis, profile=one_site, period=PERIOD)
    assert not plans["activity"]["approved"]
    assert item.id not in plans["activity"]["selected"]


def test_source_dates_in_real_cells_still_trigger_period_review(tmp_path):
    from app.monthly_report_model import ReportDraft, BlockSpec, ResolvedBlock, SectionSpec
    table = annual_site_month_table(inspected(source(tmp_path / "date.docx"))[1])
    modified = list(table.rows[0]); modified[6] = "August 2025"
    changed = replace(table, rows=(tuple(modified), *table.rows[1:]))
    profile = synthetic_profiles()[1]
    draft = ReportDraft(profile, PERIOD, "Synthetic Reviewer",
        (SectionSpec("activity", "2", "Monthly Activity Summary",
                     (BlockSpec("work_orders", "work_order_grid"),)),),
        (ResolvedBlock("work_orders", "This month", extra_tables=(changed,)),))
    assert any(c.code == "period" and c.block_key == "work_orders" for c in preflight(draft))
