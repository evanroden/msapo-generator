from io import BytesIO
from dataclasses import replace
from copy import deepcopy
import pytest
from PIL import Image
from docx import Document
from docx.oxml.ns import qn
from lxml import etree
from app.monthly_report_native_layout import build_native_docx, _contain_cover_logo, _unchanged_photo_geometry_supported
from app.monthly_report_cover_ui import replace_cover_picture
from app.monthly_report_import import inspect_docx,map_items,ImportMapping
from test_cover_logo_aspect import source,frames


def png(size,color='red'):
    out=BytesIO(); Image.new('RGBA',size,color).save(out,'PNG'); return out.getvalue()


def pics(raw):
    d=Document(BytesIO(raw)); return [(n.get(qn('r:embed')),d.part.rels[n.get(qn('r:embed'))].target_part.blob) for n in d.element.iter(qn('a:blip'))]


@pytest.mark.parametrize('role',['brand_logo','client_logo'])
@pytest.mark.parametrize('size',[(500,50),(50,500),(100,100)])
@pytest.mark.parametrize('section',[None,'cover'])
def test_replacement_selected_size_and_unselected_bytes(tmp_path,role,size,section):
    path,draft,assets,original=source(tmp_path,(400,100),role,True,repeated=True)
    replacement=replace_cover_picture(draft.blocks[0],png(size),'.png',assets)
    result=build_native_docx(replace(draft,blocks=(replacement,)),path,asset_loader=assets.__getitem__,section_key=section)
    assert frames(result)[0][0]/frames(result)[0][1] == pytest.approx(size[0]/size[1],rel=1e-5)
    assert pics(result)[1][1] == original, 'Unselected equal-byte cover occurrence must keep original bytes'
    assert frames(result)[1] == (1828800,1828800)


def test_replacement_two_selected_roles_same_source_part(tmp_path):
    path,draft,assets,original=source(tmp_path,(400,100),'brand_logo',True,repeated=True)
    inspection=inspect_docx(path); items=[i for i in inspection.items if i.kind=='image']
    mapped=map_items(path,inspection,(ImportMapping(items[0].id,'brand_logo'),ImportMapping(items[1].id,'client_logo')))
    assets=dict(mapped.assets); blocks={b.key:b for b in mapped.blocks}
    a=replace_cover_picture(blocks['brand_logo'],png((500,50),'red'),'.png',assets)
    b=replace_cover_picture(blocks['client_logo'],png((50,500),'blue'),'.png',assets)
    result=build_native_docx(replace(draft,blocks=(a,b)),path,asset_loader=assets.__getitem__)
    assert [p[1] for p in pics(result)] == [assets[a.asset_hashes[0]],assets[b.asset_hashes[0]]]
    assert frames(result)[0][0]/frames(result)[0][1] == pytest.approx(10)
    assert frames(result)[1][0]/frames(result)[1][1] == pytest.approx(.1)


@pytest.mark.parametrize('crop',[{'l':'10000','r':'15000'}, {'l':'-25000','r':'-25000','t':'-50000','b':'-50000'}])
def test_valid_crop_preserved(tmp_path,crop):
    path,draft,assets,_=source(tmp_path,(400,100),'client_logo',True)
    d=Document(path); blip=next(d.element.iter(qn('a:blip'))); rect=etree.SubElement(blip.getparent(),qn('a:srcRect'),crop)
    before=etree.tostring(d.element)
    _contain_cover_logo(blip,400,100,d.part,preserve_crop=True)
    assert etree.tostring(d.element)==before


@pytest.mark.parametrize('style', ['width:3in;height:1in', 'width:3in; height:1in ', 'width:3in;height:1in;width:2in'])
def test_vml_style_dimensions(tmp_path,style):
    path,draft,assets,_=source(tmp_path,(100,400),'client_logo',True)
    d=Document(path); rid=next(d.element.iter(qn('a:blip'))).get(qn('r:embed'))
    v='{urn:schemas-microsoft-com:vml}'
    shape=etree.Element(v+'shape',style=style); node=etree.SubElement(shape,v+'imagedata',{qn('r:id'):rid})
    assert _unchanged_photo_geometry_supported(node)
    _contain_cover_logo(node,100,400,d.part)
    styles=dict((k.strip(),v.strip()) for e in shape.get('style').split(';') if ':' in e for k,v in [e.split(':',1)])
    assert float(styles['width'][:-2])/float(styles['height'][:-2]) == pytest.approx(.25)


def test_alternate_guard_is_atomic_when_one_vml_representation_has_trailing_space(tmp_path):
    path,draft,assets,_=source(tmp_path,(100,400),'client_logo',True)
    d=Document(path); drawing=next(d.element.iter(qn('w:drawing'))); parent=drawing.getparent(); parent.remove(drawing)
    mc='{http://schemas.openxmlformats.org/markup-compatibility/2006}'; v='{urn:schemas-microsoft-com:vml}'
    alternate=etree.SubElement(parent,mc+'AlternateContent'); etree.SubElement(alternate,mc+'Choice',Requires='w').append(drawing)
    fallback=etree.SubElement(alternate,mc+'Fallback'); shape=etree.SubElement(fallback,v+'shape',style='width:3in;height:1in ')
    blip=next(drawing.iter(qn('a:blip'))); node=etree.SubElement(shape,v+'imagedata',{qn('r:id'):blip.get(qn('r:embed'))})
    for item in (blip,node): _contain_cover_logo(item,100,400,d.part)
    assert float(shape.get('style').split('width:')[1].split('in')[0]) == pytest.approx(.25)


@pytest.mark.parametrize('where',['cover','header'])
@pytest.mark.parametrize('orientation',[1,6])
def test_replaced_alternate_oriented_source_fits_both_representations(tmp_path,where,orientation):
    from app.monthly_report_model import ReportDraft,ReportPeriod,synthetic_profiles
    from docx.shared import Inches
    image=Image.new('RGB',(400,100),'green'); exif=image.getexif(); exif[274]=orientation
    data=BytesIO(); image.save(data,'JPEG',exif=exif)
    d=Document(); d.add_paragraph('Synthetic facility'); d.add_paragraph('August 2026')
    paragraph=d.add_paragraph() if where=='cover' else d.sections[0].header.paragraphs[0]
    shape=paragraph.add_run().add_picture(BytesIO(data.getvalue()),width=Inches(3),height=Inches(1))
    drawing=shape._inline.getparent(); parent=drawing.getparent(); parent.remove(drawing)
    mc='{http://schemas.openxmlformats.org/markup-compatibility/2006}'
    alternate=etree.SubElement(parent,mc+'AlternateContent')
    etree.SubElement(alternate,mc+'Choice',Requires='w').append(drawing)
    etree.SubElement(alternate,mc+'Fallback').append(deepcopy(drawing))
    path=tmp_path/'oriented.docx'; d.save(path)
    inspection=inspect_docx(path); items=[i for i in inspection.items if i.kind=='image']; assert len(items)==1
    mapped=map_items(path,inspection,(ImportMapping(items[0].id,'client_logo'),)); assets=dict(mapped.assets)
    replacement=replace_cover_picture(mapped.blocks[0],png((100,400)),'.png',assets)
    draft=ReportDraft(synthetic_profiles()[0],ReportPeriod(2026,10),'Synthetic reviewer',(),(replacement,))
    result=build_native_docx(draft,path,asset_loader=assets.__getitem__)
    output=Document(BytesIO(result)); owner=output.element if where=='cover' else output.sections[0].header._element
    ratios=[]
    for ext in owner.iter(qn('wp:extent')): ratios.append(int(ext.get('cx'))/int(ext.get('cy')))
    assert ratios == pytest.approx([.25,.25])


@pytest.mark.parametrize('replacement',[False,True])
def test_alternate_picture_branches_share_one_wp_extent(tmp_path,replacement):
    path,draft,assets,_=source(tmp_path,(100,400),'client_logo',True)
    d=Document(path); picture=next(d.element.iter(qn('pic:pic'))); parent=picture.getparent(); parent.remove(picture)
    mc='{http://schemas.openxmlformats.org/markup-compatibility/2006}'
    alternate=etree.SubElement(parent,mc+'AlternateContent')
    etree.SubElement(alternate,mc+'Choice',Requires='pic').append(picture)
    etree.SubElement(alternate,mc+'Fallback').append(deepcopy(picture))
    d.save(path); inspection=inspect_docx(path); items=[i for i in inspection.items if i.kind=='image']; assert len(items)==1
    mapped=map_items(path,inspection,(ImportMapping(items[0].id,'client_logo'),)); assets=dict(mapped.assets)
    block=mapped.blocks[0]
    if replacement: block=replace_cover_picture(block,png((100,400)),'.png',assets)
    result=build_native_docx(replace(draft,blocks=(block,)),path,asset_loader=assets.__getitem__)
    assert frames(result)==[(228600,914400),(228600,914400)]


def test_conflicting_exact_logo_destinations_do_not_choose_media_role(tmp_path):
    from app.monthly_report_native_layout import NativeLayoutError
    path,draft,assets,original=source(tmp_path,(400,100),'brand_logo',True)
    a=replace_cover_picture(draft.blocks[0],png((500,50),'red'),'.png',assets)
    b=replace_cover_picture(replace(draft.blocks[0],key='client_logo'),png((50,500),'blue'),'.png',assets)
    try: result=build_native_docx(replace(draft,blocks=(a,b)),path,asset_loader=assets.__getitem__)
    except NativeLayoutError: return
    assert pics(result)[0][1] == original, 'Conflicting exact roles cannot authorize choosing arbitrary replacement'


@pytest.mark.parametrize('alpha',[False,True])
def test_shared_outer_preserves_valid_crop_branch_when_other_branch_needs_resize(tmp_path,alpha):
    path,draft,assets,_=source(tmp_path,(400,100),'client_logo',alpha)
    d=Document(path); next(d.element.iter(qn('wp:docPr'))).set('descr','Client logo')
    picture=next(d.element.iter(qn('pic:pic'))); parent=picture.getparent(); parent.remove(picture)
    mc='{http://schemas.openxmlformats.org/markup-compatibility/2006}'
    alternate=etree.SubElement(parent,mc+'AlternateContent')
    etree.SubElement(alternate,mc+'Choice',Requires='pic').append(picture)
    etree.SubElement(alternate,mc+'Fallback').append(deepcopy(picture))
    crop=etree.SubElement(picture.find(qn('pic:blipFill')),qn('a:srcRect'),{'l':'10000','r':'15000'})
    d.save(path); inspection=inspect_docx(path); items=[i for i in inspection.items if i.kind=='image']; assert len(items)==1
    mapped=map_items(path,inspection,(ImportMapping(items[0].id,'client_logo'),)); assets=dict(mapped.assets)
    result=build_native_docx(replace(draft,blocks=mapped.blocks),path,asset_loader=assets.__getitem__)
    output=Document(BytesIO(result)); picture=next(output.element.iter(qn('pic:pic')))
    inner=picture.find(qn('pic:spPr')).find(qn('a:xfrm')).find(qn('a:ext'))
    outer=next(picture.iterancestors(qn('wp:inline'))).find(qn('wp:extent'))
    assert dict(inner.attrib)==dict(outer.attrib), 'Valid cropped branch must not inherit a resized shared frame'
    assert dict(inner.attrib)=={'cx':'2743200','cy':'914400'}


@pytest.mark.parametrize('alpha',[False,True])
def test_unchanged_imported_logo_valid_crop_survives_native_build(tmp_path,alpha):
    path,draft,assets,_=source(tmp_path,(400,100),'client_logo',alpha)
    d=Document(path); pic=next(d.element.iter(qn('pic:pic'))); blip=next(pic.iter(qn('a:blip')))
    crop=etree.Element(qn('a:srcRect'),{'l':'10000','r':'15000'}); blip.addnext(crop)
    d.save(path); inspection=inspect_docx(path); item=next(i for i in inspection.items if i.kind=='image')
    mapped=map_items(path,inspection,(ImportMapping(item.id,'client_logo'),)); assets=dict(mapped.assets)
    result=build_native_docx(replace(draft,blocks=mapped.blocks),path,asset_loader=assets.__getitem__)
    output=Document(BytesIO(result)); pic=next(output.element.iter(qn('pic:pic')))
    retained=pic.find(qn('pic:blipFill')).find(qn('a:srcRect'))
    assert retained is not None and dict(retained.attrib)=={'l':'10000','r':'15000'}
    assert frames(result)==[(2743200,914400)]
