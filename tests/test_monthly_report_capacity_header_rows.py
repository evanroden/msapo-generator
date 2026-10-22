"""Capacity headings are found without promoting a data row (Refs #286, G9-pr89-2, G9-pr89-10).

A plant name such as "Hospital boiler plant (steam)" matched both the site and
thermal vocabulary, so that data row became the header and the real
Required/Available heading above it was discarded. A CSV two-row header kept
"Column 3" and left its Required/Available row unmapped. Synthetic values only.
"""
from io import BytesIO

from openpyxl import Workbook
import pytest

from app import monthly_report_capacity as capacity, monthly_report_directory as directory
from app.monthly_report_model import Facility, ReportProfile

FACILITIES = (Facility("north", "North Hospital", ("NH",)), Facility("south", "South Hospital", ("SH",)))
PROFILE = ReportProfile("Synthetic Capacity Contract", "north", "North Hospital", FACILITIES[:1])


@pytest.fixture(autouse=True)
def isolated(tmp_path, monkeypatch):
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path))
    directory.save_directory(PROFILE.contract, tuple(directory.DirectorySite(f.key, f.title, f.aliases) for f in FACILITIES),
                             expected_revision=0, actor="Synthetic Reviewer", confirmed=True)


def csv_tables(text):
    return capacity.inspect_capacity(PROFILE, "capacity.csv", text.encode()).tables


def xlsx_tables(rows):
    book = Workbook()
    for row in rows:
        book.active.append(row)
    buffer = BytesIO()
    book.save(buffer)
    return capacity.inspect_capacity(PROFILE, "capacity.xlsx", buffer.getvalue()).tables


@pytest.mark.parametrize("required", ["40000", "NR"])
def test_a_plant_named_data_row_never_replaces_the_real_heading(required):
    table, = csv_tables("Plant,Required capacity,Available capacity,Units\n"
                        f"Hospital boiler plant (steam),{required},50000,lb/hr\n"
                        "Central chilled water,1200,1600,tons\n")
    assert table.columns == ("Plant", "Required capacity", "Available capacity", "Units")
    assert table.rows == (("Hospital boiler plant (steam)", required, "50000", "lb/hr"),
                          ("Central chilled water", "1200", "1600", "tons"))


def test_report_title_and_date_rows_above_the_heading_are_still_skipped():
    table, = xlsx_tables([["Synthetic thermal capacity report"], ["Updated 10/01/2026"],
                          ["Plant", "Required capacity", "Available capacity", "Units"],
                          ["Hospital boiler plant (steam)", 40000, 50000, "lb/hr"]])
    assert table.columns == ("Plant", "Required capacity", "Available capacity", "Units")
    assert table.rows == (("Hospital boiler plant (steam)", "40000", "50000", "lb/hr"),)


def test_csv_two_row_header_joins_its_service_and_required_available_labels():
    table, = csv_tables("Facility,Steam (lb/hr),,Chilled water (tons),\n"
                        ",Required,Available,Required,Available\n"
                        "NH,40000,50000,1200,1600\nSH,20000,24000,800,1000\n")
    assert table.columns == ("Facility", "Steam (lb/hr) / Required", "Steam (lb/hr) / Available",
                             "Chilled water (tons) / Required", "Chilled water (tons) / Available")
    assert table.rows == (("NH", "40000", "50000", "1200", "1600"), ("SH", "20000", "24000", "800", "1000"))
    assert table.site_keys == ("north", "south")


def test_a_single_csv_header_with_one_blank_column_keeps_its_rows_unchanged():
    table, = csv_tables("Facility,Required capacity,,Notes\nNH,40000,,Synthetic note\nSH,20000,,\n")
    assert table.columns == ("Facility", "Required capacity", "Column 3", "Notes")
    assert table.rows == (("NH", "40000", "", "Synthetic note"), ("SH", "20000", "", ""))
    assert table.site_keys == ("north", "south")
