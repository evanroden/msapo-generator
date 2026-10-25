"""Unresolved source drawings keep a reachable recovery control across reruns."""
from io import BytesIO

from docx import Document
from docx.oxml import parse_xml
from docx.shared import Inches
from PIL import Image
import pytest

from app import monthly_report_library as library
from app.monthly_report_import import inspect_docx
from app.monthly_report_model import synthetic_profiles
from app.monthly_report_sections import section_reviews
from tests.test_monthly_report_start_import import start_app


def drawing_report(tmp_path, section):
    document = Document()
    document.add_heading(synthetic_profiles()[0].facilities[0].title, 0)
    document.add_paragraph("September 2026")
    if section == "cover":
        image = BytesIO()
        Image.new("RGB", (600, 400), "navy").save(image, "PNG")
        document.add_picture(BytesIO(image.getvalue()), width=Inches(4))
    else:
        document.add_heading("Pending and Declined Proposals", 1)
    # A text-bearing VML shape is a real inspector input, not a fabricated
    # inspection or an approved replacement. It intentionally stays unread.
    shape = parse_xml('''<w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
        xmlns:v="urn:schemas-microsoft-com:vml"><w:r><w:pict>
        <v:shape id="invented-source-drawing" style="width:180pt;height:40pt">
        <v:textbox><w:txbxContent><w:p><w:r><w:t>Synthetic equipment decision diagram</w:t>
        </w:r></w:p></w:txbxContent></v:textbox></v:shape></w:pict></w:r></w:p>''')
    document.element.body.insert(len(document.element.body) - 1, shape)
    path = tmp_path / "invented-starting-report.docx"
    document.save(path)
    inspection = inspect_docx(path)
    portion = next(s for s in section_reviews(inspection) if s.key == section)
    assert any(i.kind == "unsupported" for i in portion.items)
    return path, portion.title


def continue_button(app):
    return next(b for b in app.button if b.label == "Continue to this month’s updates")


def test_cover_photo_gallery_does_not_approve_unresolved_native_drawing(monkeypatch, tmp_path):
    path, _ = drawing_report(tmp_path, "cover")
    original = path.read_bytes()
    app = start_app(monkeypatch, tmp_path, path)
    assert continue_button(app).disabled, "A photo choice cannot resolve an unread cover drawing."
    app.run()
    assert continue_button(app).disabled
    assert not app.exception
    assert path.read_bytes() == original
    assert not library.list_profiles(synthetic_profiles()[0].contract)


@pytest.mark.parametrize("section,upload_label", [
    ("cover", "New cover photograph"),
    ("proposals", "Updated price-free proposal image"),
])
def test_skipped_auto_preservation_has_visible_recovery_that_survives_reruns(
        monkeypatch, tmp_path, section, upload_label):
    path, title = drawing_report(tmp_path, section)
    original = path.read_bytes()
    app = start_app(monkeypatch, tmp_path, path)
    for _ in range(3):
        assert title in {box.label for box in app.expander}, "The named review has no reachable control."
        assert any("could not be preserved" in e.value for e in app.error)
        assert continue_button(app).disabled
        assert not app.exception
        app.run()
    action = next(r for r in app.radio if r.label == "What would you like to do?")
    action.set_value("Edit text or change pictures").run()
    assert upload_label in {x.proto.label for x in app.get("file_uploader")}
    assert continue_button(app).disabled
    assert path.read_bytes() == original
    assert not library.list_profiles(synthetic_profiles()[0].contract)


@pytest.mark.parametrize("approved,questions,dedicated,expected", [
    (True, (), False, False),
    (False, (), False, True),
    (True, ("A Word diagram or image needs preservation.",), False, True),
    (False, ("A Word diagram or image needs preservation.",), False, True),
    (False, ("A Word diagram or image needs preservation.",), True, False),
    (False, ("Check the table headings.",), True, True),
    (False, ("A Word diagram or image needs preservation.", "Remove pricing."), True, True),
    (False, (), True, True),
])
def test_dedicated_recovery_suppresses_only_its_own_exclusive_question(
        approved, questions, dedicated, expected):
    from app.monthly_report_section_ui import _starting_review_sections
    from app.monthly_report_sections import ImportSection
    section = ImportSection("issues", "Equipment Issues", ())
    plans = {"issues": {"approved": approved, "questions": questions}}
    failures = {"issues": "Source rendering failed"} if dedicated else {}
    assert bool(_starting_review_sections((section,), plans, failures)) is expected
    assert plans == {"issues": {"approved": approved, "questions": questions}}


def test_section_status_matches_approval_on_same_choice_without_extra_rerun(monkeypatch, tmp_path):
    """Review/Ready cannot lag behind this render's actual import decision.

    A real synthetic DOCX has an unresolved proposal drawing. Deliberately
    leaving that imported section out is permitted, but the old UI displayed
    the prior Review state until the next rerun. Returning to Keep must show
    Review immediately while the unresolved artwork still blocks Continue.
    """
    path, title = drawing_report(tmp_path, "proposals")
    original_bytes = path.read_bytes()
    app = start_app(monkeypatch, tmp_path, path)
    assert title in {item.label for item in app.expander}
    assert not app.exception
    assert "Review" in {caption.value for caption in app.caption}
    assert continue_button(app).disabled

    action = next(item for item in app.radio if item.label == "What would you like to do?")
    assert action.value == "Keep and review"
    action.set_value("Leave this section out").run()
    assert not app.exception
    # No compensating app.run() is allowed: this is the first settled action.
    captions = {item.value for item in app.caption}
    assert "Ready" in captions and "Review" not in captions

    action = next(item for item in app.radio if item.label == "What would you like to do?")
    action.set_value("Keep and review").run()
    assert not app.exception
    captions = {item.value for item in app.caption}
    assert "Review" in captions and "Ready" not in captions
    assert continue_button(app).disabled
    assert path.read_bytes() == original_bytes
    assert not library.list_profiles(synthetic_profiles()[0].contract)
