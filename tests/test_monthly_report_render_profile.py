from dataclasses import replace
from io import BytesIO
import json

from docx import Document
from docx.oxml import OxmlElement
from docx.shared import Pt
import pymupdf as fitz
from PIL import Image
import pytest

from app import monthly_report_designs as designs
from app import monthly_report_library as library
from app.monthly_report_model import default_sections, synthetic_profiles
from app.monthly_report_render_profile import DEFAULT, calibrate, validate_profile


def pair(tmp_path, origin=0, producer='Unrelated producer', title='Synthetic North September 2026'):
    image=Image.new('RGB',(300,180))
    image.putdata([(x%256,y%256,(x+y)%256)for y in range(180)for x in range(300)])
    buffer=BytesIO();image.save(buffer,format='PNG');raw=buffer.getvalue()
    doc=Document();doc.add_paragraph(title)
    doc.add_paragraph('Operations and Maintenance Monthly Review Prepared by Synthetic Team')
    inline=doc.add_paragraph().add_run().add_picture(BytesIO(raw),width=Pt(300),height=Pt(180))._inline
    inline.tag='{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}anchor'
    position=OxmlElement('wp:positionV');position.set('relativeFrom','paragraph')
    offset=OxmlElement('wp:posOffset');offset.text=str(200*12700);position.append(offset);inline.insert(0,position)
    doc.add_section()
    for section in default_sections():
        doc.add_heading(section.title,1)
    source=tmp_path/'source.docx';doc.save(source)
    pdf=fitz.open();page=pdf.new_page()
    page.insert_textbox(fitz.Rect(40,20,570,140),title+'\nOperations and Maintenance Monthly Review Prepared by Synthetic Team',fontsize=12)
    page.insert_image(fitz.Rect(72,200+origin,372,380+origin),stream=raw)
    pdf.set_metadata({'producer':producer})
    path=tmp_path/f'companion-{origin}.pdf';pdf.save(path);pdf.close()
    return source,path


def test_pair_calibrates_geometry_not_producer(tmp_path):
    source,pdf=pair(tmp_path,producer='LibreOffice pretend')
    zero=calibrate(source,pdf)
    assert zero['render_profile']['cover_zero_origin']
    assert zero['cover_measurements'][0]['origin_points']==0
    source,pdf=pair(tmp_path,origin=49,producer='Adobe PDF Library pretend')
    assert calibrate(source,pdf)['render_profile']==DEFAULT


def test_pair_rejects_wrong_identity_and_unknown_image_stays_native(tmp_path):
    source,pdf=pair(tmp_path)
    with fitz.open(pdf)as doc:
        page=doc[0];page.add_redact_annot(fitz.Rect(0,0,600,150));page.apply_redactions()
        altered=tmp_path/'other.pdf';doc.save(altered)
    with pytest.raises(ValueError,match='identity'):
        calibrate(source,altered)
    with fitz.open(pdf)as doc:
        page=doc[0];page.add_redact_annot(fitz.Rect(72,200,372,380));page.apply_redactions()
        altered=tmp_path/'no-photo.pdf';doc.save(altered)
    assert calibrate(source,altered)['render_profile']==DEFAULT


def test_paired_pin_immutable_legacy_default_and_integrity(tmp_path,monkeypatch):
    monkeypatch.setenv('EPC_DATA_DIR',str(tmp_path/'data'))
    source,pdf=pair(tmp_path)
    profile=synthetic_profiles()[0]
    legacy=designs.install(source,profile,actor='Editor',expected_revision=0)
    paired=designs.install(source,profile,actor='Editor',expected_revision=1,companion_pdf=pdf)
    assert legacy.template!=paired.template
    assert designs.render_profile_for(legacy)==DEFAULT
    assert designs.render_profile_for(paired)['cover_zero_origin']
    assert designs.source_for(paired).read_bytes()==source.read_bytes()
    assert designs.source_for(legacy).read_bytes()==source.read_bytes()
    assert designs.pin(legacy)==legacy
    meta=designs._directory(profile.contract)/'metadata'/(designs.fingerprint(paired)+'.json')
    original_metadata=meta.read_bytes()
    meta.unlink()
    with pytest.raises(ValueError,match='unavailable'):
        designs.source_for(paired)
    meta.write_bytes(original_metadata)
    value=json.loads(meta.read_text());value['render_profile']['cover_zero_origin']=False
    meta.write_text(json.dumps(value))
    with pytest.raises(ValueError,match='integrity'):
        designs.render_profile_for(paired)


def test_paired_master_versions_do_not_change_pinned_report(tmp_path,monkeypatch):
    monkeypatch.setenv('EPC_DATA_DIR',str(tmp_path/'data'))
    source,pdf=pair(tmp_path)
    old=designs.install_master(source,actor='Editor',expected_revision=0)
    profile=designs.pin(replace(synthetic_profiles()[0],template=designs.MASTER_PREFIX+old))
    new=designs.install_master(source,actor='Editor',expected_revision=1,companion_pdf=pdf)
    assert new!=old and designs.pin(profile)==profile
    latest=designs.pin(profile,latest_master=True)
    assert designs.render_profile_for(latest)['cover_zero_origin']
    assert designs.render_profile_for(profile)==DEFAULT
    assert designs.master_section_order(latest)
    metadata=library._read(library._root()/'designs/master/metadata'/(new+'.json'))
    companion=library._root()/'designs/master'/(metadata['companion_sha256']+'.pdf');companion.unlink();companion.write_bytes(b'corrupt')
    with pytest.raises(ValueError,match='integrity'):
        designs.source_for(latest)


def test_profile_flags_are_finite_typed_and_versioned():
    assert validate_profile()==DEFAULT
    for invalid in ({'version':1,'cover_zero_origin':1},{'version':True,'cover_zero_origin':False},
                    {'version':2,'cover_zero_origin':True},{**DEFAULT,'font':'Arial'}):
        with pytest.raises(ValueError):validate_profile(invalid)


def test_final_and_section_pdf_boundaries_resolve_profile_and_limit_cover_scope(tmp_path,monkeypatch):
    from app import monthly_report_docx as output, monthly_report_preview as preview
    from app.monthly_report_model import ReportPeriod, synthetic_draft
    draft=synthetic_draft(synthetic_profiles()[0],ReportPeriod(2026,9))
    profile={'version':1,'cover_zero_origin':True}
    monkeypatch.setattr(designs,'render_profile_for',lambda p:dict(profile))
    monkeypatch.setattr(output,'assemble_docx',lambda *a,**kw:b'native-download')
    monkeypatch.setattr(preview,'section_preview_docx',lambda *a,**kw:b'native-preview')
    calls=[]
    def render(raw,*,profile):
        calls.append((raw,profile));return b'conversion-only'
    monkeypatch.setattr(output,'rendering_docx',render)
    monkeypatch.setattr(preview,'rendering_docx',render)
    def convert(path):
        assert path.read_bytes()==b'conversion-only'
        result=tmp_path/'converted.pdf'
        pdf=fitz.open();pdf.new_page();pdf.save(result);pdf.close()
        return result
    monkeypatch.setattr(output.pdf_converter,'convert_to_pdf',convert)
    package=output.generate_report(draft)
    assert package.docx==b'native-download' and package.pdf
    preview.preview_section(draft,'cover')
    preview.preview_section(draft,'training')
    assert calls==[(b'native-download',profile),(b'native-preview',profile),(b'native-preview',DEFAULT)]


def test_oversized_cover_images_are_checked_before_any_pixel_decoding(monkeypatch):
    from app.monthly_report_render_profile import _pdf_evidence
    class Page:
        def get_text(self):return 'Synthetic report cover text'
        def get_images(self,full):return [(1,0,100_000,100_000)]
        def get_image_info(self,**kwargs):raise AssertionError('Oversized image decoded')
    class Pdf:
        is_encrypted=False
        metadata={'producer':'Synthetic'}
        def __enter__(self):return self
        def __exit__(self,*args):pass
        def __len__(self):return 1
        def __getitem__(self,index):return Page()
    monkeypatch.setattr(fitz,'open',lambda path:Pdf())
    assert _pdf_evidence('synthetic.pdf')['images']==[]
