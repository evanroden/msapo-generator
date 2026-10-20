"""Synthetic coverage for adjacent labels in independent Word chrome parts."""

from io import BytesIO

from docx import Document
from docx.shared import Inches
import pytest

from app.monthly_report_import import inspect_docx
from test_cover_branding_contracts import mark, photograph


def _source(tmp_path, part, label, *, after=False, pictures=1, gap=False):
    document = Document()
    document.add_paragraph("Synthetic monthly report")
    document.add_heading("Monthly Activity Summary", 1)
    document.add_paragraph("Synthetic inspection completed.")
    chrome = getattr(document.sections[0], part)
    first = chrome.paragraphs[0]
    if label and not after:
        first.add_run(label)
        if gap:
            chrome.add_paragraph("Separate synthetic text")
        picture_paragraph = chrome.add_paragraph()
    else:
        picture_paragraph = first
    for index in range(pictures):
        picture_paragraph.add_run().add_picture(
            BytesIO(mark() if index == 0 else photograph()), width=Inches(2)
        )
    if label and after:
        if gap:
            chrome.add_paragraph("Separate synthetic text")
        chrome.add_paragraph(label)
    source = tmp_path / "adjacent-chrome-label.docx"
    document.save(source)
    return source


@pytest.mark.parametrize("part", ["header", "footer"])
@pytest.mark.parametrize("after", [False, True], ids=["preceding-label", "following-label"])
@pytest.mark.parametrize("label,role", [
    ("Synthetic company logo", "brand_logo"),
    ("Synthetic client logo", "client_logo"),
])
def test_adjacent_chrome_label_is_independent_of_final_body_section(tmp_path, part, after, label, role):
    source = _source(tmp_path, part, label, after=after)
    original = source.read_bytes()
    inspection = inspect_docx(source)
    chrome_items = [item for item in inspection.items if item.part.startswith("word/" + part)]
    label_item = next(item for item in chrome_items if item.kind == "text")
    assert label_item.section == "activity", "Exercise the inherited final-body section"
    image = next(item for item in chrome_items if item.kind == "image")
    assert image.image_role == role
    assert image.suggested_slot == role
    assert source.read_bytes() == original


@pytest.mark.parametrize("part", ["header", "footer"])
@pytest.mark.parametrize("after", [False, True])
def test_adjacent_chrome_label_with_two_images_keeps_owner_ambiguous(tmp_path, part, after):
    inspection = inspect_docx(_source(tmp_path, part, "Company logo", after=after, pictures=2))
    images = [item for item in inspection.items if item.kind == "image"]
    assert len(images) == 2
    assert all(item.image_role == "logo" and not item.suggested_slot for item in images)


@pytest.mark.parametrize("part", ["header", "footer"])
@pytest.mark.parametrize("label,gap", [("", False), ("Company logo", True)])
def test_chrome_placement_or_distant_label_does_not_identify_company(tmp_path, part, label, gap):
    inspection = inspect_docx(_source(tmp_path, part, label, gap=gap))
    image = next(item for item in inspection.items if item.kind == "image")
    assert image.image_role == ("header_artwork" if part == "header" else "")
    assert image.suggested_slot != "brand_logo"


def test_body_activity_label_does_not_reclassify_section_image(tmp_path):
    document = Document()
    document.add_heading("Monthly Activity Summary", 1)
    document.add_paragraph("Company logo")
    document.add_picture(BytesIO(photograph()), width=Inches(3))
    source = tmp_path / "body-activity.docx"
    document.save(source)
    image = next(item for item in inspect_docx(source).items if item.kind == "image")
    assert image.section == "activity"
    assert image.suggested_slot == "improvements"
    assert not image.image_role
