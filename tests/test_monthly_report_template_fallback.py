"""Whole-template opt-in retains source/review without approving unread artwork."""
from dataclasses import replace
from copy import deepcopy
from docx import Document
import pytest
from app import monthly_report_designs as designs, monthly_report_library as library
from app.monthly_report_model import synthetic_profiles
from tests.test_monthly_report_starting_review_recovery import drawing_report, continue_button
from tests.test_monthly_report_start_import import start_app
from tests.test_monthly_report_designs import master_reference


def setup(monkeypatch, tmp_path, *, master=True):
    monkeypatch.setenv('EPC_DATA_DIR', str(tmp_path/'runtime'))
    if master:
        designs.install_master(master_reference(tmp_path), actor='Synthetic Editor', expected_revision=0)
    source, _ = drawing_report(tmp_path, 'cover')
    doc = Document(source)
    doc.add_heading('Activity Summary', 1)
    doc.add_paragraph('September 2026 synthetic repair.')
    doc.save(source)
    return start_app(monkeypatch, tmp_path, source), source


def choice(app):
    return next(b for b in app.button if b.label == 'Use ENFRA template and keep this review')


def plan_key(app):
    return next(k for k in app.session_state.filtered_state if k.endswith('_section_plans'))


def test_explicit_template_keeps_body_review_and_original_until_atomic_save(monkeypatch, tmp_path):
    app, source = setup(monkeypatch, tmp_path)
    original = source.read_bytes()
    assert continue_button(app).disabled
    assert not library.list_profiles(synthetic_profiles()[0].contract)
    plans = app.session_state[plan_key(app)]
    # Session plans represent already entered body edits/review decisions.
    body = plans['activity']
    item = body['selected'][0]
    body['texts'][item] = 'September 2026 synthetic human correction.'
    preserved = deepcopy(body)
    choice(app).click().run()
    assert not app.exception
    assert app.session_state[plan_key(app)]['activity'] == preserved
    assert source.read_bytes() == original
    assert not library.list_profiles(synthetic_profiles()[0].contract)
    assert not continue_button(app).disabled, [e.value for e in app.error]
    continue_button(app).click().run()
    assert not app.exception
    assert not app.error, [e.value for e in app.error]
    p = library.list_profiles(synthetic_profiles()[0].contract)[0]
    draft = library.load_imported_draft(p.contract,p.key)
    assert designs.is_master(draft.profile)
    assert any('synthetic human correction' in b.text for b in draft.blocks)
    assert source.read_bytes() == original
    stored = library._profile_path(p.contract,p.key)/'imports'/(draft.profile.imported_from+'.docx')
    assert stored.read_bytes() == original
    from io import BytesIO
    from app.monthly_report_docx import _build_docx
    output = Document(BytesIO(_build_docx(draft)))
    text = '\n'.join(p.text for p in output.paragraphs)
    assert 'synthetic human correction' in text
    assert 'Synthetic equipment decision diagram' not in text


def test_no_template_no_shortcut(monkeypatch,tmp_path):
    app,_ = setup(monkeypatch,tmp_path,master=False)
    assert not any(b.label == 'Use ENFRA template and keep this review' for b in app.button)
    assert continue_button(app).disabled


def test_cancel_restores_cover_decision_without_losing_body(monkeypatch,tmp_path):
    app,_ = setup(monkeypatch,tmp_path)
    before = deepcopy(app.session_state[plan_key(app)])
    choice(app).click().run()
    next(b for b in app.button if b.label == 'Keep uploaded report design instead').click().run()
    assert not app.exception
    assert app.session_state[plan_key(app)] == before
    assert continue_button(app).disabled


def test_corrupt_selected_master_blocks_save_without_clearing_review(monkeypatch,tmp_path):
    app,_ = setup(monkeypatch,tmp_path)
    choice(app).click().run()
    plans = deepcopy(app.session_state[plan_key(app)])
    from app.monthly_report_template_fallback import verified_master_template
    p=synthetic_profiles()[0]
    path=designs.source_for(replace(p,template=verified_master_template(p)))
    path.unlink()  # Synthetic fixture corruption; immutable original upload is untouched.
    path.write_bytes(b'invalid synthetic master')
    app.run()
    assert not app.exception
    assert continue_button(app).disabled
    after = app.session_state[plan_key(app)]
    assert after['activity'] == plans['activity']
    assert not after['cover']['approved']
    assert {k:v for k,v in after['cover'].items() if k not in {'approved','approval_stamp'}} == {k:v for k,v in plans['cover'].items() if k not in {'approved','approval_stamp'}}
    assert not library.list_profiles(p.contract)


def test_concurrent_profile_creation_cannot_be_overwritten(monkeypatch,tmp_path):
    app,_ = setup(monkeypatch,tmp_path)
    choice(app).click().run()
    from app.monthly_report_start import design_seed, save_design_start
    from app.monthly_report_model import ReportPeriod
    draft, assets=design_seed(synthetic_profiles()[0],ReportPeriod(2026,9),'Other editor')
    saved=save_design_start(draft,assets,actor='Other editor',confirmed=True)
    before=deepcopy(app.session_state[plan_key(app)])
    continue_button(app).click().run()
    assert not app.exception
    assert app.error
    assert library.load_profile(draft.profile.contract,draft.profile.key)==saved
    assert app.session_state[plan_key(app)]==before


def test_template_does_not_approve_priced_body_content(monkeypatch,tmp_path):
    app,_=setup(monkeypatch,tmp_path)
    plan=app.session_state[plan_key(app)]['activity']
    plan['texts'][plan['selected'][0]]='Repair costs $500.'
    plan['approved']=False
    plan['questions']=['Review changed content.']
    choice(app).click().run()
    assert not app.exception
    assert continue_button(app).disabled
    retained = app.session_state[plan_key(app)]['activity']
    assert not retained['approved']
    assert any('$500' in text for text in retained['texts'].values())
    assert not library.list_profiles(synthetic_profiles()[0].contract)





def test_resource_limit_explains_intake_vs_native_budget_without_lowering_intake(monkeypatch,tmp_path):
    from app import monthly_report_source_view_runtime as runtime
    from app import monthly_report_compatibility_view as compatibility
    from app.monthly_report_import import MAX_DOCX_BYTES
    monkeypatch.setenv('EPC_DATA_DIR',str(tmp_path/'runtime'))
    designs.install_master(master_reference(tmp_path), actor='Synthetic Editor', expected_revision=0)
    source,_=drawing_report(tmp_path,'cover')
    from zipfile import ZipFile, ZIP_STORED
    with ZipFile(source, 'a') as archive:
        archive.writestr('word/media/synthetic-unused.bin', b'\0' * (2 * 1024 * 1024), compress_type=ZIP_STORED)
    original=source.read_bytes()
    monkeypatch.setattr(compatibility,'MAX_VIEW_SOURCE_BYTES',1024*1024)
    def too_large(*args,**kwargs):
        raise ValueError('Native rendition is not proved: source byte budget')
    monkeypatch.setattr(runtime,'derive_import_view',too_large)
    app=start_app(monkeypatch,tmp_path,source)
    assert not app.exception
    warnings=' '.join(w.value for w in app.warning)
    assert '128 MiB upload limit' in warnings and 'limited to 1 MiB' in warnings
    assert 'not evidence that the artwork is damaged' in warnings
    assert MAX_DOCX_BYTES == 128*1024*1024
    assert source.read_bytes()==original
    assert continue_button(app).disabled
    choice(app).click().run()
    assert not app.exception
    assert not continue_button(app).disabled
    assert source.read_bytes()==original
    assert not library.list_profiles(synthetic_profiles()[0].contract)


def test_template_refuses_authored_section_it_cannot_render(monkeypatch,tmp_path):
    from app.monthly_report_template_fallback import fallback_profile, verified_master_template, validate_fallback_sections
    from app.monthly_report_model import ReportDraft, ReportPeriod, SectionSpec, BlockSpec, ResolvedBlock
    monkeypatch.setenv('EPC_DATA_DIR',str(tmp_path/'runtime'))
    designs.install_master(master_reference(tmp_path),actor='Synthetic Editor',expected_revision=0)
    p=synthetic_profiles()[0]
    p=fallback_profile(p,verified_master_template(p))
    draft=ReportDraft(p,ReportPeriod(2026,9),'Synthetic Editor',
                      (SectionSpec('accounts_receivable','','Accounts Receivable',(BlockSpec('accounts_receivable_notes','rich_text'),)),),
                      (ResolvedBlock('accounts_receivable_notes','This month',text='Synthetic pending invoice status.'),))
    with pytest.raises(ValueError,match='no layout for Accounts Receivable'):
        validate_fallback_sections(draft)
    assert draft.blocks[0].text=='Synthetic pending invoice status.'


@pytest.mark.parametrize('scope_field',range(4))
def test_stale_source_profile_or_month_choice_does_not_authorize_template(monkeypatch,tmp_path,scope_field):
    app,_=setup(monkeypatch,tmp_path)
    choice(app).click().run()
    key=next(k for k in app.session_state.filtered_state if k.endswith('_independent_template'))
    old=app.session_state[key]
    scope=list(old['scope']);scope[scope_field]='different-synthetic-scope'
    app.session_state[key]={**old,'scope':tuple(scope)}
    body=deepcopy(app.session_state[plan_key(app)]['activity'])
    app.run()
    assert not app.exception
    assert continue_button(app).disabled
    assert app.session_state[plan_key(app)]['activity']==body
    assert not library.list_profiles(synthetic_profiles()[0].contract)


def test_actual_template_output_keeps_imported_table_values(monkeypatch,tmp_path):
    from io import BytesIO
    from app.monthly_report_docx import _build_docx
    monkeypatch.setenv('EPC_DATA_DIR',str(tmp_path/'runtime'))
    master=master_reference(tmp_path)
    doc=Document(master)
    table=doc.add_table(rows=2,cols=2)
    table.cell(0,0).text='Task';table.cell(0,1).text='Status'
    next(p for p in doc.paragraphs if p.text=='Monthly Activity Summary')._p.addnext(table._tbl)
    doc.save(master)
    designs.install_master(master,actor='Synthetic Editor',expected_revision=0)
    source,_=drawing_report(tmp_path,'cover')
    doc=Document(source)
    doc.add_heading('Activity Summary',1)
    doc.add_paragraph('September 2026 synthetic work.')
    table=doc.add_table(rows=2,cols=2)
    table.cell(0,0).text='Task';table.cell(0,1).text='Status'
    table.cell(1,0).text='September 2026 synthetic pump repair';table.cell(1,1).text='Complete'
    doc.save(source)
    app=start_app(monkeypatch,tmp_path,source)
    choice(app).click().run()
    assert not continue_button(app).disabled
    continue_button(app).click().run()
    assert not app.exception and not app.error
    p=library.list_profiles(synthetic_profiles()[0].contract)[0]
    draft=library.load_imported_draft(p.contract,p.key)
    output=Document(BytesIO(_build_docx(draft)))
    values=[c.text for t in output.tables for r in t.rows for c in r.cells]
    assert 'September 2026 synthetic pump repair' in values
    assert 'Complete' in values
