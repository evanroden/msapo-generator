"""Text-frame labels remain source review evidence, not editable report prose."""
from hashlib import sha256

from docx import Document
from docx.oxml import parse_xml
import pytest

from app import monthly_report_import as importer
from app.monthly_report_sections import build_section_import, section_reviews

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"
V = "urn:schemas-microsoft-com:vml"


def box(text):
    return f'<w:txbxContent xmlns:w="{W}"><w:p><w:r><w:t>{text}</w:t></w:r></w:p></w:txbxContent>'


def source(tmp_path, markup):
    doc = Document()
    doc.add_heading("Monthly Activity Summary", 1)
    doc.add_paragraph("Ordinary work remains editable.")
    paragraph = doc.add_paragraph()
    paragraph._p.append(parse_xml(markup))
    doc.add_paragraph("Adjacent narrative remains in Activity.")
    path = tmp_path / "synthetic-text-frames.docx"
    doc.save(path)
    return path


@pytest.mark.parametrize("markup,expected", [
    (box("Synthetic grid labels"), "Synthetic grid labels"),
    (f'<w:r xmlns:w="{W}">{box("First label")}{box("Second label")}</w:r>', "First label\nSecond label"),
    (f'<v:roundrect xmlns:v="{V}"><v:textbox>{box("Rounded text frame")}</v:textbox></v:roundrect>', "Rounded text frame"),
    (f'<mc:AlternateContent xmlns:mc="{MC}"><mc:Choice Requires="wps">{box("Visible label")}</mc:Choice>'
     f'<mc:Fallback>{box("Fallback label")}</mc:Fallback></mc:AlternateContent>', "Visible label"),
])
def test_standalone_frame_has_one_source_bound_blocker_before_mapping(tmp_path, markup, expected):
    path = source(tmp_path, markup)
    original = path.read_bytes()
    inspection = importer.inspect_docx(path)
    frame = [item for item in inspection.items if item.position == 3]
    assert len(frame) == 1 and frame[0].kind == "unsupported"
    item = frame[0]
    assert item.text == expected and item.section == "activity"
    assert "text frame" in item.note and "complete page" in item.note
    assert item.part == "word/document.xml" and item.word_section == 1
    assert item.id == "item-0003" and not item.rows
    assert path.read_bytes() == original and inspection.sha256 == sha256(original).hexdigest()
    assert next(i for i in inspection.items if i.text.startswith("Adjacent")).id == "item-0004"
    assert next(s for s in section_reviews(inspection) if s.key == "activity").items
    with pytest.raises(importer.ImportError, match="Unsupported items"):
        importer.map_items(path, inspection, (importer.ImportMapping(item.id, "activity_summary"),))
    plan = {"key": "activity", "target": "activity", "approved": True,
            "selected": [item.id], "unsupported_reviewed": True}
    with pytest.raises(ValueError, match="needs a replacement"):
        build_section_import(path, inspection, (plan,))


def test_frame_label_cannot_change_the_next_paragraph_section(tmp_path):
    path = source(tmp_path, box("Training"))
    inspection = importer.inspect_docx(path)
    frame = next(item for item in inspection.items if item.position == 3)
    following = next(item for item in inspection.items if item.position == 4)
    assert frame.kind == "unsupported" and frame.section == following.section == "activity"
    assert frame.note != "Native section heading"


def test_deleted_frame_does_not_block_or_resurrect_text(tmp_path):
    path = source(tmp_path, f'<w:del xmlns:w="{W}" w:id="1">{box("Deleted frame")}</w:del>')
    inspection = importer.inspect_docx(path)
    assert not any(item.kind == "unsupported" or "Deleted frame" in item.text for item in inspection.items)
    ordinary = next(item for item in inspection.items if item.text.startswith("Ordinary"))
    mapped = importer.map_items(path, inspection, (importer.ImportMapping(ordinary.id, "activity_summary"),))
    assert mapped.blocks[0].text == ordinary.text


def test_ordinary_paragraphs_keep_existing_identity_and_remain_editable(tmp_path):
    path = source(tmp_path, f'<w:r xmlns:w="{W}"><w:t>Plain source wording</w:t></w:r>')
    inspection = importer.inspect_docx(path)
    plain = next(item for item in inspection.items if item.position == 3)
    assert plain.kind == "text" and plain.id == "item-0003"
    assert next(item for item in inspection.items if item.position == 4).id == "item-0004"
    mapped = importer.map_items(path, inspection, (importer.ImportMapping(plain.id, "activity_summary"),))
    assert mapped.blocks[0].text == "Plain source wording"
