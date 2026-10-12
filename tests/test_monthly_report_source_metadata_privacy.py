"""Synthetic previous-client accessibility metadata must not leak into output.

The visual table schema and all current technical values must remain editable;
stripping every drawing/table or all artwork is not an acceptable privacy fix.
"""
from io import BytesIO
from zipfile import ZipFile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from app.monthly_report_model import ResolvedBlock
from app.monthly_report_native_layout import (
    _clear_source_drawing_metadata,
    build_native_docx,
)
from test_monthly_report_native_layout import draft, source


CANARY = "OTHER-CLIENT-PRIVATE-CAPTION-AND-COST-98765"


def _insert_hidden_table_labels(path):
    document = Document(path)
    table_properties = document.tables[0]._tbl.tblPr
    for local in ("tblCaption", "tblDescription"):
        node = OxmlElement("w:" + local)
        node.set(qn("w:val"), CANARY)
        table_properties.append(node)
    document.save(path)


def _payload(raw):
    with ZipFile(BytesIO(raw)) as archive:
        return {name: archive.read(name) for name in archive.namelist()}


def test_imported_table_accessibility_labels_are_not_other_client_facts(tmp_path):
    path = source(tmp_path)
    _insert_hidden_table_labels(path)
    original = _payload(path.read_bytes())
    assert CANARY.encode() in original["word/document.xml"]
    current = draft(
        ResolvedBlock("activity_summary", "This month", text="Current reviewed activity."),
        ResolvedBlock("work_orders", "This month", rows=(("Inspect flowmeter", "Pending"),)),
    )
    native = build_native_docx(current, path, master=True)
    package = _payload(native)
    assert all(CANARY.encode() not in member for member in package.values())
    rendered = Document(BytesIO(native))
    assert len(rendered.tables) == 1
    assert rendered.tables[0].cell(0, 0).text == "Task"
    assert rendered.tables[0].cell(0, 1).text == "Status"
    assert rendered.tables[0].cell(1, 0).text == "Inspect flowmeter"
    assert rendered.tables[0].cell(1, 1).text == "Pending"
    assert "Current reviewed activity." in "\n".join(p.text for p in rendered.paragraphs)


def test_nonvisual_content_control_aliases_are_stripped_without_losing_text():
    document = Document()
    para = document.add_paragraph("Current editable work.")
    # These names were supplied by an old source and not by the current user.
    sdt = OxmlElement("w:sdt")
    sdtpr = OxmlElement("w:sdtPr")
    for local in ("tag", "alias"):
        node = OxmlElement("w:" + local)
        node.set(qn("w:val"), CANARY)
        sdtpr.append(node)
    sdt.append(sdtpr)
    content = OxmlElement("w:sdtContent")
    sdt.append(content)
    content.append(para._p)
    document.element.body.insert(0, sdt)
    _clear_source_drawing_metadata(document)
    assert CANARY.encode() not in document.element.xml.encode("utf-8")
    assert "Current editable work." in document.element.xml
    assert not list(document.element.iter(qn("w:tag")))
    assert not list(document.element.iter(qn("w:alias")))
