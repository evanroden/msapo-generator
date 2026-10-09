"""Synthetic layouts verify payload isolation and source geometry preservation."""
from dataclasses import replace
from io import BytesIO
from zipfile import ZipFile

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches
from PIL import Image

from app.monthly_report_model import (
    BlockSpec,
    ColumnSpec,
    Facility,
    ReportDraft,
    ReportPeriod,
    ReportProfile,
    ResolvedBlock,
    SectionSpec,
)
from app.monthly_report_native_layout import NativeLayoutError, build_native_docx


def draft(*blocks):
    sections = (
        SectionSpec('activity', '2', 'Monthly Activity Summary', (
            BlockSpec('activity_summary', 'rich_text'),
            BlockSpec('work_orders', 'table', columns=(ColumnSpec('a','Task'),ColumnSpec('b','Status'))),
            BlockSpec('improvements', 'image_grid'))),
        SectionSpec('training', '11', 'Training Summary', (BlockSpec('training_summary','rich_text'),)),
    )
    return ReportDraft(ReportProfile('Synthetic contract','site','New Site',(Facility('site','New Site'),)),
                       ReportPeriod(2026,10),'Current Editor',sections,blocks)


def source(tmp_path):
    document = Document()
    document.sections[0].left_margin = Inches(.91)
    document.add_paragraph('Old Site')
    document.add_paragraph('August 2026')
    document.add_paragraph('Prepared by: Old Editor')
    document.add_heading('MONTHLY ACTIVITY SUMMARY', 1)
    p = document.add_paragraph()
    p.paragraph_format.left_indent = Inches(.37)
    p.add_run('OLD PRIVATE ACTIVITY').bold = True
    t = document.add_table(rows=2, cols=2)
    t.cell(0,0).text='Task'; t.cell(0,1).text='Status'
    t.cell(1,0).text='OLD PRIVATE TASK'; t.cell(1,1).text='OLD STATUS'
    document.add_heading('TRAINING SUMMARY',1)
    document.add_paragraph('OLD PRIVATE TRAINING')
    path=tmp_path/'source.docx'
    document.save(path)
    return path


def test_updates_payload_preserves_geometry_and_cover(tmp_path):
    path=source(tmp_path)
    result=build_native_docx(draft(ResolvedBlock('activity_summary','This month',text='Current activity'),
                                  ResolvedBlock('work_orders','This month',rows=(('New task','Done'),))),path)
    generated=Document(BytesIO(result))
    text='\n'.join(p.text for p in generated.paragraphs)
    assert 'Current activity' in text and 'New Site' in text
    assert 'October 2026' in text and 'Current Editor' in text
    assert 'OLD PRIVATE' not in text and 'Old Site' not in text and 'Old Editor' not in text
    p=next(p for p in generated.paragraphs if p.text=='Current activity')
    assert p.runs[0].bold
    assert p.paragraph_format.left_indent == Document(path).paragraphs[4].paragraph_format.left_indent
    assert generated.sections[0].left_margin == Document(path).sections[0].left_margin
    assert generated.tables[0].cell(1,0).text=='New task'


def test_blank_and_omit_do_not_reuse_source_facts(tmp_path):
    result=build_native_docx(draft(ResolvedBlock('activity_summary','Omit',text='Omitted current text')),source(tmp_path))
    with ZipFile(BytesIO(result)) as archive:
        xml=archive.read('word/document.xml').decode()
    assert 'OLD PRIVATE' not in xml and 'OLD STATUS' not in xml
    assert 'Omitted current text' not in xml
    assert 'MONTHLY ACTIVITY SUMMARY' in xml


def test_preview_uses_logical_section_not_word_section(tmp_path):
    # Both logical sections are in one Word section.
    result=build_native_docx(draft(ResolvedBlock('training_summary','This month',text='Current training')),
                            source(tmp_path),section_key='training')
    document=Document(BytesIO(result))
    text='\n'.join(p.text for p in document.paragraphs)
    assert 'Current training' in text and 'TRAINING SUMMARY' in text
    assert 'MONTHLY ACTIVITY' not in text and 'New Site' not in text
    assert not document.tables


def test_unmatched_table_width_fails_instead_of_truncating(tmp_path):
    value=draft(ResolvedBlock('work_orders','This month',rows=(('a','b','c'),)))
    with pytest.raises(NativeLayoutError,match='columns'):
        build_native_docx(value,source(tmp_path))


def test_unknown_section_fails_explicitly(tmp_path):
    with pytest.raises(NativeLayoutError,match='does not contain'):
        build_native_docx(draft(),source(tmp_path),section_key='water')


def test_old_picture_payload_is_pruned_from_package(tmp_path):
    path=source(tmp_path)
    document=Document(path)
    image=BytesIO(); Image.new('RGB',(31,17),'magenta').save(image,format='PNG')
    p=document.add_paragraph(); p.add_run().add_picture(BytesIO(image.getvalue()))
    document.save(path)
    result=build_native_docx(draft(),path)
    with ZipFile(BytesIO(result)) as archive:
        assert image.getvalue() not in [archive.read(name) for name in archive.namelist()]


def test_global_master_never_copies_source_client_logo_or_photo(tmp_path):
    path=source(tmp_path)
    document=Document(path)
    blobs=[]
    for size,color in (((60,10),'blue'),((30,10),'orange'),((20,15),'purple')):
        data=BytesIO(); Image.new('RGB',size,color).save(data,format='PNG')
        blobs.append(data.getvalue())
        paragraph=document.add_paragraph()
        paragraph.add_run().add_picture(BytesIO(data.getvalue()))
        document.element.body.insert(0,paragraph._p)
    document.save(path)
    result=build_native_docx(draft(),path,master=True)
    with ZipFile(BytesIO(result)) as archive:
        contents=[archive.read(name) for name in archive.namelist()]
    assert blobs[0] in contents  # ENFRA's long brand frame survives.
    assert blobs[1] not in contents and blobs[2] not in contents


def test_current_logo_replaces_native_client_frame(tmp_path):
    path=source(tmp_path)
    document=Document(path)
    old=BytesIO(); Image.new('RGB',(30,10),'orange').save(old,format='PNG')
    new=BytesIO(); Image.new('RGB',(30,10),'green').save(new,format='PNG')
    paragraph=document.add_paragraph(); paragraph.add_run().add_picture(BytesIO(old.getvalue()))
    document.element.body.insert(0,paragraph._p)
    document.save(path)
    result=build_native_docx(draft(ResolvedBlock('client_logo','Library',asset_hashes=('current',))),path,
                            master=True,asset_loader=lambda _:new.getvalue())
    with ZipFile(BytesIO(result)) as archive:
        contents=[archive.read(name) for name in archive.namelist()]
    assert old.getvalue() not in contents and new.getvalue() in contents


def test_additional_text_block_uses_native_style_without_overwriting_other_block(tmp_path):
    value=draft(ResolvedBlock('activity_summary','This month',text='Main activity'),
                ResolvedBlock('extra','This month',text='Extra current text'))
    first=replace(value.sections[0],blocks=(*value.sections[0].blocks,BlockSpec('extra','rich_text')))
    value=replace(value,sections=(first,*value.sections[1:]))
    result=Document(BytesIO(build_native_docx(value,source(tmp_path))))
    text='\n'.join(p.text for p in result.paragraphs)
    assert 'Main activity' in text and 'Extra current text' in text


def test_master_table_header_does_not_carry_source_contacts(tmp_path):
    path=source(tmp_path)
    document=Document(path)
    document.tables[0].cell(0,1).text='Facility Director: Synthetic Person person@example.invalid 555-010-1234'
    document.save(path)
    result=build_native_docx(draft(),path,master=True)
    with ZipFile(BytesIO(result)) as archive:
        xml=archive.read('word/document.xml').decode()
    assert 'person@example.invalid' not in xml and 'Synthetic Person' not in xml
    assert 'Task' in xml and 'Status' in xml


def test_current_columns_adapt_native_table_without_losing_cells(tmp_path):
    value=draft(ResolvedBlock('work_orders','This month',rows=(('Alpha','Beta','Gamma'),)))
    spec=replace(value.sections[0].blocks[1],columns=(ColumnSpec('a','A'),ColumnSpec('b','B'),ColumnSpec('c','C')))
    section=replace(value.sections[0],blocks=(value.sections[0].blocks[0],spec))
    value=replace(value,sections=(section,*value.sections[1:]))
    document=Document(BytesIO(build_native_docx(value,source(tmp_path),master=True)))
    assert [cell.text for cell in document.tables[0].rows[1].cells]==['Alpha','Beta','Gamma']


def test_cleared_list_paragraph_does_not_show_empty_bullets(tmp_path):
    path=source(tmp_path)
    document=Document(path)
    paragraph=document.paragraphs[4]
    paragraph.style='List Bullet'
    document.save(path)
    generated=Document(BytesIO(build_native_docx(draft(),path)))
    paragraph=generated.paragraphs[4]
    assert not paragraph.text
    assert paragraph._p.find('./'+qn('w:pPr')+'/'+qn('w:numPr')+'/'+qn('w:numId')).get(qn('w:val'))=='0'


def test_explicit_unchanged_import_retains_native_runs(tmp_path):
    from app.monthly_report_import import inspect_docx
    path=source(tmp_path)
    document=Document(path)
    paragraph=document.paragraphs[4]
    paragraph.clear(); paragraph.add_run('Original ').italic=True
    paragraph.add_run('current report text').bold=True
    document.save(path)
    inspection=inspect_docx(path)
    item=next(item for item in inspection.items if item.text=='Original current report text')
    block=ResolvedBlock('activity_summary','This month',text=item.text,
                        references=(f'docx:{inspection.sha256}:{item.id}',))
    generated=Document(BytesIO(build_native_docx(draft(block),path)))
    paragraph=next(p for p in generated.paragraphs if p.text==item.text)
    assert len(paragraph.runs)==2
    assert paragraph.runs[0].italic and paragraph.runs[1].bold


def test_stale_import_reference_does_not_authorize_old_source_payload(tmp_path):
    from app.monthly_report_import import inspect_docx
    path=source(tmp_path)
    inspection=inspect_docx(path)
    item=next(item for item in inspection.items if item.text=='OLD PRIVATE ACTIVITY')
    block=ResolvedBlock('activity_summary','This month',text='Changed current text',
                        references=(f'docx:{inspection.sha256}:{item.id}',))
    generated=Document(BytesIO(build_native_docx(draft(block),path)))
    text='\n'.join(p.text for p in generated.paragraphs)
    assert 'Changed current text' in text and 'OLD PRIVATE ACTIVITY' not in text


def test_inspection_associates_early_divider_picture_with_following_heading(tmp_path):
    from docx.enum.section import WD_SECTION_START

    from app.monthly_report_import import inspect_docx
    document=Document()
    document.add_heading('MONTHLY ACTIVITY SUMMARY',1)
    document.add_section(WD_SECTION_START.NEW_PAGE)
    data=BytesIO(); Image.new('RGB',(25,25),'blue').save(data,format='PNG')
    paragraph=document.add_paragraph()
    shape=paragraph.add_run().add_picture(BytesIO(data.getvalue()),width=Inches(8),height=Inches(8))
    shape._inline.tag=qn('wp:anchor')
    for value in ('D6EF4B','547E7E'):
        fill=OxmlElement('a:solidFill');color=OxmlElement('a:srgbClr');color.set('val',value);fill.append(color);shape._inline.append(fill)
    for _ in range(20): document.add_paragraph()
    document.add_heading('SUB-CONTRACTOR STATUS',1)
    path=tmp_path/'early-divider.docx'; document.save(path)
    image=next(item for item in inspect_docx(path).items if item.kind=='image')
    assert image.section=='subcontractors' and image.suggested_slot=='divider_subcontractors'


def test_inspection_photo_grid_captures_cell_captions_not_work_orders(tmp_path):
    from app.monthly_report_import import inspect_docx
    document=Document()
    document.add_heading('MONTHLY ACTIVITY SUMMARY',1)
    document.add_paragraph('Improvement Highlights and Pics')
    table=document.add_table(rows=1,cols=2)
    for index,cell in enumerate(table.rows[0].cells):
        cell.text=f'Supplied caption {index+1}'
        data=BytesIO(); Image.new('RGB',(20,20),('red','green')[index]).save(data,format='PNG')
        cell.paragraphs[0].add_run().add_picture(BytesIO(data.getvalue()))
    path=tmp_path/'photo-grid.docx'; document.save(path)
    items=inspect_docx(path).items
    images=[item for item in items if item.kind=='image']
    assert [item.text for item in images]==['Supplied caption 1','Supplied caption 2']
    assert all(item.suggested_slot=='improvements' for item in images)
    assert not any(item.kind=='table' for item in items)


def test_inspection_routes_end_of_life_separately_from_capital_renewal(tmp_path):
    from app.monthly_report_import import inspect_docx
    document=Document()
    document.add_heading('PRIORITY CAPITAL RENEWAL LIST',1)
    document.add_paragraph('Contract Asset End of Useful Life Schedule')
    table=document.add_table(rows=2,cols=2)
    table.cell(0,0).text='Asset';table.cell(0,1).text='Year'
    table.cell(1,0).text='Boiler';table.cell(1,1).text='2030'
    path=tmp_path/'end-of-life.docx';document.save(path)
    item=next(item for item in inspect_docx(path).items if item.kind=='table')
    assert item.suggested_slot=='end_of_life'


def test_current_image_caption_is_visible_and_does_not_reuse_old_caption(tmp_path):
    from app.monthly_report_import import ImportMapping, inspect_docx, map_items
    document=Document()
    document.add_heading('MONTHLY ACTIVITY SUMMARY',1)
    document.add_paragraph('Improvement Highlights and Pics')
    table=document.add_table(rows=1,cols=1)
    table.cell(0,0).text='Old supplied caption'
    data=BytesIO();Image.new('RGB',(40,30),'red').save(data,format='PNG')
    table.cell(0,0).paragraphs[0].add_run().add_picture(BytesIO(data.getvalue()),width=Inches(2))
    path=tmp_path/'caption-edit.docx';document.save(path)
    inspection=inspect_docx(path)
    item=next(item for item in inspection.items if item.kind=='image')
    mapped=map_items(path,inspection,(ImportMapping(item.id,'improvements'),))
    block=replace(mapped.blocks[0],asset_captions=('Changed current caption',))
    generated=build_native_docx(draft(block),path,asset_loader=dict(mapped.assets).__getitem__)
    with ZipFile(BytesIO(generated)) as archive:
        xml=archive.read('word/document.xml').decode()
    assert 'Changed current caption' in xml and 'Old supplied caption' not in xml


def test_native_footer_uses_current_address(tmp_path):
    path=source(tmp_path)
    document=Document(path)
    document.sections[0].footer.paragraphs[0].text='Old company address | enfrasolutions.com'
    document.save(path)
    value=replace(draft(),address_line='Current company address | enfrasolutions.com')
    generated=Document(BytesIO(build_native_docx(value,path)))
    assert generated.sections[0].footer.paragraphs[0].text=='Current company address | enfrasolutions.com'


def test_large_source_photo_is_not_mistaken_for_following_divider(tmp_path):
    from app.monthly_report_import import inspect_docx
    document=Document()
    document.add_heading('MONTHLY ACTIVITY SUMMARY',1)
    data=BytesIO();Image.new('RGB',(25,25),'red').save(data,format='PNG')
    shape=document.add_paragraph().add_run().add_picture(BytesIO(data.getvalue()),width=Inches(8),height=Inches(8))
    shape._inline.tag=qn('wp:anchor')
    document.add_heading('MONTHLY SCORECARDS',1)
    path=tmp_path/'large-source-photo.docx';document.save(path)
    item=next(item for item in inspect_docx(path).items if item.kind=='image')
    assert item.section=='activity' and item.suggested_slot=='improvements'
    result=build_native_docx(draft(),path,master=True)
    with ZipFile(BytesIO(result)) as archive:
        assert data.getvalue() not in [archive.read(name) for name in archive.namelist()]


def test_source_image_fingerprint_cache_invalidates_for_new_verified_source(monkeypatch):
    from types import SimpleNamespace

    from app.monthly_report_import import ImportItem
    from app.monthly_report_native_layout import _source_image_digest
    calls=[]
    def normalize(path,item,*,line_art):
        calls.append((path,item,line_art))
        return SimpleNamespace(data=b'synthetic normalized pixels',extension='png')
    monkeypatch.setattr('app.monthly_report_import.read_import_image',normalize)
    item=ImportItem('item-1','image','word/document.xml',1,1,image_part='word/media/picture.png')
    _source_image_digest.cache_clear()
    try:
        first=_source_image_digest('/verified/source.docx','source-sha-a',item,False)
        assert _source_image_digest('/verified/source.docx','source-sha-a',item,False)==first
        assert len(calls)==1
        _source_image_digest('/verified/source.docx','source-sha-b',item,False)
        assert len(calls)==2
        _source_image_digest('/verified/source.docx','source-sha-b',item,True)
        assert len(calls)==3
    finally:
        _source_image_digest.cache_clear()


def test_native_starting_import_omits_master_dividers_but_keeps_unknown_art_questions():
    from app.monthly_report_import import DocxInspection, ImportItem
    from app.monthly_report_start_import import (
        StartingReportAnalysis,
        automatic_section_plans,
    )
    inspection=DocxInspection('a'*64,(
        ImportItem('divider','image','word/document.xml',2,1,'activity','divider_activity',image_part='word/media/art.png',image_width=1000,image_height=1000),
        ImportItem('unknown','unsupported','word/document.xml',5,2,'issues',note='Unrecognized source drawing'),
    ),(),(),2)
    analysis=StartingReportAnalysis((),'same',(),(),())
    generic=automatic_section_plans(inspection,analysis)
    native=automatic_section_plans(inspection,analysis,native_design=True)
    assert 'divider' in generic['activity']['selected']
    assert 'divider' not in native['activity']['selected']
    assert native['issues']['questions']


def test_exact_filtered_table_retains_native_grid_without_prices(tmp_path):
    from app.monthly_report_import import inspect_docx
    from app.monthly_report_sections import table_without_prices
    doc = Document()
    doc.add_heading('MONTHLY ACTIVITY SUMMARY', 1)
    table = doc.add_table(rows=2, cols=3)
    for cell, value in zip(table.rows[0].cells, ('Task', 'Cost', 'Status')):
        cell.text = value
    for cell, value in zip(table.rows[1].cells, ('Supplied task', '$123.00', 'Done')):
        cell.text = value
    table.columns[0].width = Inches(2.1)
    path = tmp_path / 'priced.docx'; doc.save(path)
    inspection = inspect_docx(path)
    item = next(item for item in inspection.items if item.kind == 'table')
    supplied, removed = table_without_prices(item)
    assert removed == (1,)
    supplied = replace(supplied, reference=f'docx:{inspection.sha256}:{item.id}')
    block = ResolvedBlock('work_orders', 'This month', extra_tables=(supplied,))
    result = Document(BytesIO(build_native_docx(draft(block), path)))
    actual = result.tables[0]
    assert len(actual.columns) == 3
    assert actual.cell(0,1).text == actual.cell(1,1).text == ''
    assert actual.cell(1,0).text == 'Supplied task'
    assert actual.cell(1,2).text == 'Done'
    assert actual.columns[0].width == Document(path).tables[0].columns[0].width


def test_footer_import_excludes_cached_page_fields(tmp_path):
    from app.monthly_report_import import inspect_docx
    doc = Document()
    p = doc.sections[0].footer.paragraphs[0]
    p.add_run('ENFRA address | enfrasolutions.com')
    for kind in ('begin', 'separate', 'end'):
        if kind == 'separate':
            instruction = OxmlElement('w:instrText'); instruction.text = ' PA'
            p.add_run()._r.append(instruction)
            instruction = OxmlElement('w:instrText'); instruction.text = 'GE '
            p.add_run()._r.append(instruction)
        node = OxmlElement('w:fldChar'); node.set(qn('w:fldCharType'), kind)
        p.add_run()._r.append(node)
        if kind == 'separate': p.add_run('2')
    simple = OxmlElement('w:fldSimple'); simple.set(qn('w:instr'), ' NUMPAGES ')
    run = OxmlElement('w:r'); value = OxmlElement('w:t'); value.text = '46'; run.append(value); simple.append(run); p._p.append(simple)
    path = tmp_path / 'footer.docx'; doc.save(path)
    inspection = inspect_docx(path)
    footer = next(item for item in inspection.items if item.suggested_slot == 'footer_text')
    assert footer.text == 'ENFRA address | enfrasolutions.com'
    from app.monthly_report_import import ImportMapping, imported_draft, map_items
    mapped = map_items(path, inspection, (ImportMapping(footer.id, 'footer_text'),))
    current = draft()
    imported = imported_draft(current.profile, current.period, current.prepared_by, mapped)
    assert imported.address_line == footer.text
    assert '2' in Document(path).sections[0].footer.paragraphs[0].text


def test_verified_cover_brand_keeps_native_header_variant_changed_brand_replaces(tmp_path):
    import hashlib

    from app.monthly_report_import import inspect_docx, read_import_image
    doc = Document()
    cover = BytesIO(); Image.new('RGB', (600, 80), 'green').save(cover, format='PNG')
    header = BytesIO(); Image.new('RGB', (600, 80), 'black').save(header, format='PNG')
    changed = BytesIO(); Image.new('RGB', (600, 80), 'red').save(changed, format='PNG')
    doc.add_picture(BytesIO(cover.getvalue()), width=Inches(3))
    doc.sections[0].header.paragraphs[0].add_run().add_picture(BytesIO(header.getvalue()), width=Inches(3))
    tagline = BytesIO(); Image.new('RGB', (600, 80), 'orange').save(tagline, format='PNG')
    doc.add_section()
    doc.add_picture(BytesIO(tagline.getvalue()), width=Inches(3))
    doc.add_heading('MONTHLY ACTIVITY SUMMARY', 1)
    path = tmp_path/'brand.docx'; doc.save(path)
    inspection = inspect_docx(path)
    item = next(i for i in inspection.items if i.kind == 'image' and i.part == 'word/document.xml')
    normalized = read_import_image(path, item, line_art=True)
    reference = hashlib.sha256(normalized.data).hexdigest()+'.'+normalized.extension
    current = draft(ResolvedBlock('brand_logo', 'Library', asset_hashes=(reference,)))
    output = Document(BytesIO(build_native_docx(current, path, asset_loader=lambda _: normalized.data)))
    assert any(p.blob == header.getvalue() for p in output.sections[0].header.part.related_parts.values())
    assert output.part.related_parts[output.inline_shapes[1]._inline.graphic.graphicData.pic.blipFill.blip.embed].blob == tagline.getvalue()
    current = draft(ResolvedBlock('brand_logo', 'Library', asset_hashes=('changed',)))
    output = Document(BytesIO(build_native_docx(current, path, asset_loader=lambda _: changed.getvalue())))
    assert any(p.blob == changed.getvalue() for p in output.sections[0].header.part.related_parts.values())
    assert not any(p.blob == header.getvalue() for p in output.sections[0].header.part.related_parts.values())
    assert output.part.related_parts[output.inline_shapes[1]._inline.graphic.graphicData.pic.blipFill.blip.embed].blob == changed.getvalue()


def test_native_master_allows_empty_and_pending_sections_without_relaxing_content_gates():
    from app.monthly_report_checks import preflight
    from app.monthly_report_model import default_sections
    current = draft()
    current = replace(current, profile=replace(current.profile, template='enfra_master:'+'a'*64),
                      sections=default_sections(), blocks=())
    assert not any(c.code == 'required' for c in preflight(current))
    pending = replace(current, blocks=(ResolvedBlock('org_chart', 'This month', text='Pending data.'),))
    assert not any(c.code == 'required' for c in preflight(pending))
    unsafe = replace(pending, blocks=(ResolvedBlock('activity_summary', 'This month', text='Cost: $500'),))
    assert any(c.code == 'pricing' and c.blocking for c in preflight(unsafe))
    legacy = replace(current, profile=replace(current.profile, template=''))
    assert any(c.code == 'required' and c.blocking for c in preflight(legacy))


@pytest.mark.parametrize('offset_kind', ['span', 'before'])
def test_price_redaction_uses_logical_grid_after_merged_or_omitted_cells(tmp_path, offset_kind):
    from app.monthly_report_import import inspect_docx
    from app.monthly_report_sections import table_without_prices
    document = Document()
    document.add_heading('MONTHLY ACTIVITY SUMMARY', 1)
    table = document.add_table(rows=2, cols=3)
    table.cell(0, 0).text = 'Task'
    table.cell(0, 1).text = 'Detail'
    table.cell(0, 2).text = 'Cost'
    if offset_kind == 'span':
        table.cell(1, 0).merge(table.cell(1, 1)).text = 'Supplied task'
    else:
        row = table.rows[1]._tr
        row.remove(row.findall(qn('w:tc'))[0])
        properties = OxmlElement('w:trPr')
        before = OxmlElement('w:gridBefore'); before.set(qn('w:val'), '1')
        properties.append(before); row.insert(0, properties)
        row.findall(qn('w:tc'))[0].find(qn('w:p')).append(OxmlElement('w:r'))
        from docx.table import _Cell
        _Cell(row.findall(qn('w:tc'))[0], table).text = 'Supplied task'
    from docx.table import _Cell
    _Cell(table.rows[1]._tr.findall(qn('w:tc'))[-1], table).text = '$123.00'
    path = tmp_path / 'logical-price.docx'; document.save(path)
    inspection = inspect_docx(path)
    item = next(item for item in inspection.items if item.kind == 'table')
    supplied, removed = table_without_prices(item)
    assert removed == (2,)
    supplied = replace(supplied, reference=f'docx:{inspection.sha256}:{item.id}')
    output = build_native_docx(draft(ResolvedBlock('work_orders', 'This month', extra_tables=(supplied,))), path)
    with ZipFile(BytesIO(output)) as archive:
        xml = archive.read('word/document.xml')
    assert b'$123.00' not in xml and b'>Cost<' not in xml
    assert b'Supplied task' in xml


def test_blank_master_does_not_keep_client_header_beside_safe_page_field(tmp_path):
    document = Document()
    document.add_heading('MONTHLY ACTIVITY SUMMARY', 1)
    paragraph = document.sections[0].header.paragraphs[0]
    paragraph.add_run('OLD PRIVATE CLIENT ')
    for kind in ('begin', 'separate', 'end'):
        marker = OxmlElement('w:fldChar'); marker.set(qn('w:fldCharType'), kind)
        paragraph.add_run()._r.append(marker)
        if kind == 'begin':
            instruction = OxmlElement('w:instrText'); instruction.text = ' PAGE '
            paragraph.add_run()._r.append(instruction)
        if kind == 'separate':
            paragraph.add_run('1')
    path = tmp_path / 'client-header.docx'; document.save(path)
    output = build_native_docx(draft(), path, master=True)
    with ZipFile(BytesIO(output)) as archive:
        header = archive.read('word/header1.xml')
    assert b'OLD PRIVATE CLIENT' not in header


def test_current_activity_keeps_period_heading_and_uses_narrative_anchor(tmp_path):
    doc = Document()
    doc.add_heading('MONTHLY ACTIVITY SUMMARY', 1)
    doc.add_paragraph('August 2026 Activity')
    doc.add_paragraph('Activity Summary')
    doc.add_paragraph('Old activity narrative')
    path = tmp_path/'activity.docx'; doc.save(path)
    current = draft(ResolvedBlock('activity_summary', 'This month', text='Current activity narrative'))
    result = Document(BytesIO(build_native_docx(current, path)))
    values = [p.text for p in result.paragraphs]
    assert 'October 2026 Activity' in values
    assert values.index('Activity Summary') < values.index('Current activity narrative')
    assert 'Old activity narrative' not in values


def test_section_preview_does_not_open_empty_trailing_word_section(tmp_path):
    from docx.enum.section import WD_SECTION_START
    doc = Document()
    doc.add_heading('MONTHLY ACTIVITY SUMMARY', 1)
    doc.add_paragraph('Previous activity')
    doc.add_section(WD_SECTION_START.NEW_PAGE)
    doc.add_heading('TRAINING SUMMARY', 1)
    doc.add_paragraph('Other section data')
    path = tmp_path/'sections.docx'; doc.save(path)
    current = draft(ResolvedBlock('activity_summary', 'This month', text='Current activity'))
    output = Document(BytesIO(build_native_docx(current, path, section_key='activity')))
    assert len(output.sections) == 1
    assert 'Current activity' in '\n'.join(p.text for p in output.paragraphs)
    assert 'Other section data' not in '\n'.join(p.text for p in output.paragraphs)


@pytest.mark.parametrize('merged_start,merged_text', [(1, 'Replace pump after inspection'), (2, '$123.00')])
def test_filtered_table_preserves_safe_cross_column_merge_geometry(tmp_path, merged_start, merged_text):
    from app.monthly_report_import import inspect_docx
    from app.monthly_report_sections import table_without_prices
    document = Document()
    document.add_heading('MONTHLY ACTIVITY SUMMARY', 1)
    table = document.add_table(rows=3, cols=4)
    for cell, value in zip(table.rows[0].cells, ('Task', 'Description', 'Cost', 'Status')):
        cell.text = value
    for cell, value in zip(table.rows[1].cells, ('Pump 1', 'Inspection', '$456.00', 'Done')):
        cell.text = value
    table.cell(2, 0).text = 'Pump 2'
    table.cell(2, merged_start).merge(table.cell(2, merged_start + 1)).text = merged_text
    table.columns[1].width = Inches(3.1)
    path = tmp_path / 'merged-priced.docx'
    document.save(path)
    inspection = inspect_docx(path)
    item = next(item for item in inspection.items if item.kind == 'table')
    supplied, removed = table_without_prices(item)
    assert removed == (2,)
    supplied = replace(supplied, reference=f'docx:{inspection.sha256}:{item.id}')
    output = build_native_docx(draft(ResolvedBlock('work_orders', 'This month', extra_tables=(supplied,))), path)
    actual = Document(BytesIO(output)).tables[0]
    assert actual._tbl.tblGrid.xml == Document(path).tables[0]._tbl.tblGrid.xml
    merged = actual.rows[2]._tr.findall(qn('w:tc'))[merged_start]
    assert merged.find('./' + qn('w:tcPr') + '/' + qn('w:gridSpan')).get(qn('w:val')) == '2'
    with ZipFile(BytesIO(output)) as archive:
        xml = archive.read('word/document.xml')
    assert b'$123.00' not in xml and b'$456.00' not in xml and b'>Cost<' not in xml
    if merged_start == 1:
        assert b'Replace pump after inspection' in xml


def test_master_clears_source_drawing_descriptions_in_body_and_header(tmp_path):
    image = BytesIO()
    Image.new('RGB', (300, 200), 'green').save(image, format='PNG')
    document = Document()
    document.add_paragraph('Old Site')
    document.add_picture(BytesIO(image.getvalue()), width=Inches(2))
    header = document.sections[0].header.paragraphs[0]
    header.add_run().add_picture(BytesIO(image.getvalue()), width=Inches(1))
    for element in (document.element, document.sections[0].header._element):
        for node in element.iter():
            if node.tag in (qn('wp:docPr'), qn('pic:cNvPr')):
                node.set('descr', 'OLD CLIENT DESCRIPTION')
                node.set('title', 'OLD CLIENT TITLE')
                node.set('name', 'OLD CLIENT FILE NAME')
    from lxml import etree
    pict = OxmlElement('w:pict')
    shape = etree.SubElement(pict, '{urn:schemas-microsoft-com:vml}rect',
                             id='safe-frame-id', style='width:20pt;height:20pt', fillcolor='#557f7f')
    shape.set('alt', 'OLD CLIENT ALT')
    shape.set('title', 'OLD CLIENT VML TITLE')
    shape.set('{urn:schemas-microsoft-com:office:office}title', 'OLD CLIENT OFFICE TITLE')
    header.add_run()._r.append(pict)
    document.add_heading('MONTHLY ACTIVITY SUMMARY', 1)
    document.add_paragraph('Old activity')
    path = tmp_path / 'described-master.docx'
    document.save(path)
    output = build_native_docx(draft(), path, master=True)
    with ZipFile(BytesIO(output)) as archive:
        all_xml = b'\n'.join(archive.read(name) for name in archive.namelist() if name.endswith('.xml'))
        assert b'OLD CLIENT' not in all_xml
        header_xml = archive.read('word/header1.xml')
        assert b'safe-frame-id' in header_xml
        assert b'width:20pt;height:20pt' in header_xml
        assert b'Drawing ' in header_xml
