"""F12: paired-reference geometric proof for the proposals divider's '10'."""
from copy import deepcopy
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.shared import Inches
from lxml import etree
import pytest

from app.monthly_report_divider_numerals import MC, VML, W, WPS, fit_proposals_numeral
from conftest import requires_libreoffice

WP = '{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}'
A = '{http://schemas.openxmlformats.org/drawingml/2006/main}'


def _paragraph(parent):
    paragraph = etree.SubElement(parent, W + 'p')
    for digit in '10':
        run = etree.SubElement(paragraph, W + 'r')
        props = etree.SubElement(run, W + 'rPr')
        etree.SubElement(props, W + 'b')
        etree.SubElement(props, W + 'color', attrib={W + 'val': 'D6EF4B'})
        etree.SubElement(props, W + 'spacing', attrib={W + 'val': '-400'})
        etree.SubElement(props, W + 'sz', attrib={W + 'val': '420'})
        etree.SubElement(run, W + 't').text = digit
    return paragraph


def _number_element():
    # Public synthetic layout: geometry only, no client-specific pictures/art.
    node = etree.Element(MC + 'AlternateContent', nsmap={
        'mc': MC[1:-1], 'w': W[1:-1], 'wp': WP[1:-1],
        'a': A[1:-1], 'wps': WPS[1:-1], 'v': VML[1:-1]})
    choice = etree.SubElement(node, MC + 'Choice', Requires='wps')
    drawing = etree.SubElement(choice, W + 'drawing')
    anchor = etree.SubElement(drawing, WP + 'anchor',
        distT='0', distB='0', distL='114300', distR='114300',
        simplePos='0', relativeHeight='251658263', behindDoc='0',
        locked='0', layoutInCell='1', allowOverlap='1')
    etree.SubElement(anchor, WP + 'simplePos', x='0', y='0')
    horizontal = etree.SubElement(anchor, WP + 'positionH', relativeFrom='column')
    etree.SubElement(horizontal, WP + 'posOffset').text = '-127000'
    vertical = etree.SubElement(anchor, WP + 'positionV', relativeFrom='page')
    etree.SubElement(vertical, WP + 'posOffset').text = '5596885'
    etree.SubElement(anchor, WP + 'extent', cx='2574290', cy='3657600')
    etree.SubElement(anchor, WP + 'effectExtent', l='0', t='0', r='0', b='0')
    etree.SubElement(anchor, WP + 'wrapNone')
    etree.SubElement(anchor, WP + 'docPr', id='999', name='Synthetic 10')
    graphic = etree.SubElement(anchor, A + 'graphic')
    data = etree.SubElement(graphic, A + 'graphicData', uri=WPS[1:-1])
    shape = etree.SubElement(data, WPS + 'wsp')
    etree.SubElement(shape, WPS + 'cNvSpPr', txBox='1')
    prop = etree.SubElement(shape, WPS + 'spPr')
    transform = etree.SubElement(prop, A + 'xfrm')
    etree.SubElement(transform, A + 'off', x='0', y='0')
    etree.SubElement(transform, A + 'ext', cx='2574290', cy='3657600')
    geometry = etree.SubElement(prop, A + 'prstGeom', prst='rect')
    etree.SubElement(geometry, A + 'avLst')
    etree.SubElement(prop, A + 'noFill')
    textbox = etree.SubElement(shape, WPS + 'txbx')
    _paragraph(etree.SubElement(textbox, W + 'txbxContent'))
    body = etree.SubElement(shape, WPS + 'bodyPr', vert='horz', wrap='square',
        lIns='0', tIns='12700', rIns='0', bIns='0', anchor='t')
    etree.SubElement(body, A + 'noAutofit')
    fallback = etree.SubElement(node, MC + 'Fallback')
    pict = etree.SubElement(fallback, W + 'pict')
    legacy = etree.SubElement(pict, VML + 'shape', id='_synthetic_number',
        type='#_x0000_t202', style=(
            'position:absolute;margin-left:-10pt;margin-top:440.7pt;'
            'width:202.7pt;height:4in;visibility:visible;v-text-anchor:top'))
    _paragraph(etree.SubElement(etree.SubElement(legacy, VML + 'textbox',
        inset='0,1pt,0,0'), W + 'txbxContent'))
    return node


def _synthetic_doc():
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.left_margin = section.right_margin = Inches(0)
    section.top_margin = section.bottom_margin = Inches(0)
    doc.add_paragraph().add_run()._r.append(_number_element())
    doc.add_paragraph('SYNTHETIC PROPOSALS DIVIDER')
    return doc


def _geometry(document):
    anchor = next(document.element.body.iter(WP + 'anchor'))
    shape = next(document.element.body.iter(VML + 'shape'))
    return (anchor.find(WP + 'extent').get('cx'),
            anchor.find(WP + 'positionV').find(WP + 'posOffset').text,
            shape.get('style'))


def test_divider_tight_two_digit_has_a_measured_full_width_and_vertical_position():
    doc = _synthetic_doc()
    assert _geometry(doc)[0] == '2574290'  # Fails pre-fix: terminal tracking clips.
    assert fit_proposals_numeral(doc, section_key='proposals')
    width, vertical, fallback = _geometry(doc)
    assert width == str(2574290 + 20 * 12700)
    assert vertical == str(5596885 + 5 * 12700)
    assert 'width:222.7pt' in fallback and 'margin-top:445.7pt' in fallback
    assert not fit_proposals_numeral(doc, section_key='proposals')  # idempotent


def test_alternate_branches_receive_identical_measured_digit_geometry():
    doc = _synthetic_doc()
    assert fit_proposals_numeral(doc, section_key='proposals')
    choice = next(doc.element.body.iter(MC + 'Choice'))
    extent = next(choice.iter(A + 'ext'))
    anchor = next(choice.iter(WP + 'extent'))
    assert extent.get('cx') == anchor.get('cx') == str(2574290 + 20 * 12700)
    assert extent.get('cy') == anchor.get('cy') == '3657600'
    assert [node.text for node in doc.element.body.iter(W + 't')].count('1') == 2
    assert [node.text for node in doc.element.body.iter(W + 't')].count('0') == 2


def test_proposals_number_is_not_rewritten_in_other_section_preview():
    doc = _synthetic_doc()
    before = BytesIO(); doc.save(before)
    assert not fit_proposals_numeral(doc, section_key='training')
    after = BytesIO(); doc.save(after)
    assert before.getvalue() == after.getvalue()


def test_ambiguous_or_nonmatching_design_retains_native_geometry():
    for mode in ('fallback-mismatch', 'font-change', 'missing-branch', 'two-live-frames'):
        doc = _synthetic_doc()
        original = _geometry(doc)
        alternate = next(doc.element.body.iter(MC + 'AlternateContent'))
        if mode == 'fallback-mismatch':
            next(alternate.iter(VML + 'shape')).set('style', 'width:244pt;height:4in')
        elif mode == 'font-change':
            next(alternate.iter(W + 'spacing')).set(W + 'val', '-200')
        elif mode == 'missing-branch':
            alternate.remove(next(alternate.iter(MC + 'Fallback')))
        else:
            doc.paragraphs[0]._p[0].append(deepcopy(alternate))
        current = etree.tostring(doc.element.body)
        assert not fit_proposals_numeral(doc, section_key='proposals')
        assert etree.tostring(doc.element.body) == current


def test_no_regular_text_or_unrelated_header_is_altered():
    doc = _synthetic_doc()
    doc.add_paragraph('Customer service request number 10, current text.')
    old = ''.join(p.text for p in doc.paragraphs)
    assert fit_proposals_numeral(doc, section_key='proposals')
    assert ''.join(p.text for p in doc.paragraphs) == old


@requires_libreoffice
def test_measured_divider_has_both_digits_in_real_pdf(tmp_path, monkeypatch):
    from app import pdf_converter
    import pymupdf
    monkeypatch.setenv('EPC_DATA_DIR', str(tmp_path / 'runtime'))
    monkeypatch.setattr(pdf_converter, 'OUTPUT_DIR', tmp_path)
    doc = _synthetic_doc()
    assert fit_proposals_numeral(doc, section_key='proposals')
    path = tmp_path / 'synthetic-num.docx'; doc.save(path)
    result = pdf_converter.convert_to_pdf(path)
    try:
        with pymupdf.open(result) as pdf:
            assert len(pdf) == 1
            text = pdf[0].get_text()
            assert '10' in text and 'SYNTHETIC PROPOSALS' in text
            assert len(pdf[0].search_for('10')) == 1
    finally:
        result.unlink(missing_ok=True)
