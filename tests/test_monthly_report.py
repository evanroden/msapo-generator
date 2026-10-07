from dataclasses import replace
from datetime import date
from io import BytesIO
from zipfile import ZipFile

from docx import Document
import pytest

from app.monthly_report_checks import placeholder_matches, preflight, stale_period_mentions
from app.monthly_report_docx import SHELL_PATH, assemble_docx, create_shell, generate_report, outline
from app.monthly_report_model import (
    BlockSpec, ColumnSpec, Facility, PLACEHOLDER_PHRASES, ReportPeriod, ReportProfile,
    ResolvedBlock, default_sections, included_sections, report_filename,
    synthetic_draft, synthetic_profiles,
)
from tests.conftest import requires_libreoffice


@pytest.fixture
def draft():
    return synthetic_draft(synthetic_profiles()[0], ReportPeriod(2026, 9))


def test_scope_is_explicit_and_aliases_do_not_expand_membership():
    profiles = synthetic_profiles()
    assert [p.scope_type for p in profiles] == ["individual", "multi_site", "regional"]
    assert [len(p.facilities) for p in profiles] == [1, 2, 2]
    assert profiles[1].facilities == profiles[2].facilities
    with pytest.raises(ValueError, match="unique"):
        replace(profiles[0], facilities=(profiles[0].facilities[0],) * 2)
    with pytest.raises(ValueError, match="exactly one"):
        replace(profiles[0], facilities=profiles[1].facilities)
    assert ReportProfile("Demo", "region", "Demonstration Region", (Facility("a", "A"),), "regional")


def test_previous_month_and_leap_year():
    assert ReportPeriod.previous(date(2026, 1, 1)) == ReportPeriod(2025, 12)
    assert ReportPeriod(2024, 2).end == date(2024, 2, 29)
    with pytest.raises(ValueError):
        ReportPeriod(2026, 13)


def test_skeleton_order_and_renumbering():
    sections = default_sections()
    assert len(sections) == 12
    assert [s.number for s in included_sections(sections)] == list(map(str, range(1, 12)))
    reordered = included_sections((sections[3], replace(sections[0], included=False), sections[1], replace(sections[-1], included=True)))
    assert [(s.key, s.number) for s in reordered] == [("mbcx", "1"), ("activity", "2"), ("rfi", "G")]


@pytest.mark.parametrize("phrase", PLACEHOLDER_PHRASES)
def test_every_instruction_phrase_is_blocked(phrase):
    assert placeholder_matches("Heading\n" + phrase.upper().replace(" ", "\n") + " here")


def test_normal_prose_is_not_placeholder_text():
    assert not placeholder_matches("The technician included the summary. Training hours not provided. Author: Synthetic Operator.")
    assert not placeholder_matches("Twenty-six hours of training.")


@pytest.mark.parametrize("text", [
    "August 2026 Activity", "September 2025 Activity", "September 10, 2025",
    "09/10/2025", "2026-08-10", "Since July. August activity.",
])
def test_stale_periods_include_dates_with_correct_month_wrong_year(text):
    assert stale_period_mentions(text, 2026, 9)


@pytest.mark.parametrize("text", [
    "September 2026 activity", "September 10, 2026", "09/10/2026", "2026-09-10",
    "Ongoing since July", "The pump may need repair.",
])
def test_current_and_local_historical_references(text):
    assert not stale_period_mentions(text, 2026, 9)


def test_required_review_and_table_cell_checks(draft):
    missing = replace(draft, blocks=())
    assert any(c.code == "required" for c in preflight(missing))
    ai = replace(draft.blocks[0], ai_written=True, text="Synthetic verified work.", references=("demo#1 p1",))
    updated = replace(draft, blocks=(ai, *draft.blocks[1:]))
    assert any(c.code == "review" for c in preflight(updated))
    reviewed = replace(ai, reviewed_fingerprint=ai.fingerprint)
    assert reviewed.reviewed
    assert not replace(reviewed, references=("demo#2 p1",)).reviewed
    assert not replace(reviewed, text="Revised.").reviewed
    assert any(c.code == "placeholder" for c in preflight(replace(draft, blocks=(
        replace(draft.blocks[0], rows=(("Author: Name",),)), *draft.blocks[1:],
    ))))


def test_generation_enforces_preflight_and_content_bound_acknowledgement(draft):
    bad = replace(draft, prepared_by="")
    with pytest.raises(ValueError, match="prepared"):
        assemble_docx(bad)
    stale = replace(draft, blocks=(replace(draft.blocks[0], text="August 2026 activity"), *draft.blocks[1:]))
    with pytest.raises(ValueError, match="Acknowledge"):
        assemble_docx(stale, acknowledged_fingerprint=draft.fingerprint)
    assert assemble_docx(stale, acknowledged_fingerprint=stale.fingerprint).startswith(b"PK")


def test_cover_toc_dividers_filename_and_deterministic_bytes(draft):
    draft = replace(draft, sections=(draft.sections[2], draft.sections[0]))
    first = assemble_docx(draft)
    assert first == assemble_docx(draft)
    document = Document(BytesIO(first))
    texts = [p.text for p in document.paragraphs]
    assert "Operations and Maintenance Monthly Review\nSeptember 2026" in texts
    assert report_filename(draft, "docx") == "Demonstration - Demonstration Individual Facility September 2026 Monthly Report.docx"
    assert [p.text for p in document.paragraphs if p.text.startswith(("1. ", "2. "))] == [
        "1. Monthly Scorecards", "2. Organizational Chart", "1. Monthly Scorecards", "2. Organizational Chart",
    ]
    assert len(document.sections) == 6  # Cover, contents, two divider/content pairs.
    for index, section in enumerate(document.sections):
        assert section.page_width.inches == 8.5 and section.page_height.inches == 11
        if index % 2 == 0:
            assert not section.header.paragraphs[0].text
            assert not section.footer.paragraphs[0].text
        else:
            assert "PAGE" in section.footer._element.xml
        assert "w:start=" not in section._sectPr.xml  # Numbering stays continuous.
    assert len(outline(draft)) == 2


def test_shell_is_reproducible_and_has_no_client_content_or_images(tmp_path):
    recreated = tmp_path / "shell.docx"
    create_shell(recreated)
    assert recreated.read_bytes() == SHELL_PATH.read_bytes()
    with ZipFile(SHELL_PATH) as archive:
        assert not any("media/" in name or "thumbnail" in name for name in archive.namelist())
        assert not any("embeddings/" in name or "vba" in name for name in archive.namelist())
        assert all(info.date_time == (2000, 1, 1, 0, 0, 0) for info in archive.infolist())
    shell = Document(SHELL_PATH)
    assert not any(p.text for p in shell.paragraphs)
    assert shell.core_properties.author == "Email Process Control"
    assert len(shell.tables) == 0


def test_table_values_are_not_interpreted_as_markup(draft):
    spec = BlockSpec("table", "table", True, columns=(ColumnSpec("item", "Item"), ColumnSpec("status", "Status")))
    section = replace(draft.sections[0], blocks=(spec,))
    block = ResolvedBlock("table", "This month", rows=(("<synthetic>", "Pending & checked"),))
    document = Document(BytesIO(assemble_docx(replace(draft, sections=(section,), blocks=(block,)))))
    assert document.tables[0].cell(1, 0).text == "<synthetic>"
    assert document.tables[0].cell(1, 1).text == "Pending & checked"


def test_pdf_failure_retains_docx_and_removes_partial_output(monkeypatch, tmp_path, draft):
    from app import pdf_converter
    paths = []

    def fail(path):
        paths.append(path)
        (tmp_path / (path.stem + ".pdf")).write_bytes(b"partial")
        raise pdf_converter.PDFConversionError("Synthetic converter failure")

    monkeypatch.setattr(pdf_converter, "OUTPUT_DIR", tmp_path)
    monkeypatch.setattr(pdf_converter, "convert_to_pdf", fail)
    package = generate_report(draft)
    assert Document(BytesIO(package.docx))
    assert package.pdf is None
    assert "DOCX is ready" in package.pdf_error
    assert not paths[0].exists()
    assert list(tmp_path.iterdir()) == []


@requires_libreoffice
def test_real_pdf_has_expected_pages_cover_and_chrome(draft):
    import fitz

    package = generate_report(draft)
    assert package.pdf, package.pdf_error
    with fitz.open(stream=package.pdf, filetype="pdf") as pdf:
        assert len(pdf) == 2 + sum(item[2] for item in outline(draft))
        assert draft.period.label in pdf[0].get_text()
        assert "SYNTHETIC" in pdf[0].get_text()
        for index, page in enumerate(pdf):
            assert "Page " not in page.get_text() if index % 2 == 0 else "Page " in page.get_text()
