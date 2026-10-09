from dataclasses import replace
from pathlib import Path

from docx import Document
import pytest

from app import monthly_report_designs as designs, monthly_report_library as library
from app.monthly_report_model import synthetic_profiles


def reference(tmp_path, name="source", title="Monthly Activity Summary"):
    path = tmp_path / (name + ".docx")
    doc = Document()
    doc.add_paragraph("Operations and Maintenance Monthly Review")
    doc.add_heading(title, 1)
    doc.add_paragraph("Source-only content")
    doc.save(path)
    return path


def test_contract_default_and_pinned_versions(tmp_path, monkeypatch):
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path / "data"))
    profile = synthetic_profiles()[0]
    source = reference(tmp_path)
    first = designs.install(source, profile, actor="Editor", expected_revision=0)
    assert designs.source_for(first).read_bytes() == source.read_bytes()
    assert designs.pin(profile) == first
    second = designs.install(reference(tmp_path, "second", "Training Summary"), profile,
                             actor="Editor", expected_revision=1)
    assert designs.fingerprint(second) != designs.fingerprint(first)
    assert designs.pin(profile) == second
    assert designs.source_for(first).name == designs.fingerprint(first) + ".docx"
    other_site = replace(profile, key="another-site")
    assert designs.fingerprint(other_site) == designs.fingerprint(first)
    assert designs.source_for(replace(other_site, contract="Other contract")) is None
    assert len(designs.state(profile.contract)["history"]) == 2


def test_revision_and_attribution_guards(tmp_path, monkeypatch):
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path / "data"))
    profile = synthetic_profiles()[0]
    source = reference(tmp_path)
    with pytest.raises(ValueError):
        designs.install(source, profile, actor="", expected_revision=0)
    designs.install(source, profile, actor="Editor", expected_revision=0)
    with pytest.raises(library.RevisionConflict):
        designs.install(source, profile, actor="Editor", expected_revision=0)
    designs.install(source, profile, actor="Editor", expected_revision=1)
    assert designs.state(profile.contract)["revision"] == 1


def test_missing_or_corrupt_pinned_reference_never_falls_back(tmp_path, monkeypatch):
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path / "data"))
    profile = synthetic_profiles()[0]
    saved = designs.install(reference(tmp_path), profile, actor="Editor", expected_revision=0)
    path = designs.source_for(saved)
    path.unlink()
    path.write_bytes(b"changed source")
    with pytest.raises(ValueError, match="integrity"):
        designs.source_for(saved)
    path.unlink()
    with pytest.raises(ValueError, match="unavailable"):
        designs.source_for(saved)


def test_old_profile_import_can_supply_pinned_design_only_within_contract(tmp_path, monkeypatch):
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path / "data"))
    import hashlib
    source = reference(tmp_path)
    digest = hashlib.sha256(source.read_bytes()).hexdigest()
    profile = replace(synthetic_profiles()[0], imported_from=digest)
    path = library._profile_path(profile.contract, profile.key) / "imports" / (digest + ".docx")
    path.parent.mkdir(parents=True)
    path.write_bytes(source.read_bytes())
    pinned = designs.pin(profile)
    assert designs.source_for(pinned) == path
    reused = replace(pinned, key="new-site", imported_from="")
    assert designs.source_for(reused) == path
    with pytest.raises(ValueError, match="unavailable"):
        designs.source_for(replace(reused, contract="Other contract"))


def test_reference_without_report_sections_is_rejected(tmp_path, monkeypatch):
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path / "data"))
    with pytest.raises(ValueError, match="recognizable"):
        designs.install(reference(tmp_path, title="Invoice"), synthetic_profiles()[0],
                        actor="Editor", expected_revision=0)
    assert designs.state(synthetic_profiles()[0].contract)["revision"] == 0


def master_reference(tmp_path, name="master", marker="First design"):
    from app.monthly_report_model import default_sections
    path = tmp_path / (name + ".docx")
    doc = Document()
    doc.add_paragraph("Operations and Maintenance Monthly Review")
    doc.add_paragraph(marker)
    for section in default_sections():
        doc.add_heading(section.title, 1)
        doc.add_paragraph("Source-only content")
    doc.save(path)
    return path


def test_master_applies_without_upload_and_versions_saved_reports(tmp_path, monkeypatch):
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path / "data"))
    first = designs.install_master(master_reference(tmp_path), actor="Editor", expected_revision=0)
    profile = replace(synthetic_profiles()[0], contract="New contract", imported_from="")
    saved = designs.pin(profile)
    assert saved.template == designs.MASTER_PREFIX + first
    assert designs.is_master(saved)
    assert designs.source_for(saved).name == first + ".docx"
    second = designs.install_master(master_reference(tmp_path, "updated", "Updated design"),
                                    actor="Editor", expected_revision=1)
    assert second != first
    assert designs.pin(profile).template == designs.MASTER_PREFIX + second
    assert designs.pin(saved) == saved
    assert designs.pin(saved, latest_master=True).template == designs.MASTER_PREFIX + second
    assert designs.source_for(saved).name == first + ".docx"
    with pytest.raises(library.RevisionConflict):
        designs.install_master(master_reference(tmp_path), actor="Editor", expected_revision=1)


def test_master_requires_all_sections(tmp_path, monkeypatch):
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path / "data"))
    with pytest.raises(ValueError, match="every standard"):
        designs.install_master(reference(tmp_path), actor="Editor", expected_revision=0)
    assert designs.master_state()["revision"] == 0


def test_output_and_live_preview_use_same_pinned_master(tmp_path, monkeypatch):
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path / "data"))
    from app.monthly_report_model import synthetic_draft, ReportPeriod
    from app.monthly_report_docx import _build_docx
    from app.monthly_report_preview import section_preview_docx, section_preview_fingerprint
    from app import monthly_report_native_layout
    designs.install_master(master_reference(tmp_path), actor="Editor", expected_revision=0)
    draft = synthetic_draft(designs.pin(synthetic_profiles()[0]), ReportPeriod(2026, 9))
    calls = []
    def build(current, source, **kwargs):
        calls.append((current, source, kwargs))
        return b"native-template"
    monkeypatch.setattr(monthly_report_native_layout, "build_native_docx", build)
    assert _build_docx(draft) == b"native-template"
    assert section_preview_docx(draft, "cover") == b"native-template"
    assert calls[0][1] == calls[1][1]
    assert calls[0][2]["master"] and calls[1][2]["master"]
    assert calls[1][2]["section_key"] == "cover"
    before = section_preview_fingerprint(draft, "cover")
    designs.install_master(master_reference(tmp_path, "second", "New master"), actor="Editor", expected_revision=1)
    assert section_preview_fingerprint(draft, "cover") == before
    assert section_preview_fingerprint(replace(draft, profile=designs.pin(draft.profile, latest_master=True)), "cover") != before
