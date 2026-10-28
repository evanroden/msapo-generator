"""Real malformed XLSX XML must reject one upload without abandoning siblings."""
from hashlib import sha256
from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook
import pytest

from app import monthly_report_library as library
from app import monthly_report_sources as sources
from app.monthly_report_model import synthetic_profiles


@pytest.fixture
def profile(monkeypatch, tmp_path):
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path))
    return synthetic_profiles()[0]


def workbook_bytes():
    workbook = Workbook()
    workbook.active.append(("Equipment", "Finding"))
    workbook.active.append(("Synthetic pump", "Valve inspected"))
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()


def broken_xml(part, *, lazy=False):
    replacement = b"<broken"
    if lazy:
        # A valid dimension lets read-only loading finish. The incomplete row
        # is encountered only when openpyxl later iterates worksheet values.
        replacement = (b'<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
                       b'<dimension ref="A1:B2"/><sheetData><row r="1"><c r="A1"')
    stream = BytesIO()
    with ZipFile(BytesIO(workbook_bytes())) as original, ZipFile(stream, "w", ZIP_DEFLATED) as changed:
        for name in original.namelist():
            changed.writestr(name, replacement if name == part else original.read(name))
    return stream.getvalue()


@pytest.mark.parametrize("part,lazy", [
    ("[Content_Types].xml", False),
    ("xl/workbook.xml", False),
    ("xl/styles.xml", False),
    ("xl/worksheets/sheet1.xml", False),
    ("xl/worksheets/sheet1.xml", True),
])
def test_actual_malformed_workbook_xml_preserves_siblings_and_retry(profile, part, lazy):
    bad = broken_xml(part, lazy=lazy)
    assert len(bad) < 10000
    first, last = b"Synthetic first accepted finding.", b"Synthetic last accepted finding."
    accepted, notices = sources.ingest_batch(profile, (
        ("first.txt", first), ("broken.xlsx", bad), ("last.txt", last)))
    assert [item.source.filename for item in accepted] == ["first.txt", "last.txt"]
    assert [sources.source_bytes(profile, item.source) for item in accepted] == [first, last]
    assert notices == ("Not read: broken.xlsx. Workbook XML could not be read. "
                       "Export a fresh XLSX or CSV copy.",)
    bad_digest = sha256(bad).hexdigest()
    assert not sources._path(profile, bad_digest, ".xlsx").exists()
    assert not sources._path(profile, bad_digest, ".json").exists()

    corrected = workbook_bytes()
    retried, notices = sources.ingest_batch(profile, (("corrected.xlsx", corrected),), accepted)
    assert not notices and retried[:2] == accepted and len(retried) == 3
    assert retried[-1].tables[0].rows == (("Synthetic pump", "Valve inspected"),)
    assert sources.source_bytes(profile, retried[-1].source) == corrected
    repeated, notices = sources.ingest_batch(profile, (("renamed.xlsx", corrected),), retried)
    assert repeated == retried and notices == ("Duplicate skipped: renamed.xlsx",)


def test_preexisting_malformed_original_is_retained(profile):
    raw = broken_xml("xl/worksheets/sheet1.xml", lazy=True)
    original = sources._path(profile, sha256(raw).hexdigest(), ".xlsx")
    library._atomic_write(original, raw)
    before = original.stat()
    accepted, notices = sources.ingest_batch(profile, (("existing.xlsx", raw),))
    assert not accepted and "Workbook XML could not be read" in notices[0]
    assert original.read_bytes() == raw and original.stat().st_ino == before.st_ino


def test_lazy_xml_failure_closes_the_actual_workbook_archive(profile, monkeypatch):
    loaded = []
    original_load = sources.load_workbook
    def load(*args, **kwargs):
        workbook = original_load(*args, **kwargs)
        loaded.append(workbook)
        return workbook
    monkeypatch.setattr(sources, "load_workbook", load)
    sources.ingest_batch(profile, (("lazy.xlsx", broken_xml("xl/worksheets/sheet1.xml", lazy=True)),))
    assert len(loaded) == 1 and loaded[0]._archive.fp is None


def test_guided_xml_failure_preserves_delivered_activity_and_reopens(monkeypatch, tmp_path):
    import streamlit as st
    from streamlit.testing.v1 import AppTest
    from test_monthly_report_activity_intake_commit import APP, activity, draft, editor, setup_app

    app, profile = setup_app(monkeypatch, tmp_path)
    manual = "  Keep this authored monthly note.\n\nDo not replace it.\n"
    first = b"Technician inspected synthetic pump P-61."
    last = b"Technician inspected synthetic fan F-82."
    originals = (("first maintenance report.txt", first),
                 ("broken.xlsx", broken_xml("xl/worksheets/sheet1.xml", lazy=True)),
                 ("last maintenance report.txt", last))
    original_uploader = st.file_uploader
    def upload(label, *args, **kwargs):
        if label == "Add source files":
            uploaded = []
            for name, raw in originals:
                stream = BytesIO(raw)
                stream.name = name
                uploaded.append(stream)
            return uploaded
        return original_uploader(label, *args, **kwargs)
    monkeypatch.setattr(st, "file_uploader", upload)
    editor(app).set_value(manual).run()
    assert not app.exception
    expected = "\n\n".join((manual, first.decode(), last.decode()))
    assert activity(app).text == editor(app).value == expected
    expected_ids = {sha256(raw).hexdigest() for raw in (first, last)}
    assert {source.id for source in draft(app).sources} == expected_ids
    messages = "\n".join(item.value for item in (*app.warning, *app.info, *app.error))
    assert "Not read: broken.xlsx. Workbook XML could not be read" in messages
    monkeypatch.setattr(st, "file_uploader", original_uploader)
    resumed = AppTest.from_string(APP, default_timeout=45).run()
    assert not resumed.exception
    assert activity(resumed).text == editor(resumed).value == expected
    assert {source.id for source in draft(resumed).sources} == expected_ids
    assert tuple(sources.source_bytes(profile, source) for source in draft(resumed).sources) == (first, last)


def test_unexpected_product_syntax_error_is_not_a_file_failure(profile, monkeypatch):
    def unexpected(*args, **kwargs):
        raise SyntaxError("Synthetic programming failure outside the workbook decoder")
    monkeypatch.setattr(sources, "suggest_facts", unexpected)
    with pytest.raises(SyntaxError, match="outside the workbook decoder"):
        sources.ingest_batch(profile, (("valid.xlsx", workbook_bytes()),))
