"""A local table edit must retain adjacent technical table and picture geometry."""
from dataclasses import replace
from hashlib import sha256
from io import BytesIO

import pytest
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Inches
from PIL import Image

from app.monthly_report_import import ImportError, ImportMapping, inspect_docx, map_items, read_import_image
from app.monthly_report_model import ReportDraft, ReportProfile, ReportPeriod, Facility, ResolvedBlock, default_sections
from app.monthly_report_native_layout import build_native_docx
from app.monthly_report_sections import table_without_prices


def mixed_tables(tmp_path, *, unmapped_diagram=False):
    doc = Document()
    doc.add_heading('Water Treatment Reports', 1)
    for index in range(2):
        table = doc.add_table(rows=3, cols=3)
        table.columns[0].width = Inches(1.2 + index)
        for cell, text in zip(table.rows[0].cells, ('Measurement', 'Value', 'Diagram')):
            cell.text = text
        table.cell(1, 0).text = f'Synthetic sample {index}'
        table.cell(1, 1).text = '7.2'
        table.cell(2, 0).merge(table.cell(2, 1)).text = 'Merged technical note'
        data = BytesIO()
        Image.new('RGB', (20, 20), ('blue', 'green')[index]).save(data, 'PNG')
        paragraph = table.cell(1, 2).paragraphs[0]
        paragraph.add_run().add_picture(BytesIO(data.getvalue()), width=Inches(.3))
        if unmapped_diagram and index == 0:
            # Local image rels are intentionally not a validated SmartArt graph.
            # Their presence must never approve the neighboring diagram canvas.
            rid = next(paragraph._p.iter(qn('a:blip'))).get(qn('r:embed'))
            drawing = OxmlElement('w:drawing')
            inline = OxmlElement('wp:inline'); drawing.append(inline)
            extent = OxmlElement('wp:extent'); extent.set('cx', '200000'); extent.set('cy', '200000'); inline.append(extent)
            props = OxmlElement('wp:docPr'); props.set('id', '200'); props.set('name', 'Unmapped diagram'); inline.append(props)
            graphic = OxmlElement('a:graphic'); inline.append(graphic)
            data_node = OxmlElement('a:graphicData'); data_node.set('uri', 'http://schemas.openxmlformats.org/drawingml/2006/diagram'); graphic.append(data_node)
            from lxml import etree
            rels = etree.SubElement(data_node, '{http://schemas.openxmlformats.org/drawingml/2006/diagram}relIds')
            for key in ('dm', 'lo', 'qs', 'cs'):
                rels.set(qn('r:' + key), rid)
            paragraph.add_run()._r.append(drawing)
        doc.add_paragraph(f'After technical table {index}')
    path = tmp_path / 'mixed.docx'
    doc.save(path)
    inspection = inspect_docx(path)
    if unmapped_diagram:
        # This table now correctly fails closed. Expose original review facts
        # instead of assuming the unsafe table was mapped into an editable grid.
        return path, inspection, path.read_bytes()
    items = [item for item in inspection.items if item.kind in ('table', 'image', 'text') and item.section == 'water' and item.kind != 'text']
    tables = []
    assets = {}
    for item in items:
        if item.kind == 'table':
            value, removed = table_without_prices(item)
            assert not removed
            tables.append(replace(value, reference=f'docx:{inspection.sha256}:{item.id}'))
        elif item.kind == 'image':
            image = read_import_image(path, item, line_art=True)
            assets[sha256(image.data).hexdigest() + '.' + image.extension] = image.data
    tables[1] = replace(tables[1], rows=(tables[1].rows[0][:1] + ('8.1',) + tables[1].rows[0][2:], *tables[1].rows[1:]))
    block = ResolvedBlock('water_reports', 'This month', extra_tables=tuple(tables), asset_hashes=tuple(assets),
                          references=tuple(f'docx:{inspection.sha256}:{item.id}' for item in items))
    draft = ReportDraft(ReportProfile('Synthetic', 'site', 'Current site', (Facility('site', 'Current site'),)),
                        ReportPeriod(2026, 9), 'Current editor', default_sections(), (block,))
    return path, draft, assets


@pytest.mark.parametrize("reordered", [False, True])
def test_changed_mixed_table_stays_at_source_anchor_without_empty_duplicates(tmp_path, reordered):
    path, draft, assets = mixed_tables(tmp_path)
    if reordered:
        draft = replace(draft, blocks=(replace(draft.blocks[0], extra_tables=tuple(reversed(draft.blocks[0].extra_tables))),))
    source = Document(path)
    output = Document(BytesIO(build_native_docx(draft, path, asset_loader=assets.__getitem__)))
    assert len(output.tables) == len(source.tables) == 2
    for original, current in zip(source.tables, output.tables):
        assert original._tbl.tblGrid.xml == current._tbl.tblGrid.xml
        assert [cell.tcPr.xml for cell in original._tbl.iter(qn('w:tc'))] == [cell.tcPr.xml for cell in current._tbl.iter(qn('w:tc'))]
        assert len(list(current._tbl.iter(qn('a:blip')))) == 1
    assert output.tables[0].cell(1, 1).text == '7.2'
    assert output.tables[1].cell(1, 1).text == '8.1'
    assert output.tables[1].cell(2, 0).text == 'Merged technical note'


@pytest.mark.parametrize('change', ['unmapped', 'replaced'])
def test_mixed_table_does_not_keep_old_picture_without_current_asset_proof(tmp_path, change):
    path, draft, assets = mixed_tables(tmp_path)
    block = draft.blocks[0]
    if change == 'unmapped':
        block = replace(block, asset_hashes=())
    else:
        new = BytesIO()
        Image.new('RGB', (20, 20), 'red').save(new, 'PNG')
        assets['replacement.png'] = new.getvalue()
        block = replace(block, asset_hashes=('replacement.png',))
    output = build_native_docx(replace(draft, blocks=(block,)), path, asset_loader=assets.__getitem__)
    from zipfile import ZipFile
    with ZipFile(path) as archive:
        old_images = [archive.read(name) for name in archive.namelist() if name.startswith('word/media/')]
    with ZipFile(BytesIO(output)) as archive:
        visible = [archive.read(name) for name in archive.namelist() if name.startswith('word/media/')]
    assert not any(value in visible for value in old_images)
    assert len(Document(BytesIO(output)).tables) == 2
    assert len(visible) == (0 if change == 'unmapped' else 1)


def test_unmapped_diagram_cannot_ride_current_mixed_table_images(tmp_path):
    path, inspection, original = mixed_tables(tmp_path, unmapped_diagram=True)
    assert path.read_bytes() == original
    assert inspection.sha256 == sha256(original).hexdigest()
    # Preserve the original explicit SmartArt inventory, not just the table
    # blocker. Both require review before any source-page replacement.
    assert any(item.note.startswith('Unmapped native SmartArt')
               for item in inspection.items)
    # A diagram inside a source table makes that whole table unsafe to flatten.
    # The image and neighboring healthy table may still be reviewed separately,
    # but neither constitutes permission to reuse the original diagram canvas.
    blocked = [item for item in inspection.items
               if item.kind == 'unsupported' and item.section == 'water'
               and 'Word table contains text boxes or unsupported drawings' in item.note]
    assert len(blocked) == 1
    assert blocked[0].part == 'word/document.xml' and blocked[0].position == 2
    assert 'Synthetic sample 0' in blocked[0].text
    assert 'Merged technical note' in blocked[0].text
    assert not any(item.kind == 'table' and item.position == blocked[0].position
                   for item in inspection.items)
    # The unrelated second table and both original picture occurrences remain
    # inspectable. No source cells are invented or silently discarded.
    healthy = [item for item in inspection.items
               if item.kind == 'table' and item.section == 'water']
    assert len(healthy) == 1 and healthy[0].position > blocked[0].position
    assert healthy[0].rows[0] == ('Measurement', 'Value', 'Diagram')
    assert healthy[0].rows[1][:2] == ('Synthetic sample 1', '7.2')
    images = [item for item in inspection.items
              if item.kind == 'image' and item.section == 'water']
    assert len(images) == 2 and all(item.image_part and item.image_digest for item in images)
    with pytest.raises(ImportError, match='Unsupported items'):
        map_items(path, inspection, (ImportMapping(blocked[0].id, 'water_reports'),))
    assert path.read_bytes() == original
