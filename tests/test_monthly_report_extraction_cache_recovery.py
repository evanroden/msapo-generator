"""Recover disposable source extraction metadata without replacing reviewed work."""
from dataclasses import replace
from hashlib import sha256
from io import BytesIO

import pytest

from app import monthly_report_library as library, monthly_report_sources as sources
from app.monthly_report_model import ReportPeriod, synthetic_draft, synthetic_profiles
from app.monthly_report_source_recovery import restore_sources


CSV = b"Equipment,Observation\nP-001,Inspected pump\nF-002,Checked fan\n"
TXT = b"Synthetic service report\nActions: Completed pump inspection.\n"


@pytest.fixture
def profile(monkeypatch, tmp_path):
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path))
    return synthetic_profiles()[0]


def _paths(profile, content):
    source = content.source
    return (sources._path(profile, source.sha256, source.suffix),
            sources._path(profile, source.sha256, ".json"))


def _damage(cache, case):
    if case == "truncated_json":
        cache.write_bytes(b'{"schema":1,"source":')
        return
    data = library._read(cache)
    if case == "missing_source":
        data.pop("source")
    elif case == "null_source":
        data["source"] = None
    elif case == "null_tables":
        data["tables"] = None
    elif case == "missing_table_name":
        data["tables"][0].pop("name")
    elif case == "columns_string":
        data["tables"][0]["columns"] = "Equipment"
    elif case == "rows_string":
        data["tables"][0]["rows"] = "AB"
    elif case == "row_string":
        data["tables"][0]["rows"][0] = "AB"
    elif case == "short_row":
        data["tables"][0]["rows"][0] = ["P-001"]
    elif case == "long_row":
        data["tables"][0]["rows"][0].append("Unexpected third cell")
    elif case == "cell_nontext":
        data["tables"][0]["rows"][0][0] = 123
    elif case == "page_texts_string":
        data["source"]["page_texts"] = "AB"
    elif case == "page_text_nontext":
        data["source"]["page_texts"] = [123]
    elif case == "selected_pages_string":
        data["source"]["selected_pages"] = "1"
    elif case == "needs_vision_boolean":
        data["source"]["needs_vision"] = [True]
    elif case == "image_items_string":
        data["image_items"] = "AB"
    elif case == "image_item_string":
        data["image_items"] = ["1a"]
    elif case == "source_id_mismatch":
        data["source"]["id"] = "a" * 64
    elif case == "source_sha_mismatch":
        data["source"]["sha256"] = "b" * 64
    elif case == "source_suffix_mismatch":
        data["source"]["suffix"] = ".txt"
    elif case == "policy_string":
        data["client_page_policy"] = "1"
    elif case == "policy_boolean":
        data["native_action_policy"] = True
    elif case == "policy_null":
        data["vendor_passage_policy"] = None
    elif case == "policy_negative":
        data["native_action_policy"] = -1
    else:
        raise AssertionError(case)
    cache.write_bytes(library._json(data))


@pytest.mark.parametrize("case", [
    "truncated_json", "missing_source", "null_source", "null_tables", "missing_table_name",
    "columns_string", "rows_string", "row_string", "short_row", "long_row", "cell_nontext", "page_texts_string",
    "page_text_nontext", "selected_pages_string", "needs_vision_boolean", "image_items_string",
    "image_item_string", "source_id_mismatch", "source_sha_mismatch", "source_suffix_mismatch",
    "policy_string", "policy_boolean", "policy_null", "policy_negative",
])
def test_corrupt_extraction_cache_rebuilds_exact_csv_and_then_stays_warm(profile, monkeypatch, case):
    original, _ = sources.ingest(profile, "synthetic.csv", CSV)
    path, cache = _paths(profile, original)
    original_stat = path.stat()
    _damage(cache, case)
    parse = sources._spreadsheet
    calls = []

    def counted(raw, suffix):
        calls.append((raw, suffix))
        return parse(raw, suffix)

    monkeypatch.setattr(sources, "_spreadsheet", counted)
    rebuilt, attachments = sources.ingest(profile, "synthetic.csv", CSV)
    assert rebuilt == original and attachments == ()
    assert calls == [(CSV, ".csv")]
    assert path.read_bytes() == CSV
    assert (path.stat().st_ino, path.stat().st_mtime_ns) == (original_stat.st_ino, original_stat.st_mtime_ns)
    assert library._read(cache)["source"]["sha256"] == sha256(CSV).hexdigest()
    monkeypatch.setattr(sources, "_spreadsheet", lambda *args: pytest.fail("repaired cache was reparsed"))
    warm, _ = sources.ingest(profile, "renamed.csv", CSV)
    assert warm == replace(original, source=replace(original.source, filename="renamed.csv"))


def _spreadsheet_bytes(suffix, *, header_only=False):
    if suffix == ".csv":
        return b"Equipment,Observation\n" if header_only else CSV
    from openpyxl import Workbook
    workbook = Workbook()
    workbook.active.append(("Equipment", "Observation"))
    if not header_only:
        workbook.active.append(("P-001", "Inspected pump"))
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


@pytest.mark.parametrize("suffix", [".csv", ".xlsx"])
@pytest.mark.parametrize("case", ["missing_tables", "empty_tables", "empty_columns"])
def test_incomplete_spreadsheet_cache_rebuilds_table(profile, monkeypatch, suffix, case):
    raw = _spreadsheet_bytes(suffix)
    original, _ = sources.ingest(profile, "synthetic" + suffix, raw)
    path, cache = _paths(profile, original)
    original_stat = path.stat()
    data = library._read(cache)
    if case == "missing_tables":
        data.pop("tables")
    elif case == "empty_tables":
        data["tables"] = []
    else:
        data["tables"][0]["columns"] = []
        data["tables"][0]["rows"] = []
    cache.write_bytes(library._json(data))
    parse, calls = sources._spreadsheet, []

    def counted(data, extension):
        calls.append((data, extension))
        return parse(data, extension)

    monkeypatch.setattr(sources, "_spreadsheet", counted)
    rebuilt, attachments = sources.ingest(profile, "synthetic" + suffix, raw)
    assert rebuilt == original and attachments == ()
    assert calls == [(raw, suffix)]
    assert path.read_bytes() == raw
    assert (path.stat().st_ino, path.stat().st_mtime_ns) == (original_stat.st_ino, original_stat.st_mtime_ns)
    monkeypatch.setattr(sources, "_spreadsheet", lambda *args: pytest.fail("repaired cache was reparsed"))
    assert sources.ingest(profile, "synthetic" + suffix, raw) == (original, ())


@pytest.mark.parametrize("suffix", [".csv", ".xlsx"])
def test_header_only_spreadsheet_cache_is_a_warm_hit(profile, monkeypatch, suffix):
    raw = _spreadsheet_bytes(suffix, header_only=True)
    original, _ = sources.ingest(profile, "headers" + suffix, raw)
    assert original.tables[0].columns == ("Equipment", "Observation")
    assert original.tables[0].rows == ()
    path, cache = _paths(profile, original)
    original_stat, cache_bytes = path.stat(), cache.read_bytes()
    monkeypatch.setattr(sources, "_spreadsheet", lambda *args: pytest.fail("valid header-only cache was reparsed"))
    assert sources.ingest(profile, "headers" + suffix, raw) == (original, ())
    assert cache.read_bytes() == cache_bytes
    assert (path.stat().st_ino, path.stat().st_mtime_ns) == (original_stat.st_ino, original_stat.st_mtime_ns)


def test_text_cache_can_omit_optional_tables_and_images(profile, monkeypatch):
    original, _ = sources.ingest(profile, "synthetic.txt", TXT)
    _, cache = _paths(profile, original)
    data = library._read(cache)
    data.pop("tables")
    data.pop("image_items")
    cache.write_bytes(library._json(data))
    cache_bytes = cache.read_bytes()
    monkeypatch.setattr(sources, "suggest_facts", lambda *args: pytest.fail("legacy text cache was reparsed"))
    assert sources.ingest(profile, "synthetic.txt", TXT) == (original, ())
    assert cache.read_bytes() == cache_bytes


def test_recovery_preserves_reviewed_sources_siblings_originals_and_history(profile):
    first, _ = sources.ingest(profile, "synthetic.csv", CSV)
    second, _ = sources.ingest(profile, "synthetic.txt", TXT)
    first_path, first_cache = _paths(profile, first)
    second_path, second_cache = _paths(profile, second)
    original_files = {path: (path.read_bytes(), path.stat().st_ino, path.stat().st_mtime_ns)
                      for path in (first_path, second_path)}
    second_cache_bytes = second_cache.read_bytes()
    reviewed = replace(first.source, vendor="Reviewed vendor", actions="", findings="Retain human findings",
                       selected_pages=(), captions=((1, "Deliberate caption"),), scope_review="retained-scope-receipt")
    draft = replace(synthetic_draft(profile, ReportPeriod(2026, 9)), sources=(reviewed, second.source))
    library.save_snapshot(draft, expected_revision=0)
    _damage(first_cache, "missing_source")
    available, missing = restore_sources(profile, draft.sources)
    assert missing == ()
    assert tuple(item.source for item in available) == draft.sources
    assert available[0].tables == first.tables
    assert second_cache.read_bytes() == second_cache_bytes
    for path, (raw, inode, mtime) in original_files.items():
        assert path.read_bytes() == raw
        assert (path.stat().st_ino, path.stat().st_mtime_ns) == (inode, mtime)
    assert library.load_snapshot(profile.contract, profile.key, draft.period).draft == draft
    assert library.load_snapshot_version(profile.contract, profile.key, draft.period, 1).draft == draft


def test_same_byte_txt_original_is_not_republished_while_cache_recovers(profile):
    content, _ = sources.ingest(profile, "synthetic.txt", TXT)
    path, cache = _paths(profile, content)
    before = path.stat()
    _damage(cache, "missing_source")
    restored, missing = restore_sources(profile, (content.source,))
    assert missing == () and restored == (content,)
    assert path.read_bytes() == TXT
    assert (path.stat().st_ino, path.stat().st_mtime_ns) == (before.st_ino, before.st_mtime_ns)


def test_bad_original_is_not_recovered_from_cache_or_allowed_to_hide_good_sibling(profile):
    bad, _ = sources.ingest(profile, "synthetic-bad.txt", TXT)
    good, _ = sources.ingest(profile, "synthetic-good.csv", CSV)
    path, cache = _paths(profile, bad)
    _damage(cache, "missing_source")
    cache_bytes = cache.read_bytes()
    path.write_bytes(b"Different original bytes")
    restored, missing = restore_sources(profile, (bad.source, good.source))
    assert restored == (good,) and missing == (bad.source,)
    assert path.read_bytes() == b"Different original bytes"
    assert cache.read_bytes() == cache_bytes


def test_explicit_same_source_reupload_can_still_restore_a_damaged_txt_original(profile):
    content, _ = sources.ingest(profile, "synthetic.txt", TXT)
    path, cache = _paths(profile, content)
    _damage(cache, "missing_source")
    path.write_bytes(b"Damaged stored copy")
    repaired, _ = sources.ingest(profile, "synthetic.txt", TXT)
    assert repaired == content
    assert path.read_bytes() == TXT


@pytest.mark.parametrize("phase", ["parse", "suggest", "publish"])
def test_failed_rebuild_preserves_preexisting_cache_and_original(profile, monkeypatch, phase):
    content, _ = sources.ingest(profile, "synthetic.csv", CSV)
    path, cache = _paths(profile, content)
    _damage(cache, "missing_source")
    cache_bytes, original_bytes = cache.read_bytes(), path.read_bytes()
    before = path.stat()
    error = RuntimeError("synthetic parser failed") if phase == "parse" else TypeError("synthetic suggestion failed")
    if phase == "publish":
        error = OSError("synthetic cache publication failed")
        original_write = library._atomic_write

        def fail_publish(target, raw):
            if target == cache:
                raise error
            return original_write(target, raw)

        monkeypatch.setattr(library, "_atomic_write", fail_publish)
    else:
        def fail(*args, **kwargs):
            raise error

        monkeypatch.setattr(sources, "_spreadsheet" if phase == "parse" else "suggest_facts", fail)
    with pytest.raises(type(error), match=str(error)):
        sources.ingest(profile, "synthetic.csv", CSV)
    assert cache.read_bytes() == cache_bytes
    assert path.read_bytes() == original_bytes
    assert (path.stat().st_ino, path.stat().st_mtime_ns) == (before.st_ino, before.st_mtime_ns)


def test_valid_legacy_cache_policies_migrate_without_parsing_or_rewriting(profile, monkeypatch):
    content, _ = sources.ingest(profile, "synthetic.txt", TXT)
    path, cache = _paths(profile, content)
    data = library._read(cache)
    data.pop("client_page_policy")
    data.pop("native_action_policy")
    data.pop("vendor_passage_policy")
    data["source"]["actions"] = "Old cached suggestion"
    data["source"]["selected_pages"] = []
    cache.write_bytes(library._json(data))
    cache_bytes = cache.read_bytes()
    original_stat = path.stat()
    monkeypatch.setattr(sources, "suggest_facts", lambda *args: pytest.fail("legacy cache was reparsed"))
    restored, _ = sources.ingest(profile, "renamed.txt", TXT)
    from app.monthly_report_content_policy import suggested_pages
    expected_actions = "\n".join(line for line, _ in sources.native_action_lines(replace(content.source, selected_pages=suggested_pages(content.source))))[:6000]
    assert restored.source.filename == "renamed.txt"
    assert restored.source.selected_pages == suggested_pages(content.source)
    assert restored.source.actions == expected_actions
    assert cache.read_bytes() == cache_bytes
    assert (path.stat().st_ino, path.stat().st_mtime_ns) == (original_stat.st_ino, original_stat.st_mtime_ns)


def test_valid_cache_policy_failure_is_not_treated_as_cache_corruption(profile, monkeypatch):
    content, _ = sources.ingest(profile, "synthetic.txt", TXT)
    _, cache = _paths(profile, content)
    data = library._read(cache)
    data["native_action_policy"] = 0
    cache.write_bytes(library._json(data))
    before = cache.read_bytes()

    def fail_policy(*args):
        raise TypeError("synthetic migration defect")

    monkeypatch.setattr(sources, "native_action_lines", fail_policy)
    monkeypatch.setattr(sources, "suggest_facts", lambda *args: pytest.fail("migration failure reparsed original"))
    with pytest.raises(TypeError, match="synthetic migration defect"):
        sources.ingest(profile, "synthetic.txt", TXT)
    assert cache.read_bytes() == before


def test_cache_scope_guard_failure_is_not_rebuilt(profile, monkeypatch):
    content, _ = sources.ingest(profile, "synthetic.csv", CSV)
    path, cache = _paths(profile, content)
    before = cache.read_bytes()
    guard = library._qa_path_guard

    def fail_scope(target):
        if target == cache:
            raise library.LibraryError("synthetic stale namespace")
        return guard(target)

    monkeypatch.setattr(library, "_qa_path_guard", fail_scope)
    monkeypatch.setattr(sources, "_spreadsheet", lambda *args: pytest.fail("scope refusal reparsed original"))
    with pytest.raises(library.LibraryError, match="synthetic stale namespace"):
        sources.ingest(profile, "synthetic.csv", CSV)
    assert cache.read_bytes() == before and path.read_bytes() == CSV


def test_namespace_change_during_original_comparison_stops_before_reparse(profile, monkeypatch):
    content, _ = sources.ingest(profile, "synthetic.csv", CSV)
    path, cache = _paths(profile, content)
    _damage(cache, "missing_source")
    cache_bytes = cache.read_bytes()
    original_stat = path.stat()
    guard, open_path = library._qa_path_guard, type(path).open
    expired = False

    def check_scope(target):
        if target == path and expired:
            raise library.LibraryError("synthetic namespace changed during read")
        return guard(target)

    class ExpiringRead:
        def __init__(self, stream):
            self.stream = stream

        def __enter__(self):
            return self

        def __exit__(self, *args):
            self.stream.close()

        def read(self, size=-1):
            nonlocal expired
            data = self.stream.read(size)
            expired = True
            return data

    def open_and_expire(target, *args, **kwargs):
        stream = open_path(target, *args, **kwargs)
        mode = args[0] if args else kwargs.get("mode", "r")
        return ExpiringRead(stream) if target == path and mode == "rb" else stream

    monkeypatch.setattr(library, "_qa_path_guard", check_scope)
    monkeypatch.setattr(type(path), "open", open_and_expire)
    monkeypatch.setattr(sources, "_spreadsheet", lambda *args: pytest.fail("stale namespace reparsed original"))
    with pytest.raises(library.LibraryError, match="synthetic namespace changed during read"):
        sources.ingest(profile, "synthetic.csv", CSV)
    assert cache.read_bytes() == cache_bytes and path.read_bytes() == CSV
    assert (path.stat().st_ino, path.stat().st_mtime_ns) == (original_stat.st_ino, original_stat.st_mtime_ns)
