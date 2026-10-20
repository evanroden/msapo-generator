"""Synthetic-only explicit compatibility receipts: no production activation."""
from dataclasses import replace
import hashlib
import json
import shutil

from docx import Document
import pytest

from app import monthly_report_library as library
from app.monthly_report_compatibility_view import derive_compatibility_view, POLICY
from app.monthly_report_import import MappedImport, imported_draft
from app.monthly_report_model import Facility, ReportPeriod, ResolvedBlock, synthetic_profiles
from test_monthly_report_source_frame_logo_role import raster, source


PERIOD = ReportPeriod(2026, 9)


@pytest.fixture
def prepared(tmp_path, monkeypatch):
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path / "store"))
    path = source(tmp_path, divergent=True)
    view = derive_compatibility_view(path.read_bytes())
    profile = replace(synthetic_profiles()[0], imported_from=view.original_inspection.sha256)
    item = next(i for i in view.original_inspection.items if i.text == "Synthetic completed inspection.")
    block = ResolvedBlock("activity_summary", "This month", text=item.text,
                          references=(f"docx:{view.original_inspection.sha256}:{item.id}",))
    draft = imported_draft(profile, PERIOD, "Synthetic Editor", MappedImport((block,), (), (), ()))
    library.save_profile(profile, expected_revision=0, actor="Synthetic Editor", confirmed=True)
    return path, view, profile, draft


def save(prepared, *, view=True, revision=1, contribution=True, **changes):
    path, expected, profile, draft = prepared
    arguments = dict(assets=(), expected_revision=revision, actor="Synthetic Editor",
                     confirmed=True, source_draft=draft if contribution else None)
    if view:
        arguments["source_view"] = expected
    arguments.update(changes)
    review = {"sha256": expected.original_inspection.sha256, "period": draft.period.key,
              "image_review_scope": "synthetic-unapproved-review", "intent": "partial"}
    return library.save_report_setup(profile, draft, path, review, **arguments)


def saved_paths(prepared):
    _, view, profile, _ = prepared
    directory = library._profile_path(profile.contract, profile.key)
    manifest = library._read(directory / "manifest.json")
    receipt = directory / "imports" / (
        view.original_inspection.sha256 + "." + view.receipt_sha256 + ".source-view.json")
    return directory, manifest, receipt


def read(prepared, operation, *, expected=True, period=PERIOD):
    _, view, profile, _ = prepared
    kwargs = {"source_view": view} if expected else {}
    if operation == "draft":
        return library.load_imported_draft(profile.contract, profile.key, **kwargs)
    if operation == "review":
        return library.imported_review_scope(profile.contract, profile.key, period, **kwargs)
    return library.starting_contribution(profile.contract, profile.key, period,
                                         view.original_inspection.sha256, **kwargs)


def test_explicit_roundtrip_keeps_original_audit_ids_and_no_new_approval(prepared):
    path, view, profile, draft = prepared
    before = path.read_bytes()
    before_inspection = view.original_inspection
    state = save(prepared)
    assert state.revision == 2
    assert read(prepared, "draft") == draft
    assert read(prepared, "replay") == (draft, True)
    assert read(prepared, "review") == "synthetic-unapproved-review"
    original = library.imported_original(profile.contract, profile.key)
    assert original.name == hashlib.sha256(before).hexdigest() + ".docx"
    assert original.read_bytes() == path.read_bytes() == before
    assert view.original_inspection == before_inspection
    assert any(i.kind == "unsupported" for i in before_inspection.items)
    assert not any(b.client_reviewed_fingerprint or b.client_asset_reviews for b in draft.blocks)
    assert read(prepared, "draft").blocks[0].references == draft.blocks[0].references
    directory, manifest, receipt = saved_paths(prepared)
    assert receipt.read_bytes() == view.receipt_bytes
    assert (directory / "history" / "00000002.json").read_bytes() == (directory / "manifest.json").read_bytes()
    assert manifest["import_review"] == profile.imported_from == view.original_inspection.sha256
    replay_key = manifest["import_source_view"]
    assert POLICY in replay_key and view.receipt_sha256 in replay_key
    assert replay_key in manifest["starting_contributions"]
    assert PERIOD.key + ":" + profile.imported_from not in manifest["starting_contributions"]
    assert manifest["source_views"][replay_key]["scope"] == {
        "schema": 1, "contract": profile.contract, "profile_key": profile.key, "period": PERIOD.key,
        "original_sha256": profile.imported_from, "receipt_sha256": view.receipt_sha256, "policy": POLICY,
    }
    assert not (directory / "imports" / (profile.imported_from + ".json")).exists()
    assert not (directory / "imports" / (view.rendition_inspection.sha256 + ".docx")).exists()


@pytest.mark.parametrize("operation", ["draft", "review", "replay"])
def test_saved_view_is_rederived_when_expected_keyword_is_omitted(prepared, operation):
    save(prepared)
    assert read(prepared, operation, expected=False) == read(prepared, operation)


@pytest.mark.parametrize("operation", ["draft", "review", "replay"])
@pytest.mark.parametrize("damage", ["receipt", "missing_receipt", "source", "scope", "scope_hash"])
def test_replay_rederives_original_and_validates_receipt_scope(prepared, operation, damage):
    save(prepared)
    directory, manifest, receipt = saved_paths(prepared)
    if damage == "receipt":
        changed = json.loads(receipt.read_bytes())
        changed["policy"] = "tampered"
        receipt.write_bytes(library._json(changed))
    elif damage == "missing_receipt":
        receipt.unlink()
    elif damage == "source":
        original = library.imported_original(prepared[2].contract, prepared[2].key)
        # Replace only this synthetic link; never mutate a shared immutable object.
        raw = original.read_bytes()
        original.unlink()
        original.write_bytes(raw + b"changed synthetic source")
    else:
        record = manifest["source_views"][manifest["import_source_view"]]
        if damage == "scope":
            record["scope"]["contract"] = "Different synthetic contract"
            record["scope_sha256"] = hashlib.sha256(library._json(record["scope"])).hexdigest()
        else:
            record["scope_sha256"] = "0" * 64
        (directory / "manifest.json").write_bytes(library._json(manifest))
    with pytest.raises(library.LibraryError):
        read(prepared, operation)


@pytest.mark.parametrize("operation", ["review", "replay"])
def test_view_receipt_does_not_cross_month(prepared, operation):
    save(prepared)
    with pytest.raises(library.LibraryError, match="scope"):
        read(prepared, operation, period=ReportPeriod(2026, 10))


@pytest.mark.parametrize("change", ["contract", "profile"])
def test_copied_profile_store_does_not_transfer_receipt_scope(prepared, change):
    save(prepared)
    _, view, profile, _ = prepared
    directory, _, _ = saved_paths(prepared)
    contract = "Different synthetic contract" if change == "contract" else profile.contract
    key = "different-profile" if change == "profile" else profile.key
    shutil.copytree(directory, library._profile_path(contract, key))
    with pytest.raises(library.LibraryError, match="scope"):
        library.starting_contribution(contract, key, PERIOD, profile.imported_from, source_view=view)


@pytest.mark.parametrize("operation", ["draft", "review", "replay"])
def test_legacy_records_cannot_be_implicitly_normalized(prepared, operation):
    save(prepared, view=False)
    assert read(prepared, operation, expected=False) is not None
    with pytest.raises(library.LibraryError, match="no saved source-view"):
        read(prepared, operation)


@pytest.mark.parametrize("field", ["receipt_bytes", "rendition_bytes", "bindings", "original_inspection"])
def test_forged_expected_view_cannot_save_or_replay(prepared, field):
    path, view, profile, draft = prepared
    save(prepared)
    changed = {"receipt_bytes": b"{}", "rendition_bytes": view.rendition_bytes + b"x",
               "bindings": (), "original_inspection": replace(view.original_inspection, notices=("Forged notice",))}
    forged = replace(view, **{field: changed[field]})
    directory, _, _ = saved_paths(prepared)
    before = (directory / "manifest.json").read_bytes()
    with pytest.raises(library.LibraryError, match="verified"):
        save((path, forged, profile, draft), revision=2)
    with pytest.raises(library.LibraryError, match="verified"):
        library.starting_contribution(profile.contract, profile.key, PERIOD,
                                      profile.imported_from, source_view=forged)
    assert (directory / "manifest.json").read_bytes() == before


def test_stale_view_and_derivative_as_original_fail_before_publication(prepared):
    path, view, profile, draft = prepared
    directory = library._profile_path(profile.contract, profile.key)
    before = (directory / "manifest.json").read_bytes()
    document = Document(path)
    document.add_paragraph("New synthetic source bytes.")
    document.save(path)
    with pytest.raises(library.LibraryError, match="changed during review"):
        save(prepared)
    path.write_bytes(view.rendition_bytes)
    with pytest.raises(library.LibraryError, match="changed during review"):
        save(prepared)
    assert (directory / "manifest.json").read_bytes() == before


@pytest.mark.parametrize("pin", ["profile", "draft", "contribution"])
def test_original_profile_pins_cannot_be_replaced_with_derivative_hash(prepared, pin):
    path, view, profile, draft = prepared
    wrong = replace(profile, imported_from=view.rendition_inspection.sha256)
    supplied = (path, view, wrong if pin == "profile" else profile,
                replace(draft, profile=wrong) if pin == "draft" else draft)
    kwargs = {"source_draft": replace(draft, profile=wrong)} if pin == "contribution" else {}
    with pytest.raises(library.LibraryError, match="original source"):
        save(supplied, **kwargs)
    assert library.load_profile(profile.contract, profile.key).revision == 1


@pytest.mark.parametrize("operation", ["draft", "review"])
def test_missing_active_pointer_never_falls_back_to_bootstrap(prepared, operation):
    save(prepared)
    directory, manifest, _ = saved_paths(prepared)
    manifest.pop("import_source_view")
    (directory / "manifest.json").write_bytes(library._json(manifest))
    with pytest.raises(library.LibraryError, match="head is missing"):
        read(prepared, operation, expected=False)


def test_missing_contribution_never_uses_scoped_bootstrap_as_approval(prepared):
    save(prepared, contribution=False)
    assert read(prepared, "draft") == prepared[3]
    with pytest.raises(library.LibraryError, match="separately acknowledged"):
        read(prepared, "replay")


def test_compatibility_save_cannot_silently_revert_to_legacy(prepared):
    save(prepared)
    directory, _, receipt = saved_paths(prepared)
    before = (directory / "manifest.json").read_bytes(), receipt.read_bytes()
    with pytest.raises(library.LibraryError, match="explicitly verified"):
        save(prepared, view=False, revision=2)
    assert ((directory / "manifest.json").read_bytes(), receipt.read_bytes()) == before


def test_interrupted_publication_preserves_old_head_original_receipt_and_history(prepared, monkeypatch, tmp_path):
    path, view, profile, draft = prepared
    save(prepared)
    directory, _, receipt = saved_paths(prepared)
    before_head = (directory / "manifest.json").read_bytes()
    before_history = (directory / "history" / "00000002.json").read_bytes()
    before_receipt, before_original = receipt.read_bytes(), path.read_bytes()
    document = Document(path)
    document.add_paragraph("Different synthetic content for second import.")
    document.save(path)
    changed_view = derive_compatibility_view(path.read_bytes())
    changed_profile = replace(profile, imported_from=changed_view.original_inspection.sha256)
    changed_draft = replace(draft, profile=changed_profile)
    real_write = library._atomic_write

    def fail_head(target, raw):
        if target == directory / "manifest.json":
            raise OSError("Synthetic interrupted publication")
        real_write(target, raw)

    monkeypatch.setattr(library, "_atomic_write", fail_head)
    with pytest.raises(OSError, match="interrupted"):
        save((path, changed_view, changed_profile, changed_draft), revision=2)
    assert (directory / "manifest.json").read_bytes() == before_head
    assert (directory / "history" / "00000002.json").read_bytes() == before_history
    assert receipt.read_bytes() == before_receipt
    assert library.imported_original(profile.contract, profile.key).read_bytes() == before_original
    assert read(prepared, "draft") == draft
    assert read(prepared, "replay") == (draft, True)
    with pytest.raises(library.LibraryError, match="no saved source-view"):
        library.starting_contribution(profile.contract, profile.key, PERIOD,
                                      changed_profile.imported_from, source_view=changed_view)


def test_existing_receipt_cannot_be_overwritten_on_resave(prepared):
    save(prepared)
    directory, _, receipt = saved_paths(prepared)
    receipt.write_bytes(b"tampered")
    before = (directory / "manifest.json").read_bytes()
    with pytest.raises(library.LibraryError, match="receipt changed"):
        save(prepared, revision=2)
    assert receipt.read_bytes() == b"tampered"
    assert (directory / "manifest.json").read_bytes() == before


def test_replay_contribution_scope_is_not_transferable(prepared):
    save(prepared)
    directory, manifest, _ = saved_paths(prepared)
    key = manifest["import_source_view"]
    reference = manifest["starting_contributions"][key]
    payload = library._read(directory / "imports" / (reference + ".contribution.json"))
    payload.pop("source_view")
    raw = library._json(payload)
    changed_reference = hashlib.sha256(raw).hexdigest()
    (directory / "imports" / (changed_reference + ".contribution.json")).write_bytes(raw)
    manifest["starting_contributions"][key] = changed_reference
    (directory / "manifest.json").write_bytes(library._json(manifest))
    with pytest.raises(library.LibraryError, match="scope changed"):
        read(prepared, "replay")


def test_same_original_distinct_months_keep_separate_receipts_and_history(prepared):
    path, view, profile, draft = prepared
    save(prepared)
    directory, first, receipt = saved_paths(prepared)
    old_history, receipt_bytes = (directory / "history" / "00000002.json").read_bytes(), receipt.read_bytes()
    following = ReportPeriod(2026, 10)
    next_draft = replace(draft, period=following)
    save((path, view, profile, next_draft), revision=2)
    _, current, _ = saved_paths(prepared)
    assert len(current["source_views"]) == len(current["starting_contributions"]) == 2
    assert first["import_source_view"] != current["import_source_view"]
    assert (directory / "history" / "00000002.json").read_bytes() == old_history
    assert receipt.read_bytes() == receipt_bytes
    assert read(prepared, "replay") == (draft, True)
    assert read(prepared, "replay", period=following) == (next_draft, True)
    assert read(prepared, "draft") == next_draft
    with pytest.raises(library.LibraryError, match="scope"):
        read(prepared, "review")


def test_first_setup_publication_preserves_source_contribution_for_automatic_reload(prepared):
    path, view, profile, draft = prepared
    fresh_profile = replace(profile, key="new-source-view")
    fresh_draft = replace(draft, profile=fresh_profile)
    fresh = (path, view, fresh_profile, fresh_draft)
    state = save(fresh, revision=0)
    assert state.revision == 1
    assert read(fresh, "replay") == (fresh_draft, True)
    assert library.imported_original(fresh_profile.contract, fresh_profile.key).read_bytes() == path.read_bytes()
    assert read(fresh, "draft", expected=False) == fresh_draft


def test_legacy_artwork_approval_cannot_transfer_to_new_receipt(prepared):
    path, view, profile, draft = prepared
    image = next(i for i in view.original_inspection.items if i.image_role == "client_logo")
    raw = raster()
    reference = library.asset_reference(raw, "png")
    artwork = ResolvedBlock("client_logo", "This month", asset_hashes=(reference,),
                            references=(f"docx:{profile.imported_from}:{image.id}",))
    approved = replace(artwork, client_reviewed_fingerprint=artwork.fingerprint)
    prior = replace(draft, blocks=(*draft.blocks, approved))
    save((path, view, profile, prior), view=False, assets=((reference, raw),))
    directory, legacy, _ = saved_paths(prepared)
    legacy_key = PERIOD.key + ":" + profile.imported_from
    legacy_contribution = legacy["starting_contributions"][legacy_key]
    old_history = (directory / "history" / "00000002.json").read_bytes()
    current = replace(draft, blocks=(*draft.blocks, artwork))
    save((path, view, profile, current), revision=2, assets=((reference, raw),))
    reloaded, verified = read(prepared, "replay")
    current_artwork = next(b for b in reloaded.blocks if b.key == "client_logo")
    assert verified and current_artwork == artwork
    assert not current_artwork.client_reviewed_fingerprint
    assert not current_artwork.client_asset_reviews
    _, manifest, _ = saved_paths(prepared)
    assert manifest["starting_contributions"][legacy_key] == legacy_contribution
    assert manifest["starting_contributions"][manifest["import_source_view"]] != legacy_contribution
    assert (directory / "history" / "00000002.json").read_bytes() == old_history
    assert library.read_asset(profile.contract, profile.key, reference) == raw
    automatic, verified = read(prepared, "replay", expected=False)
    assert verified and next(b for b in automatic.blocks if b.key == "client_logo") == artwork


@pytest.fixture
def owned(prepared):
    path, view, profile, draft = prepared
    profile = replace(profile, facilities=(Facility("source-hospital", "Synthetic Regional Community Health Hospital"),))
    return path, view, profile, replace(draft, profile=profile)


def test_render_boundary_requires_saved_same_site_and_rederives_once(owned, monkeypatch, tmp_path):
    path, view, _, draft = owned
    save(owned)
    supplied_copy = tmp_path / "unchanged-original.docx"
    supplied_copy.write_bytes(path.read_bytes())
    calls = []
    verify = library._verified_source_view

    def count_verify(original, expected):
        calls.append(original)
        return verify(original, expected)

    monkeypatch.setattr(library, "_verified_source_view", count_verify)
    assert library.verified_render_source_view(draft, supplied_copy, view) is view
    assert calls == [library.imported_original(draft.profile.contract, draft.profile.key)]
    assert supplied_copy.read_bytes() == path.read_bytes()


@pytest.mark.parametrize("legacy", [False, True])
def test_render_boundary_never_upgrades_unsaved_or_legacy_scope(owned, legacy):
    path, view, _, draft = owned
    if legacy:
        save(owned, view=False)
    with pytest.raises(library.LibraryError, match="no saved source-view"):
        library.verified_render_source_view(draft, path, view)


@pytest.mark.parametrize("change", ["contract", "profile", "month", "pin", "no_pin", "no_expected"])
def test_render_boundary_rejects_wrong_scope_or_original_pin(owned, change):
    path, view, profile, draft = owned
    save(owned)
    if change == "contract":
        draft = replace(draft, profile=replace(profile, contract="Different synthetic contract"))
    elif change == "profile":
        draft = replace(draft, profile=replace(profile, key="different-profile"))
    elif change == "month":
        draft = replace(draft, period=ReportPeriod(2026, 10))
    elif change in ("pin", "no_pin"):
        draft = replace(draft, profile=replace(profile, imported_from="0" * 64 if change == "pin" else ""))
    else:
        assert library.verified_render_source_view(draft, path) == view
        return
    with pytest.raises(library.LibraryError):
        library.verified_render_source_view(draft, path, view)


@pytest.mark.parametrize("ownership", ["unknown", "foreign", "partial"])
def test_render_boundary_requires_content_ownership_not_manual_review(owned, ownership):
    path, view, profile, draft = owned
    if ownership == "unknown":
        sites = (Facility("source-hospital", "Unknown synthetic facility"),)
    elif ownership == "foreign":
        sites = (Facility("foreign", "Another Synthetic Hospital"),)
    else:
        sites = (*profile.facilities, Facility("missing-site", "Another Synthetic Hospital"))
    profile = replace(profile, facilities=sites, scope_type="multi_site" if len(sites) > 1 else "individual")
    draft = replace(draft, profile=profile)
    changed = path, view, profile, draft
    save(changed)
    directory, manifest, _ = saved_paths(changed)
    manifest["import_source_view_review"]["site_relation"] = "same"
    manifest["import_source_view_review"]["client_approved"] = True
    (directory / "manifest.json").write_bytes(library._json(manifest))
    with pytest.raises(library.LibraryError, match="same-site ownership"):
        library.verified_render_source_view(draft, path, view)


@pytest.mark.parametrize("damage", ["caller_original", "stored_original", "receipt", "expected_view"])
def test_render_boundary_rechecks_all_source_inputs(owned, damage):
    path, view, profile, draft = owned
    save(owned)
    if damage == "caller_original":
        path.write_bytes(path.read_bytes() + b"changed synthetic bytes")
    elif damage == "stored_original":
        original = library.imported_original(profile.contract, profile.key)
        raw = original.read_bytes()
        original.unlink()
        original.write_bytes(raw + b"changed synthetic bytes")
    elif damage == "receipt":
        _, _, receipt = saved_paths(owned)
        receipt.write_bytes(b"{}")
    else:
        view = replace(view, rendition_bytes=view.rendition_bytes + b"changed synthetic bytes")
    with pytest.raises(library.LibraryError):
        library.verified_render_source_view(draft, path, view)


@pytest.mark.parametrize('change',['members','aliases','scope','stored_pin'])
def test_render_caller_cannot_replace_persisted_ownership(owned,change):
    path,view,profile,draft=owned
    save(owned)
    if change=='members':
        altered=replace(profile,facilities=(Facility('replacement','Synthetic Regional Community Health Hospital'),))
    elif change=='aliases':
        altered=replace(profile,facilities=(replace(profile.facilities[0],aliases=('Synthetic Other',)),))
    elif change=='scope':
        altered=replace(profile,scope_type='multi_site')
    else:
        directory,manifest,_=saved_paths(owned)
        manifest['profile']['imported_from']='0'*64
        (directory/'manifest.json').write_bytes(library._json(manifest))
        altered=profile
    with pytest.raises(library.LibraryError,match='persisted profile'):
        library.verified_render_source_view(replace(draft,profile=altered),path,view)


def test_render_recognizes_other_authoritative_source_sites(owned,monkeypatch):
    from app import monthly_report_site_catalog as catalog
    from docx.oxml.ns import qn
    path,_,profile,draft=owned
    document=Document(path)
    title=next(n for n in document.element.iter(qn('w:t'))
               if (n.text or '').startswith('Synthetic Regional Community'))
    title.text = title.text.replace('Hospital', 'Hospital and Independent Synthetic Children Hospital', 1)
    document.save(path)
    view=derive_compatibility_view(path.read_bytes())
    profile=replace(profile,imported_from=view.original_inspection.sha256)
    draft=replace(draft,profile=profile,blocks=())
    changed=path,view,profile,draft
    save(changed)
    monkeypatch.setattr(catalog,'site_options',lambda contract:(
        Facility('other','Independent Synthetic Children Hospital'),))
    with pytest.raises(library.LibraryError,match='same-site ownership'):
        library.verified_render_source_view(draft,path,view)
