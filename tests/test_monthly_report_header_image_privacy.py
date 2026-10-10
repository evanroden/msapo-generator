"""A fresh master cannot lend a previous client's embedded header/footer PNG."""
from io import BytesIO
from zipfile import ZipFile

from docx import Document
from docx.shared import Inches
from PIL import Image, PngImagePlugin

from app.monthly_report_model import ResolvedBlock
from app.monthly_report_native_layout import build_native_docx
from test_monthly_report_native_layout import source, draft


CANARY = "FOREIGN-CLIENT-HEADER-PNG-ZX731"


def png(color, *, note=""):
    out = BytesIO()
    kwargs = {}
    if note:
        metadata = PngImagePlugin.PngInfo()
        metadata.add_text("OtherClient", note)
        kwargs["pnginfo"] = metadata
    Image.new("RGB", (480, 90), color).save(out, format="PNG", **kwargs)
    return out.getvalue()


def package_members(raw):
    with ZipFile(BytesIO(raw)) as archive:
        return {name: archive.read(name) for name in archive.namelist()}


def test_foreign_client_header_footer_images_and_png_metadata_do_not_cross(tmp_path):
    path = source(tmp_path)
    doc = Document(path)
    foreign = png("red", note=CANARY)
    doc.sections[0].header.paragraphs[0].add_run().add_picture(BytesIO(foreign), width=Inches(1.2))
    doc.sections[0].footer.paragraphs[0].add_run().add_picture(BytesIO(foreign), width=Inches(1.2))
    doc.save(path)
    result = package_members(build_native_docx(draft(), path, master=True))
    assert foreign not in result.values()
    assert not any(CANARY.encode() in data for data in result.values())
    assert any("word/header" in name for name in result)
    assert any("word/footer" in name for name in result)
    assert b"MONTHLY ACTIVITY SUMMARY" in result["word/document.xml"]


def test_explicit_approved_current_enfra_header_art_survives(tmp_path):
    path = source(tmp_path)
    doc = Document(path)
    foreign = png("red", note=CANARY)
    approved = png("green")
    doc.sections[0].header.paragraphs[0].add_run().add_picture(BytesIO(foreign), width=Inches(2))
    doc.save(path)
    new_brand = ResolvedBlock("brand_logo", "Library", asset_hashes=("approved-header.png",))
    result = package_members(build_native_docx(draft(new_brand), path, master=True,
                           asset_loader=lambda _reference: approved))
    assert foreign not in result.values()
    assert approved in result.values()
    assert not any(CANARY.encode() in data for data in result.values())
