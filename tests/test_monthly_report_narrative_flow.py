"""Current narrative is body content, not a title or an old list canvas."""
from copy import deepcopy
from dataclasses import replace
from io import BytesIO
from zipfile import ZipFile

import pymupdf
import pytest
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from conftest import requires_libreoffice
from app import pdf_converter
from app.monthly_report_import import ImportMapping, inspect_docx, map_items
from app.monthly_report_model import (BlockSpec, Facility, ReportDraft, ReportPeriod,
    ReportProfile, ResolvedBlock, SectionSpec)
from app.monthly_report_native_layout import build_native_docx

TITLE = 'Client Utility Rate Analysis'
METRICS = 'Current September metrics: 3 total, 2 completed, 1 open; closure rate 66.7 percent.'
SECOND = 'Current second paragraph: results remain pending and no savings are asserted.'


def textbox(paragraph, texts, *, y=26.45, height=27.25, header=False, fallback=True):
    """Synthetic native absolute-positioned textbox with real compatibility XML."""
    style = 'font-size:22pt' if header else 'font-size:11pt'
    pieces = []
    for index, text in enumerate(texts):
        p = OxmlElement('w:p')
        props = OxmlElement('w:pPr'); p.append(props)
        if not header and index:
            num = OxmlElement('w:numPr')
            level = OxmlElement('w:ilvl'); level.set(qn('w:val'), '0'); num.append(level)
            identity = OxmlElement('w:numId'); identity.set(qn('w:val'), '1'); num.append(identity)
            props.append(num)
        run = OxmlElement('w:r'); p.append(run)
        rpr = OxmlElement('w:rPr'); run.append(rpr)
        font = OxmlElement('w:rFonts'); font.set(qn('w:ascii'), 'Calibri'); font.set(qn('w:hAnsi'), 'Calibri'); rpr.append(font)
        size = OxmlElement('w:sz'); size.set(qn('w:val'), '44' if header else '22'); rpr.append(size)
        t = OxmlElement('w:t'); t.text = text; run.append(t)
        pieces.append(p)
    # Same native geometry in the two representations, not two independent titles.
    ns = '''xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
      xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
      xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"
      xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape"
      xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006"
      xmlns:v="urn:schemas-microsoft-com:vml"'''
    node = parse_xml(f'''<mc:AlternateContent {ns}><mc:Choice Requires="wps"><w:drawing>
      <wp:anchor distT="0" distB="0" distL="0" distR="0" simplePos="0" relativeHeight="1" behindDoc="1" locked="0" layoutInCell="1" allowOverlap="1">
      <wp:simplePos x="0" y="0"/><wp:positionH relativeFrom="page"><wp:posOffset>457200</wp:posOffset></wp:positionH>
      <wp:positionV relativeFrom="page"><wp:posOffset>{round(y*12700)}</wp:posOffset></wp:positionV>
      <wp:extent cx="8509000" cy="{round(height*12700)}"/><wp:effectExtent l="0" t="0" r="0" b="0"/>
      <wp:wrapNone/><wp:docPr id="{24 if header else 25}" name="Synthetic text frame"/><wp:cNvGraphicFramePr/>
      <a:graphic><a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">
      <wps:wsp><wps:cNvSpPr txBox="1"/><wps:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="8509000" cy="{round(height*12700)}"/></a:xfrm>
      <a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:noFill/><a:ln><a:noFill/></a:ln></wps:spPr>
      <wps:txbx><w:txbxContent/></wps:txbx><wps:bodyPr lIns="0" tIns="0" rIns="0" bIns="0"><a:noAutofit/></wps:bodyPr>
      </wps:wsp></a:graphicData></a:graphic></wp:anchor></w:drawing></mc:Choice>
      <mc:Fallback><w:pict><v:shape id="synthetic-frame" type="#_x0000_t202" filled="f" stroked="f"
      style="position:absolute;margin-left:36pt;margin-top:{y}pt;width:670pt;height:{height}pt;mso-position-horizontal-relative:page;mso-position-vertical-relative:page;{style}">
      <v:textbox inset="0,0,0,0"><w:txbxContent/></v:textbox></v:shape></w:pict></mc:Fallback></mc:AlternateContent>''')
    for box in node.iter(qn('w:txbxContent')):
        box.extend(deepcopy(p) for p in pieces)
    if not fallback:
        node.remove(node[-1])
    paragraph.add_run()._r.append(node)


def scene(tmp_path, *, text=METRICS, issue_text='Current equipment finding.\nCurrent follow-up remains open.', header_variant='default'):
    document = Document()
    document.add_paragraph('Previous Synthetic Site')
    document.add_paragraph('August 2026')
    document.add_paragraph('Prepared by: Previous Synthetic Editor')
    document.add_section(WD_SECTION_START.NEW_PAGE)
    document.add_heading('MONTHLY SCORECARDS', 1)
    section = document.add_section(WD_SECTION_START.NEW_PAGE)
    section.page_width, section.page_height = Inches(11), Inches(8.5)
    section.top_margin = section.bottom_margin = Inches(.5)
    section.left_margin = section.right_margin = Inches(.5)
    section.header_distance = Inches(.3)
    header = {'default': section.header, 'first': section.first_page_header, 'even': section.even_page_header}[header_variant]
    header.is_linked_to_previous = False
    if header_variant == 'first':
        section.different_first_page_header_footer = True
    if header_variant == 'even':
        document.settings.odd_and_even_pages_header_footer = True
    textbox(header.paragraphs[0], (TITLE,), header=True)
    label = document.add_paragraph(TITLE + ' Author: Previous Synthetic Editor')
    level = OxmlElement('w:outlineLvl'); level.set(qn('w:val'), '0'); label._p.get_or_add_pPr().append(level)
    label.runs[0].font.size = Pt(8)
    body = document.add_paragraph('Previous utility narrative with enough words to be a body paragraph.')
    body.paragraph_format.space_before = Pt(12)
    body.runs[0].font.size = Pt(11)
    # A new section retains its own header; it must not inherit a global offset.
    divider = document.add_section(WD_SECTION_START.NEW_PAGE)
    divider.header.is_linked_to_previous = False
    document.add_heading('EQUIPMENT PERFORMANCE ISSUES', 1)
    issues = document.add_section(WD_SECTION_START.NEW_PAGE)
    issues.header.is_linked_to_previous = False
    issues.header.paragraphs[0].text = 'ENFRA'
    footer = issues.footer
    footer.is_linked_to_previous = False
    footer.paragraphs[0].text = 'enfrasolutions.com'
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE')
    footer.paragraphs[0]._p.append(field)
    textbox(document.add_paragraph(), ('ISSUES IDENTIFIED', *(f'Previous finding {i} that must not return.' for i in range(5))), y=90, height=390)
    source = tmp_path/'source.docx'; document.save(source)
    profile = ReportProfile('Synthetic Contract', 'north', 'Synthetic North', (Facility('north','Synthetic North'),))
    specs = (SectionSpec('scorecards','3','Monthly Scorecards',(BlockSpec('utility_analysis','rich_text'),)),
             SectionSpec('issues','8','Equipment Performance Issues',(BlockSpec('equipment_issues','rich_text'),)))
    draft = ReportDraft(profile, ReportPeriod(2026,9), 'Synthetic Editor', specs, (
        ResolvedBlock('utility_analysis','This month',text=text),
        ResolvedBlock('equipment_issues','This month',text=issue_text)))
    return source, draft


def output(tmp_path, *, text=METRICS, section='scorecards', header_variant='default'):
    source, draft = scene(tmp_path, text=text, header_variant=header_variant)
    raw = build_native_docx(draft,source,master=True,section_key=section)
    return source, draft, raw


def test_narrative_uses_body_typography_not_the_source_title(tmp_path):
    _, _, raw = output(tmp_path)
    doc = Document(BytesIO(raw))
    p = next(p for p in doc.paragraphs if METRICS in p.text)
    assert p.runs[0].font.size == Pt(11)
    assert p._p.find('./'+qn('w:pPr')+'/'+qn('w:outlineLvl')) is None
    assert 'Previous Synthetic Editor' not in ''.join(p.text for p in doc.paragraphs)


def test_changed_canvas_becomes_independent_current_paragraphs_without_old_bullets(tmp_path):
    _, draft, raw = output(tmp_path,section='issues')
    doc = Document(BytesIO(raw))
    paragraphs = [p for p in doc.paragraphs if 'Current ' in p.text]
    assert [p.text for p in paragraphs] == draft.blocks[1].text.splitlines()
    assert not list(doc.element.body.iter(qn('w:txbxContent')))
    assert 'Previous finding' not in doc.element.xml
    assert all(not any(t.get(qn('w:val')) != '0' for t in p._p.findall('.//'+qn('w:numId'))) for p in paragraphs)


def test_changed_body_only_gets_clearance_for_its_own_visible_header(tmp_path):
    _, _, raw = output(tmp_path)
    doc = Document(BytesIO(raw))
    assert doc.sections[-1].top_margin.pt >= 59.7
    assert doc.sections[0].top_margin.pt == 72
    _, _, raw_issues = output(tmp_path,section='issues')
    assert Document(BytesIO(raw_issues)).sections[-1].top_margin.pt == 36


def test_header_content_and_source_bytes_are_unchanged(tmp_path):
    source, draft = scene(tmp_path)
    original = source.read_bytes()
    raw = build_native_docx(draft,source,master=True,section_key='scorecards')
    after = Document(BytesIO(raw)).sections[-1].header._element.xml
    # The existing sanitizer may replace drawing names, but geometry/text stay.
    assert TITLE in after
    for tag in ('wp:positionH','wp:positionV','wp:extent','wp:wrapNone'):
        assert [(n.tag,dict(n.attrib),n.text) for n in Document(source).sections[2].header._element.iter(qn(tag))] == [(n.tag,dict(n.attrib),n.text) for n in Document(BytesIO(raw)).sections[-1].header._element.iter(qn(tag))]
    assert source.read_bytes() == original


def test_unchanged_reviewed_canvas_and_page_geometry_are_retained(tmp_path):
    source, draft = scene(tmp_path)
    inspection = inspect_docx(source)
    item = next(i for i in inspection.items if i.kind=='text' and i.suggested_slot=='equipment_issues')
    mapped = map_items(source,inspection,(ImportMapping(item.id,'equipment_issues'),))
    draft = replace(draft,blocks=mapped.blocks)
    raw = build_native_docx(draft,source,section_key='issues')
    before = list(Document(source).element.body)[item.position-1]
    after = next(e for e in Document(BytesIO(raw)).element.body if list(e.iter(qn('w:txbxContent'))))
    assert [(n.text,dict(n.attrib)) for n in before.iter(qn('w:t'))] == [(n.text,dict(n.attrib)) for n in after.iter(qn('w:t'))]
    assert len(list(after.iter(qn('w:numPr')))) == len(list(before.iter(qn('w:numPr'))))
    assert Document(BytesIO(raw)).sections[-1].top_margin.pt == 36


@requires_libreoffice
@pytest.mark.parametrize('paragraphs', [1, 3, 50])
def test_rendered_current_text_never_overlaps_the_running_utility_heading(tmp_path,monkeypatch,paragraphs):
    text='\n'.join(f'Current paragraph {n:03d}. '+METRICS+' '+SECOND for n in range(paragraphs))
    source, draft, raw = output(tmp_path,text=text)
    monkeypatch.setenv('EPC_DATA_DIR',str(tmp_path/'runtime'))
    monkeypatch.setattr(pdf_converter,'OUTPUT_DIR',tmp_path)
    path=tmp_path/'current.docx';path.write_bytes(raw)
    produced=pdf_converter.convert_to_pdf(path)
    try:
        with pymupdf.open(produced) as pdf:
            words=' '.join(page.get_text() for page in pdf)
            for n in range(paragraphs): assert words.count(f'Current paragraph {n:03d}.')==1
            for page in pdf:
                titles=page.search_for(TITLE)
                current=page.search_for('Current paragraph')
                if current:
                    assert titles
                    assert min(rect.y0 for rect in current) >= max(rect.y1 for rect in titles)+2
                    assert all(page.rect.contains(r) for r in current)
    finally:produced.unlink(missing_ok=True)


@requires_libreoffice
def test_rendered_issue_narratives_do_not_restore_empty_list_markers(tmp_path,monkeypatch):
    _,draft,raw=output(tmp_path,section='issues')
    monkeypatch.setenv('EPC_DATA_DIR',str(tmp_path/'runtime'))
    monkeypatch.setattr(pdf_converter,'OUTPUT_DIR',tmp_path)
    path=tmp_path/'issues.docx';path.write_bytes(raw)
    produced=pdf_converter.convert_to_pdf(path)
    try:
        with pymupdf.open(produced) as pdf:
            text='\n'.join(p.get_text() for p in pdf)
            assert len(pdf)==2
            for sentence in draft.blocks[1].text.splitlines():assert text.count(sentence)==1
            assert not any(mark in text for mark in ('•','\uf0b7','Previous finding'))
            assert 'enfrasolutions.com' in text
    finally:produced.unlink(missing_ok=True)


def test_body_style_does_not_inherit_an_outline_from_an_empty_derived_style():
    from docx.enum.style import WD_STYLE_TYPE
    from app.monthly_report_native_text import narrative_prototype
    doc = Document()
    base = doc.styles.add_style('SyntheticHeadingBase', WD_STYLE_TYPE.PARAGRAPH)
    level = OxmlElement('w:outlineLvl'); level.set(qn('w:val'), '0')
    base.element.get_or_add_pPr().append(level)
    derived = doc.styles.add_style('SyntheticHeadingChild', WD_STYLE_TYPE.PARAGRAPH)
    derived.base_style = base
    p = doc.add_paragraph('This is a title, not a narrative template.', style=derived)
    styles = {s.style_id: s.element for s in doc.styles}
    assert narrative_prototype(p._p, styles, set()) is None
    level = OxmlElement('w:outlineLvl'); level.set(qn('w:val'), '9')
    p._p.get_or_add_pPr().append(level)
    assert narrative_prototype(p._p, styles, set()) is not None


@pytest.mark.parametrize('branch', ['drawingml', 'vml'])
@pytest.mark.parametrize('hidden', [True, False])
def test_only_visible_unambiguous_header_geometry_reserves_space(branch, hidden):
    from app.monthly_report_native_text import reserve_running_header
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(.5)
    textbox(section.header.paragraphs[0], (TITLE,), y=45, height=30, header=True)
    header = section.header._element
    alternate = next(header.iter('{http://schemas.openxmlformats.org/markup-compatibility/2006}AlternateContent'))
    alternate.remove(alternate[-1] if branch == 'drawingml' else alternate[0])
    if hidden:
        if branch == 'drawingml':
            next(header.iter(qn('wp:docPr'))).set('hidden', '1')
        else:
            shape = next(header.iter('{urn:schemas-microsoft-com:vml}shape'))
            shape.set('style', shape.get('style') + ';visibility:hidden')
    before = header.xml
    reserve_running_header(doc, section._sectPr)
    assert section.top_margin.pt == (36 if hidden else 81)
    assert header.xml == before


@pytest.mark.parametrize('variant', ['first', 'even'])
def test_header_clearance_uses_the_actual_variant_binding(tmp_path, variant):
    source, draft, raw = output(tmp_path, header_variant=variant)
    result = Document(BytesIO(raw))
    assert result.sections[-1].top_margin.pt >= 59.7
    header = result.sections[-1].first_page_header if variant == 'first' else result.sections[-1].even_page_header
    assert TITLE in header._element.xml


def test_no_header_keeps_the_safe_utility_title_separate_from_current_text(tmp_path):
    source, draft = scene(tmp_path)
    doc = Document(source)
    header = doc.sections[2].header._element
    for paragraph in header: paragraph[:] = []
    doc.save(source)
    raw = build_native_docx(draft, source, master=True, section_key='scorecards')
    result = Document(BytesIO(raw))
    titles = [p for p in result.paragraphs if p.text == TITLE]
    assert len(titles) == 1
    assert next(p for p in result.paragraphs if METRICS in p.text).text == METRICS
    assert 'Previous Synthetic Editor' not in result.element.xml
    assert result.sections[-1].top_margin.pt == 36


@pytest.mark.parametrize('kind', ['paragraph-relative', 'full-page', 'empty', 'negative', 'malformed', 'already-clear'])
def test_header_clearance_does_not_guess_from_unsupported_or_unneeded_bounds(kind):
    from app.monthly_report_native_text import reserve_running_header
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1.5 if kind == 'already-clear' else .5)
    textbox(section.header.paragraphs[0], (TITLE,), y=45, height=30, header=True, fallback=False)
    anchor = next(section.header._element.iter(qn('wp:anchor')))
    if kind == 'paragraph-relative':
        anchor.find(qn('wp:positionV')).set('relativeFrom', 'paragraph')
    if kind == 'full-page':
        anchor.find(qn('wp:extent')).set('cy', str(Inches(11)))
    if kind in ('negative', 'malformed'):
        anchor.find('./' + qn('wp:positionV') + '/' + qn('wp:posOffset')).text = '-100000' if kind == 'negative' else 'unknown'
    if kind == 'empty':
        for node in anchor.iter(qn('w:t')): node.text = ''
    original = section._sectPr.xml
    reserve_running_header(doc, section._sectPr)
    assert section._sectPr.xml == original


@requires_libreoffice
def test_final_report_api_keeps_current_narratives_and_native_header_geometry(tmp_path, monkeypatch):
    from app import monthly_report_designs as designs
    from app.monthly_report_docx import generate_report
    monkeypatch.setenv('EPC_DATA_DIR', str(tmp_path / 'runtime'))
    monkeypatch.setattr(pdf_converter, 'OUTPUT_DIR', tmp_path)
    source, draft = scene(tmp_path)
    source_bytes = source.read_bytes()
    profile = designs.install(source, draft.profile, actor=draft.prepared_by, expected_revision=0)
    draft = replace(draft, profile=profile)
    package = generate_report(draft, acknowledged_fingerprint=draft.fingerprint)
    assert package.pdf and not package.pdf_error
    with pymupdf.open(stream=package.pdf, filetype='pdf') as pdf:
        assert len(pdf) == 5
        text = '\n'.join(page.get_text() for page in pdf)
        assert text.count(METRICS) == 1
        for sentence in draft.blocks[1].text.splitlines(): assert text.count(sentence) == 1
        assert not any(mark in text for mark in ('•', '\uf0b7', 'Previous finding'))
        page = next(page for page in pdf if page.search_for('Current September metrics'))
        assert page.search_for('Current September metrics')[0].y0 > page.search_for(TITLE)[0].y1 + 2
    assert source.read_bytes() == source_bytes
    document = Document(BytesIO(package.docx))
    assert METRICS in [p.text for p in document.paragraphs]
    assert all(not list(p._p.iter(qn('w:txbxContent'))) for p in document.paragraphs if 'Current ' in p.text)
    assert 'PAGE' in document.sections[-1].footer._element.xml


def test_changed_narrative_and_table_payloads_are_independently_retained(tmp_path):
    from app.monthly_report_model import ReportTable
    source, draft = scene(tmp_path)
    doc = Document(source)
    table = doc.add_table(rows=2, cols=2)
    for row, values in zip(table.rows, (('Equipment', 'Status'), ('Previous asset', 'Previous status'))):
        for cell, value in zip(row.cells, values): cell.text = value
    doc.save(source)
    current = replace(draft.blocks[1], extra_tables=(ReportTable(('Equipment', 'Status'), (('Synthetic pump', 'Pending'),)),))
    draft = replace(draft, blocks=(draft.blocks[0], current))
    raw = build_native_docx(draft, source, master=True, section_key='issues')
    output = Document(BytesIO(raw))
    assert [p.text for p in output.paragraphs if 'Current ' in p.text] == current.text.splitlines()
    assert [[cell.text for cell in row.cells] for row in output.tables[-1].rows] == [['Equipment', 'Status'], ['Synthetic pump', 'Pending']]
    assert 'Previous asset' not in output.element.xml


@requires_libreoffice
@pytest.mark.parametrize('kind', ['footnote', 'endnote'])
@pytest.mark.parametrize('retain', [True, False])
def test_rewriting_a_noted_paragraph_preserves_only_explicitly_retained_notes(tmp_path, monkeypatch, retain, kind):
    from lxml import etree
    source, draft = scene(tmp_path)
    with ZipFile(source) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    root = etree.fromstring(parts['word/document.xml'])
    p = next(p for p in root.iter(qn('w:p')) if 'Previous utility narrative' in ''.join(p.itertext()))
    run = etree.SubElement(p, qn('w:r'))
    etree.SubElement(run, qn('w:' + kind + 'Reference'), {qn('w:id'): '1'})
    parts['word/document.xml'] = etree.tostring(root)
    notes = etree.Element(qn('w:' + kind + 's'), nsmap={'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'})
    note = etree.SubElement(notes, qn('w:' + kind), {qn('w:id'): '1'})
    etree.SubElement(etree.SubElement(etree.SubElement(note, qn('w:p')), qn('w:r')), qn('w:t')).text = 'Explicitly retained synthetic note.'
    parts['word/' + kind + 's.xml'] = etree.tostring(notes)
    rels = etree.fromstring(parts['word/_rels/document.xml.rels'])
    etree.SubElement(rels, '{http://schemas.openxmlformats.org/package/2006/relationships}Relationship',
                     Id='rSyntheticNotes', Type='http://schemas.openxmlformats.org/officeDocument/2006/relationships/' + kind + 's', Target=kind + 's.xml')
    parts['word/_rels/document.xml.rels'] = etree.tostring(rels)
    types = etree.fromstring(parts['[Content_Types].xml'])
    etree.SubElement(types, '{http://schemas.openxmlformats.org/package/2006/content-types}Override',
                     PartName='/word/' + kind + 's.xml', ContentType='application/vnd.openxmlformats-officedocument.wordprocessingml.' + kind + 's+xml')
    parts['[Content_Types].xml'] = etree.tostring(types)
    with ZipFile(source, 'w') as archive:
        for name, data in parts.items(): archive.writestr(name, data)
    inspection = inspect_docx(source)
    note = next(i for i in inspection.items if i.note == kind + ':1')
    mapped = map_items(source, inspection, (ImportMapping(note.id, 'utility_analysis'),))
    current = replace(mapped.blocks[0], text=METRICS + ('\n' + note.text if retain else ''),
                      references=mapped.blocks[0].references if retain else ())
    draft = replace(draft, blocks=(current, draft.blocks[1]))
    raw = build_native_docx(draft, source, master=True, section_key='scorecards')
    with ZipFile(BytesIO(raw)) as archive:
        assert ('word/' + kind + 's.xml' in archive.namelist()) == retain
        body = etree.fromstring(archive.read('word/document.xml'))
        assert len(list(body.iter(qn('w:' + kind + 'Reference')))) == int(retain)
        assert METRICS in ''.join(body.itertext())
        if retain:
            assert b'Explicitly retained synthetic note.' in archive.read('word/' + kind + 's.xml')
            assert b'Explicitly retained synthetic note.' not in archive.read('word/document.xml')
    monkeypatch.setenv('EPC_DATA_DIR', str(tmp_path / 'runtime'))
    monkeypatch.setattr(pdf_converter, 'OUTPUT_DIR', tmp_path)
    output = tmp_path / 'noted.docx'; output.write_bytes(raw)
    pdf_path = pdf_converter.convert_to_pdf(output)
    try:
        with pymupdf.open(pdf_path) as pdf:
            text = '\n'.join(page.get_text() for page in pdf)
            assert text.count('Explicitly retained synthetic note.') == int(retain)
            assert text.count(METRICS) == 1
    finally:
        pdf_path.unlink(missing_ok=True)


@requires_libreoffice
@pytest.mark.parametrize('paragraphs', [2, 20])
def test_current_prose_flows_once_through_all_twelve_sections(tmp_path, monkeypatch, paragraphs):
    from app.monthly_report_model import default_sections
    from app.monthly_report_native_layout import _DEFAULT_TEXT
    source = tmp_path / 'all-sections.docx'
    doc = Document()
    doc.add_paragraph('Previous Synthetic Site')
    doc.add_paragraph('August 2026')
    sections, blocks = [], []
    expected = []
    for index, spec in enumerate(default_sections()):
        doc.add_section(WD_SECTION_START.NEW_PAGE)
        doc.add_heading(spec.title.upper(), 1)
        doc.add_section(WD_SECTION_START.NEW_PAGE)
        old = doc.add_paragraph('Previous source narrative with substantive details to replace.')
        old.runs[0].font.size = Pt(11)
        key = _DEFAULT_TEXT[spec.key]
        texts = [f'Current section {index:02d} paragraph {n:02d}. Reviewed synthetic observations and current follow-up information.' for n in range(paragraphs)]
        expected.extend(texts)
        sections.append(replace(spec, blocks=(BlockSpec(key, 'rich_text'),)))
        blocks.append(ResolvedBlock(key, 'This month', text='\n'.join(texts)))
    doc.save(source)
    profile = ReportProfile('Synthetic Contract', 'north', 'Synthetic North', (Facility('north','Synthetic North'),))
    draft = ReportDraft(profile, ReportPeriod(2026, 9), 'Synthetic Editor', tuple(sections), tuple(blocks))
    raw = build_native_docx(draft, source, master=True)
    current = Document(BytesIO(raw))
    assert [p.text for p in current.paragraphs if p.text.startswith('Current section')] == expected
    assert 'Previous source narrative' not in current.element.xml
    monkeypatch.setenv('EPC_DATA_DIR', str(tmp_path / 'runtime'))
    monkeypatch.setattr(pdf_converter, 'OUTPUT_DIR', tmp_path)
    path = tmp_path / 'current.docx'; path.write_bytes(raw)
    produced = pdf_converter.convert_to_pdf(path)
    try:
        with pymupdf.open(produced) as pdf:
            text = ' '.join(' '.join(page.get_text().split()) for page in pdf)
            for paragraph in expected:
                assert text.count(paragraph) == 1
            assert all(page.get_text().strip() for page in pdf)
    finally:
        produced.unlink(missing_ok=True)
