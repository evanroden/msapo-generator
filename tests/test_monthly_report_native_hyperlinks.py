"""Passive click targets survive only with current, source-proven visible text."""
from dataclasses import replace
from io import BytesIO
from zipfile import ZipFile
import shutil
import subprocess

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
import pytest

from app.monthly_report_import import inspect_docx, ImportMapping, map_items
from app.monthly_report_model import Facility, ReportDraft, ReportPeriod, ReportProfile, ResolvedBlock, default_sections
from app.monthly_report_native_layout import build_native_docx
from app.monthly_report_native_package import passive_docx

TARGET = 'https://example.invalid/current-procedure?revision=2#procedure'


def linked_report(tmp_path, *, target=TARGET, table=False, empty=False):
    doc = Document()
    doc.add_heading('Monthly Activity Summary', 1)
    if table:
        grid = doc.add_table(rows=2, cols=2)
        grid.cell(0, 0).text = 'Procedure'
        grid.cell(0, 1).text = 'Status'
        grid.cell(1, 1).text = 'Current'
        paragraph = grid.cell(1, 0).paragraphs[0]
    else:
        paragraph = doc.add_paragraph()
    link = OxmlElement('w:hyperlink')
    link.set(qn('r:id'), doc.part.relate_to(target, RT.HYPERLINK, is_external=True))
    run = OxmlElement('w:r'); text = OxmlElement('w:t')
    text.text = '' if empty else 'Current linked procedure'
    run.append(text); link.append(run); paragraph._p.append(link)
    path = tmp_path / 'linked.docx'
    doc.save(path)
    return path


def draft_for(path):
    inspection = inspect_docx(path)
    items = [item for item in inspection.items if item.kind in ('text', 'table') and item.section == 'activity']
    mappings = tuple(ImportMapping(item.id, 'activity_summary') for item in items if not item.note.startswith('Native section heading'))
    if any(item.kind == 'table' for item in items):
        from app.monthly_report_sections import table_without_prices
        tables = tuple(replace(table_without_prices(item)[0], reference=f'docx:{inspection.sha256}:{item.id}')
                       for item in items if item.kind == 'table')
        blocks = (ResolvedBlock('activity_summary', 'This month', extra_tables=tables,
                                references=tuple(table.reference for table in tables)),)
    else:
        blocks = map_items(path, inspection, mappings).blocks
    draft = ReportDraft(ReportProfile('Synthetic', 'site', 'Current site', (Facility('site', 'Current site'),)),
                        ReportPeriod(2026, 9), 'Current editor', default_sections(), blocks)
    return draft


def targets(raw):
    with ZipFile(BytesIO(raw)) as archive:
        return b' '.join(archive.read(name) for name in archive.namelist() if name.endswith('.rels'))


@pytest.mark.parametrize('table', [False, True])
def test_current_proven_hyperlink_remains_clickable_in_native_package(tmp_path, table):
    path = linked_report(tmp_path, table=table)
    result = build_native_docx(draft_for(path), path)
    assert TARGET.encode().replace(b'&', b'&amp;') in targets(result)


@pytest.mark.parametrize('change', ['changed', 'omitted', 'unproven'])
def test_rewritten_or_unproven_text_cannot_inherit_source_click_target(tmp_path, change):
    path = linked_report(tmp_path)
    draft = draft_for(path)
    block = next(b for b in draft.blocks if b.key == 'activity_summary')
    if change == 'changed':
        block = replace(block, text='New client procedure')
    elif change == 'omitted':
        block = replace(block, text='')
    else:
        block = replace(block, references=())
    result = build_native_docx(replace(draft, blocks=(block,)), path)
    assert TARGET.encode() not in targets(result)
    with ZipFile(BytesIO(result)) as archive:
        assert b'<w:hyperlink' not in archive.read('word/document.xml')


def test_changed_table_cell_does_not_inherit_previous_client_url(tmp_path):
    path = linked_report(tmp_path, table=True)
    draft = draft_for(path)
    block = draft.blocks[0]
    table = block.extra_tables[0]
    table = replace(table, rows=(('New client procedure', 'Current'),))
    result = build_native_docx(replace(draft, blocks=(replace(block, extra_tables=(table,)),)), path)
    assert TARGET.encode() not in targets(result)
    with ZipFile(BytesIO(result)) as archive:
        assert b'New client procedure' in archive.read('word/document.xml')


@pytest.mark.parametrize('target', ['file:///private/client.docx', r'\\server\private\client',
                                    'javascript:alert(1)', 'data:text/html,hello', 'ftp://example.invalid/file',
                                    'https://name:secret@example.invalid/', 'https://example.invalid/\nold',
                                    'https://example.invalid\\old', 'https:///no-host'])
def test_active_local_or_malformed_hyperlink_schemes_are_removed(tmp_path, target):
    raw = passive_docx(linked_report(tmp_path, target=target))
    with ZipFile(BytesIO(raw)) as archive:
        assert b'Current linked procedure' in archive.read('word/document.xml')
    assert b'TargetMode="External"' not in targets(raw)


def test_empty_hyperlink_relation_is_pruned(tmp_path):
    raw = passive_docx(linked_report(tmp_path, empty=True))
    assert TARGET.encode() not in targets(raw)


@pytest.mark.parametrize('flag', ['vanish', 'webHidden'])
def test_hidden_hyperlink_relation_is_pruned(tmp_path, flag):
    path = linked_report(tmp_path)
    doc = Document(path)
    run = next(doc.element.iter(qn('w:hyperlink'))).find(qn('w:r'))
    properties = OxmlElement('w:rPr')
    properties.append(OxmlElement('w:' + flag))
    run.insert(0, properties)
    doc.save(path)
    assert TARGET.encode() not in targets(passive_docx(path))


def test_native_current_link_reaches_exported_pdf_annotation(tmp_path):
    binary = shutil.which('libreoffice') or shutil.which('soffice')
    if binary is None:
        pytest.skip('LibreOffice is required for PDF annotation integration')
    source = linked_report(tmp_path)
    output = tmp_path / 'current.docx'
    output.write_bytes(build_native_docx(draft_for(source), source))
    result = subprocess.run([binary, '-env:UserInstallation=' + (tmp_path / 'profile').as_uri(),
                             '--headless', '--convert-to', 'pdf', '--outdir', str(tmp_path), str(output)],
                            capture_output=True, timeout=45)
    assert result.returncode == 0, result.stderr.decode(errors='replace')
    import fitz
    with fitz.open(output.with_suffix('.pdf')) as pdf:
        assert TARGET in [link.get('uri') for page in pdf for link in page.get_links()]
