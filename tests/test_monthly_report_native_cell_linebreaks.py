"""Deliberate cell line breaks are Word structure, not raw XML whitespace."""
from hashlib import sha256
from io import BytesIO

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches
from lxml import etree
import pymupdf
import pytest

from conftest import requires_libreoffice
from app.monthly_report_native_layout import _set_text


def _payload(element):
    values = []
    for node in element.iter():
        if node.tag == qn('w:t'):
            assert not any(char in (node.text or '') for char in '\r\n\t')
            values.append(node.text or '')
        elif node.getparent() is not None and node.getparent().tag == qn('w:r'):
            if node.tag == qn('w:cr') or (node.tag == qn('w:br') and
                    node.get(qn('w:type'), 'textWrapping') == 'textWrapping'):
                values.append('\n')
            elif node.tag == qn('w:tab'):
                values.append('\t')
    return ''.join(values)


@pytest.mark.parametrize('value', (
    'Inspect valve\nConfirm access',
    'Inspect valve\r\nConfirm access',
    'Inspect valve\rConfirm access',
    '\n  Keep leading spaces\n\nKeep trailing spaces  \n',
    'First\r\nSecond\rThird\n\tFourth',
))
def test_current_cell_breaks_preserve_exact_whitespace_and_native_formatting(value):
    cell = Document().add_table(rows=1, cols=1).cell(0, 0)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.left_indent = Inches(.13)
    paragraph.paragraph_format.tab_stops.add_tab_stop(Inches(.6))
    first = paragraph.add_run('OLD')
    first.bold = True
    first.add_break()
    first.add_tab()
    paragraph.add_run(' PRIVATE TEXT').italic = True
    paragraph_properties = etree.tostring(paragraph._p.pPr)
    run_properties = etree.tostring(first._r.rPr)

    _set_text(cell._tc, value)

    assert _payload(cell._tc) == value.replace('\r\n', '\n').replace('\r', '\n')
    assert etree.tostring(paragraph._p.pPr) == paragraph_properties
    assert etree.tostring(first._r.rPr) == run_properties
    assert len(paragraph.runs) == 2 and paragraph.runs[1].italic
    assert all(node.get(qn('xml:space')) == 'preserve' for node in first._r.iter(qn('w:t')))


def test_existing_raw_xml_newlines_are_repaired_even_when_inspector_text_matches():
    paragraph = Document().add_paragraph()
    text = OxmlElement('w:t')
    text.text = 'Inspect valve\nConfirm access'
    paragraph.add_run()._r.append(text)
    _set_text(paragraph._p, 'Inspect valve\nConfirm access')
    assert _payload(paragraph._p) == 'Inspect valve\nConfirm access'
    assert len(list(paragraph._p.iter(qn('w:br')))) == 1


def test_empty_and_repeated_replacement_retire_old_run_controls_without_changing_geometry():
    paragraph = Document().add_paragraph()
    paragraph.paragraph_format.tab_stops.add_tab_stop(Inches(.7))
    run = paragraph.add_run()
    run.add_break()
    run._r.append(OxmlElement('w:cr'))
    run.add_tab()
    run.add_break(WD_BREAK.PAGE)
    run.add_break(WD_BREAK.COLUMN)
    properties = etree.tostring(paragraph._p.pPr)
    for value in ('', 'Line one\nLine two', 'Single line', '', '\nAgain\n'):
        _set_text(paragraph._p, value)
        assert _payload(paragraph._p) == value
        assert etree.tostring(paragraph._p.pPr) == properties
        assert [node.get(qn('w:type')) for node in paragraph._p.iter(qn('w:br'))
                if node.get(qn('w:type')) in ('page', 'column')] == ['page', 'column']
        exact = etree.tostring(paragraph._p)
        _set_text(paragraph._p, value)
        assert etree.tostring(paragraph._p) == exact


def test_exact_single_line_replacement_does_not_use_whitespace_collapsing_equality():
    paragraph = Document().add_paragraph('  Ready  ')
    _set_text(paragraph._p, 'Ready')
    assert _payload(paragraph._p) == 'Ready'


def test_unchanged_native_lines_keep_their_distinct_run_formatting():
    paragraph = Document().add_paragraph()
    first = paragraph.add_run('Ready')
    first.bold = True
    first.add_break()
    paragraph.add_run('  For review ').italic = True
    original = etree.tostring(paragraph._p)
    _set_text(paragraph._p, 'Ready\n  For review ')
    assert etree.tostring(paragraph._p) == original


def test_blank_native_paragraph_receives_editable_lines():
    paragraph = Document().add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(.2)
    original_properties = etree.tostring(paragraph._p.pPr)
    _set_text(paragraph._p, 'First\n\nLast')
    assert _payload(paragraph._p) == 'First\n\nLast'
    assert len(list(paragraph._p.iter(qn('w:br')))) == 2
    assert etree.tostring(paragraph._p.pPr) == original_properties


def test_choice_and_fallback_receive_the_same_native_breaks_in_their_own_styled_runs():
    mc = 'http://schemas.openxmlformats.org/markup-compatibility/2006'
    root = OxmlElement('w:p')
    alternate = etree.SubElement(root, '{' + mc + '}AlternateContent')
    branches = []
    for name, emphasis in (('Choice', 'b'), ('Fallback', 'i')):
        branch = etree.SubElement(alternate, '{' + mc + '}' + name)
        paragraph = OxmlElement('w:p'); branch.append(paragraph)
        run = OxmlElement('w:r'); paragraph.append(run)
        props = OxmlElement('w:rPr'); props.append(OxmlElement('w:' + emphasis)); run.append(props)
        text = OxmlElement('w:t'); text.text = 'OLD'; run.append(text)
        run.append(OxmlElement('w:br'))
        branches.append((branch, etree.tostring(props)))
    _set_text(root, 'First\r\n\r\nLast')
    for branch, original_properties in branches:
        assert _payload(branch) == 'First\n\nLast'
        assert len(list(branch.iter(qn('w:br')))) == 2
        assert etree.tostring(next(branch.iter(qn('w:rPr')))) == original_properties


@requires_libreoffice
def test_generated_native_word_and_pdf_keep_cell_lines_distinct(tmp_path, monkeypatch):
    from app import monthly_report_designs as designs, pdf_converter
    from app.monthly_report_docx import generate_report
    from app.monthly_report_model import (
        BlockSpec, ColumnSpec, Facility, ReportDraft, ReportPeriod,
        ReportProfile, ResolvedBlock, SectionSpec,
    )
    monkeypatch.setenv('EPC_DATA_DIR', str(tmp_path / 'isolated'))
    monkeypatch.setattr(pdf_converter, 'OUTPUT_DIR', tmp_path)
    monkeypatch.setattr(pdf_converter, 'PDF_BACKEND', 'libreoffice')
    source = Document()
    source.add_heading('Pending & Declined Proposals', 1)
    table = source.add_table(rows=2, cols=2)
    for cell, value in zip(table.rows[0].cells, ('Work Description', 'Technical Notes')):
        cell.text = value
    for cell in table.rows[1].cells:
        cell.text = 'OLD PRIVATE VALUE\nOLD SECOND LINE'
    original = tmp_path / 'synthetic-native-source.docx'
    source.save(original)
    source_hash = sha256(original.read_bytes()).hexdigest()
    profile = ReportProfile('Synthetic Contract', 'site', 'Synthetic Site', (Facility('site', 'Synthetic Site'),))
    profile = designs.install(original, profile, actor='Synthetic Editor', expected_revision=0)
    spec = BlockSpec('proposals', 'table', columns=(
        ColumnSpec('work', 'Work Description'), ColumnSpec('notes', 'Technical Notes')))
    draft = ReportDraft(profile, ReportPeriod(2027, 1), 'Synthetic Editor',
        (SectionSpec('proposals', '10', 'Pending & Declined Proposals', (spec,)),),
        (ResolvedBlock('proposals', 'This month', rows=(('Inspect valve\nConfirm access', '45 gpm\r\n480 V\rReady'),)),))
    package = generate_report(draft, acknowledged_fingerprint=draft.fingerprint)
    assert package.pdf and not package.pdf_error
    with pymupdf.open(stream=package.pdf, filetype='pdf') as pdf:
        lines = [
            ''.join(span['text'] for span in line['spans']).strip()
            for page in pdf for block in page.get_text('dict')['blocks'] if 'lines' in block
            for line in block['lines']
        ]
        for expected in ('Inspect valve', 'Confirm access', '45 gpm', '480 V', 'Ready'):
            assert lines.count(expected) == 1, lines
        assert not any('OLD' in line for line in lines)
    actual = Document(BytesIO(package.docx)).tables[0]
    assert len(actual.rows) == 2 and len(actual.columns) == 2
    assert tuple(cell.text for cell in actual.rows[0].cells) == ('Work Description', 'Technical Notes')
    assert _payload(actual.cell(1, 0)._tc) == 'Inspect valve\nConfirm access'
    assert _payload(actual.cell(1, 1)._tc) == '45 gpm\n480 V\nReady'
    assert sha256(original.read_bytes()).hexdigest() == source_hash
