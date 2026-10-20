"""Independent synthetic cases for the three open PR264 contact reviews.

AppTest delivers native widget values, not actual browser timing. Persistence
and editable Word are real. The native queue journey remains a separate gate.
"""
from dataclasses import replace
from io import BytesIO

from docx import Document
import pytest
from streamlit.testing.v1 import AppTest

from app import monthly_report_active_work as active
from app.monthly_report_docx import assemble_docx
from app.monthly_report_model import BlockSpec, ColumnSpec, ReportTable, ResolvedBlock


SCRIPT = '''
from dataclasses import replace
import streamlit as st
from app.monthly_report_model import ReportDraft, ReportPeriod, default_sections, synthetic_profiles
from app.monthly_report_organization_ui import render_organization
from app.monthly_report_ui import _field
block = st.session_state.setdefault('block', st.session_state['seed_block'])
spec = st.session_state['seed_spec']
section = next(section for section in default_sections() if section.key == 'organization')
section = replace(section, blocks=(spec,))
draft = ReportDraft(synthetic_profiles()[0], ReportPeriod(2026, 9), 'Synthetic Editor', (section,), (block,))
changed = render_organization(draft, {block.key:block}, {block.key:spec}, 'contact_review', {},
    _field, lambda spec, block:block, show_preview=False)
st.session_state['block'] = changed[block.key]
st.session_state['draft'] = replace(draft, blocks=(changed[block.key],))
'''


FIRST = ReportTable(('Role', 'Name', 'Phone'), (('Operator', 'Alex Example', '202-555-0101'),), 'source-a:table-1')
SECOND = ReportTable(('Name', 'Role', 'Phone'), (('Alex Example', 'Supervisor', '202-555-0102'),), 'source-b:table-2')


def workspace(block=None, spec=None):
    app = AppTest.from_string(SCRIPT, default_timeout=30)
    app.session_state['seed_block'] = block or ResolvedBlock('contact_matrix', 'This month',
        text='Keep the ordinary contact instructions.', extra_tables=(FIRST, SECOND))
    app.session_state['seed_spec'] = spec or BlockSpec('contact_matrix', 'image_page')
    app.run()
    assert not app.exception
    return app


def button(app, label, index=0):
    return [item for item in app.button if item.label == label][index]


def field(app, label):
    return next(item for item in app.text_input if item.label == label)


@pytest.mark.parametrize('start', [0, 1])
def test_switching_imported_table_commits_pending_row_and_leaves_one_active_editor(monkeypatch, tmp_path, start):
    monkeypatch.setenv('EPC_DATA_DIR', str(tmp_path / 'data'))
    app = workspace()
    original = app.session_state['block']
    button(app, 'Edit contact', start).click().run()
    native_id = field(app, 'Name').proto.id
    field(app, 'Name').set_value('Deliberate first edit')
    next(item for item in app.text_area if item.label == 'Notes included with this part').set_value('Pending contact notes.')
    button(app, 'Edit contact').click().run()
    assert not app.exception
    result = app.session_state['block']
    assert result.extra_tables[start].rows[0][(1, 0)[start]] == 'Deliberate first edit'
    assert result.extra_tables[1-start] == original.extra_tables[1-start]
    assert result.text == 'Pending contact notes.'
    assert len([item for item in app.text_input if item.label == 'Name']) == 1
    assert field(app, 'Name').value == 'Alex Example'
    assert field(app, 'Role').value == ('Supervisor', 'Operator')[start]
    button(app, 'Edit contact').click().run()
    assert field(app, 'Name').value == 'Deliberate first edit'
    assert field(app, 'Name').proto.id == native_id
    draft = app.session_state['draft']
    saved = active.save('4' * 32, draft, expected_revision=0, base_snapshot_revision=0)
    loaded = active.load('4' * 32, draft.profile.contract, draft.profile.key, draft.period)
    assert saved.revision == loaded.revision and loaded.draft == draft
    fresh = workspace(loaded.draft.blocks[0])
    assert fresh.session_state['block'] == result and not fresh.text_input


def test_add_in_another_table_captures_the_open_row_and_cancel_drops_only_blank_draft():
    app = workspace()
    button(app, 'Edit contact').click().run()
    field(app, 'Phone').set_value('202-555-0191')
    button(app, 'Add contact', 1).click().run()
    assert not app.exception
    assert app.session_state['block'].extra_tables[0].rows[0][-1] == '202-555-0191'
    assert app.session_state['block'].extra_tables[1] == SECOND
    assert len([item for item in app.text_input if item.label == 'Name']) == 1
    assert field(app, 'Name').value == ''
    button(app, 'Close editing').click().run()
    assert not app.text_input
    assert app.session_state['block'].extra_tables[1] == SECOND


def test_an_unrelated_table_remove_and_undo_preserve_the_active_row_and_its_edit():
    app = workspace()
    button(app, 'Edit contact').click().run()
    identity = field(app, 'Name').proto.id
    field(app, 'Name').set_value('Retain with other-table removal')
    button(app, 'Remove contact', 1).click().run()
    assert not app.exception
    assert field(app, 'Name').proto.id == identity
    assert app.session_state['block'].extra_tables[0].rows[0][1] == 'Retain with other-table removal'
    assert not app.session_state['block'].extra_tables[1].rows
    button(app, 'Undo remove contact').click().run()
    assert not app.exception
    assert field(app, 'Name').proto.id == identity
    assert app.session_state['block'].extra_tables[1] == SECOND


def test_unrelated_table_action_does_not_expire_an_unchanged_row_removal_undo():
    first = replace(FIRST, rows=(*FIRST.rows, ('Assistant', 'Blair Example', '202-555-0103')))
    app = workspace(ResolvedBlock('contact_matrix', 'This month', extra_tables=(first, SECOND)))
    button(app, 'Remove contact', 1).click().run()
    button(app, 'Edit contact').click().run()
    button(app, 'Remove contact', 1).click().run()
    assert not app.exception
    assert len([item for item in app.button if item.label == 'Undo remove contact']) == 2
    button(app, 'Undo remove contact').click().run()
    assert not app.exception
    assert app.session_state['block'].extra_tables[0] == first
    assert not app.session_state['block'].extra_tables[1].rows
    button(app, 'Undo remove contact').click().run()
    assert app.session_state['block'].extra_tables == (first, SECOND)


def test_multiple_blank_imported_tables_open_only_one_uncounted_draft():
    tables = tuple(replace(table, rows=(tuple('' for _ in table.columns),)) for table in (FIRST, SECOND))
    app = workspace(ResolvedBlock('contact_matrix', 'This month', extra_tables=tables))
    assert len([item for item in app.text_input if item.label == 'Name']) == 1
    assert all(not table.rows for table in app.session_state['block'].extra_tables)
    button(app, 'Close editing').click().run()
    assert not app.text_input
    assert tuple(table.columns for table in app.session_state['block'].extra_tables) == tuple(table.columns for table in tables)


def test_later_imported_blank_table_keeps_its_first_delivered_field_open():
    app = workspace()
    button(app, 'Edit contact').click().run()
    button(app, 'Close editing').click().run()
    current = app.session_state['block']
    new_table = ReportTable(('Name', 'Role'), (('', ''),), 'source-c:table-3')
    app.session_state['block'] = replace(current, extra_tables=(*current.extra_tables, new_table))
    app.run()
    identity = field(app, 'Name').proto.id
    field(app, 'Name').set_value('New imported person').run()
    assert not app.exception
    assert field(app, 'Name').proto.id == identity
    field(app, 'Role').set_value('Technician')
    button(app, 'Close editing').click().run()
    assert app.session_state['block'].extra_tables[:2] == current.extra_tables
    assert app.session_state['block'].extra_tables[2].rows == (('New imported person', 'Technician'),)


def test_changed_same_table_stays_a_real_conflict_when_another_table_is_opened():
    app = workspace()
    button(app, 'Edit contact').click().run()
    field(app, 'Name').set_value('Pending operator wording')
    current = app.session_state['block']
    newer = replace(FIRST, rows=(('Operator', 'Newer canonical person', '202-555-0130'),))
    app.session_state['block'] = replace(current, extra_tables=(newer, SECOND))
    button(app, 'Edit contact').click().run()
    assert not app.exception
    assert app.session_state['block'].extra_tables == (newer, SECOND)
    assert any('Both versions are kept' in item.value for item in app.warning)
    assert any('Pending operator wording' in frame.value.to_string() for frame in app.dataframe)
    assert len([item for item in app.text_input if item.label == 'Name']) == 1
    assert field(app, 'Role').value == 'Supervisor'


@pytest.mark.parametrize('shape', ['imported', 'primary', 'empty_schema'])
def test_omitted_contacts_are_explicit_before_the_retained_native_cards(shape):
    spec = BlockSpec('contact_matrix', 'table', columns=(ColumnSpec('role', 'Role'), ColumnSpec('name', 'Name')))
    block = ResolvedBlock('contact_matrix', 'Omit')
    if shape == 'imported':
        block = replace(block, extra_tables=(FIRST,))
    elif shape == 'primary':
        block = replace(block, rows=(('Operator', 'Alex Example'),))
    app = workspace(block, spec)
    assert app.session_state['block'] == block
    assert any(item.value.startswith('Not included in this report.') for item in app.caption)
    assert any(item.label == ('Add contact' if shape == 'empty_schema' else 'Edit contact') for item in app.button)


@pytest.mark.parametrize('heading', ['', '   '])
def test_blank_column_gets_a_display_label_and_exact_schema_survives_word_and_reload(monkeypatch, tmp_path, heading):
    monkeypatch.setenv('EPC_DATA_DIR', str(tmp_path / 'data'))
    table = ReportTable(('Name', heading, 'Role'), (('Alex Example', 'Retained custom detail', 'Operator'),))
    block = ResolvedBlock('contact_matrix', 'This month', extra_tables=(table,))
    app = workspace(block)
    button(app, 'Edit contact').click().run()
    assert {item.label for item in app.text_input} == {'Name', 'Detail 2', 'Role'}
    field(app, 'Detail 2').set_value('Corrected custom detail')
    button(app, 'Close editing').click().run()
    assert not app.exception
    draft = app.session_state['draft']
    assert draft.blocks[0].extra_tables[0].columns == table.columns
    assert draft.blocks[0].extra_tables[0].rows == (('Alex Example', 'Corrected custom detail', 'Operator'),)
    saved = active.save('5' * 32, draft, expected_revision=0, base_snapshot_revision=0)
    loaded = active.load('5' * 32, draft.profile.contract, draft.profile.key, draft.period)
    assert loaded.draft == draft and loaded.revision == saved.revision
    word = assemble_docx(draft, acknowledged_fingerprint=draft.fingerprint)
    (tmp_path / 'custom-contact.docx').write_bytes(word)
    document = Document(BytesIO(word))
    rendered = next(t for t in document.tables if any('Corrected custom detail' in c.text for r in t.rows for c in r.cells))
    assert tuple(c.text for c in rendered.rows[0].cells) == table.columns
    assert 'Detail 2' not in '\n'.join(c.text for t in document.tables for r in t.rows for c in r.cells)


def test_old_report_action_cannot_close_or_change_a_replacement_report():
    app = workspace()
    button(app, 'Edit contact').click().run()
    old_switch = button(app, 'Edit contact')
    field(app, 'Name').set_value('Earlier report pending edit')
    replacement = ResolvedBlock('contact_matrix', 'This month', extra_tables=(SECOND,))
    app.session_state['block'] = replacement
    app.session_state['contact_review_rich_text_epoch'] = 'replacement-report'
    old_switch.click().run()
    assert not app.exception and app.session_state['block'] == replacement
    assert not app.text_input


@pytest.mark.parametrize('route', ['prefill', 'replay'])
def test_known_standing_fixture_uses_real_default_identity_and_keeps_report_only_control(monkeypatch, tmp_path, route):
    from app import monthly_report_directory as directory, monthly_report_library as library
    from app import monthly_report_qa_fixture as fixture, monthly_report_testing as qa
    from app.monthly_report_model import ReportDraft, profile_sections
    monkeypatch.setenv('EPC_DATA_DIR', str(tmp_path))
    with qa.isolated('6' * 32):
        draft = fixture.prepare(route)
        original = directory.load_directory(draft.profile.contract)
        ordinary = directory.apply_defaults(ReportDraft(draft.profile, draft.period, draft.prepared_by,
            profile_sections(draft.profile), ()))
        expected = next(block for block in ordinary.blocks if block.key == 'contact_matrix')
        contact = next(block for block in draft.blocks if block.key == 'contact_matrix')
        assert contact == expected and contact.source == 'Library'
        assert library.load_snapshot(draft.profile.contract, draft.profile.key, draft.period).draft == draft
        assert directory.load_directory(draft.profile.contract) == original
        corrected = replace(contact, source='This month', rows=(
            (*contact.rows[0][:3], '202-555-0142', contact.rows[0][4]), *contact.rows[1:]))
        # The old fixture's otherwise identical report-only state must still
        # have zero standing writes. Fixing fixture identity cannot weaken it.
        assert directory.automatically_keep_directory_contact_edits(draft.profile,
            replace(contact, source='This month'), corrected, draft.prepared_by) is None
        assert directory.load_directory(draft.profile.contract) == original
        shared = directory.automatically_keep_directory_contact_edits(draft.profile,
            contact, corrected, draft.prepared_by)
        latest = directory.load_directory(draft.profile.contract)
        assert shared is not None and shared.source == 'Library'
        assert latest.revision == original.revision + 1
        assert latest.sites[0].contacts[0].phone == '202-555-0142'
        assert latest.sites[1:] == original.sites[1:]
        assert latest.contract_contacts == original.contract_contacts

