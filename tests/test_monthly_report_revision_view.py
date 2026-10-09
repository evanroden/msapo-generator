"""Final revision view must agree in inspection, native output and PDF."""
from dataclasses import replace
from hashlib import sha256
from io import BytesIO
from zipfile import ZipFile

from docx import Document
from docx.shared import Inches
from lxml import etree as XML
import pymupdf
import pytest

from conftest import requires_libreoffice
from app.monthly_report_import import inspect_docx
from app.monthly_report_native_layout import build_native_docx
from app.monthly_report_native_package import passive_docx
from app.monthly_report_model import ResolvedBlock
from test_monthly_report_native_layout import draft, source

W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
OLD, AUTHOR = 'SYNTHETIC-REMOVED-CONTENT', 'Synthetic Prior Reviewer'


def run(parent, text, tag='t'):
    r = XML.SubElement(parent, W + 'r')
    XML.SubElement(XML.SubElement(r, W + 'rPr'), W + 'b')
    XML.SubElement(r, W + tag).text = text
    return r


def change(parent, kind, text):
    wrapper = XML.SubElement(parent, W + kind, {W + 'author': AUTHOR, W + 'id': '17'})
    run(wrapper, text, 'delText' if kind == 'del' else 't')


def annotate(p, label):
    change(p, 'ins', label + ' accepted insertion.')
    change(p, 'del', OLD + ' deleted $12,345.00')
    change(p, 'moveFrom', OLD + ' moved away.')
    change(p, 'moveTo', label + ' accepted move.')
    props = p.find(W + 'pPr')
    if props is None:
        props = XML.Element(W + 'pPr'); p.insert(0, props)
    old = XML.SubElement(props, W + 'pPrChange', {W + 'author': AUTHOR})
    XML.SubElement(XML.SubElement(old, W + 'pPr'), W + 'pStyle', {W + 'val': OLD})
    current = run(p, label + ' normal text.')
    old = XML.SubElement(current.find(W + 'rPr'), W + 'rPrChange', {W + 'author': AUTHOR})
    XML.SubElement(old, W + 'rPr')


def fixture(tmp_path, revisions=True, comments=False):
    path = source(tmp_path)
    doc = Document(path)
    p = doc.paragraphs[4]
    p.text = ''
    if revisions:
        annotate(p._p, 'Activity')
        annotate(doc.tables[0].cell(1, 0).paragraphs[0]._p, 'Table')
        annotate(doc.sections[0].header.paragraphs[0]._p, 'Header')
        annotate(doc.sections[0].footer.paragraphs[0]._p, 'Footer')
        XML.SubElement(doc.settings.element, W + 'trackRevisions')
    else:
        p.add_run('Activity normal text.').bold = True
    if comments:
        doc.add_comment(p.runs, text=OLD + ' comment', author=AUTHOR, initials='SP')
        header = doc.sections[0].header.paragraphs[0]
        XML.SubElement(header._p, W + 'commentRangeStart', {W + 'id': '0'})
        r = run(header._p, 'Current header beside comment')
        XML.SubElement(r, W + 'commentReference', {W + 'id': '0'})
        XML.SubElement(header._p, W + 'commentRangeEnd', {W + 'id': '0'})
        XML.SubElement(r, W + 'annotationRef')
    doc.save(path)
    return path


def assert_current(raw):
    banned = {'ins', 'del', 'delText', 'moveFrom', 'moveTo', 'pPrChange', 'rPrChange', 'trackRevisions',
              'commentReference', 'annotationRef', 'commentRangeStart', 'commentRangeEnd'}
    with ZipFile(BytesIO(raw)) as z:
        assert not any('comments' in n.casefold() for n in z.namelist())
        for name in z.namelist():
            if not name.endswith('.xml'):
                continue
            data = z.read(name)
            assert not any(n.tag == W + tag for n in XML.fromstring(data).iter() for tag in banned), name
            assert AUTHOR.encode() not in data and OLD.encode() not in data, name


def test_passive_accepts_current_runs_in_all_emitted_stories(tmp_path):
    path = fixture(tmp_path); original = path.read_bytes()
    raw = passive_docx(path)
    assert_current(raw)
    p = Document(BytesIO(raw)).paragraphs[4]
    assert 'Activity accepted insertion.' in p.text and 'Activity accepted move.' in p.text
    assert next(r for r in p.runs if r.text).bold
    assert p.paragraph_format.left_indent == Document(path).paragraphs[4].paragraph_format.left_indent
    assert path.read_bytes() == original


def test_inspection_current_view_preserves_original_hash_and_versions_ids(tmp_path):
    path = fixture(tmp_path); original = path.read_bytes()
    result = inspect_docx(path)
    assert OLD not in str([(i.text, i.rows) for i in result.items])
    assert 'Activity accepted move.' in str(result.items)
    assert result.sha256 == sha256(original).hexdigest()
    assert getattr(result, 'revision_view_version', 0) == 1
    assert all(i.id.startswith('current-v1-item-') for i in result.items)
    assert any('tracked' in n.lower() for n in result.notices)
    assert path.read_bytes() == original


@pytest.mark.parametrize('master', [False, True])
def test_multiline_replacements_cannot_clone_old_review_history(tmp_path, master):
    path = fixture(tmp_path); original = path.read_bytes()
    current = draft(ResolvedBlock('activity_summary', 'This month',
                                 text='Current line one.\nCurrent line two.\nCurrent line three.'))
    raw = build_native_docx(current, path, master=master)
    assert_current(raw)
    result = Document(BytesIO(raw))
    for text in ('Current line one.', 'Current line two.', 'Current line three.'):
        assert sum(text == p.text for p in result.paragraphs) == 1
    assert path.read_bytes() == original


def test_comment_cleanup_preserves_adjacent_text_and_legacy_item_ids(tmp_path):
    path = fixture(tmp_path, revisions=False, comments=True); original = path.read_bytes()
    raw = passive_docx(path)
    assert_current(raw)
    assert 'Current header beside comment' in Document(BytesIO(raw)).sections[0].header.paragraphs[0].text
    assert all(i.id.startswith('item-') for i in inspect_docx(path).items)
    assert path.read_bytes() == original


@requires_libreoffice
@pytest.mark.parametrize('master', [False, True])
def test_commented_source_final_export_and_preview_both_render(tmp_path, monkeypatch, master):
    from app import monthly_report_designs as designs
    from app.monthly_report_docx import generate_report
    from app.monthly_report_preview import preview_section
    from test_monthly_report_designs import master_reference
    monkeypatch.setenv('EPC_DATA_DIR', str(tmp_path / 'runtime'))
    path = master_reference(tmp_path)
    doc = Document(path)
    p = next(p for p in doc.paragraphs if p.text == 'Source-only content')
    doc.add_comment(p.runs, text=OLD, author=AUTHOR)
    doc.save(path)
    current = draft(ResolvedBlock('activity_summary', 'This month', text='Current line one.'))
    if master:
        designs.install_master(path, actor='Synthetic Editor', expected_revision=0)
        profile = designs.pin(current.profile)
    else:
        profile = designs.install(path, current.profile, actor='Synthetic Editor', expected_revision=0)
    current = replace(current, profile=profile)
    package = generate_report(current, acknowledged_fingerprint=current.fingerprint)
    assert package.pdf is not None, package.pdf_error
    assert package.pdf_error == ''
    assert_current(package.docx)
    preview = preview_section(current, 'activity')
    for raw in (package.pdf, preview.pdf):
        with pymupdf.open(stream=raw, filetype='pdf') as pdf:
            text = ''.join(p.get_text() for p in pdf)
            assert 'Current line one.' in text and OLD not in text


@requires_libreoffice
def test_accepted_revisions_do_not_print_as_colored_changes(tmp_path):
    from app.pdf_converter import convert_to_pdf
    path = fixture(tmp_path)
    current = draft(ResolvedBlock('activity_summary', 'This month', text='Current line one.\nCurrent line two.'))
    output = tmp_path / 'current.docx'; output.write_bytes(build_native_docx(current, path, master=True))
    rendered = convert_to_pdf(output)
    try:
        with pymupdf.open(rendered) as pdf:
            spans = [s for p in pdf for b in p.get_text('dict')['blocks'] if 'lines' in b
                     for line in b['lines'] for s in line['spans']]
            assert all(OLD not in s['text'] for s in spans)
            current_spans = [s for s in spans if 'Current line' in s['text']]
            assert len(current_spans) == 2 and all(s['color'] == 0 for s in current_spans)
    finally:
        rendered.unlink(missing_ok=True)


def test_deleted_table_row_removed_inserted_row_retained(tmp_path):
    path = fixture(tmp_path, revisions=False)
    doc = Document(path); table = doc.tables[0]
    row = table.add_row(); row.cells[0].text = OLD
    XML.SubElement(row._tr.get_or_add_trPr(), W + 'del', {W + 'author': AUTHOR})
    row = table.add_row(); row.cells[0].text = 'Inserted current row'
    XML.SubElement(row._tr.get_or_add_trPr(), W + 'ins', {W + 'author': AUTHOR})
    doc.save(path)
    raw = passive_docx(path)
    assert_current(raw)
    assert len(Document(BytesIO(raw)).tables[0].rows) == 3
    assert OLD not in str(inspect_docx(path).items)
    assert 'Inserted current row' in str(inspect_docx(path).items)


def test_deleted_only_paragraph_keeps_anchor_not_removed_text(tmp_path):
    doc = Document(); p = doc.add_paragraph(); change(p._p, 'del', OLD)
    doc.add_heading('MONTHLY ACTIVITY SUMMARY', 1); doc.add_paragraph('Current activity')
    path = tmp_path / 'deleted-only.docx'; doc.save(path)
    raw = passive_docx(path)
    assert_current(raw)
    assert len(Document(BytesIO(raw)).paragraphs) == 3
    assert next(i for i in inspect_docx(path).items if i.text == 'Current activity').position == 3


def test_unrevised_sources_retain_existing_inspection_ids(tmp_path):
    path = fixture(tmp_path, revisions=False)
    result = inspect_docx(path)
    assert getattr(result, 'revision_view_version', 0) == 0
    assert result.items[0].id == 'item-0001'
    assert_current(passive_docx(path))


@pytest.mark.parametrize('kind', ['footnote', 'endnote'])
def test_current_revision_view_in_reachable_notes(tmp_path, kind):
    from test_monthly_report_native_package import rewrite, add_relation
    path = fixture(tmp_path, revisions=False)
    def mutate(parts):
        body = XML.fromstring(parts['word/document.xml'])
        p = body.find(W + 'body/' + W + 'p')
        XML.SubElement(run(p, 'Current note anchor'), W + kind + 'Reference', {W + 'id': '1'})
        parts['word/document.xml'] = XML.tostring(body)
        notes = XML.Element(W + kind + 's', nsmap={'w': W[1:-1]})
        p = XML.SubElement(XML.SubElement(notes, W + kind, {W + 'id': '1'}), W + 'p')
        annotate(p, 'Current note')
        parts['word/' + kind + 's.xml'] = XML.tostring(notes)
        add_relation(parts, 'rNote', kind + 's', kind + 's.xml')
        types = XML.fromstring(parts['[Content_Types].xml'])
        XML.SubElement(types, '{http://schemas.openxmlformats.org/package/2006/content-types}Override',
                       PartName='/word/' + kind + 's.xml',
                       ContentType='application/vnd.openxmlformats-officedocument.wordprocessingml.' + kind + 's+xml')
        parts['[Content_Types].xml'] = XML.tostring(types)
    rewrite(path, mutate)
    raw = passive_docx(path)
    assert_current(raw)
    with ZipFile(BytesIO(raw)) as z:
        assert b'Current note accepted insertion.' in z.read('word/' + kind + 's.xml')
    result = inspect_docx(path)
    assert OLD not in str(result.items)
    assert 'Current note accepted move.' in str(result.items)


def test_revision_cleanup_in_both_textbox_compatibility_branches(tmp_path):
    from test_monthly_report_native_package import rewrite
    path = fixture(tmp_path, revisions=False)
    def mutate(parts):
        root = XML.fromstring(parts['word/document.xml'])
        p = root.find(W + 'body/' + W + 'p')
        mc = '{http://schemas.openxmlformats.org/markup-compatibility/2006}'
        alternate = XML.SubElement(p, mc + 'AlternateContent')
        for kind in ('Choice', 'Fallback'):
            branch = XML.SubElement(alternate, mc + kind)
            if kind == 'Choice': branch.set('Requires', 'w')
            box = XML.SubElement(branch, W + 'txbxContent')
            annotate(XML.SubElement(box, W + 'p'), kind)
        parts['word/document.xml'] = XML.tostring(root)
    rewrite(path, mutate)
    raw = passive_docx(path)
    assert_current(raw)
    with ZipFile(BytesIO(raw)) as z:
        data = z.read('word/document.xml')
        assert b'Choice accepted move.' in data and b'Fallback accepted move.' in data
    inspection = inspect_docx(path)
    assert OLD not in str(inspection.items)
    assert 'Fallback accepted move.' not in str(inspection.items)


def test_relationship_only_in_deleted_drawing_is_not_retained(tmp_path):
    from PIL import Image
    path = fixture(tmp_path, revisions=False)
    doc = Document(path)
    blob = BytesIO(); Image.new('RGB', (8, 8), 'red').save(blob, 'PNG')
    p = doc.add_paragraph(); picture = p.add_run(); picture.add_picture(BytesIO(blob.getvalue()))
    deleted = XML.SubElement(p._p, W + 'del', {W + 'author': AUTHOR})
    deleted.append(picture._r)
    doc.save(path)
    raw = passive_docx(path)
    assert_current(raw)
    with ZipFile(BytesIO(raw)) as z:
        assert not any(n.startswith('word/media/') for n in z.namelist())
        assert b'/image' not in z.read('word/_rels/document.xml.rels')
    assert not any(i.kind == 'image' for i in inspect_docx(path).items)


@pytest.mark.parametrize('structural', ['paragraph_mark', 'cell_delete', 'cell_merge', 'block_insert'])
def test_unsupported_structure_is_rejected_not_silently_misrendered(tmp_path, structural):
    from app.monthly_report_import import ImportError
    path = fixture(tmp_path, revisions=False)
    doc = Document(path)
    if structural == 'paragraph_mark':
        props = doc.paragraphs[4]._p.get_or_add_pPr()
        XML.SubElement(XML.SubElement(props, W + 'rPr'), W + 'del', {W + 'author': AUTHOR})
    elif structural.startswith('cell_'):
        props = doc.tables[0].cell(1, 0)._tc.get_or_add_tcPr()
        XML.SubElement(props, W + ('cellDel' if structural == 'cell_delete' else 'cellMerge'), {W + 'author': AUTHOR})
    else:
        p = doc.paragraphs[4]._p
        parent = p.getparent(); index = parent.index(p)
        wrapper = XML.Element(W + 'ins', {W + 'author': AUTHOR})
        parent.insert(index, wrapper); wrapper.append(p)
    doc.save(path); original = path.read_bytes()
    for operation in (inspect_docx, passive_docx):
        with pytest.raises(ImportError, match='structur'):
            operation(path)
    assert path.read_bytes() == original


def test_legacy_ordinal_mapping_cannot_bind_to_new_final_view(tmp_path):
    from app.monthly_report_import import ImportError, ImportMapping, map_items
    path = fixture(tmp_path)
    inspection = inspect_docx(path)
    with pytest.raises(ImportError, match='valid destination'):
        map_items(path, inspection, (ImportMapping('item-0005', 'activity_summary'),))


def test_page_counter_fields_survive_comments_and_revisions(tmp_path):
    path = fixture(tmp_path, revisions=True, comments=True)
    doc = Document(path)
    field = XML.SubElement(doc.sections[0].footer.paragraphs[0]._p, W + 'fldSimple', {W + 'instr': ' PAGE '})
    run(field, '1')
    doc.save(path)
    raw = passive_docx(path)
    assert_current(raw)
    with ZipFile(BytesIO(raw)) as z:
        footer = XML.fromstring(z.read('word/footer1.xml'))
        assert footer.find('.//' + W + 'fldSimple').get(W + 'instr') == ' PAGE '
