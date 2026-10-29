"""Selected current and historical Training rows in real plain/native Word/PDF."""
from hashlib import sha256
from io import BytesIO
import json
import os
from pathlib import Path

from docx import Document
import pymupdf
import pytest

from app import pdf_converter
from app.monthly_report_docx import assemble_docx
from app.monthly_report_native_layout import build_native_docx
from tests.conftest import requires_libreoffice
from test_training_review_flow import report, prepared, selected, publish, storage


@requires_libreoffice
@pytest.mark.parametrize("native", [False, True])
def test_reviewed_current_and_historical_rows_render_once_with_truthful_labels(tmp_path, monkeypatch, native):
    draft, review = selected(*prepared(report()))
    draft = publish(draft, review)
    source = Document()
    source.add_paragraph("Synthetic North")
    source.add_heading("12 Training Summary", 1)
    source.add_paragraph("Original synthetic placeholder")
    prototype = source.add_table(rows=2, cols=5)
    for cell, value in zip(prototype.rows[0].cells, ("Topic", "Date", "Method", "Participants", "Hours")):
        cell.text = value
    for cell, value in zip(prototype.rows[1].cells, ("Native placeholder", "", "", "", "")):
        cell.text = value
    path = tmp_path / "synthetic-training-design.docx"
    source.save(path)
    original = path.read_bytes()
    raw = build_native_docx(draft, path) if native else assemble_docx(draft, acknowledged_fingerprint=draft.fingerprint)
    document = Document(BytesIO(raw))
    tables = [tuple(tuple(c.text for c in row.cells) for row in table.rows) for table in document.tables]
    assert len(tables) == 2
    assert sum(row[0] == "Current drill" for table in tables for row in table) == 1
    assert sum(row[0] == "Prior drill" for table in tables for row in table) == 1
    assert any(table[0][0] == "Historical training" and table[1][1] == "2026-08-11" for table in tables)
    assert any(table[1][-1] == "" for table in tables if table[1][0] == "Current drill")
    word = tmp_path / "reviewed-training.docx"
    word.write_bytes(raw)
    monkeypatch.setattr(pdf_converter, "OUTPUT_DIR", tmp_path)
    pdf = pdf_converter.convert_to_pdf(word).read_bytes()
    with pymupdf.open(stream=pdf, filetype="pdf") as output:
        text = "\n".join(page.get_text() for page in output)
        assert text.count("Current drill") == text.count("Prior drill") == 1
        assert "Historical training" in text and "2026-08-11" in text
        assert "Incomplete" not in text
        folder = Path(os.environ.get("INDEPENDENT_OUTPUT", str(tmp_path))) / "training-reviewed-import" / ("native" if native else "plain")
        folder.mkdir(parents=True, exist_ok=True)
        for n, page in enumerate(output, 1):
            page.get_pixmap(dpi=100).save(folder / f"page-{n}.png")
    (folder / "report.docx").write_bytes(raw)
    (folder / "report.pdf").write_bytes(pdf)
    (folder / "manifest.json").write_text(json.dumps({"synthetic": True, "native": native,
        "docx_sha256": sha256(raw).hexdigest(), "pdf_sha256": sha256(pdf).hexdigest()}, indent=2))
    assert path.read_bytes() == original
