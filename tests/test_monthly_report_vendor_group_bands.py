"""Synthetic merged vendor headings must never become contact records."""
from dataclasses import replace
from io import BytesIO

from docx import Document

from app.monthly_report_import import ImportMapping, inspect_docx, imported_draft, map_items
from app.monthly_report_model import ReportPeriod, ReportTable, synthetic_profiles
from app.monthly_report_start_import import analyze_starting_report, automatic_section_plans, build_starting_import
from app.monthly_report_vendor_groups import grouped_vendor_contacts


PERIOD = ReportPeriod(2026, 9)
EXPECTED = (
    ("HVAC", "Synthetic Vendor One", "First Invented", "555-0101"),
    ("HVAC", "Synthetic Vendor Two", "Second Invented", "555-0102"),
    ("Electrical", "Synthetic Vendor Three", "Third Invented", "555-0103"),
)


def synthetic_vendor_source(path, *, first_band=False, ambiguous=False):
    profile = synthetic_profiles()[0]
    doc = Document()
    doc.add_heading(profile.facilities[0].title, 0)
    doc.add_paragraph("Monthly Report · September 2026")
    doc.add_heading("6 Sub-Contractor Status", 1)
    rows = (
        (("HVAC", "", ""), ("Vendor", "Contact", "Phone")) if first_band
        else (("Vendor", "Contact", "Phone"), ("HVAC", "", ""))
    )
    rows = (*rows, ("Synthetic Vendor One", "First Invented", "555-0101"),
            ("Synthetic Vendor Two", "Second Invented", "555-0102"),
            ("Electrical", "", ""), ("Synthetic Vendor Three", "Third Invented", "555-0103"))
    grid = doc.add_table(rows=len(rows), cols=3)
    for index, values in enumerate(rows):
        if index in ((0 if first_band else 1), 4):
            grid.rows[index].cells[0].merge(grid.rows[index].cells[-1]).text = values[0]
        else:
            for cell, value in zip(grid.rows[index].cells, values):
                cell.text = value
    if ambiguous:
        grid.rows[0].cells[0].text = "Opaque heading"
        grid.rows[0].cells[1].text = "Unverified"
    doc.save(path)
    return path


def find_vendor_item(inspection):
    return next(item for item in inspection.items if item.kind == "table"
                and any("Synthetic Vendor One" in cell for row in item.rows for cell in row))


def test_merge_geometry_and_three_contact_records_in_both_header_orders(tmp_path):
    for first in (False, True):
        path = synthetic_vendor_source(tmp_path / ("first-band.docx" if first else "first-header.docx"),
                                       first_band=first)
        original = path.read_bytes()
        item = find_vendor_item(inspect_docx(path))
        assert item.merged_title_rows == ((0, 4) if first else (1, 4))
        assert grouped_vendor_contacts(item) == ReportTable(
            ("Discipline", "Vendor", "Contact", "Phone"), EXPECTED)
        assert path.read_bytes() == original


def test_auto_import_and_generic_word_output_preserve_exact_contacts(tmp_path):
    path = synthetic_vendor_source(tmp_path / "report.docx")
    original = path.read_bytes()
    inspection = inspect_docx(path)
    profile = synthetic_profiles()[0]
    analysis = analyze_starting_report(inspection, profile, PERIOD)
    assert analysis.site_relation == "same"
    plans = automatic_section_plans(inspection, analysis)
    assert all(plan["approved"] for plan in plans.values())
    mapped, _ = build_starting_import(path, inspection, plans.values())
    imported = next(b for b in mapped.blocks if b.key == "subcontractor_matrix").extra_tables[0]
    assert imported.rows == EXPECTED and imported.reference.startswith("docx:")
    from app.monthly_report_docx import _build_docx
    draft = imported_draft(profile, PERIOD, "Synthetic Editor", mapped)
    word = Document(BytesIO(_build_docx(draft)))
    tables = [table for table in word.tables
              if any("Synthetic Vendor One" in cell.text
                     for row in table.rows for cell in row.cells)]
    assert len(tables) == 1
    rows = [tuple(cell.text for cell in row.cells) for row in tables[0].rows]
    assert rows == [("Discipline", "Vendor", "Contact", "Phone"), *EXPECTED]
    assert path.read_bytes() == original


def test_manual_import_and_explicit_column_selection(tmp_path):
    path = synthetic_vendor_source(tmp_path / "manual.docx", first_band=True)
    item = find_vendor_item(inspect_docx(path))
    inspection = inspect_docx(path)
    mapped = map_items(path, inspection, (ImportMapping(item.id, "subcontractor_matrix"),))
    assert mapped.blocks[0].rows == EXPECTED
    assert tuple(c.title for c in mapped.overrides[0].columns) == (
        "Discipline", "Vendor", "Contact", "Phone")
    selected = map_items(path, inspection, (
        ImportMapping(item.id, "subcontractor_matrix", columns=(0, 2)),))
    assert len(selected.blocks[0].rows) == 5
    assert len(selected.blocks[0].rows[0]) == 2


def test_ambiguous_or_unmerged_vendor_rows_do_not_become_categories(tmp_path):
    item = find_vendor_item(inspect_docx(synthetic_vendor_source(
        tmp_path / "ambiguous.docx", ambiguous=True)))
    assert grouped_vendor_contacts(item) is None
    assert grouped_vendor_contacts(replace(item, merged_title_rows=())) is None


def test_grouped_contact_cards_add_and_move_without_ghost_rows():
    from tests.test_monthly_report_compact_contacts import workspace, block
    table = ReportTable(("Discipline", "Vendor", "Contact", "Phone"), EXPECTED, "synthetic:vendors")
    app = workspace((table,))
    assert not app.exception
    assert any("Trade: HVAC" in str(label.value) for label in app.text)
    next(button for button in app.button
         if button.label == "Add vendor contact to Electrical").click().run()
    assert not app.exception
    assert next(field for field in app.text_input if field.label == "Discipline").value == "Electrical"
    next(button for button in app.button if button.label == "Close editing").click().run()
    assert block(app).extra_tables[0].rows == EXPECTED
    next(button for button in app.button
         if button.label == "Add vendor contact to Electrical").click().run()
    next(field for field in app.text_input
         if field.label == "Vendor").set_value("Synthetic Vendor Four").run()
    next(button for button in app.button if button.label == "Close editing").click().run()
    assert block(app).extra_tables[0].rows[-1] == ("Electrical", "Synthetic Vendor Four", "", "")
    next(button for button in app.button
         if button.label == "Edit vendor contact" and button.key.endswith("_open_0_0")).click().run()
    next(field for field in app.text_input if field.label == "Discipline").set_value("Electrical").run()
    next(field for field in app.text_input if field.label == "Vendor").set_value("Renamed Invented Vendor").run()
    next(button for button in app.button if button.label == "Close editing").click().run()
    result = block(app).extra_tables[0].rows
    assert len(result) == 4 and result[0][:2] == ("Electrical", "Renamed Invented Vendor")
    assert result[1:3] == EXPECTED[1:]


def manual_vendor_ui(monkeypatch, tmp_path, *, first_band=False, adjacent=None):
    from test_monthly_report_editor import app_with_library
    from app import monthly_report_import_ui as ui
    from streamlit.testing.v1 import AppTest
    _, profile = app_with_library(monkeypatch, tmp_path)
    if adjacent is not None:
        from app import monthly_report_library as library
        library.save_import(profile.contract, profile.key, (adjacent,),
                            expected_revision=1, actor="Synthetic Editor", confirmed=True)
    path = synthetic_vendor_source(tmp_path / 'manual-ui.docx', first_band=first_band)
    original_bytes = path.read_bytes()
    upload = BytesIO(original_bytes)
    upload.size, upload.name = len(original_bytes), 'manual-ui.docx'
    original_uploader = ui.st.file_uploader
    monkeypatch.setattr(ui.st, 'file_uploader', lambda label, *args, **kwargs:
        upload if label == 'Existing monthly report DOCX' else original_uploader(label, *args, **kwargs))
    app = AppTest.from_string('''
from app import monthly_report_library as library
from app.monthly_report_import_ui import render_import
from app.monthly_report_ui import _field
from app.contracts import RRH_CONTRACT
render_import(library.load_profile(RRH_CONTRACT, 'synthetic'), _field)
''', default_timeout=20).run()
    next(b for b in app.button if b.label == 'Read DOCX for mapping').click().run()
    item = find_vendor_item(inspect_docx(path))
    next(w for w in app.selectbox if w.label == 'Item to inspect').set_value(item.id).run()
    next(w for w in app.selectbox if w.label == 'Confirmed destination').set_value('subcontractor_matrix').run()
    assert not app.exception
    return app, path, item, original_bytes


def confirmed_manual_mapping(app, item):
    next(b for b in app.button if b.label == 'Add confirmed mapping').click().run()
    assert not app.exception
    mappings = next(value for key, value in app.session_state.filtered_state.items()
                    if key.endswith('_mappings'))
    return mappings[item.id]


def test_manual_ui_untouched_complete_vendor_table_groups_contacts(monkeypatch, tmp_path):
    app, path, item, original = manual_vendor_ui(monkeypatch, tmp_path, first_band=True)
    assert next(w for w in app.multiselect if w.label == 'Columns to keep in this order').value == [0, 1, 2]
    app.run()  # An unrelated rerun must not turn defaults into customization.
    mapping = confirmed_manual_mapping(app, item)
    mapped = map_items(path, inspect_docx(path), (mapping,))
    assert mapped.blocks[0].rows == EXPECTED
    assert tuple(c.title for c in mapped.overrides[0].columns) == ('Discipline', 'Vendor', 'Contact', 'Phone')
    assert path.read_bytes() == original


def test_manual_ui_customized_columns_remain_authoritative_after_restore(monkeypatch, tmp_path):
    app, path, item, original = manual_vendor_ui(monkeypatch, tmp_path)
    columns = next(w for w in app.multiselect if w.label == 'Columns to keep in this order')
    columns.set_value([2, 0]).run()
    mapping = confirmed_manual_mapping(app, item)
    mapped = map_items(path, inspect_docx(path), (mapping,))
    assert mapping.columns == (2, 0)
    assert tuple(c.title for c in mapped.overrides[0].columns) == ('Phone', 'Vendor')
    assert mapped.blocks[0].rows == tuple((row[2], row[0]) for row in item.rows[1:])
    next(w for w in app.multiselect if w.label == 'Columns to keep in this order').set_value([0, 1, 2]).run()
    app.run()
    mapping = confirmed_manual_mapping(app, item)
    assert mapping.columns == (0, 1, 2)
    assert map_items(path, inspect_docx(path), (mapping,)).blocks[0].rows == item.rows[1:]
    assert path.read_bytes() == original


def test_manual_ui_grouped_save_reopens_without_changing_adjacent_edits(monkeypatch, tmp_path):
    from app import monthly_report_library as library
    from app.contracts import RRH_CONTRACT
    from app.monthly_report_model import ResolvedBlock
    adjacent = ResolvedBlock("training_summary", "Library", text="User-edited synthetic training notes")
    app, path, item, original = manual_vendor_ui(monkeypatch, tmp_path, adjacent=adjacent)
    confirmed_manual_mapping(app, item)
    next(w for w in app.text_input if w.label == "Import editor name").set_value("Synthetic Importer").run()
    next(w for w in app.checkbox if w.label == "I confirm these facility identities, aliases and scope").check().run()
    next(w for w in app.checkbox if w.label == "I confirm these mappings and this shared save").check().run()
    save = next(b for b in app.button if b.label == "Save confirmed DOCX mappings")
    assert not save.disabled
    save.click().run()
    assert not app.exception
    state = library.load_profile(RRH_CONTRACT, "synthetic")
    assert state.revision == 3
    assert state.block("subcontractor_matrix").rows == EXPECTED
    assert state.block("training_summary") == adjacent
    assert path.read_bytes() == original
    app.run()
    assert not app.exception
    reopened = library.load_profile(RRH_CONTRACT, "synthetic")
    assert reopened.block("subcontractor_matrix").rows == EXPECTED
    assert reopened.block("training_summary") == adjacent


def test_manual_ui_no_header_and_empty_selection_are_not_grouped(monkeypatch, tmp_path):
    app, path, item, _ = manual_vendor_ui(monkeypatch, tmp_path)
    next(w for w in app.checkbox if w.label == "First row contains column headings").uncheck().run()
    mapping = confirmed_manual_mapping(app, item)
    assert mapping.columns == (0, 1, 2) and not mapping.first_row_header
    assert map_items(path, inspect_docx(path), (mapping,)).blocks[0].rows == item.rows
    next(w for w in app.multiselect if w.label == "Columns to keep in this order").set_value([]).run()
    assert next(b for b in app.button if b.label == "Add confirmed mapping").disabled
