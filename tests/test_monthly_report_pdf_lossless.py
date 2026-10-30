"""Final Monthly PDF images keep prepared input pixels; other callers opt out."""
from collections import Counter
from contextlib import contextmanager
from hashlib import sha256
from io import BytesIO
import json
from pathlib import Path
from types import SimpleNamespace

from docx import Document
from docx.shared import Inches
from PIL import Image, ImageChops, ImageDraw
import pymupdf as fitz
import pytest

from app import pdf_converter as converter, monthly_report_render_jobs as jobs
from app import monthly_report_docx as monthly
from app.monthly_report_model import synthetic_draft, synthetic_profiles, ReportPeriod
from conftest import requires_libreoffice


def identity(image):
    return image.size, sha256(image.convert('RGB').tobytes()).hexdigest()


@pytest.mark.parametrize('lossless', [False, True])
def test_writer_option_is_explicit_bounded_and_default_compatible(monkeypatch, tmp_path, lossless):
    from app import office_process
    monkeypatch.setenv('EPC_DATA_DIR', str(tmp_path/'runtime'))
    monkeypatch.setattr(converter, 'OUTPUT_DIR', tmp_path)
    monkeypatch.setattr(converter, 'PDF_BACKEND', 'libreoffice')
    monkeypatch.setattr(converter.shutil, 'which', lambda _: '/synthetic/soffice')
    path=tmp_path/'source.docx';path.write_bytes(b'unchanged Word input');before=path.read_bytes();calls=[]
    def run(command, *, timeout):
        calls.append((command,timeout))
        out=Path(command[command.index('--outdir')+1]);(out/'source.pdf').write_bytes(b'%PDF-synthetic')
        return SimpleNamespace(returncode=0,stdout='',stderr='')
    monkeypatch.setattr(office_process,'run_office',run)
    output=converter.convert_to_pdf(path,lossless_images=True) if lossless else converter.convert_to_pdf(path)
    assert output.read_bytes()==b'%PDF-synthetic' and path.read_bytes()==before
    command,timeout=calls[0];assert len(calls)==1 and timeout==120
    assert command[-1]==str(path.resolve())
    option=command[command.index('--convert-to')+1]
    if lossless:
        prefix,settings=option.split(':',2)[0:2],option.split(':',2)[2]
        assert prefix==['pdf','writer_pdf_Export']
        assert json.loads(settings)=={'UseLosslessCompression':{'type':'boolean','value':'true'},
                                      'ReduceImageResolution':{'type':'boolean','value':'false'}}
    else:assert option=='pdf:writer_pdf_Export'
    output.unlink();assert not list(tmp_path.glob('.office-*'))


@pytest.mark.parametrize('backend',['libreoffice','gotenberg','docx2pdf'])
def test_default_dispatch_preserves_existing_one_argument_backend(monkeypatch,backend):
    seen=[];monkeypatch.setattr(converter,'PDF_BACKEND',backend)
    monkeypatch.setitem(converter._BACKENDS,backend,lambda path:seen.append(path) or path)
    source=Path('source.docx');assert converter.convert_to_pdf(source)==source and seen==[source]


@pytest.mark.parametrize('backend',['gotenberg','docx2pdf'])
def test_lossless_request_cannot_silently_use_an_uncertified_backend(monkeypatch,backend):
    monkeypatch.setattr(converter,'PDF_BACKEND',backend)
    monkeypatch.setitem(converter._BACKENDS,backend,lambda *a,**kw:pytest.fail('Uncertified backend called'))
    with pytest.raises(converter.PDFConversionError,match='lossless'):
        converter.convert_to_pdf(Path('source.docx'),lossless_images=True)


@pytest.mark.parametrize('value',['false','true',1,None])
def test_invalid_lossless_option_is_not_truthy_dispatch(monkeypatch,value):
    monkeypatch.setattr(converter,'PDF_BACKEND','libreoffice')
    monkeypatch.setitem(converter._BACKENDS,'libreoffice',lambda *a,**kw:pytest.fail('Invalid flag reached backend'))
    with pytest.raises(converter.PDFConversionError,match='boolean'):
        converter.convert_to_pdf(Path('source.docx'),lossless_images=value)


@pytest.mark.parametrize('lossless',[False,True])
def test_existing_serialized_job_forwards_only_explicit_opt_in(monkeypatch,lossless):
    events=[]
    @contextmanager
    def slot(**kwargs):
        events.append(('slot',kwargs));yield;events.append('released')
    monkeypatch.setattr(jobs,'conversion_slot',slot)
    monkeypatch.setattr(converter,'convert_to_pdf',lambda path,**kwargs:events.append(('convert',path,kwargs)) or path)
    path=Path('source.docx')
    result=jobs.convert_to_pdf(path,wait_seconds=130,lossless_images=True) if lossless else jobs.convert_to_pdf(path,wait_seconds=130)
    assert result==path
    assert events==[('slot',{'wait_seconds':130}),('convert',path,{'lossless_images':True} if lossless else {}),'released']


def test_only_final_monthly_path_opts_in_and_retains_input_and_failure_contract(monkeypatch,tmp_path):
    from app import monthly_report_designs as designs
    monkeypatch.setenv('EPC_DATA_DIR',str(tmp_path/'runtime'))
    monkeypatch.setattr(converter,'OUTPUT_DIR',tmp_path)
    monkeypatch.setattr(monthly,'rendering_docx',lambda raw,**kw:raw)
    monkeypatch.setattr(designs,'source_for',lambda _:None)
    draft=synthetic_draft(synthetic_profiles()[0],ReportPeriod(2026,9))
    raw=b'unchanged native Word';seen=[]
    def convert(path,**kwargs):
        seen.append((path.read_bytes(),kwargs));raise converter.PDFConversionTimeout('Synthetic timeout')
    monkeypatch.setattr(jobs,'convert_to_pdf',convert)
    result,error=monthly._convert_report_pdf(draft,raw)
    assert result is None and 'too long' in error and 'DOCX is ready' in error
    assert seen==[(raw,{'wait_seconds':130,'lossless_images':True})]


def sample_document(tmp_path):
    # Fine RGB diagram, text and colored edges expose JPEG damage; distinct
    # source dimensions and two uses expose resampling and lost occurrences.
    # Deterministic textured RGB background prevents a newer LibreOffice from
    # choosing a tiny PNG opportunistically even with lossy export enabled.
    # Hosted 24.2 also JPEG-encoded the much simpler original chart/page inputs.
    import random
    image=Image.frombytes('RGB',(320,180),random.Random(125).randbytes(320*180*3));draw=ImageDraw.Draw(image)
    draw.rectangle((3,3,305,150),fill='white')
    draw.rectangle((10,20,100,140),fill=(0,0,255));draw.rectangle((120,20,290,145),outline=(255,0,255),width=1)
    draw.text((112,70),'Verified savings: $125',fill=(0,0,0))
    for x in range(110,300,4):draw.line((x,110,x,140),fill=(0,180,40),width=1)
    # Like the immutable prepared source pages, keep a real white border: the
    # native Writer image frame rounds its clip to a tenth of a point.
    draw.rectangle((0,0,319,179),outline='white',width=4)
    raw=BytesIO();image.save(raw,'PNG');png=raw.getvalue()
    doc=Document();doc.add_paragraph('Immutable synthetic report image.');doc.add_picture(BytesIO(png),width=Inches(2))
    doc.add_paragraph('Independent second occurrence.');doc.add_picture(BytesIO(png),width=Inches(2))
    source=tmp_path/'sample.docx';doc.save(source)
    return source,image


def inspect_pdf(raw,expected):
    images=[];occurrences=[];pages=[];paint=[]
    with fitz.open(stream=raw,filetype='pdf') as pdf:
        for xref in range(1,pdf.xref_length()):
            if pdf.xref_get_key(xref,'Subtype')[1]=='/Image':
                resource=pdf.extract_image(xref);image=Image.open(BytesIO(resource['image']))
                images.append({'xref':xref,'identity':identity(image),'filter':pdf.xref_get_key(xref,'Filter')[1],
                               'smask':resource['smask'],'max_rgb_delta':ImageChops.difference(image.convert('RGB'),expected).getextrema() if image.size==expected.size else None})
        for page in pdf:
            pages.append((tuple(page.rect),page.get_text()))
            for info in page.get_image_info(xrefs=True):
                bounds=fitz.Rect(info['bbox']);assert page.rect.contains(bounds) and bounds.width>1 and bounds.height>1
                occurrences.append((identity(Image.open(BytesIO(pdf.extract_image(info['xref'])['image']))),tuple(bounds)))
                data=BytesIO();expected.save(data,'PNG')
                with fitz.open() as clean:
                    reference=clean.new_page(width=page.rect.width,height=page.rect.height)
                    reference.insert_image(bounds,stream=data.getvalue(),keep_proportion=False)
                    a=page.get_pixmap(clip=bounds,dpi=90,alpha=False,colorspace=fitz.csRGB)
                    b=reference.get_pixmap(clip=bounds,dpi=90,alpha=False,colorspace=fitz.csRGB)
                    assert (a.width,a.height)==(b.width,b.height)
                    paint.append(ImageChops.difference(Image.frombytes('RGB',(a.width,a.height),a.samples),
                                                       Image.frombytes('RGB',(b.width,b.height),b.samples)).getextrema())
    return {'images':images,'occurrences':occurrences,'pages':pages,'paint':paint,'bytes':len(raw)}


@requires_libreoffice
def test_real_final_monthly_converter_preserves_png_pixels_two_occurrences_and_geometry(monkeypatch,tmp_path,record_property):
    from app import monthly_report_designs as designs
    monkeypatch.setenv('EPC_DATA_DIR',str(tmp_path/'runtime'));monkeypatch.setattr(converter,'OUTPUT_DIR',tmp_path)
    monkeypatch.setattr(converter,'PDF_BACKEND','libreoffice');monkeypatch.setattr(designs,'source_for',lambda _:None)
    source,image=sample_document(tmp_path);original=source.read_bytes();draft=synthetic_draft(synthetic_profiles()[0],ReportPeriod(2026,9))
    # Both conversions see the identical disposable rendering copy. Only the
    # final Monthly call requests lossless export; no arbitrary filter override.
    prepared=monthly.rendering_docx(original,profile=designs.render_profile_for(draft.profile))
    from app.monthly_report_media_privacy import sanitize_package_transparency
    prepared=sanitize_package_transparency(prepared)
    control=tmp_path/'default.docx';control.write_bytes(prepared)
    default_path=converter.convert_to_pdf(control);default=inspect_pdf(default_path.read_bytes(),image);default_path.unlink()
    raw,error=monthly._convert_report_pdf(draft,original);assert not error and raw
    (tmp_path/'final.pdf').write_bytes(raw)
    final=inspect_pdf(raw,image);record_property('converter_evidence',json.dumps({'default':default,'final':final}))
    assert source.read_bytes()==original and control.read_bytes()==prepared
    assert all(row['identity']==identity(image) and row['smask']==0 for row in final['images'])
    assert Counter(value for value,bounds in final['occurrences'])==Counter({identity(image):2})
    assert final['pages']==default['pages']
    assert [bounds for value,bounds in final['occurrences']]==[bounds for value,bounds in default['occurrences']]
    assert all(high<=3 for extrema in final['paint'] for low,high in extrema)
    assert len(raw)<1024*1024,'Small lossless line-art fixture must remain bounded'


@requires_libreoffice
def test_real_final_monthly_transparency_keeps_rgb_and_its_exact_soft_mask(monkeypatch,tmp_path,record_property):
    from app import monthly_report_designs as designs
    monkeypatch.setenv('EPC_DATA_DIR',str(tmp_path/'runtime'));monkeypatch.setattr(converter,'OUTPUT_DIR',tmp_path)
    monkeypatch.setattr(converter,'PDF_BACKEND','libreoffice');monkeypatch.setattr(designs,'source_for',lambda _:None)
    image=Image.new('RGBA',(160,90),(0,0,0,0));draw=ImageDraw.Draw(image)
    draw.rectangle((0,0,159,89),outline=(255,255,255,255),width=4)
    # Endpoint RGB channels remain exactly representable through native alpha
    # premultiplication, so this isolates image encoding and mask preservation.
    draw.rectangle((10,10,75,70),fill=(255,0,255,128));draw.rectangle((85,20,145,70),fill=(0,40,255,255))
    data=BytesIO();image.save(data,'PNG');original_image=data.getvalue()
    document=Document();document.add_picture(BytesIO(original_image),width=Inches(2));word=BytesIO();document.save(word);original=word.getvalue()
    draft=synthetic_draft(synthetic_profiles()[0],ReportPeriod(2026,9))
    raw,error=monthly._convert_report_pdf(draft,original);assert not error and raw
    with fitz.open(stream=raw,filetype='pdf') as pdf:
        infos=[(page,info) for page in pdf for info in page.get_image_info(xrefs=True)]
        assert len(infos)==1
        page,info=infos[0];resource=pdf.extract_image(info['xref']);assert resource['smask']>0 and info['has-mask']
        rgb=Image.open(BytesIO(resource['image'])).convert('RGB')
        alpha=Image.open(BytesIO(pdf.extract_image(resource['smask'])['image'])).convert('L')
        record_property('alpha_before_assertions',json.dumps({'rgb_max_delta':ImageChops.difference(rgb,image.convert('RGB')).getextrema(),
            'mask_max_delta':ImageChops.difference(alpha,image.getchannel('A')).getextrema(),
            'pixels':[(point,rgb.getpixel(point),image.getpixel(point),alpha.getpixel(point)) for point in [(0,0),(5,5),(20,20),(90,40)]]}))
        assert identity(rgb)==identity(image)
        assert alpha.size==image.size and alpha.tobytes()==image.getchannel('A').tobytes()
        image_refs={xref for xref in range(1,pdf.xref_length()) if pdf.xref_get_key(xref,'Subtype')[1]=='/Image'}
        assert image_refs=={info['xref'],resource['smask']},'No unused or unowned image/mask'
        bounds=fitz.Rect(info['bbox']);assert page.rect.contains(bounds)
        with fitz.open() as reference:
            clean=reference.new_page(width=page.rect.width,height=page.rect.height)
            clean.insert_image(bounds,stream=original_image,keep_proportion=False)
            a=page.get_pixmap(clip=bounds,dpi=90,alpha=False,colorspace=fitz.csRGB)
            b=clean.get_pixmap(clip=bounds,dpi=90,alpha=False,colorspace=fitz.csRGB)
            assert (a.width,a.height)==(b.width,b.height)
            delta=ImageChops.difference(Image.frombytes('RGB',(a.width,a.height),a.samples),Image.frombytes('RGB',(b.width,b.height),b.samples)).getextrema()
            assert all(high<=3 for low,high in delta)
        record_property('alpha_evidence',json.dumps({'bytes':len(raw),'dimensions':rgb.size,'soft_mask':resource['smask'],'paint_max_delta':delta}))
    assert word.getvalue()==original and data.getvalue()==original_image


@requires_libreoffice
def test_real_opt_in_does_not_recompress_prepared_jpeg_pixels(monkeypatch,tmp_path,record_property):
    monkeypatch.setenv('EPC_DATA_DIR',str(tmp_path/'runtime'));monkeypatch.setattr(converter,'OUTPUT_DIR',tmp_path)
    monkeypatch.setattr(converter,'PDF_BACKEND','libreoffice')
    _,image=sample_document(tmp_path);stream=BytesIO();image.save(stream,'JPEG',quality=88)
    prepared=stream.getvalue();expected=Image.open(BytesIO(prepared)).convert('RGB')
    doc=Document();doc.add_picture(BytesIO(prepared),width=Inches(2));path=tmp_path/'prepared-jpeg.docx';doc.save(path);original=path.read_bytes()
    output=converter.convert_to_pdf(path,lossless_images=True)
    with fitz.open(output) as pdf:
        infos=[info for page in pdf for info in page.get_image_info(xrefs=True)];assert len(infos)==1
        resource=pdf.extract_image(infos[0]['xref']);actual=Image.open(BytesIO(resource['image'])).convert('RGB')
        assert identity(actual)==identity(expected)
        record_property('jpeg_evidence',json.dumps({'prepared_jpeg_bytes':len(prepared),'pdf_bytes':output.stat().st_size,'filter':pdf.xref_get_key(infos[0]['xref'],'Filter')[1],'dimensions':actual.size,'rgb_sha256':identity(actual)[1]}))
    assert path.read_bytes()==original;output.unlink()
