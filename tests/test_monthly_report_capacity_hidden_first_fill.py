"""Only visible, reviewed thermal workbook facts may become standing first-fill.

All sites, amounts and workbook content here are synthetic. Hidden sheets may
hold stale, copied or other-client values that must not enter a new report.
"""
from io import BytesIO

import pytest
from openpyxl import Workbook

from app import monthly_report_capacity as capacity
from app import monthly_report_directory as directory
from app import monthly_report_library as library
from app.monthly_report_model import Facility, ReportProfile


@pytest.fixture
def profile(tmp_path, monkeypatch):
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path))
    facilities = (
        Facility("north", "Synthetic North", ("SN",)),
        Facility("south", "Synthetic South", ("SS",)),
    )
    selected = ReportProfile("Synthetic Thermal Privacy Contract", "north",
                             "Synthetic North", facilities[:1])
    directory.save_directory(
        selected.contract,
        tuple(directory.DirectorySite(f.key, f.title, f.aliases) for f in facilities),
        expected_revision=0, actor="Synthetic Reviewer", confirmed=True)
    return selected


def workbook(sheets):
    document = Workbook()
    document.remove(document.active)
    for name, state, rows in sheets:
        sheet = document.create_sheet(name)
        for row in rows:
            sheet.append(list(row))
        sheet.sheet_state = state
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def thermal_rows(required, service="Chilled water"):
    return (
        ("Facility", "Service", "Required capacity", "Available capacity", "Units"),
        ("SN", service, str(required), "NR", "tons" if service == "Chilled water" else "lb/hr"),
    )


@pytest.mark.parametrize("hidden_state", ("hidden", "veryHidden"))
def test_hidden_capacity_sheet_not_auto_imported_or_printed(profile, hidden_state):
    visible = thermal_rows(1200)
    hidden = thermal_rows(98765)
    raw = workbook((
        ("Current visible reviewed", "visible", visible),
        ("Previous client secret capacities", hidden_state, hidden),
    ))
    result = capacity.inspect_capacity(profile, "synthetic-thermal.xlsx", raw)
    assert len(result.tables) == 1
    assert result.tables[0].title == "Current visible reviewed"
    assert result.tables[0].rows == (("SN", "Chilled water", "1200", "NR", "tons"),)
    assert result.tables[0].site_keys == ("north",)
    assert all("98765" not in text for text in result.content.source.page_texts)
    assert all("Previous client" not in text for text in result.content.source.page_texts)
    assert any("Hidden thermal worksheets were excluded" in warning
               for warning in result.notices)
    stored = capacity.save_capacity(profile, result.tables, expected_revision=0,
                                    actor="Synthetic Reviewer",
                                    reviewed_fingerprint=capacity.review_fingerprint(result.tables))
    assert stored.tables[0].rows == result.tables[0].rows
    assert all("98765" not in cell for tab in stored.tables for row in tab.rows for cell in row)


def test_only_hidden_capacity_data_never_becomes_a_first_fill(profile):
    raw = workbook((
        ("Instructions", "visible", (("Report for thermal information",),)),
        ("Hidden old client", "hidden", thermal_rows(99999)),
    ))
    result = capacity.inspect_capacity(profile, "synthetic-thermal.xlsx", raw)
    assert result.tables == ()
    assert capacity.first_fill_tables(None, result.tables) == ()
    assert capacity.load_capacity(profile.contract) is None
    assert any("Hidden thermal worksheets were excluded" in x for x in result.notices)


def test_visible_new_service_is_added_without_reviving_hidden_prior_values(profile):
    raw = workbook((
        ("Cooling", "visible", thermal_rows(1200)),
        ("Old hidden steam", "hidden", thermal_rows(88888, "Steam")),
    ))
    initial = capacity.inspect_capacity(profile, "thermal.xlsx", raw)
    saved = capacity.save_capacity(profile, initial.tables, expected_revision=0,
                                   actor="Synthetic Reviewer",
                                   reviewed_fingerprint=capacity.review_fingerprint(initial.tables))
    assert len(saved.tables) == 1
    later = capacity.inspect_capacity(profile, "new-steam.csv",
        b"Facility,Service,Required capacity,Available capacity,Units\\n"
        b"SN,Steam,40000,NR,lb/hr\\n")
    first = capacity.first_fill_tables(saved, later.tables)
    assert len(first) == 1
    assert first[0].rows == (("SN", "Steam", "40000", "NR", "lb/hr"),)
    assert capacity.load_capacity(profile.contract).revision == 1
    second = capacity.save_capacity(profile, first, expected_revision=1,
                                    actor="Synthetic Reviewer",
                                    reviewed_fingerprint=capacity.review_fingerprint(first))
    assert second.revision == 2
    values = {row[1]: row[2] for table in second.tables for row in table.rows}
    assert values == {"Chilled water": "1200", "Steam": "40000"}
    assert capacity.load_capacity(profile.contract, 1) == saved


def test_two_sources_at_same_revision_cannot_silently_overwrite(profile):
    def inspect(value):
        return capacity.inspect_capacity(profile, f"thermal-{value}.csv",
            ("Facility,Service,Required capacity,Available capacity,Units\\n"
             f"SN,Steam,{value},NR,lb/hr\\n").encode())
    first = inspect(40000)
    second = inspect(50000)
    winner = capacity.save_capacity(profile, first.tables, expected_revision=0,
                                    actor="Synthetic Reviewer",
                                    reviewed_fingerprint=capacity.review_fingerprint(first.tables))
    with pytest.raises(library.RevisionConflict, match="Someone updated"):
        capacity.save_capacity(profile, second.tables, expected_revision=0,
                               actor="Second Synthetic Reviewer",
                               reviewed_fingerprint=capacity.review_fingerprint(second.tables))
    assert capacity.load_capacity(profile.contract) == winner
    assert capacity.load_capacity(profile.contract, 1) == winner
