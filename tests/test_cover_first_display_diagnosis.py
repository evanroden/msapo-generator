"""Server-side first-display controls, not native browser receipt acceptance.

AppTest supplies uploader bytes and component viewport/poll receipts at their
supported seams. The actual source gallery, pending cover payload, scheduler,
generation checks, cache and raster transport run; conversion uses a synthetic
PDF. No business click is used after the automatic source-photo suggestion.
"""

from concurrent.futures import Future
from io import BytesIO
from threading import BoundedSemaphore

from docx import Document
from docx.shared import Inches
import fitz
import pytest

from app import monthly_report_section_preview_ui as ui, monthly_report_section_ui as section_ui
from app.monthly_report_import import read_import_image
from app.monthly_report_preview import ReportPreview, section_preview_fingerprint
from report_preview_test_support import payload
from test_cover_branding_contracts import photograph
from test_cover_import_task_ui import current
from test_monthly_report_ui import monthly


@pytest.mark.parametrize("older_preview", [False, True])
def test_first_automatic_source_photo_reaches_ready_without_business_click(monkeypatch, tmp_path, older_preview):
    actual_render = ui.render_section_preview
    app = monthly(monkeypatch, tmp_path / "store")
    next(w for w in app.text_input if w.label == "Prepared by").set_value("Synthetic Reviewer").run()
    prefix = next(key.removesuffix("_draft") for key, value in app.session_state.filtered_state.items()
                  if key.startswith("report_guided_") and key.endswith("_draft") and value == current(app))
    clock, jobs, rendered = [100.0], [], []
    monkeypatch.setattr(ui, "monotonic", lambda: clock[0])
    monkeypatch.setattr(ui, "_PREVIEW_SLOT", BoundedSemaphore(1))

    class Worker:
        def submit(self, function, *args):
            future = Future()
            jobs.append((future, function, args))
            return future

    monkeypatch.setattr(ui, "_PREVIEW_WORKER", Worker())

    def convert(draft, section, loader):
        assert section == "cover", "Only the visible cover may start work"
        photo = next(block for block in draft.blocks if block.key == "cover_photo")
        raw = loader(photo.asset_hashes[0]) if photo.asset_hashes and photo.source != "Omit" else None
        rendered.append((draft, raw))
        with fitz.open() as pdf:
            page = pdf.new_page()
            page.insert_text((40, 40), "Synthetic preview transport control")
            if raw:
                page.insert_image(fitz.Rect(40, 80, 440, 380), stream=raw)
            return ReportPreview(pdf.tobytes(), 1, section_preview_fingerprint(draft, section), draft.period.label)

    monkeypatch.setattr(ui, "preview_section", convert)
    delivery = {}

    def with_receipt(draft, section, assets, report_prefix, **kwargs):
        if section == "cover" and delivery:
            state = ui._ordered_preview_state(report_prefix)
            generation = state["wanted"].get(section)
            if generation:
                key = report_prefix + "_preview_surface_" + section
                kind = delivery.pop("kind")
                if kind == "viewport":
                    ui.st.session_state[key] = {"viewport": {"generation": generation, "visible": True,
                                                               "pages": [0], "dpi": 96}}
                    ui._viewport_changed(report_prefix, section, key)
                else:
                    ui.st.session_state[key] = {"poll": {"generation": generation, "sequence": 1}}
                    ui._poll_changed(report_prefix, section, key)
        return actual_render(draft, section, assets, report_prefix, **kwargs)

    monkeypatch.setattr(ui, "render_section_preview", with_receipt)

    def receipt(kind):
        delivery["kind"] = kind
        app.run()
        assert not app.exception

    def finish_latest():
        future, function, args = jobs[-1]
        future.set_result(function(*args))
        receipt("poll")

    app.run()
    assert not app.exception and not jobs
    if older_preview:
        receipt("viewport")
        assert len(jobs) == 1
        finish_latest()
        assert payload(app, "Cover")["ready"]
        assert rendered[-1][1] is None

    draft = current(app)
    document = Document()
    document.add_paragraph(draft.profile.facilities[0].title)
    document.add_paragraph(draft.period.label)
    document.add_picture(BytesIO(photograph()), width=Inches(5))
    document.add_section()
    document.add_heading("Monthly Activity Summary", 1)
    document.add_paragraph("Synthetic fan inspection completed.")
    source = tmp_path / "synthetic-first-cover.docx"
    document.save(source)
    original = source.read_bytes()
    upload = BytesIO(original)
    upload.name, upload.size = source.name, len(original)
    real_upload = section_ui.st.file_uploader

    def uploader(label, *args, **kwargs):
        native = real_upload(label, *args, **kwargs)
        return upload if label == "Starting report" else native

    monkeypatch.setattr(section_ui.st, "file_uploader", uploader)
    app.run()
    assert not app.exception
    stage = next(value for key, value in app.session_state.filtered_state.items() if key.endswith("_stage"))
    plans = next(value for key, value in app.session_state.filtered_state.items() if key.endswith("_section_plans"))
    selected = plans["cover"]["cover_picture_id"]
    assert selected and plans["cover"]["destinations"][selected] == "cover_photo"
    expected = read_import_image(stage[1], next(item for item in stage[2].items if item.id == selected),
                                 line_art=False).data
    state = app.session_state[prefix + "_ordered_previews"]
    requested = state["requests"]["cover"]
    requested_photo = next(block for block in requested[1].blocks if block.key == "cover_photo")
    assert requested_photo.asset_hashes
    assert requested[2][requested_photo.asset_hashes[0]] == expected
    assert not next(block for block in current(app).blocks if block.key == "cover_photo").asset_hashes
    assert source.read_bytes() == original
    if older_preview:
        view = payload(app, "Cover")
        assert view["stale"] and view["pending"] and not view["ready"]
    receipt("viewport")
    clock[0] += ui.DEBOUNCE_SECONDS + .01
    receipt("poll")
    assert len(jobs) == (2 if older_preview else 1)
    finish_latest()
    view = payload(app, "Cover")
    assert view["ready"] and not view["stale"] and not view["pending"]
    assert view["status"] == "Current · 1 page" and len(view["pages"]) == 1
    assert rendered[-1][1] == expected
    assert state["cache"]["cover"]["fingerprint"] == requested[0] == view["generation"]
    assert source.read_bytes() == original
