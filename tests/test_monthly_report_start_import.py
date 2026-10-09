"""Synthetic starting reports exercise the same single-upload inference path."""

from dataclasses import replace
from io import BytesIO

from docx import Document
from docx.shared import Inches
from PIL import Image
import pytest
from streamlit.testing.v1 import AppTest

from app import monthly_report_library as library
from app.monthly_report_asset_review import pending_asset_indexes
from app.monthly_report_import import inspect_docx, imported_draft
from app.monthly_report_model import Facility, ReportPeriod, synthetic_profiles
from app.monthly_report_start_import import analyze_starting_report, automatic_section_plans, build_starting_import, mentioned_periods


PERIOD = ReportPeriod(2026, 9)


def report_file(tmp_path, *, title=None, month="September 2026", activity=(), photo=False, contact=True, table=False):
    profile = synthetic_profiles()[0]
    doc = Document()
    doc.add_heading(title or profile.facilities[0].title, 0)
    doc.add_paragraph("Monthly Report · " + month)
    if contact:
        doc.add_heading("1 Organizational Chart", 1)
        grid = doc.add_table(rows=2, cols=2)
        for row, values in zip(grid.rows, (("Role", "Name"), ("Site lead", "Synthetic Contact"))):
            for cell, value in zip(row.cells, values):
                cell.text = value
    doc.add_heading("2 Monthly Activity Summary", 1)
    for text in activity:
        doc.add_paragraph(text)
    if table:
        grid = doc.add_table(rows=3, cols=2)
        for row, values in zip(grid.rows, (("Month", "PM count"), ("July 2026", "4"), ("August 2026", "8"))):
            for cell, value in zip(row.cells, values):
                cell.text = value
    if photo:
        data = BytesIO()
        Image.new("RGB", (400, 300), "navy").save(data, "PNG")
        doc.add_picture(BytesIO(data.getvalue()), width=Inches(4))
    path = tmp_path / "synthetic-start.docx"
    doc.save(path)
    return path


def build(path, profile=None, known=()):
    profile = profile or synthetic_profiles()[0]
    inspection = inspect_docx(path)
    analysis = analyze_starting_report(inspection, profile, PERIOD, known)
    plans = automatic_section_plans(inspection, analysis)
    mapped, _ = build_starting_import(path, inspection, plans.values())
    return analysis, imported_draft(profile, PERIOD, "Synthetic Editor", mapped), plans


def test_current_half_finished_report_keeps_completed_and_undated_work_without_kind_question(tmp_path):
    path = report_file(tmp_path, activity=("September 2026 pump repair completed.", "Inspected the air handler."), photo=True)
    analysis, draft, plans = build(path)
    assert analysis.site_relation == "same" and analysis.source_period == (2026, 9)
    assert all(plan["approved"] for plan in plans.values())
    activity = next(b for b in draft.blocks if b.key == "activity_summary")
    assert "pump repair" in activity.text and "air handler" in activity.text
    picture = next(b for b in draft.blocks if b.key == "improvements")
    # Routing a picture is not a claim that its contents were human-reviewed.
    assert pending_asset_indexes(picture) == (0,)
    assert next(b for b in draft.blocks if b.key == "contact_matrix").extra_tables[0].rows


def test_earlier_report_rolls_monthly_content_forward_but_keeps_standing_contacts_and_ledger(tmp_path):
    path = report_file(tmp_path, month="July 2026", activity=("Repaired the pump.",), photo=True, table=True)
    analysis, draft, _ = build(path)
    assert analysis.source_period == (2026, 7) and analysis.older_items
    assert not any(b.key in ("activity_summary", "improvements") for b in draft.blocks)
    assert next(b for b in draft.blocks if b.key == "contact_matrix").extra_tables[0].rows
    ledger = next(b for b in draft.blocks if b.key == "work_orders").extra_tables[0]
    assert ledger.rows == (("July 2026", "4"), ("August 2026", "8"))
    assert "Repaired the pump." in " ".join(i.text for i in inspect_docx(path).items)


def test_stale_cover_does_not_discard_current_work_or_undated_pictures(tmp_path):
    path = report_file(tmp_path, month="July 2026", activity=("September 2026 repair completed.", "Undated new work.", "July 2026 old monthly update."), photo=True)
    analysis, draft, _ = build(path)
    assert analysis.current_items and analysis.older_items
    text = next(b.text for b in draft.blocks if b.key == "activity_summary")
    assert "September 2026" in text and "Undated new work" in text and "July 2026 old" not in text
    assert next(b for b in draft.blocks if b.key == "improvements").asset_hashes


def test_different_site_from_same_contract_reuses_schema_without_its_contacts_photos_or_work(tmp_path):
    north, south = synthetic_profiles()[1].facilities
    path = report_file(tmp_path, title=south.title, activity=("September 2026 South work",), photo=True, table=True)
    profile = synthetic_profiles()[0]
    analysis, draft, plans = build(path, profile, (north, south))
    assert analysis.site_relation == "other"
    assert draft.profile.facilities == (north,)
    assert not any(b.text or b.rows or b.asset_hashes or any(t.rows for t in b.extra_tables) for b in draft.blocks)
    assert next(b for b in draft.blocks if b.key == "contact_matrix").extra_tables[0].columns == ("Role", "Name")
    assert all(plan["approved"] for plan in plans.values())


def test_unknown_identity_requires_one_membership_confirmation_not_a_report_kind(tmp_path):
    path = report_file(tmp_path, title="Unnamed monthly report")
    inspection = inspect_docx(path)
    profile = synthetic_profiles()[0]
    assert analyze_starting_report(inspection, profile, PERIOD).site_relation == "unknown"
    assert analyze_starting_report(inspection, profile, PERIOD, confirmed_sites=profile.facilities).site_relation == "same"


def test_cover_matching_respects_alias_boundaries_and_saved_group_membership(tmp_path):
    north, south = synthetic_profiles()[1].facilities
    path = report_file(tmp_path, title="Demo North")
    assert analyze_starting_report(inspect_docx(path), synthetic_profiles()[0], PERIOD).site_relation == "same"
    path = report_file(tmp_path, title="Western Operations Region")
    group = replace(synthetic_profiles()[1], title="Western Operations Region")
    analysis = analyze_starting_report(inspect_docx(path), synthetic_profiles()[0], PERIOD, (north, south), known_profiles=(group,))
    assert analysis.site_relation == "other" and analysis.source_sites == (north, south)
    unity = Facility("unity", "Unity Hospital", ("Unity",))
    specialty = Facility("specialty", "Unity Specialty Hospital")
    path = report_file(tmp_path, title=specialty.title)
    target = replace(synthetic_profiles()[0], facilities=(unity,))
    assert analyze_starting_report(inspect_docx(path), target, PERIOD, (unity, specialty)).source_sites == (specialty,)


@pytest.mark.parametrize("title", ("Monthly Report", "Report", "Monthly Status Report", "ENFRA Monthly Report", "Demonstration Monthly Report", "Two-Site Monthly Report"))
@pytest.mark.parametrize("scope", (0, 1))
def test_generic_report_design_title_never_identifies_unknown_source_site(tmp_path, title, scope):
    # Actual Word inspection must not turn a generic subtitle into proof that
    # an unknown South site belongs to the saved North site/group design.
    path = report_file(tmp_path, title="Unknown South Hospital", activity=("September 2026 South-only work.",))
    doc = Document(path)
    doc.paragraphs[1].text = title + " · September 2026"
    doc.tables[0].cell(1, 1).text = "Synthetic South Contact"
    doc.save(path)
    target = replace(synthetic_profiles()[scope], title=title)
    inspection = inspect_docx(path)
    analysis = analyze_starting_report(inspection, target, PERIOD, known_profiles=(target,))
    assert analysis.site_relation == "unknown" and not analysis.source_sites
    mapped, _ = build_starting_import(path, inspection, automatic_section_plans(inspection, analysis).values())
    assert not any("Synthetic South Contact" in str(block) or "South-only work" in str(block) for block in mapped.blocks)


def test_generic_profile_title_still_allows_actual_facility_evidence(tmp_path):
    path = report_file(tmp_path)
    target = replace(synthetic_profiles()[0], title="Monthly Report")
    assert analyze_starting_report(inspect_docx(path), target, PERIOD).site_relation == "same"


def test_contract_initials_are_not_a_group_identity(tmp_path):
    path = report_file(tmp_path, title="Unknown South Hospital")
    doc = Document(path)
    doc.paragraphs[1].text = "SSC Monthly Report · September 2026"
    doc.save(path)
    target = replace(synthetic_profiles()[1], contract="Synthetic Services Contract", title="SSC Monthly Report")
    assert analyze_starting_report(inspect_docx(path), target, PERIOD).site_relation == "unknown"


def test_dates_need_explicit_year_and_understand_body_date_formats():
    assert mentioned_periods("September inspection") == set()
    assert mentioned_periods("September 3, 2026; 2026-08-31; 9/2/2026") == {(2026, 8), (2026, 9)}


def test_catalog_spelling_does_not_change_same_site_identity(tmp_path):
    profile = synthetic_profiles()[0]
    catalog = Facility("demonstration-north-facility", "Demo North", (profile.facilities[0].title,))
    path = report_file(tmp_path, title="Demo North")
    analysis = analyze_starting_report(inspect_docx(path), profile, PERIOD, (catalog,))
    assert analysis.site_relation == "same" and analysis.source_sites[0].key == "north"


def start_app(monkeypatch, tmp_path, path, *, known_sites=None):
    from app import monthly_report_section_ui as ui
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path / "runtime"))
    upload = BytesIO(path.read_bytes())
    upload.name, upload.size = path.name, len(upload.getvalue())
    real = ui.st.file_uploader
    monkeypatch.setattr(ui.st, "file_uploader", lambda label, *a, **kw: upload if label == "Starting report" else real(label, *a, **kw))
    if known_sites:
        monkeypatch.setattr(ui, "_site_options", lambda contract: known_sites)
    app = AppTest.from_string('''
from app.monthly_report_section_ui import render_section_setup
from app.monthly_report_model import synthetic_profiles, ReportPeriod
from app.monthly_report_ui import _field
p = synthetic_profiles()[0]
render_section_setup(p.contract, ReportPeriod(2026, 9), "Synthetic Editor", _field, identity=p)
''', default_timeout=30).run()
    next(b for b in app.button if b.label == "Analyze report").click().run()
    assert not app.exception
    return app


def save_start(app):
    assert not any(w.label.startswith("This section is ready") for w in app.checkbox)
    next(b for b in app.button if b.label == "Continue to this month’s updates").click().run()
    assert not app.exception and not app.error
    profile = synthetic_profiles()[0]
    return library.load_imported_draft(profile.contract, profile.key)


def test_single_upload_automatically_imports_current_content_and_saves_original(monkeypatch, tmp_path):
    path = report_file(tmp_path, activity=("September 2026 finished work",))
    original = path.read_bytes()
    app = start_app(monkeypatch, tmp_path, path)
    assert not any(w.label == "Sites shown in the starting report" for w in app.multiselect)
    draft = save_start(app)
    assert all(s.included for s in draft.sections)
    assert "finished work" in next(b.text for b in draft.blocks if b.key == "activity_summary")
    assert library.imported_original(draft.profile.contract, draft.profile.key).read_bytes() == original


def test_single_upload_other_site_keeps_target_identity_and_removes_source_contacts(monkeypatch, tmp_path):
    sites = synthetic_profiles()[1].facilities
    path = report_file(tmp_path, title=sites[1].title, activity=("South-only work.",))
    app = start_app(monkeypatch, tmp_path, path, known_sites=sites)
    draft = save_start(app)
    assert draft.profile.facilities == (sites[0],)
    assert not any("Synthetic Contact" in str(b) or "South-only work" in b.text for b in draft.blocks)


def test_unknown_source_site_blocks_start_until_membership_is_confirmed(monkeypatch, tmp_path):
    path = report_file(tmp_path, title="Report without site name", activity=("September 2026 current work",))
    app = start_app(monkeypatch, tmp_path, path)
    assert not any(b.label == "Continue to this month’s updates" for b in app.button)
    selector = next(w for w in app.multiselect if w.label == "Sites shown in the starting report")
    selector.set_value([synthetic_profiles()[0].facilities[0].key]).run()
    assert "current work" in next(b.text for b in save_start(app).blocks if b.key == "activity_summary")


def test_older_report_with_no_standing_content_can_start_blank_and_keep_original(monkeypatch, tmp_path):
    path = report_file(tmp_path, month="July 2026", contact=False, activity=("July 2026 completed work",))
    app = start_app(monkeypatch, tmp_path, path)
    draft = save_start(app)
    assert all(s.included for s in draft.sections)
    assert not any(b.text for b in draft.blocks)
