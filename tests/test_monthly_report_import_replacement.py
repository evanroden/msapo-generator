"""A rejected replacement must not discard the last usable DOCX mapping review."""
from dataclasses import replace
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED

from docx import Document
from docx.shared import Inches
from PIL import Image
import pytest
from streamlit.testing.v1 import AppTest

from app import monthly_report_import as importer, monthly_report_import_ui as ui
from app import monthly_report_library as library
from app.monthly_report_model import ReportPeriod, synthetic_profiles


APP = '''
import streamlit as st
from app import monthly_report_library as library
from app.monthly_report_import_ui import render_import
from app.monthly_report_ui import _field
profile = st.session_state["test_profile"]
state = library.load_profile(profile.contract, profile.key)
render_import(state, _field)
'''


class Upload(BytesIO):
    def __init__(self, raw, *, size=None):
        super().__init__(raw)
        self.name = "synthetic.docx"
        self.size = len(raw) if size is None else size


def source_bytes(label="Original synthetic activity"):
    doc = Document()
    doc.add_heading("2 Monthly Activity Summary", 1)
    doc.add_paragraph(label)
    table = doc.add_table(rows=3, cols=2)
    for row, values in zip(table.rows, (("Equipment", "Observation"),
                                       ("Pump 001", "Operating"), ("Pump 002", ""))):
        for cell, text in zip(row.cells, values):
            cell.text = text
    picture = BytesIO()
    Image.new("RGB", (180, 120), "navy").save(picture, "PNG")
    doc.add_picture(BytesIO(picture.getvalue()), width=Inches(2))
    output = BytesIO()
    doc.save(output)
    return output.getvalue()


def element(app, kind, label):
    return next(w for w in getattr(app, kind) if w.label == label)


def click(app, label):
    element(app, "button", label).click().run()
    assert not app.exception
    return app


def prefix(profile):
    return "report_import_" + ui._signature((profile.contract, profile.key))


def inventory(root):
    """Only durable library files; a failed read can create an empty staging dir."""
    return {str(p.relative_to(root)): p.read_bytes() for p in root.rglob("*")
            if p.is_file() and "imports" not in p.relative_to(root).parts}


@pytest.fixture
def review(tmp_path, monkeypatch):
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path / "data"))
    profile = synthetic_profiles()[0]
    library.save_profile(profile, expected_revision=0, actor="Synthetic Editor", confirmed=True)
    app = AppTest.from_string(APP, default_timeout=20)
    app.session_state["test_profile"] = profile
    original = source_bytes()
    app.session_state["test_upload"] = Upload(original)
    monkeypatch.setattr(ui.st, "file_uploader", lambda *args, **kwargs: ui.st.session_state["test_upload"])
    app.run()
    click(app, "Read DOCX for mapping")
    staged = app.session_state[prefix(profile) + "_stage"]
    text = next(i for i in staged[2].items if i.text == "Original synthetic activity")
    element(app, "selectbox", "Item to inspect").set_value(text.id).run()
    click(app, "Add confirmed mapping")
    table = next(i for i in staged[2].items if i.kind == "table")
    element(app, "selectbox", "Item to inspect").set_value(table.id).run()
    click(app, "Add confirmed mapping")
    image = next(i for i in staged[2].items if i.kind == "image")
    element(app, "selectbox", "Item to inspect").set_value(image.id).run()
    click(app, "Preview extracted image")
    element(app, "text_input", "Import editor name").set_value("Synthetic Reviewer").run()
    assert not app.exception
    yield app, profile, staged, original
    # Each retained session owns its own temporary review; no shared originals.
    current = app.session_state.filtered_state.get(prefix(profile) + "_stage")
    if current:
        current[0].cleanup()
    staged[0].cleanup()


def rejected_upload(kind, monkeypatch):
    raw = source_bytes("Replacement synthetic activity")
    if kind == "not_zip":
        return Upload(b"not a DOCX"), "DOCX"
    if kind == "bad_xml":
        output = BytesIO()
        with ZipFile(BytesIO(raw)) as source, ZipFile(output, "w", ZIP_DEFLATED) as target:
            for name in source.namelist():
                target.writestr(name, b"<broken" if name == "word/document.xml" else source.read(name))
        return Upload(output.getvalue()), "read"
    if kind == "expanded":
        monkeypatch.setattr(importer, "MAX_EXPANDED_BYTES", 1)
        return Upload(raw), "Expanded"
    if kind == "size":
        monkeypatch.setattr(ui, "MAX_DOCX_BYTES", len(raw) - 1)
        return Upload(raw), "import limit"
    if kind == "stream_size":
        monkeypatch.setattr(ui, "MAX_DOCX_BYTES", len(raw) - 1)
        return Upload(raw, size=1), "import limit"
    if kind == "io":
        def unavailable(path):
            raise OSError("Synthetic replacement read failure")
        monkeypatch.setattr(ui, "inspect_docx", unavailable)
        return Upload(raw), "Synthetic replacement read failure"
    raise AssertionError(kind)


@pytest.mark.parametrize("kind", ["not_zip", "bad_xml", "expanded", "size", "stream_size", "io"])
def test_rejected_replacement_preserves_stage_mappings_preview_and_actor(review, monkeypatch, kind):
    app, profile, staged, original = review
    key = prefix(profile)
    mappings = dict(app.session_state[key + "_mappings"])
    preview = app.session_state[key + "_preview"]
    app.session_state[key + "_saved"] = "Earlier confirmed save remains recorded."
    before = inventory(library._root())
    upload, reason = rejected_upload(kind, monkeypatch)
    incoming = upload.getvalue()
    app.session_state["test_upload"] = upload
    click(app, "Read DOCX for mapping")
    assert any(reason in error.value for error in app.error)
    assert app.session_state.filtered_state.get(key + "_stage") is staged
    assert staged[1].read_bytes() == original
    assert app.session_state[key + "_mappings"] == mappings
    assert app.session_state[key + "_preview"] == preview
    assert app.session_state[key + "_saved"] == "Earlier confirmed save remains recorded."
    assert element(app, "text_input", "Import editor name").value == "Synthetic Reviewer"
    assert len(app.get("image")) == 1
    assert any("previous DOCX review" in c.value for c in app.caption)
    assert upload.getvalue() == incoming and upload.tell() == 0
    assert inventory(library._root()) == before
    assert set((library._root() / "imports").iterdir()) == {Path(staged[0].name)}
    app.run()
    assert not app.exception
    assert app.session_state[key + "_stage"] is staged
    assert app.session_state[key + "_mappings"] == mappings
    assert element(app, "text_input", "Import editor name").value == "Synthetic Reviewer"
    assert any("last successfully read DOCX" in c.value for c in app.caption)
    assert inventory(library._root()) == before


def test_first_failed_read_leaves_no_stage_or_durable_change(tmp_path, monkeypatch):
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path / "data"))
    profile = synthetic_profiles()[0]
    library.save_profile(profile, expected_revision=0, actor="Synthetic Editor", confirmed=True)
    app = AppTest.from_string(APP)
    app.session_state["test_profile"] = profile
    app.session_state["test_upload"] = Upload(b"not a DOCX")
    monkeypatch.setattr(ui.st, "file_uploader", lambda *args, **kwargs: ui.st.session_state["test_upload"])
    before = inventory(library._root())
    app.run()
    click(app, "Read DOCX for mapping")
    assert app.error
    assert prefix(profile) + "_stage" not in app.session_state.filtered_state
    assert not any("previous DOCX review" in c.value for c in app.caption)
    assert not list((library._root() / "imports").iterdir())
    assert inventory(library._root()) == before


def test_successful_replacement_retires_old_stage_only_after_inspection(review, monkeypatch):
    app, profile, old, original = review
    key = prefix(profile)
    app.session_state[key + "_saved"] = "Earlier confirmed save."
    raw = source_bytes("Replacement synthetic activity")
    app.session_state["test_upload"] = Upload(raw)
    inspect = ui.inspect_docx
    observations = []
    def check_old_before_inspect(path):
        observations.append((old[1].read_bytes(), dict(ui.st.session_state[key + "_mappings"])))
        return inspect(path)
    monkeypatch.setattr(ui, "inspect_docx", check_old_before_inspect)
    before = inventory(library._root())
    click(app, "Read DOCX for mapping")
    assert len(observations) == 1 and observations[0][0] == original and len(observations[0][1]) == 2
    staged = app.session_state[key + "_stage"]
    assert staged is not old and staged[1].read_bytes() == raw
    assert not old[1].exists()
    assert app.session_state[key + "_mappings"] == {}
    assert key + "_preview" not in app.session_state.filtered_state
    assert key + "_saved" not in app.session_state.filtered_state
    assert len(list((library._root() / "imports").iterdir())) == 1
    assert inventory(library._root()) == before


def test_retry_succeeds_without_touching_another_profile_review(review, monkeypatch):
    app, profile, old, original = review
    key = prefix(profile)
    other_profile = replace(profile, contract="Another synthetic contract", key="other")
    other_key = prefix(other_profile)
    other = ui._stage(Upload(source_bytes("Independent review")))
    app.session_state[other_key + "_stage"] = other
    app.session_state[other_key + "_mappings"] = {"unrelated": "unchanged"}
    reads = []
    stage = ui._stage
    def record(upload):
        reads.append(upload.getvalue())
        return stage(upload)
    monkeypatch.setattr(ui, "_stage", record)
    try:
        for _ in range(2):
            app.session_state["test_upload"] = Upload(b"not a DOCX")
            click(app, "Read DOCX for mapping")
            assert app.session_state[key + "_stage"] is old
            assert old[1].read_bytes() == original
            before_reads = len(reads)
            app.run()
            assert not app.exception and len(reads) == before_reads
        assert len(reads) == 2
        raw = source_bytes("Accepted replacement")
        app.session_state["test_upload"] = Upload(raw)
        click(app, "Read DOCX for mapping")
        new = app.session_state[key + "_stage"]
        assert new[1].read_bytes() == raw and not old[1].exists()
        assert app.session_state[key + "_mappings"] == {}
        assert app.session_state[other_key + "_stage"] is other
        assert app.session_state[other_key + "_mappings"] == {"unrelated": "unchanged"}
        assert other[1].exists()
        app.run()
        assert app.session_state[key + "_stage"] is new
        assert set((library._root() / "imports").iterdir()) == {Path(new[0].name), Path(other[0].name)}
    finally:
        other[0].cleanup()


def test_failed_replacement_can_be_discarded_explicitly(review):
    app, profile, old, original = review
    key = prefix(profile)
    app.session_state["test_upload"] = Upload(b"not a DOCX")
    click(app, "Read DOCX for mapping")
    assert old[1].read_bytes() == original
    before = inventory(library._root())
    element(app, "checkbox", "Discard this staged import and unmatched review items").check().run()
    click(app, "Discard staged DOCX")
    assert not old[1].exists()
    assert all(key + suffix not in app.session_state.filtered_state
               for suffix in ("_stage", "_mappings", "_preview", "_saved"))
    assert inventory(library._root()) == before
    assert not list((library._root() / "imports").iterdir())


def test_selecting_or_clearing_unread_replacement_does_not_change_review(review):
    app, profile, old, original = review
    mappings = dict(app.session_state[prefix(profile) + "_mappings"])
    for upload in (Upload(source_bytes("Not yet read")), None):
        app.session_state["test_upload"] = upload
        app.run()
        assert not app.exception
        assert app.session_state[prefix(profile) + "_stage"] is old
        assert app.session_state[prefix(profile) + "_mappings"] == mappings
        assert old[1].read_bytes() == original
        assert element(app, "button", "Read DOCX for mapping").disabled is (upload is None)


def test_cleanup_failure_does_not_mislabel_successful_replacement(review, monkeypatch):
    app, profile, old, original = review
    cleanup = old[0].cleanup
    def fail():
        raise OSError("Synthetic previous-directory cleanup failure")
    monkeypatch.setattr(old[0], "cleanup", fail)
    try:
        raw = source_bytes("Replacement ready")
        app.session_state["test_upload"] = Upload(raw)
        click(app, "Read DOCX for mapping")
        assert not app.error
        assert any("New DOCX ready" in w.value and "could not be removed" in w.value for w in app.warning)
        new = app.session_state[prefix(profile) + "_stage"]
        assert new is not old and new[1].read_bytes() == raw
        assert app.session_state[prefix(profile) + "_mappings"] == {}
        assert old[1].read_bytes() == original
    finally:
        monkeypatch.setattr(old[0], "cleanup", cleanup)
        cleanup()


@pytest.mark.parametrize("target", ["Shared library defaults", "Prior report snapshot"])
def test_previous_confirmed_mappings_remain_saveable_after_bad_replacement(review, monkeypatch, target):
    from datetime import date
    app, profile, old, original = review
    monkeypatch.setattr(ui, "operator_today", lambda _timezone: date(2026, 10, 28))
    app.session_state["test_upload"] = Upload(b"not a DOCX")
    click(app, "Read DOCX for mapping")
    assert old[1].read_bytes() == original
    element(app, "radio", "Save confirmed content as").set_value(target).run()
    element(app, "selectbox", "Original report month").set_value(9).run()
    element(app, "selectbox", "Original report year").set_value(2026).run()
    if target == "Prior report snapshot":
        element(app, "text_input", "Original prepared by").set_value("Synthetic Preparer").run()
    assert element(app, "text_input", "Import editor name").value == "Synthetic Reviewer"
    element(app, "checkbox", "I confirm these facility identities, aliases and scope").check().run()
    element(app, "checkbox", "I confirm these mappings and this shared save").check().run()
    click(app, "Save confirmed DOCX mappings")
    assert not app.error
    assert any("Saved 2 confirmed items" in s.value for s in app.success)
    if target == "Prior report snapshot":
        saved = library.load_snapshot(profile.contract, profile.key, ReportPeriod(2026, 9))
        assert saved.revision == 1 and saved.entered_editor == "Synthetic Reviewer"
        blocks = saved.draft.blocks
    else:
        saved = library.load_profile(profile.contract, profile.key)
        assert saved.revision == 2
        blocks = tuple(saved.block(slot) for slot, _ in saved.current)
    assert any(b.text == "Original synthetic activity" for b in blocks)
    assert any(b.rows == (("Pump 001", "Operating"), ("Pump 002", "")) for b in blocks)
    assert old[1].read_bytes() == original
    assert len(app.session_state[prefix(profile) + "_mappings"]) == 2


def test_actor_edit_delivered_with_failed_read_is_retained(review):
    app, profile, old, original = review
    app.session_state["test_upload"] = Upload(b"not a DOCX")
    element(app, "text_input", "Import editor name").set_value("Edited during replacement")
    click(app, "Read DOCX for mapping")
    assert app.session_state[prefix(profile) + "_stage"] is old
    assert element(app, "text_input", "Import editor name").value == "Edited during replacement"
    assert old[1].read_bytes() == original
    app.run()
    assert element(app, "text_input", "Import editor name").value == "Edited during replacement"
