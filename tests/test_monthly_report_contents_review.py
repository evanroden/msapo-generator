"""Contents warnings describe retained source text, never certify pagination."""
from dataclasses import replace
from io import BytesIO

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import pytest

from app import monthly_report_designs as designs, monthly_report_docx as renderer
from app.monthly_report_model import (
    BlockSpec, ReportDraft, ReportPeriod, ResolvedBlock, SectionSpec, synthetic_profiles,
)


def source(tmp_path, kind):
    document = Document()
    document.add_paragraph("Synthetic Site")
    document.add_paragraph("September 2026")
    if kind == "static":
        document.add_paragraph("Table of contents")
        document.add_paragraph("Monthly Activity Summary\t3")
        document.add_paragraph("Training Summary\t4")
    elif kind == "simple":
        paragraph = document.add_paragraph()
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), ' TOC \\o "1-3" \\h ')
        run, text = OxmlElement("w:r"), OxmlElement("w:t")
        text.text = "Monthly Activity Summary 3"
        run.append(text)
        field.append(run)
        paragraph._p.append(field)
    elif kind == "complex":
        for tag, value in (("fldChar", "begin"), ("instrText", " TO"),
                           ("instrText", 'C \\o "1-3" '), ("fldChar", "separate"),
                           ("t", "Monthly Activity Summary 3"), ("fldChar", "end")):
            run = document.add_paragraph().add_run()._r
            node = OxmlElement("w:" + tag)
            if tag == "fldChar":
                node.set(qn("w:fldCharType"), value)
            else:
                node.text = value
            run.append(node)
    elif kind == "prose":
        document.add_paragraph("The contents of this report are confidential.")
        document.add_paragraph("See the table of contents for report navigation.")
    document.add_heading("MONTHLY ACTIVITY SUMMARY", 1)
    document.add_paragraph("Prior activity")
    document.add_heading("TRAINING SUMMARY", 1)
    document.add_paragraph("Prior training")
    path = tmp_path / (kind + ".docx")
    document.save(path)
    return path


def draft(tmp_path, monkeypatch, kind):
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path / "data"))
    path = source(tmp_path, kind)
    profile = designs.install(path, synthetic_profiles()[0], actor="Synthetic Editor", expected_revision=0)
    value = ReportDraft(profile, ReportPeriod(2026, 10), "Synthetic Editor", (
        SectionSpec("activity", "1", "Monthly Activity Summary", (BlockSpec("activity_summary", "rich_text"),)),
        SectionSpec("training", "2", "Training Summary", (BlockSpec("training_summary", "rich_text"),), included=False),
    ), (ResolvedBlock("activity_summary", "This month", text="Current synthetic work."),))
    return path, value


@pytest.mark.parametrize("kind,label", [("static", "static"), ("simple", "Word contents field"), ("complex", "Word contents field")])
def test_finished_package_exposes_unavailable_contents_update(tmp_path, monkeypatch, kind, label):
    path, value = draft(tmp_path, monkeypatch, kind)
    original = path.read_bytes()
    monkeypatch.setattr(renderer, "_convert_report_pdf", lambda *args: (b"%PDF-synthetic", ""))
    result = renderer.generate_report(value, acknowledged_fingerprint=value.fingerprint)
    assert label in result.contents_notice
    assert "not recalculated" in result.contents_notice
    assert "titles, order and page numbers" in result.contents_notice
    assert result.docx and result.pdf and not result.pdf_error
    assert path.read_bytes() == original
    assert designs.source_for(value.profile).read_bytes() == original
    # Retained fields must remain passive; a warning never re-enables fields.
    generated = Document(BytesIO(result.docx))
    assert not any("TOC" in (node.text or "") for node in generated.element.iter(qn("w:instrText")))


@pytest.mark.parametrize("kind", ["none", "prose"])
def test_normal_prose_and_no_contents_do_not_acquire_warning(tmp_path, monkeypatch, kind):
    _, value = draft(tmp_path, monkeypatch, kind)
    monkeypatch.setattr(renderer, "_convert_report_pdf", lambda *args: (None, "Synthetic conversion unavailable"))
    result = renderer.generate_report(value, acknowledged_fingerprint=value.fingerprint)
    assert result.contents_notice == ""


def test_contents_notice_survives_pdf_retry(tmp_path, monkeypatch):
    _, value = draft(tmp_path, monkeypatch, "static")
    monkeypatch.setattr(renderer, "_convert_report_pdf", lambda *args: (None, "Synthetic first failure"))
    package = renderer.generate_report(value, acknowledged_fingerprint=value.fingerprint)
    monkeypatch.setattr(renderer, "_convert_report_pdf", lambda *args: (b"%PDF-synthetic", ""))
    retried = renderer.retry_report_pdf(value, package)
    assert retried.contents_notice == package.contents_notice
    assert "not recalculated" in retried.contents_notice
    assert retried.docx == package.docx and retried.pdf and not retried.pdf_error


def test_selected_design_version_controls_notice_not_latest_master(tmp_path, monkeypatch):
    from app.monthly_report_contents_review import contents_notice_for
    _, value = draft(tmp_path, monkeypatch, "static")
    newer = designs.install(source(tmp_path, "none"), value.profile,
                            actor="Synthetic Editor", expected_revision=1)
    assert "static contents" in contents_notice_for(value.profile)
    assert contents_notice_for(newer) == ""


def test_lost_source_after_assembly_keeps_files_with_unavailable_notice(tmp_path, monkeypatch):
    path, value = draft(tmp_path, monkeypatch, "static")
    original = path.read_bytes()
    pinned = designs.source_for(value.profile)
    def convert(*args):
        pinned.unlink()  # Isolated race after Word was successfully assembled.
        return b"%PDF-synthetic", ""
    monkeypatch.setattr(renderer, "_convert_report_pdf", convert)
    result = renderer.generate_report(value, acknowledged_fingerprint=value.fingerprint)
    assert result.docx and result.pdf and not result.pdf_error
    assert "Contents verification is unavailable" in result.contents_notice
    assert path.read_bytes() == original


def test_bare_toc_result_and_unrelated_fields_are_not_contents():
    from app.monthly_report_contents_review import contents_kind
    document = Document()
    paragraph = document.add_paragraph()
    for tag, value in (("fldChar", "begin"), ("instrText", " PAGE "),
                       ("fldChar", "separate"), ("t", "TOC"), ("fldChar", "end")):
        node = OxmlElement("w:" + tag)
        if tag == "fldChar":
            node.set(qn("w:fldCharType"), value)
        else:
            node.text = value
        paragraph.add_run()._r.append(node)
    assert contents_kind(document.element) == ""


@pytest.mark.parametrize("pdf", [True, False])
def test_finished_files_show_contents_notice_after_shared_save_and_rerun(tmp_path, monkeypatch, pdf):
    from app import monthly_report_guided as guided
    from test_monthly_report_output_status_ui import ready, create
    from app.monthly_report_contents_review import contents_notice_for

    _, value = draft(tmp_path, monkeypatch, "static")
    notice = contents_notice_for(value.profile)
    app, calls = ready(monkeypatch, tmp_path / "ui", pdf=pdf)
    generate = guided.generate_report
    monkeypatch.setattr(guided, "generate_report", lambda *args, **kwargs:
                        replace(generate(*args, **kwargs), contents_notice=notice))
    create(app)
    assert any(item.value == notice for item in app.warning)
    assert any("[Shared version 1]" in item.value for item in app.caption)
    app.run()
    assert not app.exception
    assert sum(item.value == notice for item in app.warning) == 1
    assert [item.label for item in app.get("download_button")] == (
        ["Download DOCX", "Download PDF"] if pdf else ["Download DOCX"])
    assert len(calls) == 1
