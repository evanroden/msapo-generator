"""Imported Training values stay editable without inventing completed sessions."""
from dataclasses import replace
from datetime import date
import json

from docx import Document
import pytest
from streamlit.testing.v1 import AppTest

from app import monthly_report_library as library
from app.monthly_report_import import inspect_docx, imported_draft
from app.monthly_report_model import ReportPeriod, ReportTable, synthetic_profiles
from app.monthly_report_start_import import (
    analyze_starting_report, automatic_section_plans, build_starting_import,
)
from app.monthly_report_training import EVENT_REF, MATRIX_REF, add_event, load_training

PERIOD = ReportPeriod(2026, 9)
COLUMNS = ("Topic", "Date", "Audience", "Hours", "Source note")
ROWS = (("Current exercise", "2026-09-03", "Synthetic staff", "", "Recorded"),
        ("Prior briefing", "2026-08-01", "Other staff", "0", "Historical"),
        ("Incomplete topic", "", "", "", "Pending"))
APP = '''
import streamlit as st
from app.monthly_report_training_ui import render_training

def field(key, value):
    if key not in st.session_state:
        st.session_state[key] = value
    return key
st.session_state['blocks'] = render_training(
    st.session_state['draft'], st.session_state['blocks'], 'source_training', field)
'''


def current(app):
    return app.session_state["blocks"]["training_summary"]


def open_app(draft, block=None):
    app = AppTest.from_string(APP, default_timeout=30)
    app.session_state["draft"] = draft
    block = block or draft.blocks[0]
    app.session_state["blocks"] = {block.key: block}
    app.run()
    assert not app.exception
    return app


def source_grids(app):
    return [grid for grid in app.dataframe if "_training_source_table_" in (grid.key or "")]


def wire(app, index=None, *, edited=None, added=(), deleted=(), proto=None,
         action=None, row=0):
    """Deliver native grid protobufs alongside the ordinary sibling controls.

    AppTest's normal run omits data_editor wire state. An explicit neutral
    payload is needed to model a still-mounted native editor, not an unmount.
    """
    message = app._tree.get_widget_states()
    for grid in source_grids(app):
        try:
            delta = app.session_state[grid.key]
        except KeyError:
            delta = {"edited_rows": {}, "added_rows": [], "deleted_rows": []}
        if proto is not None and grid.proto.id == proto.id:
            continue
        if index is not None and grid is source_grids(app)[index] and proto is None:
            delta = {"edited_rows": edited or {}, "added_rows": list(added),
                     "deleted_rows": list(deleted)}
        message.widgets.add(id=grid.proto.id, string_value=json.dumps(delta))
    target = proto or (source_grids(app)[index].proto if index is not None else None)
    if proto is not None:
        message.widgets.add(id=proto.id, string_value=json.dumps({
            "edited_rows": edited or {}, "added_rows": list(added), "deleted_rows": list(deleted)}))
    if action is not None:
        message.widgets.add(id=next(iter(target.button_click_widgets.values())),
                            string_trigger_value={"data": json.dumps({"row": row, "label": action})})
    app._run(message)
    assert not app.exception
    return message


def edit(app, index=0, *, edited=None, added=(), deleted=()):
    assert len(source_grids(app)) > index, "Imported Training table has no correction control"
    return wire(app, index, edited=edited, added=added, deleted=deleted)


@pytest.fixture
def imported_training(tmp_path, monkeypatch):
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path / "runtime"))
    from app import monthly_report_directory as directory
    monkeypatch.setattr(directory, "load_directory", lambda *_: None)
    profile = synthetic_profiles()[0]
    path = tmp_path / "synthetic-training.docx"
    document = Document()
    document.add_paragraph(profile.facilities[0].title)
    document.add_paragraph("September 2026")
    document.add_heading("12 Training Summary", 1)
    table = document.add_table(rows=1 + len(ROWS), cols=len(COLUMNS))
    for target, values in zip(table.rows, (COLUMNS, *ROWS)):
        for cell, value in zip(target.cells, values):
            cell.text = value
    document.save(path)
    original = path.read_bytes()
    inspection = inspect_docx(path)
    analysis = analyze_starting_report(inspection, profile, PERIOD)
    plans = automatic_section_plans(inspection, analysis)
    assert plans['training']['approved']
    mapped, _ = build_starting_import(path, inspection, plans.values())
    draft = imported_draft(profile, PERIOD, "Synthetic Editor", mapped)
    assert len(draft.blocks) == 1
    block = draft.blocks[0]
    assert block.extra_tables[0].rows == ROWS
    assert block.extra_tables[0].reference.startswith('docx:')
    return draft, path, original


def test_actual_imported_table_is_reachable_and_passive_reruns_preserve_values(imported_training):
    draft, path, original = imported_training
    app = open_app(draft)
    assert len(source_grids(app)) == 1
    identity = source_grids(app)[0].key
    assert current(app) == draft.blocks[0]
    wire(app)
    assert source_grids(app)[0].key == identity
    assert current(app) == draft.blocks[0]
    assert path.read_bytes() == original
    assert load_training(draft.profile.contract)['revision'] == 0


@pytest.mark.parametrize('field,value', [('Hours', '1.25'), ('Date', ''),
                                        ('Source note', 'Keep [literal] text; Status: Pending')])
def test_report_only_correction_preserves_unknown_fields_and_source(imported_training, field, value):
    draft, path, original = imported_training
    app = open_app(draft)
    edit(app, edited={0: {field: value}})
    result = current(app)
    expected = list(ROWS[0]); expected[COLUMNS.index(field)] = value
    assert result.extra_tables[0].rows == (tuple(expected), *ROWS[1:])
    assert result.extra_tables[0].columns == COLUMNS
    assert result.extra_tables[0].reference == draft.blocks[0].extra_tables[0].reference
    assert result.references == draft.blocks[0].references
    assert not any(t.reference in (MATRIX_REF, EVENT_REF) for t in result.extra_tables)
    app.run()
    assert current(app) == result
    assert path.read_bytes() == original
    assert load_training(draft.profile.contract)['revision'] == 0


def test_canonical_training_without_imported_tables_adds_no_extra_editor(imported_training):
    draft, _, _ = imported_training
    matrix = ReportTable(('Team member', 'Site', 'Safety'),
                         (('Synthetic staff', draft.profile.facilities[0].title, ''),), MATRIX_REF)
    events = add_event(None, title='Explicit current event', when=date(2026, 9, 2),
                       mode='Virtual', participants=['Synthetic staff'], hours='', period=PERIOD)
    block = replace(draft.blocks[0], source='This month', extra_tables=(matrix, events))
    app = open_app(draft, block)
    assert not source_grids(app)
    assert current(app) == block


def test_source_correction_keeps_canonical_matrix_sessions_and_sibling_table(imported_training):
    draft, path, original = imported_training
    matrix = ReportTable(('Team member', 'Site', 'Safety'),
                         (('Synthetic staff', draft.profile.facilities[0].title, 'Completed'),), MATRIX_REF)
    events = add_event(None, title='Explicit current event', when=date(2026, 9, 2),
                       mode='Virtual', participants=['Synthetic staff'], hours='', period=PERIOD)
    sibling = ReportTable(('Requirement', 'Source year', 'Status'),
                         (('Annual drill', '2025', 'Completed'),), 'docx:other:item-0007')
    block = replace(draft.blocks[0], extra_tables=(*draft.blocks[0].extra_tables, sibling, matrix, events))
    app = open_app(draft, block)
    assert len(source_grids(app)) == 2
    edit(app, edited={0: {'Audience': ''}})
    result = current(app)
    assert result.extra_tables[0].rows[0][2] == ''
    assert result.extra_tables[1:] == (sibling, matrix, events)
    assert load_training(draft.profile.contract)['revision'] == 0
    assert path.read_bytes() == original


def test_same_source_occurrences_are_independent_and_literal_strings_survive(imported_training):
    draft, _, _ = imported_training
    source = draft.blocks[0].extra_tables[0]
    block = replace(draft.blocks[0], extra_tables=(source, source))
    app = open_app(draft, block)
    assert source_grids(app)[0].key != source_grids(app)[1].key
    edit(app, index=1, edited={1: {'Source note': 'Line one\nStatus: Completed; [literal](https://invalid.example)'}})
    assert current(app).extra_tables[0] == source
    assert current(app).extra_tables[1].rows[1][-1] == 'Line one\nStatus: Completed; [literal](https://invalid.example)'
    wire(app)
    assert current(app).extra_tables[0] == source


@pytest.mark.parametrize('operation', ['delete-one', 'delete-all', 'append-blank', 'insert-blank'])
def test_source_row_changes_preserve_explicit_blanks_and_do_not_restore_rows(imported_training, operation):
    draft, _, _ = imported_training
    app = open_app(draft)
    blank = ('',) * len(COLUMNS)
    if operation == 'delete-one':
        edit(app, deleted=[1]); expected = (ROWS[0], ROWS[2])
    elif operation == 'delete-all':
        edit(app, deleted=[0, 1, 2]); expected = ()
    elif operation == 'append-blank':
        edit(app, added=[dict.fromkeys(COLUMNS, '')]); expected = (*ROWS, blank)
    else:
        wire(app, 0, action='Insert blank above', row=1,
             edited={0: {'Hours': '2.5'}})
        expected = ((*ROWS[0][:3], '2.5', ROWS[0][-1]), blank, *ROWS[1:])
    assert current(app).extra_tables[0].rows == expected
    wire(app)
    assert current(app).extra_tables[0].rows == expected
    assert current(app).extra_tables[0].reference == draft.blocks[0].extra_tables[0].reference
    assert all(table.reference != EVENT_REF for table in current(app).extra_tables)


def test_empty_source_frame_never_becomes_a_completion_or_an_implicit_row(imported_training):
    draft, _, _ = imported_training
    source = replace(draft.blocks[0].extra_tables[0], rows=())
    app = open_app(draft, replace(draft.blocks[0], extra_tables=(source,)))
    wire(app)
    assert current(app).extra_tables == (source,)
    edit(app, edited={0: {'Topic': 'Unfinished imported entry'}})
    assert current(app).extra_tables[0].rows == (('Unfinished imported entry', '', '', '', ''),)
    wire(app)
    assert len(current(app).extra_tables) == 1
    assert not any(t.reference == EVENT_REF for t in current(app).extra_tables)


def test_excluded_source_values_can_be_corrected_without_reincluding_them(imported_training):
    draft, _, _ = imported_training
    block = replace(draft.blocks[0], source='Omit')
    app = open_app(draft, block)
    edit(app, edited={0: {'Hours': '2'}})
    assert current(app).source == 'Omit'
    wire(app)
    assert current(app).source == 'Omit'
    button = next(b for b in app.button if b.label == 'Include in this report')
    button.click()
    wire(app)
    assert current(app).source == 'This month'
    assert current(app).extra_tables[0].rows[0][3] == '2'


@pytest.mark.parametrize('change', ['period', 'contract', 'site', 'incarnation', 'source', 'rows', 'aba'])
def test_old_source_edit_and_insert_do_not_mutate_a_new_scope(imported_training, change):
    draft, _, _ = imported_training
    app = open_app(draft)
    old = source_grids(app)[0].proto
    if change == 'period':
        app.session_state['draft'] = replace(draft, period=ReportPeriod(2026, 10))
    elif change == 'contract':
        app.session_state['draft'] = replace(draft, profile=replace(draft.profile, contract='Unrelated synthetic contract'))
    elif change == 'site':
        from app.monthly_report_model import Facility
        app.session_state['draft'] = replace(draft, profile=replace(draft.profile,
            facilities=(Facility('other', 'Other synthetic site'),)))
    elif change == 'incarnation':
        app.session_state['source_training_rich_text_epoch'] = 'new-editor-incarnation'
    else:
        source = current(app).extra_tables[0]
        replacement = replace(source, reference='docx:replacement:item-0004') if change == 'source' else replace(source, rows=ROWS[:1])
        app.session_state['blocks'] = {'training_summary': replace(current(app), extra_tables=(replacement,))}
        if change == 'aba':
            wire(app)
            app.session_state['blocks'] = {'training_summary': draft.blocks[0]}
    wire(app)
    assert source_grids(app)[0].proto.id != old.id
    expected = current(app)
    wire(app, proto=old, edited={0: {'Hours': '99'}}, action='Insert blank above', row=0)
    assert current(app) == expected
    assert load_training(app.session_state['draft'].profile.contract)['revision'] == 0


def test_pending_note_and_source_edit_survive_unrelated_matrix_read_failure(imported_training, monkeypatch):
    from app import monthly_report_training_ui as ui
    draft, _, _ = imported_training
    block = replace(draft.blocks[0], text='Retained note')
    app = open_app(draft, block)
    for key in list(app.session_state.filtered_state):
        if key.startswith('source_training_training_') and key.endswith('_saved'):
            del app.session_state[key]
    def fail(*args):
        raise OSError('Synthetic shared requirements unavailable')
    monkeypatch.setattr(ui, 'load_training', fail)
    next(w for w in app.text_area if w.label == 'Training notes').set_value('')
    edit(app, edited={0: {'Hours': '1.75'}})
    assert any('Synthetic shared requirements unavailable' in error.value for error in app.error)
    assert current(app).text == ''
    assert current(app).extra_tables[0].rows[0][3] == '1.75'
    wire(app)
    assert current(app).text == '' and current(app).extra_tables[0].rows[0][3] == '1.75'


@pytest.mark.parametrize('table', [
    ReportTable(('Duplicate', 'Duplicate'), (('a', 'b'),), 'source:ambiguous'),
    ReportTable(('Heading', 'Other'), (('a',),), 'source:ragged'),
    ReportTable((), (('retained value',),), 'source:no-headers'),
])
def test_ambiguous_source_shape_stays_visible_but_cannot_lose_cells_in_a_grid(imported_training, table):
    draft, _, _ = imported_training
    block = replace(draft.blocks[0], extra_tables=(table,))
    app = open_app(draft, block)
    assert not source_grids(app)
    assert any('ambiguous column structure' in warning.value for warning in app.warning)
    assert current(app) == block
    app.run()
    assert current(app) == block


def test_edits_survive_snapshot_reopen_without_changing_history_or_monthly_role(imported_training):
    from app.monthly_report_setup import new_month_draft
    draft, path, original = imported_training
    baseline = library.save_snapshot(draft, expected_revision=0, entered_editor=draft.prepared_by)
    app = open_app(draft)
    edit(app, edited={0: {'Hours': '3.25'}}, deleted=[1])
    corrected = replace(draft, blocks=(current(app),))
    saved = library.save_snapshot(corrected, expected_revision=baseline.revision,
                                  entered_editor=draft.prepared_by)
    reopened = library.load_snapshot(draft.profile.contract, draft.profile.key, PERIOD)
    assert reopened == saved
    new_app = open_app(reopened.draft)
    assert current(new_app) == corrected.blocks[0]
    assert library.load_snapshot_version(draft.profile.contract, draft.profile.key, PERIOD, 1) == baseline
    for month in (10, 11):
        # Unclassified imported schemas keep their existing conservative
        # rollover: a frame, not old events promoted to current completions.
        corrected = new_month_draft(corrected, ReportPeriod(2026, month))
        carried = next(b for b in corrected.blocks if b.key == 'training_summary')
        assert carried.extra_tables == (ReportTable(COLUMNS, (), ''),)
        carried_app = open_app(corrected, carried)
        wire(carried_app)
        assert current(carried_app) == carried
    assert path.read_bytes() == original
    assert load_training(draft.profile.contract)['revision'] == 0


def test_ordinary_guided_training_section_reaches_source_editor(imported_training):
    draft, _, _ = imported_training
    app = AppTest.from_string(APP.split("st.session_state['blocks'] = render_training(")[0] + '''
from app.monthly_report_guided import _render_ordered_section
from app.monthly_report_model import BlockSpec, SectionSpec
section = SectionSpec('training', '12', 'Training Summary', (BlockSpec('training_summary', 'rich_text'),))
_, st.session_state['blocks'] = _render_ordered_section(
    st.session_state['draft'], section, st.session_state['blocks'],
    {spec.key: spec for spec in section.blocks}, None, 'source_training', {}, field)
''')
    app.session_state['draft'] = draft
    app.session_state['blocks'] = {'training_summary': draft.blocks[0]}
    app.run()
    assert not app.exception and len(source_grids(app)) == 1
    edit(app, edited={0: {'Hours': '4.5'}})
    assert current(app).extra_tables[0].rows[0][3] == '4.5'


def test_uncommitted_parent_rerun_replays_only_the_same_source_revision(imported_training):
    draft, _, _ = imported_training
    app = open_app(draft)
    edit(app, edited={0: {'Hours': '2.25'}})
    expected = current(app)
    # Model interruption after child return and before parent block commit.
    app.session_state['blocks'] = {'training_summary': draft.blocks[0]}
    wire(app)
    assert current(app) == expected
    # A real competing table edit must not accept that old child receipt.
    table = replace(expected.extra_tables[0], rows=(('Competing row', '', '', '', ''),))
    newer = replace(expected, extra_tables=(table,))
    app.session_state['blocks'] = {'training_summary': newer}
    wire(app)
    assert current(app) == newer


def test_reserved_year_matrix_is_not_offered_as_an_imported_completion_table(imported_training):
    draft, _, _ = imported_training
    annual = ReportTable(('Team member', 'Site', 'Safety'),
                         (('Synthetic staff', draft.profile.facilities[0].title, 'Completed'),),
                         'training:matrix:v1:year:2027')
    block = replace(draft.blocks[0], extra_tables=(*draft.blocks[0].extra_tables, annual))
    app = open_app(draft, block)
    assert len(source_grids(app)) == 1
    edit(app, edited={0: {'Hours': '1.25'}})
    assert current(app).extra_tables[1] == annual


def test_sibling_session_addition_preserves_pending_source_edit_and_grid_identity(imported_training):
    draft, _, _ = imported_training
    app = open_app(draft)
    identity = source_grids(app)[0].proto.id
    next(w for w in app.text_input if w.label == 'Training topic').set_value('Explicit new session')
    next(w for w in app.date_input if w.label == 'Training date').set_value(date(2026, 9, 9))
    next(w for w in app.selectbox if w.label == 'Method').set_value('Virtual')
    next(w for w in app.multiselect if w.label == 'Participants').set_value(['Synthetic staff'])
    next(w for w in app.button if w.label == 'Add completed training').click()
    edit(app, edited={0: {'Hours': '2.75'}})
    assert source_grids(app)[0].proto.id == identity
    assert current(app).extra_tables[0].rows[0][3] == '2.75'
    events = next(t for t in current(app).extra_tables if t.reference == EVENT_REF)
    assert len(events.rows) == 1 and events.rows[0][0] == 'Explicit new session'
    wire(app)
    assert source_grids(app)[0].proto.id == identity
    assert current(app).extra_tables[-1] == events
