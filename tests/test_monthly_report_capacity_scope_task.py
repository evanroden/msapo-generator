"""Capacity scope belongs to the displayed table, not any contract-wide record."""
from dataclasses import replace
from io import BytesIO

import pytest
from streamlit.testing.v1 import AppTest
from streamlit.testing.v1.element_tree import Block

from app import monthly_report_capacity as capacity, monthly_report_capacity_ui as ui
from app.monthly_report_model import ResolvedBlock
from test_monthly_report_capacity import profile as capacity_profile, CSV, draft, saved
from test_monthly_report_capacity_ui import APP


@pytest.fixture
def profile(tmp_path, monkeypatch):
    return capacity_profile.__wrapped__(tmp_path, monkeypatch)


def _path(root, target, path=()):
    if root is target:
        return path
    if isinstance(root, Block):
        for index, child in root.children.items():
            result = _path(child, target, (*path, index))
            if result is not None:
                return result
    return None


def test_manual_override_is_named_even_when_contract_capacity_exists(profile, monkeypatch):
    standing = saved(profile)
    manual = ResolvedBlock("thermal_capacity", "This month", rows=(("NH", "Steam", "NR", "lb/hr"),))
    monkeypatch.setattr(ui.st, "file_uploader", lambda *a, **kw: None)
    app = AppTest.from_string(APP)
    app.session_state["test_draft"] = draft(profile, manual)
    app.run()
    assert not app.exception
    assert any(x.value == "Capacity · this report only" for x in app.caption)
    assert not any("Saved capacity for" in x.value for x in app.caption)
    assert app.session_state["test_block"] == manual
    app.run()
    assert capacity.load_capacity(profile.contract) == standing
    future = capacity.resolve_capacity(replace(draft(profile), period=replace(draft(profile).period, month=10)))
    assert future.blocks[0].extra_tables[0].rows[1][2] == "40000"


def test_reused_scope_matches_historical_table_and_changes_after_override(profile, monkeypatch):
    standing = saved(profile)
    pinned = capacity.capacity_block(profile, standing)
    saved(profile, CSV.replace(b"1200", b"1300"), revision=1)
    monkeypatch.setattr(ui.st, "file_uploader", lambda *a, **kw: None)
    app = AppTest.from_string(APP)
    app.session_state["test_draft"] = draft(profile, pinned)
    app.run()
    assert not app.exception
    assert any(x.value == "Capacity · reused for these sites" for x in app.caption)
    assert app.session_state["test_block"] == pinned
    changed_table = replace(pinned.extra_tables[0], rows=(
        (*pinned.extra_tables[0].rows[0][:2], "NR", *pinned.extra_tables[0].rows[0][3:]),
        *pinned.extra_tables[0].rows[1:]))
    changed = replace(pinned, extra_tables=(changed_table,))
    app.session_state["test_draft"] = draft(profile, changed)
    app.run()
    assert any(x.value == "Capacity · this report only" for x in app.caption)
    assert capacity.load_capacity(profile.contract).revision == 2
    assert app.session_state["test_block"] == changed


def test_first_fill_keeps_one_native_uploader_in_the_same_subtree(profile, monkeypatch):
    upload = BytesIO(CSV)
    upload.name = "capacity.csv"
    native = ui.st.file_uploader
    def selected(*args, **kwargs):
        native(*args, **kwargs)
        return upload
    monkeypatch.setattr(ui.st, "file_uploader", selected)
    app = AppTest.from_string(APP)
    app.session_state["test_draft"] = draft(profile)
    app.run()
    assert not app.exception
    first, = app.get("file_uploader")
    before_path, before_id = _path(app._tree, first), first.id
    assert capacity.load_capacity(profile.contract).revision == 1
    app.run()
    assert not app.exception
    second, = app.get("file_uploader")
    assert second.id == before_id
    assert _path(app._tree, second) == before_path
    assert capacity.load_capacity(profile.contract).revision == 1


TASK_APP = '''
from dataclasses import replace
from types import SimpleNamespace
import streamlit as st
from app.monthly_report_guided import _render_ordered_section, _refresh_period_editors_after_commit
from app.monthly_report_model import default_sections
from app.monthly_report_ui import _field
current = st.session_state['test_draft']
section = next(section for section in default_sections() if section.key == 'scorecards')
section = replace(section, blocks=tuple(spec for spec in section.blocks if spec.key == 'thermal_capacity'))
current = replace(current, sections=(section,))
current, blocks = _render_ordered_section(current, section, {b.key: b for b in current.blocks},
    {s.key:s for s in section.blocks}, SimpleNamespace(profile=current.profile), 'test', {}, _field)
st.session_state['test_draft'] = replace(current, blocks=tuple(blocks.values()))
_refresh_period_editors_after_commit('test')
'''


def test_direct_capacity_table_precedes_its_one_source_task(profile, monkeypatch):
    pinned = capacity.capacity_block(profile, saved(profile))
    app = AppTest.from_string(TASK_APP)
    app.session_state["test_draft"] = draft(profile, pinned)
    app.run()
    assert not app.exception
    grids = [w for w in app.dataframe if w.key.startswith("test_edit_thermal_capacity_")]
    assert len(grids) == 1
    source, = [w for w in app.get("file_uploader") if w.label == "Thermal capacity file"]
    assert _path(app._tree, grids[0]) < _path(app._tree, source)
    scope = next(x for x in app.caption if x.value == "Capacity · reused for these sites")
    inventory = next(x for x in app.caption
                     if x.value.startswith("Reusable capacity ·"))
    assert "2 saved rows for these sites" in inventory.value
    assert _path(app._tree, scope) < _path(app._tree, grids[0])
    assert _path(app._tree, inventory) < _path(app._tree, grids[0])


def test_same_event_report_override_and_shared_correction_keep_both_scopes(profile, monkeypatch):
    from dataclasses import asdict
    from app.monthly_report_library import draft_from_dict
    standing = saved(profile)
    pinned = capacity.capacity_block(profile, standing)
    upload = BytesIO(CSV.replace(b"1200", b"1350"))
    upload.name = "corrected-capacity.csv"
    native = ui.st.file_uploader
    def selected(label, *args, **kwargs):
        result = native(label, *args, **kwargs)
        return upload if label == "Thermal capacity file" else result
    monkeypatch.setattr(ui.st, "file_uploader", selected)
    app = AppTest.from_string(TASK_APP)
    app.session_state["test_draft"] = draft(profile, pinned)
    app.run()
    assert not app.exception
    app.toggle("test_capacity_change").set_value(True).run()
    report_grid, = [w for w in app.dataframe if w.key.startswith("test_edit_thermal_capacity_")]
    # Native data-editor packet and native shared Apply travel in one rerun.
    app.session_state[report_grid.key] = {
        "edited_rows": {0: {"Required capacity": "NR"}}, "added_rows": [], "deleted_rows": []}
    apply, = [b for b in app.button if b.label == "Apply reviewed capacity for 2 sites"]
    apply.click().run()
    assert not app.exception
    current = app.session_state["test_draft"]
    block, = current.blocks
    assert block.extra_tables[0].rows[0][2] == "NR"
    updated = capacity.load_capacity(profile.contract)
    assert updated.revision == 2
    assert next(row for table in updated.tables for row in table.rows if row[:2] == ("NH", "Chilled water"))[2] == "1350"
    assert capacity.load_capacity(profile.contract, 1) == standing
    reloaded = draft_from_dict(asdict(current))
    assert reloaded.blocks == current.blocks
    next_month = capacity.resolve_capacity(replace(draft(profile), sections=current.sections,
        period=replace(current.period, month=10)))
    assert next_month.blocks[0].extra_tables[0].rows[0][2] == "1350"
    from docx import Document
    from app.monthly_report_docx import assemble_docx
    def output_rows(report):
        document = Document(BytesIO(assemble_docx(report, acknowledged_fingerprint=report.fingerprint)))
        return [tuple(cell.text for cell in row.cells) for table in document.tables for row in table.rows]
    assert ("NH", "Chilled water", "NR", "1600", "tons") in output_rows(reloaded)
    assert ("NH", "Chilled water", "1350", "1600", "tons") in output_rows(next_month)
    app.run()
    assert app.session_state["test_draft"].blocks == current.blocks
    assert capacity.load_capacity(profile.contract).revision == 2


def test_first_fill_keeps_existing_correction_reachable_in_same_file(profile, monkeypatch):
    original = CSV.splitlines(keepends=True)[0] + CSV.splitlines(keepends=True)[1]
    saved(profile, original)
    upload = BytesIO(CSV.replace(b"1200", b"1350"))
    upload.name = "capacity-with-new-services.csv"
    monkeypatch.setattr(ui.st, "file_uploader", lambda *a, **kw: upload)
    app = AppTest.from_string(APP)
    app.session_state["test_draft"] = draft(profile)
    app.run()
    assert not app.exception
    first = capacity.load_capacity(profile.contract)
    assert first.revision == 2
    rows = [row for table in first.tables for row in table.rows]
    assert len(rows) == 4
    assert next(row for row in rows if row[:2] == ("NH", "Chilled water"))[2] == "1200"
    app.run()
    app.toggle("test_capacity_change").set_value(True).run()
    apply = [b for b in app.button if b.label == "Apply reviewed capacity for 2 sites"]
    assert len(apply) == 1 and not apply[0].disabled
    apply[0].click().run()
    assert not app.exception
    latest = capacity.load_capacity(profile.contract)
    assert latest.revision == 3
    rows = [row for table in latest.tables for row in table.rows]
    assert len(rows) == 4
    assert next(row for row in rows if row[:2] == ("NH", "Chilled water"))[2] == "1350"
    assert capacity.load_capacity(profile.contract, 2) == first
    app.run()
    assert capacity.load_capacity(profile.contract) == latest


def test_ordered_first_fill_refreshes_the_single_report_table_after_commit(profile, monkeypatch):
    upload = BytesIO(CSV)
    upload.name = "first-capacity.csv"
    native = ui.st.file_uploader
    def selected(label, *args, **kwargs):
        result = native(label, *args, **kwargs)
        return upload if label == "Thermal capacity file" else result
    monkeypatch.setattr(ui.st, "file_uploader", selected)
    app = AppTest.from_string(TASK_APP)
    app.session_state["test_draft"] = draft(profile)
    app.run()
    assert not app.exception
    current = app.session_state["test_draft"]
    block, = current.blocks
    assert block == capacity.capacity_block(profile)
    report_grid, = [w for w in app.dataframe if w.key.startswith("test_edit_thermal_capacity_")]
    assert report_grid.value.iloc[0]["Required capacity"] == "1200"
    assert any("2 saved rows for these sites" in c.value and
               "version 1" in c.value for c in app.caption), (
        "The production ordered section must show the acknowledged first-fill "
        "capacity on its first settled run, without an extra rerun.")
    picker, = [w for w in app.get("file_uploader") if w.label == "Thermal capacity file"]
    assert _path(app._tree, report_grid) < _path(app._tree, picker)
    app.run()
    assert app.session_state["test_draft"].blocks == current.blocks
    assert capacity.load_capacity(profile.contract).revision == 1


def test_saved_capacity_summary_is_visible_for_selected_sites_only(profile, monkeypatch):
    """The visible inventory must not imply an automatic manual/report update."""
    from app.monthly_report_model import Facility
    monkeypatch.setattr(ui.st, "file_uploader", lambda *args, **kwargs: None)
    state = saved(profile)
    before = draft(profile)
    app = AppTest.from_string(APP)
    app.session_state["test_draft"] = before
    app.run()
    assert not app.exception
    lines = [c.value for c in app.caption if c.value.startswith("Reusable capacity ·")]
    assert len(lines) == 1
    line = lines[0]
    assert "2 saved rows for these sites" in line
    assert f"version {state.revision}" in line
    assert ui.literal(state.updated_at.split("T", 1)[0]) in line
    assert "capacity" in line and "source:" in line
    assert app.session_state["test_draft"] == before
    assert capacity.load_capacity(profile.contract) == state

    # A different site's saved rows are not evidence for this report. Do not
    # disclose a foreign site's source filename or count as reusable here.
    unknown = replace(profile, key="elsewhere", title="Elsewhere",
                      facilities=(Facility("elsewhere", "Elsewhere"),))
    app.session_state["test_draft"] = draft(unknown)
    app.run()
    assert not app.exception
    lines = [c.value for c in app.caption if c.value.startswith("Reusable capacity ·")]
    assert lines == ["Reusable capacity · none saved for these sites"]
    assert capacity.load_capacity(profile.contract) == state


def test_capacity_summary_empty_and_manual_scope_remain_truthful(profile, monkeypatch):
    monkeypatch.setattr(ui.st, "file_uploader", lambda *args, **kwargs: None)
    app = AppTest.from_string(APP)
    starting = draft(profile)
    app.session_state["test_draft"] = starting
    app.run()
    assert not app.exception
    assert any(c.value == "Reusable capacity · none saved for this contract" for c in app.caption)
    assert app.session_state["test_draft"] == starting
    assert capacity.load_capacity(profile.contract) is None

    state = saved(profile)
    manual = ResolvedBlock("thermal_capacity", "This month",
        text="Independent report note", rows=(("NH", "Steam", "NR", "tons"),))
    unchanged = draft(profile, manual)
    app.session_state["test_draft"] = unchanged
    app.run()
    assert not app.exception
    assert any(c.value == "Capacity · this report only" for c in app.caption)
    assert any("2 saved rows for these sites" in c.value for c in app.caption)
    assert app.session_state["test_draft"] == unchanged
    assert capacity.load_capacity(profile.contract) == state


def test_first_fill_inventory_updates_on_first_settled_render(profile, monkeypatch):
    from test_monthly_report_capacity_ui import app_for

    # This first AppTest run actually imports and commits the synthetic CSV.
    # No extra rerun, Save, or shared overwrite is permitted to see its status.
    app = app_for(monkeypatch, profile)
    assert not app.exception
    state = capacity.load_capacity(profile.contract)
    assert state.revision == 1
    assert len([c for c in app.caption if "2 saved rows for these sites" in c.value]) == 1
    assert any(f"version {state.revision}" in c.value and
               c.value.startswith("Reusable capacity ·") for c in app.caption)
    assert app.session_state["test_block"] == capacity.capacity_block(profile, state)
    app.run()
    assert capacity.load_capacity(profile.contract) == state


def test_saved_capacity_subsets_with_same_source_reference_do_not_leak_other_file(profile):
    """A table SHA/ID is not unique to one retained site-specific subset."""
    from app.monthly_report_model import Facility

    original = saved(profile)
    table, = original.tables
    north = tuple(row for row, site in zip(table.rows, table.site_keys) if site == "north")
    south = tuple(row for row, site in zip(table.rows, table.site_keys) if site == "south")
    # Retained subsets can have the same table ID/byte hash but different
    # filenames. Both references therefore match the same combined key.
    north_table = replace(table, rows=north, site_keys=("north",) * len(north),
                          source_filename="north-only.csv")
    south_table = replace(table, rows=south, site_keys=("south",) * len(south),
                          source_filename="south-only.csv")
    assert (north_table.id, north_table.source_sha256) == (
        south_table.id, south_table.source_sha256)
    shared = replace(original, tables=(north_table, south_table))
    north_summary = ui._saved_capacity_caption(profile, shared)
    assert "2 saved rows" in north_summary
    assert ui.literal("north-only.csv") in north_summary
    assert "south-only" not in north_summary

    south_profile = replace(profile, key="south", title="South Hospital",
                            facilities=(Facility("south", "South Hospital", ("SH",)),))
    south_summary = ui._saved_capacity_caption(south_profile, shared)
    assert "2 saved rows" in south_summary
    assert ui.literal("south-only.csv") in south_summary
    assert "north-only" not in south_summary
    # The regression is read-only; persisted history is untouched.
    assert capacity.load_capacity(profile.contract) == original
