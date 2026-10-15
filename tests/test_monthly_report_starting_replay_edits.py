"""Repeated Starting imports must respect saved report-only edits."""
from dataclasses import replace
from io import BytesIO

import pytest
from streamlit.testing.v1 import AppTest

from app import monthly_report_library as library
from app.monthly_report_guided import _edited_rich_text
from app.monthly_report_saved_pictures_ui import remove_saved_picture
from tests.test_monthly_report_start_import import (
    PERIOD, report_file, start_app, save_start,
)


def item(draft, key):
    return next(b for b in draft.blocks if b.key == key)


def replace_item(draft, updated):
    return replace(draft, blocks=tuple(updated if b.key == updated.key else b for b in draft.blocks))


def repeat_saved_start():
    # New session reads the actual saved snapshot and versioned profile. It
    # deliberately does not inject the prior session's draft/widget state.
    return AppTest.from_string("""
from app import monthly_report_library as library
from app.monthly_report_section_ui import render_section_setup
from app.monthly_report_ui import _field
from app.monthly_report_model import synthetic_profiles, ReportPeriod
p=synthetic_profiles()[0]
state=library.load_profile(p.contract,p.key)
render_section_setup(p.contract,ReportPeriod(2026,9),"Synthetic Editor",_field,state=state)
""", default_timeout=30).run()


@pytest.mark.parametrize("edit", ["clear_text", "remove_picture", "remove_table_row"])
def test_saved_report_only_edit_survives_reload_and_same_source_reimport(monkeypatch, tmp_path, edit):
    source = report_file(tmp_path, activity=("September 2026 source work.",), photo=True, table=True)
    raw = source.read_bytes()
    initial = save_start(start_app(monkeypatch, tmp_path, source))
    standing_before = library.load_profile(initial.profile.contract, initial.profile.key)
    library.save_snapshot(initial, expected_revision=0, entered_editor="Synthetic Editor")
    if edit == "clear_text":
        editor = AppTest.from_string("""
import streamlit as st
from app.monthly_report_guided import _edit_content
from app.monthly_report_ui import _field
st.session_state["block"] = _edit_content(
    st.session_state["spec"], st.session_state["block"], None,
    "synthetic_replay_clear", {}, _field)
""")
        editor.session_state["block"] = item(initial, "activity_summary")
        editor.session_state["spec"] = next(b for s in initial.sections for b in s.blocks if b.key == "activity_summary")
        editor.run()
        next(b for b in editor.button if b.label == "Clear Activity for this report").click().run()
        assert not editor.exception
        changed = editor.session_state["block"]
        assert changed.text == ""
    elif edit == "remove_picture":
        changed = remove_saved_picture(item(initial, "improvements"), 0)
    else:
        old = item(initial, "work_orders")
        table = old.extra_tables[0]
        changed = replace(old, extra_tables=(replace(table, rows=table.rows[1:]),))
    edited = replace_item(initial, changed)
    library.save_snapshot(edited, expected_revision=1, entered_editor="Synthetic Editor")
    assert library.load_snapshot(initial.profile.contract, initial.profile.key, PERIOD).draft == edited

    replay = repeat_saved_start()
    assert not replay.exception
    next(b for b in replay.button if b.label == "Continue to this month’s updates").click().run()
    assert not replay.exception
    assert not replay.error, [e.value for e in replay.error]
    actual = library.load_imported_draft(initial.profile.contract, initial.profile.key)
    assert item(actual, changed.key) == changed, "Reimport resurrected removed report content."
    assert source.read_bytes() == raw
    assert library.imported_original(initial.profile.contract, initial.profile.key).read_bytes() == raw
    standing_after = library.load_profile(initial.profile.contract, initial.profile.key)
    assert standing_after.current == standing_before.current
    assert standing_after.versions == standing_before.versions


def test_reimport_after_another_source_preserves_acknowledged_omission(monkeypatch, tmp_path):
    from docx import Document
    first = report_file(tmp_path, activity=("September 2026 original wording.",))
    first_raw = first.read_bytes()
    initial = save_start(start_app(monkeypatch, tmp_path, first))
    edited = replace_item(initial, _edited_rich_text(item(initial, "activity_summary"), ""))
    library.save_snapshot(edited, expected_revision=0, entered_editor="Synthetic Editor")

    doc = Document()
    doc.add_heading(initial.profile.facilities[0].title, 0)
    doc.add_paragraph("September 2026")
    doc.add_heading("11 Training Summary", 1)
    doc.add_paragraph("September 2026 client staff training pending.")
    second = tmp_path / "different-source.docx"
    doc.save(second)
    second_raw = second.read_bytes()
    from app import monthly_report_section_ui as ui
    real = ui.st.file_uploader
    current = {"raw": second_raw}
    def upload(label, *args, **kwargs):
        if label != "Starting report":
            return real(label, *args, **kwargs)
        result = BytesIO(current["raw"])
        result.name = "same-filename.docx"
        result.size = len(result.getvalue())
        return result
    monkeypatch.setattr(ui.st, "file_uploader", upload)
    added = save_start(repeat_saved_start())
    assert item(added, "activity_summary").text == ""
    assert "training pending" in item(added, "training_summary").text
    current["raw"] = first_raw
    replayed = save_start(repeat_saved_start())
    assert item(replayed, "activity_summary").text == ""
    assert item(replayed, "training_summary") == item(added, "training_summary")
    assert library.imported_original(initial.profile.contract, initial.profile.key).read_bytes() == first_raw


@pytest.mark.parametrize("stale", [False, True])
def test_unsaved_current_draft_clear_is_preserved_and_stale_revision_still_blocks(monkeypatch, tmp_path, stale):
    source = report_file(tmp_path, activity=("September 2026 source work.",))
    initial = save_start(start_app(monkeypatch, tmp_path, source))
    edited = replace_item(initial, _edited_rich_text(item(initial, "activity_summary"), ""))
    if stale:
        library.save_snapshot(initial, expected_revision=0, entered_editor="Synthetic Editor")
    before = library.load_profile(initial.profile.contract, initial.profile.key)
    app = AppTest.from_string("""
import streamlit as st
from app import monthly_report_library as library
from app.monthly_report_section_ui import render_section_setup
from app.monthly_report_ui import _field
draft=st.session_state["working"]
state=library.load_profile(draft.profile.contract,draft.profile.key)
render_section_setup(draft.profile.contract,draft.period,"Synthetic Editor",_field,
    state=state,working_draft=draft,working_revision=0)
""", default_timeout=30)
    app.session_state["working"] = edited
    app.run()
    next(b for b in app.button if b.label == "Continue to this month’s updates").click().run()
    assert not app.exception
    if stale:
        assert any("Another version was saved" in e.value for e in app.error)
        assert library.load_profile(initial.profile.contract, initial.profile.key).revision == before.revision
        assert library.load_imported_draft(initial.profile.contract, initial.profile.key) == initial
    else:
        assert not app.error
        actual = library.load_imported_draft(initial.profile.contract, initial.profile.key)
        assert item(actual, "activity_summary") == item(edited, "activity_summary")
    assert app.session_state["working"] == edited
