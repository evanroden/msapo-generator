"""Synthetic raw-unit closure: no original report XML, artwork or hashes."""
from copy import deepcopy
from io import BytesIO
from zipfile import ZipFile

from docx import Document
from docx.shared import Inches
from lxml import etree as E
import pytest

from app import monthly_report_compatibility_unit as unit
from app.monthly_report_compatibility_view import CompatibilityError
from app.monthly_report_import import _relations
from test_monthly_report_source_frame_logo_role import source, raster
from test_monthly_report_static_siblings import modern

W, A, WP, WPS, PIC, MC, R = unit.W, unit.A, unit.WP, unit.WPS, unit.PIC, unit.MC, unit.R


def fixture(tmp_path, *, picture=False, anchor=False, section=False):
    path = source(tmp_path, shared_outside=False)
    document = Document(path)
    paragraph = document.element.body[0]
    if picture:
        photo = document.add_picture(BytesIO(raster('orange')), width=Inches(2))
        drawing = photo._inline.getparent()
        old = drawing.getparent().getparent()
        paragraph.append(drawing.getparent())
        old.getparent().remove(old)
        geometry = next(drawing.iter(A+'prstGeom'))
        if not len(geometry): E.SubElement(geometry, A+'avLst')
        if anchor:
            container = drawing[0]
            container.tag = WP+'anchor'
            container.attrib.update(dict(simplePos='0',relativeHeight='2',behindDoc='0',locked='0',layoutInCell='1',allowOverlap='1'))
            for index, child in enumerate((E.Element(WP+'simplePos',x='0',y='0'), E.Element(WP+'positionH',relativeFrom='margin'), E.Element(WP+'positionV',relativeFrom='paragraph'))):
                container.insert(index, child)
                if index: E.SubElement(child,WP+'posOffset').text='1200'
            container.insert(4,E.Element(WP+'effectExtent',l='0',t='0',r='0',b='0'))
            container.insert(5,E.Element(WP+'wrapNone'))
            blip = next(drawing.iter(A+'blip'))
            extension = E.SubElement(E.SubElement(blip,A+'extLst'),A+'ext',uri=unit.DPI_URI)
            E.SubElement(extension,unit.A14+'useLocalDpi',val='0')
            fill = blip.getparent(); fill.insert(1,E.Element(A+'srcRect',t='100',b='100'))
            next(drawing.iter(WP+'docPr')).set('descr','Synthetic artwork description')
    if section:
        # Obtain a genuine python-docx-generated local header relation; all
        # content and attributes remain independently synthetic.
        document.sections[0].header.paragraphs[0].text='Synthetic header'
        props = E.Element(W+'pPr');paragraph.insert(0,props)
        runprops = E.SubElement(props,W+'rPr')
        E.SubElement(runprops,W+'rFonts',{W+'ascii':'Arial',W+'hAnsi':'Arial'})
        E.SubElement(runprops,W+'sz',{W+'val':'2'})
        sectionprops = deepcopy(document.sections[0]._sectPr)
        for child in list(sectionprops):
            if child.tag not in {W+'headerReference',W+'pgSz',W+'pgMar',W+'cols',W+'docGrid'}:sectionprops.remove(child)
        props.append(sectionprops)
    document.save(path)
    archive = ZipFile(path)
    tree = E.fromstring(archive.read('word/document.xml'))
    return archive, tree.find(W+'body')[0], _relations(archive,'word/document.xml')


def validate(data):
    archive, paragraph, relations = data
    return unit.validate_unit(paragraph, relations, archive)


def ordinary(paragraph):
    return next(d for r in paragraph if r.tag==W+'r' for d in r if d.tag==W+'drawing' and list(d.iter(PIC+'pic')))


@pytest.mark.parametrize('picture,anchor,section',[(False,False,False),(True,False,False),(True,True,True)])
def test_closed_outer_unit_is_nonmutating_and_audits_exact_drawings(tmp_path,picture,anchor,section):
    data = fixture(tmp_path,picture=picture,anchor=anchor,section=section)
    before = E.tostring(data[1]); receipt = validate(data)
    assert E.tostring(data[1])==before
    assert receipt['wrapper_count']==1
    assert len(receipt['ordinary_drawings'])==3+int(picture)
    assert all(len(x['sha256'])==64 for x in receipt['ordinary_drawings'])
    assert len(receipt['chrome_references'])==int(section)
    if picture:
        audit = receipt['ordinary_drawings'][-1]
        assert audit['class']=='ordinary-picture' and audit['media_size']>0
        assert audit['media_part'].startswith('word/media/')
        assert 'role' not in audit
    assert validate(data)==receipt


@pytest.mark.parametrize('kind',('horizontal-textframe','vertical-outlined-textframe','closed-5point-polygon'))
def test_unwrapped_static_anchor_uses_same_finite_validator(tmp_path,kind):
    data=fixture(tmp_path)
    E.SubElement(data[1],W+'r').append(deepcopy(modern(kind)[0]))
    assert validate(data)['ordinary_drawings'][-1]['class']=='unwrapped-'+kind
    next(data[1][-1].iter(WPS+'spPr')).set('unknown','payload')
    with pytest.raises(CompatibilityError):validate(data)


@pytest.mark.parametrize('tag',sorted(unit.FORBIDDEN))
@pytest.mark.parametrize('inside_wrapper',[False,True])
def test_raw_revision_command_comment_payload_rejected_even_if_view_would_erase(tmp_path,tag,inside_wrapper):
    data=fixture(tmp_path)
    target=next(data[1].iter(MC+'Fallback')) if inside_wrapper else data[1]
    E.SubElement(target,tag)
    with pytest.raises(CompatibilityError,match='revision, field or comment'):validate(data)


@pytest.mark.parametrize('kind',('outer_tag','outer_attr','run_attr','ppr_duplicate','rpr_duplicate','font_url','hidden','outer_text','unknown_frame_attr','unknown_frame_layer','frame_duplicate_body','frame_relationship','frame_nonrect','tail','comment'))
def test_closed_outer_grammar_rejects_unknown_payload(tmp_path,kind):
    data=fixture(tmp_path); p=data[1]; r=p.find(W+'r'); drawing=r.find(W+'drawing')
    if kind=='outer_tag':E.SubElement(p,'{urn:unrecognized}payload')
    elif kind=='outer_attr':p.set('unknown','payload')
    elif kind=='run_attr':r.set('onclick','command')
    elif kind=='ppr_duplicate':p.insert(0,E.Element(W+'pPr'));p.insert(0,E.Element(W+'pPr'))
    elif kind=='rpr_duplicate':r.insert(0,E.Element(W+'rPr'));r.insert(0,E.Element(W+'rPr'))
    elif kind=='font_url':E.SubElement(E.SubElement(r,W+'rPr'),W+'rFonts',{W+'ascii':'https://untrusted.invalid/font'})
    elif kind=='hidden':r.insert(0,E.Element(W+'rPr'));E.SubElement(r[0],W+'vanish')
    elif kind=='outer_text':E.SubElement(r,W+'t').text='Unclassified text'
    elif kind=='unknown_frame_attr':next(drawing.iter(WPS+'wsp')).set('mystery','payload')
    elif kind=='unknown_frame_layer':E.SubElement(next(drawing.iter(WPS+'spPr')),A+'effectLst')
    elif kind=='frame_duplicate_body':next(drawing.iter(WPS+'wsp')).append(E.Element(WPS+'bodyPr'))
    elif kind=='frame_relationship':next(drawing.iter(WPS+'bodyPr')).set(R+'id','rId1')
    elif kind=='frame_nonrect':next(drawing.iter(A+'prstGeom')).set('prst','ellipse')
    elif kind=='tail':drawing.tail='hidden text'
    elif kind=='comment':p.append(E.Comment('opaque comment'))
    with pytest.raises(CompatibilityError):validate(data)


@pytest.mark.parametrize('kind',('extra_picture','extra_layer','linked','missing','wrong_uri','external','extension_uri','extension_layer','duplicate_extension','crop','negative_extent','extent_child','action','descriptor_child','duplicate_geometry','picture_unknown_attr'))
def test_ordinary_picture_closed_geometry_and_local_resource_only(tmp_path,kind):
    data=fixture(tmp_path,picture=True,anchor=True); p=data[1]; d=ordinary(p); blip=next(d.iter(A+'blip'))
    rid=blip.get(R+'embed'); target,relkind=data[2][rid]
    if kind=='extra_picture':next(d.iter(A+'graphicData')).append(deepcopy(next(d.iter(PIC+'pic'))))
    elif kind=='extra_layer':E.SubElement(next(d.iter(PIC+'spPr')),A+'effectLst')
    elif kind=='linked':blip.set(R+'link',blip.attrib.pop(R+'embed'))
    elif kind=='missing':blip.set(R+'embed','absent')
    elif kind=='wrong_uri':data[2][rid]=(target,'https://unrecognized.invalid/image')
    elif kind=='external':data[2][rid]=('',relkind)
    elif kind=='extension_uri':next(d.iter(A+'extLst'))[0].set('uri','unrecognized')
    elif kind=='extension_layer':E.SubElement(next(d.iter(A+'extLst'))[0],unit.A14+'useLocalDpi',val='0')
    elif kind=='duplicate_extension':blip.append(deepcopy(next(d.iter(A+'extLst'))))
    elif kind=='crop':next(d.iter(A+'srcRect')).set('t','100000')
    elif kind=='negative_extent':next(d.iter(WP+'extent')).set('cx','-1')
    elif kind=='extent_child':E.SubElement(next(d.iter(A+'xfrm')).find(A+'ext'),unit.A14+'useLocalDpi',val='0')
    elif kind=='action':next(d.iter(WP+'docPr')).set('action','launch')
    elif kind=='descriptor_child':E.SubElement(next(d.iter(WP+'docPr')),A+'hlinkClick',{R+'id':rid})
    elif kind=='duplicate_geometry':next(d.iter(PIC+'spPr')).append(deepcopy(next(d.iter(A+'prstGeom'))))
    elif kind=='picture_unknown_attr':next(d.iter(PIC+'pic')).set('mystery','payload')
    with pytest.raises(CompatibilityError):validate(data)


@pytest.mark.parametrize('kind',('duplicate_type','wrong_uri','missing','external','unknown_attr','foreign_part','invalid_margin','reordered'))
def test_section_references_and_properties_are_closed(tmp_path,kind):
    data=fixture(tmp_path,section=True); sect=next(data[1].iter(W+'sectPr'));ref=sect.find(W+'headerReference');rid=ref.get(R+'id');target,relkind=data[2][rid]
    if kind=='duplicate_type':sect.insert(0,deepcopy(ref))
    elif kind=='wrong_uri':data[2][rid]=(target,'https://unrecognized.invalid/header')
    elif kind=='missing':data[2][rid]=('word/header999.xml',relkind)
    elif kind=='external':data[2][rid]=('',relkind)
    elif kind=='unknown_attr':sect.set('unknown','payload')
    elif kind=='foreign_part':data[2][rid]=('word/document.xml',relkind)
    elif kind=='invalid_margin':sect.find(W+'pgMar').set(W+'left','-1')
    elif kind=='reordered':sect.append(ref)
    with pytest.raises(CompatibilityError):validate(data)


@pytest.mark.parametrize('budget',('MAX_NODES','MAX_DEPTH','MAX_CHARS'))
def test_raw_budget_refusal_precedes_any_image_decode(tmp_path,monkeypatch,budget):
    data=fixture(tmp_path,picture=True)
    monkeypatch.setattr(unit,budget,1)
    def forbidden(*args):raise AssertionError('resource validation must not run')
    monkeypatch.setattr('app.monthly_report_compatibility_view._validate_picture_resource',forbidden)
    with pytest.raises(CompatibilityError,match='budget'):validate(data)


def test_wrapper_unknown_payload_is_not_claimed_valid_by_outer_only_api(tmp_path):
    # The outer closure deliberately leaves branch structure to _selection and
    # _branch/static validators. Its receipt is not a standalone capability.
    data=fixture(tmp_path)
    next(data[1].iter(MC+'Fallback')).set('unknown','payload')
    assert validate(data)['wrapper_count']==1
    from app.monthly_report_compatibility_view import _selection
    wrapper=next(data[1].iter(MC+'AlternateContent'))
    # _selection itself checks wrapper ordering, while branch validators check
    # the Fallback attributes; the test documents the explicit split contract.
    assert _selection(wrapper)[0].tag==MC+'Choice'
    from app.monthly_report_compatibility_view import _branch
    with pytest.raises(CompatibilityError):
        _branch(wrapper[1],data[2],data[0])


@pytest.mark.parametrize('kind',('deleted_command','deleted_text','inserted_text','unknown_tag','unknown_attr','unwrapped_unknown'))
def test_derivation_closes_the_raw_unit_before_current_view(tmp_path,kind):
    from app.monthly_report_compatibility_view import derive_compatibility_view
    path=source(tmp_path,divergent=True,shared_outside=False)
    document=Document(path); paragraph=document.element.body[0]
    if kind in ('deleted_command','deleted_text','inserted_text'):
        change=E.SubElement(paragraph,W+('ins' if kind=='inserted_text' else 'del'),{W+'id':'1',W+'author':'Synthetic Editor',W+'date':'2026-01-01T00:00:00Z'})
        run=E.SubElement(change,W+'r')
        E.SubElement(run,W+('instrText' if kind=='deleted_command' else 't')).text='Never eligible payload'
    elif kind=='unknown_tag':E.SubElement(paragraph,'{urn:unknown}payload')
    elif kind=='unknown_attr':paragraph.set('mystery','payload')
    elif kind=='unwrapped_unknown':
        next(paragraph[0].iter(WPS+'spPr')).set('mystery','payload')
    document.save(path); original=path.read_bytes()
    with pytest.raises(CompatibilityError,match='Raw native unit'):
        derive_compatibility_view(original)
    assert path.read_bytes()==original
