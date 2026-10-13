"""Canonical issue cards preserve report-only omissions and actual history."""

from dataclasses import asdict, replace
from io import BytesIO
import json

from docx import Document
import pytest
from streamlit.testing.v1 import AppTest

from app import monthly_report_library as library
from app.monthly_report_docx import assemble_docx
from app.monthly_report_followups import problems
from app.monthly_report_issues import canonical_issues, clear_issues, issue_visibility, remove_issue_statement
from app.monthly_report_model import (
    DraftParagraph, ReportDraft, ReportFollowUp, ReportPeriod, ReportSource,
    ReportTable, ResolvedBlock, default_sections, synthetic_profiles,
)
from app.monthly_report_setup import merge_drafts, new_month_draft
from app.monthly_report_sources import source_reference


def draft_for(text="- Pump vibration remains under investigation.\n- AHU belt noise."):
    return ReportDraft(synthetic_profiles()[0], ReportPeriod(2026, 9), "Synthetic Operator",
        tuple(s for s in default_sections() if s.key == "issues"),
        (ResolvedBlock("equipment_issues", "This month", text=text),), synthetic=True)


def test_migration_deduplicates_cards_and_preserves_source_images_and_provenance():
    original = draft_for()
    block = replace(original.blocks[0], asset_hashes=("picture.png",),
                    asset_captions=("Pump inspection",),
                    asset_provenance=(("picture.png", ("source:1:digest",)),))
    original = replace(original, blocks=(block,))
    canonical = canonical_issues(original)
    assert len(canonical.follow_ups) == 2
    assert canonical.blocks[0].text == ""
    assert canonical.blocks[0].asset_hashes == block.asset_hashes
    assert canonical.blocks[0].asset_captions == block.asset_captions
    assert canonical.blocks[0].asset_provenance == block.asset_provenance
    assert canonical_issues(canonical) == canonical
    assert original.blocks[0].text  # The original report was not changed.


def test_manual_correction_and_omission_survive_same_source_import_and_serialization():
    original = draft_for()
    canonical = canonical_issues(original)
    corrected = replace(canonical.follow_ups[0], text="Manual correction to the pump finding.")
    current = replace(canonical, follow_ups=(corrected, *canonical.follow_ups[1:]))
    current = issue_visibility(current, corrected.key, False)
    current = library.draft_from_dict(json.loads(json.dumps(asdict(current))))
    replay = canonical_issues(merge_drafts(current, original))
    assert len(replay.follow_ups) == 2
    assert replay.follow_ups[0].text == corrected.text
    assert not replay.follow_ups[0].included
    assert replay.follow_ups[0].status == "ongoing"
    assert replay.blocks[0].text == ""
    assert not problems(replay.follow_ups[0], replay)


def test_clear_list_survives_repeat_import_but_unresolved_items_return_next_month():
    original = draft_for()
    cleared = clear_issues(original)
    replay = canonical_issues(merge_drafts(cleared, original))
    assert len(replay.follow_ups) == 2
    assert not any(item.included for item in replay.follow_ups)
    assert all(item.status == "ongoing" for item in replay.follow_ups)
    next_month = new_month_draft(replay, ReportPeriod(2026, 10))
    assert all(item.included for item in next_month.follow_ups)
    assert [item.key for item in next_month.follow_ups] == [item.key for item in replay.follow_ups]
    assert not any(problems(item, next_month) for item in next_month.follow_ups)


def test_resolution_needs_evidence_and_does_not_reappear_as_unresolved_next_month():
    current = canonical_issues(draft_for("Pump vibration."))
    item = replace(current.follow_ups[0], status="resolved", included=False)
    assert any("evidence" in problem for problem in problems(item, current))
    item = replace(item, evidence_note="Inspection confirmed the repair and normal operation.")
    current = replace(current, follow_ups=(item,))
    assert not problems(item, current)
    assert not new_month_draft(current, ReportPeriod(2026, 10)).follow_ups
    assert current.follow_ups[0] == item


def test_unsubstantiated_resolution_cannot_silently_drop_issue_at_rollover():
    current = canonical_issues(draft_for("Pump vibration."))
    invalid = replace(current.follow_ups[0], status="resolved", included=False)
    current = replace(current, follow_ups=(invalid,))
    following = new_month_draft(current, ReportPeriod(2026, 10))
    assert following.follow_ups[0].key == invalid.key
    assert following.follow_ups[0].included
    assert any("evidence" in p for p in problems(following.follow_ups[0], following))


@pytest.mark.parametrize("statement", ["No current equipment issues.", "0 equipment issues."])
def test_explicit_no_issue_statement_stays_a_statement(statement):
    original = draft_for(statement)
    canonical = canonical_issues(original)
    assert not canonical.follow_ups
    assert canonical.blocks[0].text == statement
    assert not new_month_draft(canonical, ReportPeriod(2026, 10)).follow_ups
    removed = remove_issue_statement(canonical)
    replay = canonical_issues(merge_drafts(removed, original))
    assert not replay.follow_ups
    assert replay.blocks[0].text == ""


def test_issue_table_migration_keeps_nonprice_details_and_page_reference():
    source = ReportSource("synthetic-source", "inspection.pdf", "a" * 64, ".pdf",
                          page_texts=("Pump vibration requires investigation.",))
    reference = source_reference(source, 1)
    table = ReportTable(("Equipment", "Finding", "Cost"),
                        (("Pump", "Vibration", "$100"),), reference=reference)
    original = replace(draft_for(""), sources=(source,),
                       blocks=(ResolvedBlock("equipment_issues", "This month", extra_tables=(table,)),))
    canonical = canonical_issues(original)
    assert canonical.follow_ups[0].text == "Equipment: Pump; Finding: Vibration"
    assert canonical.follow_ups[0].references == (reference,)
    assert not canonical.blocks[0].extra_tables
    changed = replace(canonical, sources=(replace(source, findings="Source changed"),))
    assert any("source changed" in p for p in problems(canonical.follow_ups[0], changed))


def test_starting_docx_lineage_is_retained_without_pretending_it_is_new_resolution_evidence():
    original = draft_for("Pump vibration.")
    reference = "docx:" + "a" * 64 + ":paragraph-22"
    original = replace(original, blocks=(replace(original.blocks[0], references=(reference,)),))
    canonical = canonical_issues(original)
    assert canonical.blocks[0].references == (reference,)
    assert not canonical.follow_ups[0].references
    assert not problems(canonical.follow_ups[0], canonical)
    following = new_month_draft(canonical, ReportPeriod(2026, 10))
    assert following.blocks[0].references == (reference,)
    assert any("evidence" in p for p in problems(replace(following.follow_ups[0], status="resolved"), following))


def test_unreviewed_or_unsupported_source_wording_stays_in_existing_evidence_gate():
    original = draft_for("An unsupported 300 repairs.")
    block = replace(original.blocks[0], ai_written=True,
        ai_paragraphs=(DraftParagraph("An unsupported 300 repairs.", (), (), ("Unsupported number: 300",)),))
    original = replace(original, blocks=(block,))
    assert canonical_issues(original) == original
    assert not canonical_issues(original).follow_ups
    following = new_month_draft(original, ReportPeriod(2026, 10))
    assert not following.follow_ups
    assert following.blocks[0].ai_paragraphs == block.ai_paragraphs
    assert following.blocks[0].text == block.text


def test_export_contains_each_canonical_issue_once_and_omission_cannot_resurrect_prose():
    original = draft_for()
    canonical = canonical_issues(original)
    raw = assemble_docx(canonical, acknowledged_fingerprint=canonical.fingerprint)
    text = "\n".join(p.text for p in Document(BytesIO(raw)).paragraphs)
    assert text.count("Pump vibration remains under investigation.") == 1
    removed = issue_visibility(canonical, canonical.follow_ups[0].key, False)
    replay = canonical_issues(merge_drafts(removed, original))
    raw = assemble_docx(replay, acknowledged_fingerprint=replay.fingerprint)
    text = "\n".join(p.text for p in Document(BytesIO(raw)).paragraphs)
    assert "Pump vibration" not in text and text.count("AHU belt noise.") == 1


APP = '''
import streamlit as st
from app.monthly_report_issues import render_issue_cards
from app.monthly_report_ui import _field
draft = st.session_state["draft"]
st.session_state["draft"] = render_issue_cards(draft, "issue_probe", _field)
'''


def test_direct_card_edit_add_remove_clear_and_restore_have_one_canonical_list():
    app = AppTest.from_string(APP, default_timeout=15)
    app.session_state["draft"] = draft_for()
    app.run()
    assert not app.exception
    descriptions = [w for w in app.text_area if w.label == "Issue description"]
    assert len(descriptions) == 2
    descriptions[0].set_value("Corrected pump finding.").run()
    assert app.session_state["draft"].follow_ups[0].text == "Corrected pump finding."
    next(b for b in app.button if b.label == "Remove from this report").click().run()
    app.run()
    assert len([w for w in app.text_area if w.label == "Issue description"]) == 1
    assert not app.session_state["draft"].follow_ups[0].included
    next(b for b in app.button if b.label == "Include again").click().run()
    app.run()
    assert app.session_state["draft"].follow_ups[0].text == "Corrected pump finding."
    next(b for b in app.button if b.label == "Add issue").click().run()
    descriptions = [w for w in app.text_area if w.label == "Issue description"]
    descriptions[-1].set_value("New valve leak.").run()
    assert len(app.session_state["draft"].follow_ups) == 3
    next(b for b in app.button if b.label == "Clear issues from this report").click().run()
    app.run()
    assert not app.exception
    assert not any(item.included for item in app.session_state["draft"].follow_ups)
    assert all(item.status == "ongoing" for item in app.session_state["draft"].follow_ups)
    assert not any(w.label == "Issue description" for w in app.text_area)
    assert not any(b.label in ("Still open", "Still correct", "Confirm update") for b in app.button)


def test_guided_issue_omission_survives_shared_save_and_new_session_without_hidden_narrative(
        monkeypatch, tmp_path):
    from test_monthly_report_ui import ROOT, choose_report, monthly
    from app.contracts import RRH_CONTRACT
    app = monthly(monkeypatch, tmp_path)
    next(w for w in app.text_input if w.label == "Prepared by").set_value("Synthetic Editor").run()
    key, current = next((key, value) for key, value in app.session_state.filtered_state.items()
                        if key.startswith("report_guided_") and key.endswith("_draft")
                        and isinstance(value, ReportDraft))
    block = draft_for().blocks[0]
    app.session_state[key] = replace(current,
        blocks=(*tuple(b for b in current.blocks if b.key != "equipment_issues"), block))
    app.run()
    assert not app.exception
    assert len([w for w in app.text_area if w.label == "Issue description"]) == 2
    assert not any(w.label == "Equipment issues" for w in app.text_area)
    next(b for b in app.button if b.label == "Remove from this report").click().run()
    next(b for b in app.button if b.label == "Save shared version").click().run()
    assert not app.exception
    saved = library.load_snapshot(RRH_CONTRACT, current.profile.key, current.period)
    assert saved is not None and not saved.draft.follow_ups[0].included
    assert next(b for b in saved.draft.blocks if b.key == "equipment_issues").text == ""
    resumed = AppTest.from_file(ROOT / "run_web.py", default_timeout=20).run()
    resumed.segmented_control[0].set_value("Monthly report").run()
    choose_report(resumed, current.profile.facilities[0].title)
    assert not resumed.exception
    assert len([w for w in resumed.text_area if w.label == "Issue description"]) == 1
    assert next(w for w in resumed.text_area if w.label == "Issue description").value == "AHU belt noise."
    assert not any(w.label == "Equipment issues" for w in resumed.text_area)
