"""A master lends only safe table geometry, never an old client's facts."""
from dataclasses import replace
from io import BytesIO
from zipfile import ZipFile

import pytest
from docx import Document

from app.monthly_report_import import inspect_docx
from app.monthly_report_master_tables import from_inspection, initialize_master_tables
from app.monthly_report_model import (
    Facility, ReportDraft, ReportPeriod, ReportProfile, ReportTable,
    ResolvedBlock, profile_sections,
)
from app.monthly_report_native_layout import NativeLayoutError, build_native_docx


def master_source(tmp_path):
    path = tmp_path / "synthetic-original-master.docx"
    doc = Document()
    doc.add_paragraph("OLD PRIVATE CLIENT - NEVER FORWARD")
    doc.add_paragraph("August 2026")
    doc.add_heading("PENDING & DECLINED PROPOSALS", 1)
    for headings in (
        ("Project Number", "Status", "Price", "Technical Notes"),
        ("Facility", "Vendor", "Scope of Work", "Status", "Total Cost"),
    ):
        table = doc.add_table(rows=2, cols=len(headings))
        for i, label in enumerate(headings):
            table.cell(0, i).text = label
            table.cell(1, i).text = "OLD PRIVATE CLIENT - $77,777.77"
    doc.save(path)
    return path


def current(shape, *, first=True, second=True):
    profile = ReportProfile("Synthetic Independent Client", "test-two-table-master",
                            "QA North", (Facility("qa-n", "QA North"),), "individual")
    tables = (
        ReportTable(shape[0].columns,
                    (("Q-TEST-731", "Pending", "45 gpm"),) if first else (),
                    shape[0].reference),
        ReportTable(shape[1].columns,
                    (("QA North", "Synthetic Vendor", "Valve inspection", "Pending"),)
                    if second else (), shape[1].reference),
    )
    return ReportDraft(profile, ReportPeriod(2026, 10), "Synthetic QA Operator",
                       profile_sections(profile),
                       (ResolvedBlock("proposals", "This month", extra_tables=tables),),
                       synthetic=True)


def test_pinned_master_distinct_physical_tables_exclude_price_without_source_facts(tmp_path):
    path = master_source(tmp_path)
    shapes = from_inspection(inspect_docx(path))
    assert len(shapes) == 2
    assert [s.slot for s in shapes] == ["proposals", "proposals"]
    assert [s.physical_width for s in shapes] == [4, 5]
    assert [s.indices for s in shapes] == [(0, 1, 3), (0, 1, 2, 3)]
    assert [s.columns for s in shapes] == [
        ("Project Number", "Status", "Technical Notes"),
        ("Facility", "Vendor", "Scope of Work", "Status"),
    ]
    assert all("PRIVATE" not in h and "Cost" not in h and "Price" not in h
               for item in shapes for h in item.columns)

    docx = build_native_docx(current(shapes), path, master=True)
    output = Document(BytesIO(docx))
    assert [len(t.columns) for t in output.tables] == [3, 4]
    assert [[c.text for c in row.cells] for row in output.tables[0].rows] == [
        ["Project Number", "Status", "Technical Notes"],
        ["Q-TEST-731", "Pending", "45 gpm"],
    ]
    assert [[c.text for c in row.cells] for row in output.tables[1].rows] == [
        ["Facility", "Vendor", "Scope of Work", "Status"],
        ["QA North", "Synthetic Vendor", "Valve inspection", "Pending"],
    ]
    with ZipFile(BytesIO(docx)) as z:
        text = "\n".join(z.read(n).decode("utf8", "ignore") for n in z.namelist()
                         if n.endswith((".xml", ".rels", ".vml"))).casefold()
        assert "old private client" not in text
        assert "$77,777.77" not in text
        assert "total cost" not in text


def test_only_second_physical_table_is_used_when_first_is_empty(tmp_path):
    path = master_source(tmp_path)
    shapes = from_inspection(inspect_docx(path))
    docx = build_native_docx(current(shapes, first=False), path, master=True)
    tables = Document(BytesIO(docx)).tables
    assert len(tables) == 2  # Empty source regions keep their native header.
    assert len(tables[0].columns) == 3 and len(tables[0].rows) == 1
    assert len(tables[1].columns) == 4
    assert tables[1].cell(1, 0).text == "QA North"
    assert tables[1].cell(1, 2).text == "Valve inspection"
    assert "Q-TEST-731" not in " ".join(c.text for t in tables for r in t.rows for c in r.cells)


def test_rejected_forged_master_region_never_uses_an_ordinal_fallback(tmp_path):
    path = master_source(tmp_path)
    shapes = from_inspection(inspect_docx(path))
    draft = current(shapes)
    original = draft.blocks[0]
    forged = replace(original.extra_tables[1], reference="master-table:unreviewed:other-client")
    changed = replace(draft, blocks=(replace(original, extra_tables=(original.extra_tables[0], forged)),))
    with pytest.raises(NativeLayoutError, match="pinned master table reference"):
        build_native_docx(changed, path, master=True)


def test_seeding_adds_empty_unique_tables_once_and_does_not_override_work(monkeypatch, tmp_path):
    from app import monthly_report_master_tables as module
    shapes = from_inspection(inspect_docx(master_source(tmp_path)))
    monkeypatch.setattr(module, "for_profile", lambda profile: shapes)
    draft = current(shapes)
    blank = replace(draft, blocks=(ResolvedBlock("proposals", "This month"),))
    seeded = initialize_master_tables(blank)
    assert [x.reference for x in seeded.blocks[0].extra_tables] == [x.reference for x in shapes]
    assert [x.columns for x in seeded.blocks[0].extra_tables] == [x.columns for x in shapes]
    assert all(not x.rows for x in seeded.blocks[0].extra_tables)
    assert initialize_master_tables(seeded) == seeded
    assert initialize_master_tables(draft) == draft


def test_unknown_client_named_source_header_never_enters_new_table(tmp_path):
    path = tmp_path / "synthetic-name.docx"
    doc = Document()
    doc.add_heading("PENDING & DECLINED PROPOSALS", 1)
    table = doc.add_table(rows=2, cols=3)
    for i, value in enumerate(("Private Hospital", "Status", "Invoice Amount")):
        table.cell(0, i).text = value
        table.cell(1, i).text = "Some old name"
    doc.save(path)
    shapes = from_inspection(inspect_docx(path))
    assert shapes[0].columns == ("Detail 1", "Status")


def test_fresh_design_seed_uses_blank_native_regions_without_copying_source_rows(monkeypatch, tmp_path):
    from app import monthly_report_master_tables as module
    from app import monthly_report_designs as designs
    from app.monthly_report_start import design_seed
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path / "empty-test-library"))
    source = master_source(tmp_path)
    shapes = from_inspection(inspect_docx(source))
    monkeypatch.setattr(module, "for_profile", lambda _: shapes)
    monkeypatch.setattr(designs, "pin", lambda profile, **kwargs: profile)
    profile = ReportProfile("Synthetic Independent Client", "new-design",
                            "QA North", (Facility("qa-n", "QA North"),), "individual")
    seed, assets = design_seed(profile, ReportPeriod(2026, 10), "Synthetic Editor")
    proposals = next((b for b in seed.blocks if b.key == "proposals"), None)
    assert proposals is not None
    assert len(proposals.extra_tables) == 2
    assert all(not t.rows for t in proposals.extra_tables)
    assert [t.columns for t in proposals.extra_tables] == [t.columns for t in shapes]
    assert not any("OLD PRIVATE" in str(block) for block in seed.blocks)
    assert assets == {}
