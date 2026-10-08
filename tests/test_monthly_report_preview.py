from dataclasses import replace
from io import BytesIO

from docx import Document
import fitz
from PIL import Image
import pytest

from app import monthly_report_preview as preview
from app.monthly_report_docx import assemble_docx, estimate_bytes, generate_report
from app.monthly_report_model import (
    BlockSpec, ColumnSpec, ReportPeriod, ReportTable, ResolvedBlock,
    synthetic_draft, synthetic_profiles,
)
from tests.conftest import requires_libreoffice
from test_monthly_report_ui import monthly, step


def draft():
    base = synthetic_draft(synthetic_profiles()[0], ReportPeriod(2026, 9))
    return replace(base, sections=(replace(base.sections[1], blocks=(BlockSpec("activity_summary", "rich_text"),)),),
                   blocks=(ResolvedBlock("activity_summary", "This month", text="Synthetic completed maintenance."),))


def fake_converter(monkeypatch, tmp_path):
    outputs = []
    def convert(path):
        output = tmp_path / (path.stem + ".pdf")
        with fitz.open() as pdf:
            for n in range(2):
                pdf.new_page().insert_text((72, 72), "Synthetic page " + str(n + 1))
            pdf.save(output)
        outputs.append(output)
        return output
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path))
    monkeypatch.setattr(preview.pdf_converter, "convert_to_pdf", convert)
    return outputs


def test_unfinished_preview_is_watermarked_and_never_approves_final_output(monkeypatch, tmp_path):
    outputs = fake_converter(monkeypatch, tmp_path)
    unfinished = replace(draft(), prepared_by="", blocks=(replace(draft().blocks[0], ai_written=True),))
    result = preview.preview_report(unfinished)
    assert result.pages == 2 and result.period == "September 2026"
    with fitz.open(stream=result.pdf, filetype="pdf") as pdf:
        assert all("DRAFT - REVIEW ONLY" in p.get_text() and "NOT A FINISHED CLIENT REPORT" in p.get_text() for p in pdf)
    assert preview.preview_page(result, 0).startswith(b"\x89PNG")
    with pytest.raises(ValueError, match="existing preview page"):
        preview.preview_page(result, 2)
    with pytest.raises(ValueError, match="prepared|Review"):
        assemble_docx(unfinished)
    assert not any(p.exists() for p in outputs)
    assert not list((tmp_path / "monthly_reports" / "previews").glob("report-*"))


def test_preview_preserves_price_gates_bounds_and_cleans_up(monkeypatch, tmp_path):
    outputs = fake_converter(monkeypatch, tmp_path)
    priced = replace(draft(), blocks=(replace(draft().blocks[0], text="Price: $200"),))
    with pytest.raises(ValueError, match="pricing"):
        preview.preview_report(priced)
    assert not outputs
    monkeypatch.setattr(preview, "MAX_PREVIEW_PAGES", 1)
    with pytest.raises(ValueError, match="limited"):
        preview.preview_report(draft())
    assert not outputs[0].exists()
    monkeypatch.setattr(preview, "MAX_PREVIEW_PAGES", 150)
    monkeypatch.setattr(preview, "MAX_PREVIEW_BYTES", 1)
    with pytest.raises(ValueError, match="25 MB"):
        preview.preview_report(draft())
    assert not outputs[-1].exists()


def test_preview_refresh_failure_keeps_draft_and_disables_stale_download(monkeypatch, tmp_path):
    fake_converter(monkeypatch, tmp_path)
    app = monthly(monkeypatch, tmp_path)
    step(app, 2)
    next(b for b in app.button if b.label == "Create preview PDF").click().run()
    assert not app.exception
    assert not next(w for w in app.get("download_button") if w.label == "Download draft preview PDF").disabled
    next(w for w in app.text_area if w.label == "Activity summary").set_value("Synthetic updated maintenance.").run()
    assert next(w for w in app.get("download_button") if w.label == "Download draft preview PDF").disabled
    def fail(path):
        raise RuntimeError("Synthetic converter unavailable")
    monkeypatch.setattr(preview.pdf_converter, "convert_to_pdf", fail)
    next(b for b in app.button if b.label == "Refresh preview PDF").click().run()
    assert not app.exception
    assert next(w for w in app.text_area if w.label == "Activity summary").value == "Synthetic updated maintenance."
    assert next(w for w in app.get("download_button") if w.label == "Download draft preview PDF").disabled
    assert any("unchanged" in w.value for w in app.warning)


def layout_draft():
    base = draft()
    raw = BytesIO()
    Image.new("RGB", (900, 600), "#86a5b1").save(raw, "PNG")
    section = replace(base.sections[0], divider_asset="synthetic.png", blocks=(BlockSpec("visits", "table", columns=(ColumnSpec("site", "Site"), ColumnSpec("work", "Work completed"))),))
    block = ResolvedBlock("visits", "This month", rows=tuple((f"Synthetic Site {n % 3 + 1}", f"Synthetic inspection {n + 1} completed with no issues.") for n in range(70)),
                          extra_tables=(ReportTable(("Facility", "Status"), (("Synthetic North", "Complete"),)),))
    return replace(base, sections=(section,), blocks=(block,)), lambda ref: raw.getvalue()


def test_divider_geometry_table_headers_and_deterministic_docx():
    report, loader = layout_draft()
    raw = assemble_docx(report, asset_loader=loader)
    assert raw == assemble_docx(report, asset_loader=loader)
    document = Document(BytesIO(raw))
    anchors = document.element.xpath("//wp:anchor")
    assert len(anchors) == 1 and anchors[0].get("behindDoc") == "1"
    assert anchors[0].xpath("./wp:positionH")[0].get("relativeFrom") == "page"
    assert len(document.tables) == 2
    assert all(table.rows[0]._tr.xpath("./w:trPr/w:tblHeader") for table in document.tables)
    assert all(row._tr.xpath("./w:trPr/w:cantSplit") for table in document.tables for row in table.rows)
    assert estimate_bytes(report, loader) > 400_000


@requires_libreoffice
def test_real_preview_and_full_bleed_layout_render(monkeypatch, tmp_path):
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path))
    report, loader = layout_draft()
    package = generate_report(report, asset_loader=loader)
    assert package.pdf, package.pdf_error
    with fitz.open(stream=package.pdf, filetype="pdf") as pdf:
        assert "Section 1" in pdf[2].get_text()
        assert "Monthly Activity Summary" in pdf[2].get_text()
        assert tuple(pdf[2].get_image_rects(pdf[2].get_images()[0][0])[0]) == pytest.approx(tuple(pdf[2].rect), abs=0.1)
        assert len(pdf) >= 5
        assert "Site" in pdf[4].get_text() and "Work completed" in pdf[4].get_text()
        assert "Synthetic inspection 70" in "".join(p.get_text() for p in pdf)
    result = preview.preview_report(replace(report, prepared_by=""), loader)
    assert result.pages == len(fitz.open(stream=package.pdf, filetype="pdf"))
    with fitz.open(stream=result.pdf, filetype="pdf") as pdf:
        assert all("DRAFT - REVIEW ONLY" in p.get_text() for p in pdf)


def test_guided_divider_block_is_rendered_and_explicit_omission_wins():
    from zipfile import ZipFile
    image = BytesIO()
    Image.new("RGB", (800, 1000), "green").save(image, "PNG")
    base = draft()
    divider = ResolvedBlock("divider_activity", "Last month", asset_hashes=("synthetic.png",))
    divider = replace(divider, client_reviewed_fingerprint=divider.fingerprint)
    report = replace(base, blocks=(*base.blocks, divider))
    raw = assemble_docx(report, asset_loader=lambda ref: image.getvalue())
    with ZipFile(BytesIO(raw)) as package:
        assert b'behindDoc="1"' in package.read("word/document.xml")
    report = replace(report, sections=tuple(replace(s, divider_asset="synthetic.png") for s in report.sections),
                     blocks=(*base.blocks, replace(divider, source="Omit")))
    raw = assemble_docx(report, asset_loader=lambda ref: image.getvalue())
    with ZipFile(BytesIO(raw)) as package:
        assert b'behindDoc="1"' not in package.read("word/document.xml")
