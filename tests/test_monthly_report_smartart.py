from zipfile import ZipFile, ZIP_DEFLATED
from xml.etree import ElementTree as ET

from docx import Document
import pytest

from app.monthly_report_import import inspect_docx, map_items, ImportMapping
from app.monthly_report_smartart import A, R, DGM, ROOTS

REL = 'http://schemas.openxmlformats.org/package/2006/relationships'
OFFICE = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/'
W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'


def smartart_doc(tmp_path, *, drawing_text='Synthetic leader', model_text='Synthetic leader', invalid=None):
    doc = Document()
    doc.add_heading('Organizational Chart', 1)
    doc.add_paragraph('')
    doc.add_heading('Training Summary', 1)
    doc.add_paragraph('Completed training.')
    path = tmp_path / 'smartart.docx'
    doc.save(path)
    with ZipFile(path) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    root = ET.fromstring(parts['word/document.xml'])
    paragraph = list(root.find(W + 'body'))[1]
    graphic = ET.SubElement(ET.SubElement(paragraph, W + 'r'), A + 'graphicData', {'uri': DGM[1:-1]})
    ET.SubElement(graphic, DGM + 'relIds', {R + key: 'rId10' + str(n) for n, key in enumerate(('dm', 'lo', 'qs', 'cs'))})
    rels = ET.fromstring(parts['word/_rels/document.xml.rels'])
    for n, kind in enumerate(ROOTS):
        target = f'word/diagrams/{kind}.xml'
        typ = OFFICE + kind if kind != 'diagramDrawing' else 'http://schemas.microsoft.com/office/2007/relationships/diagramDrawing'
        ET.SubElement(rels, '{' + REL + '}Relationship', {'Id': 'rId10' + str(n), 'Type': typ, 'Target': 'diagrams/' + kind + '.xml'})
        node = ET.Element(ROOTS[kind])
        if kind == 'diagramData':
            ET.SubElement(node, A + 't').text = model_text
            ET.SubElement(node, '{http://schemas.microsoft.com/office/drawing/2008/diagram}dataModelExt', {'relId': 'rId104'})
        elif kind == 'diagramDrawing':
            ET.SubElement(node, A + 't').text = drawing_text
            if invalid == 'image':
                ET.SubElement(node, A + 'blip', {R + 'embed': 'image1'})
        elif kind == 'diagramLayout':
            ET.SubElement(node, DGM + 'shape', {R + 'blip': 'missing' if invalid == 'missing' else ''})
            if invalid == 'depth':
                nested = node
                for _ in range(70):
                    nested = ET.SubElement(nested, DGM + 'shape')
        parts[target] = ET.tostring(node)
    if invalid == 'external':
        external = ET.Element('{' + REL + '}Relationships')
        ET.SubElement(external, '{' + REL + '}Relationship', {'Id': 'external', 'Type': OFFICE + 'diagramData', 'Target': 'https://example.invalid/data.xml', 'TargetMode': 'External'})
        parts['word/diagrams/_rels/diagramLayout.xml.rels'] = ET.tostring(external)
    if invalid == 'missing_drawing':
        parts.pop('word/diagrams/diagramDrawing.xml')
    parts['word/document.xml'] = ET.tostring(root)
    parts['word/_rels/document.xml.rels'] = ET.tostring(rels)
    with ZipFile(path, 'w', ZIP_DEFLATED) as archive:
        for name, raw in parts.items():
            archive.writestr(name, raw)
    return path


def test_closed_smartart_text_has_stable_native_provenance_and_maps_to_chart(tmp_path):
    path = smartart_doc(tmp_path)
    inspection = inspect_docx(path)
    item = next(i for i in inspection.items if i.note == 'Validated native SmartArt text')
    assert item.kind == 'text' and item.suggested_slot == 'org_chart' and item.text == 'Synthetic leader'
    mapped = map_items(path, inspection, (ImportMapping(item.id, 'org_chart'),))
    assert mapped.blocks[0].text == item.text
    assert mapped.blocks[0].references == (f'docx:{inspection.sha256}:{item.id}',)
    assert not mapped.blocks[0].client_asset_reviews
    invalid = inspect_docx(smartart_doc(tmp_path, invalid='missing_drawing'))
    same = next(i for i in invalid.items if i.note.startswith('Unmapped native SmartArt'))
    assert same.id == item.id and same.position == item.position
    assert [(i.id, i.text) for i in inspection.items if i.section == 'training'] == [(i.id, i.text) for i in invalid.items if i.section == 'training']


@pytest.mark.parametrize('invalid', ['image', 'missing', 'external', 'missing_drawing', 'depth'])
def test_unproved_diagram_dependencies_remain_unsupported(tmp_path, invalid):
    inspection = inspect_docx(smartart_doc(tmp_path, invalid=invalid))
    assert not any(i.note == 'Validated native SmartArt text' for i in inspection.items)
    assert any(i.note.startswith('Unmapped native SmartArt') for i in inspection.items)


def test_hidden_model_text_cannot_piggyback_on_visible_cached_drawing(tmp_path):
    inspection = inspect_docx(smartart_doc(tmp_path, model_text='Synthetic leader Secret old contact'))
    assert any(i.note.startswith('Unmapped native SmartArt') for i in inspection.items)


def test_smartart_price_text_remains_visible_to_existing_content_policy(tmp_path):
    from app.monthly_report_content_policy import contains_price
    inspection = inspect_docx(smartart_doc(tmp_path, model_text='Service cost $500', drawing_text='Service cost $500'))
    item = next(i for i in inspection.items if i.note == 'Validated native SmartArt text')
    assert contains_price(item.text)


def test_smartart_part_bytes_are_bounded(tmp_path, monkeypatch):
    from app import monthly_report_smartart
    path = smartart_doc(tmp_path)
    monkeypatch.setattr(monthly_report_smartart, 'MAX_PART_BYTES', 10)
    assert any(i.note.startswith('Unmapped native SmartArt') for i in inspect_docx(path).items)
