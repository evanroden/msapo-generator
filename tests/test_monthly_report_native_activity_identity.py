"""Grounded native Activity identity; no invented dates, sites, completion or price."""
from dataclasses import replace
from io import BytesIO

from docx import Document

from app import monthly_report_sources as sources
from app import monthly_report_structured_uploads as structured
from app.monthly_report_model import ReportSource, ResolvedBlock, synthetic_profiles
from app.monthly_report_sources import SourceContent


def sample(text, *, suffix=".txt", vendor="", when="", identity="a"):
    digest = identity * 64
    report = ReportSource(digest, "vendor-evidence" + suffix, digest, suffix,
                          page_texts=(text,), vendor=vendor, service_date=when)
    return SourceContent(report)


def test_dated_txt_events_keep_each_actual_date_vendor_and_status():
    body = (
        "2026-09-04 | WO-1 | Vendor service\n"
        "Inspected air handler. Replacement was recommended, not completed.\n"
        "2026-09-09 | TRAIN-1 | Training Summary\n"
        "Completed a synthetic course.\n"
        "2026-09-20 | WO-2 | Vendor service\n"
        "Tightened a drive belt. Follow-up is pending.\n"
    )
    content = sample(body, vendor="Trane")
    result = structured.native_update((content,), "activity_summary")
    assert "2026-09-04 · Trane — Inspected air handler" in result.text
    assert "2026-09-20 · Trane — Tightened a drive belt" in result.text
    assert "Replacement was recommended, not completed." in result.text
    assert "Follow-up is pending." in result.text
    assert "Completed a synthetic course" not in result.text
    assert result.references == (sources.source_reference(content.source, 1),)
    assert structured.native_update((content,), "activity_summary", result) == result


def test_docx_source_labeled_vendor_and_date_survive_native_ingestion(monkeypatch, tmp_path):
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path))
    doc = Document()
    doc.add_paragraph("Vendor: Trane")
    doc.add_paragraph("Service date: 2026-09-12")
    doc.add_paragraph("Inspected fan motor. Repair remains pending.")
    original = BytesIO()
    doc.save(original)
    profile = synthetic_profiles()[0]
    content, _ = sources.ingest(profile, "vendor-report.docx", original.getvalue())
    assert content.source.vendor == "Trane"
    assert content.source.service_date == "2026-09-12"
    result = structured.native_update((content,), "activity_summary")
    assert "2026-09-12 · Trane — Inspected fan motor." in result.text
    assert "Repair remains pending." in result.text
    assert sources.source_bytes(profile, content.source) == original.getvalue()


def test_invalid_header_date_stays_unknown_even_if_source_has_another_date():
    body = "2026-99-99 | WO-1 | Vendor service\nInspected a synthetic motor.\n"
    content = sample(body, vendor="Trane", when="2026-09-15")
    result = structured.native_update((content,), "activity_summary")
    assert "Trane — Inspected a synthetic motor." in result.text
    assert "2026-99-99" not in result.text
    assert "2026-09-15" not in result.text


def test_metadata_correction_never_restores_edited_or_deleted_source_text():
    content = sample("Inspected AHU-7 and recorded belt wear.",
                     vendor="Trane", when="2026-09-10")
    initial = structured.native_update((content,), "activity_summary")
    assert "2026-09-10 · Trane — Inspected AHU-7" in initial.text
    corrected = replace(content, source=replace(content.source,
                                                service_date="2026-09-12", facility="demo"))
    assert corrected.source.fingerprint != content.source.fingerprint
    manually_edited = replace(initial, text="Operator's deliberately edited wording.")
    assert structured.native_update((corrected,), "activity_summary", manually_edited) == manually_edited
    cleared = replace(initial, text="")
    assert structured.native_update((corrected,), "activity_summary", cleared) == cleared


def test_changed_bytes_under_same_filename_are_a_distinct_source():
    old = sample("Inspected AHU-7.", identity="a")
    latest = sample("Replaced drive belts on AHU-7.", identity="b",
                    vendor="Trane", when="2026-09-20")
    first = structured.native_update((old,), "activity_summary")
    next_block = structured.native_update((latest,), "activity_summary", first)
    assert "Inspected AHU-7." in next_block.text
    assert "2026-09-20 · Trane — Replaced drive belts" in next_block.text
    assert len(next_block.references) == 2


def test_unknown_vendor_date_remain_uninvented_and_price_bearing_vendor_is_excluded():
    plain = sample("Inspected the pump on site.", vendor="", when="")
    assert structured.native_update((plain,), "activity_summary").text == "Inspected the pump on site."
    priced = sample("Inspected the pump on site.", vendor="Vendor fee $150",
                    when="2026-09-18")
    text = structured.native_update((priced,), "activity_summary").text
    assert text == "2026-09-18 — Inspected the pump on site."
    assert "$150" not in text


def test_csv_vendor_event_date_and_pending_status_are_source_grounded(monkeypatch, tmp_path):
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path))
    raw = (
        "Vendor,Event date,Work performed,Status,Price\n"
        "Trane,2026-09-16,Inspected chiller compressor.,Open,\"$2,000\"\n"
        "Trane,2026-09-20,Calibrated drive controller.,Completed,\"$300\"\n"
    ).encode()
    profile = synthetic_profiles()[0]
    content, _ = sources.ingest(profile, "vendor-events.csv", raw)
    result = structured.native_update((content,), "activity_summary")
    assert "2026-09-16 · Trane — Inspected chiller compressor. (Status: Open)" in result.text
    assert "2026-09-20 · Trane — Calibrated drive controller. (Status: Completed)" in result.text
    assert "$2,000" not in result.text and "$300" not in result.text
    assert len(result.references) == 1
    assert structured.native_update((content,), "activity_summary", result) == result
    assert sources.source_bytes(profile, content.source) == raw


def test_csv_ambiguous_event_date_stays_unknown_without_report_month_inference(monkeypatch, tmp_path):
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path))
    profile = synthetic_profiles()[0]
    content, _ = sources.ingest(
        profile, "ambiguous.csv",
        b"Vendor,Event date,Work performed\nTrane,09/10/2026,Inspected AHU motor.\n")
    result = structured.native_update((content,), "activity_summary")
    assert result.text == "Trane — Inspected AHU motor."
    assert "2026-09" not in result.text


def test_csv_with_mixed_sites_stays_out_of_auto_activity_until_scoped_review(monkeypatch, tmp_path):
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path))
    profile = synthetic_profiles()[0]
    content, _ = sources.ingest(
        profile, "mixed-sites.csv",
        b"Site,Vendor,Event date,Work performed\n"
        b"Foreign Site,Trane,2026-09-16,Inspected foreign AHU.\n"
        b"Demo North,Trane,2026-09-18,Inspected local AHU.\n")
    result = structured.native_update((content,), "activity_summary")
    assert not result.text and not result.references
    assert "Foreign Site" in content.source.page_texts[0]
    # The optional document reader cannot circumvent the same site-scope guard.
    # Its values are genuinely present in the original but foreign to this
    # report; a clear review action is required instead of silent insertion.
    import pytest
    with pytest.raises(ValueError, match="site-specific rows"):
        structured.reader_update((content,), "activity_summary", {
            "items": [{"source": content.source.id, "page": 1,
                       "text": "Inspected foreign AHU.", "row": []}]
        })


def test_cmms_csv_with_status_date_and_work_order_stays_on_separate_review_path(monkeypatch, tmp_path):
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path))
    profile = synthetic_profiles()[0]
    content, _ = sources.ingest(
        profile, "cmms.csv",
        b"Work Order,Status,Event date,Vendor,Work performed\n"
        b"WO-101,Completed,2026-09-18,Trane,Inspected pump.\n")
    result = structured.native_update((content,), "activity_summary")
    assert not result.text and not result.references


def test_native_vendor_csv_activity_reaches_real_word_and_pdf(monkeypatch, tmp_path):
    """Exercise the actual assembled synthetic monthly binaries, not a text stub."""
    from tests.test_monthly_report_full_synthetic_output import approved_synthetic_draft
    from app.monthly_report_docx import generate_report
    import fitz

    raw = (
        b"Vendor,Event date,Work performed,Status,Price\n"
        b"Trane,2026-09-16,Inspected chiller compressor.,Open,$1200\n"
    )
    profile = synthetic_profiles()[0]
    draft = approved_synthetic_draft(monkeypatch, tmp_path)
    content, _ = sources.ingest(profile, "vendor-technical-events.csv", raw)
    linked = structured.native_update((content,), "activity_summary")
    activity = replace(draft.blocks[0], text=linked.text)
    final = replace(draft, blocks=(activity, draft.blocks[1]))
    package = generate_report(final, acknowledged_fingerprint=final.fingerprint)
    assert package.pdf and package.docx, package.pdf_error
    word = Document(BytesIO(package.docx))
    word_text = "\n".join(p.text for p in word.paragraphs)
    assert "2026-09-16" in word_text and "Trane" in word_text
    assert "Inspected chiller compressor." in word_text
    assert "$1200" not in word_text
    with fitz.open(stream=package.pdf, filetype="pdf") as pdf:
        printed = "\n".join(page.get_text() for page in pdf)
    assert "2026-09-16" in printed and "Trane" in printed
    assert "Inspected chiller compressor." in printed
    assert "$1200" not in printed


def test_cached_reader_does_not_duplicate_dated_native_activity():
    source = sample("Completed strainer cleaning and leak inspection. Result: no leakage found.",
                    vendor="Trane", when="2026-09-03")
    native = structured.native_update((source,), "activity_summary")
    assert "2026-09-03 · Trane — Completed strainer cleaning" in native.text
    result, pages = structured.reader_update((source,), "activity_summary", {
        "items": [{"source": source.source.id, "page": 1,
                   "text": "Completed strainer cleaning and leak inspection. Result: no leakage found.",
                   "row": [], "include_page": False}],
    }, native)
    assert result.text == native.text
    assert result.references == native.references
    assert pages == {}
    corrected = replace(source, source=replace(source.source,
                                               service_date="2026-09-04"))
    repeated, _ = structured.reader_update((corrected,), "activity_summary", {
        "items": [{"source": corrected.source.id, "page": 1,
                   "text": "Completed strainer cleaning and leak inspection. Result: no leakage found.",
                   "row": []}],
    }, native)
    assert repeated.text == native.text
    # A deliberately cleared source-linked Activity is authoritative; a late
    # exact suggestion is not permission to resurrect the deleted wording.
    emptied = replace(native, text="")
    late, _ = structured.reader_update((source,), "activity_summary", {
        "items": [{"source": source.source.id, "page": 1,
                   "text": "Completed strainer cleaning and leak inspection. Result: no leakage found.",
                   "row": []}],
    }, emptied)
    assert late.text == "" and late.references == native.references
