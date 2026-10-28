"""Actual synthetic outputs contain only the selected contract's procedure pixels."""
from dataclasses import replace
from io import BytesIO
import hashlib
import json
import os
from pathlib import Path
from zipfile import ZipFile

from docx import Document
import pymupdf
import pytest
from PIL import Image

from app import monthly_report_library as library
from app import monthly_report_workflow_standards as workflows
from app.monthly_report_docx import assemble_docx, generate_report
from app.monthly_report_model import BlockSpec, SectionSpec, ResolvedBlock
from tests.conftest import requires_libreoffice
from test_monthly_report_contract_workflows import isolated, report
from test_monthly_report_workflow_standards import publish


def selected_contract(contract):
    publish('Contract A', 'blue')
    publish('Contract B', 'red')
    assets = {}
    current = workflows.apply_defaults(report(contract), assets)
    organization = SectionSpec('organization', '1', 'Organizational Chart', (
        BlockSpec('business_hours_workflow', 'image_page'), BlockSpec('contact_matrix', 'image_page')))
    contacts = ResolvedBlock('contact_matrix', 'This month', text='Synthetic contacts retained.')
    current = replace(current, profile=replace(current.profile, title=contract + ' synthetic procedures'),
                      sections=(organization,), blocks=(*current.blocks, contacts))
    return current, assets


@pytest.mark.parametrize('contract', ['Contract A', 'Contract B'])
def test_word_reopen_and_section_preview_keep_contract_picture_identity(contract):
    from app.monthly_report_preview import section_preview_docx
    current, assets = selected_contract(contract)
    saved = library.save_snapshot(current, expected_revision=0, assets=tuple(assets.items()))
    reloaded = library.load_snapshot(current.profile.contract, current.profile.key, current.period)
    assert reloaded == saved
    outputs = (assemble_docx(reloaded.draft, acknowledged_fingerprint=reloaded.draft.fingerprint,
                             asset_loader=assets.__getitem__),
               section_preview_docx(reloaded.draft, 'organization', asset_loader=assets.__getitem__))
    for raw in outputs:
        with ZipFile(BytesIO(raw)) as package:
            images = [package.read(name) for name in package.namelist() if name.startswith('word/media/')]
        def pixels(raw):
            with Image.open(BytesIO(raw)) as image:
                return image.size, image.convert("RGB").tobytes()
        # Existing output sanitization may re-encode PNGs; exact decoded
        # geometry/pixels, not compression bytes, establish picture identity.
        rendered = [pixels(raw) for raw in images]
        assert pixels(list(assets.values())[0]) in rendered
        other = workflows.load_standards('Contract B' if contract == 'Contract A' else 'Contract A')
        foreign = workflows.read_asset(other.workflows[0].block.asset_hashes[0],
                                       contract='Contract B' if contract == 'Contract A' else 'Contract A')
        assert pixels(foreign) not in rendered
        assert len(Document(BytesIO(raw)).inline_shapes) == 1
        assert 'Synthetic contacts retained.' in '\n'.join(p.text for p in Document(BytesIO(raw)).paragraphs)


@requires_libreoffice
@pytest.mark.parametrize('contract,color,foreign', [('Contract A', (0,0,255), (255,0,0)),
                                                   ('Contract B', (255,0,0), (0,0,255))])
def test_actual_pdf_contains_selected_contract_color_only(contract, color, foreign, tmp_path):
    import numpy as np
    current, assets = selected_contract(contract)
    package = generate_report(current, acknowledged_fingerprint=current.fingerprint, asset_loader=assets.__getitem__)
    assert package.pdf, package.pdf_error
    with pymupdf.open(stream=package.pdf, filetype='pdf') as pdf:
        pages = []
        for page in pdf:
            pixmap = page.get_pixmap(alpha=False)
            samples = np.frombuffer(pixmap.samples, dtype=np.uint8).reshape(-1, 3)
            counts = [int(np.all(samples == value, axis=1).sum()) for value in (color, foreign)]
            pages.append(counts)
        assert sum(count > 100 for count, _ in pages) == 1
        assert sum(count for _, count in pages) == 0
        assert 'Synthetic contacts retained.' in '\n'.join(page.get_text() for page in pdf)
    root = Path(os.environ.get('INDEPENDENT_OUTPUT', str(tmp_path))) / 'contract-workflows'
    root.mkdir(parents=True, exist_ok=True)
    name = contract.lower().replace(' ', '-')
    (root / (name + '.docx')).write_bytes(package.docx)
    (root / (name + '.pdf')).write_bytes(package.pdf)
    (root / (name + '.json')).write_text(json.dumps({
        'scope': 'Synthetic contract isolation; not native-master or genuine-report acceptance.',
        'docx_sha256': hashlib.sha256(package.docx).hexdigest(),
        'pdf_sha256': hashlib.sha256(package.pdf).hexdigest(),
        'page_color_counts': pages,
    }, indent=2))
