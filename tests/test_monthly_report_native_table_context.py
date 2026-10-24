"""Current table context keeps its source placement without borrowing old facts."""
from dataclasses import replace
from io import BytesIO

import pytest
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from app.monthly_report_import import _text, inspect_docx
from app.monthly_report_model import (
    BlockSpec, Facility, ReportDraft, ReportPeriod, ReportProfile,
    ResolvedBlock, SectionSpec,
)
from app.monthly_report_native_layout import build_native_docx
from app.monthly_report_sections import table_without_prices


TITLE = "Contract Asset List — Synthetic covered equipment"
SCOPE = "Synthetic regional campuses. Source: the approved example contract schedule."
TRAINING = "Current training information has not been supplied."
PROPOSALS = "Current proposal decisions have not been supplied."


def scenario(tmp_path, *, include_context=True):
    document = Document()
    document.add_paragraph("Synthetic facility")
    document.add_paragraph("September 2026")
    document.add_section(WD_SECTION_START.NEW_PAGE)
    document.add_heading("PENDING & DECLINED PROPOSALS", 1)
    document.add_paragraph("Synthetic proposal overview")
    for title, value in (("Rejected Quotes", "Synthetic rejected item"),
                         ("Pending Quotes", "Synthetic pending item")):
        label = document.add_paragraph(title)
        label.runs[0].bold = True
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Equipment"
        table.cell(0, 1).text = "Status"
        table.cell(1, 0).text = value
        table.cell(1, 1).text = "Historical"
    document.add_section(WD_SECTION_START.NEW_PAGE)
    document.add_heading("TRAINING SUMMARY", 1)
    document.add_paragraph("Old monthly training facts")
    document.add_section(WD_SECTION_START.NEW_PAGE)
    title = document.add_paragraph(TITLE)
    title.runs[0].bold = True
    scope = document.add_paragraph(SCOPE)
    scope.runs[0].italic = True
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Description"
    table.cell(0, 1).text = "Campus"
    table.cell(1, 0).text = "Synthetic air separator"
    table.cell(1, 1).text = "Synthetic western campus"
    path = tmp_path / "synthetic-table-context.docx"
    document.save(path)
    inspection = inspect_docx(path)

    def tables(section):
        return tuple(replace(table_without_prices(item)[0],
                             reference=f"docx:{inspection.sha256}:{item.id}")
                     for item in inspection.items if item.kind == "table" and item.section == section)

    # Context is explicitly authored current text. A prior-month import may
    # have selected only the table, without selecting the old monthly prose.
    proposals = ResolvedBlock("proposals", "This month", text="\n\n".join((
        PROPOSALS, "Synthetic proposal overview", "Rejected Quotes", "Pending Quotes")),
        extra_tables=tables("proposals"))
    training = ResolvedBlock("training_summary", "This month",
        text="\n\n".join((TRAINING, TITLE, SCOPE)) if include_context else TRAINING,
        extra_tables=tables("training"))
    profile = ReportProfile("Synthetic contract", "synthetic", "Synthetic facility",
                            (Facility("synthetic", "Synthetic facility"),))
    sections = (SectionSpec("proposals", "10", "Pending & Declined Proposals",
                            (BlockSpec("proposals", "rich_text"),)),
                SectionSpec("training", "11", "Training Summary",
                            (BlockSpec("training_summary", "rich_text"),)))
    draft = ReportDraft(profile, ReportPeriod(2026, 10), "Synthetic editor", sections, (proposals, training))
    return path, draft


def before(table):
    result = []
    node = table.getprevious()
    while node is not None and node.tag == qn("w:p"):
        if node.find(".//" + qn("w:sectPr")) is not None:
            break
        if _text(node):
            result.append(node)
        node = node.getprevious()
    return list(reversed(result))


@pytest.mark.parametrize("section", [None, "proposals"])
def test_changed_proposal_notes_do_not_duplicate_bound_table_headings(tmp_path, section):
    path, draft = scenario(tmp_path)
    original = path.read_bytes()
    document = Document(BytesIO(build_native_docx(draft, path, section_key=section)))
    text = _text(document.element)
    assert text.count("Rejected Quotes") == text.count("Pending Quotes") == 1
    for table, title in zip(document.tables, ("Rejected Quotes", "Pending Quotes")):
        label = before(table._tbl)[-1]
        assert _text(label) == title
        assert label.find(".//" + qn("w:b")) is not None
    body = list(document.element.body)
    note = next(p._p for p in document.paragraphs if p.text == PROPOSALS)
    assert body.index(note) < body.index(document.tables[0]._tbl)
    assert path.read_bytes() == original


@pytest.mark.parametrize("section", [None, "training"])
def test_explicit_appendix_context_stays_immediately_before_its_retained_table(tmp_path, section):
    path, draft = scenario(tmp_path)
    document = Document(BytesIO(build_native_docx(draft, path, section_key=section)))
    context = before(document.tables[-1]._tbl)
    assert [_text(p) for p in context[-2:]] == [TITLE, SCOPE]
    assert context[-2].find(".//" + qn("w:b")) is not None
    assert context[-1].find(".//" + qn("w:i")) is not None
    text = _text(document.element)
    assert text.count(TITLE) == text.count(SCOPE) == 1
    assert TRAINING in text and "Old monthly training facts" not in text
    assert document.tables[-1].cell(1, 1).text == "Synthetic western campus"


@pytest.mark.parametrize("page_break", [False, True])
def test_empty_source_paragraph_ends_table_context_adjacency(tmp_path, page_break):
    path, draft = scenario(tmp_path)
    document = Document(path)
    boundary = OxmlElement("w:p")
    if page_break:
        properties = OxmlElement("w:pPr")
        properties.append(OxmlElement("w:pageBreakBefore"))
        boundary.append(properties)
    document.tables[-1]._tbl.addprevious(boundary)
    document.save(path)
    original = path.read_bytes()
    inspection = inspect_docx(path)
    item = next(i for i in inspection.items if i.kind == "table" and i.section == "training")
    proposals, training = draft.blocks
    training = replace(training, extra_tables=(replace(training.extra_tables[0],
        reference=f"docx:{inspection.sha256}:{item.id}"),))
    result = Document(BytesIO(build_native_docx(replace(draft, blocks=(proposals, training)),
                                               path, section_key="training")))
    # Matching words before an empty source paragraph are not adjacent table
    # context. Keep them together in the author's current narrative instead.
    assert not any(_text(p) in (TITLE, SCOPE) for p in before(result.tables[-1]._tbl))
    lines = [p.text for p in result.paragraphs if p.text.strip()]
    start = lines.index(TRAINING)
    assert tuple(lines[start:start + 3]) == (TRAINING, TITLE, SCOPE)
    assert result.tables[-1].cell(1, 1).text == "Synthetic western campus"
    assert path.read_bytes() == original


@pytest.mark.parametrize("master", [False, True])
def test_retaining_a_table_does_not_restore_unsupplied_source_context(tmp_path, master):
    path, draft = scenario(tmp_path, include_context=False)
    document = Document(BytesIO(build_native_docx(draft, path, master=master)))
    text = _text(document.element)
    assert TITLE not in text and SCOPE not in text
    assert TRAINING in text
    assert document.tables[-1].cell(1, 1).text == "Synthetic western campus"


@pytest.mark.parametrize("kind", ["wrong-source", "edited-table", "edited-context", "omit"])
def test_unproved_or_left_out_context_never_replays_old_source_prose(tmp_path, kind):
    path, draft = scenario(tmp_path)
    proposals, training = draft.blocks
    table = training.extra_tables[0]
    if kind == "wrong-source":
        training = replace(training, extra_tables=(replace(table, reference="docx:" + "0" * 64 + ":unknown"),))
    elif kind == "edited-table":
        training = replace(training, extra_tables=(replace(table, rows=(("Current equipment", "Current campus"),)),))
    elif kind == "edited-context":
        training = replace(training, text=TRAINING + "\n\nCurrent authored scope.")
    else:
        training = replace(training, source="Omit")
    document = Document(BytesIO(build_native_docx(replace(draft, blocks=(proposals, training)), path)))
    text = _text(document.element)
    assert "Old monthly training facts" not in text
    if kind in ("edited-context", "omit"):
        assert TITLE not in text and SCOPE not in text
    elif kind == "wrong-source":
        # Current words still render, but cannot claim the old source location.
        assert TITLE in text and SCOPE in text
        assert not any(_text(p) == TITLE for p in before(document.tables[-1]._tbl))
    else:
        assert "Current equipment" in text and "Synthetic air separator" not in text
    assert "TRAINING SUMMARY" in text  # Owner-required title page survives.


def test_current_words_do_not_certify_a_hidden_source_paragraph(tmp_path):
    path, draft = scenario(tmp_path)
    document = Document(path)
    scope = next(p for p in document.paragraphs if p.text == SCOPE)
    scope.runs[0]._r.get_or_add_rPr().append(OxmlElement("w:vanish"))
    document.save(path)
    inspection = inspect_docx(path)
    item = next(item for item in inspection.items if item.kind == "table" and item.section == "training")
    proposals, training = draft.blocks
    training = replace(training, extra_tables=(replace(training.extra_tables[0],
        reference=f"docx:{inspection.sha256}:{item.id}"),))
    result = Document(BytesIO(build_native_docx(replace(draft, blocks=(proposals, training)),
                                                path, section_key="training")))
    assert SCOPE in _text(result.element)  # The author's current words survive.
    assert not any(_text(p) == SCOPE for p in before(result.tables[-1]._tbl))
    current = next(p for p in result.paragraphs if p.text == SCOPE)
    assert current._p.find(".//" + qn("w:vanish")) is None


@pytest.mark.parametrize("inherited", [False, True])
@pytest.mark.parametrize("target", [TITLE, SCOPE])
def test_matching_current_context_does_not_retain_hidden_paragraph_style(tmp_path, inherited, target):
    path, draft = scenario(tmp_path)
    document = Document(path)
    hidden = document.styles.add_style("Hidden supplied context", WD_STYLE_TYPE.PARAGRAPH)
    hidden.font.hidden = True
    if inherited:
        style = document.styles.add_style("Derived supplied context", WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = hidden
    else:
        style = hidden
    next(p for p in document.paragraphs if p.text == target).style = style
    document.save(path)
    inspection = inspect_docx(path)
    item = next(i for i in inspection.items if i.kind == "table" and i.section == "training")
    proposals, training = draft.blocks
    training = replace(training, extra_tables=(replace(training.extra_tables[0],
        reference=f"docx:{inspection.sha256}:{item.id}"),))
    result = Document(BytesIO(build_native_docx(replace(draft, blocks=(proposals, training)),
                                               path, section_key="training")))
    # The author supplied visible current words, not permission to replay a
    # hidden old paragraph. They stay in the normal authored narrative flow.
    current = next(p for p in result.paragraphs if p.text == target)
    assert current.style.name == "Normal"
    assert current.style.font.hidden is not True
    assert not any(_text(p) in (TITLE, SCOPE) for p in before(result.tables[-1]._tbl))


@pytest.mark.parametrize("inherited", [False, True])
@pytest.mark.parametrize("target", [TITLE, SCOPE])
def test_inherited_numbering_cannot_add_unsupplied_source_words(tmp_path, inherited, target):
    path, draft = scenario(tmp_path)
    document = Document(path)
    numbering = document.part.numbering_part.element
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), "90")
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    for tag, value in (("start", "1"), ("numFmt", "decimal"),
                       ("lvlText", "UNSUPPLIED SYNTHETIC PREFIX %1.")):
        node = OxmlElement("w:" + tag)
        node.set(qn("w:val"), value)
        level.append(node)
    abstract.append(level)
    numbering.append(abstract)
    number = OxmlElement("w:num")
    number.set(qn("w:numId"), "90")
    abstract_id = OxmlElement("w:abstractNumId")
    abstract_id.set(qn("w:val"), "90")
    number.append(abstract_id)
    numbering.append(number)
    numbered = document.styles.add_style("Numbered source context", WD_STYLE_TYPE.PARAGRAPH)
    properties = OxmlElement("w:numPr")
    number_id = OxmlElement("w:numId")
    number_id.set(qn("w:val"), "90")
    properties.append(number_id)
    numbered.element.get_or_add_pPr().append(properties)
    style = numbered
    if inherited:
        style = document.styles.add_style("Derived numbered context", WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = numbered
    next(p for p in document.paragraphs if p.text == target).style = style
    document.save(path)
    inspection = inspect_docx(path)
    item = next(i for i in inspection.items if i.kind == "table" and i.section == "training")
    proposals, training = draft.blocks
    training = replace(training, extra_tables=(replace(training.extra_tables[0],
        reference=f"docx:{inspection.sha256}:{item.id}"),))
    result = Document(BytesIO(build_native_docx(replace(draft, blocks=(proposals, training)),
                                               path, section_key="training")))
    current = next(p for p in result.paragraphs if p.text == target)
    assert current.style.name == "Normal"
    assert current._p.find("./" + qn("w:pPr") + "/" + qn("w:numPr")) is None
    assert not any(_text(p) in (TITLE, SCOPE) for p in before(result.tables[-1]._tbl))


@pytest.mark.parametrize("unproved", ["document-default", "missing-style", "style-cycle"])
def test_unproved_default_or_style_chain_cannot_retain_source_context(tmp_path, unproved):
    path, draft = scenario(tmp_path)
    document = Document(path)
    scope = next(p for p in document.paragraphs if p.text == SCOPE)
    style = document.styles.add_style("Supplied context", WD_STYLE_TYPE.PARAGRAPH)
    scope.style = style
    if unproved == "document-default":
        defaults = document.styles.element.find(qn("w:docDefaults"))
        properties = defaults.find("./" + qn("w:rPrDefault") + "/" + qn("w:rPr"))
        properties.append(OxmlElement("w:vanish"))
        # The fallback authoring prototype is explicitly visible; the custom
        # source style inherits the hidden document default instead.
        document.styles["Normal"].font.hidden = False
    else:
        base = OxmlElement("w:basedOn")
        base.set(qn("w:val"), "Missing context style" if unproved == "missing-style" else style.style_id)
        style.element.append(base)
    document.save(path)
    inspection = inspect_docx(path)
    item = next(i for i in inspection.items if i.kind == "table" and i.section == "training")
    proposals, training = draft.blocks
    training = replace(training, extra_tables=(replace(training.extra_tables[0],
        reference=f"docx:{inspection.sha256}:{item.id}"),))
    result = Document(BytesIO(build_native_docx(replace(draft, blocks=(proposals, training)),
                                               path, section_key="training")))
    current = next(p for p in result.paragraphs if p.text == SCOPE)
    assert current.style.name == "Normal"
    assert current.style.font.hidden is not True
    assert not any(_text(p) == SCOPE for p in before(result.tables[-1]._tbl))


@pytest.mark.parametrize("current", [
    (SCOPE, TITLE),
    (TITLE, "Current appendix-specific caveat.", SCOPE),
    (TITLE, SCOPE, "Current appendix-specific caveat."),
    ("Current corrected appendix title.", SCOPE),
    (TITLE, "Current corrected source scope."),
])
def test_changed_current_context_order_or_association_keeps_authored_flow(tmp_path, current):
    path, draft = scenario(tmp_path)
    proposals, training = draft.blocks
    expected = (TRAINING, *current)
    training = replace(training, text="\n\n".join(expected))
    result = Document(BytesIO(build_native_docx(replace(draft, blocks=(proposals, training)),
                                               path, section_key="training")))
    lines = [p.text for p in result.paragraphs if p.text.strip()]
    start = lines.index(TRAINING)
    assert tuple(lines[start:start + len(expected)]) == expected
    assert not any(_text(p) in (TITLE, SCOPE) for p in before(result.tables[-1]._tbl))


def test_quote_specific_current_notes_keep_their_authored_labels(tmp_path):
    path, draft = scenario(tmp_path)
    proposals, training = draft.blocks
    expected = ("Rejected Quotes", "Current rejection-specific note.",
                "Pending Quotes", "Current pending-specific note.")
    proposals = replace(proposals, text="\n\n".join(expected))
    result = Document(BytesIO(build_native_docx(replace(draft, blocks=(proposals, training)),
                                               path, section_key="proposals")))
    lines = [p.text for p in result.paragraphs if p.text.strip()]
    start = lines.index("Rejected Quotes")
    assert tuple(lines[start:start + len(expected)]) == expected
