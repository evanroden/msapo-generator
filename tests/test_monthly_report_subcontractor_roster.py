"""Contract-vendor reuse stays explicit, versioned and site-bound.

Only invented contracts, sites, source rows and labels are used.
"""
from dataclasses import replace
from io import BytesIO

from docx import Document
import pytest
from streamlit.testing.v1 import AppTest

from app import monthly_report_library as library
from app.monthly_report_model import (
    ReportProfile, ReportPeriod, ReportTable, ResolvedBlock,
    default_sections, synthetic_profiles,
)
from app.monthly_report_subcontractor_roster import (
    SHARED_REF_PREFIX, apply_roster, eligible_rosters, load_rosters, save_roster,
)


NATIVE_SPEC = next(s for s in default_sections()
                   if s.key == "subcontractors").blocks[0]
COLUMNS = tuple(column.title for column in NATIVE_SPEC.columns)


def profile(contract, key="a1", sites=("north",)):
    original = synthetic_profiles()[1]
    facilities = tuple(f for f in original.facilities if f.key in sites)
    return replace(original, contract=contract, key=key, title="Synthetic " + key,
                   facilities=facilities,
                   scope_type="individual" if len(facilities) == 1 else "multi_site")


def rows(*values):
    return ResolvedBlock("subcontractor_matrix", "This month", rows=tuple(values))


def row(site, vendor, status=""):
    return (site, "HVAC", vendor, "Synthetic Contact", "", status)


def publish(p, block, *, site_keys=None, revision=0, actor="Synthetic Editor"):
    return save_roster(p, block, NATIVE_SPEC,
                       site_keys=tuple(f.key for f in p.facilities) if site_keys is None else site_keys,
                       expected_revision=revision, actor=actor, confirmed=True)


def test_explicit_promotion_never_writes_before_confirmation(monkeypatch, tmp_path):
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path))
    north = profile("Synthetic Contract A")
    current = rows(row(north.facilities[0].title, "Synthetic Vendor"))
    assert load_rosters(north.contract).revision == 0
    with pytest.raises(ValueError, match="Confirm"):
        save_roster(north, current, NATIVE_SPEC, site_keys=("north",),
                    expected_revision=0, actor="Synthetic Editor")
    assert load_rosters(north.contract).revision == 0
    assert not (tmp_path / "monthly_reports" / "subcontractor_rosters").exists()


def test_two_sites_and_contracts_have_independent_approved_rosters(monkeypatch, tmp_path):
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path))
    both = profile("Synthetic Contract A", "all-sites", ("north", "south"))
    north = profile("Synthetic Contract A", "a1", ("north",))
    south = profile("Synthetic Contract A", "a2", ("south",))
    foreign = profile("Synthetic Contract B", "b1", ("north",))
    state = publish(both, rows(row(both.facilities[0].title, "Vendor A North"),
                               row(both.facilities[1].title, "Vendor A South", "Not confirmed")))
    assert state.revision == 1 and len(state.entries) == 1
    assert eligible_rosters(state, north) and eligible_rosters(state, south)
    assert not eligible_rosters(state, foreign)
    assert load_rosters(foreign.contract).entries == ()
    assert load_rosters(both.contract) == state

    north_result = apply_roster(state, state.entries[0], north,
                                ResolvedBlock("subcontractor_matrix", "Omit"),
                                NATIVE_SPEC)
    south_result = apply_roster(state, state.entries[0], south,
                                ResolvedBlock("subcontractor_matrix", "Omit"),
                                NATIVE_SPEC)
    assert north_result.rows == (row(north.facilities[0].title, "Vendor A North"),)
    assert south_result.rows == (row(south.facilities[0].title, "Vendor A South", "Not confirmed"),)
    assert north_result.source == south_result.source == "Library"
    assert north_result.references[0].startswith(SHARED_REF_PREFIX)
    assert north_result.references[0] == south_result.references[0]
    with pytest.raises(ValueError, match="selected contract and sites"):
        apply_roster(state, state.entries[0], foreign,
                     ResolvedBlock("subcontractor_matrix", "Omit"), NATIVE_SPEC)


def test_contract_wide_requires_non_site_specific_rows_and_is_explicit(monkeypatch, tmp_path):
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path))
    north = profile("Synthetic Contract A")
    south = profile("Synthetic Contract A", "a2", ("south",))
    with pytest.raises(ValueError, match="Share it for the reviewed sites"):
        publish(north, rows(row(north.facilities[0].title, "Site vendor")),
                site_keys=())
    assert load_rosters(north.contract).revision == 0
    contract = publish(north, rows(row("All sites", "Contract vendor", "")),
                       site_keys=())
    assert contract.entries[0].sites == ()
    assert eligible_rosters(contract, south)
    after = apply_roster(contract, contract.entries[0], south,
                         ResolvedBlock("subcontractor_matrix", "Omit"), NATIVE_SPEC)
    assert after.rows[0][2] == "Contract vendor"


def test_report_only_manual_override_and_vendor_picture_are_preserved(
        monkeypatch, tmp_path):
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path))
    north = profile("Synthetic Contract A")
    state = publish(north, rows(row(north.facilities[0].title, "Standing Vendor")))
    original = ResolvedBlock(
        "subcontractor_matrix", "This month", text="Keep report-only phone notes.",
        rows=(row(north.facilities[0].title, "Manually corrected vendor"),),
        asset_hashes=("synthetic-vendor-page.png",),
        asset_captions=("Keep approved report page",),
        references=("synthetic:old-table", "synthetic:vendor-page"),
    )
    with pytest.raises(ValueError, match="existing vendor rows"):
        apply_roster(state, state.entries[0], north, original, NATIVE_SPEC)
    assert original.rows[0][2] == "Manually corrected vendor"
    updated = apply_roster(state, state.entries[0], north, original,
                           NATIVE_SPEC, replace_existing=True)
    assert updated.rows[0][2] == "Standing Vendor"
    assert updated.text == original.text
    assert updated.asset_hashes == original.asset_hashes
    assert updated.asset_captions == original.asset_captions
    assert "synthetic:vendor-page" in updated.references
    assert updated.client_reviewed_fingerprint == ""
    assert load_rosters(north.contract) == state
    # Report-only edits never auto-publish the changed master revision.
    manual = replace(updated, source="This month",
                     rows=(row(north.facilities[0].title, "New report-only correction"),))
    assert manual.rows != state.entries[0].tables[0].rows
    assert load_rosters(north.contract) == state


def test_source_category_bands_zero_blanks_and_provenance_survive_site_filter(
        monkeypatch, tmp_path):
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path))
    both = profile("Synthetic Contract A", "all-sites", ("north", "south"))
    original_table = ReportTable(
        ("Trade", "Facility", "Vendor", "When to contact", "Status"),
        (("Electrical", "", "", "", ""),
         ("", both.facilities[0].title, "Vendor North", "Call after 8 am", "0"),
         ("", both.facilities[1].title, "Vendor South", "", ""),
         ("Plumbing", "", "", "", "")),
        "synthetic:word:table-1",
    )
    block = ResolvedBlock("subcontractor_matrix", "This month",
                          extra_tables=(original_table,),
                          references=("synthetic:read-source",))
    state = publish(both, block)
    stored = state.entries[0]
    assert stored.origin_references == ("synthetic:read-source", "synthetic:word:table-1")
    assert stored.tables[0].rows == original_table.rows
    north = profile("Synthetic Contract A")
    selected = apply_roster(state, stored, north,
                            ResolvedBlock("subcontractor_matrix", "Omit"), NATIVE_SPEC)
    assert selected.rows == ()
    assert len(selected.extra_tables) == 1
    assert selected.extra_tables[0].columns == original_table.columns
    assert selected.extra_tables[0].rows == (
        ("Electrical", "", "", "", ""),
        ("", both.facilities[0].title, "Vendor North", "Call after 8 am", "0"),
        ("Plumbing", "", "", "", ""),
    )
    assert selected.extra_tables[0].reference.startswith(SHARED_REF_PREFIX)
    assert "Vendor South" not in str(selected.extra_tables[0].rows)
    assert original_table.rows[2][4] == ""


def test_conflict_preserves_newer_scope_and_immutable_history(monkeypatch, tmp_path):
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path))
    both = profile("Synthetic Contract A", "all-sites", ("north", "south"))
    north = profile("Synthetic Contract A")
    south = profile("Synthetic Contract A", "a2", ("south",))
    a = publish(north, rows(row(north.facilities[0].title, "North Version 1")))
    assert a.revision == 1
    b = publish(south, rows(row(south.facilities[0].title, "South Version 1")),
                revision=1)
    assert b.revision == 2 and len(b.entries) == 2
    with pytest.raises(library.RevisionConflict, match="changed"):
        publish(north, rows(row(north.facilities[0].title, "Stale North")),
                revision=1)
    assert load_rosters(north.contract) == b
    later = publish(north, rows(row(north.facilities[0].title, "North Version 2")),
                    revision=2)
    assert later.revision == 3 and len(later.entries) == 2
    assert any(t.rows[0][2] == "South Version 1"
               for e in later.entries for t in e.tables)
    root = library._root() / "subcontractor_rosters" / library._contract_directory(north.contract)
    from json import loads
    first = loads((root / "history" / "00000001.json").read_text())
    assert first["revision"] == 1
    assert first["entries"][0]["tables"][0]["rows"][0][2] == "North Version 1"
    assert load_rosters(north.contract).revision == 3
    assert not eligible_rosters(later, both)


def test_report_snapshots_remain_pinned_after_shared_revisions(monkeypatch, tmp_path):
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path))
    north = profile("Synthetic Contract A")
    current = publish(north, rows(row(north.facilities[0].title, "Standing 1")))
    block = apply_roster(current, current.entries[0], north,
                         ResolvedBlock("subcontractor_matrix", "Omit"), NATIVE_SPEC)
    from app.monthly_report_model import ReportDraft
    draft = ReportDraft(north, ReportPeriod(2026, 9), "Synthetic Editor",
                        default_sections(), (block,))
    saved = library.save_snapshot(draft, expected_revision=0,
                                  entered_editor="Synthetic Editor")
    assert saved.draft.blocks == (block,)
    later = publish(north, rows(row(north.facilities[0].title, "Standing 2")),
                    revision=1)
    assert later.revision == 2
    reopened = library.load_snapshot(north.contract, north.key, ReportPeriod(2026, 9))
    assert reopened.draft.blocks == (block,)
    assert reopened.draft.blocks[0].rows[0][2] == "Standing 1"
    october = replace(draft, period=ReportPeriod(2026, 10))
    approved = apply_roster(later, later.entries[0], north, october.blocks[0],
                            NATIVE_SPEC, replace_existing=True)
    assert approved.rows[0][2] == "Standing 2"
    assert reopened.draft.blocks[0].rows[0][2] == "Standing 1"


def test_unapproved_roster_rejects_unknown_facility_and_prices(monkeypatch, tmp_path):
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path))
    north = profile("Synthetic Contract A")
    with pytest.raises(ValueError, match="unconfirmed site"):
        publish(north, rows(row("Unknown facility", "Vendor")))
    with pytest.raises(ValueError, match="pricing"):
        publish(north, rows(row(north.facilities[0].title, "Vendor", "Invoice $80")))
    assert not load_rosters(north.contract).entries


def test_report_word_output_has_only_selected_site_and_preserves_vendor_status(monkeypatch, tmp_path):
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path))
    from app.monthly_report_docx import _build_docx
    from app.monthly_report_model import ReportDraft
    both = profile("Synthetic Contract A", "all-sites", ("north", "south"))
    state = publish(both, rows(row(both.facilities[0].title, "Vendor North", ""),
                               row(both.facilities[1].title, "Vendor South", "Pending")))
    north = profile("Synthetic Contract A")
    adopted = apply_roster(state, state.entries[0], north,
                           ResolvedBlock("subcontractor_matrix", "Omit"), NATIVE_SPEC)
    section = next(s for s in default_sections() if s.key == "subcontractors")
    draft = ReportDraft(north, ReportPeriod(2026, 9), "Synthetic Editor",
                        (section,), (adopted,))
    raw = _build_docx(draft)
    word = Document(BytesIO(raw))
    text = "\n".join(cell.text for table in word.tables
                     for tr in table.rows for cell in tr.cells)
    assert text.count("Vendor North") == 1
    assert "Vendor South" not in text
    assert "Sub-Contractor Status" in "\n".join(p.text for p in word.paragraphs)
    assert "Pending" not in text


from tests.conftest import requires_libreoffice


@requires_libreoffice
def test_real_pdf_renders_shared_vendor_contact_once_without_other_sites(monkeypatch, tmp_path):
    import fitz
    from app import pdf_converter
    from app.monthly_report_docx import _build_docx
    from app.monthly_report_model import ReportDraft

    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path))
    monkeypatch.setattr(pdf_converter, "OUTPUT_DIR", tmp_path)
    both = profile("Synthetic Contract A", "all-sites", ("north", "south"))
    state = publish(both, rows(row(both.facilities[0].title, "Synthetic Vendor North"),
                               row(both.facilities[1].title, "Synthetic Vendor South")))
    north = profile("Synthetic Contract A")
    adopted = apply_roster(state, state.entries[0], north,
                           ResolvedBlock("subcontractor_matrix", "Omit"), NATIVE_SPEC)
    section = next(s for s in default_sections() if s.key == "subcontractors")
    draft = ReportDraft(north, ReportPeriod(2026, 9), "Synthetic Editor",
                        (section,), (adopted,))
    document = tmp_path / "synthetic-vendor.docx"
    document.write_bytes(_build_docx(draft))
    pdf_file = pdf_converter.convert_to_pdf(document)
    with fitz.open(pdf_file) as pdf:
        all_text = "\n".join(page.get_text() for page in pdf)
        assert all_text.count("Synthetic Vendor North") == 1
        assert "Synthetic Vendor South" not in all_text
        assert "Sub-Contractor Status" in all_text
