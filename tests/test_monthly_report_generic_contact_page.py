"""Generic no-master Word fallback: chart page must not crowd contact rows.

Only invented records are used. The native ENFRA master layout and genuine
owner originals require their own per-page Word/PDF review and are not fixed
or certified by these tests.
"""
from dataclasses import replace
from io import BytesIO

from docx import Document
from docx.oxml.ns import qn
import pytest

from app.monthly_report_docx import _build_docx
from app.monthly_report_model import (
    OrgChartNode, ReportPeriod, ResolvedBlock, default_sections,
    synthetic_draft, synthetic_profiles,
)
from tests.conftest import requires_libreoffice


def _draft(*, chart=True, contacts=True):
    draft = synthetic_draft(synthetic_profiles()[0], ReportPeriod(2026, 9))
    section = next(s for s in default_sections() if s.key == "organization")
    section = replace(section, blocks=tuple(
        b for b in section.blocks if b.key in {"org_chart", "contact_matrix"}))
    blocks = (
        ResolvedBlock("org_chart", "This month", org_nodes=(
            OrgChartNode("synthetic-lead", "Synthetic Site Lead", "Site Lead"),
            OrgChartNode("synthetic-tech", "Synthetic Technician", "Technician",
                         reports_to="synthetic-lead"),
        )) if chart else ResolvedBlock("org_chart", "Omit"),
        ResolvedBlock("contact_matrix", "This month", rows=(
            ("Synthetic North", "Technician", "Synthetic Contact North", "555-0100", ""),
            ("Synthetic South", "Engineer", "Synthetic Contact South", "555-0200", ""),
        )) if contacts else ResolvedBlock("contact_matrix", "Omit"),
    )
    return replace(draft, sections=(section,), blocks=blocks)


def test_generic_chart_and_contacts_have_explicit_page_boundary_and_heading():
    doc = Document(BytesIO(_build_docx(_draft())))
    heading = next(p for p in doc.paragraphs if p.text == "Facility contacts")
    assert heading.style.name == "Heading 2"
    assert heading.paragraph_format.page_break_before is True
    assert any("Synthetic Contact North" in cell.text
               for table in doc.tables for row in table.rows for cell in row.cells)
    assert len(doc.inline_shapes) == 1
    # No duplicate chart, arbitrary inserted photo, or fabricated attribution.
    assert sum(p.text == "Facility contacts" for p in doc.paragraphs) == 1


@pytest.mark.parametrize(("chart", "contacts"), [(False, True), (True, False)])
def test_generic_chart_or_contact_alone_does_not_add_phantom_page(chart, contacts):
    doc = Document(BytesIO(_build_docx(_draft(chart=chart, contacts=contacts))))
    assert not any(p.text == "Facility contacts" for p in doc.paragraphs)


@requires_libreoffice
def test_real_generic_pdf_keeps_contact_table_off_chart_page(tmp_path, monkeypatch):
    import fitz
    from app import pdf_converter

    monkeypatch.setattr(pdf_converter, "OUTPUT_DIR", tmp_path)
    source = tmp_path / "synthetic-chart-and-contacts.docx"
    source.write_bytes(_build_docx(_draft()))
    path = pdf_converter.convert_to_pdf(source)
    try:
        with fitz.open(path) as pdf:
            chart_pages = [index for index, page in enumerate(pdf) if page.get_images(full=True)
                           and index >= 3]
            # Long table cells legitimately wrap names over multiple lines in
            # the PDF text layer. Locate rows by their synthetic phone fields,
            # then prove both entries and the chart remain on separate pages.
            contacts = [index for index, page in enumerate(pdf)
                        if "555-0100" in page.get_text()]
            assert chart_pages, "Actual chart illustration was omitted from PDF."
            assert contacts, "Actual contact row was omitted from PDF."
            assert contacts[0] > chart_pages[0], "Facility contacts were squeezed beneath the chart."
            assert "Facility contacts" in pdf[contacts[0]].get_text()
            assert "555-0100" not in pdf[chart_pages[0]].get_text()
            assert "555-0200" in pdf[contacts[0]].get_text()
    finally:
        path.unlink(missing_ok=True)
