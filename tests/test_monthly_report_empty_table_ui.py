"""Empty Capital/RFI entry keeps blank native slots out of saved/exported rows."""
from dataclasses import replace
from io import BytesIO

import pytest
from docx import Document
from streamlit.testing.v1 import AppTest

from app import monthly_report_active_work as active
from app.monthly_report_checks import preflight
from app.monthly_report_docx import assemble_docx
from app.monthly_report_model import ReportTable, ResolvedBlock
from app.monthly_report_table_entry_ui import EMPTY_TABLES
from test_monthly_report_rfi_ui import SCRIPT, edit_grid, field, row_fields, button


def workspace(key):
    return AppTest.from_string(SCRIPT.replace('"rfi_matrix"', repr(key))).run()


def current(app, key):
    return next(block for block in app.session_state['draft'].blocks if block.key == key)


def grid(app, key):
    return next(name for name, value in app.session_state.filtered_state.items()
                if name.startswith('rfi_test_edit_' + key) and isinstance(value, dict) and 'edited_rows' in value)


def open_entry(app, key):
    next(button for button in app.button if button.label == EMPTY_TABLES[key][1]).click().run()


def test_open_request_removes_the_empty_sentence_and_none_cell_placeholders():
    app = workspace('rfi_matrix')
    before = app.session_state['draft']
    open_entry(app, 'rfi_matrix')
    assert not app.exception and app.session_state['draft'] == before
    assert not any(item.value in {'No requests added.', 'No information requests added.'} for item in app.caption)
    assert {item.label: item.value for item in row_fields(app)} == {
        'Facility':'', 'Requested item':'', 'Status':''}


@pytest.mark.parametrize('key', tuple(EMPTY_TABLES))
def test_fresh_entry_opens_blank_native_cells_without_stale_empty_copy_or_data(key):
    app = workspace(key)
    original = app.session_state['draft']
    assert not app.exception and not app.dataframe
    assert any(item.value == EMPTY_TABLES[key][0] for item in app.caption)
    open_entry(app, key)
    assert not app.exception
    if key == 'rfi_matrix':
        assert not app.dataframe and row_fields(app)
        identity = field(app, 'Requested item').proto.id
        assert all(item.value == '' for item in row_fields(app))
        assert button(app, 'Add request').disabled
    else:
        assert len(app.dataframe) == 1
        identity = grid(app, key)
        assert set(app.session_state[identity + '_data'][0].values()) == {''}
        assert not any(button.label == EMPTY_TABLES[key][1] for button in app.button)
    assert app.session_state['draft'] == original
    assert not any(item.value == EMPTY_TABLES[key][0] for item in app.caption)
    app.run()
    assert (field(app, 'Requested item').proto.id if key == 'rfi_matrix' else grid(app, key)) == identity
    assert app.session_state['draft'] == original


@pytest.mark.parametrize('key', tuple(EMPTY_TABLES))
def test_blank_and_removed_rows_do_not_enter_durable_data_or_actual_word_output(monkeypatch, tmp_path, key):
    monkeypatch.setenv('EPC_DATA_DIR', str(tmp_path))
    app = workspace(key)
    open_entry(app, key)
    if key == 'rfi_matrix':
        identity = field(app, 'Requested item').proto.id
        field(app, 'Requested item').set_value('   ').run()
        assert not app.exception and not current(app, key).rows
        field(app, 'Requested item').set_value('Synthetic item to remove.').run()
        assert field(app, 'Requested item').proto.id == identity
        assert current(app, key).rows[0][1] == 'Synthetic item to remove.'
        button(app, 'Close editing').click().run()
        button(app, 'Remove request').click().run()
    else:
        identity = grid(app, key)
        columns = tuple(app.session_state[identity + '_data'][0])
        edit_grid(app, identity, edited={0: {columns[1]: '   '}})
        assert not app.exception and not current(app, key).rows
        edit_grid(app, identity, edited={0: {columns[1]: 'Synthetic item to remove.'}})
        assert current(app, key).rows[0][1] == 'Synthetic item to remove.'
        assert grid(app, key) == identity
        edit_grid(app, identity, removed=[0])
    assert not app.exception and not current(app, key).rows
    draft = app.session_state['draft']
    kept = active.save('e' * 32, draft, expected_revision=0, base_snapshot_revision=0)
    loaded = active.load('e' * 32, draft.profile.contract, draft.profile.key, draft.period)
    assert loaded.draft == draft and loaded.revision == kept.revision
    output_draft = replace(draft, sections=tuple(replace(section, included=any(spec.key == key for spec in section.blocks))
                                                for section in draft.sections))
    # This intentionally empty, included optional section has only its honest
    # empty-section qualification. No site/date/source blocker is waived.
    assert all(check.code == 'empty' and not check.blocking for check in preflight(output_draft))
    document = Document(BytesIO(assemble_docx(output_draft, acknowledged_fingerprint=output_draft.fingerprint)))
    table_text = '\n'.join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
    assert 'Synthetic item to remove.' not in table_text
    assert 'None' not in table_text
    assert EMPTY_TABLES[key][0] not in table_text


@pytest.mark.parametrize('key', ('capital_renewal', 'end_of_life'))
def test_omitted_imported_empty_capital_frame_keeps_schema_until_deliberate_entry(key):
    app = workspace(key)
    original = app.session_state['draft']
    frame = ReportTable(('Source identifier', 'Asset / recommendation', 'Timing', 'Recorded status'), (('', '', '', ''),), 'native-source-frame')
    incoming = ResolvedBlock(key, 'Omit', extra_tables=(frame,), references=('synthetic-retained-source',))
    app.session_state['draft'] = replace(original, blocks=tuple(incoming if block.key == key else block for block in original.blocks))
    app.run()
    assert not app.dataframe and current(app, key) == incoming
    open_entry(app, key)
    identity = grid(app, key)
    assert current(app, key) == incoming
    edit_grid(app, identity, edited={0: {'Timing': '   '}})
    assert current(app, key) == incoming  # Whitespace never reincludes the omitted frame.
    edit_grid(app, identity, edited={0: {'Source identifier': '0007', 'Asset / recommendation': 'Replace motor', 'Timing': 'Unknown'}})
    result = current(app, key)
    assert not app.exception and result.source == 'This month'
    assert result.extra_tables[0].columns == frame.columns
    assert result.extra_tables[0].reference == frame.reference
    assert result.extra_tables[0].rows == (('0007', 'Replace motor', 'Unknown', ''),)
    assert result.references == incoming.references and grid(app, key) == identity
