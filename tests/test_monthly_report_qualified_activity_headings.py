"""Synthetic subsection headings must not inherit earlier improvement routing."""

import hashlib

from docx import Document
import pytest

from app.monthly_report_import import inspect_docx


@pytest.mark.parametrize("heading", [
    "WORK ORDER STATUS — SITE B", "PM STATUS - LOCATION #2",
    "Work Order Status Location#2", "PM Status Site 12:",
    "WORK ORDER STATUS", "PM STATUS",
])
def test_qualified_grid_after_improvements_keeps_work_order_destination(tmp_path, heading):
    document = Document()
    document.add_heading("Monthly Activity Summary", 1)
    document.add_paragraph("Work Order Status")
    first = document.add_table(rows=2, cols=2)
    first.cell(0, 0).text, first.cell(1, 0).text = "Month", "Synthetic first site"
    document.add_paragraph("Improvement Highlights")
    document.add_paragraph("Synthetic completed repair.")
    document.add_paragraph(heading)
    second = document.add_table(rows=2, cols=2)
    second.cell(0, 0).text, second.cell(1, 0).text = "Month", "Synthetic second site"
    path = tmp_path / "synthetic-qualified-headings.docx"
    document.save(path)
    original = hashlib.sha256(path.read_bytes()).hexdigest()

    items = inspect_docx(path).items
    tables = [item for item in items if item.kind == "table"]
    assert [item.suggested_slot for item in tables] == ["work_orders", "work_orders"]
    assert tables[0].rows[1][0] == "Synthetic first site"
    assert tables[1].rows[1][0] == "Synthetic second site"
    assert next(item for item in items if item.text == "Synthetic completed repair.").suggested_slot == "improvements"
    assert hashlib.sha256(path.read_bytes()).hexdigest() == original


@pytest.mark.parametrize("text", [
    "Work order status for Site B remains under review.",
    "PM status Site B is awaiting review.",
    "We discussed work order status — Site B.",
])
def test_narrative_does_not_change_subsection(tmp_path, text):
    document = Document()
    document.add_heading("Monthly Activity Summary", 1)
    document.add_paragraph("Improvement Highlights")
    document.add_paragraph(text)
    document.add_table(rows=1, cols=1).cell(0, 0).text = "Synthetic improvement"
    path = tmp_path / "synthetic-narrative.docx"
    document.save(path)
    table = next(item for item in inspect_docx(path).items if item.kind == "table")
    assert table.suggested_slot == "improvements"
