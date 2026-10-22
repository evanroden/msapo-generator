"""Synthetic heading/list roles must not turn a rewritten narrative all bold."""
from copy import deepcopy
from dataclasses import replace
from io import BytesIO

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

from app.monthly_report_import import ImportMapping, inspect_docx, map_items
from app.monthly_report_native_layout import build_native_docx
from app.monthly_report_native_text import narrative_prototype
from tests.test_monthly_report_narrative_typography import scene


def hierarchy(tmp_path, *, underlined=False, fragmented=False):
    source, draft = scene(tmp_path)
    document = Document(source)
    body = document.paragraphs[-1]
    heading = document.add_paragraph('Synthetic Operations')
    level = OxmlElement('w:outlineLvl'); level.set(qn('w:val'), '2')
    heading._p.get_or_add_pPr().append(level)
    heading.runs[0].bold = True
    title = document.add_paragraph('Campus Boiler Hub')  # A synthetic 17-character title.
    title.paragraph_format.left_indent = Pt(6)
    title.paragraph_format.space_after = Pt(8)
    title.runs[0].bold = True
    title.runs[0].font.size = Pt(14)
    title._p.get_or_add_pPr().append(deepcopy(title.runs[0]._r.rPr))
    body._p.addprevious(heading._p)
    body._p.addprevious(title._p)
    second = document.add_paragraph('A second synthetic inspection was completed and reviewed.')
    for paragraph in (body, second):
        paragraph.paragraph_format.left_indent = Pt(72)
        paragraph.paragraph_format.space_after = Pt(4)
        numbering = paragraph._p.get_or_add_pPr().get_or_add_numPr()
        numbering.get_or_add_ilvl().val = 0
        numbering.get_or_add_numId().val = 1
        run = paragraph.runs[0]
        run.font.size = Pt(12)
        run.underline = underlined
        # Complex-script weight is not Latin weight.
        run._r.get_or_add_rPr().append(OxmlElement('w:bCs'))
        run._r.get_or_add_rPr().append(OxmlElement('w:iCs'))
        paragraph._p.get_or_add_pPr().append(deepcopy(run._r.rPr))
    if fragmented:
        body.runs[0].text = 'Plain opening. '
        for text in ('plain middle. ', 'plain ending. '):
            body.add_run(text).font.size = Pt(12)
        emphasis = body.add_run('An emphasized inline phrase.')
        emphasis.bold = emphasis.underline = True
        emphasis.font.size = Pt(12)
    document.save(source)
    return source, draft, document, heading, title, body, second


@pytest.mark.parametrize('section', [None, 'activity'])
@pytest.mark.parametrize('underlined', [False, True])
@pytest.mark.parametrize('fragmented', [False, True])
def test_body_typography_does_not_import_subheading_weight_or_old_list_layout(
        tmp_path, section, underlined, fragmented):
    source, draft, *_ = hierarchy(tmp_path, underlined=underlined, fragmented=fragmented)
    before = source.read_bytes()
    output = Document(BytesIO(build_native_docx(draft, source, section_key=section)))
    paragraphs = [p for p in output.paragraphs if p.text.startswith('Current synthetic')]
    assert len(paragraphs) == 2
    for paragraph in paragraphs:
        assert paragraph.runs[0].bold is not True
        assert paragraph.runs[0].underline is underlined
        assert paragraph.runs[0].font.size == Pt(12)
        assert paragraph.paragraph_format.left_indent == Pt(6)
        # Existing page reflow normalizes all rewritten paragraphs to four points.
        assert paragraph.paragraph_format.space_after == Pt(4)
        assert not any(n.get(qn('w:val')) != '0' for n in paragraph._p.iter(qn('w:numId')))
        assert paragraph._p.find('./' + qn('w:pPr') + '/' + qn('w:rPr') + '/' + qn('w:b')) is None
    assert source.read_bytes() == before


@pytest.mark.parametrize('case', [
    'sentence', 'quoted_sentence', 'bracketed_sentence', 'ellipsis', 'colon',
    'unpunctuated_prose', 'single_word', 'all_caps',
    'mixed', 'numbered', 'no_outline', 'table', 'unlisted', 'inactive',
    'one_list_item', 'different_list', 'different_level', 'different_style', 'bold_body',
])
def test_ambiguous_or_missing_structure_keeps_existing_bold_typography(tmp_path, case):
    source, draft, document, heading, title, body, second = hierarchy(tmp_path)
    values = {'sentence': 'Inspection complete.', 'quoted_sentence': '“Inspection Completed.”',
        'bracketed_sentence': '[Inspection Completed.]', 'ellipsis': 'Inspection Completed…',
        'colon': 'The following work was completed:',
        'unpunctuated_prose': 'Work complete', 'single_word': 'Done', 'all_caps': 'WORK COMPLETE'}
    if case in values:
        title.runs[0].text = values[case]
    elif case == 'mixed':
        title.add_run(' with ordinary prose').bold = False
    elif case == 'numbered':
        title._p.get_or_add_pPr().get_or_add_numPr().get_or_add_numId().val = 1
    elif case == 'no_outline':
        heading._p.get_or_add_pPr().remove(heading._p.find('./' + qn('w:pPr') + '/' + qn('w:outlineLvl')))
    elif case == 'table':
        table = document.add_table(rows=1, cols=1)
        table.cell(0, 0).text = 'Synthetic boundary'
        title._p.addnext(table._tbl)
    elif case in ('unlisted', 'inactive'):
        for paragraph in (body, second):
            properties = paragraph._p.get_or_add_pPr()
            if case == 'unlisted':
                properties.remove(properties.numPr)
            else:
                properties.numPr.get_or_add_numId().val = 0
    elif case == 'one_list_item':
        second._p.getparent().remove(second._p)
    elif case == 'different_list':
        second._p.pPr.numPr.numId.val = 2
    elif case == 'different_level':
        second._p.pPr.numPr.ilvl.val = 1
    elif case == 'different_style':
        body.style = second.style = 'Body Text'
    elif case == 'bold_body':
        body.runs[0].bold = second.runs[0].bold = True
    document.save(source)
    before = source.read_bytes()
    output = Document(BytesIO(build_native_docx(draft, source, section_key='activity')))
    current = [p for p in output.paragraphs if p.text.startswith('Current synthetic')]
    assert len(current) == 2
    assert all(p.runs[0].bold is True for p in current)
    assert source.read_bytes() == before


@pytest.mark.parametrize('boundary', ['section', 'slot', 'word_section', 'absent_scope'])
def test_typography_context_cannot_cross_role_boundaries(tmp_path, boundary):
    _, _, document, heading, title, body, second = hierarchy(tmp_path)
    styles = {s.style_id: s.element for s in document.styles}
    scopes = {p._p: ('activity', 'activity_summary') for p in (heading, title, body, second)}
    if boundary == 'word_section':
        heading._p.get_or_add_pPr().append(deepcopy(document.sections[-1]._sectPr))
    elif boundary == 'absent_scope':
        scopes = None
    else:
        scopes[body._p] = ('other', 'activity_summary') if boundary == 'section' else ('activity', 'other')
    result = narrative_prototype(title._p, styles, set(), scopes=scopes)
    assert result.find('./' + qn('w:r') + '/' + qn('w:rPr') + '/' + qn('w:b')) is not None


def test_prototype_is_identical_before_and_after_previous_heading_text_is_cleared(tmp_path):
    _, _, document, heading, title, body, second = hierarchy(tmp_path)
    styles = {s.style_id: s.element for s in document.styles}
    scopes = {p._p: ('activity', 'activity_summary') for p in (heading, title, body, second)}
    prototype = narrative_prototype(title._p, styles, set(), scopes=scopes)
    assert prototype.pPr.spacing.get(qn('w:after')) == '160'
    before = prototype.xml
    heading.runs[0].text = ''
    assert narrative_prototype(title._p, styles, set(), scopes=scopes).xml == before


def test_author_written_bullets_and_numbers_remain_text_without_source_list_markers(tmp_path):
    source, draft, *_ = hierarchy(tmp_path)
    text = 'Plain current introduction.\n\n• User-authored item.\n1. User-authored numbered note.'
    draft = replace(draft, blocks=(replace(draft.blocks[0], text=text),))
    output = Document(BytesIO(build_native_docx(draft, source, section_key='activity')))
    for line in filter(None, text.splitlines()):
        paragraph = next(p for p in output.paragraphs if p.text == line)
        assert paragraph.runs[0].bold is not True
        assert not any(n.get(qn('w:val')) != '0' for n in paragraph._p.iter(qn('w:numId')))


def test_proven_unchanged_hierarchy_preserves_source_heading_inline_and_list_xml(tmp_path):
    source, draft, *_ = hierarchy(tmp_path, underlined=True, fragmented=True)
    before = source.read_bytes()
    inspection = inspect_docx(source)
    mappings = tuple(ImportMapping(i.id, 'activity_summary') for i in inspection.items
        if i.kind == 'text' and i.section == 'activity' and i.note != 'Native section heading')
    mapped = map_items(source, inspection, mappings)
    output = Document(BytesIO(build_native_docx(replace(draft, blocks=mapped.blocks), source, section_key='activity')))
    for original in Document(source).paragraphs:
        if not original.text or original.text == 'MONTHLY ACTIVITY SUMMARY':
            continue
        matches = [p for p in output.paragraphs if p.text == original.text]
        assert len(matches) == 1 and matches[0]._p.xml == original._p.xml
    assert source.read_bytes() == before
