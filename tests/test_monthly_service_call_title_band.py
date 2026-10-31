"""Synthetic merged service-call titles must not become editable columns."""
from dataclasses import replace
from io import BytesIO

import pytest
from docx import Document

from app.monthly_report_import import inspect_docx, map_items, ImportMapping, imported_draft
from app.monthly_report_model import ReportPeriod, synthetic_profiles
from app.monthly_report_sections import table_without_prices, section_reviews, build_section_import
from app.monthly_report_docx import _build_docx
from app.monthly_report_native_layout import build_native_docx


HEADERS = ("WO #", "Finish Date", "Area", "Task Code-Description", "Tag #")
RECORDS = (("1001", "10/05/2026", "Plant A", "Inspect pump", "P-1"),
           ("1002", "10/06/2026", "Plant B", "Test valve", "V-2"))


def source(tmp_path, *, merged=True, title="Service calls", headers=HEADERS, priced=False):
    document = Document()
    document.add_heading("5 Maintenance Schedule / In-House Maintenance", 1)
    labels = (*headers, "Cost") if priced else headers
    table = document.add_table(rows=4, cols=len(labels))
    if merged:
        table.cell(0, 0).merge(table.cell(0, len(labels) - 1)).text = title
    else:
        table.cell(0, 0).text = title
    for cell, value in zip(table.rows[1].cells, labels):
        cell.text = value
    for row, record in zip(table.rows[2:], RECORDS):
        for cell, value in zip(row.cells, (*record, "500") if priced else record):
            cell.text = value
    document.add_paragraph("After table")
    path = tmp_path / "synthetic-service-calls.docx"
    document.save(path)
    return path


@pytest.mark.parametrize("priced", [False, True])
def test_merged_title_uses_real_headers_and_keeps_all_records(tmp_path, priced):
    path = source(tmp_path, priced=priced)
    original = path.read_bytes()
    inspection = inspect_docx(path)
    item = next(i for i in inspection.items if i.kind == "table")
    table, removed = table_without_prices(item)
    assert table.columns == HEADERS
    assert table.rows == RECORDS
    assert removed == ((5,) if priced else ())
    # Audit rows and original package are immutable; title is separate context.
    assert item.rows[0][0] == "Service calls" and len(item.rows) == 4
    title = next(i for i in inspection.items if i.text == "Service calls")
    assert (title.position, title.suggested_slot) == (item.position, "service_calls")
    assert path.read_bytes() == original


@pytest.mark.parametrize("options", [
    {"merged": False}, {"title": "Work completed"},
    {"headers": ("WO #", "Finish Date", "Area", "Work performed", "Unclear")},
])
def test_ambiguous_title_is_not_promoted_to_a_header_schema(tmp_path, options):
    path = source(tmp_path, **options)
    item = next(i for i in inspect_docx(path).items if i.kind == "table")
    table, _ = table_without_prices(item)
    assert table.columns[0] == options.get("title", "Service calls")
    assert table.rows[0] == options.get("headers", HEADERS)
    assert table.rows[-2:] == RECORDS


def test_ordinary_header_and_following_item_identity_are_unchanged(tmp_path):
    path = source(tmp_path)
    before = inspect_docx(path)
    document = Document(path)
    table = document.tables[0]
    table._tbl.remove(table.rows[0]._tr)
    document.save(path)
    after = inspect_docx(path)
    normal = next(i for i in after.items if i.kind == "table")
    assert table_without_prices(normal)[0].columns == HEADERS
    assert table_without_prices(normal)[0].rows == RECORDS
    assert next(i.id for i in before.items if i.text == "After table") == next(
        i.id for i in after.items if i.text == "After table")


def test_explicit_mapping_honors_no_header_and_preserves_title_context(tmp_path):
    path = source(tmp_path)
    inspection = inspect_docx(path)
    item = next(i for i in inspection.items if i.kind == "table")
    mapped = map_items(path, inspection, (ImportMapping(item.id, "service_calls"),))
    assert tuple(c.title for c in mapped.overrides[0].columns) == HEADERS
    assert mapped.blocks[0].rows == RECORDS
    assert mapped.blocks[0].text == "Service calls"
    raw = map_items(path, inspection, (ImportMapping(item.id, "service_calls", first_row_header=False),))
    assert raw.blocks[0].rows == item.rows
    assert not raw.blocks[0].text


@pytest.mark.parametrize("native", [False, True])
@pytest.mark.parametrize("priced,edited", [(False, False), (True, False), (False, True)])
def test_section_import_exports_title_and_records_once(tmp_path, monkeypatch, native, priced, edited):
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path / "data"))
    path = source(tmp_path, priced=priced)
    inspection = inspect_docx(path)
    section = next(s for s in section_reviews(inspection) if s.key == "maintenance")
    selected = [i.id for i in section.items if i.kind == "table" or i.text == "Service calls"]
    mapped, _ = build_section_import(path, inspection, ({
        "key": "maintenance", "target": "maintenance", "approved": True, "selected": selected,
    },))
    draft = imported_draft(synthetic_profiles()[0], ReportPeriod(2026, 10), "Synthetic Editor", mapped)
    block = next(b for b in draft.blocks if b.key == "service_calls")
    assert block.text == "Service calls"
    assert block.extra_tables[0].columns == HEADERS
    assert block.extra_tables[0].rows == RECORDS
    if edited:
        table = replace(block.extra_tables[0], rows=(RECORDS[0], (*RECORDS[1][:3], "Recheck valve", "V-2")))
        draft = replace(draft, blocks=tuple(replace(b, extra_tables=(table,)) if b.key == block.key else b
                                          for b in draft.blocks))
    # Inspect actual Word output on both the generated and source-preserving path.
    raw = build_native_docx(draft, path) if native else _build_docx(draft)
    output = Document(BytesIO(raw))
    texts = [node.text for node in output.element.iter() if node.tag.endswith("}t")]
    assert texts.count("Service calls") == 1
    assert texts.count("1001") == texts.count("1002") == 1
    assert "Column 2" not in texts
    assert "500" not in texts and "Cost" not in texts
    if edited:
        assert "Recheck valve" in texts and "Test valve" not in texts
