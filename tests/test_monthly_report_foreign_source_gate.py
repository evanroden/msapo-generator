"""A warning acknowledgment cannot authorize another site's included evidence."""
from dataclasses import replace
from datetime import date
from io import BytesIO

from docx import Document
import pytest

from app.monthly_report_checks import preflight
from app.monthly_report_docx import assemble_docx
from app.monthly_report_guided import generation_block_reason
from app.monthly_report_model import (
    DraftParagraph, ReportFollowUp, ReportPeriod, ReportTable, synthetic_draft, synthetic_profiles,
)
from app.monthly_report_source_details import correct_source_details
from app.monthly_report_sources import ingest, source_bytes, source_reference


@pytest.fixture
def report(monkeypatch, tmp_path):
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path))
    profile = synthetic_profiles()[0]
    original = b"Inspected the demonstration pump. Follow-up remains pending."
    content, _ = ingest(profile, "synthetic-vendor.txt", original)
    source = replace(content.source, classification="Vendor service", service_date="2026-09-16",
                     facility="Demonstration South Facility")
    draft = synthetic_draft(profile, ReportPeriod(2026, 9))
    index = next(i for i, block in enumerate(draft.blocks) if block.key == "demo_activity")
    block = replace(draft.blocks[index], text="Inspected the demonstration pump.",
                    references=(source_reference(source, 1),))
    draft = replace(draft, blocks=(*draft.blocks[:index], block, *draft.blocks[index + 1:]),
                    sources=(source,))
    return draft, index, original


@pytest.mark.parametrize("classification,selected", [
    ("Vendor service", (1,)), ("Reference only", (1,)), ("Vendor service", ()),
])
def test_known_foreign_evidence_blocks_ui_and_actual_word_even_after_acknowledgment(report, classification, selected):
    draft, _, _ = report
    draft = replace(draft, sources=(replace(draft.sources[0], classification=classification,
                                           selected_pages=selected),))
    with pytest.raises(ValueError, match="synthetic-vendor.txt"):
        assemble_docx(draft, acknowledged_fingerprint=draft.fingerprint)
    checks = preflight(draft)
    matching = [c for c in checks if c.code == "source_facility"]
    assert len(matching) == 1 and matching[0].blocking
    assert matching[0].block_key == draft.sources[0].id
    assert "synthetic-vendor.txt" in matching[0].message
    assert generation_block_reason(conflict=None, checks=checks, warnings_ok=True)


def test_populated_imported_table_is_also_included_evidence(report):
    draft, index, _ = report
    block = replace(draft.blocks[index], references=(),
                    extra_tables=(ReportTable(("Finding",), (("Bearing inspection",),),
                                              source_reference(draft.sources[0], 1)),))
    draft = replace(draft, blocks=(*draft.blocks[:index], block, *draft.blocks[index + 1:]))
    assert any(c.code == "source_facility" and c.blocking for c in preflight(draft))


@pytest.mark.parametrize("kind", ["narrative", "table", "AI paragraph", "issue"])
def test_removed_source_cannot_authorize_orphaned_included_evidence(report, kind):
    draft, index, _ = report
    reference = source_reference(draft.sources[0], 1)
    block = draft.blocks[index]
    if kind == "table":
        block = replace(block, references=(), extra_tables=(ReportTable(
            ("Finding",), (("Bearing inspection",),), reference),))
    elif kind == "AI paragraph":
        block = replace(block, references=(), ai_paragraphs=(DraftParagraph(block.text, (), (reference,)),))
    elif kind == "issue":
        block = replace(block, references=())
        draft = replace(draft, follow_ups=(ReportFollowUp("synthetic issue", "issue", "Pump inspection",
            "2026-09", references=(reference,)),))
    draft = replace(draft, sources=(), blocks=(*draft.blocks[:index], block, *draft.blocks[index + 1:]))
    checks = preflight(draft)
    missing = [check for check in checks if check.code == "source_missing"]
    assert len(missing) == 1 and missing[0].blocking
    assert generation_block_reason(conflict=None, checks=checks, warnings_ok=True)
    with pytest.raises(ValueError, match="source file that is no longer"):
        assemble_docx(draft, acknowledged_fingerprint=draft.fingerprint)
    # Deliberately omitted report content keeps its evidence without blocking.
    omitted = replace(draft, blocks=tuple(replace(b, source="Omit") for b in draft.blocks),
        follow_ups=tuple(replace(item, included=False) for item in draft.follow_ups))
    assert not any(check.code == "source_missing" for check in preflight(omitted))


@pytest.mark.parametrize("mode", ["unused", "omitted", "section omitted", "route only", "empty table", "cleared"])
def test_unused_or_removed_evidence_does_not_create_a_required_site_block(report, mode):
    draft, index, _ = report
    block = draft.blocks[index]
    if mode == "unused":
        block = replace(block, references=(), text="Human note: " + draft.sources[0].id)
    elif mode == "omitted":
        block = replace(block, source="Omit")
    elif mode == "section omitted":
        draft = replace(draft, sections=tuple(replace(s, included=False) if s.key == "activity" else s
                                              for s in draft.sections))
    elif mode == "route only":
        block = replace(block, references=("section-source-route:v1:improvements:"
                        + draft.sources[0].id + ":" + draft.sources[0].sha256,))
    elif mode == "empty table":
        block = replace(block, references=(), extra_tables=(ReportTable(("Finding",), (),
                        source_reference(draft.sources[0], 1)),))
    else:
        block = replace(block, text="")
    draft = replace(draft, blocks=(*draft.blocks[:index], block, *draft.blocks[index + 1:]))
    assert not any(c.code == "source_facility" and c.blocking for c in preflight(draft))


def test_targeted_site_correction_clears_gate_without_rewriting_the_original(report):
    draft, _, original = report
    source = correct_source_details(draft.sources[0], draft.profile, date(2026, 9, 16), "north")
    corrected = replace(draft, sources=(source,))
    assert not any(c.code == "source_facility" for c in preflight(corrected))
    assert source_bytes(draft.profile, source) == original
    assert source.id == draft.sources[0].id and source.sha256 == draft.sources[0].sha256
    docx = assemble_docx(corrected, acknowledged_fingerprint=corrected.fingerprint)
    text = "\n".join(p.text for p in Document(BytesIO(docx)).paragraphs)
    assert "Inspected the demonstration pump." in text


@pytest.mark.parametrize("facility", ["north", "Demonstration North Facility", "Demo North"])
def test_actual_site_identifiers_and_aliases_remain_valid(report, facility):
    draft, _, _ = report
    draft = replace(draft, sources=(replace(draft.sources[0], facility=facility),))
    assert not any(c.code == "source_facility" for c in preflight(draft))


def test_combined_site_scope_does_not_treat_its_second_site_as_foreign(report):
    draft, _, _ = report
    draft = replace(draft, profile=synthetic_profiles()[1])
    assert not any(c.code == "source_facility" for c in preflight(draft))


def test_source_scope_blocker_survives_snapshot_and_correction_clears_only_its_item(report):
    from app import monthly_report_library as library

    draft, index, _ = report
    extra, _ = ingest(draft.profile, "synthetic-other-vendor.txt",
                      b"A separate synthetic pump inspection remains unresolved.")
    other = replace(extra.source, facility="Another synthetic facility",
                    classification="Vendor service", service_date="2026-09-17")
    block = replace(draft.blocks[index], references=(*draft.blocks[index].references,
                    source_reference(other, 1)))
    draft = replace(draft, sources=(*draft.sources, other),
                    blocks=(*draft.blocks[:index], block, *draft.blocks[index + 1:]))
    library.save_profile(draft.profile, expected_revision=0, actor="Synthetic Editor", confirmed=True)
    library.save_snapshot(draft, expected_revision=0, entered_editor="Synthetic Editor")
    loaded = library.load_snapshot(draft.profile.contract, draft.profile.key, draft.period).draft
    assert loaded == draft
    source = correct_source_details(loaded.sources[0], loaded.profile, date(2026, 9, 16), "north")
    corrected = replace(loaded, sources=(source, other))
    matching = [c for c in preflight(corrected) if c.code == "source_facility" and c.blocking]
    assert [c.block_key for c in matching] == [other.id]
    with pytest.raises(ValueError, match="synthetic-other-vendor.txt"):
        assemble_docx(corrected, acknowledged_fingerprint=corrected.fingerprint)
    # Correcting this browser's working source did not rewrite shared history.
    assert library.load_snapshot(draft.profile.contract, draft.profile.key, draft.period).draft == draft


def test_unresolved_site_remains_inspectable_in_private_watermarked_preview(report, monkeypatch, tmp_path):
    import fitz
    from app import monthly_report_preview as preview

    draft, _, _ = report
    converted_text = []
    def convert(path):
        converted_text.append("\n".join(p.text for p in Document(path).paragraphs))
        output = tmp_path / "synthetic-site-review.pdf"
        with fitz.open() as pdf:
            pdf.new_page().insert_text((72, 72), "Synthetic source correction preview")
            pdf.save(output)
        return output
    monkeypatch.setattr(preview.pdf_converter, "convert_to_pdf", convert)
    result = preview.preview_report(draft)
    assert "Inspected the demonstration pump." in converted_text[0]
    with fitz.open(stream=result.pdf, filetype="pdf") as pdf:
        assert "DRAFT - REVIEW ONLY" in pdf[0].get_text()
        assert "NOT A FINISHED CLIENT REPORT" in pdf[0].get_text()
    with pytest.raises(ValueError, match="synthetic-vendor.txt"):
        assemble_docx(draft, acknowledged_fingerprint=draft.fingerprint)


@pytest.mark.parametrize("kind", ["AI paragraph", "included picture", "included issue", "removed picture", "excluded issue"])
def test_provenance_fields_follow_actual_inclusion(report, kind):
    draft, index, _ = report
    reference = source_reference(draft.sources[0], 1)
    block = replace(draft.blocks[index], references=())
    if kind == "AI paragraph":
        block = replace(block, ai_paragraphs=(DraftParagraph(block.text, (), (reference,)),))
    elif "picture" in kind:
        block = replace(block, asset_hashes=("included.png",) if kind == "included picture" else (),
                        asset_provenance=(("included.png", (reference,)),))
    else:
        draft = replace(draft, follow_ups=(ReportFollowUp("synthetic issue", "issue", "Pump inspection",
                        "2026-09", references=(reference,), included=kind == "included issue"),))
    draft = replace(draft, blocks=(*draft.blocks[:index], block, *draft.blocks[index + 1:]))
    assert any(c.code == "source_facility" and c.blocking for c in preflight(draft)) == (
        kind in {"AI paragraph", "included picture", "included issue"})


@pytest.mark.parametrize("classification,selected", [("Reference only", (1,)), ("Vendor service", ())])
def test_resumed_foreign_source_keeps_one_native_correction_path(monkeypatch, tmp_path, classification, selected):
    from streamlit.testing.v1 import AppTest
    from app import monthly_report_guided as guided, monthly_report_library as library
    from app.contracts import RRH_CONTRACT
    from app.monthly_report_model import ResolvedBlock
    from test_monthly_report_ui import ROOT, choose_report, monthly

    monthly(monkeypatch, tmp_path)
    profile = library.list_profiles(RRH_CONTRACT)[0]
    period = ReportPeriod(2026, 9)
    raw = b"Synthetic inspection; follow-up remains pending."
    content, _ = ingest(profile, "synthetic-foreign-source.txt", raw)
    foreign = replace(content.source, facility="Another synthetic facility", service_date="2026-09-16",
                      classification=classification, selected_pages=selected)
    state = library.load_profile(profile.contract, profile.key)
    draft = guided._initial_draft(state, period, "Synthetic Editor")
    text = "Human-approved synthetic inspection note."
    block = ResolvedBlock("activity_summary", "This month", text=text,
                          references=(source_reference(foreign, 1),))
    draft = replace(draft, sources=(foreign,), blocks=tuple(b for b in draft.blocks if b.key != block.key) + (block,))
    library.save_snapshot(draft, expected_revision=0, entered_editor="Synthetic Editor")

    app = AppTest.from_file(ROOT / "run_web.py", default_timeout=20).run()
    app.segmented_control[0].set_value("Monthly report").run()
    choose_report(app, synthetic_profiles()[0].facilities[0].title)
    assert not app.exception
    panels = [w for w in app.expander if w.label.startswith("Correct source details · synthetic-foreign-source.txt")]
    assert len(panels) == 1 and "site does not match" in panels[0].label
    assert len([w for w in app.date_input if w.label == "Source date"]) == 1
    next(w for w in app.button if w.label == "Correct site").click().run()
    assert next(w for w in app.selectbox if w.label == "File to review").value == foreign.id
    next(w for w in app.selectbox if w.label == "Source facility").set_value(profile.facilities[0].key)
    next(w for w in app.button if w.label == "Use these source details").click().run()
    assert not app.exception
    prefix = "report_guided_" + guided._signature((RRH_CONTRACT, profile.key, period.key))
    corrected = app.session_state[prefix + "_draft"]
    assert corrected.sources[0].facility == profile.facilities[0].key
    assert not any(c.code == "source_facility" for c in preflight(corrected))
    assert next(b for b in corrected.blocks if b.key == "activity_summary").text == text
    assert source_bytes(profile, corrected.sources[0]) == raw
    assert library.load_snapshot(profile.contract, profile.key, period).draft.sources == (foreign,)
    assert any(w.label == "Done with source files" for w in app.button)
    assert not app.exception
    remove = next(w for w in app.button if w.label == "Remove source from this draft")
    assert remove.disabled
    assert "linked report content" in remove.help
