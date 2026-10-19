"""Synthetic cover identity and branding regressions from the October 19 review.

These exercise product import, persisted artwork and editable DOCX output.
They do not certify the production master, physical Safari or Microsoft Word.
"""

from dataclasses import replace
from io import BytesIO
from zipfile import ZipFile

import pytest
from PIL import Image, ImageDraw
from docx import Document
from docx.enum.section import WD_SECTION
from docx.shared import Inches

from app import monthly_report_library as library
from app.monthly_report_asset_review import approve_all_assets
from app.monthly_report_cover_ui import replace_cover_picture
from app.monthly_report_docx import assemble_docx, normalize_report_image
from app.monthly_report_import import ImportMapping, inspect_docx, map_items
from app.monthly_report_model import ReportDraft, ReportPeriod, ResolvedBlock, default_sections, synthetic_profiles
from app.monthly_report_sections import build_section_import, section_reviews
from app.monthly_report_start_import import (
    analyze_starting_report, automatic_section_plans, cover_photo_candidates,
    suggested_cover_photo, unresolved_cover_artwork,
)


def mark(color=(25, 120, 190, 255)):
    image = Image.new("RGBA", (420, 140), (0, 0, 0, 0))
    draw = ImageDraw.Draw(image)
    draw.rectangle((40, 30, 170, 110), fill=color)
    # White is intentional brand content, not a removable background.
    draw.rectangle((220, 30, 360, 110), fill=(255, 255, 255, 255))
    output = BytesIO()
    image.save(output, "PNG")
    return output.getvalue()


def photograph(size=(720, 480)):
    image = Image.new("RGB", size, (48, 112, 92))
    draw = ImageDraw.Draw(image)
    draw.ellipse((size[0] // 4, size[1] // 4, size[0] // 4 + size[1] // 2, 3 * size[1] // 4), fill=(220, 180, 60))
    output = BytesIO()
    image.save(output, "PNG")
    return output.getvalue()


def source_document(tmp_path, *, break_kind=None, contents=False, body_photo=False,
                    first_photo=False, later_photo=False):
    doc = Document()
    doc.add_paragraph(synthetic_profiles()[0].facilities[0].title)
    doc.add_paragraph("November 2026")
    if first_photo:
        doc.add_picture(BytesIO(photograph()), width=Inches(4))
    if break_kind is not None:
        doc.add_section(break_kind)
    if contents:
        doc.add_paragraph("Table of Contents")
    if body_photo:
        doc.add_heading("Monthly Activity Summary", 1)
        doc.add_paragraph("Inspected the fan drive.")
    doc.add_picture(BytesIO(photograph((1000, 700))), width=Inches(6))
    if later_photo:
        doc.add_section(WD_SECTION.NEW_PAGE)
        doc.add_picture(BytesIO(photograph((900, 500))), width=Inches(5))
    if not body_photo:
        doc.add_section(WD_SECTION.NEW_PAGE)
        doc.add_paragraph("Table of Contents")
        doc.add_heading("Monthly Activity Summary", 1)
        doc.add_paragraph("Inspected the fan drive.")
    path = tmp_path / "source.docx"
    doc.save(path)
    return path


@pytest.mark.parametrize("break_kind", [None, WD_SECTION.CONTINUOUS, WD_SECTION.NEW_PAGE])
def test_adjacent_cover_section_is_selected_and_maps_on_first_import(tmp_path, break_kind):
    path = source_document(tmp_path, break_kind=break_kind)
    before = path.read_bytes()
    inspection = inspect_docx(path)
    cover = next(s for s in section_reviews(inspection) if s.key == "cover")
    candidates = cover_photo_candidates(cover.items)
    assert len(candidates) == 1, "A title-only first Word section must not silently hide its adjacent cover picture"
    assert suggested_cover_photo(cover.items) == candidates[0]
    analysis = analyze_starting_report(inspection, synthetic_profiles()[0], ReportPeriod(2026, 11))
    assert analysis.site_relation == "same"
    plan = automatic_section_plans(inspection, analysis)["cover"]
    assert plan["approved"] and plan["destinations"][candidates[0].id] == "cover_photo"
    mapped, _ = build_section_import(path, inspection, (plan,))
    photo = next(b for b in mapped.blocks if b.key == "cover_photo")
    assert len(photo.asset_hashes) == 1
    assert photo.references == (f"docx:{inspection.sha256}:{candidates[0].id}",)
    assert path.read_bytes() == before


@pytest.mark.parametrize("options", [
    {"break_kind": WD_SECTION.NEW_PAGE, "contents": True},
    {"break_kind": WD_SECTION.NEW_PAGE, "body_photo": True},
])
def test_contents_and_activity_images_are_not_cover_candidates(tmp_path, options):
    inspection = inspect_docx(source_document(tmp_path, **options))
    assert not cover_photo_candidates(inspection.items)


def test_only_adjacent_cover_section_is_eligible_and_later_art_is_retained(tmp_path):
    inspection = inspect_docx(source_document(tmp_path, break_kind=WD_SECTION.NEW_PAGE, later_photo=True))
    images = [i for i in inspection.items if i.kind == "image"]
    assert len(images) == 2
    assert cover_photo_candidates(inspection.items) == (images[0],)
    assert images[1].word_section == 3


def test_existing_cover_photo_does_not_promote_next_image_only_section(tmp_path):
    inspection = inspect_docx(source_document(tmp_path, break_kind=WD_SECTION.NEW_PAGE, first_photo=True))
    images = [i for i in inspection.items if i.kind == "image"]
    assert cover_photo_candidates(inspection.items) == (images[0],)


@pytest.mark.parametrize("caption_position", ["before", "after", "alt"])
def test_explicit_company_and_client_logo_labels_are_not_photo_options(tmp_path, caption_position):
    doc = Document()
    doc.add_paragraph(synthetic_profiles()[0].facilities[0].title)
    doc.add_paragraph("November 2026")
    for label, color in (("Synthetic company logo", (40, 100, 160, 255)),
                         ("Synthetic client logo", (90, 160, 40, 255))):
        if caption_position == "before":
            doc.add_paragraph(label)
        shape = doc.add_picture(BytesIO(mark(color)), width=Inches(3))
        if caption_position == "after":
            doc.add_paragraph(label)
        if caption_position == "alt":
            shape._inline.docPr.set("descr", label)
    doc.add_paragraph("Facility photograph with the company logo visible on the building")
    doc.add_picture(BytesIO(photograph((160, 120))), width=Inches(3))
    path = tmp_path / "captioned.docx"
    doc.save(path)
    inspection = inspect_docx(path)
    images = [i for i in inspection.items if i.kind == "image"]
    assert [i.suggested_slot for i in images[:2]] == ["brand_logo", "client_logo"]
    assert cover_photo_candidates(inspection.items) == (images[2],)


def test_body_copy_of_identified_header_mark_cannot_become_photo(tmp_path):
    doc = Document()
    doc.add_paragraph(synthetic_profiles()[0].facilities[0].title)
    doc.add_paragraph("November 2026")
    raw = mark()
    doc.sections[0].header.paragraphs[0].add_run().add_picture(BytesIO(raw), width=Inches(2))
    doc.add_picture(BytesIO(raw), width=Inches(4))
    doc.add_picture(BytesIO(photograph()), width=Inches(5))
    path = tmp_path / "repeated-brand.docx"
    doc.save(path)
    inspection = inspect_docx(path)
    body = [i for i in inspection.items if i.kind == "image" and i.part == "word/document.xml"]
    assert body[0].suggested_slot == "brand_logo"
    assert cover_photo_candidates(inspection.items) == (body[1],)


def test_same_bytes_in_an_independent_body_chart_keep_that_section_role(tmp_path):
    doc = Document()
    doc.add_paragraph(synthetic_profiles()[0].facilities[0].title)
    raw = mark()
    doc.sections[0].header.paragraphs[0].add_run().add_picture(BytesIO(raw), width=Inches(2))
    doc.add_heading("Organizational Chart", 1)
    doc.add_picture(BytesIO(raw), width=Inches(4))
    path = tmp_path / "independent-body-art.docx"
    doc.save(path)
    inspection = inspect_docx(path)
    body = next(i for i in inspection.items if i.kind == "image" and i.part == "word/document.xml")
    assert body.suggested_slot == "org_chart" and not body.image_role
    assert not cover_photo_candidates(inspection.items)


def test_ambiguous_logo_caption_keeps_both_originals_and_a_specific_recovery(monkeypatch, tmp_path):
    from streamlit.testing.v1 import AppTest
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path / "data"))
    doc = Document()
    doc.add_paragraph(synthetic_profiles()[0].facilities[0].title)
    doc.add_paragraph("November 2026")
    doc.add_picture(BytesIO(photograph()), width=Inches(4))
    doc.add_paragraph("Logo")
    doc.add_picture(BytesIO(mark()), width=Inches(3))
    source = tmp_path / "ambiguous-caption.docx"
    doc.save(source)
    original = source.read_bytes()
    inspection = inspect_docx(source)
    cover = next(section for section in section_reviews(inspection) if section.key == "cover")
    images = tuple(item for item in cover.items if item.kind == "image")
    assert len(images) == 2 and all(item.image_role == "logo" for item in images)
    assert not cover_photo_candidates(cover.items)
    assert set(item.id for item in unresolved_cover_artwork(cover.items)) == set(item.id for item in images)
    app = AppTest.from_string('''
import streamlit as st
from app.monthly_report_section_ui import _section_card
from app.monthly_report_model import ReportPeriod
from app.monthly_report_ui import _field
plans=st.session_state.setdefault('plans', {})
_section_card(st.session_state['section'], st.session_state['source'], st.session_state['inspection'],
 ReportPeriod(2026,11), 'synthetic_ambiguous', _field, plans, True, cover_design_available=True)
''', default_timeout=20)
    app.session_state["section"] = cover
    app.session_state["source"] = source
    app.session_state["inspection"] = inspection
    app.run()
    assert not app.exception
    assert any("Other imported cover artwork" in entry.value for entry in app.markdown)
    assert any("Repair it in the original Word file and upload it again" in entry.value for entry in app.error)
    assert not app.session_state["plans"]["cover"]["approved"]
    assert not any(button.label.startswith("Choose picture") for button in app.button)
    action = next(widget for widget in app.radio if widget.label == "What would you like to do?")
    assert "Use this report’s cover design" in action.options
    action.set_value("Leave this section out").run()
    assert not app.exception
    assert app.session_state["plans"]["cover"]["approved"]
    assert not app.session_state["plans"]["cover"]["selected"]
    assert source.read_bytes() == original

    # A separate corrected copy removes the adjacency ambiguity. The genuine
    # photograph can then be used while the explicitly identified mark remains
    # excluded from both automatic and manual cover-photo destinations.
    corrected = Document()
    corrected.add_paragraph(synthetic_profiles()[0].facilities[0].title)
    corrected.add_paragraph("November 2026")
    corrected.add_picture(BytesIO(photograph()), width=Inches(4))
    corrected.add_paragraph("Facility exterior photograph")
    corrected.add_paragraph("Company logo")
    corrected.add_picture(BytesIO(mark()), width=Inches(3))
    repaired = tmp_path / "corrected-caption.docx"
    corrected.save(repaired)
    checked = inspect_docx(repaired)
    candidate, = cover_photo_candidates(checked.items)
    known_logo = next(item for item in checked.items if item.kind == "image" and item.suggested_slot == "brand_logo")
    assert candidate.image_digest == images[0].image_digest
    with pytest.raises(ValueError, match="identified as a logo"):
        map_items(repaired, checked, (ImportMapping(known_logo.id, "cover_photo"),))
    assert source.read_bytes() == original


@pytest.mark.parametrize("role", ["brand_logo", "client_logo"])
@pytest.mark.parametrize("route", ["mapping", "section"])
def test_known_logo_role_cannot_be_remapped_to_a_cover_photo(tmp_path, role, route):
    doc = Document()
    doc.add_paragraph(synthetic_profiles()[0].facilities[0].title)
    doc.add_picture(BytesIO(mark()), width=Inches(3))
    source = tmp_path / "known-logo.docx"
    doc.save(source)
    inspection = inspect_docx(source)
    image = next(i for i in inspection.items if i.kind == "image")
    # An already identified role, independently of new caption recognition.
    inspection = replace(inspection, items=tuple(replace(i, suggested_slot=role) if i.id == image.id else i
                                                for i in inspection.items))
    with pytest.raises(ValueError, match="identified as a logo"):
        if route == "mapping":
            map_items(source, inspection, (ImportMapping(image.id, "cover_photo"),))
        else:
            build_section_import(source, inspection, ({"key": "cover", "target": "cover", "approved": True,
                "selected": (image.id,), "destinations": {image.id: "cover_photo"}},))


def test_manual_import_lists_branding_destinations_without_a_cover_photo_option(tmp_path, monkeypatch):
    from streamlit.testing.v1 import AppTest
    from app import monthly_report_import_ui as import_ui
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path / "data"))
    profile = synthetic_profiles()[0]
    library.save_profile(profile, expected_revision=0, actor="Synthetic Reviewer", confirmed=True)
    doc = Document()
    doc.add_paragraph(profile.facilities[0].title)
    doc.sections[0].header.paragraphs[0].add_run().add_picture(BytesIO(mark()), width=Inches(2))
    doc.add_picture(BytesIO(photograph()), width=Inches(5))
    source = tmp_path / "logo-and-photo.docx"
    doc.save(source)
    upload = BytesIO(source.read_bytes())
    upload.name, upload.size = source.name, len(upload.getvalue())
    real_upload = import_ui.st.file_uploader
    def uploader(label, *args, **kwargs):
        native = real_upload(label, *args, **kwargs)
        return upload if label == "Existing monthly report DOCX" else native
    monkeypatch.setattr(import_ui.st, "file_uploader", uploader)
    app = AppTest.from_string('''
import streamlit as st
from app import monthly_report_library as library, monthly_report_import_ui as import_ui
from app.monthly_report_model import synthetic_profiles
profile = synthetic_profiles()[0]
def field(key, value):
    st.session_state.setdefault(key, value)
    return key
import_ui.render_import(library.load_profile(profile.contract, profile.key), field)
''').run()
    next(b for b in app.button if b.label == "Read DOCX for mapping").click().run()
    assert not app.exception
    inspection = next(v[2] for k, v in app.session_state.filtered_state.items() if k.endswith("_stage"))
    image = next(i for i in inspection.items if i.kind == "image" and i.suggested_slot == "brand_logo")
    next(w for w in app.selectbox if w.label == "Item to inspect").set_value(image.id).run()
    destination = next(w for w in app.selectbox if w.label == "Confirmed destination")
    assert "cover_photo" not in destination.options
    assert {"brand_logo", "client_logo"} <= set(destination.options)
    photo = next(i for i in inspection.items if i.kind == "image" and i.part == "word/document.xml")
    next(w for w in app.selectbox if w.label == "Item to inspect").set_value(photo.id).run()
    assert "cover_photo" in next(w for w in app.selectbox if w.label == "Confirmed destination").options


def test_unassigned_logo_is_recoverable_without_offering_it_as_photo(tmp_path):
    doc = Document()
    doc.add_paragraph(synthetic_profiles()[0].facilities[0].title)
    doc.add_paragraph("Logo")
    doc.add_picture(BytesIO(mark()), width=Inches(3))
    path = tmp_path / "unknown-logo.docx"
    doc.save(path)
    inspection = inspect_docx(path)
    image = next(i for i in inspection.items if i.kind == "image")
    assert not cover_photo_candidates(inspection.items)
    assert unresolved_cover_artwork(inspection.items) == (image,)
    plan = automatic_section_plans(inspection, analyze_starting_report(
        inspection, synthetic_profiles()[0], ReportPeriod(2026, 11)))["cover"]
    assert not plan["approved"] and plan["questions"]


def assert_mark_preserved(raw):
    image = Image.open(BytesIO(raw)).convert("RGBA")
    assert image.getpixel((0, 0))[3] == 0, "Transparent background was flattened"
    assert image.getpixel((260, 70)) == (255, 255, 255, 255), "Intentional white mark content changed"
    assert image.getpixel((80, 70)) == (25, 120, 190, 255), "Brand color changed"
    for background in ((25, 70, 150), (20, 100, 70), (0, 0, 0), (255, 255, 255)):
        composite = Image.new("RGBA", image.size, (*background, 255))
        composite.alpha_composite(image)
        assert composite.getpixel((0, 0)) == (*background, 255)
        assert composite.getpixel((260, 70)) == (255, 255, 255, 255)


@pytest.mark.parametrize("slot", ["brand_logo", "client_logo"])
def test_manual_logo_mapping_and_report_replacement_preserve_alpha(tmp_path, slot):
    raw = mark()
    doc = Document()
    doc.add_paragraph(synthetic_profiles()[0].facilities[0].title)
    doc.add_picture(BytesIO(raw), width=Inches(3))
    path = tmp_path / "manual-logo.docx"
    doc.save(path)
    inspection = inspect_docx(path)
    item = next(i for i in inspection.items if i.kind == "image")
    mapped = map_items(path, inspection, (ImportMapping(item.id, slot),))
    assert_mark_preserved(dict(mapped.assets)[mapped.blocks[0].asset_hashes[0]])
    assets = {}
    block = replace_cover_picture(ResolvedBlock(slot, "Omit"), raw, ".png", assets)
    assert_mark_preserved(assets[block.asset_hashes[0]])
    same = replace_cover_picture(approve_all_assets(block), raw, ".png", assets)
    assert same == approve_all_assets(block), "Identical replacement should preserve its exact existing review"


def test_source_photo_and_technical_page_keep_existing_white_canvas_rule():
    for line_art in (False, True):
        image = normalize_report_image(mark(), ".png", line_art=line_art)
        assert Image.open(BytesIO(image.data)).convert("RGBA").getpixel((0, 0)) == (255, 255, 255, 255)


def test_transparent_logo_survives_shared_snapshot_readback_and_docx(tmp_path, monkeypatch):
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path / "data"))
    assets = {}
    block = approve_all_assets(replace_cover_picture(ResolvedBlock("brand_logo", "Omit"), mark(), ".png", assets))
    sections = tuple(replace(s, included=s.key == "activity") for s in default_sections())
    draft = ReportDraft(synthetic_profiles()[0], ReportPeriod(2026, 11), "Synthetic Reviewer", sections,
                        (block, ResolvedBlock("activity_summary", "This month", text="Synthetic inspection completed.")))
    saved = library.save_snapshot(draft, expected_revision=0, assets=tuple(assets.items()), entered_editor="Synthetic Reviewer")
    loaded = library.load_snapshot(draft.profile.contract, draft.profile.key, draft.period)
    assert loaded == saved
    logo = next(b for b in loaded.draft.blocks if b.key == "brand_logo")
    loader = lambda ref: library.read_asset(draft.profile.contract, draft.profile.key, ref)
    assert_mark_preserved(loader(logo.asset_hashes[0]))
    output = assemble_docx(loaded.draft, acknowledged_fingerprint=loaded.draft.fingerprint, asset_loader=loader)
    with ZipFile(BytesIO(output)) as archive:
        media = [archive.read(n) for n in archive.namelist() if n.startswith("word/media/")]
    assert len(media) == 1
    assert_mark_preserved(media[0])


@pytest.mark.parametrize("slot", ["brand_logo", "client_logo"])
def test_detailed_editor_shared_logo_upload_preserves_original_alpha(tmp_path, monkeypatch, slot):
    """AppTest uploader seam through the real shared replacement/CAS path."""
    from streamlit.testing.v1 import AppTest
    from app import monthly_report_editor as editor
    monkeypatch.setenv("EPC_DATA_DIR", str(tmp_path / "data"))
    profile = synthetic_profiles()[0]
    library.save_profile(profile, expected_revision=0, actor="Synthetic Reviewer", confirmed=True)
    upload = BytesIO(mark())
    upload.name = "synthetic-logo.png"
    real_upload = editor.st.file_uploader
    def uploader(label, *args, **kwargs):
        native = real_upload(label, *args, **kwargs)
        return upload if label == "Replace " + slot.replace("_", " ") else native
    monkeypatch.setattr(editor.st, "file_uploader", uploader)
    app = AppTest.from_string('''
import streamlit as st
from app import monthly_report_editor as editor, monthly_report_library as library
from app.monthly_report_model import synthetic_profiles, layout_blocks
profile = synthetic_profiles()[0]
state = library.load_profile(profile.contract, profile.key)
def field(key, value):
    st.session_state.setdefault(key, value)
    return key
assets = st.session_state.setdefault('assets', {})
spec = next(s for s in layout_blocks() if s.key == SLOT)
st.session_state['result'] = editor._block_editor(spec, state, {}, 'Replace and save to library',
    'synthetic_logo', assets, field)
'''.replace("SLOT", repr(slot))).run()
    assert not app.exception
    block = app.session_state["result"]
    assert len(block.asset_hashes) == 1, "The actual uploader path must process the supplied bytes"
    assert_mark_preserved(app.session_state["assets"][block.asset_hashes[0]])
    next(w for w in app.text_input if w.label == "Editor name").set_value("Synthetic Reviewer").run()
    next(w for w in app.checkbox if w.label == "Save this replacement as a new shared library version").check().run()
    next(b for b in app.button if b.label == "Confirm library replacement").click().run()
    assert not app.exception
    stored = library.load_profile(profile.contract, profile.key)
    assert stored.revision == 2
    assert_mark_preserved(library.read_asset(profile.contract, profile.key, stored.block(slot).asset_hashes[0]))


@pytest.mark.parametrize("master", [False, True])
@pytest.mark.parametrize("aspect", [3, 6])
def test_native_header_uses_identified_brand_without_a_shape_guess(tmp_path, master, aspect):
    from app.monthly_report_native_layout import build_native_docx
    old = Image.new("RGB", (140 * aspect, 140), "gray")
    stream = BytesIO()
    old.save(stream, "PNG")
    doc = Document()
    doc.add_paragraph(synthetic_profiles()[0].facilities[0].title)
    doc.add_paragraph("November 2026")
    header = doc.sections[0].header.paragraphs[0]
    header.add_run().add_picture(BytesIO(stream.getvalue()), width=Inches(2))
    doc.add_section(WD_SECTION.NEW_PAGE)
    doc.add_heading("Monthly Activity Summary", 1)
    doc.add_paragraph("Synthetic inspection completed.")
    source = tmp_path / "header-source.docx"
    doc.save(source)
    # No alt text, adjacent logo caption or filename cue is supplied. The
    # established native header role identifies the mark on both geometries.
    inspected_header = next(item for item in inspect_docx(source).items
                            if item.kind == "image" and item.part.startswith("word/header"))
    assert inspected_header.suggested_slot == "brand_logo"
    assets = {}
    logo = approve_all_assets(replace_cover_picture(ResolvedBlock("brand_logo", "Omit"), mark(), ".png", assets))
    sections = tuple(replace(s, included=s.key == "activity") for s in default_sections())
    draft = ReportDraft(synthetic_profiles()[0], ReportPeriod(2026, 11), "Synthetic Reviewer", sections,
                        (logo, ResolvedBlock("activity_summary", "This month", text="Synthetic inspection completed.")))
    raw = build_native_docx(draft, source, asset_loader=assets.__getitem__, master=master)
    generated = Document(BytesIO(raw))
    pictures = [shape for part in generated.part.package.parts if str(part.partname).startswith("/word/header")
                and hasattr(part, "element") for shape in part.element.iter()
                if shape.tag.endswith("}blip")]
    assert pictures, "Replacing an approved logo must not silently remove the header mark"
    for part in generated.part.package.parts:
        if str(part.partname).startswith("/word/header") and hasattr(part, "element"):
            for node in part.element.iter():
                if node.tag.endswith("}blip"):
                    assert_mark_preserved(part.related_parts[node.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed")].blob)
