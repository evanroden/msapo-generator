"""Populate an inspected Word layout without rebuilding its page design.

The source is a geometry template, never a source of current report facts. Native
text boxes, section properties and formatting survive, but body payloads are
cleared before current values are written. Unsupported bindings fail explicitly.
"""
from __future__ import annotations

import hashlib
import math
import re
import tempfile
from collections import defaultdict
from copy import deepcopy
from functools import lru_cache
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from app.monthly_report_import import (
    _heading,
    _is_native_divider_background,
    _text,
    inspect_docx,
)


class NativeLayoutError(ValueError):
    """The source layout cannot represent the current content safely."""


_STATIC = {
    "work order status", "pm status", "activity summary", "improvement highlights and pics",
    "rate trend", "thermal services requirements and capacity", "client services calls",
    "equipment performance issues", "issues identified", "training provided",
    "rejected quotes", "pending quotes", "contract asset end of useful life schedule",
}
_DEFAULT_TEXT = {"organization": "contact_matrix", "activity": "activity_summary",
                 "scorecards": "utility_analysis", "mbcx": "mbcx_report",
                 "maintenance": "service_calls", "subcontractors": "subcontractor_matrix",
                 "water": "water_reports", "issues": "equipment_issues",
                 "capital": "capital_renewal", "proposals": "proposals",
                 "training": "training_summary", "rfi": "rfi_matrix",
                 "accounts_receivable": "accounts_receivable_notes"}
_MONTH = re.compile(r"\b(?:January|February|March|April|May|June|July|August|September|October|November|December)\s+20\d\d\b", re.IGNORECASE)


def _work_order_table_labels(original, items, section_by_pos):
    """Bind plain annual-table labels to their next physical source table."""
    labels, pending = {}, None
    for position, element in enumerate(original, 1):
        group = items[position]
        if (section_by_pos[position] != "activity"
                or element.tag == qn("w:sectPr")
                or element.find(".//" + qn("w:sectPr")) is not None):
            pending = None
            continue
        text = _text(element).casefold().strip()
        if text in ("work order status", "pm status"):
            # Never retain pictures, fields, alternate branches or uninspected
            # glyphs merely because their extracted words resemble a label.
            formatting = {qn("w:" + name) for name in (
                "p", "pPr", "pStyle", "r", "rPr", "t", "rFonts", "b", "bCs", "i", "iCs",
                "color", "sz", "szCs", "spacing", "ind", "outlineLvl", "jc", "keepNext",
                "keepLines", "contextualSpacing", "caps", "smallCaps", "u", "lang", "noProof")}
            plain = (element.tag == qn("w:p")
                     and all(node.tag in formatting for node in element.iter())
                     and all(child.tag in (qn("w:pPr"), qn("w:r")) for child in element)
                     and all(child.tag in (qn("w:rPr"), qn("w:t"))
                             for run in element.findall(qn("w:r")) for child in run))
            pending = element if plain and any(
                item.kind == "text" and item.suggested_slot == "work_orders" for item in group) else None
        elif element.tag == qn("w:tbl"):
            source_table = next((item for item in group if item.kind == "table"
                                 and item.suggested_slot == "work_orders"), None)
            if pending is not None and source_table is not None:
                labels[element] = (pending, source_table.id)
            pending = None
        elif _heading(text):
            pending = None
    return labels


def _clear_hyperlinks(element):
    """Source click destinations are facts, not reusable text formatting."""
    for link in list(element.iter(qn("w:hyperlink"))):
        parent = link.getparent()
        if parent is not None:
            index = parent.index(link)
            for child in list(link):
                parent.insert(index, child)
                index += 1
            parent.remove(link)
    for tag in (qn("a:hlinkClick"), qn("a:hlinkHover")):
        for link in list(element.iter(tag)):
            link.getparent().remove(link)


def _set_text(element, value):
    """Change run payloads while retaining paragraph/run/text-box properties."""
    _clear_hyperlinks(element)
    value = value.replace("\r\n", "\n").replace("\r", "\n")
    # The import inspector deliberately normalizes whitespace; it cannot tell
    # whether this payload already has the requested native line breaks. Keep
    # unchanged runs only when their exact text and controls really match.
    current, controls = [], []
    seen_paragraph = raw_controls = False
    for node in element.iter():
        if node.tag == qn("w:p"):
            if seen_paragraph:
                current.append("\n")
            seen_paragraph = True
        elif node.tag == qn("w:t"):
            text = node.text or ""
            current.append(text)
            raw_controls = raw_controls or any(char in text for char in "\r\n\t")
        elif node.getparent() is not None and node.getparent().tag == qn("w:r"):
            if (node.tag == qn("w:cr") or (node.tag == qn("w:br")
                    and node.get(qn("w:type"), "textWrapping") == "textWrapping")):
                current.append("\n")
                controls.append(node)
            elif node.tag == qn("w:tab"):
                current.append("\t")
                controls.append(node)
    if not raw_controls and "".join(current) == value:
        return
    # Run tabs and wrapping breaks are old text, while paragraph tab stops and
    # page/column breaks belong to the retained native layout.
    for node in controls:
        node.getparent().remove(node)
    texts = list(element.iter(qn("w:t")))
    if not texts:
        paragraph = element if element.tag == qn("w:p") else next(element.iter(qn("w:p")), None)
        if paragraph is None:
            paragraph = OxmlElement("w:p")
            element.append(paragraph)
        run = OxmlElement("w:r")
        paragraph.append(run)
        node = OxmlElement("w:t")
        run.append(node)
        texts = [node]
    groups = {}
    for node in texts:
        branch = next((ancestor for ancestor in node.iterancestors()
                       if ancestor.tag in ("{http://schemas.openxmlformats.org/markup-compatibility/2006}Choice",
                                           "{http://schemas.openxmlformats.org/markup-compatibility/2006}Fallback")), None)
        groups.setdefault(branch, []).append(node)
    for group in groups.values():
        parts = re.split(r"([\n\t])", value)
        group[0].text = parts[0]
        group[0].set(qn("xml:space"), "preserve")
        previous = group[0]
        for control, text in zip(parts[1::2], parts[2::2]):
            marker = OxmlElement("w:br" if control == "\n" else "w:tab")
            previous.addnext(marker)
            node = OxmlElement("w:t")
            node.text = text
            node.set(qn("xml:space"), "preserve")
            marker.addnext(node)
            previous = node
        for node in group[1:]:
            node.text = ""



def _remove_unapproved_master_watermark_text(document):
    """A global master cannot lend an old client's VML textpath watermark.

    WordArt/v:textpath 'string' is not a w:t run, so normal paragraph text
    replacement misses it. Clear the source-authored text value while keeping
    the drawing's shape, fill, position and unrelated ENFRA art unchanged.
    """
    watermark_tag = "{urn:schemas-microsoft-com:vml}textpath"
    for part in document.part.package.parts:
        name = str(part.partname)
        if not name.startswith(("/word/header", "/word/footer")) or not hasattr(part, "element"):
            continue
        for node in part.element.iter(watermark_tag):
            node.attrib.pop("string", None)


def _clear_source_drawing_metadata(document):
    """Discard source-only accessibility labels, never current native geometry.

    Drawings, Word table descriptions and source content-control names can
    contain previous-client names/prices without appearing in printed text.
    The current report supplies its own visible cells and captions; source
    aliases/labels have no independent authority to accompany those cells.
    """
    drawing_namespaces = {
        "http://schemas.openxmlformats.org/drawingml/2006/picture",
        "http://schemas.openxmlformats.org/drawingml/2006/main",
        "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing",
        "http://schemas.microsoft.com/office/word/2010/wordprocessingShape",
        "http://schemas.microsoft.com/office/word/2010/wordprocessingGroup",
    }
    for part in document.part.package.parts:
        if not hasattr(part, "element"):
            continue
        # Iterate a snapshot because stale nonvisual nodes are removed below.
        for node in list(part.element.iter()):
            namespace, _, local = node.tag[1:].partition("}") if node.tag.startswith("{") else ("", "", node.tag)
            if (namespace == "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
                    and local in {"tblCaption", "tblDescription", "tag", "alias"}):
                # Do not retain previous-client screen-reader descriptions or
                # content-control bookkeeping in current-client output.
                # Table grid/row/cell formatting and actual text are untouched.
                if node.getparent() is not None:
                    node.getparent().remove(node)
                continue
            if namespace in drawing_namespaces and local in {"docPr", "cNvPr"}:
                node.attrib.pop("descr", None)
                node.attrib.pop("title", None)
                # Keep the required nonvisual name without importing a client's
                # name, file name or old caption. IDs/references stay unchanged.
                if "name" in node.attrib:
                    node.set("name", "Drawing " + node.get("id", ""))
            elif namespace == "urn:schemas-microsoft-com:vml":
                for attr in list(node.attrib):
                    if attr.rsplit("}", 1)[-1] in {"alt", "title"}:
                        del node.attrib[attr]


def _remove_unapproved_header_images(part, approved_digests, approved_occurrences=()):
    """Drop only image-backed old-client header/footer art on a fresh master.

    Pure native rules/rectangles and Word geometry are left intact; approved
    current ENFRA logo bytes and explicitly preserved native occurrences remain.
    An approved native occurrence does not authorize another use of equal bytes.
    All other image relationships are facts about the imported source, even when
    PNG metadata is not visible on page.
    The package sanitizer prunes dangling image parts after saving.
    """
    vml = "{urn:schemas-microsoft-com:vml}imagedata"
    for picture in list(part.element.iter()):
        if picture.tag not in (qn("w:drawing"), qn("w:pict")):
            continue
        images = [(node, node.get(qn("r:embed")) or node.get(qn("r:id")))
                  for node in picture.iter()
                  if node.tag in (qn("a:blip"), vml)]
        images = [(node, identity) for node, identity in images if identity]
        if not images:
            continue  # Pure vector decoration, without client media.
        trusted = True
        for node, identity in images:
            rel = part.rels.get(identity)
            if rel is None or rel.is_external or not hasattr(rel.target_part, "blob"):
                trusted = False
                break
            if (node not in approved_occurrences
                    and hashlib.sha256(rel.target_part.blob).hexdigest() not in approved_digests):
                trusted = False
                break
        if not trusted and picture.getparent() is not None:
            picture.getparent().remove(picture)


def _header_image_occurrences(part, inspection, *, elements=None):
    """Bind inspected image IDs to their own header/footer XML occurrence.

    Media parts (and relationship IDs) are reusable: two independent pictures
    can contain exactly the same bytes. Never let a mapping for one occurrence
    authorize replacing every use of that media part.
    """
    name = str(part.partname).lstrip("/")
    by_position = defaultdict(list)
    for item in inspection.items:
        if item.kind == "image" and item.part == name:
            by_position[item.position].append(item)
    vml = "{urn:schemas-microsoft-com:vml}imagedata"
    mc = "{http://schemas.openxmlformats.org/markup-compatibility/2006}"

    def image_identity(identity):
        rel = part.rels.get(identity)
        if rel and not rel.is_external and hasattr(rel.target_part, "blob"):
            return hashlib.sha256(rel.target_part.blob).hexdigest()
        return None

    def logical_pictures(element):
        if element.tag in (qn("a:blip"), vml):
            attr = qn("r:embed") if element.tag == qn("a:blip") else qn("r:id")
            rel = part.rels.get(element.get(attr))
            if rel and not rel.is_external:
                return [[(element, attr, str(rel.target_part.partname).lstrip("/"))]]
            return []
        if element.tag == mc + "AlternateContent":
            selected = next((child for child in element if child.tag == mc + "Choice"), None)
            if selected is None:
                selected = next((child for child in element if child.tag == mc + "Fallback"), None)
            if selected is None:
                raise ValueError("No visible header artwork branch")
            pictures = logical_pictures(selected)
            for branch in element:
                if branch is selected:
                    continue
                alternatives = defaultdict(list)
                for group in logical_pictures(branch):
                    node, attr, _ = group[0]
                    alternatives[image_identity(node.get(attr))].append(group)
                for group in pictures:
                    node, attr, _ = group[0]
                    matches = alternatives[image_identity(node.get(attr))]
                    if not matches:
                        raise ValueError("Header artwork branches have different occurrences")
                    group.extend(matches.pop(0))
                if any(alternatives.values()):
                    raise ValueError("Header artwork branch has additional occurrences")
            return pictures
        return [group for child in element for group in logical_pictures(child)]

    bindings = []
    from app.monthly_report_native_design import equivalent_native_alternatives
    for position, element in enumerate(part.element if elements is None else elements, 1):
        # Validated Choice/Fallback pictures are representations of ONE logical
        # occurrence. Independent equal-byte drawings outside that wrapper are
        # still separate. Update every retained equivalent representation.
        if not equivalent_native_alternatives(element, image_identity):
            continue
        try:
            pictures = logical_pictures(element)
        except ValueError:
            continue  # A selected unresolved occurrence fails explicitly below.
        groups_by_target = defaultdict(list)
        for group in pictures:
            groups_by_target[group[0][2]].append(group)
        items_by_target = defaultdict(list)
        for item in by_position[position]:
            items_by_target[item.image_part].append(item)
        for target, groups in groups_by_target.items():
            source_items = items_by_target[target]
            if len(groups) != len(source_items):
                continue
            bindings.extend((node, attr, item) for group, item in zip(groups, source_items)
                            for node, attr, _ in group)
    return bindings


def _clear_diagrams(element):
    # A diagram may share a drawing canvas with a retained text-box frame.
    # Its relationship payload must still disappear when that content is cleared.
    for diagram in list(element.iter("{http://schemas.openxmlformats.org/drawingml/2006/diagram}relIds")):
        diagram.getparent().remove(diagram)


def _clear_images(element, *, preserve_decoration=False):
    _clear_diagrams(element)
    # Remove the drawing itself, not its paragraph or section break.
    for node in list(element.iter()):
        if node.tag in (qn("w:drawing"), qn("w:pict")):
            # Pure native rectangles and rules are decoration, not payloads.
            presets = list(node.iter(qn("a:prstGeom")))
            colors = {n.get("val", "").lower() for n in node.iter(qn("a:srgbClr"))}
            shapes = list(node.iter("{urn:schemas-microsoft-com:vml}shape"))
            custom = list(node.iter(qn("a:custGeom")))
            basic = bool(presets) and all(p.get("prst") in {"rect", "line", "parallelogram"} for p in presets)
            basic |= bool(custom) and colors <= {"547e7e", "557f7f", "d5ee4a", "d6ef4b"} and bool(colors) and all(
                len(list(g.iter(qn("a:pt")))) <= 6 and not list(g.iter(qn("a:cubicBezTo")))
                and not list(g.iter(qn("a:quadBezTo"))) and not list(g.iter(qn("a:arcTo"))) for g in custom)
            basic |= bool(shapes) and all(s.get("fillcolor", "").lower() in {"#547e7e", "#557f7f", "#d5ee4a", "#d6ef4b"}
                and re.fullmatch(r"[0-9,mlrxe -]{1,150}", s.get("path", "")) for s in shapes)
            if (preserve_decoration and basic
                    and not any(n.tag in (qn("w:t"), qn("a:t"), qn("a:blip"))
                                or any(k in n.attrib for k in (qn("r:embed"), qn("r:link"), qn("r:id")))
                                for n in node.iter())):
                continue
            # Text-box shapes are layout, not image payloads.
            if any(True for _ in node.iter(qn("w:txbxContent"))):
                for image in list(node.iter()):
                    image.attrib.pop(qn("r:embed"), None)
                    image.attrib.pop(qn("r:link"), None)
                    if image.tag in (qn("a:blip"), "{urn:schemas-microsoft-com:vml}imagedata"):
                        image.getparent().remove(image)
            elif node.getparent() is not None:
                node.getparent().remove(node)


def _contain_replacement_picture(element, data):
    """Keep a current improvement photograph inside its native physical frame.

    The previous asset's crop and aspect are not evidence for the current one.
    Unchanged source pictures never pass through this function.  Do not guess
    geometry for unsupported picture containers or grouped VML art.
    """
    from PIL import Image
    with Image.open(BytesIO(data)) as image:
        width, height = image.size
        # Reading orientation metadata avoids a second full image decode just to
        # measure a current photograph on Render's memory-constrained instance.
        if image.getexif().get(274, 1) in (5, 6, 7, 8):
            width, height = height, width
    if width <= 0 or height <= 0:
        raise NativeLayoutError("This replacement picture has no usable dimensions.")
    pictures = list(element.iter(qn("pic:pic")))
    for picture in pictures:
        extent = picture.find("./" + qn("pic:spPr") + "/" + qn("a:xfrm") + "/" + qn("a:ext"))
        frame = next((parent for parent in picture.iterancestors()
                      if parent.tag in (qn("wp:anchor"), qn("wp:inline"))), None)
        outer = frame.find(qn("wp:extent")) if frame is not None else None
        if extent is None or outer is None:
            raise NativeLayoutError("The native photo frame cannot be fitted safely. Add this picture to a new photo page.")
        try:
            full_width, full_height = int(outer.get("cx", "0")), int(outer.get("cy", "0"))
        except ValueError as exc:
            raise NativeLayoutError("The native photo frame dimensions are invalid.") from exc
        if min(full_width, full_height) <= 0:
            raise NativeLayoutError("The native photo frame has no usable area.")
        scale = min(full_width / width, full_height / height)
        new_width, new_height = round(width * scale), round(height * scale)
        for node in (outer, extent):
            node.set("cx", str(new_width)); node.set("cy", str(new_height))
        # The physical frame must contain every edge of the current photo.
        # A prior photo's source-rectangle values would crop those edges.
        for crop in list(picture.iter(qn("a:srcRect"))):
            crop.getparent().remove(crop)
    for shape in list(element.iter("{urn:schemas-microsoft-com:vml}shape")):
        images = list(shape.iter("{urn:schemas-microsoft-com:vml}imagedata"))
        if not images:
            continue
        if any(parent.tag == "{urn:schemas-microsoft-com:vml}group" for parent in shape.iterancestors()):
            raise NativeLayoutError("Grouped native photo frames need a reviewed replacement page; their image proportions cannot be assumed.")
        original = shape.get("style", "")
        found = {}
        for axis in ("width", "height"):
            result = re.search(r"(?:(?<=;)|^)\s*" + axis + r"\s*:\s*(\d+(?:\.\d+)?)(in|pt)(?=;|$)", original)
            if result:
                found[axis] = (result, float(result[1]) * (72 if result[2] == "in" else 1))
        if len(found) != 2:
            raise NativeLayoutError("The native photo frame dimensions cannot be fitted safely.")
        scale = min(found["width"][1] / width, found["height"][1] / height)
        target = {"width": width * scale, "height": height * scale}
        # Replace later spans first; VML permits height before width as well.
        for axis, (result, _) in sorted(found.items(), key=lambda item: item[1][0].start(), reverse=True):
            value = target[axis] / (72 if result[2] == "in" else 1)
            replacement = f"{axis}:{value:.6f}{result[2]}"
            original = original[:result.start()] + replacement + original[result.end():]
        shape.set("style", original)
        for image in images:
            for key in ("cropleft", "cropright", "croptop", "cropbottom"):
                image.attrib.pop(key, None)
    if not pictures and not any(True for _ in element.iter("{urn:schemas-microsoft-com:vml}imagedata")):
        raise NativeLayoutError("The native photo replacement lacks a supported image frame.")


def _discardable_picture_anchor(element):
    """Only an already-cleared frame paragraph may lose obsolete layout."""
    if element.tag != qn('w:p') or element.find('.//' + qn('w:sectPr')) is not None:
        return False
    if any((node.text or '').strip() for node in element.iter(qn('w:t'))):
        return False
    allowed = {qn('w:p'), qn('w:pPr'), qn('w:r'), qn('w:rPr'), qn('w:t'), qn('w:br')}
    return all(node.tag in allowed or any(parent.tag in (qn('w:pPr'), qn('w:rPr'))
               for parent in node.iterancestors()) for node in element.iter())


def _replace_image(document, element, data, *, contain=False):
    # Mixed picture/SmartArt paragraphs are cloned as image frames. Updating a
    # raster is not proof that the neighboring source diagram is current.
    _clear_diagrams(element)
    rid, _ = document.part.get_or_add_image(BytesIO(data))
    images = [n for n in element.iter() if n.tag in (qn("a:blip"), "{urn:schemas-microsoft-com:vml}imagedata")]
    if not images:
        raise NativeLayoutError("This section has no native picture frame for its current images.")
    first = images[0]
    old_rid = first.get(qn("r:embed")) or first.get(qn("r:id"))
    for node in images:
        attr = qn("r:embed") if node.tag == qn("a:blip") else qn("r:id")
        if node.get(attr) == old_rid:
            node.set(attr, rid)
            node.attrib.pop(qn("r:link"), None)
        else:
            node.getparent().remove(node)
    # Office image effects may reference the visible image in a14:imgLayer while
    # the standard blip contains a placeholder. Those payloads must update too.
    for node in element.iter():
        if node.get(qn("r:embed")):
            node.set(qn("r:embed"), rid)
        node.attrib.pop(qn("r:link"), None)
    if contain:
        _contain_replacement_picture(element, data)


def _zero_geometry_value(value):
    try:
        number = float(value[:-1] if value.endswith("%") else value)
        return math.isfinite(number) and number == 0
    except (TypeError, ValueError, OverflowError):
        return False


def _ordinary_unchanged_photo_geometry(node):
    if node.tag == qn("a:blip"):
        fill = node.getparent()
        picture = next((ancestor for ancestor in node.iterancestors() if ancestor.tag == qn("pic:pic")), None)
        frame = next((ancestor for ancestor in node.iterancestors()
                      if ancestor.tag in (qn("wp:inline"), qn("wp:anchor"))), None)
        if picture is None or frame is None or fill.tag != qn("pic:blipFill"):
            return False
        properties = picture.find(qn("pic:spPr"))
        transform = properties.find(qn("a:xfrm")) if properties is not None else None
        inner = transform.find(qn("a:ext")) if transform is not None else None
        outer = frame.find(qn("wp:extent"))
        if inner is None or outer is None:
            return False
        preset = properties.find(qn("a:prstGeom"))
        if (properties.find(qn("a:custGeom")) is not None
                or preset is not None and preset.get("prst") != "rect"
                or not _zero_geometry_value(transform.get("rot", "0"))
                or any(transform.get(key, "0") not in ("0", "false", "off") for key in ("flipH", "flipV"))):
            return False
        try:
            dimensions = [float(extent.get(axis, "0")) for extent in (inner, outer) for axis in ("cx", "cy")]
            if not all(math.isfinite(value) and value > 0 for value in dimensions):
                return False
            ratios = (dimensions[0] / dimensions[1], dimensions[2] / dimensions[3])
            if (not all(math.isfinite(value) and value > 0 for value in ratios)
                    or not math.isclose(*ratios, rel_tol=1e-6)):
                return False
        except (TypeError, ValueError, OverflowError):
            return False
        stretch = fill.find(qn("a:stretch"))
        rectangle = stretch.find(qn("a:fillRect")) if stretch is not None else None
        return not (fill.find(qn("a:tile")) is not None
                    or rectangle is not None and any(not _zero_geometry_value(value) for value in rectangle.attrib.values())
                    or any(child.tag.rsplit("}", 1)[-1] == "grpSpPr"
                           and child.find(qn("a:xfrm")) is not None
                           for ancestor in node.iterancestors() for child in ancestor))
    # Other valid VML shape subtypes have not had their effective geometry
    # proved here. Preserve them instead of making the old identity shortcut
    # newly fail or borrowing an enclosing shape's unrelated dimensions.
    shape = node.getparent()
    if shape.tag != "{urn:schemas-microsoft-com:vml}shape" or shape.get("type"):
        return False  # Referenced shapetype geometry has not been proved.
    style = dict((key.strip(), value.strip()) for entry in shape.get("style", "").split(";")
                 if ":" in entry for key, value in [entry.split(":", 1)])
    if any(key != key.casefold() for key in style if key.casefold() in ("width", "height", "rotation", "flip")):
        return False
    dimensions = []
    for axis in ("width", "height"):
        length = re.fullmatch(r"(\d+(?:\.\d*)?|\.\d+)(pt|in|cm|mm|px)", style.get(axis, ""))
        if length is None:
            return False
        points = float(length[1]) * {"pt": 1, "in": 72, "cm": 72 / 2.54, "mm": 72 / 25.4, "px": .75}[length[2]]
        if not math.isfinite(points) or points <= 0:
            return False
        dimensions.append(points)
    ratio = dimensions[0] / dimensions[1]
    if not math.isfinite(ratio) or ratio <= 0:
        return False
    return (not shape.get("path") and not shape.get("flip") and not style.get("flip")
            and _zero_geometry_value(style.get("rotation", shape.get("rotation", "0")))
            and shape.find("{urn:schemas-microsoft-com:vml}skew") is None
            and shape.find("{urn:schemas-microsoft-com:vml}path") is None
            and not any(ancestor.tag == "{urn:schemas-microsoft-com:vml}group"
                        for ancestor in node.iterancestors()))


def _unchanged_photo_geometry_supported(node):
    if not _ordinary_unchanged_photo_geometry(node):
        return False
    # Modern/fallback representations are one logical picture. If any branch
    # is outside the bounded repair, do not recrop another branch in isolation.
    for alternate in node.iterancestors("{http://schemas.openxmlformats.org/markup-compatibility/2006}AlternateContent"):
        if any(not _ordinary_unchanged_photo_geometry(image) for image in alternate.iter()
               if image.tag in (qn("a:blip"), "{urn:schemas-microsoft-com:vml}imagedata")):
            return False
    return True


def _fit_cover_photo(node, image_width, image_height, *, preserve_crop=False):
    """Fit a photo in its frame, retaining valid unchanged source crops.

    For unchanged DrawingML, this repair is deliberately limited to ordinary
    full-frame fills. Group transforms and inset/tiled destination fills need
    their effective geometry proved separately; do not introduce a new crop
    based only on their local picture extent.
    """
    if image_width <= 0 or image_height <= 0:
        raise NativeLayoutError("The cover photograph has no usable dimensions.")
    if preserve_crop and not _unchanged_photo_geometry_supported(node):
        return
    frame = None
    if node.tag == qn("a:blip"):
        for ancestor in node.iterancestors():
            transforms = ancestor.findall(qn("a:xfrm"))
            for child in ancestor:
                if child.tag.rsplit("}", 1)[-1] == "spPr":
                    transforms.extend(child.findall(qn("a:xfrm")))
            ext = next((transform.find(qn("a:ext")) for transform in transforms
                        if transform.find(qn("a:ext")) is not None), None)
            if ext is not None:
                try:
                    frame = (float(ext.get("cx", "0")), float(ext.get("cy", "0")))
                except (TypeError, ValueError, OverflowError) as exc:
                    raise NativeLayoutError("The cover photograph's native frame dimensions are unavailable.") from exc
                break
    else:
        def dimensions(element):
            values = {}
            for key, value, unit in re.findall(r"(?:^|;)\s*(width|height)\s*:\s*([\d.]+)([a-z]*)", element.get("style", "")):
                values[key] = float(value) * {"in": 72, "cm": 72 / 2.54, "mm": 72 / 25.4,
                                              "px": .75}.get(unit, 1)
            return (values.get("width", 0), values.get("height", 0))

        shape = next((ancestor for ancestor in node.iterancestors()
                      if ancestor.tag == "{urn:schemas-microsoft-com:vml}shape"), None)
        if shape is not None:
            width, height = dimensions(shape)
            for group in shape.iterancestors():
                if group.tag != "{urn:schemas-microsoft-com:vml}group":
                    continue
                coordinates = group.get("coordsize", "").split(",")
                group_width, group_height = dimensions(group)
                if len(coordinates) == 2 and all(float(value) > 0 for value in coordinates) and group_width and group_height:
                    width *= group_width / float(coordinates[0])
                    height *= group_height / float(coordinates[1])
            frame = width, height
    if not frame or not all(math.isfinite(value) and value > 0 for value in frame):
        raise NativeLayoutError("The cover photograph's native frame dimensions are unavailable.")
    frame_ratio = frame[0] / frame[1]
    image_ratio = image_width / image_height
    if preserve_crop:
        # Matching bytes prove which photo this is, not that the old Word
        # frame displayed it proportionally. Preserve an intentional source
        # crop only when its remaining pixels actually fit that same frame.
        # This also retains asymmetric crops instead of needlessly recentering.
        try:
            def fraction(value, divisor):
                if value.endswith("%"):
                    return float(value[:-1]) / 100
                if divisor == 1 and value.endswith("f"):
                    return float(value[:-1]) / 65536
                return float(value) / divisor

            if node.tag == qn("a:blip"):
                rectangle = node.getparent().find(qn("a:srcRect"))
                existing = {key: fraction(rectangle.get(key, "0"), 100000)
                            if rectangle is not None else 0 for key in ("l", "t", "r", "b")}
            else:
                existing = {key: fraction(node.get(name, "0"), 1) for key, name in {
                    "l": "cropleft", "r": "cropright", "t": "croptop", "b": "cropbottom"}.items()}
            width = 1 - existing["l"] - existing["r"]
            height = 1 - existing["t"] - existing["b"]
            if (width > 0 and height > 0
                    and all(math.isfinite(value) for value in existing.values())
                    and math.isclose(image_ratio * width / height, frame_ratio, rel_tol=.001)):
                return
        except (TypeError, ValueError, OverflowError):
            pass  # An unusable source crop cannot justify stretching a photo.
    horizontal = max(0, (1 - frame_ratio / image_ratio) / 2)
    vertical = max(0, (1 - image_ratio / frame_ratio) / 2)
    crop = {"l": horizontal, "r": horizontal, "t": vertical, "b": vertical}
    if node.tag == qn("a:blip"):
        fill = node.getparent()
        rectangle = fill.find(qn("a:srcRect"))
        if rectangle is None:
            rectangle = OxmlElement("a:srcRect")
            node.addnext(rectangle)
        for key, value in crop.items():
            rectangle.set(key, str(round(value * 100000)))
    else:
        for key, name in {"l": "cropleft", "r": "cropright", "t": "croptop", "b": "cropbottom"}.items():
            node.set(name, format(crop[key], ".8f"))


_TRACKED = {qn("w:" + name) for name in ("ins", "del", "trPrChange", "tcPrChange", "pPrChange", "rPrChange",
                                          "tblPrExChange", "cellIns", "cellDel", "cellMerge")}


def _skeleton(row):
    """A source row's geometry and formatting only, never its payload.

    Each cell keeps its properties and one paragraph whose first run's
    formatting holds an empty text node. Text, bookmarks, permissions,
    fields, notes and tracked-change authors are not copied (a master
    supplies geometry only).
    """
    row = deepcopy(row)
    for child in list(row):
        if child.tag not in (qn("w:tblPrEx"), qn("w:trPr"), qn("w:tc")):
            row.remove(child)
    for cell in row.findall(qn("w:tc")):
        source = cell.find(qn("w:p"))
        paragraph, run = OxmlElement("w:p"), OxmlElement("w:r")
        if source is not None:
            for kept, owner in ((source.find(qn("w:pPr")), paragraph),
                                (next((r.find(qn("w:rPr")) for r in source.iter(qn("w:r"))), None), run)):
                if kept is not None:
                    owner.append(kept)
        for child in list(cell):
            if child.tag != qn("w:tcPr"):
                cell.remove(child)
        run.append(OxmlElement("w:t"))
        paragraph.append(run)
        cell.append(paragraph)
    for node in [node for node in row.iter() if node.tag in _TRACKED]:
        node.getparent().remove(node)
    return row


def _without_position(row):
    """A row skeleton without Word's conditional-style position markers (first, last, band)."""
    row = _skeleton(row)
    for node in list(row.iter(qn("w:cnfStyle"))):
        node.getparent().remove(node)
    return row


def _plain_row(row):
    """A header skeleton for data: its geometry, without the header's fill or emphasis.

    Direct fills and character styles are removed, and the bold, caps and
    colour a header paragraph style would still apply are overridden on the
    run current text reuses.
    """
    row = _without_position(row)
    for tag in ("w:tblHeader", "w:shd", "w:rStyle", "w:highlight", "w:b", "w:bCs", "w:caps", "w:color"):
        for node in list(row.iter(qn(tag))):
            node.getparent().remove(node)
    for run in row.iter(qn("w:r")):
        properties = run.get_or_add_rPr()
        for name in ("b", "bCs", "caps"):
            getattr(properties, "get_or_add_" + name)().val = False
        properties.get_or_add_color().set(qn("w:val"), "auto")
    return row


def _repeats(row):
    marker = row.find(qn("w:trPr") + "/" + qn("w:tblHeader"))
    return marker is not None and marker.get(qn("w:val"), "1") not in ("0", "false", "off")


def _data_row(element, continuation):
    """The row current data copies: the table's last data row, else a plain header.

    The header is the wider of the first two rows plus any following rows
    marked to repeat as headers, so a two-row header is never taken as data.
    """
    rows = element.findall(qn("w:tr"))
    header = max(rows[:2], key=lambda row: len(row.findall(qn("w:tc"))))
    extent = rows.index(header) + 1
    while extent < len(rows) and _repeats(rows[extent]):
        extent += 1
    if continuation or extent < len(rows):
        return _without_position(rows[-1])
    return _plain_row(header)


def _table(element, columns, rows, *, physical_indices=(), continuation=False, keep_all_rows_intact=False,
           template=None):
    """Write current rows into a native table.

    ``template`` is the untouched source table when ``element`` was already
    blanked to its header, so data rows still copy a source data row.
    """
    explicit_header = bool(columns) and not continuation
    native_rows = element.findall(qn("w:tr"))
    if not native_rows:
        raise NativeLayoutError("The source table has no row geometry.")
    header = max(native_rows[:2], key=lambda row: len(row.findall(qn("w:tc"))))
    has_template = template is not None and template.find(qn("w:tr")) is not None
    prototype = _data_row(template if has_template else element, continuation)
    width = len(columns) or max((len(row) for row in rows), default=0)
    if not width or width > 100 or any(len(row) > width for row in rows):
        raise NativeLayoutError("Current table columns do not match its row values.")
    if not columns:
        columns = tuple(_text(cell) for cell in header.findall(qn("w:tc")))
        columns = (columns + ("",) * width)[:width]
    resized = width != len(header.findall(qn("w:tc")))
    grid = element.find(qn("w:tblGrid"))
    total = sum(int(col.get(qn("w:w"), "0")) for col in grid) if grid is not None else 0
    total = total or 9000
    cell_width = max(1, total // width)
    # A fresh master can exclude source pricing columns, but all other native
    # column proportions should survive rather than become equal-width cells.
    widths = ()
    if (physical_indices and grid is not None and len(physical_indices) == width
            and len(grid) > max(physical_indices)):
        original_widths = [max(1, int(col.get(qn("w:w"), "1"))) for col in grid]
        selected = [original_widths[index] for index in physical_indices]
        retained_total = sum(selected)
        widths = tuple(max(1, round(value * total / retained_total)) for value in selected)
    if grid is not None and resized:
        for col in list(grid):
            grid.remove(col)
        for index in range(width):
            col = OxmlElement("w:gridCol")
            col.set(qn("w:w"), str(widths[index] if widths else cell_width))
            grid.append(col)
    for row in native_rows:
        element.remove(row)
    # A numbered source continuation has no header in this physical frame.
    # Its first source row was DATA, not a schema to repeat into the new
    # report. An empty continuation keeps a blank shape, never an old row.
    output_rows = (rows or (("",) * width,)) if continuation else (columns, *rows)
    for index, values in enumerate(output_rows):
        row = _skeleton(header) if index == 0 and not continuation else deepcopy(prototype)
        cells = row.findall(qn("w:tc"))
        for cell in cells:
            row.remove(cell)
        if physical_indices and len(cells) > max(physical_indices):
            # A full-width source row lends each retained column its own cell,
            # never an excluded price column's formatting.
            cells = [cells[n] for n in physical_indices]
        for column, value in enumerate(tuple(values) + ("",) * (width - len(values))):
            cell = deepcopy(cells[min(column, len(cells) - 1)])
            props = cell.find(qn("w:tcPr"))
            if props is not None:
                for merged in list(props):
                    if merged.tag in (qn("w:gridSpan"), qn("w:vMerge"), qn("w:hMerge"), qn("w:noWrap")):
                        props.remove(merged)
                size = props.find(qn("w:tcW"))
                if size is not None and (resized or len(cells) != width):
                    size.set(qn("w:w"), str(widths[column] if widths else cell_width))
                    size.set(qn("w:type"), "dxa")
            _clear_images(cell)
            _set_text(cell, "" if value is None else str(value))
            row.append(cell)
        # Source fixed row heights cannot clip the user's longer current text.
        for height in list(row.iter(qn("w:trHeight"))):
            height.set(qn("w:hRule"), "atLeast")
        element.append(row)
    if explicit_header and rows:
        from app.monthly_report_table_pagination import keep_current_header_with_rows
        keep_current_header_with_rows(element, keep_all_rows_intact=keep_all_rows_intact)


def _cover(element, draft, *, preserve_issue_date=False):
    paragraphs = list(element.iter(qn("w:p")))
    for paragraph in paragraphs:
        # Nested paragraphs are handled separately; never flatten the outer shape.
        if paragraph.find('.//' + qn('w:txbxContent')) is not None:
            continue
        text = _text(paragraph)
        if not text:
            continue
        if re.search(r"prepared\s+by", text, re.IGNORECASE):
            value = "Prepared by: " + draft.prepared_by
            issued = re.search(r"\s+Date\s*:.*$", text, re.IGNORECASE)
            if issued and preserve_issue_date and not re.search(r"\bDate\s*:", value, re.IGNORECASE):
                value += issued.group()
            _set_text(paragraph, value)
        elif re.match(r"^\s*Date\s*:", text, re.IGNORECASE):
            if not preserve_issue_date:
                _set_text(paragraph, "")
        elif _MONTH.search(text):
            _set_text(paragraph, _MONTH.sub(draft.period.label, text))
        elif (text.strip().upper() not in {"ENFRA", "=", "INSERT IMAGE HERE"}
              and "operations and maintenance" not in text.casefold()
              and "monthly review" not in text.casefold()
              and "confidential" not in text.casefold()
              and "create." not in text.casefold()
              and "section" not in text.casefold()
              and "table of contents" not in text.casefold()):
            changed_title = text != draft.profile.title
            _set_text(paragraph, draft.profile.title)
            if changed_title and len(draft.profile.title) > 22:
                for run in paragraph.findall(qn("w:r")):
                    props = run.find(qn("w:rPr"))
                    if props is not None:
                        for size in props.findall(qn("w:sz")):
                            original_size = int(size.get(qn("w:val"), "48"))
                            size.set(qn("w:val"), str(max(22, int(original_size * min(1, 22 / len(draft.profile.title))))))
        elif "INSERT IMAGE" in text.upper():
            _set_text(paragraph, "")


def _native_photo_orientation_proved(node, document_part):
    from PIL import Image, UnidentifiedImageError

    representations = [node]
    for alternate in node.iterancestors("{http://schemas.openxmlformats.org/markup-compatibility/2006}AlternateContent"):
        representations.extend(image for image in alternate.iter()
                               if image.tag in (qn("a:blip"), "{urn:schemas-microsoft-com:vml}imagedata"))
    for image in set(representations):
        rid = image.get(qn("r:embed")) or image.get(qn("r:id"))
        relation = document_part.rels.get(rid)
        if relation is None or relation.is_external:
            return False
        try:
            with Image.open(BytesIO(relation.target_part.blob)) as source:
                if source.getexif().get(274, 1) != 1:
                    return False
        except (UnidentifiedImageError, OSError, ValueError, Image.DecompressionBombError):
            return False
    return True


def _contain_cover_logo(node, image_width, image_height, document_part, *, preserve_crop=False, defer=False):
    """Fit one proved logo occurrence without stretching or changing its pixels.

    Preserve already-proportional source crops. Unsupported transforms remain
    untouched rather than borrowing an enclosing text box's dimensions. The
    same conservative geometry/orientation boundary as cover photos applies.
    """
    if (not _unchanged_photo_geometry_supported(node)
            or not _native_photo_orientation_proved(node, document_part)
            or image_width <= 0 or image_height <= 0):
        return
    if node.tag == qn("a:blip"):
        shared_frame = next(p for p in node.iterancestors() if p.tag in (qn("wp:inline"), qn("wp:anchor")))
        peers = [p for p in shared_frame.iter(qn("a:blip"))
                 if next((a for a in p.iterancestors() if a.tag in (qn("wp:inline"), qn("wp:anchor"))), None) is shared_frame]
        crops = set()
        for peer in peers:
            crop = peer.getparent().find(qn("a:srcRect"))
            crops.add(tuple(crop.get(k, "0") if crop is not None else "0" for k in ("l", "t", "r", "b")))
        if len(crops) > 1:
            # Different source crops sharing one outer frame cannot each be
            # contained independently without invalidating the other's crop.
            return
    if node.tag == qn("a:blip"):
        picture = next(node.iterancestors(qn("pic:pic")))
        inner = picture.find("./" + qn("pic:spPr") + "/" + qn("a:xfrm") + "/" + qn("a:ext"))
        frame = next(p for p in picture.iterancestors() if p.tag in (qn("wp:inline"), qn("wp:anchor")))
        outer = frame.find(qn("wp:extent"))
        width, height = float(outer.get("cx")), float(outer.get("cy"))
        rectangle = node.getparent().find(qn("a:srcRect"))
        crop = {key: rectangle.get(key, "0") if rectangle is not None else "0" for key in ("l", "t", "r", "b")}
        divisor = 100000
    else:
        shape = node.getparent()
        style = shape.get("style", "")
        lengths = {}
        for axis in ("width", "height"):
            matches = list(re.finditer(r"(?:^|;)\s*" + axis + r"\s*:\s*(\d+(?:\.\d*)?|\.\d+)(pt|in|cm|mm|px)\s*(?=;|$)", style))
            match = matches[-1] if matches else None
            if match is None:
                return
            factor = {"pt": 1, "in": 72, "cm": 72 / 2.54, "mm": 72 / 25.4, "px": .75}[match[2]]
            lengths[axis] = (match, factor, float(match[1]) * factor)
        width, height = lengths["width"][2], lengths["height"][2]
        crop = {key: node.get(name, "0") for key, name in {
            "l": "cropleft", "r": "cropright", "t": "croptop", "b": "cropbottom"}.items()}
        divisor = 1
    ratio = image_width / image_height
    if not math.isfinite(ratio) or ratio <= 0:
        return
    if preserve_crop:
        try:
            def fraction(value):
                if value.endswith("%"):
                    return float(value[:-1]) / 100
                if divisor == 1 and value.endswith("f"):
                    return float(value[:-1]) / 65536
                return float(value) / divisor
            values = {key: fraction(value) for key, value in crop.items()}
            remaining_width = 1 - values["l"] - values["r"]
            remaining_height = 1 - values["t"] - values["b"]
            if (all(math.isfinite(v) for v in values.values()) and min(remaining_width, remaining_height) > 0
                    and math.isclose(ratio * remaining_width / remaining_height, width / height, rel_tol=.001)):
                return
        except (ValueError, TypeError, OverflowError):
            pass
    scale = min(width / image_width, height / image_height)
    new_width, new_height = image_width * scale, image_height * scale
    if not all(math.isfinite(v) and v > 0 for v in (new_width, new_height)):
        return
    def apply():
        if node.tag == qn("a:blip"):
            if min(round(new_width), round(new_height)) <= 0:
                return
            for extent in (inner, outer):
                extent.set("cx", str(round(new_width))); extent.set("cy", str(round(new_height)))
            if rectangle is not None:
                rectangle.getparent().remove(rectangle)
        else:
            new_style = style
            for axis, target in (("width", new_width), ("height", new_height)):
                units = {"pt": 1, "in": 72, "cm": 72 / 2.54, "mm": 72 / 25.4, "px": .75}
                new_style = re.sub(r"(?P<prefix>(?:^|;)\s*" + axis + r"\s*:\s*)(\d+(?:\.\d*)?|\.\d+)(pt|in|cm|mm|px)\s*(?=;|$)",
                               lambda m: m["prefix"] + f"{target / units[m[3]]:.8f}" + m[3], new_style)
            shape.set("style", new_style)
            for key in ("cropleft", "cropright", "croptop", "cropbottom"):
                node.attrib.pop(key, None)
    if defer:
        return apply
    apply()


def _contain_unchanged_cover_logo(node, source_path, item, document_part, *, defer=False):
    width, height = item.image_width, item.image_height
    if width <= 0 or height <= 0:
        from app.monthly_report_import import read_import_image
        image = read_import_image(Path(source_path), item, line_art=True)
        width, height = image.width, image.height
    return _contain_cover_logo(node, width, height, document_part, preserve_crop=True, defer=defer)


def _fit_unchanged_cover_photo(node, source_path, item, document_part):
    # Native renderers need not honor EXIF the same way as gallery
    # normalization. Without proving their display axes, retain all original
    # representations/crops rather than fitting one against the wrong axes.
    if (not _unchanged_photo_geometry_supported(node)
            or not _native_photo_orientation_proved(node, document_part)):
        return
    width, height = item.image_width, item.image_height
    if width <= 0 or height <= 0:
        from app.monthly_report_import import read_import_image
        normalized = read_import_image(Path(source_path), item, line_art=False)
        width, height = normalized.width, normalized.height
    _fit_cover_photo(node, width, height, preserve_crop=True)


@lru_cache(maxsize=128)
def _source_image_digest(source_path, source_digest, item, line_art, preserve_transparency=None):
    # The verified immutable source hash is part of this bounded cache key.
    from app.monthly_report_import import read_import_image
    options = {} if preserve_transparency is None else {"preserve_transparency": preserve_transparency}
    normalized = read_import_image(Path(source_path), item, line_art=line_art, **options)
    return hashlib.sha256(normalized.data).hexdigest() + "." + normalized.extension


def _unchanged_blocks(draft, inspection, source_path, *, section_key=None, source_view=None, image_digest=None):
    """Prove source facts were explicitly supplied unchanged in current inputs."""
    from app.monthly_report_operational_source import original_reference, cache_identity
    image_digest = image_digest or _source_image_digest
    by_id = {item.id: item for item in inspection.items}
    by_reference = {original_reference(inspection, item, source_view): item for item in inspection.items}
    image_cache_identity = cache_identity(inspection, source_view)
    specs = {spec.key: spec for section in draft.sections for spec in section.blocks}
    destinations = {spec.key: section.key for section in draft.sections for spec in section.blocks}
    candidates, owners = {}, {}
    for block in draft.blocks:
        if (block.source == "Omit" or block.org_nodes or block.key not in destinations
                or (section_key is not None and destinations[block.key] != section_key)):
            continue
        identities = tuple(dict.fromkeys(by_reference[ref].id for ref in block.references if ref in by_reference))
        selected = [by_id[identity] for identity in identities]
        if not selected or any(item.kind == "unsupported" for item in selected):
            continue
        if block.key in destinations and any(item.part == "word/document.xml" and item.section
                                             and item.section != destinations[block.key] for item in selected):
            continue
        text = "\n\n".join(item.text.strip() for item in selected if item.kind == "text")
        if block.text.strip() != text.strip():
            continue
        tables = [item for item in selected if item.kind == "table"]
        actual_tables = []
        if block.rows:
            spec = specs.get(block.key)
            actual_tables.append((tuple(column.title for column in spec.columns) if spec else (), block.rows))
        actual_tables.extend((table.columns, table.rows) for table in block.extra_tables)
        expected_tables = []
        from app.monthly_report_sections import table_without_prices
        for item in tables:
            table, removed = table_without_prices(item)
            if removed or table is None:
                expected_tables = None
                break
            # Import supplies display names for blank/duplicate headers. Those
            # labels do not require changing the source's actual table geometry.
            expected_tables.append((table.columns, table.rows))
        original_tables = [(item.rows[0], item.rows[1:]) for item in tables]
        if expected_tables is None or actual_tables not in (expected_tables, original_tables):
            continue
        expected_images = []
        for item in selected:
            if item.kind != "image":
                continue
            reference = image_digest(str(source_path), image_cache_identity, item,
                                             block.key not in ("cover_photo", "improvements"))
            if reference not in expected_images:
                expected_images.append(reference)
        if tuple(expected_images) != block.asset_hashes:
            continue
        expected_captions = tuple(item.text for item in selected if item.kind == "image")
        current_captions = tuple(block.asset_captions) + ("",) * max(0, len(expected_captions) - len(block.asset_captions))
        if current_captions != expected_captions:
            continue
        candidates[block.key] = set(identities)
        for identity in identities:
            owners[identity] = block.key
    by_position = defaultdict(list)
    for item in inspection.items:
        if item.part == "word/document.xml":
            by_position[item.position].append(item)
    # A native paragraph can contain several independently mapped payloads. Keep
    # it only when all are present, never because just one picture matched.
    changed = True
    while changed:
        changed = False
        supported = set().union(*candidates.values()) if candidates else set()
        for key, identities in list(candidates.items()):
            valid = True
            for identity in identities:
                item = by_id[identity]
                if item.part != "word/document.xml":
                    continue
                for neighbor in by_position[item.position]:
                    if ((neighbor.kind in ("text", "table", "image") and neighbor.id not in supported)
                            or neighbor.note.startswith("Unmapped native SmartArt")):
                        valid = False
            if not valid:
                del candidates[key]
                changed = True
    return candidates


def _patch_current_table_cells(element, source_item, expected, current):
    """Apply same-shape current cell edits without rebuilding mixed picture cells."""
    old_values = (expected.columns, *expected.rows)
    new_values = (current.columns, *current.rows)
    native_rows = element.findall(qn("w:tr"))
    if len(old_values) != len(new_values) or len(native_rows) != len(old_values):
        return False
    changes = []
    for row_index, (row, old, new) in enumerate(zip(native_rows, old_values, new_values)):
        if len(old) != len(new):
            return False
        before = row.find("./" + qn("w:trPr") + "/" + qn("w:gridBefore"))
        column = int(before.get(qn("w:val"), "0")) if before is not None else 0
        addressed = set()
        for cell in row.findall(qn("w:tc")):
            span = cell.find("./" + qn("w:tcPr") + "/" + qn("w:gridSpan"))
            width = int(span.get(qn("w:val"), "1")) if span is not None else 1
            if column >= len(old) or _text(cell) != source_item.rows[row_index][column]:
                return False
            addressed.add(column)
            if old[column] != new[column]:
                changes.append((cell, new[column]))
            column += width
        # A value cannot be written into a skipped/merged continuation column.
        if any(a != b and index not in addressed for index, (a, b) in enumerate(zip(old, new))):
            return False
    for cell, value in changes:
        _set_text(cell, value)
    return True


def build_native_docx(draft, source_path, *, asset_loader=None, section_key=None, master=False, source_view=None) -> bytes:
    """Discover saved source authority; every public operation revalidates it."""
    from app.monthly_report_source_view_runtime import render_operation
    with render_operation(draft, source_view, source_path=source_path, section_key=section_key) as operation:
        return _build_verified_native_docx(draft, source_path, asset_loader=asset_loader,
                                           section_key=section_key, master=master, operation=operation)


def _build_verified_native_docx(draft, source_path, *, asset_loader=None, section_key=None,
                                 master=False, operation):
    source_view, native_art = operation.view, operation.native_art
    if source_view is None:
        return _build_native_docx(draft, source_path, asset_loader=asset_loader, section_key=section_key, master=master)
    if master:
        raise NativeLayoutError("Experimental native source views do not support master or foreign-branding reuse.")
    from app.monthly_report_operational_source import rendition_file
    header_scope = None
    emission_proofs = []
    if native_art is not None:
        from app.monthly_report_native_art_scope import expected_header_scope, source_emission_proof
        header_scope = expected_header_scope(source_view, section_key)
    with rendition_file(source_view) as derivative:
        result = _build_native_docx(draft, derivative, asset_loader=asset_loader, section_key=section_key,
                                    master=master, source_view=source_view, original_path=Path(source_path),
                                    source_section_plan=header_scope.source_plan if header_scope else None,
                                    source_header_scope=header_scope,
                                    source_region_observer=(lambda body, pages: emission_proofs.append(
                                        source_emission_proof(header_scope, body, pages))) if header_scope else None)
    if native_art is not None:
        from app.monthly_report_pdf_image_budget import release_transient_word_memory
        release_transient_word_memory()
        from app.monthly_report_native_art_review import verify_native_art_output
        if len(emission_proofs) != 1:
            raise NativeLayoutError("Native source-region emission proof is unavailable.")
        verify_native_art_output(result, native_art, source_scope=header_scope,
                                 emission_proof=emission_proofs[0])
    return result


def _build_native_docx(draft, source_path, *, asset_loader=None, section_key=None, master=False,
                       source_view=None, original_path=None, source_section_plan=None, source_region_observer=None, source_header_scope=None) -> bytes:
    """Render current values into source geometry, or raise a visible mapping error.

    ``section_key`` selects the actual logical section's contiguous native range;
    ``cover`` selects its preface. Word's section count is not used as a page map.
    """
    from app.monthly_report_native_package import passive_docx

    source_path = Path(source_path)
    inspection = source_view.rendition_inspection if source_view is not None else inspect_docx(source_path)
    from app.monthly_report_operational_source import original_reference, cache_identity
    image_cache_identity = cache_identity(inspection, source_view)
    # Experimental operation-local cache; never retain source-view paths/items
    # across requests. Legacy source cache behavior remains unchanged.
    image_digest = (_source_image_digest if source_view is None else
                    lru_cache(maxsize=128)(_source_image_digest.__wrapped__))
    source_digest = source_view.original_inspection.sha256 if source_view is not None else inspection.sha256
    original_path = original_path or source_path
    source_ref = lambda item: original_reference(inspection, item, source_view)
    def image_matches(item, role, hashes):
        if role not in ("brand_logo", "client_logo"):
            return image_digest(str(source_path), image_cache_identity, item, role != "cover_photo") in hashes
        # Current explicit logo mapping preserves transparency; older saved
        # assets used the source suggestion's policy. Both fingerprints prove
        # this exact source artwork, not a visually guessed brand identity.
        return (image_digest(str(source_path), image_cache_identity, item, True, True) in hashes
                or image_digest(str(source_path), image_cache_identity, item, True) in hashes)

    source_ref_by_id = {item.id: source_ref(item) for item in inspection.items}
    if any(item.note.startswith("Alternate Word representations") for item in inspection.items):
        raise NativeLayoutError(
            "The native source contains unresolved modern/legacy Word artwork. "
            "A screenshot or replacement asset cannot approve a hidden alternate representation. "
            "Correct the affected source drawing before building this native report.")
    if master:
        from app.monthly_report_master_tables import from_inspection
        master_shapes = from_inspection(inspection)
        master_by_position = {shape.position: shape for shape in master_shapes}
    else:
        master_by_position = {}
    unchanged = _unchanged_blocks(draft, inspection, source_path, section_key=section_key, source_view=source_view, image_digest=image_digest)
    unchanged_items = set().union(*unchanged.values()) if unchanged else set()
    # Editing one imported table must not rebuild its unchanged neighboring
    # tables (which may have different geometry and span several pages).
    from app.monthly_report_sections import table_without_prices
    reference_prefix = f"docx:{source_digest}:"
    cover_periods = {match.group().casefold() for title in inspection.title_candidates
                     for match in _MONTH.finditer(re.split(r"\bDate\s*:", title, flags=re.IGNORECASE)[0])}
    cover_names = {re.sub(r"[^a-z0-9]+", " ", title.casefold()).strip()
                   for title in inspection.title_candidates}
    source_preparers = {re.split(r"\s+Date\s*:", re.sub(r"^\s*Prepared\s+by\s*:\s*", "", title,
                        flags=re.IGNORECASE), flags=re.IGNORECASE)[0].strip()
                        for title in inspection.title_candidates if re.match(r"^\s*Prepared\s+by\s*:", title, re.IGNORECASE)}
    current_preparer = re.split(r"\s+Date\s*:", draft.prepared_by, flags=re.IGNORECASE)[0].strip()
    preserve_issue_date = (not master and cover_periods == {draft.period.label.casefold()}
        and re.sub(r"[^a-z0-9]+", " ", draft.profile.title.casefold()).strip() in cover_names
        and current_preparer in source_preparers
        and any(ref.startswith(reference_prefix) for block in draft.blocks for ref in block.references))
    native_tables = {source_ref(item): item for item in inspection.items
                     if item.kind == "table" and item.part == "word/document.xml"}
    preserved_table_refs = set()
    redacted_table_columns = {}
    supplied_tables_by_position = {}
    retained_mixed_cells = set()
    mixed_price_headers = set()
    destinations = {spec.key: section.key for section in draft.sections for spec in section.blocks}
    for block in draft.blocks:
        if block.source == "Omit":
            continue
        for table in block.extra_tables:
            item = native_tables.get(table.reference)
            if item is not None and item.section == destinations.get(block.key):
                expected, removed = table_without_prices(item)
                if expected is not None and (table.columns, table.rows) == (expected.columns, expected.rows):
                    neighbors = [other for other in inspection.items
                                 if other.part == item.part and other.position == item.position]
                    if all(other.kind == "table" for other in neighbors):
                        preserved_table_refs.add(table.reference)
                        unchanged_items.add(item.id)
                        if removed:
                            redacted_table_columns[item.position] = removed
                            supplied_tables_by_position[item.position] = table
    # Keep independently supplied unchanged text runs even when a neighboring
    # table has an intentionally excluded pricing column.
    preserved_text = {}
    for block in draft.blocks:
        if block.source == "Omit" or block.key not in destinations:
            continue
        selected = [item for item in inspection.items if item.kind == "text"
                    and (item.section == destinations[block.key] or item.note == "Native divider artwork")
                    and source_ref(item) in block.references]
        expected = "\n\n".join(item.text.strip() for item in selected)
        if selected and block.text.strip() == expected.strip():
            preserved_text[block.key] = True
            unchanged_items.update(item.id for item in selected)
    # Current images can be unchanged even when another table in their block
    # was edited. Prove the complete selected image sequence independently.
    image_proofs = {}
    for block in draft.blocks:
        if block.source == "Omit" or block.org_nodes or block.key not in destinations:
            continue
        selected_images = [item for item in inspection.items if item.kind == "image"
                           and item.part == "word/document.xml"
                           and (item.section == destinations[block.key] or item.note == "Native divider artwork")
                           and source_ref(item) in block.references]
        image_values = [(item, image_digest(str(source_path), image_cache_identity, item,
                        block.key not in ("cover_photo", "improvements"))) for item in selected_images]
        expected_assets = tuple(dict.fromkeys(value for _, value in image_values))
        expected_captions = tuple(item.text for item in selected_images)
        captions = tuple(block.asset_captions) + ("",) * max(0, len(expected_captions) - len(block.asset_captions))
        if image_values and block.asset_hashes == expected_assets and captions == expected_captions:
            image_proofs.update({item.id: (block.key, value) for item, value in image_values})
            unchanged_items.update(item.id for item, _ in image_values)
    retained_images = defaultdict(set)
    document = Document(BytesIO(source_view.rendition_bytes if source_view is not None else passive_docx(source_path)))
    exact_client_nodes, preserved_contents_nodes = set(), set()
    source_node_items = {node:item for node, attr, item in _header_image_occurrences(
        document.part, inspection, elements=document.element.body)}
    if source_view is not None:
        from lxml import etree
        from app.monthly_report_compatibility_view import _path_node
        from app.monthly_report_operational_source import receipt_record
        root = document.element
        for record in receipt_record(source_view)["resolutions"]:
            found = [_path_node(root, record["selected_picture_path"])]
            if len(found) != 1 or found[0].tag != qn("a:blip"):
                raise NativeLayoutError("The resolved native picture occurrence cannot be bound safely.")
            original_item = next(i for i in source_view.original_inspection.items if i.id == record["original_item_id"])
            if original_item.word_section == 1:
                exact_client_nodes.add(found[0])
            elif original_item.word_section == 2:
                preserved_contents_nodes.add(found[0])
            else:
                raise NativeLayoutError("A rendition resolution is outside native cover/contents scope.")
    _clear_source_drawing_metadata(document)
    if master:
        _remove_unapproved_master_watermark_text(document)
    body = document.element.body
    original = list(body)
    # A merged technical cell may cross a removed price column. Preserve its
    # geometry only when its text is explicitly present in the supplied filtered
    # table. A cell beginning in a removed column is redacted in full.
    from app.monthly_report_content_policy import contains_price
    for position, excluded in list(redacted_table_columns.items()):
        mixed = False
        supplied = supplied_tables_by_position[position]
        source_item = next(item for item in native_tables.values() if item.position == position)
        kept_columns = [i for i in range(len(source_item.rows[0])) if i not in excluded]
        for row_index, row in enumerate(original[position - 1].findall(qn("w:tr"))):
            before = row.find("./" + qn("w:trPr") + "/" + qn("w:gridBefore"))
            column = int(before.get(qn("w:val"), "0")) if before is not None else 0
            for cell in row.findall(qn("w:tc")):
                span = cell.find("./" + qn("w:tcPr") + "/" + qn("w:gridSpan"))
                width = int(span.get(qn("w:val"), "1")) if span is not None else 1
                covered = set(range(column, column + width))
                if covered.intersection(excluded) and covered.difference(excluded) and _text(cell) and column not in excluded:
                    current_row = supplied.columns if row_index == 0 else supplied.rows[row_index - 1]
                    current_column = kept_columns.index(column) if column in kept_columns else -1
                    value = _text(cell)
                    from app.monthly_report_content_policy import logical_table_columns
                    logical_headers = logical_table_columns(source_item.rows[0])
                    if (row_index == 0 and column == 0 and width == len(logical_headers)
                            and logical_headers != tuple(source_item.rows[0])
                            and tuple(logical_headers[i] for i in kept_columns) == supplied.columns):
                        retained_mixed_cells.add((position, row_index, column))
                        mixed_price_headers.add((position, row_index, column))
                    elif (current_column >= 0 and current_column < len(current_row)
                            and current_row[current_column] == (value.strip() if row_index == 0 else value)
                            and not contains_price(value)):
                        retained_mixed_cells.add((position, row_index, column))
                    else:
                        mixed = True
                column += width
        if mixed:
            del redacted_table_columns[position]
            for reference, item in native_tables.items():
                if item.position == position:
                    preserved_table_refs.discard(reference)
                    unchanged_items.discard(item.id)
    inplace_table_positions = set()
    # Mixed technical tables retain their exact source anchor, grid, rows and
    # cell artwork only when every neighboring payload is explicitly current.
    # Bind by source reference, never by the ordinal of the remaining tables.
    for block in draft.blocks:
        if block.source == "Omit" or block.key not in destinations:
            continue
        for current in block.extra_tables:
            source_item = native_tables.get(current.reference)
            if source_item is None or source_item.section != destinations[block.key]:
                continue
            neighbors = [item for item in inspection.items
                         if item.part == source_item.part and item.position == source_item.position]
            if not any(item.kind == "image" for item in neighbors):
                continue
            if any(item.kind == "text" and item.id not in unchanged_items for item in neighbors):
                continue
            expected, removed = table_without_prices(source_item)
            if expected is None or removed:
                continue
            element = original[source_item.position - 1]
            if element.tag != qn("w:tbl") or not _patch_current_table_cells(element, source_item, expected, current):
                continue
            unchanged_items.add(source_item.id)
            preserved_table_refs.add(current.reference)
            inplace_table_positions.add(source_item.position)
    # Word inherits each header/footer kind independently. Materialize those
    # references before slicing so a section preview retains its page chrome.
    inherited = {}
    first_word_end = None
    for position, element in enumerate(original, 1):
        properties = element if element.tag == qn("w:sectPr") else element.find(".//" + qn("w:sectPr"))
        if properties is None:
            continue
        if first_word_end is None:
            first_word_end = position
        present = {}
        for node in properties:
            if node.tag in (qn("w:headerReference"), qn("w:footerReference")):
                present[(node.tag, node.get(qn("w:type"), "default"))] = node
        for key, node in inherited.items():
            if section_key and key not in present:
                properties.insert(0, deepcopy(node))
        inherited.update({key: deepcopy(node) for key, node in present.items()})
    items = defaultdict(list)
    for item in inspection.items:
        if item.part == "word/document.xml":
            items[item.position].append(item)
    section_by_pos, section, headings = {}, "cover", set()
    blocks = {block.key: block for block in draft.blocks}
    supplied_notes = {item.note: item for item in inspection.items
                      if item.note.startswith(("footnote:", "endnote:")) and item.text.strip()
                      and any(block.source != "Omit" and source_ref(item) in block.references
                              and item.text.strip() in block.text for block in draft.blocks)}
    source_titles = {title.casefold().strip() for title in inspection.title_candidates
                     if len(title) < 70 and not _MONTH.search(title)
                     and not re.search(r"prepared|section|enfra|monthly|maintenance|insert|confidential|contents", title, re.IGNORECASE)}
    current_columns = {spec.key: tuple(column.title for column in spec.columns)
                       for section in draft.sections for spec in section.blocks}

    brand = blocks.get("brand_logo")
    cover_roles = {}
    unchanged_master_brand = unchanged_client_brand = False
    if source_view is None:
        cover_roles = {}
        cover_items = [item for item in inspection.items if item.kind == "image"
                       and item.part == "word/document.xml"
                       and (not item.section or item.position <= (first_word_end or len(original)))]
        first_cover_images = [item for item in cover_items
                              if item.position <= (first_word_end or len(original))]
        primary_brand = next((item.image_part for item in first_cover_images
                              if item.image_height and item.image_width / item.image_height > 4), None)
        photo_candidates = [item for item in first_cover_images if item.image_part != primary_brand
                            and item.image_width >= 256 and item.image_height >= 256]
        primary_photo = max(photo_candidates, key=lambda item: item.image_width * item.image_height,
                            default=None)
        # Explicit import destinations are authoritative. Aspect ratio cannot tell
        # a wide hospital logo from ENFRA branding or a landscape hospital photo.
        for item in cover_items:
            for role in ("brand_logo", "client_logo", "cover_photo"):
                block = blocks.get(role)
                if block and source_ref(item) in block.references:
                    cover_roles[item.image_part] = role
                    break
        from app.monthly_report_static_branding import is_approved_company_slogan
        for item in cover_items:
            if item.image_part in cover_roles:
                continue
            # This verified company-only wordmark on the contents page is *not*
            # a second current client's logo. Source-asset identity, not aspect
            # ratio or position, establishes this narrow static-design exception.
            if is_approved_company_slogan(original_path, source_digest, item):
                cover_roles[item.image_part] = "company_slogan"
                continue
            for role in ("brand_logo", "client_logo", "cover_photo"):
                block = blocks.get(role)
                if block and block.asset_hashes and image_matches(item, role, block.asset_hashes):
                    cover_roles[item.image_part] = role
                    break
            if item.image_part not in cover_roles and item.image_height:
                ratio = item.image_width / item.image_height
                if master:
                    cover_roles[item.image_part] = ("brand_logo" if item.image_part == primary_brand else
                                                    "cover_photo" if primary_photo and item.image_part == primary_photo.image_part else
                                                    "client_logo")
                else:
                    cover_roles[item.image_part] = "brand_logo" if ratio > 4 else "client_logo" if ratio > 2 else "cover_photo"

        def unchanged_cover(role):
            block = blocks.get(role)
            return bool(block and block.source != "Omit" and block.asset_hashes and any(
                cover_roles.get(item.image_part) == role
                and image_matches(item, role, block.asset_hashes)
                for item in cover_items))

        unchanged_master_brand = unchanged_cover("brand_logo")
        unchanged_client_brand = unchanged_cover("client_logo")

    def cover_images(element, group):
        roles = cover_roles
        logo_replacements = []
        for node in list(element.iter()):
            if node.tag not in (qn("a:blip"), "{urn:schemas-microsoft-com:vml}imagedata"):
                continue
            attr = qn("r:embed") if node.tag == qn("a:blip") else qn("r:id")
            rid = node.get(attr)
            rel = document.part.rels.get(rid)
            target = str(rel.target_part.partname).lstrip("/") if rel and not rel.is_external else ""
            if source_view is not None:
                source_item = source_node_items.get(node)
                if source_item is None:
                    raise NativeLayoutError("Native cover artwork has no exact source occurrence binding.")
                reference = source_ref(source_item)
                explicit = {block.key for block in draft.blocks if reference in block.references
                            and block.key in ("brand_logo", "client_logo", "cover_photo")}
                if len(explicit) > 1 or (node in exact_client_nodes and explicit - {"client_logo"}):
                    raise NativeLayoutError("Conflicting exact native cover image destinations.")
                role = next(iter(explicit)) if explicit else "client_logo" if node in exact_client_nodes else ""
                block = blocks.get(role)
                if node in preserved_contents_nodes and not explicit:
                    continue  # Exact resolved occurrence, same-profile scope verified at entry.
                if not role:
                    continue  # Existing supported artwork is retained, never guessed by shape/hash.
                if block is None:
                    raise NativeLayoutError("The proved native client slot has no current artwork decision.")
                if block.source == "Omit":
                    node.getparent().remove(node)
                    continue
                if not block.asset_hashes:
                    raise NativeLayoutError("The exact native image replacement has no saved artwork.")
                if image_matches(source_item, role, block.asset_hashes):
                    if role == "cover_photo":
                        _fit_unchanged_cover_photo(node, source_path, source_item, document.part)
                    elif role in ("brand_logo", "client_logo"):
                        logo_replacements.append(_contain_unchanged_cover_logo(node, source_path, source_item, document.part, defer=True))
                    continue
                if asset_loader is None:
                    raise NativeLayoutError("Cover images require the saved asset resolver.")
                new_rid, image = document.part.get_or_add_image(BytesIO(asset_loader(block.asset_hashes[0])))
                node.set(attr, new_rid)
                if role == "cover_photo":
                    _fit_cover_photo(node, image.px_width, image.px_height)
                elif role in ("brand_logo", "client_logo"):
                    logo_replacements.append((node, image.px_width, image.px_height))
                continue
            # Authorize this new repair by the exact selected occurrence,
            # before the legacy shared-media role heuristic can override it.
            photo = blocks.get("cover_photo")
            occurrence = source_node_items.get(node)
            if photo and photo.source != "Omit" and photo.asset_hashes and occurrence is not None:
                reference = source_ref(occurrence)
                explicit = {block.key for block in draft.blocks if reference in block.references}
                if (explicit == {"cover_photo"}
                        and image_digest(str(source_path), image_cache_identity, occurrence, False) in photo.asset_hashes):
                    _fit_unchanged_cover_photo(node, source_path, occurrence, document.part)
                    continue
            if occurrence is not None:
                reference = source_ref(occurrence)
                mapped_roles = {block.key for block in draft.blocks if reference in block.references}
                if len(mapped_roles) == 1 and next(iter(mapped_roles)) in ("brand_logo", "client_logo"):
                    logo_role = next(iter(mapped_roles))
                    logo = blocks[logo_role]
                    if (logo.source != "Omit" and logo.asset_hashes
                            and image_matches(occurrence, logo_role, logo.asset_hashes)):
                        logo_replacements.append(_contain_unchanged_cover_logo(node, source_path, occurrence, document.part, defer=True))
                        continue
            # Explicit roles belong to occurrences, not shared relationship bytes.
            # A replacement must not overwrite an unselected same-media sibling.
            exact_roles = ({b.key for b in draft.blocks if source_ref(occurrence) in b.references}
                           if occurrence is not None else set())
            selected_logo_media = any(
                item.image_part == target and any(source_ref(item) in b.references
                    for b in draft.blocks if b.key in ("brand_logo", "client_logo"))
                for item in cover_items)
            cover_slots = {"brand_logo", "client_logo", "cover_photo"}
            if len(exact_roles) > 1 or (selected_logo_media and not exact_roles & cover_slots):
                continue
            role = next(iter(exact_roles)) if len(exact_roles) == 1 and exact_roles <= cover_slots else roles.get(target)
            if role == "company_slogan":
                continue  # Reviewed, immutable company furniture; never a client mark.
            block = blocks.get(role)
            if block and block.source != "Omit" and block.asset_hashes:
                if ((role == "client_logo" and unchanged_client_brand)
                        or (role == "brand_logo" and unchanged_master_brand)):
                    continue
                supplied = next((item for item in group if item.kind == "image" and item.image_part == target), None)
                if supplied is not None and image_matches(supplied, role, block.asset_hashes):
                    continue  # Unselected/unproved occurrences retain their original geometry.
                if asset_loader is None:
                    raise NativeLayoutError("Cover images require the saved asset resolver.")
                new_rid, image = document.part.get_or_add_image(BytesIO(asset_loader(block.asset_hashes[0])))
                node.set(attr, new_rid)
                for layer in node.iter():
                    if layer.get(qn("r:embed")):
                        layer.set(qn("r:embed"), new_rid)
                    layer.attrib.pop(qn("r:link"), None)
                if role == "cover_photo":
                    _fit_cover_photo(node, image.px_width, image.px_height)
                elif role in ("brand_logo", "client_logo"):
                    # The existing replacement path has already authorized
                    # these bytes; catalog defaults need proportional fit too.
                    logo_replacements.append((node, image.px_width, image.px_height))
            elif role != "brand_logo" and (master or block is not None):
                # A global master must never lend another client's logo/photo.
                node.getparent().remove(node)

        # All alternate representations must have their final relationships
        # before orientation/atomic geometry guards inspect their peers.
        plans = [(_contain_cover_logo(*entry, document.part, defer=True) if isinstance(entry, tuple) else entry)
                 for entry in logo_replacements]
        for plan in plans:
            if plan is not None:
                plan()

    if source_section_plan is not None:
        # This exact immutable-source plan also derives expected header closure
        # before rendering. Do not re-guess selection after changing payloads.
        section_by_pos = dict(source_section_plan.section_by_position)
        headings = set(source_section_plan.headings)
        if first_word_end != source_section_plan.first_word_end:
            raise NativeLayoutError("Native source section boundaries changed before rendering.")
    else:
        for position, element in enumerate(original, 1):
            group = items[position]
            detected = next((i.section for i in group if i.section), "")
            if detected:
                section = detected
            section_by_pos[position] = section
            if (any(i.note == "Native divider artwork" for i in group)
                    and _is_native_divider_background(element, _text(element))):
                headings.add(position)
            if any(i.note == "Native section heading" or
                   (i.kind == "text" and not i.suggested_slot and _heading(i.text)) for i in group):
                headings.add(position)
        # Divider artwork is sometimes anchored before its heading, in a separate
        # paragraph (even dozens of empty paragraphs earlier in the same Word section).
        for heading_position in sorted(headings):
            for position in range(heading_position - 1, 0, -1):
                element = original[position - 1]
                if element.find(".//" + qn("w:sectPr")) is not None:
                    break
                background = _is_native_divider_background(element, _text(element))
                if background:
                    headings.add(position)
                    for previous in range(position, heading_position):
                        section_by_pos[previous] = section_by_pos[heading_position]
                    break
    divider_word_sections = {item.word_section for position in headings for item in items[position]}
    if section_key and section_key not in set(section_by_pos.values()):
        raise NativeLayoutError(f"The source layout does not contain the {section_key} section.")
    from app.monthly_report_native_notes import note_proofs, plain_visible, without_trailing_note
    native_note_proofs = note_proofs(document, original, inspection, section_by_pos, _DEFAULT_TEXT)
    from app.monthly_report_native_table_context import supplied_table_context
    # Preserved placement also includes mixed picture tables edited in place.
    # Only exact current cell values may authorize retaining adjacent context;
    # leave the broader placement/redaction receipt unchanged.
    context_table_refs = set(preserved_table_refs)
    for block in draft.blocks:
        if block.source == "Omit":
            continue
        for table in block.extra_tables:
            if table.reference not in context_table_refs:
                continue
            expected, _ = table_without_prices(native_tables[table.reference])
            if expected is None or (table.columns, table.rows) != (expected.columns, expected.rows):
                context_table_refs.discard(table.reference)
    table_context = {} if master else supplied_table_context(
        draft, original, items, section_by_pos, native_tables, context_table_refs,
        styles=document.styles.element)
    unchanged_items.update(item.id for context, _ in table_context.values() for item in context)
    context_anchors = {key: original[context[0].position - 1]
                       for key, (context, _) in table_context.items()}
    from app.monthly_report_native_pages import NativePages
    unchanged_positions = {item.position for item in inspection.items
                           if item.part == "word/document.xml" and item.id in unchanged_items}
    page_plan = NativePages(body, original, section_by_pos, headings, unchanged_positions, fresh=master, source_header_scope=source_header_scope)
    from app.monthly_report_native_empty_activity import empty_activity_panels
    empty_panels = empty_activity_panels(draft, original, section_by_pos, unchanged_positions, headings)
    cleared_activity_paragraphs = set()
    work_order_labels = _work_order_table_labels(original, items, section_by_pos)
    from app.monthly_report_native_text import (narrative_prototype, write_narrative,
                                                 header_contains_label)
    narrative_styles = {style.style_id: style.element for style in document.styles}
    # Freeze role boundaries before source text is cleared. Typography borrowing
    # must never look across a native section or independently mapped block.
    narrative_scopes = {
        element: (section_by_pos[position], next((item.suggested_slot for item in items[position]
            if item.suggested_slot and not item.suggested_slot.startswith('divider_')),
            _DEFAULT_TEXT.get(section_by_pos[position])))
        for position, element in enumerate(original, 1)
    }
    from app.monthly_report_narrative_table_roles import table_body_roles
    narrative_exclusions, neutral_narrative_scopes = table_body_roles(
        original, narrative_styles, document.styles.element, narrative_scopes, static_labels=_STATIC)
    narrative_role_anchors = {}
    for element in original:
        if element in narrative_exclusions:
            narrative_role_anchors.setdefault(narrative_scopes[element], element)
    text_prototypes = defaultdict(list)
    image_prototypes = defaultdict(list)
    table_prototypes = defaultdict(list)
    table_by_reference = {}  # pinned physical source table, never an ordinal guess
    section_ends = {}
    section_text = defaultdict(list)
    section_images = defaultdict(list)
    all_tables = [deepcopy(e) for e in original if e.tag == qn("w:tbl")
                  and not any(True for _ in e.iter(qn("a:blip")))]
    global_images = [deepcopy(paragraph) for position, e in enumerate(original, 1)
                     if section_by_pos[position] != "cover" and position not in headings
                     for paragraph in e.iter(qn("w:p"))
                     if any(True for _ in paragraph.iter(qn("a:blip")))
                     and not any(ancestor.tag == qn("w:p") for ancestor in paragraph.iterancestors())]
    global_text = [prototype for position, e in enumerate(original, 1)
                   if section_by_pos[position] != "cover" and position not in headings
                   and e.tag == qn("w:p") and e not in narrative_exclusions
                   and (prototype := narrative_prototype(e, narrative_styles, _STATIC, scopes=narrative_scopes)) is not None]
    for position, element in enumerate(original, 1):
        sec = section_by_pos[position]
        if element.tag == qn('w:sectPr'):
            continue
        section_ends[sec] = element
        if section_key and (sec != section_key or (section_key == "cover" and position > (first_word_end or len(original)))):
            body.remove(element)
            continue
        group = items[position]
        current_payload = [item for item in group if item.kind in ("text", "table", "image")]
        if not (current_payload and all(item.id in unchanged_items for item in current_payload)
                and not any(item.note.startswith("Unmapped native SmartArt") for item in group)):
            _clear_hyperlinks(element)
        if sec == "cover":
            # Cover/company marks have a separate asset binding. Unsupported
            # SmartArt there must not lend source client facts to a fresh report.
            _clear_diagrams(element)
            _clear_hyperlinks(element)
            from app.monthly_report_native_design import clear_native_design_placeholders
            clear_native_design_placeholders(element, cover=any(
                item.kind == "native_layout" and item.note == "Native cover title and artwork" for item in group))
            if position <= (first_word_end or len(original)):
                _cover(element, draft, preserve_issue_date=preserve_issue_date)
            cover_images(element, group)
            continue
        group = items[position]
        if position in headings:
            if _is_native_divider_background(element, _text(element)):
                from app.monthly_report_native_design import clear_native_design_placeholders
                clear_native_design_placeholders(element)
                # Older saved mappings may have mistaken the images in a
                # grouped divider for organization photos. Their proven native
                # frames are already present; do not append another full-page
                # canvas for each previously selected constituent image.
                for item in group:
                    if item.kind != "image" or item.note != "Native divider artwork":
                        continue
                    for block in draft.blocks:
                        if (block.source != "Omit" and source_ref(item) in block.references
                                and block.asset_hashes):
                            digest = image_digest(str(source_path), image_cache_identity, item,
                                                         block.key not in ("cover_photo", "improvements"))
                            if digest in block.asset_hashes:
                                retained_images[block.key].add(digest)
            diagrams = [item for item in group if item.note == "Validated native SmartArt text"]
            unproved_diagram = (any(True for _ in element.iter(
                "{http://schemas.openxmlformats.org/drawingml/2006/diagram}relIds"))
                and (not diagrams or any(item.id not in unchanged_items for item in diagrams)
                     or any(item.note.startswith("Unmapped native SmartArt") for item in group)))
            if (unproved_diagram or (any(item.kind == "image" for item in group)
                    and not _is_native_divider_background(element, _text(element)))):
                _clear_images(element)
            continue
        substantive = [item for item in group if item.kind in ("text", "table", "image")]
        if (substantive and all(item.id in unchanged_items for item in substantive)
                and not any(item.note.startswith("Unmapped native SmartArt") for item in group)):
            for item in substantive:
                if item.id in image_proofs:
                    owner, digest = image_proofs[item.id]
                    retained_images[owner].add(digest)
            for row_index, row in enumerate(element.findall(qn("w:tr"))):
                before = row.find("./" + qn("w:trPr") + "/" + qn("w:gridBefore"))
                logical_column = int(before.get(qn("w:val"), "0")) if before is not None else 0
                excluded = set(redacted_table_columns.get(position, ()))
                for cell in row.findall(qn("w:tc")):
                    span = cell.find("./" + qn("w:tcPr") + "/" + qn("w:gridSpan"))
                    width = int(span.get(qn("w:val"), "1")) if span is not None else 1
                    if (excluded.intersection(range(logical_column, logical_column + width))
                            and (position, row_index, logical_column) not in retained_mixed_cells):
                        _set_text(cell, "")
                    if (position, row_index, logical_column) in mixed_price_headers:
                        for node in cell.iter(qn("w:t")):
                            if node.text:
                                node.text = re.sub(r"Cost", "", node.text, flags=re.IGNORECASE)
                    logical_column += width
            continue
        text = _text(element)
        slots = [i.suggested_slot for i in group if i.suggested_slot and not i.suggested_slot.startswith("divider_")]
        slot = slots[0] if slots else _DEFAULT_TEXT.get(sec)
        picture_table = element.tag == qn("w:tbl") and any(item.kind == "image" for item in group)
        if element.tag == qn("w:tbl") and not picture_table:
            table_slot = next((i.suggested_slot for i in group if i.kind == "table" and i.suggested_slot), slot)
            native_shape = master_by_position.get(position)
            prototype = deepcopy(element)
            table_prototypes[table_slot].append((element, prototype))
            if native_shape is not None and native_shape.slot == table_slot:
                table_by_reference[native_shape.reference] = (element, prototype, native_shape)

        if any(i.kind == "image" for i in group):
            image_slot = next((i.suggested_slot for i in group if i.kind == "image" and i.suggested_slot), slot)
            frames = [element] if element.tag == qn("w:p") else [
                paragraph for paragraph in element.iter(qn("w:p"))
                if any(True for _ in paragraph.iter(qn("a:blip")))
                and not any(ancestor.tag == qn("w:p") for ancestor in paragraph.iterancestors())]
            for frame in frames:
                image_prototypes[image_slot].append((frame, deepcopy(frame)))
                section_images[sec].append(deepcopy(frame))
        if position in inplace_table_positions:
            # The values were updated at this exact source reference. When the
            # neighboring images lack current proof, replace only those frames;
            # never erase/reappend the already current technical table itself.
            _clear_images(element)
            continue
        if element.tag == qn("w:p") and text and text.casefold().strip() not in _STATIC:
            if re.fullmatch(r"(?:Monthly Training Update [–—-] )?" + _MONTH.pattern + r"(?: Activity)?", text, re.IGNORECASE):
                _set_text(element, _MONTH.sub(draft.period.label, text))
                page_plan.support.add(element)
            else:
                label = re.fullmatch(r"(Client Utility Rate Analysis)(?:\s+Author:.*)?", text, re.IGNORECASE)
                prototype = None if label or element in narrative_exclusions else narrative_prototype(element, narrative_styles, _STATIC, scopes=narrative_scopes)
                if prototype is not None:
                    text_prototypes[slot].append((element, prototype))
                    section_text[sec].append(deepcopy(prototype))
                _set_text(element, "")
                if sec == "activity":
                    cleared_activity_paragraphs.add(element)
                if label:
                    region = page_plan.owner.get(element)
                    properties = page_plan.regions[region].properties if region is not None else document.sections[-1]._sectPr
                    if not header_contains_label(document, properties, label[1]):
                        _set_text(element, label[1])
                        page_plan.support.add(element)
                for paragraph in element.iter(qn("w:p")):
                    if not _text(paragraph):
                        properties = paragraph.find(qn("w:pPr"))
                        if properties is not None:
                            for numbering in list(properties.findall(qn("w:numPr"))):
                                properties.remove(numbering)
                            numbering = OxmlElement("w:numPr")
                            number = OxmlElement("w:numId"); number.set(qn("w:val"), "0")
                            numbering.append(number); properties.append(numbering)
        # Every non-heading body picture/table is old report content.
        _clear_images(element, preserve_decoration=any(
            item.word_section in divider_word_sections for item in group))
        if element.tag == qn("w:tbl"):
            if picture_table:
                for cell in element.iter(qn("w:tc")):
                    _set_text(cell, "")
                continue
            shape = master_by_position.get(position)
            columns = shape.columns if shape and shape.slot == table_slot else current_columns.get(table_slot)
            # For a same-width custom report design, a neutralized unsafe
            # source header can use the already-selected *current* schema's
            # non-monetary label. Otherwise retain the neutral Detail label.
            selected = current_columns.get(table_slot, ())
            if shape and len(selected) == len(columns):
                from app.monthly_report_content_policy import price_column
                columns = tuple(selected[n] if value.startswith("Detail ")
                                and not price_column(selected[n]) else value
                                for n, value in enumerate(columns))
            if master and columns:
                _table(element, columns, (), physical_indices=shape.indices if shape else (),
                       continuation=bool(shape and shape.continuation))
                continue
            native_rows = element.findall(qn("w:tr"))
            header_rows = 1
            if len(native_rows) > 1 and len(native_rows[0].findall(qn("w:tc"))) < len(native_rows[1].findall(qn("w:tc"))):
                header_rows = 2
            for index, row in enumerate(native_rows):
                for cell in row.iter(qn("w:tc")):
                    if index < header_rows:
                        header_text = _text(cell)
                        safe = (len(header_text) < 85 and header_text.casefold().strip() not in source_titles and not re.search(
                            r"@|\b(?:director|manager|price|cost|amount|total recommendations|hospital|medical|regional)\b|\d{3}[ .()-]*\d{3}[ .-]*\d{4}",
                            header_text, re.IGNORECASE))
                        if safe:
                            for text_node in cell.iter(qn("w:t")):
                                text_node.text = _MONTH.sub(draft.period.label, text_node.text or "")
                        else:
                            _set_text(cell, "")
                    else:
                        _set_text(cell, "")

    def append_native(sec, prototype, *, before=None):
        clone = deepcopy(prototype)
        # A copied payload must not introduce another Word section boundary.
        for properties in list(clone.iter(qn("w:sectPr"))):
            properties.getparent().remove(properties)
        _clear_images(clone)
        _set_text(clone, "")
        anchor = before if before is not None else section_ends[sec]
        anchor.addprevious(clone)
        page_plan.assigned(clone, anchor)
        return clone

    def text_anchor(sec, *, before=None, role=None):
        choices = [] if (sec, role) in neutral_narrative_scopes else (section_text.get(sec) or global_text)
        if choices:
            prototype = choices[0]
        else:
            prototype = OxmlElement("w:p")
        return append_native(sec, prototype, before=before), prototype

    from app.monthly_report_native_narrative_visibility import narrative_visibility_writer
    visible_narrative = narrative_visibility_writer(document)
    specs = {b.key: (s.key, b) for s in draft.sections for b in s.blocks}
    water_update_anchor = None
    water_order = next(([b.key for b in s.blocks] for s in draft.sections if s.key == "water"), [])
    water_update = blocks.get("water_monthly_update")
    if (water_update is not None and water_update.source != "Omit" and water_update.text
            and water_update.key not in unchanged and (not section_key or section_key == "water")
            and "water_monthly_update" in water_order and "water_reports" in water_order
            and water_order.index("water_monthly_update") < water_order.index("water_reports")):
        reports = blocks.get("water_reports")
        report_refs = (set(reports.references) | {t.reference for t in reports.extra_tables}
                       if reports is not None else set())
        vendor_anchor = next((element for position, element in enumerate(original, 1)
                              if section_by_pos.get(position) == "water" and position not in headings
                              and element.getparent() is body and any(
                                  item.kind in ("text", "table", "image") and
                                  (item.suggested_slot == "water_reports" or source_ref(item) in report_refs)
                                  for item in items[position])), None)
        old_anchors = text_prototypes.get("water_monthly_update", ())
        if vendor_anchor is not None and (not old_anchors or
                original.index(vendor_anchor) < original.index(old_anchors[0][0])):
            # Reserve the summary's position before vendor frames are replaced
            # or retained at their source anchors, even if draft blocks arrive
            # in a different order. Never move existing custom schema order.
            water_update_anchor = text_anchor("water", before=vendor_anchor, role="water_monthly_update")
    render_blocks = list(draft.blocks)
    if draft.follow_ups:
        from app.monthly_report_followups import report_text
        from app.monthly_report_model import ResolvedBlock
        for sec in ("issues", "proposals"):
            texts = [report_text(item) for item in draft.follow_ups
                     if item.included and ("issues" if item.category == "issue" else "proposals") == sec]
            if texts:
                key = "native_followups_" + sec
                specs[key] = (sec, type("TextSpec", (), {"type": "rich_text", "columns": ()})())
                render_blocks.append(ResolvedBlock(key, "This month", text="\n".join(texts)))
    for block in render_blocks:
        if block.source == "Omit" or block.key not in specs or block.key in unchanged:
            continue
        sec, spec = specs[block.key]
        if section_key and sec != section_key:
            continue
        if not any((block.text, block.rows, block.extra_tables, block.asset_hashes, block.org_nodes)):
            continue
        if sec not in section_ends:
            raise NativeLayoutError(f"The source has no native layout for {sec}.")
        text_value = "" if block.key in preserved_text else block.text
        if text_value and block.key in table_context:
            text_value = table_context[block.key][1]
        relocated_notes = []
        duplicate_note_lines = []
        for identity, item in supplied_notes.items():
            kind, note_id = identity.split(":", 1)
            if (source_ref(item) in block.references
                    and any(node.get(qn("w:id")) == note_id for node in body.iter(qn("w:" + kind + "Reference")))):
                # Current prose is not a bag of note substrings. Delete only
                # its exact trailing duplicate after proving the same native
                # occurrence survives in this block and remains visible.
                proof = native_note_proofs.get(identity)
                current_refs = [node for node in body.iter(qn("w:" + kind + "Reference"))
                                if node.get(qn("w:id")) == note_id]
                anchors = text_prototypes.get(block.key, ())
                owned_refs = [node for anchor, _ in anchors
                              for node in anchor.iter(qn("w:" + kind + "Reference"))
                              if node.get(qn("w:id")) == note_id]
                original_line = without_trailing_note(block.text, item.text.strip())
                residual = without_trailing_note(text_value, item.text.strip())
                if (proof is not None and proof[0] == block.key
                        and len(current_refs) == 1 and current_refs[0] is proof[1]
                        and len(owned_refs) == 1 and owned_refs[0] is proof[1]
                        and original_line is not None and residual is not None):
                    duplicate_note_lines.append(item.text.strip())
                if block.key not in preserved_text:
                    references = [node for anchor, _ in text_prototypes.get(block.key, ())
                                  for node in anchor.iter(qn("w:" + kind + "Reference"))
                                  if node.get(qn("w:id")) == note_id]
                    if references:
                        relocated_notes.append(deepcopy(references[0]))
                        for node in references:
                            node.getparent().remove(node)
        if text_value or relocated_notes:
            anchors = text_prototypes.get(block.key)
            context_anchor = context_anchors.get(block.key)
            role_anchor = narrative_role_anchors.get((sec, block.key))
            if block.key == "water_monthly_update" and water_update_anchor is not None:
                anchors = [water_update_anchor]
            elif context_anchor is not None and (not anchors or
                    original.index(context_anchor) < original.index(anchors[0][0])):
                # A newly added note still leads its section when the source's
                # only prior prose is now retained beside its exact table.
                anchors = [text_anchor(sec, before=context_anchor, role=block.key)]
            elif not anchors:
                anchors = [text_anchor(sec, before=role_anchor, role=block.key)]
            elif role_anchor is not None and original.index(role_anchor) < original.index(anchors[0][0]):
                # Borrow body typography without moving a leading note behind
                # the table that established this source caption's role.
                prototype = anchors[0][1]
                anchors = [(append_native(sec, prototype, before=role_anchor), prototype)]
            anchor, prototype = anchors[0]
            prototype = visible_narrative(prototype)
            # Context placement can select a different target than the old
            # owning paragraph. Prove the actual final narrative prototype,
            # after that selection, before removing any current note line.
            if plain_visible(prototype, document.styles.element):
                for note_line in duplicate_note_lines:
                    residual = without_trailing_note(text_value, note_line)
                    if residual is not None:
                        text_value = residual
            # Current prose is one flow, not one old canvas/page per new line.
            # Unchanged, source-proved text has already bypassed this writer.
            for index, value in enumerate(text_value.splitlines() or [""]):
                if index:
                    clone = deepcopy(prototype)
                    anchor.addnext(clone)
                    page_plan.assigned(clone, anchor)
                    anchor = clone
                write_narrative(anchor, prototype, value)
                if index == 0:
                    # Explicitly retained foot/endnotes move with rewritten
                    # prose; source-only references and reviewer history do not.
                    for note in relocated_notes:
                        run = OxmlElement("w:r")
                        properties = OxmlElement("w:rPr")
                        align = OxmlElement("w:vertAlign"); align.set(qn("w:val"), "superscript")
                        properties.append(align); run.append(properties); run.append(note)
                        anchor.append(run)
                page_plan.emitted(anchor)
        tables = []
        if block.rows:
            tables.append((tuple(c.title for c in spec.columns), block.rows, ""))
        tables.extend((t.columns, t.rows, t.reference) for t in block.extra_tables
                      if t.reference not in preserved_table_refs and (t.rows or not master))
        if tables:
            anchors = table_prototypes.get(block.key)
            if not anchors:
                width = len(tables[0][0])
                prototype = next((table for table in all_tables
                                  if len(table.find(qn("w:tr")).findall(qn("w:tc"))) == width), None)
                if prototype is None:
                    prototype = all_tables[0] if all_tables else None
                if prototype is None:
                    raise NativeLayoutError(f"The source has no native table geometry for {block.key}.")
                anchors = [(append_native(sec, prototype), prototype)]
            last = None
            emitted = set()
            for index, (columns, rows, table_reference) in enumerate(tables):
                shape = None
                if master and table_reference.startswith("master-table:"):
                    if table_reference not in table_by_reference:
                        raise NativeLayoutError("The pinned master table reference no longer identifies this source's native grid.")
                    anchor, prototype, shape = table_by_reference[table_reference]
                    if table_reference in emitted:
                        raise NativeLayoutError("The same native master table cannot receive two independent edits.")
                    emitted.add(table_reference)
                    if columns != shape.columns:
                        # A user may edit column labels, but not silently change
                        # a physical schema's width in a separate native frame.
                        if len(columns) != len(shape.columns):
                            raise NativeLayoutError("The edited master table has a different physical column count.")
                else:
                    anchor, prototype = anchors[min(index, len(anchors) - 1)]
                    if index >= len(anchors):
                        anchor = deepcopy(prototype)
                        last.addnext(anchor)
                        page_plan.assigned(anchor, last)
                _table(anchor, columns, rows, physical_indices=shape.indices if shape else (),
                       continuation=bool(shape and shape.continuation),
                       keep_all_rows_intact=(block.key == "contact_matrix"), template=prototype)
                page_plan.emitted(anchor, kind="table")
                label = work_order_labels.get(anchor)
                if label is not None and table_reference == source_ref_by_id[label[1]]:
                    # Empty-row normalization can re-emit an otherwise unchanged
                    # imported grid. Its exact source label remains useful, but
                    # an omitted/relocated table must not leave that label behind.
                    page_plan.support.add(label[0])
                from app.monthly_report_training_style import style_training_table
                style_training_table(anchor, table_reference)
                from app.monthly_report_training import matrix_title
                title = matrix_title(table_reference)
                if title:
                    caption = OxmlElement("w:p")
                    properties = OxmlElement("w:pPr")
                    properties.append(OxmlElement("w:keepNext"))
                    caption.append(properties)
                    run = OxmlElement("w:r")
                    text = OxmlElement("w:t"); text.text = title
                    run.append(text); caption.append(run)
                    anchor.addprevious(caption)
                    page_plan.assigned(caption, anchor)
                    page_plan.emitted(caption, kind="table-title")
                last = anchor
        images = []
        image_indexes = []
        if block.org_nodes:
            from app.monthly_report_visuals import org_groups, org_page
            images = [org_page(block.org_nodes, n) for n in range(len(org_groups(block.org_nodes)))]
        elif block.asset_hashes:
            if asset_loader is None:
                raise NativeLayoutError("Native images require the saved asset resolver.")
            image_indexes = [index for index, ref in enumerate(block.asset_hashes)
                             if ref not in retained_images[block.key]]
            images = [asset_loader(block.asset_hashes[index]) for index in image_indexes]
        if images:
            from app.monthly_report_asset_layout import document_page_indexes, document_caption_problems, style_document_caption
            if document_caption_problems(block, draft.sources):
                raise NativeLayoutError("Shorten the source-page caption to at most three lines before generating.")
            document_indexes = document_page_indexes(block, draft.sources)
            anchors = image_prototypes.get(block.key)
            if not anchors:
                choices = section_images.get(sec) or global_images
                if (not choices and block.key in ("improvements", "cover_photo", "brand_logo", "client_logo")
                        and not all(i in document_indexes for i in image_indexes)):
                    raise NativeLayoutError(f"The source has no native picture frame for {block.key}.")
                # Technical content needs the section's page geometry, not an
                # old photograph. Empty chart pages can accept a new inline frame.
                prototype = choices[0] if choices else OxmlElement("w:p")
                anchors = [(append_native(sec, prototype), prototype)]
            last = None
            photo_index = 0
            previous_document = False
            emitted_photo_roots = set()
            has_ordinary_photos = any(i not in document_indexes for i in image_indexes)
            unused_document_anchors = set()
            for index, data in enumerate(images):
                asset_index = image_indexes[index] if image_indexes else index
                source_page = asset_index in document_indexes
                anchor_index = photo_index if document_indexes else index
                anchor, prototype = anchors[min(anchor_index, len(anchors) - 1)]
                origin = anchor if anchor_index < len(anchors) else last
                origin_root = page_plan._root(origin)
                nested_document = source_page and origin_root is not origin
                future_photos = any((image_indexes[n] if image_indexes else n) not in document_indexes
                                    for n in range(index + 1, len(images)))
                if nested_document and (future_photos or (has_ordinary_photos and origin_root not in emitted_photo_roots)):
                    raise NativeLayoutError('A source page cannot split this native photo grid safely. Move the source page after its photos, or use a separate report-page section.')
                document_origin = (last if source_page and previous_document and last is not None
                                   and page_plan._root(last) is last else origin_root if nested_document else origin)
                technical = source_page or block.key not in ("improvements", "cover_photo", "brand_logo", "client_logo")
                if technical:
                    tail_space = (page_plan.trailing_picture_space(document,
                                  document_origin)
                                  if source_page and index == len(images) - 1 else 0)
                    has_caption = asset_index < len(block.asset_captions) and bool(block.asset_captions[asset_index].strip())
                    if source_page and has_caption:
                        tail_space += 800  # Three controlled 10pt lines plus paragraph spacing.
                    clone = page_plan.technical_picture(document, document_origin, data,
                                                        new_page=index > 0 or (source_page and bool(retained_images[block.key])),
                                                        caption=bool(block.asset_captions) and not source_page, tail_space=tail_space,
                                                        minimum_width=7200 if source_page and has_caption else 0)
                else:
                    clone = deepcopy(prototype)
                    for properties in list(clone.iter(qn("w:sectPr"))):
                        properties.getparent().remove(properties)
                    _set_text(clone, "")
                    _replace_image(document, clone, data, contain=block.key == "improvements")
                    if previous_document:
                        from docx.text.paragraph import Paragraph
                        first_paragraph = clone if clone.tag == qn("w:p") else next(clone.iter(qn("w:p")), None)
                        if first_paragraph is not None:
                            Paragraph(first_paragraph, document._body).paragraph_format.page_break_before = True
                if nested_document:
                    if previous_document and document_origin is last:
                        last.addnext(clone)
                    elif origin_root in emitted_photo_roots or anchor_index >= len(anchors):
                        origin_root.addnext(clone)
                    else:
                        origin_root.addprevious(clone)
                    page_plan.emitted(clone, origin_root, kind="picture")
                    if anchor_index < len(anchors):
                        unused_document_anchors.add(anchor)
                elif anchor_index < len(anchors):
                    anchor.addprevious(clone)
                    page_plan.emitted(clone, anchor, kind="picture")
                    # The old frame may own the section boundary. Keep that
                    # boundary until page accounting has finished its reflow.
                    if source_page:
                        unused_document_anchors.add(anchor)
                    elif anchor.find(".//" + qn("w:sectPr")) is None:
                        unused_document_anchors.discard(anchor)
                        anchor.getparent().remove(anchor)
                else:
                    last.addnext(clone)
                    page_plan.emitted(clone, last, kind="picture")
                last = clone
                if not source_page:
                    emitted_photo_roots.add(page_plan._root(clone))
                    unused_document_anchors.discard(anchor)
                    photo_index += 1
                previous_document = source_page
                caption_index = image_indexes[index] if image_indexes else index
                if caption_index < len(block.asset_captions) and block.asset_captions[caption_index].strip():
                    choices = section_text.get(sec) or global_text
                    caption = deepcopy(choices[0]) if choices and not source_page else OxmlElement("w:p")
                    for properties in list(caption.iter(qn("w:sectPr"))):
                        properties.getparent().remove(properties)
                    _clear_images(caption)
                    _set_text(caption, block.asset_captions[caption_index])
                    if source_page:
                        from docx.text.paragraph import Paragraph
                        style_document_caption(Paragraph(caption, document._body))
                    clone.addnext(caption)
                    page_plan.emitted(caption, clone, kind="caption")
                    last = caption
            # A document page did not consume a photo frame. Remove only its
            # now-unused, boundary-free cleared anchor after all photos had a
            # chance to reuse it, including in protected source regions.
            for old_anchor in unused_document_anchors:
                parent = old_anchor.getparent()
                if parent is not None and _discardable_picture_anchor(old_anchor):
                    parent.remove(old_anchor)
                    if parent.tag == qn("w:tc") and not parent.findall(qn("w:p")):
                        parent.append(OxmlElement("w:p"))
    if section_key:
        selected_positions = [position for position, sec in section_by_pos.items()
                              if sec == section_key and original[position - 1].tag != qn("w:sectPr")
                              and (section_key != "cover" or position <= (first_word_end or len(original)))]
        closing = None
        if selected_positions:
            for element in original[max(selected_positions) - 1:]:
                properties = element if element.tag == qn("w:sectPr") else element.find(".//" + qn("w:sectPr"))
                if properties is not None:
                    closing = deepcopy(properties)
                    break
        if closing is not None:
            # The selected range's final paragraph break becomes the document's
            # terminal section properties. Keeping both opens an empty next page.
            last_selected = original[max(selected_positions) - 1]
            last_boundary = last_selected.find(".//" + qn("w:sectPr"))
            if last_boundary is not None:
                last_boundary.getparent().remove(last_boundary)
            for existing in list(body.findall(qn("w:sectPr"))):
                body.remove(existing)
            body.append(closing)
            page_plan.owner[closing] = page_plan.owner[last_selected]

    contact = next((block for block in draft.blocks if block.key == "contact_matrix"), None)
    if contact is not None and not contact.rows and not any(table.rows for table in contact.extra_tables):
        # An explicitly empty current contact table must not leave an orphan
        # heading below an independently retained chart. Other source tables and
        # any boundary-bearing or emitted table keep their existing page proof.
        for anchor, _ in table_prototypes.get("contact_matrix", ()):
            if (anchor.getparent() is not None and anchor not in page_plan.live
                    and anchor.find(".//" + qn("w:sectPr")) is None):
                anchor.getparent().remove(anchor)

    # Source notes and comments are source facts, not current report content.
    explicit_notes = set(supplied_notes)
    for node in list(body.iter()):
        if node.tag in (qn("w:footnoteReference"), qn("w:endnoteReference")):
            kind = "footnote" if node.tag == qn("w:footnoteReference") else "endnote"
            if kind + ":" + node.get(qn("w:id"), "") not in explicit_notes:
                node.getparent().remove(node)
    for part in list(document.part.package.parts):
        name = str(part.partname)
        if not name.startswith(("/word/header", "/word/footer")) or not hasattr(part, "element"):
            continue
        header_occurrences = _header_image_occurrences(part, inspection)
        approved_header_nodes = set()
        header_logo_replacements = []
        _clear_diagrams(part.element)
        _clear_hyperlinks(part.element)
        for paragraph in part.element.iter(qn("w:p")):
            visible = deepcopy(paragraph)
            for fallback in list(visible.iter("{http://schemas.openxmlformats.org/markup-compatibility/2006}Fallback")):
                fallback.getparent().remove(fallback)
            from app.monthly_report_import import _remove_page_field_results
            _remove_page_field_results(visible, [])
            text = _text(visible)
            if "enfrasolutions.com" in text.casefold():
                # The presence of ENFRA's domain does not prove the rest of
                # an old header/footer belongs to the current client. A fresh
                # report without a reviewed corporate footer must not retain
                # a prior client's appended name/address in that same run.
                # Prefer the current report's explicit address; otherwise
                # retain only positively identified company/domain wording.
                safe_footer = draft.address_line.strip() or "ENFRA | enfrasolutions.com"
                if text != safe_footer:
                    _set_text(paragraph, safe_footer)
                continue
            if (text and "enfrasolutions.com" not in text.casefold()
                    and not re.fullmatch(r"(?:enfra|author:|client utility rate analysis)(?:\s+(?:author:|client utility rate analysis))*", text.strip(), re.IGNORECASE)
                    and not text.strip().isdigit()):
                _set_text(paragraph, "")
        header_logos = {key: blocks[key] for key in ("brand_logo", "client_logo")
                        if key in blocks and blocks[key].source != "Omit" and blocks[key].asset_hashes}
        if header_logos:
            if asset_loader is None:
                raise NativeLayoutError("Report branding requires the saved asset resolver.")
            header_bindings = defaultdict(list)
            for node, attr, item in header_occurrences:
                header_bindings[item.id].append((node, attr))
            for source_item in inspection.items:
                if source_item.kind != "image" or source_item.part != name.lstrip("/"):
                    continue
                reference = source_ref(source_item)
                explicit_roles = {block.key for block in draft.blocks
                                  if reference in block.references}
                role = (next(iter(explicit_roles)) if len(explicit_roles) == 1 else
                        "" if explicit_roles else source_item.image_role)
                logo = header_logos.get(role)
                if logo is None:
                    continue
                if source_item.id not in header_bindings:
                    raise NativeLayoutError(
                        "The selected header logo occurrence could not be matched safely. "
                        "Review its source artwork before generating this report.")
                if role == "brand_logo" and unchanged_master_brand and not explicit_roles:
                    # An unchanged, verified cover brand may have a distinct
                    # native header variant. Preserve that variant unless this
                    # exact header occurrence was deliberately mapped. Approval
                    # belongs to these nodes, not independent uses of equal bytes.
                    approved_header_nodes.update(node for node, _ in header_bindings[source_item.id])
                    continue
                if image_matches(source_item, role, logo.asset_hashes):
                    if explicit_roles == {role}:
                        for node, _ in header_bindings[source_item.id]:
                            header_logo_replacements.append(_contain_unchanged_cover_logo(node, source_path, source_item, part, defer=True))
                    approved_header_nodes.update(node for node, _ in header_bindings[source_item.id])
                    continue
                rid, image = part.get_or_add_image(BytesIO(asset_loader(logo.asset_hashes[0])))
                for node, attr in header_bindings[source_item.id]:
                    original_bytes = part.related_parts[node.get(attr)].blob
                    node.set(attr, rid)
                    # Replacement authority is already established here. Fit
                    # changed catalog bytes too, but retain an unmapped native
                    # crop when raw PNG bytes already match the chosen asset.
                    if explicit_roles == {role} or original_bytes != image.blob:
                        header_logo_replacements.append((node, image.px_width, image.px_height))
                    # Current client artwork is authorized at this occurrence,
                    # not at every source location containing equal bytes.
                    approved_header_nodes.add(node)
        plans = [(_contain_cover_logo(*entry, part, defer=True) if isinstance(entry, tuple) else entry)
                 for entry in header_logo_replacements]
        for plan in plans:
            if plan is not None:
                plan()
        if master:
            approved = set()
            if brand and brand.source != "Omit" and brand.asset_hashes and asset_loader:
                approved.add(hashlib.sha256(asset_loader(brand.asset_hashes[0])).hexdigest())
            _remove_unapproved_header_images(part, approved, approved_header_nodes)
    for panel in empty_panels:
        panel.remove_if_empty(body, page_plan, cleared_activity_paragraphs)
    page_plan.finish(document=document)
    from app.monthly_report_divider_numerals import fit_proposals_numeral
    fit_proposals_numeral(document, section_key=section_key)
    from app.monthly_report_native_design import name_report_in_contents_notice
    name_report_in_contents_notice(body)
    from app.monthly_report_semantic_headings import mark_native_section_headings
    # Navigation metadata belongs to the derivative only; source geometry,
    # visible style, header/artwork and report contents remain unchanged.
    navigable = {section.key for section in draft.sections if section.included}
    if section_key is not None:
        navigable.intersection_update((section_key,))
    mark_native_section_headings(document, original, headings, section_by_pos, navigable)
    if source_region_observer is not None:
        source_region_observer(body, page_plan)
    result = BytesIO()
    document.save(result)
    # Prune image relationships which became unused after old payload removal.
    with tempfile.TemporaryDirectory(prefix="native-report-") as folder:
        output = Path(folder) / "report.docx"
        output.write_bytes(result.getvalue())
        return passive_docx(output)
