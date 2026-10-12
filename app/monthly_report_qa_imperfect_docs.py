"""Deliberately imperfect invented monthly vendor and chemistry reports for isolated QA.

No authentic files, accounts or artwork. Generated on explicit QA request only,
using the libraries already installed in the Render image. Never uses a model,
network request, real contract namespace or live report data.
"""
from __future__ import annotations

from datetime import datetime
from hashlib import sha256
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from zipfile import ZipFile, ZIP_DEFLATED
import json
import subprocess

import fitz
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw

from app.monthly_report_testing import scope

RECORDS = (
    ("QA_Vendor_MixedPages", (
        ("technical", "VENDOR SERVICE REPORT | CHILLER INSPECTION", (
            "Facility: Demo North", "Vendor: SampleWorks Mechanical",
            "Service Date: 2026-09-16", "Work Order: QA-V-0916-041",
            "Equipment Tag: CH-QA-02",
            "Inspected Chiller CH-QA-02 while operating normally.",
            "Recorded oil pressure differential 25.1 PSID.",
            "Collected an oil sample; laboratory results remain pending.",
            "No repair or replacement was completed.")),
        ("blank", "", ()),
        ("legal", "GENERAL TERMS AND CONDITIONS", (
            "Governing law: fictional demonstration jurisdiction.",
            "Entire agreement: prior discussions are superseded for this invented exercise.",
            "Limitation of liability: sample disclaimer of consequential damages.",
            "Indemnification: illustrative third-party claim clause.",
            "Arbitration: disputes resolved by a fictional arbitrator.",
            "Severability: unenforceable fictional terms may be severed.")),
        ("signature", "SIGNATURE ONLY", (
            "SampleWorks Mechanical | Synthetic Technician | 2026-09-16",)),
        ("pricing", "UNAPPROVED REPAIR QUOTATION", (
            "Facility: Demo North", "Quote Reference: QA-Q-0916-09",
            "Proposed task: Replace auxiliary chiller sensor.",
            "Quote Amount: $4,875.00", "Labor Estimate: $1,350.00",
            "Approval: Pending; work NOT completed.")),
    )),
    ("QA_WaterTreatment_MixedPages", (
        ("technical", "WATER TREATMENT | COOLING TOWER ANALYSIS", (
            "Facility: Demo North", "Vendor: Example Aquatic Analytics",
            "Service Date: 2026-09-14", "Sample ID: QA-CT-0914-211",
            "Equipment: CT-QA-01 and CT-QA-02",
            "CT-QA-01 chlorides: 295 ppm (control range 100-300 ppm).",
            "CT-QA-02 chlorides: 340 ppm (control range 100-300 ppm).",
            "CT-QA-02 conductivity: 1,980 uS/cm (control range 1,300-2,000).",
            "Legionella culture: PENDING. No negative result has been established.",
            "Inspected feed pump and collected follow-up samples.")),
        ("chart", "TECHNICAL TREND | VALID FIGURE", (
            "CT-QA-02: September sample trend; the 9/14 reading exceeds the 300 ppm upper limit.",
            "Follow-up recommended: check bleed control; NOT a completed repair.")),
        ("legal", "STANDARD CONTRACT TERMS", (
            "Governing law: fictional sample territory.",
            "Arbitration: hypothetical dispute forum.",
            "Indemnification: allocation of nonexistent third-party liability.",
            "Limitation of liability: sample consequential-damage disclaimer.",
            "Entire agreement: invented agreement language only.")),
        ("signature", "RECEIPT AND SIGNATURE ONLY", (
            "Example Aquatic Analytics | Fictional Inspector | 2026-09-14",)),
        ("blank", "", ()),
    )),
    ("QA_Vendor_InvalidCompletionDate", (
        ("technical", "VENDOR INSPECTION | PLANNED WORK", (
            "Facility: Demo North", "Vendor: Invented Drive Maintenance LLC",
            "Service Date: 2026-09-18", "Completion Date: Invalid date",
            "Work Order: QA-V-0918-088", "Equipment Tag: VFD-QA-17",
            "Job Status: Planned - not completed.",
            "Observed an intermittent overtemperature indication.",
            "Diagnostic trend review recommended. No repair was performed.")),
        ("technical", "FOLLOW-UP | NOT YET COMPLETED", (
            "The next diagnostic visit is scheduled for 2026-10-03.",
            "Replacement fan request remains unapproved.",
            "Do not infer completion from the valid service date.")),
    )),
    ("QA_Chemical_DifferentSite", (
        ("technical", "BOILER CHEMISTRY | FOREIGN SITE", (
            "Facility: Demo South", "Vendor: Fictional Boiler Chemistry Co.",
            "Service Date: 2026-09-22", "Sample ID: QA-SOUTH-BOIL-322",
            "Equipment Tag: B-QA-03", "Conductivity: 3,650 uS/cm.",
            "P alkalinity: 245 ppm (target 250-350); below range.",
            "Molybdate: 12 ppm.", "Boiler remains offline.",
            "Collected sample only. Do not transfer to Demo North.")),
        ("pricing", "CHEMICAL ORDER | PRICE PAGE", (
            "Facility: Demo South", "Quote Total: $1,260.00",
            "Sample reagent order; purchase not approved.")),
        ("blank", "", ()),
    )),
)
README = """SYNTHETIC ONLY - fictitious vendors, people, sites and test measurements.
For isolated Monthly Report QA only. Never use as real client evidence.

Vendor mixed: PDF page 1 technical; 2 blank; 3 legal; 4 signature; 5 price.
Water treatment: pages 1-2 technical, including meaningful chart;
  3 legal; 4 signature; 5 blank. The elevated result and pending
  laboratory culture must not be declared normal/negative.
Invalid completion: a service date exists, but status remains PLANNED.
Different site: Demo South material cannot migrate to Demo North.
Image-only scan: page 1 needs review/approved reading; page 2 blank.

Expected: vendor work belongs to Monthly Activity, not ENFRA onsite calls.
Water readings belong to Water Treatment. Suppress blank/legal/signature
and price pages, preserve technical figures, source lineage and exclusions.
DOCX 'items' are not necessarily physical PDF page numbers.
These files cannot prove authentic-client report parity.
"""

def _page(doc: Document, kind: str, heading: str, lines: tuple[str, ...],
          chart: Path, signature: Path) -> None:
    if kind == "blank":
        doc.add_paragraph("\N{NO-BREAK SPACE}")
        return
    p = doc.add_paragraph()
    run = p.add_run("FICTIONAL SOURCE  /  SYNTHETIC QA ONLY")
    run.bold = True
    run.font.color.rgb = RGBColor(155, 50, 41)
    run.font.size = Pt(9)
    p = doc.add_paragraph()
    r = p.add_run(heading)
    r.bold = True
    r.font.size = Pt(16)
    r.font.color.rgb = RGBColor(24, 74, 73)
    doc.add_paragraph()
    if kind == "signature":
        doc.add_picture(str(signature), width=Inches(3.2))
    if kind == "chart":
        doc.add_picture(str(chart), width=Inches(6.6))
    for line in lines:
        doc.add_paragraph(line)

def _picture_files(root: Path) -> tuple[Path, Path]:
    signature = root / "signature.png"
    image = Image.new("RGB", (650, 110), "white")
    draw = ImageDraw.Draw(image)
    draw.line([(10, 70), (45, 28), (80, 72), (120, 17), (155, 66),
               (204, 35), (262, 82), (315, 28), (380, 64), (450, 28),
               (510, 74), (600, 44)], fill=(35, 89, 132), width=5)
    image.save(signature)
    chart = root / "chemical-trend.png"
    image = Image.new("RGB", (900, 330), "white")
    draw = ImageDraw.Draw(image)
    draw.line((75, 270, 850, 270), fill=(100, 120, 130), width=3)
    draw.line((75, 50, 75, 270), fill=(100, 120, 130), width=3)
    draw.line((75, 145, 850, 145), fill=(181, 61, 48), width=3)
    draw.line([(170, 200), (365, 175), (565, 98), (770, 185)],
              fill=(40, 105, 169), width=6)
    for x, y in [(170, 200), (365, 175), (565, 98), (770, 185)]:
        draw.ellipse((x-7, y-7, x+7, y+7), fill=(40, 105, 169))
    image.save(chart)
    return chart, signature

def _make_docx(path: Path, pages, chart: Path, signature: Path) -> None:
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = sec.bottom_margin = Inches(.75)
    sec.left_margin = sec.right_margin = Inches(.8)
    style = doc.styles["Normal"]
    style.font.name = "Aptos"
    style.font.size = Pt(10)
    style.paragraph_format.space_after = Pt(7)
    doc.core_properties.author = "Synthetic QA Fixture Generator"
    doc.core_properties.created = datetime(2026, 9, 30)
    doc.core_properties.modified = datetime(2026, 9, 30)
    for i, (kind, heading, lines) in enumerate(pages):
        if i:
            doc.add_page_break()
        _page(doc, kind, heading, lines, chart, signature)
    doc.save(path)

def _scan_pdf(path: Path) -> None:
    from PIL import ImageFont
    image = Image.new("RGB", (1100, 1400), "white")
    draw = ImageDraw.Draw(image)
    font_path = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
    font = ImageFont.truetype(font_path, 29) if Path(font_path).exists() else ImageFont.load_default()
    lines = (
        "IMAGE-ONLY SYNTHETIC WATER REPORT",
        "Facility: Demo North  |  Sample: QA-IMG-SEP2026",
        "Sample Date: 2026-09-21",
        "Chlorides: 340 ppm  |  Reference: 100-300",
        "Conductivity: 1980 uS/cm",
        "Legionella: Pending (not negative)",
        "No embedded PDF text layer - deliberate QA challenge",
    )
    for i, line in enumerate(lines):
        draw.text((60, 95 + 85*i), line, fill=(30, 45, 50), font=font)
    scan = path.with_suffix(".png")
    image.save(scan)
    out = fitz.open()
    out.new_page(width=612, height=792).insert_image(fitz.Rect(15, 15, 597, 777), filename=str(scan))
    out.new_page(width=612, height=792)
    out.save(str(path), garbage=4, deflate=True)
    out.close()

def build_pack() -> bytes:
    """Build a downloadable zip; call only with an active isolated QA scope."""
    if scope() is None:
        raise ValueError("Synthetic document generation requires isolated QA.")
    with TemporaryDirectory(prefix="enfra-qa-messy-") as temporary:
        root = Path(temporary)
        chart, signature = _picture_files(root)
        generated = []
        entries = []
        for name, pages in RECORDS:
            word = root / (name + ".docx")
            _make_docx(word, pages, chart, signature)
            command = [
                "libreoffice",
                "-env:UserInstallation=" + (root / "lo-profile").as_uri(),
                "--headless", "--convert-to", "pdf", "--outdir", str(root), str(word),
            ]
            subprocess.run(command, check=True, stdout=subprocess.DEVNULL,
                           stderr=subprocess.PIPE, timeout=45)
            pdf = word.with_suffix(".pdf")
            if not pdf.is_file():
                raise OSError("Synthetic PDF creation did not finish.")
            with fitz.open(pdf) as rendered:
                if len(rendered) != len(pages):
                    raise OSError("Synthetic report changed its expected page count.")
                for i, (kind, _heading, _lines) in enumerate(pages):
                    if kind == "blank" and (rendered[i].get_text().strip()
                                          or rendered[i].get_image_info()
                                          or rendered[i].get_drawings()):
                        raise OSError("The intentionally blank page is not blank.")
            for file in (word, pdf):
                generated.append(file)
                entries.append({"filename": file.name,
                                "sha256": sha256(file.read_bytes()).hexdigest(),
                                "page_kinds_in_pdf": [p[0] for p in pages]})
        scan = root / "QA_WaterTreatment_ImageOnlyScan.pdf"
        _scan_pdf(scan)
        with fitz.open(scan) as rendered:
            if len(rendered) != 2 or any(page.get_text().strip() for page in rendered):
                raise OSError("Synthetic scanned PDF has an unexpected text layer.")
        generated.append(scan)
        entries.append({"filename": scan.name, "sha256": sha256(scan.read_bytes()).hexdigest(),
                        "page_kinds_in_pdf": ["image-needs-review", "blank"]})
        manifest = json.dumps({"synthetic": True, "contract": "QA Demonstration - Synthetic Only",
                               "reporting_period": "2026-09",
                               "files": entries}, indent=2).encode()
        output = BytesIO()
        with ZipFile(output, "w", ZIP_DEFLATED, compresslevel=7) as archive:
            for file in generated:
                archive.write(file, file.name)
            archive.writestr("README.txt", README)
            archive.writestr("MANIFEST.json", manifest)
        return output.getvalue()
