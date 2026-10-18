"""Local monthly evidence ingestion, cached by digest on the runtime disk.

No model calls run here. Decode, validate, prepare thumbnails and normalize
images on the caller thread; only prepared network requests enter receipt_jobs.
"""

from __future__ import annotations

from dataclasses import asdict, dataclass, replace
from datetime import date, datetime
from email import policy
from email.parser import BytesParser
from html.parser import HTMLParser
from io import BytesIO, StringIO
import csv
import hashlib
from pathlib import Path
import re
from zipfile import BadZipFile

import fitz
from openpyxl import load_workbook

from app import monthly_report_library as library, ocr
from app.monthly_report_docx import ReportImage, normalize_report_image
from app.monthly_report_import import inspect_docx, read_import_image
from app.monthly_report_model import ReportProfile, ReportSource, ResolvedBlock


CLASSIFICATIONS = (
    "Vendor service", "Water treatment", "Improvement / training photo", "CMMS export",
    "MBCx", "Utility results", "Proposal / quote", "Capital renewal", "Training", "Reference only",
)
MAX_FILE_BYTES = 30 * 1024 * 1024
MAX_REPORT_BYTES = 180 * 1024 * 1024
MAX_SOURCES = 50
MAX_REPORT_TEXT = 800_000
MAX_PAGES = 500
MAX_EMBED_PAGES = 150
MAX_VISION_PAGES = ocr._MAX_PDF_PAGES
MAX_TABLE_CELLS = 200_000
MAX_TABLE_ROWS = 20_000
SUPPORTED = ocr.SUPPORTED_IMAGE_SUFFIXES | {".pdf", ".docx", ".xlsx", ".csv", ".eml", ".msg", ".txt"}
PAGE_DESTINATIONS = ("vendor_reports", "water_reports", "mbcx_report", "utility_analysis", "improvements")


@dataclass(frozen=True)
class SourceTable:
    name: str
    columns: tuple[str, ...]
    rows: tuple[tuple[str, ...], ...]


@dataclass(frozen=True)
class SourceContent:
    source: ReportSource
    tables: tuple[SourceTable, ...] = ()
    # DOCX raster images correspond to item indices, not invented page numbers.
    image_items: tuple[tuple[int, str], ...] = ()


def _path(profile: ReportProfile, digest: str, suffix: str) -> Path:
    if not re.fullmatch(r"[0-9a-f]{64}", digest) or suffix not in SUPPORTED | {".json"}:
        raise ValueError("Invalid monthly source reference.")
    return library._profile_path(profile.contract, profile.key) / "sources" / (digest + suffix)


def source_bytes(profile: ReportProfile, source: ReportSource) -> bytes:
    path = _path(profile, source.sha256, source.suffix)
    limit = 128 * 1024 * 1024 if source.suffix == ".docx" else MAX_FILE_BYTES
    if path.stat().st_size > limit:
        raise ValueError("Saved source exceeds the file limit.")
    raw = path.read_bytes()
    if hashlib.sha256(raw).hexdigest() != source.sha256:
        raise ValueError("A source file failed its integrity check.")
    return raw


def _date(value: str) -> str:
    for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%m/%d/%y"):
        try:
            return datetime.strptime(value.strip(), fmt).date().isoformat()
        except ValueError:
            pass
    return ""


def native_action_lines(source: ReportSource):
    """Price-free native work lines with their actual selected page origins."""
    from app.monthly_report_content_policy import contains_price, page_status
    found = {}
    for number in source.selected_pages:
        if not 1 <= number <= len(source.page_texts):
            continue
        text = source.page_texts[number - 1]
        if page_status(text, unreadable=number in source.needs_vision)[0] != "technical":
            continue
        for line in text.splitlines():
            line = line.strip()
            uncertain_work = re.search(
                r"\b(?:not|never|cannot|unable|recommend\w*|propos\w*|planned?|planning|"
                r"scheduled?|pending|awaiting|will|would|should|could|must|needs?|required|"
                r"if|unless|to be)\b|\b\w+n['’]t\b", line, re.I)
            if (re.search(r"\b(completed|repaired|replaced|inspected)\b", line, re.I)
                    and not uncertain_work and not contains_price(line)):
                found.setdefault(line, []).append(number)
    return tuple((line, tuple(dict.fromkeys(pages))) for line, pages in found.items())


def action_evidence(source: ReportSource):
    """Separate supported native action wording from edits needing source review.

    Never rewrite human fields. Unmatched wording may be a valid human summary,
    but must not acquire invented page citations merely because pages are selected.
    """
    supported = dict(native_action_lines(source))
    accepted, needs_review = [], []
    for line in dict.fromkeys(line.strip() for line in source.actions.splitlines() if line.strip()):
        if line in supported:
            accepted.append((line, supported[line]))
        else:
            needs_review.append(line)
    return tuple(accepted), tuple(needs_review)


def suggest_facts(source: ReportSource, profile: ReportProfile) -> ReportSource:
    text = "\n".join(source.page_texts)
    sample = (source.filename + "\n" + text[:30_000]).casefold()
    patterns = (("Water treatment", ("water treatment", "conductivity", "water analysis")),
                ("CMMS export", ("work order", "cmms", "task code")),
                ("MBCx", ("mbcx", "commissioning")),
                ("Utility results", ("utility results", "utility analysis", "utility consumption", "utility scorecard", "energy consumption")),
                ("Proposal / quote", ("proposal", "quotation", "quote")),
                ("Capital renewal", ("capital renewal", "end of useful life")),
                ("Training", ("training", "safety briefing")),
                ("Vendor service", ("service report", "service visit", "maintenance report")))
    hits = [kind for kind, terms in patterns if any(term in sample for term in terms)]
    classification = hits[0] if len(hits) == 1 else "Reference only"
    if not hits and source.suffix in ocr.SUPPORTED_IMAGE_SUFFIXES:
        classification = "Improvement / training photo"
    def labeled(label):
        match = re.search(r"(?im)^[ \t]*(?:" + label + r")[ \t]*[:=][ \t]*([^\r\n]{1,160})", text)
        return match[1].strip() if match else ""
    raw_date = labeled("service date|report date|completed date|date")
    date_match = re.search(r"\b(?:\d{4}-\d{1,2}-\d{1,2}|\d{1,2}/\d{1,2}/\d{2,4})\b", raw_date)
    facility = labeled("facility|site|building")
    if not facility:
        # Filenames can suggest a document type, but cannot establish which
        # site's facts its contents describe. Scope needs native source text
        # or the user's explicit item-specific correction.
        scope_text = text[:30_000].casefold()
        matches = [f.title for f in profile.facilities if any(re.search(r"(?<!\w)" + re.escape(n.casefold()) + r"(?!\w)", scope_text) for n in (f.title, *f.aliases))]
        facility = matches[0] if len(matches) == 1 else ""
    def lines(term):
        return "\n".join(line.strip() for line in text.splitlines() if re.search(term, line, re.I))[:6000]
    return replace(source, classification=classification, confidence="medium" if len(hits) == 1 else "low",
                   vendor=labeled("vendor|contractor|service company"), facility=facility,
                   service_date=_date(date_match[0]) if date_match else "", work_order=labeled("job(?: number)?|wo(?: number)?|work order"),
                   tags=tuple(v.strip() for v in re.split(r"[,;]", labeled("equipment tags?|asset tags?")) if v.strip()),
                   actions="\n".join(line for line, _ in native_action_lines(source))[:6000], findings=lines(r"\b(findings?|observed)\b"),
                   recommendations=lines(r"\brecommend"), follow_ups=lines(r"\b(follow.up|pending|scheduled)\b"),
                   quotes=lines(r"\b(quoted|quotation|quote amount)\b"), out_of_limit=lines(r"\b(out.of.(?:limit|range)|above limit|below limit)\b"))


class _PlainHTML(HTMLParser):
    def __init__(self):
        super().__init__()
        self.parts, self.hidden = [], 0
    def handle_starttag(self, tag, attrs):
        if tag in ("script", "style"):
            self.hidden += 1
        if tag in ("p", "br", "div", "tr"):
            self.parts.append("\n")
    def handle_endtag(self, tag):
        if tag in ("script", "style"):
            self.hidden = max(0, self.hidden - 1)
    def handle_data(self, data):
        if not self.hidden:
            self.parts.append(data)


def _plain_html(value: str) -> str:
    reader = _PlainHTML()
    reader.feed(value)
    return "".join(reader.parts)


def _email(raw: bytes) -> tuple[str, tuple[tuple[str, bytes], ...], tuple[str, ...]]:
    message = BytesParser(policy=policy.default).parsebytes(raw)
    body = message.get_body(preferencelist=("plain", "html"))
    text = body.get_content() if body else ""
    if body and body.get_content_type() == "text/html":
        text = _plain_html(text)
    headers = "\n".join(f"{'Message date' if key == 'Date' else key}: {message.get(key, '')}" for key in ("Subject", "Date", "From"))
    attachments, notices = [], []
    for item in message.walk():
        if item.is_multipart():
            if item.get_content_type() == "message/rfc822":
                notices.append("Nested email retained in the original; save its attachments separately.")
            continue
        filename = item.get_filename()
        if filename:
            if len(attachments) >= MAX_SOURCES:
                raise ValueError("Too many email attachments; split this message before processing.")
            payload = item.get_payload(decode=True) or b""
            attachments.append((Path(filename.replace("\\", "/")).name, payload))
    return headers + "\n\n" + str(text), tuple(attachments), tuple(notices)


def _msg(raw: bytes) -> tuple[str, tuple[tuple[str, bytes], ...], tuple[str, ...]]:
    import olefile
    with olefile.OleFileIO(BytesIO(raw), raise_defects=olefile.DEFECT_INCORRECT) as msg:
        paths = msg.listdir()
        if len(paths) > 3000:
            raise ValueError("MSG contains too many property streams.")
        def read(path):
            if not msg.exists(path):
                return b""
            if msg.get_size(path) > MAX_FILE_BYTES:
                raise ValueError("MSG attachment/property exceeds 30 MB.")
            return msg.openstream(path).read()
        def string(prefix, tag):
            unicode = read([*prefix, "__substg1.0_" + tag + "001F"])
            if unicode:
                return unicode.decode("utf-16-le", errors="replace").rstrip("\x00")
            return read([*prefix, "__substg1.0_" + tag + "001E"]).decode("cp1252", errors="replace").rstrip("\x00")
        text = string([], "1000")
        notices = []
        if not text:
            html = read(["__substg1.0_10130102"])
            if html:
                text = _plain_html(html.decode("utf-8", errors="replace"))
            else:
                notices.append("MSG has no readable plain/HTML body. Review the original or paste its text; compressed RTF is not executed.")
        text = "Subject: " + string([], "0037") + "\nSender: " + string([], "0C1A") + "\n\n" + text
        attachments = []
        folders = sorted({p[0] for p in paths if p[0].startswith("__attach_version1.0_")})
        for folder in folders:
            if len(attachments) >= MAX_SOURCES:
                raise ValueError("Too many MSG attachments; split this message before processing.")
            payload = read([folder, "__substg1.0_37010102"])
            filename = string([folder], "3707") or string([folder], "3704")
            if payload and filename:
                attachments.append((Path(filename.replace("\\", "/")).name, payload))
            else:
                notices.append("An embedded MSG object could not be read as a file attachment; review it in the original.")
        return text, tuple(attachments), tuple(notices)


def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, bool):
        return "Yes" if value else "No"
    return str(value)


def _table(name, records, *, strict_columns=False) -> SourceTable:
    rows, cells = [], 0
    for record in records:
        row = tuple(_cell(v) for v in record)
        cells += len(row)
        if len(rows) >= MAX_TABLE_ROWS or cells > MAX_TABLE_CELLS or len(row) > 200:
            raise ValueError("Spreadsheet exceeds the row/column review budget; export a smaller date range.")
        if any(c != "" for c in row):
            if strict_columns and rows and len(row) != len(rows[0]):
                # A bare '12,345.67' in a comma-separated file becomes two
                # cells and could move a price into a nonprice client column.
                # Do not guess how to rejoin columns or silently trim cells.
                raise ValueError(
                    "CSV rows have different column counts. Quote any values "
                    "containing commas (for example, a monetary amount), then upload again."
                )
            rows.append(row)
    if not rows:
        raise ValueError("Spreadsheet has no nonempty rows.")
    width = max(map(len, rows))
    columns = []
    for i in range(width):
        value = rows[0][i] if i < len(rows[0]) else ""
        title = value.strip() or f"Column {i+1}"
        base, count = title, 2
        while title in columns:
            title = f"{base} ({count})"
            count += 1
        columns.append(title)
    return SourceTable(name, tuple(columns), tuple(r + ("",) * (width - len(r)) for r in rows[1:]))


def _spreadsheet(raw: bytes, suffix: str) -> tuple[SourceTable, ...]:
    if suffix == ".csv":
        try:
            text = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = raw.decode("cp1252")
        sample = text[:8192]
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        return (_table("CSV", csv.reader(StringIO(text), dialect, strict=True),
                       strict_columns=True),)
    # Bound ZIP/XML before openpyxl reads shared strings or worksheet dimensions.
    from zipfile import ZipFile
    with ZipFile(BytesIO(raw)) as z:
        infos = z.infolist()
        if (len(infos) > 2000 or sum(i.file_size for i in infos) > 80 * 1024 * 1024
                or any(i.file_size > 32 * 1024 * 1024 for i in infos)
                or len({i.filename for i in infos}) != len(infos)):
            raise ValueError("Expanded workbook exceeds the safe parsing budget.")
    workbook = load_workbook(BytesIO(raw), read_only=True, data_only=True, keep_links=False)
    try:
        if len(workbook.worksheets) > 30:
            raise ValueError("Workbook exceeds 30 worksheets; upload the required sheets separately.")
        tables, total = [], 0
        for sheet in workbook.worksheets:
            if sheet.max_row and sheet.max_row > MAX_TABLE_ROWS or sheet.max_column and sheet.max_column > 200:
                raise ValueError("Worksheet dimensions exceed the review budget.")
            try:
                table = _table(sheet.title, sheet.iter_rows(values_only=True))
            except ValueError as exc:
                if str(exc) == "Spreadsheet has no nonempty rows.":
                    continue
                raise
            total += len(table.rows) * len(table.columns)
            if total > MAX_TABLE_CELLS:
                raise ValueError("Workbook cells exceed the total review budget.")
            tables.append(table)
        if not tables:
            raise ValueError("Workbook has no nonempty worksheets.")
        return tuple(tables)
    finally:
        workbook.close()


def ingest(profile: ReportProfile, filename: str, raw: bytes) -> tuple[SourceContent, tuple[tuple[str, bytes], ...]]:
    """Cache local extraction. Caller enforces aggregate count/byte/text budgets."""
    from app.monthly_report_content_policy import suggested_pages
    filename = Path(filename.replace("\\", "/")).name
    suffix = Path(filename).suffix.casefold()
    limit = 128 * 1024 * 1024 if suffix == ".docx" else MAX_FILE_BYTES
    if suffix not in SUPPORTED or not raw or len(raw) > limit:
        raise ValueError("Choose a supported nonempty file within its upload limit (30 MB; DOCX 128 MB).")
    digest = hashlib.sha256(raw).hexdigest()
    path, cache = _path(profile, digest, suffix), _path(profile, digest, ".json")
    attachments = ()
    if cache.exists() and path.exists() and suffix not in (".eml", ".msg"):
        data = library._read(cache)
        source = library.source_from_dict(data["source"])
        tables = tuple(SourceTable(t["name"], tuple(t["columns"]), tuple(tuple(r) for r in t["rows"])) for t in data.get("tables", ()))
        if source.suffix == suffix:
            if data.get("client_page_policy", 0) < 1 and suffix not in (".csv", ".xlsx"):
                source = replace(source, selected_pages=suggested_pages(source))
            if data.get("native_action_policy", 0) < 1:
                # This is the original extraction cache, never a reviewed draft.
                source = replace(source, actions="\n".join(line for line, _ in native_action_lines(source))[:6000])
            return SourceContent(replace(source, filename=filename), tables, tuple(tuple(p) for p in data.get("image_items", ()))), ()
    library._atomic_write(path, raw)
    page_texts, needs_vision, tables, image_items, notices = [], [], (), [], []
    if suffix == ".pdf":
        with fitz.open(stream=raw, filetype="pdf") as pdf:
            if pdf.is_encrypted and not pdf.authenticate(""):
                raise ValueError("PDF needs a password. Upload an unlocked copy.")
            if len(pdf) > MAX_PAGES:
                raise ValueError("PDF exceeds 500 pages; select a smaller document range.")
            for number, page in enumerate(pdf, 1):
                text = page.get_text().strip()
                if ocr._page_needs_ocr(page, native_text=text):
                    needs_vision.append(number)
                page_texts.append(text)
                if sum(map(len, page_texts)) > ocr._MAX_EXTRACTED_CHARS:
                    raise ValueError("PDF text exceeds 400,000 characters; select a smaller document range.")
    elif suffix in ocr.SUPPORTED_IMAGE_SUFFIXES:
        normalize_report_image(raw, suffix)  # Validate before caching or paid work.
        page_texts, needs_vision = [""], [1]
    elif suffix == ".txt":
        page_texts = [raw.decode("utf-8-sig")]
    elif suffix in (".csv", ".xlsx"):
        tables = _spreadsheet(raw, suffix)
        page_texts = ["\n".join(["\t".join(t.columns), *("\t".join(r) for r in t.rows)]) for t in tables]
        if suffix == ".xlsx":
            notices.append("Workbook values use stored formula results; formulas are never executed. Missing cached results remain blank.")
    elif suffix == ".docx":
        inspected = inspect_docx(path)
        for number, item in enumerate(inspected.items, 1):
            page_texts.append(item.text or "\n".join("\t".join(r) for r in item.rows))
            if item.kind == "unsupported":
                notices.append(f"Item {number}: {item.note or 'This original item could not be read.'}")
            elif item.note == "Unplaced package asset; confirm whether it belongs in this report.":
                notices.append(f"Item {number}: {item.note}")
            if item.kind == "image":
                image_items.append((number, item.id))
                needs_vision.append(number)
        notices.extend(inspected.notices)
        # Only this new inspection route has retained every unsupported-item
        # notice. Old cached records keep their broader artwork-review notice.
        notices.append("DOCX selections are extracted items, not Word page numbers.")
    else:
        text, attachments, extra = _email(raw) if suffix == ".eml" else _msg(raw)
        page_texts = [text]
        notices.extend(extra)
    if sum(map(len, page_texts)) > ocr._MAX_EXTRACTED_CHARS:
        raise ValueError("Source text exceeds 400,000 characters; split the input before reading.")
    source = ReportSource(digest, filename, digest, suffix, page_texts=tuple(page_texts),
                          selected_pages=tuple(range(1, len(page_texts)+1)), needs_vision=tuple(needs_vision), notices=tuple(notices))
    if suffix not in (".csv", ".xlsx"):
        source = replace(source, selected_pages=suggested_pages(source))
    source = suggest_facts(source, profile)
    content = SourceContent(source, tables, tuple(image_items))
    library._atomic_write(cache, library._json({"schema": 1, "client_page_policy": 1, "native_action_policy": 1, **asdict(content)}))
    return content, attachments


def ingest_batch(profile: ReportProfile, files: tuple[tuple[str, bytes], ...], existing: tuple[SourceContent, ...] = ()) -> tuple[tuple[SourceContent, ...], tuple[str, ...]]:
    sources = {c.source.sha256: c for c in existing}
    sizes, messages = 0, []
    for item in existing:
        # Missing originals remain in the draft inventory and its output
        # dependency checks. They cannot prevent other files from being read.
        try:
            sizes += _path(profile, item.source.sha256, item.source.suffix).stat().st_size
        except (ValueError, OSError):
            label = Path(item.source.filename.replace("\\", "/")).name
            messages.append(f"Saved source unavailable: {label}. Replace it before generating dependent content.")
    text_size = sum(len(t) for c in existing for t in c.source.page_texts)
    queue = [(name, raw, 0) for name, raw in files]
    while queue:
        filename, raw, depth = queue.pop(0)
        filename = Path(filename.replace("\\", "/")).name
        digest = hashlib.sha256(raw).hexdigest()
        if digest in sources:
            # Re-uploading the identical accepted original can repair a missing
            # or corrupted stored copy without changing the report's reviewed
            # source identity, citations, or completed historical snapshots.
            existing_source = sources[digest].source
            try:
                source_bytes(profile, existing_source)
            except (ValueError, OSError):
                try:
                    library._atomic_write(_path(profile, digest, existing_source.suffix), raw)
                    source_bytes(profile, existing_source)
                except (ValueError, OSError):
                    messages.append(f"Could not restore saved source: {filename}. The report still needs a valid copy.")
                    continue
                messages.append(f"Saved source restored: {filename}. Prior reviews and report history are unchanged.")
                sizes += len(raw)
            else:
                messages.append(f"Duplicate skipped: {filename}")
            continue
        if len(sources) >= MAX_SOURCES or sizes + len(raw) > MAX_REPORT_BYTES:
            messages.append(f"Not read: {filename}. Report limit is 50 source files / 180 MB including attachments.")
            continue
        suffix = Path(filename).suffix.casefold()
        if suffix not in SUPPORTED or depth > 1:
            messages.append(f"Not read: {filename}. Unsupported or nested attachment; retained in the original email.")
            continue
        path = _path(profile, digest, suffix)
        already_present = path.exists()
        cache = _path(profile, digest, ".json")
        existing_cache = cache.exists()
        try:
            content, attachments = ingest(profile, filename, raw)
            added = sum(map(len, content.source.page_texts))
            if text_size + added > MAX_REPORT_TEXT:
                raise ValueError("Total extracted text exceeds 800,000 characters. Remove unused sources or split the report.")
        except (ValueError, OSError, LookupError, csv.Error, RuntimeError, BadZipFile, RecursionError) as exc:
            # Only newly rejected and content-addressed bytes may be cleaned.
            # Never remove accepted originals or any pre-existing library file.
            if not already_present:
                try:
                    if path.is_file() and hashlib.sha256(path.read_bytes()).hexdigest() == digest:
                        path.unlink()
                except OSError:
                    pass
            if not existing_cache:
                try:
                    cache.unlink(missing_ok=True)
                except OSError:
                    pass
            if isinstance(exc, OSError):
                reason = "Storage or file reader was unavailable; retry this file."
            elif isinstance(exc, LookupError):
                reason = "The file uses an unknown text encoding; provide a UTF-8 copy."
            else:
                reason = str(exc)[:300]
            messages.append(f"Not read: {filename}. {reason}")
            continue
        sources[digest] = content
        sizes += len(raw)
        text_size += added
        queue.extend((name, data, depth+1) for name, data in attachments)
    return tuple(sources.values()), tuple(messages)


def page_image(profile: ReportProfile, content: SourceContent, number: int, *, preview: bool = False):
    source = content.source
    if number < 1 or number > len(source.page_texts):
        raise ValueError("Source page does not exist.")
    directory = _path(profile, source.sha256, source.suffix).parent / "prepared" / source.sha256
    # Older DOCX previews were incorrectly prepared at print resolution.
    # Keep those assets intact and give the corrected derivative its own key.
    purpose = "preview-docx-v2" if preview and source.suffix == ".docx" else "preview" if preview else "print"
    cache = directory / f"{number}-{purpose}.json"
    if cache.exists():
        saved = library._read(cache)
        if not re.fullmatch(r"[0-9a-f]{64}\.(?:jpg|png)", saved.get("asset", "")):
            raise ValueError("Invalid prepared image reference.")
        data = (directory / saved["asset"]).read_bytes()
        if library.asset_reference(data, saved["extension"]) != saved["asset"]:
            raise ValueError("A prepared source image failed its integrity check.")
        return ReportImage(data, saved["extension"], saved["width"], saved["height"])
    raw = source_bytes(profile, source) if source.suffix != ".docx" else b""
    if source.suffix == ".pdf":
        with fitz.open(stream=raw, filetype="pdf") as pdf:
            page = pdf[number-1]
            dpi = 96 if preview else 150
            width, height = page.rect.width * dpi / 72, page.rect.height * dpi / 72
            if width <= 0 or height <= 0 or width * height > ocr._MAX_PIXELS_PER_FRAME:
                raise ValueError("This PDF page is too large to rasterize safely.")
            png = page.get_pixmap(dpi=dpi, alpha=False).tobytes("png")
        result = normalize_report_image(png, ".png", line_art=True, dpi=dpi)
    elif source.suffix in ocr.SUPPORTED_IMAGE_SUFFIXES:
        result = normalize_report_image(raw, source.suffix, dpi=96 if preview else 200)
    elif source.suffix == ".docx":
        item_id = dict(content.image_items).get(number)
        if not item_id:
            raise ValueError("This DOCX item has no raster image. Use its extracted text/table instead.")
        path = _path(profile, source.sha256, source.suffix)
        inspected = inspect_docx(path)
        if inspected.sha256 != source.sha256:
            raise ValueError("A source file failed its integrity check.")
        item = next(i for i in inspected.items if i.id == item_id)
        result = read_import_image(path, item, preview=preview)
    else:
        raise ValueError("This source contains text/tables rather than image pages.")
    ref = library.asset_reference(result.data, result.extension)
    library._atomic_write(directory / ref, result.data)
    library._atomic_write(cache, library._json({"schema": 1, "asset": ref, "extension": result.extension,
                                              "width": result.width, "height": result.height}))
    return result


def source_reference(source: ReportSource, number: int) -> str:
    return f"{source.id}:{number}:{source.fingerprint}"


def image_numbers(content: SourceContent) -> tuple[int, ...]:
    if content.source.suffix == ".docx":
        return tuple(n for n, _ in content.image_items)
    if content.source.suffix == ".pdf" or content.source.suffix in ocr.SUPPORTED_IMAGE_SUFFIXES:
        return tuple(range(1, len(content.source.page_texts) + 1))
    return ()


def prepare_pages(profile: ReportProfile, contents: tuple[SourceContent, ...],
                  destinations: tuple[tuple[str, str], ...], *,
                  only_pages: frozenset[tuple[str, int]] | None = None) -> tuple[ResolvedBlock, ...]:
    """Prepare serially into content-addressed storage, with an aggregate limit.

    An optional subset restricts which pages are actually inserted without
    changing the canonical source fingerprint used in provenance references.
    No shared profile defaults, manifest, or snapshot is changed here.
    """
    from app.monthly_report_content_policy import page_allowed
    sources = {c.source.id: c for c in contents}
    if len(dict(destinations)) != len(destinations):
        raise ValueError("Each source needs one page destination.")
    allowed_pairs = set(only_pages) if only_pages is not None else None
    matched_pairs = set()
    planned = []
    for identity, slot in destinations:
        if identity not in sources or slot not in PAGE_DESTINATIONS:
            raise ValueError("Unknown page source or destination.")
        content = sources[identity]
        allowed = set(image_numbers(content))
        selected = tuple(
            n for n in content.source.selected_pages
            if n in allowed and (allowed_pairs is None or (identity, n) in allowed_pairs))
        matched_pairs.update((identity, n) for n in selected)
        if any(n < 1 or n > len(content.source.page_texts) for n in content.source.selected_pages):
            raise ValueError("Selected source page does not exist.")
        planned.extend((content, slot, n) for n in selected)
    if allowed_pairs is not None and matched_pairs != allowed_pairs:
        raise ValueError("The requested source pages changed. Review them again before inclusion.")
    if len(planned) > MAX_EMBED_PAGES:
        raise ValueError("Embedding limit is 150 selected image/PDF pages per report. Reduce the selection before preparing.")
    for content, _, number in planned:
        if not page_allowed(content.source, number):
            raise ValueError(f"Check page {number} of {content.source.filename} before including it. Prices, legal-only and blank/signature-only pages cannot be included.")
    grouped, total = {}, 0
    for content, slot, number in planned:
        normalized = page_image(profile, content, number)
        total += len(normalized.data)
        if total > 60 * 1024 * 1024:
            raise ValueError("Prepared images exceed 60 MB. Reduce the selected pages or image dimensions.")
        ref = library.asset_reference(normalized.data, normalized.extension)
        library._atomic_write(library._profile_path(profile.contract, profile.key) / "assets" / ref, normalized.data)
        caption = dict(content.source.captions).get(number, f"{content.source.filename} · page/item {number}")
        grouped.setdefault(slot, []).append((ref, caption, source_reference(content.source, number)))
    blocks = tuple(ResolvedBlock(slot, "This month", asset_hashes=tuple(v[0] for v in values),
                               asset_captions=tuple(v[1] for v in values), references=tuple(v[2] for v in values),
                               asset_provenance=tuple((reference, tuple(v[2] for v in values if v[0] == reference))
                                                      for reference in dict.fromkeys(v[0] for v in values)))
                 for slot, values in grouped.items())
    from app.monthly_report_asset_review import approve_all_assets, refresh_asset_source_context
    return tuple(approve_all_assets(refresh_asset_source_context(b, tuple(c.source for c in contents))) for b in blocks)
