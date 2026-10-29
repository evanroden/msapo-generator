"""Deterministic DOCX assembly with recoverable PDF conversion.

The committed shell contains only styles and page chrome, never client assets.
Runtime library assets are resolved by hash. Missing content fails explicitly;
it must never vanish silently from a finished report.
"""

from __future__ import annotations

from dataclasses import dataclass, replace
from collections.abc import Callable
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
import logging
import hashlib
import tempfile
from uuid import uuid4
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from PIL import Image, ImageDraw, ImageOps

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from app.monthly_report_render_compat import rendering_docx
from app import pdf_converter
from app import monthly_report_render_jobs as monthly_render_jobs
from app.ocr import SUPPORTED_IMAGE_SUFFIXES, _MAX_PIXELS_PER_FRAME
from app.monthly_report_checks import preflight
from app.monthly_report_model import ReportDraft, included_sections, report_filename, used_block_keys, used_asset_references


SHELL_PATH = Path(__file__).resolve().parents[1] / "templates" / "monthly_report" / "shell.docx"
FIXED_TIME = datetime(2000, 1, 1, tzinfo=timezone.utc)
OCEAN = "092B24"  # Matches --enfra-ocean in the existing application CSS.


@dataclass(frozen=True)
class ReportImage:
    data: bytes
    extension: str
    width: int
    height: int


def normalize_report_image(raw: bytes, suffix: str, *, line_art: bool = False,
                           frame: tuple[float, float] = (7, 9), dpi: int = 200,
                           preserve_transparency: bool = False) -> ReportImage:
    """Print normalization; vision still uses the stricter OCR request budgets.

    Line art must stay lossless, so the JPEG-only vision encoder is unsuitable
    here. Share its formats and pixel limit, and check size BEFORE decoding.
    Logos opt into their original alpha channel. Never infer transparency by
    removing white or dark pixels: those can be intentional parts of a mark.
    """
    if suffix.lower() not in SUPPORTED_IMAGE_SUFFIXES or len(raw) > 30 * 1024 * 1024:
        raise ValueError("Choose a supported image no larger than 30 MB.")
    if dpi not in (96, 150, 200) or min(frame) <= 0 or max(frame) > 11:
        raise ValueError("Invalid report image frame.")
    if suffix.lower() in {".heic", ".heif", ".hif"}:
        from pillow_heif import register_heif_opener
        register_heif_opener(thumbnails=False)
    try:
        opened = Image.open(BytesIO(raw))
    except Image.DecompressionBombError as exc:
        raise ValueError("Image is too large to decode safely. Resize it before uploading.") from exc
    with opened as image:
        if image.width * image.height > _MAX_PIXELS_PER_FRAME:
            raise ValueError("Image is too large to decode safely. Resize it before uploading.")
        try:
            frame_count = getattr(image, "n_frames", 1)
        except TypeError as exc:
            raise ValueError("Image frame information could not be read. Export a fresh single-image copy.") from exc
        if frame_count != 1:
            raise ValueError("Choose a single image; use a PDF for multiple pages.")
        oriented = ImageOps.exif_transpose(image)
        oriented.thumbnail((round(frame[0] * dpi), round(frame[1] * dpi)), Image.Resampling.LANCZOS)
        rgba = oriented.convert("RGBA")
        if preserve_transparency:
            normalized = rgba
        else:
            normalized = Image.new("RGB", rgba.size, "white")
            normalized.paste(rgba, mask=rgba.getchannel("A"))
        output = BytesIO()
        lossless = line_art or preserve_transparency
        normalized.save(output, format="PNG" if lossless else "JPEG",
                        **({"optimize": True} if lossless else {"quality": 82, "optimize": True}))
        return ReportImage(output.getvalue(), "png" if lossless else "jpg", *normalized.size)


@dataclass(frozen=True)
class ReportPackage:
    docx: bytes
    pdf: bytes | None
    docx_name: str
    pdf_name: str
    fingerprint: str
    pdf_error: str = ""
    source_view_digest: str = ""
    source_view_docx_sha256: str = ""
    source_view_art_digest: str = ""


def canonical_docx(raw: bytes) -> bytes:
    """python-docx fixes XML timestamps, but zipfile otherwise adds wall time."""
    output = BytesIO()
    with ZipFile(BytesIO(raw)) as source, ZipFile(output, "w", ZIP_DEFLATED) as target:
        for name in sorted(source.namelist()):
            info = ZipInfo(name, (2000, 1, 1, 0, 0, 0))
            info.compress_type = ZIP_DEFLATED
            info.create_system = 3
            info.external_attr = 0o600 << 16
            target.writestr(info, source.read(name))
    return output.getvalue()


def _save(document) -> bytes:
    core = document.core_properties
    core.author = "Email Process Control"
    core.last_modified_by = "Email Process Control"
    core.created = FIXED_TIME
    core.modified = FIXED_TIME
    core.revision = 1
    core.title = "Operations and Maintenance Monthly Review"
    core.subject = ""
    core.comments = ""
    # python-docx ships a stock preview thumbnail. Do not accidentally ship it
    # as a supposed branding asset in the content-free shell or generated file.
    rels = document.part.package.rels
    for key, rel in list(rels.items()):
        if rel.reltype.endswith("/thumbnail"):
            del rels[key]
    buffer = BytesIO()
    document.save(buffer)
    return canonical_docx(buffer.getvalue())


def _field(paragraph, instruction: str) -> None:
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), instruction)
    paragraph._p.append(field)


def _configure_section(section, *, content: bool, address: str = "") -> None:
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.75 if content else 0)
    section.left_margin = section.right_margin = Inches(0.75 if content else 0)
    section.header_distance = section.footer_distance = Inches(0.3)
    # Headers/footers link to the previous section by default. Explicitly
    # unlink all three variants so a divider can never inherit contact chrome.
    section.different_first_page_header_footer = False
    for part in (section.header, section.footer, section.first_page_header,
                 section.first_page_footer, section.even_page_header, section.even_page_footer):
        part.is_linked_to_previous = False
        for child in list(part._element):
            part._element.remove(child)
        part.add_paragraph()
    if content:
        section.header.paragraphs[0].text = "ENFRA | Operations and Maintenance"
        footer = section.footer.paragraphs[0]
        if address:
            footer.add_run(address + "  |  ")
        footer.add_run("Page ")
        _field(footer, "PAGE")


def create_shell(path: Path) -> None:
    """Reproducible build helper; runtime generation only reads the checked shell."""
    document = Document()
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(8)
    for style_name, size in (("Title", 30), ("Heading 1", 23), ("Heading 2", 16)):
        style = document.styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(OCEAN)
    _configure_section(document.sections[0], content=True)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(_save(document))


def outline(draft: ReportDraft) -> tuple[tuple[str, tuple[str, ...], int], ...]:
    """Include dedicated image pages; text/table overflow remains an estimate."""
    blocks = {b.key: b for b in draft.blocks if b.source != "Omit"}
    result = []
    def image_count(spec, block):
        from app.monthly_report_visuals import org_groups, photo_count
        try:
            from app.monthly_report_asset_layout import picture_runs
            return len(org_groups(block.org_nodes)) if block.org_nodes else len(tuple(picture_runs(block, draft.sources))) if spec.type == "image_grid" else len(block.asset_hashes)
        except ValueError:
            return 0  # Preflight explains the invalid chart/layout; no preview crash.
    for section in included_sections(draft.sections):
        images = sum(image_count(b, blocks[b.key]) for b in section.blocks if b.key in blocks)
        has_other = any(blocks[b.key].text or blocks[b.key].rows or blocks[b.key].extra_tables for b in section.blocks if b.key in blocks)
        result.append((f"{section.number}. {section.title}",
                       tuple(b.key for b in section.blocks if b.key in blocks), 1 + max(1, images + int(has_other))))
    return tuple(result)


def _display_page(document, *, label: str, title: str, subtitle: str = "", light: bool = False) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Inches(2.2)
    paragraph.paragraph_format.left_indent = Inches(0.75)
    run = paragraph.add_run(label)
    run.bold = True
    run.font.size = Pt(22)
    run.font.color.rgb = RGBColor.from_string("FFFFFF" if light else OCEAN)
    paragraph = document.add_paragraph(title, "Title")
    paragraph.paragraph_format.left_indent = Inches(0.75)
    paragraph.paragraph_format.right_indent = Inches(0.75)
    if light:
        for run in paragraph.runs:
            run.font.color.rgb = RGBColor(255, 255, 255)
    if subtitle:
        paragraph = document.add_paragraph(subtitle)
        paragraph.paragraph_format.left_indent = Inches(0.75)
        paragraph.paragraph_format.right_indent = Inches(0.75)


def _text(document, text: str) -> None:
    for line in text.splitlines():
        if not line.strip():
            continue
        if line.startswith(("- ", "• ")):
            document.add_paragraph(line[2:], "List Bullet")
        elif line.startswith("## "):
            document.add_paragraph(line[3:], "Heading 2")
        else:
            document.add_paragraph(line)


def estimate_bytes(draft: ReportDraft, asset_loader: Callable[[str], bytes] | None = None) -> int:
    included_keys = used_block_keys(draft)
    used_blocks = [b for b in draft.blocks if b.key in included_keys and b.source != "Omit"]
    references = used_asset_references(draft)
    from app.monthly_report_visuals import org_groups, photo_count
    # Generated page images, including editable charts/grids, occupy space too.
    generated = 0
    for block in used_blocks:
        if block.org_nodes:
            try:
                generated += 200_000 * len(org_groups(block.org_nodes))
            except ValueError:
                pass  # The content check explains an invalid chart.
        elif block.asset_hashes and any(s.key == block.key and s.type == "image_grid" for section in draft.sections for s in section.blocks):
            from app.monthly_report_asset_layout import picture_runs
            generated += 150_000 * len(tuple(picture_runs(block, draft.sources)))
    generated += 250_000 * sum(bool(s.divider_asset) for s in included_sections(draft.sections))
    return 150_000 + generated + sum(len(b.text.encode()) + sum(len(c.encode()) for row in b.rows for c in row)
                        + sum(len(c.encode()) for t in b.extra_tables for row in (t.columns, *t.rows) for c in row)
                        for b in used_blocks) + (
        sum(len(asset_loader(ref)) for ref in references) if asset_loader else 0
    )


def _picture(document, raw: bytes, *, width: float = 7, height: float = 8,
             inset: float = 0, page_break_before: bool = False) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(inset)
    paragraph.paragraph_format.space_after = Pt(0)
    if page_break_before:
        paragraph.paragraph_format.page_break_before = True
    _fit_picture(paragraph.add_run(), raw, width, height)


def _fit_picture(run, raw, width, height):
    with Image.open(BytesIO(raw)) as image:
        if image.width * image.height > _MAX_PIXELS_PER_FRAME:
            raise ValueError("Resolved image exceeds the pixel limit.")
        scale = min(width / image.width, height / image.height)
        run.add_picture(BytesIO(raw), width=Inches(image.width * scale), height=Inches(image.height * scale))


def _divider_background(document, raw):
    """Full-page photo behind editable Word text; crop only divider artwork."""
    normalized = normalize_report_image(raw, ".png", frame=(8.5, 11), dpi=150)
    with Image.open(BytesIO(normalized.data)) as image:
        page = ImageOps.fit(image.convert("RGB"), (1275, 1650), Image.Resampling.LANCZOS)
    # A solid band guarantees title contrast regardless of the supplied photo.
    draw = ImageDraw.Draw(page)
    draw.rectangle((0, 290, 1275, 1050), fill="#" + OCEAN)
    draw.rectangle((0, 290, 1275, 305), fill="#d9ee6c")
    buffer = BytesIO()
    page.save(buffer, "JPEG", quality=85, optimize=True)
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = Pt(1)
    inline = paragraph.add_run().add_picture(BytesIO(buffer.getvalue()), width=Inches(8.5), height=Inches(11))._inline
    anchor = OxmlElement("wp:anchor")
    for key, value in {"distT":"0", "distB":"0", "distL":"0", "distR":"0", "simplePos":"0", "relativeHeight":"0", "behindDoc":"1", "locked":"0", "layoutInCell":"1", "allowOverlap":"1"}.items():
        anchor.set(key, value)
    simple = OxmlElement("wp:simplePos")
    simple.set("x", "0"); simple.set("y", "0")
    anchor.append(simple)
    for axis in ("H", "V"):
        position = OxmlElement("wp:position" + axis)
        position.set("relativeFrom", "page")
        offset = OxmlElement("wp:posOffset"); offset.text = "0"
        position.append(offset); anchor.append(position)
    anchor.append(inline.extent)
    anchor.append(OxmlElement("wp:wrapNone"))
    for child in list(inline):
        anchor.append(child)
    inline.getparent().replace(inline, anchor)


def _format_table(table, has_header):
    """Keep site labels with their row and repeat column titles on overflow."""
    for index, row in enumerate(table.rows):
        properties = row._tr.get_or_add_trPr()
        properties.append(OxmlElement("w:cantSplit"))
        if index == 0 and has_header:
            properties.append(OxmlElement("w:tblHeader"))
        for cell in row.cells:
            if index == 0 and has_header:
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), OCEAN)
                cell._tc.get_or_add_tcPr().append(shading)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(4)
                paragraph.paragraph_format.keep_with_next = index == 0 and has_header
                for run in paragraph.runs:
                    run.font.size = Pt(9 if len(row.cells) > 6 else 10)
                    if index == 0 and has_header:
                        run.font.color.rgb = RGBColor(255, 255, 255)
                        run.bold = True


def assemble_docx(draft: ReportDraft, *, acknowledged_fingerprint: str = "",
                  asset_loader: Callable[[str], bytes] | None = None, source_view=None) -> bytes:
    checks = preflight(draft, estimate_bytes(draft, asset_loader))
    blocking = [c.message for c in checks if c.blocking]
    if blocking:
        raise ValueError(" ".join(blocking))
    if any(not c.blocking for c in checks) and acknowledged_fingerprint != draft.fingerprint:
        raise ValueError("Acknowledge the warnings for this version of the report.")
    return _build_docx(draft, asset_loader=asset_loader, source_view=source_view)


def _build_docx(draft, *, asset_loader=None, source_view=None):
    """Internal layout assembler. Final output always uses assemble_docx's gate."""
    from app.monthly_report_source_view_runtime import render_operation
    with render_operation(draft, source_view) as operation:
        return _build_verified_docx(draft, asset_loader=asset_loader, source_view=operation.view)


def _build_verified_docx(draft, *, asset_loader=None, source_view=None):
    from app.monthly_report_designs import source_for, is_master
    source = source_for(draft.profile)
    if source is not None:
        from app.monthly_report_native_layout import build_native_docx
        return build_native_docx(draft, source, asset_loader=asset_loader, master=is_master(draft.profile), source_view=source_view)
    if source_view is not None:
        raise ValueError("A compatibility source view requires its pinned original native design.")
    document = Document(SHELL_PATH)
    # Override the shell's stock blue theme without modifying the content-free
    # shell or importing any private branding into git.
    for node in document.styles.element.iter():
        if node.get(qn("w:themeColor")) == "accent1":
            node.attrib.pop(qn("w:themeColor"), None)
            node.attrib.pop(qn("w:themeShade"), None)
            node.attrib.pop(qn("w:themeTint"), None)
            node.set(qn("w:val") if node.tag == qn("w:color") else qn("w:color"), OCEAN)
        if node.get(qn("w:themeFill")) == "accent1":
            for name in ("themeFill", "themeFillTint", "themeFillShade"):
                node.attrib.pop(qn("w:" + name), None)
            node.set(qn("w:fill"), "EDF3F0")
    for border in document.styles["Title"].element.xpath("./w:pPr/w:pBdr"):
        border.getparent().remove(border)
    blocks = {b.key: b for b in draft.blocks}
    def cover_asset(key):
        block = blocks.get(key)
        if block and block.source != "Omit" and block.asset_hashes:
            if asset_loader is None:
                raise ValueError("Cover image requires the library resolver.")
            return asset_loader(block.asset_hashes[0])
        return None

    logo = cover_asset("brand_logo")
    def configure_content(section):
        _configure_section(section, content=True, address=draft.address_line)
        if logo:
            paragraph = section.header.paragraphs[0]
            paragraph.clear()
            _fit_picture(paragraph.add_run(), logo, 2.5, 0.3)
    _configure_section(document.sections[0], content=False)
    _display_page(document, label=draft.profile.contract, title=draft.profile.title,
                  subtitle="Operations and Maintenance Monthly Review\n" + draft.period.label)
    client_logo = cover_asset("client_logo")
    if client_logo:
        _picture(document, client_logo, width=2.5, height=0.75, inset=0.75)
    for text in ("Prepared by: " + draft.prepared_by,
                 "Facilities: " + "; ".join(f.title for f in draft.profile.facilities),
                 "SYNTHETIC DEMONSTRATION — NOT A CLIENT REPORT" if draft.synthetic else ""):
        if text:
            paragraph = document.add_paragraph(text)
            paragraph.paragraph_format.left_indent = Inches(0.75)
            paragraph.paragraph_format.right_indent = Inches(0.75)
    photo = cover_asset("cover_photo")
    if photo:
        _picture(document, photo, width=7, height=3, inset=0.75)
    configure_content(document.add_section(WD_SECTION_START.NEW_PAGE))
    document.add_heading("Table of contents", 0)
    sections = included_sections(draft.sections)
    for section in sections:
        document.add_paragraph(f"{section.number}. {section.title}")
    document.add_paragraph("Confidential — intended for the report recipients. Do not distribute without permission.")
    document.add_paragraph("Create. Sustain. Empower.")
    for section in sections:
        _configure_section(document.add_section(WD_SECTION_START.NEW_PAGE), content=False)
        divider = blocks.get("divider_" + section.key)
        divider_asset = (divider.asset_hashes[0] if divider.source != "Omit" and divider.asset_hashes else "") if divider else section.divider_asset
        if divider_asset:
            if asset_loader is None:
                raise ValueError("The divider image could not be resolved.")
            _divider_background(document, asset_loader(divider_asset))
        _display_page(document, label=f"Section {section.number}" if not section.appendix else f"Appendix {section.number}", title=section.title, light=bool(divider_asset))
        configure_content(document.add_section(WD_SECTION_START.NEW_PAGE))
        document.add_heading(f"{section.number}. {section.title}", 1)
        chart_emitted = False
        for spec in section.blocks:
            block = blocks.get(spec.key)
            if block is None or block.source == "Omit":
                continue
            if (section.key == "organization" and spec.key == "contact_matrix"
                    and chart_emitted and (block.text.strip() or block.rows or
                                          any(t.rows for t in block.extra_tables) or block.asset_hashes)):
                # A full-size chart and facility contacts have distinct jobs.
                # Never squeeze contact rows underneath the chart or split an
                # existing line into a new page without a human-readable title.
                # This is the generic fallback only: native design geometry
                # must be evaluated separately, never rewritten by this path.
                heading = document.add_heading("Facility contacts", 2)
                heading.paragraph_format.page_break_before = True
            _text(document, block.text)
            from app.monthly_report_visuals import org_groups, org_page, photo_count, photo_page
            if block.org_nodes:
                for index in range(len(org_groups(block.org_nodes))):
                    if index:
                        document.add_page_break()
                    _picture(document, org_page(block.org_nodes, index), height=7.7)
            elif spec.type == "image_grid" and block.asset_hashes:
                if asset_loader is None:
                    raise ValueError("Photo pages require the library resolver.")
                from app.monthly_report_asset_layout import picture_runs, photo_run, document_caption_problems, style_document_caption
                if document_caption_problems(block, draft.sources):
                    raise ValueError('Shorten the source-page caption to at most three lines before generating.')
                for index, (is_document, indexes) in enumerate(picture_runs(block, draft.sources)):
                    if index:
                        document.add_page_break()
                    if is_document:
                        item = indexes[0]
                        _picture(document, asset_loader(block.asset_hashes[item]), height=7.7)
                        if item < len(block.asset_captions) and block.asset_captions[item]:
                            style_document_caption(document.add_paragraph(block.asset_captions[item], style="Caption"))
                    else:
                        _picture(document, photo_page(photo_run(block, indexes), asset_loader, 0), height=7.7)
            for index, reference in enumerate(block.asset_hashes if not block.org_nodes and spec.type != "image_grid" else ()):
                if asset_loader is None:
                    raise ValueError("Image and PDF-page assets require the library resolver.")
                # The page break belongs to the next image paragraph. A
                # standalone break paragraph can consume an unwanted page
                # when a large preceding source page nearly fills the sheet.
                _picture(document, asset_loader(reference), page_break_before=index > 0)
                if index < len(block.asset_captions) and block.asset_captions[index]:
                    document.add_paragraph(block.asset_captions[index], style="Caption")
            if block.rows:
                width = len(spec.columns) or max(map(len, block.rows))
                table = document.add_table(rows=1 if spec.columns else 0, cols=width)
                table.style = "Light Shading Accent 1"
                if spec.columns:
                    for cell, col in zip(table.rows[0].cells, spec.columns):
                        cell.text = col.title
                for values in block.rows:
                    if len(values) > width:
                        raise ValueError(f"Too many table columns in {spec.key}.")
                    for cell, value in zip(table.add_row().cells, values):
                        cell.text = "" if value is None else str(value)
                _format_table(table, bool(spec.columns))
            for imported in block.extra_tables:
                if not imported.columns or not any(c is not None and str(c).strip() for row in imported.rows for c in row):
                    continue
                table = document.add_table(rows=1, cols=len(imported.columns))
                table.style = "Light Shading Accent 1"
                for cell, title in zip(table.rows[0].cells, imported.columns):
                    cell.text = title
                for values in imported.rows:
                    if len(values) > len(imported.columns):
                        raise ValueError("An imported table has more cells than columns.")
                    for cell, value in zip(table.add_row().cells, values):
                        cell.text = "" if value is None else str(value)
                _format_table(table, True)
            if spec.key == "org_chart" and (block.org_nodes or block.asset_hashes):
                chart_emitted = True
        carried = [item for item in draft.follow_ups if item.included and ("issues" if item.category == "issue" else "proposals") == section.key]
        if carried:
            from app.monthly_report_followups import report_text
            if section.key != "issues":
                document.add_heading("Carried-forward items", 2)
            for item in carried:
                _text(document, report_text(item))
    return _save(document)


def generate_report(draft: ReportDraft, *, acknowledged_fingerprint: str = "",
                    asset_loader: Callable[[str], bytes] | None = None, source_view=None) -> ReportPackage:
    try:
        with monthly_render_jobs.work_slot():
            return _generate_report(draft, acknowledged_fingerprint=acknowledged_fingerprint, asset_loader=asset_loader, source_view=source_view)
    except monthly_render_jobs.RenderBusy as exc:
        raise ValueError(str(exc)) from exc


def _generate_report(draft, *, acknowledged_fingerprint='', asset_loader=None, source_view=None):
    from app.monthly_report_source_view_runtime import render_operation
    with render_operation(draft, source_view) as operation:
        return _generate_verified_report(draft, acknowledged_fingerprint=acknowledged_fingerprint,
                                          asset_loader=asset_loader, operation=operation)


def _generate_verified_report(draft, *, acknowledged_fingerprint, asset_loader, operation):
    source_view = operation.view
    docx = assemble_docx(draft, acknowledged_fingerprint=acknowledged_fingerprint, asset_loader=asset_loader, source_view=source_view)
    # Retained native PNGs may still contain erased artwork under alpha zero.
    # This bounded derivative pass is separate from pixel-preserving metadata
    # stripping; neither operation changes the saved source/design/assets.
    from app.monthly_report_media_privacy import sanitize_package_media, sanitize_package_transparency
    docx = sanitize_package_transparency(docx)
    docx = sanitize_package_media(docx)
    pdf, error = _convert_report_pdf(draft, docx)
    return ReportPackage(docx, pdf, report_filename(draft, "docx"), report_filename(draft, "pdf"), draft.fingerprint, error,
                         source_view.receipt_sha256 if source_view is not None else "",
                         hashlib.sha256(docx).hexdigest() if source_view is not None else "",
                         operation.art_digest)


def retry_report_pdf(draft: ReportDraft, package: ReportPackage, *, source_view=None) -> ReportPackage:
    """Retry only conversion of the same current native Word bytes. Never save."""
    if package.fingerprint != draft.fingerprint:
        raise ValueError("This Word file belongs to an earlier draft. Create files for the current report first.")
    from app.monthly_report_source_view_runtime import cache_identity, render_operation
    # Preserve the legacy cached no-work path only when no scoped evidence
    # exists. Recorded/malformed evidence and even partial package tags must
    # still pass the independently verified operation below.
    if (package.pdf and source_view is None and not package.source_view_digest
            and not package.source_view_docx_sha256 and not package.source_view_art_digest
            and not cache_identity(draft)):
        return package
    try:
        with monthly_render_jobs.work_slot():
            with render_operation(draft, source_view) as operation:
                return _retry_verified_report_pdf(draft, package, operation)
    except monthly_render_jobs.RenderBusy as exc:
        raise ValueError(str(exc)) from exc


def _retry_verified_report_pdf(draft, package, operation):
    source_view = operation.view
    if package.fingerprint != draft.fingerprint:
        raise ValueError("This Word file belongs to an earlier draft. Create files for the current report first.")
    if package.source_view_digest or package.source_view_docx_sha256 or package.source_view_art_digest or source_view is not None:
        from app.monthly_report_compatibility_view import CompatibilityView
        if (not package.source_view_digest or not isinstance(source_view, CompatibilityView)
                or package.source_view_digest != source_view.receipt_sha256
                or package.source_view_docx_sha256 != hashlib.sha256(package.docx).hexdigest()
                or package.source_view_art_digest != operation.art_digest):
            raise ValueError("PDF retry requires the exact expected native source view and unchanged Word bytes.")
        from app.monthly_report_designs import source_for, is_master
        source = source_for(draft.profile)
        if source is None or is_master(draft.profile):
            raise ValueError("PDF retry requires the source view's pinned original design.")
    if package.pdf:
        return package
    pdf, error = _convert_report_pdf(draft, package.docx)
    return replace(package, pdf=pdf, pdf_error=error)


def _convert_report_pdf(draft, docx):
    pdf = None
    error = ""
    # The converter writes into the shared OUTPUT_DIR. A random stem prevents
    # collisions; both the actual and expected output paths are always cleaned.
    stem = "monthly-" + uuid4().hex
    expected_pdf = pdf_converter.OUTPUT_DIR / (stem + ".pdf")
    actual_pdf = None
    try:
        with tempfile.TemporaryDirectory(prefix="monthly-report-") as directory:
            input_path = Path(directory) / (stem + ".docx")
            from app.monthly_report_designs import render_profile_for
            # Keep the downloadable native Word file untouched. For the
            # disposable PDF-conversion copy only, bound verified photographic
            # JPEGs inherited from the pinned design. New uploaded photos,
            # diagrams, charts and non-JPEG art are not altered.
            from app.monthly_report_designs import source_for
            from app.monthly_report_pdf_image_budget import (
                compact_native_photos_for_pdf, release_transient_word_memory)
            converted_copy = rendering_docx(docx, profile=render_profile_for(draft.profile))
            pinned_source = source_for(draft.profile)
            if pinned_source is not None:
                converted_copy = compact_native_photos_for_pdf(
                    converted_copy, pinned_source=pinned_source)
            # A retry can hold Word bytes created before this privacy boundary
            # existed. Clean only its disposable PDF input, not that download.
            from app.monthly_report_media_privacy import sanitize_package_transparency
            converted_copy = sanitize_package_transparency(converted_copy)
            input_path.write_bytes(converted_copy)
            del converted_copy
            release_transient_word_memory()
            actual_pdf = monthly_render_jobs.convert_to_pdf(input_path, wait_seconds=130)
            pdf = actual_pdf.read_bytes()
            if not pdf.startswith(b"%PDF-"):
                raise ValueError("The converter returned an invalid PDF.")
    except Exception as exc:  # noqa: BLE001 - keep the finished DOCX on any backend failure
        pdf = None
        if isinstance(exc, pdf_converter.PDFConversionTimeout):
            error = "PDF conversion took too long. The DOCX is ready to download."
        else:
            error = f"PDF conversion failed ({type(exc).__name__}). The DOCX is ready to download; retry PDF generation."
    finally:
        for path in {actual_pdf, expected_pdf} - {None}:
            try:
                path.unlink(missing_ok=True)
            except OSError:
                # A cleanup failure must not discard either finished download.
                logging.getLogger(__name__).warning("Monthly-report temporary PDF cleanup failed")
    return pdf, error
